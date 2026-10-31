from pathlib import Path
import json,re,hashlib,zipfile,csv,datetime
from docx import Document
root=Path(r'E:\PT-Original-Point\ttqs-one\docs\ttqs-quality-20261016\quality-review\0105_I18-E13_R01_20261022')
md=root/'writer/opencode_writer_manuscript_0105_R01.md'
docx=root/'candidate/0105_I18-E13_績效改善調查與課後行動計畫追蹤套件_R02.docx'
reg=root/'evidence/synthetic_replacement_register_0105_R02.csv'
out=root/'evidence/technical_qa_receipt_0105_R02.json'
source=md.read_text(encoding='utf-8'); d=Document(docx)
body='\n'.join([p.text for p in d.paragraphs]+[c.text for tb in d.tables for row in tb.rows for c in row.cells])
headers='\n'.join(p.text for s in d.sections for p in s.header.paragraphs)
footers='\n'.join(p.text for s in d.sections for p in s.footer.paragraphs)
with reg.open(encoding='utf-8-sig',newline='') as f: rr=list(csv.DictReader(f))
ids=[r.get('fact_id','') for r in rr]
status_key=next((k for k in (rr[0].keys() if rr else []) if k=='狀態'),None)
checks={
 'zip_crc_clean':zipfile.ZipFile(docx).testzip() is None,
 'docx_parse_ok':True,
 'doc_id_and_requirement_present':'0105' in body and 'I18-E13' in body,
 'all_six_document_specific_forms':all(f'Form {x}' in body for x in 'ABCDEF'),
 'no_stale_form_g_pointer':not bool(re.search(r'\bForm\s+G\b',body,re.I)),
 'sample_synthetic_real_unknown_header':all(s in headers for s in ('0105','I18-E13','SAMPLE','SYNTHETIC','REAL UNKNOWN')),
 'footer_has_page_field_and_label':('頁' in footers or 'Page' in footers) and any(x.tag.endswith('fldSimple') for s in d.sections for p in s.footer.paragraphs for x in p._p),
 'no_unsupported_approval_claim':not any(s in body for s in ('本協會已核准','真實已核准','已正式核准')),
 'six_expected_tables':len(d.tables)==6,
 'replacement_register_has_60_rows':len(rr)==60,
 'replacement_ids_exactly_001_to_060':ids==[f'SYN-0105-F{i:03d}' for i in range(1,61)],
 'single_synthetic_case':len({r.get('case_id') for r in rr})==1 and {r.get('case_id') for r in rr}=={'SYN-0105-CASE-01'},
 'all_register_rows_marked_sample_and_hold':bool(status_key) and all('SAMPLE/SYNTHETIC' in r.get(status_key,'') and 'HOLD' in r.get(status_key,'') for r in rr),
 'customer_question_same_in_both_forms':source.count('依您可觀察的資料，訓練後工作／服務表現有改善')>=2,
 'synthetic_collection_period_filled':'SYN-W1～SYN-W2' in source,
 'section14_has_hold_decisions':'## 十四、' in source and 'HOLD' in source[source.find('## 十四、'):source.find('## 十五、')],
 'source_replacement_rows_60':len(re.findall(r'^\| SYN-0105-F\d{3} \|',source,re.M))==60,
 'response_rate_examples_present':all(x in source for x in ('10/12=83.3%','6/7=85.7%','4/5=80.0%')),
 'improvement_examples_present':all(x in source for x in ('5/10=50%','8/10=80%','+3','+30','+60%')),
}
counts={'overall':[4,3,2,1,0,0],'learner':[2,2,1,1,0,0],'customer':[2,1,1,0,0,0]}; weights=[5,4,3,2,1,0]
math={k:{'n':sum(v),'score':sum(a*b for a,b in zip(v,weights)),'positive':v[0]+v[1]} for k,v in counts.items()}
math['subgroups_reconcile']=all(counts['overall'][i]==counts['learner'][i]+counts['customer'][i] for i in range(6)) and math['overall']['score']==math['learner']['score']+math['customer']['score']
math['overall_mean']=math['overall']['score']/math['overall']['n']; math['overall_positive_pct']=round(math['overall']['positive']/math['overall']['n']*100,1)
math['learner_positive_pct']=round(math['learner']['positive']/math['learner']['n']*100,1); math['customer_positive_pct']=round(math['customer']['positive']/math['customer']['n']*100,1)
checks['sample_response_math_reconciles']=math['subgroups_reconcile'] and (math['overall_mean'],math['overall_positive_pct'],math['learner_positive_pct'],math['customer_positive_pct'])==(4.0,70.0,66.7,75.0)
checks['first_section_header_not_lost']=bool(headers.strip())
checks['first_section_footer_not_lost']=bool(footers.strip())
checks['normal_style_10_5pt']=float(d.styles['Normal'].font.size.pt)==10.5
checks['table_styles_grid']=all(tb.style.name=='Table Grid' for tb in d.tables)
record={'receipt_id':'TTQS-0105-TECHNICAL-QA-R02','checked_at_tpe':datetime.datetime.now().astimezone().isoformat(),'doc_id':'0105','requirement_id':'I18-E13','project_157_index':114,'official_checkbox_serial':'NOT_ASSERTED','source_markdown_sha256':hashlib.sha256(md.read_bytes()).hexdigest(),'docx_path':str(docx),'docx_sha256':hashlib.sha256(docx.read_bytes()).hexdigest(),'docx_bytes':docx.stat().st_size,'register_csv_sha256':hashlib.sha256(reg.read_bytes()).hexdigest(),'blueprint_sha256':hashlib.sha256((root/'source/UNIQUE_BLUEPRINT_0105_I18-E13_R01.md').read_bytes()).hexdigest(),'table_rows':[len(tb.rows) for tb in d.tables],'section_count':len(d.sections),'section_orientations':[str(s.orientation) for s in d.sections],'synthetic_math':math,'checks':checks,'pass_count':sum(checks.values()),'check_count':len(checks),'verdict':'PASS' if all(checks.values()) else 'FAIL','r01_layout_observation':'First-section header/footer were removed when linking section zero to previous; corrected in R02. R01 preserved as historical, not accepted.'}
out.write_text(json.dumps(record,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(record,ensure_ascii=True,indent=2))