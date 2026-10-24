from pathlib import Path
import csv,json,re,hashlib,zipfile,datetime
from docx import Document
from lxml import etree
BASE=Path(r'E:\PT-Original-Point\ttqs-one\docs\ttqs-quality-20261016\quality-review\0105_I18-E13_R01_20261022')
md=BASE/'writer/opencode_writer_manuscript_0105_R06_release.md'
docx=BASE/'candidate/0105_I18-E13_績效改善調查與課後行動計畫追蹤套件_R06.docx'
reg=BASE/'evidence/synthetic_replacement_register_0105_R06.csv'
out=BASE/'evidence/technical_qa_receipt_0105_R06.json'
source=md.read_text(encoding='utf-8'); d=Document(docx)
body='\n'.join([p.text for p in d.paragraphs]+[c.text for tb in d.tables for row in tb.rows for c in row.cells])
headers='\n'.join(p.text for s in d.sections for p in s.header.paragraphs)
footers='\n'.join(p.text for s in d.sections for p in s.footer.paragraphs)
with reg.open(encoding='utf-8-sig',newline='') as f: rr=list(csv.DictReader(f))
ids=[r.get('fact_id','') for r in rr]; status_key='狀態' if rr and '狀態' in rr[0] else None
register_tables=[tb for tb in d.tables if tb.rows and any('fact_id' in c.text for c in tb.rows[0].cells)]
register_word_ids=[]
for tb in register_tables:
 for row in tb.rows[1:]:
  m=re.search(r'SYN-0105-F(\d{3})',row.cells[0].text)
  if m: register_word_ids.append(int(m.group(1)))
validity_rows=[line for line in source.splitlines() if re.match(r'^\| SYN-RESP-[LC]\d{2} \|',line)]
parsed=[]
for line in validity_rows:
 cells=[x.strip() for x in line.strip().strip('|').split('|')]
 if len(cells)>=12:
  rid=cells[0]; role=cells[1]; numeric=int(re.sub(r'\D','',cells[3])); na_reason=int(re.sub(r'\D','',cells[4])); na_no=int(re.sub(r'\D','',cells[5])); missing=int(re.sub(r'\D','',cells[6])); verdict=cells[11].replace('*','').strip()
  parsed.append((rid,role,numeric,na_reason,na_no,missing,verdict))
valid=[x for x in parsed if x[6]=='有效']; invalid=[x for x in parsed if x[6]=='無效']
math={'responses':len(parsed),'valid':len(valid),'invalid':len(invalid),'learners':sum(x[0].startswith('SYN-RESP-L') for x in parsed),'customers':sum(x[0].startswith('SYN-RESP-C') for x in parsed),'valid_learners':sum(x[0].startswith('SYN-RESP-L') for x in valid),'valid_customers':sum(x[0].startswith('SYN-RESP-C') for x in valid),'numeric_answers':sum(x[2] for x in parsed),'na_with_reason':sum(x[3] for x in parsed),'na_without_reason':sum(x[4] for x in parsed),'missing':sum(x[5] for x in parsed)}
checks={
 'zip_crc_clean':zipfile.ZipFile(docx).testzip() is None,
 'docx_parse_ok':True,
 'doc_id_and_requirement_present':'0105' in body and 'I18-E13' in body,
 'all_six_document_specific_forms':all(f'Form {x}' in body for x in 'ABCDEF'),
 'official_checkbox_serial_not_invented':'官方勾選序號' in body and ('未提供' in body or '不主張' in body or '未宣稱' in body),
 'negative_neighbor_distinction':all(x in body for x in ('0104／I18-E11','0106／I18-E14','0095／I18-E02')) and 'I18-E12' in body,
 'sample_real_boundary':all(x in headers for x in ('SAMPLE','SYNTHETIC','REAL UNKNOWN')) and 'REAL 有效問卷' in body and 'NOT_RECEIVED' in body,
 'no_internal_revision_metadata_in_body':not any(x in body for x in ('REVISION=','R05','R06','0610f5b4d3b3ca0e35421479f25f4152e283c8f682d7a40d91594302ced68ed7','JSONL','candidate-only')),
 'validity_thresholds_and_unscored_prompts':all(x in body for x in ('ceil(11×80%)','= 9','ceil(5×80%)','= 4','非計分佐證提示題')),
 'form_b_c_item_scopes':all(x in body for x in ('11 題','5 題','1–5／N/A')),
 'validity_audit_twelve_rows':len(parsed)==12 and len(valid)==10 and len(invalid)==2,
 'validity_audit_counts_reconcile':math['learners']==7 and math['customers']==5 and math['valid_learners']==6 and math['valid_customers']==4 and math['numeric_answers']==90 and math['na_with_reason']==8 and math['na_without_reason']==3 and math['missing']==1,
 'validity_audit_case_ids_match':all(x[0] in source for x in parsed) and 'SYN-RESP-L07' in body and 'SYN-RESP-C05' in body,
 'sample_response_math_reconciles':all(x in body for x in ('12÷12×100','10÷12×100','2÷12×100','6÷7×100','4÷5×100','40÷10','4.0/5','70.0%','66.7%','75.0%')) and math['valid']+math['invalid']==math['responses'],
 'all_calculations_round_to_one_decimal':'四捨五入至**小數第 1 位**' in source and '分子' in body and '分母' in body,
 'five_result_classes':'五類' in body and '四類' not in body,
 'form_e_plan_and_evidence_hold':'預期結果' in body and '期限' in body and 'HOLD' in body and 'CLOSED_FOR_SAMPLE' in body,
 'replacement_register_csv_152_unique_contiguous':len(rr)==152 and ids==[f'SYN-0105-F{i:03d}' for i in range(1,153)],
 'replacement_register_word_152_unique_contiguous':len(register_tables)>=1 and register_word_ids==list(range(1,153)),
 'replacement_rows_sample_and_hold':bool(status_key) and all('SAMPLE/SYNTHETIC' in r.get(status_key,'') and 'HOLD' in r.get(status_key,'') for r in rr),
 'source_blueprint_hash_matches_pinned_source':hashlib.sha256((BASE/'source/UNIQUE_BLUEPRINT_0105_I18-E13_R01.md').read_bytes()).hexdigest()=='0a457287fcc005e13408ad85763f4e7065dd1915f3f9b11757928c64715cb826',
 'source_packet_hash_matches_pinned_source':hashlib.sha256((BASE/'source/SOURCE_PACKET_0105_I18-E13_R01.md').read_bytes()).hexdigest()=='166338523fad5550511cd5f16fd7b347b28a23ef295fdc9396c21bb8728ae99e',
 'normal_style_10_5pt':d.styles['Normal'].font.size.pt==10.5,
 'table_styles_grid':all(tb.style.name=='Table Grid' for tb in d.tables),
 'table_text_minimum_9_5pt':all(r.font.size is None or r.font.size.pt>=9.5 for tb in d.tables for row in tb.rows for cell in row.cells for p in cell.paragraphs for r in p.runs),
 'portrait_and_landscape_sections':any(str(s.orientation).endswith('LANDSCAPE (1)') for s in d.sections) and any(str(s.orientation).endswith('PORTRAIT (0)') for s in d.sections),
 'title_and_purpose_present':any(p.style.name=='Title' and '績效改善調查' in p.text for p in d.paragraphs),
 'footer_page_field':any(x.tag.endswith('fldSimple') for s in d.sections for p in s.footer.paragraphs for x in p._p),
 'replacement_locator_instruction_updated':'Word 頁碼及表格列索引' in body,
 'blank_version_log_not_past_claim':'（實際填寫）' in body and '版本與變更紀錄' in body,
}
with zipfile.ZipFile(docx) as z:
 ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
 styles_xml=etree.fromstring(z.read('word/styles.xml')); document_xml=etree.fromstring(z.read('word/document.xml'))
 title=styles_xml.xpath(".//w:style[@w:styleId='Title']",namespaces=ns)
 checks['title_style_has_no_theme_border']=len(title)==1 and not title[0].xpath('.//w:pBdr',namespaces=ns)
 checks['no_direct_paragraph_borders']=not document_xml.xpath('.//w:body/w:p/w:pPr/w:pBdr',namespaces=ns)
record={'receipt_id':'TTQS-0105-TECHNICAL-QA-R06','checked_at_tpe':datetime.datetime.now().astimezone().isoformat(),'doc_id':'0105','requirement_id':'I18-E13','project_157_index':114,'official_checkbox_serial':'NOT_ASSERTED','source_packet_sha256':hashlib.sha256((BASE/'source/SOURCE_PACKET_0105_I18-E13_R01.md').read_bytes()).hexdigest(),'unique_blueprint_sha256':hashlib.sha256((BASE/'source/UNIQUE_BLUEPRINT_0105_I18-E13_R01.md').read_bytes()).hexdigest(),'release_markdown_sha256':hashlib.sha256(md.read_bytes()).hexdigest(),'docx_path':str(docx),'docx_sha256':hashlib.sha256(docx.read_bytes()).hexdigest(),'docx_bytes':docx.stat().st_size,'register_sha256':hashlib.sha256(reg.read_bytes()).hexdigest(),'tables':len(d.tables),'table_rows':[len(tb.rows) for tb in d.tables],'replacement_register_rows_csv':len(rr),'replacement_register_tables':len(register_tables),'replacement_register_rows_word':len(register_word_ids),'section_count':len(d.sections),'section_orientations':[str(s.orientation) for s in d.sections],'validity_audit':math,'checks':checks,'pass_count':sum(checks.values()),'check_count':len(checks),'verdict':'PASS' if all(checks.values()) else 'FAIL','scope':'deterministic OOXML structure, content mapping, synthetic arithmetic, text/style bounds, source hashes; visual review and independent semantics are separate gates'}
out.write_text(json.dumps(record,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(record,ensure_ascii=True,indent=2))

