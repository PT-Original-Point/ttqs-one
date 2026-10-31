from __future__ import annotations
import hashlib, json, re, zipfile
from datetime import datetime
from pathlib import Path
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

BASE=Path(r'E:\PT-Original-Point\ttqs-one\docs\ttqs-quality-20261016\quality-review\0105_I18-E13_R01_20261022')
DOCX=BASE/'candidate'/'0105_I18-E13_績效改善調查與課後行動計畫追蹤套件_R07.docx'
RAW=BASE/'writer'/'opencode_writer_manuscript_0105_R07_raw.md'
RELEASE=BASE/'writer'/'opencode_writer_manuscript_0105_R07_release.md'
SOURCE=BASE/'source'/'SOURCE_PACKET_0105_I18-E13_R01.md'
BLUEPRINT=BASE/'source'/'UNIQUE_BLUEPRINT_0105_I18-E13_R01.md'
TEXT_OUT=BASE/'evidence'/'candidate_semantic_text_0105_R07.txt'
RECEIPT=BASE/'evidence'/'technical_qa_0105_R07.json'
EXPECTED_DOCX='32bc0b8d7efe767e584b2a848206507d4da73ffa344378143b22186705302027'
EXPECTED_RAW='31b7de140f4cd93169c709f30e111ca0994267b59fa6a76ec06ccc1fda5cea12'
EXPECTED_RELEASE='d580c7e14c83b09cd2c0c4608eabdf65465ef9d7856e77ac019b9195a8418a86'

def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def ordered_text(doc):
    chunks=[]
    for child in doc.element.body.iterchildren():
        if child.tag==qn('w:p'):
            chunks.append(Paragraph(child,doc).text)
        elif child.tag==qn('w:tbl'):
            table=Table(child,doc)
            for row in table.rows:
                chunks.append(' | '.join(cell.text.replace('\n',' / ') for cell in row.cells))
    return '\n'.join(chunks)

if sha(DOCX)!=EXPECTED_DOCX: raise SystemExit('Candidate SHA changed; do not bind stale QA')
if sha(RAW)!=EXPECTED_RAW or sha(RELEASE)!=EXPECTED_RELEASE: raise SystemExit('Manuscript lineage hash mismatch')
with zipfile.ZipFile(DOCX) as z:
    bad=z.testzip()
    if bad: raise SystemExit('DOCX zip integrity failed at '+bad)
doc=Document(DOCX)
text=ordered_text(doc)
TEXT_OUT.write_text(text,encoding='utf-8',newline='\n')
facts=re.findall(r'SYN-0105-F(\d{3})',text)
register_facts=[]
for table in doc.tables:
    if table.rows and table.rows[0].cells[0].text.strip().lower()=='fact_id':
        for row in table.rows[1:]:
            match=re.search(r'SYN-0105-F(\d{3})',row.cells[0].text)
            if match: register_facts.append(match.group(1))
expected=[f'{n:03d}' for n in range(1,163)]
forms={f'Form {x}':f'Form {x}' in text for x in 'ABCDEF'}
learner_items={f'B{i}':bool(re.search(rf'\bB{i}\b',text)) for i in range(1,13)}
customer_items={f'C{i}':bool(re.search(rf'\bC{i}\b',text)) for i in range(1,9)}
internal_hits={pat:bool(re.search(pat,text,re.I)) for pat in [r'\bR0[1-9]\b',r'REVISION=',r'REQUIRED_COMPANION',r'F-\d{2}',r'I18-E12',r'E:\\\\PT-Original-Point',r'\b[0-9a-f]{64}\b']}
body_sizes=[];table_sizes=[]
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        if run.text.strip() and run.font.size is not None: body_sizes.append(run.font.size.pt)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    if run.text.strip() and run.font.size is not None: table_sizes.append(run.font.size.pt)
source_sha=sha(SOURCE); blueprint_sha=sha(BLUEPRINT)
checks={
 'docx_zip_integrity':bad is None,
 'lineage_hashes':sha(RAW)==EXPECTED_RAW and sha(RELEASE)==EXPECTED_RELEASE,
 'source_and_blueprint_hashes':source_sha=='166338523fad5550511cd5f16fd7b347b28a23ef295fdc9396c21bb8728ae99e' and blueprint_sha=='0a457287fcc005e13408ad85763f4e7065dd1915f3f9b11757928c64715cb826',
 'all_forms_A_to_F':all(forms.values()),
 'question_ids_B1_B12':all(learner_items.values()),
 'question_ids_C1_C8':all(customer_items.values()),
 'register_fact_ids_F001_F162':register_facts==expected,
 'sample_synthetic_boundary':'SAMPLE / SYNTHETIC' in text and 'NOT_RECEIVED' in text and 'HOLD' in text,
 'internal_build_markers_absent':not any(internal_hits.values()),
 'doc_id_and_requirement':'0105' in text and 'I18-E13' in text,
 'minimum_body_font_10_5':bool(body_sizes) and min(body_sizes)>=10.5,
 'minimum_table_font_9_5':bool(table_sizes) and min(table_sizes)>=9.5,
 'render_and_visual_review':'PENDING',
 'independent_semantic_review':'PENDING'
}
receipt={
 'receipt_id':'TTQS-0105-TECHNICAL-QA-R07','created_at_tpe':datetime.now().astimezone().isoformat(),
 'doc_id':'0105','official_requirement_id':'I18-E13','project_157_index':114,'official_checkbox_serial':'NOT_ASSERTED',
 'candidate_docx_sha256':sha(DOCX),'candidate_docx_bytes':DOCX.stat().st_size,
 'raw_worker_manuscript_sha256':sha(RAW),'release_manuscript_sha256':sha(RELEASE),'source_packet_sha256':source_sha,'unique_blueprint_sha256':blueprint_sha,
 'tables':len(doc.tables),'ordered_text_sha256':sha(TEXT_OUT),'ordered_text_bytes':TEXT_OUT.stat().st_size,
 'fact_reference_count':len(facts),'register_row_count':len(register_facts),'register_fact_ids_unique':len(set(register_facts)),
 'form_coverage':forms,'learner_item_coverage':learner_items,'customer_item_coverage':customer_items,
 'body_min_run_font_pt':min(body_sizes) if body_sizes else None,'table_min_run_font_pt':min(table_sizes) if table_sizes else None,
 'internal_marker_hits':internal_hits,'checks':checks,
 'technical_verdict':'PASS_WITH_RENDER_AND_SEMANTIC_PENDING' if all(v is True for k,v in checks.items() if k not in ('render_and_visual_review','independent_semantic_review')) else 'FAIL',
 'formal_acceptance':'NOT_GRANTED'}
RECEIPT.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'candidate_sha256':receipt['candidate_docx_sha256'],'tables':receipt['tables'],'register_rows':receipt['register_row_count'],'unique_register_ids':receipt['register_fact_ids_unique'],'technical_verdict':receipt['technical_verdict'],'failed_checks':[k for k,v in checks.items() if v is False]},ensure_ascii=False))
