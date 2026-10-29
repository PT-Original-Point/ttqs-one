import hashlib, json, re, sys, zipfile
from datetime import datetime, timezone
from pathlib import Path
from docx import Document
from PIL import Image
base=Path(sys.argv[1]); docx=base/'candidate'/'0011_I02-E02_R09.docx'; qa=base/'qa'; render=qa/'render_R09'
b=docx.read_bytes(); sha=hashlib.sha256(b).hexdigest(); z=zipfile.ZipFile(docx); bad=z.testzip(); d=Document(str(docx))
texts=[p.text for p in d.paragraphs]
for t in d.tables:
 for row in t.rows:
  texts.extend(c.text for c in row.cells)
text='\n'.join(texts)
empty_cells=[]; sizes={'body':[],'table':[]}; unset={'body':0,'table':0}
for p in d.paragraphs:
 for r in p.runs:
  if r.text.strip():
   if r.font.size: sizes['body'].append(r.font.size.pt)
   else: unset['body']+=1
for ti,t in enumerate(d.tables,1):
 for ri,row in enumerate(t.rows,1):
  for ci,c in enumerate(row.cells,1):
   if not c.text.strip(): empty_cells.append([ti,ri,ci])
   for p in c.paragraphs:
    for r in p.runs:
     if r.text.strip():
      if r.font.size: sizes['table'].append(r.font.size.pt)
      else: unset['table']+=1
role_counts=[5,4,3]; role_total=sum(role_counts); demands=[9,8,7]; proportions=[round(x/role_total*100,1) for x in demands]
weights=[25,25,20,15,15]; options={'A':[3,3,2,5,5],'B':[4,5,5,4,3],'C':[5,4,4,3,2]}; weighted={k:sum(x*w for x,w in zip(v,weights)) for k,v in options.items()}; scores={k:weighted[k]/5 for k in options}
minutes=[20,45,60,45,50,20]; total_minutes=sum(minutes)
expected=['9 ÷ 12 × 100% = 75.0%','8 ÷ 12 × 100% = 66.7%','7 ÷ 12 × 100% = 58.3%','25 + 25 + 20 + 15 + 15 = 100%','340÷5 = 68.0','430÷5 = 86.0','380÷5 = 76.0','240 分鐘']
checks={
 'zip_integrity':bad is None,
 'python_docx_open':True,
 'table_count_27':len(d.tables)==27,
 'expected_fact_ids_F001_F083':all(f'F{i:03d}' in text for i in range(1,84)),
 'no_empty_table_cells':len(empty_cells)==0,
 'body_min_font_10_5':bool(sizes['body']) and min(sizes['body'])>=10.5 and unset['body']==0,
 'table_min_font_9_5':bool(sizes['table']) and min(sizes['table'])>=9.5 and unset['table']==0,
 'formula_strings_present':all(s in text for s in expected),
 'roles_sum_12':role_total==12,
 'sample_proportions_recomputed':proportions==[75.0,66.7,58.3],
 'weights_sum_100':sum(weights)==100,
 'weighted_matrix_recomputed':weighted=={'A':340,'B':430,'C':380} and scores=={'A':68.0,'B':86.0,'C':76.0},
 'course_minutes_recomputed':total_minutes==240,
 'sample_synthetic_boundary_visible':all(token in text for token in ['SAMPLE','SYNTHETIC','REAL','HOLD']),
 'candidate_sha_matches_build_receipt':json.loads((qa/'BUILD_RECEIPT_R09.json').read_text(encoding='utf-8'))['output_sha256']==sha,
}
page_files=sorted(render.glob('page-*.png'))
page_meta=[]
for f in page_files:
 with Image.open(f) as im: page_meta.append({'page':int(re.search(r'page-(\d+)',f.name).group(1)),'sha256':hashlib.sha256(f.read_bytes()).hexdigest(),'width':im.width,'height':im.height})
checks['28_page_renders']=len(page_files)==28 and [x['page'] for x in page_meta]==list(range(1,29))
status='TECHNICAL_PASS_SEMANTIC_PENDING' if all(checks.values()) else 'TECHNICAL_FAIL'
receipt={
 'receipt_id':'0011_I02-E02_R09_DETERMINISTIC_TECHNICAL_QA',
 'observed_at_utc':datetime.now(timezone.utc).isoformat(),
 'document_id':'0011','requirement_id':'I02-E02','artifact_state':'CANDIDATE_SAMPLE_SYNTHETIC_NOT_ACCEPTED',
 'artifact_path':str(docx),'artifact_sha256':sha,'artifact_bytes':len(b),
 'official_requirement_pdf_sha256':'805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709',
 'blueprint_file_sha256':'406d0495365d0970dcf8113a42842663b8a54372a4e204cc492b7c3f8156ef3e',
 'blueprint_row_exact_line_sha256':'115b34f20b7f8106f5aa45d06da39ecae99a85a75f5276b79970911b428b32ec',
 'checks':checks,'tables':len(d.tables),'table_rows':sum(len(t.rows) for t in d.tables),'empty_cells':empty_cells,
 'minimum_direct_run_font_pt':{'body':min(sizes['body']) if sizes['body'] else None,'table':min(sizes['table']) if sizes['table'] else None},
 'formula_recalculation':{'role_counts':role_counts,'role_total':role_total,'demand_counts':demands,'percentages':proportions,'weights':weights,'weights_total':sum(weights),'weighted_points':weighted,'standardized_scores':scores,'course_minutes':total_minutes},
 'fact_id_coverage':'F001-F083 present',
 'semantic_review':'PENDING; OpenCode 429 cooldown route; no PASS asserted',
 'visual_review':'RECORDED_SEPARATELY; all 28 pages directly inspected',
 'formal_acceptance':'UNKNOWN',
 'status':status}
out=qa/'technical_qa_R09.json'; out.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'status':status,'artifact_sha256':sha,'checks':checks,'min_font':receipt['minimum_direct_run_font_pt'],'page_count':len(page_files),'receipt':str(out)},ensure_ascii=True,indent=2))
if not all(checks.values()): sys.exit(2)

