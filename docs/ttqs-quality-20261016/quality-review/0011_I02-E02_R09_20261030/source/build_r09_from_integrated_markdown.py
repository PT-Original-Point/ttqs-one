from pathlib import Path
import hashlib, json, re, sys
from datetime import datetime, timezone
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0011_I02-E02_LUNA_TEXT_DRAFT_R01_20261029')
SRC = BASE / 'worker' / 'integrated_draft_R07.md'
WORKER_RECEIPT = BASE / 'worker' / 'luna_subagent_draft_R01.receipt.json'
INTEGRATION_RECEIPT = BASE / 'worker' / 'integration_receipt_R07.json'
OUT = BASE / 'candidate' / '0011_I02-E02_R09.docx'
QA = BASE / 'qa' / 'BUILD_RECEIPT_R09.json'
BLUEPRINT = Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\CONTROL\BLUEPRINT_CURRENT.jsonl')

FONT = 'Microsoft JhengHei'
NAVY = '17365D'
BLUE = '1F4E78'
PALE = 'EAF1F7'
ALT = 'F5F8FB'
GRID = 'B7C9D6'


def sha(data): return hashlib.sha256(data).hexdigest()
def set_east_asia(run, name=FONT):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts'); rpr.insert(0, rfonts)
    for attr in ('ascii','hAnsi','eastAsia','cs'):
        rfonts.set(qn('w:'+attr), name)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn('w:shd'))
    if shd is None: shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill); shd.set(qn('w:val'), 'clear')

def set_cell_margins(cell, top=75, start=85, bottom=75, end=85):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); mar = tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar = OxmlElement('w:tcMar'); tcPr.append(mar)
    for side, value in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = mar.find(qn('w:'+side))
        if node is None: node=OxmlElement('w:'+side); mar.append(node)
        node.set(qn('w:w'),str(value)); node.set(qn('w:type'),'dxa')

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr(); node=OxmlElement('w:tblHeader'); node.set(qn('w:val'),'true'); trPr.append(node)

def clear_paragraph_borders(ppr):
    if ppr is None: return
    for node in list(ppr.findall(qn('w:pBdr'))): ppr.remove(node)

def switch_section_orientation(doc, landscape):
    s=doc.add_section(WD_SECTION.NEW_PAGE)
    # A minimal section-break paragraph prevents a full blank page when the preceding table ends near the margin.
    bp=doc.paragraphs[-1]
    bp.paragraph_format.space_before=Pt(0); bp.paragraph_format.space_after=Pt(0)
    bp.paragraph_format.line_spacing=Pt(1); bp.paragraph_format.keep_with_next=False; bp.paragraph_format.keep_together=False
    for run in bp.runs: run.font.size=Pt(1)
    s.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    s.page_width=Inches(11.69 if landscape else 8.27)
    s.page_height=Inches(8.27 if landscape else 11.69)
    s.top_margin=Inches(.55 if landscape else .62); s.bottom_margin=Inches(.55 if landscape else .65)
    s.left_margin=Inches(.55); s.right_margin=Inches(.55)
    s.header_distance=Inches(.25); s.footer_distance=Inches(.28)
    return s

def add_page_field(paragraph):
    r=paragraph.add_run(); set_east_asia(r); r.font.size=Pt(9); r.font.color.rgb=RGBColor(90,105,120)
    begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin')
    instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text=' PAGE '
    separate=OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'),'separate')
    text=OxmlElement('w:t'); text.text='1'
    end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end')
    r._r.extend([begin,instr,separate,text,end])

def add_inline(paragraph, text, size=None, color=None):
    # Preserve simple Markdown emphasis while treating all content as literal text.
    pat=re.compile(r'(\*\*.+?\*\*|`.+?`|\*[^*]+\*)')
    pos=0
    for m in pat.finditer(text):
        if m.start()>pos: _run(paragraph,text[pos:m.start()],False,False,False,size,color)
        token=m.group(0)
        if token.startswith('**'): _run(paragraph,token[2:-2],True,False,False,size,color)
        elif token.startswith('`'): _run(paragraph,token[1:-1],False,False,True,size,color)
        else: _run(paragraph,token[1:-1],False,True,False,size,color)
        pos=m.end()
    if pos<len(text): _run(paragraph,text[pos:],False,False,False,size,color)

def _run(paragraph, text, bold=False, italic=False, code=False, size=None, color=None):
    # Markdown links become readable labels followed by their target.
    for chunk in re.split(r'(\[[^\]]+\]\([^)]+\))', text):
        if not chunk: continue
        link=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',chunk)
        value=f'{link.group(1)} ({link.group(2)})' if link else chunk
        r=paragraph.add_run(value); set_east_asia(r,'Consolas' if code else FONT)
        r.bold=bold; r.italic=italic; r.font.size=Pt(size or 10.5)
        if color: r.font.color.rgb=RGBColor.from_string(color)
        if code: r.font.color.rgb=RGBColor.from_string('4B5563')

def split_md_row(s):
    s=s.strip()
    if s.startswith('|'): s=s[1:]
    if s.endswith('|'): s=s[:-1]
    out=[]; cur=[]; escaped=False
    for ch in s:
        if ch=='|' and not escaped: out.append(''.join(cur).strip()); cur=[]
        else:
            if ch=='|' and escaped and cur: cur[-1]='|'
            else: cur.append(ch)
        escaped=(ch=='\\' and not escaped)
        if ch!='\\': escaped=False
    out.append(''.join(cur).strip())
    return out

def is_table_start(lines,i):
    if i+1>=len(lines) or not lines[i].strip().startswith('|'): return False
    sep=lines[i+1].strip()
    if not sep.startswith('|'): return False
    cells=split_md_row(sep)
    return len(cells)>=2 and all(re.fullmatch(r':?-{3,}:?',c.replace(' ','')) for c in cells)

def table_block(lines,i):
    headers=split_md_row(lines[i]); rows=[]; j=i+2
    while j<len(lines) and lines[j].strip().startswith('|'):
        vals=split_md_row(lines[j]);
        if len(vals)<len(headers): vals += ['']*(len(headers)-len(vals))
        if len(vals)>len(headers): vals=vals[:len(headers)-1]+[' | '.join(vals[len(headers)-1:])]
        rows.append(vals); j+=1
    return headers,rows,j

def style_para(p, size=10.5, before=0, after=5, line=1.12, color=None):
    p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=line
    for r in p.runs:
        set_east_asia(r); r.font.size=Pt(size)
        if color: r.font.color.rgb=RGBColor.from_string(color)

def add_table(doc, headers, rows, available_width):
    n=len(headers); t=doc.add_table(rows=1, cols=n); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    t.style='Table Grid'
    if n==1: widths=[available_width]
    elif n==2: widths=[Inches(1.7), available_width-Inches(1.7)]
    elif n==3: widths=[Inches(1.35),(available_width-Inches(1.35))/2,(available_width-Inches(1.35))/2]
    elif n==4: widths=[Inches(1.05)]+[(available_width-Inches(1.05))/3]*3
    else: widths=[Inches(0.82)]+[(available_width-Inches(0.82))/(n-1)]*(n-1)
    for idx,h in enumerate(headers):
        c=t.rows[0].cells[idx]; c.width=widths[idx]; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_shading(c,NAVY); set_cell_margins(c,85,70,85,70)
        p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
        add_inline(p,h,9.5,'FFFFFF')
        for r in p.runs: r.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    set_repeat_table_header(t.rows[0])
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for ci,val in enumerate(row):
            c=cells[ci]; c.width=widths[ci]; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP; set_cell_margins(c)
            if ri%2==1: set_cell_shading(c,ALT)
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.03
            val=val.replace('<br>','\n').replace('<br/>','\n')
            for pi,part in enumerate(val.split('\n')):
                if pi: p.add_run().add_break()
                add_inline(p,part,9.5)
    # Set table grid widths and prevent clipped edge spacing.
    tblPr=t._tbl.tblPr
    layout=tblPr.find(qn('w:tblLayout'))
    if layout is None: layout=OxmlElement('w:tblLayout'); tblPr.append(layout)
    layout.set(qn('w:type'),'fixed')
    for row in t.rows:
        trPr=row._tr.get_or_add_trPr(); cant=OxmlElement('w:cantSplit'); trPr.append(cant)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t

def main():
    src_bytes=SRC.read_bytes(); src_hash=sha(src_bytes)
    worker=json.loads(WORKER_RECEIPT.read_text(encoding='utf-8-sig'))
    integration=json.loads(INTEGRATION_RECEIPT.read_text(encoding='utf-8-sig'))
    if integration.get('integrated_draft_sha256')!=src_hash: raise SystemExit('integration receipt source SHA mismatch')
    if worker.get('output_sha256')!=integration.get('source_draft_sha256'): raise SystemExit('worker source lineage mismatch')
    lines=SRC.read_text(encoding='utf-8-sig').splitlines()
    blueprint_bytes=BLUEPRINT.read_bytes(); blueprint_hash=sha(blueprint_bytes)
    bp_rows=[json.loads(x) for x in BLUEPRINT.read_text(encoding='utf-8-sig').splitlines() if x.strip()]
    row=next(x for x in bp_rows if x.get('deliverable_id')=='0011')
    if row.get('requirement_id')!='I02-E02' or row.get('official_requirement_text')!='與訓練政策有關之訓練產品與市場定位。':
        raise SystemExit('current unique Blueprint mapping mismatch')
    row_line=next(x.encode('utf-8') for x in BLUEPRINT.read_text(encoding='utf-8-sig').splitlines() if json.loads(x).get('deliverable_id')=='0011')
    row_line_hash=sha(row_line)
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
    sec.top_margin=Inches(.62); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.55); sec.right_margin=Inches(.55)
    sec.header_distance=Inches(.28); sec.footer_distance=Inches(.30)
    available=sec.page_width-sec.left_margin-sec.right_margin
    styles=doc.styles
    clear_paragraph_borders(styles['Title']._element.pPr)
    normal=styles['Normal']; normal.font.name=FONT; normal.font.size=Pt(10.5); normal.font.color.rgb=RGBColor.from_string('202A33')
    normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.12
    normal._element.rPr.rFonts.set(qn('w:eastAsia'),FONT)
    for name,size,col in [('Title',22,NAVY),('Heading 1',15,NAVY),('Heading 2',12,BLUE),('Heading 3',11,'385D7A')]:
        s=styles[name]; s.font.name=FONT; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(col); s._element.rPr.rFonts.set(qn('w:eastAsia'),FONT)
        s.paragraph_format.keep_with_next=True; s.paragraph_format.space_before=Pt(10 if name!='Title' else 0); s.paragraph_format.space_after=Pt(5)
    doc.core_properties.title='訓練產品與市場定位管理'
    doc.core_properties.subject='0011 I02-E02 SAMPLE SYNTHETIC 制度草案'
    doc.core_properties.author='TTQS ONE candidate build'
    doc.core_properties.keywords='TTQS, I02-E02, SAMPLE, SYNTHETIC, CANDIDATE'
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    add_inline(footer,'0011  SAMPLE／SYNTHETIC 候選  ·  頁 ',9,'5A6978'); add_page_field(footer)
    # Preserve source order and content; use a clean title derived from the source heading.
    i=0; table_count=0; paragraph_count=0; headings=[]; first_heading=True; table_col_counts=[]; landscape_form=False; current_form=None
    while i<len(lines):
        line=lines[i].rstrip()
        if not line.strip(): i+=1; continue
        if is_table_start(lines,i):
            h,r,j=table_block(lines,i)
            # Keep the identified table headers with their first data row; R07 visual QA found these two orphan headers.
            normalized=tuple(x.strip() for x in h)
            if normalized == ('分級','判定條件','立即措施','關閉條件'):
                doc.add_page_break()
            current=doc.sections[-1]; current_width=current.page_width-current.left_margin-current.right_margin
            add_table(doc,h,r,current_width); table_count+=1; table_col_counts.append(len(h)); i=j
            completed_form=current_form; current_form=None
            if landscape_form and completed_form=='E':
                switch_section_orientation(doc,False); landscape_form=False
            continue
        m=re.match(r'^(#{1,6})\s+(.*)$',line)
        if m:
            level=len(m.group(1)); text=m.group(2).strip()
            if level==1 and first_heading:
                title=text
                title=re.sub(r'^SAMPLE\s*制度草案\s*[｜|]\s*','',title)
                p=doc.add_paragraph(style='Title'); clear_paragraph_borders(p._p.get_or_add_pPr()); add_inline(p,title,22,NAVY)
                sub=doc.add_paragraph(); sub.paragraph_format.space_after=Pt(11)
                add_inline(sub,'SAMPLE SYNTHETIC 制度草案',11,BLUE)
                for r0 in sub.runs: r0.bold=True
                first_heading=False
            else:
                if level==2 and text.startswith('8.'):
                    current_form='section8'
                    if not landscape_form:
                        switch_section_orientation(doc,True); landscape_form=True
                elif level>=3 and text.startswith(('表單 A','表單 B','表單 C','表單 D','表單 E')):
                    current_form=text[3] if len(text)>3 else None
                    if current_form in ('B','C') and not landscape_form:
                        switch_section_orientation(doc,True); landscape_form=True
                    p=doc.add_paragraph(style=f'Heading {min(level,4)-1}')
                    p.paragraph_format.page_break_before=True
                    add_inline(p,text,15 if min(level,4)==2 else 12 if min(level,4)==3 else 11,NAVY if min(level,4)==2 else BLUE)
                    headings.append(text); paragraph_count+=1; i+=1; continue
                wlevel=min(level,4) # ## => Word Heading 1, ### => Heading 2
                p=doc.add_paragraph(style=f'Heading {wlevel-1}')
                add_inline(p,text,15 if wlevel==2 else 12 if wlevel==3 else 11,NAVY if wlevel==2 else BLUE)
                headings.append(text)
            paragraph_count+=1; i+=1; continue
        if line.strip() == '**表 7-5｜需求—職能—目標—評量示例**':
            # Keep the table caption with its table; R08 left this caption alone at the page foot.
            doc.add_page_break()
        if line.strip() in ('---','***','___'):
            i+=1; continue
        if line.lstrip().startswith('>'):
            qtext=line.lstrip()[1:].strip(); p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.22); p.paragraph_format.right_indent=Inches(.12); p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(7)
            add_inline(p,qtext,10.5,'385D7A'); paragraph_count+=1; i+=1; continue
        b=re.match(r'^(\s*)[-*+]\s+(.*)$',line)
        num=re.match(r'^\s*\d+[.)]\s+(.*)$',line)
        if b or num:
            content=b.group(2) if b else num.group(1)
            p=doc.add_paragraph(style='List Bullet' if b else 'List Number'); p.paragraph_format.left_indent=Inches(.28+(len(b.group(1))//2*.18 if b else 0)); p.paragraph_format.first_line_indent=Inches(-.15); p.paragraph_format.space_after=Pt(3)
            add_inline(p,content,10.5); paragraph_count+=1; i+=1; continue
        p=doc.add_paragraph(); add_inline(p,line,10.5); paragraph_count+=1; i+=1
    if table_count!=27: raise SystemExit(f'expected 27 source tables, parsed {table_count}')
    OUT.parent.mkdir(parents=True,exist_ok=True); doc.save(OUT)
    out_bytes=OUT.read_bytes()
    receipt={
      'status':'DOCX_BUILT_CANDIDATE_NOT_ACCEPTED',
      'deliverable_id':'0011','requirement_id':'I02-E02',
      'official_requirement_text':row['official_requirement_text'],
      'official_source_pdf_sha256':worker['official_source_pdf_sha256'],
      'worker_source_path':str(BASE / 'worker' / 'luna_subagent_draft_R01.md'),'worker_source_sha256':worker['output_sha256'],
      'worker_receipt_sha256':sha(WORKER_RECEIPT.read_bytes()),
      'integrated_draft_path':str(SRC),'integrated_draft_sha256':src_hash,
      'integration_receipt_sha256':sha(INTEGRATION_RECEIPT.read_bytes()),
      'integration_owner':'Codex integrator; completed form examples use existing CASE-S-001 Fact IDs only',
      'blueprint_path':str(BLUEPRINT),'blueprint_file_sha256':blueprint_hash,
      'blueprint_row_exact_line_sha256':row_line_hash,
      'blueprint_receipt_row_sha256_reported_by_worker':worker['source_blueprint_row_sha256'],
      'blueprint_receipt_row_hash_method_match':'NOT_ASSERTED',
      'output_path':str(OUT),'output_sha256':sha(out_bytes),'output_bytes':len(out_bytes),
      'builder_script_sha256':sha(Path(__file__).read_bytes()),
      'build_time_utc':datetime.now(timezone.utc).isoformat(),
      'format':'A4 portrait; Section 8 forms A-E in one landscape block; each reusable form starts on a fresh landscape page; Microsoft JhengHei; body 10.5pt; table 9.5pt; repeated table header; page field footer; built-in Title border removed',
      'source_parser_counts':{'markdown_tables':table_count,'table_columns':table_col_counts,'paragraphs_and_headings':paragraph_count,'headings':headings},
      'independent_semantic_review':'PENDING',
      'visual_qa':'PENDING',
      'formal_acceptance':'UNKNOWN'
    }
    QA.parent.mkdir(parents=True,exist_ok=True); QA.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'output_sha256':receipt['output_sha256'],'output_bytes':len(out_bytes),'table_count':table_count,'paragraph_count':paragraph_count,'blueprint_sha256':blueprint_hash,'blueprint_row_line_sha256':row_line_hash,'build_receipt':str(QA)},ensure_ascii=False))

if __name__=='__main__': main()
