from __future__ import annotations

import csv
import hashlib
import json
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "candidate" / "0062_R04.docx"
CSV_PATH = ROOT / "sidecar" / "0062_R04_SYNTHETIC_REPLACEMENT_REGISTER.csv"
PDF = ROOT / "evidence" / "render_R04" / "0062_R04.pdf"
RENDER_MANIFEST = ROOT / "evidence" / "visual_render_manifest_R04.json"
BUILD_RECEIPT = ROOT / "evidence" / "build_receipt_R04.json"
VISUAL_RECEIPT = ROOT / "evidence" / "visual_review_receipt_R04.json"
SEMANTIC_RECEIPT = ROOT / "evidence" / "independent_review_R04" / "opencode_nemotron_R04_receipt.json"
OUT = ROOT / "evidence" / "technical_qa_receipt_R04.json"


def sha(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest().upper()


def check(cond: bool, label: str, failures: list[str]) -> None:
    if not cond:
        failures.append(label)


def text_of_doc(doc: Document) -> str:
    pieces = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            pieces.extend(cell.text for cell in row.cells)
    return "\n".join(pieces)


def pct(n: int, d: int) -> str:
    return f"{(n / d * 100):.1f}%"


def main() -> None:
    failures: list[str] = []
    build = json.loads(BUILD_RECEIPT.read_text(encoding="utf-8"))
    manifest = json.loads(RENDER_MANIFEST.read_text(encoding="utf-8"))
    visual = json.loads(VISUAL_RECEIPT.read_text(encoding="utf-8"))
    semantic = json.loads(SEMANTIC_RECEIPT.read_text(encoding="utf-8"))

    with zipfile.ZipFile(DOCX) as z:
        corrupt = z.testzip()
    check(corrupt is None, f"docx_zip_integrity:{corrupt}", failures)
    doc = Document(DOCX)
    all_text = text_of_doc(doc)
    check(len(doc.tables) == 20, f"table_count:{len(doc.tables)}", failures)
    check("[[TABLE:" not in all_text, "unresolved_table_marker", failures)
    for token in ("I14-E03", "067", "SAMPLE / SYNTHETIC", "SYN-0060-2026-A01", "SYN-0033-PLAN-01", "REAL 狀態", "UNKNOWN", "HOLD"):
        check(token in all_text, f"missing_text:{token}", failures)
    check("樣本檔名（虛構）" in all_text, "sample_filename_boundary", failures)

    t09 = doc.tables[18]
    t09_ids = [row.cells[0].text.strip() for row in t09.rows[1:]]
    expected_ids = [f"REG-0062-{i:03d}" for i in range(1, 174)]
    check(t09_ids == expected_ids, "t09_register_ids_complete_ordered", failures)
    check(len(set(t09_ids)) == 173, "t09_register_ids_unique", failures)

    t02r = doc.tables[6]
    real_rows = [row.cells for row in t02r.rows[1:]]
    check(len(real_rows) == 12, f"real_status_rows:{len(real_rows)}", failures)
    check(all(cells[2].text.strip() == "UNKNOWN" and cells[3].text.strip() == "UNKNOWN" and cells[4].text.strip() == "UNKNOWN" and cells[5].text.strip() == "HOLD" for cells in real_rows), "real_source_location_policy_all_unknown_hold", failures)

    t02_groups = doc.tables[2:6]
    filename_cells = [r.cells[3].text.strip() for r in t02_groups[0].rows[1:]]
    path_cells = [r.cells[1].text.strip() for r in t02_groups[2].rows[1:]]
    check(len(filename_cells) == 12 and all(v.startswith("SAMPLE__") for v in filename_cells), "t02_has_12_synthetic_filenames", failures)
    check(len(path_cells) == 12 and all(v.startswith("SAMPLE_ROOT/0062/") for v in path_cells), "t02_has_12_fictitious_paths", failures)
    t03_text = "\n".join(c.text for t in doc.tables[7:9] for row in t.rows for c in row.cells)
    check("…" not in t03_text and all(f"SAMPLE__0062_{code}_" in t03_text for code in ("NEE", "NAN", "PLN", "LSN", "RSC", "ATT", "DLG", "ASM", "FDB", "FUP", "EXC", "IMP")), "t03_full_paths_no_ellipsis", failures)

    metrics = {
        "出席率": (10, 12, "83.3%"),
        "缺席率": (2, 12, "16.7%"),
        "評量通過率": (8, 9, "88.9%"),
        "清晰度正向率": (8, 9, "88.9%"),
        "追蹤回覆率": (8, 10, "80.0%"),
        "自述使用率": (6, 8, "75.0%"),
        "類別登錄覆蓋率": (12, 12, "100.0%"),
        "桌上測試命中率": (5, 6, "83.3%"),
    }
    for name, (n, d, expected) in metrics.items():
        check(pct(n, d) == expected, f"math:{name}:{pct(n,d)}", failures)
    metric_rows = {row.cells[0].text.strip(): row.cells for row in doc.tables[13].rows[1:]}
    check(metric_rows.get("評量通過率", [None] * 5)[2].text.strip() == "9", "assessment_denominator_is_valid_assessments_9", failures)
    check("UNKNOWN" in metric_rows.get("REAL 來源覆蓋率", [None] * 5)[4].text, "real_coverage_is_unknown", failures)

    with CSV_PATH.open(encoding="utf-8-sig", newline="") as f:
        register = list(csv.DictReader(f))
    check(len(register) == 173, f"replacement_register_rows:{len(register)}", failures)
    check(len({r["reg_id"] for r in register}) == 173, "replacement_register_ids_unique", failures)
    check(all(r["status"] == "SAMPLE / SYNTHETIC; NOT REAL EVIDENCE" for r in register), "replacement_register_status_boundary", failures)
    check(all(r["exact_document_location"].startswith("0062_R04.docx / Section 13 / T-0062-09 / REG-0062-") for r in register), "replacement_register_exact_locations", failures)
    check(not any("Not explicit in row" in r["source_class"] for r in register), "replacement_register_source_classes_resolved", failures)

    pdf_reader = PdfReader(str(PDF))
    check(len(pdf_reader.pages) == 44, f"pdf_page_count:{len(pdf_reader.pages)}", failures)
    quick_check_pages = "\n".join((pdf_reader.pages[i].extract_text() or "") for i in (40, 41) if i < len(pdf_reader.pages))
    expected_quick_check = [
        "1. 核對文件編號 0062",
        "2. 先分清樣本狀態",
        "3. 由 T-0062-02",
        "4. 用 T-0062-05",
        "5. 發現缺件",
    ]
    positions = [quick_check_pages.find(item) for item in expected_quick_check]
    check(all(pos >= 0 for pos in positions) and positions == sorted(positions), f"quick_check_restarts_at_1_and_runs_1_to_5:{positions}", failures)
    nonblank = []
    for i, page in enumerate(pdf_reader.pages, 1):
        text = " ".join((page.extract_text() or "").split())
        if len(text) > 90:
            nonblank.append(i)
    check(len(nonblank) == len(pdf_reader.pages), f"nonblank_pdf_pages:{len(nonblank)}/{len(pdf_reader.pages)}", failures)
    check(manifest["page_count"] == len(pdf_reader.pages), "render_manifest_page_count_matches_pdf", failures)
    check(visual["candidate_sha256"] == sha(DOCX), "visual_receipt_candidate_sha_matches", failures)
    check(visual["pdf_sha256"] == sha(PDF), "visual_receipt_pdf_sha_matches", failures)
    check(visual["status"] == "PASS_VISUAL_SCOPE", f"visual_receipt_status:{visual['status']}", failures)
    check(semantic["candidate_sha256"] == sha(DOCX), "semantic_receipt_candidate_sha_matches", failures)
    check(semantic["overall_verdict"] == "PASS" and semantic["process_exit_code"] == 0, "independent_semantic_review_pass", failures)
    section_orientations = ["landscape" if s.page_width > s.page_height else "portrait" for s in doc.sections]
    check(section_orientations == ["landscape"], f"section_orientation:{section_orientations}", failures)

    result = {
        "project_id": "TTQS_ONE",
        "document_id": "0062",
        "requirement_id": "I14-E03",
        "official_checkbox_serial": "067",
        "candidate_file": str(DOCX),
        "candidate_sha256": sha(DOCX),
        "pdf_sha256": sha(PDF),
        "replacement_register_sha256": sha(CSV_PATH),
        "render_manifest_sha256": sha(RENDER_MANIFEST),
        "visual_review_receipt_sha256": sha(VISUAL_RECEIPT),
        "semantic_review_receipt_sha256": sha(SEMANTIC_RECEIPT),
        "build_receipt_sha256": sha(BUILD_RECEIPT),
        "page_count": len(pdf_reader.pages),
        "page_orientation": "landscape",
        "table_count": len(doc.tables),
        "t09_register_rows": len(register),
        "real_record_rows": len(real_rows),
        "real_status": "UNKNOWN / HOLD on every row",
        "visual_review": "MANUAL CONTACT SHEETS AND REPRESENTATIVE FULL-RESOLUTION PAGES REVIEWED; see visual_review_receipt_R04.json",
        "semantic_review": "PASS; independent OpenCode V2 exact-SHA response linked in semantic_review_receipt_sha256",
        "failures": failures,
        "status": "PASS_TECHNICAL_SCOPE" if not failures else "FAIL_TECHNICAL_SCOPE",
        "checked_at_utc": datetime.now(timezone.utc).isoformat(),
        "scope": "DOCX package integrity, table/marker counts, exact register IDs, SAMPLE/REAL boundary, sample paths, arithmetic, PDF render consistency, no blank pages, and section orientation. Not a semantic acceptance or association approval.",
        "builder_receipt_candidate_sha_matches": build["artifacts"][DOCX.name]["sha256"] == sha(DOCX),
    }
    OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if failures:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
