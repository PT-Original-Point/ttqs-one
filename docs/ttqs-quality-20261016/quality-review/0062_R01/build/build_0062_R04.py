from __future__ import annotations

import csv
import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "inputs"
EVIDENCE = ROOT / "evidence"
BODY_PATH = INPUT / "0062_BODY_R01_HUMAN_READY.md"
BLUEPRINT_PATH = INPUT / "0062_UNIQUE_BLUEPRINT_R04.md"
TASK_PATH = INPUT / "0062_DOCUMENT_WRITER_TASK_R01.md"
NORMALIZATION_PATH = EVIDENCE / "response_normalization_R01.json"
CANDIDATE_DIR = ROOT / "candidate"
SIDECAR_DIR = ROOT / "sidecar"
OUTPUT_DOCX = CANDIDATE_DIR / "0062_R04.docx"
OUTPUT_REGISTER = SIDECAR_DIR / "0062_R04_SYNTHETIC_REPLACEMENT_REGISTER.csv"
OUTPUT_RECEIPT = EVIDENCE / "build_receipt_R04.json"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest().upper()


def fresh_numbering_instance(doc: Document, style_name: str = "List Number", start: int = 1) -> int:
    style_ppr = doc.styles[style_name]._element.find(qn("w:pPr"))
    style_numpr = style_ppr.find(qn("w:numPr")) if style_ppr is not None else None
    style_numid = style_numpr.find(qn("w:numId")) if style_numpr is not None else None
    if style_numid is None:
        raise ValueError(f"Numbering style has no numId: {style_name}")
    base_num_id = style_numid.get(qn("w:val"))
    numbering = doc.part.numbering_part.element
    base_num = next((n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == base_num_id), None)
    if base_num is None:
        raise ValueError(f"Numbering instance not found: {base_num_id}")
    abstract_id = base_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    new_id = max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_id)
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), str(start))
    override.append(start_override)
    num.append(override)
    numbering.append(num)
    return new_id


def set_paragraph_numbering_id(paragraph, num_id: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    previous = ppr.find(qn("w:numPr"))
    if previous is not None:
        ppr.remove(previous)
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numpr.append(ilvl)
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(num_id))
    numpr.append(numid)
    ppr.append(numpr)


def clean_inline(value: str) -> str:
    value = value.strip()
    value = re.sub(r"\*\*(.*?)\*\*", r"\1", value)
    value = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"\1", value)
    value = value.replace("`", "")
    value = value.replace("\\|", "|")
    return value


def split_md_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    cells: list[str] = []
    buf: list[str] = []
    escaped = False
    for ch in line:
        if ch == "|" and not escaped:
            cells.append(clean_inline("".join(buf)))
            buf = []
        else:
            buf.append(ch)
        escaped = (ch == "\\" and not escaped)
        if ch != "\\":
            escaped = False
    cells.append(clean_inline("".join(buf)))
    return cells


def is_separator_row(line: str) -> bool:
    cells = split_md_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{2,}:?", c.replace(" ", "")) for c in cells)


def parse_table_blocks(lines: list[str]) -> list[tuple[list[str], list[list[str]]]]:
    blocks: list[tuple[list[str], list[list[str]]]] = []
    i = 0
    while i < len(lines):
        if not lines[i].lstrip().startswith("|"):
            i += 1
            continue
        raw: list[str] = []
        while i < len(lines) and lines[i].lstrip().startswith("|"):
            raw.append(lines[i])
            i += 1
        if len(raw) < 2 or not is_separator_row(raw[1]):
            continue
        header = split_md_row(raw[0])
        rows = [split_md_row(line) for line in raw[2:] if not is_separator_row(line)]
        normalized = []
        for row in rows:
            if len(row) < len(header):
                row += [""] * (len(header) - len(row))
            elif len(row) > len(header):
                row = row[: len(header) - 1] + [" | ".join(row[len(header) - 1 :])]
            normalized.append(row)
        blocks.append((header, normalized))
    return blocks


def section_lines(text: str, heading_pattern: str, end_pattern: str | None = None) -> list[str]:
    lines = text.splitlines()
    start = next((i for i, line in enumerate(lines) if re.search(heading_pattern, line)), None)
    if start is None:
        raise ValueError(f"Blueprint section not found: {heading_pattern}")
    end = len(lines)
    if end_pattern:
        for i in range(start + 1, len(lines)):
            if re.search(end_pattern, lines[i]):
                end = i
                break
    return lines[start:end]


def first_table(text: str, heading_pattern: str, end_pattern: str | None = None) -> tuple[list[str], list[list[str]]]:
    blocks = parse_table_blocks(section_lines(text, heading_pattern, end_pattern))
    if not blocks:
        raise ValueError(f"No table found in section: {heading_pattern}")
    return blocks[0]


def extract_t09(blueprint: str) -> tuple[list[dict[str, str]], list[tuple[str, list[str], list[list[str]]]]]:
    lines = section_lines(blueprint, r"^## 十三、.*T-0062-09", r"^## 十四、")
    subheading = "13.0 未分類"
    records: list[dict[str, str]] = []
    display_blocks: list[tuple[str, list[str], list[list[str]]]] = []
    i = 0
    while i < len(lines):
        if lines[i].startswith("### "):
            subheading = clean_inline(lines[i][4:])
            i += 1
            continue
        if lines[i].lstrip().startswith("|"):
            raw: list[str] = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                raw.append(lines[i])
                i += 1
            if len(raw) < 2 or not is_separator_row(raw[1]):
                continue
            header = split_md_row(raw[0])
            rows = [split_md_row(line) for line in raw[2:] if not is_separator_row(line)]
            normalized: list[list[str]] = []
            for row in rows:
                if len(row) < len(header):
                    row += [""] * (len(header) - len(row))
                elif len(row) > len(header):
                    row = row[: len(header) - 1] + [" | ".join(row[len(header) - 1 :])]
                normalized.append(row)
                item = {header[j].strip().lower(): row[j] for j in range(len(header))}
                reg_id = next((v for k, v in item.items() if "reg_id" in k or "reg id" in k), "")
                if not reg_id:
                    raise ValueError(f"T-0062-09 row missing register ID: {row}")
                kind = next((v for k, v in item.items() if k in {"種類", "kind", "類型"}), "")
                value = next((v for k, v in item.items() if k in {"值", "value"}), "")
                source = next((v for k, v in item.items() if "來源" in k or k == "source"), "")
                source_display, source_class = classify_t09_source(subheading, source)
                records.append({
                    "reg_id": reg_id,
                    "subgroup": subheading,
                    "kind": kind,
                    "value": value,
                    "source": source,
                    "source_display": source_display,
                    "source_class": source_class,
                })
            display_blocks.append((subheading, header, normalized))
            continue
        i += 1
    records.sort(key=lambda r: int(re.search(r"(\d+)$", r["reg_id"]).group(1)))
    return records, display_blocks


def classify_t09_source(subgroup: str, source: str) -> tuple[str, str]:
    if source.strip():
        if "SRC" in source:
            return source, "SOURCE-DERIVED; 0060 candidate data or explicitly identified adoption"
        if "SYN" in source:
            return source, "SYNTHETIC"
        if "0062 推導" in source or "0062推導" in source:
            return source, "DERIVED; calculated or indexed in 0062"
        return source, "SOURCE REFERENCE; retain its stated scope"
    text = subgroup.casefold()
    if "syn" in text or "合成" in text:
        return "SYN（依分組標示）", "SYNTHETIC"
    if "src" in text or "來源" in text:
        return "SRC（依分組標示）", "SOURCE-DERIVED; scope is limited to the named source field"
    if "桌上測試" in text:
        return "SYN（合成桌上測試）", "SYNTHETIC TABLETOP TEST"
    if "衍生計算" in text or "數量" in text:
        return "0062 推導／彙總；不得回寫為 0060 原始事實", "DERIVED SUMMARY; based on synthetic 0062 records"
    return "依 Blueprint 分組說明；須維持樣本界線", "SUBGROUP-CONTEXT; verify before replacement"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=85, bottom=70, end=85) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def set_font(run, size: float, color: str = "243447", bold: bool | None = None) -> None:
    run.font.name = "Microsoft JhengHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def add_rich_text(paragraph, text: str, size: float = 10.5, color: str = "243447") -> None:
    token_re = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*)")
    pos = 0
    for match in token_re.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            set_font(run, size, color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, size, color, True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, size, "315B7D")
        else:
            run = paragraph.add_run(token[1:-1])
            set_font(run, size, color)
            run.italic = True
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_font(run, size, color)


def set_section(section, landscape: bool) -> None:
    if landscape:
        section.page_width = Inches(11.69)
        section.page_height = Inches(8.27)
        section.left_margin = Inches(0.42)
        section.right_margin = Inches(0.42)
        section.top_margin = Inches(0.52)
        section.bottom_margin = Inches(0.52)
    else:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.left_margin = Inches(0.78)
        section.right_margin = Inches(0.78)
        section.top_margin = Inches(0.72)
        section.bottom_margin = Inches(0.72)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.32)


def add_section_header(section) -> None:
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("原始點協會｜TTQS 文件工作稿　·　SAMPLE / SYNTHETIC")
    set_font(r, 8.2, "587085")
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("候選示範｜不得作為真實紀錄　　·　頁 ")
    set_font(r, 8.2, "587085")
    add_page_field(fp)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, 10, "164A72", True)


def table_rows_for(doc: Document, headers: list[str], rows: list[list[str]], caption: str,
                   weights: list[float] | None = None) -> None:
    if not headers:
        raise ValueError(f"Empty table headers for {caption}")
    add_caption(doc, caption)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    section = doc.sections[-1]
    usable_width = section.page_width - section.left_margin - section.right_margin
    if not weights or len(weights) != len(headers):
        weights = [1.0] * len(headers)
    total = sum(weights)
    widths = [usable_width * (w / total) for w in weights]
    for idx, (cell, value) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = widths[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "164A72")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_together = True
        r = p.add_run(value)
        set_font(r, 9.5, "FFFFFF", True)
    repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for ridx, row_data in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for idx in range(len(headers)):
            cell = row.cells[idx]
            cell.width = widths[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
            if ridx % 2:
                set_cell_shading(cell, "F1F5F8")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.04
            value = row_data[idx] if idx < len(row_data) else ""
            r = p.add_run(str(value))
            set_font(r, 9.5, "243447")


def table_section(doc: Document, tables: list[tuple[str, list[str], list[list[str]], list[float] | None]]) -> None:
    for caption, headers, rows, weights in tables:
        table_rows_for(doc, headers, rows, caption, weights)


def select_rows(headers: list[str], rows: list[list[str]], columns: list[str]) -> list[list[str]]:
    indices = [headers.index(column) for column in columns]
    return [[row[i] for i in indices] for row in rows]


def require_columns(headers: list[str], columns: list[str], table_name: str) -> None:
    missing = [c for c in columns if c not in headers]
    if missing:
        raise ValueError(f"{table_name} missing columns: {missing}; found {headers}")


def build_tables(blueprint: str) -> tuple[dict[str, list[tuple[str, list[str], list[list[str]], list[float] | None]]], list[dict[str, str]]]:
    table_specs: dict[str, list[tuple[str, list[str], list[list[str]], list[float] | None]]] = {}

    h, r = first_table(blueprint, r"^## 五、.*T-0062-01", r"^## 六、")
    a = ["列ID", "階段", "來源doc ID", "類別代碼來源", "類別描述來源", "附帶數據 fact"]
    b = ["列ID", "0060 row location", "下游 artifact_id", "擁有功能", "樣本狀態", "樣本決策"]
    require_columns(h, a + b[1:], "T-0062-01")
    table_specs["T-0062-01"] = [
        ("T-0062-01A｜需求階段、來源與事實依據", a, select_rows(h, r, a), [0.8, 1.4, 0.7, 1.35, 1.35, 1.5]),
        ("T-0062-01B｜來源定位、責任功能與樣本判定", b, select_rows(h, r, b), [0.8, 2.1, 1.2, 1.45, 0.8, 0.8]),
    ]

    h2, r2 = first_table(blueprint, r"^## 六、.*T-0062-02", r"^### `T-0062-02R`")
    cols2 = ["列ID", "artifact_id", "類別", "樣本檔名（虛構）", "格式", "代碼字串（來源）", "代碼fact", "描述fact", "樣本日期（來源）", "來源定位參照", "樣本存放路徑（虛構）", "ACCESS樣本碼", "RET樣本碼", "樣本驗證範圍", "樣本狀態", "樣本決策", "拒絕條件"]
    require_columns(h2, cols2, "T-0062-02")
    groups2 = [
        (["列ID", "artifact_id", "類別", "樣本檔名（虛構）", "格式"], [0.55, 0.95, 0.7, 3.4, 0.65], "A｜索引與樣本檔名"),
        (["列ID", "代碼字串（來源）", "代碼fact", "描述fact", "樣本日期（來源）", "來源定位參照"], [0.55, 3.1, 0.9, 0.9, 1.25, 1.4], "B｜代碼、來源事實與日期"),
        (["列ID", "樣本存放路徑（虛構）", "ACCESS樣本碼", "RET樣本碼", "樣本驗證範圍"], [0.55, 3.8, 0.8, 0.65, 2.3], "C｜虛構定位與驗證範圍"),
        (["列ID", "樣本狀態", "樣本決策", "拒絕條件"], [0.7, 1.2, 1.2, 5.0], "D｜樣本判定與退件條件"),
    ]
    table_specs["T-0062-02"] = [(f"T-0062-02{label}", cols, select_rows(h2, r2, cols), weights) for cols, weights, label in groups2]

    h2r, r2r = first_table(blueprint, r"^### `T-0062-02R`", r"^`T-0062-02` 拒絕條件")
    table_specs["T-0062-02R"] = [("T-0062-02R｜真實來源尚未核實時逐列維持 HOLD", h2r, r2r, [0.7, 1.1, 1.5, 1.5, 1.3, 2.0])]

    h3, r3 = first_table(blueprint, r"^## 七、.*T-0062-03", r"^## 八、")
    path_by_artifact = {row[1]: row[10] for row in r2}
    r3 = [row.copy() for row in r3]
    for row in r3:
        if "…" in row[4] and row[0] in path_by_artifact:
            row[4] = path_by_artifact[row[0]]
    t3a = ["artifact_id", "節點型別", "依據", "樣本檔名（虛構）", "路徑（虛構）", "版本"]
    t3b = ["artifact_id", "上游", "關係", "更正/備註", "樣本狀態"]
    require_columns(h3, t3a + t3b[1:], "T-0062-03")
    table_specs["T-0062-03"] = [
        ("T-0062-03A｜樣本描述節點與完整虛構定位", t3a, select_rows(h3, r3, t3a), [0.9, 1.25, 1.0, 2.5, 4.2, 0.6]),
        ("T-0062-03B｜上下游、關係與版本註記", t3b, select_rows(h3, r3, t3b), [1.0, 1.5, 1.3, 4.2, 0.9]),
    ]

    h4, r4 = first_table(blueprint, r"^## 八、.*T-0062-04", r"^## 九、")
    t4a = ["artifact_id", "樣本檔名（虛構）", "樣本存放路徑（虛構）", "定位型別", "ACCESS 樣本碼", "RET 樣本碼"]
    t4b = ["artifact_id", "樣本驗證功能", "樣本驗證日", "REAL 定位", "REAL 政策", "REAL 狀態"]
    require_columns(h4, t4a + t4b[1:], "T-0062-04")
    table_specs["T-0062-04"] = [
        ("T-0062-04A｜虛構檔名、定位與樣本控制碼", t4a, select_rows(h4, r4, t4a), [0.9, 2.6, 4.0, 1.0, 1.0, 0.8]),
        ("T-0062-04B｜樣本驗證與真實欄位狀態", t4b, select_rows(h4, r4, t4b), [1.0, 2.0, 1.2, 1.5, 1.5, 1.0]),
    ]

    h5, r5 = first_table(blueprint, r"^## 九、.*T-0062-05", r"^## 十、")
    t5a = ["測試ID", "測試日", "artifact_id", "預期樣本路徑（虛構）", "實測取回"]
    t5b = ["測試ID", "秒數", "結果", "限制"]
    require_columns(h5, t5a + t5b[1:], "T-0062-05")
    table_specs["T-0062-05"] = [
        ("T-0062-05A｜合成桌上取回測試定位與結果", t5a, select_rows(h5, r5, t5a), [1.15, 1.0, 1.0, 5.8, 1.0]),
        ("T-0062-05B｜測試時間與限制", t5b, select_rows(h5, r5, t5b), [1.3, 0.9, 1.0, 5.0]),
    ]

    h6, r6 = first_table(blueprint, r"^## 十、.*T-0062-06", r"^## 十一、")
    t6a = ["指標", "分子", "分母", "公式", "值"]
    t6b = ["指標", "單位", "來源 fact", "限制"]
    require_columns(h6, t6a + t6b[1:], "T-0062-06")
    table_specs["T-0062-06"] = [
        ("T-0062-06A｜逐項指標、分子分母與計算", t6a, select_rows(h6, r6, t6a), [1.7, 0.65, 0.65, 2.4, 1.0]),
        ("T-0062-06B｜單位、來源與適用限制", t6b, select_rows(h6, r6, t6b), [1.8, 0.7, 2.0, 3.5]),
    ]

    h7, r7 = first_table(blueprint, r"^## 十一、.*T-0062-07", r"^## 十二、")
    table_specs["T-0062-07"] = [("T-0062-07｜別名、事實值、來源與限制", h7, r7, [1.45, 2.25, 0.75, 1.1, 1.65, 2.0])]

    h8, r8 = first_table(blueprint, r"^## 十二、.*T-0062-08", r"^## 十三、")
    t8a = ["local_event_id", "source_scenario_id", "branch_code", "artifact_id", "原因", "功能"]
    t8b = ["local_event_id", "期限", "樣本狀態", "樣本決策", "解除條件", "拒絕條件"]
    require_columns(h8, t8a + t8b[1:], "T-0062-08")
    table_specs["T-0062-08"] = [
        ("T-0062-08A｜本地事件、來源情境與分支碼（欄位不得混用）", t8a, select_rows(h8, r8, t8a), [1.45, 1.5, 1.5, 1.25, 3.2, 1.4]),
        ("T-0062-08B｜期限、樣本狀態與處理門檻", t8b, select_rows(h8, r8, t8b), [1.6, 1.0, 1.0, 1.0, 3.0, 3.0]),
    ]

    t09_records, t09_blocks = extract_t09(blueprint)
    if len(t09_records) != 173 or len({x["reg_id"] for x in t09_records}) != 173:
        raise ValueError(f"T-0062-09 requires 173 unique rows, found {len(t09_records)}")
    expected_ids = [f"REG-0062-{i:03d}" for i in range(1, 174)]
    if [x["reg_id"] for x in t09_records] != expected_ids:
        raise ValueError("T-0062-09 IDs are not a complete 001..173 sequence")
    display_rows = [[x["reg_id"], x["subgroup"], x["kind"], x["value"], x["source_display"]] for x in t09_records]
    table_specs["T-0062-09"] = [("T-0062-09｜173 列合成替換登錄（每列一事實）", ["登錄ID", "子類別", "資料種類", "樣本值", "來源／備註"], display_rows, [1.35, 1.7, 1.35, 3.2, 2.3])]

    h10, r10 = first_table(blueprint, r"^## 十四、.*T-0062-10", r"^## 十五、")
    table_specs["T-0062-10"] = [("T-0062-10｜真實資料進件檢核（進件前不得解除 HOLD）", h10, r10, [1.25, 2.5, 1.35, 2.1, 2.3, 0.8])]
    return table_specs, t09_records


def replacement_action(kind: str, value: str) -> str:
    k = kind.lower()
    if any(x in k for x in ("檔名", "路徑", "定位")):
        return "以實際原件及有權來源核對後逐列替換；未核實前維持 SAMPLE／HOLD。"
    if any(x in k for x in ("計算", "率", "數量")):
        return "取得真實原始分子、分母與來源後重新計算；不得沿用樣本數值或推論成效。"
    if any(x in k for x in ("日期", "期限", "版本", "測試", "秒數", "結果", "事件", "情境")):
        return "以可追溯的真實形成紀錄逐項替換；不得將合成情境或桌上測試冒充實際執行。"
    if any(x in k for x in ("功能", "標籤", "角色", "權限")):
        return "依有權核定的職掌與權限文件替換；不得推定任命或授權。"
    if any(x in k for x in ("artifact_id", "代碼", "編號", "ID")):
        return "依協會核定的真實編碼規則重新對應，保留舊值與新值的可追溯關係。"
    if "來源" in k or "政策" in k or "保存" in k:
        return "以有權來源原件核實並記錄版本、日期與定位；未知時維持 UNKNOWN／HOLD。"
    return "以真實、可追溯原件逐項替換；無來源證據時維持 SAMPLE／HOLD。"


def write_replacement_register(records: list[dict[str, str]], blueprint_hash: str) -> None:
    SIDECAR_DIR.mkdir(parents=True, exist_ok=True)
    with OUTPUT_REGISTER.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=[
            "reg_id", "subgroup", "kind", "sample_value", "source_reference",
            "source_class", "exact_document_location", "status", "replacement_action",
        ])
        writer.writeheader()
        for record in records:
            source = record["source"]
            source_class = record["source_class"]
            writer.writerow({
                "reg_id": record["reg_id"],
                "subgroup": record["subgroup"],
                "kind": record["kind"],
                "sample_value": record["value"],
                "source_reference": source or f"0062_UNIQUE_BLUEPRINT_R04; subsection {record['subgroup']}; source classification derived from that named subgroup",
                "source_class": source_class,
                "exact_document_location": f"0062_R04.docx / Section 13 / T-0062-09 / {record['reg_id']}",
                "status": "SAMPLE / SYNTHETIC; NOT REAL EVIDENCE",
                "replacement_action": replacement_action(record["kind"], record["value"]),
            })


def add_table_for_marker(doc: Document, marker: str, tables: dict[str, list[tuple[str, list[str], list[list[str]], list[float] | None]]]) -> None:
    key = marker.strip()
    if key not in tables:
        raise ValueError(f"No table spec for marker {key}")
    table_section(doc, tables[key])


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(76)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TTQS 訓練機構版｜佐證工作稿")
    set_font(r, 13, "587085", True)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("完整的訓練文件建檔資料")
    set_font(r, 25, "164A72", True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("單一課程的文件清冊、定位與追溯示範")
    set_font(r, 14, "315B7D")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("SAMPLE / SYNTHETIC　｜　非協會真實紀錄")
    set_font(r, 18, "A33B35", True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("官方對照：I14-E03　｜　方框流水號：067　｜　文件編號：0062")
    set_font(r, 11, "243447")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(32)
    p.paragraph_format.left_indent = Inches(0.85)
    p.paragraph_format.right_indent = Inches(0.85)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("本工作稿以合成案例展示建檔方式。樣本檔名與 SAMPLE_ROOT 路徑均為虛構；12 類真實來源、定位、政策及權限仍為 UNKNOWN／HOLD。不得將本件作為實際開課、簽到、評量、核准或歸檔證據。")
    set_font(r, 12, "243447")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("候選示範文件｜正式接受與真實資料替換尚未完成")
    set_font(r, 10, "587085")
    doc.add_page_break()


def set_document_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft JhengHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("243447")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, color in (("Heading 1", 16, "164A72"), ("Heading 2", 13.2, "1E6683"), ("Heading 3", 11.5, "315B7D")):
        st = styles[name]
        st.font.name = "Microsoft JhengHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(10 if name == "Heading 1" else 7)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True


def build() -> dict[str, object]:
    for path in (BODY_PATH, BLUEPRINT_PATH, TASK_PATH, NORMALIZATION_PATH):
        if not path.is_file():
            raise FileNotFoundError(path)
    body = BODY_PATH.read_text(encoding="utf-8-sig")
    blueprint = BLUEPRINT_PATH.read_text(encoding="utf-8-sig")
    tables, t09_records = build_tables(blueprint)
    markers = re.findall(r"\[\[TABLE:(T-0062-[^\]]+)\]\]", body)
    expected = ["T-0062-01", "T-0062-02", "T-0062-02R", "T-0062-03", "T-0062-04", "T-0062-05", "T-0062-06", "T-0062-07", "T-0062-08", "T-0062-09", "T-0062-10"]
    if markers != expected:
        raise ValueError(f"Body table markers changed or reordered: {markers}")

    CANDIDATE_DIR.mkdir(parents=True, exist_ok=True)
    EVIDENCE.mkdir(parents=True, exist_ok=True)
    doc = Document()
    set_document_styles(doc)
    # Keep the data-heavy document in one landscape section. Narrative paragraphs
    # use indents to preserve a portrait-like reading measure without section gaps.
    set_section(doc.sections[0], True)
    add_section_header(doc.sections[0])
    props = doc.core_properties
    props.title = "完整的訓練文件建檔資料｜0062｜SAMPLE / SYNTHETIC"
    props.subject = "I14-E03 / official checkbox 067; candidate demonstration only"
    props.author = "原始點協會 TTQS 工作稿"
    props.keywords = "TTQS, 0062, I14-E03, 067, SAMPLE, SYNTHETIC"
    add_cover(doc)

    lines = body.splitlines()
    i = 0
    pending: list[str] = []
    quick_check_restart_pending = False
    quick_check_num_id: int | None = None

    def flush_paragraph() -> None:
        nonlocal pending
        text = " ".join(x.strip() for x in pending).strip()
        if text:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(1.05)
            p.paragraph_format.right_indent = Inches(1.05)
            add_rich_text(p, text, 10.5)
        pending = []

    while i < len(lines):
        raw = lines[i]
        stripped = raw.strip()
        if not stripped:
            flush_paragraph()
            i += 1
            continue
        marker_match = re.fullmatch(r"\[\[TABLE:(T-0062-[^\]]+)\]\]", stripped)
        if marker_match:
            flush_paragraph()
            add_table_for_marker(doc, marker_match.group(1), tables)
            i += 1
            continue
        if stripped.startswith("#"):
            flush_paragraph()
            n = len(stripped) - len(stripped.lstrip("#"))
            title = stripped[n:].strip()
            level = min(max(n, 1), 3)
            if level == 2:
                if "16. 快速查核順序" in title:
                    quick_check_restart_pending = True
                    quick_check_num_id = None
                else:
                    quick_check_restart_pending = False
                    quick_check_num_id = None
            p = doc.add_paragraph(style=f"Heading {level}")
            p.paragraph_format.left_indent = Inches(1.05)
            p.paragraph_format.right_indent = Inches(1.05)
            add_rich_text(p, title, {1: 16, 2: 13.2, 3: 11.5}[level], "164A72" if level == 1 else "1E6683")
            i += 1
            continue
        if stripped.startswith("- ") or re.match(r"^\d+[.)、]\s+", stripped):
            flush_paragraph()
            style = "List Bullet" if stripped.startswith("- ") else "List Number"
            value = stripped[2:] if stripped.startswith("- ") else re.sub(r"^\d+[.)、]\s+", "", stripped)
            p = doc.add_paragraph(style=style)
            if style == "List Number" and quick_check_restart_pending:
                quick_check_num_id = fresh_numbering_instance(doc, style, start=1)
                quick_check_restart_pending = False
            if style == "List Number" and quick_check_num_id is not None:
                set_paragraph_numbering_id(p, quick_check_num_id)
            p.paragraph_format.left_indent = Inches(1.25)
            p.paragraph_format.right_indent = Inches(1.05)
            add_rich_text(p, value, 10.2)
            i += 1
            continue
        if stripped == "---":
            flush_paragraph()
            i += 1
            continue
        pending.append(stripped)
        i += 1
    flush_paragraph()

    doc.save(OUTPUT_DOCX)
    write_replacement_register(t09_records, sha256(BLUEPRINT_PATH))
    docx_sha = sha256(OUTPUT_DOCX)
    register_sha = sha256(OUTPUT_REGISTER)
    receipt = {
        "project_id": "TTQS_ONE",
        "document_id": "0062",
        "official_requirement_id": "I14-E03",
        "official_checkbox_serial": "067",
        "official_evidence_title": "完整的訓練文件建檔資料",
        "candidate_status": "CANDIDATE; HUMAN_ACCEPTANCE_PENDING; REAL_EVIDENCE_NOT_ESTABLISHED",
        "builder": "Deterministic python-docx renderer from OpenCode v2 writer response and approved document-specific normalized body/Blueprint",
        "build_time_utc": datetime.now(timezone.utc).isoformat(),
        "inputs": {
            str(BODY_PATH.name): sha256(BODY_PATH),
            str(BLUEPRINT_PATH.name): sha256(BLUEPRINT_PATH),
            str(TASK_PATH.name): sha256(TASK_PATH),
            str(NORMALIZATION_PATH.name): sha256(NORMALIZATION_PATH),
        },
        "artifacts": {
            str(OUTPUT_DOCX.name): {"path": str(OUTPUT_DOCX), "sha256": docx_sha, "size_bytes": OUTPUT_DOCX.stat().st_size},
            str(OUTPUT_REGISTER.name): {"path": str(OUTPUT_REGISTER), "sha256": register_sha, "row_count": len(t09_records)},
        },
        "table_markers": markers,
        "unique_blueprint_tables_rendered": sorted(tables.keys()),
        "synthetic_register_rows": len(t09_records),
        "claims": {
            "semantic_review": "PENDING",
            "render_visual_review": "PENDING",
            "technical_validation": "PENDING",
            "human_acceptance": "PENDING",
            "real_association_evidence": "UNKNOWN; all 12 source records remain HOLD",
        },
    }
    OUTPUT_RECEIPT.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return receipt


if __name__ == "__main__":
    result = build()
    print(json.dumps(result, ensure_ascii=False, indent=2))
