from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor, Twips
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
import re, hashlib, json, datetime
root=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0131_I19-E20_SOURCE_PREP_R01_20261024')
src=root/'candidate'/'0131_I19-E20_BODY_R01C.md'
out=root/'candidate'/'0131_I19-E20_CANDIDATE_R01D.docx'
text=src.read_text(encoding='utf-8')
doc=Document()
sec=doc.sections[0]
sec.orientation=WD_ORIENT.LANDSCAPE
sec.page_width=Cm(29.7); sec.page_height=Cm(21.0)
sec.top_margin=Cm(1.55); sec.bottom_margin=Cm(1.45); sec.left_margin=Cm(1.35); sec.right_margin=Cm(1.35)
sec.header_distance=Cm(0.65); sec.footer_distance=Cm(0.65)
usable_w=sec.page_width-sec.left_margin-sec.right_margin
usable_twips=int(usable_w/635)
# Core type system
styles=doc.styles
normal=styles['Normal']; normal.font.name='Microsoft JhengHei'; normal.font.size=Pt(10.3); normal.font.color.rgb=RGBColor(40,48,58)
normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei')
normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.12
for name,size,color in [('Title',24,'17365D'),('Subtitle',12,'355070'),('Heading 1',17,'17365D'),('Heading 2',13,'1F4E79'),('Heading 3',11,'365F91')]:
    st=styles[name]; st.font.name='Microsoft JhengHei'; st.font.size=Pt(size); st.font.bold=name!='Subtitle'; st.font.color.rgb=RGBColor.from_string(color); st._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei'); st.paragraph_format.keep_with_next=True; st.paragraph_format.space_before=Pt(10 if name.startswith('Heading') else 0); st.paragraph_format.space_after=Pt(5)
# Helper XML formatting
def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)
def cell_margins(cell,top=55,start=65,bottom=55,end=65):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); mar=tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar=OxmlElement('w:tcMar'); tcPr.append(mar)
    for tag,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=mar.find(qn('w:'+tag))
        if node is None: node=OxmlElement('w:'+tag); mar.append(node)
        node.set(qn('w:w'),str(val)); node.set(qn('w:type'),'dxa')
def keep_row(row):
    pr=row._tr.get_or_add_trPr(); el=OxmlElement('w:cantSplit'); pr.append(el)
def repeat_header(row):
    pr=row._tr.get_or_add_trPr(); el=OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); pr.append(el)
def set_cell_font(cell,size=8.8,bold=False,color='263238'):
    for p in cell.paragraphs:
        p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
        for r in p.runs:
            r.font.name='Microsoft JhengHei'; r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei'); r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=RGBColor.from_string(color)
def add_field(paragraph,field):
    run=paragraph.add_run(); begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin')
    instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text=field
    separate=OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'),'separate')
    textel=OxmlElement('w:t'); textel.text='1'
    end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end')
    run._r.extend([begin,instr,separate,textel,end])
# Header/footer on every page carry the sample boundary.
header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.CENTER; header.paragraph_format.space_after=Pt(2)
r=header.add_run('SAMPLE / SYNTHETIC — NOT ASSOCIATION PERFORMANCE    |    TTQS ONE · 0131 · CANDIDATE R01')
r.font.name='Microsoft JhengHei'; r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei'); r.font.size=Pt(8); r.font.bold=True; r.font.color.rgb=RGBColor(120,40,40)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('REAL 結果＝UNKNOWN；正式接受未授予    ·    第 '); r.font.name='Microsoft JhengHei'; r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei'); r.font.size=Pt(8)
add_field(footer,'PAGE'); r=footer.add_run(' 頁'); r.font.size=Pt(8)
# Inline Markdown styling: strong, inline code, italic.
def inline(paragraph,value,base_size=10.3):
    value=re.sub(r'\$([^$]+)\$',r'\1',value)
    value=re.sub(r"\[([^\]]+)\]\([^)]+\)",r"\1",value)
    token=re.compile(r'(\*\*.*?\*\*|`[^`]+`|\*[^*]+\*)')
    pos=0
    for m in token.finditer(value):
        if m.start()>pos: paragraph.add_run(value[pos:m.start()])
        t=m.group(0)
        if t.startswith('**'): run=paragraph.add_run(t[2:-2]); run.bold=True
        elif t.startswith('`'): run=paragraph.add_run(t[1:-1]); run.font.name='Cascadia Mono'; run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei'); run.font.size=Pt(base_size-0.4); run.font.color.rgb=RGBColor(75,83,95)
        else: run=paragraph.add_run(t[1:-1]); run.italic=True
        pos=m.end()
    if pos<len(value): paragraph.add_run(value[pos:])
    for run in paragraph.runs:
        if run.font.size is None: run.font.size=Pt(base_size)
        if run.font.name is None: run.font.name='Microsoft JhengHei'; run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei')
def clean_cell(v):
    v=v.strip().replace('<br>','\n')
    v=re.sub(r'\*\*(.*?)\*\*',r'\1',v)
    v=re.sub(r'`([^`]+)`',r'\1',v)
    v=re.sub(r'\*(.*?)\*',r'\1',v)
    return v
# Table width derived from content but bounded to avoid pathological columns.
def add_table(rows):
    parsed=[[clean_cell(c) for c in row.strip().strip('|').split('|')] for row in rows]
    n=len(parsed[0]); parsed=[r if len(r)==n else (r+['']*n)[:n] for r in parsed]
    table=doc.add_table(rows=0,cols=n); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
    table.style='Table Grid'
    weights=[]
    for j in range(n):
        length=max([len(re.sub(r'\s+','',r[j])) for r in parsed[1:]]+[len(parsed[0][j])])
        weights.append(max(7,min(32,length)))
    total=sum(weights); widths=[int(usable_twips*w/total) for w in weights]
    # adjust rounding to exact page width
    widths[-1]+=usable_twips-sum(widths)
    for i,rowdata in enumerate(parsed):
        row=table.add_row(); keep_row(row)
        if i==0: repeat_header(row)
        for j,val in enumerate(rowdata):
            cell=row.cells[j]; cell.width=Twips(widths[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell_margins(cell)
            p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.LEFT; inline(p,val,8.8)
            if i==0:
                shade(cell,'17365D')
                for rr in p.runs: rr.font.bold=True; rr.font.color.rgb=RGBColor(255,255,255); rr.font.size=Pt(8.8)
            else:
                if i%2==0: shade(cell,'F1F5F9')
                for rr in p.runs: rr.font.size=Pt(8.7); rr.font.color.rgb=RGBColor(38,50,56)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
# parse Markdown stream
lines=text.splitlines(); i=0; in_code=False; code=[]; first_h1=True
while i<len(lines):
    line=lines[i]
    if line.strip().startswith('```'):
        if not in_code:
            in_code=True; code=[]
        else:
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(.25); p.paragraph_format.right_indent=Cm(.25); p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(7)
            pPr=p._p.get_or_add_pPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'F2F4F7'); pPr.append(shd)
            run=p.add_run('\n'.join(code)); run.font.name='Cascadia Mono'; run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei'); run.font.size=Pt(8.8); run.font.color.rgb=RGBColor(45,55,65)
            in_code=False
        i+=1; continue
    if in_code: code.append(line); i+=1; continue
    if line.startswith('|') and '|' in line[1:]:
        rows=[]
        while i<len(lines) and lines[i].startswith('|') and '|' in lines[i][1:]:
            # Markdown separator rows are represented by Word header styling, not body rows.
            if not re.match(r'^\|\s*:?-{2,}',lines[i]): rows.append(lines[i])
            i+=1
        add_table(rows); continue
    if not line.strip() or line.strip()=='---': i+=1; continue
    if line.startswith('# '):
        title=line[2:].strip()
        if title.startswith('SAMPLE / SYNTHETIC'):
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(9)
            inline(p,title,10); 
            for rr in p.runs: rr.bold=True; rr.font.color.rgb=RGBColor(155,40,40)
        elif first_h1:
            p=doc.add_paragraph(style='Title'); inline(p,title,24); first_h1=False
        else:
            p=doc.add_paragraph(style='Heading 1'); inline(p,title,17)
        i+=1; continue
    if line.startswith('## '):
        p=doc.add_paragraph(style='Heading 1'); inline(p,line[3:].strip(),17); i+=1; continue
    if line.startswith('### '):
        p=doc.add_paragraph(style='Heading 2'); inline(p,line[4:].strip(),13); i+=1; continue
    if line.startswith('#### '):
        p=doc.add_paragraph(style='Heading 3')

        inline(p,line[5:].strip(),11); i+=1; continue
    if line.startswith('>'):
        content=line[1:].strip(); p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(.3); p.paragraph_format.right_indent=Cm(.2); p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(6)
        pPr=p._p.get_or_add_pPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'FFF4E5'); inline(p,content,9.7)
        for rr in p.runs: rr.font.color.rgb=RGBColor(93,65,23)
        i+=1; continue
    numbered=re.match(r'^(\d+)\.\s+(.*)',line)
    if numbered:
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(.35); p.paragraph_format.first_line_indent=Cm(-.35); inline(p,numbered.group(1)+'. '+numbered.group(2),10.3); i+=1; continue
    if line.startswith('- '):
        p=doc.add_paragraph(style='List Bullet'); inline(p,line[2:],10.1); i+=1; continue
    p=doc.add_paragraph(); inline(p,line.strip(),10.3); i+=1
# document properties
props=doc.core_properties; props.title='學員證照通過與就業追蹤統計表｜SAMPLE/SYNTHETIC'; props.subject='0131 / I19-E20 / project_157_index 144 / official p.8'; props.author='TTQS ONE candidate production'; props.comments='候選 R01；所有數據為合成測試資料；REAL 結果 UNKNOWN；正式接受未授予。'
doc.save(out)
print(json.dumps({'output':str(out),'sha256':hashlib.sha256(out.read_bytes()).hexdigest(),'bytes':out.stat().st_size,'source_markdown_sha256':hashlib.sha256(src.read_bytes()).hexdigest()},ensure_ascii=True))



