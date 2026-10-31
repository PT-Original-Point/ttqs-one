from __future__ import annotations

import csv
import re
import sys
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[5]
PACKAGE = ROOT / "docs" / "ttqs-quality-20261016" / "quality-review" / "0067_R01"
SOURCE = PACKAGE / "worker" / "opencode_writer_R01_body.md"
DATASET = PACKAGE / "inputs" / "SYNTHETIC_DATASET_R01.csv"
OUTPUT = PACKAGE / "candidate" / "0067_R02.docx"

NAVY = "17365D"
TEAL = "176B72"
PALE = "EAF1F7"
ALTERNATE = "F5F8FB"
TEXT = "243447"
MUTED = "5B6875"
SAMPLE = "8A3B12"
FONT = "Microsoft JhengHei"
BT = chr(96)


def set_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def style_font(style, size, color=TEXT, bold=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    rpr = style._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + key), FONT)


def add_markup(paragraph, text, base_size=None):
    token_re = re.compile(r"(\*\*.*?\*\*|" + re.escape(BT) + r".*?" + re.escape(BT) + r")")
    lines = text.split("\n")
    for li, line in enumerate(lines):
        if li:
            paragraph.add_run().add_break()
        pos = 0
        for match in token_re.finditer(line):
            if match.start() > pos:
                run = paragraph.add_run(line[pos:match.start()])
                set_font(run, size=base_size)
            token = match.group(0)
            if token.startswith("**"):
                run = paragraph.add_run(token[2:-2])
                set_font(run, size=base_size, bold=True)
            else:
                run = paragraph.add_run(token[1:-1])
                set_font(run, size=base_size, color=TEAL)
            pos = match.end()
        if pos < len(line):
            run = paragraph.add_run(line[pos:])
            set_font(run, size=base_size)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=85, bottom=70, end=85):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "C7D2DE")


def mark_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, txt, end])
    set_font(run, size=8, color=MUTED)


def configure_section(section, landscape=False):
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11.69)
        section.page_height = Inches(8.27)
        section.left_margin = Inches(0.47)
        section.right_margin = Inches(0.47)
        section.top_margin = Inches(0.52)
        section.bottom_margin = Inches(0.55)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.left_margin = Inches(0.69)
        section.right_margin = Inches(0.69)
        section.top_margin = Inches(0.62)
        section.bottom_margin = Inches(0.62)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.30)
    section.different_first_page_header_footer = False
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_markup(p, "TTQS ONE  |  DOCX 0067  |  SAMPLE / SYNTHETIC  |  第 ", 8)
    set_page_field(p)
    add_markup(p, " 頁", 8)


def markdown_cells(line):
    return [x.strip() for x in line.strip().strip("|").split("|")]


def is_separator(line):
    cells = markdown_cells(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells)


def collect_table(lines, index):
    raw = []
    while index < len(lines) and lines[index].strip().startswith("|"):
        if not is_separator(lines[index]):
            raw.append(markdown_cells(lines[index]))
        index += 1
    if not raw:
        return [], index
    width = len(raw[0])
    rows = [r for r in raw if len(r) == width]
    return rows, index


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, size=9, bold=True, color=TEAL)


def add_table(doc, headers, rows, widths=None, font_size=8.4, first_col_bold=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    if widths is None:
        section = doc.sections[-1]
        available = section.page_width.inches - section.left_margin.inches - section.right_margin.inches
        weights = [max(1.0, min(3.2, max(len(str(h)), 5) / 9)) for h in headers]
        total = sum(weights)
        widths = [available * w / total for w in weights]
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Inches(widths[i])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        add_markup(p, str(text), font_size)
        for r in p.runs:
            set_font(r, size=font_size, bold=True, color="FFFFFF")
    mark_repeat_header(table.rows[0])
    keep_row_together(table.rows[0])
    for ridx, values in enumerate(rows):
        cells = table.add_row().cells
        keep_row_together(table.rows[-1])
        for i, value in enumerate(values):
            cell = cells[i]
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=65, start=80, bottom=65, end=80)
            if ridx % 2 == 1:
                set_cell_shading(cell, ALTERNATE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if i > 0 and str(value).strip().replace(",", "").replace(".", "").replace("%", "").replace("−", "").isdigit():
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_markup(p, str(value), font_size)
            if first_col_bold and i == 0:
                for r in p.runs:
                    r.bold = True
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def make_section(doc, landscape):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(sec, landscape=landscape)


def add_section5_tables(doc, rows):
    headers = rows[0]
    records = [dict(zip(headers, row)) for row in rows[1:]]
    total = next((r for r in records if r.get("班次") == "全體"), None)
    cohorts = [r for r in records if r.get("班次") != "全體"]
    add_caption(doc, "表 5A　班次識別與出席概況")
    cols = ["班次", "案例代碼", "課程（合成）", "日期", "登錄", "到課"]
    display = [[r.get(c, "") for c in cols] for r in cohorts]
    if total:
        display.append([total.get("班次", ""), "—", total.get("課程（合成）", ""), total.get("日期", ""), total.get("登錄", ""), total.get("到課", "")])
    add_table(doc, cols, display, [0.68, 1.65, 4.15, 1.27, 1.0, 1.0], 8.8, True)
    add_caption(doc, "表 5B　缺席、完成與評量分母橋接")
    cols = ["班次", "缺席", "完成", "有效評量", "達標", "輸入 fact 範圍"]
    display = [[r.get(c, "") for c in cols] for r in records]
    add_table(doc, cols, display, [2.2, 1.45, 1.45, 1.75, 1.25, 2.4], 8.8, True)
    add_caption(doc, "表 5C　問卷回收與正向回覆")
    cols = ["班次", "問卷發出", "問卷有效", "正向", "輸入 fact 範圍"]
    display = [[r.get(c, "") for c in cols] for r in records]
    add_table(doc, cols, display, [2.3, 1.65, 1.65, 1.4, 3.5], 8.8, True)
    add_caption(doc, "表 5D　計畫、支出與差額")
    source_plan = "計畫 (TWD)"
    source_spend = "支出 (TWD)"
    cols = ["班次", "計畫 (TWD)", "支出 (TWD)", "差額（支出減計畫）", "使用率", "輸入 fact 範圍"]
    display = []
    for r in records:
        plan = Decimal(r[source_plan].replace(",", ""))
        spend = Decimal(r[source_spend].replace(",", ""))
        rate = (spend / plan * Decimal(100)).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)
        variance = spend - plan
        display.append([r["班次"], f"{int(plan):,}", f"{int(spend):,}", f"{int(variance):+,}", f"{rate:.1f}%", r.get("輸入 fact 範圍", "")])
    add_table(doc, cols, display, [1.45, 1.65, 1.65, 2.2, 1.25, 2.3], 8.5, True)


def add_appendix_table(doc, rows):
    headers = rows[0]
    index = {name: headers.index(name) for name in headers}
    out = []
    for row in rows[1:]:
        ident = "\n".join([
            row[index["fact_id"]],
            row[index["班次"]] + "　" + row[index["欄位"]],
            "值：" + row[index["值"]] + ("　" + row[index["單位"]] if row[index["單位"]] else ""),
            "狀態：" + row[index["狀態"]],
        ])
        out.append([ident, row[index["原始定位"]], row[index["待替換來源"]], row[index["合成原因"]]])
    cols = ["事實識別與數值", "正文／來源定位", "改用 REAL 時須取得", "SAMPLE／SYNTHETIC 原因"]
    add_table(doc, cols, out, [2.15, 2.05, 3.05, 3.35], 7.4, True)


def main():
    if not SOURCE.exists():
        raise SystemExit(f"Missing writer body: {SOURCE}")
    if OUTPUT.exists():
        raise SystemExit(f"Refusing to overwrite existing candidate: {OUTPUT}")
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.splitlines()
    data_rows = list(csv.DictReader(DATASET.open(encoding="utf-8-sig", newline="")))
    if len(data_rows) != 3:
        raise SystemExit("Expected exactly three synthetic cohorts")

    doc = Document()
    doc.core_properties.title = "訓練執行評估指標分析及相關統計分析"
    doc.core_properties.subject = "DOCX 0067 | I15-E03 | 官方方框 073 | SAMPLE / SYNTHETIC 候選文件"
    doc.core_properties.category = "候選文件"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "TTQS ONE, 0067, I15-E03, 073, SAMPLE, SYNTHETIC"
    doc.settings.odd_and_even_pages_header_footer = False
    configure_section(doc.sections[0], landscape=False)

    styles = doc.styles
    style_font(styles["Normal"], 10.2, TEXT)
    styles["Normal"].paragraph_format.space_after = Pt(4)
    styles["Normal"].paragraph_format.line_spacing = 1.17
    style_font(styles["Title"], 22, NAVY, True)
    title_style_ppr = styles["Title"]._element.pPr
    if title_style_ppr is not None:
        title_border = title_style_ppr.find(qn("w:pBdr"))
        if title_border is not None:
            title_style_ppr.remove(title_border)
    styles["Title"].paragraph_format.space_after = Pt(5)
    styles["Title"].paragraph_format.keep_with_next = True
    style_font(styles["Heading 1"], 15, NAVY, True)
    styles["Heading 1"].paragraph_format.space_before = Pt(12)
    styles["Heading 1"].paragraph_format.space_after = Pt(5)
    styles["Heading 1"].paragraph_format.keep_with_next = True
    style_font(styles["Heading 2"], 11.7, TEAL, True)
    styles["Heading 2"].paragraph_format.space_before = Pt(8)
    styles["Heading 2"].paragraph_format.space_after = Pt(3)
    styles["Heading 2"].paragraph_format.keep_with_next = True
    style_font(styles["List Bullet"], 10.1, TEXT)
    style_font(styles["List Number"], 10.1, TEXT)
    style_font(styles["Caption"], 9, TEAL, True)

    title = doc.add_paragraph(style="Title")
    title_ppr = title._p.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_markup(title, "訓練執行評估指標分析及相關統計分析", 22)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(8)
    add_markup(sub, "跨班訓練 KPI 統計分析報告", 12)
    for run in sub.runs:
        set_font(run, size=12, color=TEAL, bold=True)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(4)
    add_markup(meta, "DOCX 0067  |  I15-E03  |  官方方框 073  |  候選文件", 9)
    for run in meta.runs:
        set_font(run, size=9, color=MUTED)
    sample = doc.add_paragraph()
    sample.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sample.paragraph_format.space_after = Pt(14)
    add_markup(sample, "全篇 SAMPLE / SYNTHETIC　REAL 成效：UNKNOWN / HOLD", 10)
    for run in sample.runs:
        set_font(run, size=10, color=SAMPLE, bold=True)

    i = 1
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            i += 1
            continue
        if line.startswith("## 3."):
            make_section(doc, False)
        elif line.startswith("## 5."):
            make_section(doc, True)
        elif line.startswith("## 6."):
            make_section(doc, False)
        elif line.startswith("## 7."):
            make_section(doc, True)
        elif line.startswith("### 7.2 "):
            make_section(doc, False)
        elif line.startswith("## 附錄 A") or line.startswith("## 附录 A"):
            make_section(doc, True)
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            heading_text = line[level:].strip()
            p = doc.add_paragraph(style="Heading 1" if level == 2 else "Heading 2")
            add_markup(p, heading_text)
            i += 1
            continue
        if line.startswith("|"):
            rows, nxt = collect_table(lines, i)
            if not rows:
                i = max(nxt, i + 1)
                continue
            header = rows[0]
            if header and header[0] == "班次" and len(header) >= 15:
                add_section5_tables(doc, rows)
            elif header and header[0] == "fact_id":
                add_appendix_table(doc, rows)
            else:
                add_table(doc, header, rows[1:], font_size=8.25, first_col_bold=True)
            i = nxt
            continue
        if line.startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Inches(0.24)
            p.paragraph_format.first_line_indent = Inches(-0.14)
            add_markup(p, line[2:])
        elif re.match(r"^\d+\.\s+", line):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            add_markup(p, line)
        else:
            p = doc.add_paragraph()
            add_markup(p, line)
        i += 1

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"DOCX={OUTPUT}")
    print(f"BYTES={OUTPUT.stat().st_size}")


if __name__ == "__main__":
    main()
