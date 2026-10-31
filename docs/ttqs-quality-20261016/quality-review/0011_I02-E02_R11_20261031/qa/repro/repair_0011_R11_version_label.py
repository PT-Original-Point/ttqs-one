from pathlib import Path
from docx import Document
import hashlib,json,datetime,sys
base=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0011_I02-E02_LUNA_TEXT_DRAFT_R01_20261029')
src=base/'candidate'/'0011_I02-E02_R10.docx'; dst=base/'candidate'/'0011_I02-E02_R11.docx'
if dst.exists(): raise SystemExit('R11_EXISTS_NO_OVERWRITE')
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
old_sha=sha(src); doc=Document(src)
matched=[p for p in doc.paragraphs if '文件代碼：0011／I02-E02' in p.text and '版本：R09' in p.text]
if len(matched)!=1: raise SystemExit(f'EXPECTED_ONE_VERSION_HEADER_GOT_{len(matched)}')
p=matched[0]; old=p.text
changed=0
for run in p.runs:
 if '版本：R09' in run.text:
  run.text=run.text.replace('版本：R09','版本：R11',1); changed+=1
if changed!=1: raise SystemExit(f'EXPECTED_ONE_RUN_PATCH_GOT_{changed}')
if '版本：R11' not in p.text or '版本：R09' in p.text: raise SystemExit('PATCH_VERIFY_FAILED')
doc.save(dst)
reopened=Document(dst)
headers=[x.text for x in reopened.paragraphs if '文件代碼：0011／I02-E02' in x.text]
if len(headers)!=1 or '版本：R11' not in headers[0]: raise SystemExit('SAVED_DOCX_REOPEN_VERIFY_FAILED')
receipt={'schema':'ttqs-one-minimal-candidate-repair-v1','work_id':'TTQS-0011-R11-VERSION-LABEL-REPAIR-20261031-01','doc_id':'0011','requirement_id':'I02-E02','status':'PATCH_APPLIED_CANDIDATE','observed_at_tpe':datetime.datetime.now().astimezone().isoformat(timespec='seconds'),'source_candidate_path':str(src),'source_candidate_sha256':old_sha,'source_candidate_preserved':True,'new_candidate_path':str(dst),'new_candidate_sha256':sha(dst),'change_locator':'首頁第1段（文件代碼／版本行）','old_text':old,'new_text':headers[0],'changed_paragraph_count':1,'changed_run_count':changed,'mutation_scope':'僅版本標籤 R09→R11；未重建正文；其他內容待確定性比較。','source_control':'local observed bytes only; no formal acceptance or REAL evidence claimed'}
receipt_path=base/'qa'/'PATCH_RECEIPT_R11.json'; receipt_path.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'status':receipt['status'],'source_sha256':old_sha,'candidate_sha256':receipt['new_candidate_sha256'],'receipt_path':str(receipt_path),'receipt_sha256':sha(receipt_path),'header':headers[0]},ensure_ascii=False))
