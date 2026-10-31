from pathlib import Path
from docx import Document
import hashlib,json,re,zipfile,sys,csv
from datetime import datetime
B=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0011_I02-E02_LUNA_TEXT_DRAFT_R01_20261029')
P=B/'candidate'/'0011_I02-E02_R11.docx'; OLD=B/'candidate'/'0011_I02-E02_R10.docx'
PDF=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\INPUTS\R04\06_OFFICIAL_SOURCES\TTQS查核佐證文件資料表-訓練機構版 (2).pdf')
BP=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\CONTROL\BLUEPRINT_CURRENT.jsonl')
CW=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\CONTROL\TTQS_ONE_129_REQUIREMENT_TO_142_DELIVERABLE_CROSSWALK_R01.csv')
OUT=B/'qa'/'technical_qa_R11.json'
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def parse(p):
 d=Document(str(p)); return d,[x.text for x in d.paragraphs],[[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
def compact(s): return re.sub(r'\s+','',s)
new,paras,tables=parse(P); old,oldparas,oldtables=parse(OLD)
with zipfile.ZipFile(P) as z: zip_ok=z.testzip() is None
body='\n'.join(paras)+'\n'+'\n'.join('\n'.join(' | '.join(r) for r in t) for t in tables)
oldbody='\n'.join(oldparas)+'\n'+'\n'.join('\n'.join(' | '.join(r) for r in t) for t in oldtables)
empty=[(ti,ri,ci) for ti,t in enumerate(tables,1) for ri,row in enumerate(t,1) for ci,v in enumerate(row,1) if not v.strip()]
body_fonts=[r.font.size.pt for p in new.paragraphs for r in p.runs if r.font.size is not None]
table_fonts=[r.font.size.pt for t in new.tables for row in t.rows for c in row.cells for p in c.paragraphs for r in p.runs if r.font.size is not None]
header=[x for x in paras if '文件代碼：0011／I02-E02' in x]
paragraph_delta=[]
for i,(a,b) in enumerate(zip(oldparas,paras)):
 if a!=b: paragraph_delta.append({'index':i,'old':a,'new':b})
blueprint_rows=[]
for line in BP.read_text(encoding='utf-8-sig').splitlines():
 try:
  r=json.loads(line)
  if r.get('deliverable_id')=='0011': blueprint_rows.append((line,r))
 except Exception: pass
with CW.open(encoding='utf-8-sig',newline='') as f: cross=[r for r in csv.DictReader(f) if r.get('requirement_id')=='I02-E02']
weights=[25,25,20,15,15]; scores={'A':[3,3,2,5,5],'B':[4,5,5,4,3],'C':[5,4,4,3,2]}; calc={k:sum(a*b for a,b in zip(v,weights)) for k,v in scores.items()}; standard={k:v/5 for k,v in calc.items()}
source_phrase='與訓練政策有關之訓練產品與市場定位。'; pdf_phrase=False; pdf_extract='UNKNOWN'
try:
 from pypdf import PdfReader
 txt=PdfReader(str(PDF)).pages[0].extract_text() or ''; pdf_phrase=compact(source_phrase) in compact(txt); pdf_extract='pypdf page 1 whitespace-normalized exact phrase match'
except Exception as e: pdf_extract='UNAVAILABLE:'+type(e).__name__
checks={
 'zip_integrity':zip_ok,'python_docx_open':True,'candidate_sha_computed':P.exists(),
 'correct_doc_and_requirement_binding':'文件代碼：0011／I02-E02' in '\n'.join(paras),
 'revision_label_matches_candidate':len(header)==1 and '版本：R11' in header[0] and '版本：R09' not in header[0],
 'only_one_expected_paragraph_changed':len(paras)==len(oldparas) and len(paragraph_delta)==1 and '版本：R09' in paragraph_delta[0]['old'] and '版本：R11' in paragraph_delta[0]['new'],
 'all_tables_identical_to_R10':tables==oldtables,
 'exactly_27_tables':len(tables)==27,'no_empty_table_cells':not empty,
 'body_min_font_ge_10_5':bool(body_fonts) and min(body_fonts)>=10.5,
 'table_min_font_ge_9_5':bool(table_fonts) and min(table_fonts)>=9.5,
 'formula_recomputed':sum(weights)==100 and calc=={'A':340,'B':430,'C':380} and standard=={'A':68.0,'B':86.0,'C':76.0},
 'formula_outputs_present':all(v in body for v in ['340','430','380','68.0','86.0','76.0']),
 'fact_ids_unchanged_from_R10':set(re.findall(r'\bF\d{3}\b',body))==set(re.findall(r'\bF\d{3}\b',oldbody)),
 'unique_fact_ids_83':len(set(re.findall(r'\bF\d{3}\b',body)))==83,
 'attachment_index_present':'12.1 附件索引' in body and '沒有另外附掛的外部檔案' in body,
 'official_source_sha_expected':sha(PDF)=='805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709',
 'official_requirement_page1_phrase_match':pdf_phrase,
 'blueprint_file_sha_expected':sha(BP)=='406d0495365d0970dcf8113a42842663b8a54372a4e204cc492b7c3f8156ef3e',
 'blueprint_unique_row_matches':len(blueprint_rows)==1 and blueprint_rows[0][1].get('requirement_id')=='I02-E02' and blueprint_rows[0][1].get('final_filename')=='0011__與訓練政策有關之訓練產品與市場定位_制度文件.docx',
 'crosswalk_sha_expected':sha(CW)=='1ddfbb608bcad25170683ddb42a1dbb27bafa6f1f35ff0498f1b8d75b24f0679',
 'crosswalk_maps_only_0011':len(cross)==1 and cross[0].get('deliverable_ids')=='0011' and cross[0].get('deliverable_count')=='1'
}
row_line,row=blueprint_rows[0] if blueprint_rows else ('',{})
receipt={'schema':'ttqs-one-technical-qa-r11-v1','receipt_id':'0011_I02-E02_R11_DETERMINISTIC_TECHNICAL_QA','observed_at_tpe':datetime.now().astimezone().isoformat(timespec='seconds'),'doc_id':'0011','requirement_id':'I02-E02','official_checkbox_serial':'UNKNOWN','official_requirement_text':source_phrase,'candidate_path':str(P),'candidate_sha256':sha(P),'candidate_bytes':P.stat().st_size,'base_candidate_R10_sha256':sha(OLD),'change_receipt_path':str(B/'qa'/'PATCH_RECEIPT_R11.json'),'change_receipt_sha256':sha(B/'qa'/'PATCH_RECEIPT_R11.json'),'blueprint_file_sha256':sha(BP),'blueprint_row_sha256_raw_line':hashlib.sha256(row_line.encode('utf-8')).hexdigest() if row_line else None,'blueprint_row_sha256_canonical_sorted_compact':hashlib.sha256(json.dumps(row,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode('utf-8')).hexdigest() if row else None,'crosswalk_file_sha256':sha(CW),'official_source_pdf_sha256':sha(PDF),'official_phrase_extraction':pdf_extract,'checks':checks,'failed_checks':[k for k,v in checks.items() if v is not True],'tables':len(tables),'table_rows_including_headers':sum(len(t) for t in tables),'empty_cells':empty,'paragraph_delta':paragraph_delta,'body_min_font_pt':min(body_fonts) if body_fonts else None,'table_min_font_pt':min(table_fonts) if table_fonts else None,'formula_recalculation':{'weights':weights,'weights_total':sum(weights),'weighted_points':calc,'standardized_scores':standard},'fact_ids_unique_count':len(set(re.findall(r'\bF\d{3}\b',body))),'semantic_review':'PENDING_FULL_EXACT_SHA_REVIEW','visual_qa':'PENDING_R11_WORD_NATIVE_RENDER','formal_acceptance':'NOT_GRANTED','candidate_state':'CANDIDATE_SAMPLE_SYNTHETIC','real_evidence':'UNKNOWN_NOT_CLAIMED'}
receipt['status']='TECHNICAL_PASS' if not receipt['failed_checks'] else 'TECHNICAL_FAIL'
OUT.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'status':receipt['status'],'candidate_sha256':receipt['candidate_sha256'],'candidate_bytes':receipt['candidate_bytes'],'checks':checks,'failed_checks':receipt['failed_checks'],'tables':len(tables),'rows':sum(len(t) for t in tables),'fact_ids':receipt['fact_ids_unique_count'],'body_min_font_pt':receipt['body_min_font_pt'],'table_min_font_pt':receipt['table_min_font_pt'],'blueprint_raw_sha':receipt['blueprint_row_sha256_raw_line'],'blueprint_canonical_sha':receipt['blueprint_row_sha256_canonical_sorted_compact'],'official_phrase_extraction':pdf_extract,'receipt_sha256':sha(OUT)},ensure_ascii=True))
if receipt['status']!='TECHNICAL_PASS': sys.exit(2)
