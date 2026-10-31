from pathlib import Path
import csv
from docx import Document
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'candidate' / '0055_R02.docx'
REG = ROOT / 'evidence' / 'SYNTHETIC_REPLACEMENT_REGISTER_0055_R02.csv'
OUT.parent.mkdir(parents=True, exist_ok=True)
REG.parent.mkdir(parents=True, exist_ok=True)

NAVY = '17365D'
BLUE = 'D9EAF7'
PALE = 'EAF2F8'
GOLD = 'FCE4A6'
RED = 'FCE4D6'
GRAY = 'F2F4F7'
GREEN = 'E2F0D9'
INK = '243247'
FONT = 'Microsoft JhengHei'

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), color)

def set_cell_margins(cell, top=75, start=100, bottom=75, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    trPr.append(el)

def keep_row(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:cantSplit'))

def set_cell_text(cell, text, bold=False, color=INK, size=9.5, align=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    if align is not None:
        p.alignment = align
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)

def set_repeat_table_borders(table, color='B7C9D6', size='4'):
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tblPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), size)
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)

def table(doc, headers, rows, widths=None, font_size=8.2, header_color=NAVY, first_col_bold=False):
    # Keep dense evaluator-facing tables readable at normal print scale.
    font_size = max(font_size, 9.5)
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_repeat_table_borders(t)
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]
        set_cell_text(c, h, True, 'FFFFFF', font_size, WD_ALIGN_PARAGRAPH.CENTER)
        shade(c, header_color)
    repeat_header(t.rows[0])
    keep_row(t.rows[0])
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for j, value in enumerate(row):
            set_cell_text(cells[j], value, first_col_bold and j == 0, size=font_size)
            if ridx % 2 == 1:
                shade(cells[j], 'F7F9FB')
        keep_row(t.rows[-1])
    if widths:
        for row in t.rows:
            for j, width in enumerate(widths):
                row.cells[j].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def set_run(run, size=10, bold=False, color=INK, italic=False):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)

def para(doc, text='', style=None, size=10, bold=False, color=INK, before=0, after=5, align=None, italic=False):
    size = max(size, 10.5)
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run(r, size, bold, color, italic)
    return p

def heading(doc, text, level=1, page_break_before=False):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break_before
    p.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run(r, 16 if level == 1 else 12 if level == 2 else 10.5, True, NAVY)
    return p

def bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    set_run(r, 10.5)
    return p

def numbered(doc, text):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    set_run(r, 10.5)
    return p

def callout(doc, label, text, color=PALE):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_repeat_table_borders(t, '9BB4C7', '6')
    keep_row(t.rows[0])
    c = t.cell(0, 0)
    shade(c, color)
    c.text = ''
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label + '｜')
    set_run(r, 9.5, True, NAVY)
    r = p.add_run(text)
    set_run(r, 9.5)
    set_cell_margins(c, 120, 150, 120, 150)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def page_field(p):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('講師遴選與聘任流程範例　•　SAMPLE / SYNTHETIC　•　第 ')
    set_run(r, 8, False, '64748B')
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    p._p.append(fld)
    r = p.add_run(' 頁')
    set_run(r, 8, False, '64748B')

def add_candidate_score(doc, code, scores, rationale):
    criteria = [
        ('專業內容正確性', 25), ('教學示範與引導', 25),
        ('成人志工適配與可近性', 20), ('實務應用與移轉', 20),
        ('個資、角色與安全界線', 10),
    ]
    rows = []
    weighted_sum = 0
    for (name, weight), triplet in zip(criteria, scores):
        mean = sum(triplet) / 3
        weighted = mean * weight / 100
        weighted_sum += weighted
        rows.append((name, str(weight), ' / '.join(map(str, triplet)), f'{mean:.0f}', f'{weighted:.1f}'))
    table(doc, ['評分面向', '權重%', '評審甲／乙／丙', '平均', '加權分'], rows,
          [2.15, .65, 1.35, .65, .72], 8.0)
    para(doc, f'{code} 合計 = ' + ' + '.join(r[4] for r in rows) + f' = {weighted_sum:.1f} 分。', size=9.2, bold=True, color=NAVY)
    para(doc, rationale, size=9.2)
    return weighted_sum

doc = Document()
sec = doc.sections[0]
sec.page_width = Mm(210)
sec.page_height = Mm(297)
sec.top_margin = Inches(.62)
sec.bottom_margin = Inches(.62)
sec.left_margin = Inches(.7)
sec.right_margin = Inches(.7)
sec.header_distance = Inches(.28)
sec.footer_distance = Inches(.28)
styles = doc.styles
for name in ('Normal', 'Body Text', 'List Bullet', 'List Bullet 2', 'List Number'):
    st = styles[name]
    st.font.name = FONT
    st._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    st.font.size = Pt(10.5)
    st.font.color.rgb = RGBColor.from_string(INK)
    st.paragraph_format.space_after = Pt(4)
for n, sz in [('Heading 1',16),('Heading 2',12),('Heading 3',10.5)]:
    st = styles[n]
    st.font.name = FONT
    st._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    st.font.size = Pt(sz)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(NAVY)
    st.paragraph_format.keep_with_next = True
header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = header.add_run('原始點協會｜講師遴選與聘任流程範例　•　示範文件，非正式任命紀錄')
set_run(r, 8, False, '64748B')
page_field(sec.footer.paragraphs[0])

# Cover and operational quick path
para(doc, '講師遴選與聘任流程範例', size=24, bold=True, color=NAVY, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
para(doc, '從課程需求、評選量尺、同條件試教、三人比較，到委員會決定與聘書處理的完整示範', size=11.5, color='356A91', after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
callout(doc, 'SAMPLE / SYNTHETIC｜未核准、未發聘、不可作為真實任命證明',
        '本範例所有人選代碼、分數、評審、資格核驗狀態、會議、決議、日期、費用及聘書文字均為合成示範。原始點協會目前是否設有有效訓練管理委員會、其權限、真實講師與資格、真實核准及聘任均為 UNKNOWN。', GOLD)
para(doc, '本範例供課程需求提出人、遴選承辦人、評審及具有效授權的決策者使用，說明如何為一門特定課程訂出講師能力需求、以一致且可重算的量尺比較候選人、記錄試教與利益衝突、核對資格與授權，並在取得有權決定後才處理聘書；正文依下方作業順序所標示的章節閱讀。', size=10.5, after=4)
para(doc, '案例 SYN-0033-PLAN-01｜志工服務紀錄、交接與工作界線實務，合成規劃 360 分鐘（90+90+120+60）；沿用既有合成規劃資料，不表示課程已開辦。', size=10.5, after=4)
table(doc, ['遴選作業順序與正文位置', '形成資料與放行條件'], [
    ('建立需求與核對權限（第 2–3 節）', '先確認課程職能缺口與決策授權；委員會、評審或簽發權任何一項缺少來源即 HOLD。'),
    ('設定量尺並查核候選人（第 4–5、8 節）', '使用同一版標準；完成資格、授權與利益衝突檢查，缺原件或未揭露即 HOLD。'),
    ('安排同條件試教並比較（第 6–7 節）', '保存觀察依據、三人評分、可重算總分及排序；示例分數只作推薦參考，不等於錄用。'),
    ('有權決議後處理聘書（第 9–10 節），再依第 12 節逐筆替換 SAMPLE', '先核對有效權限、出席、表決與附件；正式決議、簽發與接受未完成前，不得把聘書當作已發出。'),
], [2.3, 4.5], 9.5, first_col_bold=True)
para(doc, '官方來源定位：訓練品質評核表第 4 頁指標 12c（需求代碼 I12-E03；官方方框 058）：「講師的遴聘依師資遴選標準及流程評選並經訓練管理委員會通過後發給聘書。」', size=10.5, after=4)

heading(doc, '1｜需求、用途與相鄰需求界線', 1)
para(doc, '這份文件要回答：為這一門特定課程，需要什麼講師能力？如何用一致、可重算的標準比較候選人？哪個機關有權決定？核准後如何形成聘任文件？它不是講師授課後的評估、學員招生，也不是設備或場地檢查。', size=10.5)
table(doc, ['相鄰需求（需求代碼／官方方框）', '本包如何區分'], [
    ('訓前五項適配檢視與調整（I12-E01／方框 056）', '檢視學員、教材、講師、方法與環境／設備的適配並調整；不能代替講師候選人比較與任命。'),
    ('學員遴選與教材審查（I12-E02／方框 057）', '入班遴選條件與教材審查；不判斷講師資格、試教分數或聘書。'),
    ('本包：講師遴選與聘任（I12-E03／官方方框 058）', '本包：標準、三人比較、試教、資格／授權／利益衝突核驗、委員會決定、核准後聘書處理。'),
    ('設備檢查與場地保險（I12-E05／方框 059）', '設備清點與場地投保證明；聘任資料不能代替場地／設備證據。'),
    ('單元目標與評量設計（I11-E02／方框 052）', '同案例課程單元與評量設計可供對照；不能證明任何一位講師實際授課或表現。'),
], [2.35, 4.45], 9.5, first_col_bold=True)

heading(doc, '2｜合成課程需求與可驗證能力', 1)
table(doc, ['需求／職能落差（合成）', '本課程講師須可示範的行為', '可收集的原始依據'], [
    ('志工紀錄欄位易漏、事實與推測混寫', '示範逐欄填寫；能辨認觀察、引述、推論；使用情境題修正錯誤紀錄。', '課程需求摘要、講師既有作品或試教講義；取用與個資授權紀錄。'),
    ('交接時責任、期限、未結事項不清', '以交接清單說明待辦、責任人、期限、風險和回覆確認。', '試教評分表、情境演練觀察紀錄。'),
    ('志工容易越過角色界線或傳遞敏感資訊', '能辨認何時停止、保護資料並轉介；不以志工課程取代專業服務。', '邊界情境試教、資料保護口頭提問及答案紀錄。'),
    ('成人學員背景差異大、部分以手機參與', '指令分段、使用易讀材料、提供口頭與書面步驟及替代參與方式。', '試教可近性觀察表、教材版本與無障礙需求紀錄。'),
], [1.7, 2.7, 2.35], 8.0)
callout(doc, '設計決策', '本示範將講師效能拆成五項可觀察能力；講師履歷年資本身不直接等同教學能力。試教使用同一題目、同時長、同一量尺與相同提示，避免只憑印象排序。')

heading(doc, '3｜權限、角色與不可跳過的關卡', 1)
table(doc, ['角色', '工作與留存紀錄', '獨立性／停止條件'], [
    ('課程需求提出人', '說明職能落差、目標、對象及時數；簽名／日期須用真實來源替換。', '不得自行宣稱具委員會核准權。'),
    ('遴選承辦人', '收件、編碼、檢核完整性、安排一致試教、保管檔案與版本。', '若同時是候選人或利害關係人，須揭露並退出受影響步驟。'),
    ('評審甲／乙／丙', '按同一評分表獨立評分，寫具體觀察，揭露利益衝突。', '合成案例中的代碼非真實人員；真實名單與資格尚未確認。'),
    ('委員會主席／記錄', '先核對章程、有效任期、出席及表決規則；記錄討論、決定、反對／迴避、責任與期限。', '目前協會是否有該委員會或授權均 UNKNOWN。證據未齊就不得表決或把草稿標為已核准。'),
    ('聘任承辦／核准人', '依有權決議核定課程、期間、費用及條件，出具並送達聘書，取得回簽。', '不可倒填日期；未核准不得發書；真實簽章須由權責者親簽。'),
], [1.35, 3.3, 2.1], 7.9, first_col_bold=True)
para(doc, '最小權限檢查清單（逐項勾選並附來源）：□ 現行章程／組織規程；□ 委員會設置與任期；□ 委員名冊與迴避規則；□ 決議門檻；□ 課程與費用核定權；□ 聘書簽發權；□ 文件保存與個資存取規則。任何一格缺來源，狀態維持 HOLD。', size=9)

heading(doc, '4｜標準、量尺、通過門檻與退件規則', 1, page_break_before=False)
table(doc, ['評分面向', '權重', '0 分錨點', '60 分錨點', '100 分錨點'], [
    ('專業內容正確性', '25%', '多項關鍵資訊錯誤，或不承認未知。', '主要概念正確，仍有部分例子／界線需釐清。', '內容準確、說明證據界線，能辨認不確定並正確轉介。'),
    ('教學示範與引導', '25%', '未交代目標、步驟或檢核理解。', '有基本說明與一個練習，回饋較單向。', '目標可觀察、步驟清楚，提問／練習／回饋形成閉環。'),
    ('成人志工適配與可近性', '20%', '假設所有人背景相同，材料不可讀或參與門檻高。', '可跟上一般學員，少量調整可提供。', '節奏、語言、字級、替代參與及休息安排均具體。'),
    ('實務應用與移轉', '20%', '只有抽象講述，沒有工作情境演練。', '有示範與練習，但回到工作後的步驟不完整。', '以真實可替換情境練習，提供可執行檢核及回到工作後的追蹤。'),
    ('個資、角色與安全界線', '10%', '洩露敏感資訊、越權承諾或鼓勵危險做法。', '能說明基本注意事項，但停止／轉介條件不完整。', '最小必要資料、保密、角色界線及停止／轉介點均說清楚。'),
], [1.25, .55, 1.65, 1.65, 1.65], 7.5)
para(doc, '評分公式：面向平均＝三位評審原始分數加總 ÷ 3；加權分＝面向平均 × 權重 ÷ 100；總分＝五項加權分總和。面向分數取整數記錄；加權分至小數一位；總分依未四捨五入的平均值計算至小數一位。權重合計 25+25+20+20+10=100%。', size=9.1, bold=True)
table(doc, ['狀態', '規則', '處置'], [
    ('通過門檻（合成示例）', '總分 ≥80；專業內容正確性、教學示範與引導、成人志工適配與可近性各 ≥75；個資、角色與安全界線 ≥80；無未解利益衝突；資格、授權及原始資料檢核完成；委員會有權且合法開會。', '列入可推薦名單；仍須經有權委員會決定。'),
    ('HOLD', '原件缺漏／未驗證、權限未知、利益衝突未解、評分證據不一致、合理調整未確認。', '暫停排序或聘任；列明缺件、責任人及期限；不得先發聘書。'),
    ('REJECT', '明顯不實資料；重大專業錯誤；安全／個資紅線；門檻未達且無可補正方案。', '記錄具體證據與告知／申覆途徑；不得私下刪除不利紀錄。'),
], [1.25, 3.7, 1.8], 8.1, first_col_bold=True)
callout(doc, '非補償性條件', '總分再高也不能抵銷資格未核、未授權、無效委員會、未揭露衝突或安全紅線。通過分數只形成推薦依據，不自動等於聘任。', RED)

heading(doc, '5｜候選人收件、資格與衝突核驗', 1)
table(doc, ['核驗欄位', 'C-01（合成）', 'C-02（合成）', 'C-03（合成）', '真實使用時的原始證據'], [
    ('履歷／相關經驗', '樣本履歷索引 S-C01', '樣本履歷索引 S-C02', '樣本履歷索引 S-C03', '本人提交的履歷、服務證明；逐項標來源及核驗日期。'),
    ('專業資格', '樣本核驗標記：PASS', '樣本核驗標記：PASS', '樣本核驗標記：PASS', '證照影本與發證機構查驗紀錄；不得只靠自述。'),
    ('教學授權／可用時段', '樣本：可配合', '樣本：可配合', '樣本：可配合', '本人同意、課程需求與排程確認；不可代替正式聘任。'),
    ('利益衝突', '樣本：無衝突申報', '樣本：無衝突申報', '樣本：無衝突申報', '本人簽署申報及評審迴避紀錄；未申報先 HOLD。'),
    ('個資及作品使用', '樣本：同意限本次遴選', '樣本：同意限本次遴選', '樣本：同意限本次遴選', '有日期、用途、保存期與撤回管道的真實授權。'),
], [1.2, 1.15, 1.15, 1.15, 2.0], 7.6)
para(doc, '上表的 PASS 都只是虛構流程欄位值，沒有任何真實證照或查詢結果附在本包。真實承辦人須把狀態改為「已驗證」並附原件位置、驗證人、日期、方式與結果；若不能核對，填「未驗證／HOLD」，不得照抄 PASS。', size=9.1, bold=True, color='9C3A22')

heading(doc, '6｜一致試教與觀察紀錄', 1)
table(doc, ['項目', '三位候選人共用條件（合成）'], [
    ('題目與時長', '12 分鐘微型試教：「如何把一筆志工服務紀錄寫成可交接、可查核且不洩露多餘個資的紀錄」。每人同一題目、同一預備時間。'),
    ('必備教學元素', '開場說明學習目標；用一份合成錯誤紀錄示範修正；帶學員完成一次交接欄位練習；用兩題口頭檢核確認事實／推測及停止／轉介界線。'),
    ('共同觀察環境', '相同教室／視訊設定、相同投影設備、同一套大字版樣本、相同提問卡；合理調整以不改變評量目的為原則並記錄。'),
    ('評審紀錄', '三人先獨立填表，記錄可觀察行為或原話；評分後才討論。不得把個人好惡寫成能力結論。'),
    ('試教後確認', '詢問候選人如何處理錯誤答案、敏感資料、學員不同步及課後移轉；將回答摘要與面向分數分開保存。'),
], [1.25, 5.5], 8.1, first_col_bold=True)
table(doc, ['觀察提示', '應記錄的具體證據（例）', '不可直接推論'], [
    ('目標', '候選人是否用可觀察動詞說明學員課後要完成什麼。', '「看起來很專業」不能代替目標證據。'),
    ('示範', '是否指出欄位、事實來源、交接責任與完成期限。', '投影片漂亮不能代替示範。'),
    ('回饋', '是否指出錯誤位置、原因及修正方法。', '只說「很好」不能證明理解。'),
    ('安全界線', '是否說明停止蒐集、保護資料並轉介的具體條件。', '沒有提到風險，不得記成「無風險」。'),
], [1.2, 3.0, 2.55], 7.8)

heading(doc, '7｜三人三評審比較與可重算結果', 1, page_break_before=False)
para(doc, '以下是依相同量尺完成的虛構評分示例。評審甲、乙、丙為示範角色代碼，不代表協會人員。各面向原始分數範圍 0–100。', size=9.2)
scores = {
    'C-01': [(85,86,87),(80,83,83),(78,76,80),(82,80,78),(90,92,88)],
    'C-02': [(76,78,80),(90,86,88),(74,76,75),(78,80,82),(84,86,85)],
    'C-03': [(90,92,88),(84,86,88),(88,90,86),(84,86,85),(92,90,88)],
}
totals = {}
for code in ('C-01','C-02','C-03'):
    heading(doc, f'{code}｜候選評分分解（SAMPLE）', 2)
    rationale = {
        'C-01': '試教步驟清楚，學習者適配尚可；評審認為需增加多元學員的替代練習與回到工作後的追蹤提示。界線面向高分只表示示範情境符合量尺，不代表真實資格核驗。',
        'C-02': '試教引導分數高，能帶動短練習；專業內容正確性與學員適配仍需在教材可讀性及例題涵蓋面補強。安全界線分數達示例門檻，資格仍須原件確認。',
        'C-03': '示範能連接紀錄、交接和界線情境，並為不同背景提供替代參與方式。以加權總分排序第一；推薦仍受資格原件、利益衝突及正式權限核驗條件拘束。',
    }[code]
    totals[code] = add_candidate_score(doc, code, scores[code], rationale)
assert abs(totals['C-01'] - 82.6) < 1e-9
assert abs(totals['C-02'] - 81.0) < 1e-9
assert abs(totals['C-03'] - 87.6) < 1e-9
table(doc, ['排序', '候選代碼', '總分', '分數門檻', '示例決定'], [
    ('1', 'C-03', '87.6', '達標', '可提出條件式推薦；先核原件及權限。'),
    ('2', 'C-01', '82.6', '達標', '備選；依相同條件完成核驗。'),
    ('3', 'C-02', '81.0', '達標', '備選；其學員適配補強事項列入條件。'),
], [.7, 1.15, .8, 1.0, 3.1], 8.2, first_col_bold=True)
para(doc, '排序核對：C-03 87.6 > C-01 82.6 > C-02 81.0。各候選的專業內容正確性、教學示範與引導、成人志工適配與可近性，以及個資、角色與安全界線均達示例門檻；若真實原件或權限有任何一項無法核驗，該候選立即 HOLD，不因排序高而例外。', size=10.5, bold=True)
heading(doc, '評審觀察摘要（合成）', 2)
table(doc, ['候選', '支持分數的具體觀察', '待確認／改善事項'], [
    ('C-01', '使用三欄示例指出事實、來源、交接責任，練習後要求學員復述。', '增加較大字版及口頭替代回覆；補上課後交接檢核。'),
    ('C-02', '先示範錯誤紀錄，再邀請學員提出修改，並用問題確認學習。', '教學材料需補不同閱讀程度例子；授權與原件仍須核驗。'),
    ('C-03', '用同一案例完成紀錄與交接，明確指出不記錄推測及停止／轉介點。', '取得真實資格原件、確認授課時段與費用授權後才可聘任。'),
], [0.8, 3.0, 2.95], 7.9, first_col_bold=True)

doc.add_page_break()
heading(doc, '8｜利益衝突處理與推薦規則', 1)
table(doc, ['狀況', '處理方式', '紀錄欄位'], [
    ('評審與候選人有近親、財務、督導、競爭或其他可能影響公正的關係', '評審於看分前申報；承辦判斷迴避範圍；需要時由有權者指定替代評審。', '申報人、關係類型、評估者、處置、替代人、日期。'),
    ('候選人資料涉及不必要個資', '只保留評選必要欄位；限制存取；依組織正式保存規則處理到期刪除。', '資料項目、用途、存取人、保存期限、刪除核准。'),
    ('資料不一致或資格查不到', '先凍結該候選評分結論；要求補件或以獨立來源查驗；保留往返紀錄。', '缺件、查驗來源、聯絡日期、期限、結果。'),
], [1.65, 3.35, 1.75], 7.8)
para(doc, '本合成情境設定三名評審均申報「無已知衝突」，僅供示範表單填寫。真實評審必須分別簽署日期與申報內容；空白不等於無衝突。', size=9, bold=True)

heading(doc, '9｜訓練管理委員會議程與決定紀錄（模擬）', 1)
callout(doc, '完全合成的會議情境', '以下會議日期、出席、權限、討論與決定均為虛構。文件以「假設權限已由真實文件證明」示範決議記錄形式，並不證明原始點協會已成立委員會或作出決議。', GOLD)
table(doc, ['欄位', 'SAMPLE 填寫示例／真實替換要求'], [
    ('會議基本資料', '合成日期：2026-11-10（未來示例）；14:00–15:00；會議方式：合成實體會議；地點：SAMPLE。正式使用需填實際日期、場地、通知與簽到原件。'),
    ('出席及法定人數', '主席 P-01、委員 M-01、M-02、M-03、紀錄員 S-01（全為合成代碼）；示範設定 3/3 有表決權委員出席。須附真實委員任期、簽到及法定人數算法。'),
    ('議程一：權限確認', '示範決定前先檢視章程、委員會任期、決議門檻、聘書簽發權與費用核定權；本樣本假設附件均已驗證，真實情形目前 UNKNOWN。'),
    ('議程二：需求與流程', '確認課程需求摘要、評選量尺、同題試教、三位評審的利益衝突申報及資格查驗清單。'),
    ('議程三：比較結果', '審閱 C-01 82.6、C-02 81.0、C-03 87.6；討論 C-03 教學適配較完整，但須在實際聘任前完成原件查驗及排程確認。'),
    ('合成決議', 'SAMPLE：在假設委員會權限有效、表決程序完備且原始資格查驗通過的前提下，決議優先推薦 C-03；承辦人應先完成所有 HOLD 項目，主席確認無未解異議後，始得依法定授權簽發聘書。此為教學用模擬文字，並非協會真實決議。'),
    ('反對／迴避／附件', '示例：無（合成設定）。真實會議須逐一記錄反對票、棄權、迴避、附件版本及頁數，不可用「無」代替未記錄。'),
    ('責任與期限', '承辦 A-01（合成）於 2026-11-17 前核對 C-03 原件、授課時段與費用授權；主席 P-01（合成）於核驗後依真實權限決定是否簽發；未完成則維持 HOLD。'),
], [1.45, 5.35], 7.8, first_col_bold=True)
para(doc, '決議確認欄（真實使用）：主席簽名／日期：________________；紀錄員簽名／日期：________________；附件清單：________________。此處空白且不得由本樣本代簽。', size=9)

heading(doc, '10｜聘書處理流程與未發出聘書樣本', 1)
table(doc, ['步驟', '必須先有的證據', '不得越過的關卡'], [
    ('1. 形成推薦', '完整評分、試教、資格／授權／衝突檢核。', '未達門檻或有 HOLD 不推薦。'),
    ('2. 委員會決定', '有效權限、通知、簽到、議事規則、逐項決議及附件。', '草稿、口頭意見或不明權限不能當正式核准。'),
    ('3. 核定條件', '課程版本、日期、總時數、費用來源、支付條件、教材及個資條款。', '未經授權不得自行承諾費用或期間。'),
    ('4. 簽發與送達', '正式聘書編號、核准人簽章、送達證明。', '不得預簽、倒填、代簽或以此樣本冒充原件。'),
    ('5. 接受與歸檔', '講師簽收／接受、版本、保存位置及權限。', '未接受不視為已完成任命。'),
], [1.0, 3.6, 2.2], 7.9, first_col_bold=True)
callout(doc, 'SAMPLE｜NOT ISSUED｜NOT VALID', '以下文字只示範聘書欄位與條件。不得列印後簽發；不得使用協會真實印信；日期、費用與角色代碼全為虛構。正式發出前須替換並核實所有資料，取得有權委員會決議及核定人簽章。', RED)
table(doc, ['聘書欄位', '示範文字（全部為合成草稿）'], [
    ('受聘人／編號', 'C-03（虛構代碼；非真實講師姓名）；聘書草稿編號 SAMPLE-APPT-001。'),
    ('課程與工作範圍', '課程：志工服務紀錄、交接與工作界線實務；四單元，共 360 分鐘。工作：依核定教案授課、帶領練習、回報學習狀況、完成必要教材修訂及課後紀錄。'),
    ('日期與地點', '樣本期間：2026-12-05；日期為虛構未來日期，實際開課時間與地點待核定。'),
    ('費用條件', '示例酬金 NT$4,800（360 分鐘 × 合成每小時 NT$800）。此數字未經協會核准，不代表支付承諾；正式使用需有預算、稅務與支付規則來源。'),
    ('履約條件', '遵守核定課程目標、角色界線、個資保護與場地安全規則；未經同意不得另行蒐集或公開學員資料；教材授權範圍以真實契約為準。'),
    ('取消／變更', '日期、時數、費用或教材有變更時，由雙方以可追溯方式確認並依授權程序核准；不可單方補填。'),
    ('簽章欄', '委任單位有權代表：________________／日期：____；受聘人接受：________________／日期：____。本樣本未簽發、未送達、未接受。'),
], [1.35, 5.45], 8.0, first_col_bold=True)

heading(doc, '11｜授課後回饋與改進閉環（尚未發生）', 1)
para(doc, '此節是任命後的追蹤表單設計，不是本案例已授課或已有成效。實際授課後應把觀察與學員回饋連回遴選面向，並作成續聘、補強或停止合作的有權決定。', size=9.2)
table(doc, ['追蹤時點', '收集資料', '判讀問題', '責任／處置'], [
    ('課前', '核定課程、教材版本、講師資格原件、場地與個資說明。', '人、課程、版本和核准條件是否一致？', '承辦核對；任一不一致先暫停開課。'),
    ('課中', '出席、授課完成度、學員練習、重大偏差或安全事件。', '教學是否達標；是否需即時調整？', '授課承辦記錄事實，不把未觀察寫成通過。'),
    ('課後 5 個工作日（示例時限）', '學員回饋、成果作品、講師自評及改進建議。', '目標達成情形如何；資料是否足以支持結論？', '承辦彙整；不足時標 UNKNOWN 並補證據。'),
    ('下次遴選前', '歷次回饋、改進閉環、申訴或事件處理。', '是否調整量尺、教材或聘任條件？', '有權者核定變更並保留版次。'),
], [1.25, 2.05, 1.8, 1.7], 7.7)

heading(doc, '12｜真實使用替換清冊與完整證據鏈', 1)
para(doc, '本節隨附一份逐筆替換清冊，逐筆列出合成資料、假設理由、跨文件引用、應取得的真實原始來源、負責核實角色、真實／合成狀態及 HOLD 條件，共 52 筆；不得直接把示範值複製為真實紀錄。下表列出主要證據鏈。', size=10.5)
table(doc, ['順序', '真實原始證據', '形成／保管責任', '缺件時狀態'], [
    ('1', '現行章程、委員會設置／任期／權限及聘書簽發授權。', '秘書／有權治理者核對正本或受控副本。', 'HOLD：不得開會作成正式決議。'),
    ('2', '有日期的職能需求、課程目標、對象與課程計畫版次。', '需求提出人與課程承辦；版本連結至課程檔案。', 'HOLD：不得以本合成需求代替。'),
    ('3', '候選人本人履歷、資格原件及獨立查驗紀錄。', '遴選承辦；保留來源、日期、方式與核驗人。', 'HOLD：未查驗不得列為資格 PASS。'),
    ('4', '三位評審原始評分、試教觀察、利益衝突申報。', '各評審獨立簽署；承辦保留原始版本。', 'HOLD：不可補寫一致意見掩蓋差異。'),
    ('5', '合法通知、簽到、會議記錄、附件、逐項決議及表決結果。', '有權主席與紀錄員；保存更正軌跡。', 'HOLD：不得稱已核准。'),
    ('6', '有權核定的聘書、簽發與送達證明、講師接受／回簽。', '授權簽發人及聘任承辦；登錄原件位置。', 'HOLD：不得授課或支付，除非另有有效授權依據。'),
    ('7', '課前核對、課程執行、課後評估及改進紀錄。', '課程承辦與評估責任人。', 'UNKNOWN：未實際發生前不可填結果。'),
], [.55, 3.1, 2.05, 1.1], 7.5)

heading(doc, '13｜文件狀態與驗收控制', 1)
table(doc, ['欄位', '狀態／內容'], [
    ('範例定位', '講師遴選與聘任流程範例；對應官方指標 12c（需求代碼 I12-E03；官方方框 058）。'),
    ('候選狀態', '候選示範文件；不是協會正式核准、任命或真實證據。'),
    ('來源依據', '官方訓練品質評核表第 4 頁指標 12c；寫作手冊指標 10 與 12；官方方框 058 的內容方向。來源版本資訊另存於驗收附屬檔案。'),
    ('一致性核對', '課程 360 分鐘＝90+90+120+60；五權重合計 100%；加權總分 C-01 82.6、C-02 81.0、C-03 87.6；排序一致。'),
    ('形式與樣本界線', '全文件及表單標為 SAMPLE/SYNTHETIC；無真實人名、證照影本、核准簽章或任命事實。聘書標 NOT ISSUED / NOT VALID。'),
    ('正式使用前的驗收要求', '須完成技術版面檢查、全頁目視、獨立語意複核與人工驗收；本文件目前僅為候選示範，尚未取得任何正式驗收。'),
], [1.5, 5.25], 9.5, first_col_bold=True)

doc.core_properties.title = '講師遴選與聘任證據包｜DOCX 0055 R02｜SAMPLE/SYNTHETIC'
doc.core_properties.subject = 'I12-E03 / official checkbox 058; instructor selection and appointment evidence pack'
doc.core_properties.author = 'TTQS ONE candidate document builder'
doc.core_properties.keywords = 'SAMPLE; SYNTHETIC; I12-E03; checkbox 058; 0055'
doc.save(OUT)

# A 52-row actionable replacement register. Every row carries the full Mission §4
# field set: fact_id, case_id, doc_id, a printable locator that exists in the DOCX
# body, the synthetic value/field, the assumption rationale, cross-document
# references, the REAL source to obtain, the accountable verifier, the
# REAL/SYNTHETIC status and the HOLD condition.
CASE = 'SYN-0033-PLAN-01'
DOCID = '0055'
CASE_DOCS = '0033 R03／0049 R01／0050 R01／0052 R02（同一合成案例）'
NEW_SYN = '無（本件新增合成，無跨文件事實引用）'
UNK = 'SYNTHETIC／REAL 未知'
register = [
 ('FAC-0055-001','FAC-0055-001',CASE,DOCID,'封面案例行「案例 SYN-0033-PLAN-01」；第 13 節一致性核對列',
  '合成案例識別碼','沿用既有四份同案例文件的合成案例碼，僅作跨文件對齊',CASE_DOCS,
  '正式課程計畫與核准版次','課程承辦','SYNTHETIC','未取得正式計畫不得沿用此 ID'),
 ('FAC-0055-002','FAC-0055-002',CASE,DOCID,'封面案例行；第 10 節聘書表「課程與工作範圍」列',
  '合成課程名稱','沿用同案例課程名稱，供招募與聘書草案一致',CASE_DOCS,
  '正式課程計畫','課程承辦','SYNTHETIC','名稱或範圍不一致即 HOLD'),
 ('FAC-0055-003','FAC-0055-003',CASE,DOCID,'封面案例行「合成規劃 360 分鐘」；第 10 節聘書表「課程與工作範圍」列',
  '合成總時數 360 分鐘','沿用同案例 360 分鐘，使聘任範圍完整',CASE_DOCS,
  '核定課程計畫及時程','課程承辦','SYNTHETIC','未核定時數不得寫入聘書'),
 ('FAC-0055-004','FAC-0055-004',CASE,DOCID,'封面案例行「（90+90+120+60）」；第 13 節一致性核對列',
  '合成單元時數 90+90+120+60','沿用同案例四單元時數，證明 90+90+120+60=360',CASE_DOCS,
  '核定課程計畫各單元頁','課程設計者','SYNTHETIC','總和不等於 360 即 HOLD'),
 ('FAC-0055-005','FAC-0055-005',CASE,DOCID,'第 2 節需求表「志工紀錄欄位易漏、事實與推測混寫」列',
  '合成需求缺口：紀錄欄位漏填','自同案例需求轉寫為可觀察的講師職能落差','0033 R03 需求調查；0049 R01 計畫',
  '需求調查原表及彙整方法','需求提出人','SYNTHETIC','無可追溯需求來源則重做需求分析'),
 ('FAC-0055-006','FAC-0055-006',CASE,DOCID,'第 2 節需求表「志工紀錄欄位易漏、事實與推測混寫」列',
  '合成需求缺口：事實與推測混寫','同上，用於設計事實／推測辨識的試教情境','0033 R03 需求調查；0049 R01 計畫',
  '去識別化需求案例及分析','需求提出人','SYNTHETIC','不得把虛構案例標為真實個案'),
 ('FAC-0055-007','FAC-0055-007',CASE,DOCID,'第 2 節需求表「交接時責任、期限、未結事項不清」列',
  '合成需求缺口：交接責任與期限不清','自同案例交接需求轉寫，供試教評分情境','0033 R03 需求調查；0049 R01 計畫',
  '需求訪談／工作流程紀錄','需求提出人','SYNTHETIC','缺原始資料則標 UNKNOWN'),
 ('FAC-0055-008','FAC-0055-008',CASE,DOCID,'第 2 節需求表「志工容易越過角色界線或傳遞敏感資訊」列',
  '合成需求缺口：角色界線與敏感資訊','自同案例角色界線需求轉寫，支撐安全面向量尺','0033 R03 需求調查；0049 R01 計畫',
  '正式服務規範與匿名化需求資料','治理責任人','SYNTHETIC','未核准規範不得當協會政策'),
 ('FAC-0055-009','FAC-0055-009',CASE,DOCID,'第 2 節需求表「成人學員背景差異大、部分以手機參與」列',
  '合成需求缺口：成人學員可近性','示範成人學員可近性面向，非真實學員資料','0033 R03 需求調查；0049 R01 計畫',
  '學員需求／合理調整紀錄','課程承辦','SYNTHETIC','未確認需求不得宣稱已滿足'),
 ('FAC-0055-010','FAC-0055-010',CASE,DOCID,'第 4 節量尺表「權重」欄；第 7 節各候選分解表「權重%」欄',
  '合成評選權重 25/25/20/20/10','示範可重算的加權量尺，權重合計 100%',NEW_SYN,
  '正式核定遴選辦法','權責主管／委員會','SYNTHETIC','未核定不得用於正式遴選'),
 ('FAC-0055-011','FAC-0055-011',CASE,DOCID,'第 4 節狀態表「通過門檻（合成示例）」列',
  '合成遴選門檻總分 80','設定示例總分門檻，使推薦有明確判準',NEW_SYN,
  '正式核定遴選辦法','權責主管／委員會','SYNTHETIC','門檻未授權不得宣稱合格'),
 ('FAC-0055-012','FAC-0055-012',CASE,DOCID,'第 4 節狀態表「通過門檻（合成示例）」列',
  '合成個別面向最低分 75','設定示例面向門檻，避免單項高分掩蓋弱項',NEW_SYN,
  '正式核定量尺','權責主管／委員會','SYNTHETIC','不可事後為候選人改門檻'),
 ('FAC-0055-013','FAC-0055-013',CASE,DOCID,'第 4 節狀態表「通過門檻（合成示例）」列',
  '合成安全界線最低分 80','設定示例安全門檻，配合非補償性條件',NEW_SYN,
  '正式核定量尺及風險規則','權責主管／委員會','SYNTHETIC','安全低於門檻即 HOLD/REJECT'),
 ('FAC-0055-014','FAC-0055-014',CASE,DOCID,'第 3 節角色表「評審甲／乙／丙」列；第 7 節分解表「評審甲／乙／丙」欄',
  '虛構評審代碼 評審甲／乙／丙','用角色代碼代替真實評審姓名',NEW_SYN,
  '真實委任／評審名單與簽到','遴選承辦','SYNTHETIC','未確認身份不得把代碼當實際評審'),
 ('FAC-0055-015','FAC-0055-015',CASE,DOCID,'第 5 節候選核驗表「C-01（合成）」欄',
  '虛構候選代碼 C-01','用候選代碼代替真實候選人，避免識別',NEW_SYN,
  '候選人正式申請與同意','遴選承辦','SYNTHETIC','無本人申請與同意即不得列名'),
 ('FAC-0055-016','FAC-0055-016',CASE,DOCID,'第 5 節候選核驗表「C-02（合成）」欄',
  '虛構候選代碼 C-02','用候選代碼代替真實候選人，避免識別',NEW_SYN,
  '候選人正式申請與同意','遴選承辦','SYNTHETIC','無本人申請與同意即不得列名'),
 ('FAC-0055-017','FAC-0055-017',CASE,DOCID,'第 5 節候選核驗表「C-03（合成）」欄',
  '虛構候選代碼 C-03','用候選代碼代替真實候選人，避免識別',NEW_SYN,
  '候選人正式申請與同意','遴選承辦','SYNTHETIC','無本人申請與同意即不得列名'),
 ('FAC-0055-018','FAC-0055-018',CASE,DOCID,'第 7 節「C-01｜候選評分分解（SAMPLE）」表「專業內容正確性」列',
  'C-01 專業內容正確性評分 85/86/87','以合成評分向量示範可重算平均與加權',NEW_SYN,
  '三份原始評分表及試教紀錄','各評審','SYNTHETIC','無原始分數不得保留平均'),
 ('FAC-0055-019','FAC-0055-019',CASE,DOCID,'第 7 節「C-01｜候選評分分解（SAMPLE）」表「教學示範與引導」列',
  'C-01 教學示範與引導評分 80/83/83','以合成評分向量示範可重算平均與加權',NEW_SYN,
  '三份原始評分表及試教紀錄','各評審','SYNTHETIC','無原始分數不得保留平均'),
 ('FAC-0055-020','FAC-0055-020',CASE,DOCID,'第 7 節「C-01｜候選評分分解（SAMPLE）」表「成人志工適配與可近性」列',
  'C-01 成人志工適配與可近性評分 78/76/80','以合成評分向量示範可重算平均與加權',NEW_SYN,
  '三份原始評分表及可近性觀察','各評審','SYNTHETIC','無原始分數不得保留平均'),
 ('FAC-0055-021','FAC-0055-021',CASE,DOCID,'第 7 節「C-01｜候選評分分解（SAMPLE）」表「實務應用與移轉」列',
  'C-01 實務應用與移轉評分 82/80/78','以合成評分向量示範可重算平均與加權',NEW_SYN,
  '三份原始評分與演練紀錄','各評審','SYNTHETIC','無原始分數不得保留平均'),
 ('FAC-0055-022','FAC-0055-022',CASE,DOCID,'第 7 節「C-01｜候選評分分解（SAMPLE）」表「個資、角色與安全界線」列',
  'C-01 個資、角色與安全界線評分 90/92/88','以合成評分向量示範可重算平均與加權',NEW_SYN,
  '三份原始評分及風險情境答覆','各評審','SYNTHETIC','無原始分數不得保留平均'),
 ('FAC-0055-023','FAC-0055-023',CASE,DOCID,'第 7 節「C-02｜候選評分分解（SAMPLE）」表「專業內容正確性」列',
  'C-02 專業內容正確性評分 76/78/80','以合成評分向量示範可重算平均與加權',NEW_SYN,
  '三份原始評分表及試教紀錄','各評審','SYNTHETIC','無原始分數不得保留平均'),
 ('FAC-0055-024','FAC-0055-024',CASE,DOCID,'第 7 節「C-02｜候選評分分解（SAMPLE）」表「教學示範與引導」列',
  'C-02 教學示範與引導評分 90/86/88','以合成評分向量示範可重算平均與加權',NEW_SYN,
  '三份原始評分表及試教紀錄','各評審','SYNTHETIC','無原始分數不得保留平均'),
 ('FAC-0055-025','FAC-0055-025',CASE,DOCID,'第 7 節「C-02｜候選評分分解（SAMPLE）」表「成人志工適配與可近性」列',
  'C-02 成人志工適配與可近性評分 74/76/75','以合成評分向量示範可重算平均與加權',NEW_SYN,
  '三份原始評分表及可近性觀察','各評審','SYNTHETIC','無原始分數不得保留平均'),
 ('FAC-0055-026','FAC-0055-026',CASE,DOCID,'第 7 節「C-02｜候選評分分解（SAMPLE）」表「實務應用與移轉」列',
  'C-02 實務應用與移轉評分 78/80/82','以合成評分向量示範可重算平均與加權',NEW_SYN,
  '三份原始評分與演練紀錄','各評審','SYNTHETIC','無原始分數不得保留平均'),
 ('FAC-0055-027','FAC-0055-027',CASE,DOCID,'第 7 節「C-02｜候選評分分解（SAMPLE）」表「個資、角色與安全界線」列',
  'C-02 個資、角色與安全界線評分 84/86/85','以合成評分向量示範可重算平均與加權',NEW_SYN,
  '三份原始評分及風險情境答覆','各評審','SYNTHETIC','無原始分數不得保留平均'),
 ('FAC-0055-028','FAC-0055-028',CASE,DOCID,'第 7 節「C-03｜候選評分分解（SAMPLE）」表「專業內容正確性」列',
  'C-03 專業內容正確性評分 90/92/88','以合成評分向量示範可重算平均與加權',NEW_SYN,
  '三份原始評分表及試教紀錄','各評審','SYNTHETIC','無原始分數不得保留平均'),
 ('FAC-0055-029','FAC-0055-029',CASE,DOCID,'第 7 節「C-03｜候選評分分解（SAMPLE）」表「教學示範與引導」列',
  'C-03 教學示範與引導評分 84/86/88','以合成評分向量示範可重算平均與加權',NEW_SYN,
  '三份原始評分表及試教紀錄','各評審','SYNTHETIC','無原始分數不得保留平均'),
 ('FAC-0055-030','FAC-0055-030',CASE,DOCID,'第 7 節「C-03｜候選評分分解（SAMPLE）」表「成人志工適配與可近性」列',
  'C-03 成人志工適配與可近性評分 88/90/86','以合成評分向量示範可重算平均與加權',NEW_SYN,
  '三份原始評分表及可近性觀察','各評審','SYNTHETIC','無原始分數不得保留平均'),
 ('FAC-0055-031','FAC-0055-031',CASE,DOCID,'第 7 節「C-03｜候選評分分解（SAMPLE）」表「實務應用與移轉」列',
  'C-03 實務應用與移轉評分 84/86/85','以合成評分向量示範可重算平均與加權',NEW_SYN,
  '三份原始評分與演練紀錄','各評審','SYNTHETIC','無原始分數不得保留平均'),
 ('FAC-0055-032','FAC-0055-032',CASE,DOCID,'第 7 節「C-03｜候選評分分解（SAMPLE）」表「個資、角色與安全界線」列',
  'C-03 個資、角色與安全界線評分 92/90/88','以合成評分向量示範可重算平均與加權',NEW_SYN,
  '三份原始評分及風險情境答覆','各評審','SYNTHETIC','無原始分數不得保留平均'),
 ('FAC-0055-033','FAC-0055-033',CASE,DOCID,'第 5 節候選核驗表「利益衝突」列；第 8 節說明',
  '合成評審無衝突申報','示範衝突申報欄位填寫，非真實申報結果',NEW_SYN,
  '各評審本人簽署申報與處理決定','評審／遴選承辦','SYNTHETIC','空白或未申報即 HOLD'),
 ('FAC-0055-034','FAC-0055-034',CASE,DOCID,'第 5 節候選核驗表「專業資格」列',
  '合成資格查核 PASS 標記','示範資格核驗欄位值，無任何真實證照',NEW_SYN,
  '資格原件及發證單位查驗紀錄','遴選承辦',UNK,'無原件和獨立查驗即 HOLD'),
 ('FAC-0055-035','FAC-0055-035',CASE,DOCID,'第 5 節候選核驗表「教學授權／可用時段」列',
  '合成候選人可授課時段','示範可用時段欄位，非真實承諾',NEW_SYN,
  '候選人書面同意與核定時程','課程承辦','SYNTHETIC','未確認時段不得承諾'),
 ('FAC-0055-036','FAC-0055-036',CASE,DOCID,'第 5 節候選核驗表「個資及作品使用」列',
  '合成限本次遴選的資料同意','示範個資授權欄位，非真實同意書',NEW_SYN,
  '本人告知、同意與撤回紀錄','資料管理責任人','SYNTHETIC','無合法依據不得保存或公開'),
 ('FAC-0055-037','FAC-0055-037',CASE,DOCID,'封面 SAMPLE 說明框；第 9 節「完全合成的會議情境」說明框',
  '模擬委員會存在與權限（明示為假設）','明示委員會權限為假設，避免假稱已成立',NEW_SYN,
  '章程、設置文件、任期及授權原件','治理責任人',UNK,'權限 UNKNOWN 即不得正式開會決議'),
 ('FAC-0055-038','FAC-0055-038',CASE,DOCID,'第 9 節議程表「出席及法定人數」列',
  '合成出席 3/3','示範法定人數記錄格式，非真實簽到',NEW_SYN,
  '真實通知、簽到及法定人數規則','委員會紀錄員',UNK,'無原始簽到／法定人數來源即 HOLD'),
 ('FAC-0055-039','FAC-0055-039',CASE,DOCID,'第 9 節議程表「會議基本資料」列',
  '模擬會議日期 2026-11-10','用未來合成日期示範會議基本資料欄位',NEW_SYN,
  '實際會議通知、簽到與議事錄','委員會紀錄員',UNK,'不可倒填或沿用示範日期'),
 ('FAC-0055-040','FAC-0055-040',CASE,DOCID,'第 9 節議程表「合成決議」列；第 7 節排序表',
  '模擬推薦 C-03 決議','示範附條件的模擬決議，非真實核准',NEW_SYN,
  '有權委員會逐項決議及表決紀錄','委員會主席／紀錄員',UNK,'無有效決議不得稱已核准'),
 ('FAC-0055-041','FAC-0055-041',CASE,DOCID,'第 9 節議程表「責任與期限」列',
  '虛構責任人代碼 A-01','用代碼示範責任人欄位，非真實指派',NEW_SYN,
  '實際分工授權與責任確認','課程主管',UNK,'未指派不得假填責任完成'),
 ('FAC-0055-042','FAC-0055-042',CASE,DOCID,'第 9 節議程表「責任與期限」列',
  '虛構期限 2026-11-17','示範追蹤期限欄位，非真實決議期限',NEW_SYN,
  '實際會議決議及追蹤紀錄','委員會紀錄員',UNK,'不可冒用或倒填期限'),
 ('FAC-0055-043','FAC-0055-043',CASE,DOCID,'第 10 節聘書表「受聘人／編號」列',
  '虛構聘書草稿編號 SAMPLE-APPT-001','示範聘書編號欄位，非正式發文編號',NEW_SYN,
  '正式文件編碼台帳','文件管制人',UNK,'未核准不得登錄為正式發文'),
 ('FAC-0055-044','FAC-0055-044',CASE,DOCID,'第 10 節聘書表「日期與地點」列',
  '虛構聘書日期 2026-12-05','用未來合成日期示範聘書日期欄位',NEW_SYN,
  '正式決議、實際發文與送達證明','授權簽發人',UNK,'不可預簽、倒填或提前發出'),
 ('FAC-0055-045','FAC-0055-045',CASE,DOCID,'第 10 節聘書表「費用條件」列',
  '合成酬金 NT$4,800','示範酬金欄位與計算，明示未經核准',NEW_SYN,
  '核准預算、契約、稅務及支付依據','財務／授權主管',UNK,'無核准預算與授權不得承諾或支付'),
 ('FAC-0055-046','FAC-0055-046',CASE,DOCID,'第 10 節聘書表「費用條件」列',
  '合成費率 NT$800／小時','示範費率欄位，由 360 分鐘×800 得 4800',NEW_SYN,
  '核准費率表／契約','財務／授權主管',UNK,'不得將示例數字當協會標準'),
 ('FAC-0055-047','FAC-0055-047',CASE,DOCID,'第 10 節聘書表「履約條件」列',
  '示範教材授權範圍','示範教材授權條款，真實範圍以待契約',NEW_SYN,
  '真實授權契約與教材權利來源','課程承辦／權責主管','SYNTHETIC','未授權不得重製或公開'),
 ('FAC-0055-048','FAC-0055-048',CASE,DOCID,'第 8 節衝突處理表「候選人資料涉及不必要個資」列；第 10 節聘書表「取消／變更」列',
  '未指定的資料保存期限','示範保存期限欄位，真實規則未知',NEW_SYN,
  '現行保存規則及核定表單','資料管理責任人',UNK,'規則未知時限制存取且不得任意銷毀'),
 ('FAC-0055-049','FAC-0055-049',CASE,DOCID,'第 10 節聘書表「取消／變更」列',
  '示範取消與變更條件','示範變更條件條款，非真實契約',NEW_SYN,
  '雙方有效契約及核准變更紀錄','契約承辦','SYNTHETIC','無授權條款不得主張已同意'),
 ('FAC-0055-050','FAC-0055-050',CASE,DOCID,'第 11 節追蹤表「課中」列',
  '實際授課完成（未發生／UNKNOWN）','明示課程尚未發生，避免假造成果','0033 R03／0049 R01／0050 R01／0052 R02（同一課程鏈）',
  '出席、課程紀錄、講師回報及成果','課程承辦',UNK,'未發生前不得填成果或評語'),
 ('FAC-0055-051','FAC-0055-051',CASE,DOCID,'第 11 節追蹤表「課後 5 個工作日（示例時限）」列',
  '課後回饋與效能評估（待未來資料）','示範課後回饋工具，非真實回收資料','0033 R03／0049 R01／0050 R01／0052 R02（同一課程鏈）',
  '匿名回饋、作業、評估工具與判讀紀錄','評估責任人',UNK,'無資料時結論 UNKNOWN'),
 ('FAC-0055-052','FAC-0055-052',CASE,DOCID,'第 11 節追蹤表「下次遴選前」列',
  '改進／續聘／停聘決定（待未來有權決定）','示範改善／續聘判讀欄位，非真實決定','0033 R03／0049 R01／0050 R01／0052 R02（同一課程鏈）',
  '課後證據、改善追蹤及正式核准','有權委員會／主管',UNK,'不得預先宣稱續聘或成效'),
]
with REG.open('w', encoding='utf-8-sig', newline='') as f:
    w = csv.writer(f)
    w.writerow(['register_id','fact_id','case_id','doc_id','docx_locator',
                'synthetic_value_or_field','assumption_rationale','cross_doc_reference',
                'real_source_to_obtain','responsible_verifier','evidence_status','hold_condition'])
    w.writerows(register)

print(f'DOCX={OUT}')
print(f'REGISTER={REG}')
print(f'REGISTER_ROWS={len(register)}')
print(f'SCORES={totals}')
