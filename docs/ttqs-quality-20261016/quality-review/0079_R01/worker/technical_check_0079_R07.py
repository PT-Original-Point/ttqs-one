import csv, hashlib, json, re, zipfile
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
DOCX=ROOT/'candidate/0079_R07.docx'
BODY=ROOT/'candidate/0079_R07_body.md'
OLD=ROOT/'candidate/0079_R05.docx'
REGISTER=ROOT/'inputs/SYNTHETIC_REPLACEMENT_REGISTER_0079_R01.csv'
EXTRACT=ROOT/'evidence/docx_extracted_R07.md'
RECEIPT=ROOT/'evidence/technical_review_R07.json'
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def norm(s): return re.sub(r'\s+',' ',s).strip()
def main():
    with zipfile.ZipFile(DOCX) as z:
        bad=z.testzip()
        missing=sorted({'[Content_Types].xml','word/document.xml','word/styles.xml'}-set(z.namelist()))
    d=Document(DOCX); paras=[]; tables=[]
    for p in d.paragraphs:
        if norm(p.text): paras.append(norm(p.text))
    for t in d.tables:
        rows=[[norm(c.text) for c in r.cells] for r in t.rows]
        tables.append(rows)
    text='\n'.join(paras+['\n'.join(' | '.join(r) for r in t) for t in tables])
    EXTRACT.write_text(text+'\n',encoding='utf-8',newline='\n')
    body=BODY.read_text(encoding='utf-8-sig')
    reg=list(csv.DictReader(REGISTER.open(encoding='utf-8-sig',newline='')))
    allcells=[r for t in tables for r in t]
    neighbor80=any(len(r)>=5 and r[0]=='0080' and 'I17B-E02' in r[1] and '085' in r[1] for r in allcells)
    neighbor81=any(len(r)>=5 and r[0]=='0081' and 'I17B-E03' in r[1] and '086' in r[1] for r in allcells)
    old_wrong_rows=any(len(r)>=2 and r[0] in {'085','086'} and r[1] in {'085','086'} for r in allcells)
    checks={
      'zip_integrity':bad is None,
      'required_ooxml_parts':not missing,
      'docx_sha_matches_build_receipt':sha(DOCX)==json.loads((ROOT/'evidence/build_receipt_R07.json').read_text(encoding='utf-8-sig'))['docx_sha256'],
      'document_title':d.core_properties.title=='訓後個別職涯與學習規劃建議表 0079 I17B-E01',
      'primary_identity': 'I17B-E01' in text and '084' in text,
      'neighbor_doc_0080_maps_to_checkbox_085':neighbor80,
      'neighbor_doc_0081_maps_to_checkbox_086':neighbor81,
      'old_docid_checkbox_collision_absent':not old_wrong_rows,
      'no_misnumbered_0085_course_review':'0085 課程評量檢討' not in text and '0086 結案行動計畫查核' not in text,
      'candidate_and_real_boundary':'SAMPLE/SYNTHETIC' in text and 'REAL=UNKNOWN/HOLD' in text and 'NOT_GRANTED' in text,
      'composite_checkbox_084_held':'084' in text and 'PARTIAL/HOLD' in text,
      'replacement_register_preserved':len(reg)==104,
      'old_candidate_unchanged':sha(OLD)=='d6689cef08d43f92b7110f3371856ad49991db69147334d5a494b1f7e5637cd2',
      'body_revision_R07':'候選修訂**：R07' in body,
    }
    receipt={'schema':'TTQS_ONE_0079_TECHNICAL_REVIEW_R07.v1','document_id':'0079','requirement_id':'I17B-E01','official_checkbox_serial':'084','candidate_docx_path':'candidate/0079_R07.docx','candidate_docx_sha256':sha(DOCX),'candidate_docx_bytes':DOCX.stat().st_size,'body_sha256':sha(BODY),'unique_blueprint_sha256':sha(ROOT/'inputs/UNIQUE_BLUEPRINT_0079_R02.md'),'source_packet_sha256':sha(ROOT/'inputs/SOURCE_PACKET_0079_R02.md'),'replacement_register_rows':len(reg),'paragraphs':len(paras),'tables':len(tables),'checks':checks,'zip_bad_member':bad,'missing_ooxml_parts':missing,'status':'PASS' if all(checks.values()) else 'FAIL','formal_acceptance':'NOT_GRANTED','real_evidence':'UNKNOWN/HOLD'}
    RECEIPT.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')
    print(json.dumps({'sha256':receipt['candidate_docx_sha256'],'status':receipt['status'],'failed':[k for k,v in checks.items() if not v],'tables':len(tables)},ensure_ascii=False))
if __name__=='__main__': main()

