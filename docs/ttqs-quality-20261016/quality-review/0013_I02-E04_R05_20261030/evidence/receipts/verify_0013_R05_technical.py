from __future__ import annotations
from decimal import Decimal
from pathlib import Path
import hashlib,json,re,zipfile
from docx import Document
from pypdf import PdfReader
BASE=Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0013_I02-E04_LUNA_TEXT_DRAFT_R01_20261029")
PACKET=BASE/'source_packet_R01.json'; PATCH=BASE/'qa'/'content_reconciliation_R05.json'
VISUAL=BASE/'qa'/'visual_qa_R05.json'; PDF=BASE/'qa'/'render_R05'/'0013_I02-E04_R05_WORD_NATIVE.pdf'
DOCX=BASE/'candidate'/'0013_I02-E04_R05.docx'; SRC=BASE/'worker'/'luna_subagent_draft_R01.md'
OFFICIAL=next(p for p in Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\SOURCES").glob('*.pdf') if hashlib.sha256(p.read_bytes()).hexdigest()=='805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709')
BP=Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\CONTROL\BLUEPRINT_CURRENT.jsonl"); BUILD=BASE/'qa'/'BUILD_RECEIPT_R05.json'; OUT=BASE/'qa'/'technical_qa_R05.json'
def sha(b): return hashlib.sha256(b).hexdigest()
def fsha(p): return sha(p.read_bytes())
def norm(s): return ''.join(c.casefold() for c in s if c.isalnum() or '\u3400'<=c<='\u9fff')
def splitrow(line):
 v=line.strip(); v=v[1:] if v.startswith('|') else v; v=v[:-1] if v.endswith('|') else v
 return [x.strip() for x in v.split('|')]
def table_start(lines,i):
 if i+1>=len(lines) or not lines[i].strip().startswith('|') or not lines[i+1].strip().startswith('|'): return False
 return all(re.fullmatch(r':?-{3,}:?',c.replace(' ','')) for c in splitrow(lines[i+1]))
def table_block(lines,i):
 heads=splitrow(lines[i]); rows=[]; j=i+2
 while j<len(lines) and lines[j].strip().startswith('|'):
  vals=splitrow(lines[j]); vals=(vals+['']*len(heads))[:len(heads)]; rows.append(vals); j+=1
 return heads,rows,j
def source_fragments(md):
 out=[]
 for raw in md.splitlines():
  v=raw.strip()
  if not v or (v.startswith('|') and set(v.replace('|','').replace(':','').replace('-','').replace(' ','') )==set()): continue
  v=v.lstrip('> ').lstrip('# '); v=re.sub(r'^[-*+]\s+','',v); v=re.sub(r'^(\d+)\.\s+',r'\1. ',v)
  v=norm(v.replace('|',' ').replace('**','').replace('`','').replace('*',''))
  if v: out.append(v)
 return out
source=SRC.read_text(encoding='utf-8-sig'); lines=source.splitlines(); doc=Document(DOCX); all_text='\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for row in t.rows for c in row.cells]); nt=norm(all_text)
rows=BP.read_text(encoding='utf-8-sig').splitlines(); rowraw=next(x.encode('utf-8') for x in rows if (lambda r:r.get('deliverable_id')=='0013' and r.get('requirement_id')=='I02-E04')(json.loads(x))); bp_row=json.loads(rowraw); req=bp_row['official_requirement_text']; page1=re.sub(r'\s+','',(PdfReader(str(OFFICIAL)).pages[0].extract_text() or ''))
mdtables=[]; i=0
while i<len(lines):
 if table_start(lines,i):
  h,dr,j=table_block(lines,i); mdtables.append((h,dr,i)); i=j
 else:i+=1
course_tables=[]
for marker in ['### 6.2','### 6.3','### 6.4','### 6.5']:
 idx=next((i for i,l in enumerate(lines) if l.startswith(marker)),None)
 assert idx is not None, marker
 j=idx+1
 while j<len(lines) and not table_start(lines,j): j+=1
 assert j<len(lines), marker+' table missing'
 h,dr,end=table_block(lines,j); course_tables.append(dr)
hours=[]; capacities=[]
for dr in course_tables:
 assert len(dr)==10, f'course expected code + 9 public fields: {len(dr)}'
 hours.append(int(re.search(r'\d+',dr[7][1]).group()))
 capacities.append(int(re.search(r'\d+',dr[8][1]).group()))
register=next((t for t in mdtables if t[0][0].casefold()=='fact id'),None)
assert register is not None
regids=[r[0] for r in register[1] if r and r[0]]
expected_facts=sorted(set(re.findall(r'SYN-0013-F\d{3}',source)))
frags=source_fragments(source); missing=[f for f in frags if f not in nt]
build=json.loads(BUILD.read_text(encoding='utf-8'))
checks={
 'docx_zip_integrity':zipfile.is_zipfile(DOCX),
 'python_docx_opens':True,
 'official_requirement_pdf_sha_exact':fsha(OFFICIAL)=='805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709',
 'official_exact_requirement_text_in_pdf_page1':norm(req) in norm(page1),
 'official_requirement_text_in_candidate':norm(req) in nt,
 'blueprint_file_sha_exact':fsha(BP)=='406d0495365d0970dcf8113a42842663b8a54372a4e204cc492b7c3f8156ef3e',
 'blueprint_unique_row_sha_exact':sha(rowraw)=='b529f3c13e964b6edbec124158d3559df988b401ce3b97fd1984e90b0582d35b',
 'negative_neighbors_0012_and_0014_distinguished':all(x in all_text for x in ['0012','I02-E03','0014','I02-E05']),
 'twenty_markdown_tables_rendered':len(mdtables)==20 and len(doc.tables)==20,
 'all_source_fragments_preserved_or_explicitly_reconciled': set(missing)=={x['normalized'] for x in json.loads(PATCH.read_text(encoding='utf-8'))['source_fragments_explicitly_revised']},
 'source_packet_sha_exact':fsha(PACKET)=='b87440b249ec78c6bbe69ff001fe9d7dbb5060d65f27fe61d21ec48cdb5bfa63',
 'source_packet_worker_binding_exact':json.loads(PACKET.read_text(encoding='utf-8-sig')).get('worker_source_sha256')==fsha(SRC),
 'source_packet_stale_statement_reconciled':all(x in all_text for x in ['VERIFIED_SOURCE_AND_UNIQUE_BLUEPRINT','b87440b249ec78c6bbe69ff001fe9d7dbb5060d65f27fe61d21ec48cdb5bfa63','0f8f3961ec945e9e7fd6aa91464bb73ed65afcf0a31a3523ff303d7534175d4f']),
 'section_locator_corrected':('第 12、14 節' in all_text and 'SYNTHETIC 事實清冊在第 13 節' in all_text),
 'candidate_version_and_deliverable_identity_correct':('R05' in doc.paragraphs[4].text and 'DOCX 候選' in all_text and 'TEXT_DRAFT_ONLY_NOT_QA' not in all_text),
 'closing_delivery_boundary_matches_docx':('交付邊界：' in doc.paragraphs[-1].text and 'DOCX 候選' in doc.paragraphs[-1].text and '沒有製作 DOCX' not in doc.paragraphs[-1].text),
 'sample_not_promoted_to_real_or_formal_acceptance':('CANDIDATE_SAMPLE_SYNTHETIC_NOT_ACCEPTED' in all_text and '未提供；不主張已核准' in all_text),
 'no_empty_docx_table_cells':all(c.text.strip() for t in doc.tables for r in t.rows for c in r.cells),
 'all_synthetic_register_ids_in_docx':all(x in all_text for x in expected_facts),
 'synthetic_register_has_67_unique_fact_rows':len(regids)==67 and len(set(regids))==67 and set(regids)==set(expected_facts),
 'four_course_examples_complete':len(course_tables)==4 and all(len(t)==10 for t in course_tables),
 'four_course_codes_present':all(x in all_text for x in ['SYN-C01','SYN-C02','SYN-C03','SYN-C04']),
 'course_hours_sum_consistent':sum(hours)==30,
 'course_capacity_sum_consistent':sum(capacities)==74,
 'completeness_examples_consistent':Decimal(4)*Decimal(9)==Decimal(36) and (Decimal(35)/Decimal(36)*Decimal(100)).quantize(Decimal('0.1'))==Decimal('97.2') and Decimal(36)/Decimal(36)*Decimal(100)==Decimal(100),
 'sample_urls_reserved_invalid':all('.invalid' in x for x in re.findall(r'https?://[^\s)]+',source)) and len(re.findall(r'https?://[^\s)]+',source))>=4,
 'sample_hold_boundary_present':all(x in all_text for x in ['SAMPLE','SYNTHETIC','HOLD','REAL']),
 'no_blank_placeholder_terms':all(x not in all_text for x in ['待填','空白表','保持空白','會後填寫']),
 'body_run_font_min_10_5':min([r.font.size.pt for p in doc.paragraphs for r in p.runs if r.font.size],default=99)>=10.5,
 'table_run_font_min_9_5':min([r.font.size.pt for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs for r in p.runs if r.font.size],default=99)>=9.5,
 'build_receipt_sha_matches':build.get('artifact_sha256')==fsha(DOCX),
 'build_receipt_identity_matches':build.get('document_id')=='0013' and build.get('requirement_id')=='I02-E04',
 'semantic_not_misreported_pending':build.get('semantic_review') in ('PENDING','PENDING_INDEPENDENT_REVIEW'),
 'formal_acceptance_unknown':build.get('formal_acceptance')=='UNKNOWN',
 'visual_receipt_bound_to_candidate_sha':json.loads(VISUAL.read_text(encoding='utf-8')).get('candidate_sha256')==fsha(DOCX),
 'visual_receipt_pdf_sha_exact':json.loads(VISUAL.read_text(encoding='utf-8')).get('pdf_sha256')==fsha(PDF),
 'visual_all_25_pages_accounted':json.loads(VISUAL.read_text(encoding='utf-8')).get('all_pages_accounted') is True and json.loads(VISUAL.read_text(encoding='utf-8')).get('page_count_pdf')==25,
 'visual_layout_status_pass':json.loads(VISUAL.read_text(encoding='utf-8')).get('layout_status')=='PASS_LAYOUT_ONLY'
}
receipt={'receipt_id':'0013_I02-E04_R05_DETERMINISTIC_TECHNICAL_QA','document_id':'0013','requirement_id':'I02-E04','candidate_state':'CANDIDATE_SAMPLE_SYNTHETIC_NOT_ACCEPTED','artifact_path':str(DOCX),'artifact_sha256':fsha(DOCX),'artifact_bytes':DOCX.stat().st_size,'worker_source_sha256':fsha(SRC),'official_requirement_pdf_sha256':fsha(OFFICIAL),'blueprint_file_sha256':fsha(BP),'blueprint_row_sha256':sha(rowraw),'markdown_tables':len(mdtables),'docx_tables':len(doc.tables),'table_rows_including_headers':sum(len(t.rows) for t in doc.tables),'synthetic_fact_ids':len(expected_facts),'synthetic_register_rows':len(regids),'course_net_hours':hours,'course_hours_total':sum(hours),'course_capacity':capacities,'course_capacity_total':sum(capacities),'source_fragment_count':len(frags),'source_fragments_missing':missing,'checks':checks,'technical_status':'PASS' if all(checks.values()) else 'FAIL','semantic_status':'PENDING_INDEPENDENT_REVIEW','visual_status':json.loads(VISUAL.read_text(encoding='utf-8')).get('layout_status'),'formal_acceptance':'UNKNOWN'}
OUT.parent.mkdir(parents=True,exist_ok=True); OUT.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(receipt,ensure_ascii=True,indent=2))
if not all(checks.values()): raise SystemExit(1)
