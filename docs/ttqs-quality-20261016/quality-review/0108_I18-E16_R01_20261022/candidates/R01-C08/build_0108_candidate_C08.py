from pathlib import Path
import re, sys, json, hashlib, zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=Path(r'E:\TTQS\_r09_opencode_0108_text_only_20261022_R01')
SRC=BASE/'candidate'/'0108_I18-E16_candidate_R01-C08.md'
OUT=BASE/'candidate'/'0108_I18-E16_candidate_R01-C08.docx'
TEXT=SRC.read_text(encoding='utf-8')
NAVY='18344F'; TEAL='176B73'; LIGHT='EAF2F5'; PALE='F5F8FA'; GREY='64748B'; DARK='1E293B'; GOLD='FFF4D6'; RED='8A2C2C'
FONT='Microsoft JhengHei'

def set_run_font(run,size=10.5,bold=None,color=None,italic=None,name=FONT):
    run.font.name=name
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),name)
    run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    if italic is not None: run.italic=italic
    if color: run.font.color.rgb=RGBColor.from_string(color)

def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)

def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)

def cant_split(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:cantSplit'); trPr.append(el)

def set_cell_margins(cell,top=80,start=90,bottom=80,end=90):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); mar=tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar=OxmlElement('w:tcMar'); tcPr.append(mar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        el=mar.find(qn('w:'+m))
        if el is None: el=OxmlElement('w:'+m); mar.append(el)
        el.set(qn('w:w'),str(v)); el.set(qn('w:type'),'dxa')

def add_inline(p,text,size=10.5,color=DARK):
    parts=re.split(r'(\*\*.*?\*\*|`[^`]+`)',text)
    for part in parts:
        if not part: continue
        if part.startswith('**') and part.endswith('**'):
            r=p.add_run(part[2:-2]); set_run_font(r,size,bold=True,color=color)
        elif part.startswith('`') and part.endswith('`'):
            r=p.add_run(part[1:-1]); set_run_font(r,max(size,10.5),color='375B68',name='Consolas')
        else:
            r=p.add_run(part); set_run_font(r,size,color=color)

def parse_row(line):
    x=line.strip().strip('|')
    return [c.strip().replace('<br>','\n') for c in x.split('|')]

def is_separator(line):
    return bool(re.fullmatch(r'\|?\s*:?-+:?\s*(?:\|\s*:?-+:?\s*)+\|?\s*',line.strip()))

def style_cell_text(cell,text,header=False,label=False):
    cell.text=''
    p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0); p.paragraph_format.line_spacing=1.0
    add_inline(p,text,9.5,color=('FFFFFF' if header else (TEAL if label else DARK)))
    if header:
        for r in p.runs:r.bold=True
    if label: shade(cell,LIGHT)
    elif header: shade(cell,NAVY)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)

def add_grid_table(rows,caption=None):
    n=len(rows[0]) if rows else 0
    t=doc.add_table(rows=1,cols=n)
    t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'; t.autofit=True
    for j,val in enumerate(rows[0]): style_cell_text(t.rows[0].cells[j],val,header=True)
    set_repeat_table_header(t.rows[0]); cant_split(t.rows[0])
    for ri,row in enumerate(rows[1:],1):
        cells=t.add_row().cells; cant_split(t.rows[-1])
        for j,val in enumerate(row):
            style_cell_text(cells[j],val)
            if ri%2==0: shade(cells[j],PALE)
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(1); p.paragraph_format.space_before=Pt(0)
    return t

def add_wide_cards(rows):
    headers=rows[0]
    for ri,row in enumerate(rows[1:],1):
        key=row[0] if row else f'項目 {ri}'
        p=doc.add_paragraph(style='Heading 3'); p.paragraph_format.keep_with_next=True; p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(2)
        add_inline(p,key,10.5,color=TEAL)
        t=doc.add_table(rows=0,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'; t.autofit=False
        for j in range(1,len(headers)):
            cells=t.add_row().cells; cant_split(t.rows[-1]); cells[0].width=Inches(1.65); cells[1].width=Inches(5.2)
            style_cell_text(cells[0],headers[j],label=True)
            style_cell_text(cells[1],row[j] if j<len(row) else 'UNKNOWN')
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(1); p.paragraph_format.space_before=Pt(0)

def add_table_block(block):
    rows=[]
    for line in block:
        if is_separator(line): continue
        row=parse_row(line)
        if rows and len(row)!=len(rows[0]):
            raise ValueError(f'Unequal table width {len(rows[0])} vs {len(row)}: {line}')
        rows.append(row)
    if not rows:return
    if len(rows[0])>4: add_wide_cards(rows)
    else: add_grid_table(rows)

def add_callout(text):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.12); p.paragraph_format.right_indent=Inches(.1); p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(7)
    p.paragraph_format.keep_together=True
    ppr=p._p.get_or_add_pPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),GOLD if 'SAMPLE / SYNTHETIC' in text or 'NOT 原始點' in text else LIGHT); ppr.append(shd)
    add_inline(p,text.replace('>','',1).strip(),10.5,color=RED if 'NOT 原始點' in text else DARK)

def add_page_field(p):
    r=p.add_run(); set_run_font(r,8.0,color=GREY)
    begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin')
    instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text=' PAGE '
    separate=OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'),'separate')
    value=OxmlElement('w:t'); value.text='1'
    end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end')
    for e in (begin,instr,separate,value,end):r._r.append(e)

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
sec.top_margin=Inches(.62); sec.bottom_margin=Inches(.62); sec.left_margin=Inches(.68); sec.right_margin=Inches(.68)
sec.header_distance=Inches(.28); sec.footer_distance=Inches(.28)
# Normal and headings
normal=doc.styles['Normal']; normal.font.name=FONT; normal._element.rPr.rFonts.set(qn('w:eastAsia'),FONT); normal.font.size=Pt(10.5); normal.font.color.rgb=RGBColor.from_string(DARK)
normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.12
for nm,size,color in [('Title',24,NAVY),('Subtitle',12,TEAL),('Heading 1',16,NAVY),('Heading 2',12.5,TEAL),('Heading 3',10.8,DARK)]:
 st=doc.styles[nm]; st.font.name=FONT; st._element.rPr.rFonts.set(qn('w:eastAsia'),FONT); st.font.size=Pt(size); st.font.bold=nm!='Subtitle'; st.font.color.rgb=RGBColor.from_string(color); st.paragraph_format.keep_with_next=True
 st.paragraph_format.space_before=Pt(10 if nm.startswith('Heading') else 0); st.paragraph_format.space_after=Pt(4)
# Document metadata
props=doc.core_properties; props.title='薪資變動佐證紀錄'; props.subject='薪資變動佐證紀錄 SAMPLE SYNTHETIC 示範稿'; props.author=''; props.keywords='薪資變動佐證紀錄; SAMPLE; SYNTHETIC'
# Header/footer
h=sec.header.paragraphs[0]; h.alignment=WD_ALIGN_PARAGRAPH.RIGHT; h.paragraph_format.space_after=Pt(0)
add_inline(h,'SAMPLE / SYNTHETIC 示範資料',8,color=GREY)
f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER; f.paragraph_format.space_before=Pt(0)
add_inline(f,'示範資料  ·  第 ',8,color=GREY); add_page_field(f); add_inline(f,' 頁',8,color=GREY)
lines=TEXT.splitlines(); i=0
# Title block is polished independently; keep the source heading text in body otherwise.
title_p=doc.add_paragraph(style='Title'); title_p.paragraph_format.space_after=Pt(2); add_inline(title_p,'薪資變動佐證紀錄',24,color=NAVY)
sub=doc.add_paragraph(style='Subtitle'); sub.paragraph_format.space_after=Pt(3); add_inline(sub,'薪資變動佐證紀錄 SAMPLE SYNTHETIC 示範稿',12,color=TEAL)
status=doc.add_paragraph(); status.paragraph_format.space_after=Pt(10)
add_inline(status,'全部案例與金額均為 SAMPLE / SYNTHETIC，不是協會實績或真實薪資證明。',10.5,color=RED)
# Skip markdown title/subtitle; process content from the first blockquote onward.
while i<len(lines) and (not lines[i].strip() or lines[i].startswith('#')): i+=1
while i<len(lines):
 line=lines[i].rstrip()
 if not line.strip(): i+=1; continue
 if line.startswith('|'):
  block=[]
  while i<len(lines) and lines[i].startswith('|'):
   block.append(lines[i]); i+=1
  add_table_block(block); continue
 if line.startswith('>'):
  q=[]
  while i<len(lines) and lines[i].startswith('>'):
   q.append(lines[i].lstrip('>').strip()); i+=1
  add_callout(' '.join(q)); continue
 if line.startswith('#'):
  level=len(line)-len(line.lstrip('#')); text=line[level:].strip()
  style='Heading 1' if level==2 else ('Heading 2' if level==3 else 'Heading 3')
  p=doc.add_paragraph(style=style); add_inline(p,text,16 if level==2 else (12.5 if level==3 else 10.8),color=NAVY if level==2 else TEAL)
  i+=1; continue
 if re.match(r'^\s*[-*]\s+',line):
  p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.left_indent=Inches(.25); p.paragraph_format.first_line_indent=Inches(-.12); add_inline(p,re.sub(r'^\s*[-*]\s+','',line),10.5)
  i+=1; continue
 m=re.match(r'^\s*(\d+)\.\s+(.*)$',line)
 if m:
  p=doc.add_paragraph(style='List Number'); p.paragraph_format.left_indent=Inches(.28); p.paragraph_format.first_line_indent=Inches(-.16); add_inline(p,m.group(2),10.5)
  i+=1; continue
 p=doc.add_paragraph(); add_inline(p,line,10.5)
 i+=1
OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
# Basic structural verification and exact output digest
with zipfile.ZipFile(OUT) as z:
 bad=z.testzip()
 if bad: raise RuntimeError('ZIP integrity failure '+bad)
check=Document(OUT)
alltext='\n'.join([p.text for p in check.paragraphs]+[c.text for t in check.tables for row in t.rows for c in row.cells])
required=['薪資變動佐證紀錄','SYN-P01','SYN-P08','205,500','212,500','3.406%','SAMPLE','SYNTHETIC','REAL','EXCLUDE_WITH_REASON','REJECTED_RECORD']
missing=[x for x in required if x not in alltext]
if missing: raise RuntimeError('Missing DOCX text '+repr(missing))
b=OUT.read_bytes()
print(json.dumps({'docx':str(OUT),'bytes':len(b),'sha256':hashlib.sha256(b).hexdigest(),'paragraphs':len(check.paragraphs),'tables':len(check.tables),'sections':len(check.sections),'docx_text_checks_missing':missing,'zip_integrity':'PASS'},ensure_ascii=False))