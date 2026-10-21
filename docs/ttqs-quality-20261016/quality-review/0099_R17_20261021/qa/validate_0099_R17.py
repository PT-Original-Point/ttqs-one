from pathlib import Path
from docx import Document
from zipfile import ZipFile
import csv, hashlib, json, re, sys
root=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0099_R01_BLUEPRINT_R01')
docx=root/'candidate'/'0099_R17.docx'; md=root/'candidate'/'0099_R14_manuscript.md'
reg=root/'evidence'/'SYNTHETIC_REPLACEMENT_REGISTER_0099_R14.csv'; fixture=root/'evidence'/'fixture_0099_R01.csv'
bp=root/'inputs'/'UNIQUE_BLUEPRINT_0099_R02.md'; pdf=root/'render'/'0099_R17_word.pdf'
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
checks={}
for p in (docx,md,reg,fixture,bp,pdf):
    if not p.is_file(): raise SystemExit(f'MISSING:{p}')
with ZipFile(docx) as z:
    bad=z.testzip(); checks['docx_zip_integrity']=bad is None
    if bad: raise SystemExit('DOCX_ZIP_BAD:'+bad)
d=Document(docx)
body='\n'.join([p.text for p in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells])
with fixture.open(encoding='utf-8-sig',newline='') as f: data=list(csv.DictReader(f))
ids=[r['participant_code'] for r in data]
assert ids==[f'A{i:02d}' for i in range(1,9)],ids
score_fields=[f'{stage}_{dim}' for stage in ('pre','post') for dim in 'ABCDE']
raw=[int(r[k]) for r in data for k in score_fields]
assert len(raw)==80 and all(0<=x<=4 for x in raw)
for r in data:
    assert sum(int(r[f'pre_{x}']) for x in 'ABCDE')==int(r['pre_total'])
    assert sum(int(r[f'post_{x}']) for x in 'ABCDE')==int(r['post_total'])
    assert int(r['post_total'])-int(r['pre_total'])==int(r['gain'])
pre=sum(int(r['pre_total']) for r in data); post=sum(int(r['post_total']) for r in data)
gain=post-pre; passn=sum(r['sample_threshold_result']=='PASS_SAMPLE' for r in data)
assert (pre,post,gain,passn)==(75,117,42,6),(pre,post,gain,passn)
checks['fixture_80_cells']=True; checks['fixture_row_totals_8_of_8']=True; checks['fixture_aggregate_75_117_42']=True; checks['sample_threshold_6_of_8']=True
assert len(d.tables)>=10
scoretab=d.tables[8]
assert len(scoretab.rows)==10 and len(scoretab.columns)==8
for i,r in enumerate(data,1):
    cells=scoretab.rows[i].cells
    doc_pre=[int(x.strip()) for x in cells[1].text.split('/')]
    doc_post=[int(x.strip()) for x in cells[3].text.split('/')]
    assert doc_pre==[int(r[f'pre_{x}']) for x in 'ABCDE'],(i,doc_pre)
    assert doc_post==[int(r[f'post_{x}']) for x in 'ABCDE'],(i,doc_post)
    assert int(cells[2].text.strip())==int(r['pre_total'])
    assert int(cells[4].text.strip())==int(r['post_total'])
    assert cells[7].text.strip()==('達門檻' if r['sample_threshold_result']=='PASS_SAMPLE' else '未達：總分')
assert int(scoretab.rows[9].cells[2].text.strip())==75
assert int(scoretab.rows[9].cells[4].text.strip())==117
assert scoretab.rows[9].cells[7].text.strip()=='達標率 75%'
checks['docx_80_cells_match_fixture']=True; checks['docx_individual_results_match_fixture']=True
# Rounding and normalized aggregate statements present in the rendered document.
for term in ('80 格','75','117','42','5.250','6 人','26.250','56.0%','目標客戶','UNKNOWN/HOLD','SAMPLE / SYNTHETIC'):
    assert term in body, 'MISSING_BODY_TERM:'+term
assert '160 個原始評分格' not in body
assert not re.search(r'\b[0-9a-fA-F]{64}\b',body)
assert not re.search(r'(?i)(?:[A-Z]:\\|E:/TTQS|C:/Users/|github\.com/)',body)
checks['body_no_sha_or_internal_path']=True
# No stale engineer-only R11 marker in the actual Word text.
assert 'R11' not in body
checks['no_stale_R11_label']=True
# Unique synthetic fact register integrity.
with reg.open(encoding='utf-8-sig',newline='') as f: rows=list(csv.DictReader(f))
assert len(rows)==383
key='fact_id' if 'fact_id' in rows[0] else None
assert key is not None
keys=[r[key] for r in rows]
assert len(set(keys))==383 and all(keys)
checks['register_383_unique_fact_ids']=True
# R02 source linkage.
bpt=bp.read_text(encoding='utf-8-sig')
assert 'I18-E07' in bpt and '80' in bpt and 'R02' in bpt
checks['blueprint_R02_binding']=True
# DOCX title style must not carry a bottom border.
title=d.styles['Title']._element.get_or_add_pPr()
from docx.oxml.ns import qn
assert title.find(qn('w:pBdr')) is None
checks['cover_title_rule_removed']=True
receipt={
 'schema':'ttqs_one_0099_r17_technical_visual_receipt.v1','doc_id':'0099','requirement_id':'I18-E07',
 'printed_official_checkbox_serial':'UNKNOWN_NOT_PRINTED_IN_READ_SOURCE','project_157_index':108,
 'blueprint_revision':'R02','blueprint_sha256':sha(bp),'fixture_sha256':sha(fixture),
 'candidate_revision':'R17','candidate_docx_sha256':sha(docx),'candidate_docx_bytes':docx.stat().st_size,
 'manuscript_sha256':sha(md),'register_sha256':sha(reg),'register_rows':len(rows),'unique_fact_ids':len(set(keys)),
 'word_export_pdf_sha256':sha(pdf),'word_page_count':37,
 'fixture_raw_score_cells':len(raw),'fixture_pre_total':pre,'fixture_post_total':post,'fixture_gain_total':gain,
 'fixture_mean_gain':gain/8,'sample_threshold_pass_count':passn,'sample_threshold_n':8,
 'technical_checks':checks,
 'visual_scope':{'status':'PASS_SCOPED','all_page_overview':'37/37 contact-sheet pages visually inspected',
   'pixel_comparison_base':'R15 Word PDF rendered at same Poppler version and 85dpi','changed_pages':[1,22,23],
   'full_size_inspected_pages':[1,22,23,24],'unchanged_pages_pixel_identical':34,
   'page_23_table_10_2_caption_with_table':'PASS','page_24_repeating_table_header':'PASS',
   'rationale':'R17 vs R15 pixel comparison found only pages 1,22,23 changed; all 37 pages had overview inspection; changed pages and continuation page inspected at full size.'},
 'semantic_review':'PENDING_EXACT_SHA_INDEPENDENT_REVIEW','artifact_state':'CANDIDATE_NOT_FORMALLY_ACCEPTED',
 'real_evidence':'NOT_PROVIDED_SAMPLE_ONLY','human_operability_test':'AI proxy only; real novice test NOT_TESTED',
 'result':'PASS_TECHNICAL_AND_SCOPED_VISUAL'
}
out=root/'evidence'/'technical_visual_receipt_0099_R17.json'
out.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'receipt':str(out),'receipt_sha256':sha(out),'candidate_docx_sha256':sha(docx),'pdf_sha256':sha(pdf),'checks':checks,'result':receipt['result']},ensure_ascii=False))

