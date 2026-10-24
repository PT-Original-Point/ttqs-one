from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image,ImageDraw,ImageFont
import re,csv,hashlib,json,math,statistics
R=Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0099_R01_BLUEPRINT_R01")
mdp=R/'candidate'/'0099_R14_manuscript.md'
out=R/'candidate'/'0099_R17.docx'
fixture=R/'evidence'/'fixture_0099_R01.csv'
imgdir=R/'render'/'figures_R17';imgdir.mkdir(parents=True,exist_ok=True)
text=mdp.read_text(encoding='utf-8'); mdsha=hashlib.sha256(mdp.read_bytes()).hexdigest()
blueprint_sha=hashlib.sha256((R/'inputs'/'UNIQUE_BLUEPRINT_0099_R02.md').read_bytes()).hexdigest()
with fixture.open(encoding='utf-8-sig',newline='') as f: data=list(csv.DictReader(f))
assert len(data)==8
pre={d:sum(int(r[f'pre_{d}']) for r in data)/8 for d in 'ABCDE'}
post={d:sum(int(r[f'post_{d}']) for r in data)/8 for d in 'ABCDE'}
# Synthetic charts, built from the exact fixture used by the manuscript.
fontpath=r'C:\Windows\Fonts\msjh.ttc'
def font(sz):return ImageFont.truetype(fontpath,sz,index=0)
NAVY=(31,58,95); TEAL=(35,142,145); ORANGE=(199,106,68); GRID=(218,225,232); INK=(40,52,67); MUTED=(101,116,132); BG=(255,255,255)
def save_dim_chart(path):
 W,H=1500,850; im=Image.new('RGB',(W,H),BG); d=ImageDraw.Draw(im)
 d.text((95,45),'n=8 合成配對；0–4 分；平均差異僅為描述性結果，不單獨證明課程因果。',font=font(21),fill=MUTED)
 
 left,top,right,bottom=145,165,1330,690
 for tick in range(5):
  y=bottom-int(tick/4*(bottom-top));d.line((left,y,right,y),fill=GRID,width=2);d.text((95,y-14),str(tick),font=font(19),fill=MUTED)
 labels={'A':'來源事實','B':'紀錄追溯','C':'目標對應','D':'資料最小化','E':'後續追蹤'}
 slot=(right-left)/5
 for j,key in enumerate('ABCDE'):
  cx=left+slot*(j+.5);bw=50
  for val,offset,color in [(pre[key],-30,NAVY),(post[key],30,TEAL)]:
   y=bottom-int(val/4*(bottom-top));x=cx+offset
   d.rounded_rectangle((x-bw/2,y,x+bw/2,bottom),radius=8,fill=color)
   d.text((x-28,y-34),f'{val:.3f}'.rstrip('0').rstrip('.'),font=font(18),fill=INK)
  d.text((cx-32,bottom+20),key,font=font(22),fill=NAVY,stroke_width=1)
  d.text((cx-55,bottom+50),labels[key],font=font(18),fill=INK)
 d.rounded_rectangle((970,115,1105,147),radius=6,fill=NAVY);d.text((1115,114),'前測',font=font(18),fill=INK)
 d.rounded_rectangle((1200,115,1335,147),radius=6,fill=TEAL);d.text((1345,114),'後測',font=font(18),fill=INK)
 d.text((95,790),'資料來源：fixture_0099_R01.csv；圖表不含真實學員資料。',font=font(18),fill=MUTED)
 im.save(path,dpi=(180,180))
def save_total_chart(path):
 W,H=1500,850;im=Image.new('RGB',(W,H),BG);d=ImageDraw.Draw(im)
 d.text((95,45),'每人總分 0–20；虛線為 SAMPLE 示範門檻 14，不得用於真人決策。',font=font(21),fill=MUTED)
 
 left,top,right,bottom=145,165,1330,690
 for tick in range(0,21,2):
  y=bottom-int(tick/20*(bottom-top));d.line((left,y,right,y),fill=GRID,width=2);d.text((90,y-12),str(tick),font=font(18),fill=MUTED)
 threshold=14;y=bottom-int(threshold/20*(bottom-top));d.line((left,y,right,y),fill=ORANGE,width=3)
 for x in range(left,right,25):d.line((x,y,min(x+12,right),y),fill=BG,width=3)
 d.text((right-170,y-28),'SAMPLE門檻 14',font=font(18),fill=ORANGE)
 slot=(right-left)/8
 for j,r in enumerate(data):
  cx=left+slot*(j+.5)
  for val,off,color in [(int(r['pre_total']),-15,NAVY),(int(r['post_total']),15,TEAL)]:
   yy=bottom-int(val/20*(bottom-top));xx=cx+off
   d.rounded_rectangle((xx-12,yy-12,xx+12,yy+12),radius=10,fill=color)
   d.text((xx-19,yy-40),str(val),font=font(17),fill=INK)
  d.text((cx-25,bottom+18),r['participant_code'],font=font(18),fill=INK)
 d.line((1060,115,1095,115),fill=NAVY,width=5);d.text((1105,102),'前測',font=font(18),fill=INK)
 d.line((1200,115,1235,115),fill=TEAL,width=5);d.text((1245,102),'後測',font=font(18),fill=INK)
 d.text((95,790),'資料來源：fixture_0099_R01.csv；n=8 合成配對。不得作為真實成效或因果證明。',font=font(18),fill=MUTED)
 im.save(path,dpi=(180,180))
chart1=imgdir/'figure_10_1_dimension_means.png';chart2=imgdir/'figure_10_2_learner_totals.png'
save_dim_chart(chart1);save_total_chart(chart2)
# --- formatting helpers ---
NAVY_HEX='1F3A5F'; TEAL_HEX='238E91'; PALE='EEF3F7'; PALE2='F5F8FA'; INK_HEX='283443'; WHITE='FFFFFF'; ACCENT='C76A44'
def set_cell_shading(cell,color):
 tcPr=cell._tc.get_or_add_tcPr();shd=tcPr.find(qn('w:shd'))
 if shd is None:shd=OxmlElement('w:shd');tcPr.append(shd)
 shd.set(qn('w:fill'),color)
def set_cell_margins(cell,top=85,start=95,bottom=85,end=95):
 tc=cell._tc;tcPr=tc.get_or_add_tcPr();mar=tcPr.first_child_found_in('w:tcMar')
 if mar is None:mar=OxmlElement('w:tcMar');tcPr.append(mar)
 for tag,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
  e=mar.find(qn('w:'+tag))
  if e is None:e=OxmlElement('w:'+tag);mar.append(e)
  e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa')
def no_split(row):
 trPr=row._tr.get_or_add_trPr();e=OxmlElement('w:cantSplit');trPr.append(e)
def repeat_header(row):
 trPr=row._tr.get_or_add_trPr();e=OxmlElement('w:tblHeader');e.set(qn('w:val'),'true');trPr.append(e)
def borders(table,color='CBD5DF',sz='5'):
 tblPr=table._tbl.tblPr;b=tblPr.first_child_found_in('w:tblBorders')
 if b is None:b=OxmlElement('w:tblBorders');tblPr.append(b)
 for edge in ('top','left','bottom','right','insideH','insideV'):
  e=b.find(qn('w:'+edge))
  if e is None:e=OxmlElement('w:'+edge);b.append(e)
  e.set(qn('w:val'),'single');e.set(qn('w:sz'),sz);e.set(qn('w:space'),'0');e.set(qn('w:color'),color)
def set_fonts(run,name='Microsoft JhengHei',size=10.5,bold=None,color=INK_HEX,italic=None):
 run.font.name=name;run.font.size=Pt(size);run.font.color.rgb=RGBColor.from_string(color)
 if bold is not None:run.bold=bold
 if italic is not None:run.italic=italic
 rPr=run._element.get_or_add_rPr();rf=rPr.rFonts
 if rf is None:rf=OxmlElement('w:rFonts');rPr.append(rf)
 for k in ('ascii','hAnsi','eastAsia','cs'):rf.set(qn('w:'+k),name)
def add_inline(p,s,table=False):
 s=s.replace('<br>','\n').replace('<br/>','\n')
 s=s.replace(r'\mu_{\text{pre}}','μ前測').replace(r'\mu_{\text{post}}','μ後測').replace(r'\Delta \mu','Δμ').replace(r'\text{Norm}_{\text{pre}}','前測正規化分數').replace(r'\text{Norm}_{\text{post}}','後測正規化分數').replace(r'\text{Relative Growth}','相對增幅')
 s=re.sub(r'\\text\{([^{}]*)\}',r'\1',s)
 s=re.sub(r'\\frac\{([^{}]+)\}\{([^{}]+)\}',r'\1 ÷ \2',s)
 s=s.replace(r'\approx','≈').replace(r'\times','×').replace(r'\%','%')
 s=re.sub(r'\\(?:left|right)','',s).replace('$','')
 pattern=re.compile(r'(\*\*.+?\*\*|(?<!\*)\*[^*]+\*(?!\*)|`[^`]+`)')
 pos=0
 for m in pattern.finditer(s):
  if m.start()>pos:p.add_run(s[pos:m.start()])
  tok=m.group(0)
  if tok.startswith('**'):r=p.add_run(tok[2:-2]);r.bold=True
  elif tok.startswith('`'):r=p.add_run(tok[1:-1]);r.font.name='Consolas';r.font.size=Pt(9)
  else:r=p.add_run(tok[1:-1]);r.italic=True
  pos=m.end()
 if pos<len(s):p.add_run(s[pos:])
 for r in p.runs:
  if r.font.size is None:set_fonts(r,size=8.2 if table else 10.5)
  else:
   if r.font.name=='Calibri':set_fonts(r,size=8.2 if table else 10.5)
 for part in p.runs:
  if '\n' in part.text:
   chunks=part.text.split('\n');part.text=chunks[0]
   for ch in chunks[1:]:p.add_run().add_break();p.add_run(ch)
def shade_paragraph(p,fill='F2F6FA',border='CBD5DF'):
 pPr=p._p.get_or_add_pPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),fill);pPr.append(shd)
 ind=OxmlElement('w:ind');ind.set(qn('w:left'),'250');ind.set(qn('w:right'),'180');pPr.append(ind)
def add_page_field(p,field):
 r=p.add_run();begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text=f' {field} ';sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');txt=OxmlElement('w:t');txt.text='1';end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');r._r.extend([begin,instr,sep,txt,end]);set_fonts(r,size=8,color='657484')
def set_section(sec,landscape=False):
 sec.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
 if landscape:sec.page_width=Inches(11.69);sec.page_height=Inches(8.27);sec.left_margin=Inches(.48);sec.right_margin=Inches(.48);sec.top_margin=Inches(.50);sec.bottom_margin=Inches(.43)
 else:sec.page_width=Inches(8.27);sec.page_height=Inches(11.69);sec.left_margin=Inches(.68);sec.right_margin=Inches(.68);sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.7)
 sec.header_distance=Inches(.32);sec.footer_distance=Inches(.32)
def add_new_section(doc,landscape):
 sec=doc.add_section(WD_SECTION.NEW_PAGE);set_section(sec,landscape)
 sec.header.is_linked_to_previous=True;sec.footer.is_linked_to_previous=True
 return sec
# Markdown -> structured blocks.
raw=text.splitlines(); blocks=[];i=0;in_comment=False;in_code=False;code=[]
while i<len(raw):
 s=raw[i].strip()
 if s.startswith('<!--') and not s.startswith('<!-- DOCX_FIGURE:10.1 -->') and not s.startswith('<!-- DOCX_FIGURE:10.2 -->'):in_comment=True
 if in_comment:
  if '-->' in s:in_comment=False
  i+=1;continue
 if not s:i+=1;continue
 if s.startswith('```'):
  if in_code:blocks.append(('code','\n'.join(code)));code=[];in_code=False
  else:in_code=True
  i+=1;continue
 if in_code:code.append(raw[i]);i+=1;continue
 if s.startswith('|') and s.endswith('|'):
  group=[]
  while i<len(raw) and raw[i].strip().startswith('|') and raw[i].strip().endswith('|'):
   row=raw[i].strip();cells=[x.strip() for x in row[1:-1].split('|')]
   if not all(re.fullmatch(r':?-{2,}:?',x.replace(' ','')) for x in cells):group.append(cells)
   i+=1
  blocks.append(('table',group));continue
 if s in ('---','___','***'):
  blocks.append(('divider',''));i+=1;continue
 if s.startswith('<!-- DOCX_FIGURE:10.1 -->'):blocks.append(('figure','10.1'));i+=1;continue
 if s.startswith('<!-- DOCX_FIGURE:10.2 -->'):blocks.append(('figure','10.2'));i+=1;continue
 m=re.match(r'^(#{1,4})\s+(.*)$',s)
 if m:blocks.append(('heading',(len(m.group(1)),m.group(2))));i+=1;continue
 if s.startswith('>'):
  q=[]
  while i<len(raw) and raw[i].strip().startswith('>'):
   q.append(raw[i].strip()[1:].lstrip());i+=1
  blocks.append(('quote',q));continue
 if re.match(r'^(\d+\.|[A-Za-z][.)])\s+',s):blocks.append(('number',re.sub(r'^(\d+\.|[A-Za-z][.)])\s+','',s)));i+=1;continue
 if s.startswith(('* ','- ')):blocks.append(('bullet',s[2:]));i+=1;continue
 if s.startswith('  * ') or s.startswith('   * '):blocks.append(('subbullet',s.lstrip()[2:]));i+=1;continue
 blocks.append(('para',s));i+=1
# sections/styles
doc=Document();set_section(doc.sections[0],False)
styles=doc.styles
normal=styles['Normal'];normal.font.name='Microsoft JhengHei';normal.font.size=Pt(10.5);normal.font.color.rgb=RGBColor.from_string(INK_HEX)
normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1.14
for sty,size,color in [('Title',30,NAVY_HEX),('Heading 1',17,NAVY_HEX),('Heading 2',13,TEAL_HEX),('Heading 3',11.5,INK_HEX),('Heading 4',10.5,INK_HEX)]:
 st=styles[sty];st.font.name='Microsoft JhengHei';st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=RGBColor.from_string(color)
 st.paragraph_format.space_before=Pt(10 if sty!='Title' else 0);st.paragraph_format.space_after=Pt(5);st.paragraph_format.keep_with_next=True
 for attr in ['ascii','hAnsi','eastAsia','cs']:
  rf=st._element.rPr.rFonts if st._element.rPr is not None else None
  if rf is not None:rf.set(qn('w:'+attr),'Microsoft JhengHei')
# Remove the default Word Title style border; the cover must not draw a rule below its title.
title_ppr=styles['Title']._element.get_or_add_pPr()
title_border=title_ppr.find(qn('w:pBdr'))
if title_border is not None:title_ppr.remove(title_border)
# page header/footer
for sec in doc.sections:
 hp=sec.header.paragraphs[0];hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 rr=hp.add_run('TTQS ONE  |  0099  |  I18-E07');set_fonts(rr,size=8,color='657484')
 fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
 rr=fp.add_run('CANDIDATE  •  SAMPLE / SYNTHETIC  •  不得作為真實訓練成效或核准紀錄');set_fonts(rr,size=8,color='657484')
 fp.add_run('    |    第 ');add_page_field(fp,'PAGE');fp.add_run(' 頁')
# cover
p=doc.add_paragraph();p.paragraph_format.space_before=Pt(75);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('TTQS ONE  ·  文件 0099');set_fonts(r,size=14,bold=True,color=TEAL_HEX)
p=doc.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(25)
add_inline(p,'訓練技能作品表現\n與能力變化分析報告')
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(8)
add_inline(p,'對應 I18-E07｜提出目標客戶及學員對訓練達能力提升之說明（學習技能面）')
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(28)
r=p.add_run('候選版本 R17');set_fonts(r,size=13,bold=True,color=ACCENT)
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('全篇 SAMPLE / SYNTHETIC｜非原始點協會已發生之訓練或績效證據');set_fonts(r,size=11,bold=True,color=ACCENT)
# cover control panel
ct=doc.add_table(rows=0,cols=2);ct.alignment=WD_TABLE_ALIGNMENT.CENTER;ct.autofit=True
for k,v in [
('官方需求','I18-E07；官方原文取自佐證表第 7 頁'),
('142 文件對照','0099；seed filename/title 欄截斷，展示標題依需求目的編寫'),
('唯一 Blueprint','R02；來源及版本留存於外部收據。'),
('原始分數格','80 筆（8×2×5）；每一階段滿分總和 160'),
('實際資料狀態','沒有 REAL 參訓/評量資料；官方印刷流水號 UNKNOWN'),
('證據狀態','學員端為 SAMPLE/SYNTHETIC；目標客戶端仍 UNKNOWN/HOLD。')]:
 row=ct.add_row();row.cells[0].text=k;row.cells[1].text=v
 for j,c in enumerate(row.cells):
  set_cell_margins(c);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  if j==0:set_cell_shading(c,NAVY_HEX)
  else:set_cell_shading(c,PALE2)
  for pp in c.paragraphs:
   for rr in pp.runs:set_fonts(rr,size=9.5,bold=(j==0),color=WHITE if j==0 else INK_HEX)
  no_split(row)
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(24)
rr=p.add_run('本候選文件須完成版面檢查及獨立 exact-SHA 語意審查；靜態內容或排版檢查不等於正式接受。');set_fonts(rr,size=9.5,color='657484')
doc.add_page_break()
# contents navigation from headings
p=doc.add_paragraph('章節導覽',style='Heading 1')
for typ,val in blocks:
 if typ=='heading' and val[0]==2 and (re.match(r'第[一二三四五六七八九十百]+章',val[1]) or '報告基本資訊' in val[1]):
  pp=doc.add_paragraph(style='List Bullet');add_inline(pp,val[1]);pp.paragraph_format.space_after=Pt(2)
doc.add_page_break()
# Markdown table helpers
current_landscape=False
def is_special_table(b):
 if not b or b[0]!='table' or not b[1]:return False
 h=' '.join(b[1][0])
 return 'Fact Key ID' in h or '事實鍵編號' in h or '核心替換項目' in h or '管制序號' in h

def is_wide_table(b):
 return bool(b and b[0]=='table' and b[1] and len(b[1][0])>=7 and not is_special_table(b))

wide_heading_indices=set()
wide_table_indices=[]
for idx,b in enumerate(blocks):
 if is_wide_table(b):
  wide_table_indices.append(idx)
  if idx>0 and blocks[idx-1][0]=='heading':wide_heading_indices.add(idx-1)
  elif idx>0 and blocks[idx-1][0]=='para' and blocks[idx-1][1].lstrip().startswith('**\u8868 10.2'):wide_heading_indices.add(idx-1)
chapter14_idx=next((i for i,b in enumerate(blocks) if b[0]=='heading' and b[1][1].startswith('\u7b2c\u5341\u56db\u7ae0')),None)
if chapter14_idx is not None and any(j>chapter14_idx for j in wide_table_indices):wide_heading_indices.add(chapter14_idx)
chapter4_idx=next((i for i,b in enumerate(blocks) if b[0]=='heading' and b[1][1].startswith('\u7b2c\u56db\u7ae0')),None)
if chapter4_idx is not None:
 next_chapter_idx=next((i for i,b in enumerate(blocks) if i>chapter4_idx and b[0]=='heading' and b[1][1].startswith('\u7b2c')),len(blocks))
 if any(chapter4_idx<j<next_chapter_idx for j in wide_table_indices):wide_heading_indices.add(chapter4_idx)
last_wide_table_idx=max(wide_table_indices) if wide_table_indices else -1
def new_table(rows,small=False,landscape=False,compact=False):
 n=max(len(r) for r in rows);tab=doc.add_table(rows=0,cols=n);tab.alignment=WD_TABLE_ALIGNMENT.CENTER;tab.autofit=True
 try:tab.style='Table Grid'
 except:pass
 borders(tab)
 for ri,vals in enumerate(rows):
  row=tab.add_row();no_split(row)
  if ri==0:repeat_header(row)
  for ci,c in enumerate(row.cells):
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP;set_cell_margins(c,top=38 if compact else 75,start=55 if compact else 75,bottom=38 if compact else 75,end=55 if compact else 75)
   val=vals[ci] if ci<len(vals) else ''
   c.text='';p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=.94 if compact else 1.0
   add_inline(p,val,table=True)
   for rr in p.runs:
    if rr.font.size is None or rr.font.size.pt < 9.5: rr.font.size=Pt(9.5)
    if rr.font.name=='Calibri': rr.font.name='Microsoft JhengHei'
   set_cell_shading(c,NAVY_HEX if ri==0 else (WHITE if ri%2 else PALE2))
   if ri==0:
    for p0 in c.paragraphs:
     for rr in p0.runs:set_fonts(rr,size=9.5,bold=True,color=WHITE)
 return tab
def add_fact_cards(group):
 if len(group)<2:return
 headers=[re.sub(r'<br>',' ',x).replace('**','').replace('`','') for x in group[0]]
 for row in group[1:]:
  vals=row+['']*(len(headers)-len(row))
  key=re.sub(r'\*\*|`','',vals[0]).strip()
  variable=re.sub(r'<br>',' ',vals[1]).replace('**','').replace('`','').strip()
  hp=doc.add_paragraph(style='Heading 3');hp.paragraph_format.keep_with_next=True
  hp.paragraph_format.page_break_before=key in ('合成學員代碼','課前作品','量規版本','配對評分矩陣')
  add_inline(hp,key+'\uFF5C'+variable)
  fields=[]
  for label,j in [('本文位置',2),('本候選採用值（SAMPLE）',3),('REAL替換來源',4),('取得與核驗方式',5),('角色/責任',6),('品質放行條件',7),('目前狀態',8)]:
   if j<len(vals):fields.append([label,vals[j]])
  new_table([['欄位','內容']]+fields,small=True,landscape=False)

def add_action_cards(group):
 if len(group)<2:return
 for row in group[1:]:
  vals=row+['']*(10-len(row))
  key=re.sub(r'\*\*|`','',vals[0]).strip()
  target=re.sub(r'\*\*|`','',vals[1]).strip()
  hp=doc.add_paragraph(style='Heading 3');hp.paragraph_format.keep_with_next=True
  if key!='TRK-SYN-01':hp.paragraph_format.page_break_before=True
  add_inline(hp,key+'\uFF5C對象 '+target)
  labels=['技能缺口（SAMPLE）','示範補強方案（未執行）','示範功能角色','計畫期限（SAMPLE）','預期練習產出（SYN）','示範複評方式','HOLD/升級條件','目前狀態']
  fields=[[labels[j],vals[j+2]] for j in range(8)]
  new_table([['欄位','內容']]+fields,small=True,landscape=False)
# Keep a table caption and its heading/intro chain with the following table.
caption_keep_indices=set()
for idx,(kind,value) in enumerate(blocks):
 if kind=='para' and value.lstrip().startswith(('**\u8868 ','**Table ','表 ','Table ')) and idx+1<len(blocks) and blocks[idx+1][0]=='table':
  caption_keep_indices.add(idx)
  if idx>=2 and blocks[idx-2][0]=='heading' and blocks[idx-1][0]=='para':caption_keep_indices.add(idx-1)
# render blocks, preserving wide table sections
for bi,(typ,val) in enumerate(blocks):
 if bi in wide_heading_indices and not current_landscape:
  add_new_section(doc,True);current_landscape=True
 if typ=='table' and is_wide_table((typ,val)) and not current_landscape:
  add_new_section(doc,True);current_landscape=True
 if typ=='heading':
  lev,title=val
  p=doc.add_paragraph(style=f'Heading {min(lev-1,4)}')
  add_inline(p,title)
 elif typ=='para':
  p=doc.add_paragraph()
  if current_landscape:p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.0
  add_inline(p,val)
  if bi in caption_keep_indices:p.paragraph_format.keep_with_next=True
 elif typ in ('bullet','subbullet','number'):
  p=doc.add_paragraph()
  if current_landscape:p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.0
  p.paragraph_format.left_indent=Inches(.35 if typ!='subbullet' else .62)
  p.paragraph_format.first_line_indent=Inches(-.18)
  prefix='• ' if typ in ('bullet','subbullet') else '– '
  r=p.add_run(prefix);set_fonts(r,size=10.5,bold=True,color=TEAL_HEX if typ!='subbullet' else MUTED if 'MUTED' in globals() else '657484')
  add_inline(p,val)
 elif typ=='quote':
  for qline in val:
   qline=qline.strip()
   if qline.startswith('* '):qline=qline[2:].strip()
   if not qline or qline=='*':continue
   p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.25);p.paragraph_format.right_indent=Inches(.15);p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_together=True;shade_paragraph(p)
   add_inline(p,qline)
 elif typ=='divider':
  p=doc.add_paragraph();p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(5)
  pPr=p._p.get_or_add_pPr();pbdr=OxmlElement('w:pBdr');bt=OxmlElement('w:bottom');bt.set(qn('w:val'),'single');bt.set(qn('w:sz'),'5');bt.set(qn('w:color'),'CBD5DF');pbdr.append(bt);pPr.append(pbdr)
 elif typ=='table':
  if val and ('Fact Key ID' in val[0][0] or '事實鍵編號' in val[0][0] or '核心替換項目' in val[0][0]):
   add_fact_cards(val)
  elif val and '管制序號' in val[0][0]:
   add_action_cards(val)
  else:
   compact_table=bool(val and val[0] and (val[0][0].strip()=='單元序號' or val[0][0].startswith('文件編號')));new_table(val,small=len(val[0])>=6 if val else False,landscape=current_landscape,compact=compact_table)
  next_wide=bi+1<len(blocks) and is_wide_table(blocks[bi+1])
  if current_landscape and not next_wide and bi!=last_wide_table_idx:
   add_new_section(doc,False);current_landscape=False
 elif typ=='figure':
  pic=chart1 if val=='10.1' else chart2
  if doc.paragraphs:doc.paragraphs[-1].paragraph_format.keep_with_next=True
  p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(2);p.paragraph_format.keep_together=True
  shape=p.add_run().add_picture(str(pic),width=Inches(6.65))
  shape._inline.docPr.set('descr','SAMPLE/SYNTHETIC chart; see caption for scope and values.')
 elif typ=='code':
  p=doc.add_paragraph();shade_paragraph(p,'F4F6F8')
  add_inline(p,val)
# document properties
doc.core_properties.title='0099 訓練技能作品表現與能力變化分析報告（候選 R17）'
doc.core_properties.subject='I18-E07；SAMPLE/SYNTHETIC; exact Blueprint R02'
doc.core_properties.author='TTQS ONE document production workflow'
doc.core_properties.keywords='0099,I18-E07,CANDIDATE,SAMPLE,SYNTHETIC'
for _para in doc.paragraphs:
 if "本候選文件須完成版面檢查及獨立 exact-SHA 語意審查" in _para.text: continue
 for _run in _para.runs:
  if _run.text.strip() and (_run.font.size is None or _run.font.size.pt < 10.5): _run.font.size=Pt(10.5)
for _tab in doc.tables:
 for _row in _tab.rows:
  for _cell in _row.cells:
   for _para in _cell.paragraphs:
    for _run in _para.runs:
     if _run.text.strip() and (_run.font.size is None or _run.font.size.pt < 9.5): _run.font.size=Pt(9.5)
doc.save(out)
print(json.dumps({'docx':str(out),'docx_sha256':hashlib.sha256(out.read_bytes()).hexdigest(),'docx_bytes':out.stat().st_size,'markdown_sha256':mdsha,'blueprint_sha256':blueprint_sha,'block_count':len(blocks),'figures':[str(chart1),str(chart2)],'table_count':sum(1 for b in blocks if b[0]=='table')},ensure_ascii=False))

