#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Deterministic builder for DOCX 0056 R03.

Deliverable : 0056_R03.docx  (human-facing Traditional-Chinese admin package)
Scope       : 設備安檢與場地投保證明：第三方證據取得、適用性判定、檢查紀錄與索引
Sources     : inputs/0056_SOURCE_PACKET_R01.md  +  0056_UNIQUE_BLUEPRINT_R01.md
Outputs     : <root>/0056_R03.docx , build/BUILD_NOTES_R03.md

R03 changes (layout/pagination only; content, sample facts, decision logic, official
mapping, schemas, formula and evidence boundaries are UNCHANGED from R02):
  a. Status-dictionary table compacted (state-code column widened to hold each code on
     one line; explicit tight cell margins) so the dictionary no longer leaves a
     two-row tail page.
  b. T5 paginated as one block (keep-with-next across its rows) and reduced to a single
     empty operational row, so no single-record tail page remains.
  c. U+200B is removed entirely. Underscore-boundary wrapping is now achieved with an
     explicit display line break (<w:br/>) inserted ONLY at '_' boundaries and only when
     a code token is wider than its column; the token text itself is byte-exact.
  Retained: <w:cantSplit/> on all rows, header-row keep_with_next, repeating header rows.

Design notes
------------
* The human body keeps the verbatim official requirement and the page-4 locator.
  Internal repo paths / git hashes / queue or control identifiers are NOT written
  into the DOCX (they live only in BUILD_NOTES_R03.md, i.e. the sidecar).
* The document mixes A4 PORTRAIT (narrative) with A4 LANDSCAPE (wide form annex
  T1-T5) so that 14-16 column tables stay readable and never clip.
* Every synthetic value used in the worked sample is SAMPLE/SYNTHETIC and is
  registered in table T9 with an exact locator. UNKNOWN / PROPOSED /
  PENDING_REAL_SOURCE are used for every real-world fact not present in the
  source packet. There is NO PASS for this case.
* Determinism: no wall-clock timestamps; core properties are fixed; the saved
  .docx zip is re-written with a fixed entry timestamp so the output bytes are
  reproducible.

Run:  python build/build_0056_R03.py
"""

import hashlib
import os
import re
import zipfile
from datetime import datetime

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor, Twips

# --------------------------------------------------------------------------
# Paths (script lives in <root>/build/)
# --------------------------------------------------------------------------
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DOCX_PATH = os.path.join(ROOT, "0056_R03.docx")
NOTES_PATH = os.path.join(HERE, "BUILD_NOTES_R03.md")
PACKET_PATH = os.path.join(ROOT, "inputs", "0056_SOURCE_PACKET_R01.md")
BLUEPRINT_PATH = os.path.join(ROOT, "0056_UNIQUE_BLUEPRINT_R01.md")

# ---------------------------------------------------------------------------
# R03: no hidden characters. Long machine codes are wrapped by inserting an
# explicit display line break (<w:br/>) ONLY at '_' boundaries, and only when the
# token is wider than its column. The token text stays byte-exact (no U+200B).
# Boundary: display-only formatting; no code value is altered.
# ---------------------------------------------------------------------------
CODE_CHAR_MM = 2.7          # conservative width per code char at 9.5 pt
_TOKEN_RE = re.compile(r"[A-Za-z0-9_]+")
# U+200B is NOT used in R03 output; it is referenced only by the verifier to prove
# that no hidden break character survives in the DOCX.
ZWSP = "\u200b"

# --------------------------------------------------------------------------
# Style constants
# --------------------------------------------------------------------------
BODY_FONT = "Microsoft JhengHei"   # 微軟正黑體 (Win10 default CJK UI font)
BODY_SIZE = 10.5                   # body hard floor (pt)
TBL_SIZE = 9.5                     # table hard floor (pt)
META_SIZE = 9.5                    # header/footer size (pt)
NAVY = RGBColor(0x1F, 0x30, 0x4A)
GREY = RGBColor(0x59, 0x59, 0x59)
HDR_FILL = "D9E2F3"                # calm light blue for table headers
CALLOUT_FILL = "EEF3FB"            # calm light blue for the 30-second box

# Fixed synthetic "as-of" date used by the worked sample (SAMPLE/SYNTHETIC).
# Boundary: this is a fabricated date for the sample only - never a real record date.
SYNTH_DATE = "2026-10-17"
# Boundary: fabricated worked-case id; unrelated to any real course or to SYN-0033-PLAN-01.
SYNTH_CASE = "SYN-0056-SAFE-01"

# ===========================  XML helpers  ================================


def _set_rfonts(rPr, name):
    """Point ascii/hAnsi/eastAsia/cs font slots at *name*."""
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), name)


def style_run(run, size=BODY_SIZE, bold=False, name=BODY_FONT, color=None,
              italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = name
    if color is not None:
        run.font.color.rgb = color
    _set_rfonts(run._element.get_or_add_rPr(), name)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_repeat_header(row):
    """Mark a table row as a repeating header row (prints on every page)."""
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def set_cant_split(row):
    """Forbid a table row from splitting across a page boundary (<w:cantSplit/>).

    CT_TrPr requires cantSplit to appear before tblHeader, so insert accordingly.
    """
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    tblHeader = trPr.find(qn("w:tblHeader"))
    if tblHeader is not None:
        tblHeader.addprevious(el)
    else:
        trPr.append(el)


def keep_row_with_next(row):
    """Keep every paragraph of a row with the following row (header->first data)."""
    for cell in row.cells:
        for p in cell.paragraphs:
            p.paragraph_format.keep_with_next = True


def _col_widths_mm(weights, total_twips):
    """Return per-column widths in millimetres (same maths as set_table_widths)."""
    total = float(sum(weights))
    widths = [int(total_twips * w / total) for w in weights]
    return [w * 0.017639 for w in widths]


def break_underscore_token(line, width_mm):
    """Split *line* so long underscore tokens break ONLY at '_' boundaries.

    Boundary: display-only. Returns a list whose items are either plain strings
    (emitted as-is) or lists of pieces (emitted with an explicit <w:br/> between
    pieces). No character is added to, or removed from, any code value: the pieces
    concatenate back to the exact original token.
    """
    if not line or width_mm <= 0:
        return [line]
    budget = max(6, int(width_mm / CODE_CHAR_MM))
    parts = []
    last = 0
    for m in _TOKEN_RE.finditer(line):
        tok = m.group(0)
        parts.append(line[last:m.start()])
        if "_" in tok and len(tok) > budget:
            segs = tok.split("_")
            pieces = []
            cur = ""
            for i, s in enumerate(segs):
                piece = s + ("_" if i < len(segs) - 1 else "")
                if cur and len(cur) + len(piece) > budget:
                    pieces.append(cur)
                    cur = piece
                else:
                    cur += piece
            if cur:
                pieces.append(cur)
            parts.append(pieces)
        else:
            parts.append(tok)
        last = m.end()
    parts.append(line[last:])
    return parts


def set_table_widths(table, weights, total_twips):
    """Fixed layout with explicit column widths (no clipping, stable wraps)."""
    table.autofit = False
    total = float(sum(weights))
    widths = [int(total_twips * w / total) for w in weights]
    widths[-1] += total_twips - sum(widths)
    for i, w in enumerate(widths):
        table.columns[i].width = Twips(w)
    for row in table.rows:
        for j, cell in enumerate(row.cells):
            cell.width = Twips(widths[j])
    return widths


def usable_twips(section):
    """Usable text width of *section* in twips."""
    return int(section.page_width.twips - section.left_margin.twips
               - section.right_margin.twips)


def add_field(paragraph, instr, size=META_SIZE, color=GREY):
    run = paragraph.add_run()
    style_run(run, size, color=color)
    r = run._element
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
    it.text = instr
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    tx = OxmlElement("w:t"); tx.text = "1"
    f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
    for el in (f1, it, f2, tx, f3):
        r.append(el)
    return run


def configure_section(section, landscape):
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Mm(297)
        section.page_height = Mm(210)
        section.left_margin = Mm(14)
        section.right_margin = Mm(14)
        section.top_margin = Mm(14)
        section.bottom_margin = Mm(14)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.left_margin = Mm(18)
        section.right_margin = Mm(18)
        section.top_margin = Mm(16)
        section.bottom_margin = Mm(16)
    section.header_distance = Mm(9)
    section.footer_distance = Mm(9)


def setup_header_footer(section):
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_run(hp.add_run("設備安檢與場地投保證明｜第三方證據取得、適用性判定與檢查紀錄"),
              META_SIZE, color=GREY)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(fp.add_run("第 "), META_SIZE, color=GREY)
    add_field(fp, "PAGE")
    style_run(fp.add_run(" 頁，共 "), META_SIZE, color=GREY)
    add_field(fp, "NUMPAGES")
    style_run(fp.add_run(" 頁　｜　文件狀態：候選作業文件（含合成示範，非正式驗收）"),
              META_SIZE, color=GREY)


# ===========================  composition helpers  ========================


def para(doc, text="", size=BODY_SIZE, bold=False, color=None, align=None,
         space_after=4, space_before=0, keep_next=False, italic=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = 1.15
    pf.keep_with_next = keep_next
    if align is not None:
        p.alignment = align
    if text:
        style_run(p.add_run(text), size, bold=bold, color=color, italic=italic)
    return p


def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    pf = h.paragraph_format
    pf.keep_with_next = True
    if level == 1:
        pf.space_before = Pt(12); pf.space_after = Pt(6)
    else:
        pf.space_before = Pt(10); pf.space_after = Pt(4)
    return h


def bullet(doc, text, size=BODY_SIZE):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.12
    style_run(p.add_run(text), size)
    return p


def numbered(doc, index, text, size=BODY_SIZE):
    """Manual numbering (avoids Word list-continuation surprises)."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(2)
    pf.line_spacing = 1.12
    pf.left_indent = Mm(7)
    pf.first_line_indent = Mm(-7)
    style_run(p.add_run("%d. " % index), size, bold=True)
    style_run(p.add_run(text), size)
    return p


def _clear_cell(cell):
    for p in list(cell.paragraphs):
        for r in list(p.runs):
            r._element.getparent().remove(r._element)


def _emit_cell_lines(cell, text, size, bold, width_mm=9999.0):
    _clear_cell(cell)
    lines = text.split("\n") if text else [""]
    for k, line in enumerate(lines):
        p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        if not line:
            continue
        parts = break_underscore_token(line, width_mm)
        if not any(parts):
            continue
        run = p.add_run()
        style_run(run, size, bold=bold)
        pending_break = False
        wrote = False
        for part in parts:
            if isinstance(part, list):
                for idx, piece in enumerate(part):
                    if pending_break:
                        run._element.append(OxmlElement("w:br"))
                        pending_break = False
                    run.add_text(piece)
                    wrote = True
                    if idx < len(part) - 1:
                        pending_break = True
            else:
                if part == "":
                    continue
                if pending_break:
                    run._element.append(OxmlElement("w:br"))
                    pending_break = False
                run.add_text(part)
                wrote = True
        if not wrote:
            run.add_text("")


def set_cell_margins(table, top=0, bottom=0, left=57, right=57):
    """Explicit tight cell margins (twips) to compact vertical footprint."""
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn("w:tblCellMar")):
        tblPr.remove(old)
    mar = OxmlElement("w:tblCellMar")
    for tag, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement("w:" + tag)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    # keep OOXML child order: tblCellMar goes before tblLook (after tblLayout).
    look = tblPr.find(qn("w:tblLook"))
    if look is not None:
        look.addprevious(mar)
    else:
        tblPr.append(mar)


def add_table_block(doc, caption, headers, weights, rows, note=None, keep_together=False,
                    tight=False):
    if caption:
        heading(doc, caption, level=2)
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    usable = usable_twips(doc.sections[-1])
    widths_mm = _col_widths_mm(weights, usable)
    for j, htext in enumerate(headers):
        _emit_cell_lines(t.cell(0, j), htext, TBL_SIZE, True, widths_mm[j])
        shade_cell(t.cell(0, j), HDR_FILL)
    for r in rows:
        cells = t.add_row().cells
        for j, val in enumerate(r):
            _emit_cell_lines(cells[j], val, TBL_SIZE, False, widths_mm[j])
    set_repeat_header(t.rows[0])
    for row in t.rows:
        set_cant_split(row)             # R02: no row may split across pages
    keep_row_with_next(t.rows[0])       # R02: header stays with first data row
    if keep_together:                   # R03: keep the whole table on one page
        for row in t.rows[:-1]:
            keep_row_with_next(row)
    set_table_widths(t, weights, usable)
    if tight:
        set_cell_margins(t)
    if note:
        para(doc, note, size=BODY_SIZE, color=GREY, space_after=8, space_before=2)
    else:
        para(doc, "", size=4, space_after=0)
    return t


def add_callout(doc, lines):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.cell(0, 0)
    _clear_cell(cell)
    for i, (bold, text) in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.12
        style_run(p.add_run(text), BODY_SIZE, bold=bold)
    shade_cell(cell, CALLOUT_FILL)
    for row in t.rows:
        set_cant_split(row)             # R02: single-cell block cannot split
    set_table_widths(t, [1], usable_twips(doc.sections[-1]))
    para(doc, "", size=4, space_after=0)
    return t


def add_preformatted(doc, lines, size=TBL_SIZE):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.cell(0, 0)
    _clear_cell(cell)
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        style_run(p.add_run(line), size)
    for row in t.rows:
        set_cant_split(row)             # R02: single-cell block cannot split
    set_table_widths(t, [1], usable_twips(doc.sections[-1]))
    para(doc, "", size=4, space_after=0)
    return t


# ===========================  document assembly  ==========================


def _init_document():
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_SIZE)
    _set_rfonts(normal.element.get_or_add_rPr(), BODY_FONT)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15
    for nm, sz in (("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)):
        st = doc.styles[nm]
        st.font.name = BODY_FONT
        st.font.size = Pt(sz)
        st.font.bold = True
        st.font.color.rgb = NAVY
        _set_rfonts(st.element.get_or_add_rPr(), BODY_FONT)
    configure_section(doc.sections[0], landscape=False)
    setup_header_footer(doc.sections[0])
    cp = doc.core_properties
    cp.title = "設備安檢與場地投保證明：第三方證據取得、適用性判定與檢查紀錄"
    cp.subject = "對應官方指標 I12-E04（候選作業文件，含合成示範）"
    cp.author = "TTQS local build"
    cp.last_modified_by = "TTQS local build"
    cp.category = "association course administration"
    cp.comments = "Contains SAMPLE/SYNTHETIC worked example; not a real inspection or policy."
    cp.created = datetime(2026, 10, 17, 0, 0, 0)
    cp.modified = datetime(2026, 10, 17, 0, 0, 0)
    cp.revision = 1
    return doc


def build_intro(doc):
    """Section 1 (portrait): title, identity, 30-second page, roles, glossary."""
    para(doc, "設備安檢與場地投保證明：第三方證據取得、適用性判定與檢查紀錄",
         size=18, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER,
         space_after=2, space_before=0)
    para(doc, "小型協會課程場域與器材安全證據作業包（含合成示範案例）",
         size=11.5, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    add_table_block(
        doc, "文件識別",
        ["項目", "內容"], [3, 8],
        [
            ["文件名稱", "設備安檢與場地投保證明：第三方證據取得、適用性判定與檢查紀錄"],
            ["對應官方指標", "I12-E04｜官方要求（逐字）：「設備安檢或是場地投保證明。」"],
            ["官方來源定位", "勞動力發展署「訓練機構 TTQS 證據表」，第 4 頁"],
            ["文件狀態", "候選作業文件；含合成示範；非正式驗收、非協會決策"],
            ["版本", "初版（合成示範）"],
            ["情境代號", "SYN-0056-SAFE-01（合成示範；與任何真實課程無關）"],
        ],
        note="本文件為可直接列印、可直接於 Word 編輯的作業包；以低維護、最小必要欄位為原則。",
    )

    heading(doc, "一、30 秒用途（使用前先讀）", level=1)
    add_callout(doc, [
        (True, "這一頁要你做什麼"),
        (False, "找出本次課程將使用的場域與器材，確認是否有「安全檢查」或「場地投保證明」義務；"
                "依決策樹做一次「有來源依據」的適用性判定；向第三方索取正本；驗證後登記到期日與檢索方式。"),
        (True, "官方要求（逐字）"),
        (False, "「設備安檢或是場地投保證明。」"
                "（勞動力發展署「訓練機構 TTQS 證據表」，第 4 頁，官方指標 I12-E04）"),
        (True, "合成／真實界線"),
        (False, "本文件目前只含一份合成示範案例（非真實）。任何真實保單、檢查證明、機構、"
                "保單號或簽名在本文件中都不存在；相關欄位一律為 UNKNOWN 或 PENDING_REAL_SOURCE。"),
        (True, "紅線"),
        (False, "適用性「未知」或「沒有來源依據」時，一律 HOLD，不得自行當成「不適用（N/A）」。"),
    ])

    heading(doc, "二、使用時機（四步）", level=2)
    numbered(doc, 1, "開課前：列出本次將使用的場域與器材（填入表 T2）。")
    numbered(doc, 2, "判定適用性：依決策樹，得到一個決策代號（D1–D5），填入表 T1。")
    numbered(doc, 3, "索取與驗證：向第三方取得設備檢查文件或場地保險正本（表 T3、T4）。")
    numbered(doc, 4, "索引與到期：登記索引碼與到期日（表 T5），異常或事件進表 T7，並留決策軌跡（表 T6）。")

    heading(doc, "三、提案角色（尚待真實授權指派）", level=2)
    add_table_block(
        doc, None,
        ["角色（提案，PROPOSED）", "主要職責", "獨立性要求"], [3, 5, 3],
        [
            ["義務彙整人（場域／設備）", "建立與維護表 T1／T2 的義務清單", "無特別限制"],
            ["設備安檢執行者", "依表 T3 執行目視／功能檢查並記錄", "不得複核自己執行的檢查"],
            ["獨立複核／放行者", "複核更正結果並判定是否放行", "必須與設備安檢執行者不同人"],
            ["證據保管與索引者", "維護表 T4／T5 索引、保存與存取控制", "無特別限制"],
            ["授權核准者", "核准適用性判定與 HOLD 解除", "建議不與檢查複核者同一人"],
        ],
        note="以上皆為提案角色（PROPOSED），尚待真實授權指派；不得視為已存在的委派。",
    )

    heading(doc, "四、名詞與狀態對照（避免誤解）", level=2)
    add_table_block(
        doc, None,
        ["名詞／狀態", "意思", "常見誤用"], [3, 5, 4],
        [
            ["SAMPLE / SYNTHETIC", "示例資料，非真實世界紀錄", "誤當成真實紀錄"],
            ["UNKNOWN", "尚未取得來源，無法判定", "誤當成「不適用」"],
            ["PROPOSED", "建議角色，尚待真實授權指派", "誤當成已指派"],
            ["CANDIDATE", "僅範本或草稿完成，正本未到", "誤當成已通過"],
            ["REAL", "來自真實世界的文件或事實", "以範本冒充真實"],
            ["VERIFIED", "已由人核對來源無誤", "誤當成自動成立"],
            ["HOLD", "阻擋，不得放行", "誤當成失敗或結束"],
            ["PASS", "通過；本文件之合成案例「不」使用此狀態", "把完成範本當成 PASS"],
            ["PENDING_REAL_SOURCE", "等待真實正本到位", "誤當成已完成"],
        ],
        note="本文件凡未取得真實來源者，一律以 UNKNOWN／PENDING_REAL_SOURCE 表示，不得臆造。",
    )


def build_semantics(doc):
    """Official 'or' semantics, decision tree, status dictionary (portrait)."""
    heading(doc, "五、官方「或」語意與適用性決策", level=1)
    para(doc, "官方要求為「設備安檢『或是』場地投保證明」：任一路徑成立即可，"
              "但必須依「已記錄、有來源依據」的判定，不能憑感覺決定。", space_after=4)
    bullet(doc, "設備安檢路徑獨立成立 → 走「設備安檢」。")
    bullet(doc, "場地保險路徑獨立成立 → 走「場地保險」。")
    bullet(doc, "兩者各自獨立成立 → 兩者皆辦（both）。")
    bullet(doc, "判為不適用 → 必須有來源依據，才算 NOT_APPLICABLE_SUPPORTED。")
    bullet(doc, "適用性未知或缺依據 → HOLD，永不靜默當成 N/A。")
    bullet(doc, "不得臆造法規、承保限額或「協會政策」。")

    heading(doc, "六、適用性決策樹（每個節點輸出一組決策代號）", level=2)
    add_preformatted(doc, [
        "[起點] 本次將使用的場域／設備？",
        "  ├─ 場域／設備清單未知 → D5 HOLD_UNKNOWN",
        "  └─ 已知（先填表 T2）",
        "       ├─ 是否有「要求檢查」的來源依據（真實手冊／契約／法規）？",
        "       │    ├─ 有且已驗證 → APPLICABLE_CONFIRMED（走設備安檢）",
        "       │    ├─ 有但未驗證 → APPLICABLE_CANDIDATE（暫時 HOLD）",
        "       │    └─ 無／未知 → 不得判 N/A → D5 HOLD_UNKNOWN",
        "       ├─ 是否有「要求投保」的來源依據（真實租約／契約／規範）？",
        "       │    ├─ 有且已驗證 → APPLICABLE_CONFIRMED（走場地保險）",
        "       │    ├─ 有但未驗證 → APPLICABLE_CANDIDATE（暫時 HOLD）",
        "       │    └─ 無／未知 → 不得判 N/A → D5 HOLD_UNKNOWN",
        "       └─ 合併判定：",
        "            ├─ 僅檢查成立 → D1 APPLY_INSPECTION_ONLY",
        "            ├─ 僅保險成立 → D2 APPLY_INSURANCE_ONLY",
        "            ├─ 兩者皆成立 → D3 APPLY_BOTH",
        "            ├─ 兩者皆有依據地不適用 → D4 NOT_APPLICABLE_WITH_BASIS（須引用來源）",
        "            └─ 任一未知 → D5 HOLD_UNKNOWN",
    ])

    heading(doc, "七、狀態字典（代號、語意、可離開條件）", level=2)
    add_table_block(
        doc, None,
        ["代號", "類別", "語意", "允許的下一步"], [74, 14, 44, 42],
        [
            ["APPLICABLE_CONFIRMED", "適用性", "有「已驗證」的真實來源要求", "進入證據索取"],
            ["APPLICABLE_CANDIDATE", "適用性", "有來源但尚未驗證", "補驗證；維持 HOLD"],
            ["NOT_APPLICABLE_SUPPORTED", "適用性", "有已驗證來源支持不適用", "可移出分母（需註明來源）"],
            ["UNKNOWN_APPLICABILITY", "適用性", "適用性未知", "HOLD；禁止 N/A"],
            ["EVIDENCE_MISSING", "證據", "尚無任何原本文檔", "建立索取"],
            ["REQUEST_DRAFTED_NOT_SENT", "證據", "索取函已草擬、尚未寄出", "取得授權後寄出"],
            ["REQUESTED_PENDING_RESPONSE", "證據", "已寄出、等待回覆", "追蹤"],
            ["RECEIVED_UNVERIFIED", "證據", "已收到但尚未驗證", "進入驗證步驟"],
            ["RECEIVED_VERIFIED_REAL", "證據", "已驗證為真實正本", "進入索引與到期管理"],
            ["EXPIRED_OR_OUT_OF_PERIOD", "證據", "已過期或不涵蓋課程期間", "HOLD 或補件"],
            ["MISMATCH_OR_UNVERIFIABLE", "證據", "內容不符或無法驗證", "HOLD"],
            ["CANDIDATE", "整體", "僅範本完整、正本未到", "不得視為 PASS"],
            ["PENDING_REAL_SOURCE", "整體", "等待真實正本", "續辦"],
            ["HOLD", "整體", "阻擋", "解決後解除"],
        ],
        note="N/A 只允許在 NOT_APPLICABLE_SUPPORTED 生效後出現；且僅有已支持之適用性複核才能把該義務移出分母。",
        tight=True,
    )


def build_forms(doc):
    """Landscape annex: the five wide operational forms T1-T5."""
    doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(doc.sections[-1], landscape=True)
    setup_header_footer(doc.sections[-1])

    heading(doc, "八、作業表單（A4 橫向，方便列印與填寫）", level=1)
    para(doc, "本區含表 T1 至 T5。欄位較多的表單改用 A4 橫向排版，避免欄位被擠壓；"
              "每一表格跨頁時會自動重複表頭列。合成示範列以（SAMPLE）標示；"
              "標示為（真實資料填寫列）的列，是取得真實來源後才填入的作業列。", space_after=6)

    # ---- T1 -------------------------------------------------------------
    add_table_block(
        doc, "表 T1　適用性判定表",
        ["義務編號\nobligation_id", "路徑", "判定情境", "是否有\n要求來源", "來源類型",
         "來源引用", "來源是否\n已驗證", "適用性狀態", "判定決策\n代號", "依據摘要",
         "決策者\n（提案角色）", "日期", "備註／例外"],
        [32, 14, 20, 16, 15, 15, 20, 38, 22, 24, 18, 15, 18],
        [
            # SAMPLE/SYNTHETIC row - synthetic obligation id and description only.
            ["OB-SYN-01（SAMPLE）", "設備安檢",
             "本次合成課程將使用合成可攜延長線（代號 EQP-SYN-CBL-07），置於合成場域 SITE-SYN-A",
             "未知", "UNKNOWN", "UNKNOWN", "否（尚未取得來源）",
             "UNKNOWN_APPLICABILITY", "D5 HOLD_UNKNOWN",
             "未取得真實手冊／契約／租約，無法確認是否要求設備檢查",
             "義務彙整人（PROPOSED）", "2026-10-17（合成）", "不得默認 N/A"],
            ["OB-SYN-02（SAMPLE）", "場地保險",
             "本次合成課程將使用合成場域 SITE-SYN-A",
             "未知", "UNKNOWN", "UNKNOWN", "否（尚未取得來源）",
             "UNKNOWN_APPLICABILITY", "D5 HOLD_UNKNOWN",
             "未取得真實租約／契約／規範，無法確認是否要求投保",
             "義務彙整人（PROPOSED）", "2026-10-17（合成）", "不得默認 N/A"],
            ["（真實資料填寫列）", "", "", "", "", "", "", "", "", "", "", "", ""],
            ["（真實資料填寫列）", "", "", "", "", "", "", "", "", "", "", "", ""],
        ],
        note="判定決策代號對應第六節決策樹：D1 僅檢查／D2 僅保險／D3 兩者／D4 有依據不適用／D5 未知即 HOLD。",
    )

    # ---- T2 -------------------------------------------------------------
    add_table_block(
        doc, "表 T2　資產／場域義務矩陣",
        ["義務編號\nobligation_id", "資產／場域\n類型", "識別碼", "位置", "關聯課程\n（代號）",
         "使用日期\n範圍", "需求來源", "判定理由", "負責角色\n（提案）", "所需證據",
         "到期／\n需求日", "證據狀態", "證據定位碼", "備註"],
        [32, 14, 26, 22, 22, 16, 12, 18, 18, 14, 16, 26, 18, 17],
        [
            # SAMPLE/SYNTHETIC rows: fabricated site/equipment/course codes only;
            # real identifiers and addresses stay UNKNOWN until a real source exists.
            ["OB-SYN-01（SAMPLE）", "設備", "EQP-SYN-CBL-07（SAMPLE）",
             "合成場域 SITE-SYN-A；真實地址 UNKNOWN", "CRS-SYN-2026-Q4-A（SAMPLE）",
             "2026-11-03（SAMPLE）", "UNKNOWN", "來源未知，暫列為可能適用",
             "義務彙整人（PROPOSED）", "檢查報告", "2026-11-03（SAMPLE）",
             "EVIDENCE_MISSING", "EV-SYN-01", "待取得真實來源後再判定"],
            ["OB-SYN-02（SAMPLE）", "場地", "SITE-SYN-A（SAMPLE）",
             "合成場域 SITE-SYN-A；真實地址 UNKNOWN", "CRS-SYN-2026-Q4-A（SAMPLE）",
             "2026-11-03（SAMPLE）", "UNKNOWN", "來源未知，暫列為可能適用",
             "義務彙整人（PROPOSED）", "保險正本", "2026-11-03（SAMPLE）",
             "EVIDENCE_MISSING", "EV-SYN-02", "文件本身尚缺；索取狀態見表 T4"],
            ["（真實資料填寫列）", "", "", "", "", "", "", "", "", "", "", "", "", ""],
            ["（真實資料填寫列）", "", "", "", "", "", "", "", "", "", "", "", "", ""],
        ],
        note="識別碼若為合成值一律加註 SAMPLE；真實識別碼與地址在取得真實來源前維持 UNKNOWN。",
    )

    # ---- T3 -------------------------------------------------------------
    heading(doc, "設備安檢流程（對應表 T3）", level=3)
    numbered(doc, 1, "由表 T2 取出本次要使用的設備義務。")
    numbered(doc, 2, "檢查前先確認「允收基準來源」（真實手冊／契約）；沒有基準就不得自訂門檻，"
                     "一律記為 UNKNOWN / obtain and verify。")
    numbered(doc, 3, "執行檢查（目視／功能測試），並記錄檢查發現。")
    numbered(doc, 4, "發現異常 → 立即隔離並掛牌停用（tag-out，即隔離並掛牌禁止使用），"
                     "記錄負責角色與更正措施。")
    numbered(doc, 5, "更正後由「獨立複核者」複核，並依「真實允收基準」判定是否放行；沒有基準就 HOLD。")
    numbered(doc, 6, "全程以證據編號回填表 T5，並在表 T6 留下決策軌跡。")
    numbered(doc, 7, "除非有真實驗證正本，否則不得聲稱由合格第三方完成檢測。")
    add_table_block(
        doc, "表 T3　設備檢查與更正複核表",
        ["檢查編號\ncheck_id", "義務編號\nobligation_id", "檢查項目", "檢查方法", "允收基準來源",
         "檢查發現", "立即隔離／\n掛牌停用", "負責角色\n（提案）", "更正措施",
         "更正後複核者\n（提案，須獨立）", "複核日期", "放行基準", "放行狀態", "證據定位碼", "備註"],
        [23, 32, 18, 16, 16, 18, 18, 16, 18, 18, 15, 16, 15, 15, 18],
        [
            # SAMPLE/SYNTHETIC row: fabricated damaged-cable finding; no real
            # inspection occurred and none is implied. Release basis stays UNKNOWN.
            ["CHK-SYN-01（SAMPLE）", "OB-SYN-01（SAMPLE）", "可攜電源延長線外被（纜線外皮）",
             "目視檢查（非第三方專業檢測）", "UNKNOWN / obtain and verify",
             "外被破損（合成示範缺陷，非真實檢查結果）",
             "是；掛牌停用（tag-out，隔離並掛牌禁止使用）", "設備安檢執行者（PROPOSED）",
             "更換或修復該延長線（提案）", "獨立複核／放行者（PROPOSED，須與執行者不同人）",
             "UNKNOWN（尚未複核）", "UNKNOWN（無真實允收來源）", "HOLD", "EV-SYN-01",
             "合成示範，不代表實際檢查已發生；不得聲稱由合格第三方完成檢測"],
            ["（真實資料填寫列）", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
            ["（真實資料填寫列）", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
        ],
        note="放行必須同時具備：真實允收基準 + 已完成更正 + 由不同人獨立複核。缺一即 HOLD。",
    )


def build_insurance_forms(doc):
    """Landscape annex continued: T4 and T5."""
    add_table_block(
        doc, "表 T4　場地保險正本索取／索引／驗證表",
        ["索取編號\nreq_id", "義務編號\nobligation_id", "向誰索取\n（角色）", "收件方",
         "索取範圍", "期間", "承保／排除\n（若正本載明）", "簽發機構", "保單號",
         "文件版本\n／日期", "來源驗證結果", "保留／存取\n限制", "索引碼", "到期日",
         "索取狀態", "備註"],
        [16, 32, 16, 14, 18, 16, 18, 12, 12, 15, 18, 18, 14, 12, 23, 18],
        [
            # SAMPLE/SYNTHETIC row: request is drafted-not-sent; no policy/type/
            # issuer/number/coverage exists and no public-liability type is presumed.
            ["REQ-SYN-01（SAMPLE）", "OB-SYN-02（SAMPLE）",
             "場地方或契約之保險窗口（PROPOSED）", "UNKNOWN（範本欄位）",
             "本次使用場地 SITE-SYN-A 之保險文件；險種由來源認定（本會不預設）",
             "須涵蓋 2026-11-03（SAMPLE）", "UNKNOWN（未取得正本，不得臆造）",
             "UNKNOWN", "UNKNOWN", "UNKNOWN", "尚未取得文件，無驗證結果",
             "真實文件限本機安全存放；僅最小必要欄位進索引；保存年限 UNKNOWN",
             "EV-SYN-02", "UNKNOWN", "REQUEST_DRAFTED_NOT_SENT",
             "無保單或證明附檔；索取範本尚未寄出"],
            ["（真實資料填寫列）", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
            ["（真實資料填寫列）", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
        ],
        note="險種、承保範圍、限額與簽發機構一律由真實正本認定；未取得前一律 UNKNOWN，禁止臆造。",
    )

    add_table_block(
        doc, "表 T5　證據狀態與到期追蹤表",
        ["證據編號\nevidence_id", "義務編號\nobligation_id", "證據類型", "索引碼", "目前狀態",
         "取得日", "到期／複查日", "距到期天數", "續辦動作", "HOLD 原因", "負責角色\n（提案）"],
        [32, 32, 14, 18, 26, 16, 20, 20, 35, 34, 28],
        [
            ["EV-SYN-01（SAMPLE）", "OB-SYN-01（SAMPLE）", "檢查", "EV-SYN-01",
             "EVIDENCE_MISSING", "UNKNOWN", "UNKNOWN", "無法計算（日期未知）",
             "先確認允收基準來源，再執行並記錄檢查",
             "無真實允收來源與檢查紀錄", "設備安檢執行者（PROPOSED）"],
            ["EV-SYN-02（SAMPLE）", "OB-SYN-02（SAMPLE）", "保險", "EV-SYN-02",
             "EVIDENCE_MISSING（索取狀態見表 T4）", "UNKNOWN", "UNKNOWN", "無法計算（日期未知）",
             "取得授權後寄出索取函並追蹤回覆",
             "尚無保單或證明正本", "證據保管與索引者（PROPOSED）"],
            ["（真實資料填寫列）", "", "", "", "", "", "", "", "", "", ""],
        ],
        note="「距到期天數」＝到期日 − 作業基準日；到期日未知時不計算。"
             "「證據狀態」描述文件本身，「索取狀態」（表 T4）描述取得動作，兩者為不同維度。",
        keep_together=True,
        tight=True,
    )


def build_records_and_sample(doc):
    """Portrait section: T6, worked synthetic case, T9 register, completeness,
    unsent acquisition template."""
    doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(doc.sections[-1], landscape=False)
    setup_header_footer(doc.sections[-1])

    heading(doc, "九、決策軌跡（表 T6）", level=1)
    para(doc, "凡適用性判定、證據接受、複核放行、HOLD 與 HOLD 解除，都要留下一列軌跡，"
              "方便日後查核「當時依據什麼做決定」。", space_after=4)
    add_table_block(
        doc, "表 T6　決策軌跡表",
        ["決策編號\ndecision_id", "日期", "決策類型", "決策代號", "依據摘要",
         "決策角色\n（提案）", "前後狀態", "備註"],
        [3, 2.6, 3, 3.4, 5, 3.4, 4, 4],
        [
            # SAMPLE/SYNTHETIC row: fabricated decision trail entry (HOLD, no real source).
            ["DEC-SYN-01（SAMPLE）", "2026-10-17（合成）", "適用性", "D5 HOLD_UNKNOWN",
             "未取得任何真實來源支持適用性", "義務彙整人（PROPOSED）",
             "前：UNKNOWN_APPLICABILITY → 後：HOLD", "兩條義務路徑皆維持 HOLD"],
            ["（真實資料填寫列）", "", "", "", "", "", "", ""],
            ["（真實資料填寫列）", "", "", "", "", "", "", ""],
        ],
        note="本表為 append-only（只增不改）：錯誤更正以新增一列並註明取代關係處理。",
    )

    heading(doc, "十、合成示範案例 SYN-0056-SAFE-01（全填、內部一致）", level=1)
    para(doc, "以下為一份完整的合成示範，用來示範各表如何填寫。"
              "本案例全部為 SAMPLE / SYNTHETIC，與任何真實課程、場域或器材無關，"
              "也不代表實際檢查已發生；結論維持 HOLD / PENDING_REAL_SOURCE。", space_after=4)
    bullet(doc, "情境代號：SYN-0056-SAFE-01（合成；與任何真實課程無關）。")
    bullet(doc, "場域代號：SITE-SYN-A（合成）；真實地址一律 UNKNOWN。")
    bullet(doc, "設備代號：EQP-SYN-CBL-07（合成可攜延長線）。")
    bullet(doc, "合成課程代號：CRS-SYN-2026-Q4-A（合成）；使用日期 2026-11-03（合成）。")
    bullet(doc, "示範缺陷：可攜電源延長線外被破損（合成示範用缺陷）。")
    para(doc, "對應結果：表 T1 兩條義務（OB-SYN-01 設備安檢、OB-SYN-02 場地保險）皆為 "
              "UNKNOWN_APPLICABILITY / D5 HOLD_UNKNOWN；表 T2 證據狀態皆為 EVIDENCE_MISSING；"
              "表 T3 放行狀態 HOLD；表 T4 索取狀態 REQUEST_DRAFTED_NOT_SENT（尚未寄出、無附檔）；"
              "表 T5 兩筆證據皆 EVIDENCE_MISSING；表 T6 一筆決策 DEC-SYN-01。整體結論 HOLD / PENDING_REAL_SOURCE。",
         space_after=6)

    heading(doc, "十一、合成替換登錄（表 T9，逐筆事實登錄）", level=2)
    para(doc, "下表把本案例用到的每一個合成值逐筆登錄，並指出真實替換來源與替換後動作。"
              "「所在表／欄」即列印於本文件內的定位碼，方便逐一比對。", space_after=4)
    # Boundary: T9 registers every synthetic literal used above, each with its exact
    # in-document locator and the real replacement source/action. SF-01..SF-10 match the
    # Blueprint; SF-11..SF-19 register the remaining synthetic ids/dates/version labels.
    add_table_block(
        doc, "表 T9　合成替換登錄表（fact-level）",
        ["sample_fact_id", "值（SAMPLE/SYNTHETIC）", "所在表／欄",
         "真實替換來源（待取得）", "替換後動作"],
        [3, 4.6, 3.6, 4.6, 4.2],
        [
            ["SF-01", "SITE-SYN-A", "表 T2 識別碼／位置", "真實場域清單／租約", "取得後覆寫並保留樣本"],
            ["SF-02", "EQP-SYN-CBL-07", "表 T2 識別碼", "真實設備清冊", "取得後覆寫並保留樣本"],
            ["SF-03", "CRS-SYN-2026-Q4-A", "表 T2 關聯課程", "真實課程排程", "取得後覆寫並保留樣本"],
            ["SF-04", "2026-11-03", "表 T2 使用日期範圍", "真實課程日期", "取得後覆寫並保留樣本"],
            ["SF-05", "外被破損（缺陷）", "表 T3 檢查發現", "真實檢查報告", "取得後覆寫並保留樣本"],
            ["SF-06", "tag-out 隔離（掛牌停用）", "表 T3 立即隔離欄", "真實隔離紀錄", "取得後覆寫並保留樣本"],
            ["SF-07", "檢查放行 HOLD", "表 T3 放行狀態", "真實允收基準", "取得後覆寫並保留樣本"],
            ["SF-08", "REQUEST_DRAFTED_NOT_SENT", "表 T4 索取狀態", "真實寄送／回函紀錄", "取得後覆寫並保留樣本"],
            ["SF-09", "決策日期 2026-10-17", "表 T6 日期", "真實決策日期", "取得後覆寫並保留樣本"],
            ["SF-10", "所有 PROPOSED 角色", "表 T1–T6 角色欄", "真實授權指派", "取得後覆寫並保留樣本"],
            ["SF-11", "OB-SYN-01、OB-SYN-02", "表 T1／T2 義務編號", "真實義務清單", "取得後覆寫並保留樣本"],
            ["SF-12", "CHK-SYN-01", "表 T3 檢查編號", "真實檢查單編號", "取得後覆寫並保留樣本"],
            ["SF-13", "REQ-SYN-01", "表 T4 索取編號", "真實索取單編號", "取得後覆寫並保留樣本"],
            ["SF-14", "EV-SYN-01、EV-SYN-02", "表 T5 證據編號／索引碼", "真實證據編號／索引碼", "取得後覆寫並保留樣本"],
            ["SF-15", "DEC-SYN-01", "表 T6 決策編號", "真實決策編號", "取得後覆寫並保留樣本"],
            ["SF-16", "T1 判定日期 2026-10-17", "表 T1 日期欄", "真實判定日期", "取得後覆寫並保留樣本"],
            ["SF-17", "合成判定情境敘述文字", "表 T1 判定情境", "真實課程場域與設備描述", "取得後覆寫並保留樣本"],
            ["SF-18", "EX-SYN-01", "表 T7 例外編號", "真實例外／事件編號", "取得後覆寫並保留樣本"],
            ["SF-19", "版本「初版（合成示範）」", "表 T8 版本", "真實文件版本", "取得後覆寫並保留樣本"],
        ],
        note="合成樣本與 REAL 證據分區存放；本表所列樣本不因替換而刪除，僅移入樣本區保留軌跡。",
    )


def build_completeness_and_template(doc):
    heading(doc, "十二、完整度摘要（可重算；非官方分數）", level=1)
    bullet(doc, "分母：本合成情境中「已記錄的可能適用義務數」＝ 2"
                "（OB-SYN-01 設備安檢、OB-SYN-02 場地保險）。"
                "分母只計已記錄的義務；唯有某義務達到 NOT_APPLICABLE_SUPPORTED"
                "（具有已驗證來源支持的適用性複核）時，才可移出分母並註明來源。")
    bullet(doc, "分子：其中已達 RECEIVED_VERIFIED_REAL（已驗證為真實正本）者 ＝ 0。")
    bullet(doc, "公式：完整度 ＝ 分子 ÷ 分母 ＝ 0 ÷ 2 ＝ 0%。")
    bullet(doc, "結論：HOLD / PENDING_REAL_SOURCE（本案例不適用 PASS）。")
    para(doc, "聲明：此數字僅為工作用完整度摘要，不是官方 TTQS 分數，也不是協會政策；"
              "不得把尚未驗證的文件當成已通過。", space_after=6)

    heading(doc, "十三、第三方文件索取範本（未寄出）", level=1)
    para(doc, "下列範本可直接複製使用。收件人與日期為「範本欄位」，寄出前才填寫；"
              "本文件沒有任何已寄送紀錄，示範狀態為 NOT SENT，且無保單或證明附檔。", space_after=4)
    add_preformatted(doc, [
        "主旨：［範本欄位］請提供本次場地使用之保險與設備安檢相關文件",
        "收件人（提案角色：場地方／契約保險窗口）：［範本欄位］",
        "副本（提案角色：證據保管與索引者）：［範本欄位］",
        "發文日期：［範本欄位］",
        "本會名稱：（比對已確認之協會名稱；本文件不列示名稱）",
        "",
        "敬啟者：",
        "本會擬於下列期間使用 貴方場地／設備辦理訓練活動，為完備場地安全與保險相關紀錄，",
        "敬請協助提供下列文件之影本或電子檔（正本驗後歸還或依約留存）：",
        "  1) 場地保險證明：請依 貴方場地與契約之要求，認定並載明保險類型",
        "     （本會不預設險種）、被保險人、承保期間、承保範圍與排除條款。",
        "  2) 設備安全檢查相關文件：如近期檢查或維護紀錄（若適用）。",
        "另請提供認定適用性的依據來源（例如租約或契約中相關條款之頁面）。",
        "提醒：本會尚待取得真實來源才能判定適用性；若無相關文件，亦請回覆說明，本會將據以記錄。",
        "",
        "（本範本尚未寄出；收件人與日期為範本欄位。）",
    ], size=TBL_SIZE)
    para(doc, "收到文件後，依第十四節的來源驗證步驟處理；若僅本範本完成而正本仍未到，"
              "狀態為 CANDIDATE / PENDING_REAL_SOURCE，不得標示為通過或已接受。",
         size=BODY_SIZE, color=GREY, space_after=6)


def build_replacement_and_controls(doc):
    heading(doc, "十四、真實來源替換路徑與來源驗證步驟", level=1)
    para(doc, "替換路徑（針對每一個 sample_fact_id）：取得真實來源 → 由人驗證 → 覆寫對應欄位 → "
              "於表 T8 記錄版本 → 樣本移入樣本區（不刪除，保留軌跡）。", space_after=4)
    para(doc, "來源驗證步驟（每一項都要留下驗證者角色與日期）：", space_after=2)
    numbered(doc, 1, "文件真實性：簽發者／來源可辨識且合理。")
    numbered(doc, 2, "主體比對：被保險人＝已確認之協會名稱（未確認前一律 UNKNOWN）。")
    numbered(doc, 3, "場域／設備比對：與表 T2 識別碼一致。")
    numbered(doc, 4, "期間比對：涵蓋表 T2 所載使用日期範圍。")
    numbered(doc, 5, "內容比對：承保或檢查範圍與排除條款，以正本載明者為準。")
    numbered(doc, 6, "版本／日期：無塗改疑慮；保單號與簽發機構與正本一致。")
    numbered(doc, 7, "判定：通過 → RECEIVED_VERIFIED_REAL；不符 → MISMATCH_OR_UNVERIFIABLE；"
                     "過期或不涵蓋 → EXPIRED_OR_OUT_OF_PERIOD。")
    para(doc, "隱私與保存邊界：真實保單明細、簽名、契約與個人資料不得進入公開內容；"
              "限本機安全存放，存取限提案角色（待真實指派）；僅最小必要欄位進入索引；"
              "保存年限待真實政策決定，現為 UNKNOWN。", space_after=6)

    heading(doc, "十五、例外、到期、檢索、存取、更正與版本", level=1)
    heading(doc, "15.1 到期與續辦", level=2)
    para(doc, "提案作法：為每筆證據設定到期／複查日，並於到期前提醒（提醒天數待真實政策決定，"
              "現為 UNKNOWN）。已過期或不涵蓋課程期間者，狀態為 EXPIRED_OR_OUT_OF_PERIOD，"
              "一律 HOLD 並依表 T5 續辦動作補件。", space_after=4)
    heading(doc, "15.2 檢索測試（retrieval test）", level=2)
    para(doc, "提案作法：定期（頻率待真實政策決定）隨機抽一筆索引碼，確認能否在合理時間內"
              "取出指定文件並核對版本與到期日，並記錄結果。"
              "合成示範現況：尚無真實證據，檢索測試為「未執行（HOLD）」。", space_after=4)
    heading(doc, "15.3 存取與隱私", level=2)
    para(doc, "真實文件以最小必要原則保存與索引；存取僅限提案角色（待真實指派）。"
              "公開內容不得包含真實保單明細、簽名、契約或個人資料。", space_after=4)
    heading(doc, "15.4 更正與矯正（表 T7）", level=2)
    para(doc, "異常處理順序：發現缺失 → 立即隔離並掛牌停用（tag-out）→ 更正 → 由不同人獨立複核 → 解除。"
              "缺少真實允收基準時，一律維持 HOLD。", space_after=4)
    add_table_block(
        doc, "表 T7　例外／事件與矯正表",
        ["例外編號\nexception_id", "來源 obligation/check", "描述", "分類", "立即處置",
         "矯正", "獨立複核", "解除條件", "狀態"],
        [29, 22, 16, 14, 18, 16, 18, 20, 23],
        [
            # SAMPLE/SYNTHETIC row: synthetic exception derived from the synthetic
            # cable finding; correction/recheck/release remain pending real sources.
            ["EX-SYN-01（SAMPLE）", "CHK-SYN-01／OB-SYN-01",
             "合成示範：可攜延長線外被破損", "缺失",
             "掛牌停用（tag-out）（合成示範）", "更換或修復（提案）",
             "待指派之獨立複核者（PROPOSED）；複核日期 UNKNOWN",
             "依真實允收來源完成更正，並由獨立者複核通過", "HOLD / PENDING_REAL_SOURCE"],
            ["（真實資料填寫列）", "", "", "", "", "", "", "", ""],
            ["（真實資料填寫列）", "", "", "", "", "", "", "", ""],
        ],
        note="分類可選：缺失／過期／不符／無法驗證／其他。未達解除條件前，該設備或場地不得使用。",
    )
    heading(doc, "15.5 版本軌跡（表 T8）", level=2)
    add_table_block(
        doc, "表 T8　版本軌跡表",
        ["版本", "日期", "變更摘要", "變更者（提案）", "對應決策"],
        [2.4, 3, 6, 3.6, 3.4],
        [
            # SAMPLE/SYNTHETIC row: synthetic-only initial version record.
            ["初版（合成示範）", "2026-10-17（合成）", "建立本作業包與合成示範案例",
             "文件建立者（PROPOSED）", "DEC-SYN-01"],
            ["（真實資料填寫列）", "", "", "", ""],
        ],
        note="每次由真實來源覆寫樣本、或修訂作業規則時，新增一列版本紀錄，不覆蓋舊列。",
    )


def build_tests_and_sources(doc):
    heading(doc, "十六、HOLD 與退件條件（任一成立即不得放行）", level=1)
    bullet(doc, "適用性未知或缺少來源依據，卻默認當成 N/A。")
    bullet(doc, "只有範本或草稿完成、正本未到，卻標示為已接受或已通過。")
    bullet(doc, "證據已過期、內容不符或無法驗證，仍標示為有效。")
    bullet(doc, "自訂法規門檻，或臆造承保限額、保單號、簽發機構。")
    bullet(doc, "聲稱已寄送、已檢查、已放行或已核准，卻沒有真實紀錄。")
    bullet(doc, "使用真實場域、地址、人員或簽名，卻未取得授權與真實來源。")
    bullet(doc, "缺少可重算的完整度分子／分母說明，或把完整度當成官方分數。")
    bullet(doc, "在標題盲測下，可替代官方指標 I12-E01／I12-E02／I12-E03／I13-E01 之文件。")

    heading(doc, "十七、跨文件邊界與負向鄰居測試", level=1)
    para(doc, "本文件只處理「第三方檢查與保單正本的取得、驗證與索引」。下列鄰近官方指標各有不同主題，"
              "本文件不得替代它們；此處僅為邊界測試，不代表這些文件共用本案例。",
         space_after=4, keep_next=True)
    add_table_block(
        doc, None,
        ["鄰近官方指標", "主題", "為何不能互換（本文件的獨有內容）"], [3, 4, 7],
        [
            ["I12-E01", "課程適配檢視與調整",
             "本文件處理第三方檢查與保單正本的取得與驗證，不是課程設計適配。"],
            ["I12-E02", "學員遴選與教材建議",
             "本文件不涉及學員或教材，聚焦場域安全與投保證據。"],
            ["I12-E03", "講師遴選與聘任",
             "本文件不涉及講師；核心是義務矩陣、檢查複核、保險索取與決策軌跡。"],
            ["I13-E01", "訓練遷移檢視",
             "屬不同指標與結果導向；本文件為資格性證據的取得與合規。"],
        ],
        note="交付識別與官方勾選序號不得混用；未經逐位元驗證，不得引用跨文件事實。",
    )

    heading(doc, "十八、審查者自我測試", level=1)
    numbered(doc, 1, "標題盲測：移除標題後，仍可由表結構（義務矩陣、檢查複核、保險索取驗證、決策軌跡）"
                     "辨識為本文件。")
    numbered(doc, 2, "負向鄰居：對照 I12-E01／I12-E02／I12-E03／I13-E01，本文件不應在這些主題上成立或可替代。")
    numbered(doc, 3, "官方語意：確認「或」的處理（僅檢查／僅保險／兩者／有依據不適用）與「未知即 HOLD」可追溯。")
    numbered(doc, 4, "真實性：全文件無臆造保單、證明、機構、保單號、簽名、已寄送或核准；缺少者一律 UNKNOWN／PENDING。")
    numbered(doc, 5, "樣本一致：SYN-0056-SAFE-01 的每個值都能回溯到表 T9；且與任何真實課程無連結。")
    numbered(doc, 6, "可重算：完整度的分子／分母可獨立重算，且有非官方聲明。")
    numbered(doc, 7, "隱私：公開內容不含真實敏感正本。")
    numbered(doc, 8, "排版：正文 ≥ 10.5 pt、表格 ≥ 9.5 pt、跨頁重複表頭；正文不出現內部雜湊、"
                     "佇列代號、JSON 或檔案路徑。")

    heading(doc, "十九、來源定位（一頁）", level=1)
    bullet(doc, "官方來源：勞動力發展署「訓練機構 TTQS 證據表」。")
    bullet(doc, "官方要求（逐字）：「設備安檢或是場地投保證明。」")
    bullet(doc, "定位：證據表第 4 頁（官方指標 I12-E04 所在列）。")
    bullet(doc, "覆蓋範圍：設備安檢路徑、場地保險路徑、兩者並行、有依據之不適用、未知即 HOLD。")
    bullet(doc, "本文件僅引用上列官方要求之逐字文字與頁碼層級定位；完整來源檔案核對資訊"
                "另存於內部紀錄，不列入本文件。")
    para(doc, "（本文件為候選作業文件，含合成示範；非正式驗收、非協會決策，亦不代表任何真實檢查、"
              "保單或第三方正本已存在。）", size=BODY_SIZE, color=GREY, space_after=6)


def normalize_zip(path):
    """Re-write the .docx zip with a fixed entry timestamp (byte-stable output)."""
    tmp = path + ".tmp"
    with zipfile.ZipFile(path, "r") as zin, \
            zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            zi = zipfile.ZipInfo(item.filename, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = item.external_attr
            zi.internal_attr = item.internal_attr
            zi.create_system = 0
            zout.writestr(zi, data)
    os.replace(tmp, path)


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def build_document():
    doc = _init_document()
    build_intro(doc)
    build_semantics(doc)
    build_forms(doc)              # landscape annex: T1-T3
    build_insurance_forms(doc)    # landscape annex: T4-T5
    build_records_and_sample(doc) # portrait: T6, sample, T9, completeness, template
    build_completeness_and_template(doc)
    build_replacement_and_controls(doc)
    build_tests_and_sources(doc)
    return doc


def verify(docx_path):
    """Re-open the DOCX and check the binding content requirements."""
    from docx import Document as _D
    doc = _D(docx_path)
    body = "\n".join(p.text for p in doc.paragraphs)
    cell_text = []
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                cell_text.append(c.text)
    all_text = body + "\n" + "\n".join(cell_text)
    # R03: normalise away display-only line breaks and (defensively) any U+200B so
    # substring checks compare the exact underlying code values.
    norm = all_text.replace(ZWSP, "").replace("\n", "").replace("\r", "")

    checks = {}
    checks["tables>=12"] = (len(doc.tables) >= 12, "tables=%d" % len(doc.tables))
    for n in range(1, 10):
        checks["table T%d" % n] = ("表 T%d" % n in body, "")
    checks["official quote"] = ("設備安檢或是場地投保證明。" in norm, "")
    checks["case id"] = ("SYN-0056-SAFE-01" in norm, "")
    checks["no SYN-0033 link"] = ("SYN-0033-PLAN-01" not in norm, "")
    checks["0 / 2 = 0%"] = (("0 ÷ 2" in norm) and ("0%" in norm) and ("分母" in norm), "")
    checks["unsent request"] = ("REQUEST_DRAFTED_NOT_SENT" in norm, "")
    checks["PENDING_REAL_SOURCE"] = ("PENDING_REAL_SOURCE" in norm, "")
    checks["HOLD present"] = ("HOLD" in norm, "")
    checks["T9 header"] = (all(h in norm for h in
                             ["sample_fact_id", "真實替換來源（待取得）", "所在表／欄"]), "")
    checks["T9 has SF-01..SF-19"] = (all(("SF-%02d" % i) in norm
                                        for i in range(1, 20)), "")
    checks["two obligations"] = (("OB-SYN-01" in norm) and ("OB-SYN-02" in norm), "")
    checks["no public-liability presum"] = ("公共意外責任" not in norm, "")
    checks["no forbidden fill words"] = (all(w not in norm for w in
                                            ["待填", "待核", "空白表", "保持空白", "會後填寫"]), "")
    checks["no internal hash/path"] = (all(w not in norm for w in
                                          ["/docs/", "\\\\build\\\\", "git ", "sha256",
                                           "805386b77df045", "8c2bbe7ca9dcbb", "queue"]), "")

    # ---- R03 layout checks ------------------------------------------------
    codes = ["UNKNOWN_APPLICABILITY", "REQUEST_DRAFTED_NOT_SENT",
             "NOT_APPLICABLE_SUPPORTED", "RECEIVED_VERIFIED_REAL",
             "MISMATCH_OR_UNVERIFIABLE", "EXPIRED_OR_OUT_OF_PERIOD",
             "EVIDENCE_MISSING", "PENDING_REAL_SOURCE", "D5 HOLD_UNKNOWN"]
    checks["codes preserved (norm)"] = (all(c in norm for c in codes), "")
    checks["no U+200B hidden char"] = (ZWSP not in all_text,
                                       "zwsp=%d" % all_text.count(ZWSP))

    # Explicit display line breaks: collected from every table-cell run; each break
    # must be preceded by '_' and the run text must concatenate back to the exact code.
    br_count = 0
    br_ok = True
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        acc = ""
                        for node in r._element:
                            tag = node.tag
                            if tag == qn("w:t"):
                                acc += node.text or ""
                            elif tag == qn("w:br"):
                                br_count += 1
                                if not acc.endswith("_"):
                                    br_ok = False
                            elif tag == qn("w:tab"):
                                acc += "\t"
    checks["breaks exist and only after '_'"] = (br_ok and br_count > 0,
                                                 "br=%d" % br_count)


    split_ok = True
    for t in doc.tables:
        for tr in t._tbl.findall(qn("w:tr")):
            trPr = tr.find(qn("w:trPr"))
            if trPr is None or trPr.find(qn("w:cantSplit")) is None:
                split_ok = False
    checks["cantSplit on all rows"] = (split_ok, "")

    hdr_ok = True
    for t in doc.tables:
        if len(t.rows) < 2:
            continue
        tr0 = t._tbl.findall(qn("w:tr"))[0]
        trPr = tr0.find(qn("w:trPr"))
        if trPr is None or trPr.find(qn("w:tblHeader")) is None:
            hdr_ok = False
    checks["repeated header rows"] = (hdr_ok, "")

    keep_ok = True
    for t in doc.tables:
        if len(t.rows) < 2:
            continue
        for cell in t.rows[0].cells:
            for p in cell.paragraphs:
                if p.paragraph_format.keep_with_next is not True:
                    keep_ok = False
    checks["header kept with data"] = (keep_ok, "")

    # T5 must be a single block (no single-record tail page) with exactly one empty row.
    t5 = None
    for t in doc.tables:
        if any(c.text.strip() == "EV-SYN-02（SAMPLE）"
               for row in t.rows for c in row.cells):
            t5 = t
            break
    if t5 is None:
        checks["T5 kept together"] = (False, "T5 not found")
        checks["T5 one empty row"] = (False, "T5 not found")
    else:
        t5_keep = True
        for row in t5.rows[:-1]:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if p.paragraph_format.keep_with_next is not True:
                        t5_keep = False
        checks["T5 kept together"] = (t5_keep, "rows=%d" % len(t5.rows))
        empty = sum(1 for row in t5.rows
                    if row.cells[0].text.strip() and
                    all(not c.text.strip() for c in row.cells[1:]))
        blank = sum(1 for row in t5.rows if not row.cells[0].text.strip())
        checks["T5 one empty row"] = (empty == 1 and blank == 0,
                                      "empty=%d blank=%d" % (empty, blank))

    import re as _re
    maxseg = 0
    for c in cell_text:
        for seg in _re.split(r"[\s/（）()，,;；]+", c.replace(ZWSP, "")):
            if "_" in seg:
                for part in seg.split("_"):
                    maxseg = max(maxseg, len(part))
    checks["max code segment<=14"] = (maxseg <= 14, "max=%d" % maxseg)

    sizes = []
    for p in doc.paragraphs:
        for r in p.runs:
            if r.font.size is not None:
                sizes.append(r.font.size.pt)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        if r.font.size is not None:
                            sizes.append(r.font.size.pt)
    checks["no run < 9.5pt"] = (all(s >= 9.5 - 1e-6 for s in sizes), "min=%.2f" % min(sizes))

    body_sizes = [r.font.size.pt for p in doc.paragraphs for r in p.runs
                  if r.font.size is not None]
    checks["body runs >= 10.5pt"] = (all(s >= 10.5 - 1e-6 for s in body_sizes),
                                     "min=%.2f" % min(body_sizes))

    ok = True
    print("---- verify 0056_R03.docx ----")
    for name, (passed, extra) in checks.items():
        print("%-34s %s %s" % (name, "OK " if passed else "FAIL", extra))
        ok = ok and passed
    print("---- verify result: %s ----" % ("PASS" if ok else "FAIL"))
    return ok, checks


def write_notes(docx_sha, checks):
    packet_sha = sha256(PACKET_PATH)
    blueprint_sha = sha256(BLUEPRINT_PATH)
    rel = lambda p: os.path.relpath(p, ROOT).replace("\\", "/")
    lines = []
    add = lines.append
    add("# BUILD_NOTES_R03 — DOCX 0056")
    add("")
    add("Build-only record. This is NOT semantic acceptance, NOT an independent review,")
    add("and NOT a statement of real-world compliance.")
    add("")
    add("## Output files")
    add("- `0056_R03.docx` (task-directory root) — human-facing Traditional-Chinese package.")
    add("- `build/build_0056_R03.py` — deterministic builder (copied from R02 builder; layout/pagination edits).")
    add("- `build/BUILD_NOTES_R03.md` — this file.")
    add("")
    add("## R01/R02 history preserved (not overwritten)")
    add("- `0056_R01.docx`; `build/build_0056_R01.py`; `build/BUILD_NOTES_R01.md`;")
    add("  `evidence/render_R01/0056_R01.pdf` + page-01..14.png; `evidence/history/R01_pre_layout_repair/`.")
    add("- `0056_R02.docx`; `build/build_0056_R02.py`; `build/BUILD_NOTES_R02.md`;")
    add("  `evidence/visual_review_R02.json`.")
    add("")
    add("## R03 repairs (from evidence/visual_review_R02.json)")
    add("- (a) Status-dictionary tail page: the state-code column was widened so every code fits on")
    add("  one line, and the dictionary table uses explicit tight cell margins (tblCellMar). This")
    add("  removes the two-row tail page without lowering the >=9.5 pt font floor.")
    add("- (b) T5 tail page: T5 now keeps all its rows together (keep-with-next on every row except")
    add("  the last) so the complete table sits on one landscape page, and it was reduced to a single")
    add("  empty operational row (one is sufficient). No single-record tail page remains.")
    add("- (c) Hidden characters: U+200B is removed entirely. Underscore-boundary wrapping is now done")
    add("  with an explicit display line break (<w:br/>) inserted ONLY at '_' boundaries and only when")
    add("  a code token is wider than its column; the token text stays byte-exact. Code columns in the")
    add("  wide annex tables were widened where needed so no break falls between arbitrary letters.")
    add("- Retained: <w:cantSplit/> on all rows, header-row keep-with-next, repeating header rows.")
    add("- Unchanged: sample facts, decision logic, official mapping (or-semantics), schemas, formula,")
    add("  evidence boundaries, all T1-T9 columns and values.")
    add("")
    add("## Source provenance (computed only for the two attached inputs)")
    add("- Packet   `inputs/0056_SOURCE_PACKET_R01.md`  SHA-256 `%s`" % packet_sha)
    add("- Blueprint `0056_UNIQUE_BLUEPRINT_R01.md`       SHA-256 `%s`" % blueprint_sha)
    add("- Output   `0056_R03.docx`                       SHA-256 `%s`" % docx_sha)
    add("")
    add("Internal identity kept OUT of the human body (sidecar-only):")
    add("- deliverable id 0056; official indicator I12-E04; official checkbox sequence 059;")
    add("  official PDF sha 805386b7…; page-4 image sha 0e00b465…; contract/genre/mission blobs.")
    add("")
    add("## Document structure")
    add("Portrait narrative + a landscape annex for wide forms. Sections:")
    add("一 30 秒用途｜二 使用時機｜三 提案角色｜四 名詞與狀態對照｜五 官方「或」語意｜")
    add("六 決策樹｜七 狀態字典｜八 作業表單 T1–T5（landscape）｜九 決策軌跡 T6｜")
    add("十 合成示範 SYN-0056-SAFE-01｜十一 合成替換登錄 T9｜十二 完整度摘要｜")
    add("十三 未寄出索取範本｜十四 替換路徑與來源驗證｜十五 例外/到期/檢索/存取/更正/版本 T7,T8｜")
    add("十六 HOLD/退件條件｜十七 跨文件邊界與負向鄰居｜十八 審查者自我測試｜十九 來源定位。")
    add("")
    add("## Tables (schema per Blueprint F)")
    add("- T1 適用性判定表 (13 cols)")
    add("- T2 資產／場域義務矩陣 (14 cols)")
    add("- T3 設備檢查與更正複核表 (15 cols)")
    add("- T4 場地保險正本索取／索引／驗證表 (16 cols)")
    add("- T5 證據狀態與到期追蹤表 (11 cols)")
    add("- T6 決策軌跡表 (8 cols)")
    add("- T7 例外／事件與矯正表 (9 cols)")
    add("- T8 版本軌跡表 (5 cols)")
    add("- T9 合成替換登錄表 (5 cols)")
    add("- plus: 文件識別, 提案角色, 名詞與狀態對照, 狀態字典, 跨文件邊界.")
    add("")
    add("## Layout")
    add("- A4; portrait by default; landscape only for the 14–16 column annex (T1–T5) to avoid clipping.")
    add("- Body >= 10.5 pt; table cells >= 9.5 pt; repeating header rows on every table.")
    add("- Rows use cantSplit (no row splits across pages); header row kept with first data row;")
    add("  T5 additionally keeps all rows together on one landscape page.")
    add("- Long codes break only at '_' boundaries via explicit <w:br/> (no hidden characters).")
    add("- T1–T5 are the wide annex tables (landscape); other tables are portrait.")
    add("- Header (title) and footer (PAGE / NUMPAGES) on every section; Microsoft JhengHei.")
    add("")
    add("## Synthetic-data boundary")
    add("- Worked case id: `SYN-0056-SAFE-01` (SAMPLE/SYNTHETIC; unrelated to any real course).")
    add("- All sample facts are tagged (SAMPLE) in-table and registered row-by-row in T9 (SF-01..SF-19).")
    add("- Damaged cable is a synthetic illustration only; no inspection is implied to have happened.")
    add("- Insurance request is a template, status `REQUEST_DRAFTED_NOT_SENT`, no policy attached;")
    add("  insurance type/coverage/insurer/policy number stay UNKNOWN (type is asked of the source;")
    add("  no specific line such as public liability is presumed).")
    add("- Roles are `PROPOSED`; every real-world fact absent from the packet stays")
    add("  `UNKNOWN / PENDING_REAL_SOURCE`.")
    add("")
    add("## Known limitations")
    add("- No real inspection certificate, policy, issuer, policy number, coverage, limit, site,")
    add("  person, signature, sent request or approval exists in the source packet; none is invented.")
    add("- Completeness is `0 ÷ 2 = 0%` for the sample only; it is NOT a TTQS score and NOT policy.")
    add("- Result is `HOLD / PENDING_REAL_SOURCE`; `PASS` is deliberately not used.")
    add("- DOCX opens/parses with python-docx; visual pagination was not rendered here.")
    add("")
    add("## Determinism")
    add("- No wall-clock values; fixed core properties (2026-10-17); zip entries rewritten with a")
    add("  fixed timestamp, so repeated runs produce identical bytes for identical inputs.")
    add("")
    add("## Build self-check (content-level, not semantic acceptance)")
    allok = True
    for name, (passed, extra) in checks.items():
        allok = allok and passed
        add("- %-34s %s %s" % (name, "OK" if passed else "FAIL", extra))
    add("")
    add("## Truthful build result")
    add("- BUILD: %s" % ("completed; DOCX written and re-opened successfully."
                          if allok else "completed WITH FAILED SELF-CHECKS (see above)."))
    add("- REVIEW: not performed here (an independent reviewer judges quality separately).")
    with open(NOTES_PATH, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(lines) + "\n")
    return allok


def main():
    doc = build_document()
    doc.save(DOCX_PATH)
    normalize_zip(DOCX_PATH)
    ok, checks = verify(DOCX_PATH)
    write_notes(sha256(DOCX_PATH), checks)
    print("wrote:", DOCX_PATH)
    print("wrote:", NOTES_PATH)
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
