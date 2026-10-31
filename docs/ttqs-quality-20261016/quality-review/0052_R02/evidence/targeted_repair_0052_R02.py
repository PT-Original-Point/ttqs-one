from pathlib import Path
from docx import Document
from docx.shared import RGBColor
import hashlib

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / '0052_R01' / 'candidate' / '0052_R01.docx'
OUT = ROOT / '0052_R02' / 'candidate' / '0052_R02.docx'
EXPECTED_SOURCE_SHA = '406d1ef2c6e9b6266ea1adaeadedc5dfac77964fe76f73c2471168220ee6d455'

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def replace_paragraph(paragraph, text):
    runs = paragraph.runs
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)

def replace_cell(doc, table, row, col, text):
    cell = doc.tables[table].cell(row, col)
    if len(cell.paragraphs) != 1:
        raise RuntimeError(f'Unexpected paragraph count in table {table}, row {row}, col {col}')
    replace_paragraph(cell.paragraphs[0], text)

if sha(SRC) != EXPECTED_SOURCE_SHA:
    raise RuntimeError(f'R01 source SHA mismatch: {sha(SRC)}')
if OUT.exists() and sha(OUT) != 'c5cb59eaf0cbf850808c6bb387a786f164a124de0fbfdb7823af92454ba33091':
    raise RuntimeError(f'Refusing to overwrite unexpected existing output {OUT}: {sha(OUT)}')
OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_bytes(SRC.read_bytes())
doc = Document(str(OUT))

# Replace only engineering identifiers and source-control prose with reader-facing scope and use guidance.
replace_paragraph(doc.paragraphs[2], '範例涵蓋意見蒐集、資料檢核、方案取捨、執行責任與下一輪驗證。')
replace_paragraph(doc.paragraphs[8], '本件說明課程計畫如何吸收在訓及結訓學員意見，經資料檢核、分析取捨、授權及下一輪驗證形成可追蹤修訂。它不取代訓練需求調查、課程需求與排課計畫、單元教案與評量設計、正式會議決議或課程適切性檢討；各項原始紀錄應分別保存。')
replace_paragraph(doc.paragraphs[11], '本計畫依課程意見回流的官方要求編製，示範如何將在訓學員及結訓學員意見連結至資料檢核、分析取捨、課程修訂、授權確認及下一輪驗證；本文件不取代需求調查、課程計畫、單元教案、正式會議決議或課程適切性檢討的原始紀錄。')

# Table 1: remove the internal revision code from the reader-facing status text.
replace_cell(doc, 1, 1, 2, '本表是完整的合成示範；實際核定權限與課程變更決定，須由協會依現行授權程序確認。')

# Table 2: keep the human-readable boundaries while removing DOCX / indicator IDs.
replace_cell(doc, 2, 1, 0, '課程意見回流與修訂計畫')
replace_cell(doc, 2, 1, 1, '訓練需求調查；課程需求與排課計畫；單元設計與評量')
replace_cell(doc, 2, 2, 0, '課程產出討論與確認紀錄')
replace_cell(doc, 2, 3, 0, '課程適切性檢討與調整紀錄')
replace_cell(doc, 2, 3, 2, '課程適切性檢討依其專屬設計要求處理；本件只說明學員與結訓學員意見如何回饋課程規劃。')

# Table 3: retain the sample boundary and provenance purpose without hashes, paths, or filenames.
for col, value in enumerate(('資料類別', '本文件內容', '使用限制')):
    replace_cell(doc, 3, 0, col, value)
rows = {
    1: ('官方要求', '依「規劃課程依學員及結訓學員意見訂定訓練課程，且流程有完整紀錄」編製。', '正式使用前，請按當期適用的官方資料表確認要求。'),
    2: ('示範資料', '本文日期、角色、課程、回覆、追蹤及計算均為 SAMPLE / SYNTHETIC。', '不得作為協會已辦訓、已調查、已追蹤或已核准的證明。'),
    3: ('同一示範案例', '沿用同一示範課程及四個單元的規劃假設，保持學員意見與課程修訂前後一致。', '課程及 360 分鐘均為合成規劃，不是實際排課或課程紀錄。'),
    4: ('操作原則', '以課後意見、結訓追蹤、方案比較、修訂責任和成果驗證形成完整流程。', '原始問卷、分析、核定與成果紀錄應分別保存。'),
    5: ('來源與版本', '官方要求、範例來源及各項假設均列於配套追溯資料。', '追溯資料供查核，不可替代協會的實際原件。'),
}
for row, values in rows.items():
    for col, value in enumerate(values):
        replace_cell(doc, 3, row, col, value)

# Remove document-number pointers from the instructional examples while preserving their operational meaning.
replace_cell(doc, 16, 1, 2, '下一期沿用本問卷的服務交接題目；以五題至少四題達標且沒有越權承諾作為示範門檻，並逐次記錄有效分母與結果。')
replace_cell(doc, 19, 1, 2, '最高分，但與方案 C 僅差 0.05，屬示範評分的小幅差異；35/20/35 為本件合成比較基準，不代表協會既有紀錄。')
replace_cell(doc, 21, 0, 1, '維持的既有目標（同一示範課程）')
replace_cell(doc, 22, 4, 1, '同一套示範量規逐項達成數 ÷ 有效受測數')
replace_cell(doc, 22, 4, 3, '沿用同一案例的示範量尺；正式執行時另存實際版本與去識別評量原始紀錄。')
replace_cell(doc, 24, 1, 1, '同一示範案例的需求調查資料；案例代碼 SYN-0033-PLAN-01。')
replace_cell(doc, 24, 1, 2, '本文件僅承接同一示範案例的合成需求背景，不重貼需求調查表，也不代表已完成真人調查。')
replace_cell(doc, 24, 2, 1, '同一示範案例的課程需求及計畫形成日。')
replace_cell(doc, 24, 2, 2, '沿用同一案例的課程名稱、12 人規劃假設、四個單元及 360 分鐘示範時數；日期為合成示例，不是協會實際排課。')
replace_cell(doc, 26, 5, 1, '本節及 §1；依當期官方資料表的課程規劃要求核對。')
replace_cell(doc, 26, 5, 2, '以課程規劃及學員意見回流的官方要求為準；本表協助核對正文是否完整，不取代官方原件。')
replace_cell(doc, 29, 5, 1, '沿用本文件所列回饋題、結訓追蹤題及同一示範量尺；期限須於正式排課後確認。')

# Replace internal status jargon with plain instructions for association staff.
replace_cell(doc, 4, 0, 0, '真實資料查核結果：截至本示範編製時，尚未取得本要求專屬且可驗證的協會學員意見、結訓追蹤、核定修訂、學員名冊或授權紀錄。因此本文問卷回覆、日期、角色、分母、分析與修訂均為合成示例；協會實際權責、保存期限、課程實施及核定情形尚待核實。')
role_texts = {
    (1, 1): '課程規劃承辦（示範角色；實際職務與人選由協會確認）',
    (1, 3): '待協會確認；此為建議功能，不代表現行組織職稱或已指派人員。',
    (2, 1): '資料與隱私檢核者（示範角色；任命與權限由協會確認）',
    (2, 3): '待協會確認；此為建議功能，不代表現行組織職稱或已指派人員。',
    (3, 1): '獨立覆核者（示範角色；實際指派由協會確認）',
    (3, 3): '待協會確認；此為建議功能，不代表現行組織職稱或已指派人員。',
    (4, 1): '有權核定者（須依章程或有效授權文件確認；目前尚未核實）',
    (4, 3): '待協會確認；不得視為現行組織職稱或已指派人員。',
    (5, 1): '授課與演練引導者（示範角色；實際師資由協會確認）',
    (5, 3): '待協會確認；此為建議功能，不代表現行組織職稱或已指派人員。',
}
for (row, col), value in role_texts.items():
    replace_cell(doc, 6, row, col, value)
replace_cell(doc, 21, 2, 3, '未核實的事實不得以推測補填；應標示待查，並記錄查核來源與責任人。')
replace_cell(doc, 23, 5, 3, '如需向正式治理來源確認權限，應記錄查核依據、日期及後續處置。')
replace_cell(doc, 24, 7, 0, '下一期（日期由協會確認）')

# Rewrite the last section as a practical user checklist, not an engineering receipt.
replace_paragraph(doc.paragraphs[75], '15｜正式使用前的資料與權限確認')
replace_cell(doc, 30, 0, 0, '使用前項目')
replace_cell(doc, 30, 0, 1, '確認事項')
last_rows = {
    1: ('本文件用途', '課後意見及結訓追蹤的課程修訂示範；實際使用前應替換全部合成資料。'),
    2: ('官方要求', '依學員與結訓學員意見規劃或修訂訓練課程，並保留完整作業紀錄。'),
    3: ('範例資料', '日期、樣本、課程、角色、分數及決定均為 SAMPLE / SYNTHETIC。'),
    4: ('真實資料', '使用前應取得可核驗的問卷、追蹤、辦訓與課程修訂資料；本示範未附這些原件。'),
    5: ('授權與核定', '由協會確認資料使用權限並依內部程序審核；本示範不構成協會核准。'),
    6: ('來源紀錄', '官方要求、範例來源與版本核對資料另行保存，供主管或評核者查核。'),
}
for row, values in last_rows.items():
    for col, value in enumerate(values):
        replace_cell(doc, 30, row, col, value)

# Preserve white text on the dark header rows we changed.
for table_index, columns in ((3, range(3)), (21, (1,)), (30, range(2))):
    for column in columns:
        cell = doc.tables[table_index].cell(0, column)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)

# Keep the page-number field intact while changing repeated header/footer labels.
for section in doc.sections:
    if section.header.paragraphs:
        replace_paragraph(section.header.paragraphs[0], '原始點協會｜課程意見回流與課程規劃')
    if section.footer.paragraphs and section.footer.paragraphs[0].runs:
        section.footer.paragraphs[0].runs[0].text = 'SAMPLE / SYNTHETIC 示範｜未經協會核定｜'

doc.save(str(OUT))

# Validate there are no internal doc IDs, blueprint names, source paths, raw hashes, or build commit IDs in body/header/footer.
doc = Document(str(OUT))
texts = [p.text for p in doc.paragraphs]
for table in doc.tables:
    texts.extend(cell.text for row in table.rows for cell in row.cells)
for section in doc.sections:
    texts.extend(p.text for p in section.header.paragraphs)
    texts.extend(p.text for p in section.footer.paragraphs)
body = '\n'.join(texts)
for forbidden in ('DOCX', 'I11-E04', 'Unique Blueprint', 'Blueprint', 'SHA-256', 'BLUEPRINT_CURRENT', 'checkbox', 'crosswalk', 'SOURCE_MAP', 'bdfa5a675ca77ff460530b96971b26ff44c455c4', 'fe5f3cfd84c761da86435a71e0fd899daa725be78628c4d2d3f2324c1fdb5f4f', 'e8212788b1f7b201ca4af4f4c5fb83ab131627dbfa815987f7cf6344edfda727', 'd0d37c2c69b161a2cdd37ee3097306722c4a65a1afd844ae92e1dff58b43381a', '0052 R01', '0049 R01', '0050 R01', '0033 R03', '0051 R01', '0053 R01', 'UNKNOWN'):
    if forbidden.casefold() in body.casefold():
        raise RuntimeError(f'Forbidden internal string remains in reader-facing DOCX: {forbidden}')
print(f'OUTPUT={OUT}')
print(f'BYTES={OUT.stat().st_size}')
print(f'SHA256={sha(OUT)}')
print(f'PARAGRAPHS={len(doc.paragraphs)} TABLES={len(doc.tables)}')