from pathlib import Path
import hashlib
from docx import Document

root = Path(__file__).resolve().parent
source = root.parent / "0136_I19-E26_SOURCE_PREP_R01_20261025" / "source"
docx_path = root / "0136_I19-E26_R01_CANDIDATE_20261025.docx"
out = root / "REVIEW_INPUT_0136_R01_EXACT_SHA.md"
sha = hashlib.sha256(docx_path.read_bytes()).hexdigest()
doc = Document(docx_path)
parts = []
for p in doc.paragraphs:
    if p.text.strip():
        parts.append(p.text)
for i, table in enumerate(doc.tables, 1):
    parts.append(f"\n[TABLE {i}]")
    for row in table.rows:
        cells = [c.text.replace("\n", " / ").strip() for c in row.cells]
        parts.append(" | ".join(cells))

sections = [
    "# Independent semantic review input",
    "",
    "This packet is read-only review material. Do not edit files or claim association approval.",
    f"Exact target DOCX SHA-256: {sha}",
    f"DOCX bytes: {docx_path.stat().st_size}",
    "Review scope: semantic content only. Visual QA is recorded separately against this exact SHA.",
    "",
    "## Official source excerpt (source packet)",
    (source / "OFFICIAL_SOURCE_P8_P9_TEXT.txt").read_text(encoding="utf-8"),
    "",
    "## Current source packet",
    (source / "SOURCE_PACKET_0136_I19-E26_R01.md").read_text(encoding="utf-8"),
    "",
    "## Unique Blueprint",
    (source / "UNIQUE_BLUEPRINT_0136_I19-E26_R01.md").read_text(encoding="utf-8"),
    "",
    "## Project index and neighboring requirements",
    (source / "PROJECT_157_INDEX150_AND_NEIGHBORS.md").read_text(encoding="utf-8"),
    "",
    "## Candidate text extracted from the exact target DOCX",
    "\n".join(parts),
]
out.write_text("\n\n".join(sections), encoding="utf-8")
print(f"output={out}")
print(f"target_docx_sha256={sha}")
print(f"review_input_sha256={hashlib.sha256(out.read_bytes()).hexdigest()}")
print(f"review_input_bytes={out.stat().st_size}")
