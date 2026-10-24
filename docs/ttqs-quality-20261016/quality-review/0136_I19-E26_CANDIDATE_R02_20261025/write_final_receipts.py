from pathlib import Path
from datetime import datetime, timezone
import hashlib, json, zipfile
from docx import Document

root = Path(__file__).resolve().parent
docx = root / "0136_I19-E26_R02_CANDIDATE_20261025.docx"
render = root / "render_r04"
sha = hashlib.sha256(docx.read_bytes()).hexdigest()
with zipfile.ZipFile(docx) as zf:
    zip_test = zf.testzip()
d = Document(docx)
table_sizes, body_sizes = [], []
for table in d.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    if run.text.strip():
                        table_sizes.append(run.font.size.pt if run.font.size else None)
for p in d.paragraphs:
    for run in p.runs:
        if run.text.strip():
            body_sizes.append(run.font.size.pt if run.font.size else None)
pages = sorted(render.glob("page-*.png"))
pdf = render / "0136_R02.pdf"
assert len(pages) == 14, f"expected 14 rendered pages, got {len(pages)}"
assert zip_test is None, f"zip test failed: {zip_test}"
assert min(x for x in table_sizes if x is not None) >= 9.5
assert all(x is not None for x in table_sizes)
assert min(x for x in body_sizes if x is not None) >= 10.5
assert all(x is not None for x in body_sizes)
layout = {
    "project_id": "TTQS_ONE",
    "doc_id": "0136",
    "official_requirement_id": "I19-E26",
    "project_157_index": 150,
    "official_printed_checkbox_serial": "UNKNOWN",
    "candidate_revision": "R02_CANDIDATE_20261025",
    "docx_filename": docx.name,
    "docx_sha256": sha,
    "docx_bytes": docx.stat().st_size,
    "zip_integrity": "PASS",
    "paragraph_count": len(d.paragraphs),
    "table_count": len(d.tables),
    "table_text_runs": len(table_sizes),
    "minimum_explicit_table_font_pt": min(table_sizes),
    "body_text_runs": len(body_sizes),
    "minimum_explicit_body_font_pt": min(body_sizes),
    "renderer": "Microsoft Word COM, version 16.0, hidden read-only export",
    "rasterizer": "Poppler pdftoppm 110 dpi",
    "pdf_filename": pdf.name,
    "pdf_sha256": hashlib.sha256(pdf.read_bytes()).hexdigest(),
    "pdf_bytes": pdf.stat().st_size,
    "page_count": len(pages),
    "visually_inspected_pages_at_original_resolution": [p.name for p in pages],
    "visual_review": "PASS",
    "visual_findings": [
        "R01 predecessor SHA 21e208652930c491e2a93c07e829b101a7fb34b77151caa4ef1784dbbe3697e4 had an empty rendered page 11.",
        "R02 changed the section 10 heading to paragraph-level page_break_before; no content was removed.",
        "R02 renders to 14 pages with no blank page; all tables, margins, headings, and footer numbering were visually checked."
    ],
    "page_images": [
        {"filename": p.name, "sha256": hashlib.sha256(p.read_bytes()).hexdigest(), "bytes": p.stat().st_size}
        for p in pages
    ],
    "examiner": "Codex visual inspection; not an independent semantic reviewer",
    "independent_semantic_review": "UNKNOWN_BLOCKED_PROVIDER_FREE_LIMIT",
    "formal_acceptance": "NOT_GRANTED",
    "observed_at_utc": datetime.now(timezone.utc).isoformat()
}
(root / "LAYOUT_QA_RECEIPT_0136_R02.json").write_text(json.dumps(layout, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
state = {
    "project_id": "TTQS_ONE",
    "doc_id": "0136",
    "official_requirement_id": "I19-E26",
    "project_157_index": 150,
    "official_printed_checkbox_serial": "UNKNOWN",
    "candidate_revision": "R02_CANDIDATE_20261025",
    "candidate_sha256": sha,
    "candidate_state": "CANDIDATE_WAITING_INDEPENDENT_SEMANTIC_REVIEW",
    "technical_integrity": "PASS",
    "full_page_visual_review": "PASS_14_OF_14",
    "semantic_review": "UNKNOWN_PROVIDER_FREE_LIMIT_429",
    "real_evidence": "UNKNOWN_NO_REAL_SOURCE_ATTACHED",
    "formal_acceptance": "NOT_GRANTED",
    "antigravity_output_adopted": False,
    "opencode_writer_output_adopted": False,
    "next_gate": "Independent semantic review bound to this exact SHA; preserve the DOCX unless a concrete semantic finding requires a change."
}
(root / "CANDIDATE_STATE_0136_R02.json").write_text(json.dumps(state, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps({"docx_sha256": sha, "pages": len(pages), "pdf_sha256": layout["pdf_sha256"], "layout": layout["visual_review"], "semantic": state["semantic_review"], "table_font_min_pt": min(table_sizes), "body_font_min_pt": min(body_sizes)}, ensure_ascii=False))
