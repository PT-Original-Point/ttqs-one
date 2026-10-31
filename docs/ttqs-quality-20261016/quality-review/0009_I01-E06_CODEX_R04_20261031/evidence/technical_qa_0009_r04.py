from pathlib import Path
from datetime import datetime
import csv,hashlib,json,re,zipfile
from docx import Document
ROOT=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0009_I01-E06_LUNA_TEXT_DRAFT_R01_20261029\0009_I01-E06_R04_SCOPED_SEMANTIC_REPAIR_20261031')
DOC=ROOT/'candidate'/'0009_I01-E06_R04.docx'; MD=ROOT/'source'/'MANUSCRIPT_0009_I01-E06_R04.md'; CSV=ROOT/'source'/'SAMPLE_DATASET_0009_R02.csv'; OUT=ROOT/'evidence'/'TECHNICAL_QA_0009_R04.json'
def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest().upper()
checks={}
with zipfile.ZipFile(DOC) as z: checks['docx_zip_integrity']=z.testzip() is None
D=Document(DOC); checks['python_docx_open']=True
# Verify the added content blocks appear in proper document order, guarding against orphan/reversed insertion.
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
sequence=[]
for el in D.element.body.iterchildren():
    if isinstance(el,CT_P): sequence.append(('p',Paragraph(el,D).text.strip()))
    elif isinstance(el,CT_Tbl):
        tt=Table(el,D); sequence.append(('t',tt.rows[0].cells[0].text.strip()))
def seqpos(kind,prefix): return next((i for i,(k,v) in enumerate(sequence) if k==kind and v.startswith(prefix)), -1)
checks['target_section_table_order']=0<=seqpos('p','2.2')<seqpos('t','目標ID')<seqpos('p','3. 計算公式')
checks['threshold_section_table_order']=0<=seqpos('p','3.3')<seqpos('t','行動ID')<seqpos('p','4. 2025')
checks['register_section_table_order']=0<=seqpos('p','9.1')<seqpos('t','替換鍵')<seqpos('p','10. 官方')
with CSV.open(encoding='utf-8-sig',newline='') as f: data=list(csv.DictReader(f))
checks['source_dataset_rows_119']=len(data)==119
checks['docx_tables_23']=len(D.tables)==23
register=[t for t in D.tables if t.rows and t.rows[0].cells[0].text.strip()=='替換鍵']
checks['record_level_register_table_unique']=len(register)==1
checks['record_level_register_rows_119']=len(register)==1 and len(register[0].rows)==120
if len(register)==1:
 rows=register[0].rows[1:]
 ids=[r.cells[0].text.strip() for r in rows]
 srcids=[r['item_id'] for r in data]
 checks['unique_register_keys']=ids==[f'S{i:04d}' for i in range(1,120)]
 checks['item_ids_in_source_order']=all(srcids[i] in rows[i].cells[1].text for i in range(119))
 checks['each_csv_line_locator']=all(f'CSV第{i+2}資料列' in rows[i].cells[1].text for i in range(119))
 checks['all_data_status_synthetic']=all('SAMPLE_SYNTHETIC' in r.cells[4].text for r in rows)
checks['target_setting_rows_4']=any(t.rows and t.rows[0].cells[0].text.strip()=='目標ID' and len(t.rows)==5 for t in D.tables)
checks['q1_threshold_rows_4']=any(t.rows and t.rows[0].cells[0].text.strip()=='行動ID' and len(t.rows)==5 for t in D.tables)
# Recompute annual and Q1 counts from the source dataset.
groups={x['dataset']:[x for x in data if x['dataset']==x['dataset']] for x in []}
from collections import defaultdict
by=defaultdict(list)
for r in data: by[r['dataset']].append(r)
def total(k): return sum(int(r['numerator_flag']) for r in by[k]),sum(int(r['denominator_flag']) for r in by[k])
checks['annual_D1_5_of_6']=total('D1_annual_activity')==(5,6)
checks['annual_D2_45_of_50']=total('D2_annual_admin')==(45,50)
checks['annual_D3_17_of_20']=total('D3_valid_request')==(17,20)
checks['annual_D4_3_of_4']=total('D4_quarterly_review')==(3,4)
for aid,n,d,threshold in [('A01',3,3,1.0),('A02',24,25,.95),('A03',9,10,.90),('A04',1,1,1.0)]:
 rr=[r for r in by['Q1_action_followup'] if r['item_id'].startswith(aid+'-')]
 actual=(sum(int(r['numerator_flag']) for r in rr),sum(int(r['denominator_flag']) for r in rr))
 checks[f'q1_{aid}_count_threshold']=actual==(n,d) and n/d>=threshold
action_table=next((t for t in D.tables if t.rows and t.rows[0].cells[0].text.strip()=='行動 ID'),None)
checks['all_action_rows_reference_preset_thresholds']=bool(action_table) and all('§3.3' in action_table.rows[i].cells[4].text for i in range(1,5))
text=MD.read_text(encoding='utf-8')
for token,key in [('2024-12-15','target_setting_date_present'),('2025-12-15','q1_threshold_date_present'),('§2.2','target_setting_cross_reference'),('§3.3','threshold_cross_reference'),('§9.1','row_register_reference')]: checks[key]=token in text
checks['r03_partial_review_carried']= (ROOT/'review'/'DARWIN_SEMANTIC_REVIEW_0009_R03_PARTIAL.md').exists() and 'PARTIAL' in (ROOT/'review'/'DARWIN_SEMANTIC_REVIEW_0009_R03_PARTIAL.md').read_text(encoding='utf-8')
min_body=1000; min_table=1000
for p in D.paragraphs:
 for r in p.runs:
  if r.text.strip() and r.font.size: min_body=min(min_body,r.font.size.pt)
for t in D.tables:
 for row in t.rows:
  for cell in row.cells:
   for p in cell.paragraphs:
    for r in p.runs:
     if r.text.strip() and r.font.size: min_table=min(min_table,r.font.size.pt)
checks['body_min_pt_ge_10_5']=min_body>=10.5
checks['tables_min_pt_ge_9_5']=min_table>=9.5
result={'doc_id':'0009','official_requirement_id':'I01-E06','revision':'R04','docx_sha256':sha(DOC),'manuscript_sha256':sha(MD),'dataset_sha256':sha(CSV),'checks':checks,'technical_state':'PASS' if all(checks.values()) else 'FAIL','visual_state':'PENDING','semantic_state':'PENDING','formal_acceptance':'NOT_GRANTED','real_association_data':'UNKNOWN','min_body_font_pt':min_body,'min_table_font_pt':min_table,'table_count':len(D.tables),'dataset_rows':len(data),'register_rows':len(register[0].rows)-1 if register else 0,'checked_at_local':datetime.now().astimezone().isoformat(timespec='seconds')}
OUT.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'technical_state':result['technical_state'],'sha256':result['docx_sha256'],'checks_passed':sum(checks.values()),'checks_total':len(checks),'failed':[k for k,v in checks.items() if not v],'register_rows':result['register_rows'],'min_body_pt':min_body,'min_table_pt':min_table},ensure_ascii=False))
if not all(checks.values()): raise SystemExit(2)
