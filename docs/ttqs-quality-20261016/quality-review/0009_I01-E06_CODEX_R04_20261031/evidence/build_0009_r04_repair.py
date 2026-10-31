from pathlib import Path
from datetime import datetime
import csv, hashlib, json, re, shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0009_I01-E06_LUNA_TEXT_DRAFT_R01_20261029')
OUT=BASE/'0009_I01-E06_R04_SCOPED_SEMANTIC_REPAIR_20261031'
R03=BASE/'candidate'/'0009_I01-E06_R03.docx'
R03MD=BASE/'source'/'MANUSCRIPT_0009_I01-E06_R03_CODEX_INTEGRATION.md'
CSV=OUT/'source'/'SAMPLE_DATASET_0009_R02.csv'
DOCX=OUT/'candidate'/'0009_I01-E06_R04.docx'
MD=OUT/'source'/'MANUSCRIPT_0009_I01-E06_R04.md'

def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest().upper()
def req_replace(text,old,new,label):
    n=text.count(old)
    if n!=1: raise ValueError(f'{label}: expected 1 occurrence, got {n}: {old[:80]}')
    return text.replace(old,new,1)
def md_cell(x): return str(x).replace('|','／').replace('\n',' ')

# 1) Build exact manuscript delta from the reviewed R03 manuscript.
md=R03MD.read_text(encoding='utf-8-sig')
md=req_replace(md,'**文件編號：0009｜對應指標：I01-E06｜文類：年度績效目標分析紀錄**','**文件編號：0009｜對應指標：I01-E06｜文類：年度績效目標分析紀錄｜候選修訂：R04**','revision marker')
targets=[
 ['G01','2024-12-15','2024 基線 5／6＝83.3%；2025 目標 6／6＝100.0%，比基線增加 1 場、16.7 個百分點。六場均來自本例事前年度清單；選全數完成是避免已排定工作被事後剔除。','SAMPLE假設年度清單固定 6 場；角色 A 在各場前 30 日確認主場地與備援，角色 B 每季抽核。這是可行性假設，不是協會實際人力或場地證據。','此日期、清單、基線和能力均為 SYNTHETIC；無真實核定者。REAL須附年度目標版本、採定日期、權責來源與六場完整計畫。'],
 ['G02','2024-12-15','2024 基線 42／50＝84.0%；2025 目標 48／50＝96.0%，要求比前期多 6 項按期完成。選值保留 2 項可供例外分析，不把 100% 當作無證據的承諾。','SAMPLE假設角色 A 維護 50 項母體、每月一次檢視責任人／原期限／交付物，並為資料延遲與責任不清設升級提醒；不代表真實工時或資源已核准。','此設定理由與資源是假設；REAL須附前期同口徑基線、年度責任清冊、目標設定版本及授權紀錄。'],
 ['G03','2024-12-15','2024 基線 15／20＝75.0%；2025 目標 18／20＝90.0%，要求比前期增加 3 件於五個工作日內首次實質回覆。五日是本例服務時限，不是官方門檻。','SAMPLE假設角色 A 每週一次檢視新件、登錄收件及首次實質回覆時間，待內部資料者保留時計並升級；以 20 件年母體示範可追蹤流程，不宣稱實際承載量。','設定日期、母體、服務時限與資源假設均為 SYNTHETIC；REAL須附去識別化收件／回覆時間、有效件與去重規則及權責來源。'],
 ['G04','2024-12-15','2024 基線 3／4＝75.0%；2025 目標 4／4＝100.0%，較基線多完成一次完整季度檢討。選值依預先定義四季母體，完整紀錄須含目標、差異、原因、責任與追蹤。','SAMPLE假設每季預留一個檢討時段及一個備援時段；角色 B 彙總並追蹤缺席資料。資源僅示範安排方式，無實際核准或會議紀錄。','此日期與排程均為 SYNTHETIC；REAL須附年度行事曆、目標核定版本及四季逐次檢討／追蹤紀錄。'],
]
target_block='''### 2.2 2025年度目標事前設定示範（SAMPLE）

下表補足「如何訂出目標」的實際示範。設定日期均假設為 2024-12-15，早於 2025 衡量期間；基線與能力條件也都是合成假設，不是協會實際目標、資源承諾或核准。正式採用時須以真正的事前版本及權責紀錄替換。

| 目標 ID | SAMPLE 設定日 | 基線、目標及選值理由 | 資源／能力假設 | 證據界線與 REAL 替換 |
|---|---|---|---|---|
'''
for row in targets: target_block+='| '+' | '.join(md_cell(x) for x in row)+' |\n'
target_block+='\n'
md=req_replace(md,'## 3. 計算公式與判定門檻',target_block+'## 3. 計算公式與判定門檻','annual target table insertion')
q1=[
 ['A01','100%（3/3應完成）','2025-12-15，2026 Q1開始前；要求所有預定活動均完成30日前場地／備援核對，否則不能稱流程執行完整。','3／3＝100.0%','100.0% ≥ 100.0%，符合。','分母缺列或活動遭事後移除即HOLD；SAMPLE，不是協會核准門檻。'],
 ['A02','≥95%','2025-12-15，2026 Q1開始前；允許最多1/25例外以顯示可追蹤容差，並保留逾期原期限。','24／25＝96.0%','96.0% ≥ 95.0%，符合。','未達時逐項列明逾期與交付物；分母不完整即HOLD。'],
 ['A03','≥90%（五工作日內）','2025-12-15，2026 Q1開始前；以本例年度服務目標90%作短期檢核底線，保留單件時計與逾期原因。','9／10＝90.0%','90.0% ≥ 90.0%，符合。','工作日須按真實日曆重算；此門檻非官方要求。'],
 ['A04','100%（當季到期檢討）','2025-12-15，2026 Q1開始前；所有當季到期的檢討均須有完整目標、差異、原因、責任及追蹤。','1／1＝100.0%','100.0% ≥ 100.0%，符合。','1/1僅代表Q1到期母體，不外推全年四季；缺完整內容即不計分子。'],
]
q1_block='''### 3.3 2026 Q1 行動檢查門檻（事前 SAMPLE 設定）

本例假設於 2025-12-15、Q1開始前設定下列門檻，故 §5 表 5-1 的「符合」可據此重算。這些是示範內部管理門檻，不是 TTQS 官方評分線或原始點協會核准標準。母體或來源不完整時一律 HOLD，不以抽樣率代替完整母體。

| 行動 ID | 事前合格門檻 | 設定時點與理由 | Q1 SAMPLE實績 | 判定重算 | HOLD／REAL注意事項 |
|---|---|---|---|---|---|
'''
for row in q1:q1_block+='| '+' | '.join(md_cell(x) for x in row)+' |\n'
q1_block+='\n'
md=req_replace(md,'## 4. 2025 年度 SAMPLE 個案與完整計算',q1_block+'## 4. 2025 年度 SAMPLE 個案與完整計算','Q1 threshold insertion')
for old,new in [
 ('3／3＝100.0%；達到本例行動檢查目標','3／3＝100.0%；符合 §3.3 表3-3 事前 100% 門檻'),
 ('24／25＝96.0%；達到本例行動檢查目標','24／25＝96.0%；符合 §3.3 表3-3 事前 ≥95% 門檻'),
 ('9／10＝90.0%；達到本例行動檢查目標','9／10＝90.0%；符合 §3.3 表3-3 事前 ≥90% 門檻'),
 ('Q1 為 1／1＝100.0%；僅代表第一季到期檢討完成','Q1 為 1／1＝100.0%；符合 §3.3 表3-3 事前 100% 門檻；僅代表第一季到期檢討完成')]: md=req_replace(md,old,new,'action verdict cross-reference')
# Mark the five old aggregate appendix-register facts as summary indexes and point to new row-level mappings.
for fid in ('F0016','F0017','F0018','F0019','F0020'):
    matches=[line for line in md.splitlines() if line.startswith('| '+fid+' |')]
    if len(matches)!=1: raise ValueError(f'aggregate row not unique: {fid}')
    line=matches[0]; cells=line.split('|'); cells[2]=cells[2].rstrip()+'；逐筆對應見 §9.1 替換鍵 S0001–S0119'
    md=md.replace(line,'|'.join(cells),1)
new_facts=[
 ['F0022','§2.2／表2-2／G01：SAMPLE設定日2024-12-15；基線5／6＝83.3%；目標6／6＝100.0%；選值理由為事前列入六場均需完成；容量假設為逐場30日前核對場地與備援','示範從基線、既定母體及可執行控制推導目標；不聲稱核准','REAL年度目標版本、事前採定日期、六場原始清冊、權責依據及場地／備援資源紀錄','SYNTHETIC'],
 ['F0023','§2.2／表2-2／G02：SAMPLE設定日2024-12-15；基線42／50＝84.0%；目標48／50＝96.0%；要求增加6項按期；月檢50項責任／期限／交付物是假設控制','呈現目標增幅及月度管控能力假設','REAL前期同口徑基線、年度責任清冊、目標版本、月檢人力／權責來源','SYNTHETIC'],
 ['F0024','§2.2／表2-2／G03：SAMPLE設定日2024-12-15；基線15／20＝75.0%；目標18／20＝90.0%；增加3件五工作日內回覆；每週分派檢視為假設資源','示範服務時限與回覆能力假設的關係','REAL去識別化收件／回覆紀錄、有效件及去重規則、工作日曆、服務目標採定紀錄','SYNTHETIC'],
 ['F0025','§2.2／表2-2／G04：SAMPLE設定日2024-12-15；基線3／4＝75.0%；目標4／4＝100.0%；四季檢討時段及備援時段為假設','展示完整年度四季母體與事前排程','REAL年度行事曆、目標版本、四季完整檢討及責任追蹤紀錄','SYNTHETIC'],
 ['F0026','§3.3／表3-3／A01：SAMPLE設定2025-12-15；Q1標準100%；3／3＝100%','所有預定活動均應於30日前完成場地／備援核對','REAL Q1完整活動母體、核對日期及來源紀錄、門檻事前採定依據','SYNTHETIC'],
 ['F0027','§3.3／表3-3／A02：SAMPLE設定2025-12-15；Q1標準≥95%；24／25＝96%','事前容許最多一項例外並保留完整母體','REAL Q1行政事項母體、原期限／完成證明與門檻採定紀錄','SYNTHETIC'],
 ['F0028','§3.3／表3-3／A03：SAMPLE設定2025-12-15；Q1標準≥90%；9／10＝90%；五工作日時計','以明確等於線作為合格判定並保留逐件時戳','REAL有效需求完整清冊、真實工作日曆、分派／回覆時戳及門檻來源','SYNTHETIC'],
 ['F0029','§3.3／表3-3／A04：SAMPLE設定2025-12-15；Q1到期檢討門檻100%；1／1＝100%','只檢查Q1到期母體且不外推全年','REAL季度行事曆、完整檢討內容、後續追蹤及事前門檻紀錄','SYNTHETIC'],
]
fact_rows='\n'.join('| '+' | '.join(md_cell(x) for x in r)+' |' for r in new_facts)+'\n'
md=req_replace(md,'## 10. 官方要求與文件證據追溯',fact_rows+'\n## 10. 官方要求與文件證據追溯','new target/threshold register facts')
# Detailed 119-row record-level replacement mapping bound to the existing appendix item IDs and source CSV line.
with CSV.open(encoding='utf-8-sig',newline='') as f: data=list(csv.DictReader(f))
if len(data)!=119: raise ValueError(f'expected 119 synthetic rows, got {len(data)}')
appendix={'D1_annual_activity':'附錄A.1','D2_annual_admin':'附錄A.2','D3_valid_request':'附錄A.3','D4_quarterly_review':'附錄A.4'}
real={'D1_annual_activity':'真實事前年度活動清單／核准變更／完成或取消原件；逐場比對日期及分母。','D2_annual_admin':'真實行政責任清冊／原期限／交付物／完成日；核對完整年度母體。','D3_valid_request':'去識別化有效需求清冊／收件及首次實質回覆時戳／實際工作日曆與去重依據。','D4_quarterly_review':'事前季度行事曆／每季完整檢討紀錄／追蹤責任與日期。','Q1_action_followup':'真實行動指派及Q1母體／逐筆核對原期限、完成證據、分子分母與門檻事前採定來源。'}
register=[]
for ix,r in enumerate(data,1):
    ds=r['dataset']; item=r['item_id']
    if ds=='Q1_action_followup':
        action=item.split('-')[0]; loc={'A01':'附錄A.5','A02':'附錄A.6','A03':'附錄A.7','A04':'附錄A.8'}[action]
    else: loc=appendix[ds]
    date_parts=[f"期間={r['period']}"]
    if r['planned_or_due_date']: date_parts.append(f"計畫／到期／收件日={r['planned_or_due_date']}")
    if r['actual_date']: date_parts.append(f"實際／回覆日={r['actual_date']}")
    else: date_parts.append('實際／回覆日=未完成')
    den=int(r['denominator_flag'] or 0); num=int(r['numerator_flag'] or 0)
    outcome=f"分子／分母旗標={num}/{den}"
    if den: outcome+=f"；本列率={num/den*100:.1f}%"
    else: outcome+='；分母為0，率不計'
    if r['workdays_to_response']: outcome+=f"；工作日間隔={r['workdays_to_response']}"
    note=r['sample_reason_or_exception'] or '無另列例外；僅為合成記錄'
    key=f"S{ix:04d}"
    register.append([key,f"{loc}／{item}；CSV第{ix+1}資料列；{r['synthetic_source_key']}", '；'.join(date_parts), outcome, f"{note}；{r['data_status']}", real[ds]])
if len({r[0] for r in register})!=119 or len({r[1].split('；')[0] for r in register})!=119: raise ValueError('row fact IDs / locators not unique')
reg_md='''### 9.1 逐筆合成事實替換索引（119筆）

本表以唯一 `S0001`–`S0119` 替換鍵逐筆連結附錄 A 的每一個樣本明細列及原始合成 CSV 資料列。日期、分子旗標、分母旗標、比率、工作日、例外敘述及 SAMPLE 狀態均逐列列出；REAL 替換時不得將這些合成值沿用為事實。`CSV第N資料列`以標頭為第1行計，指向隨附 `SAMPLE_DATASET_0009_R02.csv`。\n\n| 替換鍵 | 附錄定位／合成ID／CSV資料列 | SAMPLE日期／期間 | 本列數值與結果 | 合成例外／狀態 | REAL替換原件與逐筆核驗 |\n|---|---|---|---|---|---|\n'''
for row in register: reg_md+='| '+' | '.join(md_cell(x) for x in row)+' |\n'
reg_md+='\n'
md=req_replace(md,'## 10. 官方要求與文件證據追溯',reg_md+'## 10. 官方要求與文件證據追溯','detail register insert')
md=req_replace(md,'§2表2-1；§4表4-1','§2表2-1與§2.2表2-2；§3.3表3-3；§4表4-1','trace updated')
MD.write_text(md,encoding='utf-8',newline='\n')

# 2) Make an exact localized DOCX revision from R03, retaining all untouched content.
doc=Document(R03)
# Add explicit R04 version marker to the existing subtitle, without changing other title content.
for p in doc.paragraphs:
    if '文類：年度績效目標分析紀錄' in p.text:
        p.text=p.text+'｜候選修訂：R04'
        break
else: raise ValueError('title mapping paragraph missing')
# Find update anchors before appending new tables.
def para_start(prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix): return p
    raise ValueError('paragraph anchor missing: '+prefix)
anchor3=para_start('3. 計算公式與判定門檻')
anchor4=para_start('4. 2025 年度 SAMPLE 個案與完整計算')
anchor10=para_start('10. 官方要求與文件證據追溯')
# Patch action status wording inside the existing action table.
for t in doc.tables:
    if t.rows and '行動 ID' in t.rows[0].cells[0].text:
        thresholds={'A01':'100%','A02':'≥95%','A03':'≥90%','A04':'100%'}
        for row in t.rows[1:]:
            aid=row.cells[0].text.strip()
            if aid in thresholds:
                cell=row.cells[4]
                if aid=='A04':
                    cell.text=cell.text.replace('Q1 為 1／1＝100.0%；僅代表第一季到期檢討完成','Q1 為 1／1＝100.0%；符合 §3.3 表3-3 事前 100% 門檻；僅代表第一季到期檢討完成')
                else:
                    cell.text=cell.text.replace('達到本例行動檢查目標',f'符合 §3.3 表3-3 事前 {thresholds[aid]} 門檻')
# Register fact summary locators and new facts.
for t in doc.tables:
    if t.rows and t.rows[0].cells[0].text.strip()=='fact_id':
        byid={row.cells[0].text.strip():row for row in t.rows[1:]}
        for fid in ('F0016','F0017','F0018','F0019','F0020'):
            row=byid.get(fid)
            if row:
                row.cells[1].text=row.cells[1].text+'；逐筆對應見 §9.1 替換鍵 S0001–S0119'
        for fact in new_facts:
            cells=t.add_row().cells
            for c,v in zip(cells,fact): c.text=v
        break
else: raise ValueError('fact register table missing')
# Update trace table locator.
for t in doc.tables:
    if t.rows and '官方要求／追溯項' in t.rows[0].cells[0].text:
        for row in t.rows[1:]:
            if 'I01-E06年度目標及分析' in row.cells[0].text:
                row.cells[2].text=row.cells[2].text.replace('§2表2-1；§4表4-1','§2表2-1與§2.2表2-2；§3.3表3-3；§4表4-1')
        break
# Formatting and table insertion helpers.
def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def set_repeat(row):
    trPr=row._tr.get_or_add_trPr(); e=OxmlElement('w:tblHeader'); e.set(qn('w:val'),'true'); trPr.append(e)
def no_split(row):
    trPr=row._tr.get_or_add_trPr(); e=OxmlElement('w:cantSplit'); trPr.append(e)
def style_table(t, widths):
    t.style='Table Grid'; t.autofit=False
    layout=OxmlElement('w:tblLayout'); layout.set(qn('w:type'),'fixed'); t._tbl.tblPr.append(layout)
    for i,w in enumerate(widths):
        t.columns[i].width=Inches(w)
    for ri,row in enumerate(t.rows):
        no_split(row)
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
                for run in p.runs:
                    run.font.name='Microsoft JhengHei'; run._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei'); run.font.size=Pt(9.5)
                    if ri==0: run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255)
            if ri==0: shade(cell,'17365D')
    set_repeat(t.rows[0])
def make_table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers))
    for c,h in zip(t.rows[0].cells,headers): c.text=h
    for vals in rows:
        cells=t.add_row().cells
        for c,v in zip(cells,vals): c.text=str(v)
    style_table(t,widths)
    return t

def make_heading(text):
    p=doc.add_paragraph(style='Heading 2'); p.add_run(text); p.paragraph_format.keep_with_next=True
    return p
def make_body(text):
    p=doc.add_paragraph(text,style='Normal'); p.paragraph_format.space_after=Pt(4); return p
def make_caption(text):
    p=doc.add_paragraph(text,style='Caption'); p.paragraph_format.keep_with_next=True; return p
def insert_group(anchor,items):
    for item in items: anchor._p.addprevious(item._p if hasattr(item,'_p') else item._tbl)

block1=[make_heading('2.2 2025年度目標事前設定示範（SAMPLE）'),make_body('下表補足如何訂出目標。設定日均為合成假設，早於2025衡量期間；基線、母體、選值理由與能力條件不代表協會實際目標、資源承諾或核准。正式採用須以真實事前版本及權責紀錄替換。'),make_caption('表2-2　年度目標事前設定與能力假設'),make_table(['目標ID','SAMPLE設定日','基線、目標及選值理由','資源／能力假設','證據界線與REAL替換'],targets,[0.62,0.95,2.95,3.05,2.93])]
insert_group(anchor3,block1)
h_q1=make_heading('3.3 2026 Q1 行動檢查門檻（事前 SAMPLE 設定）')
h_q1.paragraph_format.page_break_before=True
block2=[h_q1,make_body('本例假設於2025-12-15、Q1開始前設定下列門檻，故 §5 表5-1的「符合」可據此重算。這些是示範內部管理門檻，不是TTQS官方評分線或原始點協會核准標準。母體或來源不完整時一律HOLD。'),make_caption('表3-3　Q1行動指標預設門檻及實績判定'),make_table(['行動ID','事前合格門檻','設定時點與理由','Q1 SAMPLE實績','判定重算','HOLD／REAL注意事項'],q1,[0.62,0.95,2.55,1.1,2.25,3.03])]
insert_group(anchor4,block2)
block3=[make_heading('9.1 逐筆合成事實替換索引（119筆）'),make_body('本表以唯一S0001–S0119鍵逐筆連結附錄A每一個樣本明細列及隨附合成CSV資料列。日期、分子、分母、比率、工作日、例外敘述與SAMPLE狀態均列示；REAL替換時不得沿用合成值。CSV第N資料列以標頭為第1行計。'),make_caption('表9-1　119筆逐筆替換鍵及REAL原件核對'),make_table(['替換鍵','附錄定位／合成ID／CSV列','SAMPLE日期／期間','本列數值與結果','合成例外／狀態','REAL替換原件與逐筆核驗'],register,[0.68,1.45,1.73,1.22,2.15,3.27])]
insert_group(anchor10,block3)
# Save as new candidate only.
doc.core_properties.title='0009 I01-E06 年度績效目標訂定與分析紀錄 R04'
doc.core_properties.subject='Localized repair of independent semantic findings; CANDIDATE only'
doc.save(DOCX)

# 3) Verify manuscript/table construction basics, capture immutable build provenance.
assert len(data)==119 and len(register)==119
assert all(row[-2].endswith('SAMPLE_SYNTHETIC') for row in register)
assert all(x in md for x in ['## 2.2 2025年度目標事前設定示範','### 3.3 2026 Q1 行動檢查門檻','### 9.1 逐筆合成事實替換索引（119筆）','S0119'])
receipt={
 'doc_id':'0009','official_requirement_id':'I01-E06','checkbox_serial':'I01-E06','revision':'R04','state':'CANDIDATE_NOT_ACCEPTED',
 'repair_scope':['add dated and reasoned sample target-setting records for G01-G04','predefine Q1 action thresholds and recompute A01-A04 status','add unique fact-level mapping for all 119 sample dataset rows'],
 'input_r03_docx_sha256':sha(R03),'input_r03_manuscript_sha256':sha(R03MD),'sample_dataset_sha256':sha(CSV),
 'official_source_pdf_sha256':'805386B77DF045B396FABAFED7A47D12491C25A0B5744EDD9A04876CAEE34709',
 'unique_blueprint_row_sha256':'4D00595289266D8126E8E890FC9E503037ACC43D296C0A176DB5025EC01BC3F0',
 'r04_manuscript_sha256':sha(MD),'candidate_path':str(DOCX),'candidate_sha256':sha(DOCX),'candidate_bytes':DOCX.stat().st_size,
 'synthetic_register_row_count':len(register),'sample_dataset_rows':len(data),'annual_baseline_target_records':4,'q1_threshold_records':4,
 'formal_acceptance':'NOT_GRANTED','semantic_review':'PENDING','real_association_facts':'UNKNOWN',
 'created_at_tpe':datetime.now().astimezone().isoformat(timespec='seconds')
}
(OUT/'evidence'/'BUILD_RECEIPT_0009_R04.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'candidate':str(DOCX),'sha256':receipt['candidate_sha256'],'bytes':receipt['candidate_bytes'],'manuscript_sha256':receipt['r04_manuscript_sha256'],'register_rows':len(register),'dataset_rows':len(data),'tables':len(Document(DOCX).tables),'receipt':str(OUT/'evidence'/'BUILD_RECEIPT_0009_R04.json')},ensure_ascii=False))
