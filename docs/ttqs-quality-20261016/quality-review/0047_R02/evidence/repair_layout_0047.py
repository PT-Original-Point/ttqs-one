from __future__ import annotations

import hashlib
import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[5]
PACKAGE = Path(__file__).resolve().parents[1]
SOURCE = PACKAGE.parent / "0047_R01" / "candidate" / "0047_R01.docx"
DEST = PACKAGE / "candidate" / "0047_R02.docx"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = tr_pr.find(qn("w:tblHeader"))
    if node is None:
        node = OxmlElement("w:tblHeader")
        tr_pr.append(node)
    node.set(qn("w:val"), "1")


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = tr_pr.find(qn("w:cantSplit"))
    if node is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=75, start=90, bottom=75, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "5")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "D9D9D9")


def set_fixed_widths(table, widths) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    for i, width in enumerate(widths):
        table.columns[i].width = Inches(width)
        table._tbl.tblGrid.gridCol_lst[i].set(qn("w:w"), str(int(Inches(width).twips)))
        for row in table.rows:
            row.cells[i].width = Inches(width)


def set_run_font(run, size=10, bold=False, color=None) -> None:
    run.font.name = "Microsoft JhengHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def fill_table(table, headers, records, widths, centered_columns=()) -> None:
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_fixed_widths(table, widths)
    set_borders(table)
    rows = [headers, *records]
    for ri, values in enumerate(rows):
        row = table.rows[0] if ri == 0 else table.add_row()
        set_cant_split(row)
        if ri == 0:
            set_repeat_header(row)
        for ci, value in enumerate(values):
            cell = row.cells[ci]
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if ri == 0:
                set_cell_shading(cell, "1F4E78")
            elif ri % 2 == 0:
                set_cell_shading(cell, "F2F6FA")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER if ci in centered_columns or ri == 0
                    else WD_ALIGN_PARAGRAPH.LEFT
                )
                for run in paragraph.runs:
                    set_run_font(run, size=10, bold=(ri == 0), color=("FFFFFF" if ri == 0 else None))
    set_fixed_widths(table, widths)


def add_heading(doc, text: str):
    p = doc.add_paragraph(style="Heading 3")
    p.add_run(text)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    return p


if not SOURCE.is_file():
    raise FileNotFoundError(SOURCE)
if DEST.exists():
    raise FileExistsError(f"Refusing to overwrite an existing candidate: {DEST}")

source_sha_before = sha256(SOURCE)
shutil.copy2(SOURCE, DEST)
doc = Document(DEST)

# Locate and validate the exact existing section and 13-column sample table.
dictionary_heading = next(p for p in doc.paragraphs if p.text.strip() == "資料欄位、定義、分母、來源及限制字典")
annual_heading = next(p for p in doc.paragraphs if p.text.strip() == "年度案例明細與逐步計算")
old_annual_table = doc.tables[2]
original = [[cell.text.strip() for cell in row.cells] for row in old_annual_table.rows]
expected_header = [
    "示範年度", "案例", "需求與事前規準分數", "選擇理由可重現分數", "契約及驗收分數",
    "履約檢討改善分數", "四項權重", "加權計算", "加權總分／5", "履約檢核完成／應完成",
    "履約達成率", "判讀", "資料狀態",
]
if original[0] != expected_header or len(original) != 3:
    raise AssertionError("R01 source table differs from the observed candidate; refusing to repair an unexpected version.")
if original[1][2:6] != ["4", "4", "3", "2"] or original[2][2:6] != ["4", "4", "4", "4"]:
    raise AssertionError("Observed scoring values changed; refusing to carry forward stale repair assumptions.")
if original[1][8] != "3.25" or original[1][9:11] != ["3／4", "75.0%"]:
    raise AssertionError("2024 sample calculations do not match the audited R01 baseline.")
if original[2][8] != "4.00" or original[2][9:11] != ["4／4", "100.0%"]:
    raise AssertionError("2025 sample calculations do not match the audited R01 baseline.")

# Keep the dictionary intact, but give its heading and full table a clean page start;
# repeat the header if Word still needs to flow the table to a continuation page.
dictionary_heading.paragraph_format.page_break_before = True
dictionary_heading.paragraph_format.keep_with_next = True
dictionary_table = doc.tables[1]
set_repeat_header(dictionary_table.rows[0])
for row in dictionary_table.rows:
    set_cant_split(row)

# Normalize the over-wide table into three small, readable tables. All source values
# remain unchanged; the split is strictly presentational and is bound to R01 SHA.
score_headers = ["示範年度", "案例", "評分構面", "分數（1–5）", "權重"]
score_records = []
criteria = [
    ("需求與事前規準", 2, 0),
    ("選擇理由可重現", 3, 1),
    ("契約及驗收", 4, 2),
    ("履約檢討改善", 5, 3),
]
for row in original[1:]:
    year, case = row[0], row[1]
    for label, score_col, _ in criteria:
        score_records.append([year, case, label, row[score_col], "25%"])

calc_headers = ["示範年度", "案例", "加權計算式", "加權總分／5"]
calc_records = [[row[0], row[1], row[7], row[8]] for row in original[1:]]

tracking_headers = ["示範年度", "案例", "完成／應完成", "履約達成率", "判讀", "資料狀態"]
tracking_records = [[row[0], row[1], row[9], row[10], row[11], row[12]] for row in original[1:]]

elements = []
tables = []

add_heading(doc, "評分構面與示範分數")
elements.append(doc.paragraphs[-1]._p)
t = doc.add_table(rows=1, cols=len(score_headers))
fill_table(t, score_headers, score_records, [0.82, 0.70, 3.52, 0.90, 0.76], centered_columns=(0, 1, 3, 4))
elements.append(t._tbl)
tables.append(t)

add_heading(doc, "加權計算與總分")
elements.append(doc.paragraphs[-1]._p)
t = doc.add_table(rows=1, cols=len(calc_headers))
fill_table(t, calc_headers, calc_records, [0.82, 0.70, 4.28, 0.90], centered_columns=(0, 1, 3))
elements.append(t._tbl)
tables.append(t)

add_heading(doc, "履約檢核結果與判讀")
elements.append(doc.paragraphs[-1]._p)
t = doc.add_table(rows=1, cols=len(tracking_headers))
fill_table(t, tracking_headers, tracking_records, [0.78, 0.66, 0.86, 0.88, 2.58, 0.92], centered_columns=(0, 1, 2, 3))
elements.append(t._tbl)
tables.append(t)

annual_heading.paragraph_format.page_break_before = True
annual_heading.paragraph_format.keep_with_next = True
anchor = old_annual_table._tbl.getnext()
if anchor is None:
    raise AssertionError("Could not locate the next section to place repaired tables.")
old_annual_table._tbl.getparent().remove(old_annual_table._tbl)
for element in elements:
    element.getparent().remove(element)
    anchor.addprevious(element)

doc.core_properties.comments = "R02: presentation-only repair of the orphaned dictionary row and 13-column annual table; all values retained from exact R01 candidate."
doc.save(DEST)

# Re-open and verify shape, content, arithmetic source values, repeated headers, and parent immutability.
check = Document(DEST)
if len(check.tables) != 11:
    raise AssertionError(f"Expected 11 tables after replacing one table with three; found {len(check.tables)}")
new_tables = [check.tables[i] for i in (2, 3, 4)]
if [len(t.columns) for t in new_tables] != [5, 4, 6]:
    raise AssertionError("Repaired annual tables have an unexpected column count.")
if any(t.rows[0]._tr.trPr.find(qn("w:tblHeader")) is None for t in new_tables):
    raise AssertionError("One or more new table headers are not marked as repeating.")
flat = "\n".join(cell.text for t in new_tables for row in t.rows for cell in row.cells)
for expected in ("4×25%＋4×25%＋3×25%＋2×25%", "3.25", "75.0%", "4×25%＋4×25%＋4×25%＋4×25%", "4.00", "100.0%", "示範資料｜非本會實際紀錄"):
    if expected not in flat:
        raise AssertionError(f"Value lost during table restructuring: {expected}")
if sha256(SOURCE) != source_sha_before:
    raise AssertionError("R01 parent bytes changed during R02 creation.")

receipt = {
    "doc_id": "0047",
    "revision": "R02",
    "change_scope": "PRESENTATION_ONLY",
    "parent_candidate": str(SOURCE),
    "parent_sha256": source_sha_before,
    "candidate": str(DEST),
    "candidate_sha256": sha256(DEST),
    "layout_fixes": [
        "Start the data dictionary section on a clean page and repeat its table header.",
        "Replace the observed 13-column annual detail table with three readable tables: scores, weighted calculation, and fulfillment interpretation.",
        "Mark all repaired table header rows to repeat and table rows not to split across pages.",
    ],
    "preserved_sample_values": [
        "2024 course A: scores 4,4,3,2; equal 25% weights; weighted score 3.25/5; fulfillment 3/4=75.0%.",
        "2025 course B: scores 4,4,4,4; equal 25% weights; weighted score 4.00/5; fulfillment 4/4=100.0%.",
        "All source judgment text and SAMPLE/SYNTHETIC labels retained verbatim from R01 table.",
    ],
    "post_save_checks": "PASS: python-docx reopened R02; table counts and column counts match; repeating headers present; key numeric values retained; R01 SHA unchanged.",
    "independent_semantic_review": "PENDING: historical R01 review is not inherited by R02 SHA.",
}
(PACKAGE / "evidence" / "R02_LAYOUT_REPAIR_RECEIPT.json").write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(receipt, ensure_ascii=False, indent=2))
