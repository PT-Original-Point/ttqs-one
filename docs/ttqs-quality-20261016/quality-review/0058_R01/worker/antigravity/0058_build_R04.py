#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
0058_build_R04.py
TTQS 文件流水號 0058 (I13-E02 / 方框 062) 專屬 Word (DOCX) 建置器（R04 真實性與提案邊界校正版）

【功能說明】
- 本腳本專門讀取同目錄下的 0058_CONTENT_R04.json
- 使用 python-docx 建置高品質、排版成熟的繁體中文 Word 候選文件：candidate/0058_R04.docx
- 恪遵 TTQS 品質硬規則與 R04 校正重點：
  * 正文字級 10.5 pt，表格字級 9.5 pt，嚴禁縮字擠頁
  * 表格跨頁自動重複標題列 (tblHeader)，列不跨頁斷裂 (cantSplit)
  * 第一頁具備實質性導讀與正式來源／示範資料邊界明確宣告
  * 包含第 5 節表 5-2 之完整單筆示範案例操作紀錄 (REC-0058-A01-DEMO，註明為 2/2 之一)
  * 第七節空白操作表與已填示範表清楚隔離，並標示「實際使用時另開新表」且落實受控代碼
  * 第八節待取得佐證原件採用 REQ-0058-01 至 08 編號，標明目前狀態 UNKNOWN
  * 第九節定位為供審議之示範審查/退件矩陣提案
  * 頁首頁尾包含文件識別與頁碼（第 X 頁共 Y 頁）
- 僅讀取同目錄之 0058_CONTENT_R04.json，不連網、不讀取外部 DOCX

【執行方式】
由整合者於專案工作包目錄下執行：
  python worker/antigravity/0058_build_R04.py
"""

import os
import sys
import json
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

FONT_NAME = "微軟正黑體"
COLOR_PRIMARY = RGBColor(31, 73, 125)     # 深藍 (#1F497D)
COLOR_SECONDARY = RGBColor(44, 62, 80)   # 深灰藍 (#2C3E50)
COLOR_TEXT = RGBColor(51, 51, 51)        # 正文深灰 (#333333)
COLOR_MUTED = RGBColor(102, 102, 102)    # 註釋次灰 (#666666)
COLOR_WARNING = RGBColor(180, 40, 40)    # 警示暗紅 (#B42828)

def set_run_font(run, font_name=FONT_NAME, size_pt=10.5, bold=False, italic=False, color_rgb=COLOR_TEXT):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:eastAsia="{font_name}"/>')
    rPr.append(rFonts)

def set_cell_shading(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

def set_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="B0C4DE", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_page_number_to_run(run):
    fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
    fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    fldChar3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def add_numpages_to_run(run):
    fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> NUMPAGES </w:instrText>')
    fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    fldChar3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def build_docx():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    json_path = os.path.join(script_dir, "0058_CONTENT_R04.json")
    
    if not os.path.exists(json_path):
        print(f"[ERROR] 找不到 R04 內容資料檔：{json_path}", file=sys.stderr)
        sys.exit(1)
        
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)
        
    metadata = data.get("metadata", {})
    front_matter = data.get("front_matter", {})
    sections = data.get("sections", [])
    
    # 決定輸出目錄 candidate/0058_R04.docx
    package_root = os.path.abspath(os.path.join(script_dir, "..", ".."))
    candidate_dir = os.path.join(package_root, "candidate")
    os.makedirs(candidate_dir, exist_ok=True)
    out_docx_path = os.path.join(candidate_dir, "0058_R04.docx")
    
    doc = Document()
    
    # 設定標準邊界 (A4, 2.54 cm)
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        
        # 頁首設定
        header = section.header
        header_p = header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        h_run = header_p.add_run("社團法人屏東縣原始點關懷協會｜TTQS 評核佐證文件 0058（I13-E02）")
        set_run_font(h_run, size_pt=9.0, color_rgb=COLOR_MUTED)
        
        # 頁尾設定
        footer = section.footer
        footer_p = footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        f_run1 = footer_p.add_run("【完整示範文件 SAMPLE / SYNTHETIC｜非協會真實開課或核定實績】  第 ")
        set_run_font(f_run1, size_pt=9.0, color_rgb=COLOR_MUTED)
        f_run_page = footer_p.add_run()
        set_run_font(f_run_page, size_pt=9.0, color_rgb=COLOR_MUTED)
        add_page_number_to_run(f_run_page)
        f_run2 = footer_p.add_run(" 頁，共 ")
        set_run_font(f_run2, size_pt=9.0, color_rgb=COLOR_MUTED)
        f_run_numpages = footer_p.add_run()
        set_run_font(f_run_numpages, size_pt=9.0, color_rgb=COLOR_MUTED)
        add_numpages_to_run(f_run_numpages)
        f_run3 = footer_p.add_run(" 頁")
        set_run_font(f_run3, size_pt=9.0, color_rgb=COLOR_MUTED)

    # ------------------ 文件首頁：大標題與識別區塊 ------------------
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(6)
    title_p.paragraph_format.space_after = Pt(4)
    run_org = title_p.add_run("社團法人屏東縣原始點關懷協會\n")
    set_run_font(run_org, size_pt=14.0, bold=True, color_rgb=COLOR_SECONDARY)
    run_main_title = title_p.add_run(metadata.get("title", "TTQS 訓練成果移轉與成效追蹤工具（示範提案）"))
    set_run_font(run_main_title, size_pt=17.0, bold=True, color_rgb=COLOR_PRIMARY)
    
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(12)
    meta_text = (
        f"文件編號：TTQS-0058  |  指標代碼：{metadata.get('requirement_id', 'I13-E02')}  |  "
        f"官方方框：062（資料表第4頁）  |  版次：R04（真實性與提案邊界校正版）  |  "
        f"示範案號：{metadata.get('case_id', 'SYN-0058-202608')}"
    )
    run_sub = sub_p.add_run(meta_text)
    set_run_font(run_sub, size_pt=9.5, bold=True, color_rgb=COLOR_MUTED)

    # 導讀與真實邊界特別提示框（第一頁實質導讀）
    lead_box_table = doc.add_table(rows=1, cols=1)
    lead_box_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = lead_box_table.cell(0, 0)
    set_cell_shading(cell, "F2F4F7")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    p_box = cell.paragraphs[0]
    p_box.paragraph_format.line_spacing = 1.2
    p_box.paragraph_format.space_after = Pt(4)
    r_bh = p_box.add_run("【文件用途與審查導讀核心指引】\n")
    set_run_font(r_bh, size_pt=11.0, bold=True, color_rgb=COLOR_PRIMARY)
    r_bp = p_box.add_run(front_matter.get("plain_language_purpose", "") + "\n\n")
    set_run_font(r_bp, size_pt=10.0, color_rgb=COLOR_TEXT)
    
    r_br = p_box.add_run("【主要讀者與操作對象】\n")
    set_run_font(r_br, size_pt=10.5, bold=True, color_rgb=COLOR_SECONDARY)
    r_br_txt = p_box.add_run(front_matter.get("reader", "") + "\n\n")
    set_run_font(r_br_txt, size_pt=10.0, color_rgb=COLOR_TEXT)
    
    r_bs = p_box.add_run("【資料真實性與邊界聲明 (SAMPLE BOUNDARY)】\n")
    set_run_font(r_bs, size_pt=10.5, bold=True, color_rgb=COLOR_WARNING)
    r_bs_txt = p_box.add_run(front_matter.get("sample_boundary", "") + "\n\n")
    set_run_font(r_bs_txt, size_pt=10.0, color_rgb=COLOR_TEXT)

    r_bu = p_box.add_run("【建議使用順序】\n")
    set_run_font(r_bu, size_pt=10.5, bold=True, color_rgb=COLOR_SECONDARY)
    r_bu_txt = p_box.add_run(front_matter.get("use_sequence", ""))
    set_run_font(r_bu_txt, size_pt=10.0, color_rgb=COLOR_TEXT)
    
    set_table_borders(lead_box_table, color="BDC3C7", sz="6", val="single")

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ------------------ 渲染十個實質章節 ------------------
    for sec_idx, sec in enumerate(sections, start=1):
        # 章節標題 (Heading 1)
        h1_p = doc.add_paragraph()
        h1_p.paragraph_format.space_before = Pt(14)
        h1_p.paragraph_format.space_after = Pt(6)
        h1_p.paragraph_format.keep_with_next = True
        h1_run = h1_p.add_run(sec.get("heading", ""))
        set_run_font(h1_run, size_pt=13.5, bold=True, color_rgb=COLOR_PRIMARY)
        
        # 章節段落
        for para_text in sec.get("paragraphs", []):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.25
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run = p.add_run(para_text)
            set_run_font(run, size_pt=10.5, color_rgb=COLOR_TEXT)

        # 章節表格
        for tbl_data in sec.get("tables", []):
            caption = tbl_data.get("caption", "")
            columns = tbl_data.get("columns", [])
            rows = tbl_data.get("rows", [])
            
            # 表格標題 Caption
            cap_p = doc.add_paragraph()
            cap_p.paragraph_format.space_before = Pt(8)
            cap_p.paragraph_format.space_after = Pt(3)
            cap_p.paragraph_format.keep_with_next = True
            cap_run = cap_p.add_run(caption)
            set_run_font(cap_run, size_pt=10.5, bold=True, color_rgb=COLOR_SECONDARY)
            
            # 建立表格
            table = doc.add_table(rows=len(rows) + 1, cols=len(columns))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(table, color="BDC3C7", sz="4", val="single")
            
            # 表頭列
            hdr_cells = table.rows[0].cells
            set_repeat_table_header(table.rows[0])
            set_row_cant_split(table.rows[0])
            for col_idx, col_name in enumerate(columns):
                hdr_cells[col_idx].text = col_name
                set_cell_shading(hdr_cells[col_idx], "D9E1F2") # 典雅淡藍表頭
                set_cell_margins(hdr_cells[col_idx], top=110, bottom=110, left=130, right=130)
                p_hdr = hdr_cells[col_idx].paragraphs[0]
                p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_hdr.paragraph_format.space_after = Pt(0)
                for r in p_hdr.runs:
                    set_run_font(r, size_pt=9.5, bold=True, color_rgb=COLOR_SECONDARY)
            
            # 資料列
            for r_idx, row_values in enumerate(rows, start=1):
                row = table.rows[r_idx]
                set_row_cant_split(row)
                row_cells = row.cells
                
                # 交替背景色 (斑馬紋)
                bg_color = "F9FAFC" if (r_idx % 2 == 0) else "FFFFFF"
                
                # 若為彙整統計列或特殊宣告列，給予加深標註
                is_summary_row = False
                if len(row_values) > 1 and any(keyword in str(row_values[0]) or keyword in str(row_values[1]) for keyword in ["彙整統計", "整體成效", "組織整體層級", "方案狀態宣告"]):
                    bg_color = "EAEEF5"
                    is_summary_row = True
                
                # 若為第 5 節表 5-2 的完整案例紀錄表，給予專屬高雅淡綠/米白微調
                if "表 5-2" in caption:
                    if r_idx == 1:
                        bg_color = "EEF4F8"
                    elif r_idx == len(rows):
                        bg_color = "FDF6E2"
                
                for c_idx, cell_value in enumerate(row_values):
                    row_cells[c_idx].text = str(cell_value)
                    set_cell_shading(row_cells[c_idx], bg_color)
                    set_cell_margins(row_cells[c_idx], top=90, bottom=90, left=120, right=120)
                    
                    p_cell = row_cells[c_idx].paragraphs[0]
                    p_cell.paragraph_format.space_after = Pt(0)
                    p_cell.paragraph_format.line_spacing = 1.15
                    
                    # 判斷文字對齊：短代碼/項次置中，長文字靠左
                    txt = str(cell_value).strip()
                    if len(txt) <= 10 or txt in ["達標", "未達標", "0.0 pp", "REJECT (退件提案)", "HOLD (暫緩提案)"]:
                        p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        p_cell.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        
                    for r in p_cell.runs:
                        is_bold = is_summary_row or txt in ["達標", "REJECT (退件提案)"]
                        font_color = COLOR_WARNING if "未達標" in txt or "REJECT" in txt else COLOR_TEXT
                        set_run_font(r, size_pt=9.5, bold=is_bold, color_rgb=font_color)
            
            # 表格後間距
            doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 儲存輸出檔案
    doc.save(out_docx_path)
    file_size = os.path.getsize(out_docx_path)
    print(f"[SUCCESS] 0058 R04 DOCX 建置完成！")
    print(f"檔案路徑：{out_docx_path}")
    print(f"檔案大小：{file_size:,} bytes")
    print(f"章節數量：{len(sections)} 節")
    print(f"表格總數：{sum(len(s.get('tables', [])) for s in sections)} 張")

if __name__ == "__main__":
    build_docx()
