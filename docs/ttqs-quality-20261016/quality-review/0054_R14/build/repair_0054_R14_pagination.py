from __future__ import annotations
import argparse
import hashlib
import json
import zipfile
from datetime import datetime
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

MONITOR_HEADING = '遴選與教材監測指標（單梯示範）'
MINUTES_HEADING = '會議議程、意見與決議'

def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()

def body_signature(doc: Document) -> str:
    blocks = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn('w:p'):
            text = ''.join(t.text or '' for t in child.iter(qn('w:t')))
            blocks.append(['p', text])
        elif child.tag == qn('w:tbl'):
            rows = []
            for row in child.iter(qn('w:tr')):
                cells = []
                for cell in row.iter(qn('w:tc')):
                    cells.append(''.join(t.text or '' for t in cell.iter(qn('w:t'))))
                rows.append(cells)
            blocks.append(['tbl', rows])
    return sha(json.dumps(blocks, ensure_ascii=False, separators=(',', ':')).encode('utf-8'))

def table_flag(tr_pr, name: str, value: str = '1'):
    el = tr_pr.find(qn(name))
    if el is None:
        el = OxmlElement(name)
        tr_pr.append(el)
    el.set(qn('w:val'), value)

ap = argparse.ArgumentParser()
ap.add_argument('--source', required=True)
ap.add_argument('--output', required=True)
ap.add_argument('--receipt', required=True)
a = ap.parse_args()
source = Path(a.source).resolve()
output = Path(a.output).resolve()
receipt_path = Path(a.receipt).resolve()
source_bytes = source.read_bytes()
source_sha = sha(source_bytes)
expected_source_sha = '7b14bcfc09ad38c42b73ffebbfaee1c09939e7b9f7ae578240168cd1aa13b6ea'
if source_sha.lower() != expected_source_sha:
    raise SystemExit(f'Unexpected source SHA: {source_sha}')
source_doc = Document(source)
text_sig_before = body_signature(source_doc)
source_zip = zipfile.ZipFile(source)
source_parts = {name: source_zip.read(name) for name in source_zip.namelist()}

for table in source_doc.tables:
    header_tr = table.rows[0]._tr
    header_pr = header_tr.get_or_add_trPr()
    table_flag(header_pr, 'w:tblHeader')
    for row in table.rows:
        row_pr = row._tr.get_or_add_trPr()
        table_flag(row_pr, 'w:cantSplit')
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_with_next = True

monitor_matches = [p for p in source_doc.paragraphs if p.text.strip() == MONITOR_HEADING]
if len(monitor_matches) != 1:
    raise SystemExit(f'Expected exactly one monitoring heading; found {len(monitor_matches)}')
monitor_matches[0].paragraph_format.page_break_before = False
monitor_matches[0].paragraph_format.keep_with_next = True

minutes_matches = [i for i, p in enumerate(source_doc.paragraphs) if p.text.strip() == MINUTES_HEADING]
if len(minutes_matches) != 1:
    raise SystemExit(f'Expected exactly one minutes heading; found {len(minutes_matches)}')
minutes_index = minutes_matches[0]
for paragraph in source_doc.paragraphs[minutes_index + 1:]:
    paragraph.paragraph_format.space_after = Pt(0)

output.parent.mkdir(parents=True, exist_ok=True)
source_doc.save(output)
# Restore original core properties; the repair changes pagination controls only.
tmp = output.with_name(output.name + '.tmp')
with zipfile.ZipFile(output, 'r') as zin:
    with zipfile.ZipFile(tmp, 'w') as zout:
        for info in zin.infolist():
            data = source_parts['docProps/core.xml'] if info.filename == 'docProps/core.xml' else zin.read(info.filename)
            zout.writestr(info, data)
tmp.replace(output)
output_doc = Document(output)
text_sig_after = body_signature(output_doc)
if text_sig_before != text_sig_after:
    raise SystemExit('Semantic text signature changed during layout repair')
if len(output_doc.tables) != 11:
    raise SystemExit(f'Unexpected table count: {len(output_doc.tables)}')
for index, table in enumerate(output_doc.tables):
    tr_pr = table.rows[0]._tr.trPr
    if tr_pr is None or tr_pr.find(qn('w:tblHeader')) is None:
        raise SystemExit(f'Table {index} header is not repeated')
    for ri, row in enumerate(table.rows):
        row_pr = row._tr.trPr
        if row_pr is None or row_pr.find(qn('w:cantSplit')) is None:
            raise SystemExit(f'Table {index} row {ri} can still split')
with zipfile.ZipFile(output) as zf:
    bad = zf.testzip()
    if bad:
        raise SystemExit(f'Invalid DOCX ZIP member: {bad}')
    output_parts = {name: zf.read(name) for name in zf.namelist()}
changed_parts = sorted(name for name in set(source_parts) | set(output_parts) if source_parts.get(name) != output_parts.get(name))
receipt = {
    'schema': 'ttqs.docx_layout_repair_receipt.v1',
    'project_id': 'TTQS_ONE',
    'doc_id': '0054',
    'requirement_id': 'I12-E02',
    'official_checkbox_sequence': '054',
    'source_revision': 'R13',
    'source_sha256': source_sha,
    'candidate_revision': 'R14',
    'candidate_path': str(output),
    'candidate_sha256': sha(output.read_bytes()),
    'candidate_bytes': output.stat().st_size,
    'semantic_text_signature_before': text_sig_before,
    'semantic_text_signature_after': text_sig_after,
    'semantic_text_unchanged': text_sig_before == text_sig_after,
    'table_count': len(output_doc.tables),
    'repeated_header_tables': 11,
    'non_splitting_rows': sum(len(t.rows) for t in output_doc.tables),
    'pagination_repairs': [
        'Set the first row of each of 11 tables to repeat as a header.',
        'Set every table row to stay on one page; no row content was edited.',
        'Keep header-cell paragraphs with following content to avoid orphan table headers.',
        'Allow the monitoring section to flow after the exception table; keep its heading with the next paragraph to avoid an orphan heading.',
        'Remove 4 pt after-spacing from paragraphs in the final meeting-minutes block so the handoff paragraph can flow without a sparse last page.'
    ],
    'changed_docx_parts': changed_parts,
    'zip_integrity': 'PASS',
    'built_at_local': datetime.now().astimezone().isoformat(timespec='seconds')
}
receipt_path.parent.mkdir(parents=True, exist_ok=True)
receipt_path.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
print(json.dumps(receipt, ensure_ascii=False, indent=2))
