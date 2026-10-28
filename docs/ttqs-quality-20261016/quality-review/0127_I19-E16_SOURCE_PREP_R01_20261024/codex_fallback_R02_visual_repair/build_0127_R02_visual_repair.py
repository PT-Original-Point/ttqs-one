from pathlib import Path
from datetime import datetime
import csv,hashlib,json
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION,WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
root=Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0127_I19-E16_SOURCE_PREP_R01_20261024")
src=root/"source";out=root/"codex_fallback_R02_visual_repair";assets=out/"assets";out.mkdir(parents=True,exist_ok=True);assets.mkdir(exist_ok=True)
docx=out/"0127__目標市場與顧客價值量化效益分析表_R02_CODEX_CANDIDATE.docx"
md=out/"manuscript_0127_I19-E16_R02_CODEX_CANDIDATE.md";register=out/"SYNTHETIC_REPLACEMENT_REGISTER_0127_R02.csv"
receipt=out/"technical_candidate_receipt_0127_R02.json"
def sha(b):return hashlib.sha256(b).hexdigest()
pins={"SOURCE_PACKET_0127_I19-E16_R01.md":"058663e200fd0207196c08a02e4d159c6e3eaa7f167094cd7b409b4050b18675","UNIQUE_BLUEPRINT_0127_I19-E16_R01.md":"9a249a2e398407469bd82c9bf5d7a6188c9e3d28be79172bd4cc33b213c700ff","SAMPLE_MARKET_PANEL_R01.csv":"9d188e1c8c4b2a488cac2a989a029358dfae8c022ac30a9298518126d137fa36","SAMPLE_CUSTOMER_VALUE_PANEL_R01.csv":"ca7d92fc02a98ae32e0d306c0a42ff2fd9d6cd58a25b53cb34e89203325b15b6"}
for f,h in pins.items():
 if sha((src/f).read_bytes())!=h:raise SystemExit("source changed "+f)
m=list(csv.DictReader((src/"SAMPLE_MARKET_PANEL_R01.csv").open(encoding="utf-8-sig",newline="")))
c=list(csv.DictReader((src/"SAMPLE_CUSTOMER_VALUE_PANEL_R01.csv").open(encoding="utf-8-sig",newline="")))
mk=("qualified_unique_target_prospects_A","qualified_unique_target_prospects_B","first_purchase_conversions_A","first_purchase_conversions_B")
ck=("usefulness_A_1to5","convenience_A_1to5","benefit_vs_price_time_A_1to5","usefulness_B_1to5","convenience_B_1to5","benefit_vs_price_time_B_1to5")
for r in m:
 for k in mk:r[k]=int(r[k])
for r in c:
 for k in ck:r[k]=int(r[k])
assert (len(m),len(c))==(8,24)
a=sum(r[mk[0]] for r in m);b=sum(r[mk[1]] for r in m);ca=sum(r[mk[2]] for r in m);cb=sum(r[mk[3]] for r in m)
assert (a,b,ca,cb)==(32,48,8,15)
keys=[("usefulness","需求效用"),("convenience","使用便利"),("benefit_vs_price_time","相對價格／時間效益")]
st={}
for k,nm in keys:
 x=sum(r[k+"_A_1to5"] for r in c);y=sum(r[k+"_B_1to5"] for r in c);st[k]=(x,y,x/24,y/24,(y-x)/24)
totA=sum(r[k] for r in c for k in ck[:3]);totB=sum(r[k] for r in c for k in ck[3:])
assert (totA,totB)==(234,288)
meanA=totA/72;meanB=totB/72;rateA=ca/a;rateB=cb/b
# Create two charts with two independent units for the market chart.
font=r"C:\Windows\Fonts\msjh.ttc";F=lambda n:ImageFont.truetype(font,n)
nav=(31,56,78);teal=(43,121,118);gold=(205,145,52);grid=(215,224,229);ink=(45,55,65);red=(164,75,65)
def ctr(d,x,y,t,f,col=ink):
 q=d.textbbox((0,0),t,font=f);d.text((x-(q[2]-q[0])/2,y),t,font=f,fill=col)
def chart(path,title,groups,lo,hi,note):
 im=Image.new("RGB",(1600,620),"white");d=ImageDraw.Draw(im);ctr(d,800,18,title,F(36),nav);L=110;R=1490;T=160;B=440
 for i in range(6):
  v=lo+(hi-lo)*i/5;y=B-(B-T)*i/5;d.line((L,y,R,y),fill=grid,width=2);d.text((45,y-12),f"{v:g}",font=F(24),fill=ink)
 d.rectangle((600,91,626,117),fill=teal);d.text((638,87),"M1",font=F(25),fill=ink);d.rectangle((735,91,761,117),fill=gold);d.text((773,87),"M4",font=F(25),fill=ink)
 g=(R-L)/len(groups)
 for i,(label,x,yv) in enumerate(groups):
  cx=L+g*(i+.5)
  for j,v in enumerate((x,yv)):
   xx=cx+(j-.5)*54;h=(B-T)*(v-lo)/(hi-lo);d.rounded_rectangle((xx-21,B-h,xx+21,B),radius=5,fill=teal if j==0 else gold);ctr(d,xx,B-h-27,f"{v:.2f}",F(23))
  ctr(d,cx,B+14,label,F(25))
 ctr(d,800,540,note,F(26),red);im.save(path,dpi=(180,180))
def chart_market(path):
 im=Image.new("RGB",(1600,620),"white");d=ImageDraw.Draw(im);ctr(d,800,18,"圖1｜合格客群觸及與首購轉換（不同尺度分圖）",F(36),nav)
 panels=[(70,780,"合格且去重潛客（人）",[32,48],60),(820,1530,"同窗首購轉換率（%）",[25,31.25],40)]
 for x0,x1,title,vals,scale in panels:
  d.rounded_rectangle((x0,100,x1,500),radius=16,fill=(250,252,252),outline=grid,width=2);ctr(d,(x0+x1)//2,113,title,F(31),nav)
  L=x0+75;R=x1-35;T=190;B=418
  for i in range(6):
   v=scale*i/5;y=B-(B-T)*i/5;d.line((L,y,R,y),fill=grid,width=2);d.text((x0+16,y-11),f"{v:g}",font=F(15),fill=ink)
  for j,v in enumerate(vals):
   cx=L+(R-L)*(.34 if j==0 else .72);h=(B-T)*v/scale;d.rounded_rectangle((cx-54,B-h,cx+54,B),radius=8,fill=teal if j==0 else gold);ctr(d,cx,B-h-29,f"{v:g}"+("%" if "%" in title else ""),F(26));ctr(d,cx,B+12,("M1示範窗","M4示範窗")[j],F(24))
 ctr(d,800,540,"全部為SAMPLE/SYNTHETIC；M1、M4是固定30日代碼窗，非真實日期。",F(26),red);im.save(path,dpi=(180,180))
f1=assets/"figure1_market_sample.png";f2=assets/"figure2_customer_sample.png"
chart_market(f1)
chart(f2,"圖2｜顧客價值分題與複合均分",[("需求效用",st["usefulness"][2],st["usefulness"][3]),("使用便利",st["convenience"][2],st["convenience"][3]),("價格/時間",st["benefit_vs_price_time"][2],st["benefit_vs_price_time"][3]),("複合",meanA,meanB)],1,5,"24組合成配對×3題；不是協會實績或真實回覆率")

d=Document();N="1F384E";TE="2B7976";PA="F3F7F8";RE="A44B41";WH="FFFFFF";IN="2D3741"
def fr(r,size=10.5,b=False,col=IN):
 r.font.name="Microsoft JhengHei";r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"Microsoft JhengHei");r._element.rPr.rFonts.set(qn("w:ascii"),"Microsoft JhengHei");r._element.rPr.rFonts.set(qn("w:hAnsi"),"Microsoft JhengHei");r.font.size=Pt(size);r.bold=b;r.font.color.rgb=RGBColor.from_string(col)
for nm,sz,col in [("Normal",10.5,IN),("Title",25,N),("Heading 1",16,N),("Heading 2",13,TE)]:
 s=d.styles[nm];s.font.name="Microsoft JhengHei";s._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft JhengHei");s.font.size=Pt(sz);s.font.color.rgb=RGBColor.from_string(col)
def p(t,size=10.5,b=False,col=IN,center=False,style=None):
 x=d.add_paragraph(style=style);x.paragraph_format.space_after=Pt(5);x.paragraph_format.line_spacing=1.13
 if center:x.alignment=WD_ALIGN_PARAGRAPH.CENTER
 fr(x.add_run(str(t)),size,b,col);return x
def h(t,l=1):return p(t,16 if l==1 else 13,True,N if l==1 else TE,False,"Heading "+str(l))
def bullet(t):
 x=d.add_paragraph(style="List Bullet");x.paragraph_format.space_after=Pt(3);fr(x.add_run(t))
def shade(c,col):
 e=OxmlElement("w:shd");e.set(qn("w:fill"),col);c._tc.get_or_add_tcPr().append(e)
def tbl(cap,heads,rows,fs=9.2,widths=None):
 if isinstance(fs,(list,tuple)):widths,fs=fs,9.2
 fs=max(float(fs),9.5)
 p(cap,size=10.5,b=True,col=TE);t=d.add_table(rows=1,cols=len(heads));t.style="Table Grid";t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");t.rows[0]._tr.get_or_add_trPr().append(e)
 for j,v in enumerate(heads):
  c0=t.rows[0].cells[j];shade(c0,N);z=c0.paragraphs[0];z.alignment=WD_ALIGN_PARAGRAPH.CENTER;z.paragraph_format.space_after=Pt(0);fr(z.add_run(str(v)),fs,True,WH)
 for i,row in enumerate(rows):
  cs=t.add_row().cells;t.rows[-1]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
  for j,v in enumerate(row):
   c0=cs[j];c0.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   if i%2:shade(c0,PA)
   z=c0.paragraphs[0];z.paragraph_format.space_after=Pt(0);z.paragraph_format.line_spacing=1.0;fr(z.add_run(str(v)),fs)
 if widths:
  for row in t.rows:
   for j,w in enumerate(widths):row.cells[j].width=Inches(w)
 d.add_paragraph();return t
def note(t):
 t0=d.add_table(rows=1,cols=1);c0=t0.cell(0,0);shade(c0,"FCEEEB");fr(c0.paragraphs[0].add_run(t),10.5,True,RE);d.add_paragraph()
def fig(path,cap,alt):
 x=d.add_paragraph();x.alignment=WD_ALIGN_PARAGRAPH.CENTER;sh=x.add_run().add_picture(str(path),width=Inches(6.55))
 try:sh._inline.docPr.set("descr",alt)
 except:pass
 p(cap,size=10.5,col="52616D",center=True)
def fields(par,code):
 r=par.add_run();fr(r,8.5,False,"52616D")
 for typ,val in [("begin",None),("instr",code),("sep",None),("text","1"),("end",None)]:
  if typ in ("begin","sep","end"):
   z=OxmlElement("w:fldChar");z.set(qn("w:fldCharType"),{"begin":"begin","sep":"separate","end":"end"}[typ]);r._r.append(z)
  elif typ=="instr":
   z=OxmlElement("w:instrText");z.set(qn("xml:space"),"preserve");z.text=" "+val+" ";r._r.append(z)
  else:z=OxmlElement("w:t");z.text=val;r._r.append(z)
def running(s):
 x=s.header.paragraphs[0];x.alignment=WD_ALIGN_PARAGRAPH.RIGHT;fr(x.add_run("原始點協會｜候選示範　0127 / I19-E16"),8.3,False,"52616D")
 x=s.footer.paragraphs[0];x.alignment=WD_ALIGN_PARAGRAPH.CENTER;fr(x.add_run("SAMPLE / SYNTHETIC ONLY｜非協會實績｜第 "),8.3,True,RE);fields(x,"PAGE");fr(x.add_run(" 頁 / 共 "),8.3);fields(x,"NUMPAGES");fr(x.add_run(" 頁"),8.3)
for s in d.sections:
 s.page_width=Inches(8.27);s.page_height=Inches(11.69);s.top_margin=Inches(.65);s.bottom_margin=Inches(.62);s.left_margin=Inches(.68);s.right_margin=Inches(.68);running(s)

# Cover / decision first
x=d.add_paragraph(style="Title");x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.space_before=Pt(55);fr(x.add_run("目標市場與顧客價值\n量化效益分析表"),25,True,N)
p("訓練機構評核佐證候選｜完整合成資料演算示範",13,True,TE,True)
note("SAMPLE / SYNTHETIC ONLY｜本檔企業、顧客、課程情境、期間、數字及結論均為合成示範，不是原始點協會或會員實績；尚未正式接受，禁止作為真實成果送審。")
tbl("文件身分與用途",["欄位","內容"],[
["文件ID","0127"],["官方要求","I19-E16；唯一Blueprint 0127"],["官方位置","查核佐證資料表第8頁"],["文類","分析／評估報告：市場與顧客價值量化分析"],["讀者","協會主管、訓練行政、資料覆核者"],["候選狀態","完整內容候選；獨立語意及全頁Word視覺驗收待完成"],["REAL資料","無協會課程、事業、顧客、交易、問卷、核准或成果原件"]
],[1.55,4.85])
p("操作：先核來源、分母與期間，再重算表2和表4；權限未明、資料缺漏或期間不可比即標HOLD，不插補。")
d.add_page_break()
h("執行摘要")
p("本報告分開量化目標市場與顧客價值。市場面看事前定義客群的去重觸及及同窗首購轉換；顧客面看配對顧客對需求效用、使用便利、相對價格／時間效益的三題評分。全部結果是SAMPLE CSV的演算，不能證明原始點協會已提升。")
tbl("雙構面摘要（SAMPLE / SYNTHETIC ONLY）",["構面","M1","M4","變化","可支持／不可支持"],[
["合格潛客觸及","32人","48人","+16；+50.00%","示範觸及提高；非市場規模/市占"],
["首購轉換","8/32=25.00%","15/48=31.25%","+6.25pp；相對+25.00%","同窗首購比例；非營收/獲利"],
["顧客價值","234/72=3.25","288/72=4.00","+0.75；相對+23.08%","主觀評分；非客觀福利"]
],[1.15,1.05,1.05,1.25,1.9],9)
note("合成資料兩構面均上升。協會REAL成果證據仍為UNKNOWN；沒有真實來源、比較設計與混雜因素控制，不能寫成「訓練造成提升」。")
h("30秒操作指引",2)
for x in ["用途：查看市場觸及/首購及配對顧客感知價值。","時點：僅於兩期定義、期間、資料抽取與問卷版本可比時比較。","操作：核來源、分母、N，再計算；缺件即HOLD。","輸出：記差異、限制、建議角色和KPI；姓名/期限由協會核定。","界線：全為SAMPLE/SYNTHETIC，無REAL原件或正式核准。"]:bullet(x)
h("1. 官方要求與範圍")
p("官方佐證資料表第8頁 I19-E16 原文：「有量化之效益可呈現目標市場及顧客價值均有一定程度提升的效果。」兩構面均需獨立呈現；不得以一項代另一項，也不得擴大為市場總規模、市占、薪資或個人工作績效。")
tbl("納入與排除",["本件納入","排除/不得替代"],[
["合格潛客觸及、首購率、三題配對顧客評價、公式與限制。","0126/I19-E14訓後薪資；本件無個人薪資。"],
["M1/M4固定30日示範窗、資料品質、REAL替換。","0128/I19-E17個人訓後績效；不得以市場數字替代。"],
["決策建議、覆核與退件門檻。","歷史0127正文；禁止當模板。"]
],[3.2,3.2])
h("2. 指標定義與資料品質")
p("合成課程情境為「數位顧客經營與目標市場拓展示範工作坊」，不表示原始點協會或會員曾開課。市場CSV含8個虛構事業；顧客CSV含24個虛構配對ID、每案3人。M1/M4是固定30日代碼窗，非真實日期。")
tbl("表1　指標字典",["指標","定義/公式","分母/單位","資料/缺漏","限制"],[
["合格觸及","符合事前條件、去重且有接觸紀錄","人；按窗分開","REAL需名冊及日誌","非總市場/市占"],
["觸及變化","(M4−M1)/M1×100%","M1觸及","M1=0時不定義","受季節/宣傳影響"],
["首購率","同窗首購/同窗合格潛客","潛客；%","交易日期須可追溯","非營收/利潤/續購"],
["轉換差","M4率−M1率","百分點","各期先除分母","不可混稱相對百分比"],
["顧客價值","三題1–5分逐題均分及複合","有效題N","空白不計0、不插補","主觀感知"],
["複合均分","三題總分/有效題數","每期最多72題","核配對與量尺版本","不是因果效果"]
],[1.0,1.55,1.1,1.4,1.2],8.7)
for x in ["8案合計：M1潛客32、M4 48、首購8與15；逐案首購均≤同窗潛客。","24個不重複合成顧客ID，每案3人；六個分數欄均在1–5，資料集無缺值。","每期24×3=72題，總分234及288；完整值只屬資料集構造，不是REAL回覆率。","REAL最低完整配對門檻須預先核定；來源未授權門檻數值，本報告不自行設值。"]:bullet(x)
h("3. 目標市場結果")
p("逐案加總後以總分子除總分母。總體轉換率不是各案率的簡單平均；觸及數不是市場總規模。")
fig(f1,"圖1　市場觸及人數與首購轉換率分圖呈現；兩者單位不同。SAMPLE/SYNTHETIC。","M1/M4觸及32/48人，首購率25/31.25%，分開尺度。")
mr=[]
for r in m:
 x=r[mk[0]];y=r[mk[1]];u=r[mk[2]];v=r[mk[3]]
 mr.append([r["case_id"]+" / "+r["replacement_key"],r["segment_SAMPLE"],f"{x} / {u} ({u/x*100:.1f}%)",f"{y} / {v} ({v/y*100:.1f}%)"])
tbl("表2　8案市場明細（潛客 / 首購 / 率）",["合成案例/鍵","合成類別","M1","M4"],mr,8.8,[1.5,1.6,1.55,1.55])
tbl("表2A　總體重算",["項目","計算","結果","控制"],[
["M1觸及","3+4+2+5+3+4+6+5","32人","8案加總"],
["M4觸及","5+6+4+7+6+6+8+6","48人","8案加總"],
["觸及變化","(48−32)/32×100%","50.00%","+16人"],
["M1首購率","8/32×100%","25.00%","總體加權率"],
["M4首購率","15/48×100%","31.25%","總體加權率"],
["百分點差","31.25−25.00","+6.25pp","非相對%"],
["相對率變化","6.25/25×100%","+25.00%","只作合成描述"]
],[1.2,2.2,1.05,1.85],8.9)
p("S-O07個案轉換率由33.33%降至25.00%；其餘合成案例持平或提高。總體上升不能掩蓋個別下滑。")
h("4. 顧客價值問卷與分析")
p("Q1需求效用：「此產品／服務有助於解決我原有需求。」Q2便利：「取得／使用此產品／服務的流程便利。」Q3相對效益：「相較付出的價格或時間，我感受到的效益值得。」量尺1=非常不同意/效益極低，5=非常同意/效益很高；空白不得記0。")
tbl("表3　逐題及複合均分",["面向","M1總分/N/均分","M4總分/N/均分","差"],[
["Q1 需求效用",f'{st["usefulness"][0]}/24/{st["usefulness"][2]:.2f}',f'{st["usefulness"][1]}/24/{st["usefulness"][3]:.2f}',f'+{st["usefulness"][4]:.2f}'],
["Q2 便利",f'{st["convenience"][0]}/24/{st["convenience"][2]:.2f}',f'{st["convenience"][1]}/24/{st["convenience"][3]:.2f}',f'+{st["convenience"][4]:.2f}'],
["Q3 相對效益",f'{st["benefit_vs_price_time"][0]}/24/{st["benefit_vs_price_time"][2]:.2f}',f'{st["benefit_vs_price_time"][1]}/24/{st["benefit_vs_price_time"][3]:.2f}',f'+{st["benefit_vs_price_time"][4]:.2f}'],
["三題複合",f"{totA}/72/{meanA:.2f}",f"{totB}/72/{meanB:.2f}",f"+{meanB-meanA:.2f}"]
],8.8,[1.5,2.05,2.05,1.0])
fig(f2,"圖2　三個分題與複合均分；縱軸1–5分，全部SAMPLE/SYNTHETIC。","M1/M4需求效用、便利、相對效益與複合均分。")
tbl("表3A　複合均分重算",["步驟","M1","M4","差"],[
["配對×題數","24×3=72","24×3=72","同一合成配對組"],
["總分","234","288","+54"],
["均分","234/72=3.25","288/72=4.00","+0.75"],
["相對變化","—","—","0.75/3.25=23.08%"]
],[1.45,1.55,1.55,1.75],9)
h("5. 雙構面結論與限制")
tbl("表4　可支持的範圍",["構面","合成觀察","支持範圍","不可主張/待補"],[
["市場觸及","32→48；+50%","合成觸及提高","非市場規模/市占；REAL要名冊/接觸"],
["首購","25→31.25%；+6.25pp","同窗合成比例提高","非利潤/因果；REAL要交易和分母"],
["顧客價值","3.25→4.00；+0.75","主觀評分提高","非客觀福利；REAL要問卷/同意/N"],
["協會成果","無REAL原件","只示範計算","協會已改善或訓練造成=UNKNOWN"]
],[1.05,1.45,1.75,2.05],8.7)
for x in ["所有數據均SYNTHETIC；協會實績證據UNKNOWN。","價格、季節、宣傳、通路、選樣及資料記錄變化可能影響兩期差異；無比較組或混雜因子控制。","只能說合成資料兩期不同，不可主張訓練因果、協會成效、市占或財務收益。","REAL若任一構面、分母、期間或合法來源缺失，整體HOLD，不以另一構面補足。"]:bullet(x)
h("6. 品質、缺漏及未回覆控制")
tbl("表5　檢核與異常處置",["檢核","通過條件","不符處置"],[
["客群","事前資格、排除與去重規則一致","HOLD，重建範圍"],
["分子/分母","同窗首購可追溯且≤潛客","未解前不發布率"],
["問卷","同版題目/量尺；逐題報N","空白NA、不插補"],
["配對","受控去識別ID；門檻預先核定","配對不足即HOLD"],
["期間","起訖日及定義可比","不可比不下趨勢結論"],
["未回覆","有邀請母數、回覆N與處置","母數缺則回覆率UNKNOWN"],
["權限","合法同意、存取與最小化","權限不明即停個資連結"]
],[1.05,3.0,2.55],8.8)
p("REAL公開版只保留必要彙總。沒有合法同意或資料權限時，不連結個人層級資料；使用經核准匿名彙總或維持HOLD。")
h("7. 建議行動（尚未核定/執行）")
tbl("表6　主管可採取的下一步",["建議","操作/KPI","角色","期限","停止條件"],[
["固定客群/視窗","記資格、30日窗、唯一鍵；定義一致率","市場資料承辦(待指派)","核定後","缺母數停比"],
["核驗首購","接觸ID、交易日抽查；首購≤潛客","資料管理(待指派)","核定後","來源不符HOLD"],
["配對調查","核題目、量尺、邀請母數、配對門檻、同意","調查/隱私角色(待指派)","發放前","逐題N/配對不足暫緩"],
["查異質性","逐案檢視下滑與替代原因","分析覆核(待指派)","分析後","無來源不推定因果"],
["送出前覆核","核原件、公式、清冊、個資、版本","主管/權責者(待指派)","送出前","原件未齊維持SAMPLE"]
],[1.05,2.2,1.3,1.0,1.7],8.4)
h("8. REAL替換與退件門檻")
for x in ["取得協會指定之資料權責者、真實課程/班次、核定範圍與日期。","先凍結客群、首購、問卷版本/量尺、配對門檻、缺漏規則。","依法取用最小必要去識別市場接觸/交易及問卷，保存同意和邀請/回覆母數。","另一位授權覆核者抽查分子、分母、去重、配對和結論定位。","逐鍵更新附錄清冊；原件、權限或矛盾未解，保留SAMPLE/UNKNOWN。"]:bullet(x)
tbl("表7　REJECT / HOLD判定",["狀態","觸發","處置/解除"],[
["REJECT","缺一構面；錯公式/分母；不可重算；合成冒稱REAL；誇大因果/市占/財務","停用錯誤結論，修正來源與計算後重審"],
["HOLD","權限、原件、期間、客群、回覆母數、配對門檻或覆核未齊","列缺件/責任/解除條件；原件到位重算"],
["SAMPLE ONLY","只有合成資料","只示範，不升REAL或正式接受"],
["候選QA條件","映射、兩構面、重算、清冊、逐頁視覺及獨立語意皆完成","只綁定exact SHA；QA不等於正式接受"]
],[1.0,2.65,2.95],8.8)
h("9. 官方追溯與相鄰要求")
tbl("表8　需求—正文定位—目前證據",["官方需求/檢核","正文位置","支持程度","REAL缺件"],[
["I19-E16第8頁原文","本節、表2/3/4/7、圖1/2","完整合成示範；協會成果UNKNOWN","客群、交易、問卷、覆核原件"],
["市場量化","第3節/表2/2A/圖1","8案可重算，非市占","名冊、定義、接觸及交易"],
["顧客量化","第4節/表3/圖2/附錄A","24配對×3題合成","核定問卷、同意、母數、有效N"],
["限制/決策","第5–8節","限制、行動、拒收條件具體","真實期間、決議與追蹤"]
],[1.65,1.65,2.0,1.65],8.5)
tbl("鄰件排除",["文件","專屬焦點","本件不含"],[
["0126/I19-E14","訓後薪資變化","個人薪資與薪酬時序"],
["0128/I19-E17","個人訓後績效","職務、個人目標與主管評分"]
],[1.7,2.0,3.3])
p("正式核准、簽章、送審及最終接受均未發生。建置收據不等於獨立驗收；所有QA須綁定最終SHA。")
h("10. 來源與版本")
tbl("來源清冊",["來源","SHA/定位","使用界線"],[
["官方查核佐證PDF","805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709；第8頁","官方原文，非REAL成果"],
["BLUEPRINT_CURRENT","406d0495365d0970dcf8113a42842663b8a54372a4e204cc492b7c3f8156ef3e；唯一列SHA 8e0f7650307842de1c98513f9a056db74c1169e271249220c7e206b7d70d2a77","需求拆解；歷史正文禁止"],
["142-ID crosswalk","1ddfbb608bcad25170683ddb42a1dbb27bafa6f1f35ff0498f1b8d75b24f0679","I19-E16→0127"],
["R01來源包/Blueprint/CSV","各檔hash見本地技術收據","CSV全為SAMPLE，無REAL原件"]
],[1.5,3.75,1.8],8.5)
s=d.add_section(WD_SECTION.NEW_PAGE);s.orientation=WD_ORIENT.LANDSCAPE;s.page_width=Inches(11.69);s.page_height=Inches(8.27);s.top_margin=Inches(.48);s.bottom_margin=Inches(.48);s.left_margin=Inches(.52);s.right_margin=Inches(.52);running(s)
h("附錄A　24組顧客配對原始資料")
p("每列保留兩期三題原始值；單人均分=(Q1+Q2+Q3)/3，差=B−A。全部SYNTHETIC。")
rows=[]
for r in c:
 x=[r[k] for k in ck[:3]];y=[r[k] for k in ck[3:]]
 rows.append([r["respondent_id_SAMPLE"],r["case_id_SAMPLE"],*x,f"{sum(x)/3:.2f}",*y,f"{sum(y)/3:.2f}",f"{sum(y)/3-sum(x)/3:+.2f}",r["replacement_key"]])
tbl("附錄表A1　24列配對問卷（SAMPLE / SYNTHETIC ONLY）",["顧客ID","事業","A-Q1","A-Q2","A-Q3","A均分","B-Q1","B-Q2","B-Q3","B均分","差","鍵"],rows,8.4)
p("合計M1=234/72=3.25，M4=288/72=4.00；分題均分變化Q1 +1.00、Q2 +0.75、Q3 +0.50。")
h("附錄B　逐項合成替換清冊")
reg=[]
for r in m:reg.append([r["replacement_key"],"合成事業/市場列",f'{r["case_id"]}；{r["segment_SAMPLE"]}；M1 {r[mk[0]]}/{r[mk[2]]}；M4 {r[mk[1]]}/{r[mk[3]]}',"附錄C1市場列","去識別事業、客群規則、接觸/首購原件及授權"])
for r in c:reg.append([r["replacement_key"],"合成顧客配對",f'{r["respondent_id_SAMPLE"]}/{r["case_id_SAMPLE"]}；A {r[ck[0]]},{r[ck[1]]},{r[ck[2]]}；B {r[ck[3]]},{r[ck[4]]},{r[ck[5]]}',"附錄A同ID列","去識別配對、原始問卷、同意、邀請/回覆N及權限"])
for row in [
["RR-COURSE-01","虛構課程","數位顧客經營與目標市場拓展示範工作坊","封面/範圍","真實課程、班次、核定日期"],
["RR-WINDOW-01/02","虛構期間","M1/M4固定30日代碼，無真實日期","表2–4","真實起訖日、時區、可比證據"],
["RR-ITEM-Q1/Q2/Q3","合成題目","需求效用、便利、價格/時間效益","第4節","核定問卷版本及發放原件"],
["RR-SCALE-01","合成量尺","1–5；空白非0","第4節","核定量尺和缺漏規則"],
["RR-FORMULA-01","合成公式","觸及變化、首購率、百分點差","表1/2A","資料欄位、去重、首購定義"],
["RR-FORMULA-02","合成公式","逐題與複合均分","表3/附錄A","原問卷、有效N、配對規則"],
["RR-RESULT-01","衍生市場結果","32→48；25→31.25%；+6.25pp","摘要/表2","REAL名冊與交易重算"],
["RR-RESULT-02","衍生顧客結果","234→288；3.25→4.00；+0.75","摘要/表3","REAL逐題、N、配對覆核"],
["RR-INTERPRETATION-01","衍生解讀","非因果/非市占/非協會實績","摘要/第5節","REAL設計及替代因素"],
["RR-ACTION-01","建議行動","角色/期限待指定、尚未執行","表6","協會決議、承辦與追蹤"],
["RR-FIGURE-01/02","合成圖表","所有數值、長條與標籤","圖1/2","REAL核准資料重繪"]
]:reg.append(row)
tbl("附錄表B1　原始列與衍生事實替換鍵",["鍵","類型","目前SAMPLE值","定位","REAL原件/控制"],reg,7.9)
with register.open("w",encoding="utf-8-sig",newline="") as f:
 w=csv.writer(f);w.writerow(["replacement_key","type","sample_value","location","real_replacement_required"]);w.writerows(reg)
h("附錄C　候選覆核清單")
tbl("附錄表C1",["檢查","條件","目前狀態"],[
["身分","0127/I19-E16/唯一Blueprint","來源映射確認"],
["雙構面","市場與顧客各自量化","完整SAMPLE示範"],
["數學","分母/百分點/均分可重算","摘要計算一致；外部QA待完成"],
["真實性","無原件不稱實績","SAMPLE/REAL分界明確；REAL UNKNOWN"],
["獨立QA","語意及Word逐頁視覺綁exact SHA","PENDING"],
["正式接受","協會權責程序","NOT GRANTED"]
],[1.2,4.0,3.1],8.6)

for i,s in enumerate(d.sections):
 if i==0:s.orientation=WD_ORIENT.PORTRAIT;s.page_width=Inches(8.27);s.page_height=Inches(11.69);s.left_margin=Inches(.68);s.right_margin=Inches(.68)
 else:s.orientation=WD_ORIENT.LANDSCAPE;s.page_width=Inches(11.69);s.page_height=Inches(8.27);s.left_margin=Inches(.52);s.right_margin=Inches(.52)
d.core_properties.title="目標市場與顧客價值量化效益分析表";d.core_properties.subject="0127/I19-E16；SAMPLE/SYNTHETIC候選；未正式接受";d.core_properties.author="TTQS ONE候選施工"
d.save(docx)
# Linear manuscript includes every paragraph and table in order.
lines=["# 目標市場與顧客價值量化效益分析表","","0127 / I19-E16｜SAMPLE/SYNTHETIC候選 R02；未正式接受","","所有示範事實均為合成資料。",""]
for child in d.element.body.iterchildren():
 if child.tag==qn("w:p"):
  z=Paragraph(child,d);txt=z.text.strip()
  if txt:
   style_name=z.style.name if z.style else ""
   lines += [("## " if style_name.startswith("Heading") else "")+txt,""]
 elif child.tag==qn("w:tbl"):
  t=Table(child,d)
  for row in t.rows:lines.append("| "+" | ".join(cell.text.replace("|","\\|").replace("\n"," / ") for cell in row.cells)+" |")
  lines.append("")
md.write_text("\n".join(lines),encoding="utf-8")
rec={"project_id":"TTQS_ONE","doc_id":"0127","requirement":"I19-E16","official_page":8,"candidate":"CANDIDATE_NOT_FORMALLY_ACCEPTED","real_evidence":"UNKNOWN_UNTIL_SOURCE","builder":"Codex deterministic layout repair from pinned R01 source; no new model writing","docx_sha256":sha(docx.read_bytes()),"docx_bytes":docx.stat().st_size,"manuscript_sha256":sha(md.read_bytes()),"replacement_register_sha256":sha(register.read_bytes()),"figure1_sha256":sha(f1.read_bytes()),"figure2_sha256":sha(f2.read_bytes()),"source_hashes":{f:sha((src/f).read_bytes()) for f in pins},"math":{"M1_reach":a,"M4_reach":b,"reach_change_pct":50,"M1_purchases":ca,"M4_purchases":cb,"M1_rate_pct":25,"M4_rate_pct":31.25,"delta_pp":6.25,"customer_pairs":24,"items_per_period":72,"M1_total":totA,"M4_total":totB,"M1_mean":meanA,"M4_mean":meanB,"delta":meanB-meanA,"relative_pct":round((meanB-meanA)/meanA*100,2),"items":st},"replacement_rows":len(reg),"render":"PENDING","independent_semantic":"PENDING","formal_acceptance":"NOT_GRANTED","built_at":datetime.now().astimezone().isoformat()}
receipt.write_text(json.dumps(rec,ensure_ascii=False,indent=2),encoding="utf-8")
print(json.dumps({"docx":str(docx),"sha256":rec["docx_sha256"],"bytes":rec["docx_bytes"],"manuscript_sha256":rec["manuscript_sha256"],"register_sha256":rec["replacement_register_sha256"],"figure_sha256":[rec["figure1_sha256"],rec["figure2_sha256"]],"market":rec["math"],"replacement_rows":len(reg)},ensure_ascii=False))
