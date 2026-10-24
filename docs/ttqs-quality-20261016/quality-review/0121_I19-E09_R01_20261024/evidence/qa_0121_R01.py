from pathlib import Path
from zipfile import ZipFile
import re,json,hashlib,statistics
import pypdfium2 as pdfium
from docx import Document
ROOT=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0121_I19-E09_SOURCE_PREP_R01_20261024')
DOC=ROOT/'candidate_R01'/'0121_學員滿意度與學習評量報告_R01.docx'
PDF=ROOT/'candidate_R01'/'0121_R01_word_render.pdf'
RENDER=ROOT/'candidate_R01'/'render_R01_20261024_1340'
MD=ROOT/'manuscript_R01_revised.md'
BLUEPRINT_SHA='99320fdc9078615ad04087c43bcbd22c5bbf751a22127c51e048e909726a5fca'
SOURCE_PACKET_SHA='fbeb11b96648fff29a49a0fca74bf0cc3b3d1d29934f07dfaa3cf7e410a557fd'
MAPPING_SHA='fc9e9704748e58847c14c707078acf66ed07f024e2b83c67201b2372ead8208d'
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def parse_table(lines):
 out=[]
 for line in lines:
  if not line.strip().startswith('|'):continue
  c=[x.strip() for x in line.strip().strip('|').split('|')]
  if c and all(re.fullmatch(r':?-{3,}:?',x.replace(' ','')) for x in c):continue
  out.append(c)
 return out
def table_after(text,marker):
 lines=text.splitlines();st=next(i for i,x in enumerate(lines) if marker in x)
 st=next(i for i in range(st+1,len(lines)) if lines[i].strip().startswith('|'))
 en=st
 while en+1<len(lines) and lines[en+1].strip().startswith('|'):en+=1
 return parse_table(lines[st:en+1])
def norm(x):
 x=re.sub(r'\\text\{([^{}]*)\}',r'\1',x)
 return re.sub(r'[*`$]','',x).replace('{','').replace('}','').replace(r'\_','_').strip()
md=MD.read_text(encoding='utf-8')
# Satisfaction response matrix and item summary must reconcile independently.
sur=table_after(md,'### 6.2 代碼化學員滿意度逐筆填答矩陣')
rows={norm(r[0]):r for r in sur[1:] if re.fullmatch(r'P(?:0[1-9]|10)',norm(r[0]))}
assert len(rows)==10,('survey row count',len(rows))
item_summary=table_after(md,'### 9.1 滿意度各題項統計摘要表')
sumrows={norm(r[0]):r for r in item_summary[1:] if re.fullmatch(r'Q\d{2}',norm(r[0]))}
assert len(sumrows)==10,('summary rows',len(sumrows))
items=[]
for j in range(1,11):
 vals=[int(norm(rows[f'P{k:02d}'][j])) for k in range(1,11) if norm(rows[f'P{k:02d}'][j]).isdigit()]
 s=sum(vals);n=len(vals);pos=sum(v>=4 for v in vals);q=f'Q{j:02d}';r=sumrows[q]
 assert int(norm(r[2]))==n and int(norm(r[4]))==s and int(norm(r[7]))==pos,(q,n,s,pos,r)
 shown_mean=float(norm(r[5]));shown_rate=float(norm(r[8]).split('%')[0])
 assert round(s/n+1e-9,2)==shown_mean and round(pos/n*100+1e-9,1)==shown_rate,(q,shown_mean,shown_rate)
 items.append({'item':q,'n':n,'sum':s,'mean':round(s/n,2),'positive_n':pos,'positive_rate_pct':round(pos/n*100,1)})
assert [x['n'] for x in items]==[9,9,9,9,9,9,8,8,8,8]
assert [x['sum'] for x in items]==[40,39,40,34,31,37,34,33,35,32]
# Composite mean from complete responses P01-P08 only.
composite=[sum(int(norm(rows[f'P{k:02d}'][j])) for j in range(1,11)) for k in range(1,9)]
assert composite==[46,40,42,38,48,38,44,37],composite
assert sum(composite)==333 and 8*10==80 and sum(composite)/80==4.1625
# Skills pre/post source tables and each row sum.
pre=table_after(md,'### 8.1 技能實作前後測逐筆評分資料矩陣')
post=table_after(md,'**後測與配對結果**')
pr={norm(r[0]):r for r in pre[1:] if re.fullmatch(r'P(?:0[1-9]|10)',norm(r[0]))}
po={norm(r[0]):r for r in post[1:] if re.fullmatch(r'P(?:0[1-9]|10)',norm(r[0]))}
assert len(pr)==10 and len(po)==10
pre_tot={};post_tot={};paired=[]
for pid in [f'P{k:02d}' for k in range(1,11)]:
 v=[norm(x) for x in pr[pid][1:6]];total=norm(pr[pid][6])
 if pid!='P10':
  assert all(x.isdigit() for x in v) and sum(map(int,v))==int(total),(pid,v,total)
  pre_tot[pid]=int(total)
 w=[norm(x) for x in po[pid][1:6]];total2=norm(po[pid][6])
 if pid in [f'P{k:02d}' for k in range(1,9)]:
  assert all(x.isdigit() for x in w) and sum(map(int,w))==int(total2),(pid,w,total2)
  assert int(norm(po[pid][7]).replace('+',''))==int(total2)-pre_tot[pid]
  paired.append(pid);post_tot[pid]=int(total2)
assert sum(pre_tot[p] for p in paired)==310 and sum(post_tot.values())==690
assert sum(post_tot[p]-pre_tot[p] for p in paired)==380
assert (310/8,690/8,380/8)==(38.75,86.25,47.5)
pass_n=sum(v>=70 for v in post_tot.values());assert pass_n==7 and len(paired)==8
# Per-task skill totals and full-mark counts.
tasktab=table_after(md,'### 9.3.3 各項技能任務分項成長比對表')
per_task=[]
for k in range(5):
 pre_sum=sum(int(norm(pr[pid][k+1])) for pid in paired)
 post_sum=sum(int(norm(po[pid][k+1])) for pid in paired)
 perfect=sum(int(norm(po[pid][k+1]))==20 for pid in paired)
 r=tasktab[k+1]
 assert abs(float(norm(r[1]).split('=')[-1]) - pre_sum/8)<1e-8
 assert abs(float(norm(r[2]).split('=')[-1]) - post_sum/8)<1e-8
 assert abs(float(norm(r[3]).replace('+',''))-(post_sum-pre_sum)/8)<1e-8
 assert f'{perfect}/8' in norm(r[4])
 per_task.append({'task':f'T0{k+1}','pre_sum':pre_sum,'post_sum':post_sum,'pre_mean':pre_sum/8,'post_mean':post_sum/8,'mean_delta':(post_sum-pre_sum)/8,'perfect_n':perfect})
assert (sum(pre_tot.values()),sum(post_tot.values()),sum(post_tot.values())-sum(pre_tot[p] for p in paired))==(320,690,380)
# Flow reconciliation, replacement keys, and real/sample boundary.
flow=table_after(md,'### 6.1 樣本回收與有效性流向勾稽表')
flowtext=' '.join(' '.join(norm(x) for x in r) for r in flow)
assert 'N_COMPOSITE (8)' in flowtext and 'N_PARTIAL (1)' in flowtext and 'N_NON_RESP (1)' in flowtext
reg=table_after(md,'### 12.2 逐項事實替換登記表')
keys=[norm(r[0]) for r in reg[1:]]
assert len(keys)==len(set(keys)) and len(keys)>=18,(len(keys),len(set(keys)))
assert 'SAMPLE / SYNTHETIC' in md and 'REAL evidence 未建立' in md and 'NOT_GRANTED' in md
# DOCX archive, extracted text and native render checks.
with ZipFile(DOC) as z:bad=z.testzip();assert bad is None,bad
word=Document(str(DOC));text='\n'.join(p.text for p in word.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in r.cells) for t in word.tables for r in t.rows)
for bad in ['\\frac','N_RETURNN_INVITED','fracX_i','dots','來源稿原列','課堂發放之無記名']:
 assert bad not in text,bad
assert 'I19-E09' in text and '0121' in text and 'SAMPLE / SYNTHETIC' in text and 'NOT_GRANTED' in text
assert len(word.tables)==17
pdf=pdfium.PdfDocument(str(PDF));assert len(pdf)==23,len(pdf)
imgs=sorted(RENDER.glob('page-*.png'));assert len(imgs)==23,len(imgs)
page_meta=[]
for n,p in enumerate(imgs,1):
 im=__import__('PIL.Image',fromlist=['Image']).open(p)
 assert im.size==(1403,993),(n,im.size)
 page_meta.append({'page':n,'png_sha256':sha(p),'bytes':p.stat().st_size,'size':list(im.size)})
assert pdf[len(pdf)-1].get_textpage().get_text_range().strip(), 'last page is blank'
record={
 'document_id':'0121','official_requirement_id':'I19-E09','official_source':'official_source_0108_r11.pdf p.8',
 'source_packet_sha256':SOURCE_PACKET_SHA,'unique_blueprint_sha256':BLUEPRINT_SHA,'mapping_sha256':MAPPING_SHA,
 'candidate_path':str(DOC),'candidate_sha256':sha(DOC),'candidate_bytes':DOC.stat().st_size,
 'manuscript_sha256':sha(MD),'word_zip_integrity':'PASS','docx_tables':len(word.tables),
 'word_native_render':'PASS','pdf_path':str(PDF),'pdf_sha256':sha(PDF),'pdf_pages':len(pdf),
 'visual_render_dir':str(RENDER),'visual_render_pages':len(imgs),'visual_render_size_px':[1403,993],
 'visual_page_sha256':page_meta,'visual_inspection':'23/23 pages inspected; unchanged rasters were hash-matched to prior inspected pages; page 24 blank removed.',
 'math_checks':{'survey_items':items,'complete_response_totals':composite,'composite_sum':333,'composite_valid_answers':80,'composite_mean':4.1625,'pre_sum_paired':310,'pre_mean_paired':38.75,'post_sum_paired':690,'post_mean_paired':86.25,'paired_delta_sum':380,'paired_mean_delta':47.5,'pass_n':pass_n,'paired_n':len(paired),'per_task':per_task,'flow_balance':'10=8+1+1'},
 'register_keys_count':len(keys),'register_keys_unique':True,
 'sample_real_boundary':'PASS - all examples synthetic; REAL evidence not provided; formal acceptance NOT_GRANTED.',
 'technical_layout_status':'PASS for the listed checks','independent_semantic_review':'PENDING exact SHA; this receipt is not semantic acceptance.',
 'formal_acceptance':'NOT_GRANTED'
}
out=ROOT/'candidate_R01'/'technical_visual_QA_R01.json';out.write_text(json.dumps(record,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'receipt':str(out),'receipt_sha256':sha(out),'candidate_sha256':record['candidate_sha256'],'pdf_sha256':record['pdf_sha256'],'pages':record['pdf_pages'],'survey_n':[x['n'] for x in items],'survey_sums':[x['sum'] for x in items],'composite':record['math_checks']['composite_mean'],'paired_change':record['math_checks']['paired_mean_delta'],'pass_rate':f'{pass_n}/8','per_task':per_task,'zip_integrity':'PASS','independent_semantic_review':'PENDING'},ensure_ascii=False,indent=2))
