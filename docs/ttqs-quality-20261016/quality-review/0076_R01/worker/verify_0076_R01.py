from __future__ import annotations
import csv, hashlib, json, re, zipfile
from pathlib import Path
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
BODY=ROOT/"candidate/0076_R01_body.md"
DOCX=ROOT/"candidate/0076_R01.docx"
REGISTER=ROOT/"inputs/SYNTHETIC_REPLACEMENT_REGISTER_0076_R01.csv"
DATA=ROOT/"inputs/SYNTHETIC_ASSESSMENT_DATA_0076_R01.csv"
ITEMS=ROOT/"inputs/SYNTHETIC_ITEM_ANALYSIS_0076_R01.csv"
OUT=ROOT/"evidence/technical_review_R01.json"
KEYS=list("ACBDACBDCADBCADBACDB")
POST=[70,75,85,80,75,90,65,80,85,70,90,75,80,90,70,85]
PRACT=[(5,4,4,5),(4,4,3,4),(4,3,3,4),(4,4,4,4),(4,5,4,4),(5,5,4,5),(3,4,3,3),(5,4,4,5),(4,4,3,4),(4,3,3,4),(5,5,5,5),(4,4,3,4),(4,4,4,4),(4,5,4,4),(4,4,3,3),(3,3,3,3)]
EXPECTED_REGISTER=["fact_id","case_id","document_locator","field_or_claim","synthetic_value_or_status","required_real_replacement","verification_role","hold_or_rejection_rule","source_link"]

def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def main():
    errors=[]; checks={}
    for p in (BODY,DOCX,REGISTER,DATA,ITEMS):
        if not p.is_file(): errors.append("MISSING "+str(p))
    if errors: raise SystemExit(json.dumps({"errors":errors}))
    body=BODY.read_text(encoding="utf-8-sig")
    for token in ["SYN-0073-2026-A01","2026-08-15","1265","79.0625","54.375","24.6875","0077","0078","0079","REAL=UNKNOWN/HOLD","SAMPLE/SYNTHETIC"]:
        checks["body_has_"+re.sub(r"[^A-Za-z0-9]+","_",token).strip("_")]=token.casefold() in body.casefold()
        if not checks["body_has_"+re.sub(r"[^A-Za-z0-9]+","_",token).strip("_")]: errors.append("body missing "+token)
    checks["body_binds_checkbox_084"]="084" in body and ("官方勾選" in body or "勾選 084" in body)
    if not checks["body_binds_checkbox_084"]: errors.append("body missing official checkbox serial 084")
    with DATA.open(encoding="utf-8-sig",newline="") as f:
        rows=list(csv.DictReader(f))
    checks["assessment_rows"]=len(rows)
    if len(rows)!=16: errors.append("assessment row count !=16")
    if [r["learner_code"] for r in rows]!=[f"P{i:02d}" for i in range(1,17)]: errors.append("learner IDs mismatch")
    item_correct=[0]*20
    for i,r in enumerate(rows):
        post=int(r["posttest_total_from_0074_synthetic"])
        if post!=POST[i]: errors.append(f"post score mismatch P{i+1:02d}")
        if int(r["pretest_total_from_0074_synthetic"]) not in range(0,101): errors.append("pretest range")
        responses=[r[f"Q{q:02d}_response"] for q in range(1,21)]
        if len(responses)!=20: errors.append(f"response width P{i+1:02d}")
        correct=sum(a==KEYS[q] for q,a in enumerate(responses))
        if correct!=post//5: errors.append(f"response score mismatch P{i+1:02d}: {correct} != {post//5}")
        for q,a in enumerate(responses):
            if a==KEYS[q]: item_correct[q]+=1
        ps=[int(r[k]) for k in ["communication_0_5","privacy_boundary_0_5","record_accuracy_0_5","handoff_followup_0_5"]]
        if ps!=list(PRACT[i]): errors.append(f"practical subscore mismatch P{i+1:02d}")
        if sum(ps)!=int(r["practical_total_0_20"]): errors.append(f"practical total mismatch P{i+1:02d}")
    post_sum=sum(int(r["posttest_total_from_0074_synthetic"]) for r in rows)
    pre_sum=sum(int(r["pretest_total_from_0074_synthetic"]) for r in rows)
    pract_sum=sum(int(r["practical_total_0_20"]) for r in rows)
    both=sum(int(r["posttest_total_from_0074_synthetic"])>=70 and int(r["practical_total_0_20"])>=15 for r in rows)
    checks.update({"pre_sum":pre_sum,"pre_mean":pre_sum/16,"post_sum":post_sum,"post_mean":post_sum/16,"mean_gain":post_sum/16-pre_sum/16,"post_ge_80_n":sum(int(r["posttest_total_from_0074_synthetic"])>=80 for r in rows),"practical_sum":pract_sum,"practical_mean":pract_sum/16,"practical_ge_15_n":sum(int(r["practical_total_0_20"])>=15 for r in rows),"both_pass_n":both,"item_correct_n":item_correct})
    if (pre_sum,post_sum,pract_sum,both)!=(870,1265,253,11): errors.append("aggregate totals mismatch")
    with ITEMS.open(encoding="utf-8-sig",newline="") as f: ir=list(csv.DictReader(f))
    checks["item_analysis_rows"]=len(ir)
    if len(ir)!=20: errors.append("item-analysis rows !=20")
    if [int(x["correct_n"]) for x in ir]!=item_correct: errors.append("item analysis mismatches response matrix")
    with REGISTER.open(encoding="utf-8-sig",newline="") as f:
        reader=csv.DictReader(f); register=list(reader); headers=reader.fieldnames
    checks["register_rows"]=len(register)
    if headers!=EXPECTED_REGISTER: errors.append("replacement register headers mismatch")
    if len(register)<70: errors.append("replacement register below 70 rows")
    ids=[r.get("fact_id","") for r in register]
    if ids!=[f"SYN-0076-FACT-{i:04d}" for i in range(1,len(ids)+1)]: errors.append("register IDs not unique contiguous")
    for n,r in enumerate(register,1):
        label=(r.get("synthetic_value_or_status","")+" "+r.get("source_link","")).casefold()
        if "sample/synthetic" not in label or "real=unknown/hold" not in label: errors.append(f"register sample/real tag missing row {n}")
        if any(not r.get(k,"").strip() for k in EXPECTED_REGISTER): errors.append(f"empty register field row {n}")
    with zipfile.ZipFile(DOCX) as z:
        bad=z.testzip()
        if bad: errors.append("bad DOCX zip entry "+bad)
        if "word/document.xml" not in z.namelist(): errors.append("document.xml missing")
        external=any("TargetMode=\"External\"" in z.read(n).decode("utf-8","ignore") for n in z.namelist() if n.endswith(".rels"))
    d=Document(DOCX)
    txt="\n".join(p.text for p in d.paragraphs)+"\n"+"\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)
    checks["docx_sha256"]=sha(DOCX); checks["docx_bytes"]=DOCX.stat().st_size
    checks["docx_paragraphs"]=len(d.paragraphs); checks["docx_tables"]=len(d.tables)
    checks["docx_contains_real_boundary"]="REAL=UNKNOWN/HOLD" in txt
    if not checks["docx_contains_real_boundary"]: errors.append("DOCX boundary text missing")
    if len(d.tables)<10: errors.append("unexpectedly few tables")
    # Body and table type sizes; builder contract is 10.5 pt body and 9.5 pt tables.
    body_sizes=[]; table_sizes=[]
    for p in d.paragraphs:
        for run in p.runs:
            if run.text.strip() and run.font.size: body_sizes.append(run.font.size.pt)
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for run in p.runs:
                        if run.text.strip() and run.font.size: table_sizes.append(run.font.size.pt)
    checks["min_body_run_pt"]=min(body_sizes) if body_sizes else None
    checks["min_table_run_pt"]=min(table_sizes) if table_sizes else None
    if body_sizes and min(body_sizes)<10.5: errors.append("body run under 10.5pt")
    if table_sizes and min(table_sizes)<9.5: errors.append("table run under 9.5pt")
    checks["docx_has_external_relationships"]=external
    if checks["docx_has_external_relationships"]: errors.append("external relationship found")
    receipt={"schema":"TTQS_ONE_0076_TECHNICAL_REVIEW_R01.v1","document_id":"0076","requirement_id":"I17B-E01","official_checkbox_serial":"084","docx_sha256":checks["docx_sha256"],"checks":checks,"errors":errors,"technical_status":"PASS" if not errors else "FAIL","formal_acceptance":"NOT_GRANTED","real_status":"UNKNOWN/HOLD"}
    OUT.parent.mkdir(parents=True,exist_ok=True); OUT.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(receipt,ensure_ascii=False))
    if errors: raise SystemExit(1)
if __name__=="__main__": main()
