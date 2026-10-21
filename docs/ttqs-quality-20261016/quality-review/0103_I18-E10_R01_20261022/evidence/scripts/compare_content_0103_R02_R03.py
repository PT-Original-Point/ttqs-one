from pathlib import Path
import hashlib, json
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
root=Path(__file__).resolve().parents[2]
files={
 "R02":root/"candidate"/"0103__I18-E10__感謝函取得與真偽查核工作紀錄_R02.docx",
 "R03":root/"candidate"/"0103__I18-E10__感謝函取得與真偽查核工作紀錄_R03.docx"
}
def extract(path):
 d=Document(path); out=[]
 for node in d.element.body.iterchildren():
  if node.tag==qn("w:p"):
   value=Paragraph(node,d).text
   if value.strip(): out.append(["P",value])
  elif node.tag==qn("w:tbl"):
   t=Table(node,d)
   for row in t.rows: out.append(["T",[c.text for c in row.cells]])
 return out
values={k:extract(v) for k,v in files.items()}
serial={k:json.dumps(v,ensure_ascii=False,separators=(",",":")) for k,v in values.items()}
sha={k:hashlib.sha256(v.encode("utf-8")).hexdigest() for k,v in serial.items()}
same=serial["R02"]==serial["R03"]
diff=[]
if not same:
 for i,(a,b) in enumerate(zip(values["R02"],values["R03"])):
  if a!=b:
   diff.append({"index":i,"R02":a,"R03":b})
   if len(diff)>=8: break
 print(json.dumps({"same":same,"R02_blocks":len(values["R02"]),"R03_blocks":len(values["R03"]),"R02_text_sha256":sha["R02"],"R03_text_sha256":sha["R03"],"first_differences":diff},ensure_ascii=False))
 raise SystemExit(1)
print(json.dumps({"same":same,"R02_blocks":len(values["R02"]),"R03_blocks":len(values["R03"]),"R02_text_sha256":sha["R02"],"R03_text_sha256":sha["R03"],"paragraphs_R02":len(Document(files["R02"]).paragraphs),"paragraphs_R03":len(Document(files["R03"]).paragraphs),"tables_R02":len(Document(files["R02"]).tables),"tables_R03":len(Document(files["R03"]).tables)},ensure_ascii=False))