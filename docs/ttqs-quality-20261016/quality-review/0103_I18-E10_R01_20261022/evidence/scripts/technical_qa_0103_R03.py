from pathlib import Path
import hashlib,json,zipfile
from lxml import etree
from docx import Document
root=Path(__file__).resolve().parents[2]
docx=root/"candidate"/"0103__I18-E10__感謝函取得與真偽查核工作紀錄_R03.docx"
z=zipfile.ZipFile(docx)
bad=z.testzip()
xml=[n for n in z.namelist() if n.endswith(".xml") or n.endswith(".rels")]
parse_ok=True
for n in xml:
 try: etree.fromstring(z.read(n))
 except Exception: parse_ok=False
d=Document(docx)
sha=hashlib.sha256(docx.read_bytes()).hexdigest()
receipt={
 "receipt_id":"TTQS-0103-I18-E10-R03-TECHNICAL-QA",
 "doc_id":"0103","requirement_id":"I18-E10","official_checkbox":111,"revision":"R03-LAYOUT-EXPERIMENT",
 "candidate_docx_sha256":sha,"bytes":docx.stat().st_size,
 "checks":{"zip_crc_ok":bad is None,"xml_parts_parse":parse_ok,"python_docx_opened":True,"paragraphs":len(d.paragraphs),"tables":len(d.tables),"sections":len(d.sections),"layout_marker_not_visible":all("LANDSCAPE_CONTEXT" not in p.text for p in d.paragraphs)},
 "technical_status":"PASS_PACKAGE_ONLY" if bad is None and parse_ok else "FAIL",
 "visual_status":"FAIL_REPAIRABLE_LAYOUT",
 "semantic_status":"NOT_REVIEWED_ON_R03_EXACT_SHA",
 "candidate_disposition":"NOT_ADOPTED; INTERNAL_LAYOUT_EXPERIMENT",
 "formal_acceptance":"NOT_GRANTED"
}
out=root/"evidence"/"technical_qa_0103_R03.json"
out.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
print(json.dumps({"sha256":sha,"bytes":docx.stat().st_size,"paragraphs":len(d.paragraphs),"tables":len(d.tables),"sections":len(d.sections),"zip_crc":bad,"xml_parts":len(xml),"technical_status":receipt["technical_status"]},ensure_ascii=False))