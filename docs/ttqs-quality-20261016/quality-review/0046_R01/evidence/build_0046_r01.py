from __future__ import annotations
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
import csv, hashlib, json, re, zipfile
from datetime import datetime, timezone
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'candidate' / '0046_R01.docx'
REGISTER = ROOT / 'evidence' / 'SYNTHETIC_REPLACEMENT_REGISTER_R01.csv'
NAVY = '244B63'
PALE = 'F1F5F7'
GRID = 'AEBBC4'
FONT = 'Microsoft JhengHei'

facts=[]
def fact(value, location, replace_source, role, shared='0045 R01 同案合成資料；不得視為 REAL'):
    fid=f'SYN-046-{len(facts)+1:03d}'
    facts.append({'fact_id':fid,'value_synthetic':str(value),'type':'SAMPLE_SYNTHETIC','document_location':location,'shared_sample_reference':shared,'real_replacement_source_required':replace_source,'responsible_role':role,'status':'REPLACE_BEFORE_REAL_USE'})
    return fid

def set_font(run, size=10.5, bold=None, color=None, name=FONT):
    run.font.name=name
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), name)
    run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    if color: run.font.color.rgb=RGBColor.from_string(color)

def set_cell_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)

def set_cell_border(cell, color=GRID, size='6'):
    tcPr=cell._tc.get_or_add_tcPr(); borders=tcPr.first_child_found_in('w:tcBorders')
    if borders is None: borders=OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag='w:'+edge; el=borders.find(qn(tag))
        if el is None: el=OxmlElement(tag); borders.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),size); el.set(qn('w:color'),color)

def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr(); tblHeader=OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'),'true'); trPr.append(tblHeader)

def set_cant_split(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:cantSplit'); trPr.append(el)

def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next=True

def add_p(doc, text='', bold_prefix=None, size=10.5, color=None, after=4, before=0, align=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(after); p.paragraph_format.space_before=Pt(before); p.paragraph_format.line_spacing=1.12
    if align is not None: p.alignment=align
    if bold_prefix and text.startswith(bold_prefix):
        r=p.add_run(bold_prefix); set_font(r,size,True,color)
        r=p.add_run(text[len(bold_prefix):]); set_font(r,size,False,color)
    else:
        r=p.add_run(text); set_font(r,size,color=color)
    return p

def add_h(doc, text, level=1):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(7 if level==1 else 5); p.paragraph_format.space_after=Pt(4); set_keep_with_next(p)
    r=p.add_run(text); set_font(r,14 if level==1 else 11.5,True,NAVY)
    return p

def add_table(doc, headers, rows, widths=None, font_size=9.5, repeat=True, header=True, spacer=True):
    t=doc.add_table(rows=1 if header else 0, cols=len(headers) if header else len(rows[0]))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    if header:
        for i,h in enumerate(headers):
            c=t.rows[0].cells[i]; c.text=''; p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.05
            r=p.add_run(str(h)); set_font(r,font_size,True,'FFFFFF'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_shading(c,NAVY); set_cell_border(c,NAVY)
        set_cant_split(t.rows[0])
        if repeat: set_repeat_table_header(t.rows[0])
    for ri,row in enumerate(rows):
        cells=t.add_row().cells; set_cant_split(t.rows[-1])
        for ci,val in enumerate(row):
            c=cells[ci]; c.text=''; p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.08
            r=p.add_run(str(val)); set_font(r,font_size,False)
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2==1: set_cell_shading(c,PALE)
            set_cell_border(c)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
        for i,w in enumerate(widths): t.columns[i].width=Inches(w)
    if spacer:
        doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def add_page(doc): doc.add_page_break()

def add_footer(doc):
    sec=doc.sections[0]
    p=sec.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0)
    r=p.add_run('0046 R01  |  SAMPLE / SYNTHETIC  |  第 '); set_font(r,8.5,False,'687782')
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); p._p.append(fld)
    r=p.add_run(' 頁'); set_font(r,8.5,False,'687782')

def add_header(doc):
    sec=doc.sections[0]
    p=sec.header.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; p.paragraph_format.space_after=Pt(2)
    r=p.add_run('SAMPLE / SYNTHETIC  |  未生效示範表'); set_font(r,8.5,True,'65737D')

def setup(doc):
    sec=doc.sections[0]; sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
    sec.left_margin=Inches(.67); sec.right_margin=Inches(.67); sec.top_margin=Inches(.72); sec.bottom_margin=Inches(.62)
    sec.header_distance=Inches(.32); sec.footer_distance=Inches(.3)
    normal=doc.styles['Normal']; normal.font.name=FONT; normal._element.rPr.rFonts.set(qn('w:eastAsia'),FONT); normal.font.size=Pt(10.5)
    normal.paragraph_format.space_after=Pt(4); normal.paragraph_format.line_spacing=1.12
    add_header(doc); add_footer(doc)

def add_title(doc):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(3); set_keep_with_next(p)
    r=p.add_run('師資教材廠商評選紀錄表'); set_font(r,20,True,'183B50')
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(10); set_keep_with_next(p)
    r=p.add_run('單案需求、候選比較、逐項評分、核定與交付追蹤'); set_font(r,11,False,'52636E')

def add_fact_row(value, location, replacement, role, shared=None):
    return fact(value,location,replacement,role,shared or '0045 R01 §10.1–§10.4 同案 SAMPLE/SYNTHETIC 值')

# Shared case values from 0045; they remain SAMPLE/SYNTHETIC here.
case='SYN-045-01'; formid='SYN-046-01'; course='志工服務溝通與紀錄整理工作坊'
request_date='2026-05-04'; request_ref='SAMPLE-LOC/SYN-045-01/01-demand'
participants=12; gap_service=(8,12); gap_records=(7,12); hours_each=6; sessions=2; total_hours=12; budget=15000
vendors=[
 {'id':'A','name':'方案 A','scope':'師資 A、教材 A；主題符合，個案案例較少；回覆兩次授課時段及教材交付。','price':14000,'tier':2,'scores':[5,4,4,4,2,4], 'raw_e01':[5,4,3,4,2,5], 'raw_e02':[5,4,5,4,2,3]},
 {'id':'B','name':'方案 B','scope':'師資 B 與共同備課教材；逐項回應兩項能力落差，列出流程、教材樣稿、修訂節點及授課時段。','price':12000,'tier':3,'scores':[4,5,5,5,3,4], 'raw_e01':[4,5,5,5,3,4], 'raw_e02':[4,5,5,5,3,4]},
 {'id':'C','name':'方案 C','scope':'師資 C 與教材套件；成本最低，紀錄練習充足，但個案情境及課後修訂說明較少。','price':9000,'tier':5,'scores':[4,4,3,4,5,3], 'raw_e01':[3,4,2,5,5,3], 'raw_e02':[5,4,4,3,5,3]}
]
criteria=[('需求符合度',30),('師資與教學',25),('排程與回應',15),('教材與可近用',10),('含稅總價',10),('風險與服務',10)]

# PAGE 1 — form identity, use point, and sample boundary
D=Document(); setup(D); add_title(D)
add_p(D,'用途：在訓練需求與同一規格的候選方案已備妥後，逐案記錄師資、教材或合作廠商的可比評選、理由、覆核及後續採購／交付追蹤。這是單案紀錄表；制度流程見 0045，不是會議原始紀錄或年度檢討。',size=10.5)
add_p(D,'本頁及全表為 SAMPLE / SYNTHETIC 示範。所有角色、日期、候選、報價、分數、決議、附件代碼與結果均為合成內容；不得作為原始點協會曾評選、核准、簽約、付款或辦訓的證明。',size=9.5,color='8B3B32')
add_h(D,'1. 單案識別與使用控制')
identity=[]
for v,loc,source,role in [
 (formid,'§1 識別表／紀錄編號','正式案號及表單編號','需求承辦'),
 (case,'§1 識別表／需求案號','正式核定需求案號及來源','需求承辦'),
 (course,'§1 識別表／課程名稱','核定課程計畫及正式名稱','課程負責人'),
 (request_date,'§1 識別表／需求盤點日','原始需求調查及日期','需求承辦'),
 ('2026-05-12','§1 識別表／評選日','實際評選紀錄及日期','評選召集／紀錄人'),
 ('R01、E01、E02、A01','§1 識別表／角色代碼','真實任命或職務分工資料','填表人／權責者'),
 ('I10-E01；官方方框 047','§1 識別表／官方追溯','官方核對表與核實的 Unique Blueprint','品質承辦')]:
    fid=add_fact_row(v,loc,source,role)
    identity.append((fid,loc.split('／')[-1],v,role+'；核對來源與必填'))
add_table(D,['合成事實ID／欄位','欄位用途','完整填寫示例','填寫與驗證責任'],identity,[1.55,1.20,2.25,1.65],9.5)
add_h(D,'2. 何時填、怎麼使用')
add_table(D,['工作時點','本表操作','完成輸出與檢查'],[
 ('承諾／採購前','先凍結需求與評分權重；將同規格邀請、收件與資格證據依序編號。','候選可以比較；超預算或必要資格不符者不得靠加權分數補救。'),
 ('候選評分時','評選人先各自評分，再記錄原分、證據、理由、差異與平均。','兩人分數不可事後協商改寫；差異與原始值保留。'),
 ('建議／核定時','列出排序、門檻、推薦理由、限制、核定權限與日期。','評選建議與正式核准分開；沒有權限或原件時不代表核准。'),
 ('履約／結案時','連結契約、交付、缺失、補正、複驗及成效材料。','每個缺失有責任人、期限、證據位置與關閉條件。')],[1.25,3.55,1.85],9.5,spacer=False)
# PAGE 2 — need/specification fields and validation
p=add_p(D,'複製本表作新案使用時，建立新紀錄編號、保留同規格比較原則，替換所有 SYNTHETIC 值及附件定位，並由有權限者按真實制度核定；不可直接把本示例改名後當成真實評選紀錄。',size=9.5); p.paragraph_format.page_break_before=True; add_h(D,'3. 需求與評選規格（候選邀請前固定）')
need_rows=[]
need_values=[
 ('需求來源與日期',request_date+'；需求盤點資料',request_ref,'需求承辦'),
 ('規劃人數',f'{participants} 人；分母為預定參訓總數', 'SAMPLE-LOC/SYN-045-01/01-demand','需求承辦'),
 ('能力落差 A',f'{gap_service[0]}/{gap_service[1]} 人；{gap_service[0]/gap_service[1]*100:.1f}%回報志工服務情境練習需求','SAMPLE-LOC/SYN-045-01/01-demand','需求承辦'),
 ('能力落差 B',f'{gap_records[0]}/{gap_records[1]} 人；{gap_records[0]/gap_records[1]*100:.1f}%回報個案紀錄判斷需求','SAMPLE-LOC/SYN-045-01/01-demand','需求承辦'),
 ('課程設計規格',f'{sessions} 次 × 每次 {hours_each} 小時 = {total_hours} 小時；包含情境演練、回饋與修訂','SAMPLE-LOC/SYN-045-01/02-spec','課程負責人'),
 ('預算上限',f'新臺幣 {budget:,} 元（含稅及必要交付成本）；示範上限，非協會核定預算','SAMPLE-LOC/SYN-045-01/03-budget','預算權責者'),
]
for label,val,source,role in need_values:
    add_fact_row(val,'§3需求表／'+label,'真實需求調查、核定課程計畫、預算核定原件',role)
    need_rows.append((label,val,source,role+'／核日期分母'))
add_table(D,['需求／規格欄位','完整示例值','資料來源／附件定位（示例碼）','填寫／驗證'],need_rows,[1.25,2.6,1.65,1.15],9.5)
add_page(D); add_h(D,'3.1 同規格邀請內容及可比性檢查',2)
spec_rows=[
 ('師資資格與教學','資格、相關教學方式、兩項能力落差如何練習與回饋。','資格及經驗附件；逐項標頁碼／檔名。','缺少必要資格即不合格；不得以低價抵銷。'),
 ('教材與可近用','樣稿、版次、格式、可讀性、修訂與使用授權。','教材樣稿及授權範圍。','版次／數量／格式清楚，權利可核對。'),
 ('排程與交付','兩個授課時段、教材交付日期、補正節點及回應方式。','逐項回覆表及時程。','同一時數、日期窗與交付條件才能比較。'),
 ('含稅完整總價','師資、教材、必要交通／交付及修訂費用。','正式報價原件；示例碼只作定位。','超預算或總價範圍不明者先退回補正。'),
 ('風險、資料與服務','資料處理、保密、取消、替代師資、客訴與補正責任。','服務說明、資料處理條款。','不明確時不得推定已承諾。')]
for x in spec_rows: add_fact_row(x[0]+'；'+x[1],'§3.1同規格欄位／'+x[0],'正式需求規格與候選原件', '課程負責人／評選人')
add_table(D,['可比項目','邀請規格／必須回答','候選附件定位方式','驗證規則'],spec_rows,[1.15,2.4,1.55,1.55],9.5)
add_p(D,'驗證示例：若一方案只報講師費，其他方案報含教材與交付總價，先要求補成同範圍；未補齊不得把三個金額直接排名。分母、人數、時數、含稅條件須一致。',size=9.5)

# PAGE 3 — candidate and gate register
add_page(D); add_h(D,'4. 候選方案來源與資格門檻')
add_p(D,'本示例假設於 2026-05-06 發出同一規格邀請，三個合成方案於 2026-05-11 前回覆，並於 2026-05-12 評選。候選 A/B/C 及資料代碼不是實際廠商或報價。',size=9.5)
for v in [3,'2026-05-06','2026-05-11','2026-05-12']:
    add_fact_row(v,'§4候選總表／邀請、截止與評選欄','真實邀請寄送紀錄、原始收件時間及評選紀錄','需求承辦／紀錄人')
offers=[]
for v in vendors:
    src=f'SAMPLE-LOC/SYN-045-01/04-offer-{v["id"]}'
    add_fact_row(v['name'],'§4候選表／方案 '+v['id'],'真實候選名稱及原始報價／資格文件','需求承辦')
    add_fact_row(v['scope'],'§4候選表／方案 '+v['id']+' 資源組合','正式候選書面回覆及附件','評選人')
    add_fact_row(v['price'],'§4候選表／方案 '+v['id']+' 含稅總價','正式報價原件、含稅及交付費明細','財務覆核者')
    add_fact_row(v['tier'],'§4候選表／方案 '+v['id']+'價格級距分','正式核定的評分尺及可核對總價','評選人')
    offers.append((v['name'],v['scope'],f'NT$ {v["price"]:,}',f'級距 {v["tier"]} 分',src))
add_table(D,['候選','資源組合與回覆摘要','含稅總價','價格級距分','附件定位（合成碼）'],offers,[.7,2.45,.85,.9,1.75],9.5)
add_h(D,'4.1 先資格、後加權（硬門檻不可補分）',2)
gates=[]
for v in vendors:
    for gate in ['師資資格符合','時程可交付','規格及教材樣稿具備','含稅總價不超預算','資料／服務責任有說明']:
        add_fact_row(f'{v["id"]}／{gate}=PASS（合成）','§4.1資格表／'+v['id']+'／'+gate,'真實資格、規格、報價及服務原件','承辦人初檢；評選召集複核')
        gates.append((v['id'],gate,'PASS（合成）','依示例附件定位表查得；非真實查核結果'))
add_table(D,['方案','硬門檻','示例判定','必留查核證據'],gates,[.7,1.75,1.05,3.15],9.5)
add_p(D,'例外分支：資格未過、超預算、總價不可比、缺必要附件或來源不明 → 標記 HOLD／退補，停止該方案加權；記錄缺件、通知日、期限、補件原件與覆核者。候選數不足三件時依 0045 程序先停案並取得有權者事前處理，不得自行降低資格或追認。',size=9.5)

# PAGE 4 — conflict declaration and scale
add_page(D); add_h(D,'5. 利益衝突聲明與評分尺')
add_p(D,'評分前由每位評選人各自閱讀同一規格、申報關係並保存原始分數。出現本人、近親、財務、僱傭、密切合作或其他可能影響公正之關係，先停止該人評分，交由無衝突且具權限者判斷迴避與替補；不得由當事人自行認定無影響後繼續評分。',size=9.5)
conflicts=[]
for eid in ['E01','E02']:
    add_fact_row(f'{eid}；示例聲明「未發現須迴避關係」','§5利益聲明／'+eid,'真實利益關係申報、覆核與迴避紀錄','評選人；召集人複核')
    conflicts.append((eid,'示例：未發現須迴避關係','未發現（合成）','評分前；非真實人員聲明'))
add_table(D,['評選角色碼','利益關係聲明示例','處置','必備真實證據'],conflicts,[1.05,2.1,1.3,2.2],9.5)
add_p(D,'角色碼 E01、E02 是示範代碼，不是自然人姓名；本表沒有真實簽名或利益申報。正式資料應保存申報原文、申報日、覆核者、迴避／替補決定及替代評選者的獨立評分。',size=9.5)
add_h(D,'5.1 事前評分量尺（1–5；2、4為相鄰錨點間的中間判斷）',2)
rubric=[
 ('需求符合度 30%','1–2：未回應多項明確落差；3：主要需求有回應；4–5：逐項對應兩項落差並有可檢視練習／評量。'),
 ('師資與教學 25%','1–2：資格或教法證明不足；3：資格與基本教學流程符合；4–5：教法、教材、練習、回饋及目標連結完整。'),
 ('排程與回應 15%','1–2：時段／交付不確定；3：可排程但修訂回應有限；4–5：兩次時段、交付及修訂節點清楚。'),
 ('教材與可近用 10%','1–2：內容／格式或使用條件不符；3：符合基本規格；4–5：可讀、可用、版次可追且有品質／可近用安排。'),
 ('含稅總價 10%','超上限或總價不可比：不合格；其餘依事前核定互斥級距給 5 至 2 分；低價不等於最佳。'),
 ('風險與服務 10%','1–2：資料處理／責任未說明；3：基本保密、變更與補正方式清楚；4–5：版本、存取、事件通報、替代及補正責任可執行。')]
for name,desc in rubric: add_fact_row(name+'；'+desc,'§5.1評分尺／'+name,'正式核定的評選標準與授權紀錄','有權者核定；評選人依同尺評分')
add_table(D,['面向／權重','可觀察評分錨點'],rubric,[1.55,5.1],9.5)
add_p(D,'同分順序：需求符合度加權分 → 師資與教學 → 排程與回應；仍同分則暫停推薦，交有權者依事前公告條件處理，不得事後抽籤或新增偏好。此示例無同分。',size=9.5)

# PAGE 5 — raw independent scoring
add_page(D); add_h(D,'6. 兩名評選者逐項原始分數與理由')
add_p(D,'E01、E02 各自依同一評分尺獨立給 1–5 分。表列均分 = (E01 + E02) ÷ 2；平均只作後續加權輸入，不覆蓋原始分。每項理由須指向同一方案的合成附件定位，正式使用時替換成可追溯原件頁碼。',size=9.5)
score_rows=[]
for v in vendors:
    reasons=[
      ['兩項需求均有對應練習與教材。','資格及教學流程符合規格。','兩次時段與交付可行，修訂節點說明一般。','格式符合需求，版本與可近用安排一般。','14,000 元在 13,000–15,000 元級距。','資料處理與補正有說明；細節判斷有差異。'],
      ['主要需求均有回應，尚有一項可再補強。','教法、教材與評量完整連結目標。','兩次時段、交付及修訂節點清楚。','樣稿符合格式並列明版本修訂。','12,000 元在 11,000–13,000 元級距。','資料界線與補正責任可執行。'],
      ['有紀錄練習但個案情境較少，評分有差異。','資格及基本教學流程符合。','時段可行但修訂與回覆節點較少。','教材內容充足，可近用及版本維護說明較少。','9,000 元落在最低價格級距。','基本服務說明存在，補正細節有限。']
    ][['A','B','C'].index(v['id'])]
    for i,(criterion,weight) in enumerate(criteria):
        e1=v['raw_e01'][i]; e2=v['raw_e02'][i]; mean=(Decimal(e1+e2)/2).quantize(Decimal('0.0'))
        loc=f'§6評分表／{v["id"]}／{criterion}／E01、E02、均分、理由'
        for evaluator,score in [('E01',e1),('E02',e2)]: fact(score,loc+'／'+evaluator,'原始個別評分表、評分人及其依據附件','評選者本人；召集人核對')
        fact(mean,loc+'／平均','由可查原分重算；核實計算規則','紀錄人重算；第二人覆核')
        score_rows.append((v['id'],f'{criterion} {weight}%',f'{e1}／{e2}',str(mean),reasons[i]))
add_table(D,['方案','評分面向／權重','E01／E02','平均','具體理由（示範）'],score_rows,[.48,1.65,.75,.55,3.2],9.5)
add_p(D,'評分完成檢查：六項權重總計 100%；每個方案均有 6 組 E01/E02 原分，共 36 個原始分；任一分數不得空缺、代填或事後覆寫。差異由兩人各自補充理由，仍保留原分及差異。',size=9.5)

# PAGE 6 — calculation, ranking, recommendation and decision
add_page(D); add_h(D,'7. 加權計算、排序與建議')
add_p(D,'計算式：各面向貢獻 = 權重 × ((E01 原分 + E02 原分) ÷ 2) ÷ 5；總分 = 六面向貢獻加總。權重以百分點計，結果四捨五入至小數一位。示範門檻 70 分、價格級距及上限都是合成參數，須由有權者核定後才可用於真實案件。',size=9.5)
for x in [100,70,82,88,77]: fact(x,'§7計算與排序／權重、門檻或候選總分','事前核定評分尺及獨立原始分數；以公式重新計算','紀錄人計算；覆核者重算')
summary=[]
for v,total in zip(vendors,[82,88,77]): summary.append((v['id'],f'{v["price"]:,} 元；價格級距 {v["tier"]}',f'{total:.1f}', '達示範門檻 70' if total>=70 else '未達', '合格；依總分排序'))
add_table(D,['候選','含稅總價／價格分','加權總分','示範門檻','排序'],summary,[.7,1.7,1.0,1.3,1.95],9.5)
add_h(D,'7.1 可重算算例與結論',2)
add_p(D,'方案 B：30×4÷5 + 25×5÷5 + 15×5÷5 + 10×5÷5 + 10×3÷5 + 10×4÷5 = 24 + 25 + 15 + 10 + 6 + 8 = 88.0。方案 A = 30 + 20 + 12 + 8 + 4 + 8 = 82.0；方案 C = 24 + 20 + 9 + 8 + 10 + 6 = 77.0。權重總計 100%，三案均達示範門檻。',size=10)
add_h(D,'7.2 建議與有權核定欄',2)
decision=[
 ('評選排序','B（88.0）→ A（82.0）→ C（77.0）；同分條件未觸發。','評分總表及逐項原分'),
 ('建議理由','示範建議 B：需求、教學、排程、教材與補正安排整體較完整；C 的低價不單獨決定結果。','同規格方案表、逐項理由'),
 ('示範權限欄','A01 為合成角色碼；欄位示意核定角色、權限依據、核定日期與決定內容。','正式授權文件、真實核定原件'),
 ('狀態與限制','未經協會正式核定；無真人簽名、無實際採購、無契約或付款證明。','有權者實際核定後始能更新狀態')]
for row in decision: fact(row[1],'§7.2核定欄／'+row[0],'正式評選原始資料、權限依據及有權者簽核原件','承辦人；權責核定者')
add_table(D,['欄位','完整示例內容','應附／應查原件'],decision,[1.2,3.35,2.1],9.5)

# PAGE 7 — procurement, contract, delivery, follow-up
add_page(D); add_h(D,'8. 採購、契約、交付與履約追蹤')
add_p(D,'推薦、核定、簽約、交付、驗收及付款是不同狀態。下列日期／結果皆是合成流程示範，不是協會執行紀錄；表單承辦須以真實原件填寫每個日期、文件定位、缺失、補正與覆核。',size=9.5)
timeline=[
 ('2026-05-15','示範核定欄由合成角色 A01 記錄推薦方案 B 與 15,000 元示範上限。','無真人核定、簽章或組織授權'),
 ('2026-05-18','合成契約條件 SYN-CT-045-01：總價 12,000 元；兩次各 6 小時；12 份教材；開課前修訂。','不是報價、契約、發票、付款或真實用印'),
 ('2026-05-29','教材初檢：2 份索引頁碼未更新；記錄缺失並通知合成方案 B 修正。','示範缺失；不得當成真實驗收'),
 ('2026-05-30','示範複驗：12 份最終教材頁碼一致；課程流程與個案情境已納入。','非真實收件／驗收證明'),
 ('2026-06-01、2026-06-08','兩次各 6 小時；合計 12 小時；示範完成 12/12 預定小時。','非真實出席、授課或成果證明'),
 ('2026-06-22','10/12 有效回饋；4 人評 5、4 人評 4、2 人評 3；平均 4.2/5。','合成問卷；沒有真實參與者資料')]
for row in timeline:
    fact(row[0],'§8履約時間表／'+row[1],'真實核定／契約／簽收／授課／問卷原件，依欄位取得','承辦／交付驗收者／課程負責人')
    fact(row[1],'§8履約時間表／'+row[0],'真實來源原件及責任人確認','承辦及相應責任角色')
add_table(D,['合成日期','紀錄內容及完整示例','證據狀態限制'],timeline,[1.35,3.35,1.95],9.5)
add_h(D,'8.1 履約驗收欄與缺失關閉條件',2)
accept=[
 ('師資／服務','依契約核對授課人員、時數、課程內容及實際授課紀錄。','契約列明項目逐項符合；不足記缺失、責任與期限。'),
 ('教材／資源','核對版次、數量、內容、格式、可讀性及使用授權。','示範初檢 2 份索引待修；複驗 12/12 一致後才關閉示例缺失。'),
 ('時程','比較約定日及實際日，標明延期天數、原因及核定。','不得只寫準時；實際日期缺失時標 UNKNOWN，不可推估。'),
 ('品質／回饋','記錄有效回覆分子、邀請分母、題目量尺、缺漏與計算。','10/12=83.3%；(4×5+4×4+2×3)/10=4.2/5；示範目標4.0/5。'),
 ('付款／結案','依真實契約、請款原件、核定及財務流程處理。','本示例沒有付款資料，不能據此表示已付款或已結案。')]
for row in accept: fact(row[1],'§8.1驗收表／'+row[0],'契約、實物／授課原件、真實問卷及有權核定紀錄','交付驗收者及覆核者')
add_table(D,['驗收面向','欄位與查驗動作','合格／異常處理'],accept,[1.1,2.6,2.95],9.5)

# PAGE 8 — exceptions, evidence mapping, handoff, neighbor distinction
add_page(D); add_h(D,'9. 缺漏、錯誤及例外處理')
exceptions=[
 ('合成例：教材索引','2026-05-29；2份索引頁碼未更新。','通知補正，原檢查結果及版本留存；2026-05-30重驗12份一致後關閉示例缺失。'),
 ('缺報價／附件','收件人、時間、缺少欄位、通知期限、補件原件及覆核者。','未補齊前 HOLD；不以口頭摘要補成報價。'),
 ('利益衝突','申報人、關係、暫停時間、迴避決定、替補權限與新評分原件。','存在疑慮即暫停相關評分；有權者處理後再續。'),
 ('超預算／不可比','總價範圍、稅費、必要交付、版本及預算核定差異。','排除或退補；未取得事前權限不得追認例外。'),
 ('無權核定／簽約','缺少授權來源、核定或簽章原件及待補責任人。','不得把推薦改寫成已核定；不簽約、不付款，依權限升級。')]
for row in exceptions: fact(row[0]+'；'+row[1]+'；'+row[2],'§9例外表／'+row[0],'真實缺失通報、補正、覆核及有權決定原件','承辦人記錄；權責者處理')
add_table(D,['問題類型','必記事實','處置、覆核及關閉條件'],exceptions,[1.2,2.55,2.9],9.5)
add_h(D,'10. 官方要求、表單欄位及附件追溯')
trace=[
 ('需求與規格','§3；同規格邀請與需求附件','需求調查、核定課程計畫、預算來源'),
 ('師資／教材／廠商比較','§4–§6；候選、資格、利益申報、原始分數及逐項理由','各候選原件、共同規格、評選人原分'),
 ('評估與推薦','§7；公式、門檻、排序、理由及權限欄','事前規則、重算表、正式核定原件'),
 ('採購／簽約／後續管理','§8–§9；契約條件、交付、缺失、複驗、回饋與結案','真實契約、簽收／授課／問卷／請款資料'),
 ('官方 I10-E01','完整文書手冊與紀錄分析相關資料；於課程設計時評估選擇師資、教材或合作廠商，並有評估、採購、簽約與後續管理流程。','本表呈現單一案件紀錄；流程規範見 0045；本表示例不代替任何原件。')]
for row in trace: fact(row[0]+'；'+row[1],'§10追溯表／'+row[0],'正式來源及原始附件位置；官方來源使用核定頁次','填表承辦；覆核者')
add_table(D,['追溯對象','本表定位','正式附件／責任'],trace,[1.35,3.05,2.25],9.5)
add_h(D,'10.1 相關文件與文類界線',2)
add_p(D,'0045／I10-E01《資源採購與甄選管理辦法》規範跨案件流程；本表只記單案。0044／I09-E04 保留會議或電子郵件原始討論與決議，不可由本表代造。0047／I10-E02 分析跨年度趨勢與制度改善，不由本單案表取代。',size=9.5)
add_h(D,'11. 表單完成條件與版本',2)
add_p(D,'表單完成條件：需求與同規格已核實；三案資格與總價可比；利益申報完成；每項原分、理由、計算與門檻可重算；推薦與正式核定分開；契約、交付、缺失、複驗、回饋及附件可追溯；必填欄位無空缺。缺原件時標 UNKNOWN／HOLD 並指定責任人，不補造證據。',size=9.5)
add_table(D,['版本','日期','修訂說明','狀態'],[('R01','2026-10-17','依 I10-E01 Unique Blueprint 建立完整單案表單與合成走讀。','SAMPLE候選；未核定、未簽章')],[.8,1.2,3.9,1.1],9.5)

# Generate replacement register: every synthetic value, role, decision, score, date and sample locator.
with REGISTER.open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.DictWriter(f,fieldnames=['fact_id','value_synthetic','type','document_location','shared_sample_reference','real_replacement_source_required','responsible_role','status'],lineterminator='\n')
    w.writeheader(); w.writerows(facts)
D.core_properties.title='師資教材廠商評選紀錄表'
D.core_properties.subject='I10-E01 / official sequence 047 / SAMPLE-SYNTHETIC operational record form'
D.core_properties.author='TTQS ONE document production'
D.core_properties.keywords='0046; I10-E01; SAMPLE; SYNTHETIC'
D.save(OUT)
sha=hashlib.sha256(OUT.read_bytes()).hexdigest()
with zipfile.ZipFile(OUT) as z:
    bad=z.testzip()
    if bad: raise SystemExit('DOCX ZIP error '+bad)
# exact-source and math receipts
blueprint=ROOT/'inputs'/'UNIQUE_BLUEPRINT_0046.jsonl'
blueprint_obj=json.loads(blueprint.read_text(encoding='utf-8-sig'))
now=datetime.now(timezone.utc).isoformat(timespec='seconds').replace('+00:00','Z')
source_map={
 'deliverable_id':'0046','requirement_id':'I10-E01','official_checkbox_sequence':'047',
 'official_requirement_text':blueprint_obj['official_requirement_text'],
 'official_source':'E:\\TTQS\\TTQS_ONE_CODEX_RUNTIME\\SOURCES\\DRIVE_READBACKS\\DRIVE_OFFICIAL_EVIDENCE_TABLE.pdf',
 'official_source_page':'3','official_source_sha256':'805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709',
 'unique_blueprint_source':'E:\\TTQS\\TTQS_ONE_CODEX_RUNTIME\\CONTROL\\BLUEPRINT_CURRENT.jsonl',
 'unique_blueprint_source_sha256':'406d0495365d0970dcf8113a42842663b8a54372a4e204cc492b7c3f8156ef3e',
 'unique_blueprint_row_sha256':'6d2c7798cda52923f9bec34cac347bae43c1ea56883b90d3d412845d80988f05',
 'source_commit':'d68e0f784525b3e94733cc40b239f87b3f620f8b',
 'sibling_document':'0045 / I10-E01 / procedure, distinct from this single-case form',
 'negative_neighbors':'0044 I09-E04 meeting/email originals; 0047 I10-E02 multi-period review',
 'synthetic_case_link':'SYN-045-01 in 0045 R01; shared scenario only, no REAL claim',
 'created_at_utc':now}
(ROOT/'inputs'/'SOURCE_MAP.json').write_text(json.dumps(source_map,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
# calculate scores from the raw data defined above
weighted={}
for v in vendors:
    subtotal=Decimal('0')
    for i,(_,weight) in enumerate(criteria):
        mean=(Decimal(v['raw_e01'][i]+v['raw_e02'][i])/2)
        subtotal += Decimal(weight)*mean/Decimal(5)
    weighted[v['id']]=str(subtotal.quantize(Decimal('0.1'),rounding=ROUND_HALF_UP))
receipt={
 'deliverable_id':'0046','requirement_id':'I10-E01','official_checkbox_sequence':'047','candidate_revision':'R01',
 'candidate_path':'candidate/0046_R01.docx','candidate_sha256':sha,'candidate_bytes':OUT.stat().st_size,
 'blueprint_source_sha256':source_map['unique_blueprint_source_sha256'],'blueprint_row_sha256':source_map['unique_blueprint_row_sha256'],
 'official_source_sha256':source_map['official_source_sha256'],'source_commit':source_map['source_commit'],
 'synthetic_register_rows':len(facts),'weights_sum':sum(w for _,w in criteria),'recomputed_weighted_scores':weighted,
 'semantic_review':'PENDING','visual_review':'PENDING','promotion':'NOT_PROMOTED','approval':'NONE_REPRESENTED',
 'generated_utc':now,'limitations':['All populated transaction, role, date, offer, score, decision, contract, delivery, and outcome facts are SAMPLE/SYNTHETIC.',
 'No Association approval, signature, real offer, procurement, contract, invoice, payment, attendance, delivery, or actual training is asserted.',
 'Weights, price tiers, threshold, budget cap, and retention assumptions require authorized Association review before real use.',
 'Independent exact-SHA semantic review remains pending.']}
(ROOT/'evidence'/'build_receipt_R01.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
progress={
 'project_id':'TTQS_ONE','document_id':'0046','requirement_id':'I10-E01','official_checkbox_sequence':'047',
 'state':'CANDIDATE_BUILT_QA_PENDING','candidate_sha256':sha,'source_commit':source_map['source_commit'],
 'synthetic_register_rows':len(facts),'promotion':'NOT_PROMOTED','independent_semantic_review':'PENDING',
 'next_actions':['Render and inspect every page','Freeze exact SHA and technical receipt','Obtain valid independent exact-SHA semantic review','Sync complete package to existing GitHub branch and read back bytes']}
(ROOT/'PROGRESS.json').write_text(json.dumps(progress,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'docx_sha256':sha,'bytes':OUT.stat().st_size,'synthetic_register_rows':len(facts),'weighted_scores':weighted,'page_break_sections':8,'created_utc':now},ensure_ascii=True,indent=2))
