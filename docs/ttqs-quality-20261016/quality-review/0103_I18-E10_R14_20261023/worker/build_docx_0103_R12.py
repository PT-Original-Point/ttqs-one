from __future__ import annotations
import argparse
import re
from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

NAVY = "1F4E78"
INK = "222222"
PALE = "EAF1F7"
ALT = "F5F8FB"
BORDER = "C7D2DC"


def set_east_asia(run, font="Microsoft JhengHei"):
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), font)


def set_style_font(style, size, bold=False, color=INK):
    style.font.name = "Microsoft JhengHei"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), "Microsoft JhengHei")


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=80, start=85, bottom=80, end=85):
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    tc_mar = tcpr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tcpr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row):
    trpr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trpr.append(tbl_header)


def keep_row_together(row):
    trpr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trpr.append(cant)


def add_field(paragraph, name):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {name} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_east_asia(run)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("65717C")


def apply_section(sec, landscape=False):
    sec.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    if landscape:
        sec.page_width, sec.page_height = Mm(297), Mm(210)
        sec.left_margin = sec.right_margin = Mm(15)
        sec.top_margin = Mm(15)
        sec.bottom_margin = Mm(15)
    else:
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin = sec.right_margin = Mm(17)
        sec.top_margin = Mm(17)
        sec.bottom_margin = Mm(16)
    sec.header_distance = Mm(8)
    sec.footer_distance = Mm(8)


def install_header_footer(sec):
    sec.header.is_linked_to_previous = False
    hp = sec.header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run("感謝函取得與真偽查核工作紀錄　｜　0103")
    set_east_asia(hr)
    hr.font.size = Pt(8)
    hr.font.color.rgb = RGBColor.from_string("65717C")

    sec.footer.is_linked_to_previous = False
    fp = sec.footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    fr = fp.add_run("真實原件未提供　｜　正式驗收未授予　｜　第 ")
    set_east_asia(fr)
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor.from_string("65717C")
    add_field(fp, "PAGE")
    fr2 = fp.add_run(" 頁")
    set_east_asia(fr2)
    fr2.font.size = Pt(8)
    fr2.font.color.rgb = RGBColor.from_string("65717C")


def add_inline(paragraph, text, base_size=None, color=None):
    token_re = re.compile(r"(\*\*.+?\*\*|`[^`]+`|(?<!\*)\*[^*]+\*(?!\*))")
    pos = 0
    for m in token_re.finditer(text):
        if m.start() > pos:
            r = paragraph.add_run(text[pos:m.start()])
            set_east_asia(r)
            if base_size: r.font.size = Pt(base_size)
            if color: r.font.color.rgb = RGBColor.from_string(color)
        token = m.group(0)
        if token.startswith("**"):
            r = paragraph.add_run(token[2:-2]); r.bold = True
        elif token.startswith("`"):
            r = paragraph.add_run(token[1:-1]); r.font.name = "Consolas"
            r.font.size = Pt((base_size or 10.5) - 0.4)
            r.font.color.rgb = RGBColor.from_string("1F3A52")
            rpr = r._element.get_or_add_rPr()
            rfonts = rpr.rFonts
            if rfonts is None:
                rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
            rfonts.set(qn("w:ascii"), "Consolas"); rfonts.set(qn("w:hAnsi"), "Consolas")
            rfonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        else:
            r = paragraph.add_run(token[1:-1]); r.italic = True
        set_east_asia(r)
        if base_size and not token.startswith("`"): r.font.size = Pt(base_size)
        if color: r.font.color.rgb = RGBColor.from_string(color)
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:]); set_east_asia(r)
        if base_size: r.font.size = Pt(base_size)
        if color: r.font.color.rgb = RGBColor.from_string(color)


def parse_cells(line):
    raw = line.strip()
    if raw.startswith("|"): raw = raw[1:]
    if raw.endswith("|"): raw = raw[:-1]
    cells=[]; current=[]; escaped=False
    for ch in raw:
        if ch == "|" and not escaped:
            cells.append("".join(current).strip()); current=[]
        else:
            current.append(ch)
        escaped = (ch == "\\" and not escaped)
        if ch != "\\": escaped=False
    cells.append("".join(current).strip())
    return [c.replace("\\|", "|") for c in cells]


def is_table_delimiter(line):
    return bool(re.match(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?$", line.strip()))


def heading_text(raw):
    s = re.sub(r"^#+\s*", "", raw).strip()
    s = re.sub(r"^\d+(?:\.\d+)*[.．、]?\s*", "", s)
    if s == "SAMPLE TEMPLATE - NOT SENT":
        return s
    s = s.replace("：", " ").replace(":", " ").replace("｜", " ").replace("|", " ")
    s = s.replace("-", " ").replace("／", " ").replace("/", " ")
    s = re.sub(r"[。；;，,：:（）()【】\[\]<>「」『』]+", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def add_markdown_table(doc, rows, font_size, available_width):
    if not rows: return
    ncols = max(len(r) for r in rows)
    normalized = [r + [""] * (ncols-len(r)) for r in rows]
    table = doc.add_table(rows=len(normalized), cols=ncols)
    table.style = "Table Grid"
    table.autofit = False
    each = int(available_width / ncols)
    for col in table.columns:
        col.width = each
    for ri, row_data in enumerate(normalized):
        row = table.rows[ri]
        keep_row_together(row)
        for ci, value in enumerate(row_data):
            cell = row.cells[ci]
            cell.width = each
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
            if ri == 0:
                set_cell_shading(cell, NAVY)
            elif ri % 2 == 0:
                set_cell_shading(cell, ALT)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.03
            add_inline(p, value.replace("<br>", "\n").replace("<br/>", "\n"), base_size=font_size, color=("FFFFFF" if ri == 0 else INK))
            for run in p.runs:
                run.font.size = Pt(font_size)
                if ri == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255,255,255)
        if ri == 0: repeat_table_header(row)
    if ncols >= 6 and len(table.rows) >= 3:
        # Keep the final two rows together so a short final exception row cannot orphan.
        penultimate_last = table.rows[-2].cells[-1].paragraphs[-1]
        penultimate_last.paragraph_format.keep_with_next = True
    # Keep table width within the active section and prevent accidental overflow.
    tblpr = table._tbl.tblPr
    layout = tblpr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout"); tblpr.append(layout)
    layout.set(qn("w:type"), "fixed")
    return table


def compute_landscape_regions(lines):
    """Return layout ranges based on wide tables and only-effective context markers."""
    table_spans=[]
    i=0
    while i < len(lines)-1:
        if lines[i].strip().startswith("|") and is_table_delimiter(lines[i+1]):
            start=i
            rows=[parse_cells(lines[i])]
            i+=2
            while i<len(lines) and lines[i].strip().startswith("|"):
                rows.append(parse_cells(lines[i])); i+=1
            table_spans.append((start,i-1,max(map(len,rows))))
        else:
            i+=1
    markers=[i for i,line in enumerate(lines) if line.strip()=="<!-- LANDSCAPE_CONTEXT -->"]
    regions=[]
    effective_markers=set()
    associated_wide_starts=set()
    for marker in markers:
        next_table=next((t for t in table_spans if t[0]>marker),None)
        if next_table and next_table[2]>=6:
            first_content=marker+1
            while first_content<len(lines) and (not lines[first_content].strip() or lines[first_content].strip()=="---" or lines[first_content].strip()=="<!-- LANDSCAPE_CONTEXT -->"):
                first_content+=1
            target_heading=[n for n in range(first_content,next_table[0]) if re.match(r"^#{2,3}\s+",lines[n].strip())]
            # Start at the heading that directly owns the wide table. This keeps
            # preceding portrait-only narrative (for example 9.4) out of the table page.
            start=target_heading[-1] if target_heading else first_content
            region_end=next_table[1]
            # Keep the explanation of section 8's metrics with its table,
            # and the short post-table rule in 9.5 with the landscape exception
            # table. Stop at each subsection boundary so the next topic begins
            # cleanly in portrait orientation.
            owner_heading=lines[start].strip()
            stop_pattern=None
            if re.match(r"^##\s*8\.",owner_heading):
                stop_pattern=r"^##\s+"
            elif re.match(r"^###\s*9\.5\b",owner_heading):
                stop_pattern=r"^#{1,3}\s+"
            if stop_pattern:
                following_heading=next((n for n in range(region_end+1,len(lines))
                                         if re.match(stop_pattern,lines[n].strip())),None)
                if following_heading is not None:
                    region_end=following_heading-1
            regions.append((start,region_end))
            effective_markers.add(marker)
            associated_wide_starts.add(next_table[0])
    for table_start,table_end,cols in table_spans:
        if cols<6 or table_start in associated_wide_starts:
            continue
        previous_end=max((end for begin,end,_ in table_spans if end<table_start),default=-1)
        headings=[n for n in range(previous_end+1,table_start) if re.match(r"^#{2,3}\s+",lines[n].strip())]
        start=headings[-1] if headings else table_start
        regions.append((start,table_end))
    mask=[False]*len(lines)
    starts=set()
    for start,end in regions:
        start=max(0,start); end=min(len(lines)-1,end)
        starts.add(start)
        for n in range(start,end+1): mask[n]=True
    return mask,starts,effective_markers


def build(markdown_path: Path, output_path: Path, dry_run=False):
    text = markdown_path.read_text(encoding="utf-8-sig")
    lines = text.splitlines()
    tables=[]; scan=0
    while scan < len(lines)-1:
        if lines[scan].strip().startswith("|") and is_table_delimiter(lines[scan+1]):
            start=scan; rows=[parse_cells(lines[scan])]; scan += 2
            while scan < len(lines) and lines[scan].strip().startswith("|"):
                rows.append(parse_cells(lines[scan])); scan += 1
            tables.append({"line":start,"columns":max(map(len,rows)),"body_rows":len(rows)-1})
        else: scan += 1
    if dry_run:
        print(f"chars={len(text)} tables={len(tables)} wide_tables={sum(t['columns']>=6 for t in tables)}")
        for t in tables: print(t)
        return

    doc=Document()
    for style_name, size, bold, color in (("Normal",10.5,False,INK),("Title",21,False,INK),("Heading 1",15,True,NAVY),("Heading 2",12.8,True,NAVY),("Heading 3",11.5,True,"38566F")):
        st=doc.styles[style_name]; set_style_font(st,size,bold,color)
        st.paragraph_format.space_after=Pt(5 if style_name=="Normal" else 7)
        if style_name=="Normal": st.paragraph_format.line_spacing=1.12
        if style_name.startswith("Heading"):
            st.paragraph_format.space_before=Pt(12 if style_name=="Heading 1" else 8)
            st.paragraph_format.keep_with_next=True
    sec=doc.sections[0]; apply_section(sec,False); install_header_footer(sec)
    # Remove title-style border residue if present.
    title_style=doc.styles["Title"].element
    ppr=title_style.find(qn("w:pPr"))
    if ppr is not None:
        pbr=ppr.find(qn("w:pBdr"))
        if pbr is not None: ppr.remove(pbr)
    landscape_mask,landscape_starts,effective_markers=compute_landscape_regions(lines)
    current_landscape=False
    caption_pending=False
    compact_section=False
    compact_table_pending=False
    last_heading=""
    i=0
    while i < len(lines):
        line=lines[i].rstrip()
        if line.strip() == "<!-- LANDSCAPE_CONTEXT -->":
            if i in effective_markers:
                caption_pending=True
            i+=1
            continue
        if not line.strip() or line.strip()=="---": i+=1; continue
        wanted_landscape=landscape_mask[i]
        if wanted_landscape and not current_landscape:
            sec=doc.add_section(WD_SECTION_START.NEW_PAGE); apply_section(sec,True); install_header_footer(sec); current_landscape=True
            caption_pending=True
        elif not wanted_landscape and current_landscape:
            sec=doc.add_section(WD_SECTION_START.NEW_PAGE); apply_section(sec,False); install_header_footer(sec); current_landscape=False
            caption_pending=False
        if line.strip().startswith("|") and i+1<len(lines) and is_table_delimiter(lines[i+1]):
            rows=[parse_cells(line)]; i+=2
            while i<len(lines) and lines[i].strip().startswith("|"):
                rows.append(parse_cells(lines[i])); i+=1
            cols=max(map(len,rows))
            wide=cols>=6
            if wide and not current_landscape:
                sec=doc.add_section(WD_SECTION_START.NEW_PAGE); apply_section(sec,True); install_header_footer(sec); current_landscape=True
                caption_pending=True
            if not wide and current_landscape:
                sec=doc.add_section(WD_SECTION_START.NEW_PAGE); apply_section(sec,False); install_header_footer(sec); current_landscape=False
                caption_pending=False
            if wide and caption_pending:
                cap=doc.add_paragraph()
                cap.paragraph_format.space_after=Pt(4)
                cap.paragraph_format.keep_with_next=True
                rr=cap.add_run((last_heading or "表格")+"　寬版表格")
                set_east_asia(rr); rr.font.size=Pt(9); rr.font.italic=True; rr.font.color.rgb=RGBColor.from_string("65717C")
                caption_pending=False
            sec_now=doc.sections[-1]
            available = sec_now.page_width - sec_now.left_margin - sec_now.right_margin
            fs=7.9 if cols>=9 else (8.2 if cols>=7 else (8.7 if cols>=5 else 9.0))
            add_markdown_table(doc, rows, fs, available)
            if compact_table_pending:
                for row in doc.tables[-1].rows:
                    for cell in row.cells:
                        set_cell_margins(cell,top=36,start=85,bottom=36,end=85)
                        for cell_p in cell.paragraphs:
                            cell_p.paragraph_format.space_after=Pt(0.5)
                            cell_p.paragraph_format.line_spacing=1.0
                compact_table_pending=False
            p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(1); p.paragraph_format.space_before=Pt(0)
            continue
        if line.strip() == "---": i+=1; continue
        if line.startswith("# "):
            p=doc.add_paragraph(style="Title"); p.paragraph_format.space_after=Pt(9)
            add_inline(p,heading_text(line),base_size=21,color=INK)
            i+=1; continue
        hm=re.match(r"^(#{2,4})\s+(.*)$",line)
        if hm:
            level=len(hm.group(1)); heading=heading_text(line); last_heading=heading
            if level==2:
                compact_section=heading.startswith("9.")
            if level==3 and heading.startswith("9.1"):
                compact_table_pending=True
            style_name={2:"Heading 1",3:"Heading 2",4:"Heading 3"}[level]
            p=doc.add_paragraph(style=style_name)
            if level==4:
                p.paragraph_format.space_before=Pt(4)
                p.paragraph_format.space_after=Pt(2)
            if compact_section:
                p.paragraph_format.space_before=Pt(4)
                p.paragraph_format.space_after=Pt(2)
            font_size={2:15,3:12.8,4:11.5}[level]
            heading_color=("38566F" if level==4 else NAVY)
            add_inline(p,heading,base_size=font_size,color=heading_color)
            i+=1; continue
        if line.lstrip().startswith(">"): 
            quote=line.lstrip()[1:].strip()
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Mm(4); p.paragraph_format.right_indent=Mm(2)
            p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(6)
            ppr=p._p.get_or_add_pPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),PALE); ppr.append(shd)
            add_inline(p,quote,base_size=10.2,color=INK)
            i+=1; continue
        bullet=re.match(r"^\s*-\s+(.*)$",line)
        if bullet:
            item=bullet.group(1)
            if item.startswith("[ ] "): item="[ ]　"+item[4:]
            elif item.startswith("[x] "): item="[x]　"+item[4:]
            p=doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent=Mm(7); p.paragraph_format.first_line_indent=Mm(-3)
            p.paragraph_format.space_after=Pt(0.5 if compact_section else 3)
            if compact_section: p.paragraph_format.line_spacing=1.0
            add_inline(p,item,base_size=10.2)
            i+=1; continue
        ordered=re.match(r"^\s*(\d+)[.)]\s+(.*)$",line)
        if ordered:
            p=doc.add_paragraph()
            p.paragraph_format.left_indent=Mm(7); p.paragraph_format.first_line_indent=Mm(-7)
            p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.1
            n=p.add_run(ordered.group(1)+". "); n.bold=True; n.font.color.rgb=RGBColor.from_string(NAVY); set_east_asia(n)
            add_inline(p,ordered.group(2),base_size=10.2)
            i+=1; continue
        p=doc.add_paragraph()
        p.paragraph_format.space_after=Pt(2 if compact_section else 4)
        p.paragraph_format.line_spacing=(1.05 if compact_section else 1.12)
        add_inline(p,line,base_size=10.5)
        i+=1

    # Add consistent document properties without internal build or repository metadata.
    props=doc.core_properties
    props.title="感謝函取得與真偽查核工作紀錄"
    props.subject="0103 I18-E10 感謝函真實來源取得與查核作業"
    props.author=""
    props.last_modified_by=""
    output_path.parent.mkdir(parents=True,exist_ok=True)
    doc.save(output_path)
    print(f"saved={output_path}")
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)}")
    print(f"wide_tables={sum(t['columns']>=6 for t in tables)} source_chars={len(text)}")


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("input",type=Path)
    ap.add_argument("output",type=Path)
    ap.add_argument("--dry-run",action="store_true")
    args=ap.parse_args()
    build(args.input,args.output,args.dry_run)

if __name__ == "__main__": main()
