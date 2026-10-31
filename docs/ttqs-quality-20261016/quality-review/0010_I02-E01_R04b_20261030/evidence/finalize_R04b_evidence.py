from __future__ import annotations

import hashlib
import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from docx import Document

ROOT = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0010_I02-E01_ASSESSMENT_STANDARD_R04_20261030")
R03 = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0010_I02-E01_ASSESSMENT_STANDARD_R03_20261029")
DOCX = ROOT / "candidate" / "0010_I02-E01_R04b.docx"
PDF = ROOT / "qa" / "render_R04b_native" / "0010_I02-E01_R04b.pdf"
RENDER_RECEIPT = ROOT / "qa" / "render_R04b_native" / "WORD_NATIVE_RENDER_RECEIPT_R04b.json"
TEXT = ROOT / "worker" / "0010_I02-E01_R04b.md"
REGISTER = ROOT / "worker" / "SYNTHETIC_REGISTER_R04.md"
BINDING = ROOT / "source" / "SOURCE_BINDING_R04.json"
REQ = ROOT / "source" / "OFFICIAL_REQUIREMENT_EXTRACT_R04.txt"
BLUEPRINT = ROOT / "source" / "BLUEPRINT_ROW_0010_I02-E01_R04.json"
TECH = ROOT / "qa" / "TECHNICAL_QA_R04b.json"
CHECK_SCRIPT = ROOT / "qa" / "technical_visual_checks_R04b.py"
BUILD_SCRIPT = ROOT / "qa" / "build_0010_r04b_docx.py"
OLD_START = ROOT / "reviewer" / "OPENCODE_R04" / "DISPATCH_START_R04.json"
OLD_END = ROOT / "reviewer" / "OPENCODE_R04" / "DISPATCH_TERMINAL_R04.json"
OLD_ERR = ROOT / "reviewer" / "OPENCODE_R04" / "RAW_STDERR_R04.txt"
OLD_OUT = ROOT / "reviewer" / "OPENCODE_R04" / "RAW_STDOUT_R04.jsonl"
REVIEW_DIR = ROOT / "reviewer" / "OPENCODE_R04b"
REVIEW_DIR.mkdir(parents=True, exist_ok=True)


def sha_bytes(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


def sha(p: Path) -> str:
    return sha_bytes(p.read_bytes())


def write_json(path: Path, obj: dict):
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")


now = datetime.now().astimezone()
docx_sha, pdf_sha, text_sha, register_sha = map(sha, [DOCX, PDF, TEXT, REGISTER])
binding_obj = json.loads(BINDING.read_text(encoding="utf-8"))
render_obj = json.loads(RENDER_RECEIPT.read_text(encoding="utf-8"))
tech_obj = json.loads(TECH.read_text(encoding="utf-8"))
old_end = json.loads(OLD_END.read_text(encoding="utf-8"))
retry_ms = None
try:
    raw_out_text = OLD_OUT.read_text(encoding="utf-8")
    old_event = json.loads(raw_out_text.splitlines()[0])
    retry_ms = old_event.get("error", {}).get("retryAfterMs")
except Exception:
    old_event = None
retry_after = None
if retry_ms is not None:
    retry_after = datetime.fromtimestamp(old_end["ended_at"] if isinstance(old_end["ended_at"], (int, float)) else
        datetime.fromisoformat(old_end["ended_at"].replace("Z", "+00:00")).timestamp(), timezone.utc) + timedelta(milliseconds=retry_ms)

# Stable, plain-text extraction of the exact Word candidate for a text-only reviewer.
doc = Document(DOCX)
extract_lines = ["# Extracted text from exact Word candidate", f"docx_sha256={docx_sha}", ""]
for i, p in enumerate(doc.paragraphs, 1):
    if p.text.strip():
        extract_lines.append(f"[P{i}] {p.text}")
for ti, table in enumerate(doc.tables, 1):
    extract_lines.extend(["", f"[TABLE {ti}: {len(table.rows)} rows x {len(table.columns)} columns]"])
    for ri, row in enumerate(table.rows, 1):
        extract_lines.append(f"[T{ti}R{ri}] " + " | ".join(c.text.replace("\n", " / ").strip() for c in row.cells))
extract_path = REVIEW_DIR / "DOCX_TEXT_EXTRACT_R04b.txt"
extract_path.write_text("\n".join(extract_lines) + "\n", encoding="utf-8", newline="\n")
extract_sha = sha(extract_path)

png_dir = ROOT / "qa" / "render_R04b_native" / "png"
contact_dir = ROOT / "qa" / "render_R04b_native" / "contact"
pngs = sorted(png_dir.glob("page-*.png"), key=lambda p: int(p.stem.split("-")[-1]))
contacts = sorted(contact_dir.glob("sheet-*.jpg"))
selected_pages = [5, 14, 16, 24, 27, 28]
visual = {
    "receipt_type": "visual_review_R04b",
    "doc_id": "0010",
    "requirement_id": "I02-E01",
    "candidate_sha256": docx_sha,
    "native_pdf_sha256": pdf_sha,
    "native_render_receipt_sha256": sha(RENDER_RECEIPT),
    "page_count": len(pngs),
    "all_page_rasters": [{"page": int(p.stem.split("-")[-1]), "sha256": sha(p)} for p in pngs],
    "contact_sheets": [{"name": p.name, "sha256": sha(p)} for p in contacts],
    "inspection_scope": {
        "contact_sheet_pages": "1-28, all seven contact sheets inspected",
        "full_resolution_pages": selected_pages,
        "observations": [
            "Page headings, repeated register headers, table boundaries and footers are consistent across the 28-page native Word render.",
            "Dense rubric and six-column replacement-register tables remain within page bounds; sampled full-resolution cells are legible and not clipped.",
            "Page 28 continues the last two synthetic-register rows under a repeated header. It has substantial lower-page whitespace but is not blank; both rows are readable and not split or clipped.",
            "Pages 6 and 11 have short section content with visible whitespace; no overflow, overlap or clipped content was observed."
        ]
    },
    "visual_verdict": "PASS_WITH_LAYOUT_NOTE",
    "reviewer_role": "primary visual inspection; not an independent semantic review",
    "checked_at_taipei": now.isoformat(timespec="seconds")
}
visual_path = ROOT / "qa" / "VISUAL_QA_R04b.json"
write_json(visual_path, visual)

build_receipt = {
    "receipt_type": "word_build_receipt_R04b",
    "doc_id": "0010",
    "requirement_id": "I02-E01",
    "source_markdown_sha256": text_sha,
    "builder_script_sha256": sha(BUILD_SCRIPT),
    "candidate_docx_sha256": docx_sha,
    "candidate_bytes": DOCX.stat().st_size,
    "native_pdf_sha256": pdf_sha,
    "native_page_count": render_obj.get("page_count"),
    "built_at_taipei": now.isoformat(timespec="seconds"),
    "candidate_state": "CANDIDATE_SAMPLE_SYNTHETIC_NOT_ACCEPTED"
}
build_path = ROOT / "qa" / "BUILD_RECEIPT_R04b.json"
write_json(build_path, build_receipt)

request = f"""# 0010 / I02-E01 R04b exact-SHA independent semantic review

Requested route: OpenCode v2 CLI `opencode/deepseek-v4.1-flash-free` when its recorded provider cooldown has expired. Review only the attached text packet; no tools, browsing, file edits, or claims of unperformed local checks. Return one JSON object matching the schema below.

Exact Word candidate SHA-256: `{docx_sha}`.
Source Markdown SHA-256: `{text_sha}`.
Word text extraction SHA-256: `{extract_sha}`.
SAMPLE/SYNTHETIC register SHA-256: `{register_sha}`.
Official requirement: `{binding_obj['official_requirement']}` (printed p.1, indicator 2).
Unique Blueprint file SHA-256: `{binding_obj['blueprint_file_sha256']}`; row SHA-256: `{binding_obj['blueprint_row_sha256']}`.

Review the complete R04b body, extracted Word text and all 157 register rows. Verify R03 semantic findings and R04 changes remain repaired; question/LO allocation, 0–4 rubric, case and retest math, thresholds, atomic fact locators, case/retest provenance, title-blind genre and rejection of negative neighbors 0009/I01-E06 and 0011/I02-E02; practical use by a small volunteer association; consistency between Markdown, extracted Word text, register and exact candidate SHA; and SAMPLE/REAL boundary. Identify precise section/table/field locators for every finding. Do not promote unresolved Blueprint outward-policy/public-disclosure scope beyond PARTIAL/UNKNOWN. No actual learner, roster, authorizer, assessment or association approval evidence was supplied. Do not call synthetic material REAL, adopted, effective or formally accepted.

Technical/visual receipts are supplied for context but do not substitute for this independent semantic review. Do not perform layout review unless a concrete text-layout issue appears in the extracted content; the Word render is separately bound to the exact DOCX SHA.

JSON schema:
{{
  \"receipt_type\":\"independent_exact_sha_semantic_review\",
  \"review_id\":\"0010-I02-E01-R04b\",
  \"reviewed_at_utc\":\"ISO-8601\",
  \"candidate_sha256\":\"{docx_sha}\",
  \"source_markdown_sha256\":\"{text_sha}\",
  \"extracted_text_sha256\":\"{extract_sha}\",
  \"register_sha256\":\"{register_sha}\",
  \"verdict\":\"PASS|PARTIAL|FAIL\",
  \"scope\":\"...\",
  \"checks\":[{{\"id\":\"...\",\"result\":\"PASS|PARTIAL|FAIL\",\"locator\":\"...\",\"reason\":\"...\"}}],
  \"findings\":[{{\"id\":\"...\",\"severity\":\"critical|major|moderate|minor\",\"locator\":\"...\",\"evidence\":\"...\",\"required_repair\":\"...\"}}],
  \"limitations\":[\"...\"],
  \"does_not_claim\":[\"REAL evidence\",\"formal acceptance\",\"unresolved Blueprint scope\"]
}}
"""
request_path = REVIEW_DIR / "REVIEW_REQUEST_R04b.md"
request_path.write_text(request, encoding="utf-8", newline="\n")

parts = [
    ("Review request bound to exact R04b Word SHA", request_path),
    ("Official requirement extract", ROOT / "source" / "OFFICIAL_REQUIREMENT_EXTRACT_R04.txt"),
    ("Exact current unique Blueprint row", ROOT / "source" / "BLUEPRINT_ROW_0010_I02-E01_R04.json"),
    ("Current source binding", ROOT / "source" / "SOURCE_BINDING_R04.json"),
    ("R03 independent semantic review", R03 / "reviewer" / "SEMANTIC_PRE_REVIEW_R03.json"),
    ("R03 deterministic QA", R03 / "qa" / "DETERMINISTIC_MATH_AND_REGISTER_R03.json"),
    ("R04 deterministic QA", ROOT / "qa" / "DETERMINISTIC_QA_R04.json"),
    ("R04 repair receipt", ROOT / "qa" / "R04_REPAIR_RECEIPT.json"),
    ("R04 writer receipt reconciliation", ROOT / "worker" / "WRITER_RECEIPT_RECONCILIATION_R04.json"),
    ("Exact R04b source body", TEXT),
    ("Exact R04b Word extraction", extract_path),
    ("Complete SAMPLE/SYNTHETIC register sidecar", REGISTER),
    ("Candidate-specific deterministic QA", TECH),
    ("Candidate-specific visual QA", visual_path),
]
packet_parts = ["# 0010/I02-E01 R04b independent semantic review packet", ""]
for title, path in parts:
    raw = path.read_bytes()
    body = raw.decode("utf-8")
    packet_parts.extend([f"## {title}", f"source_path={path}", f"source_sha256={sha_bytes(raw)}", f"source_bytes={len(raw)}", "", body.rstrip("\r\n"), ""])
packet_path = REVIEW_DIR / "REVIEW_PACKET_R04b.md"
packet_path.write_text("\n".join(packet_parts), encoding="utf-8", newline="\n")

route_status = {
    "receipt_type": "review_route_status_R04b",
    "work_id": "TTQS-0010-I02-E01-R04b-SEMANTIC",
    "doc_id": "0010",
    "requirement_id": "I02-E01",
    "candidate_sha256": docx_sha,
    "review_packet_sha256": sha(packet_path),
    "current_review_request_sha256": sha(request_path),
    "route": "OpenCode v2 CLI opencode/deepseek-v4.1-flash-free",
    "historical_attempt": {
        "work_id": old_end.get("work_id"),
        "candidate_markdown_sha256": "131b703267feebfbef3ec49fc8aa9beaab2ded0818415e9e5d86771b84925d90",
        "packet_sha256": old_end.get("packet_sha256"),
        "exit_code": old_end.get("exit_code"),
        "stdout_sha256": old_end.get("stdout_sha256"),
        "stderr_sha256": old_end.get("stderr_sha256"),
        "stderr_bytes": old_end.get("stderr_bytes"),
        "provider_error": old_event.get("error") if old_event else "UNKNOWN",
        "applies_to_R04b": False
    },
    "current_review_status": "PENDING_NOT_DISPATCHED_R04b",
    "reason": "The prior single attempt was for predecessor R04, returned provider HTTP 429/free-limit with retryAfterMs=30381000 and no semantic response. Its packet was bound to an older text SHA; it cannot review R04b.",
    "earliest_retry_after_utc": retry_after.isoformat() if retry_after else "UNKNOWN",
    "same_input_retry_performed": False,
    "fallback_or_paid_route_used": False,
    "checked_at_taipei": now.isoformat(timespec="seconds")
}
route_path = REVIEW_DIR / "REVIEW_ROUTE_STATUS_R04b.json"
write_json(route_path, route_status)

all_outputs = [
    TECH, ROOT / "qa" / "VISUAL_QA_R04b.json", build_path, extract_path,
    request_path, packet_path, route_path
]
print(json.dumps({
    "docx_sha256": docx_sha,
    "text_sha256": text_sha,
    "register_sha256": register_sha,
    "pdf_sha256": pdf_sha,
    "visual_receipt_sha256": sha(visual_path),
    "build_receipt_sha256": sha(build_path),
    "review_packet_sha256": sha(packet_path),
    "review_route_status": route_status["current_review_status"],
    "earliest_retry_after_utc": route_status["earliest_retry_after_utc"],
    "outputs": [{"path": str(p), "sha256": sha(p), "bytes": p.stat().st_size} for p in all_outputs]
}, ensure_ascii=False, indent=2))
