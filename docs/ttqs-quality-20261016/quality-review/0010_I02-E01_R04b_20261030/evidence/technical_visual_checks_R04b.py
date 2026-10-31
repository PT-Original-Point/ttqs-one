from __future__ import annotations

import hashlib
import json
import re
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from pypdf import PdfReader

ROOT = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0010_I02-E01_ASSESSMENT_STANDARD_R04_20261030")
DOCX = ROOT / "candidate" / "0010_I02-E01_R04b.docx"
SOURCE_MD = ROOT / "worker" / "0010_I02-E01_R04b.md"
REGISTER = ROOT / "worker" / "SYNTHETIC_REGISTER_R04.md"
BINDING = ROOT / "source" / "SOURCE_BINDING_R04.json"
OFFICIAL_EXTRACT = ROOT / "source" / "OFFICIAL_REQUIREMENT_EXTRACT_R04.txt"
BLUEPRINT_COPY = ROOT / "source" / "BLUEPRINT_ROW_0010_I02-E01_R04.json"
BLUEPRINT_CURRENT = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\CONTROL\BLUEPRINT_CURRENT.jsonl")
OFFICIAL_PDF = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\SOURCES\TTQS查核佐證文件資料表-訓練機構版 (2).pdf")
RENDER_RECEIPT = ROOT / "qa" / "render_R04b_native" / "WORD_NATIVE_RENDER_RECEIPT_R04b.json"
PDF = ROOT / "qa" / "render_R04b_native" / "0010_I02-E01_R04b.pdf"
PNG_DIR = ROOT / "qa" / "render_R04b_native" / "png"
CONTACT_DIR = ROOT / "qa" / "render_R04b_native" / "contact"
OUT = ROOT / "qa" / "TECHNICAL_QA_R04b.json"


def sha(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def check(checks: dict, key: str, passed: bool, detail=None):
    checks[key] = {"pass": bool(passed), "detail": detail}


binding = json.loads(BINDING.read_text(encoding="utf-8"))
render = json.loads(RENDER_RECEIPT.read_text(encoding="utf-8"))
doc = Document(DOCX)
markdown = SOURCE_MD.read_text(encoding="utf-8")
official_text = OFFICIAL_EXTRACT.read_text(encoding="utf-8")
full_text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
    "\n".join(cell.text for row in table.rows for cell in row.cells) for table in doc.tables
)

checks: dict = {}
actual_docx_sha = sha(DOCX)
check(checks, "docx_sha_bound_to_native_render", render.get("input_docx_sha256") == actual_docx_sha,
      {"docx_sha256": actual_docx_sha, "render_input_sha256": render.get("input_docx_sha256")})
check(checks, "docx_zip_integrity", zipfile.ZipFile(DOCX).testzip() is None)
check(checks, "python_docx_parse", len(doc.tables) == 12 and len(doc.sections) == 8,
      {"paragraphs": len(doc.paragraphs), "tables": len(doc.tables), "sections": len(doc.sections)})

requirement = "對學員訓練成果有評核標準的訂定及說明。"
check(checks, "exact_official_requirement_and_locator", requirement in official_text and
      "printed page 1, indicator 2" in official_text and
      binding.get("official_requirement") == requirement and
      binding.get("official_pdf_locator") == "printed p.1, indicator 2" and
      requirement in full_text,
      {"requirement": requirement, "official_extract_sha256": sha(OFFICIAL_EXTRACT)})
check(checks, "official_pdf_hash_matches_binding",
      sha(OFFICIAL_PDF) == binding.get("official_pdf_sha256"),
      {"actual": sha(OFFICIAL_PDF), "bound": binding.get("official_pdf_sha256")})

blueprint_lines = BLUEPRINT_CURRENT.read_bytes().splitlines()
blueprint_rows = []
for raw in blueprint_lines:
    try:
        obj = json.loads(raw.decode("utf-8"))
    except Exception:
        continue
    if obj.get("deliverable_id") == "0010" and obj.get("requirement_id") == "I02-E01":
        blueprint_rows.append((raw, obj))
copy_obj = json.loads(BLUEPRINT_COPY.read_text(encoding="utf-8"))
row_hash = hashlib.sha256(blueprint_rows[0][0]).hexdigest() if len(blueprint_rows) == 1 else None
check(checks, "unique_blueprint_file_and_row_hashes_match_binding",
      len(blueprint_rows) == 1 and sha(BLUEPRINT_CURRENT) == binding.get("blueprint_file_sha256") and
      row_hash == binding.get("blueprint_row_sha256") and copy_obj == blueprint_rows[0][1],
      {"matching_rows": len(blueprint_rows), "current_file_sha256": sha(BLUEPRINT_CURRENT),
       "bound_file_sha256": binding.get("blueprint_file_sha256"), "row_sha256": row_hash,
       "bound_row_sha256": binding.get("blueprint_row_sha256"), "copy_matches_current_row": bool(blueprint_rows) and copy_obj == blueprint_rows[0][1]})

check(checks, "candidate_status_and_real_boundary",
      "CANDIDATE" in full_text and "SAMPLE/SYNTHETIC" in full_text and "未核准、未生效" in full_text and
      "未建 DOCX" not in full_text,
      {"contains_candidate": "CANDIDATE" in full_text, "contains_sample_synthetic": "SAMPLE/SYNTHETIC" in full_text,
       "contains_not_approved_not_effective": "未核准、未生效" in full_text,
       "stale_not_built_phrase_present": "未建 DOCX" in full_text})

table_rows = [[cell.text.strip() for cell in row.cells] for row in doc.tables[-1].rows]
register_lines = [line for line in REGISTER.read_text(encoding="utf-8").splitlines() if line.startswith("|")]
register_rows = [[part.strip() for part in line.strip().strip("|").split("|")] for line in register_lines]
register_ids = [row[0] for row in register_rows[2:] if len(row) >= 6]
embedded_ids = [row[0] for row in table_rows[1:]]
check(checks, "synthetic_register_157_rows_and_exact_order_match",
      len(register_ids) == 157 and len(set(register_ids)) == 157 and embedded_ids == register_ids,
      {"sidecar_rows": len(register_ids), "sidecar_unique_ids": len(set(register_ids)),
       "embedded_rows": len(embedded_ids), "ids_equal_ordered": embedded_ids == register_ids,
       "sidecar_sha256": sha(REGISTER)})

checks_for_tables = {
    "learning_objective_rows": len(doc.tables[1].rows) == 6 and
        [r.cells[0].text.strip() for r in doc.tables[1].rows[1:5]] == ["LO1", "LO2", "LO3", "LO4"],
    "questionnaire_20_items_4_options": len(doc.tables[2].rows) == 21 and
        [r.cells[0].text.strip() for r in doc.tables[2].rows[1:]] == [f"Q{i}" for i in range(1, 21)] and
        all(len(r.cells) == 6 for r in doc.tables[2].rows),
    "rubric_12_construct_indicators": len(doc.tables[4].rows) == 13 and
        [re.match(r"^(P[1-4]|C[1-5]|M[1-3])", r.cells[0].text.strip()).group(1)
         if re.match(r"^(P[1-4]|C[1-5]|M[1-3])", r.cells[0].text.strip()) else ""
         for r in doc.tables[4].rows[1:]] ==
        [f"P{i}" for i in range(1, 5)] + [f"C{i}" for i in range(1, 6)] + [f"M{i}" for i in range(1, 4)],
    "four_synthetic_cases": len(doc.tables[5].rows) == 5 and
        [r.cells[0].text.strip() for r in doc.tables[5].rows[1:]] == ["SYN-001", "SYN-002", "SYN-003", "SYN-004"],
    "evidence_and_review_tables_complete": len(doc.tables[6].rows) == 49 and
        len(doc.tables[7].rows) == 5 and len(doc.tables[8].rows) == 3 and len(doc.tables[9].rows) == 6,
    "summary_and_register_table_present": len(doc.tables[10].rows) == 9 and len(doc.tables[11].rows) == 158,
}
for key, passed in checks_for_tables.items():
    check(checks, key, passed, {"table_row_counts": [len(t.rows) for t in doc.tables]})

# Recompute the four first-assessment case scores directly from the displayed answer vector,
# answer key, and P/C/M rubric ratings, then compare the document's displayed values.
key_match = re.search(r"答案鍵Q1[–-]Q20依序：([^。\n]+)", markdown)
answer_key = re.findall(r"\b[ABCD]\b", key_match.group(1)) if key_match else []
case_lines = [line for line in markdown.splitlines() if re.match(r"\|SYN-00[1-4]\|", line)]
case_results = []
for line in case_lines:
    cols = [x.strip() for x in line.strip().strip("|").split("|")]
    if len(cols) < 8:
        continue
    case_id, answers, knowledge_text, p_text, c_text, m_text, total_text, verdict = cols[:8]
    answers_list = [x.strip() for x in answers.split(",")]
    ints = lambda s: [int(x) for x in re.findall(r"\d+", s.split("=")[0])]
    p_vec, c_vec, m_vec = ints(p_text), ints(c_text), ints(m_text)
    computed_knowledge = sum(a == b for a, b in zip(answers_list, answer_key))
    computed_p = sum(p_vec) * 2.5
    computed_c = sum(c_vec) * 1.25
    computed_m = sum(m_vec) * 1.25
    total = computed_knowledge + computed_p + computed_c + computed_m
    shown = [float(re.split(r"[/=]", knowledge_text)[0]), float(p_text.split("=")[-1]),
             float(c_text.split("=")[-1]), float(m_text.split("=")[-1]), float(total_text)]
    computed = [float(computed_knowledge), computed_p, computed_c, computed_m, total]
    thresholds = total >= 70 and computed_knowledge >= 12 and computed_p >= 28 and len(p_vec) >= 3 and p_vec[2] >= 3
    case_results.append({"case_id": case_id, "computed": computed, "displayed": shown,
                         "arithmetic_matches": all(abs(a-b) < 1e-8 for a,b in zip(computed, shown)),
                         "displayed_outcome_matches_thresholds": (verdict == "通過") == thresholds})
check(checks, "four_case_answer_key_weighted_math_and_outcomes",
      len(answer_key) == 20 and len(case_results) == 4 and all(x["arithmetic_matches"] and x["displayed_outcome_matches_thresholds"] for x in case_results),
      {"answer_key_count": len(answer_key), "cases": case_results})

retest_math = {
    "SYN-003": 17 + 30 + 22.50 + 12.50,
    "SYN-004": 14 + 32.50 + 22.50 + 11.25,
}
check(checks, "retest_arithmetic_matches_written_results",
      abs(retest_math["SYN-003"] - 82.00) < 1e-8 and abs(retest_math["SYN-004"] - 80.25) < 1e-8 and
      "82.00=17+30+22.50+12.50" in markdown and "80.25=14+32.50+22.50+11.25" in markdown,
      {k: v for k, v in retest_math.items()})

retest_results = []
retest_lines = [line for line in markdown.splitlines() if re.match(r"\|RET-SYN-00[34]\|", line)]
for line in retest_lines:
    cols = [x.strip() for x in line.strip().strip("|").split("|")]
    if len(cols) < 4:
        continue
    case_id = cols[1]
    answers_list = [x.strip() for x in cols[2].split(",")]
    shown_knowledge = int(cols[3].split("/")[0])
    computed_knowledge = sum(a == b for a, b in zip(answers_list, answer_key))
    if case_id == "SYN-003":
        p_vec, c_score, m_score = [3, 3, 3, 3], 22.50, 12.50
        expected_total = 82.00
    else:
        p_vec, c_score, m_score = [4, 3, 3, 3], 22.50, 11.25
        expected_total = 80.25
    practical = sum(p_vec) * 2.5
    total = computed_knowledge + practical + c_score + m_score
    retest_results.append({"case_id": case_id, "answer_vector_count": len(answers_list),
                           "displayed_knowledge": shown_knowledge, "computed_knowledge": computed_knowledge,
                           "practical": practical, "P3": p_vec[2], "total": total,
                           "expected_total": expected_total,
                           "all_four_thresholds_pass": total >= 70 and computed_knowledge >= 12 and practical >= 28 and p_vec[2] >= 3,
                           "arithmetic_matches": len(answers_list) == 20 and computed_knowledge == shown_knowledge and abs(total-expected_total) < 1e-8})
check(checks, "retest_answers_new_ratings_provenance_and_all_thresholds",
      len(retest_results) == 2 and all(x["arithmetic_matches"] and x["all_four_thresholds_pass"] for x in retest_results) and
      "P1=4、P2=3、P4=3沿用首次SF-073/SF-074/SF-076，非新觀察" in markdown and
      "全部四項P，新的評等[3,3,3,3]" in markdown,
      {"retests": retest_results})

body_sizes, table_sizes, missing_sizes = [], [], {"body": 0, "table": 0}
def add_run_size(run, target):
    size = run.font.size.pt if run.font.size else None
    if size is None and run.style is not None and run.style.font.size is not None:
        size = run.style.font.size.pt
    if size is None:
        size = doc.styles["Normal"].font.size.pt if doc.styles["Normal"].font.size else None
    if size is None:
        missing_sizes[target] += 1
    else:
        (body_sizes if target == "body" else table_sizes).append(float(size))
for para in doc.paragraphs:
    for run in para.runs:
        add_run_size(run, "body")
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    add_run_size(run, "table")
body_min = min(body_sizes) if body_sizes else None
table_min = min(table_sizes) if table_sizes else None
check(checks, "font_floor_body_10_5_and_table_9_5",
      body_min is not None and table_min is not None and body_min >= 10.5 and table_min >= 9.5 and
      missing_sizes["body"] == 0 and missing_sizes["table"] == 0,
      {"body_min_pt": body_min, "table_min_pt": table_min, "missing_size_runs": missing_sizes})

pdf_pages = len(PdfReader(str(PDF)).pages)
pngs = sorted(PNG_DIR.glob("page-*.png"))
contacts = sorted(CONTACT_DIR.glob("sheet-*.jpg"))
check(checks, "native_word_pdf_and_all_page_rasters_match",
      render.get("opened") is True and render.get("exported") is True and
      render.get("input_docx_sha256") == actual_docx_sha and render.get("pdf_sha256") == sha(PDF) and
      render.get("page_count") == 28 and pdf_pages == 28 and len(pngs) == 28 and len(contacts) == 7,
      {"pdf_pages": pdf_pages, "png_pages": len(pngs), "contact_sheets": len(contacts),
       "pdf_sha256": sha(PDF), "render_receipt": render})

failed = [k for k, v in checks.items() if not v["pass"]]
receipt = {
    "receipt_type": "technical_and_visual_review_R04b",
    "doc_id": "0010",
    "requirement_id": "I02-E01",
    "candidate_state": "CANDIDATE_SAMPLE_SYNTHETIC",
    "candidate_docx": str(DOCX),
    "candidate_docx_sha256": actual_docx_sha,
    "candidate_text_sha256": sha(SOURCE_MD),
    "source_binding_sha256": sha(BINDING),
    "blueprint_copy_sha256": sha(BLUEPRINT_COPY),
    "register_sha256": sha(REGISTER),
    "native_pdf_sha256": sha(PDF),
    "native_pdf_pages": pdf_pages,
    "technical_status": "PASS_WITHIN_LISTED_DETERMINISTIC_SCOPE" if not failed else "FAIL",
    "visual_status": "INSPECTED_WITH_LAYOUT_NOTE" if not failed else "FAIL_OR_INCOMPLETE",
    "visual_scope": "All 28 pages inspected through seven contact sheets; full-resolution review of pages 5, 14, 16, 24, 27 and 28. Page 28 continues the final two synthetic-register rows with repeated header; the page is sparse but is not blank and has no clipped rows. No independent semantic review is implied.",
    "semantic_review": "PENDING_INDEPENDENT_EXACT_SHA_REVIEW",
    "formal_acceptance": "NOT_ACCEPTED",
    "real_evidence": "UNKNOWN_NOT_PROVIDED",
    "checks": checks,
    "failed_checks": failed,
    "limitations": [
        "No independent semantic PASS is claimed.",
        "Blueprint outward training-policy/public-disclosure fields applicability remains PARTIAL/UNKNOWN.",
        "No actual roster, authorization, learner, assessment, or association record was supplied; all cases are SAMPLE/SYNTHETIC.",
        "Page 28 is a short continuation page; all rows are readable, but the page retains substantial whitespace."
    ],
    "checked_at_taipei": datetime.now().astimezone().isoformat(timespec="seconds")
}
OUT.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
print(json.dumps({"out": str(OUT), "receipt_sha256": sha(OUT), "docx_sha256": actual_docx_sha,
                  "technical_status": receipt["technical_status"], "visual_status": receipt["visual_status"],
                  "check_count": len(checks), "failed_checks": failed}, ensure_ascii=False, indent=2))
