from pathlib import Path
import hashlib, json, csv, re, zipfile
from copy import deepcopy
from docx import Document
from docx.shared import Pt, RGBColor

ROOT = Path(r'E:\TTQS\_TTQS_ONE_REVIEW_SYNC_0036_R08_20261016\docs\ttqs-quality-20261016\quality-review')
R05 = ROOT / '0040_R05'
R06 = ROOT / '0040_R06'
SOURCE = R05 / 'candidate/0040_R05.docx'
OUTPUT = R06 / 'candidate/0040_R06.docx'
REGISTER_IN = R05 / 'evidence/SYNTHETIC_REPLACEMENT_REGISTER_R05.csv'
REGISTER_OUT = R06 / 'evidence/SYNTHETIC_REPLACEMENT_REGISTER_R06.csv'
EXPECTED_PARENT = '95531ac5f8a79eed7c605f3ab1f42da3948e8f15316499d325b606bd6d434384'
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
if sha(SOURCE) != EXPECTED_PARENT:
    raise SystemExit('R05 parent bytes differ from the reviewed exact SHA; refusing build')

def set_cell(cell, text, header=False):
    cell.text = text
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            run.font.name = 'Microsoft JhengHei'
            run.font.size = Pt(9.5)
            run.bold = bool(header)
            if header:
                run.font.color.rgb = RGBColor(255, 255, 255)

doc = Document(SOURCE)
if len(doc.tables) != 15:
    raise SystemExit(f'Unexpected table count {len(doc.tables)}')

# R05-F01: give every K row an actual evidence identifier/locator field while
# explicitly keeping the displayed pointer synthetic and unattached.
evidence_cells = [
    '索引：SYN-EV-0039-01／K1（本示範、未附外件）；正式填入受控檔案ID、附件/頁碼、版次；另記去識別碼、K1觀察事實、分數、日期及評量角色。',
    '索引：SYN-EV-0039-01／K2（本示範、未附外件）；正式填入受控檔案ID、附件/頁碼、版次；另記去識別碼、K2觀察事實、分數、日期及評量角色。',
    '索引：SYN-EV-0039-01／K3（本示範、未附外件）；正式填入受控檔案ID、附件/頁碼、版次；另記去識別碼、K3觀察事實、分數、日期及評量角色。',
    '索引：SYN-EV-0039-01／K4（本示範、未附外件）；正式填入受控檔案ID、附件/頁碼、版次；另記去識別碼、K4觀察事實、分數、日期及評量角色。'
]
matrix = doc.tables[4]
if len(matrix.rows) != 5 or len(matrix.columns) != 4:
    raise SystemExit('Unexpected p.3 K1-K4 matrix shape')
for row, text in zip(matrix.rows[1:], evidence_cells):
    set_cell(row.cells[3], text)

# R05-F01: demonstrate where the source evaluation record is located without
# pretending that an external original exists.
d04 = doc.tables[8]
if len(d04.rows) != 7 or len(d04.columns) != 3:
    raise SystemExit('Unexpected D04 sample table shape')
new_xml = deepcopy(d04.rows[-1]._tr)
d04._tbl.append(new_xml)
new = d04.rows[-1]
set_cell(new.cells[0], '原始評量證據定位')
set_cell(new.cells[1], 'SYN-EV-0039-01／D04後測 K1–K4欄（SAMPLE索引；未附外部原件）。')
set_cell(new.cells[2], '正式填用時填受控檔案ID、附件/頁碼及版次；缺件標「未提供／待補」，不可用本工具頁碼代替原件位置。')

# R05-F03: correct package metadata to the actual candidate revision.
doc.core_properties.description = '候選修訂 R06。SAMPLE/SYNTHETIC only. Not approved or effective.'
doc.core_properties.revision = 6
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
# python-docx in the bundled runtime does not expose core-properties description;
# patch only that exact OOXML element and assert a single replacement.
temporary = OUTPUT.with_suffix('.metadata.tmp')
with zipfile.ZipFile(OUTPUT, 'r') as zin, zipfile.ZipFile(temporary, 'w', zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == 'docProps/core.xml':
            xml = data.decode('utf-8')
            xml, count = re.subn(r'(<dc:description>).*?(</dc:description>)', r'\g<1>候選修訂 R06。SAMPLE/SYNTHETIC only. Not approved or effective.\g<2>', xml, count=1)
            if count != 1:
                raise SystemExit('Expected exactly one core-properties description element')
            data = xml.encode('utf-8')
        zout.writestr(item, data)
temporary.replace(OUTPUT)

# Build a new exact-SHA-bound register. Keep the optional 30-day rule and the
# D04 rescore status as separate facts with separate printed locators.
with REGISTER_IN.open(encoding='utf-8-sig', newline='') as f:
    reader = csv.DictReader(f)
    fields = reader.fieldnames
    rows = list(reader)
by_id = {row['fact_id']: row for row in rows}
if len(by_id) != len(rows) or '0040-SYN-009' not in by_id or '0040-SYN-037' not in by_id:
    raise SystemExit('R05 register facts do not match expected reviewed structure')
rule = by_id['0040-SYN-009']
rule['fact_type'] = 'optional_followup_authorization_rule'
rule['synthetic_value'] = '示範規則：課後30日追蹤須另行核准。'
rule['body_locator'] = '列印第1頁步驟表「4 定時點」列／操作欄「30日追蹤另行核准」。本列只定位授權規則，不表示本案例曾執行或未執行30日追蹤。'
rule['assumption_reason'] = '把30日追蹤設為選配且須先核准，避免將其當成預設課務；本示範沒有30日追蹤實施資料。'
rule['real_source_to_replace'] = '若正式採用，依經核准的課程評估方案或追蹤授權文件核實；未經核准不實施。'
rule['status'] = 'SAMPLE RULE; NOT ASSOCIATION RECORD'

# R05-F01 evidence pointer is itself a synthetic fact; no external evidence file is fabricated.
rows.append({
    'fact_id': '0040-SYN-039', 'doc_id': '0040', 'case_id': 'SAMPLE-0039-01',
    'fact_type': 'synthetic_evidence_locator_example',
    'synthetic_value': 'SYN-EV-0039-01 是本工具中的合成證據索引示例，不對應外部檔案，且未附原件；正式紀錄填受控檔案ID、附件/頁碼及版次。',
    'body_locator': '列印第3頁 K1-K4矩陣末欄各列；列印第5頁 D04示例表「原始評量證據定位」列。',
    'assumption_reason': '提供可照填的識別碼與定位樣式，同時明示未附外部原件，避免把示範索引誤當真實佐證。',
    'real_source_to_replace': '有權使用的實際評量表、受控檔案識別碼、附件/頁碼、工具版次及存取權限。',
    'real_owner_role': '評量角色及受控資料保管角色；實際人員與權責待協會核實。',
    'status': 'SAMPLE/SYNTHETIC; NO EXTERNAL ORIGINAL',
    'bound_candidate_sha256': ''
})
candidate_sha = sha(OUTPUT)
for row in rows:
    row['bound_candidate_sha256'] = candidate_sha
if len(rows) != 40 or len({r['fact_id'] for r in rows}) != 40:
    raise SystemExit('R06 register must contain 40 unique facts')
with REGISTER_OUT.open('w', encoding='utf-8-sig', newline='') as f:
    writer = csv.DictWriter(f, fieldnames=fields, lineterminator='\n')
    writer.writeheader()
    writer.writerows(rows)
print(json.dumps({
    'parent_r05_sha256': sha(SOURCE), 'candidate_r06_sha256': candidate_sha,
    'register_rows': len(rows), 'register_sha256': sha(REGISTER_OUT),
    'repair_findings': ['R05-F01 evidence location fields', 'R05-F02 split locator facts', 'R05-F03 core properties revision']
}, ensure_ascii=False, indent=2))
