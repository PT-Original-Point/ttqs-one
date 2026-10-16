from pathlib import Path
from zipfile import ZipFile
import csv, hashlib, json, re, xml.etree.ElementTree as ET
from docx import Document
from pypdf import PdfReader

ROOT=Path(r'E:\TTQS\_TTQS_ONE_REVIEW_SYNC_0036_R08_20261016\docs\ttqs-quality-20261016\quality-review')
P=ROOT/'0040_R06'; R05=ROOT/'0040_R05'
docx=P/'candidate/0040_R06.docx'; pdf=P/'evidence/render/candidate.pdf'
register=P/'evidence/SYNTHETIC_REPLACEMENT_REGISTER_R06.csv'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
norm=lambda s:re.sub(r'[\W_]+','',s,flags=re.UNICODE).lower()
with ZipFile(docx) as z:
    corrupt=z.testzip(); xmlnames=[n for n in z.namelist() if n.endswith('.xml')]; xmlerrors=[]
    for n in xmlnames:
        try: ET.fromstring(z.read(n))
        except Exception as e: xmlerrors.append({'part':n,'error':str(e)})
d=Document(docx); reader=PdfReader(str(pdf)); pages=[' '.join((p.extract_text() or '').split()) for p in reader.pages]
with register.open(encoding='utf-8-sig',newline='') as f: rows=list(csv.DictReader(f))
by_id={r['fact_id']:r for r in rows}; candidate_sha=sha(docx)
binding=all(r['bound_candidate_sha256'].lower()==candidate_sha for r in rows)

# Review-specific locator checks bind evidence instructions to actual R06 printed content.
matrix=d.tables[4]
matrix_values=[r.cells[3].text for r in matrix.rows[1:]]
matrix_checks={f'K{i}':('SYN-EV-0039-01' in matrix_values[i-1] and f'K{i}' in matrix_values[i-1] and '受控檔案ID' in matrix_values[i-1] and '附件/頁碼' in matrix_values[i-1] and '未附外件' in matrix_values[i-1]) for i in range(1,5)}
d04=d.tables[8]
evidence_row=[r for r in d04.rows if r.cells[0].text.strip()=='原始評量證據定位']
d04_evidence_check=(len(evidence_row)==1 and all(token in ' '.join(c.text for c in evidence_row[0].cells) for token in ['SYN-EV-0039-01','受控檔案ID','附件/頁碼','未提供／待補']))
rule=by_id['0040-SYN-009']; rescore=by_id['0040-SYN-037']; new_fact=by_id.get('0040-SYN-039')
register_separation={
  'optional_rule_locator_only':'第1頁' in rule['body_locator'] and '第6頁' not in rule['body_locator'],
  'optional_rule_scope_clear':'不表示本案例曾執行或未執行30日追蹤' in rule['body_locator'],
  'd04_status_locator_separate':'第7頁' in rescore['body_locator'] and '本例尚未執行' in rescore['synthetic_value'],
  'evidence_pointer_registered':bool(new_fact and 'SYN-EV-0039-01' in new_fact['synthetic_value'] and '不對應外部檔案' in new_fact['synthetic_value'] and '未附原件' in new_fact['synthetic_value'])
}
expected_page_refs={
 '0040-REF-001':(1,),
 **{f'0040-SYN-{i:03d}':(1,) for i in (1,2,3,7,8,9)},
 '0040-SYN-004':(3,),'0040-SYN-005':(3,4),'0040-SYN-006':(4,5),'0040-SYN-010':(3,4,6),
 '0040-SYN-BL-K1':(3,),'0040-SYN-BL-K2':(3,),'0040-SYN-BL-K3':(3,),'0040-SYN-BL-K4':(3,),
 **{f'0040-SYN-PRE-D{i:02d}':(5,) for i in range(1,6)},
 **{f'0040-SYN-POST-D{i:02d}':(5,) for i in range(1,6)},
 **{f'0040-SYN-CALC-{i:03d}':(5,) for i in range(2,9)},
 '0040-SYN-CALC-001':(2,),
 '0040-SYN-CALC-009':(6,7),
 **{f'0040-SYN-{i:03d}':(7,) for i in range(34,39)},
 '0040-SYN-039':(3,6)
}
def pages_covered(locator):
    covered=set(map(int,re.findall(r'第(\d+)頁',locator)))
    for first,last in re.findall(r'第(\d+)[–-](\d+)頁',locator):
        covered.update(range(int(first),int(last)+1))
    return covered
locator_page_checks={fid:set(refs).issubset(pages_covered(by_id[fid]['body_locator'])) for fid,refs in expected_page_refs.items()}
page_anchor_checks={
 'p3_matrix_all_baselines':all(f'K{i}' in pages[2] for i in range(1,5)) and all(v in pages[2] for v in ('1.0','0.6','1.2','0.8','1.4')) and 'SYN-EV-0039-01' in pages[2],
 'p4_matrix_and_scoring':'K3' in pages[3] and 'K4' in pages[3] and '計分' in pages[3],
 'p5_five_examples':'D01' in pages[4] and 'D05' in pages[4],
 'p6_d04_form':'D04' in pages[5] and '原始評量證據定位' in pages[5],
 'p7_d04_followup':'D04' in pages[6] and '尚未執行' in pages[6]
}
with ZipFile(docx) as z:
    core=ET.fromstring(z.read('docProps/core.xml'))
    ns={'cp':'http://schemas.openxmlformats.org/package/2006/metadata/core-properties','dc':'http://purl.org/dc/elements/1.1/'}
    description=core.findtext('dc:description',default='',namespaces=ns)
    revision=core.findtext('cp:revision',default='',namespaces=ns)
metadata={'description_marks_r06':'R06' in description,'revision_is_6':revision=='6','description':description,'revision':revision}

# Recompute five displayed paired examples, preserving R05's assessed math.
score_rows=[]
for row in d.tables[6].rows[1:]:
    code=row.cells[0].text.strip()
    parse=lambda s: tuple(map(int,re.search(r'([0-2])/([0-2])/([0-2])/([0-2])\s*[；;]\s*(\d+)/8',s).groups()))
    pre=parse(row.cells[1].text); post=parse(row.cells[2].text)
    if sum(pre[:4])!=pre[4] or sum(post[:4])!=post[4]: raise ValueError('Displayed total mismatch '+code)
    delta=int(row.cells[3].text.replace('+','').strip())
    if post[4]-pre[4]!=delta: raise ValueError('Displayed change mismatch '+code)
    passed=post[4]>=6 and post[2]>=1
    if row.cells[4].text.strip().startswith('通過：')!=passed: raise ValueError('Displayed threshold verdict mismatch '+code)
    score_rows.append({'case':code,'pre_total':pre[4],'post_total':post[4],'delta':delta,'pass':passed})
pre_sum=sum(x['pre_total'] for x in score_rows); post_sum=sum(x['post_total'] for x in score_rows)
change_sum=sum(x['delta'] for x in score_rows); pass_count=sum(x['pass'] for x in score_rows)

font_sizes=[]
with ZipFile(docx) as z:
    root=ET.fromstring(z.read('word/document.xml')); ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    body=root.find('w:body',ns)
    for child in list(body):
        is_table=child.tag.endswith('}tbl')
        for run in child.findall('.//w:r',ns):
            size=run.find('./w:rPr/w:sz',ns)
            if size is not None:
                try: font_sizes.append((int(size.attrib.get('{'+ns['w']+'}val'))/2,is_table))
                except Exception: pass
min_body=min((v for v,t in font_sizes if not t),default=None)
min_table=min((v for v,t in font_sizes if t),default=None)
pngs=sorted((P/'evidence/render').glob('page-*.png'))
page_png_sha={str(i):sha(P/f'evidence/render/page-{i}.png') for i in range(1,len(pngs)+1)}
report={
 'project_id':'TTQS_ONE','doc_id':'0040','revision':'R06','requirement_id':'I08-E03','official_checkbox_sequence':'040',
 'candidate_sha256':candidate_sha,'parent_r05_sha256':sha(R05/'candidate/0040_R05.docx'),
 'parent_r05_verdict':'PARTIAL','register_sha256':sha(register),'register_rows':len(rows),
 'register_unique_fact_ids':len({r['fact_id'] for r in rows}),'register_all_bound_to_candidate_sha':binding,
 'locator_repairs':{'matrix_k1_k4':matrix_checks,'d04_evidence_row':d04_evidence_check,'separate_register_facts':register_separation,'register_page_refs':locator_page_checks,'rendered_page_anchors':page_anchor_checks},
 'core_properties':metadata,'docx_zip_integrity':'PASS' if corrupt is None else 'FAIL',
 'xml_part_count':len(xmlnames),'xml_parse_errors':xmlerrors,'table_count':len(d.tables),
 'pdf_sha256':sha(pdf),'page_count_pdf':len(reader.pages),'page_count_png':len(pngs),'page_png_sha256':page_png_sha,
 'score_recalculation':{'rows':score_rows,'pre_total':pre_sum,'post_total':post_sum,'paired_change_total':change_sum,'pass_count':pass_count,'denominator':len(score_rows),'pre_mean':pre_sum/len(score_rows),'post_mean':post_sum/len(score_rows),'paired_mean_change':change_sum/len(score_rows),'pass_rate_percent':pass_count/len(score_rows)*100},
 'font_minima_pt':{'body':min_body,'tables':min_table},
 'status':'PASS_TECHNICAL' if corrupt is None and not xmlerrors and len(rows)==40 and len({r['fact_id'] for r in rows})==40 and binding and all(matrix_checks.values()) and d04_evidence_check and all(register_separation.values()) and all(locator_page_checks.values()) and all(page_anchor_checks.values()) and all(metadata[k] for k in ['description_marks_r06','revision_is_6']) and len(reader.pages)==9 and len(pngs)==9 and min_body>=10.5 and min_table>=9.5 and pre_sum==17 and post_sum==31 and change_sum==14 and pass_count==4 else 'FAIL'
}
out=P/'evidence/qa/0040_technical_review_r06_exact.json'
out.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'status':report['status'],'candidate_sha256':candidate_sha,'register_sha256':sha(register),'pdf_sha256':sha(pdf),'pages':len(reader.pages),'pngs':len(pngs),'matrix_checks':matrix_checks,'d04_evidence_row':d04_evidence_check,'register_split':register_separation,'locator_pages':locator_page_checks,'rendered_page_anchors':page_anchor_checks,'metadata':metadata,'font_minima':report['font_minima_pt'],'math':report['score_recalculation']},ensure_ascii=True,indent=2))
