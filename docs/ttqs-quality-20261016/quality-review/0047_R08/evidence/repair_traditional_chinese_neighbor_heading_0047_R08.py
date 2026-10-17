from pathlib import Path
from shutil import copy2
import hashlib, json
from datetime import datetime, timezone
from docx import Document
root=Path('docs/ttqs-quality-20261016/quality-review')
src=root/'0047_R07/candidate/0047_R07.docx'
dst=root/'0047_R08/candidate/0047_R08.docx'
script_path=root/'0047_R08/evidence/repair_traditional_chinese_neighbor_heading_0047_R08.py'
old_bytes=src.read_bytes()
old_sha=hashlib.sha256(old_bytes).hexdigest()
copy2(src,dst)
doc=Document(str(dst))
old_heading='相邻要求邊界與負向近鄰拒收規則'
new_heading='相鄰要求邊界與負向近鄰拒收規則'
found=[]
for p in doc.paragraphs:
    if p.text == old_heading:
        found.append(p)
if len(found)!=1:
    raise SystemExit(f'expected exactly one target heading, found {len(found)}')
p=found[0]
if len(p.runs)!=1:
    raise SystemExit(f'expected one run to preserve heading formatting safely, got {len(p.runs)}')
p.runs[0].text=new_heading
doc.save(str(dst))
new_doc=Document(str(dst))
other_paras=[x.text for x in Document(str(src)).paragraphs if x.text!=old_heading]
new_other=[x.text for x in new_doc.paragraphs if x.text!=new_heading]
def tables_text(d):
    return [[[''.join(c.text) for c in row.cells] for row in table.rows] for table in d.tables]
if other_paras!=new_other:
    raise SystemExit('non-target paragraph content changed')
if tables_text(Document(str(src)))!=tables_text(new_doc):
    raise SystemExit('table text changed')
new_sha=hashlib.sha256(dst.read_bytes()).hexdigest()
receipt={
 'doc_id':'0047','revision':'R08','status':'CANDIDATE_ONLY',
 'source_candidate':'0047_R07/candidate/0047_R07.docx','source_sha256':old_sha,
 'candidate':'0047_R08/candidate/0047_R08.docx','candidate_sha256':new_sha,
 'change_scope':'TARGETED_TRADITIONAL_CHINESE_HEADING_CORRECTION',
 'old_heading':old_heading,'new_heading':new_heading,
 'paragraphs_other_than_target_unchanged':True,'all_table_text_unchanged':True,
 'sample_numbers_and_calculations_changed':False,'render_visual_review':'PENDING',
 'independent_semantic_review':'PENDING_EXACT_SHA','formal_acceptance':'NOT_CLAIMED',
 'created_utc':datetime.now(timezone.utc).isoformat()
}
out=root/'0047_R08/evidence/R08_TITLE_REPAIR_RECEIPT.json'
out.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
script_path.write_text(Path(__file__).read_text(encoding='utf-8'),encoding='utf-8')
print(json.dumps(receipt,ensure_ascii=False,indent=2))