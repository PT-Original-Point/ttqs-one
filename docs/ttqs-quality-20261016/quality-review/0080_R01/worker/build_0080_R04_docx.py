from __future__ import annotations

import hashlib
import json
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
BODY = ROOT / "candidate" / "0080_R04_body.md"
BLUEPRINT = ROOT / "inputs" / "UNIQUE_BLUEPRINT_0080_R03.md"
REGISTER = ROOT / "inputs" / "SYNTHETIC_REPLACEMENT_REGISTER_0080_R01.csv"
SOURCE_PACKET = ROOT / "inputs" / "SOURCE_PACKET_0080_R03.md"
OUTPUT = ROOT / "candidate" / "0080_R04.docx"
RECEIPT = ROOT / "evidence" / "build_receipt_R04.json"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=75, start=85, bottom=75, end=85) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, size=10.2, bold=None, color=None, name="Microsoft JhengHei") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*)")


def add_inline(paragraph, text: str, size=10.5, color=None) -> None:
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size=size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, color="135B73", name="Consolas")
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, color=color)
            run.italic = True
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=size, color=color)


def parse_table_row(line: str) -> list[str]:
    text = line.strip()
    if text.startswith("|"):
        text = text[1:]
    if text.endswith("|"):
        text = text[:-1]
    cells = re.split(r"(?<!\\)\|", text)
    return [cell.replace(r"\|", "|").strip() for cell in cells]


def is_table_separator(cells: list[str]) -> bool:
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells)


def set_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])



def add_card_table(doc: Document, title: str, field_pairs: list[tuple[str, str]]) -> None:
    outer = doc.add_table(rows=1, cols=1)
    outer.autofit = False
    outer_row = outer.rows[0]
    prevent_row_split(outer_row)
    outer_cell = outer_row.cells[0]
    set_cell_margins(outer_cell, top=0, start=0, bottom=0, end=0)
    tc_pr = outer_cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    tc_pr.append(borders)
    initial_paragraph = outer_cell.paragraphs[0]._element
    initial_paragraph.getparent().remove(initial_paragraph)
    table = outer_cell.add_table(rows=1, cols=2)
    table.autofit = False
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    head = table.rows[0]
    prevent_row_split(head)
    set_repeat_table_header(head)
    merged = head.cells[0].merge(head.cells[1])
    merged.width = Inches(6.86)
    set_cell_shading(merged, "17365D")
    set_cell_margins(merged, top=80, start=105, bottom=80, end=105)
    merged.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    hp = merged.paragraphs[0]
    hp.paragraph_format.space_after = Pt(0)
    hp.paragraph_format.line_spacing = 1.0
    add_inline(hp, title, size=9.8, color="FFFFFF")
    for run in hp.runs:
        run.bold = True
    for idx, (label, value) in enumerate(field_pairs):
        row = table.add_row()
        prevent_row_split(row)
        row.cells[0].width = Inches(1.42)
        row.cells[1].width = Inches(5.44)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell, top=68, start=82, bottom=68, end=82)
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            cell.paragraphs[0].paragraph_format.space_before = Pt(0)
            cell.paragraphs[0].paragraph_format.line_spacing = 1.0
        set_cell_shading(row.cells[0], "E8F0F4")
        add_inline(row.cells[0].paragraphs[0], label, size=9.3, color="17365D")
        for run in row.cells[0].paragraphs[0].runs:
            run.bold = True
        if idx % 2 == 1:
            set_cell_shading(row.cells[1], "F7F9FA")
        add_inline(row.cells[1].paragraphs[0], value, size=9.4, color="222222")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

def build_wide_markdown_table_as_cards(doc: Document, rows: list[list[str]]) -> tuple[int, int] | None:
    if not rows:
        return None
    headers = rows[0]
    if headers == ["編號", "建議（提案）", "負責角色（提案）", "下一步（提案）", "樣本目標日", "查核證據（未來）", "狀態"]:
        labels = ["建議內容（提案）", "負責角色（提案）", "下一步（提案）", "樣本目標日", "查核證據（未來）", "狀態"]
        for row in rows[1:]:
            if len(row) != 7:
                raise ValueError(f"Action card has {len(row)} cells; expected 7")
            add_card_table(doc, f"改善建議 {row[0]}", list(zip(labels, row[1:])))
        return len(rows), len(headers)
    if headers == ["文件", "要求編號", "勾選序號", "指標", "不同功能", "與本文件之關係"]:
        for row in rows[1:]:
            if len(row) != 6:
                raise ValueError(f"Cross-document card has {len(row)} cells; expected 6")
            title = f"關聯文件｜{row[0]}｜要求 {row[1]}｜勾選 {row[2]}"
            pairs = [("指標", row[3]), ("文件功能", row[4]), ("與本文件的關係", row[5])]
            add_card_table(doc, title, pairs)
        return len(rows), len(headers)
    if len(headers) >= 6:
        for row in rows[1:]:
            if len(row) != len(headers):
                raise ValueError(f"Wide record has {len(row)} cells; expected {len(headers)}")
            record_id = row[0].strip()
            title = f"{headers[0]}｜{record_id}" if record_id else headers[0]
            pairs = list(zip(headers[1:], row[1:]))
            add_card_table(doc, title, pairs)
        return len(rows), len(headers)
    return None

def build_table(doc: Document, rows: list[list[str]]) -> tuple[int, int]:
    card_result = build_wide_markdown_table_as_cards(doc, rows)
    if card_result is not None:
        return card_result
    if len(rows) < 2:
        raise ValueError("Markdown table has no body rows")
    cols = len(rows[0])
    bad = [(i + 1, len(row)) for i, row in enumerate(rows) if len(row) != cols]
    if bad:
        raise ValueError(f"Inconsistent table column counts: expected {cols}, got {bad[:8]}")
    table = doc.add_table(rows=1, cols=cols)
    locator_table = rows[0] == ["類別", "定位", "SHA-256／狀態"]
    table.autofit = not locator_table
    table.allow_autofit = not locator_table
    if locator_table:
        locator_widths = [Inches(1.30), Inches(2.35), Inches(3.01)]
        for column, width in zip(table.columns, locator_widths):
            column.width = width
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    font_size = 8.7 if locator_table else 9.5
    for r_i, values in enumerate(rows):
        row = table.rows[0] if r_i == 0 else table.add_row()
        prevent_row_split(row)
        for c_i, value in enumerate(values):
            cell = row.cells[c_i]
            if locator_table:
                cell.width = locator_widths[c_i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=48 if locator_table else 75, start=60 if locator_table else 85, bottom=48 if locator_table else 75, end=60 if locator_table else 85)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            if r_i == 0:
                set_cell_shading(cell, "17365D")
                add_inline(paragraph, value, size=font_size, color="FFFFFF")
                for run in paragraph.runs:
                    run.bold = True
            else:
                if r_i % 2 == 0:
                    set_cell_shading(cell, "F0F5F8")
                add_inline(paragraph, value, size=font_size, color="222222")
    set_repeat_table_header(table.rows[0])
    if rows[0] == ["#", "核對項", "對應紀錄", "狀態"]:
        for cell in table.rows[0].cells:
            cell.paragraphs[0].paragraph_format.keep_with_next = True
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return len(rows), cols


def add_quote(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.13)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_together = True
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "EAF3F6")
    p_pr.append(shd)
    border = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "20")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), "1B7285")
    border.append(left)
    p_pr.append(border)
    add_inline(p, text.strip())


def main() -> None:
    body = BODY.read_text(encoding="utf-8-sig")
    lines = body.splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(0.68)
    sec.bottom_margin = Inches(0.66)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)
    sec.header_distance = Inches(0.30)
    sec.footer_distance = Inches(0.30)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft JhengHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("242B31")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.18

    for style_name, size, color, before, after in [
        ("Title", 22, "000000", 0, 13),
        ("Heading 1", 15, "000000", 16, 7),
        ("Heading 2", 12, "000000", 11, 5),
        ("Heading 3", 10.8, "000000", 8, 3),
        ("Heading 4", 10.2, "000000", 6, 2),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Microsoft JhengHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("TTQS ONE  |  DOCX 0080  |  I17B-E02  |  官方勾選 085")
    set_run_font(run, size=8, color="657986")
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    f = footer.add_run("SAMPLE/SYNTHETIC  |  CANDIDATE ONLY  |  REAL=UNKNOWN/HOLD    •    第 ")
    set_run_font(f, size=8, color="657986")
    set_page_field(footer)
    f2 = footer.add_run(" 頁")
    set_run_font(f2, size=8, color="657986")

    tables = 0
    table_rows = 0
    headings = 0
    paragraphs = 0
    i = 0
    first_content = True
    while i < len(lines):
        raw = lines[i]
        text = raw.strip()
        if not text or text == "---":
            i += 1
            continue
        if text.startswith("|"):
            raw_rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw_rows.append(parse_table_row(lines[i]))
                i += 1
            if len(raw_rows) > 1 and is_table_separator(raw_rows[1]):
                raw_rows.pop(1)
            n, _ = build_table(doc, raw_rows)
            tables += 1
            table_rows += n - 1
            first_content = False
            continue
        if text.startswith("> "):
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote_lines.append(re.sub(r"^>\s?", "", lines[i].strip()))
                i += 1
            add_quote(doc, " ".join(quote_lines))
            paragraphs += 1
            first_content = False
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)$", text)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2).strip()
            style_name = "Title" if first_content and level == 1 else f"Heading {min(level,4)}"
            p = doc.add_paragraph(style=style_name)
            add_inline(p, title, size={"Title": 22, "Heading 1": 15, "Heading 2": 12, "Heading 3": 10.8, "Heading 4": 10.2}[style_name])
            if title == "附錄 A 來源定位 Locator":
                p.paragraph_format.page_break_before = True
            headings += 1
            paragraphs += 1
            first_content = False
            i += 1
            continue
        if re.fullmatch(r"[-*_]{3,}", text):
            i += 1
            continue
        bullet = re.match(r"^\s*[-*+]\s+(.*)$", raw)
        number = re.match(r"^\s*(\d+[.)])\s+(.*)$", raw)
        if bullet or number:
            if bullet:
                content = bullet.group(1)
                p = doc.add_paragraph(style="List Bullet")
            else:
                # Preserve Markdown's explicit sequence numbers. Word's List Number
                # style continues numbering across unrelated sections and can make
                # independent lists appear to start at 8 or 9 in the final DOCX.
                content = f"{number.group(1)} {number.group(2)}"
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.28)
                p.paragraph_format.first_line_indent = Inches(-0.28)
            p.paragraph_format.space_after = Pt(2.5)
            p.paragraph_format.line_spacing = 1.12
            add_inline(p, content)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.widow_control = True
            # Bold lead-ins immediately before a table label the following form
            # section and must not be stranded at the end of the prior page.
            next_block = i + 1
            while next_block < len(lines) and not lines[next_block].strip():
                next_block += 1
            if text.startswith("**") and text.endswith("**") and next_block < len(lines) and lines[next_block].strip().startswith("|"):
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.keep_together = True
            add_inline(p, text)
        paragraphs += 1
        first_content = False
        i += 1

    # The builder emits a spacer paragraph after every table. Keep the final
    # spacer to one point so Word does not push it onto a blank extra page.
    if doc.paragraphs and not doc.paragraphs[-1].text.strip():
        trailing = doc.paragraphs[-1]
        trailing.paragraph_format.space_before = Pt(0)
        trailing.paragraph_format.space_after = Pt(0)
        trailing.paragraph_format.line_spacing = Pt(1)
        trailing.paragraph_format.keep_together = False
        for run in trailing.runs:
            run.font.size = Pt(1)

    doc.core_properties.title = "課程評量與學員建議檢討紀錄 0080 I17B-E02"
    doc.core_properties.subject = "SAMPLE/SYNTHETIC 候選示範；REAL=UNKNOWN/HOLD"
    doc.core_properties.author = "TTQS ONE candidate document production"
    doc.core_properties.keywords = "0080, I17B-E02, checkbox 085, SAMPLE, SYNTHETIC"
    doc.save(OUTPUT)

    receipt = {
        "schema": "TTQS_ONE_0080_R04_BUILD_RECEIPT.v1",
        "document_id": "0080",
        "requirement_id": "I17B-E02",
        "official_checkbox_serial": "085",
        "output": str(OUTPUT),
        "docx_sha256": sha256(OUTPUT),
        "inputs": {
            str(BODY.relative_to(ROOT)): sha256(BODY),
            str(BLUEPRINT.relative_to(ROOT)): sha256(BLUEPRINT),
            str(REGISTER.relative_to(ROOT)): sha256(REGISTER),
            str(SOURCE_PACKET.relative_to(ROOT)): sha256(SOURCE_PACKET),
        },
        "layout_transform": {
            "source_body_unchanged": True,
            "markdown_tables_rendered_as_two_column_cards": [
                "Any six-or-more-column Markdown table is transformed into one readable field/value card per record; source fields are preserved.",
                "Smaller tables remain native tables with repeating header rows."
            ],
            "reason": "Preserve every source field while avoiding unreadably compressed wide tables",
            "title_rule_removed": True,
            "wide_record_cards_kept_together": True,
            "blank_record_id_card_title_has_no_trailing_separator": True,
            "bold_lead_in_before_table_kept_with_section": True,
            "final_table_spacer_is_minimal": True,
            "appendix_locator_columns_balanced_to_avoid_orphaned_last_row": True,
            "appendix_locator_reduced_cell_padding_and_font": True,
            "appendix_locator_heading_and_complete_table_start_on_new_page": True
        },
        "formatting": {
            "page": "A4 portrait; 0.72in side margins; 0.68/0.66in top/bottom",
            "font": "Microsoft JhengHei",
            "heading_palette": ["000000"],
            "table_header": "17365D with white text; alternating light rows; repeating header row",
            "footer": "SAMPLE/SYNTHETIC; CANDIDATE ONLY; REAL=UNKNOWN/HOLD; page field",
        },
        "parsed": {"headings": headings, "paragraphs": paragraphs, "tables": tables, "table_body_rows": table_rows},
        "built_at_local": datetime.now().astimezone().isoformat(),
        "status": "BUILT_UNVERIFIED",
        "formal_acceptance": "NOT_GRANTED",
        "real_evidence": "UNKNOWN/HOLD",
    }
    RECEIPT.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"docx": str(OUTPUT), "docx_sha256": receipt["docx_sha256"], "parsed": receipt["parsed"], "status": receipt["status"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()


