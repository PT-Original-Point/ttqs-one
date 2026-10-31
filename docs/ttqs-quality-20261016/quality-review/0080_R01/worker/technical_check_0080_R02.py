from pathlib import Path
import csv,hashlib,json,zipfile,re,statistics,collections
from docx import Document
ROOT=Path(r'docs/ttqs-quality-20261016/quality-review/0080_R01')
DOCX=ROOT/'candidate/0080_R02.docx'; BODY=ROOT/'candidate/0080_R02_body.md'; BP=ROOT/'inputs/UNIQUE_BLUEPRINT_0080_R02.md'; PACKET=ROOT/'inputs/SOURCE_PACKET_0080_R02.md'; REG=ROOT/'inputs/SYNTHETIC_REPLACEMENT_REGISTER_0080_R01.csv'; SUMMARY=ROOT/'inputs/SYNTHETIC_REVIEW_DATA_0080_R01.csv'; RECEIPT=ROOT/'evidence/build_receipt_R02.json'; OUT=ROOT/'evidence/technical_review_R02.json'; EXTRACT=ROOT/'evidence/docx_extracted_R02.md'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest(); norm=lambda s:re.sub(r'\s+',' ',s).strip()
body=BODY.read_text(encoding='utf-8-sig'); bp=BP.read_text(encoding='utf-8-sig'); packet=PACKET.read_text(encoding='utf-8-sig')
with zipfile.ZipFile(DOCX) as z: bad=z.testzip(); missing=sorted({'[Content_Types].xml','word/document.xml','word/styles.xml'}-set(z.namelist()))
d=Document(DOCX); chunks=[]
for p in d.paragraphs:
 if norm(p.text): chunks.append(norm(p.text))
for section in d.sections:
 for p in list(section.header.paragraphs)+list(section.footer.paragraphs):
  if norm(p.text): chunks.append(norm(p.text))
tables=[]
for t in d.tables:
 rows=[[norm(c.text) for c in row.cells] for row in t.rows]; tables.append(rows); chunks.extend([' | '.join(x) for x in rows])
text='\n'.join(chunks); EXTRACT.write_text(text+'\n',encoding='utf-8',newline='\n')
with (ROOT.parent/'0076_R01/inputs/SYNTHETIC_ASSESSMENT_DATA_0076_R01.csv').open(encoding='utf-8-sig',newline='') as f: assessment=list(csv.DictReader(f))
with (ROOT.parent/'0076_R01/inputs/SYNTHETIC_ITEM_ANALYSIS_0076_R01.csv').open(encoding='utf-8-sig',newline='') as f: items=list(csv.DictReader(f))
with (ROOT.parent/'0077_R01/inputs/REFLECTION_SYNTHETIC_DATA_0077_R01.csv').open(encoding='utf-8-sig',newline='') as f: reflections=list(csv.DictReader(f))
with REG.open(encoding='utf-8-sig',newline='') as f: registers=list(csv.DictReader(f))
with SUMMARY.open(encoding='utf-8-sig',newline='') as f: summary=list(csv.DictReader(f))
pre=sum(int(x['pretest_total_from_0074_synthetic']) for x in assessment); post=sum(int(x['posttest_total_from_0074_synthetic']) for x in assessment); practical=sum(int(x['practical_total_0_20']) for x in assessment)
keys={x['item_id']:x['answer_key_letter'] for x in items}; correct={q:sum(1 for x in assessment if x[f'{q}_response']==key) for q,key in keys.items()}
accurate=all(int(x['practical_total_0_20'])==sum(int(x[k]) for k in ['communication_0_5','privacy_boundary_0_5','record_accuracy_0_5','handoff_followup_0_5']) for x in assessment)
themes=collections.Counter(x['primary_theme_code'] for x in reflections)
wrong='3bdce0a1675a5b377dfafe2d7519f9f877e51530f5edc4ba3551a116bc6e985'; actual='3bdce0a1675a5b377dfafe2d75139d1f877e51530f5edc4ba3551a116bc6e985'
build=json.loads(RECEIPT.read_text(encoding='utf-8-sig'))
checks={
 'zip_integrity':bad is None,'required_ooxml_parts':not missing,'docx_sha_matches_build_receipt':sha(DOCX)==build['docx_sha256'],
 'correct_identity':'0080' in text and 'I17B-E02' in text and '085' in text,
 'corrected_current_source_hash_in_body':actual in body and wrong not in body and actual in bp and actual in packet and wrong not in bp and wrong not in packet,
 'appendix_points_to_exact_r02_input_hashes':sha(BP) in body and sha(PACKET) in body,
 'real_candidate_boundary':all(x in text for x in ['SAMPLE/SYNTHETIC','REAL=UNKNOWN/HOLD','NOT_GRANTED','PROPOSED','NOT_EXECUTED']),
 'meeting_is_explicitly_synthetic_and_unheld':all(x in text for x in ['2026-08-29','未實際召開','無任何真實與會者']),
 'adjacent_requirement_boundaries':all(x in text for x in ['0081','0082','I17B-E03','I17C-E01']),
 'source_assessment_rows_16':len(assessment)==16,
 'source_reflection_rows_16':len(reflections)==16,
 'source_replacement_register_23':len(registers)==23,
 'source_review_summary_26':len(summary)==26,
 'pretest_math':pre==870 and pre/16==54.375 and '54.375' in body,
 'posttest_math':post==1265 and post/16==79.0625 and '79.0625' in body,
 'gain_math':post/16-pre/16==24.6875 and '+24.6875' in body,
 'practical_math':practical==253 and practical/16==15.8125 and '15.8125' in body and accurate,
 'item_analysis_recomputed':all(int(x['correct_n'])==correct[x['item_id']] for x in items) and all(f'{q}／{correct[q]}／16' in body or f'{q}／{correct[q]}／16' in text for q in ['Q01','Q17','Q18','Q19']),
 'body_length_target':7000<=len(re.findall(r'[\u4e00-\u9fff]',body))<=9500,
 'status_candidate_only':build['status']=='BUILT_UNVERIFIED' and build['formal_acceptance']=='NOT_GRANTED' and build['real_evidence']=='UNKNOWN/HOLD'
}
receipt={'schema':'TTQS_ONE_0080_TECHNICAL_REVIEW_R02.v1','document_id':'0080','requirement_id':'I17B-E02','official_checkbox_serial':'085','candidate_docx_path':'candidate/0080_R02.docx','candidate_docx_sha256':sha(DOCX),'candidate_docx_bytes':DOCX.stat().st_size,'body_sha256':sha(BODY),'unique_blueprint_sha256':sha(BP),'source_packet_sha256':sha(PACKET),'source_csv_sha256':{'assessment':sha(ROOT.parent/'0076_R01/inputs/SYNTHETIC_ASSESSMENT_DATA_0076_R01.csv'),'item_analysis':sha(ROOT.parent/'0076_R01/inputs/SYNTHETIC_ITEM_ANALYSIS_0076_R01.csv'),'reflections':sha(ROOT.parent/'0077_R01/inputs/REFLECTION_SYNTHETIC_DATA_0077_R01.csv')},'recomputed_metrics':{'rows':len(assessment),'pretest_total':pre,'pretest_mean':pre/16,'posttest_total':post,'posttest_mean':post/16,'mean_gain':post/16-pre/16,'practical_total':practical,'practical_mean':practical/16,'item_correct_counts':correct,'reflection_theme_counts':dict(themes),'practical_subscores_equal_total':accurate},'cjk_characters':len(re.findall(r'[\u4e00-\u9fff]',body)),'tables':len(tables),'checks':checks,'status':'PASS' if all(checks.values()) and bad is None and not missing else 'FAIL','formal_acceptance':'NOT_GRANTED','real_evidence':'UNKNOWN/HOLD','independent_semantic_review':'PENDING'}
OUT.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'docx_sha256':receipt['candidate_docx_sha256'],'checks':checks,'status':receipt['status'],'metrics':receipt['recomputed_metrics'],'cjk':receipt['cjk_characters'],'tables':len(tables)},ensure_ascii=False))
