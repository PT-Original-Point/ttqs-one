from pathlib import Path
import csv,hashlib,json,zipfile,re,collections,xml.etree.ElementTree as ET
from docx import Document
R=Path(r'docs/ttqs-quality-20261016/quality-review/0080_R01')
DOCX=R/'candidate/0080_R05.docx'; BODY=R/'candidate/0080_R05_body.md'; BP=R/'inputs/UNIQUE_BLUEPRINT_0080_R03.md'; PACK=R/'inputs/SOURCE_PACKET_0080_R03.md'; REG=R/'inputs/SYNTHETIC_REPLACEMENT_REGISTER_0080_R01.csv'; SUMMARY=R/'inputs/SYNTHETIC_REVIEW_DATA_0080_R01.csv'; BUILD=R/'evidence/build_receipt_R05.json'; OUT=R/'evidence/technical_review_R05.json'; EXTRACT=R/'evidence/docx_extracted_R05.md'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest(); norm=lambda s:re.sub(r'\s+',' ',s).strip()
body=BODY.read_text(encoding='utf-8-sig');bp=BP.read_text(encoding='utf-8-sig');pack=PACK.read_text(encoding='utf-8-sig')
with zipfile.ZipFile(DOCX) as z:
 bad=z.testzip();missing=sorted({'[Content_Types].xml','word/document.xml','word/styles.xml'}-set(z.namelist()));root=ET.fromstring(z.read('word/document.xml'))
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
xmltext='\n'.join(''.join(p.itertext()) for p in root.findall('.//w:p',ns))
xmltext=norm(xmltext)
EXTRACT.write_text(xmltext+'\n',encoding='utf-8',newline='\n')
d=Document(DOCX); alltexts=[norm(p.text) for p in d.paragraphs if norm(p.text)]
for sec in d.sections:
 for p in list(sec.header.paragraphs)+list(sec.footer.paragraphs):
  if norm(p.text): alltexts.append(norm(p.text))
for table in d.tables:
 for row in table.rows:
  alltexts.append(' | '.join(norm(c.text) for c in row.cells))
for x in root.findall('.//w:t',ns):
 if x.text: alltexts.append(norm(x.text))
text='\n'.join(alltexts)
with (R.parent/'0076_R01/inputs/SYNTHETIC_ASSESSMENT_DATA_0076_R01.csv').open(encoding='utf-8-sig',newline='') as f: assess=list(csv.DictReader(f))
with (R.parent/'0076_R01/inputs/SYNTHETIC_ITEM_ANALYSIS_0076_R01.csv').open(encoding='utf-8-sig',newline='') as f: itemrows=list(csv.DictReader(f))
with (R.parent/'0077_R01/inputs/REFLECTION_SYNTHETIC_DATA_0077_R01.csv').open(encoding='utf-8-sig',newline='') as f: refs=list(csv.DictReader(f))
with REG.open(encoding='utf-8-sig',newline='') as f: reg=list(csv.DictReader(f))
with SUMMARY.open(encoding='utf-8-sig',newline='') as f: summ=list(csv.DictReader(f))
pre=sum(int(x['pretest_total_from_0074_synthetic']) for x in assess); post=sum(int(x['posttest_total_from_0074_synthetic']) for x in assess); pract=sum(int(x['practical_total_0_20']) for x in assess)
keys={x['item_id']:x['answer_key_letter'] for x in itemrows}; correct={q:sum(x[f'{q}_response']==key for x in assess) for q,key in keys.items()}
themes=collections.Counter(x['primary_theme_code'] for x in refs)
wrong='3bdce0a1675a5b377dfafe2d7519f9f877e51530f5edc4ba3551a116bc6e985'; actual='3bdce0a1675a5b377dfafe2d75139d1f877e51530f5edc4ba3551a116bc6e985';sha79='0361cf65557733cb3e1af93d5e59197aa96d307622a840800d008d8c580c66a7'
build=json.loads(BUILD.read_text(encoding='utf-8-sig')); cjk=len(re.findall(r'[\u4e00-\u9fff]',body))
subscores_ok=all(int(x['practical_total_0_20'])==sum(int(x[k]) for k in ['communication_0_5','privacy_boundary_0_5','record_accuracy_0_5','handoff_followup_0_5']) for x in assess)
checks={
'zip_integrity':bad is None,'required_ooxml_parts':not missing,'docx_sha_matches_build_receipt':sha(DOCX)==build['docx_sha256'],
'identity_0080_I17B_E02_085':all(x in xmltext for x in ['0080','I17B-E02','085']),
'official_neighbor_0081_identity_086':all(x in xmltext for x in ['0081','I17B-E03','086']),
'official_neighbor_0082_identity_088':all(x in xmltext for x in ['0082','I17C-E01','088']),
'current_0079_R07_exact_sha_bound':sha79 in body and sha79 in pack and sha79 in bp and sha79 in xmltext,
'corrected_0077_source_sha_consistent':actual in body and actual in pack and actual in bp and wrong not in body+pack+bp,
'appendix_input_hashes_match':sha(BP) in body and sha(PACK) in body,
'sample_real_boundary':all(x in xmltext for x in ['SAMPLE/SYNTHETIC','REAL=UNKNOWN/HOLD','NOT_GRANTED','PROPOSED','NOT_EXECUTED']),
'no_real_meeting_claim':all(x in xmltext for x in ['2026-08-29','未實際召開','無任何真實與會者']),
'upstream_rows_16_each':len(assess)==16 and len(refs)==16,
'replacement_register_23_rows':len(reg)==23,
'summary_dataset_26_rows':len(summ)==26,
'pretest_sum_mean':pre==870 and pre/16==54.375 and '54.375' in xmltext,
'posttest_sum_mean':post==1265 and post/16==79.0625 and '79.0625' in xmltext,
'gain_mean':post/16-pre/16==24.6875 and '+24.6875' in xmltext,
'practical_total_and_mean':pract==253 and pract/16==15.8125 and '15.8125' in xmltext and subscores_ok,
'item_counts_recomputed_and_reported':all(int(x['correct_n'])==correct[x['item_id']] for x in itemrows) and 'Q01／Q17／Q18／Q19' in body and '各 11／16＝0.6875' in body and 'Q02／Q03' in body and '各 12／16＝0.75' in body,
'reflection_theme_counts_recomputed':sorted(themes.values())==[1,3,3,4,5] and all(x in xmltext for x in ['需求澄清','隱私用途同意','最小必要','閉環交接','角色界線']),
'cjk_body_size_7000_9500':7000<=cjk<=9500,
'candidate_not_formally_accepted':build['status']=='BUILT_UNVERIFIED' and build['formal_acceptance']=='NOT_GRANTED' and build['real_evidence']=='UNKNOWN/HOLD'
}
receipt={'schema':'TTQS_ONE_0080_TECHNICAL_REVIEW_R03.v1','document_id':'0080','requirement_id':'I17B-E02','official_checkbox_serial':'085','candidate_docx_path':'candidate/0080_R05.docx','candidate_docx_sha256':sha(DOCX),'candidate_docx_bytes':DOCX.stat().st_size,'body_sha256':sha(BODY),'unique_blueprint_sha256':sha(BP),'source_packet_sha256':sha(PACK),'rendered_ooxml_text_sha256':hashlib.sha256(xmltext.encode()).hexdigest(),'source_csv_sha256':{'assessment':sha(R.parent/'0076_R01/inputs/SYNTHETIC_ASSESSMENT_DATA_0076_R01.csv'),'item_analysis':sha(R.parent/'0076_R01/inputs/SYNTHETIC_ITEM_ANALYSIS_0076_R01.csv'),'reflections':sha(R.parent/'0077_R01/inputs/REFLECTION_SYNTHETIC_DATA_0077_R01.csv')},'recomputed_metrics':{'pre_rows':len(assess),'pre_total':pre,'pre_mean':pre/16,'post_total':post,'post_mean':post/16,'mean_gain':post/16-pre/16,'practical_total':pract,'practical_mean':pract/16,'item_correct_counts':correct,'reflection_theme_counts':dict(themes),'practical_subscores_match_total':subscores_ok},'body_cjk_characters':cjk,'top_level_tables':len(d.tables),'checks':checks,'status':'PASS' if bad is None and not missing and all(checks.values()) else 'FAIL','formal_acceptance':'NOT_GRANTED','real_evidence':'UNKNOWN/HOLD','independent_semantic_review':'PENDING'}
OUT.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'docx_sha256':receipt['candidate_docx_sha256'],'status':receipt['status'],'checks':checks,'metrics':receipt['recomputed_metrics'],'cjk':cjk,'nested_mapping_present':all(x in xmltext for x in ['I17B-E03','I17C-E01'])},ensure_ascii=False))
