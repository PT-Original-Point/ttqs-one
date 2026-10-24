import re, sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

src=Path(sys.argv[1]); out=Path(sys.argv[2]); text=src.read_text(encoding="utf-8")
doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
sec.top_margin=Inches(0.68); sec.bottom_margin=Inches(0.67)
sec.left_margin=Inches(0.72); sec.right_margin=Inches(0.72)
sec.header_distance=Inches(0.3); sec.footer_distance=Inches(0.3)
FONT="Microsoft JhengHei"; NAVY="1F3B5B"; BLUE="DCE8F2"; LIGHT="F3F6F9"; GRAY="5C6770"; BORDER="B8C6D1"

def set_run_font(run,size=10.5,bold=None,color=None,italic=None):
    run.font.name=FONT; run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),FONT); run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    if italic is not None: run.italic=italic
    if color: run.font.color.rgb=RGBColor.from_string(color)

def set_cell_shading(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn("w:shd"))
    if shd is None: shd=OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"),fill)

def set_cell_margins(cell,top=75,start=95,bottom=75,end=95):
    tcPr=cell._tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in("w:tcMar")
    if tcMar is None: tcMar=OxmlElement("w:tcMar"); tcPr.append(tcMar)
    for m,v in (("top",top),("start",start),("bottom",bottom),("end",end)):
        el=tcMar.find(qn("w:"+m))
        if el is None: el=OxmlElement("w:"+m); tcMar.append(el)
        el.set(qn("w:w"),str(v)); el.set(qn("w:type"),"dxa")

def add_border(cell):
    tcPr=cell._tc.get_or_add_tcPr(); borders=tcPr.first_child_found_in("w:tcBorders")
    if borders is None: borders=OxmlElement("w:tcBorders"); tcPr.append(borders)
    for edge in ("top","left","bottom","right","insideH","insideV"):
        tag="w:"+edge; el=borders.find(qn(tag))
        if el is None: el=OxmlElement(tag); borders.append(el)
        el.set(qn("w:val"),"single"); el.set(qn("w:sz"),"4"); el.set(qn("w:space"),"0"); el.set(qn("w:color"),BORDER)

def add_inline(p,value,size=10.5,color=None):
    for part in re.split(r"(\*\*.*?\*\*)",value):
        if not part: continue
        bold=part.startswith("**") and part.endswith("**")
        val=part[2:-2] if bold else part
        run=p.add_run(val); set_run_font(run,size=size,bold=True if bold else None,color=color)

def add_page_field(p):
    run=p.add_run(); set_run_font(run,9,color=GRAY)
    begin=OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"),"begin")
    instr=OxmlElement("w:instrText"); instr.set(qn("xml:space"),"preserve"); instr.text=" PAGE "
    sep=OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"),"separate")
    textel=OxmlElement("w:t"); textel.text="1"
    end=OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"),"end")
    for el in (begin,instr,sep,textel,end):run._r.append(el)

normal=doc.styles["Normal"]; normal.font.name=FONT; normal._element.rPr.rFonts.set(qn("w:eastAsia"),FONT); normal.font.size=Pt(10.5); normal.font.color.rgb=RGBColor.from_string("202B35"); normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.12
for name,size in (("Title",24),("Heading 1",15.5),("Heading 2",12.5),("Heading 3",11)):
    s=doc.styles[name]; s.font.name=FONT; s._element.rPr.rFonts.set(qn("w:eastAsia"),FONT); s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string("202B35" if name=="Title" else NAVY); s.paragraph_format.keep_with_next=True; s.paragraph_format.space_before=Pt(10); s.paragraph_format.space_after=Pt(5)
# The built-in Word Title style may carry a paragraph border. Remove it so the
# document title stays plain and does not acquire an unintended blue rule.
title_ppr=doc.styles["Title"]._element.get_or_add_pPr()
title_border=title_ppr.find(qn("w:pBdr"))
if title_border is not None: title_ppr.remove(title_border)
doc.styles["Title"].paragraph_format.space_before=Pt(24); doc.styles["Title"].paragraph_format.space_after=Pt(8)
header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT; add_inline(header,"原始點協會｜DOCX 0123｜候選示範文件",9,color=GRAY)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; add_inline(footer,"REAL 個案資料未提供  ·  正式接受未授權  ·  第 ",9,color=GRAY); add_page_field(footer); add_inline(footer," 頁",9,color=GRAY)
upd=OxmlElement("w:updateFields"); upd.set(qn("w:val"),"true"); doc.settings.element.append(upd)

def table_rows(lines,start):
    rows=[]; i=start
    while i<len(lines) and lines[i].strip().startswith("|"):
        cells=[c.strip().replace(r"\|","|") for c in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?",c.replace(" ","")) for c in cells): rows.append(cells)
        i+=1
    return rows,i

def make_table(rows):
    if not rows:return
    cols=max(len(r) for r in rows); rows=[r+[""]*(cols-len(r)) for r in rows]
    table=doc.add_table(rows=len(rows),cols=cols); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
    width=7.0
    headers=[c.strip() for c in rows[0]]
    ratios={1:[1.0],2:[0.28,0.72],3:[0.21,0.39,0.40],4:[0.18,0.30,0.32,0.20],5:[0.15,0.23,0.24,0.25,0.13],6:[0.11,0.18,0.17,0.21,0.15,0.18]}
    if cols==4 and headers[0].startswith("替換鍵"):
        ratios[4]=[0.20,0.25,0.37,0.18]
    elif cols==2 and headers[0]=="核對項目":
        ratios[2]=[0.35,0.65]
    elif cols==5 and headers[0]=="項次":
        ratios[5]=[0.08,0.27,0.30,0.24,0.11]
    elif cols==5 and headers[0]=="行動":
        ratios[5]=[0.23,0.18,0.15,0.22,0.22]
    weights=ratios.get(cols,[1/cols]*cols)
    if len(weights)!=cols: weights=[1/cols]*cols
    for j,w in enumerate(weights): table.columns[j].width=Inches(width*w)
    for ri,row in enumerate(rows):
        trPr=table.rows[ri]._tr.get_or_add_trPr(); trPr.append(OxmlElement("w:cantSplit"))
        if ri==0: th=OxmlElement("w:tblHeader"); th.set(qn("w:val"),"true"); trPr.append(th)
        for ci,val in enumerate(row):
            cell=table.cell(ri,ci); cell.width=Inches(width*weights[ci]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cell); add_border(cell)
            if ri==0:set_cell_shading(cell,NAVY)
            elif ri%2==0:set_cell_shading(cell,LIGHT)
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
            # Long underscore runs in narrow form cells wrap into fragmented
            # lines in Word. Keep a visible, usable short fill line instead.
            display_val=re.sub(r"_{12,}","__________",val) if cols>=5 else val
            add_inline(p,display_val,size=9.5,color=("FFFFFF" if ri==0 else None))
            if ri==0:
                for r in p.runs:r.bold=True
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

lines=text.splitlines(); i=0; first_title=True
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith("|"):
        rows,i=table_rows(lines,i);make_table(rows);continue
    if line.startswith("# "):
        p=doc.add_paragraph(style="Title" if first_title else "Heading 1"); first_title=False
        if p.style.name=="Title":
            ppr=p._p.get_or_add_pPr(); pborder=ppr.find(qn("w:pBdr"))
            if pborder is not None:ppr.remove(pborder)
        add_inline(p,line[2:].strip(),size=24 if p.style.name=="Title" else 15.5,color="202B35" if p.style.name=="Title" else NAVY);i+=1;continue
    if line.startswith("## "):
        p=doc.add_paragraph(style="Heading 1");add_inline(p,line[3:].strip(),size=15.5,color=NAVY);i+=1;continue
    if line.startswith("### "):
        p=doc.add_paragraph(style="Heading 2");add_inline(p,line[4:].strip(),size=12.5,color=NAVY);i+=1;continue
    if line.startswith("#### "):
        p=doc.add_paragraph(style="Heading 3");add_inline(p,line[5:].strip(),size=11,color=NAVY);i+=1;continue
    if line.startswith("> "):
        q=[]
        while i<len(lines) and lines[i].strip().startswith(">"):q.append(lines[i].strip()[1:].strip());i+=1
        p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.18);p.paragraph_format.right_indent=Inches(.12);p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(7)
        shd=OxmlElement("w:shd");shd.set(qn("w:fill"),BLUE);p._p.get_or_add_pPr().append(shd);add_inline(p," ".join(q),size=10.5,color=NAVY);continue
    if re.match(r"^\d+\.\s+",line):
        p=doc.add_paragraph(style="List Number");add_inline(p,re.sub(r"^\d+\.\s+","",line),size=10.5);i+=1;continue
    if line.startswith("- "):
        p=doc.add_paragraph(style="List Bullet");add_inline(p,line[2:],size=10.5);i+=1;continue
    p=doc.add_paragraph();add_inline(p,line,size=10.5);i+=1

doc.core_properties.title="會員訓後事業成就個案紀錄"
doc.core_properties.subject="DOCX 0123 / I19-E11 / SAMPLE-SYNTHETIC candidate R01"
doc.core_properties.author="TTQS ONE candidate builder"
doc.save(out)
print("CREATED="+str(out));print("PARAGRAPHS="+str(len(doc.paragraphs))+" TABLES="+str(len(doc.tables)))
