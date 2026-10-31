from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[4]
BASE = ROOT / "docs" / "ttqs-quality-20261016" / "quality-review" / "0059_R01"
DOCX = BASE / "0059_R03.docx"
CONTENT = BASE / "worker" / "opencode" / "0059_DOCUMENT_BODY_R03_FINAL.md"
PDF = BASE / "evidence" / "render_final_0263017F" / "0059_R03.pdf"
RENDER = PDF.parent
OUT = BASE / "evidence" / "technical_layout_qa_R03.json"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest().upper()


def text(cell) -> str:
    return " ".join(cell.text.split())


def table_by_header(doc, first: str, second: str | None = None):
    normalize = lambda s: "".join(s.split())
    for table in doc.tables:
        if not table.rows:
            continue
        row = [text(c) for c in table.rows[0].cells]
        if row and normalize(first) in normalize(row[0]) and (second is None or any(normalize(second) in normalize(x) for x in row[1:])):
            return table
    raise AssertionError(f"table not found: {first} / {second}")


def effective_run_size(run, paragraph) -> float | None:
    if run.font.size is not None:
        return run.font.size.pt
    style = paragraph.style
    while style is not None:
        if style.font.size is not None:
            return style.font.size.pt
        style = style.base_style
    return None


doc = Document(DOCX)
body_parts: list[str] = []
for p in doc.paragraphs:
    if p.text.strip():
        body_parts.append(p.text.strip())
for table in doc.tables:
    for row in table.rows:
        body_parts.append(" | ".join(text(cell) for cell in row.cells))
all_text = "\n".join(body_parts)

errors: list[str] = []
expected_docx_sha = "0263017F16D81B1121BB97410DE9895B43B7A42F75632BF71A82B86426FBC2B0"
actual_docx_sha = sha256(DOCX)
if actual_docx_sha != expected_docx_sha:
    errors.append("DOCX SHA differs from reviewed candidate SHA")
if len(doc.tables) != 59:
    errors.append(f"expected 59 tables, found {len(doc.tables)}")

required_strings = [
    "I13-E03", "勾選單 063", "SAMPLE-SYN-113", "NOT ADOPTED", "HOLD", "UNKNOWN",
    "提供行動計畫、競賽、回訓、學習心得交流的平台、證照、實習或其他相關機制的建議或協助",
    "四機制", "退補件率（退補件÷投稿件）≤10%",
]
normalized_all_text = "".join(all_text.split())
missing_strings = [s for s in required_strings if "".join(s.split()) not in normalized_all_text]
if missing_strings:
    errors.append("missing required body strings: " + "; ".join(missing_strings))

for forbidden in ["SYN-ISSUER", "SYN-機構", "理ubectl", "文件結束"]:
    if forbidden in all_text:
        errors.append(f"forbidden residue found: {forbidden}")

# Recompute all ten first-test records directly from the DOCX table.
first_test = table_by_header(doc, "學員", "五構面分數")
first_rows = [[text(c) for c in row.cells] for row in first_test.rows[1:]]
scores: dict[str, list[int]] = {}
for row in first_rows:
    if len(row) >= 4 and re.fullmatch(r"S-\d{2}", row[0]):
        values = [int(x) for x in row[1].split("/")]
        if len(values) != 5:
            errors.append(f"wrong rubric dimension count for {row[0]}")
        scores[row[0]] = values
score_totals = {person: sum(values) for person, values in scores.items()}
dimension_totals = [sum(scores[p][i] for p in scores) for i in range(5)] if scores else []
first_test_pass = [p for p, total in score_totals.items() if total >= 11 and scores[p][4] > 0]
if len(scores) != 10 or len(first_test_pass) != 8 or sum(score_totals.values()) != 116 or dimension_totals != [22, 22, 20, 23, 29]:
    errors.append("first-test rubric arithmetic mismatch")

retest = table_by_header(doc, "學員", "首測分數／結果")
retest_rows = {text(row.cells[0]): [text(c) for c in row.cells] for row in retest.rows[1:]}
if set(retest_rows) != {"S-15", "S-16"} or "11" not in retest_rows.get("S-15", [""])[-1] or "HOLD" not in retest_rows.get("S-16", [""])[-1]:
    errors.append("retest population/result mismatch")

survey_tables = []
for table in doc.tables:
    if not table.rows:
        continue
    header = [text(c) for c in table.rows[0].cells]
    if header and header[0] == "代碼" and any(x.startswith(("Q1", "Q5", "Q9")) for x in header[1:]):
        survey_tables.append(table)
if len(survey_tables) != 3:
    errors.append(f"expected 3 raw-response survey tables, found {len(survey_tables)}")
survey: dict[str, list[int]] = {f"Q{i}": [] for i in range(1, 11)}
for table in survey_tables:
    headers = [text(c) for c in table.rows[0].cells]
    for row in table.rows[1:]:
        for heading, cell in zip(headers[1:], row.cells[1:]):
            q = re.match(r"(Q\d+)", heading)
            val = text(cell)
            if q and val != "N/A":
                survey[q.group(1)].append(int(val))
survey_stats = {q: {"n": len(v), "sum": sum(v), "mean": round(sum(v) / len(v), 1) if v else None} for q, v in survey.items()}
expected_survey = {
    "Q1": (15, 63, 4.2), "Q2": (12, 48, 4.0), "Q3": (15, 60, 4.0), "Q4": (9, 38, 4.2),
    "Q5": (15, 57, 3.8), "Q6": (15, 57, 3.8), "Q7": (15, 57, 3.8), "Q8": (15, 57, 3.8),
    "Q9": (15, 60, 4.0), "Q10": (15, 60, 4.0),
}
for q, expected in expected_survey.items():
    actual = (survey_stats[q]["n"], survey_stats[q]["sum"], survey_stats[q]["mean"])
    if actual != expected:
        errors.append(f"survey arithmetic mismatch {q}: {actual} != {expected}")

cross_table = table_by_header(doc, "學員", "A/B/D 去重")
cross_rows = [[text(c) for c in row.cells] for row in cross_table.rows[1:]]
people = [r[0] for r in cross_rows if len(r) >= 6 and any("●" in r[i] for i in (1, 2, 4))]
if len(set(people)) != 16 or set(people) != {f"S-{i:02d}" for i in range(1, 17)}:
    errors.append(f"cross-mechanism deduplication mismatch: {sorted(set(people))}")

pdf_sha = sha256(PDF)
page_count = len(PdfReader(str(PDF)).pages)
page_files = sorted(RENDER.glob("page-*.png"), key=lambda p: int(p.stem.split("-")[-1]))
if len(page_files) != page_count or page_count != 34:
    errors.append(f"render page count mismatch: pdf={page_count}, png={len(page_files)}")

paragraph_sizes = [effective_run_size(run, p) for p in doc.paragraphs for run in p.runs if run.text.strip()]
table_sizes = [effective_run_size(run, p) for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs for run in p.runs if run.text.strip()]
paragraph_sizes = [x for x in paragraph_sizes if x is not None]
table_sizes = [x for x in table_sizes if x is not None]
min_body_pt = min(paragraph_sizes) if paragraph_sizes else None
min_table_pt = min(table_sizes) if table_sizes else None
if min_body_pt is not None and min_body_pt < 10.5:
    errors.append(f"body font below 10.5pt: {min_body_pt}")
if min_table_pt is not None and min_table_pt < 9.5:
    errors.append(f"table font below 9.5pt: {min_table_pt}")

receipt = {
    "status": "PASS" if not errors else "FAIL",
    "candidate_state": "CANDIDATE / SAMPLE-SYNTHETIC; not formally accepted",
    "docx": {"path": str(DOCX), "sha256": actual_docx_sha, "bytes": DOCX.stat().st_size, "tables": len(doc.tables)},
    "source_content": {"path": str(CONTENT), "sha256": sha256(CONTENT)},
    "render": {"route": "Microsoft Word COM 16.0 to PDF, then Poppler pdftoppm", "pdf_sha256": pdf_sha, "page_count": page_count, "png_count": len(page_files), "page_png_sha256": {p.name: sha256(p) for p in page_files}},
    "typography": {"minimum_body_run_pt": min_body_pt, "minimum_table_run_pt": min_table_pt},
    "first_test": {"records": len(scores), "pass_count": len(first_test_pass), "score_sum": sum(score_totals.values()), "dimension_sums": dimension_totals, "person_scores": score_totals},
    "retest": {person: {"record": row} for person, row in retest_rows.items()},
    "survey": survey_stats,
    "cross_mechanism_union": {"count": len(set(people)), "participants": sorted(set(people)), "rate": "16/18=88.9%"},
    "independent_visual_review": "See separately authored visual receipt; this script does not claim image inspection.",
    "errors": errors,
}
OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps({"status": receipt["status"], "docx_sha256": actual_docx_sha, "pdf_sha256": pdf_sha, "page_count": page_count, "tables": len(doc.tables), "min_body_pt": min_body_pt, "min_table_pt": min_table_pt, "first_test_pass": len(first_test_pass), "retest_ids": sorted(retest_rows), "survey": survey_stats, "cross_union": len(set(people)), "errors": errors}, ensure_ascii=False, indent=2))
if errors:
    raise SystemExit(1)
