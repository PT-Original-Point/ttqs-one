from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table
from pypdf import PdfReader
import hashlib, json, re, datetime
root=Path(r'E:\TTQS\_github_sync_0064\docs\ttqs-quality-20261016\quality-review\0070_R01')
docx=root/'candidate'/'0070_R01.docx'
body=root/'candidate'/'0070_R01_body.md'
official=root/'inputs'/'official_checklist_source.pdf'
blueprint=root/'inputs'/'UNIQUE_BLUEPRINT_0070_R01.md'
official_extract=root/'inputs'/'OFFICIAL_PAGE5_EXTRACT_R01.md'
source_facts=root/'inputs'/'SOURCE_FACTS_R01.md'
render=root/'evidence'/'render_word_R03_final'
pdf=render/'candidate.pdf'
imgs=sorted(render.glob('page-*.png'))
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def extract(d):
    doc=Document(d); out=[]
    for node in doc.element.body.iterchildren():
        if node.tag.endswith('}p'):
            t=Paragraph(node,doc).text.strip()
            if t: out.append(t)
        elif node.tag.endswith('}tbl'):
            table=Table(node,doc); out.append('[TABLE]')
            for row in table.rows: out.append(' | '.join(' '.join(c.text.split()) for c in row.cells))
            out.append('[/TABLE]')
    return doc,out
D,lines=extract(docx)
text='\n'.join(lines)
(root/'evidence'/'docx_extracted_R03.md').write_text(f'# Exact candidate extraction\n\nDOCX SHA-256: `{sha(docx)}`\n\n'+'\n'.join(lines)+'\n',encoding='utf-8')
with ZipFile(docx) as z: bad=z.testzip(); members=z.namelist()
body_runs=[r for p in D.paragraphs for r in p.runs if r.text.strip() and r.font.size]
table_runs=[r for t in D.tables for row in t.rows for c in row.cells for p in c.paragraphs for r in p.runs if r.text.strip() and r.font.size]
min_body=min((r.font.size.pt for r in body_runs),default=None)
min_table=min((r.font.size.pt for r in table_runs),default=None)
blueprint_sha=sha(blueprint); official_sha=sha(official); official_extract_sha=sha(official_extract); source_facts_sha=sha(source_facts)
checks={
 'docx_zip_integrity': bad is None,
 'docx_xml_parse': D is not None,
 'document_identity_0070_I16_E02_checkbox_077': all(s in text for s in ['DOCX 0070','I16-E02','077']),
 'official_source_identity': official_sha=='805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709',
 'blueprint_and_source_map_present': all(p.exists() for p in [blueprint,official_extract,source_facts]),
 'sample_synthetic_boundary': 'SAMPLE' in text and 'SYNTHETIC' in text,
 'real_unknown_hold_boundary': 'UNKNOWN' in text and 'HOLD' in text,
 'case_id_present': 'SYN-0070-CASE-01' in text,
 'date_and_sequence_consistency': all(s in text for s in ['2026-09-25','08:42','08:44','08:48','08:55','09:30','11:00','2026-10-02','2026-10-09','2026-10-16','2026-10-17']),
 'notification_math': all(s in text for s in ['10 ÷ 12 | 83.3%','2 ÷ 2 | 100%','12 ÷ 12 | 100%']),
 'effect_math': all(s in text for s in ['9 ÷ 9 | 100%','3 ÷ 3 | 100%']),
 'cross_document_boundaries': all(s in text for s in ['0069','0071','0072','0033','case_id']),
 'reject_hold_criteria': 'C1' in text and 'C10' in text and 'H1' in text and 'H7' in text,
 'local_path_absent_from_candidate': r'E:\TTQS' not in text,
 'technical_or_model_metadata_absent_from_candidate': not any(s in text.lower() for s in ['sha-256','opencode','antigravity cli','codex model']),
 'body_font_floor_met': min_body is not None and min_body>=10.5,
 'table_font_floor_met': min_table is not None and min_table>=9.5,
 'render_pdf_22_pages': pdf.exists() and len(PdfReader(str(pdf)).pages)==22,
 'render_png_22_pages': len(imgs)==22,
 'render_receipt_matches_candidate': json.loads((render/'export_receipt.json').read_text(encoding='utf-8-sig'))['source_sha256_after']==sha(docx),
}
receipt={
 'project_id':'TTQS_ONE','doc_id':'0070','requirement_id':'I16-E02','official_checkbox_serial':'077','revision':'R03_FINAL_CANDIDATE',
 'candidate_state':'CANDIDATE_SAMPLE_SYNTHETIC_NOT_FORMALLY_ACCEPTED','docx_sha256':sha(docx),'docx_bytes':docx.stat().st_size,
 'body_sha256':sha(body),'extracted_candidate_text_sha256':hashlib.sha256(text.encode()).hexdigest(),
 'source_hashes':{'official_pdf_sha256':official_sha,'official_extract_sha256':official_extract_sha,'unique_blueprint_sha256':blueprint_sha,'source_facts_sha256':source_facts_sha},
 'structure':{'tables':len(D.tables),'paragraphs_top_level':len(D.paragraphs),'sections':len(D.sections),'body_explicit_font_min_pt':min_body,'table_explicit_font_min_pt':min_table,'body_runs_checked':len(body_runs),'table_runs_checked':len(table_runs),'zip_members':len(members)},
 'render':{'pdf_sha256':sha(pdf),'pdf_pages':len(PdfReader(str(pdf)).pages),'png_pages':len(imgs),'renderer':'Microsoft Word COM ExportAsFixedFormat + Poppler pdftoppm','dpi':140},
 'checks':checks,'failed_checks':[k for k,v in checks.items() if not v],
 'technical_status':'PASS_TECHNICAL_SCOPE' if all(checks.values()) else 'FAIL_TECHNICAL_SCOPE',
 'semantic_acceptance':'SEPARATE_INDEPENDENT_REVIEW_RECEIPT_REQUIRED','visual_acceptance':'SEPARATE_VISUAL_RECEIPT',
 'association_approval':'NOT_CLAIMED','checked_at_local':datetime.datetime.now().astimezone().isoformat(timespec='seconds')}
(root/'evidence'/'technical_review_R03.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'docx_sha256':receipt['docx_sha256'],'docx_bytes':receipt['docx_bytes'],'extracted_text_sha256':receipt['extracted_candidate_text_sha256'],'checks':checks,'technical_status':receipt['technical_status'],'tables':len(D.tables),'body_min_pt':min_body,'table_min_pt':min_table,'render_pages':len(imgs)},ensure_ascii=False,indent=2))
