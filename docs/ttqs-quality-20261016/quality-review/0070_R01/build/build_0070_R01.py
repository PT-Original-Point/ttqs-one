#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Reproducible local build for DOCX 0070 R01.

Reads:  candidate/0070_R01_body.md
Writes: candidate/0070_R01.docx
        qa/build_receipt_R01.json

Deterministic, offline, no network, no Word rendering, no desktop UI.
Enforces: body font floor >= 10.5 pt, table font floor >= 9.5 pt,
stable fixed table widths, repeated header rows, no row splitting,
and explicit page breaks from the [[PAGEBREAK]] directive.
"""

import hashlib
import json
import re
import struct
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ---------------------------------------------------------------- constants
ROOT = Path(__file__).resolve().parents[1]
BODY_MD = ROOT / "candidate" / "0070_R01_body.md"
OUT_DOCX = ROOT / "candidate" / "0070_R01.docx"
OUT_RECEIPT = ROOT / "qa" / "build_receipt_R01.json"

BODY_FLOOR_PT = 10.5
TABLE_FLOOR_PT = 9.5
BODY_SIZE_PT = 11.0
TABLE_SIZE_PT = 10.0

CJK_FONT = "Microsoft JhengHei"
LATIN_FONT = "Microsoft JhengHei"

CONTENT_WIDTH_CM = 17.4  # A4 (21cm) minus 2 x 1.8cm margins


# ---------------------------------------------------------------- helpers
def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def _set_run_font(run, size_pt, bold=None, italic=None):
    run.font.name = LATIN_FONT
    run.font.size = Pt(size_pt)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), LATIN_FONT)
    rfonts.set(qn("w:hAnsi"), LATIN_FONT)
    rfonts.set(qn("w:eastAsia"), CJK_FONT)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic


def _add_inline(paragraph, text, size_pt, base_bold=False):
    """Add text supporting **bold** spans into an existing paragraph."""
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            _set_run_font(run, size_pt, bold=True)
        else:
            run = paragraph.add_run(part)
            _set_run_font(run, size_pt, bold=base_bold)


def _shade(element, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def _set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    tr_pr.append(el)


def _set_fixed_layout(table):
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    width = OxmlElement("w:tblW")
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(int(CONTENT_WIDTH_CM * 567)))
    tbl_pr.append(width)


def _set_col_widths(table, ncols):
    total_twips = int(CONTENT_WIDTH_CM * 567)
    share = total_twips // ncols
    widths = [share] * ncols
    widths[-1] += total_twips - sum(widths)
    # tblGrid
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for gridcol, w in zip(grid.findall(qn("w:gridCol")), widths):
            gridcol.set(qn("w:w"), str(w))
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Pt(w / 20.0)


def _cell_text(cell, text, size_pt, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(1)
    _add_inline(paragraph, text, size_pt, base_bold=bold)
    return paragraph


def _parse_row(line):
    inner = line.strip()
    if inner.startswith("|"):
        inner = inner[1:]
    if inner.endswith("|"):
        inner = inner[:-1]
    return [c.strip() for c in inner.split("|")]


def _is_separator(line):
    cells = _parse_row(line)
    return all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c != "") and len(cells) > 0


# ---------------------------------------------------------------- document
def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(BODY_SIZE_PT)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15

    heading_sizes = {"Heading 1": 18, "Heading 2": 14, "Heading 3": 12, "Title": 24}
    for name, size in heading_sizes.items():
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        st.font.name = LATIN_FONT
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0x1F, 0x33, 0x64)
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        rfonts.set(qn("w:eastAsia"), CJK_FONT)


def add_table(doc: Document, rows):
    header = rows[0]
    body = rows[1:]
    ncols = len(header)
    table = doc.add_table(rows=1 + len(body), cols=ncols)
    table.style = "Table Grid"
    table.autofit = False
    _set_fixed_layout(table)

    for j, text in enumerate(header):
        _cell_text(table.rows[0].cells[j], text, TABLE_SIZE_PT, bold=True)
        _shade(table.rows[0].cells[j]._tc.get_or_add_tcPr(), "D9E2F3")
    _set_repeat_header(table.rows[0])
    _set_cant_split(table.rows[0])

    for i, r in enumerate(body, start=1):
        for j in range(ncols):
            value = r[j] if j < len(r) else ""
            _cell_text(table.rows[i].cells[j], value, TABLE_SIZE_PT)
        _set_cant_split(table.rows[i])

    _set_col_widths(table, ncols)
    return table


def build(md_text: str) -> Document:
    doc = Document()
    configure_document(doc)

    lines = md_text.splitlines()
    i = 0
    total_tables = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped == "[[PAGEBREAK]]":
            p = doc.add_paragraph()
            p.add_run().add_break(WD_BREAK.PAGE)
            i += 1
            continue

        if stripped == "":
            i += 1
            continue

        if stripped.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            rows = []
            for b in block:
                if _is_separator(b):
                    continue
                rows.append(_parse_row(b))
            if rows:
                add_table(doc, rows)
                total_tables += 1
            continue

        if stripped.startswith("### "):
            h = doc.add_heading(level=3)
            _add_inline(h, stripped[4:], 12.0, base_bold=True)
            i += 1
            continue

        if stripped.startswith("## "):
            h = doc.add_heading(level=2)
            _add_inline(h, stripped[3:], 14.0, base_bold=True)
            i += 1
            continue

        if stripped.startswith("# "):
            h = doc.add_heading(level=1)
            _add_inline(h, stripped[2:], 18.0, base_bold=True)
            i += 1
            continue

        if stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.4)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(6)
            _shade(p._p.get_or_add_pPr(), "FFF2CC")
            _add_inline(p, stripped[2:], BODY_SIZE_PT, base_bold=True)
            i += 1
            continue

        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            _add_inline(p, stripped[2:], BODY_SIZE_PT)
            i += 1
            continue

        p = doc.add_paragraph()
        _add_inline(p, stripped, BODY_SIZE_PT)
        i += 1

    return doc


# ---------------------------------------------------------------- self-check
def collect_font_sizes(doc: Document):
    body_sizes, table_sizes = [], []

    def run_size(run, default):
        if run.font.size is not None:
            return run.font.size.pt
        style = run.style
        while style is not None:
            if style.font.size is not None:
                return style.font.size.pt
            style = style.base_style
        return default

    for p in doc.paragraphs:
        for run in p.runs:
            if run.text.strip():
                body_sizes.append(run_size(run, BODY_SIZE_PT))

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        if run.text.strip():
                            table_sizes.append(run_size(run, TABLE_SIZE_PT))

    return body_sizes, table_sizes


def main():
    md_text = BODY_MD.read_text(encoding="utf-8")
    doc = build(md_text)
    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT_DOCX))

    # reload for accurate structural counts
    check = Document(str(OUT_DOCX))
    body_sizes, table_sizes = collect_font_sizes(check)
    table_para_count = sum(
        len(cell.paragraphs) for t in check.tables for row in t.rows for cell in row.cells
    )

    body_min = min(body_sizes) if body_sizes else None
    table_min = min(table_sizes) if table_sizes else None

    receipt = {
        "artifact": "0070_R01",
        "doc_id": "0070",
        "requirement_id": "I16-E02",
        "checkbox_serial": "077",
        "generated_utc": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "builder": "scripts/build_0070_R01.py",
        "inputs": {
            "body_md": "candidate/0070_R01_body.md",
            "body_md_sha256": sha256_file(BODY_MD),
            "body_md_bytes": BODY_MD.stat().st_size,
        },
        "outputs": {
            "docx": "candidate/0070_R01.docx",
            "docx_sha256": sha256_file(OUT_DOCX),
            "docx_bytes": OUT_DOCX.stat().st_size,
        },
        "counts": {
            "tables": len(check.tables),
            "body_paragraphs_top_level": len(check.paragraphs),
            "table_cell_paragraphs": table_para_count,
            "total_paragraphs_including_tables": len(check.paragraphs) + table_para_count,
            "sections": len(check.sections),
        },
        "font_floor_checks": {
            "body_floor_pt": BODY_FLOOR_PT,
            "table_floor_pt": TABLE_FLOOR_PT,
            "body_min_pt_observed": body_min,
            "table_min_pt_observed": table_min,
            "body_floor_met": bool(body_min is not None and body_min >= BODY_FLOOR_PT - 1e-6),
            "table_floor_met": bool(table_min is not None and table_min >= TABLE_FLOOR_PT - 1e-6),
            "body_runs_checked": len(body_sizes),
            "table_runs_checked": len(table_sizes),
        },
        "structural_checks": {
            "fixed_table_layout": True,
            "repeat_header_rows": True,
            "prevent_row_split": True,
            "page_break_directives_in_body": md_text.count("[[PAGEBREAK]]"),
            "cjk_font": CJK_FONT,
            "latin_font": LATIN_FONT,
        },
        "acceptance": {
            "independent_semantic_acceptance": "NOT CLAIMED",
            "independent_visual_acceptance": "NOT CLAIMED",
            "association_approval": "NOT CLAIMED",
            "note": "Writer self-check only. Receipt records build facts; it is not a PASS.",
        },
    }

    OUT_RECEIPT.parent.mkdir(parents=True, exist_ok=True)
    OUT_RECEIPT.write_text(
        json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    print(json.dumps({
        "docx_sha256": receipt["outputs"]["docx_sha256"],
        "docx_bytes": receipt["outputs"]["docx_bytes"],
        "tables": receipt["counts"]["tables"],
        "paragraphs_total": receipt["counts"]["total_paragraphs_including_tables"],
        "body_min_pt": body_min,
        "table_min_pt": table_min,
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
