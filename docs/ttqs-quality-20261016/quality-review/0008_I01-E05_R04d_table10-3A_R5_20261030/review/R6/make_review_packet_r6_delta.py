from pathlib import Path
import hashlib
import json
import re

from docx import Document


ROOT = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0008_I01-E05_W1_TABLE10-3A_R04d_20261030")
QA = ROOT / "qa" / "reviewer_r5_readonly"
DOCX = ROOT / "candidate" / "0008_I01-E05_R04d_table10-3A_r5.docx"
STAGING = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\GITHUB_STAGING\0008_I01-E05_R04c_layoutfix_r2_20261030\source")
R5_JSON = QA / "SEMANTIC_REVIEW_R5.json"
PACKET = QA / "SEMANTIC_REVIEW_PACKET_R6_DELTA.md"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def clean(value: str) -> str:
    return " ".join(value.replace("\xa0", " ").split())


docx_sha = sha256(DOCX)
r5 = json.loads(R5_JSON.read_text(encoding="utf-8-sig"))
if r5.get("candidate_sha256", "").lower() != docx_sha:
    raise SystemExit("R5 review is not bound to the current DOCX bytes")

mapping_path = STAGING / "OFFICIAL_BLUEPRINT_MAPPING_R04c.json"
blueprint_path = STAGING / "BLUEPRINT_ROW_I01-E05.jsonl"
crosswalk_path = STAGING / "CROSSWALK_ROW_I01-E05.csv"
mapping = json.loads(mapping_path.read_text(encoding="utf-8-sig"))
blueprint = blueprint_path.read_text(encoding="utf-8-sig").strip()
crosswalk = crosswalk_path.read_text(encoding="utf-8-sig").strip()

document = Document(str(DOCX))
fact_table = document.tables[22]
fact_rows = []
for row in fact_table.rows[1:]:
    values = [clean(cell.text) for cell in row.cells]
    match = re.fullmatch(r"SF-(\d{5})", values[0] if values else "")
    if match and 689 <= int(match.group(1)) <= 717:
        fact_rows.append(values)

expected_ids = [f"SF-{number:05d}" for number in range(689, 718)]
actual_ids = [row[0] for row in fact_rows]
if actual_ids != expected_ids or len(set(actual_ids)) != 29:
    raise SystemExit(
        f"Expected 29 sequential unique rows, got {len(actual_ids)}: {actual_ids}"
    )

table_10_3a = document.tables[18]
table_lines = []
for row in table_10_3a.rows:
    values = [clean(cell.text).replace("|", "／") for cell in row.cells]
    table_lines.append("| " + " | ".join(values) + " |")

fact_lines = []
for row in fact_rows:
    values = [value.replace("|", "／") for value in row]
    fact_lines.append("| " + " | ".join(values) + " |")

parts = [
    "# TTQS ONE｜0008 R5 exact-SHA semantic addendum R6",
    "",
    "PROJECT_ID=TTQS_ONE",
    "DOC_ID=0008",
    "OFFICIAL_REQUIREMENT_ID=I01-E05",
    "OFFICIAL_CHECKBOX=I01-E05",
    "UNIQUE_BLUEPRINT_ID=I01-E05",
    f"CANDIDATE_DOCX_SHA256={docx_sha}",
    f"R5_REVIEW_JSON_SHA256={sha256(R5_JSON)}",
    f"R5_REVIEWER_RESPONSE_SHA256={sha256(QA / 'ORIGINAL_REVIEWER_RESPONSE_R5_RAW.txt')}",
    f"R5_PACKET_SHA256={sha256(QA / 'SEMANTIC_REVIEW_PACKET_R5.md')}",
    f"OFFICIAL_MAPPING_SHA256={sha256(mapping_path)}",
    f"UNIQUE_BLUEPRINT_ROW_SHA256={sha256(blueprint_path)}",
    f"CROSSWALK_ROW_SHA256={sha256(crosswalk_path)}",
    "CANDIDATE_STATE=CANDIDATE; FORMAL_ACCEPTANCE=NOT_GRANTED; REAL_EVIDENCE=NOT_VERIFIED",
    "",
    "## Addendum boundary",
    "R5 reviewed the main body, functional tables 0–21, and only SF-00689–SF-00699. R5 returned PARTIAL because SF-00700–SF-00717 were absent from its packet. This addendum repairs that review-packet omission; it does not claim to repeat R5's broad body review or to inspect Word layout.",
    "",
    "## Official mapping and source anchors",
    json.dumps(mapping, ensure_ascii=False, indent=2),
    "",
    "### Unique Blueprint row",
    blueprint,
    "",
    "### Historical 129-to-142 crosswalk row",
    crosswalk,
    "",
    "## Exact candidate Table 10-3A (DOCX table index 18)",
    *table_lines,
    "",
    "## Complete 29-row fact-register delta (DOCX table index 22)",
    "The rows below were extracted from the exact candidate bytes. The required IDs are SF-00689 through SF-00717 inclusive; all 29 are supplied once, in order.",
    "| fact_id | case_id | fact_type | SAMPLE/SYNTHETIC value | exact locator | REAL source/check key | role |",
    *fact_lines,
    "",
    "## Prior R5 review evidence (included for bounded combined adjudication)",
    json.dumps(r5, ensure_ascii=False, indent=2),
    "",
    "## Independent review request",
    "Independently verify the 29 supplied fact rows against Table 10-3A, the mapping, the unique Blueprint, and the R5 findings. Confirm sequential/unique IDs; source keys and exact table-cell locators; all arithmetic and cumulative carry-forward; Q4 N/A treatment; annual OPEN status; five OPEN status facts; and SAMPLE/REAL boundary. Specifically inspect the previously omitted SF-00700–SF-00717. Decide whether those facts close R5's high-severity packet-scope gap. For R5's four low-severity clarity observations, state whether each is blocking or nonblocking based on the evidence included; do not invent DOCX edits or claim a new visual inspection.",
    "",
    "Return one compact JSON object only with: addendum_verdict PASS/PARTIAL/FAIL; combined_semantic_verdict PASS/PARTIAL/FAIL; exact candidate SHA; reviewed IDs/count; missing/duplicate IDs; table-to-register result; arithmetic result; Q4 N/A and OPEN status result; R5 packet-gap disposition; each prior low finding's blocking/nonblocking disposition; specific locators; remaining limitations; concise rationale. A combined PASS must be supported by R5's exact-SHA broad review plus this exact-SHA delta review, with no material unresolved finding. Do not modify files, use tools, claim visual review, or grant REAL/formal acceptance.",
]

PACKET.write_text("\n".join(parts) + "\n", encoding="utf-8")
print(
    json.dumps(
        {
            "packet": str(PACKET),
            "packet_sha256": sha256(PACKET),
            "candidate_sha256": docx_sha,
            "rows_supplied": len(fact_rows),
            "first_id": actual_ids[0],
            "last_id": actual_ids[-1],
            "table_10_3a_rows": len(table_10_3a.rows),
            "prior_r5_sha256": sha256(R5_JSON),
        },
        ensure_ascii=True,
    )
)
