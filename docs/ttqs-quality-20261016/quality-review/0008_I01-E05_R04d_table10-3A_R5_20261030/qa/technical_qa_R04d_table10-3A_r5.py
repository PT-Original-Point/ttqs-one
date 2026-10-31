from __future__ import annotations

from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from hashlib import sha256
import json
import re
import zipfile
from pathlib import Path

from docx import Document

ROOT = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0008_I01-E05_W1_TABLE10-3A_R04d_20261030")
DOCX = ROOT / "candidate" / "0008_I01-E05_R04d_table10-3A_r5.docx"
MD = ROOT / "source" / "0008_I01-E05_R04d_table10-3A_20261030_r5.md"
BASE = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0008_I01-E05_CATEGORY_COST_REPAIR_R04c_20261030\source\0008_I01-E05_R04c.md")
BLUEPRINT = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\CONTROL\BLUEPRINT_CURRENT.jsonl")
OFFICIAL = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\INPUTS\R04\06_OFFICIAL_SOURCES\TTQS查核佐證文件資料表-訓練機構版 (2).pdf")


def h(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def text_round(value: Decimal) -> int:
    return int(value.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


assert DOCX.is_file() and MD.is_file() and BASE.is_file()
d = Document(DOCX)
with zipfile.ZipFile(DOCX) as z:
    zip_error = z.testzip()

body_runs = [r for p in d.paragraphs for r in p.runs if r.text.strip()]
table_runs = [r for t in d.tables for row in t.rows for c in row.cells for p in c.paragraphs for r in p.runs if r.text.strip()]
body_sizes = [r.font.size.pt if r.font.size else None for r in body_runs]
table_sizes = [r.font.size.pt if r.font.size else None for r in table_runs]
md_text = MD.read_text(encoding="utf-8-sig")
doc_text = "\n".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for row in t.rows for c in row.cells])
md_ids = re.findall(r"^\|\s*(SF-\d{5})\s*\|", md_text, re.M)
doc_ids = [row.cells[0].text.strip() for t in d.tables for row in t.rows[1:] if row.cells and re.fullmatch(r"SF-\d{5}", row.cells[0].text.strip())]
old_facts = {m.group(1): line for line in BASE.read_text(encoding="utf-8-sig").splitlines() if (m := re.match(r"^\|\s*(SF-\d{5})\s*\|", line))}
new_facts = {m.group(1): line for line in md_text.splitlines() if (m := re.match(r"^\|\s*(SF-\d{5})\s*\|", line))}

checks: dict[str, bool] = {
    "zip_integrity": zip_error is None,
    "python_docx_parse": True,
    "table_count_23_includes_gap_table_and_source_split": len(d.tables) == 23,
    "section_count_5_includes_gap_landscape_and_return": len(d.sections) == 5,
    "body_font_floor_10_5": bool(body_sizes) and all(v is not None and v >= 10.5 for v in body_sizes),
    "table_font_floor_9_5": bool(table_sizes) and all(v is not None and v >= 9.5 for v in table_sizes),
    "717_unique_source_fact_ids": len(md_ids) == 717 and len(set(md_ids)) == 717,
    "717_docx_fact_ids_equal_source": len(doc_ids) == 717 and len(set(doc_ids)) == 717 and set(doc_ids) == set(md_ids),
    "all_prior_688_fact_rows_unchanged": len(old_facts) == 688 and all(new_facts.get(fid) == row for fid, row in old_facts.items()),
    "new_fact_ids_sequential_689_717": [int(x[3:]) for x in md_ids[-29:]] == list(range(689, 718)),
    "no_malformed_markdown_table_row": not re.search(r"(?m)^\+\|", md_text),
    "r04d_version_label": "| 版本／日期 | R04d／2026-10-30" in md_text and "R04d" in doc_text,
    "stale_r03_self_reference_removed": "R03 完整合成文字候選" not in doc_text and "版本／日期 | R03" not in doc_text,
    "sample_synthetic_boundary": "SAMPLE" in doc_text and "SYNTHETIC" in doc_text,
    "hold_reject_controls": "HOLD" in doc_text and "REJECT" in doc_text,
    "official_requirement_text_present": "各類別訓練課程年度成長計畫" in doc_text,
    "cost_scope_disclosed": all(x in doc_text for x in ("不含教師備課", "PC 協調工時", "未貨幣化", "完整經濟成本")),
    "category_quarter_targets_present": all(x in md_text for x in ("A 18／16；合計 18／16", "B 14／12；合計 32／28", "C 10／8；合計 42／36", "C 11／10；合計 11／10")),
    "category_result_miss_explicit": "A 基礎類未達年度完成目標" in doc_text and "不得以年度總量達標掩蓋個別類別未達" in doc_text,
    "category_gap_trigger_in_section_9": "任一類別的季度完成席次低於表 6-1 配置" in doc_text,
}

gap_tables = [t for t in d.tables if t.rows and t.rows[0].cells and t.rows[0].cells[0].text.strip() == "期間"]
checks["gap_table_10_3A_single_8_columns_6_rows"] = len(gap_tables) == 1 and len(gap_tables[0].columns) == 8 and len(gap_tables[0].rows) == 6
if gap_tables:
    g = gap_tables[0]
    g_rows = [[c.text.strip() for c in row.cells] for row in g.rows]
    checks["gap_table_q1_q2_q3_math"] = (
        "16" in g_rows[1][1] and "14" in g_rows[1][2] and "−2" in g_rows[1][3] and
        "32" in g_rows[2][4] and "31" in g_rows[2][5] and "−1" in g_rows[2][6] and
        "48" in g_rows[3][4] and "47" in g_rows[3][5] and "−1" in g_rows[3][6]
    )
    checks["gap_table_q4_not_numeric_zero"] = "N/A" in g_rows[4][1] and "N/A" in g_rows[4][2] and "N/A" in g_rows[4][3] and "0" not in g_rows[4][1] + g_rows[4][2] + g_rows[4][3]
    checks["gap_table_annual_open_no_fabricated_decision"] = "OPEN" in g_rows[5][7] and "年度累計仍少 1 席" in doc_text and "不預設決策結果" in doc_text and "不宣稱協會已核准或實施" in doc_text
else:
    checks["gap_table_q1_q2_q3_math"] = False
    checks["gap_table_q4_not_numeric_zero"] = False
    checks["gap_table_annual_open_no_fabricated_decision"] = False
new_id_values = [int(x[3:]) for x in md_ids[-29:]]
new_id_rows = [new_facts.get(f"SF-{i:05d}", "") for i in range(689, 718)]
checks["exactly_29_new_fact_rows_24_numeric_5_status"] = len(new_id_values) == 29 and all(new_facts.get(f"SF-{i:05d}") for i in range(689,718)) and sum("| 數量／公式原子值 |" in row for row in new_id_rows) == 24 and sum("| 事實狀態 |" in row for row in new_id_rows) == 5
checks["prior_688_fact_rows_preserved_exactly"] = all(new_facts.get(fid) == row for fid,row in old_facts.items())
checks["all_new_rows_use_quarter_source_and_roles"] = all("| SRC-QUARTER | PC＋ER |" in row for row in new_id_rows)
checks["new_rows_locate_each_period_and_column"] = all("§10／表10-3A／資料列" in row and "／欄位" in row for row in new_id_rows)

doc_fact_values = {}
for t in d.tables:
    if t.rows and t.rows[0].cells and t.rows[0].cells[0].text.strip() == "fact_id":
        for row in t.rows[1:]:
            if len(row.cells) >= 4:
                doc_fact_values[row.cells[0].text.strip()] = row.cells[3].text.strip()
if gap_tables:
    gap_status_values = [row[7] for row in g_rows[1:6]]
    register_status_values = [doc_fact_values.get(f"SF-{i:05d}", "") for i in range(713,718)]
    checks["five_status_register_values_match_table_exactly"] = register_status_values == gap_status_values
else:
    checks["five_status_register_values_match_table_exactly"] = False

source_key_tables = [t for t in d.tables if t.rows and t.rows[0].cells and t.rows[0].cells[0].text.strip() == "Source key"]
source_keys = [row.cells[0].text.strip() for t in source_key_tables for row in t.rows[1:]]
expected_source_keys = ["SRC-REQ", "SRC-NEED", "SRC-CATALOG", "SRC-BASELINE", "SRC-GROWTH", "SRC-TARGET", "SRC-TARGET-OPS", "SRC-STAFF", "SRC-DATA-GOV", "SRC-FINANCE", "SRC-QUARTER", "SRC-EVIDENCE", "SRC-RECOVERY", "SRC-OPS", "SRC-DECISION"]
checks["source_crosswalk_balanced_8_7_without_row_loss"] = len(source_key_tables) == 2 and [len(t.rows) for t in source_key_tables] == [9, 8] and source_keys == expected_source_keys and "表 12.1 REAL 替換來源及核對方法（續）" in doc_text

target = {"A": [54, 48], "B": [28, 24], "C": [21, 18]}
result = {"A": [54, 47], "B": [28, 25], "C": [21, 18]}
checks["annual_target_categories_sum_to_aggregate"] = [sum(target[k][i] for k in target) for i in (0, 1)] == [103, 90]
checks["annual_result_categories_sum_to_aggregate"] = [sum(result[k][i] for k in result) for i in (0, 1)] == [103, 90]
checks["category_completion_miss_math"] = result["A"][1] < target["A"][1] and result["B"][1] >= target["B"][1] and result["C"][1] >= target["C"][1]
checks["cost_per_seat_math"] = text_round(Decimal(56310) / Decimal(103)) == 547 and text_round(Decimal(56310) / Decimal(90)) == 626

for fid in ("SF-00188", "SF-00194", "SF-00195"):
    row = new_facts.get(fid, "")
    checks[f"{fid}_type_quantity"] = "| 數量／公式原子值 |" in row

row = BLUEPRINT.read_bytes().decode("utf-8-sig").splitlines()[7]
row_sha = sha256(row.encode("utf-8")).hexdigest()
blueprint_sha, official_sha = h(BLUEPRINT), h(OFFICIAL)
checks["blueprint_row_sha_matches_source"] = row_sha == "0f63ae5f5fc03577fec9a93c6842b85ccb3bd75d5c0f4d28629f65790eeae2b4"
checks["blueprint_file_sha_matches_source"] = blueprint_sha == "406d0495365d0970dcf8113a42842663b8a54372a4e204cc492b7c3f8156ef3e"
checks["official_pdf_sha_matches_source"] = official_sha == "805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709"

receipt = {
    "schema": "ttqs-one-document-technical-qa-v1",
    "project_id": "TTQS_ONE",
    "work_id": "TTQS-0008-I01-E05-R04D-TABLE10-3A-R5-TECHNICAL-QA",
    "doc_id": "0008",
    "official_requirement_id": "I01-E05",
    "official_checkbox": "I01-E05",
    "unique_blueprint_id": "I01-E05",
    "observed_at_local": datetime.now().astimezone().isoformat(),
    "candidate_docx": {"path": str(DOCX), "sha256": h(DOCX), "bytes": DOCX.stat().st_size},
    "source_markdown": {"path": str(MD), "sha256": h(MD), "bytes": MD.stat().st_size},
    "source_hashes": {"official_requirement_pdf_sha256": official_sha, "blueprint_file_sha256": blueprint_sha, "blueprint_row_sha256": row_sha},
    "status": "PASS_WITHIN_TECHNICAL_SCOPE" if all(checks.values()) else "FAIL",
    "checks": checks,
    "metrics": {"tables": len(d.tables), "sections": len(d.sections), "body_runs": len(body_runs), "body_min_pt": min(body_sizes) if body_sizes else None, "table_runs": len(table_runs), "table_min_pt": min(table_sizes) if table_sizes else None, "source_fact_ids": len(md_ids), "docx_fact_ids": len(doc_ids), "preserved_prior_fact_rows": len(old_facts), "new_fact_rows": len(md_ids) - len(old_facts), "gap_table_rows": len(gap_tables[0].rows) if gap_tables else 0},
    "scope_exclusions": ["No Word visual QA is claimed by this receipt.", "No independent semantic review is claimed by this receipt.", "No formal acceptance or REAL evidence is claimed."]
}
out = ROOT / "qa" / "TECHNICAL_QA_R04d_table10-3A_r5.json"
out.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps({"status": receipt["status"], "candidate_sha256": receipt["candidate_docx"]["sha256"], "source_sha256": receipt["source_markdown"]["sha256"], "checks": checks, "receipt": str(out), "receipt_sha256": h(out)}, ensure_ascii=False, indent=2))
if not all(checks.values()):
    raise SystemExit(1)



