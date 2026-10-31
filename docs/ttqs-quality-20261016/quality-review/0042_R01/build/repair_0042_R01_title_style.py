from __future__ import annotations

import hashlib
import json
import os
import re
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor


PKG = Path(__file__).resolve().parents[1]
candidate_dir = PKG / "candidate"
candidates = list(candidate_dir.glob("*.docx"))
if len(candidates) != 1:
    raise SystemExit(f"Expected one 0042 R01 candidate; found {len(candidates)}")
candidate = candidates[0]

document = Document(candidate)
for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
    style = document.styles[style_name]
    style.font.color.rgb = RGBColor(0, 0, 0)
    ppr = style._element.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)

for paragraph in document.paragraphs:
    if paragraph.style.name == "Title":
        ppr = paragraph._p.get_or_add_pPr()
        border = ppr.find(qn("w:pBdr"))
        if border is not None:
            ppr.remove(border)
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(0, 0, 0)

# Keep the subtitle and running header within the document style's black-text rule.
if len(document.paragraphs) > 1:
    for run in document.paragraphs[1].runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
for section in document.sections:
    for paragraph in section.header.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(0, 0, 0)

fd, temp_name = tempfile.mkstemp(prefix="0042_R01_style_", suffix=".docx", dir=candidate.parent)
os.close(fd)
temp_path = Path(temp_name)
try:
    document.save(temp_path)
    os.replace(temp_path, candidate)
finally:
    if temp_path.exists():
        temp_path.unlink()

digest = hashlib.sha256(candidate.read_bytes()).hexdigest()
receipt_path = PKG / "evidence" / "technical_integrity_receipt_R01.json"
receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
receipt["candidate_sha256"] = digest
receipt["status"] = "CANDIDATE_PENDING_POST_REPAIR_RENDER_AND_INDEPENDENT_REVIEW"
receipt["layout_repair"] = {
    "scope": "Remove inherited blue bottom rule under Title; set title, subtitle, heading styles, and running header text to black.",
    "reason": "Confirmed rendered layout/style defect; content and tables were not changed.",
    "prior_sha256": "9133d34ec2d2f5a1595856bb43159e0e182332423cf7269fa3b4202bae4330df",
    "repaired_sha256": digest,
    "post_repair_render": "PENDING",
}
receipt_path.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

progress_path = PKG / "PROGRESS.json"
progress = json.loads(progress_path.read_text(encoding="utf-8"))
progress["candidate_sha256"] = digest
progress["layout_status"] = "PENDING_POST_REPAIR_WORD_RENDER_AND_PAGE_REVIEW"
progress["independent_semantic_review"] = "PENDING"
progress["remote_readback"] = "PENDING"
progress.setdefault("layout_repair", {})
progress["layout_repair"].update({
    "defect": "Inherited blue Title-style bottom rule and non-black heading/header colors.",
    "repair": "Removed Title paragraph border and set title, subtitle, headings, and running header to black.",
    "post_repair_sha256": digest,
    "post_repair_visual_review": "PENDING",
})
progress_path.write_text(json.dumps(progress, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

readme_path = PKG / "README.md"
readme = readme_path.read_text(encoding="utf-8")
readme = re.sub(r"候選 DOCX SHA-256：[0-9a-f]{64}", f"候選 DOCX SHA-256：{digest}", readme)
readme = readme.replace(
    "完成狀態：技術封裝檢查 PASS；Word 原生 PDF/全頁視覺檢查與獨立 exact-SHA 語意審查待完成。",
    "完成狀態：技術封裝檢查 PASS；已修正 Title 樣式底線與標題字色，修復後 Word 渲染/全頁視覺檢查及獨立 exact-SHA 語意審查待完成。",
)
readme_path.write_text(readme, encoding="utf-8")

print(json.dumps({"candidate": str(candidate), "sha256": digest, "prior_sha256": receipt["layout_repair"]["prior_sha256"], "status": "RENDER_REQUIRED"}, ensure_ascii=False, indent=2))
