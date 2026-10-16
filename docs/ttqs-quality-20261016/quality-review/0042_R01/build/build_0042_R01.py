from __future__ import annotations

import csv
import hashlib
import json
import re
import subprocess
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader


PKG = Path(__file__).resolve().parents[1]
REPO = PKG.parents[3]
DOCS = REPO / "docs" / "ttqs-quality-20261016"
QR = DOCS / "quality-review"
RUNTIME = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME")
OFFICIAL_PDF = QR / "0041_R04" / "inputs" / "official_evidence_table.pdf"
OFFICIAL_PAGE_IMAGE = QR / "0043_R04" / "inputs" / "official_source_page_03.png"
BLUEPRINT_CURRENT = RUNTIME / "CONTROL" / "BLUEPRINT_CURRENT.jsonl"
CROSSWALK = RUNTIME / "CONTROL" / "TTQS_ONE_129_REQUIREMENT_TO_142_DELIVERABLE_CROSSWALK_R01.csv"
SOURCE_MAP = RUNTIME / "CONTROL" / "SOURCE_REQUIREMENT_MAP_WAVE005_R03.csv"
BLUEPRINT_SEED = RUNTIME / "CONTROL" / "BLUEPRINT_SEED_142.csv"
OUTPUT_DOCX = PKG / "candidate" / "0042__利害關係人參與課程設計與審查執行紀錄_R01.docx"

for folder in (
    PKG / "build",
    PKG / "candidate",
    PKG / "inputs",
    PKG / "evidence",
    PKG / "evidence" / "render",
    PKG / "review",
):
    folder.mkdir(parents=True, exist_ok=True)

if OUTPUT_DOCX.exists():
    raise SystemExit(f"Refusing to overwrite existing candidate: {OUTPUT_DOCX}")


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def git_text(*args: str) -> str:
    return subprocess.check_output(["git", "-C", str(REPO), *args], text=True, encoding="utf-8").strip()


def json_write(path: Path, data: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


if not all(p.exists() for p in (OFFICIAL_PDF, OFFICIAL_PAGE_IMAGE, BLUEPRINT_CURRENT, CROSSWALK, SOURCE_MAP, BLUEPRINT_SEED)):
    raise SystemExit("A required official, Blueprint, crosswalk, or source-map input is missing.")
if sha256_file(OFFICIAL_PDF) != "805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709":
    raise SystemExit("The official checklist PDF SHA does not match the recorded source.")
official_page_text = PdfReader(str(OFFICIAL_PDF)).pages[2].extract_text() or ""
official_page_normalized = re.sub(r"\s+", "", official_page_text)
official_text_normalized = re.sub(
    r"\s+",
    "",
    "執行本項目與作業流程相關文件資料（展現利益關係人參與課程設計與審查，包括出席、發言、意見溝通、核准、認同、支持…等各種形式的參與）。",
)
if official_text_normalized not in official_page_normalized:
    raise SystemExit("The official requirement could not be read from printed page 3.")

official_text = (
    "執行本項目與作業流程相關文件資料（展現利益關係人參與課程設計與審查，"
    "包括出席、 發言、意見溝通、核准、認同、 支持…等各種形式的參與）。"
)

blueprint_rows = []
blueprint_raw_rows = []
for raw in BLUEPRINT_CURRENT.read_text(encoding="utf-8").splitlines():
    try:
        row = json.loads(raw)
    except json.JSONDecodeError:
        continue
    if str(row.get("deliverable_id", "")).zfill(4) in {"0041", "0042", "0043", "0044"}:
        blueprint_rows.append(row)
        blueprint_raw_rows.append(raw)
blueprint_by_id = {str(x.get("deliverable_id", "")).zfill(4): x for x in blueprint_rows}
if "0042" not in blueprint_by_id:
    raise SystemExit("Current BLUEPRINT_CURRENT has no DOCX 0042 row.")
bp = blueprint_by_id["0042"]
bp_raw = next(
    raw for raw in blueprint_raw_rows
    if str(json.loads(raw).get("deliverable_id", "")).zfill(4) == "0042"
)

crosswalk_rows = list(csv.DictReader(CROSSWALK.open(encoding="utf-8-sig", newline="")))
crosswalk_match = next((r for r in crosswalk_rows if r.get("requirement_id") == "I09-E02"), None)
crosswalk_deliverables = str(crosswalk_match.get("deliverable_ids", "")).replace(" ", "").split("|") if crosswalk_match else []
if not crosswalk_match or "0042" not in crosswalk_deliverables:
    raise SystemExit("The 129-to-142 crosswalk does not map I09-E02 to DOCX 0042.")

source_map_rows = list(csv.DictReader(SOURCE_MAP.open(encoding="utf-8-sig", newline="")))
source_map_match = next((r for r in source_map_rows if r.get("doc_id") == "0042"), None)
seed_rows = list(csv.DictReader(BLUEPRINT_SEED.open(encoding="utf-8-sig", newline="")))
seed_match = next((r for r in seed_rows if r.get("deliverable_id") == "0042"), None)
if not seed_match or seed_match.get("requirement_id") != "I09-E02":
    raise SystemExit("The 142-document seed row does not map DOCX 0042 to I09-E02.")

branch = git_text("branch", "--show-current")
head = git_text("rev-parse", "HEAD")
docx_blob = f"docs/ttqs-quality-20261016/quality-review/0042_R01/candidate/{OUTPUT_DOCX.name}"
spec_files = [
    "docs/ttqs-quality-20261016/01_HUMAN_MISSION_AND_ACCEPTANCE.md",
    "docs/ttqs-quality-20261016/02_GENRE_WRITING_MANUAL.md",
    "docs/ttqs-quality-20261016/03_157_CHECKBOX_BLUEPRINTS.md",
    "docs/ttqs-quality-20261016/04_CODEX_GOAL_24X7_EXECUTION.md",
    "docs/ttqs-quality-20261016/05_HOURLY_NATIVE_WAKE_RESUME.md",
    "docs/ttqs-quality-20261016/06_EXTERNAL_AUDIT_AND_CONTINUATION_20261016.md",
]
spec_blob_ids = {p: git_text("rev-parse", f"{head}:{p}") for p in spec_files}

copy2(OFFICIAL_PDF, PKG / "inputs" / "official_checklist_full_source.pdf")
copy2(OFFICIAL_PAGE_IMAGE, PKG / "inputs" / "official_requirement_page_03.png")
(PKG / "inputs" / "unique_blueprint_row_0042.json").write_text(
    json.dumps(bp, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
)
(PKG / "inputs" / "neighbor_blueprint_rows.json").write_text(
    json.dumps(
        {k: blueprint_by_id[k] for k in ("0041", "0043", "0044") if k in blueprint_by_id},
        ensure_ascii=False,
        indent=2,
    ) + "\n",
    encoding="utf-8",
)
(PKG / "inputs" / "unique_blueprint_source_line.jsonl").write_text(bp_raw + "\n", encoding="utf-8")
(PKG / "inputs" / "crosswalk_row_I09-E02.csv").write_text(
    ",".join(crosswalk_match.keys()) + "\n"
    + ",".join('"' + str(v).replace('"', '""') + '"' for v in crosswalk_match.values()) + "\n",
    encoding="utf-8-sig",
)

source_index = {
    "project_id": "TTQS_ONE",
    "doc_id": "0042",
    "requirement_id": "I09-E02",
    "official_checkbox_sequence": "043",
    "source_commit": head,
    "source_branch": branch,
    "source_spec_blob_ids": spec_blob_ids,
    "official_source": {
        "path": "quality-review/0041_R04/inputs/official_evidence_table.pdf",
        "sha256": sha256_file(OFFICIAL_PDF),
        "printed_page": 3,
        "exact_requirement_text": official_text,
    },
    "current_blueprint": {
        "path": "E:/TTQS/TTQS_ONE_CODEX_RUNTIME/CONTROL/BLUEPRINT_CURRENT.jsonl",
        "file_sha256": sha256_file(BLUEPRINT_CURRENT),
        "unique_row_sha256": hashlib.sha256((bp_raw + "\n").encode("utf-8")).hexdigest(),
        "unique_row": bp,
    },
    "crosswalk": {
        "path": "E:/TTQS/TTQS_ONE_CODEX_RUNTIME/CONTROL/TTQS_ONE_129_REQUIREMENT_TO_142_DELIVERABLE_CROSSWALK_R01.csv",
        "file_sha256": sha256_file(CROSSWALK),
        "row": crosswalk_match,
    },
    "142_document_seed": {
        "path": "E:/TTQS/TTQS_ONE_CODEX_RUNTIME/CONTROL/BLUEPRINT_SEED_142.csv",
        "file_sha256": sha256_file(BLUEPRINT_SEED),
        "row": seed_match,
    },
    "legacy_source_map": {
        "path": "E:/TTQS/TTQS_ONE_CODEX_RUNTIME/CONTROL/SOURCE_REQUIREMENT_MAP_WAVE005_R03.csv",
        "file_sha256": sha256_file(SOURCE_MAP),
        "doc_0042_row_present": source_map_match is not None,
        "interpretation": (
            "No direct row for DOCX 0042 in this legacy Wave005 map; mapping is taken from the current unique Blueprint, "
            "the 129-to-142 crosswalk, and the official checklist page 3. No missing row is invented."
        ),
    },
    "mapping_basis": (
        "DOCX 0042 maps to I09-E02 in the current Blueprint and 129-to-142 crosswalk. "
        "The official page 3 item order and adjacent mappings 0041=042 and 0043=044 place this item at checkbox 043."
    ),
    "body_source_policy": "FROM_SOURCE_ONLY; no historical DOCX body used as a writing template.",
}
json_write(PKG / "inputs" / "source_index.json", source_index)
(PKG / "inputs" / "official_requirement_extract.txt").write_text(
    "官方文件：《TTQS 查核佐證文件資料表—訓練機構版》\n"
    "頁次：印刷頁 3\n"
    "文件映射：DOCX 0042 = I09-E02 = 官方方框 043\n"
    "官方要求：\n" + official_text + "\n"
    "辨識邊界：0041 = I09-E01／042；0043 = I09-E03／044；0044 = I09-E04。\n",
    encoding="utf-8",
)


participants = [
    {
        "id": "P01", "role": "受訓者代表（示範角色）", "invite_date": "2026-09-02",
        "delivery": "已送達（示範）", "response": "有效回覆", "channel": "會議出席、口頭發言",
        "consent": "同意示範記錄", "mode": "到場、發言", "validity": "納入",
        "reason": "有可定位的課程版本與具體設計意見", "comment": "希望縮短單向講授，增加安全的情境演練。",
        "disposition": "採納：增加引導練習及情境演練時間", "evidence": "SYN-MIN-01／SYN-ATT-01",
        "formal_approval": "未證明",
    },
    {
        "id": "P02", "role": "講師（示範角色）", "invite_date": "2026-09-02",
        "delivery": "已送達（示範）", "response": "有效回覆", "channel": "會議出席、口頭及書面補充",
        "consent": "同意示範記錄", "mode": "到場、發言、補充意見", "validity": "納入",
        "reason": "意見連結課程目標、教學方式及回饋", "comment": "演練後須保留講師回饋與重試時間，避免只做情境表演。",
        "disposition": "採納：保留回饋修正時段；講師能力條件列入待遴選項目",
        "evidence": "SYN-MIN-01／SYN-WR-02", "formal_approval": "未證明",
    },
    {
        "id": "P03", "role": "訓練行政協力者（示範角色）", "invite_date": "2026-09-02",
        "delivery": "已送達（示範）", "response": "有效回覆", "channel": "會議出席、口頭發言",
        "consent": "同意示範記錄", "mode": "到場、發言", "validity": "納入",
        "reason": "提出教材版本與可讀性控制", "comment": "參與者需要可列印的簡明版，教材應有版本日期與替代格式。",
        "disposition": "採納：規劃可列印純文字版及版本標示",
        "evidence": "SYN-MIN-01／SYN-ATT-01", "formal_approval": "未證明",
    },
    {
        "id": "P04", "role": "需求／服務使用單位代表（示範角色）", "invite_date": "2026-09-02",
        "delivery": "已送達（示範）", "response": "有效回覆", "channel": "書面意見（未寄出的示範稿）",
        "consent": "同意示範記錄", "mode": "非同步書面意見", "validity": "納入",
        "reason": "意見對應工作交接能力及課程評量", "comment": "目標應能看出交接資訊是否完整，不以課後滿意度取代技能檢核。",
        "disposition": "調整：加入交接清單操作及觀察式量尺；不外推為用人單位正式核准",
        "evidence": "SYN-WR-04", "formal_approval": "未證明；授權來源不存在",
    },
    {
        "id": "P05", "role": "課程內容顧問（示範角色）", "invite_date": "2026-09-02",
        "delivery": "已送達（示範）", "response": "有效回覆", "channel": "會議出席、口頭發言",
        "consent": "同意示範記錄", "mode": "到場、發言", "validity": "納入",
        "reason": "指出隱私界線與案例風險", "comment": "案例不得帶入可辨識的服務對象資料；需明確說明可記錄與不可記錄的內容。",
        "disposition": "採納：以完全合成案例示範，增加隱私界線檢核",
        "evidence": "SYN-MIN-01／SYN-ATT-01", "formal_approval": "未證明",
    },
    {
        "id": "P06", "role": "社區服務協力志工（示範角色）", "invite_date": "2026-09-02",
        "delivery": "已送達（示範）", "response": "有效回覆", "channel": "電話摘要（未實際通話）",
        "consent": "同意示範記錄", "mode": "電話溝通摘要", "validity": "納入",
        "reason": "提出演練後需要清楚的回饋方式", "comment": "支持練習方向，但希望每次演練都有講師回饋，且不強迫分享個人經驗。",
        "disposition": "採納條件：安排講師回饋，允許使用者選擇觀察角色",
        "evidence": "SYN-PH-06", "formal_approval": "支持不等於正式核准",
    },
    {
        "id": "P07", "role": "受訓者代表（示範角色）", "invite_date": "2026-09-02",
        "delivery": "已送達（示範）", "response": "有效回覆", "channel": "結構化意見表（未實際發放）",
        "consent": "同意示範記錄", "mode": "非同步書面意見", "validity": "納入",
        "reason": "具體說明參與障礙與可行替代", "comment": "不希望被要求演出個人真實經驗；應可使用虛構情境或改以觀察方式參與。",
        "disposition": "採納：所有演練提供虛構情境與觀察替代，不要求揭露個人經驗",
        "evidence": "SYN-FB-07", "formal_approval": "未證明",
    },
    {
        "id": "P08", "role": "受訓者代表（示範角色）", "invite_date": "2026-09-02",
        "delivery": "已送達（示範）", "response": "撤回記錄同意；排除內容",
        "channel": "撤回狀態通知（僅示範）", "consent": "撤回",
        "mode": "不保留意見內容，只記最小必要狀態", "validity": "排除",
        "reason": "撤回同意後不保存、不分析原意見", "comment": "內容不保留。",
        "disposition": "排除並在示範批次中計為一筆排除回覆",
        "evidence": "SYN-WD-08", "formal_approval": "未證明",
    },
    {
        "id": "P09", "role": "另一位受訓者代表（示範角色）", "invite_date": "2026-09-02",
        "delivery": "已送達（示範）", "response": "截止時未回覆",
        "channel": "邀請管道（合成）", "consent": "未取得",
        "mode": "無參與紀錄", "validity": "不納入回覆分母以外的內容分析",
        "reason": "未回覆不等於反對、支持或缺席會議", "comment": "無意見資料。",
        "disposition": "只列未回覆；不推定立場", "evidence": "SYN-INV-09",
        "formal_approval": "未證明",
    },
    {
        "id": "P10", "role": "服務使用單位代表（示範角色）", "invite_date": "2026-09-02",
        "delivery": "未送達（示範退信）", "response": "無回覆",
        "channel": "邀請管道（合成）", "consent": "未取得",
        "mode": "無參與紀錄", "validity": "不納入回覆分母以外的內容分析",
        "reason": "邀請未送達，不能記為已接觸", "comment": "無意見資料。",
        "disposition": "列為送達失敗；改由正式承辦人核實聯絡方式後再邀請",
        "evidence": "SYN-INV-10", "formal_approval": "未證明",
    },
]

agenda = [
    ("14:00–14:05", "開場、說明紀錄範圍與自願提供意見", "記錄同意範圍；不錄音、不代簽"),
    ("14:05–14:15", "看同一版課程草案與官方要求", "確認討論版本為 SYN-V0.2"),
    ("14:15–14:30", "檢視學習目標與參與條件", "記下需求、不同意見與未決事項"),
    ("14:30–14:50", "比較教學方法與時程", "逐案記錄支持理由、限制及取捨"),
    ("14:50–15:05", "檢視師資能力、教材、環境", "不指定真實人員或場地"),
    ("15:05–15:20", "檢視評量方法與合格門檻", "確認能連回學習目標"),
    ("15:20–15:25", "提出隱私及參與壓力異議", "保留異議，不以沉默推定認同"),
    ("15:25–15:30", "逐項回讀處置、責任與期限", "回讀的是示範摘要，不是正式核准"),
]

design_decisions = [
    {
        "id": "D01", "input": "P01：希望少講授、多練習；P02：需保留回饋及重試",
        "disposition": "採納兩項意見；不以縮短課程總長度解決",
        "change": "講授 55→40 分鐘；引導練習 35→50 分鐘；保留 35 分鐘回饋",
        "owner": "示範課程設計承辦角色；實際人選待核",
        "due": "2026-09-10（合成）",
        "check": "版本表 V0.2／V0.3 時數重算",
        "formal": "示範處置，非正式核准",
    },
    {
        "id": "D02", "input": "P05：案例不得含可辨識資料；P07：不要求分享真實經驗",
        "disposition": "採納；將隱私檢核與參與替代納入",
        "change": "情境演練使用全合成案例，提供觀察角色，不要求個人揭露",
        "owner": "示範教材編製角色；實際人選待核",
        "due": "2026-09-11（合成）",
        "check": "案例卡與參與說明逐項檢查",
        "formal": "示範處置，非正式核准",
    },
    {
        "id": "D03", "input": "P03：需可列印簡明教材及清楚版本日期",
        "disposition": "採納",
        "change": "規劃可列印純文字版；教材頁首顯示版本及更新日期",
        "owner": "示範行政協力角色；實際人選待核",
        "due": "2026-09-11（合成）",
        "check": "列印版與電子版內容核對",
        "formal": "示範處置，非正式核准",
    },
    {
        "id": "D04", "input": "P04：目標要檢查交接資訊完整，不以滿意度代替技能",
        "disposition": "調整",
        "change": "目標加入四步交接清單操作；採觀察式量尺並回連學習目標",
        "owner": "示範課程設計承辦角色；實際人選待核",
        "due": "2026-09-10（合成）",
        "check": "評量題目與四項量尺對照",
        "formal": "示範處置；P04 核准權限未證明",
    },
    {
        "id": "D05", "input": "P06：支持練習，但要求講師回饋；不強迫分享個人經驗",
        "disposition": "有條件採納",
        "change": "每輪演練後由教學者回饋；參與者可選觀察或虛構角色",
        "owner": "示範課程設計承辦角色；實際人選待核",
        "due": "2026-09-11（合成）",
        "check": "教學流程與評量量尺確認回饋時間",
        "formal": "支持不等於正式核准",
    },
    {
        "id": "D06", "input": "P07：反對被要求演出個人真實經驗",
        "disposition": "保留異議並採用替代參與",
        "change": "提供觀察、虛構案例角色及拒絕分享選項",
        "owner": "示範帶領角色；實際人選待核",
        "due": "2026-09-11（合成）",
        "check": "邀請及課前說明列出自願參與界線",
        "formal": "異議未刪除；示範處置",
    },
    {
        "id": "D07", "input": "P08 撤回記錄同意；原意見不可保留",
        "disposition": "排除其意見內容",
        "change": "只保留最小必要撤回狀態及排除計數",
        "owner": "示範資料管理角色；實際人選待核",
        "due": "2026-09-09（合成）",
        "check": "分析表不含其意見、不將其納入有效數",
        "formal": "資料處理示範；非真人紀錄",
    },
    {
        "id": "D08", "input": "無正式授權文件或實際核准原件",
        "disposition": "暫停正式發布；不推定核准",
        "change": "真實使用前取得授權依據及正式核准紀錄",
        "owner": "正式授權人待依章程及授權文件確認",
        "due": "正式實施前；不得以示範日期代替",
        "check": "核章原件、授權範圍、核准版本及檔案位置",
        "formal": "NOT EVIDENCED",
    },
]

plan_v02 = [
    ("開場與學習目標", 15),
    ("溝通步驟講授", 55),
    ("引導練習", 35),
    ("情境演練", 55),
    ("回饋與修正", 35),
    ("能力檢核", 25),
    ("轉用行動", 10),
    ("休息", 10),
]
plan_v03 = [
    ("開場與學習目標", 15),
    ("溝通步驟講授", 40),
    ("引導練習", 50),
    ("情境演練（含隱私界線案例）", 60),
    ("回饋與修正", 35),
    ("能力檢核", 20),
    ("轉用行動", 10),
    ("休息", 10),
]
active_v02 = sum(m for name, m in plan_v02 if name != "休息")
active_v03 = sum(m for name, m in plan_v03 if name != "休息")
total_v02 = sum(m for _, m in plan_v02)
total_v03 = sum(m for _, m in plan_v03)
assert (active_v02, total_v02, active_v03, total_v03) == (230, 240, 230, 240)

invited = len(participants)
delivered = sum(p["delivery"].startswith("已送達") for p in participants)
undelivered = invited - delivered
responses = sum(p["response"] in {"有效回覆", "撤回記錄同意；排除內容"} for p in participants)
valid = sum(p["validity"] == "納入" for p in participants)
excluded = sum(p["validity"] == "排除" for p in participants)
nonresponse_delivered = sum(p["delivery"].startswith("已送達") and p["response"] == "截止時未回覆" for p in participants)
duplicate_returns = 0
assert (invited, delivered, undelivered, responses, valid, excluded, nonresponse_delivered) == (10, 9, 1, 8, 7, 1, 1)
assert invited == delivered + undelivered
assert delivered == responses + nonresponse_delivered
assert responses == valid + excluded
assert invited == valid + excluded + nonresponse_delivered + undelivered
contact_rate = delivered / invited
response_rate_invited = responses / invited
response_rate_delivered = responses / delivered
valid_rate = valid / responses
valid_share_invited = valid / invited
assert round(contact_rate * 100, 1) == 90.0
assert round(response_rate_invited * 100, 1) == 80.0
assert round(response_rate_delivered * 100, 1) == 88.9
assert round(valid_rate * 100, 1) == 87.5
assert round(valid_share_invited * 100, 1) == 70.0


facts: list[dict[str, str]] = []


def add_fact(
    fact_id: str,
    locator: str,
    field: str,
    value: object,
    rationale: str,
    replacement: str,
    source_type: str = "SAMPLE_SYNTHETIC",
) -> None:
    facts.append({
        "fact_id": fact_id,
        "case_id": "SYN-0042-A01",
        "doc_id": "0042",
        "requirement_id": "I09-E02",
        "locator": locator,
        "field": field,
        "synthetic_value": str(value),
        "source_type": source_type,
        "assumption_or_rationale": rationale,
        "cross_document_reference": "無共用案例；僅與 0041/I09-E01/042、0043/I09-E03/044 作工作物邊界對照",
        "required_real_source_for_replacement": replacement,
        "responsible_verifier": "正式承辦人及有權覆核者；身分須依章程與授權文件確認",
        "status": "PENDING_REAL_SOURCE",
    })


add_fact("SYN-C01-F01", "封面／§2", "案例代碼", "SYN-0042-A01", "為展示完整執行追溯流程而建立的虛構案例", "實際課程/案件正式編碼及核准索引")
add_fact("SYN-C01-F02", "封面／§2", "課程名稱", "社區服務溝通與安全界線工作坊", "以小型社區服務訓練情境說明參與設計", "實際核准課程名稱及課程計畫")
add_fact("SYN-C01-F03", "§2／§4", "初稿版本", "SYN-V0.2", "凍結供示範對照的假設版本", "實際送審草案檔、版次與建立時間")
add_fact("SYN-C01-F04", "§2／§4", "修訂版本", "SYN-V0.3", "以參與意見演示如何追溯設計修訂", "實際修訂檔、版本差異與核准記錄")
add_fact("SYN-C01-F05", "§3／表3", "討論日期", "2026-09-09", "示範時間順序", "實際會議通知、出席或經授權的電子紀錄")
add_fact("SYN-C01-F06", "§3／表3", "討論時段", "14:00–15:30（90 分鐘）", "示範逐項紀錄及回讀流程", "實際議程、簽到/登入及會議紀錄")
add_fact("SYN-C01-F07", "§2／§5", "受邀數", invited, "完整批次分母示例", "實際去識別邀請清冊及送達狀態")
add_fact("SYN-C01-F08", "§5／表5", "送達數", delivered, "由十筆示範邀請狀態重算", "實際送達或聯絡原始紀錄")
add_fact("SYN-C01-F09", "§5／表5", "有效回覆數", valid, "依同意、版本及內容完整度判定", "實際回覆、同意與逐筆有效性覆核")
add_fact("SYN-C01-F10", "§5／表5", "排除數", excluded, "撤回同意的示範排除狀態", "實際撤回與資料刪除/排除紀錄")

for p in participants:
    rid = p["id"]
    loc = f"§5 表「十筆完整流向紀錄與納入決定」／{rid}"
    if p["validity"] in {"納入", "排除"}:
        loc += "；表「七筆有效意見與一筆排除狀態（逐筆內容）」／" + rid
    for suffix, field, value, replacement in [
        ("01", "示範個案代碼", rid, "實際去識別代碼及授權保管索引"),
        ("02", "示範利害關係角色", p["role"], "實際角色依組織分工或服務關係核實，不公開個資"),
        ("03", "邀請日期", p["invite_date"], "實際邀請原始資料與時區"),
        ("04", "送達狀態", p["delivery"], "實際送達回條/退信或聯絡紀錄"),
        ("05", "回覆狀態", p["response"], "實際回覆時間與原始紀錄；撤回者依同意規則處理"),
        ("06", "來源管道", p["channel"], "實際會議、書面、電話或電子來源；敏感原件留在授權處"),
        ("07", "記錄同意狀態", p["consent"], "實際同意/撤回原件或合法電子紀錄"),
        ("08", "有效性判定", p["validity"], "覆核者依實際納入/排除規則的判定"),
        ("09", "判定理由", p["reason"], "實際覆核理由及資料品質依據"),
        ("10", "設計意見摘要", p["comment"], "有同意且可公開的原始意見；否則只留受控摘要"),
        ("11", "設計處置示例", p["disposition"], "實際處置決定、負責人、期限及課程版次差異"),
        ("12", "證據代碼", p["evidence"], "正式檔案索引、頁/欄/版本與存取權限"),
        ("13", "正式核准狀態", p["formal_approval"], "核准人的授權依據、簽核原件與核准版本"),
    ]:
        add_fact(
            f"SYN-{rid}-{suffix}", loc, field, value,
            "逐筆合成參與資料，僅展示欄位與判定方法，不代表真人或協會事件",
            replacement,
        )

for index, (time_range, item, control) in enumerate(agenda, start=1):
    add_fact(f"SYN-AG-{index:02d}-F01", f"§3 表「示範會議議程及紀錄控制」／議程{index}", "議程時段", time_range, "示範90分鐘會議時序", "實際核准議程與會議時間紀錄")
    add_fact(f"SYN-AG-{index:02d}-F02", f"§3 表「示範會議議程及紀錄控制」／議程{index}", "議程內容及控制", f"{item}；{control}", "示範逐項討論與證據控制", "實際議程、會議紀錄及決議追蹤")

for index, d in enumerate(design_decisions, start=1):
    for suffix, field, value, replacement in [
        ("01", "參與輸入", d["input"], "同意保存的真實意見原件及來源版本"),
        ("02", "示範處置", d["disposition"], "實際審查紀錄及處置理由"),
        ("03", "課程設計變更", d["change"], "修訂前後課程計畫及差異核對"),
        ("04", "責任角色", d["owner"], "正式指派紀錄、真實授權與承辦人確認"),
        ("05", "期限", d["due"], "實際追蹤日及核准排程"),
        ("06", "驗收方法", d["check"], "真實完成證據與覆核紀錄"),
        ("07", "核准界線", d["formal"], "正式授權、核准人及核准版本原件"),
    ]:
        add_fact(
            f"SYN-{d['id']}-{suffix}", f"§7 表「意見—處置—設計變更—追蹤」／{d['id']}", field, value,
            "示範如何將參與意見連到取捨、課程修訂及查核，不表示真實決議",
            replacement,
        )

for version, plan in (("V0.2", plan_v02), ("V0.3", plan_v03)):
    for index, (unit, minutes) in enumerate(plan, start=1):
        add_fact(
            f"SYN-PLAN-{version}-{index:02d}",
            f"§8 表「課程時間設計版本比較」／{version}／單元{index}",
            "單元與分鐘數",
            f"{unit}：{minutes} 分鐘",
            "示範課程設計版本及合計，不是已排課/已授課事實",
            "實際核准課程表、師資/場地可行性及版本紀錄",
        )

for metric, value, formula in [
    ("接觸率", f"{delivered}/{invited}={contact_rate:.1%}", "送達÷受邀"),
    ("受邀者回覆率", f"{responses}/{invited}={response_rate_invited:.1%}", "回覆÷受邀"),
    ("已送達者回覆率", f"{responses}/{delivered}={response_rate_delivered:.1%}", "回覆÷送達"),
    ("有效率", f"{valid}/{responses}={valid_rate:.1%}", "有效÷回覆"),
    ("受邀者有效占比", f"{valid}/{invited}={valid_share_invited:.1%}", "有效÷受邀"),
]:
    add_fact(
        f"SYN-CALC-{metric}", "§6 表「逐步計算與讀法」／計算列", metric, value,
        f"由本件合成流向數據計算；公式為{formula}",
        "由真實邀請/回收/排除清冊重算並由第二人核對",
        source_type="DERIVED_FROM_SAMPLE_SYNTHETIC",
    )

for idx, criterion in enumerate(
    ["確認需求", "整理重點", "隱私界線", "確認下一步"], start=1
):
    add_fact(
        f"SYN-RUBRIC-{idx:02d}",
        "§8 表「示範觀察量尺」",
        "評量項目",
        criterion,
        "合成課程設計用量尺，非協會核定標準",
        "經授權核定的課程目標、評分規準及版本",
        source_type="SAMPLE_DESIGN",
    )
rubric_anchors = [
    ("確認需求", ["未確認即處置", "只確認部分目的", "確認對象可理解的需求與範圍"]),
    ("整理重點", ["資訊散亂或遺漏", "有摘要但缺一項必要資訊", "按清單清楚整理必要資訊"]),
    ("隱私界線", ["揭露不必要可識別資訊", "察覺風險但需提示才修正", "只使用必要且合宜資訊"]),
    ("確認下一步", ["沒有確認承接或紀錄", "指出下一步但未確認承接", "確認責任、時點與紀錄方式"]),
]
for item, anchors in rubric_anchors:
    for score, anchor in enumerate(anchors):
        add_fact(
            f"SYN-RUBRIC-ANCHOR-{item}-{score}",
            "§8 表「示範觀察量尺（設計用，不代表已施測）」／" + item,
            f"{score}分行為錨點",
            anchor,
            "本文件擬議的合成評分描述，不是協會核定標準或已施測結果",
            "正式核定的評量規準、版本及授權來源",
            source_type="SAMPLE_DESIGN",
        )
for idx, score in enumerate([2, 2, 1, 1], start=1):
    add_fact(
        f"SYN-SCORE-{idx:02d}",
        "§8 表「評分計算示例」/ SYN-LEARNER-01",
        "合成示範分數",
        score,
        "展示量尺如何記分，並非真人學習成果",
        "如真實施測，須取得有同意且可核對的評量原件與覆核",
        source_type="DERIVED_FROM_SAMPLE_SYNTHETIC",
    )
add_fact(
    "SYN-SCORE-TOTAL",
    "§8 表「評分計算示例」/ SYN-LEARNER-01",
    "合成示範總分與判定",
    "2+2+1+1=6/8；僅作門檻演算示例",
    "由四項合成分數加總；不是學員成績或訓練結果",
    "若真實施測，使用正式評量表、評分者及複核紀錄",
    source_type="DERIVED_FROM_SAMPLE_SYNTHETIC",
)
add_fact(
    "SYN-ATTEND-COUNT",
    "§5 表「十筆完整流向紀錄與納入判定」/ P01、P02、P03、P05",
    "示範會議到場筆數",
    "4",
    "由四筆合成參與方式標記為到場的列計算",
    "實際簽到/登入來源及逐筆出席覆核",
    source_type="DERIVED_FROM_SAMPLE_SYNTHETIC",
)
add_fact(
    "SYN-COR-01-F01",
    "§10 表「修正記錄示範」/ SYN-COR-01",
    "合成原始分類",
    "P06初始摘要分類曾把有條件支持標成核准",
    "虛構的錯誤分類演練，用以說明更正軌跡",
    "真實原值、來源、操作者與覆核者；不得把演練當成實際事故",
)
add_fact(
    "SYN-COR-01-F02",
    "§10 表「修正記錄示範」/ SYN-COR-01",
    "合成更正後分類",
    "有條件支持；正式核准未證明",
    "演示支持與核准應分欄",
    "實際意見來源、核准權限依據及更正審核紀錄",
)
add_fact(
    "SYN-COR-01-F03",
    "§10 表「修正記錄示範」/ SYN-COR-01",
    "示範數量影響",
    "有效樣本數不變；需重新檢查核准欄",
    "分類更正不新增或刪除一筆回覆",
    "更正前後的真實流向清冊與重算結果",
    source_type="DERIVED_FROM_SAMPLE_SYNTHETIC",
)
for index, (topic, owner, due, completion, evidence, release) in enumerate([
    ("修訂 V0.3 時程、情境及評量草案", "課程設計承辦角色（未指派真人）", "2026-09-10", "示範完成", "時數核算 230+10=240；SYN-REV-01", "否；仍需真實覆核"),
    ("製作可列印教材及版本標示", "行政協力角色（未指派真人）", "2026-09-11", "示範草稿完成", "SYN-MAT-01；未實際發放", "否；未核准"),
    ("核對授權人與正式核准版本", "正式授權人待依章程/授權確認", "正式使用前", "未完成", "需真實權限依據及簽核原件", "否；NOT EVIDENCED"),
    ("核對場地無障礙及實際學員條件", "實際承辦角色待指定", "開課前", "未執行", "需實地查核、正式招生條件及需求資料", "否；不得推定"),
], start=1):
    add_fact(
        f"SYN-FOLLOWUP-{index:02d}",
        "§9 表「示範後續追蹤表」／項目" + str(index),
        "追蹤內容、責任、期限、狀態與發布判定",
        f"{topic}；{owner}；{due}；{completion}；{evidence}；可正式發布：{release}",
        "合成追蹤示例，未指派真人且不代表已核准或實際完成",
        "真實指派、期限、完成證據、課程版本覆核與正式發布授權",
    )
for code, locator, value in [
    ("SYN-MIN-01", "§3、§5–§8", "合成會議紀錄索引"),
    ("SYN-ATT-01", "§5", "合成出席索引；無簽名"),
    ("SYN-WR-04", "§5", "未寄出/未收受的合成書面意見索引"),
    ("SYN-PH-06", "§5", "未實際通話的合成摘要索引"),
    ("SYN-FB-07", "§5", "未實際發放的合成意見表索引"),
    ("SYN-REV-01", "§9", "合成版本修訂索引"),
    ("SYN-MAT-01", "§9", "未實際發放的合成教材索引"),
]:
    add_fact(
        f"SYN-ART-{code}",
        locator,
        "合成示範來源代碼",
        value,
        "僅是正文內的示範索引，不代表存在真實附件",
        "真實且經授權的原始來源及受控檔案索引",
    )

replacement_path = PKG / "evidence" / "SYNTHETIC_REPLACEMENT_REGISTER.csv"
with replacement_path.open("w", encoding="utf-8-sig", newline="") as f:
    writer = csv.DictWriter(f, fieldnames=list(facts[0].keys()))
    writer.writeheader()
    writer.writerows(facts)


NAVY = "1F4E79"
TEAL = "23666B"
PALE = "EAF2F5"
PALE_ALT = "F5F8FA"
TEXT = "202A33"
MUTED = "5C6873"
RED = "9B1C1C"
WHITE = "FFFFFF"
PAGE_WIDTH_IN = 8.27
CONTENT_WIDTH_IN = 6.95


def set_run_font(run, size: float, bold: bool = False, color: str = TEXT, italic: bool = False) -> None:
    run.font.name = "Microsoft JhengHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=85, start=95, bottom=85, end=95) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_text(cell, text: object, size: float, color: str = TEXT, bold: bool = False) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(str(text))
    set_run_font(run, size, bold=bold, color=color)
    set_cell_margins(cell)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(0.62)
section.bottom_margin = Inches(0.62)
section.left_margin = Inches(0.66)
section.right_margin = Inches(0.66)
section.header_distance = Inches(0.28)
section.footer_distance = Inches(0.28)

normal = doc.styles["Normal"]
normal.font.name = "Microsoft JhengHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(TEXT)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.12
for name, size, color in (("Title", 24, "000000"), ("Heading 1", 15, "000000"), ("Heading 2", 12, "000000"), ("Heading 3", 10.5, "000000")):
    style = doc.styles[name]
    style.font.name = "Microsoft JhengHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    style.font.size = Pt(size)
    style.font.bold = name != "Title"
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(9 if name != "Title" else 0)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.keep_with_next = True

# Word's built-in Title style carries a blue bottom border in the default template.
# Remove it so the document title relies on spacing and typography only.
title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.paragraph_format.space_after = Pt(0)
r = header.add_run("SAMPLE / SYNTHETIC　｜　非原始點協會實際紀錄")
set_run_font(r, 8.5, bold=True, color="000000")
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.space_before = Pt(0)
fr = footer.add_run("候選工作稿　｜　未證明真人參與、正式核准或實際執行　｜　第 ")
set_run_font(fr, 8.5, color=MUTED)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
fr2 = footer.add_run(" 頁")
set_run_font(fr2, 8.5, color=MUTED)

doc.core_properties.title = "利害關係人參與課程設計與審查執行紀錄"
doc.core_properties.subject = "TTQS ONE DOCX 0042／I09-E02／官方方框043"
doc.core_properties.author = "TTQS ONE 候選文件"
doc.core_properties.keywords = "SAMPLE; SYNTHETIC; I09-E02; 043"

table_counter = 0


def add_body(text: str, bold_prefix: str | None = None, color: str = TEXT) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.widow_control = True
    if bold_prefix and text.startswith(bold_prefix):
        a = p.add_run(bold_prefix)
        set_run_font(a, 10.5, bold=True, color=TEAL)
        b = p.add_run(text[len(bold_prefix):])
        set_run_font(b, 10.5, color=color)
    else:
        r = p.add_run(text)
        set_run_font(r, 10.5, color=color)


def add_bullet(text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    set_run_font(r, 10.5)


def add_heading(text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_table(caption: str, headers: list[str], rows: list[list[object]], widths: list[float] | None = None) -> str:
    global table_counter
    table_counter += 1
    key = f"T{table_counter:02d}"
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(f"表 {key}　{caption}")
    set_run_font(run, 9.5, bold=True, color=NAVY)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header_row = table.rows[0]
    repeat_table_header(header_row)
    prevent_row_split(header_row)
    for i, name in enumerate(headers):
        set_cell_text(header_row.cells[i], name, 9.5, color=WHITE, bold=True)
        set_cell_shading(header_row.cells[i], NAVY)
    for ri, row_data in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for ci, value in enumerate(row_data):
            set_cell_text(row.cells[ci], value, 9.5)
            if ri % 2 == 1:
                set_cell_shading(row.cells[ci], PALE_ALT)
    if widths and len(widths) == len(headers):
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    for row in table.rows:
        for idx, width in enumerate(widths or [CONTENT_WIDTH_IN / len(headers)] * len(headers)):
            row.cells[idx].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return key


def add_notice(text: str) -> None:
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell = t.cell(0, 0)
    set_cell_shading(cell, "FCEEEE")
    set_cell_text(cell, text, 9.5, color=RED, bold=True)
    cell.width = Inches(CONTENT_WIDTH_IN)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


# Cover and 30-second orientation
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = p.add_run("利害關係人參與課程設計與審查執行紀錄")
set_run_font(r, 24, bold=False, color="000000")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
r = p.add_run("從邀請、出席、發言與意見溝通，到處置追蹤與核准界線")
set_run_font(r, 13, bold=True, color="000000")
add_notice("完整示範資料　SAMPLE / SYNTHETIC　非協會真實會議、真人參與、正式核准、簽署、信件或實際辦訓紀錄")
add_body("本件供訓練承辦人逐次記錄不同利害關係人如何看過同一課程版本、以何種方式提供意見、意見如何被採納或保留，以及哪些事項仍不能發布。以下所有案例代碼、角色、日期、意見、數字、課程安排及處置均為合成示範。")
add_table(
    "30秒使用導引",
    ["你要完成的工作", "在本件找什麼", "完成後留下什麼"],
    [
        ["安排參與", "§4 的流程、角色及邀請/同意欄位", "能辨認誰收到哪一版、由何管道參與"],
        ["記錄實際意見", "§5–§6 的逐筆流向、會議與意見處置", "出席、發言、書面、電話、支持、異議各有來源"],
        ["修訂課程", "§7–§8 的設計版本差異與追蹤", "每項意見對應採納、調整、暫緩或不採納理由"],
        ["確認能否發布", "§9–§12 的核准、隱私、退件與替換清單", "未核實授權或原件時維持暫停正式發布"],
    ],
    [1.55, 2.55, 2.85],
)
add_table(
    "官方映射與工作物邊界",
    ["識別", "本件對應", "本件證明的工作"],
    [
        ["DOCX 0042", "I09-E02／官方方框 043／官方表第3頁", "執行參與流程及其紀錄；同時記出席、發言、溝通、支持/認同與核准證據狀態"],
        ["相鄰 0041", "I09-E01／042", "課程規劃書如何納入不同利害關係人觀點；不替代本件的實際參與來源紀錄"],
        ["相鄰 0043", "I09-E03／044", "需求單位訪談與統計分析；不把本件意見流向計數冒充需求調查分析"],
        ["相鄰 0044", "I09-E04", "完整討論/會議/E-mail檔案與簽核歸檔；本件示範紀錄不冒充該分件的正式原件"],
    ],
    [1.3, 1.9, 3.75],
)
add_body("使用時點：課程草案有可辨識版次後、正式定稿或發布前。一次只追蹤一個課程案例與一個草案版本。若同時有不同課程或不同版本，另開獨立案件，不混計分母。")
doc.add_page_break()

add_heading("一、用途、範圍與證據界線", 1)
add_body("官方要求是留存與課程設計及審查相關的參與資料，參與方式可以是到場、發言、意見溝通、核准、認同、支持等。此文件的操作重點，是把每個參與行為連到邀請對象、被檢視的課程版次、同意範圍、可核對來源、處置理由與後續修訂。")
add_body("本件不以「有開會」推定每位受邀者都出席，也不以未回覆推定支持或反對。支持、認同與正式核准分開記錄；只有查到有權限者針對明確版本的真實核准原件，才可把核准欄由 NOT EVIDENCED 改為已核實。")
add_bullet("不把參與者的意見當成需求單位訪談統計；該工作物屬 0043／I09-E03。")
add_bullet("不把本件一段示範性討論紀錄當成完整會議與電子郵件原件封包；該工作物屬 0044／I09-E04。")
add_bullet("不以本件取代 0041／I09-E01 課程規劃書。課程計畫的修訂結果可由版本與欄位互相查核，但不複製其正文或案例。")
add_bullet("本件樣本案例與 0041、0043 的樣本互不連結；除非將來有真實共同來源並能核對，不共用 case_id 或捏造跨文件關聯。")
add_notice("本次示範沒有真實原始點協會參與紀錄。正文中的「會議」「回覆」「處置」均為虛構案例，不得填入 REAL 清冊或宣稱本會已核定。")

add_heading("二、示範案例與納入規則", 1)
add_body("示範案件 SYN-0042-A01 設定為一門假設課程「社區服務溝通與安全界線工作坊」。這只是展示如何收集不同觀點並把意見連到課程設計的合成情境，不表示協會目前開設此課、已有這些角色或接受此課程方向。")
add_body("邀請框架採目的性涵蓋不同觀點，不宣稱統計代表性。小型協會實務中可按個案選擇受訓者/學員、講師、行政協力者、需求或服務使用單位代表、部門主管、內容專家等合適角色；實際是否存在、由誰擔任及其核准權限都必須查正式來源。")
add_table(
    "一次批次的納入與排除規則",
    ["判定", "可納入的最低條件", "不符合時的處理"],
    [
        ["納入有效意見", "有案件/課程及草案版次；意見管道與日期可追；同意記錄範圍明確；內容可對應設計或審查議題", "列明缺項；取得補正前不納入分析，不補造原始內容"],
        ["出席或發言", "有真實出席來源；發言依同意範圍記摘要；能區分出席與發言", "簽到不等於發言；不以會議邀請或名冊代替實際出席"],
        ["支持或認同", "保留本人表達的原始依據及所支持的版本/項目", "不能推定全體同意；不能把支持改寫成正式核准"],
        ["正式核准", "核准人身分及授權、核准日期、明確版本、原件位置均可核實", "任一要件缺失即記 NOT EVIDENCED 並暫停正式發布"],
        ["排除/撤回", "記錄必要狀態、原因類別及覆核；保留資料依同意與適用規範處理", "撤回內容不進分析；不因需要分母而保留已撤回的意見正文"],
    ],
    [1.25, 3.1, 2.6],
)

add_heading("三、會議時序與逐項紀錄做法", 1)
add_body("以下 90 分鐘議程是合成執行例。真正使用時，議程需在會前與參與者分享；每個議題記錄所看版本、不同意見、回應、暫定處置與未解事項。主持人逐項回讀摘要，紀錄者請與會者更正可辨認的誤記；不把沒有出聲寫成同意。")
agenda_rows = [[t, item, control] for t, item, control in agenda]
add_table("示範會議議程及紀錄控制", ["時間（合成）", "議題與設計問題", "現場控制及留下的結果"], agenda_rows, [1.05, 2.7, 3.2])
add_table(
    "執行步驟、角色與輸出",
    ["階段", "責任角色（實際人選待核）", "操作與判斷", "輸出/控制點"],
    [
        ["1. 定範圍", "課程設計承辦角色", "指定一個課程、草案版次、欲取得的觀點；連結官方 I09-E02", "範圍表；不得將別件資料併入"],
        ["2. 建框架", "承辦角色＋紀錄者", "選擇必要利害關係群；只收作設計所需的最少資訊", "代碼清冊；實際名單受權限控管"],
        ["3. 邀請", "承辦角色", "說明目的、版本、方式、期限、可拒絕範圍及資料用途", "邀請/送達狀態；未送達不可記接觸成功"],
        ["4. 取得意見", "主持人/承辦角色", "以會議、書面、電話摘要或其他同意方式；分開記出席、發言、支持及核准", "來源代碼、日期、版本及同意狀態"],
        ["5. 檢核品質", "紀錄者＋第二覆核者", "檢查必要欄位、重複、撤回、來源對版次、意見是否能定位", "納入、排除或待補；不補寫來源沒有的內容"],
        ["6. 作處置", "課程設計承辦角色", "逐項列採納、調整、暫緩、不採納及理由，保留異議", "設計差異表、責任、期限、未決事項"],
        ["7. 回讀與修訂", "紀錄者＋提意見者（如同意）", "回覆處置及版本差異；依原同意範圍回讀，不擴張公開", "更正軌跡及修訂版次"],
        ["8. 授權與封存", "正式授權人待查；資料管理者", "查授權後才核准；敏感原件留在合法權限來源", "核准原件/存取索引；缺件則維持 NOT EVIDENCED"],
    ],
    [0.8, 1.65, 2.6, 1.9],
)

add_heading("四、參與紀錄欄位與可直接採用的提示", 1)
add_body("每次執行先凍結課程草案版次，再收意見。下表可複製到協會允許的表單；只收完成追溯所需欄位，姓名、聯絡方式及敏感原件保留在有權限的正式來源，不要複製到公開報告。")
add_table(
    "最小記錄欄位與定義",
    ["欄位", "填寫內容", "檢查問題"],
    [
        ["案件/課程/版本", "案件代碼、課程名稱、草案版本、檢視日期", "提意見的人是否看同一版本？"],
        ["角色群與參與方式", "受訓者、講師、行政、需求/服務單位等；出席、發言、書面、電話、支持、認同或核准", "到場、發言、支持、核准是否分開？"],
        ["來源與同意", "原始紀錄類型、來源代碼、同意記錄/分析/分享範圍", "撤回後是否已停止使用並適當排除？"],
        ["意見及定位", "去識別摘要、對應設計議題、原件所在的受控索引", "摘要是否保留不同意見而未改成共識？"],
        ["判定與處置", "納入/排除/待補、理由、設計變更、不採納原因", "是否能回到具體版本差異？"],
        ["責任與期限", "承辦角色、期限、追蹤結果、覆核人", "真實責任人是否依授權指定？"],
        ["核准欄", "核准人/權限依據/核准日期/明確版次/原件位置，或 NOT EVIDENCED", "只有全部要件可核時才標已核准"],
    ],
    [1.45, 2.8, 2.7],
)
add_table(
    "意見提示表（使用前依課程調整）",
    ["提示", "可選或填答方式", "怎麼避免誤判"],
    [
        ["你檢視的是哪個課程及版本？", "顯示課程名與版次；受訪者回讀確認", "版本未知時先不把意見連到設計決策"],
        ["哪個目標/單元最需調整？為什麼？", "勾選目標/方法/時程/師資/教材/入班條件/環境/評量；補充理由", "不把沒有說明的選項當成原因"],
        ["你建議保留、調整、暫緩或不採納什麼？", "逐項記原意與處置結果", "保留相異意見和未決事項"],
        ["是否同意記錄、分析或分享這項意見？", "分別選擇；可限制用途或撤回", "不同意某用途時不得擴大使用"],
        ["你表達的是支持、認同、反對或正式核准嗎？", "明確選擇並說明對象及版本", "支持/認同不可自動轉成正式核准"],
        ["如有正式核准，授權依據與原件在哪？", "填受控索引及核准版本；無則記 NOT EVIDENCED", "不得用口頭印象或匿名回覆推定權限"],
    ],
    [2.0, 2.45, 2.5],
)

add_heading("五、完整示範批次：邀請、回收與納入判定", 1)
add_notice("下列10筆均為合成角色代碼及狀態。沒有真實邀請、信件、電話、簽到、意見表或服務對象。")
contact_rows = []
for p in participants:
    response_text = p["response"]
    contact_rows.append([
        f"{p['id']}｜{p['role']}",
        f"{p['invite_date']} 邀請；{p['delivery']}；{response_text}",
        f"{p['channel']}；同意：{p['consent']}",
        f"{p['validity']}；{p['reason']}；來源 {p['evidence']}",
    ])
contact_key = add_table(
    "十筆完整流向紀錄與納入決定",
    ["代碼與示範角色", "邀請/送達/回覆", "參與管道與同意", "有效判定及理由"],
    contact_rows,
    [1.45, 1.75, 1.85, 1.9],
)
add_body("逐筆解讀：P01、P02、P03、P05 為示範會議到場者；P04、P07 以非同步書面方式參與；P06 為電話摘要。P08 是一次回覆狀態，但已撤回記錄同意，內容排除。P09 收到邀請未回覆；P10 邀請未送達。未回覆與未送達都不是反對、支持或缺席的證據。")

valid_rows = []
for p in participants:
    if p["validity"] == "納入":
        valid_rows.append([
            f"{p['id']}｜{p['channel']}",
            f"看過版本 SYN-V0.2；同意：{p['consent']}；方式：{p['mode']}",
            p["comment"],
            f"{p['disposition']}；{p['evidence']}",
        ])
    elif p["validity"] == "排除":
        valid_rows.append([
            f"{p['id']}｜撤回狀態",
            "撤回同意；不保留原意見內容",
            "不刊載、不分析原意見；僅計一筆排除狀態",
            "排除；SYN-WD-08",
        ])
responses_key = add_table(
    "七筆有效意見與一筆排除狀態（逐筆內容）",
    ["代碼/來源", "同意/版次/方式", "示範意見摘要", "示範處置與來源索引"],
    valid_rows,
    [1.45, 1.7, 2.1, 1.7],
)
add_body("表內敘述是為本文件撰寫的合成摘要，不是任何人的原話。P08 只示範撤回狀態，不保留被撤回的意見正文；在真實流程中，撤回處理應依事前告知、同意範圍及適用規範記錄。")

add_heading("六、流向計算、限制與資料品質處理", 1)
add_body("本批次的「回覆」定義為收到內容或收到撤回通知；撤回通知可計入回覆流向，但不能計入有效意見。所有比例只描述這個合成邀請批次，不推論原始點協會、課程市場或利害關係人的實際態度。")
flow_rows = [
    ["受邀", str(invited), "10筆合成邀請"],
    ["成功送達", str(delivered), "9筆；P10送達失敗"],
    ["未送達", str(undelivered), "1筆；不得列為已接觸"],
    ["已送達者未回覆", str(nonresponse_delivered), "1筆；P09，不推定立場"],
    ["回覆總數", str(responses), "8筆；包含P08撤回通知"],
    ["有效意見", str(valid), "7筆；含出席發言、書面及電話摘要"],
    ["排除", str(excluded), "1筆；P08撤回記錄同意"],
    ["重複回覆", str(duplicate_returns), "0筆；用案件代碼/來源鍵檢查"],
]
flow_key = add_table(
    "流向數量與分母",
    ["狀態", "數量", "核對說明"],
    flow_rows,
    [1.55, 0.85, 4.55],
)
calc_rows = [
    ["送達率", f"{delivered} ÷ {invited} = {contact_rate:.1%}", "成功送達 ÷ 受邀；本例 9/10=90.0%"],
    ["受邀者回覆率", f"{responses} ÷ {invited} = {response_rate_invited:.1%}", "收到回覆（含撤回通知）÷ 受邀；8/10=80.0%"],
    ["已送達者回覆率", f"{responses} ÷ {delivered} = {response_rate_delivered:.1%}", "回覆 ÷ 成功送達；8/9=88.9%"],
    ["有效率（Blueprint口徑）", f"{valid} ÷ {responses} = {valid_rate:.1%}", "有效意見 ÷ 回覆；7/8=87.5%"],
    ["受邀者有效占比", f"{valid} ÷ {invited} = {valid_share_invited:.1%}", "有效意見 ÷ 受邀；7/10=70.0%"],
]
calc_key = add_table("逐步計算與讀法", ["指標", "代入及結果", "分母及解讀"], calc_rows, [1.65, 2.1, 3.2])
add_body("流向平衡檢查：受邀 10 = 送達 9 + 未送達 1；送達 9 = 回覆 8 + 已送達未回覆 1；回覆 8 = 有效 7 + 排除 1；受邀 10 = 有效 7 + 排除 1 + 未回覆 1 + 未送達 1。各式皆相等。")
add_body("分母為 0 時不得寫 0% 或自行填分母；應寫 N/A，說明原因並停止該比例解讀。若同一人由不同管道重複回覆，先以受控索引確認是否同一案件，再保留合乎同意範圍的版本；未能釐清前標待查，不把兩份當兩人。")
add_body("本批只有七筆有效的示範意見，且以目的性邀請建立；回覆率不能證明所有群體都被適當代表，也不能把同意/支持比例當作正式核准。正式使用需附真實抽樣框、邀請送達來源、排除理由及第二人重算紀錄。")

add_heading("七、意見如何改變課程設計", 1)
add_body("本節是示範決策追溯：先保留意見的不同，再寫出處置與具體課程欄位。V0.2 與 V0.3 都是虛構的計畫版本，不是已核定課程，也不是協會已選定的講師、場地、教材或學員資格。")
decision_rows = [
    [d["id"], d["input"], d["disposition"] + "；" + d["change"], f"{d['owner']}；期限：{d['due']}；核對：{d['check']}；{d['formal']}"]
    for d in design_decisions
]
decision_key = add_table(
    "意見—處置—設計變更—追蹤",
    ["項目", "來源意見及代碼", "示範處置及課程變更", "責任、期限、核對與核准界線"],
    decision_rows,
    [0.62, 1.85, 2.55, 1.93],
)
add_body("未採納或保留的理由必須對提意見者可回讀，並保留未解議題。示範案例中 P07 對被要求演出真實經驗的疑慮沒有被改寫成「全體同意」；處置是提供虛構情境及觀察替代。P04 的建議被納入設計草案，但其正式授權未證明，不能被寫成用人單位核准。")

add_heading("八、課程草案版本差異及評量設計", 1)
add_table(
    "課程時間設計版本比較（分鐘）",
    ["單元", "SYN-V0.2", "SYN-V0.3", "變更理由/來源"],
    [
        [name,
         str(dict(plan_v02)["情境演練"] if name.startswith("情境演練") else dict(plan_v02)[name]),
         str(dict(plan_v03)[name]),
         "依 P01/P02/P05/P07 的示範意見調整；" + ("增加隱私情境" if "情境演練" in name else "維持或重分配時數")]
        for name, _ in plan_v03
    ] + [
        ["教學與評量時間小計", str(active_v02), str(active_v03), "休息時間另計，不列教學時數"],
        ["休息", "10", "10", "示範安排"],
        ["總時數", str(total_v02), str(total_v03), "兩版皆為4小時；未驗證場地、師資及實際可行性"],
    ],
    [1.7, 1.1, 1.1, 3.05],
)
add_body("示範可測目標：在虛構情境中，學習者能依「確認需求—整理重點—守住隱私界線—確認下一步」完成交接，並留下適當紀錄。這是供課程設計討論的目標草案，不是協會既定職能標準。")
add_table(
    "示範觀察量尺（設計用，不代表已施測）",
    ["可觀察項目", "0分", "1分", "2分"],
    [
        ["確認需求", "未確認即處置", "只確認部分目的", "確認對象可理解的需求與範圍"],
        ["整理重點", "資訊散亂或遺漏", "有摘要但缺一項必要資訊", "按清單清楚整理必要資訊"],
        ["隱私界線", "揭露不必要可識別資訊", "察覺風險但需提示才修正", "只使用必要且合宜資訊"],
        ["確認下一步", "沒有確認承接或紀錄", "指出下一步但未確認承接", "確認責任、時點與紀錄方式"],
    ],
    [1.55, 1.75, 1.8, 1.85],
)
add_body("示範門檻：4項各0–2分，滿分8分；總分至少6分且隱私界線項不得出現關鍵錯誤，才列達到本例的學習檢核門檻。若總分高於6但有關鍵隱私錯誤，仍判未達並進行補強。這是本案例的擬議量尺，真實門檻須由有權者核定。")
add_table(
    "評分計算示例（合成作答，不是學員成績）",
    ["項目", "確認需求", "整理重點", "隱私界線", "下一步", "總分/判定"],
    [["SYN-LEARNER-01", "2", "2", "1", "1", "6/8；本例達總分門檻，仍需確認無關鍵隱私錯誤"]],
    [1.4, 1.05, 1.05, 1.05, 1.05, 1.35],
)

add_heading("九、追蹤、回讀與正式核准狀態", 1)
add_table(
    "示範後續追蹤表",
    ["追蹤事項", "示範責任角色", "期限/狀態", "完成判定與證據", "是否可正式發布"],
    [
        ["修訂 V0.3 時程、情境及評量草案", "課程設計承辦角色（未指派真人）", "2026-09-10；示範完成", "時數核算 230+10=240；SYN-REV-01", "否；仍需真實覆核"],
        ["製作可列印教材及版本標示", "行政協力角色（未指派真人）", "2026-09-11；示範草稿完成", "SYN-MAT-01；未實際發放", "否；未核准"],
        ["核對授權人與正式核准版本", "正式授權人待依章程/授權確認", "正式使用前；未完成", "需真實權限依據及簽核原件", "否；NOT EVIDENCED"],
        ["核對場地無障礙及實際學員條件", "實際承辦角色待指定", "開課前；未執行", "需實地查核、正式招生條件及需求資料", "否；不得推定"],
    ],
    [1.6, 1.55, 1.25, 1.65, 0.9],
)
add_notice("正式核准狀態：NOT EVIDENCED。合成支持、認同、示範處置或版本差異都不是協會核准。取得有權者對明確課程版本的真實核准文件前，不得標示已核定或對外發布。")

add_heading("十、資料同意、保管、偏差與修正", 1)
add_table(
    "資料品質與隱私控制",
    ["情況", "當下處理", "留下的最小紀錄"],
    [
        ["未同意某種使用", "只做同意範圍內處理；限制分享；不得把意見轉作其他文件", "同意範圍、限制、停止使用位置"],
        ["撤回同意", "停止未授權用途；依告知及規範刪除/排除；重新核算分母", "撤回日期、已採動作與覆核；避免保留意見正文"],
        ["重複回覆", "用受控索引去重；不能判斷則標待查", "原始列、合併理由、保留版本及覆核者"],
        ["版次不明", "不把意見套到任何課程版次；請求確認", "原始回覆、缺失欄及補正狀態"],
        ["更正紀錄", "不可覆蓋原值；記原值/新值/理由/日期/操作者/覆核者", "更正軌跡及受影響計算重算結果"],
        ["個資或敏感內容", "移出公開副本；縮限存取；只保留必要去識別摘要", "權限索引與刪除/遮蔽核對，不公開原件"],
        ["核准權限不明", "保持 NOT EVIDENCED，暫停正式發布", "待補的授權來源、版本與責任角色"],
    ],
    [1.35, 3.15, 2.55],
)
add_table(
    "修正記錄示範（純粹展示控制步驟）",
    ["更正代碼", "原記錄", "更正後", "原因與影響", "核對狀態"],
    [[
        "SYN-COR-01",
        "P06初始摘要分類曾把「有條件支持」標成「核准」",
        "更正為「有條件支持；正式核准未證明」",
        "將支持與授權分欄；不改有效樣本數，重新檢查核准表",
        "示範完成；非真實更正紀錄",
    ]],
    [0.95, 1.5, 1.55, 1.9, 1.15],
)
add_body("此更正示例為本文件撰寫的合成錯誤演練，不表示曾有真人或協會紀錄被誤填。真實更正必須保留原值並依有權覆核程序確認，不能只改摘要。")
add_table(
    "資料存取及保存設定",
    ["資料層", "建議控制", "正式使用前須確認"],
    [
        ["姓名/聯絡方式/簽到原件", "僅存於有權限的正式來源；公開副本以代碼替代", "誰可讀寫、合法目的、保存期限與刪除責任"],
        ["意見原件及電子郵件", "依同意範圍及正式制度保存；不複製到公開儲存庫", "原始寄件/收件人、授權分享及存取紀錄"],
        ["去識別分析表", "只留必要欄位，限制下載及轉寄", "重識別風險、版本、覆核及保存規則"],
        ["本公開示範包", "只包含合成案例、公開官方來源與去識別QA", "提交前檢查不得混入真實個資或受保護原件"],
    ],
    [1.55, 2.8, 2.6],
)

add_heading("十一、正式使用前的真實資料替換清單", 1)
add_body("使用前先另存一份正式工作副本，保留本合成版本作教學參考。實際參與人數、日期、角色、出席、意見、核准及附件位置均須從原始來源逐項取得；不因本件已有填入示範值而複製成 REAL。")
add_table(
    "必須逐件取得的原始來源",
    ["待替換內容", "可接受的真實來源例", "核對/保留方式"],
    [
        ["正式課程名稱與草案版本", "有權承辦人保存的課程計畫、版本紀錄", "記檔名、版次、日期及受控位置"],
        ["受邀對象/角色與送達", "依法可用的邀請清冊、寄送/簽收/退信紀錄", "私密聯絡資料留在正式系統，公開副本去識別"],
        ["出席與發言", "經同意的簽到/登入、會議紀錄或其他可核來源", "分開標出席與發言；簽名不得仿製"],
        ["書面、電話或其他溝通", "原始意見、同意範圍及經核對的紀錄摘要", "以來源代碼定位；不得製造郵件或通話"],
        ["納入、排除與撤回", "真實回收清冊、同意/撤回及覆核記錄", "保留判定理由；排除/去重後重算"],
        ["設計處置及版本差異", "會議處置、課程計畫前後版、回覆及追蹤紀錄", "一項意見對應一個可核對變更或理由"],
        ["正式核准", "核准權限依據及有權者對明確版本的正式簽核原件", "無核准原件即 NOT EVIDENCED，不發布"],
        ["責任與期限", "實際指派及追蹤資料", "不得照抄示範代碼/日期；簽核者依章程確認"],
    ],
    [1.45, 2.9, 2.6],
)
add_body("若真實來源含會員、學員、服務對象、電話、電子郵件、簽名或受保護文件，原件留在經授權的正式來源；公開稽核包只保留可公開的去識別摘要、不可逆摘要值及受控來源索引。摘要不等於已公開驗證原件。")

add_heading("十二、退件/暫停條件與證據索引", 1)
add_table(
    "遇到下列情形即退回補正或暫停",
    ["缺陷", "處置"],
    [
        ["課程或被審版本不清，意見無法連到版本", "標待補；不將意見計入某一版的設計決策"],
        ["只有邀請/出席名單，無實際出席或意見來源", "不得宣稱已參與；補取得可核原始紀錄"],
        ["回覆分母、撤回、排除或重複未交代，公式不可重算", "退回資料清理；重算全部受影響比例"],
        ["支持或認同被寫成正式核准，或核准權限/版本不明", "改為 NOT EVIDENCED；暫停正式發布"],
        ["摘要把異議改成共識，或未回覆被判為支持/反對", "退回逐筆核對，恢復原差異與限制"],
        ["個資未同意、來源不明、公開副本含敏感原件", "停止分享；依授權範圍處理並記錄修正"],
        ["決議無負責角色、期限、設計變更或不採納理由", "補完整處置追蹤後再結案"],
    ],
    [3.0, 3.95],
)
add_table(
    "本示範包的來源與核對索引",
    ["代碼", "本件位置/用途", "真實狀態"],
    [
        ["OFF-03", "官方資料表印刷頁3；I09-E02參與紀錄文字", "公開官方來源；未證明協會事實"],
        ["BP-0042", "唯一 Blueprint；指定 execution_record 及欄位/計算", "需求拆解依據，不是官方核准"],
        ["SYN-MIN-01", "§3議程與§5–§8示範處置", "合成會議摘要，非真實紀錄"],
        ["SYN-ATT-01", "§5示範出席代碼", "無真人姓名/簽章；非真實簽到"],
        ["SYN-WR-04", "§5服務單位書面意見示例", "未寄出/未收受之合成示例"],
        ["SYN-PH-06", "§5電話摘要示例", "未實際通話之合成示例"],
        ["SYN-FB-07", "§5結構化意見表示例", "未實際發放之合成示例"],
        ["SYN-REV-01", "§7–§9設計版次及追蹤", "示範修訂，不是正式課程核定"],
        ["REAL-SOURCE", "授權正式來源位置由協會依法指定", "目前 UNKNOWN / NOT EVIDENCED"],
    ],
    [1.2, 3.45, 2.3],
)
add_body("開始正式執行時，紀錄者在受控工作副本填入真實日期、可核角色/代碼、課程版本、參與方式、同意狀態及證據索引；第二位覆核者重算邀請/送達/回覆/有效/排除流向。缺少核准原件或授權依據時，保留未核准狀態，不補作紀錄。")

add_heading("附錄A｜逐件覆核清單", 1)
for line in [
    "□ 每位提意見者都看過可辨識的同一課程版本，並有來源可核對。",
    "□ 出席、發言、書面/電話溝通、支持/認同及正式核准分欄記錄。",
    "□ 邀請、送達、回覆、有效、排除、未回覆與未送達數量可逐筆重算。",
    "□ 撤回同意者的內容未被保留或分析；未回覆不被推定為任何立場。",
    "□ 每項意見都有處置、理由、責任、期限及課程版本差異或暫緩理由。",
    "□ 反對/保留意見仍可辨認；支持不被擴寫成共識或核准。",
    "□ 正式核准有可核身份、授權、日期、版本及原件位置。",
    "□ 公開副本不含未授權個資、真實簽名或受保護原件。",
    "□ 0041、0043、0044只作工作物邊界引用，未混成同一份案例證據。",
]:
    add_body(line)

add_body("文件狀態：R01 候選工作稿；本件完整案例均為 SAMPLE/SYNTHETIC。正式協會實際參與、核准及 TTQS 接受均未證明。")

doc.save(OUTPUT_DOCX)

# Technical checks are bounded to this exact output and its registered sources.
with zipfile.ZipFile(OUTPUT_DOCX) as zf:
    bad_member = zf.testzip()
    members = zf.namelist()
if bad_member:
    raise SystemExit(f"DOCX package integrity failed at {bad_member}")

docx_sha = sha256_file(OUTPUT_DOCX)
body_sizes = [r.font.size.pt for p in doc.paragraphs for r in p.runs if r.text.strip() and r.font.size]
table_sizes = [
    r.font.size.pt
    for table in doc.tables
    for row in table.rows
    for cell in row.cells
    for paragraph in cell.paragraphs
    for r in paragraph.runs
    if r.text.strip() and r.font.size
]
technical = {
    "project_id": "TTQS_ONE",
    "doc_id": "0042",
    "requirement_id": "I09-E02",
    "official_checkbox_sequence": "043",
    "candidate_path": str(OUTPUT_DOCX),
    "candidate_sha256": docx_sha,
    "source_commit": head,
    "source_branch": branch,
    "zip_integrity": "PASS",
    "zip_member_count": len(members),
    "paragraph_count": len(doc.paragraphs),
    "table_count": len(doc.tables),
    "replacement_register_rows": len(facts),
    "body_font_min_pt": min(body_sizes) if body_sizes else None,
    "table_font_min_pt": min(table_sizes) if table_sizes else None,
    "arithmetic": {
        "invited": invited,
        "delivered": delivered,
        "undelivered": undelivered,
        "responses": responses,
        "valid": valid,
        "excluded": excluded,
        "nonresponse_delivered": nonresponse_delivered,
        "duplicate_returns": duplicate_returns,
        "checks": {
            "invited_equals_delivered_plus_undelivered": invited == delivered + undelivered,
            "delivered_equals_responses_plus_nonresponse": delivered == responses + nonresponse_delivered,
            "responses_equals_valid_plus_excluded": responses == valid + excluded,
            "invited_flow_balances": invited == valid + excluded + nonresponse_delivered + undelivered,
            "course_v02_minutes": total_v02,
            "course_v03_minutes": total_v03,
            "course_v02_instruction_minutes": active_v02,
            "course_v03_instruction_minutes": active_v03,
            "rubric_max": 8,
            "sample_score": 6,
        },
    },
    "status": "CANDIDATE_PENDING_RENDER_AND_INDEPENDENT_REVIEW",
    "formal_acceptance": "NOT_GRANTED",
    "created_at_utc": datetime.now(timezone.utc).isoformat(),
}
json_write(PKG / "evidence" / "technical_integrity_receipt_R01.json", technical)

progress = {
    "project_id": "TTQS_ONE",
    "doc_id": "0042",
    "revision": "R01",
    "requirement_id": "I09-E02",
    "official_checkbox_sequence": "043",
    "candidate_sha256": docx_sha,
    "candidate_state": "CANDIDATE_NOT_PROMOTED",
    "genre": "execution_record",
    "sample_scope": "SAMPLE_SYNTHETIC_ONLY",
    "technical_status": "PASS",
    "layout_status": "PENDING_WORD_RENDER_AND_PAGE_REVIEW",
    "independent_semantic_review": "PENDING",
    "formal_acceptance": "NOT_GRANTED",
    "source_head": head,
    "source_index_sha256": sha256_file(PKG / "inputs" / "source_index.json"),
    "replacement_register_sha256": sha256_file(replacement_path),
    "replacement_register_rows": len(facts),
    "remote_readback": "PENDING",
    "notes": [
        "Freshly authored from official page 3, current unique Blueprint row and 129-to-142 crosswalk.",
        "No historical DOCX body was used as a writing template.",
        "No real association approval, attendance, participant, email, signature or training fact is asserted.",
    ],
}
json_write(PKG / "PROGRESS.json", progress)

readme = """# DOCX 0042 R01 — 利害關係人參與課程設計與審查執行紀錄

證據狀態：候選／SAMPLE-SYNTHETIC only；尚未獨立驗收或正式接受。

映射：DOCX 0042 = I09-E02 = 官方方框 043。官方來源是訓練機構版查核佐證文件第3頁。相鄰件 0041/I09-E01/042、0043/I09-E03/044、0044/I09-E04 為不同工作物。

本文從邀請、送達、出席、發言、書面/電話溝通、同意、有效性、排除、意見處置、設計版本差異、追蹤與核准界線形成完整合成案例。所有角色、日期、意見、數字、課程版本、處置與結果均非原始點協會真實事件。沒有真人參與、正式核准、簽名、信件或第三方原件。

寫作只使用官方來源、目前唯一 Blueprint、129-to-142 crosswalk 及公開品質契約；未用任何歷史 DOCX 正文作模板。逐事實替換清冊位於 evidence/SYNTHETIC_REPLACEMENT_REGISTER.csv。

候選 DOCX SHA-256：""" + docx_sha + """
來源 branch/head：""" + branch + " / " + head + """

完成狀態：技術封裝檢查 PASS；Word 原生 PDF/全頁視覺檢查與獨立 exact-SHA 語意審查待完成。不得據此宣稱正式接受或真實 TTQS 證據。
"""
(PKG / "README.md").write_text(readme, encoding="utf-8")

print(json.dumps({
    "candidate": str(OUTPUT_DOCX),
    "sha256": docx_sha,
    "paragraphs": len(doc.paragraphs),
    "tables": len(doc.tables),
    "replacement_register_rows": len(facts),
    "body_font_min_pt": min(body_sizes) if body_sizes else None,
    "table_font_min_pt": min(table_sizes) if table_sizes else None,
    "source_commit": head,
    "crosswalk_row": crosswalk_match,
    "source_map_row_present": source_map_match is not None,
    "status": "CANDIDATE_PENDING_LAYOUT_AND_INDEPENDENT_REVIEW",
}, ensure_ascii=False, indent=2))
