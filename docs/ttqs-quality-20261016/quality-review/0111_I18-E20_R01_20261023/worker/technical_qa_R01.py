from __future__ import annotations

import hashlib
import json
import sys
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "candidate" / "0111_I18-E20_特殊績效及社會面成果證據登錄表_R01.docx"
MANUSCRIPT = ROOT / "manuscript" / "0111_I18-E20_特殊績效及社會面成果證據登錄表_R01.md"
OUT = ROOT / "evidence" / "technical_qa_0111_R01.json"
EXPECTED_SHA = "31fc783e6021a6f19ae2fa316a0b8bb28b0cf058508da46bcb312ea7787090fd"
EXPECTED_ROWS = [7, 8, 3, 5, 11, 7, 9, 9, 11, 3, 8, 8, 6, 27, 8, 6]
EXPECTED_COLS = [3, 5, 8, 5, 7, 7, 8, 7, 3, 7, 6, 5, 6, 6, 5, 5]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def direct_size_pt(run) -> float | None:
    if run.font.size is not None:
        return float(run.font.size.pt)
    rpr = run._element.rPr
    if rpr is not None and rpr.sz is not None and rpr.sz.val is not None:
        return float(rpr.sz.val) / 2
    return None


def main() -> int:
    actual_sha = sha256(DOCX)
    source_sha = sha256(MANUSCRIPT)
    checks: dict[str, bool] = {}
    zip_error = None
    with zipfile.ZipFile(DOCX) as zf:
        bad_member = zf.testzip()
        required = {"[Content_Types].xml", "word/document.xml", "word/styles.xml"}
        checks["docx_zip_integrity"] = bad_member is None
        checks["required_ooxml_parts"] = required.issubset(set(zf.namelist()))
        if bad_member:
            zip_error = bad_member

    doc = Document(DOCX)
    rows = [len(t.rows) for t in doc.tables]
    cols = [len(t.columns) for t in doc.tables]
    table_sizes = [
        direct_size_pt(run)
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
        for para in cell.paragraphs
        for run in para.runs
        if run.text.strip()
    ]
    unspecified_table_runs = sum(v is None for v in table_sizes)
    explicit_table_sizes = [v for v in table_sizes if v is not None]
    min_table_size = min(explicit_table_sizes) if explicit_table_sizes else None
    body_sizes = [
        direct_size_pt(run)
        for para in doc.paragraphs
        for run in para.runs
        if run.text.strip()
    ]
    unspecified_body_runs = sum(v is None for v in body_sizes)

    all_text = "\n".join(
        [p.text for p in doc.paragraphs]
        + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    )
    required_phrases = {
        "sample_boundary": "SAMPLE / SYNTHETIC",
        "real_zero_hold": "REAL 原件為 0/2",
        "e0_hold": "E0／HOLD",
        "denominator_math": "6/8",
        "baseline_math": "3/8",
        "official_requirement": "I18-E20",
    }

    checks["candidate_exact_sha"] = actual_sha == EXPECTED_SHA
    checks["manuscript_present"] = MANUSCRIPT.is_file()
    checks["expected_table_count"] = len(doc.tables) == 16
    checks["expected_table_rows"] = rows == EXPECTED_ROWS
    checks["expected_table_columns"] = cols == EXPECTED_COLS
    checks["table_text_min_9_5pt"] = min_table_size is not None and min_table_size >= 9.5 and unspecified_table_runs == 0
    checks["sample_hold_and_math_text"] = all(v in all_text for v in required_phrases.values())
    checks["paragraphs_nonempty"] = len(doc.paragraphs) >= 100
    checks["sections_present"] = len(doc.sections) >= 6

    receipt = {
        "receipt_id": "TTQS-0111-TECHNICAL-QA-R01",
        "observed_at_utc": datetime.now(timezone.utc).isoformat(),
        "doc_id": "0111",
        "requirement_id": "I18-E20",
        "official_checkbox_serial": "UNKNOWN_NOT_ASSERTED",
        "candidate_docx": str(DOCX),
        "candidate_docx_sha256": actual_sha,
        "candidate_docx_bytes": DOCX.stat().st_size,
        "source_manuscript": str(MANUSCRIPT),
        "source_manuscript_sha256": source_sha,
        "checks": checks,
        "ooxml_bad_member": zip_error,
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "table_row_counts": rows,
        "table_column_counts": cols,
        "table_font_direct_min_pt": min_table_size,
        "table_font_runs_without_direct_size": unspecified_table_runs,
        "body_font_runs_without_direct_size": unspecified_body_runs,
        "required_content_markers": required_phrases,
        "result": "PASS_TECHNICAL" if all(checks.values()) else "FAIL_TECHNICAL",
        "scope_limit": "Deterministic OOXML/package, structure, direct table-font and marker checks only; not independent semantic review or visual inspection.",
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(receipt, ensure_ascii=False, indent=2))
    return 0 if all(checks.values()) else 2


if __name__ == "__main__":
    sys.exit(main())
