from __future__ import annotations

import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "candidate" / "0111_I18-E20_特殊績效及社會面成果證據登錄表_R01.docx"
SOURCE_PACKET = ROOT / "inputs" / "source_packet_0111_R01.md"
BLUEPRINT = ROOT / "inputs" / "UNIQUE_BLUEPRINT_0111_R01.json"
OFFICIAL_TEXT = ROOT / "source" / "official_source_page7_text.txt"
MANUSCRIPT = ROOT / "manuscript" / "0111_I18-E20_特殊績效及社會面成果證據登錄表_R01.md"
TECH_QA = ROOT / "evidence" / "technical_qa_0111_R01.json"
VISUAL_QA = ROOT / "evidence" / "visual_qa_receipt_0111_R01.json"
OUT = ROOT / "review" / "semantic_review_packet_0111_R01.md"
MANIFEST = ROOT / "review" / "semantic_review_input_manifest_0111_R01.json"
EXPECTED_DOCX_SHA = "31fc783e6021a6f19ae2fa316a0b8bb28b0cf058508da46bcb312ea7787090fd"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> int:
    candidate_sha = sha256(DOCX)
    if candidate_sha != EXPECTED_DOCX_SHA:
        raise SystemExit(f"CANDIDATE_SHA_CHANGED expected={EXPECTED_DOCX_SHA} actual={candidate_sha}")
    doc = Document(DOCX)
    paras = [p.text for p in doc.paragraphs]
    tables = []
    for i, table in enumerate(doc.tables, 1):
        rows = []
        for ri, row in enumerate(table.rows, 1):
            rows.append(" | ".join(cell.text.replace("\n", " / ").strip() for cell in row.cells))
        tables.append(f"## Table T{i:02d} ({len(table.rows)} rows × {len(table.columns)} columns)\n" + "\n".join(f"T{i:02d}-R{ri:02d}: {row}" for ri, row in enumerate(rows, 1)))

    inputs = {
        "candidate_docx": DOCX,
        "source_packet": SOURCE_PACKET,
        "unique_blueprint": BLUEPRINT,
        "official_page7_text": OFFICIAL_TEXT,
        "source_manuscript": MANUSCRIPT,
        "technical_qa_receipt": TECH_QA,
        "visual_qa_receipt": VISUAL_QA,
    }
    hashes = {name: {"path": str(path), "sha256": sha256(path), "bytes": path.stat().st_size} for name, path in inputs.items()}
    extracted = "\n".join(f"P{i:03d}: {text}" for i, text in enumerate(paras, 1) if text.strip())
    body = [
        "# Independent semantic review packet — 0111 / I18-E20 / R01",
        "",
        "This packet is a read-only review request. Do not edit files, invoke Computer Use, generate a replacement manuscript, or claim formal association acceptance.",
        "Review the actual content extracted from the candidate DOCX and the authoritative requirement/Unique Blueprint below. Bind your verdict to the candidate DOCX SHA exactly; if you cannot independently read its bytes, say so explicitly and scope your review to the supplied extraction.",
        "",
        "## Candidate identity and evidence inputs",
        f"- Candidate DOCX SHA-256: `{candidate_sha}`",
        f"- DOCX bytes: {DOCX.stat().st_size}",
        f"- Requirement: I18-E20; document ID 0111; official checkbox serial intentionally UNKNOWN_NOT_ASSERTED.",
        f"- Extraction: {len(paras)} paragraphs and {len(tables)} tables, extracted directly from the exact DOCX above using python-docx.",
        "- Technical and visual receipts are included as bounded scope evidence; their contents do not predetermine your semantic verdict.",
        "",
        "## Review questions",
        "1. Does the artifact have the correct I18-E20 social/other special-outcome register/report genre and directly answer the exact official checkbox?",
        "2. Is the unique I18-E20 Blueprint followed, including required scope, population/denominators, field dictionary, formulas, filled case, group comparison, limitations, evidence strength, action trace, exception/hold workflow, roles/deadlines, replacement register and version controls?",
        "3. Are customer and learner samples distinct, internally complete, mathematically consistent, and clearly descriptive rather than causal? Recompute the 3/8→6/8 change, 6/8 learner application, original acquisition 0/2 and all relevant denominators from displayed rows.",
        "4. Is SAMPLE/SYNTHETIC distinct from REAL throughout; are all third-party originals, consent, approvals, real outcomes and formal acceptance withheld where unsupported?",
        "5. Is the replacement register sufficiently fact-level, with usable locators/replacement evidence; identify any material uncovered synthetic facts or ambiguous consent/control claims.",
        "6. Does title-blind reading distinguish I18-E20 from 0110/I18-E19 survey/action-plan work, 0112/I18-E21 licensing/return-training work, and financial performance reporting?",
        "7. Could association training staff and an evaluator actually use this complete artifact? Identify missing fields, unusable procedures, contradictions or unsupported claims with paragraph/table-row anchors.",
        "8. Return one verdict: PASS, PARTIAL, or FAIL. List exact blocking and nonblocking findings, anchors, calculation checks, review scope/exclusions, and whether the SHA was independently rehashed. Never turn the visual QA receipt into a semantic finding or claim formal acceptance.",
        "",
        "## Exact official requirement and source packet",
        SOURCE_PACKET.read_text(encoding="utf-8"),
        "",
        "### Official page-7 extracted text",
        "```text",
        OFFICIAL_TEXT.read_text(encoding="utf-8").rstrip(),
        "```",
        "",
        "## Unique Blueprint JSON",
        "```json",
        BLUEPRINT.read_text(encoding="utf-8").rstrip(),
        "```",
        "",
        "## Candidate DOCX body — numbered paragraphs",
        extracted,
        "",
        "## Candidate DOCX tables — exact extracted cells",
        "\n\n".join(tables),
        "",
        "## Technical QA receipt (evidence, not a semantic verdict)",
        "```json",
        TECH_QA.read_text(encoding="utf-8").rstrip(),
        "```",
        "",
        "## Visual QA scope (evidence, not a semantic verdict)",
        "```json",
        VISUAL_QA.read_text(encoding="utf-8").rstrip(),
        "```",
        "",
        "## Input SHA-256 manifest",
        "```json",
        json.dumps(hashes, ensure_ascii=False, indent=2),
        "```",
        "",
        "## Source manuscript (line-level comparison aid only; the Word extraction above is the artifact under review)",
        MANUSCRIPT.read_text(encoding="utf-8").rstrip(),
        "",
    ]
    packet_text = "\n".join(body)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(packet_text, encoding="utf-8")
    manifest = {
        "work_id": "TTQS-0111-I18-E20-INDEPENDENT-SEMANTIC-REVIEW-R01-20261023",
        "doc_id": "0111",
        "requirement_id": "I18-E20",
        "official_checkbox_serial": "UNKNOWN_NOT_ASSERTED",
        "candidate_docx_sha256": candidate_sha,
        "candidate_docx_bytes": DOCX.stat().st_size,
        "packet_path": str(OUT),
        "packet_sha256": sha256(OUT),
        "packet_bytes": OUT.stat().st_size,
        "input_hashes": hashes,
        "paragraphs_extracted": len(paras),
        "tables_extracted": len(tables),
        "prepared_at_utc": datetime.now(timezone.utc).isoformat(),
        "result": "READY_FOR_INDEPENDENT_EXACT_SHA_TEXT_REVIEW",
    }
    MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    main()
