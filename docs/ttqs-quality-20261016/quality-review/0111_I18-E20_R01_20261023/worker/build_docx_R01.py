from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "manuscript" / "0111_I18-E20_特殊績效及社會面成果證據登錄表_R01.md"
OUT_DIR = ROOT / "candidate"
OUT = OUT_DIR / "0111_I18-E20_特殊績效及社會面成果證據登錄表_R01.docx"
RECEIPT = ROOT / "evidence" / "docx_build_receipt_0111_R01.json"

NAVY = "17365D"
TEAL = "0F6B78"
PALE = "EAF2F8"
SAMPLE_BG = "FCE4D6"
SAMPLE_RED = "9C0006"
GRAY = "666666"
FONT = "Microsoft JhengHei"


def set_run_font(run, size=10.5, bold=None, color=None, italic=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, inches):
    cell.width = Inches(inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def add_field(paragraph, field):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {field} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))
    set_run_font(run, 8, color=GRAY)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|(?<!\*)\*[^*]+\*)")
LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def add_inline(paragraph, text, base_size=10.5, color=None):
    text = LINK.sub(lambda m: f"{m.group(1)} ({m.group(2)})", text)
    pos = 0
    for match in INLINE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, base_size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, base_size, bold=True, color=color)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, max(9.5, base_size - 0.2), color=TEAL)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, base_size, italic=True, color=color)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, base_size, color=color)


def split_pipe_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    cells, buf, escaped = [], [], False
    for char in line:
        if char == "|" and not escaped:
            cells.append("".join(buf).strip().replace("\\|", "|"))
            buf = []
        else:
            buf.append(char)
        escaped = char == "\\" and not escaped
        if char != "\\":
            escaped = False
    cells.append("".join(buf).strip().replace("\\|", "|"))
    return cells


def is_separator(line):
    cells = split_pipe_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells)


def parse_markdown(text):
    lines = text.splitlines()
    items = []
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith("|"):
            raw = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                raw.append(lines[i].strip())
                i += 1
            if len(raw) >= 2 and is_separator(raw[1]):
                rows = [split_pipe_row(raw[0])] + [split_pipe_row(x) for x in raw[2:]]
                width = max(map(len, rows))
                rows = [row + [""] * (width - len(row)) for row in rows]
                items.append(("table", rows))
            else:
                items.extend(("paragraph", r) for r in raw)
            continue
        if line.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1 if i < len(lines) else 0
            items.append(("code", "\n".join(code)))
            continue
        if line.startswith(">"):
            items.append(("quote", line[1:].lstrip()))
            i += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            items.append(("heading", (len(heading.group(1)), heading.group(2).strip())))
            i += 1
            continue
        bullet = re.match(r"^(\s*)([-*+]\s+)(.*)$", line)
        if bullet:
            items.append(("bullet", (len(bullet.group(1)), bullet.group(3))))
            i += 1
            continue
        numbered = re.match(r"^(\s*)(\d+)[.)]\s+(.*)$", line)
        if numbered:
            items.append(("number", (len(numbered.group(1)), int(numbered.group(2)), numbered.group(3))))
            i += 1
            continue
        paragraph_lines = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(("#", ">", "|", "```")) and not re.match(r"^\s*(?:[-*+]\s+|\d+[.)]\s+)", lines[i]):
            paragraph_lines.append(lines[i].strip())
            i += 1
        items.append(("paragraph", " ".join(paragraph_lines)))
    return items


def setup_section(section, landscape=False):
    if landscape:
        section.orientation = 1
        section.page_width = Inches(11.7)
        section.page_height = Inches(8.3)
        section.left_margin = Inches(0.52)
        section.right_margin = Inches(0.52)
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
    else:
        section.orientation = 0
        section.page_width = Inches(8.3)
        section.page_height = Inches(11.7)
        section.left_margin = Inches(0.72)
        section.right_margin = Inches(0.72)
        section.top_margin = Inches(0.72)
        section.bottom_margin = Inches(0.72)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.32)


def add_header_footer(section):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header = section.header.paragraphs[0]
    for p in section.header.paragraphs:
        for run in list(p.runs):
            p._p.remove(run._r)
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("0111  |  特殊績效及社會面成果證據登錄表  |  SAMPLE 示範")
    set_run_font(r, 8, color=GRAY)
    footer = section.footer.paragraphs[0]
    for p in section.footer.paragraphs:
        for run in list(p.runs):
            p._p.remove(run._r)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("候選示範文件  •  REAL 原件 0 件  •  正式驗收未授予    第 ")
    set_run_font(r, 8, color=GRAY)
    add_field(footer, "PAGE")
    r = footer.add_run(" 頁")
    set_run_font(r, 8, color=GRAY)


def style_document(doc):
    section = doc.sections[0]
    setup_section(section, False)
    for name, size, color in (("Normal", 10.5, "222222"), ("Title", 23, NAVY), ("Heading 1", 16, NAVY), ("Heading 2", 12.5, TEAL), ("Heading 3", 11, NAVY)):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        if name != "Normal":
            style.font.bold = True
            style.paragraph_format.keep_with_next = True
    title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    normal = doc.styles["Normal"]
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15
    for s in doc.sections:
        add_header_footer(s)


def table_widths(rows, usable):
    cols = len(rows[0])
    weights = []
    for c in range(cols):
        samples = [rows[0][c]] + [r[c] for r in rows[1:]]
        header_len = min(len(samples[0]), 24)
        avg_len = sum(min(len(x), 60) for x in samples[1:]) / max(1, len(samples) - 1)
        weights.append(max(1.0, header_len * 1.25 + min(avg_len, 34) * 0.55))
    total = sum(weights)
    widths = [usable * w / total for w in weights]
    min_w = 0.62 if cols >= 7 else 0.72
    if any(w < min_w for w in widths):
        widths = [max(min_w, w) for w in widths]
        scale = usable / sum(widths)
        widths = [w * scale for w in widths]
    return widths


def add_markdown_table(doc, rows, landscape):
    cols = len(rows[0])
    usable = 10.65 if landscape else 6.86
    tbl = doc.add_table(rows=1, cols=cols)
    tbl.style = "Table Grid"
    tbl.autofit = False
    tblPr = tbl._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    widths = table_widths(rows, usable)
    for ridx, values in enumerate(rows):
        row = tbl.rows[0] if ridx == 0 else tbl.add_row()
        if ridx == 0:
            set_repeat_header(row)
        prevent_row_split(row)
        for cidx, value in enumerate(values):
            cell = row.cells[cidx]
            set_cell_width(cell, widths[cidx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER if ridx == 0 else WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if ridx == 0:
                shade(cell, NAVY)
                add_inline(p, value, 9.5, "FFFFFF")
                for run in p.runs:
                    run.bold = True
            else:
                if ridx % 2 == 0:
                    shade(cell, "F3F6F9")
                add_inline(p, value, 9.5, "222222")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


def next_wide_table(items, start):
    for kind, value in items[start + 1:]:
        if kind == "table":
            return len(value[0]) >= 7
        if kind in ("heading", "paragraph", "quote", "bullet", "number", "code"):
            return False
    return False


def deferred_table_context(items):
    deferred = {}
    skip = set()
    for i, (kind, value) in enumerate(items):
        if kind != "table":
            continue
        caption_i = i - 1
        if caption_i < 0:
            continue
        prev_kind, prev_value = items[caption_i]
        if prev_kind != "paragraph" or not re.match(r"^\*\*表\s*\d+", prev_value):
            continue
        context = []
        heading_i = caption_i - 1
        while heading_i >= 0 and items[heading_i][0] == "heading":
            context.append(items[heading_i])
            skip.add(heading_i)
            heading_i -= 1
        context.reverse()
        context.append(items[caption_i])
        skip.add(caption_i)
        deferred[i] = context
    return deferred, skip


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    source_bytes = SOURCE.read_bytes()
    text = source_bytes.decode("utf-8")
    items = parse_markdown(text)
    table_context, skip_items = deferred_table_context(items)
    doc = Document()
    style_document(doc)
    landscape = False

    for idx, (kind, value) in enumerate(items):
        if idx in skip_items:
            continue
        if kind == "table":
            want_landscape = len(value[0]) >= 7
            orientation_change = want_landscape != landscape
            if want_landscape != landscape:
                section = doc.add_section(WD_SECTION.NEW_PAGE)
                setup_section(section, want_landscape)
                add_header_footer(section)
                landscape = want_landscape
            context_items = table_context.get(idx, [])
            annex_index_table = any(kind == "paragraph" and "附件索引" in content for kind, content in context_items)
            complete_cohort_table = len(value) >= 9 and len(value[0]) >= 7
            force_table_page = (complete_cohort_table or annex_index_table) and not orientation_change
            if force_table_page and not context_items:
                p = doc.add_paragraph()
                p.paragraph_format.page_break_before = True
            for context_index, (context_kind, context_value) in enumerate(context_items):
                if context_kind == "heading":
                    level, title = context_value
                    p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
                    p.paragraph_format.keep_with_next = True
                    if force_table_page and context_index == 0:
                        p.paragraph_format.page_break_before = True
                    add_inline(p, title, {2: 12.5, 3: 11, 4: 10.5, 5: 10.5, 6: 10.5}.get(level, 10.5), TEAL if level == 2 else NAVY)
                else:
                    cap = doc.add_paragraph()
                    cap.paragraph_format.keep_with_next = True
                    cap.paragraph_format.space_after = Pt(5)
                    if force_table_page and context_index == 0:
                        cap.paragraph_format.page_break_before = True
                    add_inline(cap, context_value, 10.5, NAVY)
            add_markdown_table(doc, value, landscape)
        elif kind == "heading":
            level, title = value
            if level == 1:
                p = doc.add_paragraph(style="Title")
                p.paragraph_format.space_after = Pt(12)
                add_inline(p, title, 23, NAVY)
            else:
                p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
                p.paragraph_format.keep_with_next = True
                add_inline(p, title, {2: 12.5, 3: 11, 4: 10.5, 5: 10.5, 6: 10.5}.get(level, 10.5), TEAL if level == 2 else NAVY)
        elif kind == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.right_indent = Inches(0.12)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            p_pr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), SAMPLE_BG)
            p_pr.append(shd)
            add_inline(p, value, 10.5, SAMPLE_RED)
            for run in p.runs:
                run.bold = True
        elif kind == "bullet":
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Inches(0.28 + min(value[0] // 2, 3) * 0.18)
            add_inline(p, value[1], 10.5)
        elif kind == "number":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.28 + min(value[0] // 2, 3) * 0.18)
            add_inline(p, f"{value[1]}. ", 10.5)
            add_inline(p, value[2], 10.5)
        elif kind == "code":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.right_indent = Inches(0.18)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(7)
            for j, line in enumerate(value.splitlines()):
                if j:
                    p.add_run().add_break()
                r = p.add_run(line)
                set_run_font(r, 9.5, color=TEAL)
        else:
            p = doc.add_paragraph()
            if value.startswith("**"):
                p.paragraph_format.space_before = Pt(3)
            add_inline(p, value, 10.5)

    doc.core_properties.title = "0111 特殊績效及社會面成果證據登錄表"
    doc.core_properties.subject = "I18-E20 候選示範文件；SAMPLE／SYNTHETIC；REAL E0／HOLD"
    doc.core_properties.author = "TTQS ONE 候選文件施工"
    doc.core_properties.keywords = "0111, I18-E20, SAMPLE, SYNTHETIC, HOLD"
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    RECEIPT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    sha = hashlib.sha256(OUT.read_bytes()).hexdigest()
    receipt = {
        "work_id": "TTQS-0111-I18-E20-DOCX-BUILD-R01-20261023",
        "doc_id": "0111",
        "requirement_id": "I18-E20",
        "official_checkbox_serial": "UNKNOWN_NOT_ASSERTED",
        "source_manuscript": str(SOURCE),
        "source_manuscript_sha256": hashlib.sha256(source_bytes).hexdigest(),
        "docx_path": str(OUT),
        "docx_sha256": sha,
        "docx_bytes": OUT.stat().st_size,
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "table_row_counts": [len(t.rows) for t in doc.tables],
        "table_column_counts": [len(t.columns) for t in doc.tables],
        "style": "A4 portrait narrative with next-page landscape sections for tables of 7+ columns; repeated header rows; minimum table text 9.5 pt; body 10.5 pt; SAMPLE/HOLD footer.",
        "result": "BUILT_CANDIDATE_PENDING_TECHNICAL_RENDER_VISUAL_AND_INDEPENDENT_SEMANTIC_REVIEW",
    }
    RECEIPT.write_text(json.dumps(receipt, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(receipt, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
