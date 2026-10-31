from __future__ import annotations
import re
from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0008_I01-E05_LUNA_TEXT_DRAFT_R03_20261030")
SOURCE = ROOT / "worker" / "0008_I01-E05_R03.md"
OUTPUT = ROOT / "candidate" / "0008_I01-E05_R03c.docx"
FONT = "Microsoft JhengHei"
NAVY, PALE, BORDER = "17365D", "EAF1F8", "D9D9D9"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)

def font(run, size=10.5, bold=None, color="000000"):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    for key in ("eastAsia", "ascii", "hAnsi"):
        rpr.rFonts.set(qn("w:" + key), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold

def inline(p, s, size=10.5, bold=False, color="000000"):
    rx = re.compile(r"(\*\*.*?\*\*|\*[^*]+\*)")
    at = 0
    for m in rx.finditer(s):
        if m.start() > at:
            font(p.add_run(s[at:m.start()]), size, bold, color)
        token = m.group(0)
        if token.startswith("**"):
            font(p.add_run(token[2:-2]), size, True, color)
        else:
            r = p.add_run(token[1:-1])
            font(r, size, bold, color)
            r.italic = True
        at = m.end()
    if at < len(s):
        font(p.add_run(s[at:]), size, bold, color)

def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    e = tcpr.find(qn("w:shd"))
    if e is None:
        e = OxmlElement("w:shd")
        tcpr.append(e)
    e.set(qn("w:fill"), fill)

def cell_margins(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = tcpr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcpr.append(mar)
    for edge in ("top", "start", "bottom", "end"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:w"), "90")
        e.set(qn("w:type"), "dxa")
        mar.append(e)

def table_borders(table):
    pr = table._tbl.tblPr
    b = pr.first_child_found_in("w:tblBorders")
    if b is None:
        b = OxmlElement("w:tblBorders")
        pr.append(b)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), BORDER)
        b.append(e)

def repeat_header(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))

def add_page_field(p):
    r = p.add_run()
    font(r, 8.5, color="536273")
    begin, instr, sep, value, end = [OxmlElement("w:" + t) for t in ("fldChar", "instrText", "fldChar", "t", "fldChar")]
    begin.set(qn("w:fldCharType"), "begin")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    sep.set(qn("w:fldCharType"), "separate")
    value.text = "1"
    end.set(qn("w:fldCharType"), "end")
    r._r.extend([begin, instr, sep, value, end])

def configure_section(sec, landscape=False):
    sec.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    sec.header.is_linked_to_previous = False
    sec.footer.is_linked_to_previous = False
    if landscape:
        sec.page_width, sec.page_height = Inches(11.69), Inches(8.27)
        sec.left_margin = sec.right_margin = Inches(0.48)
        sec.top_margin = sec.bottom_margin = Inches(0.62)
    else:
        sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
        sec.left_margin = sec.right_margin = Inches(0.66)
        sec.top_margin, sec.bottom_margin = Inches(0.72), Inches(0.70)
    sec.header_distance = sec.footer_distance = Inches(0.32)
    h = sec.header.paragraphs[0]
    h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(h.add_run("TTQS ONE  |  0008  |  CANDIDATE  |  SAMPLE / SYNTHETIC"), 8.5, color="536273")
    f = sec.footer.paragraphs[0]
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(f.add_run("非核定示例  ·  頁 "), 8.5, color="536273")
    add_page_field(f)

def widths(n, landscape, header):
    usable = 10.62 if landscape else 6.95
    if n == 7 and header and header[0].lower() == "fact_id":
        # Give the fact identifier enough width to keep SF-00001 on one line.
        # Recover the 0.20 inches from the long locator column; total width stays 10.73 in.
        return [0.90, 1.55, 1.15, 0.72, 4.35, 1.10, 0.96]
    patterns = {
        1: [1.0], 2: [0.28, 0.72], 3: [0.19, 0.38, 0.43],
        4: [0.15, 0.29, 0.34, 0.22], 5: [0.14, 0.19, 0.23, 0.27, 0.17],
        6: [0.13, 0.16, 0.19, 0.20, 0.20, 0.12],
    }
    if n in patterns:
        return [usable * x for x in patterns[n]]
    weights = [max(1, len(x)) for x in header]
    return [usable * x / sum(weights) for x in weights]

def parse_row(s):
    return [x.strip().replace(r"\|", "|") for x in s.strip().strip("|").split("|")]

def is_separator(s):
    return bool(re.match(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)+\|?\s*$", s))

def add_table(doc, rows, landscape):
    if not rows:
        return
    n = max(map(len, rows))
    rows = [r + [""] * (n-len(r)) for r in rows]
    t = doc.add_table(rows=1, cols=n)
    t.alignment, t.autofit, t.style = WD_TABLE_ALIGNMENT.CENTER, False, "Table Grid"
    table_borders(t)
    ws = widths(n, landscape, rows[0])
    for j, w in enumerate(ws):
        t.columns[j].width = Inches(w)
    repeat_header(t.rows[0])
    for ri, values in enumerate(rows):
        row = t.rows[0] if ri == 0 else t.add_row()
        for j, value in enumerate(values):
            c = row.cells[j]
            c.width = Inches(ws[j])
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.04
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 or (j == 0 and n > 2) else WD_ALIGN_PARAGRAPH.LEFT
            inline(p, value, size=9.5, bold=(ri == 0), color=("FFFFFF" if ri == 0 else "000000"))
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(c)
            if ri == 0:
                shade(c, NAVY)
            elif ri % 2 == 0:
                shade(c, PALE)
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

def main():
    if OUTPUT.exists():
        raise FileExistsError(OUTPUT)
    lines = SOURCE.read_text(encoding="utf-8-sig").splitlines()
    d = Document()
    normal = d.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.10
    for name, size in (("Title", 23), ("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)):
        st = d.styles[name]
        st.font.name = FONT
        st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.bold = name != "Title"
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.space_before = Pt(10)
        st.paragraph_format.space_after = Pt(5)
    d.styles["Caption"].font.name = FONT
    d.styles["Caption"].font.size = Pt(10.5)
    d.styles["Caption"].font.bold = True
    d.styles["Caption"].font.color.rgb = RGBColor(0, 0, 0)
    d.styles["Caption"].paragraph_format.keep_with_next = True
    configure_section(d.sections[0], False)
    landscape = False
    title_done = False
    code_mode = False
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith(chr(96) * 3):
            code_mode = not code_mode
            i += 1
            continue
        if code_mode:
            p = d.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            font(p.add_run(line), 9.5, color="27364B")
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        if line.startswith("### 12.2 原子合成事實列"):
            configure_section(d.add_section(WD_SECTION.NEW_PAGE), True)
            landscape = True
        elif line.startswith("## 13. REAL 替換"):
            configure_section(d.add_section(WD_SECTION.NEW_PAGE), False)
            landscape = False
        if line.lstrip().startswith("|"):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                s = lines[i].strip()
                if not is_separator(s):
                    rows.append(parse_row(s))
                i += 1
            add_table(d, rows, landscape)
            continue
        if line.startswith("# "):
            title = re.sub(r"^#\s*", "", line)
            title = re.sub(r"^0008[｜|]\s*", "", title).replace("（", " ").replace("）", "").replace("：", " ")
            p = d.add_paragraph(style="Title" if not title_done else "Heading 1")
            inline(p, ("0008 " + title) if not title_done else title, size=23 if not title_done else 15, bold=True)
            title_done = True
            i += 1
            continue
        h = re.match(r"^(#{2,4})\s+(.*)$", line)
        if h:
            level = len(h.group(1))
            style = {2:"Heading 1", 3:"Heading 2", 4:"Heading 3"}[level]
            p = d.add_paragraph(style=style)
            inline(p, h.group(2), size={"Heading 1":15,"Heading 2":12.5,"Heading 3":11}[style], bold=True)
            i += 1
            continue
        if re.match(r"^\s*[-*]\s+", line):
            p = d.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(4)
            inline(p, re.sub(r"^\s*[-*]\s+", "", line))
            i += 1
            continue
        if re.match(r"^\s*\d+\.\s+", line):
            p = d.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(4)
            inline(p, re.sub(r"^\s*\d+\.\s+", "", line))
            i += 1
            continue
        if line.startswith("**表 ") and line.endswith("**"):
            p = d.add_paragraph(style="Caption")
            inline(p, line[2:-2], size=10.5, bold=True)
        else:
            p = d.add_paragraph()
            inline(p, line)
        i += 1
    if not title_done:
        raise ValueError("No title")
    d.core_properties.title = "0008 各類別訓練課程年度成長計畫 SAMPLE"
    d.core_properties.subject = "I01-E05 R03 候選；未核定、未實施"
    d.core_properties.author = "TTQS ONE candidate"
    d.core_properties.keywords = "TTQS, 0008, I01-E05, SAMPLE, SYNTHETIC, CANDIDATE"
    d.save(OUTPUT)
    print(f"created={OUTPUT}")
    print(f"tables={len(d.tables)}")
    print(f"sections={len(d.sections)}")
    print(f"bytes={OUTPUT.stat().st_size}")
    print(f"source_lines={len(lines)}")

if __name__ == "__main__":
    main()
