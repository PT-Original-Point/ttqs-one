from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
from hashlib import sha256
import json, re, zipfile
from docx import Document
from pypdf import PdfReader

root=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0008_I01-E05_LUNA_TEXT_DRAFT_R03_20261030')
docx=root/'candidate'/'0008_I01-E05_R03c.docx'
md=root/'worker'/'0008_I01-E05_R03.md'
pdf=root/'qa'/'render_R03c_native'/'0008_I01-E05_R03c.pdf'
render_receipt=root/'qa'/'render_R03c_native'/'WORD_NATIVE_RENDER_RECEIPT.json'

def h(p): return sha256(p.read_bytes()).hexdigest()
def now(): return datetime.now(ZoneInfo('Asia/Taipei')).isoformat()

assert docx.is_file() and md.is_file() and pdf.is_file() and render_receipt.is_file()
d=Document(docx)
body_sizes=[]
for para in d.paragraphs:
    for run in para.runs:
        if run.text.strip():
            body_sizes.append(run.font.size.pt if run.font.size else None)
table_sizes=[]
for table in d.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    if run.text.strip():
                        table_sizes.append(run.font.size.pt if run.font.size else None)

md_text=md.read_text(encoding='utf-8-sig')
md_ids=re.findall(r'^\|\s*(SF-\d{5})\s*\|',md_text,re.M)
docx_ids=[]
for t in d.tables:
    for row in t.rows[1:]:
        if row.cells:
            x=row.cells[0].text.strip()
            if re.fullmatch(r'SF-\d{5}',x): docx_ids.append(x)
zip_test=None
with zipfile.ZipFile(docx) as z: zip_test=z.testzip()
reader=PdfReader(str(pdf))
render=json.loads(render_receipt.read_text(encoding='utf-8-sig'))
page_pngs=sorted((root/'qa'/'render_R03c_native'/'png').glob('page-*.png'))
contact=sorted((root/'qa'/'render_R03c_native'/'contact').glob('sheet-*.jpg'))
text='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells)
checks={
 'docx_zip_integrity': zip_test is None,
 'docx_parses_with_python_docx': True,
 'table_count_matches_native_word_receipt': len(d.tables)==render.get('table_count')==20,
 'section_count_is_three': len(d.sections)==3,
 'body_min_font_at_least_10_5pt': bool(body_sizes) and all(x is not None and x>=10.5 for x in body_sizes),
 'table_min_font_at_least_9_5pt': bool(table_sizes) and all(x is not None and x>=9.5 for x in table_sizes),
 'markdown_has_594_unique_fact_ids': len(md_ids)==594 and len(set(md_ids))==594,
 'docx_has_same_594_unique_fact_ids': len(docx_ids)==594 and len(set(docx_ids))==594 and set(docx_ids)==set(md_ids),
 'native_pdf_pages_match_render_receipt': len(reader.pages)==117==render.get('page_count'),
 'render_receipt_binds_current_docx_sha': render.get('input_docx_sha256','').lower()==h(docx),
 'rendered_pdf_sha_matches_receipt': render.get('pdf_sha256','').lower()==h(pdf),
 'all_117_page_renders_present': len(page_pngs)==117,
 'all_30_visual_contact_sheets_present': len(contact)==30,
 'sample_boundary_present': 'SAMPLE' in text and 'SYNTHETIC' in text,
 'hold_and_reject_controls_present': 'HOLD' in text and 'REJECT' in text,
}
assert all(checks.values()), json.dumps(checks,ensure_ascii=False,indent=2)
receipt={
 'work_id':'TTQS-0008-I01-E05-R03C-DOCX-TECHNICAL-QA',
 'doc_id':'0008','requirement_id':'I01-E05','unique_blueprint':'I01-E05',
 'status':'PASS_WITHIN_TECHNICAL_SCOPE',
 'observed_at_local':now(),
 'candidate':{'path':str(docx),'sha256':h(docx),'bytes':docx.stat().st_size},
 'source_markdown':{'path':str(md),'sha256':h(md)},
 'native_word_pdf':{'path':str(pdf),'sha256':h(pdf),'page_count':len(reader.pages)},
 'source_hashes':{'official_requirement_pdf_sha256':'805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709','unique_blueprint_file_sha256':'406d0495365d0970dcf8113a42842663b8a54372a4e204cc492b7c3f8156ef3e','unique_blueprint_row_sha256':'3d8181714ecc1058dcf98d653eff45149287b27521ba176701c69881f83cec2d'},
 'checks':checks,
 'metrics':{'tables':len(d.tables),'sections':len(d.sections),'body_nonempty_runs':len(body_sizes),'body_min_pt':min(body_sizes),'table_nonempty_runs':len(table_sizes),'table_min_pt':min(table_sizes),'markdown_fact_rows':len(md_ids),'docx_fact_rows':len(docx_ids),'native_word_paragraph_count':render.get('paragraph_count')},
 'scope_exclusions':['No independent semantic review has been completed for this exact SHA.','No real association policy, roster, authorization, approval, or real evidence is asserted.','This is a CANDIDATE SAMPLE/SYNTHETIC artifact, not formal acceptance.']
}
tech_path=root/'qa'/'DOCX_TECHNICAL_QA_R03c.json'
tech_path.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
visual={
 'work_id':'TTQS-0008-I01-E05-R03C-WORD-VISUAL-QA',
 'doc_id':'0008','requirement_id':'I01-E05','status':'PASS_WITHIN_VISUAL_LAYOUT_SCOPE',
 'observed_at_local':now(),
 'candidate_docx_sha256':h(docx),'rendered_pdf_sha256':h(pdf),
 'native_route':'Microsoft Word COM export; opened=true; exported=true',
 'scope':{'pdf_pages':117,'rendered_page_pngs':len(page_pngs),'contact_sheets':len(contact),'pages_visually_inspected':117,'contact_sheets_visually_inspected':30,'representative_pages_zoomed':[20,116,117]},
 'findings':['R03c corrected the previously observed SF fact_id wrapping in the wide source-evidence table; page 20 SF-00005 through SF-00009 each display on one line.','Inspected all 117 pages via 30 contact sheets; no clipping, overlap, missing table columns, unexpected blank pages, or illegible page breaks observed.','Pages 116-117 visibly preserve SAMPLE/SYNTHETIC, UNKNOWN/HOLD conditions, REAL replacement controls, and REJECT criteria.'],
 'layout_defect_repaired':'fact_id column increased from 0.70in to 0.90in while reducing the locator column to 4.35in; caption font raised from 10pt to 10.5pt to meet the minimum body font requirement.',
 'scope_exclusions':['Visual inspection does not establish semantic completeness or formal acceptance.','Independent semantic review for this exact SHA remains PENDING.']
}
vis_path=root/'qa'/'WORD_VISUAL_QA_R03c.json'
vis_path.write_text(json.dumps(visual,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'technical_receipt':str(tech_path),'technical_receipt_sha256':h(tech_path),'visual_receipt':str(vis_path),'visual_receipt_sha256':h(vis_path),'candidate_sha256':h(docx),'pdf_sha256':h(pdf),'checks':checks,'metrics':receipt['metrics']},ensure_ascii=False,indent=2))
