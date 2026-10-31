from __future__ import annotations

from datetime import datetime
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile
import json
import re

from docx import Document

ROOT = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0010_I02-E01_ASSESSMENT_STANDARD_R04_20261030")
DOCX = ROOT / "candidate" / "0010_I02-E01_R18_R03.docx"
SOURCE = ROOT / "worker" / "0010_I02-E01_R18_R02.md"
REGISTER = ROOT / "worker" / "SYNTHETIC_REGISTER_R17_R02.md"
BLUEPRINT = ROOT / "source" / "BLUEPRINT_ROW_0010_I02-E01_R06.json"
BINDING = ROOT / "source" / "SOURCE_BINDING_R06.json"
BUILDER = ROOT / "qa" / "build_0010_r18_r03_docx.py"
PREDECESSOR = ROOT / "candidate" / "0010_I02-E01_R17.docx"
OUT = ROOT / "qa" / "TECHNICAL_QA_R18_R03.json"


def digest(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def document_text(doc: Document) -> str:
    result = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            result.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(result)


def fact_rows(text: str) -> list[str]:
    lines = text.splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith("|Fact ID|"))
    result = []
    for line in lines[start + 1:]:
        if not line.startswith("|"):
            break
        if not line.startswith("|---"):
            result.append(line)
    return result


if OUT.exists():
    raise SystemExit(f"Refusing to overwrite existing receipt: {OUT}")

with ZipFile(DOCX) as archive:
    zip_ok = archive.testzip() is None and "word/document.xml" in archive.namelist()

doc = Document(DOCX)
full = document_text(doc)
source = SOURCE.read_text(encoding="utf-8-sig")
register = REGISTER.read_text(encoding="utf-8-sig")
blueprint = json.loads(BLUEPRINT.read_text(encoding="utf-8-sig"))
binding = json.loads(BINDING.read_text(encoding="utf-8-sig"))
source_rows = fact_rows(source)
register_rows = fact_rows(register)
checks: dict[str, bool] = {}

checks["zip_integrity"] = zip_ok
checks["python_docx_open_and_structure"] = bool(doc.paragraphs) and len(doc.tables) == 22
ev_table = next((t for t in doc.tables if t.rows and t.rows[0].cells[0].text.strip().lower() == "ev id"), None)
checks["ev_behavior_column_width_over_4in"] = ev_table is not None and ev_table.columns[3].width.inches >= 4.0
checks["cover_title_docid_and_r18"] = "0010 學員訓練成果評核標準與說明" in full and "版本：R18；狀態：SAMPLE/SYNTHETIC 候選稿" in full
checks["official_mapping_i02_e01_checkbox_008"] = "對學員訓練成果有評核標準的訂定及說明。" in full and "I02-E01" in full and "008" in full
checks["blueprint_r06_binding"] = blueprint.get("blueprint_revision") == "R06" and binding.get("unique_blueprint_path") == BLUEPRINT.name and binding.get("unique_blueprint_sha256") == digest(BLUEPRINT)
checks["blueprint_count_reconciled_to_12"] = "12 個行為指標量規" in " ".join(blueprint.get("acceptance_tests", [])) and "13 個行為指標量規" not in " ".join(blueprint.get("acceptance_tests", [])) and "knowledge test separate" in blueprint.get("count_reconciliation", {}).get("correction", "")
checks["sample_real_scope_boundary"] = "SAMPLE/SYNTHETIC" in full and "不是真人授權、簽章或原件" in full and "任一缺漏或衝突即停止並HOLD" in full
checks["retest_only_auth_card"] = "AUTH-RET-SYN-01" in full and "首次評量中授權UNKNOWN" in full and "SYS-RET-SYN-01" in full
checks["minimum_necessary_and_restricted_fields"] = "收件狀態及下一步欄位" in full and "禁止查看其他服務對象紀錄" in full
checks["synthetic_facts_162_and_exact_parity"] = len(source_rows) == 162 and len(register_rows) == 162 and source_rows == register_rows
checks["new_fact_ids_present"] = all(any(row.startswith(f"|SF-{i}|") for row in register_rows) for i in range(114, 120))
checks["p3_003_observation_supports_level_3"] = all(x in full for x in ("EV-SYN-003-P3-R1", "AUTH-RET-SYN-003", "未觸及禁限欄", "非關鍵回覆管道欄"))
checks["p3_004_observation_supports_level_3"] = all(x in full for x in ("EV-SYN-004-P3-R1", "AUTH-RET-SYN-004", "看到待安排即停止", "未查他案或禁限資料"))
checks["p3_prompt_status_and_rating"] = all(any(row.cells[0].text == ev and row.cells[4].text == "提示前；未使用" and row.cells[5].text == "3" for table in doc.tables for row in table.rows) for ev in ("EV-SYN-003-P3-R1", "EV-SYN-004-P3-R1"))
checks["retest_form_and_answer_vectors_preserved"] = all(x in full for x in ("FORM-SYN-A v2", "KEY-SYN-A-v2", "17/20", "14/20"))
checks["retest_003_formula_and_result"] = "79.50=17+27.50+22.50+12.50" in full and "27.50/40<28" in full and "79.50/100；實作未達門檻，NOT_PASS（完整證據、可重算）" in full
checks["retest_004_formula_and_limited_scope"] = "80.25=14+32.50+22.50+11.25" in full and "80.25/100" in full and "PASS（範圍受限複評" in full
checks["hold_vs_not_pass_boundary"] = "資料完整且可重算，但實作門檻未達" in full and "不足以可靠判讀" in full and "HOLD" in full and "NOT_PASS" in full
checks["20_question_and_balanced_key"] = all(f"Q{i}" in full for i in range(1, 21)) and "A=5、B=5、C=5、D=5" in full
checks["body_min_font_10_5pt"] = bool([r.font.size.pt for p in doc.paragraphs for r in p.runs if r.font.size]) and min(r.font.size.pt for p in doc.paragraphs for r in p.runs if r.font.size) >= 10.5
checks["table_min_font_9_5pt"] = bool([r.font.size.pt for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs for r in p.runs if r.font.size]) and min(r.font.size.pt for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs for r in p.runs if r.font.size) >= 9.5
checks["no_unlabeled_placeholder_cells"] = not any(c.text.strip() in {"待填", "待核", "空白表", "會後填寫"} for t in doc.tables for row in t.rows for c in row.cells)
checks["builder_bound_to_r18_source_and_output"] = "0010_I02-E01_R18_R02.md" in BUILDER.read_text(encoding="utf-8-sig") and "0010_I02-E01_R18_R03.docx" in BUILDER.read_text(encoding="utf-8-sig")
checks["r17_predecessor_unchanged"] = digest(PREDECESSOR) == "57e48e29be753d029d18c45bc86706866bdb98a9257061509b213cc3897eac70"

result = {
    "work_id": "TTQS-0010-R18-R03-TECHNICAL-QA-20261030",
    "doc_id": "0010",
    "requirement_id": "I02-E01",
    "official_checkbox_serial": "008",
    "candidate_status": "CANDIDATE_NOT_FORMALLY_ACCEPTED",
    "candidate_docx_sha256": digest(DOCX),
    "source_sha256": digest(SOURCE),
    "synthetic_register_sha256": digest(REGISTER),
    "unique_blueprint_sha256": digest(BLUEPRINT),
    "source_binding_sha256": digest(BINDING),
    "builder_sha256": digest(BUILDER),
    "checks": checks,
    "all_checks_pass": all(checks.values()),
    "table_count": len(doc.tables),
    "paragraph_count": len(doc.paragraphs),
    "synthetic_fact_rows": len(register_rows),
    "body_min_explicit_font_pt": min(r.font.size.pt for p in doc.paragraphs for r in p.runs if r.font.size),
    "table_min_explicit_font_pt": min(r.font.size.pt for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs for r in p.runs if r.font.size),
    "formula_check": "SYN-003 79.50=17+27.50+22.50+12.50; 27.50/40<28 => NOT_PASS. SYN-004 80.25=14+32.50+22.50+11.25; numeric thresholds met, limited-scope PASS only because P3 level-3 evidence is now explicit. Missing/conflicting REAL evidence => HOLD.",
    "semantic_review": "PENDING; deterministic technical result is not an independent semantic verdict.",
    "formal_acceptance": "NOT_ACCEPTED",
    "checked_at_taipei": datetime.now().astimezone().isoformat(),
}
OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(result, ensure_ascii=False, indent=2))
if not result["all_checks_pass"]:
    raise SystemExit(2)
