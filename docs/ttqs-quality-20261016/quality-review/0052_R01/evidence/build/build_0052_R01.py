from __future__ import annotations

import csv
import hashlib
import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


PKG = Path(__file__).resolve().parents[2]
REPO = PKG.parents[3]
RUNTIME = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME")
DOCX = PKG / "candidate" / "0052_R01.docx"
REGISTER = PKG / "evidence" / "SYNTHETIC_REPLACEMENT_REGISTER_R01.csv"
SOURCE_MAP = PKG / "inputs" / "SOURCE_MAP_R01.md"
BLUEPRINT_COPY = PKG / "inputs" / "UNIQUE_BLUEPRINT_0052.jsonl"

OFFICIAL_TEXT = "規劃課程依學員及結訓學員意見訂定訓練課程，且流程有完整紀錄。"
COURSE = "志工服務紀錄、交接與工作界線實務"
CASE = "SYN-0033-PLAN-01"
BUILD_SOURCE_HEAD = "bdfa5a675ca77ff460530b96971b26ff44c455c4"

LEARNER_Q = [
    "課程內容與志工服務中的角色及任務相符。",
    "我能說明課程目標、適用範圍及何時應停止自行判斷。",
    "課程提供足夠的引導練習，讓我能完成服務交接步驟。",
    "遇到資料不明時，我能標示未知、停止推測並依正式路徑查證。",
    "案例與服務紀錄、需求確認及交接情境相關。",
    "我知道只記錄完成交接所需的最少資料。",
    "我知道如何把未解問題交給下一個適當責任角色處理。",
]
LEARNER_ROWS = [
    [2, 2, 1, 2, 3, 3, 2],
    [3, 3, 2, 3, 3, 4, 3],
    [3, 3, 2, 3, 4, 4, 3],
    [4, 4, 3, 3, 4, 4, 4],
    [4, 4, 3, 4, 4, 4, 4],
    [4, 4, 3, 4, 4, 5, 4],
    [5, 5, 5, 4, 5, 5, 4],
    [5, 5, 4, 5, 5, 5, 5],
]
GRAD_Q = [
    "我能在來源不明時辨認應交由哪一責任角色查證。",
    "我能在紀錄及交接中採取最小必要資料原則。",
    "我能把未解事項、下一步及接手責任記錄清楚。",
]
GRAD_ROWS = [
    {"id": "G01", "used": "是", "reason": "不適用－曾使用", "ratings": [2, 3, 2], "needs": ["快速參考卡", "情境角色演練"]},
    {"id": "G02", "used": "是", "reason": "不適用－曾使用", "ratings": [3, 4, 3], "needs": ["快速參考卡", "來源查證情境"]},
    {"id": "G03", "used": "是", "reason": "不適用－曾使用", "ratings": [3, 4, 3], "needs": ["快速參考卡", "情境角色演練"]},
    {"id": "G04", "used": "否", "reason": "未遇到相符任務", "ratings": [4, 4, 4], "needs": ["快速參考卡", "情境角色演練", "來源查證情境"]},
    {"id": "G05", "used": "否", "reason": "尚無使用機會", "ratings": [4, 5, 5], "needs": ["目前不需要"]},
]
NEED_OPTIONS = ["快速參考卡", "情境角色演練", "來源查證情境", "目前不需要"]
LIKERT = ["非常不同意", "不同意", "普通／不確定", "同意", "非常同意"]

WEIGHTS = [30, 25, 20, 15, 10]
CRITERIA = [
    "與學員／結訓意見配適",
    "能增加可觀察練習",
    "時間與學員負擔",
    "師資／教材可行性",
    "修訂與追蹤可留證",
]
OPTIONS = {
    "A｜同時數增練與工具": [4, 4, 5, 4, 5],
    "B｜由 G3 移時數給 G1": [5, 4, 3, 4, 5],
    "C｜總時數增加 15 分": [5, 5, 2, 4, 5],
}
ROLES = {
    "P1": "課程規劃承辦（建議功能代碼；協會實際職務與人選 UNKNOWN）",
    "P2": "資料／隱私檢核者（建議功能代碼；任命與權限 UNKNOWN）",
    "P3": "獨立覆核者（建議功能代碼；實際指派 UNKNOWN）",
    "P4": "有權核定者（必須由章程或有效授權原件確認；目前 UNKNOWN）",
    "P5": "授課／演練引導者（建議功能代碼；實際師資 UNKNOWN）",
}


def sha(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def ensure_dirs() -> None:
    for d in [DOCX.parent, REGISTER.parent, SOURCE_MAP.parent, PKG / "evidence" / "build"]:
        d.mkdir(parents=True, exist_ok=True)


def load_blueprint() -> dict:
    source = RUNTIME / "CONTROL" / "BLUEPRINT_CURRENT.jsonl"
    rows = []
    with source.open("r", encoding="utf-8-sig") as f:
        for line in f:
            try:
                obj = json.loads(line)
            except json.JSONDecodeError:
                continue
            if str(obj.get("deliverable_id", obj.get("doc_id", ""))).zfill(4) == "0052":
                rows.append((line.rstrip("\r\n"), obj))
    if len(rows) != 1:
        raise RuntimeError(f"Expected one current blueprint row for 0052; found {len(rows)}")
    raw, obj = rows[0]
    BLUEPRINT_COPY.write_text(raw + "\n", encoding="utf-8")
    return obj


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = margins.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            margins.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def set_cell_width(cell, width_cm: float) -> None:
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def set_run_font(run, size=10.5, bold=False, color="243746") -> None:
    run.font.name = "Microsoft JhengHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_text(cell, text: str, *, header=False, size=9.5, color=None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(str(text))
    set_run_font(r, size=size, bold=header, color=color or ("FFFFFF" if header else "243746"))
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    if header:
        set_cell_shading(cell, "164E63")


def set_table_borders(table, color="A9BBC4", sz="5") -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def add_table(doc, headers, rows, widths=None, *, size=9.5, highlight_first=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    if widths is None:
        widths = [17.0 / len(headers)] * len(headers)
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, header=True, size=size)
        set_cell_width(table.rows[0].cells[i], widths[i])
    repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for ci, value in enumerate(row):
            set_cell_text(cells[ci], value, size=size)
            set_cell_width(cells[ci], widths[ci])
            if highlight_first and ci == 0:
                set_cell_shading(cells[ci], "E8F2F4")
        if ri % 2 == 1:
            for cell in cells:
                if cell._tc.tcPr.find(qn("w:shd")) is None:
                    set_cell_shading(cell, "F6F9FA")
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_para(doc, text="", *, size=10.5, bold=False, color="243746", align=None, italic=False, after=5):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)
    r.italic = italic
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    set_run_font(r, 10.5)
    return p


def add_h(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(9 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, 15 if level == 1 else 12, bold=True, color="164E63" if level == 1 else "0F6B78")
    return p


def add_callout(doc, text, fill="FFF2CC", size=10.5):
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    t.columns[0].width = Cm(17)
    c = t.cell(0, 0)
    set_cell_text(c, text, size=size)
    set_cell_shading(c, fill)
    set_cell_margins(c, top=130, start=150, bottom=130, end=150)
    set_table_borders(t, color="D8B85A", sz="8")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = paragraph.add_run("第 ")
    set_run_font(r, 8.5, color="617783")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    r2 = paragraph.add_run(" 頁")
    set_run_font(r2, 8.5, color="617783")


facts: list[dict] = []


def fact(fid, value, category, locator, source, replace_with, owner, note=""):
    facts.append({
        "fact_id": fid,
        "sample_value": str(value),
        "fact_category": category,
        "exact_document_locator": locator,
        "source_status": source,
        "required_real_replacement": replace_with,
        "replacement_responsible_role": owner,
        "notes": note,
    })


def add_base_facts():
    base = [
        ("CASE-ID", CASE, "case identifier", "Cover; §§7-10", "SYNTHETIC", "verified source case ID and permitted linkage", "P1", "Same invented case identifier used by 0033 R03, 0049 R01 and 0050 R01."),
        ("COURSE-NAME", COURSE, "course title", "Cover; §§7-10", "SYNTHETIC", "authorized current course plan and version", "P1", "No real course delivery is established."),
        ("CLASS-DATE", "2026-08-28", "date", "§§7,10", "SYNTHETIC", "actual approved roster/course record date", "P1", "Illustrative course date only."),
        ("LEARNER-INVITED", 12, "denominator", "§6", "SYNTHETIC", "actual number invited/eligible with roster source", "P1", "Scenario denominator, not attendance evidence."),
        ("LEARNER-VALID", 8, "denominator", "§6", "SYNTHETIC", "actual valid anonymous feedback forms", "P1", "Eight complete synthetic rows."),
        ("GRAD-FOLLOWUP-DATE", "2026-09-27", "date", "§§7,8", "SYNTHETIC", "actual dated follow-up records", "P1", "Thirty-day illustrative follow-up."),
        ("ANALYSIS-DATE", "2026-09-30", "date", "§§8,12", "SYNTHETIC", "actual analysis record date", "P1", "Illustrative only."),
        ("PLAN-REVISION-DATE", "2026-10-02", "date", "§§9,10,12", "SYNTHETIC", "authorized controlled plan revision record", "P1/P4", "Draft proposal date; not approval."),
        ("LEARNER-RESPONSE-RATE", "8/12 = 66.7%", "derived proportion", "§6", "DERIVED FROM SYNTHETIC", "recalculate from actual invite and valid-form logs", "P3", "Rounded to one decimal; denominator is synthetic invite count."),
        ("GRAD-RESPONSE-RATE", "5/12 = 41.7%", "derived proportion", "§8", "DERIVED FROM SYNTHETIC", "recalculate from actual graduate invite and valid response logs", "P3", "Multi-select needs are not mutually exclusive."),
        ("TRIGGER-RULE", "mean < 3.5/5 or any privacy/authorization concern triggers review", "sample-only rule", "§6", "SAMPLE POLICY EXAMPLE", "replace with formally authorized local decision rule", "P4", "Not an official TTQS minimum or association policy."),
        ("ROLE-P1", ROLES["P1"], "proposed function", "§§3,10,12", "PROPOSED; NOT CURRENT-STATE FACT", "verify actual assigned position and authority", "P4", "No real person or office is named."),
        ("ROLE-P2", ROLES["P2"], "proposed function", "§§3,10,12", "PROPOSED; NOT CURRENT-STATE FACT", "verify actual privacy/data custodian role", "P4", "No appointment is asserted."),
        ("ROLE-P3", ROLES["P3"], "proposed function", "§§3,10,12", "PROPOSED; NOT CURRENT-STATE FACT", "verify actual second-reviewer assignment", "P4", "No review event is asserted as real."),
        ("ROLE-P4", ROLES["P4"], "authority boundary", "§§3,10,12", "UNKNOWN; GATE NOT PASSED", "obtain governing authority/delegation and actual approval record", "P4", "No formal approval or signatory is fabricated."),
        ("ROLE-P5", ROLES["P5"], "proposed function", "§§3,10,12", "PROPOSED; NOT CURRENT-STATE FACT", "verify actual instructor and assignment", "P4", "No real instructor is named."),
        ("COURSE-DURATION", "90+90+120+60 = 360 teaching minutes", "inherited synthetic plan figure", "§§2,9,10", "SYNTHETIC; LINKED TO 0049/0050", "verify approved course plan and actual agenda", "P1/P3", "Breaks/admin excluded; not a delivered course record."),
        ("PLANNING-THRESHOLD-Q3", "learner Q3 mean >= 3.5/5", "proposed sample target", "§10", "SAMPLE POLICY EXAMPLE", "approve or replace with authorized local target", "P4", "Not an official TTQS cutoff."),
    ]
    for row in base:
        fact(*row)


def aggregate(values):
    n = len(values)
    total = sum(values)
    counts = [values.count(i) for i in range(1, 6)]
    return n, total, total / n if n else None, counts


def build_register():
    add_base_facts()
    for ri, row in enumerate(LEARNER_ROWS, 1):
        for qi, value in enumerate(row, 1):
            fact(f"SYN-0052-L{ri:02d}-Q{qi:02d}", value, "learner rating", f"§5, Appendix A, row L{ri:02d}, Q{qi:02d}", "SYNTHETIC RESPONSE", "actual anonymized response from authorized course feedback instrument", "P1/P2", "No name or personal identifier is included." )
    for gi, row in enumerate(GRAD_ROWS, 1):
        fact(f"SYN-0052-G{gi:02d}-USE", row["used"], "graduate self-report", f"§7, Appendix B, row {row['id']}, use", "SYNTHETIC RESPONSE", "actual dated follow-up response with consent/authority basis", "P1/P2", "Self-report is not independently verified behavior." )
        fact(f"SYN-0052-G{gi:02d}-USE-REASON", row["reason"], "graduate use/reason response", f"§7, Appendix B, row {row['id']}, use reason", "SYNTHETIC RESPONSE", "actual response using the approved reason choices", "P1/P2", "For yes, reason is structurally not applicable; for no, the synthetic example records one reason." )
        for qi, value in enumerate(row["ratings"], 3):
            fact(f"SYN-0052-G{gi:02d}-Q{qi:02d}", value, "graduate rating", f"§7, Appendix B, row {row['id']}, Q{qi:02d}", "SYNTHETIC RESPONSE", "actual anonymized follow-up response", "P1/P2", "Rating uses the same five-point response scale." )
        for ni, option in enumerate(NEED_OPTIONS, 1):
            selected = "是" if option in row["needs"] else "否"
            fact(f"SYN-0052-G{gi:02d}-NEED{ni:02d}", selected, "graduate multi-select", f"§7, Appendix B, row {row['id']}, {option}", "SYNTHETIC RESPONSE", "actual multi-select response", "P1/P2", "The three support needs may overlap; '目前不需要' is exclusive." )
    # Register every reported item aggregate and weighted comparison as a derived fact.
    for qi, q in enumerate(LEARNER_Q, 1):
        values = [r[qi - 1] for r in LEARNER_ROWS]
        n, total, mean, counts = aggregate(values)
        fact(f"SYN-0052-LQ{qi:02d}-N", n, "derived denominator", f"§6, learner summary Q{qi:02d}", "DERIVED", "recalculate from actual valid responses", "P3")
        fact(f"SYN-0052-LQ{qi:02d}-SUM", total, "derived sum", f"§6, learner summary Q{qi:02d}", "DERIVED", "recalculate from actual response values", "P3")
        fact(f"SYN-0052-LQ{qi:02d}-MEAN", f"{mean:.3f}", "derived mean", f"§6, learner summary Q{qi:02d}", "DERIVED", "recalculate sum/valid N; report rounding", "P3")
        fact(f"SYN-0052-LQ{qi:02d}-COUNTS", "/".join(map(str, counts)), "derived distribution", f"§6, learner summary Q{qi:02d}", "DERIVED", "recount actual ratings 1 through 5", "P3")
    for qi in range(3):
        vals = [r["ratings"][qi] for r in GRAD_ROWS]
        n, total, mean, counts = aggregate(vals)
        fact(f"SYN-0052-GQ{qi+3:02d}-N", n, "derived denominator", f"§8, graduate summary GQ{qi+3:02d}", "DERIVED", "recalculate from actual valid responses", "P3")
        fact(f"SYN-0052-GQ{qi+3:02d}-SUM", total, "derived sum", f"§8, graduate summary GQ{qi+3:02d}", "DERIVED", "recalculate from actual response values", "P3")
        fact(f"SYN-0052-GQ{qi+3:02d}-MEAN", f"{mean:.3f}", "derived mean", f"§8, graduate summary GQ{qi+3:02d}", "DERIVED", "recalculate sum/valid N; report rounding", "P3")
        fact(f"SYN-0052-GQ{qi+3:02d}-COUNTS", "/".join(map(str, counts)), "derived distribution", f"§8, graduate summary GQ{qi+3:02d}", "DERIVED", "recount actual ratings 1 through 5", "P3")
    for ni, option in enumerate(NEED_OPTIONS, 1):
        n = sum(option in r["needs"] for r in GRAD_ROWS)
        fact(f"SYN-0052-GNEED{ni:02d}-N", n, "derived multi-select count", f"§8, graduate needs", "DERIVED", "recount actual selections; denominator is valid respondents", "P3")
        fact(f"SYN-0052-GNEED{ni:02d}-PCT", f"{n}/5 = {n/5*100:.1f}%", "derived multi-select proportion", f"§8, graduate needs", "DERIVED", "recalculate; do not sum percentages", "P3")
    for oi, (name, scores) in enumerate(OPTIONS.items(), 1):
        products = [w * s for w, s in zip(WEIGHTS, scores)]
        total = sum(products)
        fact(f"SYN-0052-OPT{oi}-SCORES", "/".join(map(str, scores)), "synthetic option ratings", f"§9, option matrix {name}", "SYNTHETIC JUDGMENT", "obtain documented actual decision criteria/scores", "P1/P3", "Ratings are illustrative, not measured outcomes." )
        fact(f"SYN-0052-OPT{oi}-PRODUCTS", "+".join(map(str, products)), "derived weighted products", f"§9, option matrix {name}", "DERIVED", "recalculate each weight x score", "P3")
        fact(f"SYN-0052-OPT{oi}-TOTAL", total, "derived weighted total", f"§9, option matrix {name}", "DERIVED", "recalculate; weights must total 100", "P3")
        fact(f"SYN-0052-OPT{oi}-MEAN", f"{total/100:.2f}", "derived weighted mean", f"§9, option matrix {name}", "DERIVED", "recalculate total/100", "P3")
    for i, (weight, name) in enumerate(zip(WEIGHTS, CRITERIA), 1):
        fact(f"SYN-0052-W{i}", f"{weight}%｜{name}", "synthetic decision weight", "§9, criteria", "SYNTHETIC JUDGMENT", "authorize criteria and weights before real use", "P4")
    fact("SYN-0052-CHOICE", "方案 A 暫列示範選擇；待正式權限確認", "synthetic proposed choice", "§9 and §12 process trail", "SYNTHETIC; NOT APPROVED", "actual authorized documented decision", "P4", "No real approval is claimed." )
    fact("SYN-0052-DURATION-SUM", "90+90+120+60=360", "derived duration sum", "§9 course plan", "DERIVED FROM LINKED SYNTHETIC PLAN", "verify minutes against current approved plan", "P3", "No additional time is added in sample option A." )

    headers = ["fact_id", "sample_value", "fact_category", "exact_document_locator", "source_status", "required_real_replacement", "replacement_responsible_role", "notes"]
    with REGISTER.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=headers)
        w.writeheader()
        w.writerows(facts)


def setup_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.7)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
    sec.header_distance = Cm(0.8)
    sec.footer_distance = Cm(0.8)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft JhengHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("243746")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12
    for name, sz, col in [("Heading 1", 15, "164E63"), ("Heading 2", 12, "0F6B78"), ("Heading 3", 11, "164E63")]:
        st = styles[name]
        st.font.name = "Microsoft JhengHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        st.font.size = Pt(sz)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(col)
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    rr = header.add_run("原始點協會｜TTQS 候選工作文件｜I11-E04")
    set_run_font(rr, 8.5, color="617783")
    footer = sec.footer.paragraphs[0]
    rr = footer.add_run("SAMPLE / SYNTHETIC 示範｜候選未核定｜0052 R01     ")
    set_run_font(rr, 8.5, color="9A4A00")
    add_page_number(footer)
    return doc


def build_doc(blueprint):
    doc = setup_doc()
    props = doc.core_properties
    props.title = "課程意見回流與課程規劃閉環｜0052 R01 SAMPLE/SYNTHETIC"
    props.subject = "I11-E04／官方第054方框；策略／計畫文件候選"
    props.author = "TTQS ONE 候選文件建置"
    props.keywords = "SAMPLE; SYNTHETIC; I11-E04; 0052; course feedback plan"

    add_para(doc, "課程意見回流與課程規劃閉環", size=24, bold=True, color="164E63", align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_para(doc, "依學員及結訓學員意見形成可追溯的課程修訂計畫", size=14, color="0F6B78", align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
    add_para(doc, "DOCX 0052｜Unique Blueprint I11-E04｜官方勾選流水號 054｜官方資料表第 4 頁", size=10.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    add_callout(doc, "本文件為 SAMPLE / SYNTHETIC 完整示範，並非社團法人屏東縣原始點關懷協會的調查、上課、結訓追蹤、會議、核准或實績紀錄。文件可供承辦者複製流程及表單；正式使用前必須以可核實原件替換本文件登錄的每項合成事實，並由真正有權者核定。", "FFF0CC", 11)
    add_h(doc, "30 秒使用說明", 1)
    add_table(doc, ["讀者／使用時點", "操作", "交付與狀態"], [
        ["課程規劃承辦、資料檢核者及授權覆核者；課程後回收學員意見、結訓約 30 日追蹤後使用。", "依第 2 節流程收集與核驗；照第 4、7 節問卷及量尺取得資料；按第 6、8 節列分母、公式及限制；依第 9、10 節比較方案並提出下一版課程計畫。", "輸出是具來源、取捨、工作責任、期限、指標及覆核欄位的課程修訂提案。此 R01 只是一個合成候選，核定權限 UNKNOWN，沒有 REAL 批准。"]
    ], [4.0, 7.0, 6.0], size=9.5)
    add_h(doc, "本件對應與範圍", 1)
    add_para(doc, "官方原文：" + OFFICIAL_TEXT, size=11, bold=True, color="164E63")
    add_para(doc, "本件要證明的是：課程計畫如何把在訓學員與結訓學員意見，連到資料檢核、分析取捨、課程設計修訂、授權邊界及下一輪驗證。它不是需求調查報告（0033／I07-E01／官方流水號 033），也不重做課程需求與計畫（0049／I11-E01／官方流水號 051）；0051／I11-E03／官方流水號 053 是課程產出討論與確認紀錄，本件不代替該件的會議或核准原件。0053／I12-E01 則是另一項課程切合性檢討。這些文件各自對應不同官方工作物。")
    add_table(doc, ["本件", "同一案例的上游／鄰件", "不可互相替代的差異"], [
        ["0052｜I11-E04｜checkbox 054｜課程意見回流計畫", "0033 R03 I07-E01（需求調查）；0049 R01 I11-E01（課程需求與計畫）；0050 R01 I11-E02（單元設計與評量）", "本件新增課後意見及結訓追蹤、明示分母與方案取捨、提出下一版課程計畫；不重複貼入上游問卷、單元教案或假會議核准。"],
        ["負向鄰項 0051｜I11-E03", "課程產出經討論確認與留存", "會議／產出確認的原始紀錄另有其件；本件不代簽、不替代決議原件。"],
        ["負向鄰項 0053｜I12-E01", "訓練課程切合性檢討與調整", "課程適配檢討需依其獨特 Blueprint；本件只處理學員／結訓學員意見如何更新課程規劃。"]
    ], [4.1, 5.5, 7.4], size=9.5)

    add_h(doc, "1｜來源、決策邊界與合成案例", 1)
    add_para(doc, "官方要求依官方訓練品質評核資料表 PDF 第 4 頁第 11 項第 4 方框核對；流水號是 054。DOCX ID 0052 與 checkbox 054 是兩套不同編號。本件以正式 Unique Blueprint 的 doc_id 0052／I11-E04／genre strategy_plan 為唯一文件身份。")
    add_table(doc, ["來源／版本", "定位與核對值", "可支持的結論／限制"], [
        ["官方評核資料表 PDF", "第 4 頁；PDF SHA-256 805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709", "確認官方原文及方框位置；不證明協會已辦訓。"],
        ["官方方框頁影像", "inputs/official_evidence_table_page4.png；SHA-256 0e00b46586ad0ee181046d4a13d0023654bb07537f7eebc3429dad3bfa29d35f", "可直接目視定位第 054 方框。"],
        ["目前 Unique Blueprint", "BLUEPRINT_CURRENT.jsonl；完整來源檔 SHA-256 406d0495365d0970dcf8113a42842663b8a54372a4e204cc492b7c3f8156ef3e；本件原始列另存 inputs。", "定義策略／計畫文類、專屬章節、必備表格、逐事實替換簿、負向鄰項及驗收條件。"],
        ["正式寫作手冊", "02_GENRE_WRITING_MANUAL.md｜指標 11：結訓意見→課程變更→驗證方式，形成時間一致閉環。", "說明文類與鄰件邊界；不替代官方要求。"],
        ["合成案例上游", "0033 R03 SHA fe5f3cfd84c761da86435a71e0fd899daa725be78628c4d2d3f2324c1fdb5f4f；0049 R01 SHA e8212788b1f7b201ca4af4f4c5fb83ab131627dbfa815987f7cf6344edfda727；0050 R01 SHA d0d37c2c69b161a2cdd37ee3097306722c4a65a1afd844ae92e1dff58b43381a", "只承接同一個 SAMPLE 案例 SYN-0033-PLAN-01、課程名稱及四單元 360 教學分鐘；其數據與課程都不是協會 REAL 活動。"],
    ], [3.2, 6.6, 7.2], size=9.5)
    add_callout(doc, "真實資料查核結果：目前沒有找到本要求專屬、可驗證的協會學員意見原件、結訓學員追蹤原件、已核定之課程修訂紀錄、真實學員名冊或授權核定紀錄。故所有問卷回覆、日期、角色、分母、取捨及計畫修訂均明標合成；實際權責、保存期限、課程實施與核定狀態為 UNKNOWN。", "FCE4D6", 10)

    add_h(doc, "2｜計畫的運作流程", 1)
    add_para(doc, "流程每步均要留可讀回的輸出。若來源、資料權限、核准者或真實課程版本不明，承辦者只可記錄未知與待取得項目，不得把示範值轉成正式事實。")
    add_table(doc, ["步驟", "輸入／動作", "檢核與停止條件", "留存輸出"], [
        ["1 啟動", "課程完成後發出學員意見工具；結訓後約 30 日發出追蹤工具。", "先核對已核定課程版本、適用對象、資料權限、告知方式；未確認就不發送 REAL 問卷。", "版本、適用名單來源、邀請分母、發送及回收時間。"],
        ["2 去識別與有效性", "以無法直接識別個人的代碼收件；依預先定義的完整性規則標記有效／缺答。", "不收不必要姓名、身分證、健康或個案敏感資料；疑似個資外洩立即停用副本並走正式事件流程。", "匿名代碼、有效／缺答判定及第二人覆核。"],
        ["3 分開分析兩種聲音", "在訓學員回饋反映課程體驗；結訓學員追蹤提供後續使用自陳。", "不得把自陳等同觀察到的工作成效；清楚報 N、分母、缺答、題目與時間點。", "分題頻數、分母、平均數、開放意見分類及限制。"],
        ["4 設定議題", "依已核定的內部觸發規則標出需檢討的題目及安全／授權疑慮。", "本文件 3.5/5 只為合成規則，不是官方門檻；任何隱私或越權疑慮優先處理。", "議題清單、證據定位、不能下的結論。"],
        ["5 比較方案", "至少比較保留現制、同時數重新配置、增加時數等不同可選方案。", "明列權重、分數、理由、資源和限制；第二人重算；分差小不得說成統計上的優越。", "可重算的加權表及異議／假設。"],
        ["6 擬定計畫", "把選項寫成單元、時數、方法、評量、資源、責任、風險、里程碑及檢核。", "核定角色與權限必須由真實章程／授權證明；示例代碼不能核准。", "版本化的課程計畫草案與變更對照。"],
        ["7 核定後執行", "只有真正有權者書面核定後，才可轉為有效課程版並發出教材。", "未取得核定原件時維持草案，不得排課、公告或稱正式實施。", "授權原件、正式版本、教材清單及發行紀錄。"],
        ["8 下一輪驗證", "於下一期課程收相同核心題、學習評量及 30 日追蹤，與本次基準並列。", "保留相同題目及算法，報樣本差異；不得由前後差異逕稱因果。", "前後版本、原始匿名列、重算結果及再修訂決定。"],
    ], [1.7, 5.3, 6.1, 3.9], size=9.5)

    p = add_h(doc, "3｜責任與權限配置（待真實授權替換）", 1)
    p.paragraph_format.page_break_before = True
    add_table(doc, ["功能代碼", "建議責任", "工作輸出", "權限狀態"], [[k, v, "依第 2 節步驟及第 10 節時程執行；署名或任命前須核實。", "目前 UNKNOWN；不可視為現行組織職稱或已指派人員。"] for k, v in ROLES.items()], [1.5, 5.2, 5.4, 4.9], size=9.5)
    add_para(doc, "建議最小分工：P1 不得單獨核定自己的課程修訂；P2 檢查個資與去識別；P3 重算分母和加權表、核對跨文件版本；P4 必須具實際授權；P5 按已核定計畫授課並保留評量。若協會現況沒有這些職務，不可在文件中寫成既有委員會或制度。")

    doc.add_page_break()
    add_h(doc, "4｜學員課後意見工具（可直接改作正式問卷）", 1)
    add_para(doc, "示範發送點：課程結束後；填答約 3–5 分鐘；不要求姓名。使用前核對告知、保存與刪除政策。採 1–5 單選量尺，1=非常不同意、2=不同意、3=普通／不確定、4=同意、5=非常同意；另設「不適用／不願回答」為缺答代碼，不納入平均，必須報數。")
    add_table(doc, ["代碼", "可直接使用題目", "對應課程設計問題", "資料型態"], [[f"Q{i:02d}", q, link, "1–5 單選；不適用另記 NA"] for i, (q, link) in enumerate(zip(LEARNER_Q, ["角色／任務配適", "目標與邊界", "練習時間", "未知與查證", "案例關聯", "資料最小化", "下一責任角色" ]), 1)], [1.2, 8.4, 4.0, 3.4], size=9.5)
    add_h(doc, "問卷操作欄位", 2)
    add_table(doc, ["欄位", "填寫／驗證規則"], [
        ["課程版本／日期", "由發送者從已核定版本登錄；本示範 2026-08-28 為 SYNTHETIC。"],
        ["邀請與回收分母", "受邀數、收到數、有效數、不完整數分開計；不得把回收率寫成滿意度。"],
        ["匿名代碼", "使用隨機碼；不得保留可回推個人的對照表於公開資料夾。"],
        ["開放意見（選填）", "提示不得輸入姓名、服務對象、病史或其他可識別個資；彙整時先去識別。"],
        ["缺答處理", "每題列有效 n 及缺答 n；平均只用有效答案，不用 0 補值。"],
    ], [4.0, 13.0], size=9.5)

    add_h(doc, "5｜學員完整合成資料列", 1)
    add_callout(doc, "以下 8 列為逐人合成答卷，並非真人。案例假設 12 個邀請名額、8 份有效完整答卷；這只用來示範分母與重算，不能證明 12 人實際上課或 8 人實際回覆。", "FFF0CC", 10)
    learner_rows = [[f"L{i:02d}"] + row for i, row in enumerate(LEARNER_ROWS, 1)]
    add_table(doc, ["代碼", "Q1", "Q2", "Q3", "Q4", "Q5", "Q6", "Q7"], learner_rows, [1.3] + [2.24] * 7, size=9.5)
    add_para(doc, "完整計算示例：Q3 練習充分度的 8 個分數為 1,2,2,3,3,3,5,4；合計 23，平均=23÷8=2.875，顯示到小數三位。五級頻數由低至高為 1/2/3/1/1。不可把平均 2.875 說成 287.5%，也不可把頻數當成百分比。", size=10.5, bold=True)
    add_h(doc, "6｜學員意見彙總與判讀", 1)
    summary_labels = ["角色／任務配適", "目標／邊界清楚", "引導練習充分", "未知／查證路徑", "案例相關性", "最小必要資料", "未解事項責任路徑"]
    learner_summary = []
    for i, q in enumerate(LEARNER_Q, 1):
        vals = [r[i - 1] for r in LEARNER_ROWS]
        n, total, mean, counts = aggregate(vals)
        learner_summary.append([f"Q{i:02d}", summary_labels[i-1], n, total, f"{mean:.3f}", "/".join(map(str, counts))])
    add_table(doc, ["題號", "題目簡稱", "有效 n", "分數合計", "平均／5", "頻數 1/2/3/4/5"], learner_summary, [1.1, 7.4, 1.4, 1.6, 1.8, 3.7], size=9.5)
    add_para(doc, "樣本回覆率（示範）=8 份有效答卷÷12 個邀請名額×100=66.7%；若實際邀請分母不同，必須重算。 Q3 平均 2.875 低於本例預先設定的 3.5 觸發線，列為優先改進；Q4 恰為 3.500，不低於觸發線，但不能因此忽視個別安全意見。Q5 4.000、Q6 4.250 是合成的感受評分，不代表客觀行為或隱私合規已通過。")
    add_para(doc, "小樣本僅做描述性整理。沒有母體抽樣框、代表性設計或推論檢定，不對所有志工、全協會或未來班級推論。此數據也不能證明課程造成工作表現變化。任何可識別開放留言需另外去識別並依真實保存規則處理。")

    add_h(doc, "7｜結訓學員 30 日追蹤工具與完整合成資料", 1)
    add_para(doc, "追蹤工具以課程結束日為 D0，在約 D+30 發送。請受訪者先回覆是否有機會在相關任務中使用課程方法；沒有遇到相符任務時選「沒有機會／不適用」，不得當成不會使用。若答否或沒有機會，再選一項原因。後續三題沿用 1–5 量尺；三項支持需求可複選，「目前不需要」為互斥選項。")
    add_table(doc, ["欄位", "可直接使用題目／選項", "計數規則"], [
        ["G1 使用情形", "過去約 30 日是否曾在相符任務使用課程方法？是／否／沒有機會／不願回答。若答否或沒有機會，原因選一：未遇到相符任務／尚無使用機會／其他（不寫個資）。", "逐項報是、否、沒有機會、缺答；只有是/否者可計入已使用比例，分母要註明。合成資料逐列呈現原因。"],
        ["G2 後續所需支持（複選）", "快速參考卡；再一次情境角色演練；來源查證情境；其他（去識別文字）；目前不需要（選此項不再勾其他支持）。", "三項支持可重疊，逐項算人數／有效追蹤者；「目前不需要」為互斥選項。"],
        ["G3 未知來源責任路徑", GRAD_Q[0], "1–5 單選；NA 為缺答。"],
        ["G4 最小必要資料", GRAD_Q[1], "1–5 單選；感受自陳，不等於紀錄抽查。"],
        ["G5 未解事項紀錄", GRAD_Q[2], "1–5 單選；感受自陳，不等於真實工作紀錄查核。"],
    ], [2.0, 10.0, 5.0], size=9.5)
    grad_table = []
    for row in GRAD_ROWS:
        grad_table.append([row["id"], f"{row['used']}／{row['reason']}", *row["ratings"], "、".join(row["needs"]) if row["needs"] else "未勾支持需求"])
    add_table(doc, ["代碼", "G1 使用／原因", "G3", "G4", "G5", "G2 支持需求"], grad_table, [1.1, 3.3, 1.1, 1.1, 1.1, 9.3], size=9.5)
    add_para(doc, "合成追蹤受邀分母假設 12 人，5 份完整有效追蹤；5÷12×100=41.7%。其中 3/5=60% 自陳曾使用，2/5 自陳未使用；沒有外部工作紀錄可驗證其陳述，亦不推論未使用原因。")
    add_h(doc, "8｜結訓意見分析及與課程修訂的連結", 1)
    grad_summary = []
    for qi, label in enumerate(GRAD_Q, 3):
        values = [r["ratings"][qi - 3] for r in GRAD_ROWS]
        n, total, mean, counts = aggregate(values)
        grad_summary.append([f"G{qi}", label, n, total, f"{mean:.3f}", "/".join(map(str, counts))])
    add_table(doc, ["題號", "追蹤題目", "n", "合計", "平均／5", "頻數 1/2/3/4/5"], grad_summary, [1.1, 7.5, 1.2, 1.4, 1.7, 4.2], size=9.5)
    needs_rows = []
    for label in NEED_OPTIONS:
        count = sum(label in r["needs"] for r in GRAD_ROWS)
        read_rule = "互斥；選此項不再勾支持" if label == "目前不需要" else "複選；可與其他支持重疊"
        needs_rows.append([label, count, 5, f"{count}/5 = {count/5*100:.1f}%", read_rule])
    add_table(doc, ["支持需求", "選取人數", "有效 n", "比例算式", "讀法"], needs_rows, [4.0, 2.0, 2.0, 3.2, 5.8], size=9.5)
    add_para(doc, "合成訊號：G3=16÷5=3.200、G5=17÷5=3.400，低於本例 3.5 觸發線；4/5=80% 勾選快速參考卡，3/5 勾選角色演練，2/5 勾選來源查證情境。三項可重疊支持需求共 9 次勾選；另 1/5=20% 選「目前不需要」，此項與其他支持互斥。不可把多選比例相加當人口比例。G4=20÷5=4.000 只表示樣本自評較高，仍不能推翻個資風險檢核。")
    add_table(doc, ["訊號", "課程計畫調整", "如何驗證，而非只宣稱已改善"], [
        ["Q3 練習充分度 2.875", "提案在單元 G1 的 90 分鐘內增加 15 分鐘引導角色演練；不增加總教學時數。", "下一期沿用 Q3；另按 0050 R01 的 G1 5 題至少 4 題且無越權承諾的示範量尺記錄分母與結果。"],
        ["G3 / G5 3.200 / 3.400；4/5 要快速卡", "製作一頁式交接參考卡，讓「未知／來源待查／下一責任角色／下一步／期限」清楚可辨；嵌入 G2、G3 的既有練習，不額外加時。", "下一期用 G3 五欄完整度及 G4 整合檢核；追蹤問卷重問責任路徑與未解事項。"],
        ["Q6 自評 4.250；仍須安全控制", "卡片不加收個資；由 P2 對欄位、教材例子及作業答案作資料最小化檢查。", "任何不必要敏感資訊或越權處理為立即例外，不能被平均分數抵銷。"],
    ], [4.2, 6.5, 6.3], size=9.5)

    add_h(doc, "9｜方案比較與示範取捨", 1)
    add_para(doc, "比較尺度為 1 至 5 分，5 分表示依本項準則最有利。權重及分數均為合成規劃判斷，不是受訪者量測或正式協會決策。每一格加權分 = 準則權重 × 方案評分；總分除以 100 得加權平均。")
    add_table(doc, ["準則", "權重", "評分錨點（摘要）"], [[CRITERIA[i], f"{WEIGHTS[i]}%", ["是否回應學員／結訓明確訊號", "是否增可觀察演練並保留評量", "是否維持原計畫總負擔", "教材、師資與現有單元能否承接", "是否留版本、責任及下一輪核對"] [i]] for i in range(5)], [4.3, 1.5, 11.2], size=9.5)
    option_rows = []
    for name, scores in OPTIONS.items():
        products = [w * s for w, s in zip(WEIGHTS, scores)]
        total = sum(products)
        option_rows.append([name] + scores + [" + ".join(map(str, products)) + f" = {total}", f"{total}/100 = {total/100:.2f}"])
    add_table(doc, ["方案", "F1\n30%", "F2\n25%", "F3\n20%", "F4\n15%", "F5\n10%", "加權分 (w×s)", "平均／5"], option_rows, [2.5, 1.1, 1.1, 1.1, 1.1, 1.1, 5.8, 2.0], size=9.5)
    add_table(doc, ["方案", "內容", "限制與依據"], [
        ["A｜4.30", "維持 360 分鐘；在 G1 原 90 分鐘內把提案分段由合成基準 35/20/35 改為 20/20/50（說明/示範/引導練習），增加 15 分鐘演練；G2/G3 加入快速參考卡使用，不加總時數。", "最高分，但與 C 僅差 0.05，屬主觀細差；35/20/35 是本件合成比較基準，不是 0049/0050 已存在的內部分鐘紀錄。"],
        ["B｜4.20", "把 15 分鐘由 G3 挪到 G1，G1 105 分、G3 105 分，G2/G4 不變，總時數仍 360。", "較直接增練，但壓縮 G3 的最小必要紀錄練習；可能與學員想要快速工具的訊號不如 A 平衡。"],
        ["C｜4.25", "在原 360 分鐘上新增 15 分鐘，總教學時數 375 分鐘。", "練習時間足，但增加學員與師資負擔；場地、志工時間與預算尚無 REAL 證據。"],
    ], [2.0, 9.0, 7.0], size=9.5)
    add_callout(doc, "示範暫選 A：加權值 430/100=4.30，高於 C 的 4.25 與 B 的 4.20。差距只有 0.05–0.10，且評分是合成判斷；這只展示如何記錄取捨，不表示真正有權機關已選定 A，也不表示統計上 A 較優。正式採用須由 P4 依真實授權決定。", "E2F0D9", 10)

    add_h(doc, "10｜課程修訂提案與執行計畫（草案，待授權）", 1)
    add_table(doc, ["單元／合成時數", "維持的既有目標（連結 0050）", "本件提案變更", "評量與預定留存"], [
        ["G1｜90 分", "辨認越出受訓角色範圍的請求，停止越權承諾並按正式路徑交接。", "合成分段提案 20 分重點說明+20 分示範+50 分引導演練；對原合成比較基準把 15 分說明轉為練習。", "五個情境至少 4 題判斷正確且無越權承諾；留匿名作答與錯誤類型。"],
        ["G2｜90 分", "開放提問、重述、界線辨識、確認下一步，四項至少三項達標且無關鍵隱私違規。", "把「來源不明→標示未知→停止推測→確認下一責任角色與期限」納入既有角色演練；總時數不變。", "觀察規準及補練／再評；每個未確認事實保持 UNKNOWN。"],
        ["G3｜120 分", "完成任務、狀態、未結事項、下一步、接手角色五欄；不記不必要個資。", "加一頁參考卡：來源/核對狀態、未解事項、下一步、接手角色、期限；用匿名合成案例練習。", "沿用五欄完整度；P2 檢核不必要資料，留版本與去識別作品。"],
        ["G4｜60 分", "把界線、需求確認、紀錄、責任及未知資訊串成完整模擬交接。", "整合演練時強制說出「未確認來源」及正式追問路徑；未達者按同一檢核點補練。", "四項整合檢核全完成；未達及未重評者不得登錄通過。"],
        ["合計｜360 分", "90+90+120+60=360 教學分鐘；休息及行政排除。", "示範 A 不改總分鐘；所有分段與卡片均是提案，不構成核定課表。", "下一版需核對實際已核定總時數、師資、教材、場地與版本。"],
    ], [2.0, 5.0, 6.1, 4.9], size=9.5)
    add_para(doc, "設計衝突處理：本方案把練習加在 G1，並在 G2/G3 的既有演練中嵌入卡片，不從 G3 挪時；如果真實教師或課程版本無法在現有單元完成，應重新比較方案，不可默默把 360 改成 375。任何變更先由 P3 重算與核對，再送真實 P4 核定。")

    add_h(doc, "下一期指標、公式與預先判定規則", 2)
    add_table(doc, ["指標", "分子／分母或公式", "合成計畫門檻／判定", "資料來源與限制"], [
        ["學員練習充分度 Q3", "有效 Q3 分數合計 ÷ 該題有效 n", "本例建議目標 ≥3.5/5；若 n<5 不做平均比較，只列原始 n 及意見。", "課後匿名問卷；自評感受，不是能力驗收。3.5 為示範門檻。"],
        ["畢業後責任路徑 G3", "有效 G3 分數合計 ÷ 有效 n", "示範目標 ≥3.5/5；同時報追蹤有效 n / 邀請 n。", "D+30 自陳；不代表查核過真實服務紀錄。"],
        ["未解事項 G5", "有效 G5 分數合計 ÷ 有效 n", "示範目標 ≥3.5/5；若任何個案涉及安全或授權疑慮立即轉正式處理。", "D+30 自陳；不可用平均抵銷單一安全事件。"],
        ["課程示範能力評量", "依 0050 同版量規逐項達成數 / 有效受測數", "G1≥4/5 且無越權承諾；G2≥3/4 且無關鍵隱私違規；G3 五欄全備；G4 四項全備。", "引用 0050 R01 的同案示範量規；正式執行要另存實際版本與匿名評量原始記錄。"],
        ["任何授權／隱私例外", "事件件數及受影響流程，不與平均合併", "0 件是目標；任何事件先停用相關資料／教材並依正式制度處理。", "需有真實事件原件及受權處理者；本示例沒有真實事件。"],
    ], [3.0, 3.8, 5.0, 6.2], size=9.5)

    doc.add_page_break()
    add_h(doc, "11｜例外、修正及覆核追蹤", 1)
    add_table(doc, ["例外編號", "SAMPLE 情境／識別", "立即處置及決策", "責任／期限", "關閉證據"], [
        ["EX-S01", "問卷出現姓名或服務個案細節（本例僅作假設，不是已發生事件）。", "限制存取、不要複製到公開報告；由真實資料保護責任者判斷是否需事件程序。", "建議 P2；若權限未知，立即升報正式授權者。", "存取限制、去識別版本、依正式制度的處置紀錄；缺少時保持 OPEN。"],
        ["EX-S02", "回覆分母或課程版本無法核對。", "停止計算回覆率或合併版本；補回邀請紀錄、版本與缺答資料。", "P1 補來源；P3 重算；期限依正式工作安排確認。", "來源定位、重算表、第二人覆核。"],
        ["EX-S03", "平均值達標但出現一項越權或隱私安全問題。", "安全例外優先，不讓高平均抵銷；按正式制度暫停相關教材／活動。", "P2 立即通報有權角色；真實權限待核。", "可追溯事件處理原件及恢復授權；未完成不得關閉。"],
        ["EX-S04", "方案權重、課程分鐘或教材資源資料不完整。", "不把缺值當 0；明示假設，先比較可確認部分，必要時維持未決。", "P1 補資料、P3 重算、P4 核定。", "修訂版方案矩陣及理由／核定原件。"],
        ["EX-S05", "未取得 P4 真正核定權限的原件。", "版本維持 DRAFT / CANDIDATE；不可公告、開課或宣稱有效。", "P1 向正式治理來源查權限；日期 UNKNOWN。", "章程／授權文件與正式決議；沒有就不轉正式版。"],
    ], [1.4, 4.3, 5.2, 2.9, 3.7], size=9.5)
    add_para(doc, "本表列出的 EX-S01 至 EX-S05 是風險情境與控制示例，不表示原始點協會曾發生事故。真實例外須另用正式受限紀錄，避免公開個資；公開分支只放去識別的必要摘要。")

    add_h(doc, "12｜完整合成閉環紀錄與版本追蹤", 1)
    add_table(doc, ["示範日期", "記錄／責任角色", "可核對內容與判定", "狀態／下一步"], [
        ["2026-08-03 至 08-14", "0033 R03 同案需求調查；合成案例 SYN-0033-PLAN-01。", "僅引用 0033 R03 已登錄需求來源；該文件描述合成需求調查，不在本件重貼其問卷或聲稱為真人調查。", "SYNTHETIC 上游。"],
        ["2026-08-20", "0049 R01 同案課程需求及計畫形成日。", "同一合成課程名稱、12 人規劃假設、四單元與 360 分鐘。日期為 0049 示例日期，不是協會排課。", "SYNTHETIC 上游計畫。"],
        ["2026-08-28", "本件假設課後回饋情境；功能角色 P1/P2 均為示例代碼。", "假設邀請數 12、有效學員問卷 8；按 Q1-Q7 匿名逐列統計。沒有真實出席／回收原件。", "SAMPLE；計算見 §§5-6。"],
        ["2026-09-27", "假設結訓 D+30 追蹤。", "假設向 12 人邀請、5 份有效；3/5 自陳使用，需求多選與 G3-G5 分數逐列列於 §7。", "SAMPLE；沒有真實服務紀錄。"],
        ["2026-09-30", "P1 彙總，P3 示範重算。", "Q3=2.875；G3=3.200；G5=3.400；快速卡 4/5。按示範觸發規則比較 A/B/C。", "SYNTHETIC 派生數字，非實際審查會議。"],
        ["2026-10-02", "P1 起草方案 A；P3 只示範重算；P4 實際權限未知。", "方案 A 4.30/5；建議 G1 加 15 分鐘練習、在 G2/G3 嵌入快速卡、總時數仍 360。分數接近，取捨屬合成判斷。", "僅草案；未核定、未發布、未開課。"],
        ["下一期（日期 UNKNOWN）", "真實授權者核定後才排程；P1/P2/P3/P5 依正式指派工作。", "以原始匿名答卷、課程版本、量規、D+30 追蹤及例外處理驗證，並記錄未回覆分母。", "未取得核准及真實來源前不得標示已完成。"],
    ], [2.4, 4.2, 7.4, 3.5], size=9.5)
    add_para(doc, "草案變更對照：來源意見（Q3 低、畢業追蹤的 G3/G5 低且快速卡需求多）→方案 A（同時數內增加 G1 引導練習，G2/G3 加工具與未知／責任欄）→驗證（同版 Q3、既有 G1-G4 量規、D+30 G3/G5）→再議（若未改善或出現安全例外，開新一輪比較）。只有真實核定版才可生成正式發布紀錄。")
    add_h(doc, "13｜正式替換與上線檢核", 1)
    add_table(doc, ["步驟", "必須取得的 REAL 證據", "拒收／停止條件"], [
        ["1. 核對課程身份", "正式課程名稱、有效版本、課程日期／梯次、可核實需求來源與其授權。", "僅有歷史草稿、未核定課表或無法定位附件時不可當 REAL。"],
        ["2. 核對學員與追蹤", "依法授權的邀請分母、去識別問卷原始列、回收與缺答、30 日追蹤來源與發送紀錄。", "缺分母、個資權限不明、把 NA 當 0、只給平均而無原始資料。"],
        ["3. 核對計算與選擇", "題目版本、公式、完整頻數、方案準則／權重、異議及正式決策依據。", "不能重算、把小樣本說成統計代表性、把建議門檻說成官方規定。"],
        ["4. 核對權限與變更", "有權核定者的有效授權依據、正式核定原件、版本差異、發布及撤回紀錄。", "AI 文字、示範角色代碼、電子簽名字樣或未核准會議紀錄不算核定。"],
        ["5. 下一輪閉環", "實際使用的教材／課程版本、評量原始紀錄、學員意見與結訓追蹤及修訂記錄。", "未實施就維持未執行；無跟進證據不能寫成「改善完成」。"],
    ], [2.6, 8.7, 6.1], size=9.5)

    add_h(doc, "14｜官方要求—操作欄位—附件定位追溯", 1)
    add_table(doc, ["官方要求／子義務", "本件可核對欄位或段落", "示範證據及限制"], [
        ["依學員意見訂訓練課程", "§4 問卷 Q1-Q7；§5 逐列資料；§6 分母及分析。", "8 列全為合成，無真人問卷。"],
        ["依結訓學員意見訂訓練課程", "§7 G1-G5 與複選需求；§8 使用比例、平均及限制。", "5 列全為合成自陳，未驗證實際行為。"],
        ["意見轉成實際課程規劃取捨", "§9 準則、權重、A/B/C 公式；§10 單元、時數、方法與評量修訂。", "權重、分數與選擇是示範判斷，未經核准。"],
        ["流程有完整紀錄", "§2 SOP；§3 角色權限；§11 異常；§12 時間軸／版本；§13 替換與拒收。", "時間軸中的課程與追蹤是完整合成流程，不能冒稱協會原始紀錄。"],
        ["官方來源定位", "本節、§1；inputs/official_evidence_table_page4.png。", "官方資料表 PDF p.4、checkbox 054；不以 DOCX 編號替代官方流水號。"],
    ], [4.4, 6.0, 7.0], size=9.5)

    add_h(doc, "附錄 A｜學員課後問卷可複製表", 1)
    add_para(doc, "基本欄：課程代碼______ 課程版次______ 課程日期______（由核定版本填）；匿名代碼______；本問卷告知版本______。請逐題圈選 1–5 或 NA。請勿填寫姓名或服務對象可識別資訊。")
    add_table(doc, ["代碼", "題目", "1", "2", "3", "4", "5", "NA"], [[f"Q{i:02d}", q, "□", "□", "□", "□", "□", "□"] for i, q in enumerate(LEARNER_Q, 1)], [1.0, 8.3, 1.2, 1.2, 1.2, 1.2, 1.2, 1.7], size=9.5)
    add_para(doc, "可選意見（不得填個人或服務對象可識別資料）：最有幫助的情境______；仍不確定的步驟______；建議增加／調整的練習______。正式使用時另附告知、資料保存及申訴說明。")

    add_h(doc, "附錄 B｜結訓學員 30 日追蹤表", 1)
    add_para(doc, "課程版本______；結訓日期______；追蹤日期______；匿名代碼______。G1：曾使用／未使用／沒有機會／不願回答；未使用或沒有機會者選一原因。G2 多選：快速卡／角色演練／來源查證／其他去識別建議；「目前不需要」須單獨選擇。G3-G5 各選 1–5 或 NA。沒有使用機會者不納入「已使用」比例分母。")
    add_table(doc, ["題號", "題目", "1", "2", "3", "4", "5", "NA"], [[f"G{idx+3}", q, "□", "□", "□", "□", "□", "□"] for idx, q in enumerate(GRAD_Q)], [1.0, 8.3, 1.2, 1.2, 1.2, 1.2, 1.2, 1.7], size=9.5)

    add_h(doc, "附錄 C｜課程修訂追蹤記錄（合成範例已填）", 1)
    add_table(doc, ["欄位", "SAMPLE / SYNTHETIC 填寫示例", "REAL 替換要求"], [
        ["來源意見", "Q3 2.875；G3 3.200；G5 3.400；快速參考卡 4/5。", "連結原始匿名列及該版問卷；不可只貼彙總。"],
        ["課程版本差異", "提案：G1 同時數增加 15 分鐘引導練習；G2/G3 使用一頁卡；360 分鐘維持。", "取得前版／新版本課程計畫逐欄差異。"],
        ["替代方案與理由", "A=4.30；B=4.20；C=4.25；分數為合成判斷。", "保留每方案分數、權重來源、反對意見與重算簽核。"],
        ["授權與狀態", "待核定；沒有核定人或正式會議紀錄。", "附有效權限依據、真實決議原件與正式版本識別。"],
        ["下一期驗證", "沿用 Q3/G3/G5 及 0050 量規；期限待正式排課。", "填實際梯次、日期、邀請／有效分母、評量原件及結論。"],
    ], [3.2, 6.9, 7.5], size=9.5)

    add_h(doc, "15｜版本、限制及候選狀態", 1)
    add_table(doc, ["項目", "狀態"], [
        ["本文件", "DOCX 0052 R01；候選文件，非正式核定計畫。"],
        ["官方映射", "I11-E04；官方頁 4；方框 054；crosswalk deliverable_count=1、deliverable_ids=0052。"],
        ["樣本與日期", "全為 SAMPLE / SYNTHETIC；逐事實登錄 evidence/SYNTHETIC_REPLACEMENT_REGISTER_R01.csv。"],
        ["REAL 原件", "目前未取得本要求專屬協會真實意見、追蹤、實際辦訓或正式核准原件；保持 UNKNOWN。"],
        ["正式接受", "未授予；需獨立 exact-SHA 語意／視覺驗收及真正權限者後續接受。"],
        ["建置來源", f"本地建置讀取分支來源 HEAD {BUILD_SOURCE_HEAD}；官方與 Blueprint 精確雜湊、鄰件 SHA 及限制見 SOURCE_MAP_R01.md。"],
    ], [3.3, 13.7], size=9.5)
    add_callout(doc, "可直接操作的內容包括兩份問卷、逐列合成答卷、分母與計算、選項評分、下一版單元設計、責任分工、風險處理及版本追蹤。正式轉 REAL 的必要條件是取得並定位真實來源，完成去識別、授權、重算及正式核定；未完成前不得改標籤。", "DDEBF7", 10.2)

    DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX)


def write_source_map(blueprint):
    runtime_blueprint = RUNTIME / "CONTROL" / "BLUEPRINT_CURRENT.jsonl"
    official_img = PKG / "inputs" / "official_evidence_table_page4.png"
    official_img_hash = sha(official_img) if official_img.exists() else "NOT_COPIED"
    row_hash = sha(BLUEPRINT_COPY)
    docx_hash = sha(DOCX)
    register_hash = sha(REGISTER)
    lines = [
        "# DOCX 0052 R01｜來源與範圍圖",
        "",
        "## 身分與官方要求",
        f"- doc_id: 0052; Unique Blueprint requirement_id: I11-E04; 官方 checkbox: 054; 官方資料表頁: 4。",
        f"- 官方原文: {OFFICIAL_TEXT}",
        "- 文類: strategy_plan（課程意見回流與課程修訂策略／計畫），不是 0051 的產出審查紀錄，也不是 0053 的課程切合性檢討。",
        "- official source PDF SHA-256: 805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709。",
        f"- 官方 p.4 方框影像 SHA-256: {official_img_hash}。",
        f"- BLUEPRINT_CURRENT.jsonl 全檔 SHA-256: {sha(runtime_blueprint)}。",
        f"- 本件 Unique Blueprint 原始列副本 SHA-256: {row_hash}。",
        "- 工作分支於建置前讀回 HEAD: " + BUILD_SOURCE_HEAD + "。",
        "",
        "## 唯一合成案例與跨文件事實界線",
        f"- Case ID: {CASE}；課程名稱: {COURSE}。所有 0052 問卷列、追蹤、日期、分母、評分、取捨和版本變更均為 SYNTHETIC。",
        "- 0033 R03 I07-E01 / checkbox 033 DOCX SHA: fe5f3cfd84c761da86435a71e0fd899daa725be78628c4d2d3f2324c1fdb5f4f。僅作同案例的合成需求上游，不在本件複製其問卷或將其結果當 REAL。",
        "- 0049 R01 I11-E01 / checkbox 051 DOCX SHA: e8212788b1f7b201ca4af4f4c5fb83ab131627dbfa815987f7cf6344edfda727。來源只支持合成課程計畫、四單元與 12 人規劃假設，不證明實際排課。",
        "- 0050 R01 I11-E02 / checkbox 052 DOCX SHA: d0d37c2c69b161a2cdd37ee3097306722c4a65a1afd844ae92e1dff58b43381a。來源只支持同案示範單元目標與評量，不證明課程實施或學員表現。",
        "- 0051 R01 I11-E03 / checkbox 053 DOCX SHA: 6f575ebb41c6c26b4d6b36262fb4d9a5abdae53cbdc1ff4e75acda46d5f0b495。該件處理課程產出討論／確認；本件不代替會議或核准原件。",
        "- 0053 I12-E01 是另一官方要求；本件不處理課程切合性檢討。",
        "",
        "## 資料來源與方法",
        "- 合成案例受訪數：學員 8 列，假設邀請 12；結訓追蹤 5 列，假設邀請 12。每個原始回答和衍生頻數、平均、比例、選項分數均在 SYNTHETIC_REPLACEMENT_REGISTER_R01.csv 以唯一 fact_id 登錄。",
        "- 學員題目 1–7 為五點同意量尺；平均=有效分數合計÷該題有效 n。缺答不補 0。",
        "- 結訓追蹤 G3–G5 同量尺；三項支持需求分母為有效回答者、選項可重疊；「目前不需要」與支持項互斥。自陳不是客觀行為稽核。",
        "- 本例內部觸發線 mean<3.5/5 與權重 30/25/20/15/10 均為 SAMPLE POLICY EXAMPLE，不是官方門檻或協會政策。",
        "- 加權評分公式: sum(weight_percent × score_1_to_5)/100。三案 A/B/C 的原始權重、評分、乘積與合計均在正文及替換簿可重算。",
        "",
        "## 未取得／不得聲稱",
        "- 未取得本要求專屬協會真實學員意見或結訓追蹤原件；未證明有真實課程、學員、受邀分母、課後回覆、服務使用或辦訓成果。",
        "- 未取得協會核定權限、實際責任人、個資保存政策、正式課程修訂核准或正式會議紀錄。這些狀態維持 UNKNOWN。",
        "- 本件不得冒稱 REAL、正式接受或已核定。人名、簽章、核准、第三方原件及學員個資均未偽造。",
        "",
        "## Build output",
        f"- Candidate: candidate/0052_R01.docx; SHA-256 {docx_hash}。",
        f"- Synthetic replacement register: evidence/SYNTHETIC_REPLACEMENT_REGISTER_R01.csv; SHA-256 {register_hash}; rows {len(facts)}（含表頭前資料列數）。",
        "- Independent exact-SHA reviewer and final visual receipts are stored separately under review/ and evidence/qa/ after completion.",
    ]
    SOURCE_MAP.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main():
    ensure_dirs()
    blueprint = load_blueprint()
    build_register()
    build_doc(blueprint)
    write_source_map(blueprint)
    metadata = {
        "project_id": "TTQS_ONE",
        "doc_id": "0052",
        "revision": "R01",
        "requirement_id": "I11-E04",
        "official_checkbox_sequence": "054",
        "genre": "strategy_plan",
        "candidate_state": "CANDIDATE_NOT_PROMOTED",
        "sample_state": "SAMPLE_SYNTHETIC_ONLY_PER_REGISTER",
        "source_head_at_build": BUILD_SOURCE_HEAD,
        "candidate_sha256": sha(DOCX),
        "source_map_sha256": sha(SOURCE_MAP),
        "blueprint_row_copy_sha256": sha(BLUEPRINT_COPY),
        "replacement_register_sha256": sha(REGISTER),
        "replacement_register_data_rows": len(facts),
        "formal_acceptance": "NOT_GRANTED",
        "real_data_status": "NO_REQUIREMENT_SPECIFIC_REAL_LEARNER_OR_GRADUATE_SOURCE_LOCATED",
        "authorization_status": "UNKNOWN_NOT_GRANTED",
    }
    meta_path = PKG / "evidence" / "build" / "BUILD_METADATA_R01.json"
    meta_path.write_text(json.dumps(metadata, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(metadata, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

