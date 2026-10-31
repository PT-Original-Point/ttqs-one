from pathlib import Path
from hashlib import sha256
import json, subprocess, zipfile, io, re, sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0137_I19-E27_CANDIDATE_R01_20261025')
REPO=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\GITHUB_SYNC_0131_R01D_20261024')
REF=subprocess.check_output(['git','-c',f'safe.directory={REPO}','-C',str(REPO),'rev-parse','FETCH_HEAD'],text=True).strip()
DOCX=ROOT/'0137_I19-E27_媒體採訪與正面報導原件查驗表_R01.docx'

def git_bytes(path):
    return subprocess.check_output(['git','-c',f'safe.directory={REPO}','-C',str(REPO),'show',f'{REF}:{path}'])
def blob(path):
    return subprocess.check_output(['git','-c',f'safe.directory={REPO}','-C',str(REPO),'rev-parse',f'{REF}:{path}'],text=True).strip()
def h(data): return sha256(data).hexdigest()
def write_text(path, text):
    Path(path).write_text(text,encoding='utf-8',newline='\n')

# Preserve independently sourced inputs from the current PR tree.
source_dir=ROOT/'source'; source_dir.mkdir(exist_ok=True)
input_paths={
 'official_source.pdf':'docs/ttqs-quality-20261016/quality-review/0037_R07/inputs/official_source.pdf',
 'official_crosswalk.csv':'docs/ttqs-quality-20261016/quality-review/0136_I19-E26_CANDIDATE_R02_20261025/source/official_129_to_142_crosswalk.csv',
 'project_index149_151.md':'docs/ttqs-quality-20261016/quality-review/0136_I19-E26_CANDIDATE_R02_20261025/source/PROJECT_157_INDEX150_AND_NEIGHBORS.md',
 'official_source_pages8_9_excerpt.pdf':'docs/ttqs-quality-20261016/quality-review/0136_I19-E26_CANDIDATE_R02_20261025/source/OFFICIAL_SOURCE_P8_P9_EXTRACT.pdf',
 'official_source_pages8_9_text.txt':'docs/ttqs-quality-20261016/quality-review/0136_I19-E26_CANDIDATE_R02_20261025/source/OFFICIAL_SOURCE_P8_P9_TEXT.txt',
}
source_manifest={'project_id':'TTQS_ONE','doc_id':'0137','requirement_id':'I19-E27','official_checkbox_serial':'UNKNOWN','project_157_index':151,'github_ref':REF,'files':{}}
for name,path in input_paths.items():
    data=git_bytes(path); (source_dir/name).write_bytes(data)
    source_manifest['files'][name]={'github_path':path,'git_blob':blob(path),'bytes':len(data),'sha256':h(data)}
# Contract snapshots are source hashes plus narrowly scoped excerpt; full docs stay on PR.
control_paths={
 'mission':'docs/ttqs-quality-20261016/01_HUMAN_MISSION_AND_ACCEPTANCE.md',
 'genre_manual':'docs/ttqs-quality-20261016/02_GENRE_WRITING_MANUAL.md',
 '157_blueprints':'docs/ttqs-quality-20261016/03_157_CHECKBOX_BLUEPRINTS.md',
 '04_goal_execution':'docs/ttqs-quality-20261016/04_CODEX_GOAL_24X7_EXECUTION.md',
 '05_native_wake':'docs/ttqs-quality-20261016/05_HOURLY_NATIVE_WAKE_RESUME.md',
 '06_external_continuation':'docs/ttqs-quality-20261016/06_EXTERNAL_AUDIT_AND_CONTINUATION_20261016.md',
 'AGENTS':'AGENTS.md'}
source_manifest['control_blobs']={k:{'path':p,'git_blob':blob(p)} for k,p in control_paths.items()}
source_manifest['binding']={'official_requirement_id':'I19-E27','doc_id':'0137','project_157_index':151,'printed_checkbox_serial':'UNKNOWN','mapping_row':'I19-E27,1,0,0137,媒體採訪與正面報導原件查驗表'}
source_manifest['candidate_state']='SOURCE_PREP_AND_CANDIDATE; NOT FORMALLY ACCEPTED'
write_text(ROOT/'SOURCE_MANIFEST_0137_R01.json',json.dumps(source_manifest,ensure_ascii=False,indent=2))

source_packet='''# 0137 I19-E27 正式來源包 R01\n\n## 身分與狀態\n- doc_id：0137；官方要求：I19-E27；157 專案索引：151；官方印刷方框流水號：UNKNOWN。\n- 本候選只提供媒體原件取得、來源核驗、內容摘錄、事實對照及索引流程。原始點協會真實媒體報導尚未提供，REAL_MEDIA_SOURCE=UNKNOWN/HOLD，正式接受=NOT_GRANTED。\n- PR #92 分支及施工來源：`'+REF+'`。本來源包的精確檔案雜湊見 `SOURCE_MANIFEST_0137_R01.json`。\n\n## 官方對應\n- 官方查核資料表第 9 頁列出：「提出報章雜誌採訪或電子媒體訪問之正面消息（其他）」。此為佐證類型，不表示本會已有採訪、報導、正面評價或任何 TTQS 分數。\n- 最新 129→142 crosswalk 對應 `I19-E27,1,0,0137,媒體採訪與正面報導原件查驗表`。\n- 157 項寫作藍圖索引 151：真實報導來源與截頁索引，需有特定來源定位、核心內容或示例、決策或結果、與其他文件的事實連結。第三方原件取得類文件可保留待取得欄位，但不能以空欄冒充已取得原件。\n- 指標 19 寫作手冊：媒體/社會影響佐證要有來源索引及合法取得工具；沒有真實第三方來源，不得假造採訪、報導、成果或評分。\n\n## 使用來源與界線\n1. 官方 PDF：現行分支 `quality-review/0037_R07/inputs/official_source.pdf`；原始 PDF bytes 保存在本目錄。第 9 頁是 I19-E27 媒體來源項目的頁面。\n2. 129→142 對照：`official_crosswalk.csv`。\n3. 157 項 149–151 摘錄：`project_index149_151.md`；本件使用索引 151，並將 150 的公益/表揚/產學社會影響及 149 的市場價值趨勢列為負向鄰件。\n4. Mission、寫作手冊、157 Blueprint、GOAL/喚醒/外部延續規則的現行 blob 見來源 manifest；人類接受規則、反造假及 candidate 狀態不因 PR 上傳改變。\n\n## 文件專屬問題\n如何確認一篇媒體報導的發布來源、完整上下文、日期、作者、標題及原件位置；如何拆分媒體實際報導與協會可證明的事實；如何判斷正面訊息與訓練/公益相關性；如何留存可回查截頁並處理轉載、修正、授權與隱私；如何在來源、權限、內容或歸屬不足時 HOLD/REJECT。\n\n## 範例資料\n文件中的 `SYN-0137-CASE-01`、媒體名稱、網址、日期、受訪內容、數字、原件與內部紀錄均為 SAMPLE/SYNTHETIC。`example.invalid` 是保留的不可解析示範網域。不得將示例剪貼或改名後當作本會的真實報導。\n'''
write_text(ROOT/'SOURCE_PACKET_0137_I19-E27_R01.md',source_packet)

blueprint='''# UNIQUE BLUEPRINT 0137 I19-E27 R01\n\n## 唯一用途\n- doc_id：0137；requirement_id：I19-E27；157 index：151；官方印刷流水號：UNKNOWN。\n- 文類：第三方媒體報導原件查驗、內容對照與截頁索引。\n- 主要讀者：小型志工協會承辦人、原件保管人及獨立覆核者。\n- 不可把本件改寫成協會新聞稿、成效評估、得獎證明、合作契約、會員成長分析或市場趨勢報告。\n\n## 官方與跨件定位\n- 官方來源：查核資料表第 9 頁 I19 媒體項目；原件頁面和 crosswalk 的 raw SHA/Git blob 記於 SOURCE_MANIFEST。\n- 129→142：I19-E27→0137。\n- 157 index 151：正面媒體報導的真實來源與截頁索引。\n- Negative neighbors：0136/I19-E26（公益/表揚/產學或研發社會影響原件）；0135/I19-E25（市場價值趨勢）。媒體提及其他證據不會自動替代它們。\n\n## 必要內容\n1. 開篇說明文件功能、目前 REAL 狀態、官方項目界線。\n2. 以發布機構原始頁面/完整刊物/節目存檔核對來源，不以截圖、搜尋摘要或協會稿單獨冒充媒體原件。\n3. 保存標題、媒體、作者、發布/修正時間、URL/頁次/時間碼、取得時間、原件 byte/hash、上下文及版本。\n4. 把報導文字的每項可檢驗聲稱拆成 claim_id，連到協會內部或第三方原始佐證，逐項標記 supported/unsupported/out-of-scope；不從正面語氣推導真實成效或因果。\n5. 說明正面訊息、訓練/公益關聯、媒體獨立性、轉載重複、修正、著作權/引用、隱私與保存方式。\n6. 提供完整 SAMPLE/SYNTHETIC worked case、欄位表、索取稿、replacement register、HOLD/REJECT 分流與可供人員照做的覆核清單。\n7. Cross-doc facts 只有 same case_id 且原始證據互相核對後才能連結；不同案例不得假裝連續。\n\n## 專用決策\n- RECORDABLE_LIMITED：原始媒體來源已核實，且資料只被用於證明報導存在與其實際文字；不自動證明報導中的協會成果。\n- HOLD：來源、完整上下文、授權、歸屬、內部事實連結或修正歷史仍待查；不得計入已驗證原件數。\n- REJECT：偽造/竄改原件、錯誤機構、單純協會新聞稿假扮第三方報導、重複轉載被重複計數、無關或負面內容被冒稱正面、或故意省略重大上下文。\n\n## Sample arithmetic\n- 示範課程：24 人登記、22 人出席、20 人符合 6/8 小時完成規則、2 人未達完成門檻。出席率=22/24=91.7%；出席者完成率=20/22=90.9%。\n- 這些數字只驗證合成案例的內部一致性。新聞文字不能作為 attendance source；如新聞把活動效果寫成「顯著提升」，而沒有可檢查的成效測量，該效果 claim 必須 HOLD/不得採用。\n\n## Reject conditions\n不提供完整原件也不標明缺失；網址/媒體/日期/作者互相矛盾；來源為協會自有宣傳卻申報媒體採訪；媒體文字外推為正式評分、組織成果或因果；cross-doc case_id/數據不一致；把 SAMPLE 變更名稱後當 REAL；未完成逐頁/逐區目視或 exact-SHA 獨立語意審查卻宣稱接受。\n'''
write_text(ROOT/'UNIQUE_BLUEPRINT_0137_I19-E27_R01.md',blueprint)

manuscript='''# 0137 正文提綱與合成示範\n\n候選只在正式第三方媒體原件、協會內部事實來源與必要授權都可回查時，才可把個案轉成 REAL。沒有來源時，操作人員可以使用本文件的空白採集欄位，但須把狀態留在 HOLD；示例不能補足本會實績。\n\n`SYN-0137-CASE-01` 是完整的虛構教學案例：虛構媒體「示例社區報」、保留網域 `example.invalid`、假設日期與假設的防災課程數字。虛構報導表示某工作坊有 22 人出席、20 人完成，並使用「增強社區應變」的廣泛描述。合成內部工作表顯示 24 人登記、22 人出席、20 人完成；因此兩個可計數聲稱在示範資料內相符。廣泛成效描述沒有合成前後測或獨立觀察資料支持，須 HOLD，不能當成可歸因成果。\n\n實際使用時依序保存原始發布位置、完整上下文、版本與 hash；拆解聲稱；向協會授權保管的原始課程/活動紀錄逐項核對；評估來源獨立性、修正及重複轉載；再由未撰寫紀錄者覆核。只有媒體原件可證明報導存在和媒體如何描述；出席、學習、社會影響及協會正式決議要各自回到其原始來源。\n'''
write_text(ROOT/'MANUSCRIPT_0137_R01.md',manuscript)

# DOCX formatting helpers.
def set_cell_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); mar=tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar=OxmlElement('w:tcMar'); tcPr.append(mar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        e=mar.find(qn('w:'+m))
        if e is None: e=OxmlElement('w:'+m); mar.append(e)
        e.set(qn('w:w'),str(v)); e.set(qn('w:type'),'dxa')
def set_borders(table):
    tblPr=table._tbl.tblPr; borders=tblPr.first_child_found_in('w:tblBorders')
    if borders is None: borders=OxmlElement('w:tblBorders'); tblPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag='w:'+edge; e=borders.find(qn(tag))
        if e is None: e=OxmlElement(tag); borders.append(e)
        e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:space'),'0'); e.set(qn('w:color'),'D9D9D9')
def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr(); tblHeader=OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'),'true'); trPr.append(tblHeader)
def set_row_no_split(row):
    trPr=row._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))
def style_run(run,size=10.5,bold=False,color='000000'):
    run.font.name='Microsoft JhengHei'; run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei'); run._element.rPr.rFonts.set(qn('w:ascii'),'Microsoft JhengHei'); run._element.rPr.rFonts.set(qn('w:hAnsi'),'Microsoft JhengHei'); run.font.size=Pt(size); run.font.bold=bold; run.font.color.rgb=RGBColor.from_string(color)
def set_cell_text(cell,text,header=False):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.08
    r=p.add_run(str(text)); style_run(r,9.5,header, 'FFFFFF' if header else '000000')
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cell)
def add_table(doc,headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; t.style='Table Grid'; set_borders(t); set_repeat_table_header(t.rows[0])
    for i,x in enumerate(headers): set_cell_text(t.rows[0].cells[i],x,True); set_cell_shading(t.rows[0].cells[i],'365F7D')
    for ridx,row in enumerate(rows):
        cells=t.add_row().cells; set_row_no_split(t.rows[-1])
        for i,x in enumerate(row):
            set_cell_text(cells[i],x)
            if ridx%2: set_cell_shading(cells[i],'F1F5F8')
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t
def add_p(doc,text='',bold_lead=None,style=None):
    p=doc.add_paragraph(style=style) if style else doc.add_paragraph()
    p.paragraph_format.space_after=Pt(5); p.paragraph_format.line_spacing=1.12
    if bold_lead and text.startswith(bold_lead):
        a=p.add_run(bold_lead); style_run(a,10.5,True); b=p.add_run(text[len(bold_lead):]); style_run(b,10.5)
    else:
        r=p.add_run(text); style_run(r,10.5)
    return p
def add_h(doc,text,level=1):
    p=doc.add_paragraph(style=f'Heading {level}'); p.paragraph_format.keep_with_next=True; p.paragraph_format.space_before=Pt(10 if level==1 else 6); p.paragraph_format.space_after=Pt(4)
    r=p.add_run(text); style_run(r,15 if level==1 else 12.2,True)
    return p
def add_bullet(doc,text):
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.08; r=p.add_run(text); style_run(r,10.5); return p

doc=Document(); sec=doc.sections[0]; sec.page_width=Inches(8.27); sec.page_height=Inches(11.69); sec.top_margin=Inches(.68); sec.bottom_margin=Inches(.68); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
styles=doc.styles
for stylename,size in [('Normal',10.5),('Title',22),('Heading 1',15),('Heading 2',12.2)]:
    s=styles[stylename]; s.font.name='Microsoft JhengHei'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei'); s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)
    if stylename.startswith('Heading'): s.font.bold=True
# Title block
p=doc.add_paragraph(style='Title'); p.paragraph_format.space_after=Pt(6); r=p.add_run('媒體採訪與正面報導原件查驗表'); style_run(r,22,True)
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(9); r=p.add_run('來源核驗流程與截頁索引'); style_run(r,14,True)
add_table(doc,['文件識別','目前狀態'],[
 ('doc_id／官方要求／157索引','0137 ／ I19-E27 ／ 151；官方印刷流水號 UNKNOWN'),
 ('候選版本','R01；候選文件，不是協會正式核定或 TTQS 評分'),
 ('真實媒體原件','UNKNOWN / HOLD：尚未提供可核驗的本會媒體原件'),
 ('示範案例','SAMPLE / SYNTHETIC：虛構媒體、網址、人物角色及數字，不可引用為本會事實'),
], [2.1,4.9])
add_p(doc,'本表用來登錄媒體原件、核對發布來源與上下文、拆解報導聲稱並追溯各聲稱的真正證據。媒體報導只能證明報導內容如何呈現；它本身不能取代出席、訓練成效、合作、得獎或社會影響原始紀錄。')
add_p(doc,'目前沒有可回查的原始點協會媒體報導原件。本候選只建立可操作工具並示範核驗方式；REAL_MEDIA_SOURCE=UNKNOWN/HOLD，正式接受=NOT_GRANTED。',bold_lead='目前沒有可回查的原始點協會媒體報導原件。')

add_h(doc,'1. 收到報導後先走這七步')
steps=[
'先登錄媒體、標題、作者、發布時間、原始網址或刊物頁次；若是電視/廣播，記下節目、播出日期與時間碼。',
'從發布媒體或出版單位取得完整頁面、刊物或節目存檔。搜尋摘要、社群轉貼、裁切截圖及協會新聞稿都不能單獨當成媒體原件。',
'保存原始檔及完整上下文，記錄取得時間、檔案大小、SHA-256、版本/更正狀態與受控保存位置。',
'將報導拆成一列一項的 claim_id；逐項對照課程、活動、出席、完成或結果的原始來源。',
'判斷媒體是否獨立、內容是否實際提及本會與訓練/公益活動、正面訊息來自何段原文；不要從語氣推導成果。',
'檢查轉載/更正、引用授權、個資遮罩與保存期限；需要擴散前由授權人確認引用範圍。',
'由未建立本筆紀錄的人覆核原件、claim 對照及 HOLD/REJECT 分流，再決定是否可列入佐證索引。'
]
for s in steps: add_bullet(doc,s)

add_h(doc,'2. 官方要求與文件邊界')
add_p(doc,'157 專案藍圖索引 151 的專屬要求是「正面媒體報導：真實報導來源與截頁索引」；文件要定位特定來源，呈現核心內容或完整 SAMPLE/SYNTHETIC 示範，留下判讀結果，並建立與其他文件的事實連結。')
add_p(doc,'官方查核資料表第 9 頁列出「提出報章雜誌採訪或電子媒體訪問之正面消息（其他）」。這是可供查核的佐證類型，不表示原始點協會已受訪、已被報導、已獲正面評價或已達成任何分數。')
add_table(doc,['邊界項目','0137 的處理方式','不能推論的內容'],[
('官方 crosswalk','I19-E27 → DOCX 0137「媒體採訪與正面報導原件查驗表」。','不能把內部序號或檔名當作官方方框印刷流水號。'),
('157 專案索引 151','必須定位真實報導來源、核心內容/示例、判讀結果及與其他文件的事實連結。','不能只交一張空白登錄表，也不能把示範改名當成真實報導。'),
('0136／I19-E26','記公益、表揚、產學/研發社會影響的原件與來源。','0137 的新聞內容不會取代獎狀、合作文件或社會影響原始證據。'),
('0135／I19-E25','記市場價值及長期趨勢。','新聞提到市場或服務，不會自動變成量化市場趨勢。'),
('媒體報導可證明','經核實的出版/播出來源、時間、上下文及報導實際文字。','報導文字不自動證明訓練造成改善、組織績效或正式核定。'),
], [1.25,3.0,2.75])

add_h(doc,'3. 核驗角色與逐關流程')
add_table(doc,['關卡','承辦人要留下的紀錄','通過條件','未通過處理'],[
('G1 來源定位','媒體/出版者、URL/刊名頁次/節目時間碼、作者與發布時間。','來源可回到媒體或出版者原始發布處。','來源不明先 HOLD；查明偽造則 REJECT。'),
('G2 原件完整','完整頁面/刊物/節目存檔、上下文、原檔 hash、取得時間。','原件足以辨認標題、發布者、日期及報導內容。','只剩截圖或摘要則 HOLD；竄改則 REJECT。'),
('G3 媒體獨立性','報導與新聞稿/付費廣告/轉載的關係、編輯來源及更正史。','來源性質如實標示，未把自有宣傳偽裝成第三方報導。','關係待核先 HOLD；故意誤標則 REJECT。'),
('G4 對象與正面訊息','報導中本會、訓練/公益活動及正面內容的精確位置。','正文可定位支持所述範圍。','對象或正面內容不清則 HOLD；無關或負面被冒稱正面則 REJECT。'),
('G5 逐項事實對照','claim_id 對應內部/第三方原始紀錄與範圍限制。','各 claim 分開核實；媒體文字不被當作成果底稿。','缺佐證的 claim 保持 HOLD 或從主張中刪除。'),
('G6 權限與隱私','引用/保存權限、個資遮罩、存取人及保存期限。','僅在授權範圍使用必要摘錄。','權限未知先 HOLD；不得公開擴散。'),
('G7 獨立覆核','覆核者、日期、核對原件與 claim 的結果。','覆核者未建立原始登錄，且留有可回查定位。','未覆核不得記為已核驗原件。'),
], [0.8,2.05,2.25,1.9])
add_p(doc,'三種決策狀態要分開記錄：RECORDABLE_LIMITED 表示來源已核、只可用來證明媒體曾如此報導；HOLD 表示仍缺來源、權限、上下文或事實核對；REJECT 表示原件偽造、冒稱媒體、錯誤對象、重複計數或故意省略重大上下文。RECORDABLE_LIMITED 不等於成果已證明。')

add_h(doc,'4. 可供實際採集的原件登錄表')
add_p(doc,'以下欄位留給真實第三方原件取得與索引。每個欄位必須填實際可回查資訊，或明確記 UNKNOWN / HOLD / NOT APPLICABLE；不得用推測值填補。')
add_table(doc,['欄位','實際登錄內容'],[
('case_id / media_id / claim_id',''),('媒體/出版者與媒體類型',''),('標題／作者／發布日期與時間／修正時間',''),('原始 URL／刊名卷期頁次／節目與時間碼',''),('取得人角色／取得日期與時間／方法',''),('原件檔名／檔案類型／大小／SHA-256',''),('完整上下文保存位置／存取權限／保存期限',''),('是否轉載、付費、贊助或協會自有稿',''),('本會/課程/公益活動出現位置與逐字/摘要定位',''),('引用授權範圍／個資遮罩與權利限制',''),('claim 對照來源／支持範圍／未支持事項',''),('決策：RECORDABLE_LIMITED／HOLD／REJECT',''),('承辦角色／覆核角色／日期／覆核結果',''),
], [2.4,4.6])

add_h(doc,'5. 完整合成示範個案')
add_p(doc,'SAMPLE / SYNTHETIC ONLY：本節媒體、文章、活動、角色、日期及數據全部虛構，只演示查驗流程，不可引用為本會事實。')
add_p(doc,'本節所有內容都是虛構的教學資料，不是本會報導、活動、課程、媒體、學員或成效紀錄。示例故意保留一項未能由證據支持的廣泛成效說法，示範如何限制引用。')
add_table(doc,['SAMPLE/SYNTHETIC 欄位','虛構示範值'],[
('case_id','SYN-0137-CASE-01'),('媒體與類型','虛構《示例社區報》；電子媒體；不是實際出版物。'),('示範網址','https://news.example.invalid/2026/05/18/community-drill'),('標題與時間','「防災工作坊完成訓練並提出演練建議」；假設發布 2026-05-18，活動日 2026-05-10。'),('原件狀態','SAMPLE/SYNTHETIC 教學頁；不存在可供本會正式引用的真實媒體原件。'),('合成活動數據','登記 24 人；出席 22 人；完成 20 人；未達完成門檻 2 人。完成門檻為全程 8 小時至少參與 6 小時。'),('計算','出席率 22/24=91.7%；出席者完成率 20/22=90.9%。這只檢查示範內部算術，不是協會績效。'),('狀態','Sample test：資料鏈示範可追溯。真實協會資料仍 UNKNOWN/HOLD；不得把示範結論移作 REAL。'),
], [2.0,5.0])

add_h(doc,'6. 逐項 claim 對照與示範判讀')
add_p(doc,'SAMPLE / SYNTHETIC ONLY：以下 claim_id 與原始來源定位屬虛構測試資料，沒有任何真實媒體或協會記錄。')
add_table(doc,['claim_id','示範報導聲稱','合成原始來源定位','核對結果與允許說法'],[
('SYN-C01','活動於 2026-05-10 舉行。','SYN-TRN-01 課程計畫表列日期；SYN-TRN-02 場地簽到記錄列活動日。','兩筆合成資料一致。只可示範如何核日期；不能證明真實活動。'),
('SYN-C02','有 22 人出席。','SYN-TRN-02：22 個唯一學員代碼；去除重複簽到後仍為 22。','22/24=91.7% 出席率。媒體文字不是出席名冊的替代來源。'),
('SYN-C03','20 人完成課程。','SYN-TRN-03：每人 8 小時簽到時數；規則為至少 6 小時，20 人符合。','20/22=90.9% 出席者完成率。分母是 22 名出席者，不是 24 名登記者。'),
('SYN-C04','課程「增強社區應變能力」。','示例沒有可重算的前後測、情境演練量尺或獨立觀察紀錄。','HOLD。不得將新聞用語寫成已證明的組織或社會成效，也不得作因果推論。'),
('SYN-C05','媒體為獨立第三方發布。','示例原始來源頁標示虛構出版者；真實案例需核出版者、稿件來源、贊助/付費與轉載關係。','只在合成案例內作流程演練；本會真實媒體獨立性 UNKNOWN。'),
], [0.78,1.7,2.35,2.17])
add_p(doc,'示範索引結論：若真實原件與內部資料均依 G1–G7 完成核驗，媒體報導可列為 RECORDABLE_LIMITED，僅證明發布內容與來源；SYN-C04 仍應 HOLD，直到有適當成效來源。此結論僅為規則演示，不代表本會已取得報導。')

add_h(doc,'7. 截頁、原件保存與權利處理')
add_table(doc,['控制項','實際做法','不合格訊號'],[
('原始頁面','保存媒體原始 URL 的完整頁面或出版者提供的原始檔；記錄取得日期、時間與登入/公開狀態。','只存搜尋摘要、轉貼或沒有網址列的局部圖片。'),
('上下文','保留標題、署名、日期、前後段落、圖片說明及更正/更新註記；紙本記封面、卷期與連續頁次。','剪掉贊助聲明、前後文或更正欄。'),
('位元組與版本','原檔唯讀保存，計算 SHA-256；修正版另存新版本並記錄兩個 hash，不覆蓋舊版。','只記檔名，或只保存可任意覆寫的雲端連結。'),
('引用與著作權','先核定可引用範圍；工作底稿只摘錄必要內容，完整原件存於受控位置。','將整篇文章公開轉貼或未核授權就散布。'),
('個資與存取','僅保留核驗所需個資；評核展示使用代碼或遮罩；限制可存取角色與保存期限。','將未遮罩名單或聯絡方式放入公開 PR/附件。'),
('更正與撤稿','定期核對原始來源是否更新；更正、撤稿或網址失效時保留時間戳與結果，重新判斷可用範圍。','用已撤稿版本宣稱當前有效，或不記錄版本差異。'),
], [1.1,3.55,2.35])

add_h(doc,'8. 與其他文件連結')
add_table(doc,['文件/requirement','可共享的欄位','連結成立條件','禁止的替代'],[
('0136／I19-E26','case_id、活動/合作/表揚識別、報導 source_id。','同一事件可由媒體原件和相應原始活動/合作/表揚原件共同定位；兩份各保留自己的來源 hash。','媒體文章不能取代獎狀、合作協議或公益活動原件。'),
('0135／I19-E25','市場/服務主題、觀察期間、來源識別。','同一 market fact_id 必須在報導與可重算的市場資料都存在，且範圍一致。','單篇正面新聞不能當作長期趨勢或市場價值數據。'),
('0110／I18-E19','只有 same_case_id 確認後，才可能對照課程設計/評估。','先用同一真實來源記錄核對課程、日期、對象；目前示例 case_id 為 SYN-0137-CASE-01，不與 0110 真實個案共用。','不得因兩文件相鄰或題材相似就拼接不同案例。'),
], [1.2,1.45,2.55,1.8])

add_h(doc,'9. 媒體原件索取文字範本')
add_p(doc,'以下只是一段可供授權人員改寫的索取稿，不表示本會已聯絡任何媒體。寄送前須由有權限者確認收件人、引用範圍及個資。')
add_table(doc,['索取稿欄位','範本文字'],[
('主旨','請協助提供報導完整來源與引用範圍'),
('正文','您好：本會正在核對一項媒體佐證的完整來源。若貴媒體曾刊載/播出與［活動或課程名稱］相關內容，敬請提供可回查的原始網址、刊物名稱/卷期頁次或節目與時間碼、作者/發布日期、完整上下文及更正版本資訊，並告知可引用的文字或截圖範圍。此為待授權人員確認的範本，不代表已寄出或已獲回覆。謝謝。'),
('寄送前必核','收件對象真實；本會有權提出請求；不夾帶學員名單；寄件人使用授權管道；寄出後才記錄日期、回覆及原件。'),
], [1.4,5.6])

add_h(doc,'10. HOLD 與 REJECT 的分流')
add_table(doc,['情境','決策','理由與下一步'],[
('只有截圖，缺完整媒體頁、日期或上下文。','HOLD','索取完整原件與發布定位；補齊前不列為已核原件。'),
('來源頁有效，但報導中的數字尚未與本會原始紀錄核對。','HOLD claim；來源可暫登錄','逐 claim 對照來源；媒體原文不能代替活動、出席或成效紀錄。'),
('原件經媒體來源核實，但授權範圍或個資遮罩未完成。','HOLD 使用','完成授權/遮罩與存取控管前，不對外傳送。'),
('協會自己的新聞稿/付費稿被標成獨立媒體採訪。','REJECT 該媒體獨立性主張','保留來源性質，不能冒稱第三方採訪；如是付費內容需明確標示。'),
('偽造頁面、篡改日期/文字、錯認另一機構或蓄意裁掉更正資訊。','REJECT','封存差異證據、停止引用，交由有權限角色依內部程序處理。'),
('相同文章被多站轉載並重複計為多篇獨立報導。','REJECT 重複計數','以原始稿件/媒體編輯來源合併為一件，保留轉載清單但不增加獨立件數。'),
('報導非正面或與訓練/公益無關，卻被寫成 I19-E27 支持。','REJECT 該用途','可以依政策保留為背景資料，但不得冒充本要求的正面媒體佐證。'),
], [2.4,1.15,3.45])

add_h(doc,'11. 真實案例的逐件覆核清單')
checks=[
'□ 我能從媒體/出版者/節目原始來源回到同一篇報導。',
'□ 媒體、標題、作者、日期/時間、網址/頁次/時間碼和更正歷史已記錄。',
'□ 原件完整上下文已保存；SHA-256、檔案版本、取得時間與保管位置可重查。',
'□ 每個 claim_id 都已對照真正適用的活動/課程/結果原始資料，並列出未支持主張。',
'□ 正面訊息實際出現在原文；媒體關係、付費/贊助及轉載已揭露。',
'□ 著作權/引用與個資遮罩經授權角色核定；公開包沒有受保護原件或敏感資料。',
'□ HOLD 和 REJECT 有明確理由；未核資料沒有被算入「已驗證原件」。',
'□ 獨立覆核者核對同一原件與同一版本；記錄 exact SHA 和覆核範圍。',
'□ 跨文件僅連結相同 case_id 與可互相核驗的原始來源；未把媒體描述當作正式成果。',
'□ 最終結論仍明確區分 SAMPLE、REAL、candidate、Human formal acceptance。'
]
for c in checks:add_bullet(doc,c)

add_h(doc,'12. SAMPLE/SYNTHETIC 替換清冊')
register_rows=[
('SYN-0137-F01','SYN-0137-CASE-01','媒體名、類型與獨立性','虛構《示例社區報》；所有媒體身份均合成。','媒體官方刊物/網站/節目存檔及出版者來源。','未取得真實媒體原件前 HOLD。'),
('SYN-0137-F02','SYN-0137-CASE-01','標題、作者、日期與網址','虛構標題、署名、2026-05-18與 example.invalid。','完整原始報導、刊物卷期頁次或節目時間碼。','不得把保留網址改寫成真實來源。'),
('SYN-0137-F03','SYN-0137-CASE-01','活動日期與地點','假設活動日2026-05-10。','正式課程計畫/活動紀錄與場地來源。','需核同一 case_id。'),
('SYN-0137-F04','SYN-0137-CASE-01','登記/出席/完成數','24登記、22出席、20完成、2未達門檻。','經授權保存的匿名簽到/時數/完成名冊及計算表。','先核去重、定義及分母。'),
('SYN-0137-F05','SYN-0137-CASE-01','22/24及20/22比率','91.7%出席；90.9%出席者完成。','核對 F04 的原始列與公式重算。','不得把新聞文字當計算底稿。'),
('SYN-0137-F06','SYN-0137-CASE-01','廣泛能力提升說法','假設報導使用「增強社區應變」描述。','適用的前後測、演練量尺或獨立觀察紀錄。','沒有結果來源則 HOLD/刪除成效主張。'),
('SYN-0137-F07','SYN-0137-CASE-01','媒體截頁、hash、取得時間','假設的來源副本與 SHA，不是真實檔案。','原始檔 bytes、取得時戳、雜湊及受控保存紀錄。','本清冊中不提供假 hash。'),
('SYN-0137-F08','SYN-0137-CASE-01','引用/個資/保存權限','角色與授權狀態均合成。','實際著作權/媒體授權、個資遮罩與保存規則。','授權未知即 HOLD。'),
('SYN-0137-F09','SYN-0137-CASE-01','覆核與最終用途','示範覆核流程，沒有真實簽核人或核准。','獨立覆核紀錄及有權限者對引用用途的確認。','不得把示例決策記成協會決議。'),
]
add_table(doc,['fact_id／case_id','正文內容','SAMPLE/SYNTHETIC 值','REAL 替換所需來源','狀態'],[
(r[0]+'\n'+r[1],r[2],r[3],r[4],r[5]) for r in register_rows
], [1.25,1.4,1.8,1.8,1.0])

add_h(doc,'13. 可供主管抽查的結果摘要')
add_table(doc,['抽查欄位','目前候選值','使用規則'],[
('文件與官方要求','0137 / I19-E27 / index 151；official serial UNKNOWN。','依正式 crosswalk，不由標題推定方框流水號。'),
('真實報導原件數','UNKNOWN；目前未提供可核實的本會媒體原件。','沒有原件不得填 0 篇已查獲或 1 篇已完成。'),
('真實 claim 核對數','UNKNOWN/HOLD。','逐項標示已支持、未支持或不適用；未查 claim 不算完成。'),
('合成案例算式','22/24=91.7%；20/22=90.9%。','僅為樣本內算術核對，不是協會實績或官方分數。'),
('候選及接受','CANDIDATE；formal acceptance NOT_GRANTED。','GitHub同步或技術/版面PASS不代表正式接受。'),
], [1.55,2.35,3.35])
add_p(doc,'這份表完成的是 I19-E27 原件查驗方法與完整 SAMPLE 教學案例。要轉為協會真實佐證，必須由授權人員提供媒體原件、可追溯的協會事實來源與引用依據，完成逐 claim 獨立覆核，再由 Human 決定正式接受或送審。')

# Header/footer and document properties.
header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT; rr=header.add_run('TTQS ONE  |  DOCX 0137  |  I19-E27  |  CANDIDATE'); style_run(rr,8.5,False,'555555')
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('SAMPLE / SYNTHETIC 示範與 REAL 未核狀態須分開　　頁 '); style_run(r,8.5,False,'555555')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
doc.core_properties.title='媒體採訪與正面報導原件查驗表'
doc.core_properties.subject='TTQS I19-E27 DOCX 0137 候選 R01'
doc.core_properties.author='TTQS ONE candidate preparation'
doc.core_properties.comments='SAMPLE/SYNTHETIC candidate; real media source not supplied; not formally accepted.'
doc.save(DOCX)

# Deterministic direct content/math/technical checks; semantic review remains independent and pending.
d=Document(DOCX)
alltext='\n'.join([p.text for p in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells])
checks={
 'docx_zip_crc':zipfile.ZipFile(DOCX).testzip() is None,
 'docx_opens':True,
 'doc_id_requirement_bound':'0137' in alltext and 'I19-E27' in alltext,
 'printed_checkbox_not_invented':'UNKNOWN' in alltext,
 'official_quote_present':'提出報章雜誌採訪或電子媒體訪問之正面消息' in alltext,
 'index_151_source_requirement_present':all(x in alltext for x in ['151','真實報導來源與截頁索引','核心內容或完整 SAMPLE/SYNTHETIC 示範','與其他文件的事實連結']),
 'negative_neighbors_separated':'0136／I19-E26' in alltext and '0135／I19-E25' in alltext,
 'sample_boundary_present':all(x in alltext for x in ['SAMPLE / SYNTHETIC ONLY','SAMPLE/SYNTHETIC','不可引用為本會事實','REAL_MEDIA_SOURCE=UNKNOWN/HOLD']),
 'real_media_unknown_hold':'UNKNOWN/HOLD' in alltext,
 'no_formal_acceptance_claim':'NOT_GRANTED' in alltext,
 'reserved_invalid_domain_present':'example.invalid' in alltext,
 'sample_arithmetic_22_24':round(22/24*100,1)==91.7 and '91.7%' in alltext,
 'sample_arithmetic_20_22':round(20/22*100,1)==90.9 and '90.9%' in alltext,
 'hold_reject_recordable_boundaries':all(x in alltext for x in ['RECORDABLE_LIMITED','HOLD','REJECT']),
 'third_party_request_sample_present':'索取文字範本' in alltext,
 'replacement_register_facts_complete':all('SYN-0137-F'+f'{i:02d}' in alltext for i in range(1,10)),
 'cross_document_same_case_rule':'same_case_id' in alltext or '同一 case_id' in alltext,
 'no_real_association_claim':'沒有可回查的原始點協會媒體報導原件' in alltext,
}
# Extract font sizes and table dimensions.
from docx.oxml.ns import qn as _qn
body_sizes=[]; table_sizes=[]
for p in d.paragraphs:
 for run in p.runs:
  if run.text.strip() and run.font.size: body_sizes.append(run.font.size.pt)
for t in d.tables:
 for row in t.rows:
  for cell in row.cells:
   for p in cell.paragraphs:
    for run in p.runs:
     if run.text.strip() and run.font.size: table_sizes.append(run.font.size.pt)
checks['body_font_floor_10_5']=min(body_sizes,default=0)>=10.5
checks['table_font_floor_9_5']=min(table_sizes,default=0)>=9.5
with zipfile.ZipFile(DOCX) as z: zip_bytes=len(z.read('word/document.xml'))
summary={'project_id':'TTQS_ONE','doc_id':'0137','requirement_id':'I19-E27','candidate_state':'CANDIDATE_NOT_ACCEPTED','official_checkbox_serial':'UNKNOWN','docx_path':str(DOCX),'docx_bytes':DOCX.stat().st_size,'docx_sha256':h(DOCX.read_bytes()),'paragraph_count':len(d.paragraphs),'table_count':len(d.tables),'body_font_min_pt':min(body_sizes,default=None),'table_font_min_pt':min(table_sizes,default=None),'document_xml_bytes':zip_bytes,'checks':checks,'technical_result':'PASS' if all(checks.values()) else 'FAIL','visual_status':'PENDING_WORD_NATIVE_RENDER_AND_PAGE_INSPECTION','semantic_status':'PENDING_INDEPENDENT_EXACT_SHA_REVIEW','real_media_source':'UNKNOWN/HOLD','formal_acceptance':'NOT_GRANTED','source_packet_sha256':h((ROOT/'SOURCE_PACKET_0137_I19-E27_R01.md').read_bytes()),'blueprint_sha256':h((ROOT/'UNIQUE_BLUEPRINT_0137_I19-E27_R01.md').read_bytes()),'source_manifest_sha256':h((ROOT/'SOURCE_MANIFEST_0137_R01.json').read_bytes()),'source_commit':REF}
write_text(ROOT/'TECHNICAL_QA_0137_R01.json',json.dumps(summary,ensure_ascii=False,indent=2))
print(json.dumps(summary,ensure_ascii=False,indent=2))