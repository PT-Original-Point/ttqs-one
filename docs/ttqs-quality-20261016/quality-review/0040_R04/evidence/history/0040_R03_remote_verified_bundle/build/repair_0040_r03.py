from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

src=Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0040_R02\candidate\0040__訓練評估方法與成效指標對照表_R02.docx")
out=Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0040_R03\candidate\0040__訓練評估方法與成效指標對照表_R03.docx")
doc=Document(src)

replacements=[
 ("R05 示範評分角色覆核","評量覆核角色（示範職務）負責覆核"),
 ("R05 示範評分角色","評量角色（示範職務）"),
 ("R05 評量角色（示範代碼）","評量角色（示範職務）"),
 ("R03 課程設計角色（示範代碼）","課程設計角色（示範職務）"),
 ("R03 示範課程角色安排","課程設計角色安排"),
 ("R05示範評分角色","評量角色（示範職務）"),
]
paragraphs=list(doc.paragraphs)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            paragraphs.extend(cell.paragraphs)
for p in paragraphs:
    original=p.text
    updated=original
    for old,new in replacements:
        updated=updated.replace(old,new)
    if updated!=original:
        runs=p.runs
        if runs:
            runs[0].text=updated
            for run in runs[1:]:
                run.text=""
        else:
            p.add_run(updated)

remaining=[p.text for p in paragraphs if "R03" in p.text or "R05" in p.text]
if remaining:
    raise SystemExit("Unresolved R03/R05 codes in body: "+repr(remaining))

target=next((p for p in doc.paragraphs if p.text.startswith("官方要求：")),None)
if target is None:
    raise SystemExit("Could not locate closing official requirement paragraph")

def paragraph_before(ref, text, style=None, keep=False):
    element=OxmlElement("w:p")
    ref._p.addprevious(element)
    p=Paragraph(element,ref._parent)
    if style:
        p.style=style
    run=p.add_run(text)
    p.paragraph_format.keep_with_next=keep
    return p

heading=paragraph_before(target,"未達門檻的改善追蹤示例（SAMPLE / SYNTHETIC）","Heading 2",True)
note=paragraph_before(target,"下列 D04 只示範如何把未達門檻轉為有風險、責任角色、期限及後續交接的可追蹤安排；日期與角色皆為合成資料，未經協會核准且尚未執行。","Normal",True)
note.paragraph_format.space_after=Pt(4)

table=doc.add_table(rows=2,cols=5)
table.style="Table Grid"
table.alignment=WD_TABLE_ALIGNMENT.CENTER
table.autofit=False
table._tbl.addprevious(target._p)
widths=[1.22,1.42,1.18,1.10,2.19]
headers=["差距／未達門檻","示範風險","責任角色","示範期限／狀態","交接資料／結案規則"]
values=[
 "D04：後測 5/8；低於 6/8 門檻。",
 "K2 澄清摘要及 K4 交接未完整；屬訓練練習風險，不是實際服務事件。",
 "課程設計角色修訂練習；評量覆核角色以同版量尺評新情境。",
 "2026-06-19（SAMPLE 目標期限）；本例尚未執行，沒有重評結果。",
 "移交修訂教材、責任與期限紀錄、等難度重評表；四項完整、總分至少 6/8 且 K3 至少 1 分才可結案。"
]
for i,w in enumerate(widths):
    table.columns[i].width=Inches(w)
    for row in table.rows:
        row.cells[i].width=Inches(w)
for row_index,row in enumerate(table.rows):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for col_index,cell in enumerate(row.cells):
        cell.width=Inches(widths[col_index])
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text=headers[col_index] if row_index==0 else values[col_index]
        tcPr=cell._tc.get_or_add_tcPr()
        margins=OxmlElement("w:tcMar")
        for side,value in [("top","70"),("bottom","70"),("start","75"),("end","75")]:
            node=OxmlElement("w:"+side)
            node.set(qn("w:w"),value)
            node.set(qn("w:type"),"dxa")
            margins.append(node)
        tcPr.append(margins)
        if row_index==0:
            shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"17365D");tcPr.append(shade)
        else:
            shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"F2F6FA");tcPr.append(shade)
        for p in cell.paragraphs:
            p.paragraph_format.space_before=Pt(0)
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1.0
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if row_index==0 else WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name="Microsoft JhengHei"
                run.font.size=Pt(9.5)
                run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"Microsoft JhengHei")
                if row_index==0:
                    run.font.bold=True
                    run.font.color.rgb=__import__("docx").shared.RGBColor(255,255,255)
if table.rows:
    header=OxmlElement("w:tblHeader")
    header.set(qn("w:val"),"true")
    table.rows[0]._tr.get_or_add_trPr().append(header)

doc.core_properties.title="訓練評估方法與成效指標對照表"
doc.core_properties.subject="可操作評量矩陣；SAMPLE / SYNTHETIC 示範"
doc.save(out)
print("OUTPUT",out)
print("TABLES",len(doc.tables))
print("INSERTED_ROW", " | ".join(values))