from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
p=Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0040_R04\candidate\0040__訓練評估方法與成效指標對照表_R04.docx")
d=Document(p)
control=next(t for t in d.tables if t.rows and t.rows[0].cells[0].text.strip()=="控制欄位")
if not any(r.cells[0].text.strip()=="官方映射" for r in control.rows[1:]):
    r=control.add_row()
    r.cells[0].text="官方映射"
    r.cells[1].text="DOCX 0040 → I08-E03 → 官方 checkbox 040；公開來源為官方查核佐證表 PDF 第 3 頁，SHA-256 805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709。此為來源定位，不是協會佐證。"
    r._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for ci,cell in enumerate(r.cells):
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tcPr=cell._tc.get_or_add_tcPr()
        shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),"F2F6FA"); tcPr.append(shd)
        margins=OxmlElement("w:tcMar")
        for side,amount in (("top","55"),("bottom","55"),("start","65"),("end","65")):
            n=OxmlElement("w:"+side); n.set(qn("w:w"),amount); n.set(qn("w:type"),"dxa"); margins.append(n)
        tcPr.append(margins)
        for para in cell.paragraphs:
            para.paragraph_format.space_before=Pt(0); para.paragraph_format.space_after=Pt(1)
            for run in para.runs:
                run.font.name="Microsoft JhengHei"; run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"Microsoft JhengHei"); run.font.size=Pt(9.5)
index=next(t for t in d.tables if t.rows and t.rows[0].cells[0].text.strip()=="證據目的／內容")
index.rows[1].cells[1].text="第1頁「官方要求（逐字引用）」及「來源」段；本頁「目前文件控制」表之「官方映射」列。"
for para in index.rows[1].cells[1].paragraphs:
    for run in para.runs:
        run.font.name="Microsoft JhengHei"; run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"Microsoft JhengHei"); run.font.size=Pt(9.5)
d.save(p)
print("PATCHED_OFFICIAL_MAPPING_ROW",len(control.rows))
