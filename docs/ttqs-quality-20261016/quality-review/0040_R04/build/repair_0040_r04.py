from pathlib import Path
from shutil import copy2
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import csv, sys

src_doc = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0040_R03\candidate\0040__訓練評估方法與成效指標對照表_R03.docx")
src_reg = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0040_R03\sources\0040_SOURCES_AND_SYNTHETIC_REPLACEMENT_REGISTER_R03.csv")
out_dir = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0040_R04")
out_doc = out_dir / "candidate" / "0040__訓練評估方法與成效指標對照表_R04.docx"
out_reg = out_dir / "sources" / "0040_SOURCES_AND_SYNTHETIC_REPLACEMENT_REGISTER_R04.csv"
script_path = out_dir / "build" / "repair_0040_r04.py"

if out_doc.exists():
    raise SystemExit("Refusing to overwrite an existing R04 candidate.")
if not src_doc.is_file() or not src_reg.is_file():
    raise SystemExit("R03 source candidate or register missing.")
for folder in ("candidate", "sources", "qa", "render", "history", "build"):
    (out_dir / folder).mkdir(parents=True, exist_ok=True)
copy2(src_doc, out_doc)
copy2(src_reg, out_reg)

doc = Document(out_doc)
sample_para = next((p for p in doc.paragraphs if p.text.startswith("示範界線：")), None)
if sample_para is None:
    raise SystemExit("Could not locate the existing sample-boundary paragraph.")
if "SAMPLE-0039-01" not in sample_para.text:
    sample_para.add_run(" 該示範案例代碼：SAMPLE-0039-01。")

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    trPr.append(node)

def add_table(rows, widths, header_fill="17365D", font_size=9.5):
    table = doc.add_table(rows=1, cols=len(widths))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, width in enumerate(widths):
        table.columns[i].width = Inches(width)
    for ri, values in enumerate(rows):
        row = table.rows[0] if ri == 0 else table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for ci, value in enumerate(values):
            cell = row.cells[ci]
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = str(value)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.name = "Microsoft JhengHei"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
                    run.font.size = Pt(font_size)
                    if ri == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            tcPr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side, amount in (("top","55"),("bottom","55"),("start","65"),("end","65")):
                node = OxmlElement("w:" + side)
                node.set(qn("w:w"), amount)
                node.set(qn("w:type"), "dxa")
                margins.append(node)
            tcPr.append(margins)
            shade(cell, header_fill if ri == 0 else ("F2F6FA" if ri % 2 else "FFFFFF"))
    set_repeat_header(table.rows[0])
    return table

doc.add_page_break()
h = doc.add_paragraph("查核索引與文件版次紀錄", style="Heading 1")
h.paragraph_format.keep_with_next = True
p = doc.add_paragraph("本頁只管理本候選工具的版本、效力狀態及內容定位；不構成原始點協會核准、採用或課程執行紀錄。")
p.paragraph_format.keep_with_next = True

doc.add_paragraph("目前文件控制", style="Heading 2").paragraph_format.keep_with_next = True
control_rows = [
    ["控制欄位", "目前記錄"],
    ["文件版次", "R04（候選修訂）；前版 R03 SHA-256：c84e200ff11d2e9f46360b4e03b7867e22fb03ac4ec298d022c9d1f3d76ffdf8。"],
    ["評量工具版次", "SAMPLE v1.0；僅供本文件合成案例示範，尚未核准、未生效，不能直接代表協會正式量尺。"],
    ["本次修訂日期", "2026-10-17（本地候選施工日期；不是協會核准或採用日期）。"],
    ["正式效力／核准", "無。正式使用前須由具權責人員依真實需求、課程及資料核實並完成核准；目前不得宣稱已採用。"],
    ["維護責任", "文件維護角色（實際人員與授權待核）；資料持有人與評量覆核角色須依真實權責表確認。"],
    ["本次修訂依據", "既有獨立 R03 exact-SHA 審查為 PARTIAL：補足可見版本／查核索引，並將替換清冊改為可在列印頁面直接找到的定位；既有評分案例與數據未改。"],
]
add_table(control_rows, [1.35, 5.35])

doc.add_paragraph("修訂歷史（均為候選版本，不代表正式採用）", style="Heading 2").paragraph_format.keep_with_next = True
history_rows = [
    ["版本／記錄日", "修訂內容與稽核依據", "效力狀態"],
    ["R01／2026-10-16（審查記錄日）", "初版候選；R01 獨立審查所列來源定位、內部代碼、版面與證據界線問題，詳本包 R01 歷史收據。", "歷史候選；未接受"],
    ["R02／2026-10-17（審查記錄日）", "依 R01 修正來源、正文角色表述與版面；獨立審查仍要求補上 D04 差距—風險—責任—期限—交接—結案單列。", "PARTIAL；未接受"],
    ["R03／2026-10-17", "加入 D04 完整合成追蹤列並統一角色文字；獨立 exact-SHA 審查發現版次／索引缺漏及 34/39 筆虛構 § 定位。", "PARTIAL；未接受"],
    ["R04／2026-10-17", "新增本頁版本控制與查核索引；在案例界線列印案例碼；將清冊改為頁碼、原文表頭、列標籤與欄位定位；修正清冊中過時角色代碼。", "候選；待獨立 exact-SHA 複核"],
]
add_table(history_rows, [1.35, 4.1, 1.25])

doc.add_paragraph("內容查核索引", style="Heading 2").paragraph_format.keep_with_next = True
index_rows = [
    ["證據目的／內容", "列印頁面定位（依本版 Word 頁碼）", "直接核對方式"],
    ["官方要求與來源", "第1頁「官方要求（逐字引用）」及「來源」段；本頁「目前文件控制」與下列「官方映射」索引列。", "核對官方逐字要求、來源頁次與本文件用途；不得推論協會已達標。"],
    ["使用方式、資料界線與案例", "第1頁「使用者／使用時點」段、「示範界線」段及文件控制表。", "案例代碼 SAMPLE-0039-01；課名、示範日期、SAMPLE/SYNTHETIC 狀態及未取得真實資料的限制。"],
    ["評量工作順序", "第1頁表「步驟｜操作｜完成後留下的可查結果」，列1至列5。", "依序核需求、定目標、定工具、定時點、重算與改善；核對輸入與輸出。"],
    ["同案與不同件邊界", "第2頁表「文件／資料｜本件如何引用｜不得推論」，列「配套會議示範案例」及其他列。", "確認本矩陣只承接評量工具，不代替會議紀錄、需求調查、系統改善或完整課程計畫。"],
    ["設計流程與覆蓋率", "第2頁表「流程節點｜可核對輸出」；其後「設計覆蓋率」段。", "核對需求至課程修訂六節點；4/4=100% 只表示設計覆蓋。"],
    ["目標、方法、時點、門檻、證據與責任", "第3頁表「差距與目標｜直接評量／工具｜時點與門檻｜證據欄位、定位及責任」。", "逐列查 K1、K2、K3、K4；核對基線、差距、觀察、角色及證據欄位。"],
    ["評分量尺", "第3頁表「構面｜0分：未達｜1分：部分達成｜2分：完整達成｜直接觀察重點」。", "核對 K1–K4 0/1/2 行為錨點與 K3 安全關鍵門檻；其後「計分」段。"],
    ["五筆示範資料與計算", "第4頁表「合成代碼｜前測…｜後測…｜總分變化｜依門檻判定」；表「指標｜定義／分母｜代入合成資料｜結果及解讀限制」。", "重算 D01–D05、前後測平均、配對平均變化、達標率及構面平均。"],
    ["個案評分與異常處理", "第5頁表「欄位｜完整合成填寫示例｜實際使用時檢查」；表「觸發狀況｜當下處理｜統計／結案規則」。", "逐欄查 D04 個案及缺漏、隱私、評分差異、版本不一致、N=0 與門檻修訂處理。"],
    ["交接與真實化清單", "第5–6頁表「交接對象／步驟｜必交資料｜未完成時」。", "核對角色交接、配套會議資料、真實來源、授權、原始評分、版本、核准與附件定位。"],
    ["差距到結案的追蹤", "第6頁表「差距／未達門檻｜示範風險｜責任角色｜示範期限／狀態｜交接資料／結案規則」，D04 列。", "驗證風險、責任、SAMPLE 期限、交接物與結案門檻完整連結；本例未執行。"],
    ["版次與修訂歷史", "本頁「目前文件控制」表及「修訂歷史」表。", "確認候選版次、工具版次、修訂依據、未核准狀態及 R01–R04 差異。"],
]
add_table(index_rows, [1.25, 3.15, 2.3], font_size=9.5)

doc.core_properties.title = "訓練評估方法與成效指標對照表"
doc.core_properties.subject = "R04 候選；SAMPLE/SYNTHETIC 評量矩陣工具；未核准未生效"
doc.core_properties.revision = 4
doc.core_properties.comments = "候選修訂 R04。SAMPLE/SYNTHETIC only. Not approved or effective."
doc.save(out_doc)

print("OUTPUT_DOCX", out_doc)
print("OUTPUT_REGISTER", out_reg)
print("DOC_TABLES", len(doc.tables))
print("PARAGRAPH_CASE_CODE", "SAMPLE-0039-01" in sample_para.text)
