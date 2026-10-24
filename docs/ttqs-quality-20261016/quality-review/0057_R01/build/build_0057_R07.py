from __future__ import annotations

import hashlib
import json
import shutil
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
R06 = ROOT / "candidate" / "0057_R06.docx"
R07 = ROOT / "candidate" / "0057_R07.docx"
CSV_R06 = ROOT / "sidecar" / "0057_R06_SYNTHETIC_REPLACEMENT_REGISTER.csv"
JSON_R06 = ROOT / "sidecar" / "0057_R06_SYNTHETIC_REPLACEMENT_REGISTER.json"
CSV_R07 = ROOT / "sidecar" / "0057_R07_SYNTHETIC_REPLACEMENT_REGISTER.csv"
JSON_R07 = ROOT / "sidecar" / "0057_R07_SYNTHETIC_REPLACEMENT_REGISTER.json"
RECEIPT = ROOT / "build" / "0057_R07_build_receipt.md"
EXPECTED_R06 = "3D4CA6C5D1BD0262A5CE47A1339021C4F7563D9E9F1D06C82676D9BB002DD472"
OLD = "且不列入 D；兩者均留在分母 C。"
NEW = "且不列入 D。"
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def sha(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest().upper()


def main() -> None:
    if sha(R06) != EXPECTED_R06:
        raise RuntimeError(f"R06 exact SHA mismatch: {sha(R06)}")
    if any(path.exists() for path in (R07, CSV_R07, JSON_R07, RECEIPT)):
        raise RuntimeError("R07 output already exists; refusing to overwrite")

    part_changes = 0
    package_changes: list[str] = []
    with zipfile.ZipFile(R06, "r") as source, zipfile.ZipFile(R07, "w") as output:
        names = source.namelist()
        if len(names) != len(set(names)):
            raise RuntimeError("Source package contains duplicate ZIP members")
        if "word/document.xml" not in names:
            raise RuntimeError("word/document.xml is missing")
        old_bytes = OLD.encode("utf-8")
        new_bytes = NEW.encode("utf-8")
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "word/document.xml":
                count = data.count(old_bytes)
                if count != 1:
                    raise RuntimeError(f"Expected exactly one duplicate Appendix C phrase; found {count}")
                data = data.replace(old_bytes, new_bytes, 1)
                part_changes += 1
                package_changes.append(info.filename)
            output.writestr(info, data)

    shutil.copyfile(CSV_R06, CSV_R07)
    shutil.copyfile(JSON_R06, JSON_R07)

    with zipfile.ZipFile(R06, "r") as before, zipfile.ZipFile(R07, "r") as after:
        if before.namelist() != after.namelist():
            raise RuntimeError("R06→R07 package member inventory changed")
        for name in before.namelist():
            if name == "word/document.xml":
                if before.read(name).replace(OLD.encode("utf-8"), NEW.encode("utf-8"), 1) != after.read(name):
                    raise RuntimeError("word/document.xml contains changes beyond the exact duplicate removal")
            elif before.read(name) != after.read(name):
                raise RuntimeError(f"Unchanged package part differs: {name}")
        if after.testzip() is not None:
            raise RuntimeError("R07 DOCX ZIP integrity failed")
        footer_files = [n for n in after.namelist() if n.startswith("word/footer") and n.endswith(".xml")]
        if len(footer_files) != 5:
            raise RuntimeError(f"Expected five footer parts; found {len(footer_files)}")
        for name in footer_files:
            root = ET.fromstring(after.read(name))
            page_fields = 0
            for field in root.findall(".//w:fldSimple", NS):
                if (field.get(f"{{{NS['w']}}}instr") or "").strip().upper() == "PAGE":
                    page_fields += 1
            for instr in root.findall(".//w:instrText", NS):
                if (instr.text or "").strip().upper() == "PAGE":
                    page_fields += 1
            if page_fields != 1:
                raise RuntimeError(f"Expected one PAGE field in {name}; found {page_fields}")

    doc = Document(R07)
    all_text = "\n".join(p.text for p in doc.paragraphs)
    if OLD in all_text or all_text.count(NEW) < 1:
        raise RuntimeError("Duplicate phrase remains or corrected sentence is missing")
    if len(doc.tables) != 22 or len(doc.sections) != 5 or len(doc.inline_shapes) != 2:
        raise RuntimeError("Table, section, or inline-image inventory changed")
    if sha(CSV_R06) != sha(CSV_R07) or sha(JSON_R06) != sha(JSON_R07):
        raise RuntimeError("R07 replacement register copies are not byte-identical")

    receipt = (
        "# 0057 R07 bounded copy-edit receipt\n\n"
        "## Identity\n"
        "- Document: DOCX 0057; official requirement I13-E01; checkbox 061\n"
        "- Source candidate: `candidate/0057_R06.docx`\n"
        f"- Source SHA-256: `{sha(R06)}`\n"
        "- Output candidate: `candidate/0057_R07.docx`\n"
        f"- Output SHA-256: `{sha(R07)}`\n"
        "- Source Blueprint: `inputs/0057_UNIQUE_BLUEPRINT_R09.md`\n"
        "- Blueprint SHA-256: `A6602B5B37250F03914A075CBAD909D469C9AAE71383BEAD85006DF334899E02`\n"
        "- Bounded patcher: `build/build_0057_R07.py`\n"
        f"- Bounded patcher SHA-256: `{sha(Path(__file__))}`\n\n"
        "## Change and checks\n"
        "- Removed the repeated final clause `兩者均留在分母 C。` from Appendix C after the same sentence had already stated that both groups remain in C and are excluded from D.\n"
        "- Only `word/document.xml` changed; every other uncompressed OOXML package part is byte-identical to R06. No values, dates, tables, charts, styles, sections, or decisions changed.\n"
        f"- R07 CSV SHA-256: `{sha(CSV_R07)}`; byte-identical to R06.\n"
        f"- R07 JSON SHA-256: `{sha(JSON_R07)}`; byte-identical to R06.\n"
        "- DOCX integrity, 22 tables, two inline images, five sections, five independent footer parts and one PAGE field per footer passed this bounded check.\n\n"
        "## Limits\n"
        "This receipt records a bounded copy edit only. It is not independent technical acceptance, semantic acceptance, visual approval, or formal association acceptance. Render and inspect exact R07 before review.\n"
    )
    RECEIPT.write_text(receipt, encoding="utf-8", newline="\n")
    print(json.dumps({
        "source_sha256": sha(R06),
        "candidate_sha256": sha(R07),
        "patcher_sha256": sha(Path(__file__)),
        "changed_package_parts": package_changes,
        "replacement_count": part_changes,
        "tables": len(doc.tables),
        "sections": len(doc.sections),
        "inline_shapes": len(doc.inline_shapes),
        "sidecar_csv_sha256": sha(CSV_R07),
        "sidecar_json_sha256": sha(JSON_R07),
        "status": "BOUNDED_COPY_EDIT_COMPLETE_NOT_ACCEPTED",
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
