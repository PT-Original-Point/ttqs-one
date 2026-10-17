from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
from pathlib import Path
import csv, json, hashlib, shutil

ROOT=Path(r"E:\TTQS\_TTQS_ONE_REVIEW_SYNC_0036_R08_20261016\docs\ttqs-quality-20261016\quality-review\0038_R02")
BLUEPRINT=Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\CONTROL\BLUEPRINT_CURRENT.jsonl")
OFFICIAL_DIR=Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0038_PREFLIGHT")
DOCX=ROOT/"candidate"/"0038_R02.docx"
ASSETS=ROOT/"evidence"/"assets"
INPUTS=ROOT/"inputs"
REGISTER=ROOT/"evidence"/"SYNTHETIC_REPLACEMENT_REGISTER.csv"
SPEC_HEAD="56de53f2dfb46aa3513716b6ad595d39e4b207de"
OFFICIAL_SHA="805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709"
BP_FILE_SHA="406d0495365d0970dcf8113a42842663b8a54372a4e204cc492b7c3f8156ef3e"
BP_ROW_SHA="2504485c1698a5c83ac2205938c05b532edc8d697f2441ad1d6a0c324897b73e"
SPEC_BLOBS={"01":"538bedb9d87e8465f43151933b4a5c1bfa4bbc4e","02":"4ff863d9f534b6359fc1c7c1685bd8cce02da886","03":"8c2bbe7ca9dcbb03088a48934b9e57e7147ede63","04":"62e3dd7dfd8579b6bc97952ffa6236e7494f3b4d","05":"60ae010c58140e1710ed24016085c6dfaf6ffaee"}
for p in (DOCX.parent,ASSETS,INPUTS,ROOT/"evidence"/"build"): p.mkdir(parents=True,exist_ok=True)
if DOCX.exists(): raise SystemExit("Refusing to overwrite existing 0038_R02.docx")

A=[
("A01",(1,1,1,0),(1,1,1,1)),("A02",(1,1,1,1),(1,1,1,2)),
("A03",(0,1,1,0),(1,1,1,0)),("A04",(1,1,1,2),(2,1,1,2)),
("A05",(1,1,1,0),(1,1,1,1)),("A06",(1,1,1,1),(1,1,1,2)),
("A07",(1,1,1,2),(2,1,1,2)),("A08",(0,1,1,0),(1,1,1,0)),
("A09",(1,1,1,1),(1,1,1,2)),("A10",(1,1,1,0),(1,1,1,1)),
("A11",(1,1,1,2),(2,1,2,2)),("A12",(1,1,1,1),(1,1,1,2))]
B=[
("B01",(1,0,1,1),(2,2,2,1)),("B02",(1,2,0,1),(2,2,1,2)),
("B03",(0,1,1,0),(1,2,1,2)),("B04",(1,1,1,2),(2,2,2,2)),
("B05",(0,1,1,1),(1,2,1,2)),("B06",(2,1,0,1),(2,1,1,2)),
("B07",(2,1,1,1),(2,1,2,2)),("B08",(0,1,0,1),(1,1,1,2)),
("B09",(1,1,1,1),(2,2,2,2)),("B10",(1,1,0,1),(1,2,1,2)),
("B11",(1,2,1,1),(2,2,1,2)),("B12",(1,0,2,1),(1,1,2,2))]
def sm(v): return sum(v)
def vec(v): return "/".join(str(x) for x in v)
def ok(v): return sm(v)>=6 and v[2]>0 and v[3]>0
ap=sum(sm(x[1]) for x in A); aq=sum(sm(x[2]) for x in A)
bp=sum(sm(x[1]) for x in B); bq=sum(sm(x[2]) for x in B)
ac=sum(ok(x[2]) for x in A); bc=sum(ok(x[2]) for x in B)
az=sum(x[2][2]==0 or x[2][3]==0 for x in A); bz=sum(x[2][2]==0 or x[2][3]==0 for x in B)
assert (ap,aq,bp,bq,ac,bc,az,bz)==(44,57,44,79,3,11,2,0)

font_path=r"C:\Windows\Fonts\msjh.ttc"
def ft(n): return ImageFont.truetype(font_path,size=n,index=0)
def center(draw,box,text,fnt,fill):
    x0,y0,x1,y1=box
    bb=draw.multiline_textbbox((0,0),text,font=fnt,spacing=8,align="center")
    draw.multiline_text((x0+(x1-x0-(bb[2]-bb[0]))/2,y0+(y1-y0-(bb[3]-bb[1]))/2),text,font=fnt,fill=fill,spacing=8,align="center")
def flow_image(path):
    im=Image.new("RGB",(1900,420),"white"); d=ImageDraw.Draw(im)
    d.text((25,24),"0038 專屬流程｜設計缺口 → 選擇 → 執行 → 成效 → 決定",font=ft(30),fill="#172B35")
    names=["確認缺口\n量尺不足","比較方案\n選定改法","修訂設計\n固定評量","分組試行\n完整留痕","核對資料\n重算結果","判讀限制\n決定續行"]
    xs=[24,340,656,972,1288,1604]; y0,y1=105,315; w=255
    colors=["#EAF1F5","#EAF1F5","#DFEEF0","#DFEEF0","#EAF1F5","#F2EEE5"]
    for i,(x,name,col) in enumerate(zip(xs,names,colors)):
        d.rounded_rectangle((x,y0,x+w,y1),radius=18,fill=col,outline="#456475",width=4)
        center(d,(x+8,y0+8,x+w-8,y1-8),name,ft(35),"#172B35")
        if i<5:
            ax=x+w+10; ay=(y0+y1)//2
            d.line((ax,ay,ax+32,ay),fill="#456475",width=6)
            d.polygon([(ax+32,ay-12),(ax+52,ay),(ax+32,ay+12)],fill="#456475")
    im.save(path)
def chart_image(path):
    im=Image.new("RGB",(1700,760),"white"); d=ImageDraw.Draw(im)
    d.text((55,28),"SAMPLE/SYNTHETIC｜同量尺前後測平均分（0–8）",font=ft(34),fill="#172B35")
    l,r,t,b=420,1515,150,652
    for k in range(5):
        v=k*2; x=l+(r-l)*v/8
        d.line((x,t,x,b),fill="#D9E0E4",width=2); d.text((x-16,b+15),str(v),font=ft(22),fill="#53656D")
    series=[("講授為主 A",ap/12,aq/12,"#7793A2","#315B72"),("情境修訂 B",bp/12,bq/12,"#80AAA3","#2E756B")]
    for (label,pre,post,c1,c2),y in zip(series,(245,465)):
        d.text((55,y+40),label,font=ft(27),fill="#172B35")
        for off,val,col,tag in ((0,pre,c1,"前測"),(82,post,c2,"後測")):
            end=l+(r-l)*val/8
            d.rounded_rectangle((l,y+off,r,y+off+54),radius=9,fill="#F3F5F6")
            d.rounded_rectangle((l,y+off,end,y+off+54),radius=9,fill=col)
            d.text((end+10,y+off+8),f"{tag} {val:.2f}",font=ft(21),fill="#172B35")
    d.text((55,710),"兩組非隨機試驗；差異只供示範判讀，不是協會實績或因果證明。",font=ft(21),fill="#53656D")
    im.save(path)
flow=ASSETS/"design_improvement_flow.png"; chart=ASSETS/"synthetic_prepost_comparison.png"
flow_image(flow); chart_image(chart)
for nm in ("official-2.png","official-3.png"): shutil.copy2(OFFICIAL_DIR/nm,INPUTS/nm)
bp_row=next(x for x in BLUEPRINT.read_text(encoding="utf-8-sig").splitlines() if '"deliverable_id":"0038"' in x)
if hashlib.sha256(bp_row.encode("utf-8")).hexdigest()!=BP_ROW_SHA: raise SystemExit("Blueprint row SHA mismatch")
(INPUTS/"BLUEPRINT_CURRENT_0038.jsonl").write_text(bp_row+"\n",encoding="utf-8")
source_map=f"""# 0038 R02 source map

- Project: TTQS_ONE; DOCX 0038; requirement I08-E02; official checkbox sequence 039.
- Official wording: 針對訓練方案的系統設計能適時檢討及改善，執行後有良好效果。
- Official locator: 勞動力發展署《TTQS 查核佐證文件資料表—訓練機構版》，I08 p.2; continuation context p.3.
- Official PDF SHA-256: {OFFICIAL_SHA}; locally inspected page images: official-2.png and official-3.png.
- Current spec branch/head: work/ttqs-one-document-quality-contract-20261016 / {SPEC_HEAD}.
- Spec blob IDs: 01 {SPEC_BLOBS["01"]}; 02 {SPEC_BLOBS["02"]}; 03 {SPEC_BLOBS["03"]}; 04 {SPEC_BLOBS["04"]}; 05 {SPEC_BLOBS["05"]}.
- Runtime blueprint file SHA-256: {BP_FILE_SHA}; exact 0038 row SHA-256: {BP_ROW_SHA}.
- Historical 142-ID crosswalk maps 0038 to I08-E02 and identifies the prior body as incomplete; old body use is prohibited.
- Negative neighbors: 0037/I08-E01 is multi-period design review; 0039/I08-E03 is a planning-meeting record. This sample is independent and claims no shared case.
- No REAL association execution, approval, participants, measurement, or outcome has been authenticated. Dates, roles, course, options, scores and decisions are SAMPLE/SYNTHETIC.
"""
(INPUTS/"SOURCE_MAP.md").write_text(source_map,encoding="utf-8")

reg=[]
def addreg(fid,loc,val,real):
    reg.append([fid,"SYN-I08E02-2026-01","0038",loc,val,"完整示範假設，非協會事實",real,"待協會核實之授權角色（身分 UNKNOWN）","SYNTHETIC_REPLACE_BEFORE_REAL"])
addreg("SYN-CASE-001","首頁案例識別","假設課程：社區關懷志工服務溝通與轉介界線","正式課程代碼、計畫、適用對象")
addreg("SYN-PERIOD-001","§2 執行期間","比較組 2026-09-08；修訂組 2026-09-15","真實課程日期、簽到及執行紀錄")
addreg("SYN-POP-001","§2 樣本框","假設24名符合條件者，A/B各12名","真實招生名冊、納入規則及去識別抽樣紀錄")
addreg("SYN-FLOW-001","§5 樣本流向表","接觸24、邀請24、成對回收24、有效24、排除0","邀請、同意、前後測原件及排除紀錄")
addreg("SYN-ROLE-001","§2 角色流程表","示範引導者、觀察評分者、資料保管者；非現役人員","正式職務分工及授權資料")
addreg("SYN-BASELINE-001","§3 比較基準","120分鐘講授為主；非已查證的協會課程","原版課程計畫、教材、評量及版本核准紀錄")
addreg("SYN-CHANGE-001","§3–§4 修訂設計","總時數同為120分鐘；以情境演練、量尺與回饋取代35分鐘部分講授","修訂計畫、版次差異及執行證據")
addreg("SYN-OPTIONS-001","§3 方案評分表","權重40/25/20/15；加權總分2.95/4.75/4.10","需求來源、評分者記錄及選案依據")
addreg("SYN-RUBRIC-001","§5 評量量尺","D1–D4各0–2；總分0–8；達標≥6且D3、D4>0","量尺版次、評分工具與試評者校準")
addreg("SYN-THRESHOLD-001","§8 示範判讀規則","續行建議門檻：B組達標率≥80%且關鍵失誤=0","事前核准的評估方案與負責人簽認")
addreg("SYN-STORAGE-001","§5/§9 資料處理","假設使用匿名代碼；實際保存位置與權限 UNKNOWN","真實資料存取、保存、刪除和備份規則")
addreg("SYN-DECISION-001","§8 示範決定","建議下一輪小規模試行；非協會核定或採行","真實審查、核定、課程版控及後續追蹤")
addreg("SYN-ANOMALY-001","§7 B08覆核","後測1/1/1/2加總5；保留未達標，不排除","原始評量表及評分複核軌跡")
for group,rows,day,design in (("A",A,"2026-09-08","講授為主對照"),("B",B,"2026-09-15","情境練習修訂")):
    for i,(pid,pre,post) in enumerate(rows,1):
        addreg(f"SYN-{group}-ROW-{i:02d}",f"§6/§7及附表A {group}{i:02d}逐筆列",
            f"日期={day};組別={design};渠道=合成紙本情境卡/匿名評分表;同意=是（合成假設）;納入=是（完成前後測且D1-D4均有評分，符合事前完整配對規則）;缺漏=0;重複=0;更正=0（未更動原值）；前測={vec(pre)}合計{sm(pre)};後測={vec(post)}合計{sm(post)};差={sm(post)-sm(pre)};判定={'達標' if ok(post) else '未達標'}",
            "去識別前後測、課程簽到、同意及資料清理原件")
with REGISTER.open("w",encoding="utf-8-sig",newline="") as f:
    w=csv.writer(f); w.writerow(["fact_id","case_id","doc_id","document_locator","synthetic_value","assumption_reason","required_real_source_for_replacement","responsible_role_to_verify","status"]); w.writerows(reg)

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
sec.top_margin=Inches(.62); sec.bottom_margin=Inches(.62); sec.left_margin=Inches(.68); sec.right_margin=Inches(.68)
sec.header_distance=Inches(.28); sec.footer_distance=Inches(.28)
INK="172B35"; NAVY="234B62"; MUTED="53656D"; PALE="F5F8F9"; BORDER="D9D9D9"
def sf(run,size=11,bold=False,color=INK):
    run.font.name="Microsoft JhengHei"; run.font.size=Pt(size); run.font.bold=bold; run.font.color.rgb=RGBColor.from_string(color)
    rpr=run._element.get_or_add_rPr()
    rf=rpr.rFonts
    if rf is None: rf=OxmlElement("w:rFonts"); rpr.insert(0,rf)
    for k in ("ascii","hAnsi","eastAsia","cs"): rf.set(qn("w:"+k),"Microsoft JhengHei")
normal=doc.styles["Normal"]; normal.font.name="Microsoft JhengHei"; normal.font.size=Pt(11); normal.font.color.rgb=RGBColor.from_string(INK)
normal._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft JhengHei")
normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.14
for name,size in (("Title",24),("Heading 1",16),("Heading 2",12.5),("Heading 3",11.5)):
    st=doc.styles[name]; st.font.name="Microsoft JhengHei"; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string("000000")
    st._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft JhengHei")
    st.paragraph_format.space_before=Pt(8); st.paragraph_format.space_after=Pt(5)
for style_name in ("Title",):
    ppr=doc.styles[style_name]._element.get_or_add_pPr()
    border=ppr.find(qn("w:pBdr"))
    if border is not None: ppr.remove(border)
def P(text="",lead=None,size=11,color=INK,after=5,style=None):
    p=doc.add_paragraph(style=style) if style else doc.add_paragraph()
    p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.14
    if lead and text.startswith(lead):
        sf(p.add_run(lead),size,True,NAVY); sf(p.add_run(text[len(lead):]),size,False,color)
    else: sf(p.add_run(text),size,False,color)
    return p
pending_break=[False]
def H(text,level=1):
    p=doc.add_paragraph(style=f"Heading {level}")
    sf(p.add_run(text),{1:16,2:12.5,3:11.5}[level],True,"000000")
    p.paragraph_format.keep_with_next=True
    if pending_break[0]:
        p.paragraph_format.page_break_before=True
        pending_break[0]=False
def shade(cell,color):
    tc=cell._tc; pr=tc.get_or_add_tcPr(); sh=pr.find(qn("w:shd"))
    if sh is None: sh=OxmlElement("w:shd"); pr.append(sh)
    sh.set(qn("w:fill"),color)
def margins(cell):
    pr=cell._tc.get_or_add_tcPr(); node=pr.first_child_found_in("w:tcMar")
    if node is None: node=OxmlElement("w:tcMar"); pr.append(node)
    for edge,val in (("top",90),("start",100),("bottom",90),("end",100)):
        e=node.find(qn("w:"+edge))
        if e is None: e=OxmlElement("w:"+edge); node.append(e)
        e.set(qn("w:w"),str(val)); e.set(qn("w:type"),"dxa")
def borders(t):
    pr=t._tbl.tblPr; bd=pr.first_child_found_in("w:tblBorders")
    if bd is None: bd=OxmlElement("w:tblBorders"); pr.append(bd)
    for edge in ("top","left","bottom","right","insideH","insideV"):
        e=bd.find(qn("w:"+edge))
        if e is None: e=OxmlElement("w:"+edge); bd.append(e)
        e.set(qn("w:val"),"single"); e.set(qn("w:sz"),"5"); e.set(qn("w:color"),BORDER)
def table(headers,rows,widths,size=9.6,center_cols=()):
    size=max(size,9.5)
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; borders(t)
    for j,w in enumerate(widths): t.columns[j].width=Inches(w)
    hdr=t.rows[0]; trp=hdr._tr.get_or_add_trPr(); rep=OxmlElement("w:tblHeader"); rep.set(qn("w:val"),"true"); trp.append(rep)
    for j,val in enumerate(headers):
        c=hdr.cells[j]; c.width=Inches(widths[j]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(c); shade(c,NAVY)
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.03
        sf(p.add_run(str(val)),size,True,"FFFFFF")
    for ri,row in enumerate(rows):
        cells=t.add_row().cells; trp=t.rows[-1]._tr.get_or_add_trPr(); trp.append(OxmlElement("w:cantSplit"))
        for j,val in enumerate(row):
            c=cells[j]; c.width=Inches(widths[j]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(c)
            if ri%2: shade(c,PALE)
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.02
            if j in center_cols: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            sf(p.add_run(str(val)),size,False,INK)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t
def img(path,width):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(4); p.add_run().add_picture(str(path),width=Inches(width))
def br(): pending_break[0]=True
header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT; header.paragraph_format.space_after=Pt(0)
sf(header.add_run("0038｜I08-E02｜完整示範資料，不代表協會實績"),8.5,False,MUTED)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
sf(footer.add_run("候選文件・SYNTHETIC　｜　第 "),8.5,False,MUTED)
field=OxmlElement("w:fldSimple"); field.set(qn("w:instr"),"PAGE"); footer._p.append(field); sf(footer.add_run(" 頁"),8.5,False,MUTED)

# 1. Reader-facing opening.
title=doc.add_paragraph(style="Title"); title.paragraph_format.space_after=Pt(3); sf(title.add_run("訓練方案系統設計改善執行與成效示範紀錄"),24,True,"000000")
P("文件 0038　｜　I08-E02　｜　官方具名方框序號 039",size=12,color=NAVY,after=8)
P("SAMPLE / SYNTHETIC COMPLETE EXAMPLE｜非原始點協會實際辦訓、核准、簽署或成效紀錄",size=10.2,color="8A4B16",after=9)
H("這份紀錄供誰使用",2)
P("供小型訓練機構承辦人檢討一項課程設計、記錄實際修改，並用同一量尺核對執行結果。行政人員可按本件順序完成找出缺口、比較方案、修訂課程、記錄樣本、重算結果及決定下一步；評核委員可查看官方需求、逐筆資料和判讀限制。")
H("本例的示範結論",2)
P("完全虛構的兩組各12人資料中，講授為主組後測達標3/12（25.0%）；加入情境演練與回饋組達標11/12（91.7%）。修訂組平均分增加2.92/8分，對照組增加1.08/8分，兩組增益差1.83分。這些數字只示範記錄與計算，既非協會實績，也不能證明因果效果。")
P("唯一命題：訓練設計改善是否有事前假設、可比較選項、具體改動、可重算執行結果及有限度的後續決定。",lead="唯一命題：",size=10.5)
P("案例：SYN-I08E02-2026-01；假設課程「社區關懷志工服務溝通與轉介界線」。",lead="案例：",size=10.5)
P("資料狀態：全篇日期、角色、課程、樣本、量尺、分數、決定均為合成示範；沒有任何真人或協會原始紀錄。",lead="資料狀態：",size=10.5)
P("來源定位：官方訓練機構版佐證表 I08 p.2；p.3為相鄰項續頁。本件對應 I08-E02。",lead="來源定位：",size=10.5)
H("使用方式",2)
P("實際使用時，將 SAMPLE 欄位換成經授權可查的課程版本、執行與評量資料，再保存附件位置、去識別方式及覆核軌跡。沒有真實資料時，保留本件為示範稿，不得勾選成協會已執行或正式通過。")

# 2. Requirement and genre boundary.
br(); H("1　官方要求與本件工作邊界",1)
P("官方文字：「針對訓練方案的系統設計能適時檢討及改善，執行後有良好效果。」本件以一項課程設計的缺口、替代方案、設計變更、執行記錄和結果判讀回應此命題；不以會議形式取代實際改善證據。")
table(["對照項","本件如何處理","查核位置"],[
["官方定位","I08-E02；官方方框序號039；資料表p.2，續頁p.3。","本頁及來源索引"],
["0038文類","執行與樣本紀錄；呈現操作流向、逐筆示範、計算和後續處置。","§2、§5–§9"],
["0037負向鄰項","I08-E01重點是跨期設計變化；本件記錄一項改善的執行與效果，不代替跨期檢討。","差異說明"],
["0039負向鄰項","I08-E03重點是規劃會議議程、意見、決議及追蹤；本件沒有假造會議或核准。","差異說明"],
["真實界線","協會原始執行未驗證；案例獨立，不宣稱與0037或0039共用事實。","替換清冊與來源圖"]
],[1.05,4.15,1.7],9.4)
H("合成案例的前提",2)
P("本例假設承辦人發現講授與知識測驗不足以觀察志工的溝通、隱私及轉介界線，因此提出加入短情境演練的假設。這不是對原始點協會現行課程、服務對象、工作角色或既往成效的陳述。")
img(flow,6.8)

# 3. Problem and alternatives.
br(); H("2　設計缺口、觀察與方案比較",1)
H("示範問題與可檢驗假設",2)
P("合成基準設計假設為120分鐘講授為主，評量集中於知識題。它能檢查記憶，卻未直接觀察學習者是否說明服務目的、取得同意、保護資訊、守住非醫療界線並作適當轉介。此處只是示範設計假設，不是查得的協會課程缺陷。")
P("假設：總時數、課程目標、量尺和評量時點相同時，以35分鐘情境演練與回饋取代部分講授，會提高學習者在等難度情境中展現目標行為的比例。")
H("比較方案與示範評分",2)
P("各項1–5分，5分代表較符合；加權分＝Σ（分數×權重），總分上限5.00。權重先於評分設定，合計100%。")
table(["準則／權重","A 增加講義","B 情境演練＋回饋","C 外聘全日實作"],[
["可觀察行為證據 40%","2×.40=.80","5×.40=2.00","5×.40=2.00"],
["小型課程可執行 25%","5×.25=1.25","4×.25=1.00","2×.25=.50"],
["隱私／非醫療界線 20%","3×.20=.60","5×.20=1.00","5×.20=1.00"],
["結果可追溯 15%","2×.15=.30","5×.15=.75","4×.15=.60"],
["加權總分","2.95/5.00","4.75/5.00","4.10/5.00"]
],[1.5,1.55,2.1,1.75],9.2,center_cols=(1,2,3))
P("示範取捨：B得分最高且不增加總時數，以可觀察行為與回饋替代部分講授。A成本較低，但仍不能直接檢查行為；C觀察力較高，卻需更多經費及外部師資。上述分數與選擇是假設值，沒有協會正式決議。")
H("從缺口到可驗證結果",2)
table(["設計缺口","具體修改","可查結果","事前判定規則"],[
["只用知識題","加入情境演練，評D1、D2。","前後測向量及評分表。","總分≥6/8。"],
["服務界線混淆","加入界線示範、兩情境練習和回饋。","D3、D4分數及異常覆核。","D3、D4均>0。"],
["改善無固定口徑","兩組同時數、同量尺、配對前後測。","逐筆資料、有效數和增益。","B達標率≥80%，關鍵失誤=0才建議續試。"]
],[1.35,2.0,1.75,1.9],9.1)
P("本例沒有真實需求調查、職能盤點或先前實績支持此假設。真實執行時，須先附需求來源及版本差異再確認是否試作。",size=9.6,color=MUTED)

# 4. Detailed design.
br(); H("3　修訂後課程設計與執行安排",1)
P("示範課程：「社區關懷志工服務溝通與轉介界線」。對象是假設的社區關懷志工學習者；不得推定協會已開辦。兩組總時數同為120分鐘。")
H("120分鐘教學安排",2)
table(["時間","活動與方法","引導者操作","學習者產出／檢核"],[
["00–10","情境前測","說明匿名代碼及評分；發等難度情境卡。","完成四項行為前測。"],
["10–20","目的、同意與拒絕權","示範先說明目的並詢問是否願意繼續。","重述可停止及改約界線。"],
["20–40","開放提問與中性摘要","示範開放問句，區分聆聽、推測和承諾。","完成摘要並確認理解。"],
["40–55","隱私及非醫療界線","比較必要資訊與多餘細節；示範停止並轉介。","指出何時不記、不建議、不代替專業判斷。"],
["55–80","兩人一組演練","輪替服務者與觀察者，每人至少完整演練一次。","依D1–D4記錄可見行為。"],
["80–105","個別後測情境","用另一等難度情境評分，不提示答案。","四項分數及簡短觀察事實。"],
["105–115","回饋與重試","說明具體行為後給一次練習；正式分數不覆寫。","完成回饋練習，保留首測值。"],
["115–120","轉用計畫","提示選一項下次可採行的服務行為。","記錄一項可觀察行動。"]
],[.68,1.35,2.65,2.22],9.15)
H("目標與證據對應",2)
table(["可觀察目標","教學活動","評量證據"],[
["說明目的並取得同意","目的／拒絕權示範和演練","D1：先說明並取得同意"],
["以開放問題理解需求","問句、摘要及理解確認","D2：開放提問與中性摘要"],
["保護隱私並守住界線","必要資訊及非醫療界線情境","D3：隱私與角色界線"],
["辨識超出角色問題並轉介","轉介流程示範及後測","D4：合宜路徑且不作保證"]
],[1.65,2.7,2.55],9.3)

# 5. Instrument and controls.
br(); H("4　評量工具、量尺與資料控制",1)
P("前後測用難度相當但不同的情境，避免直接重複答案；兩組使用同一量尺。每項0–2分，總分0–8；正式判定同時符合總分和關鍵項條件。")
table(["項目","0分","1分","2分","關鍵"],[
["D1 目的與同意","直接追問或假定同意。","說明目的或詢問意願，未同時做到。","先說明目的，再詢問是否願意並接受拒絕。","否"],
["D2 開放提問與摘要","誘導或封閉主導，未確認理解。","做到開放提問或摘要其中一項。","開放提問、中性摘要並確認理解。","否"],
["D3 隱私與服務界線","不必要蒐集／揭露，或提供超出角色判斷。","保護部分資訊，但界線不完整。","只記必要資訊，明確不診斷、不承諾醫療效果。","是"],
["D4 適當轉介與後續","未轉介，或作未確認保證。","指出資源但未說明下一步。","說明適當聯繫及可確認後續，不保證結果。","是"]
],[1.12,1.48,1.48,2.0,.82],9.0)
P("達標規則：總分至少6/8，且D3>0、D4>0。未達標保留原始分數，另記補強與重測；不得刪列或改寫為通過。",lead="達標規則：",size=10)
H("等難度情境與執行資料",2)
table(["情境","前測示例","後測示例","同等要求"],[
["目的與隱私","活動後談到家庭近況；志工不確定能否記錄及轉述。","電話關懷中對方談及不願公開的生活資訊並詢問保密。","說明目的、詢問同意、只記必要資訊、說清保密界線。"],
["角色與轉介","對方問志工能否判斷身體不適原因。","對方請志工保證某項服務一定有效。","不診斷、不保證；說明適當專業聯繫路徑。"]
],[.8,2.05,2.05,1.8],9.0)
H("缺漏、重複與資料保存",2)
table(["檢查點","本示範記錄","真實執行控制"],[
["同意與代碼","24列假設代碼；無姓名或聯絡方式。","未同意者不納入評量，只列合計與排除原因。"],
["配對與重複","前後測配對24；重複代碼0。","以原件和日期核對；更正留軌跡，無法判定者列理由。"],
["缺測與評分","缺測0；不以未達標作排除理由。","漏評不得補分；依事前規則列缺漏並報有效分母。"],
["權限與保存","實際位置及存取權限UNKNOWN。","依核實保存期、權限、刪除和備份規則辦理。"]
],[1.0,2.55,3.15],9.05)

# 6. Sample frame and flow.
br(); H("5　樣本批次、納入規則與資料流向",1)
P("本批次所有資料都是SAMPLE/SYNTHETIC。假設樣本框為24名符合條件的學習者，A、B兩組各12名；以下日期、渠道、同意狀態及排除數均非實際協會紀錄。")
table(["欄位","全批次示範值","真實替換與核對來源"],[
["母體與樣本框","假設24名參加者；A與B各12名。","正式招生範圍、課程對象及去識別樣本框。"],
["納入與排除規則","預先假設：同意參與、完成前後兩次評量、D1–D4均可評才納入有效配對。未同意、缺任一測或缺評項須記理由。","納入規則版本、同意紀錄、逐筆缺漏及排除清單。"],
["日期與執行渠道","A組2026-09-08；B組2026-09-15。模擬紙本情境卡及匿名評分表。","實際日期、課程簽到、使用工具及原始評量表。"],
["代碼與同意","A01–A12、B01–B12為合成代碼；假設同意24/24；不含姓名、電話或簽名。","合法同意及代碼對照保管方式；公開版只留去識別資料。"],
["缺漏、重複與更正","有效24；缺漏0；重複0；排除0。B08加總覆核見§7；沒有其他更正。","原始列、重複檢查、差異版次、更正人及日期。"],
["保存與存取","本候選無個資；協會實際保存位置、權限和期限均UNKNOWN。","協會核實的權限表、保存年限、刪除和備份紀錄。"]
],[1.18,3.55,2.17],9.5)
H("樣本流向與分母",2)
table(["階段","筆數","計算／判定","說明"],[
["接觸合格樣本框","24","—","假設樣本框，不是協會名冊。"],
["發出邀請","24","24/24","全數受邀，尚非回收數。"],
["收到成對前後測","24","24/24","A組12、B組12；以代碼配對。"],
["有效成對資料","24","有效率=24/24=100%","有效率=有效數÷回收成對數；分母是回收數。"],
["排除／缺漏","0","0筆","沒有因未達標而排除的列。"]
],[1.32,.58,1.62,3.38],9.5,center_cols=(1,))
P("若成對回收數為0，有效率應寫N/A並註明無回收分母，不可填0%。24/24有效率只表示資料可計算，不等於達標率或學習效果。",size=9.8,color=MUTED)
P("共同欄位定義：每列保留case_id、合成日期、組別、渠道、同意／納入判定、D1–D4前後向量、缺漏／重複狀態及更正理由。真實執行時，原始附件路徑及覆核角色須可追查。",size=9.8)

# 7. A cohort.
br(); H("6　完整示範批次：比較組 A",1)
P("A組示範日期為2026-09-08，渠道為模擬紙本情境評量；各列同意／納入與缺漏、重複、更正狀態見附表A。講授為主設計並非協會既有課程。向量順序D1/D2/D3/D4，合計與差值可直接重算。")
H("A組逐筆結果｜講授為主",2)
rows=[]
for pid,pre,post in A: rows.append([pid,vec(pre),sm(pre),vec(post),sm(post),f"{sm(post)-sm(pre):+d}","達標" if ok(post) else "未達"])
table(["代碼","前測 D1/D2/D3/D4","前測","後測 D1/D2/D3/D4","後測","差值","判定"],rows,[.58,1.38,.5,1.38,.5,.55,.85],9.5,center_cols=(0,2,4,5,6))
P("A組小計：前測總分44、後測總分57；平均由44/12=3.67升至57/12=4.75，平均增益13/12=1.08分；後測3/12達標，D3或D4為0者2/12。",lead="A組小計：",size=9.6)
# 7. B cohort and anomaly.
br(); H("7　完整示範批次：修訂組 B",1)
P("B組是假設的同為120分鐘、加入情境演練、行為量尺與回饋的修訂設計。總時數、前後測計分及達標門檻與A組一致；兩組並非隨機分派。各列渠道、同意／納入與資料狀態見附表A。")
H("B組逐筆結果｜情境練習與回饋",2)
rows=[]
for pid,pre,post in B: rows.append([pid,vec(pre),sm(pre),vec(post),sm(post),f"{sm(post)-sm(pre):+d}","達標" if ok(post) else "未達"])
table(["代碼","前測 D1/D2/D3/D4","前測","後測 D1/D2/D3/D4","後測","差值","判定"],rows,[.58,1.38,.5,1.38,.5,.55,.85],8.9,center_cols=(0,2,4,5,6))
P("B組小計：前測總分44、後測總分79；平均由44/12=3.67升至79/12=6.58，平均增益35/12=2.92分；後測11/12達標，D3或D4為0者0/12。",lead="B組小計：",size=9.6)
H("異常覆核 SYN-ANOMALY-001",2)
P("B08後測向量1/1/1/2，加總為1+1+1+2=5分，低於6分門檻。覆核後與表列總分一致；D3=1、D4=2，關鍵界線通過。保留「未達標」，不刪除此筆、不補分，也不以回饋後重試分數覆寫正式後測。")
table(["異常","核對","處置","資料影響"],[
["B08總分不足","原始向量重算5，與列值一致。","保留有效列並標未達；補練後須另建重測紀錄。","分母仍24；B組達標分子仍11。"]
],[1.12,2.05,2.45,1.55],9.0)

# 8. Calculations and limits.
br(); H("8　計算、結果與可歸因界線",1)
table(["指標","講授為主 A","修訂設計 B","重算方式"],[
["有效成對樣本","12/12","12/12","各組有效列數。"],
["平均前測（0–8）","44/12=3.67","44/12=3.67","前測各列向量加總÷12。"],
["平均後測（0–8）","57/12=4.75","79/12=6.58","後測各列向量加總÷12。"],
["平均增益","(57−44)/12=1.08","(79−44)/12=2.92","平均後測−平均前測，單位分。"],
["後測達標率","3/12=25.0%","11/12=91.7%","總分≥6且D3、D4均>0。"],
["關鍵項為0","2/12=16.7%","0/12=0%","D3或D4任一為0；分母各組12。"]
],[1.25,1.4,1.4,2.85],9.1)
img(chart,6.7)
H("兩組差異怎麼讀",2)
P("平均增益差=(79−44)/12−(57−44)/12=22/12=1.83分（0–8量尺）。後測達標率差=11/12−3/12=8/12=66.7個百分點。有效率24/24=100%是資料可用比例，不是學習成效；達標率按各組有效樣本計算。")
P("示範判讀：B組達標率高於事前設定的80%，且關鍵項為0者為0；本案例內可記為「建議下一輪小規模續試，並對B08補強」。這不是協會已核准或已採行。")
H("限制與退件條件",2)
P("兩組都是合成小樣本，未隨機分派，也沒有真實出席、需求、評分者一致性或外部來源。基線平均相同只是示範資料設計，不能推論改善造成差異、母群效果或協會實際成效。")
P("分母不明、量尺不一致、逐筆加總錯誤、以不及格為由刪列、關鍵項為0仍算通過、重複或缺漏未揭露、合成事實未登錄、虛構核准，任一出現即退回受影響內容。",size=10)

# 9. Next steps and source index.
br(); H("9　後續處理、真實替換與查核索引",1)
H("本例示範的後續工作",2)
table(["時點","工作","責任角色（示範）","完成判定"],[
["執行前","取得正式課程版次，確認對象、評量用途及量尺版本。","課程承辦角色","有版本、來源和授權依據。"],
["每次執行","記錄邀請、同意、配對、缺漏、排除及異常。","執行／資料角色","分母可重算，逐筆值可追源。"],
["課後覆核","非原評分者抽核原始表與加總；更正留初值及理由。","覆核角色","範圍、差異、處置留痕。"],
["下一輪前","比較真實基線與修訂結果；說明組成和評分差異。","承辦角色","限制、失敗個案及再測時點明確。"]
],[.82,2.65,1.55,1.88],9.0)
H("改成真實紀錄前的替換清單",2)
for s in [
"核對正式機構名稱、課程代碼、計畫版次及具權限的承辦角色。",
"取得原版與修訂版課程計畫、需求／職能來源及變更記錄。",
"以真實日期、納入規則、邀請／同意、前後測原表、出席及排除紀錄替換樣本。",
"核驗評分者權限、量尺版次、校準、原始值、更正軌跡和資料保管權限。",
"重算有效率、平均分、增益、達標率、關鍵項與限制；缺正式核准或簽章時保留UNKNOWN。",
"逐筆更新 replacement register，記錄原件位置、去識別方式和覆核人。"
]:
    p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(3); sf(p.add_run(s),10.2,False,INK)
H("官方與專案來源索引",2)
table(["來源","定位與用途"],[
["官方評核依據","訓練機構版佐證表 I08 p.2、續頁 p.3；PDF SHA記於SOURCE_MAP.md。"],
["正式項次與藍圖","0038→I08-E02→官方序號039；對照目前Blueprint row、142序列及01/02/03規格。"],
["相鄰要求","0037/I08-E01為多期設計檢討；0039/I08-E03為規劃會議。本案例不宣稱同案。"],
["合成事實清冊","SYNTHETIC_REPLACEMENT_REGISTER.csv含case_id、doc_id、定位、REAL原件及核實角色。"],
 ["官方要求逐項對照","附表B連結官方命題、紀錄欄位、來源、判定及覆核位置。"],
["版本收據","SOURCE_MAP.md、BLUEPRINT_CURRENT_0038.jsonl、PROGRESS.json及manifest.json。"]
],[1.4,5.5],9.1)
H("目前狀態",2)
P("本件為CANDIDATE示範文件；逐項驗收狀態以隨附收據為準，協會正式接受未授與。此件只回應0038/I08-E02，不取代真實課程原件、0037跨期檢討或0039會議紀錄。",size=10)
H("一次執行的控制順序（角色均為合成示範）",2)
table(["階段","負責角色（均為示範）","操作與控制點","輸出"],[
["定位","承辦角色","記錄官方要求、課程版次與缺口；推測不可寫成事實。","問題與來源"],
["比較","設計覆核角色","用事前權重評分三方案並保留取捨。","方案比較表"],
["修訂","課程引導角色","將可觀察目標對到活動、評分行為和通過規則。","變更表"],
["執行","引導者／觀察者","同時數、同量尺完成兩組示範；記錄樣本流向。","逐筆樣本"],
["覆核","資料覆核角色","核對加總、配對、缺漏、排除及關鍵界線。","覆核軌跡"],
["判讀","承辦角色","計算比例與增益，列限制及下一步。","成效判讀"]
],[.7,1.35,3.15,1.7],9.5)
P("正式使用時須依協會實際章程、分工及資料權限確認角色；本表不構成現職指派或授權證據。",size=9.6,color=MUTED)

# 10. Exact row-level synthetic execution trace and requirement-to-record index.
br(); H("附表A　24筆合成樣本逐列追溯",1)
P("本表逐列連結§6、§7的A01–A12、B01–B12分數紀錄。渠道、同意及納入均為合成假設，並非真人同意或協會實際執行證明；各列均符合本例預設完整配對條件。")
trace_rows=[]
for group,rows in (("A",A),("B",B)):
    for pid,pre,post in rows:
        trace_rows.append([pid,"模擬紙本情境卡／匿名評分表","同意=是（合成假設）；納入=是，原因：完成前後測且D1–D4均有評分","缺漏=0；重複=0；更正=0（未更動原值）"])
table(["case_id","渠道（SYNTHETIC）","同意／納入判定（SYNTHETIC）","缺漏／重複／更正（SYNTHETIC）"],trace_rows,[.72,1.55,2.45,1.65],9.5)
P("狀態定義：缺漏=未完成前後測或量尺項；重複=重複代碼列；更正=對原始值有更動。真實執行不得把未知狀態填0，須保留原值、原因、核對人及日期。每列REAL替換所需原件為合法同意／納入依據、課程簽到、匿名化前後測原表及更正／覆核軌跡；缺一項即維持UNKNOWN。",size=9.8)

br(); H("附表B　官方要求、紀錄欄位與覆核位置對照",1)
P("本表把官方命題拆成可查的紀錄鏈。SAMPLE資料只示範欄位與計算；REAL欄列出升格前必須取得的原件，不代表原始點協會已持有或核准。")
table(["官方命題／查核問題","本件紀錄欄位與位置","SAMPLE來源／REAL替換原件","判定與限制","覆核位置"],[
["訓練方案系統設計是否適時檢討並改善？","設計缺口、可檢驗假設、加權方案、修訂活動、目標與評量對應。","SAMPLE：§2–§4；REAL：原版／修訂課程計畫、需求或職能來源、版本差異及授權依據。","示範方案B加權4.75/5.00；此為假設選擇，非協會決議。","§2方案比較表；§3時程與目標表；§9替換清單。"],
["執行後是否有良好效果？","樣本框、前後測量尺、24筆逐列結果、有效分母、達標與增益。","SAMPLE：§5–§8及附表A；REAL：出席、合法同意、匿名前後測原表、評分者權限與覆核紀錄。","A 3/12、B 11/12達標；合成小樣本、非隨機，不能推論因果或協會成效。","§5納入／樣本流向；§6–§7逐筆表；§8計算與限制。"],
["每列資料能否追到執行狀態？","渠道、同意、納入理由、缺漏、重複及更正狀態。","SAMPLE：附表A逐列欄位與替換清冊；REAL：課程簽到、同意依據、逐筆原表、差異更正軌跡。","本例24列皆假設同意、納入，缺漏／重複／更正皆0；實際未知不得填0。","附表A；replacement register各SYN-A/B-ROW列；§5處理規則。"],
["判定是否有來源且可覆核？","來源、公式、判定門檻、異常及限制。","SAMPLE：逐筆表、量尺、替換清冊；REAL：正式計畫、原始評量、核准量尺及資料權限記錄。","公式可重算；正式來源、核准與協會實績尚未驗證，維持SAMPLE。","§4量尺；§8彙總；§9來源索引；附表B。"]
],[1.42,1.65,1.65,1.42,1.15],9.5)

doc.core_properties.title="訓練方案系統設計改善執行與成效示範紀錄"
doc.core_properties.subject="TTQS ONE 0038 / I08-E02 / 官方序號039 / SAMPLE-SYNTHETIC"
doc.core_properties.author="TTQS ONE candidate document production"
doc.save(DOCX)
docsha=hashlib.sha256(DOCX.read_bytes()).hexdigest()
progress={
"project":"TTQS_ONE","document_id":"0038","requirement_id":"I08-E02","official_checkbox_sequence":"039","candidate":"0038_R02",
"candidate_path":"candidate/0038_R02.docx","candidate_sha256":docsha,
"status":"BUILT_CANDIDATE; independent_semantic_review=PENDING; formal_acceptance=NOT_GRANTED",
"source_spec_branch":"work/ttqs-one-document-quality-contract-20261016","source_spec_head":SPEC_HEAD,"source_spec_blob_ids":SPEC_BLOBS,
"official_pdf_sha256":OFFICIAL_SHA,"blueprint_file_sha256":BP_FILE_SHA,"blueprint_row_sha256":BP_ROW_SHA,
"case_id":"SYN-I08E02-2026-01","sample_boundary":"All dates, roles, course facts, options, participant codes, scores, calculations and decisions are SYNTHETIC.",
"sample_calculation":{"group_a":{"n":12,"pre_sum":ap,"post_sum":aq,"gain":aq-ap,"pass":ac,"critical_zero":az},
"group_b":{"n":12,"pre_sum":bp,"post_sum":bq,"gain":bq-bp,"pass":bc,"critical_zero":bz},
"valid_pairs":24,"returned_pairs":24,"valid_rate":"24/24=100%","gain_difference":(bq-bp-aq+ap)/12},
"review":{"semantic":"PENDING","technical":"BUILT_NOT_VERIFIED","visual":"PENDING_RENDER","exact_sha_bound_review":False},
"supersedes":{"candidate":"0038_R01","sha256":"75065568b801506d736297441ee30abd3eaccc4fa3748ef6da2a11c67715088a","independent_review":"PARTIAL","resolved_findings":["row-level execution trace fields","official-requirement-to-record/source/decision/review crosswalk"]},
"remote":{"status":"NOT_SYNCED","github_path":None,"verified_readback":False},
"created_local":"2026-10-17"}
(ROOT/"PROGRESS.json").write_text(json.dumps(progress,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
(ROOT/"README.md").write_text("""# 0038 R02 candidate

此候選對應 DOCX 0038 / I08-E02 / 官方方框序號039。文類為系統設計改善的執行與樣本紀錄，不是0037跨期檢討，也不是0039會議紀錄。R02依R01 exact-SHA獨立審查的兩項PARTIAL內容缺口補入逐列追溯附表及官方要求對照表。

所有日期、角色、情境、學習者代碼、分數、計算和決定均為SAMPLE/SYNTHETIC，不證明原始點協會曾開課、核准或取得成效。狀態以PROGRESS.json為準；沒有綁定精確SHA的獨立驗收前，不得標為品質接受。
""",encoding="utf-8")
print(json.dumps({"docx":str(DOCX),"sha256":docsha,"register_rows":len(reg),"groupA":{"pre":ap,"post":aq,"pass":ac},"groupB":{"pre":bp,"post":bq,"pass":bc}},ensure_ascii=False))
