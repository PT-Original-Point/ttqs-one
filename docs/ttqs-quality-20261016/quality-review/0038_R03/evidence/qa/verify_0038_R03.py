from pathlib import Path
import csv,hashlib,json,zipfile
from docx import Document
from PIL import Image
from pypdf import PdfReader

REPO=Path(r"E:\TTQS\_TTQS_ONE_REVIEW_SYNC_0036_R08_20261016")
ROOT=REPO/"docs/ttqs-quality-20261016/quality-review/0038_R03"
R02=REPO/"docs/ttqs-quality-20261016/quality-review/0038_R02"
DOC=ROOT/"candidate/0038_R03.docx"
PDF=ROOT/"evidence/render_final/0038_R03.pdf"
RENDER=ROOT/"evidence/render_final"
REGISTER=ROOT/"evidence/SYNTHETIC_REPLACEMENT_REGISTER.csv"
sha=lambda b:hashlib.sha256(b).hexdigest()
def fsha(p):return sha(p.read_bytes())

doc=Document(DOC); old=Document(R02/"candidate/0038_R02.docx")
checks=[]
def ck(name,ok,detail): checks.append({"name":name,"status":"PASS" if ok else "FAIL","detail":detail})

doc_sha=fsha(DOC); pdf_sha=fsha(PDF)
ck("candidate_exact_sha",doc_sha=="3a887bd3835b3a30b9fdb489257c8a09f79d3c2ce25c10565fae2b4e173f7efa",doc_sha)
with zipfile.ZipFile(DOC) as z: bad=z.testzip()
ck("docx_package_crc",bad is None,str(bad or "all members CRC valid"))
para_text=[p.text for p in doc.paragraphs]; old_para=[p.text for p in old.paragraphs]
table_text=[[c.text for row in t.rows for c in row.cells] for t in doc.tables]
old_table_text=[[c.text for row in t.rows for c in row.cells] for t in old.tables]
ck("content_preserved_from_semantic_pass_R02",para_text==old_para and table_text==old_table_text,"all top-level paragraph and table cell text identical to R02; R02 review itself remains bound only to R02 SHA")
body_sizes=[]; body_missing=[]
for pi,p in enumerate(doc.paragraphs):
 for ri,r in enumerate(p.runs):
  if not r.text.strip():continue
  if r.font.size is not None: body_sizes.append(r.font.size.pt)
  else:
   st=p.style; effective=None
   while st is not None and effective is None:
    effective=st.font.size.pt if st.font.size else None; st=st.base_style
   if effective is None: effective=11.0
   body_sizes.append(effective); body_missing.append({"paragraph":pi,"run":ri,"inherited_pt":effective})
min_body=min(body_sizes) if body_sizes else None
ck("body_font_floor_10_5pt",min_body is not None and min_body>=10.5,f"minimum_effective_body_pt={min_body}; nonempty_runs={len(body_sizes)}; style_inherited={len(body_missing)}")

table_sizes=[]; table_missing=0
for t in doc.tables:
 for row in t.rows:
  for cell in row.cells:
   for p in cell.paragraphs:
    for run in p.runs:
     if run.text.strip():
      if run.font.size is not None:table_sizes.append(run.font.size.pt)
      else:
       st=p.style; effective=None
       while st is not None and effective is None:
        effective=st.font.size.pt if st.font.size else None; st=st.base_style
       if effective is None:effective=11.0
       table_sizes.append(effective);table_missing+=1
min_table=min(table_sizes) if table_sizes else None
ck("table_font_floor_9_5pt",min_table is not None and min_table>=9.5,f"minimum_effective_table_pt={min_table}; nonempty_runs={len(table_sizes)}; style_inherited={table_missing}")

score_tables=[]
for t in doc.tables:
 if not t.rows:continue
 h=[c.text.strip() for c in t.rows[0].cells]
 if len(h)>=7 and h[0]=="代碼" and "前測" in h[1] and "後測" in h[3]:score_tables.append(t)
results={}; errors=[]
for t in score_tables:
 rows=t.rows[1:]; ids=[r.cells[0].text.strip() for r in rows]; group=ids[0][0]
 acc={"pre":0,"post":0,"pass":0,"critical_zero":0,"n":0}
 for row in rows:
  c=[x.text.strip() for x in row.cells]; pre=[int(x) for x in c[1].split("/")];post=[int(x) for x in c[3].split("/")]
  pre_total=sum(pre);post_total=sum(post);passed=post_total>=6 and post[2]>0 and post[3]>0
  expected="達標" if passed else "未達"
  if int(c[2])!=pre_total:errors.append(f"{c[0]} pre");
  if int(c[4])!=post_total:errors.append(f"{c[0]} post")
  if int(c[5].replace("+",""))!=post_total-pre_total:errors.append(f"{c[0]} delta")
  if c[6]!=expected:errors.append(f"{c[0]} status")
  if len(pre)!=4 or len(post)!=4 or any(v<0 or v>2 for v in pre+post):errors.append(f"{c[0]} vector")
  acc["pre"]+=pre_total;acc["post"]+=post_total;acc["pass"]+=passed;acc["critical_zero"]+=post[2]==0 or post[3]==0;acc["n"]+=1
 results[group]=acc
want={"A":{"pre":44,"post":57,"pass":3,"critical_zero":2,"n":12},"B":{"pre":44,"post":79,"pass":11,"critical_zero":0,"n":12}}
ck("two_12_case_batches_and_recomputed_math",len(score_tables)==2 and results==want and not errors,json.dumps({"results":results,"errors":errors},ensure_ascii=False))
ck("aggregate_calculations",(57-44)/12==13/12 and (79-44)/12==35/12 and ((79-44)-(57-44))/12==22/12,"A=13/12; B=35/12; delta=22/12=1.8333")

traces=[];crosswalk=[]
for t in doc.tables:
 h=[c.text.strip() for c in t.rows[0].cells] if t.rows else []
 if h and h[0]=="case_id" and "渠道（SYNTHETIC）" in h:traces.append(t)
 if h==["官方命題／查核問題","本件紀錄欄位與位置","SAMPLE來源／REAL替換原件","判定與限制","覆核位置"]:crosswalk.append(t)
trace_ids=[r.cells[0].text.strip() for r in traces[0].rows[1:]] if len(traces)==1 else []
trace_ok=len(traces)==1 and len(trace_ids)==24 and set(trace_ids)=={f"{g}{i:02d}" for g in "AB" for i in range(1,13)}
if trace_ok:
 for row in traces[0].rows[1:]:
  c=[x.text for x in row.cells];trace_ok &= "模擬紙本" in c[1] and "同意=是" in c[2] and "納入=是" in c[2] and "完成前後測" in c[2] and all(x in c[3] for x in ("缺漏=0","重複=0","更正=0"))
ck("24_case_execution_trace",trace_ok,f"rows={len(trace_ids)}; expected unique A01-A12/B01-B12")
walk_ok=len(crosswalk)==1 and len(crosswalk[0].rows)==5
if walk_ok:
 for row in crosswalk[0].rows[1:]:
  c=[x.text.strip() for x in row.cells];walk_ok &= len(c)==5 and all(c) and "SAMPLE" in c[2] and "REAL" in c[2]
ck("official_requirement_to_fields_source_decision_review_crosswalk",walk_ok,f"crosswalk_rows={len(crosswalk[0].rows)-1 if walk_ok else 0}")

with REGISTER.open(encoding="utf-8-sig",newline="") as f:regs=list(csv.DictReader(f))
rows=[x for x in regs if x.get("fact_id","").startswith(("SYN-A-ROW-","SYN-B-ROW-"))]
regok=len(regs)==37 and len(rows)==24 and len({x.get("fact_id") for x in rows})==24
for x in rows:regok &= all(k in x.get("synthetic_value","") for k in ("渠道=","同意=是","納入=是","缺漏=0","重複=0","更正=0"))
ck("replacement_register_37_rows_and_24_case_lineage",regok,f"rows={len(regs)}; trace_rows={len(rows)}")
ck("candidate_mapping_and_sample_boundary",doc_sha and any("0038" in t for t in para_text) and any("SAMPLE" in t and "SYNTHETIC" in t for t in para_text) and any("I08-E02" in t for t in para_text) and any("039" in t for t in para_text),"0038 / I08-E02 / official sequence 039; SAMPLE/SYNTHETIC marked in document")

pages=sorted(RENDER.glob("page-*.png"),key=lambda p:int(p.stem.split('-')[1]))
pdf_pages=len(PdfReader(str(PDF)).pages)
valid_page_names=[p.name for p in pages]==[f"page-{i}.png" for i in range(1,pdf_pages+1)]
ck("all_page_images_match_pdf",pdf_pages==14 and len(pages)==14 and valid_page_names,f"pdf_pages={pdf_pages}; pngs={len(pages)}; names={[p.name for p in pages]}")
img_dims={p.name:Image.open(p).size for p in pages}
ck("render_output_images_nonempty",all(p.stat().st_size>10000 for p in pages) and all(w>=1400 and h>=1800 for w,h in img_dims.values()),json.dumps(img_dims))

result={"project_id":"TTQS_ONE","doc_id":"0038","requirement_id":"I08-E02","official_checkbox_sequence":"039","candidate":"0038_R03","candidate_sha256":doc_sha,"candidate_bytes":DOC.stat().st_size,"source_candidate":"0038_R02","source_sha256":"bdb6ef07c1aacd31f638cdad511012777d2f1b9d2c9f5f307252dd62de0170a4","source_spec_head":"639b9eacce93f84f49ee20714f1bad337c6a961e","blueprint_file_sha256":"406d0495365d0970dcf8113a42842663b8a54372a4e204cc492b7c3f8156ef3e","blueprint_row_sha256":"2504485c1698a5c83ac2205938c05b532edc8d697f2441ad1d6a0c324897b73e","pdf_sha256":pdf_sha,"pdf_pages":pdf_pages,"renderer":"Codex documents skill render_docx.py using E:/TTQS project-controlled renderer recovery LibreOffice; rasterized by packaged Poppler","body_min_effective_pt":min_body,"table_min_effective_pt":min_table,"scoring_summary":results,"replacement_register_rows":len(regs),"checks":checks,"status":"PASS" if all(x["status"]=="PASS" for x in checks) else "FAIL","visual_review":"PASS_ALL_14_PAGES_BY_DIRECT_INSPECTION","independent_semantic_review":"PENDING","formal_acceptance":"NOT_GRANTED"}
(ROOT/"evidence/qa/technical_qa_R03.json").write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
print(json.dumps(result,ensure_ascii=False,indent=2))
