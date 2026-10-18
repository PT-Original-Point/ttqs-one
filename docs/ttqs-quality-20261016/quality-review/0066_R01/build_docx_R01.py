from pathlib import Path
import hashlib,re,json
from docx import Document
from docx.shared import Pt,RGBColor,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START,WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(r'E:\TTQS\_TTQS_ONE_REVIEW_SYNC_0036_R08_20261016\docs\ttqs-quality-20261016\quality-review\0066_R01')
SRC=ROOT/'inputs'/'0066_BODY_R01_HUMAN_READY.md';OUT=ROOT/'candidate'/'0066_R01.docx';FONT='Microsoft JhengHei';NAVY='164E63';TEAL='0F766E';ALT='F3F7FA';INK='172B4D';MUTED='52606D'
def shade(c,fill):
 p=c._tc.get_or_add_tcPr();x=p.find(qn('w:shd'))
 if x is None:x=OxmlElement('w:shd');p.append(x)
 x.set(qn('w:fill'),fill)
def margins(c,top=65,start=75,bottom=65,end=75):
 p=c._tc.get_or_add_tcPr();x=p.first_child_found_in('w:tcMar')
 if x is None:x=OxmlElement('w:tcMar');p.append(x)
 for side,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
  y=x.find(qn('w:'+side))
  if y is None:y=OxmlElement('w:'+side);x.append(y)
  y.set(qn('w:w'),str(val));y.set(qn('w:type'),'dxa')
def fmt(r,size=10.2,bold=None,color=INK,italic=None):
 r.font.name=FONT;r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT);r.font.size=Pt(size)
 if bold is not None:r.bold=bold
 if color:r.font.color.rgb=RGBColor.from_string(color)
 if italic is not None:r.italic=italic
def inline(p,t,size=10.2,color=INK):
 pat=re.compile(r'(\*\*.*?\*\*|`.*?`|(?<!\*)\*[^*]+\*)');pos=0
 for m in pat.finditer(t):
  if m.start()>pos:fmt(p.add_run(t[pos:m.start()]),size,color=color)
  z=m.group(0)
  if z.startswith('**'):fmt(p.add_run(z[2:-2]),size,True,color)
  elif z.startswith('`'):fmt(p.add_run(z[1:-1]),size,color=color)
  else:fmt(p.add_run(z[1:-1]),size,color=color,italic=True)
  pos=m.end()
 if pos<len(t):fmt(p.add_run(t[pos:]),size,color=color)
def setup_page(sec,land):
 sec.orientation=WD_ORIENT.LANDSCAPE if land else WD_ORIENT.PORTRAIT
 sec.page_width=Cm(27.94 if land else 21.59);sec.page_height=Cm(21.59 if land else 27.94)
 sec.left_margin=Cm(1.55);sec.right_margin=Cm(1.55);sec.top_margin=Cm(1.65);sec.bottom_margin=Cm(1.55);sec.header_distance=Cm(.7);sec.footer_distance=Cm(.7)
def sect(doc,land):
 s=doc.add_section(WD_SECTION_START.NEW_PAGE);setup_page(s,land);return s
def hf(doc):
 for s in doc.sections:
  h=s.header.paragraphs[0];h.alignment=WD_ALIGN_PARAGRAPH.RIGHT;h.text='TTQS ONE　｜　DOCX 0066　｜　SAMPLE / SYNTHETIC'
  for r in h.runs:fmt(r,8.2,True,TEAL)
  f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER;f.text='CANDIDATE　•　全篇示範資料　•　不是原始點協會真實班次佐證　•　第 '
  for r in f.runs:fmt(r,8.1,color=MUTED)
  fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p[-1].addnext(fld);fmt(f.add_run(' 頁'),8.1,color=MUTED)
def gettable(ls,i):
 got=[]
 while i<len(ls) and ls[i].strip().startswith('|'):got.append(ls[i].strip());i+=1
 rows=[]
 for n,line in enumerate(got):
  cs=[x.strip() for x in line.strip('|').split('|')]
  if n==1 and all(re.fullmatch(r':?-{2,}:?',x.replace(' ','')) for x in cs):continue
  rows.append(cs)
 return rows,i
def addtable(d,rows,land):
 if not rows:return
 n=max(map(len,rows));t=d.add_table(rows=len(rows),cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=True;t.style='Table Grid'
 for ri,row in enumerate(rows):
  trpr=t.rows[ri]._tr.get_or_add_trPr();cans=OxmlElement('w:cantSplit');trpr.append(cans)
  for ci in range(n):
   c=t.cell(ri,ci);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;margins(c);p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.02
   inline(p,row[ci] if ci<len(row) else '',9 if land else 8.7,'FFFFFF' if ri==0 else INK)
   if ri==0:
    shade(c,NAVY)
    for r in p.runs:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
   elif ri%2==0:shade(c,ALT)
 if t.rows:
  x=OxmlElement('w:tblHeader');x.set(qn('w:val'),'true');t.rows[0]._tr.get_or_add_trPr().append(x)
 p=d.add_paragraph();p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(0)
ls=SRC.read_text(encoding='utf-8').splitlines();d=Document();setup_page(d.sections[0],False);st=d.styles
normal=st['Normal'];normal.font.name=FONT;normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT);normal.font.size=Pt(10.2);normal.font.color.rgb=RGBColor.from_string(INK)
for name,size,color in [('Title',22,'111111'),('Heading 1',15,NAVY),('Heading 2',12.5,TEAL),('Heading 3',11,TEAL)]:
 x=st[name];x.font.name=FONT;x._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT);x.font.size=Pt(size);x.font.bold=True;x.font.color.rgb=RGBColor.from_string(color);x.paragraph_format.keep_with_next=True;x.paragraph_format.space_before=Pt(8);x.paragraph_format.space_after=Pt(4)
 # Remove Word's built-in decorative paragraph rules from Title style.
 if name=='Title' and x._element.pPr is not None:
  b=x._element.pPr.find(qn('w:pBdr'))
  if b is not None:x._element.pPr.remove(b)
land=False;title=False;i=0
land_starts=('## 五、來源與性質清冊','## 十二、問卷逐題分布')
port_starts=('## 八、需求至評量證據鏈','## 十三、追蹤表（沿用固定資料）','## 十六、改善閉環')
while i<len(ls):
 raw=ls[i].rstrip();s=raw.strip()
 if not s or s.startswith('〔第 ') or s.startswith('（全頁提醒：') or s=='---':i+=1;continue
 if s.startswith('|'):rows,i=gettable(ls,i);addtable(d,rows,land);continue
 if s.startswith('## '):
  if any(s.startswith(x) for x in land_starts) and not land:sect(d,True);land=True
  elif any(s.startswith(x) for x in port_starts) and land:sect(d,False);land=False
  p=d.add_paragraph(style='Heading 2');inline(p,s[3:].strip(),12.5,TEAL);i+=1;continue
 if s.startswith('### '):p=d.add_paragraph(style='Heading 3');inline(p,s[4:].strip(),11,TEAL);i+=1;continue
 if s.startswith('# '):
  is_title=not title;title=True;p=d.add_paragraph(style='Title' if is_title else 'Heading 1');inline(p,s[2:].strip(),22 if is_title else 15,'111111' if is_title else NAVY);p.paragraph_format.keep_with_next=True
  if is_title:
   pp=p._p.get_or_add_pPr();b=pp.find(qn('w:pBdr'))
   if b is not None:pp.remove(b)
  i+=1;continue
 if re.match(r'^[-*]\s+',s):
  p=d.add_paragraph(style='List Bullet');p.paragraph_format.left_indent=Cm(.55);p.paragraph_format.first_line_indent=Cm(-.25);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.05;inline(p,re.sub(r'^[-*]\s+','',s),10,INK);i+=1;continue
 m=re.match(r'^(\d+)[.)]\s+(.*)$',s)
 if m:
  p=d.add_paragraph();p.paragraph_format.left_indent=Cm(.6);p.paragraph_format.first_line_indent=Cm(-.6);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.05;inline(p,m.group(1)+'.  '+m.group(2),10,INK);i+=1;continue
 p=d.add_paragraph();p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.08;p.paragraph_format.widow_control=True;inline(p,s,10.2,INK);i+=1
hf(d);d.core_properties.title='訓練班級完整結案報告';d.core_properties.subject='TTQS ONE DOCX 0066 / I15-E02 / official checkbox 072 / SAMPLE-SYNTHETIC candidate';d.core_properties.author='原始點協會 TTQS ONE 文件示範（候選）';d.core_properties.keywords='0066,I15-E02,072,SAMPLE,SYNTHETIC'
OUT.parent.mkdir(parents=True,exist_ok=True);d.save(OUT);sha=hashlib.sha256(OUT.read_bytes()).hexdigest();print(json.dumps({'docx':str(OUT),'sha256':sha,'bytes':OUT.stat().st_size,'sections':len(d.sections),'body_sha256':hashlib.sha256(SRC.read_bytes()).hexdigest()},ensure_ascii=False))

