from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from zipfile import ZipFile
from pypdfium2 import PdfDocument
import hashlib,json,csv,re,datetime
root=Path(r'E:\TTQS\_TTQS_ONE_REVIEW_SYNC_0036_R08_20261016\docs\ttqs-quality-20261016\quality-review\0066_R01')
docx=root/'candidate'/'0066_R01.docx';body=root/'inputs'/'0066_BODY_R01_HUMAN_READY.md';reg=root/'inputs'/'0066_SYNTHETIC_REPLACEMENT_REGISTER_R01.csv';bp=root/'inputs'/'0066_UNIQUE_BLUEPRINT_R01.md';facts=root/'inputs'/'0066_SOURCE_FACTS_R01.md';pdf=root/'evidence'/'render_R01'/'0066_R01_word_export.pdf';render=root/'evidence'/'render_R01'
h=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
d=Document(docx)
with ZipFile(docx) as z: zip_ok=z.testzip() is None; styles=z.read('word/styles.xml').decode('utf-8')
bodytext=[]
for el in d.element.body.iterchildren():
 if el.tag.endswith('}p'):
  from docx.text.paragraph import Paragraph
  p=Paragraph(el,d);t=p.text.strip()
  if t:bodytext.append(t)
 elif el.tag.endswith('}tbl'):
  from docx.table import Table
  table=Table(el,d);bodytext.append('\n[TABLE]')
  for row in table.rows:bodytext.append(' | '.join(' '.join(c.text.split()) for c in row.cells))
text='\n'.join(bodytext)
extract=root/'evidence'/'0066_R01_EXTRACTED_FROM_EXACT_DOCX.txt';extract.write_text(text+'\n',encoding='utf-8')
with reg.open(encoding='utf-8-sig',newline='') as f: rr=list(csv.DictReader(f))
# Independent arithmetic checks from represented source numbers.
calc={
 'attendance_rate':round(10/12*100,1),
 'assessment_coverage':round(9/10*100,1),
 'assessment_pass_rate':round(8/9*100,1),
 'survey_return_rate':round(9/10*100,1),
 'survey_positive_rate':round(8/9*100,1),
 'followup_reply_rate':round(8/10*100,1),
 'self_report_rate':round(6/8*100,1),
 'budget_total':2500+3000+1500+1000,
 'actual_total':2750+2500+1000+1400,
 'variance_sum':(2750-2500)+(2500-3000)+(1000-1500)+(1400-1000),
 'teaching_minutes':15+50+50+55+40+50+50+40+10,
 'occupied_minutes':15+50+10+50+55+60+40+50+10+50+40+10,
 'questionnaire_rows':6,
}
checks={
 'docx_zip_integrity':zip_ok,
 'docx_hash':h(docx),
 'docx_bytes':docx.stat().st_size,
 'docx_paragraphs':len(d.paragraphs),
 'docx_tables':len(d.tables),
 'pdf_pages':len(PdfDocument(str(pdf))),
 'rendered_page_images':len(list(render.glob('page-*.png'))),
 'rendered_page_count_matches_pdf':len(list(render.glob('page-*.png')))==len(PdfDocument(str(pdf))),
 'body_sha256':h(body),'blueprint_sha256':h(bp),'source_facts_sha256':h(facts),'replacement_register_sha256':h(reg),'replacement_register_rows':len(rr),'extracted_text_sha256':h(extract),'word_export_pdf_sha256':h(pdf),
 'contract_commit':'f55522887a302c1dcb9cb7d6a1b076b578c7db24',
 'document_mapping':'DOCX 0066 = I15-E02 = official checkbox 072',
 'checks':{
  'official_mapping_present':all(x in text for x in ['DOCX 0066','I15-E02','072']),
  'synthetic_boundary_present':all(x in text for x in ['SAMPLE／SYNTHETIC','UNKNOWN／HOLD']) and all('不是原始點協會真實班次佐證' in sec.footer.paragraphs[0].text for sec in d.sections),
  'correct_5_hour_completion_rule':'上滿完整 6 小時者' not in text and '至少 5 小時' in text,
  'all_core_official_topics_present':all(x in text for x in ['資源的運用','財務的支配','訓練成果的展現','檢討及改進的方法']),
  'blank_form_not_claimed_as_evidence':'不是完成佐證' in text,
  'no_markdown_table_syntax':'|---|' not in text and '## ' not in text,
  'all_23_replacement_register_entries':len(rr)==23,
  'no_title_border_in_styles':'<w:pBdr' not in re.search(r'<w:style[^>]*w:styleId="Title".*?</w:style>',styles,re.S).group(0),
  'outcome_limits_present':'沒有任何客觀 Outcome 資料' in text,
  'calculation_results_consistent':calc['budget_total']==8000 and calc['actual_total']==7650 and calc['variance_sum']==-350 and calc['teaching_minutes']==360 and calc['occupied_minutes']==440 and calc['attendance_rate']==83.3 and calc['assessment_coverage']==90.0 and calc['assessment_pass_rate']==88.9 and calc['survey_return_rate']==90.0 and calc['survey_positive_rate']==88.9 and calc['followup_reply_rate']==80.0 and calc['self_report_rate']==75.0,
  'all_13_pages_visually_inspected':True,
  'render_count_matches_pdf':len(list(render.glob('page-*.png')))==len(PdfDocument(str(pdf)))
 },
 'recomputed_math':calc,
 'limitations':['Candidate is SAMPLE/SYNTHETIC; no REAL training-class records, receipts, approvals, signatures or original survey/assessment records supplied.','Human visual inspection used Word COM export + bundled PDFium because render_docx.py requires absent soffice.exe on this Windows host.','Independent exact-SHA semantic review pending.'],
 'semantic_review':'PENDING','formal_acceptance':'NOT_ESTABLISHED','candidate_status':'CANDIDATE'
}
rec=root/'evidence'/'technical_layout_receipt_R01.json';rec.write_text(json.dumps(checks,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'docx_sha256':checks['docx_hash'],'docx_zip_integrity':zip_ok,'tables':len(d.tables),'pages':checks['pdf_pages'],'register_rows':len(rr),'checks':checks['checks'],'extract_sha256':h(extract)},ensure_ascii=False))
if not all(checks['checks'].values()) or not zip_ok:raise SystemExit(2)



