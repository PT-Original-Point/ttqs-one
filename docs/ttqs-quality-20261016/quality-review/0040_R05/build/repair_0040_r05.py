from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.shared import Pt

ROOT = Path(r'E:\TTQS\_TTQS_ONE_REVIEW_SYNC_0036_R08_20261016\docs\ttqs-quality-20261016\quality-review')
SOURCE = ROOT / '0040_R04' / 'candidate' / '0040_R04.docx'
OUTPUT = ROOT / '0040_R05' / 'candidate' / '0040_R05.docx'


def replace_paragraph(paragraph, new_text):
    paragraph.clear()
    run = paragraph.add_run(new_text)
    run.font.name = 'Microsoft JhengHei'


def set_cell(cell, text, header=False):
    cell.text = text
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.05
        for run in paragraph.runs:
            run.font.name = 'Microsoft JhengHei'
            run.font.size = Pt(9.5)
            run.bold = bool(header)
            if header:
                run.font.color.rgb = __import__('docx').shared.RGBColor(255, 255, 255)


def replace_table(table, headers, data):
    if len(headers) != len(table.columns) or any(len(row) != len(headers) for row in data):
        raise ValueError('Table column mismatch')
    original_rows = list(table._tbl.tr_lst)
    if len(original_rows) < 2:
        raise ValueError('Expected formatted header and body rows')
    body_templates = [deepcopy(original_rows[1])]
    if len(original_rows) > 2:
        body_templates.append(deepcopy(original_rows[2]))
    for row in list(table._tbl.tr_lst[1:]):
        table._tbl.remove(row)
    for idx, text in enumerate(headers):
        set_cell(table.rows[0].cells[idx], text, header=True)
    for i, values in enumerate(data):
        row_xml = deepcopy(body_templates[i % len(body_templates)])
        table._tbl.append(row_xml)
        row = table.rows[-1]
        for cell, value in zip(row.cells, values):
            set_cell(cell, value)


doc = Document(SOURCE)
paragraph_replacements = {
    '查核索引與文件版次紀錄': '評量工具使用方式與版本狀態',
    '本頁只管理本候選工具的版本、效力狀態及內容定位；不構成原始點協會核准、採用或課程執行紀錄。':
        '本頁說明示範工具的使用時點與正式採用前要核實的資料、權限及責任。案例、評分與結果均為 SAMPLE / SYNTHETIC。',
    '目前文件控制': '工具版本與使用狀態',
    '修訂歷史（均為候選版本，不代表正式採用）': '本版調整摘要',
    '內容查核索引': '依工作時點操作導覽',
}
seen = set()
for paragraph in doc.paragraphs:
    old = paragraph.text.strip()
    if old in paragraph_replacements:
        replace_paragraph(paragraph, paragraph_replacements[old])
        seen.add(old)
if seen != set(paragraph_replacements):
    missing = set(paragraph_replacements) - seen
    raise ValueError(f'Missing expected headings: {sorted(missing)}')

replace_table(doc.tables[12], ['使用者要確認的項目', '目前狀態與實際使用方式'], [
    ['示範版本與資料性質', '示範工具 v1.1。案例代碼、日期、角色、評分、計算及結果均為 SAMPLE / SYNTHETIC，不能當作原始點協會的營運、課程或評量紀錄。'],
    ['正式效力', '尚未由協會具權責人員核准或採用；不代表已辦訓、已評量、成效達標或已取得外部證明。'],
    ['適用時點', '課程設計定案前，核對需求、目標、量尺、評量時點、門檻及證據欄位；課後只用已核實且有權處理的真實評量資料分析實際成效。'],
    ['正式使用前核實', '確認需求來源與日期、適用課程和對象、目標與任務難度、門檻及缺漏規則、參與同意、資料最小化、存取與保存、覆核方式。門檻或題目變更須先核定版本，不得依結果回頭修改。'],
    ['角色與權責', '「課程設計角色、評量角色、覆核角色」只表示工作功能。實際人員、授權與交接依協會有效章程、權責資料及原件確認，不預設現行職稱或人員。'],
    ['替換示範資料', '取得合法可追溯的真實來源並確認使用權限後，逐筆替換相應示範事實；保存原始評分、工具版本及覆核紀錄。不得仿製簽名、核准、第三方文件或真人參與紀錄。'],
])
replace_table(doc.tables[13], ['版本', '本版調整', '使用效力'], [
    ['示範工具 v1.1', '改以使用者操作導覽說明工作步驟，並列明正式使用前須核實的資料、授權與責任。原有示範分數、門檻、計算及 D04 追蹤案例未變。', 'SAMPLE；尚未核准或生效'],
])
replace_table(doc.tables[14], ['工作時點／目的', '操作位置', '實際操作與應保留的結果'], [
    ['課程規劃前：確認需求與目標', '第 1–3 頁', '先核對需求來源、日期、課程及對象，再把差距連到可觀察行為、評量方法、時點、門檻、證據欄位與責任角色。來源未核實時標 SAMPLE／待驗證。'],
    ['設定評量規則', '第 3 頁', '本示範採 0／1／2 量尺；示範通過條件為總分至少 6/8 且 K3 至少 1 分。正式課程須由有權責者事前核定自己的量尺與門檻，不得看完結果後回改。'],
    ['登錄前後測與檢查完整性', '第 4–5 頁', '使用去識別代碼、同版量尺及相近難度任務；四構面完整才計總分。缺漏不填 0、不補值；有效樣本數為 0 時報 N/A 並交代原因。'],
    ['解讀設計覆蓋與成效指標', '第 2、4 頁', '覆蓋率＝有直接評量的目標數÷全部目標數；示範 4/4 只代表設計覆蓋。平均、配對變化與達標率各依表列有效分母重算；合成結果不可當作本會成效或因果證明。'],
    ['處理評量異常', '第 5 頁', '依表處理缺欄、K3 安全界線、評分差異、任務版本不同或 N=0；保留原始評分與理由，不用口頭印象、補值或覆寫結果。'],
    ['追蹤未達項目與交接', '第 5–6 頁', '把差距連到示範風險、責任角色、期限、交接資料、重評條件及結案規則。D04 的補強與重評尚未執行，不能改寫為已完成或通過。'],
    ['替換為正式工作資料', '第 6–7 頁', '逐項取得有權使用的原始需求、課程、評量與核准資料，核實來源、責任、同意、隱私、版本及計算後再替換 SAMPLE。未取得前不得以本示範或空表作為真實佐證。'],
    ['區分相鄰工作文件', '第 2 頁', '會議紀錄記載真實討論及決議；需求調查呈現需求資料；課程計畫載明課務安排；系統檢討記錄改善。本矩陣只協助設計與核對評量方法，不代替上述文件。'],
])

# Keep the existing section, page furniture, and all substantive pages 1-6 intact.
doc.core_properties.title = '訓練評估方法與成效指標對照表（示範工具）'
doc.core_properties.subject = '使用者版次與操作導覽；SAMPLE / SYNTHETIC'
doc.core_properties.author = ''
doc.core_properties.last_modified_by = ''
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(f'Created targeted R05 candidate from R04; tables changed=12,13,14; paragraphs changed={len(seen)}; output={OUTPUT}')
