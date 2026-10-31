from pathlib import Path
from zipfile import ZipFile
import csv,hashlib,json,re,xml.etree.ElementTree as ET
from pypdf import PdfReader
from docx import Document

ROOT=Path(r'E:\TTQS\_TTQS_ONE_REVIEW_SYNC_0036_R08_20261016\docs\ttqs-quality-20261016\quality-review')
P=ROOT/'0040_R05'; P04=ROOT/'0040_R04'
docx=P/'candidate/0040_R05.docx'; pdf=P/'evidence/render/candidate.pdf'
register=P/'evidence/SYNTHETIC_REPLACEMENT_REGISTER_R05.csv'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
norm=lambda s:re.sub(r'[\W_]+','',s,flags=re.UNICODE).lower()
with ZipFile(docx) as z:
    corrupt=z.testzip(); xmlnames=[n for n in z.namelist() if n.endswith('.xml')]; xmlerrors=[]
    for n in xmlnames:
        try: ET.fromstring(z.read(n))
        except Exception as e: xmlerrors.append({'part':n,'error':str(e)})
d=Document(docx); reader=PdfReader(str(pdf)); pages=[' '.join((p.extract_text() or '').split()) for p in reader.pages]
with register.open(encoding='utf-8-sig',newline='') as f: rows=list(csv.DictReader(f))
ids=[r['fact_id'] for r in rows]; candidate_sha=sha(docx)
binding=all(r['bound_candidate_sha256'].lower()==candidate_sha for r in rows)
anchors={
'0040-REF-001':[(1,'官方要求（逐字引用）'),(1,'來源'),(1,'PDF 第 3 頁')],
'0040-SYN-001':[(1,'SAMPLE-0039-01')],
'0040-SYN-002':[(1,'社區服務志工接待與活動引導')],
'0040-SYN-003':[(1,'示範資料日期'),(1,'2026-06-12')],
'0040-SYN-004':[(3,'K1'),(3,'K2'),(3,'K3'),(3,'K4')],
'0040-SYN-005':[(3,'0 分：未達'),(3,'1 分：部分達成'),(3,'2 分：完整達成')],
'0040-SYN-006':[(3,'總分至少 6 分'),(3,'K3 至少 1 分'),(4,'前測 K1/K2/K3/K4；總分'),(4,'後測 K1/K2/K3/K4；總分'),(4,'依門檻判定')],
'0040-SYN-007':[(1,'開課前做基線')],
'0040-SYN-008':[(1,'最後一次練習後、結訓前做後測')],
'0040-SYN-009':[(1,'30 日追蹤另行核准'),(6,'本例尚未執行，沒有重評結果')],
'0040-SYN-010':[(3,'評量角色'),(3,'評量覆核角色'),(5,'課程設計角色')],
'0040-SYN-BL-K1':[(3,'K1 開場與角色說明'),(3,'基線 1.0，差距 1.0')],
'0040-SYN-BL-K2':[(3,'K2 澄清與摘要確認'),(3,'基線 0.6，差距 1.4')],
'0040-SYN-BL-K3':[(3,'K3 同意與隱私界線'),(3,'基線 1.2，差距 0.8')],
'0040-SYN-BL-K4':[(3,'K4 收尾與交接'),(3,'基線 0.6，差距 1.4')],
'0040-SYN-CALC-001':[(2,'4 ÷ 4＝100%')],
'0040-SYN-CALC-002':[(4,'平均前測'),(4,'17 ÷ 5'),(4,'3.4/8')],
'0040-SYN-CALC-003':[(4,'平均後測'),(4,'31 ÷ 5'),(4,'6.2/8')],
'0040-SYN-CALC-004':[(4,'配對平均變化'),(4,'14 ÷ 5'),(4,'2.8 分/8')],
'0040-SYN-CALC-005':[(4,'前測 K1/K2/K3/K4；總分'),(4,'後測 K1/K2/K3/K4；總分'),(4,'D01'),(4,'D02'),(4,'D03'),(4,'D05'),(4,'後測達標率')],
'0040-SYN-CALC-006':[(4,'後測達標率'),(4,'4 ÷ 5'),(4,'80%')],
'0040-SYN-CALC-007':[(4,'前測 K1–K4'),(4,'1.0/0.6/1.2/0.6')],
'0040-SYN-CALC-008':[(4,'後測 2.0/1.6/1.6/1.0')],
'0040-SYN-CALC-009':[(5,'後續行動（示範待辦）'),(6,'D04')],
'0040-SYN-034':[(6,'D04'),(6,'5/8'),(6,'低於 6/8')],
'0040-SYN-035':[(6,'訓練練習風險'),(6,'不是實際服務事件')],
'0040-SYN-036':[(6,'課程設計角色'),(6,'評量覆核角色')],
'0040-SYN-037':[(6,'2026-06-19'),(6,'SAMPLE 目標期限')],
'0040-SYN-038':[(6,'交接資料／結案規則'),(6,'總分至少 6/8'),(6,'K3 至少 1 分')]
}
pre_vectors=['1/0/1/0','0/1/1/1','1/1/2/0','1/0/1/1','2/1/1/1']
post_vectors=['2/1/2/1','2/2/2/1','2/2/1/1','2/1/1/1','2/2/2/1']
for i in range(1,6):
    anchors[f'0040-SYN-PRE-D{i:02d}']=[(4,f'D{i:02d}'),(4,pre_vectors[i-1])]
    anchors[f'0040-SYN-POST-D{i:02d}']=[(4,f'D{i:02d}'),(4,post_vectors[i-1])]
def has(pg,token): return norm(token) in norm(pages[pg-1])
anchor_results={fid:all(has(pg,tok) for pg,tok in checks) for fid,checks in anchors.items()}
# Recalculate all five displayed pairs from the document table itself.
score_rows=[]
for row in d.tables[6].rows[1:]:
    code=row.cells[0].text.strip()
    parse=lambda s: tuple(map(int,re.search(r'([0-2])/([0-2])/([0-2])/([0-2])\s*[；;]\s*(\d+)/8',s).groups()))
    pre=parse(row.cells[1].text); post=parse(row.cells[2].text)
    if sum(pre[:4])!=pre[4] or sum(post[:4])!=post[4]: raise ValueError('Displayed total mismatch '+code)
    delta=int(row.cells[3].text.replace('+','').strip())
    if post[4]-pre[4]!=delta: raise ValueError('Displayed change mismatch '+code)
    pass_expected=post[4]>=6 and post[2]>=1
    if row.cells[4].text.strip().startswith('通過：')!=pass_expected: raise ValueError('Displayed threshold verdict mismatch '+code)
    score_rows.append({'case':code,'pre_total':pre[4],'post_total':post[4],'delta':delta,'pass':pass_expected})
if len(score_rows)!=5: raise ValueError('Expected five score rows')
pre_sum=sum(x['pre_total'] for x in score_rows); post_sum=sum(x['post_total'] for x in score_rows); change_sum=sum(x['delta'] for x in score_rows)
pass_count=sum(x['pass'] for x in score_rows)
# Text identity of p.1-p.6 is checked after normalization against the reviewed R04 PDF.
r04pdf=P04/'evidence/render/0040_R04_final_word_export.pdf'
r04pages=[' '.join((p.extract_text() or '').split()) for p in PdfReader(str(r04pdf)).pages]
first_six_text_equal={str(i):norm(pages[i-1])==norm(r04pages[i-1]) for i in range(1,7)}
# Body pages 7-8 must contain no internal routing, SHA, or reviewer-history data.
internal_patterns=[r'(?i)\b[0-9a-f]{64}\b',r'(?i)\bDOCX\s*0040\b',r'(?i)I08\s*[-–]\s*E03',r'(?i)checkbox\s*040',r'(?i)SHA\s*[-–]?\s*256',r'(?i)R0[1-4]\b',r'(?i)PARTIAL']
pages_78=' '.join(pages[6:8])
leaks=[pat for pat in internal_patterns if re.search(pat,pages_78)]
new_sections={'user_facing_status':has(7,'工具版本與使用狀態'),'practical_index':has(7,'依工作時點操作導覽'),'sample_boundary':has(7,'SAMPLE / SYNTHETIC'),'not_approved':has(7,'尚未由協會具權責人員核准或採用'),'n0_na':has(8,'有效樣本數為 0 時報 N/A'),'coverage_is_design_only':has(8,'只代表設計覆蓋'),'separates_neighbor_documents':has(8,'本矩陣只協助設計與核對評量方法')}
# OOXML direct font sizes.
font_sizes=[]
with ZipFile(docx) as z:
 root=ET.fromstring(z.read('word/document.xml')); ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
 body=root.find('w:body',ns)
 for child in list(body):
  is_table=child.tag.endswith('}tbl')
  for run in child.findall('.//w:r',ns):
   size=run.find('./w:rPr/w:sz',ns)
   if size is not None:
    try: font_sizes.append((int(size.attrib.get('{'+ns['w']+'}val'))/2,is_table))
    except Exception: pass
min_body=min((v for v,t in font_sizes if not t),default=None); min_table=min((v for v,t in font_sizes if t),default=None)
render_dir=P/'evidence/render'; pngs=sorted(render_dir.glob('page-*.png'))
png_hashes={str(i):sha(render_dir/f'page-{i}.png') for i in range(1,len(pngs)+1)}
report={
 'project_id':'TTQS_ONE','doc_id':'0040','revision':'R05','requirement_id':'I08-E03','official_checkbox_sequence':'040',
 'candidate_sha256':candidate_sha,'parent_candidate_r04_sha256':sha(P04/'candidate/0040_R04.docx'),
 'register_sha256':sha(register),'register_rows':len(rows),'register_unique_fact_ids':len(set(ids)),
 'register_all_bound_to_candidate_sha':binding,'register_rows_with_validated_printed_anchors':sum(anchor_results.values()),'locator_anchor_checks':anchor_results,
 'docx_zip_integrity':'PASS' if corrupt is None else 'FAIL','xml_parts':len(xmlnames),'xml_parse_errors':xmlerrors,'table_count':len(d.tables),
 'rendered_pdf_sha256':sha(pdf),'rendered_pages':len(reader.pages),'page_png_sha256':png_hashes,
 'page_1_to_6_extracted_text_equal_to_R04':first_six_text_equal,'sample_score_recalculation':{'rows':score_rows,'pre_total':pre_sum,'post_total':post_sum,'paired_change_total':change_sum,'pass_count':pass_count,'denominator':len(score_rows),'pre_mean':pre_sum/len(score_rows),'post_mean':post_sum/len(score_rows),'paired_mean_change':change_sum/len(score_rows),'pass_rate_percent':pass_count/len(score_rows)*100},
 'min_direct_body_font_pt':min_body,'min_direct_table_font_pt':min_table,'pp7_8_internal_content_patterns_found':leaks,'new_section_checks':new_sections,
 'r04_reviewed_state':'PARTIAL; exact R04 preserved; repair limited to review findings','status':'PASS_TECHNICAL' if corrupt is None and not xmlerrors and len(reader.pages)==8 and len(pngs)==8 and len(rows)==39 and len(set(ids))==39 and binding and all(anchor_results.values()) and all(first_six_text_equal.values()) and not leaks and all(new_sections.values()) and min_body>=10.5 and min_table>=9.5 and pre_sum==17 and post_sum==31 and change_sum==14 and pass_count==4 else 'FAIL'
}
out=P/'evidence/qa/0040_technical_review_r05_exact.json'; out.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'status':report['status'],'candidate_sha256':candidate_sha,'pdf_sha256':sha(pdf),'register_sha256':sha(register),'rows':len(rows),'anchors_pass':sum(anchor_results.values()),'anchors_total':len(anchor_results),'pages':len(reader.pages),'p1_p6_equal':first_six_text_equal,'leaks':leaks,'font_min':[min_body,min_table],'calculations':report['sample_score_recalculation']},ensure_ascii=False,indent=2))
