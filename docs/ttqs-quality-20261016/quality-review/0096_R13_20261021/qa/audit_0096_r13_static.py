from pathlib import Path
from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table
from pypdf import PdfReader
import csv,hashlib,json,re,zipfile

R=Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0096_R01_ANTIGRAVITY_WRITER_R01")
D=R/"output"/"0096_R13.docx"; R12PREV=R/"output"/"0096_R12.docx"
MD=R/"source"/"0096_R13_CANDIDATE_MANUSCRIPT.md"; CSV=R/"source"/"0096_R13_SYNTHETIC_REPLACEMENT_REGISTER.csv"
PDF=next((R/"qa"/"rendered_r13").glob("candidate_*.pdf"))
SNAP=next((R/"qa"/"rendered_r13").glob("candidate_*.docx"))
OFF=R/"inputs"/"official_source_page6_context.txt"
BP=Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\CONTROL\BLUEPRINT_CURRENT.jsonl")
CW=Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\CONTROL\TTQS_ONE_129_REQUIREMENT_TO_142_DELIVERABLE_CROSSWALK_R01.csv")
SYNC=R/"qa"/"0096_R13_register_sync_receipt.json"
VIS=R/"qa"/"0096_R13_visual_review_receipt.json"
R12REVIEW=R/"review"/"opencode_0096_R12_independent"/"0096_R12_independent_semantic_review.md"
H=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
docsha=H(D); r12prevsha=H(R12PREV); mdsha=H(MD); csvsha=H(CSV); pdfsha=H(PDF)
sync=json.loads(SYNC.read_text(encoding="utf-8")); vis=json.loads(VIS.read_text(encoding="utf-8"))
with zipfile.ZipFile(D) as z:
    zip_error=z.testzip(); names=z.namelist(); xml=z.read("word/document.xml").decode("utf-8")
doc=Document(D)
parts=[]
for el in doc.element.body.iterchildren():
    kind=el.tag.rsplit("}",1)[-1]
    if kind=="p":
        p=Paragraph(el,doc)
        if p.text.strip(): parts.append(p.text.strip())
    elif kind=="tbl":
        t=Table(el,doc)
        parts.append("[TABLE]\n"+"\n".join(" | ".join(c.text.strip().replace("\n"," / ") for c in row.cells) for row in t.rows))
body="\n\n".join(parts)
reviewdir=R/"review"/"opencode_0096_R13_independent"
reviewdir.mkdir(parents=True,exist_ok=True)
SEM_REVIEW=reviewdir/"0096_R13_independent_semantic_review.md"
semreview_sha=H(SEM_REVIEW)
bodyfile=reviewdir/"0096_R13_docx_body_reading_order.txt"
bodyfile.write_text(body,encoding="utf-8",newline="\n")
rows=list(csv.DictReader(CSV.open(encoding="utf-8-sig",newline="")))
ids=[r["fact_id"] for r in rows]
regtable=next((t for t in doc.tables if t.rows and t.rows[0].cells[0].text.strip()=="事實代碼"),None)
def norm(x):
    x=x.replace("<br>"," ").replace("<br/>"," ").replace("**","").replace(chr(96),"")
    return re.sub(r"\s+"," ",x).strip()
embedded=[[norm(c.text) for c in row.cells] for row in regtable.rows[1:]] if regtable else []
keys=["fact_id","synthetic_value","doc_section_location","real_replacement_source","responsible_role","post_replacement_verification"]
csvnorm=[[norm(r[k]) for k in keys] for r in rows]
mismatches=[i for i,(a,b) in enumerate(zip(embedded,csvnorm),1) if a!=b]
codes=set(re.findall(r"\b(?:SAMPLE|HOLD|PENDING)-[A-Z0-9]+(?:-[A-Z0-9]+)*\b",body))
missing=sorted(codes-set(ids))
mdtext=MD.read_text(encoding="utf-8")
csv_digest_match=re.search(r"0096_R13_SYNTHETIC_REPLACEMENT_REGISTER\.csv；SHA-256：([0-9a-f]{64})",mdtext)
r13bodycodes=sorted(set(re.findall(r"\b(?:SAMPLE|HOLD|PENDING)-[A-Z0-9]+(?:-[A-Z0-9]+)*\b",body)))
body_min=[run.font.size.pt for p in doc.paragraphs for run in p.runs if run.font.size]
table_min=[run.font.size.pt for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs for run in p.runs if run.font.size]
pdfpages=len(PdfReader(str(PDF)).pages)
BPtext=BP.read_text(encoding="utf-8-sig").splitlines()
bprows=[json.loads(x) for x in BPtext if x.strip() and json.loads(x).get("deliverable_id")=="0096" and json.loads(x).get("requirement_id")=="I18-E03"]
bpraw=next(x for x in BPtext if x.strip() and json.loads(x).get("deliverable_id")=="0096" and json.loads(x).get("requirement_id")=="I18-E03")
cwrows=[x for x in csv.DictReader(CW.open(encoding="utf-8-sig",newline="")) if x.get("requirement_id")=="I18-E03"]
neighbor_rows=[x for x in csv.DictReader(CW.open(encoding="utf-8-sig",newline="")) if x.get("deliverable_ids") in ("0095","0096","0097","0098")]
assert len(bprows)==1 and len(cwrows)==1 and cwrows[0].get("deliverable_ids")=="0096"
assert H(R12PREV)=="1c4626130b160cf64f90e8c8f507f01bafbc09b70f7f1f4510051c7b11241d33"
assert H(R12REVIEW)=="cf13db9f899dff1818780c8d2cb85dbd5b77bd4f2e3b38ad05dc6b195e2a016a"
assert len(rows)==len(set(ids))==len(embedded)==121 and not mismatches and not missing
assert sync["embedded_csv_exact_match"] is True and sync["stale_row_status"]=="ALREADY_ABSENT_IN_R12"
assert sync["r13_csv_sha256"]==csvsha and sync["r13_manuscript_sha256"]==mdsha
assert vis["candidate_docx_sha256"]==docsha and vis["render_pdf_sha256"]==pdfsha and pdfpages==69
assert H(SNAP)==docsha and csv_digest_match and csv_digest_match.group(1)==csvsha
section9={
"SAMPLE-CASE-002":"第九節「合成更正與退件狀態演練（非實際案件）」下表資料列 1〈演練追蹤編號〉",
"SAMPLE-OLD-CATEGORY-01":"第九節「合成更正與退件狀態演練（非實際案件）」下表資料列 3〈原登錄值（舊值）〉",
"SAMPLE-NEW-CATEGORY-01":"第九節「合成更正與退件狀態演練（非實際案件）」下表資料列 4〈更正後值（新值）〉",
"SAMPLE-CORRECTION-REASON-01":"第九節「合成更正與退件狀態演練（非實際案件）」下表資料列 5〈更正原因〉",
"SAMPLE-ROLE-HANDLER-02":"第九節「合成更正與退件狀態演練（非實際案件）」下表資料列 7〈承辦角色〉",
"SAMPLE-ROLE-REVIEWER-02":"第九節「合成更正與退件狀態演練（非實際案件）」下表資料列 8〈覆核角色／利益衝突〉"}
csv_by_id={r["fact_id"]:r for r in rows}
exact_locators=all(csv_by_id.get(fid,{}).get("doc_section_location")==loc for fid,loc in section9.items())
tables_text=["\n".join(" | ".join(c.text for c in row.cells) for row in t.rows) for t in doc.tables]
case003_table=next((t for t in doc.tables if any("SAMPLE-ROLE-REVIEWER-03" in " | ".join(c.text for c in row.cells) for row in t.rows) and any("核心欄位" in c.text for c in t.rows[0].cells)),None)
case003_row11=" | ".join(c.text for c in case003_table.rows[11].cells) if case003_table and len(case003_table.rows)>11 else ""
sample_start=body.find("SAMPLE／SYNTHETIC 感謝函回覆範例")
sample_end=body.find("第五節",sample_start+1) if sample_start>=0 else -1
sample=body[sample_start:sample_end if sample_end>sample_start else len(body)] if sample_start>=0 else ""
checks={
 "official_doc_id_requirement_and_checkbox_distinct":"DOCX 0096" in body and "I18-E03" in body and "感謝函。" in body,
 "official_checkbox_text_exact_in_official_source":"感謝函。" in OFF.read_text(encoding="utf-8") and "感謝函。" in body,
 "unique_blueprint_title_mapped":"感謝函_可直接使用成品" in body,
 "neighbor_items_remain_distinct":all(x in body for x in ["0095","I18-E02","0097","I18-E05","0098","I18-E06"]),
 "sample_real_and_pending_boundary":("SAMPLE／SYNTHETIC" in body or "SAMPLE / SYNTHETIC" in body) and "生效日未設定" in body and "REAL 核實仍為 0／12" in body,
 "no_unsupported_17d_claim":"17d" not in body,
 "sample_does_not_fabricate_association_issuer":"原始點協會" not in sample,
 "general_survey_routed_to_0098":"0098" in body and ("不收集" in body or "不蒐集" in body) and ("一般課程意見" in body or "一般課程回饋" in body),
 "state_precedence_and_role_glossary":all(x in body for x in ["HOLD","INVESTIGATE","REJECT","經指定承辦角色","經指定覆核角色","經指定核准角色"]),
 "required_complete_examples_present":all(x in body for x in ["SAMPLE-CASE-003","SAMPLE-CASE-004","SAMPLE-CASE-INV-05"]),
 "sample_rate_arithmetic_present":all(x in body for x in ["12 ÷ 12 × 100%＝100%","0 ÷ 12 × 100%＝0%"]),
 "CASE001_fields_1_2_identified_as_sample":"第 1 追蹤編號及第 2 收件日期／途徑僅為 SAMPLE" in body,
 "CASE003_post_listing_followup_and_escalation":all(x in body for x in ["無核心缺證","RESTRICTED 列件後續追蹤","2026-11-26","2026-11-27","提報經指定核准角色"]),
 "CASE003_status_row_excludes_handler":case003_table is not None and "SAMPLE-ROLE-HANDLER-03" not in case003_row11 and "SAMPLE-ROLE-REVIEWER-03" in case003_row11,
 "CASE004_followup_roles_deadline_and_pause_rule":all(x in body for x in ["SAMPLE-ACTION-04","2026-10-27","SAMPLE-ROLE-HANDLER-04","SAMPLE-ROLE-REVIEWER-04","若未確認，立即暫停受限使用","提報經指定核准角色"]),
 "EX02_EX05_HOLD_timing_clarified":"由覆核角色立案時" not in body and "於登錄／建立追蹤時" in body,
 "Section9_locators_exact":exact_locators,
 "stale_orphan_fact_removed":"SAMPLE-DUE-RULE-INV-05" not in body and "SAMPLE-DUE-RULE-INV-05" not in mdtext,
 "register_csv_body_checksum_matches":bool(csv_digest_match) and csv_digest_match.group(1)==csvsha and len(rows)==121 and not mismatches,
 "build_snapshot_identical":H(SNAP)==docsha,
 "word_pdf_and_visual_receipt_exact_sha":vis["candidate_docx_sha256"]==docsha and vis["render_pdf_sha256"]==pdfsha and vis["page_count"]==pdfpages==69,
 "all_document_pages_visually_covered":vis["all_page_pngs_present"] is True and vis["all_page_contact_sheet_coverage"]["pages"]==list(range(1,70)),
 "docx_package_and_control_markup":zip_error is None and not re.search(r"<w:(ins|del|commentRangeStart)(?:\s|>)",xml) and "word/comments.xml" not in names,
 "minimum_explicit_body_font":bool(body_min) and min(body_min)>=10.5,
 "minimum_explicit_table_font":bool(table_min) and min(table_min)>=9.5,
 "candidate_revision_metadata":"R13" in body and "R13" in doc.core_properties.title and "R13" in doc.core_properties.subject,
 "review_F1_unsigned_is_missing_not_investigate":"未簽名文件先記缺件並維持 HOLD" in body and "正式立案後如有未簽名文件" not in body and "正式立案後發現未簽名文件" not in body,
 "review_F1_EX08_tracking_start_correct":"於登錄／建立追蹤時記錄期限、起算點及責任角色" in body and "立案時記錄期限、起算點及責任角色" not in body,
 "review_F3_INV05_value_matches_field1":any(r["fact_id"]=="SAMPLE-CASE-INV-05" and norm(r["synthetic_value"])==norm("SAMPLE-CASE-INV-05｜演練追蹤碼；非實際收文流水號") for r in rows) and "| SAMPLE-CASE-INV-05 | SAMPLE-CASE-INV-05｜演練追蹤碼；非實際收文流水號 |" in mdtext,
 "review_N1_official_trace_table_present":any(len(t.rows)>=5 and [norm(c.text) for c in t.rows[0].cells]==["官方要求／正式來源","表單欄位","文件內證據位置","覆核與狀態"] and any("I18-E03" in c.text for row in t.rows[1:] for c in row.cells) for t in doc.tables)
}
repairlog={
 "project_id":"TTQS_ONE","doc_id":"0096","requirement_id":"I18-E03","official_checkbox_serial":"099","revision":"R13",
 "candidate_state":"CANDIDATE_NOT_FORMALLY_ACCEPTED","predecessor_R12_sha256":r12prevsha,"candidate_docx_sha256":docsha,
 "source_manuscript_sha256":mdsha,"synthetic_register_sha256":csvsha,
 "changes":[
  {"finding":"F1","location":"第3節流程三第6點、第3節流程四第6點","repair":"未簽名文件明確先列缺件並維持HOLD；只有另有具體身分疑點、代簽疑慮或其他可指認矛盾且依核定程序正式立案後才記INVESTIGATE。"},
  {"finding":"F1","location":"第8節EX-08","repair":"期限、起算點與責任角色記錄時點改為登錄／建立追蹤時。"},
  {"finding":"F3","location":"第6節SAMPLE-CASE-INV-05欄1、Table 12、R13 CSV","repair":"清冊顯示值與欄位1同步為SAMPLE-CASE-INV-05｜演練追蹤碼；非實際收文流水號。"},
  {"finding":"N1","location":"第11節","repair":"新增四欄官方要求—表單欄位—文件內證據位置—覆核與狀態追溯表，僅映射第6頁I18-E03感謝函。"}],
 "prior_R12_partial_findings_preserved":True,"R12_candidate_modified":False,"no_real_records_or_approvals_added":True}
REPAIRLOG=R/"qa"/"0096_R13_body_repair_log.json"
REPAIRLOG.write_text(json.dumps(repairlog,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
bp_sha=hashlib.sha256(bpraw.encode("utf-8")).hexdigest()
offpdf_sha="805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709"
prhead="d5c4f0b73da4c98c9ba724424e05d3c8c549fb18"
binding={
 "project_id":"TTQS_ONE","doc_id":"0096","requirement_id":"I18-E03","official_checkbox_serial":"099",
 "revision":"R13","candidate_state":"CANDIDATE_NOT_FORMALLY_ACCEPTED","queue_state":"PENDING","queue_phase":"PENDING_HUMAN_FORMAL_ACCEPTANCE",
 "official_checkbox_text":"感謝函。","internal_scope_gloss":"索取稿及原件登錄，不假裝已收函","official_requirement_text":"感謝函。","official_source_page":"6","official_source_readback_status":"PDF_TEXT_EXACT_PAGE_MATCH",
 "official_source_page6_extract_sha256":H(OFF),"official_source_pdf_sha256":offpdf_sha,
 "quality_contract_source_commit":prhead,"quality_contract_03_blob_sha":"8c2bbe7ca9dcbb03088a48934b9e57e7147ede63","quality_contract_06_blob_sha":"263aaca84891bba8dca087875804343c3670e921",
 "unique_blueprint_sha256":bp_sha,"blueprint_control_source":"CONTROL/BLUEPRINT_CURRENT.jsonl","blueprint_source_file_sha256":H(BP),
 "crosswalk_source":"CONTROL/TTQS_ONE_129_REQUIREMENT_TO_142_DELIVERABLE_CROSSWALK_R01.csv","crosswalk_source_file_sha256":H(CW),"crosswalk_row":cwrows[0],"neighbor_crosswalk_rows":neighbor_rows,
 "candidate_docx_sha256":docsha,"candidate_docx_bytes":D.stat().st_size,"candidate_pdf_sha256":pdfsha,"candidate_pdf_pages":pdfpages,
 "source_manuscript_sha256":mdsha,"synthetic_replacement_register_sha256":csvsha,"register_rows":len(rows),
 "docx_body_extraction_sha256":H(bodyfile),"build_script_sha256":H(R/"build_docx_0096_r13.py"),
 "register_sync_script_sha256":H(R/"qa"/"build_0096_r13_register.py"),"register_sync_receipt_sha256":H(SYNC),"body_repair_log_sha256":H(REPAIRLOG),
 "static_technical_receipt":"SEE_0096_R13_STATIC_TECHNICAL_AUDIT_JSON","body_repair_log_sha256":H(REPAIRLOG),"visual_review_receipt_sha256":H(VIS),
 "predecessor_R12_candidate_sha256":r12prevsha,"predecessor_R12_review_response_sha256":H(R12REVIEW),
 "document_genre":"感謝函作業指引、索取稿與收件查驗表單（候選稿）",
 "repair_scope":"R13 addresses only exact-SHA R12 PARTIAL findings F1/F3 and Blueprint finding N1: unsigned-only remains HOLD in Process 3/4; EX-08 timing starts at registration/tracking; INV-05 field-1 display value is synchronized across case, Table 12 and CSV; add four-column page-6 I18-E03 trace table for form fields 1-13. R12 and its evidence remain unchanged.",
 "semantic_review_status":"PASS_EXACT_SHA_TEXT_ONLY_WITH_NONBLOCKING_EDITORIAL_NOTES","semantic_review_response_sha256":semreview_sha,"semantic_review_scope":"Text-only exact-SHA review of body and CSV; reviewer explicitly did not inspect DOCX binary or visual pages.","visual_review_status":"PASS_CONTACT_SHEETS_AND_TARGETED_FULL_PAGE"
}
bindpath=R/"inputs"/"source_binding_0096_R13.json"
bindpath.write_text(json.dumps(binding,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
build={
 "project_id":"TTQS_ONE","doc_id":"0096","requirement_id":"I18-E03","official_checkbox_serial":"099","revision":"R13",
 "candidate_docx_sha256":docsha,"candidate_docx_bytes":D.stat().st_size,"source_manuscript_sha256":mdsha,"register_sha256":csvsha,
 "build_script_sha256":H(R/"build_docx_0096_r13.py"),"predecessor_R12_sha256":r12prevsha,"predecessor_R12_unchanged":r12prevsha=="1c4626130b160cf64f90e8c8f507f01bafbc09b70f7f1f4510051c7b11241d33",
 "candidate_state":"CANDIDATE_NOT_FORMALLY_ACCEPTED","build_output_path":str(D)
}
(R/"qa"/"0096_R13_build_receipt.json").write_text(json.dumps(build,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
receipt={
 "project_id":"TTQS_ONE","doc_id":"0096","requirement_id":"I18-E03","official_checkbox_serial":"099","revision":"R13",
 "candidate_docx_sha256":docsha,"candidate_docx_bytes":D.stat().st_size,"predecessor_R12_sha256":r12prevsha,"predecessor_R12_unchanged":r12prevsha=="1c4626130b160cf64f90e8c8f507f01bafbc09b70f7f1f4510051c7b11241d33",
 "candidate_pdf_sha256":pdfsha,"pdf_pages":pdfpages,"docx_zip_integrity":"PASS" if zip_error is None else "FAIL:"+str(zip_error),
 "docx_table_count":len(doc.tables),"register_rows":len(rows),"register_unique_ids":len(set(ids)),"embedded_register_rows":len(embedded),
 "embedded_register_csv_exact_match":bool(regtable and not mismatches and len(embedded)==len(rows)),"embedded_register_mismatch_rows":mismatches,
 "body_code_count":len(codes),"body_codes_missing_from_register":missing,"minimum_explicit_body_font_pt":min(body_min) if body_min else None,
 "minimum_explicit_table_font_pt":min(table_min) if table_min else None,"checks":checks,"all_static_checks_pass":all(checks.values()),
 "current_06_blob_sha":"263aaca84891bba8dca087875804343c3670e921","current_03_blob_sha":"8c2bbe7ca9dcbb03088a48934b9e57e7147ede63","current_pr_head":prhead,
 "source_binding_sha256":H(bindpath),"visual_receipt_sha256":H(VIS),"register_sync_receipt_sha256":H(SYNC),"body_repair_log_sha256":H(REPAIRLOG),
 "review_status":"NOT_ASSESSED_BY_TECHNICAL_AUDIT","independent_semantic_review_status":"PASS_EXACT_SHA_TEXT_ONLY_WITH_NONBLOCKING_EDITORIAL_NOTES","independent_semantic_review_response_sha256":semreview_sha,
 "technical_scope":"Exact DOCX identity and package integrity, Word PDF render and page count, full page raster coverage, embedded Table 12 versus CSV row-by-row equality, source-code coverage, targeted case/locator/role/status invariants, official/Blueprint/crosswalk binding, document font minima, and no tracked changes/comments. This receipt is not an independent semantic verdict or organizational approval."
}
out=R/"qa"/"0096_R13_static_technical_audit.json"
out.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
print(json.dumps({"docx_sha256":docsha,"pdf_pages":pdfpages,"rows":len(rows),"font_min_body":min(body_min) if body_min else None,"font_min_table":min(table_min) if table_min else None,"checks":checks,"all_static_checks_pass":all(checks.values()),"static_receipt_sha256":H(out),"binding_sha256":H(bindpath)},ensure_ascii=True,indent=2))

