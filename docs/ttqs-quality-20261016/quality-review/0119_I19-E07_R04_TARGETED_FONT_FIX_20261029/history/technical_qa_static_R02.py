import sys, json, csv, hashlib, re, zipfile
from pathlib import Path
from docx import Document

DOCX=Path(sys.argv[1]); REGISTER=Path(sys.argv[2]); EXPECTED=sys.argv[3].lower(); OUT=Path(sys.argv[4])
def sha(p):
    h=hashlib.sha256()
    with p.open('rb') as f:
        for block in iter(lambda:f.read(1024*1024),b''): h.update(block)
    return h.hexdigest()
def style_size(style):
    while style is not None:
        try:
            if style.font.size is not None: return style.font.size.pt
            style=style.base_style
        except Exception: break
    return None
errors=[]
actual=sha(DOCX)
if actual != EXPECTED: errors.append('candidate_sha_mismatch')
with zipfile.ZipFile(DOCX) as z:
    bad=z.testzip()
    if bad is not None: errors.append('zip_crc_failure:'+bad)
doc=Document(str(DOCX))
paras=[p for p in doc.paragraphs if p.text.strip()]
tables=doc.tables
all_text=[]
all_text.extend(p.text for p in doc.paragraphs)
for ti,t in enumerate(tables):
    for row in t.rows:
        all_text.append(' | '.join(c.text for c in row.cells))
text='\n'.join(all_text)
flat=re.sub(r'\s+','',text).lower()
body_sizes=[]; table_sizes=[]; unresolved=[]
def collect(paragraphs, target, scope):
    for pi,p in enumerate(paragraphs):
        if not p.text.strip(): continue
        for ri,r in enumerate(p.runs):
            if not r.text.strip(): continue
            pt=r.font.size.pt if r.font.size is not None else style_size(r.style or p.style)
            if pt is None: unresolved.append(f'{scope}:{pi}:{ri}')
            else: target.append(float(pt))
for p in doc.paragraphs: collect([p],body_sizes,'body')
for ti,t in enumerate(tables):
    for ri,row in enumerate(t.rows):
        for ci,cell in enumerate(row.cells):
            for p in cell.paragraphs: collect([p],table_sizes,'table')
if not body_sizes: errors.append('no_body_font_sizes_resolved')
if not table_sizes: errors.append('no_table_font_sizes_resolved')
body_min=min(body_sizes) if body_sizes else None
table_min=min(table_sizes) if table_sizes else None
if body_min is not None and body_min < 10.5: errors.append('body_font_below_10_5pt')
if table_min is not None and table_min < 9.5: errors.append('table_font_below_9_5pt')
if len(tables)!=9: errors.append('table_count_not_9')
required={
 'market_baseline_8_20_40':('8/20','40.0%'),
 'market_compare_12_25_48':('12/25','48.0%'),
 'market_delta_8pp':('8.0','百分點'),
 'value_baseline_10_16_62_5':('10/16','62.5%'),
 'value_compare_15_20_75':('15/20','75.0%'),
 'value_delta_12_5pp':('12.5','百分點'),
 'baseline_response_16_20_80':('16/20','80.0%'),
 'compare_response_20_25_80':('20/25','80.0%')}
number_checks={k:all(v in flat for v in vals) for k,vals in required.items()}
for k,ok in number_checks.items():
    if not ok: errors.append('missing_or_unmatched_number:'+k)
truth_tokens=['sample','synthetic','示範','合成']
truth_count=sum(flat.count(t) for t in truth_tokens)
if truth_count < 4: errors.append('insufficient_sample_synthetic_markers')
with REGISTER.open('r',encoding='utf-8-sig',newline='') as f:
    reader=csv.DictReader(f); rows=list(reader); headers=reader.fieldnames or []
if len(rows)!=50: errors.append('register_row_count_not_50')
if not headers: errors.append('register_missing_header')
first_col=headers[0] if headers else None
ids=[(r.get(first_col,'') or '').strip() for r in rows] if first_col else []
if first_col and (any(not x for x in ids) or len(set(ids))!=len(ids)): errors.append('register_first_column_empty_or_duplicate')
reg_text=' '.join(' '.join(str(v or '') for v in r.values()) for r in rows).lower()
if len(rows) and not any(x in reg_text for x in ['sample','synthetic','示範','合成']): errors.append('register_not_labeled_synthetic')
result={
 'schema':'TTQS_ONE_0119_STATIC_TECHNICAL_QA_R02',
 'doc_id':'0119','requirement_id':'I19-E07','candidate_path':str(DOCX),
 'candidate_sha256':actual,'expected_sha256':EXPECTED,'candidate_bytes':DOCX.stat().st_size,
 'docx_zip_crc':'PASS' if bad is None else 'FAIL','python_docx_reopen':'PASS',
 'paragraph_count':len(doc.paragraphs),'nonblank_paragraph_count':len(paras),'table_count':len(tables),
 'body_min_font_pt':body_min,'table_min_font_pt':table_min,'unresolved_run_font_count':len(unresolved),
 'number_checks':number_checks,'sample_synthetic_marker_count':truth_count,
 'register_path':str(REGISTER),'register_sha256':sha(REGISTER),'register_row_count':len(rows),'register_headers':headers,'register_identifier_unique_nonblank':(first_col is not None and len(set(ids))==len(rows) and all(ids)),
 'errors':errors,'technical_result':'PASS_DETERMINISTIC_SCOPE' if not errors else 'FAIL',
 'scope_limit':'Static package integrity, text presence, numeric token cross-check, table count, register row identity, and resolvable font sizes only. No visual inspection, independent semantic acceptance, official/Human acceptance, or REAL evidence is asserted.'}
OUT.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(result,ensure_ascii=False,indent=2))