import sys,zipfile,hashlib,json
from pathlib import Path
from lxml import etree
from docx import Document
from docx.shared import Pt
src=Path(sys.argv[1]);dst=Path(sys.argv[2]);receipt=Path(sys.argv[3])
def sha(b):return hashlib.sha256(b).hexdigest()
source_sha=sha(src.read_bytes())
if source_sha!='1d626fdbcceb06668777315b73535799b28aadd51e3b32bf3be010f35557e5c0':raise SystemExit('SOURCE_SHA_MISMATCH')
doc=Document(str(src));p=doc.paragraphs[63];r=p.runs[0];old_pt=r.font.size.pt if r.font.size is not None else None
if old_pt!=10.0:raise SystemExit('TARGET_FONT_NOT_10PT')
text=r.text; text_sha=sha(text.encode('utf-8')); ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with zipfile.ZipFile(src,'r') as zin:
 infos=zin.infolist(); entries={i.filename:zin.read(i.filename) for i in infos}
 if 'word/document.xml' not in entries:raise SystemExit('DOCUMENT_XML_MISSING')
 root=etree.fromstring(entries['word/document.xml']); pars=root.xpath('./w:body/w:p',namespaces=ns)
 if len(pars)<=63:raise SystemExit('PARAGRAPH_INDEX_NOT_FOUND')
 runs=pars[63].xpath('./w:r',namespaces=ns)
 if len(runs)<=0:raise SystemExit('RUN_INDEX_NOT_FOUND')
 run=runs[0]; rtext=''.join(run.xpath('.//w:t/text()',namespaces=ns))
 if rtext!=text:raise SystemExit('XML_RUN_TEXT_MISMATCH')
 rpr=run.find('{%s}rPr'%ns['w'])
 if rpr is None:
  rpr=etree.Element('{%s}rPr'%ns['w']);run.insert(0,rpr)
 sz=rpr.find('{%s}sz'%ns['w'])
 old_val=sz.get('{%s}val'%ns['w']) if sz is not None else None
 if old_val not in ('20',None):raise SystemExit('UNEXPECTED_XML_FONT_SIZE:'+str(old_val))
 if sz is None:
  sz=etree.SubElement(rpr,'{%s}sz'%ns['w'])
 sz.set('{%s}val'%ns['w'],'21')
 entries['word/document.xml']=etree.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True)
with zipfile.ZipFile(dst,'w') as zout:
 for info in infos:zout.writestr(info,entries[info.filename])
with zipfile.ZipFile(src,'r') as a,zipfile.ZipFile(dst,'r') as b:
 if a.testzip() is not None or b.testzip() is not None:raise SystemExit('ZIP_CRC_FAIL')
 if a.namelist()!=b.namelist():raise SystemExit('ZIP_ENTRY_LIST_CHANGED')
 changed=[n for n in a.namelist() if a.read(n)!=b.read(n)]
 if changed!=['word/document.xml']:raise SystemExit('UNEXPECTED_PACKAGE_PARTS_CHANGED:'+','.join(changed))
new=Document(str(dst));newrun=new.paragraphs[63].runs[0];new_pt=newrun.font.size.pt if newrun.font.size is not None else None
if new_pt!=10.5 or newrun.text!=text:raise SystemExit('POST_PATCH_ASSERTION_FAILED')
res={'schema':'TTQS_ONE_0119_TARGETED_FONT_PATCH_R01','doc_id':'0119','requirement_id':'I19-E07','source_candidate_revision':'R03_WRITER_CANDIDATE_R01','source_candidate_sha256':source_sha,'new_candidate_revision':'R04_TARGETED_FONT_FIX_R02','new_candidate_path':str(dst),'new_candidate_sha256':sha(dst.read_bytes()),'new_candidate_bytes':dst.stat().st_size,'source_packet_sha256':'935a5cbc91c7dadfdcd0d9f3a7a45f56dcc5013ee9dd81fef1cfedd6b5b73656','unique_blueprint_sha256':'d097c7491235e323ebca4a31ff147ee5be2e6e5f893041087fd3076664aa0473','locator':{'top_level_paragraph_index':63,'run_index':0,'text_sha256':text_sha},'change':{'font_pt_before':old_pt,'font_pt_after':new_pt,'word_xml_old_w_sz':old_val,'word_xml_new_w_sz':'21'},'document_text_unchanged':True,'zip_entry_list_unchanged':True,'only_package_part_changed':changed,'zip_crc':'PASS','python_docx_reopen':'PASS','semantic_status':'PENDING_INDEPENDENT_EXACT_SHA_REVIEW','visual_status':'PENDING_WORD_NATIVE_RENDER_AND_VISUAL_QA','formal_acceptance':'NOT_GRANTED','real_evidence':'UNKNOWN/HOLD'}
receipt.write_text(json.dumps(res,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({k:v for k,v in res.items() if k not in ('new_candidate_path',)},ensure_ascii=True,indent=2))