from pathlib import Path
import csv, hashlib, json
from datetime import datetime
from docx import Document
root=Path(__file__).resolve().parents[2]
doc=root/'candidate'/'0051_R01.docx'
reg=root/'evidence'/'SYNTHETIC_REPLACEMENT_REGISTER_R01.csv'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
rows=list(csv.DictReader(reg.open(encoding='utf-8-sig',newline='')))
ids=[r['synthetic_fact_id'] for r in rows]
expected=[f'F{i:03d}' for i in range(1,228)]
assert len(rows)==227 and ids==expected and len(set(ids))==227
assert all(all((r.get(k) or '').strip() for k in ['document_location','real_replacement_source_required','replacement_responsible_role','status']) for r in rows)
assert all(r['status']=='SAMPLE_ONLY_NOT_VERIFIED_AS_REAL' for r in rows)
byid={r['synthetic_fact_id']:r for r in rows}
loc='§3表6、§5表7及表9–10、§6表12'
for fid in ['F024','F025','F026','F027']:
    assert byid[fid]['document_location']==loc, (fid,byid[fid]['document_location'])
d=Document(doc)
texts=[p.text for p in d.paragraphs]
for t in d.tables:
    texts.extend(c.text for row in t.rows for c in row.cells)
body='\n'.join(texts)
assert '\u8868\u0031\u0032' in body
for target in ['G1','G2','G3','G4']:
    assert target in body, target
out={
 'result':'PASS',
 'check_scope':'Post-review register-only locator correction; no semantic re-review claimed.',
 'candidate_docx_sha256':sha(doc),
 'candidate_docx_unchanged':sha(doc)=='6f575ebb41c6c26b4d6b36262fb4d9a5abdae53cbdc1ff4e75acda46d5f0b495',
 'current_register_sha256':sha(reg),
 'row_count':len(rows),'unique_ids':len(set(ids)),'expected_id_range':'F001-F227',
 'all_required_fields_nonblank':True,'all_status_sample_only_unverified':True,
 'corrected_fact_ids':['F024','F025','F026','F027'],
 'corrected_locator':loc,
 'candidate_body_contains_table_12_reference':True,
 'candidate_body_contains_all_four_objective_codes':True,
 'reviewer_scope_note':'OpenCode v2 semantic review passed exact DOCX SHA and reviewed preceding register SHA cddc262b205850d9bbdc9297c316d901783b56b66d03d05e3e7af33889d94266; this receipt verifies the later register-only correction.'
}
(root/'evidence'/'qa'/'REGISTER_LOCATOR_QA_R01.json').write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(out,ensure_ascii=False,indent=2))
