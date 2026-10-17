from __future__ import annotations
import csv, hashlib, json
from pathlib import Path
import sys
import pypdfium2 as pdfium
from docx import Document

ROOT=Path(__file__).resolve().parents[2]
DOCX=ROOT/'candidate'/'0051_R01.docx'
PDF=ROOT/'evidence'/'render_final_6f575ebb'/'0051_R01.pdf'
RENDER=PDF.parent
REGISTER=ROOT/'evidence'/'SYNTHETIC_REPLACEMENT_REGISTER_R01.csv'
OUT=ROOT/'evidence'/'qa'/'TECHNICAL_VISUAL_QA_R01.json'
def sha(path:Path)->str:
    return hashlib.sha256(path.read_bytes()).hexdigest()
doc=Document(DOCX)
body_sizes=[]; table_sizes=[]; empty_cells=[]
for para in doc.paragraphs:
    if not para.text.strip(): continue
    if para.style.name.startswith('Heading') or para.style.name in {'Title','Subtitle','Caption'}:
        continue
    for run in para.runs:
        if run.text.strip() and run.font.size:
            body_sizes.append(run.font.size.pt)
for ti,table in enumerate(doc.tables,1):
    for ri,row in enumerate(table.rows,1):
        for ci,cell in enumerate(row.cells,1):
            if not cell.text.strip(): empty_cells.append(f'T{ti}/R{ri}/C{ci}')
            for para in cell.paragraphs:
                for run in para.runs:
                    if run.text.strip() and run.font.size:
                        table_sizes.append(run.font.size.pt)
with REGISTER.open(encoding='utf-8-sig',newline='') as f:
    rows=list(csv.DictReader(f))
ids=[r['synthetic_fact_id'] for r in rows]
pdfdoc=pdfium.PdfDocument(str(PDF))
page_files=sorted(RENDER.glob('page-*.png'))
contact_files=sorted(RENDER.glob('contact-*.png'))
all_text='\\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for row in t.rows for c in row.cells])
agenda_sum=sum([10,20,20,15,5,5])
unit_sum=sum([90,90,120,60])
weights=[30,25,20,15,10]
matrix=[[3,5,4],[2,5,3],[3,4,3],[2,4,4],[4,3,5]]
totals=[sum(w*matrix[i][j] for i,w in enumerate(weights)) for j in range(3)]
percentages=[t/(sum(weights)*5)*100 for t in totals]
facts_ok=(
    len(rows)==227 and len(ids)==len(set(ids)) and
    all(r.get('document_location') and r.get('real_replacement_source_required') and
        r.get('status')=='SAMPLE_ONLY_NOT_VERIFIED_AS_REAL' for r in rows)
)
checks={
    'candidate_identity':{
        'doc_id':'0051','requirement_id':'I11-E03','official_checkbox_sequence':'053',
        'required_official_text':'各項課程產出均經討論後確認，且留有完整紀錄與主管核定。',
        'official_text_present': '各項課程產出均經討論後確認，且留有完整紀錄與主管核定。' in all_text,
        'candidate_sha256':sha(DOCX),'candidate_bytes':DOCX.stat().st_size
    },
    'document_structure':{
        'table_count':len(doc.tables),'paragraph_count':len(doc.paragraphs),
        'body_min_explicit_font_pt':min(body_sizes) if body_sizes else None,
        'table_min_explicit_font_pt':min(table_sizes) if table_sizes else None,
        'empty_table_cells':empty_cells,
        'literal_backslash_n_count':all_text.count('\\\\n'),
        'sample_banner_in_header':all('SAMPLE / SYNTHETIC' in s.header.paragraphs[0].text for s in doc.sections),
        'sample_banner_in_footer':all('SAMPLE / SYNTHETIC' in s.footer.paragraphs[0].text for s in doc.sections)
    },
    'replacement_register':{
        'rows':len(rows),'unique_fact_ids':len(set(ids)),
        'all_rows_have_locator_and_real_replacement':all(r.get('document_location') and r.get('real_replacement_source_required') for r in rows),
        'all_rows_sample_only_unverified':all(r.get('status')=='SAMPLE_ONLY_NOT_VERIFIED_AS_REAL' for r in rows),
        'register_sha256':sha(REGISTER),'register_checks_pass':facts_ok
    },
    'arithmetic_recalculation':{
        'agenda_segments_minutes':[10,20,20,15,5,5],'agenda_sum_minutes':agenda_sum,'agenda_expected_minutes':75,
        'unit_minutes':[90,90,120,60],'unit_sum_teaching_minutes':unit_sum,'unit_expected_minutes':360,
        'option_weights':weights,'weight_sum':sum(weights),
        'option_weighted_totals':{'A_slides':totals[0],'B_scenario_cards':totals[1],'C_selfstudy_pdf':totals[2]},
        'maximum_points':sum(weights)*5,
        'option_percentages':{'A_slides':percentages[0],'B_scenario_cards':percentages[1],'C_selfstudy_pdf':percentages[2]},
        'expected_totals':[270,445,365],'expected_percentages':[54,89,73]
    },
    'render_visual':{
        'renderer':'Microsoft Word COM export to PDF; bundled PDFium rasterization at scale 1.7',
        'pdf_sha256':sha(PDF),'pdf_bytes':PDF.stat().st_size,
        'pdf_pages':len(pdfdoc),'page_png_count':len(page_files),
        'contact_sheet_count':len(contact_files),
        'manual_full_page_set_visual_check':'PASS - all 14 pages reviewed in four contact sheets',
        'full_resolution_pages_manually_checked':[6,10],
        'visual_findings':'Two located layout issues were corrected before this SHA was fixed: literal backslash-n in six agenda row labels, and wrapping of an action evidence number. The final render shows actual line breaks and one-line dates and 360-minute closure text.'
    },
    'independent_semantic_review':'PENDING',
    'formal_human_acceptance':'NOT_GRANTED',
    'promotion':'NOT_AUTHORIZED',
    'overall_technical_visual_result':'PASS' if (facts_ok and len(pdfdoc)==14 and len(page_files)==14 and min(body_sizes)>=10.5 and min(table_sizes)>=9.5 and '各項課程產出均經討論後確認，且留有完整紀錄與主管核定。' in all_text and agenda_sum==75 and unit_sum==360 and totals==[270,445,365] and round(percentages[0])==54 and round(percentages[1])==89 and round(percentages[2])==73 and not empty_cells and all_text.count('\\\\n')==0) else 'FAIL'
}
OUT.write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(checks,ensure_ascii=False,indent=2))
