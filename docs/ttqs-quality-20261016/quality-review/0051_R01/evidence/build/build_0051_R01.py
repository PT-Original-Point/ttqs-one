from __future__ import annotations
import csv, hashlib, json, sys
from datetime import datetime
from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'candidate'/'0051_R01.docx'
REG=ROOT/'evidence'/'SYNTHETIC_REPLACEMENT_REGISTER_R01.csv'
FONT='Microsoft JhengHei'
CASE='SYN-0033-PLAN-01'
COURSE='志工服務紀錄、交接與工作界線實務'
facts=[]
def fact(fid,value,kind,loc,source,replace_by,role='經查證之課程規劃責任角色'):
    facts.append({'synthetic_fact_id':fid,'synthetic_value':str(value),'fact_type':kind,
      'document_location':loc,'source_or_provenance':source,'real_replacement_source_required':replace_by,
      'replacement_responsible_role':role,'status':'SAMPLE_ONLY_NOT_VERIFIED_AS_REAL'})
def item(value,kind,loc,source,replace_by,role='經查證之課程規劃責任角色'):
    fact(f'F{len(facts)+1:03}',value,kind,loc,source,replace_by,role)

# Fact-level replacement register. Each invented date, time, role, score,
# decision, version, action and retention value is a separate row. Shared
# upstream facts remain synthetic and identify their exact candidate source.
invented='本件為SAMPLE/SYNTHETIC情境設定；不是原始點協會真實會議或營運事實'
meeting_source='本件合成會議設定；沒有真實通知、出席、錄音或原始紀錄'
meeting_replace='真實會議通知、簽到／系統出席紀錄及經確認的會議紀錄'
item(CASE,'案例識別','封面、§3','沿用0050 R01合成案例；0050 SHA-256 d0d37c2c69b161a2cdd37ee3097306722c4a65a1afd844ae92e1dff58b43381a', '正式課程與需求追溯碼')
item(COURSE,'課程名稱','封面、§3','沿用0050 R01合成課名；不得視為實際開課', '現行核定課程計畫與課程清冊')
item('SYN-MIN-I11E03-20260518-01','會議紀錄碼','封面、§3',invented,'協會正式編碼規則及受控紀錄')
item('2026-05-18','會議日期','§3表5、§8表20',meeting_source,meeting_replace)
item('09:00','會議開始時間','§3表5、表6',meeting_source,meeting_replace)
item('10:15','會議結束時間','§3表5、表6',meeting_source,meeting_replace)
item('75分鐘','會議總長','§3表5、表6；§4', '由六段合成議程加總；不是實際會議時長',meeting_replace)
item('SYN-ROOM-01','會議地點代碼','§3表5',invented,'真實場地或合法會議系統紀錄')
roles=[('SYN-R01','主持兼課程規劃角色','主持、確認範圍及統整決議'),('SYN-R02','授課者代表','檢視教案、教材、評量與可教性'),('SYN-R03','學員／志工代表','提出閱讀負擔及使用情境意見'),('SYN-R04','紀錄覆核角色','核版本、數值、附件位置與紀錄完整度'),('SYN-SUP-01','模擬主管角色','只示範決定紀錄欄位，不代表協會職權')]
for code,title,duty in roles:
    item(code,'合成角色代碼','§2表4、§3表5、§5–§8',invented,'真實姓名、職掌、合法參與身分、出席及授權證據')
    item(title,'合成角色職稱','§2表4、§3表5',invented,'協會正式組織、職掌及任務指派資料')
    item(duty,'合成角色責任','§2表4、§5–§8',invented,'真實分工及會議紀錄原件')

# Individual course facts inherited only from the synthetic same-case source.
source0050='沿用0050 R01合成課程材料；DOCX SHA-256 d0d37c2c69b161a2cdd37ee3097306722c4a65a1afd844ae92e1dff58b43381a'
for goal in ['G1 角色界線與適切交接','G2 需求確認與下一步','G3 最小必要紀錄','G4 整合交接演練']:
    item(goal,'合成單元目標','§3表6、§5表7及表9–10、§6表12',source0050,'正式核定課程計畫、職能需求及服務流程')
for unit,duration in enumerate([90,90,120,60],1):
    item(f'第{unit}單元 {duration} 分鐘','合成單元時數','§3表6、§5表7、§5表10', '沿用0049/0050同案合成課表；不可當成實際班次', '正式核定課表、課程計畫及分段時間表')
item('360教學分鐘','四單元時數合計','§3表6、§4表10','90+90+120+60的合成課表計算','正式核定課表及獨立重算')
item('12人','規劃學員數','§3表6','沿用0050 R01示範規劃值；不是實際班級或出席人數','核定班級計畫；實際人數須另有合法名冊')
item('5張情境卡','合成教材數量','§5表9、§7表18','沿用0050 R01示範教案；不是正式教材清單','核准教材清冊及實際版次')
item('5欄交接練習','合成表單欄數','§5表9、§6表16','沿用0050 R01示範教案；欄位為合成設計','正式核准表單及資料最小化規則')
for label,value in [('前測題數','10題'),('後測題數','10題'),('每題配分','1分'),('示範參考門檻','8/10')]:
    item(value,'合成評量數值','§5表10、§6表15、§7表18、§11表27',source0050,'核准評量工具、計分規則、門檻依據及適用決議')

# Meeting agenda: each time boundary, duration, topic and subject decision is traceable.
agenda=[
 ('09:00','09:10',10,'核對要求、版本及範圍','確認只審合成輸出；不把需求調查或教案說成真實活動'),
 ('09:10','09:30',20,'比較輸出組合及教學可用性','選短講解、單頁情境卡、雙人演練及授課者評分錨點'),
 ('09:30','09:50',20,'核對四單元輸出、時數和附件','保留G1–G4順序；評量嵌入單元，不重複計時'),
 ('09:50','10:05',15,'檢查演練、評量、個資及課室條件','使用虛構代碼；學員版與答案鍵分開'),
 ('10:05','10:10',5,'確認版本、責任、期限、覆核和保存欄','六輸出及六行動逐項留痕'),
 ('10:10','10:15',5,'示範主管決定及未決真實資料','只作合成附條件核准示範，真實權限及原件仍缺')]
for n,(start,end,mins,topic,outcome) in enumerate(agenda,1):
    loc='§3表6；§5表9–11'
    item(start,'合成議程開始時間',loc,meeting_source,meeting_replace)
    item(end,'合成議程結束時間',loc,meeting_source,meeting_replace)
    item(f'{mins}分鐘','合成議程段長',loc,'合成時段設定；不是實際會議紀錄',meeting_replace)
    item(topic,'合成議程議題',loc,meeting_source,meeting_replace)
    item(outcome,'合成討論結論／輸出',loc,meeting_source,meeting_replace)
item('10+20+20+15+5+5=75分鐘','議程加總','§3表6','由六段合成議程重算；無重疊假設也屬示範',meeting_replace)

# Option model: each weight, raw score, contribution, subtotal, denominator and ranking is a separate fact.
criteria=[('單元目標對齊',30,[3,5,4]),('現場可操作性',25,[2,5,3]),('志工易理解',20,[3,4,3]),('資料最小化與安全',15,[2,4,4]),('版本維護容易度',10,[4,3,5])]
options=['A 投影片','B 情境卡','C 自學PDF']
item('1–5分','評分量尺','§4表7–8',invented,'經核准的評分規則、定義及會議採用依據')
for crit_i,(criterion,weight,scores) in enumerate(criteria,1):
    item(criterion,'合成比較準則','§4表7',invented,'有權責者核准的方案比較準則及需求依據')
    item(f'{weight}%','該準則合成權重','§4表7',invented,'有權責者的真實決策依據及核准權重')
    for oi,(opt,score) in enumerate(zip(options,scores),1):
        item(f'{opt}：{score}分','合成方案原始分數','§4表7',invented,'原始討論材料、評分者身分、評分依據及複核紀錄')
        item(f'{opt}：{weight}×{score}={weight*score}點','合成加權貢獻','§4表7–8','按本例權重與分數計算；須可獨立重算','真實評分表與獨立計算核對')
for opt,total,pct,rank in [('A 投影片',270,54,3),('B 情境卡',445,89,1),('C 自學PDF',365,73,2)]:
    item(f'{opt}：{total}點','合成加權總分','§4表7–8','五項合成加權貢獻加總','真實評分表與獨立計算核對')
    item('500點','合成最高總分分母','§4表7–8','權重總和100×最高分5','經核准權重／量尺及重算')
    item(f'{opt}：{pct}%','合成百分比','§4表7–8',f'{total}/500×100%；需重算','原始評分表及獨立重算')
    item(f'{opt}：第{rank}名','合成排序','§4表7',invented,'經核實評分總分與核准決策規則')
item('30+25+20+15+10=100','合成權重總和','§4表7–8','五項明列權重相加；須重算','核准權重表與獨立重算')
item('B 情境卡方案','合成選擇','§4表8、§5表9',invented,'正式會議決議、選擇依據及有權核定者')
item('短講解＋單頁情境卡＋雙人演練＋授課者評分錨點','合成採用組合','§5表9',invented,'正式採用教材清單、會議決議與版本原件')
item('情境卡可能增加閱讀負擔','合成保留意見','§4表8、§5表9',meeting_source,'真實發言紀錄及處理結論')
item('短句、單頁一情境、試用時記理解困難','合成緩解決定','§4表8、§5表9',meeting_source,'核准教材修訂及實際試用紀錄')

# Six output records: individual codes, titles, version endpoints, modifications and actions.
outputs=[
('單元目標對照表','G1–G4增加教學、作品、判準和附件位置','SYN-R01','SYN-R02','SYN-R04'),
('四單元課程進度','列四單元時間並標明評量已含在單元、不重複計時','SYN-R01','SYN-R04','SYN-R04'),
('情境卡與學員練習','使用角色碼並加入未知來源停用提示','SYN-R02','SYN-R03','SYN-R04'),
('前後測與答案鍵','分離A/B學員卷與教師答案鍵並綁G1–G4','SYN-R02','SYN-R04','SYN-R04'),
('授課者指引與學員材料','拆教師版／學員版，補活動及補練步驟','SYN-R02','SYN-R03','SYN-R04'),
('版本與保存索引','補來源SHA、核定原件、保存與調閱索引','SYN-R04','SYN-R01','SYN-R01')]
for i,(title,change,owner,checker,reviewer) in enumerate(outputs,1):
    code=f'SYN-OUT-{i:02}'
    item(code,'合成輸出識別碼','§6表12–17、§7表19',invented,'正式課程輸出清冊與受控編碼')
    item(title,'合成輸出名稱',f'§6表{11+i}',invented,'現行課程文件清冊與受控標題')
    item(f'{code} v0.1','合成原版次','§6表12–17、§7表19',invented,'實際輸入文件、版次及建立紀錄')
    item(f'{code} v0.2','合成修訂版次','§6表12–17、§7表19',invented,'實際修訂文件及版次控制紀錄')
    item(change,'合成修訂內容',f'§6表{11+i}、§7表18–19',invented,'實際差異檔、修改理由及會議核定紀錄')
    item(f'ACT-{i:02}','合成行動識別碼','§7表18',invented,'正式行動追蹤編碼')
    item(owner,'合成責任角色','§6表12–17、§7表18',invented,'真實指派者與任務授權')
    item(checker,'合成交叉核對角色','§6表12–17、§7表18',invented,'實際覆核人及覆核軌跡')
    item(reviewer,'合成覆核角色','§6表12–17、§7表18',invented,'真實獨立覆核人及覆核紀錄')
    item('2026-05-22','合成行動期限','§6表12–17、§7表18',invented,'真實決議期限及通知紀錄')
    item(f'ACT-{i:02} 示範完成','合成行動狀態','§7表18',invented,'實際修訂版本、完成證據與覆核結果')

# Approval, retention, privacy boundaries, and worked rejection cases.
item('SYN-APP-053-01','合成決議識別碼','§5表11、§8表20',invented,'正式決議編碼與受控紀錄')
item('SAMPLE_SIMULATED_APPROVAL_WITH_REVISIONS','合成模擬決定狀態','§5表11、§8表20',invented,'真實有權主管的核定決議、範圍、條件與原件')
item('附條件核准示範輸出包','合成決定內容','§5表11、§8表20',invented,'正式核定內容與完成條件')
item('SYN-SUP-01','合成決定者角色碼','§5表11、§8表20',invented,'核定者身分、職權、授權及原件')
item('未提供簽章或印章','合成簽核狀態','封面、§8表20',invented,'合法真實簽核／正式系統紀錄；不得仿造')
item('不得發給學員或當成協會核准教材','合成使用限制','§8表20、§10表24',invented,'正式核准版、生效條件與發布證據')
item('2026-06-01','合成課程結束日','§9表22',invented,'正式核定課表或結訓紀錄')
item('SYN-POL-RET-01','合成保存政策代碼','§9表22',invented,'協會現行核准保存政策及版次')
item('3年','合成保存期間','§9表22',invented,'現行核准政策及適用範圍')
item('2029-06-01','合成期限算例','§9表22','2026-06-01加3年；示例算法，非法律或協會政策','核准保存規則、正式起算日與獨立重算')
item('SYNTHETIC_ROOT/TTQS/SYN-0033-PLAN-01','合成檔案路徑','§9表22',invented,'有權限的正式文件庫路徑、權限及版本記錄')
item('沒有真人姓名、聯絡方式或真實學習結果','資料界線聲明','封面、§2–§11','本候選僅使用角色碼與示範結果','實際資料授權、必要性、最小化及保存依據')
for case,inp,verdict in [
 ('主管授權缺失','只有SYN-SUP-01、沒有正式章程或授權原件','HOLD，不得改成REAL核定'),
 ('時數不一致','合成課表354分鐘、核定計畫360分鐘','REJECT，停止受影響版'),
 ('附件無定位','只稱教材已確認，未列名稱及版次','PARTIAL，補件前不得關閉'),
 ('答案鍵外流','學員版附教師答案','HOLD，分版並查傳閱範圍'),
 ('示例誤作真實','合成日期或角色碼進正式紀錄','更正並保留稽核軌跡')]:
    item(case,'合成退件測例名稱','§11表27',invented,'實際缺失紀錄、受影響版次及處置證據')
    item(inp,'合成退件測例輸入','§11表27',invented,'可重現的原始文件、版本和具體欄位')
    item(verdict,'合成退件預期處置','§11表27',invented,'有權責者批准的處置規則與結案證據')

for i in range(1,7):
    item(f'SYN-OUT-{i:02}→v0.2','合成版本轉換','§7表19',invented,'真實差異檔、理由、修訂者、覆核及核定原件')

with REG.open('w',newline='',encoding='utf-8-sig') as f:
    w=csv.DictWriter(f,fieldnames=list(facts[0]));w.writeheader();w.writerows(facts)
if '--register-only' in sys.argv:
    metadata_path=ROOT/'evidence'/'build'/'BUILD_METADATA_R01.json'
    metadata=json.loads(metadata_path.read_text(encoding='utf-8'))
    metadata['synthetic_fact_register_rows']=len(facts)
    metadata['synthetic_replacement_register_sha256']=hashlib.sha256(REG.read_bytes()).hexdigest()
    metadata['register_updated_at']=datetime.now().astimezone().isoformat()
    metadata_path.write_text(json.dumps(metadata,ensure_ascii=False,indent=2),encoding='utf-8')
    print(f'REGISTER_ONLY rows={len(facts)} sha256={metadata["synthetic_replacement_register_sha256"]}')
    raise SystemExit(0)

D=Document(); sec=D.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.68);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.72);sec.right_margin=Inches(.72)
sec.header_distance=Inches(.28);sec.footer_distance=Inches(.28)
BLACK=RGBColor(0,0,0);FONT_FILL='E7E7E7'
for name in ('Normal','Body Text','List Bullet','List Number'):
    s=D.styles[name];s.font.name=FONT;s.font.size=Pt(10.5);s.font.color.rgb=BLACK
    s._element.rPr.rFonts.set(qn('w:eastAsia'),FONT);s.paragraph_format.space_after=Pt(4);s.paragraph_format.line_spacing=1.12
for name,size in [('Title',23),('Subtitle',13),('Heading 1',15),('Heading 2',12),('Heading 3',10.5),('Caption',9.5)]:
    s=D.styles[name];s.font.name=FONT;s.font.size=Pt(size);s.font.bold=name not in ('Subtitle','Caption');s.font.color.rgb=BLACK
    s._element.rPr.rFonts.set(qn('w:eastAsia'),FONT);s.paragraph_format.space_before=Pt(7);s.paragraph_format.space_after=Pt(4);s.paragraph_format.keep_with_next=True
def run(r,size=10.5,bold=False):
    r.font.name=FONT;r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=BLACK;r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT)
header=sec.header.paragraphs[0];header.alignment=WD_ALIGN_PARAGRAPH.RIGHT;run(header.add_run('TTQS ONE  |  候選文件 0051  |  SAMPLE / SYNTHETIC'),8.5)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
run(footer.add_run('SAMPLE / SYNTHETIC  |  非真實會議或核定  |  0051  |  第 '),8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld);run(footer.add_run(' 頁'),8)
def p(text='',style=None,bold=False):
    x=D.add_paragraph(style=style);x.paragraph_format.widow_control=True;run(x.add_run(text),10.5,bold);return x
def h(text,n=1):
    x=D.add_paragraph(style=f'Heading {n}');run(x.add_run(text),15 if n==1 else 12,True);return x
def cap(text):
    x=D.add_paragraph(style='Caption');x.paragraph_format.keep_with_next=True;run(x.add_run(text),9.5,True)
def shade(cell):
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),FONT_FILL);cell._tc.get_or_add_tcPr().append(sh)
def table(headers,rows,widths=None):
    t=D.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=widths or [Inches(6.83/len(headers))]*len(headers)
    for i,w in enumerate(widths):t.columns[i].width=w
    for i,z in enumerate(headers):
        c=t.rows[0].cells[i];c.width=widths[i];c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;shade(c)
        r=c.paragraphs[0].add_run(str(z));run(r,9.5,True)
    repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for i,z in enumerate(row):
            c=cells[i];c.width=widths[i];c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            mar=OxmlElement('w:tcMar')
            for side,val in [('top','70'),('start','85'),('bottom','70'),('end','85')]:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),val);el.set(qn('w:type'),'dxa');mar.append(el)
            c._tc.get_or_add_tcPr().append(mar)
            r=c.paragraphs[0].add_run(str(z));c.paragraphs[0].paragraph_format.space_after=Pt(0);c.paragraphs[0].paragraph_format.line_spacing=1.02;run(r,9.5)
    D.add_paragraph().paragraph_format.space_after=Pt(0)
    return t
def kv(title,rows):
    cap(title);table(['記錄欄位','SAMPLE / SYNTHETIC 完整示例'],rows,[Inches(1.45),Inches(5.38)])
def page():D.add_page_break()

# Cover
x=D.add_paragraph();x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.space_before=Pt(55);run(x.add_run('SAMPLE / SYNTHETIC'),18,True)
x=D.add_paragraph(style='Title');x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.space_before=Pt(18);run(x.add_run('課程產出審議紀錄與核定留存示範'),23,True)
x=D.add_paragraph(style='Subtitle');x.alignment=WD_ALIGN_PARAGRAPH.CENTER;run(x.add_run('課程輸出逐項討論、修訂、決議與保存的操作紀錄'),13)
table(['文件識別','對照資料'],[
('DOCX 0051 R01','Unique Blueprint I11-E03；官方方框053'),
('官方佐證原文','各項課程產出均經討論後確認，且留有完整紀錄與主管核定。'),
('文件文類','專屬佐證紀錄；以完整合成會議紀錄示範實際操作'),
('合成案例',f'{CASE}；課名承接0049 R01／0050 R01的合成案例'),
('候選狀態','CANDIDATE ONLY；沒有協會真實會議、簽核或正式接受'),
('使用界線','本件所有日期、角色、討論、評分、決議、核准及保存規則均為合成')],[Inches(1.4),Inches(5.43)])
p('本例中的主管決定是明確標示的模擬情境，不代表原始點協會曾召開會議、指派核定者或核准教材。文件沒有仿造真人簽名、職章或會議簽到。')
page()

h('1  30秒使用說明')
p('用途：在課程教材、教案、評量工具及其他課程產出送主管核定前，逐項記錄討論依據、不同意見、修改內容、責任期限、覆核結果、核定狀態及保存位置。適用者為課程規劃人員、授課者、參與討論者、紀錄者及有權核定的主管。')
p('使用時點：課程輸出形成初稿後、實際使用前；課程目標、教材、評量、時數或適用條件有重大改動時，另開新版本紀錄。會議完成不等於核定完成，仍須保存可驗證的主管權限與決定原件。')
cap('表1  官方要求與本件證明範圍')
table(['項目','本件界定'],[
('DOCX／要求','0051對應I11-E03、官方方框053。官方原文為「各項課程產出均經討論後確認，且留有完整紀錄與主管核定。」'),
('本件直接證據','記錄課程輸出如何提出、討論、比較、修改、確認、核定及保存。'),
('不由本件代替','0050／I11-E02證明單元內容與目標的教學對照；0052／I11-E04證明學員意見如何回饋課程規劃。'),
('目前真實狀態','未提供0051真實會議紀錄或主管核定原件；本候選只展示完整做法。'),
('讀者30秒內應取得','誰在何時審查哪些輸出、討論依據與取捨、修改內容、責任、核定者及真實原件位置。')],[Inches(1.4),Inches(5.43)])
cap('表2  與最近鄰要求的排除檢查')
table(['文件與要求','本件不代替的工作','正文可查的界線'],[
('0049／I11-E01／方框051','需求如何形成整體課程計畫。','只引用同一合成計畫，不重做需求調查。'),
('0050／I11-E02／方框052','單元目標、教學活動、練習和評量是否一致。','把0050當會議審查材料，焦點是討論、修訂、決議和核定。'),
('0052／I11-E04／方框054','學員與結訓學員意見如何影響課程。','不聲稱收集真實學員回饋；回饋議題依0052另記。')],[Inches(1.4),Inches(2.65),Inches(2.78)])

page();h('2  紀錄規則與責任分工')
p('會議紀錄須把輸入證據、議題、分歧、決議、修訂、覆核、主管核定和保存位置連成可追溯順序。只寫「討論通過」或只附最後教材，無法證明各項輸出曾經討論及核定。')
cap('表3  欄位定義與填寫控制')
table(['欄位','填寫與驗證規則','本例做法'],[
('輸出識別','每項輸出有唯一代碼、名稱、版本及頁節／附件定位。','用SYN-OUT-01至06；只屬合成情境。'),
('輸入依據','記來源、版本、日期和位置；未知來源先標缺證。','列0049、0050候選SHA，不當成REAL。'),
('討論紀錄','留議題、選項、支持與疑慮、不同意見及處理。','按時間記角色碼、理由、決議與風險。'),
('修改軌跡','列舊版、新版、改動、理由、責任、期限、覆核。','以SYN-OUT版本碼示範；未改動0050 DOCX。'),
('核定控制','核對核定者身分、職掌、授權、日期、範圍及原件。','SYN-SUP-01只模擬流程；真實權限缺證。'),
('保存與調閱','記受控位置、版本、存取角色、保存政策和調閱。','合成路徑和期限不代表協會政策。')],[Inches(1.05),Inches(3.1),Inches(2.68)])
cap('表4  角色、責任及真實替換條件')
table(['示範角色','本例責任','轉為REAL前必須核實'],[
('SYN-R01','主持討論、確認輸出、控制議程。','姓名、正式職掌、主持權限及出席證據。'),
('SYN-R02','授課者代表；檢視教案、教材和評量。','真實師資資格、任務關係及實際意見。'),
('SYN-R03','學員／志工代表；指出文字和表單負擔。','參與身分、合法同意及意見原件。'),
('SYN-R04','紀錄覆核；核對版次、數值、附件位置。','覆核者身分及覆核軌跡。'),
('SYN-SUP-01','示範主管決定的紀錄位置。','章程、授權、真實核定者和決定原件。')],[Inches(1.05),Inches(2.65),Inches(3.13)])
p('控制方法：紀錄者在會中同步記錄，會後逐項比對附件版次與決議；若口頭結論和文件不一致，先標未核定，交有權者確認。這是建議做法，不表示協會已正式採納。')

page();h('3  完整合成會議基本資料')
p('下列紀錄只示範如何填寫I11-E03。日期、出席、角色、發言、議程及任何「同意」都屬合成設定，並非原始點協會會議紀錄。')
cap('表5  SAMPLE / SYNTHETIC會議識別')
table(['欄位','合成紀錄值'],[
('紀錄碼','SYN-MIN-I11E03-20260518-01'),('日期與時間','2026-05-18，09:00–10:15，共75分鐘。'),
('合成地點','SYN-ROOM-01線上會議室；沒有真實網址或登入紀錄。'),
('合成案件',f'{CASE}；課程「{COURSE}」。'),
('審議目的','在示範課程輸出進入使用前，檢查來源、討論、修訂、使用條件及核定狀態。'),
('輸入材料','0049 R01合成課程計畫，SHA-256 e8212788b1f7b201ca4af4f4c5fb83ab131627dbfa815987f7cf6344edfda727；0050 R01合成單元教學設計，SHA-256 d0d37c2c69b161a2cdd37ee3097306722c4a65a1afd844ae92e1dff58b43381a。'),
('上游需求線索','0033 R03合成需求調查候選SHA-256 fe5f3cfd84c761da86435a71e0fd899daa725be78628c4d2d3f2324c1fdb5f4f；不把合成回覆寫成會員意見。'),
('紀錄者與狀態','SYN-R04虛構角色碼；沒有真實邀請、簽到、錄音或主管簽章。')],[Inches(1.3),Inches(5.53)])
cap('表6  議程與時間核對')
table(['時間','分鐘','議題','輸入與預期輸出'],[
('09:00–09:10','10','核對要求、版本及範圍。','官方方框053、0049及0050；確認只審合成輸出。'),
('09:10–09:30','20','比較輸出組合及教學可用性。','方案A/B/C評分；保留意見和取捨。'),
('09:30–09:50','20','核對四單元輸出、時數和附件。','G1–G4、90+90+120+60分鐘；修訂清單。'),
('09:50–10:05','15','檢查演練、評量、個資及課室條件。','情境卡、前後測、答案鍵；停止／補正條件。'),
('10:05–10:10','5','確認版本、責任、期限、覆核和保存欄。','SYN-OUT-01至06、ACT-01至06。'),
('10:10–10:15','5','示範主管決定及未決真實資料。','模擬決定；不構成協會核定。'),
('合計','75','10+20+20+15+5+5=75分鐘。','時段首尾相接，無重疊。')],[Inches(1.05),Inches(.55),Inches(2.1),Inches(3.13)])

page();h('4  課程輸出方案比較與計算')
p('本例比較三種可製作的輸出包，再按五個條件加權。權重及分數是合成假設，不是學員調查、真實會議投票或協會既定標準。')
cap('表7  合成評分權重與方案評分')
table(['評估條件','權重','A 投影片','B 情境卡','C 自學PDF'],[
('單元目標對齊','30','3分 / 90','5分 / 150','4分 / 120'),
('現場可操作性','25','2分 / 50','5分 / 125','3分 / 75'),
('志工易理解','20','3分 / 60','4分 / 80','3分 / 60'),
('資料最小化與安全','15','2分 / 30','4分 / 60','4分 / 60'),
('版本維護容易度','10','4分 / 40','3分 / 30','5分 / 50'),
('加權點數合計','100','270 / 500','445 / 500','365 / 500'),
('換算分數','—','54%','89%','73%'),('排序','—','第3','第1','第2')],[Inches(1.62),Inches(.55),Inches(1.45),Inches(1.58),Inches(1.63)])
cap('表8  計分步驟與決策限制')
table(['計算或判讀','示範結果'],[
('權重總和','30+25+20+15+10=100。'),
('方案B加權點數','30×5 + 25×5 + 20×4 + 15×4 + 10×3 = 150+125+80+60+30 = 445。'),
('換算','445÷(100×5)×100%=89%；A=270÷500=54%；C=365÷500=73%。'),
('選擇結果','選B作為示範方案，理由是演練、可觀察作品和回饋能對應G1–G4。'),
('保留意見','SYN-R03認為情境卡可能增加閱讀負擔；決議改短句、單頁一情境，試用時記理解困難。'),
('限制','權重不是經驗證的偏好調查；分數不能代表因果效果、課程成效或真實核定。')],[Inches(1.45),Inches(5.38)])
p('退件條件：計算不符、權重不是100、評分依據不明、評分者不明，或未說明不同意見如何處理時，方案比較不得標成完成。')

page();h('5  逐項議程紀錄與討論結果')
p('逐項保留時間、依據、不同意見、決定和責任。實際紀錄須依授權保留真實出席資料；本例只使用合成角色碼。')
cap('表9  議題一與二')
table(['議題與時間','依據及討論','合成決定、保留意見與後續'],[
('一｜09:00–09:10\n確認來源與範圍','SYN-R01核對方框053；SYN-R04核對0049與0050案例碼均為SYN-0033-PLAN-01。各角色確認材料是SAMPLE候選，沒有真實課程或會議原件。','只審示範輸出是否形成完整紀錄；不把需求調查或單元教案寫成真實活動。ACT-01建立來源與版本索引。若真實課名、服務界線或授權不同，停止套用本示例。'),
('二｜09:10–09:30\n比較輸出組合','SYN-R02認為只有投影片不易留下練習證據；SYN-R03希望表單簡短。依表7比較A/B/C，B為89%，A為54%，C為73%。','選「短講解＋單頁情境卡＋雙人演練＋授課者評分錨點」。保留閱讀負擔疑慮；ACT-02縮短文字並保留大字版。決定屬合成示範。')],[Inches(1.4),Inches(2.7),Inches(2.73)])
cap('表10  議題三與四')
table(['議題與時間','依據及討論','合成決定、保留意見與後續'],[
('三｜09:30–09:50\n單元輸出與時數','SYN-R02檢查G1–G4；SYN-R04核對0050表4、表5及單元1–4。四單元為90、90、120、60分鐘；評量嵌入單元，不另加一次。','保留四單元順序，90+90+120+60=360教學分鐘；休息與行政不計。每單元須有目標、教學、練習作品、判準及位置索引。ACT-03逐表核對。'),
('四｜09:50–10:05\n情境、評量與資料','SYN-R03指出卡片不得要求學員猜真人身分或聯絡資訊。SYN-R02提出前後測答案鍵。0050示範為A/B各10題、每題1分，後測參考點8/10。','教材只用虛構代碼；學員版與授課者答案鍵分開。8/10只屬示範門檻，不是資格或真人成果。ACT-04補隱私提示；若答案與正式程序衝突，停用題目並查證。')],[Inches(1.4),Inches(2.7),Inches(2.73)])
cap('表11  議題五與六')
table(['議題與時間','依據及討論','合成決定、保留意見與後續'],[
('五｜10:05–10:10\n版本與保存','SYN-R01要求各輸出有版號、理由、責任、期限及覆核。SYN-R04指出沒有政策原件就不能填成協會規定。','使用SYN-OUT版本碼和ACT-01至06。保存年限只用虛構規則示範；實際政策保持待查證。'),
('六｜10:10–10:15\n示範主管決定','SYN-SUP-01按表12核對六項輸出來源、修訂、覆核和未決風險。角色權限是虛構設定，沒有簽名或印章。','決議碼SYN-APP-053-01：合成情境內「附條件核准示範輸出包」。不是協會核定。真實授權、日期、決議原件和核定範圍均待替換。')],[Inches(1.4),Inches(2.7),Inches(2.73)])

page();h('6  六項課程輸出的審查與處置')
p('每項保留輸入、發現、討論、決定、修改、負責、期限、覆核和退件條件。所有代碼與修訂只屬合成會議。')
kv('表12  輸出一｜單元目標對照表',[
('材料','SYN-OUT-01 v0.1；參照0050 R01表4。'),
('討論','G1–G4已列出；SYN-R02要求標明練習產物，避免只靠口頭連結。'),
('決定與修改','保留目標文字，新增教學活動、可觀察產物、判準及附件位置；v0.1→v0.2。未改動0050 DOCX。'),
('責任期限','SYN-R01整理、SYN-R02逐單元核對；示範期限2026-05-22。'),
('覆核','SYN-R04逐項確認四單元都有可觀察證據和判準；合成情境內完成。'),
('退件條件','單元只列主題、沒有作品／判準，或附件版次不符即標未核定。')])
kv('表13  輸出二｜四單元課程進度',[
('材料','SYN-OUT-02 v0.1；同案計畫承接0049 R01。'),
('討論','測驗已排在單元一和單元四；在總表外重列會重複計時。'),
('決定與修改','G1 90、G2 90、G3 120、G4 60分鐘；評量嵌入單元；v0.1→v0.2增加分段總計。'),
('責任期限','SYN-R01重算、SYN-R04獨立重算；2026-05-22。'),
('覆核','90+90+120+60=360教學分鐘；休息和行政另列，不算教學分鐘。'),
('退件條件','分段不合、評量重算，或與正式課程計畫衝突，停用受影響版。')])

page();h('6  六項課程輸出的審查與處置（續）')
kv('表14  輸出三｜情境卡與學員練習',[
('材料','SYN-OUT-03 v0.1；五張合成情境卡與練習單。'),
('討論','SYN-R03要求任務卡不逼學員猜真人資料或未核實的服務窗口。'),
('決定與修改','五張獨立虛構情境，使用角色代碼；未知流程改寫為查正式來源／交有權角色；v0.1→v0.2。'),
('責任期限','SYN-R02修改、SYN-R03試讀、SYN-R04核對；2026-05-22。'),
('覆核','五張卡不含真人姓名、地址、電話或可回推身分的組合。'),
('退件條件','教導未查證外部流程、含真實個資或缺少教師判讀錨點。')])
kv('表15  輸出四｜前後測與答案鍵',[
('材料','SYN-OUT-04 v0.1；前測A、後測B及教師答案鍵。'),
('討論','兩卷需覆蓋同一目標而使用不同情境；答案鍵不應出現在學員版。'),
('決定與修改','各10題、每題1分、0–10分；8/10只作示範參考點；v0.1→v0.2分離版本並綁G1–G4。'),
('責任期限','SYN-R02核答案、SYN-R04核計分；2026-05-22。'),
('覆核','兩卷共20題，總分口徑一致；沒有真實作答或成效資料。'),
('退件條件','答案鍵不一致、兩卷目標不同，或把8/10寫成協會核定標準。')])

page();h('6  六項課程輸出的審查與處置（續）')
kv('表16  輸出五｜授課者指引與學員材料',[
('材料','SYN-OUT-05 v0.1；教師指引、五欄交接示例和課室清單。'),
('討論','只有學員練習單時，教師缺少一致回饋尺度；答案若混入學員版會影響獨立作答。'),
('決定與修改','新增教師版答案錨點、活動流程及補練步驟；教師版與學員版分開；v0.1→v0.2。'),
('責任期限','SYN-R02整理，SYN-R03試讀，SYN-R04核版次；2026-05-22。'),
('覆核','兩版有獨立代碼；教師答案不在學員版。'),
('退件條件','答案鍵混入學員版、指引缺時間或引用未授權流程。')])
kv('表17  輸出六｜版本與保存索引',[
('材料','SYN-OUT-06 v0.1；會議決議、修訂、核定和保存索引。'),
('討論','實際雲端路徑、保存年限或職稱可能被誤認為正式政策。'),
('決定與修改','只用SYNTHETIC_ROOT與SYN-POL-RET-01；列入正式政策、路徑、權限、核定原件替換欄；v0.1→v0.2。'),
('責任期限','SYN-R04整理、SYN-R01確認範圍；2026-05-22。'),
('覆核','只連結公開合成材料，沒有私密原件位置。'),
('退件條件','路徑無權限、保存期限無依據，或核定原件無法綁定版本。')])

page();h('7  修改、責任與覆核軌跡')
p('行動狀態只描述合成案例內的完成情況，不表示協會人員已執行或課程輸出已核定。')
cap('表18  SAMPLE / SYNTHETIC行動清冊')
table(['行動','輸出與具體修訂','責任與覆核','期限','合成關閉證據'],[
('ACT-01','SYN-OUT-01加G1–G4教學、作品、判準和位置。','SYN-R01整理；R02核對；R04覆核。','2026-05-22','四列均有目標、活動及評量位置。'),
('ACT-02','SYN-OUT-02列四單元時數及評量已含在單元。','R01重算；R04獨立重算。','2026-05-22','合計360教學分。'),
('ACT-03','SYN-OUT-03移除識別資料，加入未知來源停用提示。','R02修改；R03試讀；R04檢查。','2026-05-22','五張合成卡逐張核對。'),
('ACT-04','SYN-OUT-04分離A/B卷與答案鍵並綁目標。','R02核答案；R04核10題與門檻。','2026-05-22','A/B各10題、1分；門檻只是示範。'),
('ACT-05','SYN-OUT-05拆分教師版／學員版，補活動和補練步驟。','R02完成；R03試讀；R04核版本。','2026-05-22','學員版不含答案；版本可識別。'),
('ACT-06','SYN-OUT-06補來源SHA、核定、保存和調閱索引。','R04整理；R01確認範圍。','2026-05-22','索引完整；真實權限、政策、原件缺證。')],[Inches(.62),Inches(2.25),Inches(1.56),Inches(1.00),Inches(1.40)])
cap('表19  版本差異索引')
table(['輸出','示範版次','改動及理由','覆核'],[
('SYN-OUT-01','v0.1→v0.2','G1–G4增加教學、作品、判準和位置。','合成核對完成。'),
('SYN-OUT-02','v0.1→v0.2','列四單元時間及評量是否重複計時。','360分鐘重算。'),
('SYN-OUT-03','v0.1→v0.2','使用角色碼，加入未知來源停止條件。','五張卡核對。'),
('SYN-OUT-04','v0.1→v0.2','分離學員卷和答案鍵，綁G1–G4。','20題與計分核對。'),
('SYN-OUT-05','v0.1→v0.2','分教師版／學員版並新增補練步驟。','用途和版本核對。'),
('SYN-OUT-06','v0.1→v0.2','加入SHA、核定原件及保存索引。','真實欄位標待替換。')],[Inches(1.13),Inches(1.25),Inches(3.06),Inches(1.39)])

page();h('8  示範主管決定與停止條件')
p('本節核定欄只示範需要保存的資訊。它不能證明協會存在同名職務、已授權該角色、召開會議或核准教材。正式使用前須以章程、職掌及實際核定原件重新填寫。')
cap('表20  SAMPLE / SYNTHETIC示範核定紀錄')
table(['核定欄位','合成示範紀錄','轉為REAL前證據'],[
('核定範圍','SYN-OUT-01至06示範輸出包。','正式輸出清冊、版次及核定範圍。'),
('決定角色','SYN-SUP-01；虛構主管角色碼。','姓名、職務、授權來源及權限。'),
('決定時間','2026-05-18 10:10–10:15；合成。','真實會議或核簽時間。'),
('決定內容','SYN-APP-053-01：模擬「附條件核准示範輸出包」。','正式決議、條件、完成證據與核定原件。'),
('簽章及核准','未提供；不仿造簽名、印章或機構核准。','合法原始簽核或正式系統紀錄。'),
('實際使用狀態','不得發給學員或作為協會已核准教材。','核定後實際版次發布及生效條件。')],[Inches(1.2),Inches(2.88),Inches(2.75)])
cap('表21  退件、補正與停止使用條件')
table(['檢查情況','判定','處置','重新開放條件'],[
('只寫討論通過，無輸出、版次、依據。','退件','補議題、決議、附件位置及責任期限。','表格與來源可互相定位。'),
('核定者身分或權限不可驗證。','暫停核定','不得標主管核定；查正式授權來源。','權限與決議綁定版本。'),
('決議、課表、答案鍵或教材版本不一致。','停止受影響版','記差異並交責任角色核實，不猜測有效版。','權責確認唯一版並覆核。'),
('SAMPLE被當作真實出席、結果或核准。','撤下誤用內容','記誤用範圍並更正傳閱版本。','取得真實原件並留更正軌跡。'),
('時數或評分算式無法重算。','退回覆核','逐式重算，補位置和來源。','第二人重算留痕。'),
('含未授權個資或第三方私密原件。','停止公開及使用','限制存取並記影響範圍。','權責確認合法處理方式。')],[Inches(1.75),Inches(.88),Inches(2.45),Inches(1.75)])

page();h('9  保存、調閱與真實資料替換')
p('本例用虛構規則示範如何記期限，不宣稱法律或協會保存年限。正式紀錄須引用現行核准政策；沒有政策原件時應標待查證，不可沿用示例期限。')
cap('表22  合成保存與存取索引')
table(['欄位','合成示例','真實替換及控制'],[
('保存規則','SYN-POL-RET-01；假設課程結束後保存3年。','取得現行核准政策、版次及權責依據。'),
('課程結束日','2026-06-01；合成。','正式核定課表或結訓紀錄。'),
('期限算例','2026-06-01 + 3年 = 2029-06-01。','按核准政策重算並記覆核者。'),
('檔案位置','SYNTHETIC_ROOT/TTQS/SYN-0033-PLAN-01。','正式授權文件庫路徑及權限。'),
('調閱紀錄','示例沒有真人調閱；實用時記日期、角色、目的、版次及歸還／撤權。','核准的存取日誌。'),
('更正規則','留舊版、修訂理由、修訂者、覆核者、核定狀態和時間。','依協會核定的文件控制規則。')],[Inches(1.25),Inches(2.65),Inches(2.93)])
cap('表23  轉為REAL前的原件與替換清單')
table(['需替換項目','必須取得的真實來源','狀態與責任'],[
('會議身份與出席','通知、簽到或正式系統紀錄；姓名與角色依法核實。','未提供；授權紀錄者登錄，覆核者核來源。'),
('討論與異議','原始紀錄、合法錄音摘要或逐項討論材料。','未提供；紀錄者按原始決議定位。'),
('課程輸出版本','計畫、教案、教材、題卷、答案鍵及附件版次。','0049/0050是合成候選，非核定原件。'),
('主管核定','章程／授權、核定者、決議、日期及原件位置。','未提供；有權者依正式流程核定。'),
('行動完成證據','修訂文件、版次差異、責任者完成及覆核紀錄。','合成狀態不是實際完成；逐項取原件。'),
('保存與調閱','保存政策、權限清冊、實際庫位及調閱紀錄。','合成規則不可移植；由文件管理角色替換。')],[Inches(1.32),Inches(2.83),Inches(2.68)])

page();h('10  結案核對與跨文件追溯')
cap('表24  I11-E03結案核對表')
table(['核對項目','本例定位','合成判定','真實關閉證據'],[
('官方文字與識別','封面、表1；I11-E03／053。','PASS for sample mapping','核對現行方框和Blueprint。'),
('輸出代碼、版本與位置','§6、表12–17、表19。','PASS for sample completeness','附件可開啟且版次相符。'),
('議題、不同意見及取捨','§4–5、表7–11。','PASS for sample completeness','真實紀錄及來源原件。'),
('責任、期限和覆核','§7、表18–19。','PASS only in synthetic scenario','真實修訂完成與覆核軌跡。'),
('主管身分、權限及原件','§8、表20。','REAL evidence missing','正式授權及核定原件。'),
('時數與評分可重算','§4、表7–8；§6。','PASS for listed sample arithmetic','第二人按原始資料重算。'),
('SAMPLE／REAL分界','封面、頁尾、§3、§8–9、替換簿。','PASS for sample labeling','替換合成事實並留痕。'),
('最近鄰要求已排除','表2、表25。','PASS for sample scope','以最新Blueprint重新核對。')],[Inches(1.42),Inches(1.55),Inches(1.76),Inches(2.1)])
cap('表25  文件間來源與責任界線')
table(['文件','本案用途與可引用內容','本件不得宣稱'],[
('0033 R03','上游合成需求示例；SHA fe5f3cfd...。','不是真實會員或結訓學員意見。'),
('0049 R01／I11-E01／051','同案合成課程計畫；SHA e8212788...。','不是真實需求確認、班次或核准。'),
('0050 R01／I11-E02／052','四單元教案、G1–G4、90+90+120+60=360；SHA d0d37c2c...。','不等於本件會議紀錄或主管核定。'),
('0051 R01／I11-E03／053','本件記錄輸出討論、修訂、責任、模擬核定及保存。','不是協會真實會議、正式接受或核准。'),
('0052／I11-E04／054','鄰項處理學員回饋對課程的影響。','本件無真實回饋統計，不代替其流程。')],[Inches(1.68),Inches(2.78),Inches(2.37)])
p('跨文件核對結果：本件只承接0050 R01已載明的課名、案例碼、G1–G4、四單元時數、12人示範規劃值、五張情境卡、五欄練習、兩份10題測驗及8/10參考點。未借用真實班級、學員資料、會議出席、成績或主管核准。')

page();h('11  正式紀錄的操作步驟與退件測例')
cap('表26  可直接照做的會議紀錄流程')
table(['步驟','操作','通過條件','未通過處置'],[
('1 收件','列待審輸出、版號、來源和待決事項。','檔案可開啟，唯一版本可定位。','缺來源或版號不進核定。'),
('2 確認權限','確認參與者職責及主管核定依據。','姓名、角色、出席和授權可驗證。','權限未明不得寫主管核定。'),
('3 逐項討論','記議題、證據、方案、分歧、理由及決議。','讀者能理解採用或退回原因。','只寫已討論視為紀錄不足。'),
('4 指派修訂','定位修改檔案、版次、欄位，指定責任與期限。','修訂能檢查並有覆核人。','無責任、期限、版本差異則未結。'),
('5 覆核核定','核附件和決議一致，再由有權者作決定。','決定者、範圍、日期、條件、原件完整。','原件缺失只標待核定。'),
('6 保存調閱','依核准政策封存紀錄、附件、核定原件和歷史版。','路徑、權限、政策和日誌可追溯。','不套用本例虛構路徑或年限。')],[Inches(.55),Inches(2.2),Inches(2.23),Inches(1.85)])
cap('表27  退件測例')
table(['測例','輸入狀態','預期判定','應留下的紀錄'],[
('主管授權缺失','只有SYN-SUP-01，沒有正式章程或授權原件。','HOLD；不得改成REAL核定。','缺證、索取路徑、責任及再審條件。'),
('時數不一致','課表354分鐘，核定計畫360分鐘。','REJECT；停止受影響版。','差異算式、版次、修訂與覆核。'),
('附件無定位','只寫教材已確認，未列名稱和版次。','PARTIAL；補資料前不能關閉。','缺件、來源及補齊責任。'),
('答案鍵外流','學員版附教師答案。','HOLD；分版並查傳閱情況。','版本處置、影響範圍、重核紀錄。'),
('示例誤作真實','合成日期或角色碼進入正式紀錄。','更正並保留稽核軌跡。','誤用位置、撤回責任、真實來源。')],[Inches(1.18),Inches(2.03),Inches(1.38),Inches(2.24)])
p('本例完成界線：本文含完整合成討論、方案比較、逐項輸出審查、修訂行動、模擬主管決定及保存示例。REAL資料、真實會議原件、主管授權與簽核未提供。因此它是可操作的SAMPLE候選，不是協會實績或核准證據。')

D.core_properties.title='課程產出審議紀錄與核定留存示範'
D.core_properties.subject='DOCX 0051 / I11-E03 / official checkbox 053 / candidate sample'
D.core_properties.author='SAMPLE / SYNTHETIC'
D.core_properties.keywords='TTQS ONE, 0051, I11-E03, SAMPLE, SYNTHETIC, CANDIDATE ONLY'
OUT.parent.mkdir(parents=True,exist_ok=True);D.save(OUT)
sha=hashlib.sha256(OUT.read_bytes()).hexdigest()
(ROOT/'evidence'/'build'/'BUILD_METADATA_R01.json').write_text(json.dumps({
 'project_id':'TTQS_ONE','doc_id':'0051','requirement_id':'I11-E03','official_checkbox_sequence':'053','revision':'R01',
 'candidate_path':'candidate/0051_R01.docx','candidate_sha256':sha,'candidate_bytes':OUT.stat().st_size,
 'case_id':CASE,'sample_only':True,'real_meeting_or_approval_claimed':False,'synthetic_fact_register_rows':len(facts),
 'built_at':datetime.now().astimezone().isoformat(),'builder_script':'evidence/build/build_0051_R01.py'
},ensure_ascii=False,indent=2),encoding='utf-8')
print(f'BUILT {OUT}\nSHA256 {sha}\nBYTES {OUT.stat().st_size}\nREGISTER_ROWS {len(facts)}')
