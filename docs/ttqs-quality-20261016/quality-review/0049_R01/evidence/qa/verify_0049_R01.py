from __future__ import annotations

import csv
import hashlib
import json
import re
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from pypdf import PdfReader
from PIL import Image


ROOT = Path(r"E:\TTQS\_TTQS_ONE_REVIEW_SYNC_0036_R08_20261016")
PKG = ROOT / "docs/ttqs-quality-20261016/quality-review/0049_R01"
DOCX = PKG / "candidate/0049_R01.docx"
REGISTER = PKG / "evidence/SYNTHETIC_REPLACEMENT_REGISTER_R01.csv"
RENDER = PKG / "evidence/render_R01_e8212788"
PDF = RENDER / "0049_R01.pdf"
SOURCE = ROOT / "docs/ttqs-quality-20261016/quality-review/0033_R03/candidate/0033_R03.docx"
SOURCE_REGISTER = ROOT / "docs/ttqs-quality-20261016/quality-review/0033_R03/evidence/SYNTHETIC_REPLACEMENT_REGISTER.csv"
BLUEPRINT = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\CONTROL\BLUEPRINT_CURRENT.jsonl")
OFFICIAL = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\SOURCES\DRIVE_READBACKS\DRIVE_OFFICIAL_EVIDENCE_TABLE.pdf")
OFFICIAL_PAGE = PKG / "source/official_evidence_table_page4.png"
EXPECTED = {
    "docx": "e8212788b1f7b201ca4af4f4c5fb83ab131627dbfa815987f7cf6344edfda727",
    "register": "c79dd292c3679b83d334f5773694d2d9089dce7c7d532197d287e3c5ad5be501",
    "source_docx": "fe5f3cfd84c761da86435a71e0fd899daa725be78628c4d2d3f2324c1fdb5f4f",
    "source_register": "d114f36d9c4852cda2239bf22fcc709bdc0a40aad62418e076c6b9c2ae523247",
    "blueprint": "406d0495365d0970dcf8113a42842663b8a54372a4e204cc492b7c3f8156ef3e",
    "official_pdf": "805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709",
}


def sha(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1_048_576), b""):
            h.update(chunk)
    return h.hexdigest()


hashes = {
    "docx": sha(DOCX),
    "register": sha(REGISTER),
    "source_docx": sha(SOURCE),
    "source_register": sha(SOURCE_REGISTER),
    "blueprint": sha(BLUEPRINT),
    "official_pdf": sha(OFFICIAL),
    "official_page4_png": sha(OFFICIAL_PAGE),
    "render_pdf": sha(PDF),
}
for key, value in EXPECTED.items():
    assert hashes[key] == value, f"{key} SHA mismatch: {hashes[key]}"

with zipfile.ZipFile(DOCX) as archive:
    assert archive.testzip() is None, "DOCX ZIP CRC failure"
    assert {"[Content_Types].xml", "word/document.xml", "word/styles.xml"}.issubset(archive.namelist())

doc = Document(str(DOCX))
table_text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
body_text = "\n".join([*(p.text for p in doc.paragraphs), table_text])
for s in doc.sections:
    body_text += "\n" + "\n".join([*(p.text for p in s.header.paragraphs), *(p.text for p in s.footer.paragraphs)])
for required in (
    "DOCX 0049", "I11-E01", "官方方框 051",
    "訓練課程計畫配合業主、受訓者以及職能需求等調查規劃的證明。",
    "SYN-0033-PLAN-01", "DOCX 0033 R03", "DOCX 0050", "DOCX 0048",
    "SAMPLE", "SYNTHETIC", "業主", "受訓者", "職能需求", "方案甲", "方案乙", "方案丙",
    "360分鐘", "8/10", "正式採用前四道檢核門", "UNKNOWN", "NOT APPROVED", "NOT IMPLEMENTED",
    "第4.1節方案比較", "0033 R03表8", "0033 R03表12", "0033 R03表13", "示範班級12人", "建議14席",
):
    assert required in body_text, f"missing required content: {required}"
for stale in ("同案Q04重要度資料列在表3", "0033 R03表14", "0033 R03表16", "第5節方案比較", "示範班級8人", "建議10席"):
    assert stale not in body_text, f"stale reviewed defect remains: {stale}"
assert len(doc.tables) == 31, f"unexpected table count: {len(doc.tables)}"

with REGISTER.open(encoding="utf-8-sig", newline="") as f:
    register = list(csv.DictReader(f))
assert len(register) == 158
assert len({row["fact_id"] for row in register}) == len(register)
assert all(row["doc_id"] == "0049" and row["case_id"] == "SYN-0033-PLAN-01" for row in register)
assert all("SAMPLE" in row["verification_status"] or "SYNTHETIC" in row["provenance_status"] for row in register)

source_doc = Document(str(SOURCE))
importance = [[cell.text.strip() for cell in row.cells] for row in source_doc.tables[8].rows]
confidence = [[cell.text.strip() for cell in row.cells] for row in source_doc.tables[9].rows]
alternatives = [[cell.text.strip() for cell in row.cells] for row in source_doc.tables[13].rows]
modules = [[cell.text.strip() for cell in row.cells] for row in source_doc.tables[15].rows]
importance_sums = [sum(int(row[col]) for row in importance[1:]) for col in range(2, 6)]
confidence_sums = [sum(int(row[col]) for row in confidence[1:]) for col in range(2, 6)]
importance_means = [round(value / 9, 2) for value in importance_sums]
confidence_means = [round(value / 9, 2) for value in confidence_sums]
gaps = [round(a - b, 2) for a, b in zip(importance_means, confidence_means)]
weights = [0.4, 0.3, 0.2, 0.1]
alternative_scores = [round(sum(int(row[col]) * weights[col - 2] for col in range(2, 6)), 2) for row in alternatives[1:]]
module_minutes = [int(re.search(r"\d+", row[3]).group()) for row in modules[1:]]
assert importance_sums == [41, 38, 35, 39]
assert confidence_sums == [24, 25, 27, 21]
assert importance_means == [4.56, 4.22, 3.89, 4.33]
assert confidence_means == [2.67, 2.78, 3.00, 2.33]
assert gaps == [1.89, 1.44, 0.89, 2.00]
assert alternative_scores == [4.80, 4.00, 3.50]
assert module_minutes == [90, 90, 120, 60] and sum(module_minutes) == 360

font_sizes = {"body": [], "table": [], "header": [], "footer": []}
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        if run.text.strip() and run.font.size:
            font_sizes["body"].append(run.font.size.pt)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    if run.text.strip() and run.font.size:
                        font_sizes["table"].append(run.font.size.pt)
for section in doc.sections:
    for kind, paragraphs in (("header", section.header.paragraphs), ("footer", section.footer.paragraphs)):
        for paragraph in paragraphs:
            for run in paragraph.runs:
                if run.text.strip() and run.font.size:
                    font_sizes[kind].append(run.font.size.pt)
minimum_fonts = {key: min(values) if values else None for key, values in font_sizes.items()}
assert minimum_fonts["body"] >= 10.5, minimum_fonts
assert minimum_fonts["table"] >= 9.5, minimum_fonts

pages = sorted(RENDER.glob("page-*.png"), key=lambda p: int(re.search(r"\d+", p.stem).group()))
assert len(pages) == 15 and len(PdfReader(str(PDF)).pages) == 15
page_receipts = []
for number, page in enumerate(pages, 1):
    with Image.open(page) as image:
        assert image.width >= 1100 and image.height >= 1500
        page_receipts.append({"page": number, "width": image.width, "height": image.height, "bytes": page.stat().st_size, "sha256": sha(page)})

receipt = {
    "project_id": "TTQS_ONE", "doc_id": "0049", "revision": "R01", "requirement_id": "I11-E01",
    "official_checkbox_sequence": "051", "candidate_state": "CANDIDATE_ONLY", "formal_acceptance": "NOT_GRANTED",
    "independent_semantic_review": "PENDING", "sha256": hashes,
    "docx_bytes": DOCX.stat().st_size, "paragraph_count": len(doc.paragraphs), "table_count": len(doc.tables),
    "body_text_characters": len(body_text), "sample_register_rows": len(register), "sample_register_fact_ids_unique": True,
    "case_id": "SYN-0033-PLAN-01", "source_0033_exact_sha_verified": True,
    "calculation_recheck": {
        "importance_sums": importance_sums, "confidence_sums": confidence_sums,
        "importance_means": importance_means, "confidence_means": confidence_means, "gaps": gaps,
        "weighted_alternatives": alternative_scores, "module_minutes": module_minutes,
        "module_total_minutes": sum(module_minutes),
    },
    "render": {"pdf_pages": 15, "page_image_count": len(pages), "pages": page_receipts},
    "font_minimums_pt": minimum_fonts, "automated_technical_checks": "PASS",
    "renderer": "Microsoft Word native PDF export; rendered to PNG with bundled Poppler because LibreOffice is not installed on this Windows host.",
    "visual_review_scope": "All 15 final page images individually reviewed at full resolution by primary reviewer; final contact sheet reviewed. No clipping, overlap, broken cell, or orphan table header observed. Section 10.1 begins with its header and four gate rows together; Chapter 5 starts on a fresh page as specified by the document structure.",
    "reviewer_route": "OpenCode CLI v2 independent exact-SHA semantic and visual review pending",
    "recorded_local_time": datetime.now().astimezone().isoformat(),
}
output = PKG / "evidence/qa/TECHNICAL_VISUAL_QA_R01.json"
output.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps({key: receipt[key] for key in ("sha256", "docx_bytes", "paragraph_count", "table_count", "body_text_characters", "sample_register_rows", "calculation_recheck", "render", "font_minimums_pt", "automated_technical_checks")}, ensure_ascii=False, indent=2))
