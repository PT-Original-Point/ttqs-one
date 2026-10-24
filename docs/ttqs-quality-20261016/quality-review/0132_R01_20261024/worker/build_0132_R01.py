from __future__ import annotations
import hashlib, json, os, re
from pathlib import Path
os.environ['MPLBACKEND'] = 'Agg'
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
BODY = ROOT / 'candidate' / '0132_I19-E21_BODY_R01.md'
OUT = ROOT / 'candidate' / '0132_I19-E21_CANDIDATE_R01.docx'
CHARTS = ROOT / 'render' / 'R01'
CHARTS.mkdir(parents=True, exist_ok=True)
FONT = 'Microsoft JhengHei'

# Reuse the previously generated chart images: the six batch values did not change in this body repair.
for asset in [CHARTS/'pooled_rate.png', CHARTS/'period_trend.png']:
    if not asset.is_file() or asset.stat().st_size < 1000:
        raise FileNotFoundError(f'validated chart asset missing or empty: {asset}')


def set_run_font(run, size=10.5, bold=None, color=None, font=FONT):
    run.font.name = font
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font)
    run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = RGBColor(*color)


def add_inline(p, text, size=10.5, color=None):
    pat = re.compile(r'(\*\*.+?\*\*|`[^`]+`)')
    for piece in pat.split(text):
        if not piece: continue
        if piece.startswith('**') and piece.endswith('**'):
            run=p.add_run(piece[2:-2]); set_run_font(run,size,bold=True,color=color)
        elif piece.startswith('`') and piece.endswith('`'):
            run=p.add_run(piece[1:-1]); set_run_font(run, max(8.5,size-1),font='Consolas',color=(70,70,70))
        else:
            run=p.add_run(piece); set_run_font(run,size,color=color)


def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)


def set_cell_margins(cell, top=70, start=75, bottom=70, end=75):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); mar=tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar=OxmlElement('w:tcMar'); tcPr.append(mar)
    for side,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=mar.find(qn('w:'+side))
        if node is None: node=OxmlElement('w:'+side); mar.append(node)
        node.set(qn('w:w'),str(val)); node.set(qn('w:type'),'dxa')


def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)


def prevent_row_split(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:cantSplit'); trPr.append(el)


def parse_table_row(line):
    line=line.strip().strip('|')
    return [c.strip() for c in line.split('|')]


def add_md_table(doc, lines, usable_width):
    rows=[parse_table_row(x) for x in lines]
    if len(rows)>1 and all(re.fullmatch(r':?-{3,}:?',c.replace(' ','')) for c in rows[1]):
        rows.pop(1)
    cols=max(len(r) for r in rows)
    table=doc.add_table(rows=0, cols=cols)
    table.style='Table Grid'; table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
    width=usable_width/cols
    for ri,data in enumerate(rows):
        cells=table.add_row().cells
        for ci in range(cols):
            cell=cells[ci]; cell.width=Inches(width); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            cell.text=''
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
            value=data[ci] if ci<len(data) else ''
            add_inline(p,value,9.5,color=(255,255,255) if ri==0 else None)
            if ri==0:
                shade(cell,'244A68')
                for run in p.runs: run.bold=True
            elif ri%2==0: shade(cell,'F2F6F8')
        prevent_row_split(table.rows[-1])
    set_repeat_table_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return table


def add_page_field(p):
    r=p.add_run(); set_run_font(r,9,color=(90,90,90))
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); r._r.addnext(fld)


def setup(doc):
    sec=doc.sections[0]
    sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
    sec.top_margin=Inches(.68); sec.bottom_margin=Inches(.66); sec.left_margin=Inches(.62); sec.right_margin=Inches(.62)
    sec.header_distance=Inches(.28); sec.footer_distance=Inches(.3)
    styles=doc.styles
    normal=styles['Normal']; normal.font.name=FONT; normal._element.rPr.rFonts.set(qn('w:eastAsia'),FONT); normal.font.size=Pt(10.5)
    normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.10
    for name,size,col in [('Title',23,(31,70,97)),('Heading 1',15,(31,70,97)),('Heading 2',12.7,(47,98,131)),('Heading 3',11.2,(55,83,101))]:
        st=styles[name]; st.font.name=FONT; st._element.rPr.rFonts.set(qn('w:eastAsia'),FONT); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor(*col); st.paragraph_format.space_before=Pt(10); st.paragraph_format.space_after=Pt(4); st.paragraph_format.keep_with_next=True
    title_ppr=styles['Title']._element.pPr
    if title_ppr is not None:
        for border in list(title_ppr.findall(qn('w:pBdr'))): title_ppr.remove(border)
    for nm in ['List Bullet','List Number']:
        st=styles[nm]; st.font.name=FONT; st._element.rPr.rFonts.set(qn('w:eastAsia'),FONT); st.font.size=Pt(10.5)
    header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    add_inline(header,'TTQS ONE　｜　0132　｜　SAMPLE / SYNTHETIC',8.5,color=(90,103,112))
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    add_inline(footer,'CANDIDATE　•　REAL EVIDENCE UNKNOWN　•　Page ',8.5,color=(90,103,112)); add_page_field(footer)
    doc.core_properties.title='0132 個人產品不良件數與訓後改善分析表'
    doc.core_properties.subject='TTQS ONE I19-E21 | SAMPLE/SYNTHETIC candidate'
    doc.core_properties.author='TTQS ONE candidate production'


def add_chart(doc, marker, usable_width):
    if marker=='POOLED_RATE':
        file=CHARTS/'pooled_rate.png'; cap='圖 1｜訓前與訓後加權不良率比較（SAMPLE／SYNTHETIC）。'
    else:
        file=CHARTS/'period_trend.png'; cap='圖 2｜六批逐期不良率；虛線為合成訓練與同時工具／材料變更界線，不是因果證據。'
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(1)
    pic=p.add_run().add_picture(str(file),width=Inches(min(6.15,usable_width-.2)))
    pic._inline.docPr.set('descr',cap)
    c=doc.add_paragraph(); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.space_after=Pt(7)
    add_inline(c,cap,9.2,color=(73,83,90))


def build():
    text=BODY.read_text(encoding='utf-8')
    doc=Document(); setup(doc)
    lines=text.splitlines(); i=0; usable=8.27-.62-.62
    while i<len(lines):
        line=lines[i].rstrip()
        if not line.strip(): i+=1; continue
        m=re.fullmatch(r'\[\[CHART:(POOLED_RATE|PERIOD_TREND)\]\]',line.strip())
        if m: add_chart(doc,m.group(1),usable); i+=1; continue
        if line.startswith('|'):
            block=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                block.append(lines[i].strip()); i+=1
            if len(block)>=2: add_md_table(doc,block,usable)
            continue
        if line.strip() in ['### SAMPLE 後續查核排序（不是風險等級）','## 8. 可直接使用的逐批記錄表']:
            doc.add_page_break()
        if line.startswith('### '):
            p=doc.add_paragraph(style='Heading 3'); add_inline(p,line[4:],11.2); i+=1; continue
        if line.startswith('## '):
            p=doc.add_paragraph(style='Heading 2'); add_inline(p,line[3:],12.7); i+=1; continue
        if line.startswith('# '):
            p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; add_inline(p,line[2:],23); i+=1; continue
        if line.startswith('> '):
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.18); p.paragraph_format.right_indent=Inches(.18); p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(7)
            pPr=p._p.get_or_add_pPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'FFF2D5'); pPr.append(shd)
            add_inline(p,line[2:],10,color=(101,66,16)); i+=1; continue
        if line.startswith('- [ ] '):
            p=doc.add_paragraph(style='List Bullet'); add_inline(p,'☐ '+line[6:],10.5); i+=1; continue
        if line.startswith('- '):
            p=doc.add_paragraph(style='List Bullet'); add_inline(p,line[2:],10.5); i+=1; continue
        if re.match(r'^\d+\.\s',line):
            p=doc.add_paragraph(style='List Number'); add_inline(p,re.sub(r'^\d+\.\s','',line),10.5); i+=1; continue
        if line.strip()=='---':
            p=doc.add_paragraph(); pPr=p._p.get_or_add_pPr(); bdr=OxmlElement('w:pBdr'); bottom=OxmlElement('w:bottom'); bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),'6'); bottom.set(qn('w:color'),'8FAFC2'); bdr.append(bottom); pPr.append(bdr); i+=1; continue
        p=doc.add_paragraph(); add_inline(p,line,10.5); i+=1
    doc.save(OUT)
    digest=hashlib.sha256(OUT.read_bytes()).hexdigest()
    receipt={'doc_id':'0132','internal_requirement_id':'I19-E21','source_body_sha256':hashlib.sha256(BODY.read_bytes()).hexdigest(),'candidate_path':str(OUT),'candidate_sha256':digest,'candidate_bytes':OUT.stat().st_size,'builder':'custom deterministic python-docx markdown table renderer','builder_script_sha256':hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),'charts':{'pooled_rate.png':hashlib.sha256((CHARTS/'pooled_rate.png').read_bytes()).hexdigest(),'period_trend.png':hashlib.sha256((CHARTS/'period_trend.png').read_bytes()).hexdigest()},'font_body_pt':10.5,'font_table_pt':9.5,'real_evidence':'UNKNOWN','semantic_review':'PENDING','formal_acceptance':'NOT_GRANTED'}
    (ROOT/'qa'/'BUILD_RECEIPT_0132_R01.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(receipt,ensure_ascii=False))

if __name__=='__main__': build()