from pathlib import Path
from hashlib import sha256
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
from docx.text.paragraph import Paragraph
from docx.table import Table
import json
root=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0132_I19-E21_SOURCE_PREP_R01_20261024')
rev=root/'review'/'W2_EXACT_SHA_20261026'
docx=root/'candidate'/'0132_I19-E21_CANDIDATE_R01.docx'
doc=Document(str(docx))
parts=[]; tn=0; pn=0
for child in doc.element.body.iterchildren():
    if isinstance(child, CT_P):
        p=Paragraph(child,doc)
        if p.text.strip(): parts.append(p.text.strip()); pn+=1
    elif isinstance(child, CT_Tbl):
        tn+=1; parts.append(f'\n[TABLE {tn}]'); t=Table(child,doc)
        for row in t.rows: parts.append(' | '.join(c.text.replace('\n',' / ').strip() for c in row.cells))
body='\n'.join(parts)+'\n'
ns={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
def vals(node):
    if node is None:return []
    out=[]
    for pt in node.xpath('.//c:pt',namespaces=ns):
        ix=pt.get('idx'); v=pt.xpath('./c:v/text()',namespaces=ns)
        if v: out.append((ix,v[0]))
    if not out:
        out=[(str(i),v) for i,v in enumerate(node.xpath('.//c:v/text()',namespaces=ns))]
    return out
charts=[]
with ZipFile(docx) as z:
    names=set(z.namelist())
    for name in sorted(n for n in names if n.startswith('word/charts/chart') and n.endswith('.xml')):
        x=etree.fromstring(z.read(name)); title=' '.join(x.xpath('.//c:title//a:t/text()',namespaces=ns))
        series=[]
        for ser in x.xpath('.//c:ser',namespaces=ns):
            tx=ser.find('c:tx',namespaces=ns)
            label=' '.join(tx.xpath('.//c:v/text()|.//c:pt//c:v/text()',namespaces=ns)) if tx is not None else ''
            cat=ser.find('c:cat',namespaces=ns); val=ser.find('c:val',namespaces=ns)
            cats=vals(cat); vs=vals(val)
            series.append({'series':label,'categories':cats,'values':vs})
        charts.append({'part':name,'title':title,'series':series})
headers=[]; footers=[]
for si,s in enumerate(doc.sections,1):
    for kind,obj,target in [('header',s.header,headers),('footer',s.footer,footers)]:
        text=' | '.join(p.text.strip() for p in obj.paragraphs if p.text.strip())
        for ti,t in enumerate(obj.tables,1):
            for row in t.rows: text+=' | '+' | '.join(c.text.replace('\n',' / ').strip() for c in row.cells)
        if text: target.append({'section':si,'text':text})
chart_text=json.dumps(charts,ensure_ascii=False,indent=2)
extra='\n[HEADERS]\n'+json.dumps(headers,ensure_ascii=False,indent=2)+'\n[FOOTERS]\n'+json.dumps(footers,ensure_ascii=False,indent=2)+'\n[EMBEDDED_CHART_CACHED_DATA]\n'+chart_text+'\n'
payload=body+extra
meta=(f'DOC_ID=0132\nINTERNAL_REQUIREMENT=I19-E21\nPROJECT_157_INDEX=145 (I19-L3)\nOFFICIAL_TEXT=個人產品不良件數有下降趨勢。\nCANDIDATE_DOCX_SHA256={sha256(docx.read_bytes()).hexdigest()}\nEXTRACTION_METHOD=python-docx body paragraphs/tables + OOXML chart cached categories/values + section headers/footers\nEXTRACTED_PAYLOAD_SHA256={sha256(payload.encode()).hexdigest()}\nPARAGRAPH_BLOCKS={pn}\nTABLES={tn}\nCHARTS={len(charts)}\nVISUAL_REVIEW=not part of semantic payload; existing 10-page Word visual QA is separate\nSTATUS=CANDIDATE; all sample values synthetic; REAL evidence UNKNOWN; formal acceptance NOT_GRANTED\n\n')
out=rev/'0132_R01_EXACT_SHA_SEMANTIC_PAYLOAD_V2.txt'
out.write_text(meta+payload,encoding='utf-8',newline='\n')
print(json.dumps({'path':str(out),'docx_sha256':sha256(docx.read_bytes()).hexdigest(),'payload_sha256':sha256(payload.encode()).hexdigest(),'payload_bytes':len(payload.encode()),'paragraph_blocks':pn,'tables':tn,'charts':len(charts),'chart_titles':[c['title'] for c in charts],'chart_values':[s for c in charts for s in c['series']]},ensure_ascii=False))