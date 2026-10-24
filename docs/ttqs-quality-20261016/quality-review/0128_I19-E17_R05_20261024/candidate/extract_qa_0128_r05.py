from pathlib import Path
from zipfile import ZipFile
from datetime import datetime
from docx import Document
from docx.oxml.ns import qn
import hashlib,json,re
root=Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0128_I19-E17_SOURCE_PREP_R01_20261024")
p=root/"candidate"/"0128_I19-E17_R05.docx"
sha=hashlib.sha256(p.read_bytes()).hexdigest()
doc=Document(p)
parts=[]
for item in doc.iter_inner_content():
    if hasattr(item,"rows"):
        parts.append("\n[TABLE]\n")
        for row in item.rows:
            parts.append(" | ".join(cell.text.replace("\n"," / ") for cell in row.cells)+"\n")
        parts.append("[/TABLE]\n")
    else:
        if item.text.strip(): parts.append(item.text.strip()+"\n")
extracted="\n".join(parts)
out=root/"candidate"/"0128_I19-E17_R05_EXTRACTED_TEXT.txt"
out.write_text(extracted,encoding="utf-8",newline="")
with ZipFile(p) as z:
    crc=z.testzip()
    members=z.namelist()
orientation=[]
for s in doc.sections: orientation.append({"width_in":round(s.page_width.inches,2),"height_in":round(s.page_height.inches,2),"orientation":"landscape" if s.page_width>s.page_height else "portrait"})
textchecks={"all_six_cases":all(f"S-EMP-{i:02d}" in extracted for i in range(1,7)),"weighted_aggregates":all(x in extracted for x in ["93 / 120","101 / 120","77.50%","84.17%","+6.67 pp"]),"contains_markdown_artifacts":any(x in extracted for x in ["**","|:---","[!CAUTION]","```","\\text{"]),"internal_path_leak":bool(re.search(r"[A-Z]:\\\\|github\.com/",extracted,re.I)),"synthetic_boundary":"SAMPLE / SYNTHETIC" in extracted or "SAMPLE" in extracted}
q={"project_id":"TTQS_ONE","work_id":"TTQS-0128-I19-E17-TECH-QA-R05","doc_id":"0128","official_requirement_id":"I19-E17","candidate_path":str(p),"docx_sha256":sha,"docx_bytes":p.stat().st_size,"extracted_text_path":str(out),"extracted_text_sha256":hashlib.sha256(out.read_bytes()).hexdigest(),"extracted_chars":len(extracted),"package_crc":"PASS" if crc is None else "FAIL","section_layouts":orientation,"table_count":len(doc.tables),"paragraph_count":len(doc.paragraphs),"text_checks":textchecks,"visual_render":"BLOCKED_NO_LIBREOFFICE_ON_PATH; Word COM not retried because prior bounded attempt timed out and renderer state has not changed","semantic_review":"PENDING_FRESH_REVIEW","real_evidence":"UNKNOWN_UNTIL_SOURCE","formal_acceptance":"NOT_GRANTED","checked_at_local":datetime.now().astimezone().isoformat()}
qpath=root/"candidate"/"TECHNICAL_QA_R05.json";qpath.write_text(json.dumps(q,ensure_ascii=False,indent=2),encoding="utf-8")
rr={"work_id":"TTQS-0128-I19-E17-VISUAL-RENDER-R05","docx_sha256":sha,"renderer":"documents skill render_docx.py with bundled Python 3.12.14","command_exit_code":1,"failure":"LibreOffice soffice.exe not found on PATH","renderer_stdout_excerpt":"[render_docx] NOTE: multiple section page sizes/orientations detected; DPI computed from first section; FileNotFoundError: LibreOffice soffice.exe was not found on PATH","retry_decision":"No retry: no renderer-state change; previous bounded Word COM export attempt for another DOCX timed out; do not relaunch existing WINWORD automation processes","status":"VISUAL_QA_BLOCKED_NOT_PASS","checked_at_local":datetime.now().astimezone().isoformat()}
rpath=root/"candidate"/"VISUAL_RENDER_FAILURE_R05.json";rpath.write_text(json.dumps(rr,ensure_ascii=False,indent=2),encoding="utf-8")
print(json.dumps({"docx_sha256":sha,"extracted_sha256":hashlib.sha256(out.read_bytes()).hexdigest(),"extracted_chars":len(extracted),"textchecks":textchecks,"sections":orientation,"tables":len(doc.tables),"qa_path":str(qpath),"qa_sha256":hashlib.sha256(qpath.read_bytes()).hexdigest(),"visual_receipt_sha256":hashlib.sha256(rpath.read_bytes()).hexdigest()},ensure_ascii=False))