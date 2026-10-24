from __future__ import annotations

import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "candidate" / "0065_R04_body.md"
OUTPUT = ROOT / "candidate" / "0065_R04.docx"
RECEIPT = ROOT / "evidence" / "build_layout_receipt_R04.json"
FONT = "Microsoft JhengHei"
NAVY = "17365D"
BLUE = "D9EAF7"
PALE = "F3F6F8"
AMBER = "FFF2CC"
GRAY = "667085"
PAGE_WIDTH = 8.27
LEFT_MARGIN = 0.66
RIGHT_MARGIN = 0.66
CONTENT_WIDTH = PAGE_WIDTH - LEFT_MARGIN - RIGHT_MARGIN


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=75, start=95, bottom=75, end=95):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    tr_pr.append(el)


def set_cell_width(cell, inches: float):
    cell.width = Inches(inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_run_font(run, size=10.5, bold=None, color=None, italic=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_inline(paragraph, text: str, default_size=10.5, default_color=None):
    # Preserve content while rendering common Markdown emphasis/code.
    token = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*)")
    pos = 0
    for m in token.finditer(text):
        if m.start() > pos:
            r = paragraph.add_run(text[pos:m.start()])
            set_run_font(r, default_size, color=default_color)
        t = m.group(0)
        if t.startswith("**"):
            r = paragraph.add_run(t[2:-2])
            set_run_font(r, default_size, bold=True, color=default_color)
        elif t.startswith("`"):
            r = paragraph.add_run(t[1:-1])
            set_run_font(r, default_size, color="365F91")
        else:
            r = paragraph.add_run(t[1:-1])
            set_run_font(r, default_size, italic=True, color=default_color)
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:])
        set_run_font(r, default_size, color=default_color)


def parse_cells(line: str):
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [x.strip() for x in s.split("|")]


def is_table_separator(line: str):
    cells = parse_cells(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells)


def add_paragraph(doc, text, style=None, quote=False, bullet=None, level=0):
    p = doc.add_paragraph(style=style)
    if bullet is not None:
        p.paragraph_format.left_indent = Inches(0.22 + level * 0.18)
        p.paragraph_format.first_line_indent = Inches(-0.16)
        r = p.add_run(("• " if bullet == "bullet" else f"{bullet} "))
        set_run_font(r, 10.5, bold=(bullet == "bullet"), color=NAVY)
    add_inline(p, text, 10.5, "333333")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.12
    p.paragraph_format.keep_together = True
    if quote:
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.right_indent = Inches(0.12)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(7)
        p_pr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), AMBER if "SAMPLE／SYNTHETIC 全篇聲明" in text else PALE)
        p_pr.append(shd)
        for r in p.runs:
            if "SAMPLE／SYNTHETIC 全篇聲明" in text:
                r.font.color.rgb = RGBColor.from_string("7F6000")
                if "SAMPLE／SYNTHETIC" in r.text:
                    r.bold = True
    return p


def add_table(doc, rows):
    if not rows:
        return
    header = rows[0]
    data = rows[1:]
    n = len(header)
    # Very wide Markdown tables become readable record cards in Word; every original
    # header/value is retained, while avoiding tiny landscape text and clipped columns.
    if n > 5:
        for idx, row in enumerate(data, 1):
            row = (row + [""] * n)[:n]
            nonempty = [v for v in row if v.strip()]
            if not nonempty:
                title = f"填寫列 {idx}"
            else:
                title = None
                if "指標" in [h.strip() for h in header] and "#" in [h.strip() for h in header]:
                    number_index = [h.strip() for h in header].index("#")
                    metric_index = [h.strip() for h in header].index("指標")
                    if row[number_index].strip():
                        title = f"指標 #{row[number_index].strip()}｜{row[metric_index].strip() or '待填指標名稱'}"
                for j, h in enumerate(header):
                    if title is not None:
                        break
                    if h.strip() in {"#", "編號", "序號"} and row[j].strip():
                        title = f"{h.strip()} {row[j].strip()}"
                        break
                if title is None:
                    for j, h in enumerate(header):
                        if h.strip() in {"指標", "議題", "項目", "欄位"} and row[j].strip():
                            title = row[j].strip()
                            break
                title = title or f"記錄 {idx}"
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(7)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            rr = p.add_run(title)
            set_run_font(rr, 10.5, bold=True, color=NAVY)
            table = doc.add_table(rows=0, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            lower_headers = [h.strip() for h in header]
            groups = None
            if {"分子", "分母", "結果", "期間／日期", "來源（示範）", "示範門檻", "落差", "限制", "決策問題"}.issubset(set(lower_headers)):
                ix = {h: lower_headers.index(h) for h in lower_headers}
                def field(key, placeholder="__________"):
                    value = row[ix[key]].strip()
                    return value if value else placeholder
                formula = f"分子 {field('分子')} ÷ 分母 {field('分母')} = 結果 {field('結果')}"
                period_source = f"期間：{field('期間／日期')}；來源：{field('來源（示範）')}"
                target_gap = f"示範門檻：{field('示範門檻')}；落差：{field('落差')}"
                limit_question = f"限制：{field('限制')}；決策問題：{field('決策問題')}"
                groups = [
                    ("分子／分母／結果", formula, ["分子", "分母", "結果"]),
                    ("期間／來源", period_source, ["期間／日期", "來源（示範）"]),
                    ("示範門檻／落差", target_gap, ["示範門檻", "落差"]),
                    ("限制／決策問題", limit_question, ["限制", "決策問題"]),
                ]
            elif {"發現（證據ID）", "根因假設 vs 已確認事實", "具體措施", "負責功能", "擬議期限", "可量測驗證方式", "應留存來源", "狀態", "未授權時"}.issubset(set(lower_headers)):
                ix = {h: lower_headers.index(h) for h in lower_headers}
                groups = [
                    ("發現／證據", row[ix["發現（證據ID）"]], ["發現（證據ID）"]),
                    ("根因假設／已確認事實", row[ix["根因假設 vs 已確認事實"]], ["根因假設 vs 已確認事實"]),
                    ("措施／負責功能／期限", f"{row[ix['具體措施']]}；負責：{row[ix['負責功能']]}；期限：{row[ix['擬議期限']]}", ["具體措施", "負責功能", "擬議期限"]),
                    ("驗證方式／留存來源", f"{row[ix['可量測驗證方式']]}；留存：{row[ix['應留存來源']]}", ["可量測驗證方式", "應留存來源"]),
                    ("狀態／未授權處置", f"{row[ix['狀態']]}；{row[ix['未授權時']]}", ["狀態", "未授權時"]),
                ]
            if groups is None:
                groups = []
                for j, (h, val) in enumerate(zip(header, row)):
                    if h.strip() in {"#", "編號", "序號"} and val.strip():
                        continue
                    groups.append((h.strip(), val.strip() or "________________", [h.strip()]))
            for j, (label, val, original_fields) in enumerate(groups):
                cells = table.add_row().cells
                set_cell_width(cells[0], 1.52)
                set_cell_width(cells[1], CONTENT_WIDTH - 1.52)
                cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_shading(cells[0], BLUE)
                set_cell_shading(cells[1], "FFFFFF" if j % 2 == 0 else PALE)
                for c in cells:
                    set_cell_margins(c)
                p0 = cells[0].paragraphs[0]
                p0.paragraph_format.space_after = Pt(0)
                r0 = p0.add_run(label)
                set_run_font(r0, 9.5, bold=True, color=NAVY)
                p1 = cells[1].paragraphs[0]
                p1.paragraph_format.space_after = Pt(0)
                add_inline(p1, val.strip() or "________________", 9.5, "333333")
                if j < len(groups) - 1:
                    p0.paragraph_format.keep_with_next = True
                    p1.paragraph_format.keep_with_next = True
                set_row_cant_split(table.rows[-1])
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
        return

    table = doc.add_table(rows=1, cols=n)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    widths = [CONTENT_WIDTH / n] * n
    if n == 5 and any("時段" in h for h in header):
        widths = [0.90, 0.55, 2.30, 1.32, CONTENT_WIDTH - 5.07]
    elif n == 4 and any("預算" in h for h in header):
        widths = [2.20, 1.45, 1.45, CONTENT_WIDTH - 5.10]
    elif n == 3:
        widths = [1.35, 2.55, CONTENT_WIDTH - 3.90]
    for j, cell in enumerate(table.rows[0].cells):
        set_cell_width(cell, widths[j])
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, 90, 95, 90, 95)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_inline(p, header[j], 9.5, "FFFFFF")
        for run in p.runs:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string("FFFFFF")
    set_repeat_table_header(table.rows[0])
    set_row_cant_split(table.rows[0])
    for i, vals in enumerate(data):
        vals = (vals + [""] * n)[:n]
        cells = table.add_row().cells
        set_row_cant_split(table.rows[-1])
        for j, (cell, val) in enumerate(zip(cells, vals)):
            set_cell_width(cell, widths[j])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, "FFFFFF" if i % 2 == 0 else PALE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            add_inline(p, val, 9.5, "333333")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_page_field(paragraph):
    r = paragraph.add_run()
    set_run_font(r, 8.5, color=GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r._r.extend([begin, instr, separate, text, end])


def configure(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(PAGE_WIDTH)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(0.62)
    sec.bottom_margin = Inches(0.62)
    sec.left_margin = Inches(LEFT_MARGIN)
    sec.right_margin = Inches(RIGHT_MARGIN)
    sec.header_distance = Inches(0.28)
    sec.footer_distance = Inches(0.30)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("333333")
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, color in [("Heading 1", 16, NAVY), ("Heading 2", 13.5, NAVY), ("Heading 3", 11.5, "365F91")]:
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10 if name != "Heading 1" else 0)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    hdr = sec.header.paragraphs[0]
    hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hdr.add_run("DOCX 0065  |  I15-E01  |  SAMPLE／SYNTHETIC")
    set_run_font(r, 8, bold=True, color=GRAY)
    ftr = sec.footer.paragraphs[0]
    ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = ftr.add_run("示範文件，不是真實協會紀錄   •   第 ")
    set_run_font(r, 8.5, color=GRAY)
    add_page_field(ftr)
    r = ftr.add_run(" 頁")
    set_run_font(r, 8.5, color=GRAY)
    doc.core_properties.title = "訓練後定期綜合檢討／管理審查會議紀錄（單一合成案例）"
    doc.core_properties.subject = "DOCX 0065｜I15-E01｜官方方框 071｜SAMPLE/SYNTHETIC"
    doc.core_properties.author = "原始點協會訓練行政文件示範"
    doc.core_properties.keywords = "TTQS; I15-E01; SAMPLE; SYNTHETIC"


def build():
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()
    configure(doc)
    blocks = {"headings": 0, "tables": 0, "wide_tables_stacked": 0, "quotes": 0, "list_items": 0}
    i = 0
    first_heading = True
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p_pr = p._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "4")
            bottom.set(qn("w:color"), "D9E2F3")
            bottom.set(qn("w:space"), "1")
            borders.append(bottom)
            p_pr.append(borders)
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and is_table_separator(lines[i + 1].strip()):
            rows = [parse_cells(line)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(parse_cells(lines[i].strip()))
                i += 1
            blocks["tables"] += 1
            if len(rows[0]) > 5:
                blocks["wide_tables_stacked"] += 1
            add_table(doc, rows)
            continue
        if line.startswith(">"):
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                q = lines[i].strip()[1:].lstrip()
                if q.startswith("## "):
                    q = q[3:]
                quote_lines.append(q)
                i += 1
            p = add_paragraph(doc, "\n".join(quote_lines), quote=True)
            blocks["quotes"] += 1
            if "SAMPLE／SYNTHETIC 全篇聲明" in p.text:
                for r in p.runs:
                    r.bold = True
                    r.font.color.rgb = RGBColor.from_string("7F6000")
            continue
        if line.startswith("#"):
            m = re.match(r"^(#{1,4})\s+(.*)$", line)
            level, title = len(m.group(1)), m.group(2)
            if level == 1:
                p = doc.add_paragraph(style="Title")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(12)
                p.paragraph_format.keep_with_next = True
                add_inline(p, title, 20, NAVY)
                for r in p.runs:
                    r.bold = True
                first_heading = False
            else:
                p = doc.add_paragraph(style=f"Heading {min(level - 1, 3)}")
                if re.match(r"## (4\.|5\.|6\.|7\.|10\.)", line) or "可直接使用之空白真實會議紀錄表" in title or title.startswith("附錄：HOLD／退回條件"):
                    p.paragraph_format.page_break_before = True
                add_inline(p, title, 16 if level == 2 else (13.5 if level == 3 else 11.5), NAVY)
                blocks["headings"] += 1
            i += 1
            continue
        if re.match(r"(?:[-*])\s+|\d+\.\s+", line):
            while i < len(lines):
                item = lines[i].strip()
                bm = re.match(r"^[-*]\s+(.*)$", item)
                nm = re.match(r"^(\d+)\.\s+(.*)$", item)
                if bm:
                    add_paragraph(doc, bm.group(1), bullet="bullet")
                elif nm:
                    add_paragraph(doc, nm.group(2), bullet=nm.group(1) + ".")
                else:
                    break
                blocks["list_items"] += 1
                i += 1
            continue
        para = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt or nxt == "---" or nxt.startswith("#") or nxt.startswith(">") or nxt.startswith("|") or re.match(r"(?:[-*])\s+|\d+\.\s+", nxt):
                break
            para.append(nxt)
            i += 1
        p = add_paragraph(doc, " ".join(para))
    doc.save(OUTPUT)
    sha = hashlib.sha256(OUTPUT.read_bytes()).hexdigest()
    receipt = {
        "document_id": "0065",
        "revision": "R04",
        "requirement_id": "I15-E01",
        "official_checkbox_serial": "071",
        "status": "CANDIDATE_NOT_FORMALLY_ACCEPTED",
        "source_body": str(SOURCE),
        "source_body_sha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
        "docx_path": str(OUTPUT),
        "docx_bytes": OUTPUT.stat().st_size,
        "docx_sha256": sha,
        "generated_utc": datetime.now(timezone.utc).isoformat(),
        "layout_transformation": "Markdown tables with more than five columns are rendered as readable two-column record cards; all headers and cell values are retained.",
        "revision_parent": {"revision": "R03", "docx_sha256": "7a1c70120da229cac5691aab3dbf34959598735cf4d5191d1b8a465e6de0ad7f", "independent_semantic_verdict": "PASS for candidate status; exact appendix label reference and page-split layout correction pending R04 verification"},
        "revision_change_summary": "Aligned Section 6 appendix reference with its exact heading and forced the HOLD appendix to a new page after R03 visual review found an orphaned heading/item at the bottom of page 21; no case facts or calculations changed.",
        "layout_iteration_history": [
            {
                "iteration": 1,
                "docx_sha256": "d3fe88674a3115f05de3f0f2f5ee4801701ac9b7db099df9fb91cbf74347c744",
                "finding": "Visual inspection found action A-03 and blank-form record rows could split across pages; the wide dashboard/action register also used unnecessarily tall individual field rows.",
                "disposition": "Preliminary build, not externally reviewed or a stable candidate."
            },
            {
                "iteration": 2,
                "layout_fixes": ["group related dashboard fields while preserving their original values", "compact action cards without removing facts", "keep each record card together where page space permits", "raise standard table text to 9.5 pt"],
                "disposition": "Current layout iteration; requires fresh full render and inspection."
            },
            {
                "iteration": 3,
                "finding": "Inspection of iteration 2 found grouped wide-table cards displayed the final source-header label repeatedly instead of each grouped label.",
                "layout_fixes": ["render each group label from its explicit display label", "reinspect metric, action-register and blank-form cards"],
                "disposition": "Current iteration; requires fresh full render and inspection."
            },
            {
                "iteration": 4,
                "finding": "Technical style audit raised standard table header text from 9.4 pt to 9.5 pt to match the document's table readability floor.",
                "layout_fixes": ["standard table headers now use 9.5 pt"],
                "disposition": "R03 rendered and reviewed; see R03 exact-SHA receipts."
            },
            {
                "iteration": 5,
                "finding": "R03 visual inspection found appendix heading and item 1 orphaned at the end of page 21, with items 2-10 alone on page 22.",
                "layout_fixes": ["start the HOLD/return appendix on a new page; align cross-reference with exact heading"],
                "disposition": "Current R04 build; requires full final render, visual inspection, and exact-SHA independent review."
            }
        ],
        "block_counts": blocks,
        "independent_semantic_review": "PENDING",
        "formal_human_acceptance": "PENDING",
    }
    RECEIPT.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(receipt, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    build()
