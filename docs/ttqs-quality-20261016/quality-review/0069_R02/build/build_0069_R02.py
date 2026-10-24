from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import hashlib, json, re

PKG=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0069_R01')
SRC=PKG/'candidate'/'0069_R02_body.md'
OUT=PKG/'candidate'/'0069_R02.docx'
text=SRC.read_text(encoding='utf-8-sig')
doc=Document()
sec=doc.sections[0]
sec.orientation=WD_ORIENT.LANDSCAPE
sec.page_width=Inches(11.69); sec.page_height=Inches(8.27)
sec.top_margin=Inches(.62); sec.bottom_margin=Inches(.62)
sec.left_margin=Inches(.62); sec.right_margin=Inches(.62)
sec.header_distance=Inches(.28); sec.footer_distance=Inches(.28)

CJK='Microsoft JhengHei'
def set_font(run,size,bold=None,italic=None,color=None,name=CJK):
    run.font.name=name; run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    if italic is not None: run.italic=italic
    if color: run.font.color.rgb=RGBColor(*color)
    rpr=run._element.get_or_add_rPr()
    rfonts=rpr.rFonts
    if rfonts is None:
        rfonts=OxmlElement('w:rFonts'); rpr.insert(0,rfonts)
    for attr in ('ascii','hAnsi','eastAsia','cs'): rfonts.set(qn('w:'+attr),name)

def set_cell_shading(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)

def set_cell_margins(cell,top=80,start=85,bottom=80,end=85):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn('w:'+m))
        if node is None: node=OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

def add_inline(p,s,size,bold=False,italic=False,color=None):
    # Preserve readable emphasis for bold, italic and inline-code runs.
    pat=re.compile(r'(\*\*.*?\*\*|(?<!\*)\*[^*]+\*(?!\*)|`[^`]+`)')
    pos=0
    for m in pat.finditer(s):
        if m.start()>pos:
            r=p.add_run(s[pos:m.start()]); set_font(r,size,bold,italic,color)
        tok=m.group(0)
        if tok.startswith('**'):
            r=p.add_run(tok[2:-2]); set_font(r,size,True,italic,color)
        elif tok.startswith('`'):
            r=p.add_run(tok[1:-1]); set_font(r,size,bold,italic,(60,75,95))
        else:
            r=p.add_run(tok[1:-1]); set_font(r,size,bold,True,color)
        pos=m.end()
    if pos<len(s):
        r=p.add_run(s[pos:]); set_font(r,size,bold,italic,color)

def style_paragraph(p,size=10.5,after=4,before=0,spacing=1.08):
    pf=p.paragraph_format; pf.space_after=Pt(after); pf.space_before=Pt(before); pf.line_spacing=spacing
    for r in p.runs: set_font(r,size)

def parse_row(line):
    s=line.strip()
    if s.startswith('|'): s=s[1:]
    if s.endswith('|'): s=s[:-1]
    return [c.strip().replace('\\|','|') for c in s.split('|')]

def add_table(rows):
    if not rows: return
    ncols=max(len(r) for r in rows)
    rows=[r+['']*(ncols-len(r)) for r in rows]
    table=doc.add_table(rows=1,cols=ncols)
    table.style='Table Grid'; table.autofit=False
    tblPr=table._tbl.tblPr
    layout=tblPr.find(qn('w:tblLayout'))
    if layout is None: layout=OxmlElement('w:tblLayout'); tblPr.append(layout)
    layout.set(qn('w:type'),'fixed')
    usable=11.69-.62-.62
    maxlens=[]
    for c in range(ncols):
        vals=[re.sub(r'[`*_]','',r[c]) for r in rows]
        maxlens.append(max(5,min(48,max((len(v) for v in vals),default=8)))**0.48)
    floor=min(.52,usable/ncols*.82)
    remain=usable-floor*ncols
    weights=[max(.25,w-.48) for w in maxlens]; sw=sum(weights)
    widths=[floor+remain*w/sw for w in weights]
    for row_idx,vals in enumerate(rows):
        row=table.rows[0] if row_idx==0 else table.add_row()
        trPr=row._tr.get_or_add_trPr(); cant=OxmlElement('w:cantSplit'); trPr.append(cant)
        if row_idx==0:
            repeat=OxmlElement('w:tblHeader'); repeat.set(qn('w:val'),'true'); trPr.append(repeat)
        for i,val in enumerate(vals):
            cell=row.cells[i]; cell.width=Inches(widths[i]); table.columns[i].width=Inches(widths[i])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(1); p.paragraph_format.line_spacing=1.0
            add_inline(p,val,9.5,bold=(row_idx==0))
            if row_idx==0:
                set_cell_shading(cell,'D9EAF7')
            elif row_idx%2==0:
                set_cell_shading(cell,'F5F8FA')
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

def add_heading(level,s):
    p=doc.add_paragraph(style=f'Heading {min(level,3)}')
    p.paragraph_format.keep_with_next=True
    sizes={1:17,2:14,3:11.5,4:10.5}
    add_inline(p,s,sizes[level],bold=True,color=(31,78,121))
    p.paragraph_format.space_before=Pt(9 if level>1 else 3); p.paragraph_format.space_after=Pt(4)
    return p

# Global styles and accessible running identification.
normal=doc.styles['Normal']; normal.font.name=CJK; normal.font.size=Pt(10.5)
normal._element.rPr.rFonts.set(qn('w:eastAsia'),CJK)
for nm,size in [('Heading 1',17),('Heading 2',14),('Heading 3',11.5)]:
    st=doc.styles[nm]; st.font.name=CJK; st.font.size=Pt(size); st.font.bold=True
    st._element.rPr.rFonts.set(qn('w:eastAsia'),CJK)
header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
add_inline(header,'TTQS ONE｜DOCX 0069｜候選 R02｜I16-E01／方框 076',8.5,color=(90,100,110))
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('候選 SAMPLE／SYNTHETIC　　頁 '); set_font(r,8.5,color=(90,100,110))
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

lines=text.splitlines(); i=0; in_code=False; code=[]; table_count=0; paragraph_count=0
while i<len(lines):
    raw=lines[i]; line=raw.strip()
    if line.startswith('```'):
        if not in_code: in_code=True; code=[]
        else:
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.16); p.paragraph_format.space_after=Pt(6)
            r=p.add_run('\n'.join(code)); set_font(r,10.5,name='Microsoft JhengHei')
            in_code=False; paragraph_count+=1
        i+=1; continue
    if in_code: code.append(raw); i+=1; continue
    if not line or line=='---': i+=1; continue
    if line.startswith('|'):
        block=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            row=parse_row(lines[i])
            if row and not all(re.fullmatch(r':?-{2,}:?',x.replace(' ','')) for x in row): block.append(row)
            i+=1
        add_table(block); table_count+=1; continue
    m=re.match(r'^(#{1,4})\s+(.*)$',line)
    if m: add_heading(len(m.group(1)),m.group(2)); i+=1; continue
    if line.startswith('>'):
        content=line[1:].strip(); p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.18); p.paragraph_format.right_indent=Inches(.1)
        p.paragraph_format.space_after=Pt(3); add_inline(p,content,10.5,color=(75,75,75)); paragraph_count+=1; i+=1; continue
    if re.match(r'^[-*]\s+',line):
        p=doc.add_paragraph(style='List Bullet'); add_inline(p,re.sub(r'^[-*]\s+','',line),10.5); style_paragraph(p,10.5,2); paragraph_count+=1; i+=1; continue
    if re.match(r'^\d+\.\s+',line):
        p=doc.add_paragraph(style='List Number'); add_inline(p,re.sub(r'^\d+\.\s+','',line),10.5); style_paragraph(p,10.5,2); paragraph_count+=1; i+=1; continue
    p=doc.add_paragraph(); add_inline(p,line,10.5); style_paragraph(p,10.5,4); paragraph_count+=1; i+=1

if in_code: raise RuntimeError('Unclosed fenced code block')
# Core properties distinguish the candidate from accepted association evidence.
doc.core_properties.title='訓練過程管制程序與實施紀錄｜SAMPLE／SYNTHETIC 候選'
doc.core_properties.subject='DOCX 0069 / I16-E01 / official checkbox 076'
doc.core_properties.author='TTQS ONE candidate builder'
doc.core_properties.comments='SAMPLE/SYNTHETIC only; not association approval or real evidence.'
doc.save(OUT)
sha=hashlib.sha256(OUT.read_bytes()).hexdigest()
receipt={'doc_id':'0069','revision':'R02','input_body_sha256':hashlib.sha256(SRC.read_bytes()).hexdigest(),'docx_path':str(OUT),'docx_sha256':sha,'docx_bytes':OUT.stat().st_size,'tables':table_count,'paragraphs':paragraph_count,'orientation':'A4 landscape','min_body_font_pt':10.5,'min_table_font_pt':9.5,'render_engine_prepared':'Microsoft Word COM + Poppler; rendering receipt pending'}
(PKG/'qa'/'build_receipt_R02.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(receipt,ensure_ascii=False,indent=2))
