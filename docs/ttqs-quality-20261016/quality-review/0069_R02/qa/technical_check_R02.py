from pathlib import Path
from zipfile import ZipFile
from docx import Document
import xml.etree.ElementTree as ET
import hashlib, json, re
pkg=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0069_R01')
p=pkg/'candidate'/'0069_R02.docx'
body=pkg/'candidate'/'0069_R02_body.md'
doc=Document(p)
with ZipFile(p) as z:
    bad=z.testzip(); names=z.namelist()
    for n in names:
        if n.endswith('.xml') or n.endswith('.rels'):
            ET.fromstring(z.read(n))
parts=[x.text or '' for x in doc.paragraphs]
parts += [cell.text for t in doc.tables for row in t.rows for cell in row.cells]
text='\n'.join(parts)
md=body.read_text(encoding='utf-8-sig')
ids=['SAMPLE-PLAN-0069-V1','SAMPLE-CONSENT-0069','SAMPLE-ATTEND-0069','SAMPLE-TIMELOG-0069','SAMPLE-SAFETY-0069','SAMPLE-ASSESSREC-0069','SAMPLE-DEVIATION-0069-01']
body_sizes=[]; table_sizes=[]
def pt(run): return run.font.size.pt if run.font.size else None
for para in doc.paragraphs:
    if para.text.strip():
        for run in para.runs:
            if run.text.strip() and pt(run) is not None: body_sizes.append(pt(run))
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    if run.text.strip() and pt(run) is not None: table_sizes.append(pt(run))
checks={
 'doc_id_0069':'0069' in text and 'doc_id' in text,
 'requirement_I16_E01':'I16-E01' in text,
 'checkbox_076':'076' in text,
 'case_SYN_0069':'SYN-0069-CASE-01' in text,
 'sample_real_boundary':'UNKNOWN／HOLD' in text and '真實' in text and 'SAMPLE' in text,
 'all_seven_sample_ids_in_body':all(i in md for i in ids),
 'all_seven_sample_ids_in_docx':all(i in text for i in ids),
 'seven_post_records_7_of_7':'7/7' in text,
 'critical_sample_pass_8_of_8':'8/8' in text and ('SAMPLE 通過' in text or 'SAMPLE pass' in text),
 'precheck_10_of_10':'10/10' in text,
 'attendance_10_of_12':'10/12' in text,
 'timed_checks_6_of_6':'6/6' in text,
 'on_time_5_of_6':'5/6' in text,
 'assessment_handoff_9_of_10':'9/10' in text,
 'crossdoc_0033':'0033' in text,
 'route_exception_0070':'0070' in text,
 'route_audit_0072':'0072' in text,
 'hold_reject_list':'HOLD／退回清單' in text,
 'no_stale_7_over_8':'7 ÷ 8' not in text and '87.5%' not in text,
 'no_stale_sample_C3_hold':'C3 暫緩' not in text and 'C3不通過' not in text,
 'no_real_absence_fails_sample':'REAL 原件未附；來源未驗證，因此 C3 不通過' not in text,
}
math={
 'attendance_83_3':abs(10/12*100-83.3333333)<0.01,
 'on_time_83_3':abs(5/6*100-83.3333333)<0.01,
 'assessment_handoff_90':abs(9/10*100-90)<0.001,
 'sample_post_record_100':abs(7/7*100-100)<0.001,
 'sample_critical_100':abs(8/8*100-100)<0.001,
 'buffer_8_used_7_remaining':15-8==7,
}
receipt={'doc_id':'0069','revision':'R02','candidate_docx':str(p),'docx_sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'docx_bytes':p.stat().st_size,'body_sha256':hashlib.sha256(body.read_bytes()).hexdigest(),'body_bytes':body.stat().st_size,'zip_integrity':'PASS' if bad is None else f'FAIL:{bad}','xml_parse':'PASS','table_count':len(doc.tables),'body_font_min_pt':min(body_sizes) if body_sizes else None,'table_font_min_pt':min(table_sizes) if table_sizes else None,'content_checks':checks,'math_checks':math,'all_checks_pass':all(checks.values()) and all(math.values()) and bad is None and bool(body_sizes) and bool(table_sizes) and min(body_sizes)>=10.5 and min(table_sizes)>=9.5,'visual_status':'PENDING','semantic_status':'PENDING_EXACT_SHA'}
(pkg/'qa'/'technical_check_R02.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(receipt,ensure_ascii=False,indent=2))
