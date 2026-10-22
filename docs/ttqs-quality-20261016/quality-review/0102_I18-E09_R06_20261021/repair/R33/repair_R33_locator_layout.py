from pathlib import Path
from copy import deepcopy
from hashlib import sha256
from datetime import datetime
from zoneinfo import ZoneInfo
import json, zipfile
from docx import Document

root = Path(r'E:\TTQS\_r09_sync_0104_20261022\docs\ttqs-quality-20261016\quality-review\0102_I18-E09_R06_20261021')
source = root/'candidate'/'0102__技能競賽獲獎證據取得與查核工作紀錄_R32.docx'
out = root/'candidate'/'0102__技能競賽獲獎證據取得與查核工作紀錄_R33.docx'
prev_receipt = root/'repair_R32_font_floor_20261023'/'repair_receipt_R32.json'
receipt_path = root/'repair_R33_locator_layout_20261023'/'repair_receipt_R33.json'
script_path = root/'repair_R33_locator_layout_20261023'/'repair_R33_locator_layout.py'
assert source.is_file() and not out.exists(), 'R32 must exist and R33 output must be new'
def digest(p): return sha256(p.read_bytes()).hexdigest().upper()
source_sha = digest(source)
assert source_sha == '2322EF89812DB2249B2A98729D4370C82E38F3BA646F340B6C4AD36B57615FB9', source_sha
base_receipt=json.loads(prev_receipt.read_text(encoding='utf-8'))
doc=Document(source)

def set_cell_text(cell, value):
    assert len(cell.paragraphs)==1, f'Unexpected multi-paragraph cell: {cell.text!r}'
    p=cell.paragraphs[0]
    if not p.runs: p.add_run()
    p.runs[0].text=value
    for run in p.runs[1:]: run.text=''

# Replace only table-location text in the verified replacement-register rows.
register=doc.tables[39]
updates={
  3:'Form A Table 9／§4.2 已填 SAMPLE：CASE-01 欄（learner_ref）／§12.1 案例登錄表（case_id × CASE-01）',
  23:'Form A Table 9／§12.1 案例登錄表（internal_review_date × CASE-01）／§10.1 CASE-01',
  24:'§4／§12.1 案例登錄表（owner_role × CASE-01）／§10.1 CASE-01',
  29:'Form A Table 9／§12.1 案例登錄表（internal_review_date × CASE-02）／§10.2 CASE-02',
  30:'§4／§12.1 案例登錄表（owner_role × CASE-02）／§10.2 CASE-02',
  38:'Form A Table 9 row 17／§12.1 案例登錄表（evidence_state × CASE-01）',
  44:'Form A Table 9 row 17／§12.1 案例登錄表（evidence_state × CASE-02）',
  45:'Form F Table 28／Table 34 row 7／§12.1 案例登錄表（case_status × CASE-01）／§10.1',
  46:'Form F Table 29／Table 34 row 7／§12.1 案例登錄表（case_status × CASE-02）／§10.2',
}
assert len(register.rows)>max(updates)
for ri,new in updates.items():
    cells=register.rows[ri].cells
    fact_id=cells[0].text.strip()
    assert fact_id.startswith('F-0102-'), (ri,fact_id)
    set_cell_text(cells[3],new)

# Update current-candidate labels only; retain earlier R32 history as history.
set_cell_text(doc.tables[0].rows[5].cells[1], doc.tables[0].rows[5].cells[1].text.replace('R32','R33'))
set_cell_text(doc.tables[0].rows[9].cells[1], doc.tables[0].rows[9].cells[1].text.replace('R32','R33'))
version_table=doc.tables[45]
assert version_table.rows[-1].cells[0].text.strip()=='R32'
version_table._tbl.append(deepcopy(version_table.rows[-1]._tr))
newrow=version_table.rows[-1]
assert len(newrow.cells)==5
new_values=[
 'R33','2026-10-23',
 '修正清冊中案例登錄表的表號／列號定位，改用 §12.1 欄位 × 個案語意座標；調整 15.5 末段分頁。',
 '文件維護角色',
 '只修定位與分頁；未變更案例值、日期、狀態、公式或判定；SAMPLE/SYNTHETIC 候選；未正式接受。'
]
for cell,value in zip(newrow.cells,new_values): set_cell_text(cell,value)

# Keep the final heading and its opening sentence with the closing checklist on one page.
heading=[p for p in doc.paragraphs if p.text.strip()=='15.5 最後提醒']
intro=[p for p in doc.paragraphs if p.text.strip()=='這份文件最重要的不是表格，而是停下來的地方：']
assert len(heading)==1 and len(intro)==1
heading[0].paragraph_format.page_break_before=True
heading[0].paragraph_format.keep_with_next=True
intro[0].paragraph_format.keep_with_next=True

# Refresh current revision metadata/footer; preserve historical entries.
for section in doc.sections:
    for p in section.footer.paragraphs:
        for run in p.runs:
            run.text=run.text.replace('候選 R32','候選 R33')
props=doc.core_properties
props.title=props.title.replace('R32','R33')
props.subject=props.subject.replace('R32','R33')
doc.save(out)

# Reopen, validate targeted changes, and preserve values outside the changed locators.
check=Document(out)
assert check.tables[0].rows[5].cells[1].text.find('R33')>=0
assert check.tables[0].rows[9].cells[1].text.find('R33')>=0
assert check.tables[45].rows[-1].cells[0].text.strip()=='R33'
assert len(check.tables[45].rows)==len(version_table.rows)
assert all(check.tables[39].rows[ri].cells[3].text==value for ri,value in updates.items())
assert check.tables[39].rows[3].cells[4].text=='P-SYN-01'
assert check.tables[39].rows[23].cells[4].text=='CASE-01 internal_review_date：2026-10-26（SAMPLE；未送通知，獨立設定的內部追蹤日，非外部期限）'
assert check.tables[39].rows[29].cells[4].text=='CASE-02 internal_review_date：2026-10-20（SAMPLE；未送通知，獨立設定的內部追蹤日，非外部期限）'
assert check.tables[39].rows[38].cells[4].text=='CASE-01 evidence_state：RECEIVED_UNVERIFIED（SAMPLE）'
assert check.tables[39].rows[45].cells[4].text=='案件決定：HOLD'
assert check.tables[39].rows[46].cells[4].text=='案件決定：HOLD'
assert next(p for p in check.paragraphs if p.text.strip()=='15.5 最後提醒').paragraph_format.page_break_before
with zipfile.ZipFile(out) as z: assert z.testzip() is None

# Font-floor scan: direct sizes; inherited styles were covered in R32 and only cloned rows are new.
body_sizes=[]; table_sizes=[]
for p in check.paragraphs:
    for r in p.runs:
        if r.text.strip() and r.font.size: body_sizes.append(r.font.size.pt)
for t in check.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    if r.text.strip() and r.font.size: table_sizes.append(r.font.size.pt)
assert min(body_sizes)>=10.5, min(body_sizes)
assert min(table_sizes)>=9.5, min(table_sizes)

sha=digest(out)
receipt={
 'schema':'ttqs.docx_targeted_locator_layout_repair.v1',
 'work_id':'0102_R33_LOCATOR_AND_CLOSING_LAYOUT_REPAIR',
 'doc_id':'0102','official_requirement_id':'I18-E09','project_blueprint_index':110,
 'printed_official_checkbox_serial':'UNKNOWN',
 'source_candidate':'R32','source_docx_sha256':source_sha,
 'source_receipt_sha256':digest(prev_receipt),
 'candidate_revision':'R33','candidate_path':str(out),'candidate_docx_sha256':sha,
 'source_control_commit':base_receipt['source_control_commit'],
 'control_blobs':base_receipt['control_blobs'],
 'source_packet_sha256':base_receipt['source_packet_sha256'],
 'unique_blueprint_sha256':base_receipt['unique_blueprint_sha256'],
 'repair':{
  'incorrect_locator_rows':{str(k):{'fact_id':register.rows[k].cells[0].text.strip(),'new_locator':v} for k,v in updates.items()},
  'locator_policy':'Use section + named field × case column; remove ambiguous global table ordinal and row number.',
  'closing_layout':'Page break before 15.5 so heading, opening sentence, and closing checklist begin together.',
  'preserved':'All synthetic values, dates, status, formulas, case conclusions, official mapping and all other text; R32 history retained.',
  'body_minimum_pt':min(body_sizes),'table_minimum_pt':min(table_sizes),
  'zip_integrity':'PASS','docx_reopen':'PASS'
 },
 'synthetic_boundary':base_receipt['synthetic_boundary'],
 'technical_status':'OOXML_REOPEN_PASS; ZIP_INTEGRITY_PASS; TEXT_SCOPE_ASSERTIONS_PASS; FONT_FLOOR_PASS; PENDING_NATIVE_WORD_EXPORT_AND_VISUAL_QA',
 'independent_semantic_review':'PENDING_EXACT_R33_SHA','formal_acceptance':'NOT_GRANTED',
 'script_sha256':digest(script_path),
 'created_at_local':datetime.now(ZoneInfo('Asia/Taipei')).isoformat()
}
receipt_path.write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'candidate':str(out),'sha256':sha,'receipt':str(receipt_path),'tables':len(check.tables),'paragraph_min_pt':min(body_sizes),'table_min_pt':min(table_sizes),'zip':'PASS','text_scope':'PASS'},ensure_ascii=False))
