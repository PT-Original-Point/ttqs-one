from pathlib import Path
import re, json, hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

lane=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0102_I18-E09_R01_20261021')
source=lane/'writer'/'opencode_writer_manuscript_0102_R01_repaired.md'
outdir=lane/'candidate'; outdir.mkdir(exist_ok=True)
out=outdir/'0102__技能競賽獲獎證據取得與查核工作紀錄_R02.docx'
if out.exists(): raise SystemExit('REFUSE overwrite existing candidate')
text=source.read_text(encoding='utf-8')
doc=Document()
sec=doc.sections[0]
sec.orientation=WD_ORIENT.LANDSCAPE
sec.page_width=Inches(11.69); sec.page_height=Inches(8.27)
sec.top_margin=Inches(0.72); sec.bottom_margin=Inches(0.70); sec.left_margin=Inches(0.72); sec.right_margin=Inches(0.72)
sec.header_distance=Inches(0.30); sec.footer_distance=Inches(0.30)

def set_font(run, name='Microsoft JhengHei', size=10.2, color='202A35', bold=None, italic=None):
    run.font.name=name; run.font.size=Pt(size); run.font.color.rgb=RGBColor.from_string(color)
    if bold is not None: run.bold=bold
    if italic is not None: run.italic=italic
    rpr=run._element.get_or_add_rPr(); rfonts=rpr.rFonts
    if rfonts is None:
        rfonts=OxmlElement('w:rFonts'); rpr.insert(0,rfonts)
    rfonts.set(qn('w:eastAsia'),name); rfonts.set(qn('w:ascii'),name); rfonts.set(qn('w:hAnsi'),name)

def set_style_font(style, size, bold=False, color='202A35'):
    style.font.name='Microsoft JhengHei'; style.font.size=Pt(size); style.font.bold=bold; style.font.color.rgb=RGBColor.from_string(color)
    rpr=style._element.get_or_add_rPr(); rfonts=rpr.rFonts
    if rfonts is None: rfonts=OxmlElement('w:rFonts'); rpr.insert(0,rfonts)
    rfonts.set(qn('w:eastAsia'),'Microsoft JhengHei'); rfonts.set(qn('w:ascii'),'Microsoft JhengHei'); rfonts.set(qn('w:hAnsi'),'Microsoft JhengHei')

normal=doc.styles['Normal']; set_style_font(normal,10.2); normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.10
for nm,sz in [('Title',24),('Heading 1',17),('Heading 2',13),('Heading 3',11.5),('Heading 4',10.5)]:
    st=doc.styles[nm]; set_style_font(st,sz,True,'000000'); st.paragraph_format.keep_with_next=True; st.paragraph_format.space_before=Pt(12 if nm=='Heading 1' else 8); st.paragraph_format.space_after=Pt(5)
# Remove the built-in blue bottom border from Word's Title style.
title_style=doc.styles['Title']._element
title_ppr=title_style.find(qn('w:pPr'))
if title_ppr is not None:
    for child in list(title_ppr):
        if child.tag==qn('w:pBdr'): title_ppr.remove(child)
# Body headings remain plain black, with no underline or decorative rule.

header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=header.add_run('SAMPLE / SYNTHETIC  |  示範資料，非原始點協會真實紀錄'); set_font(r,size=8.3,color='9C2F2F',bold=True)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('0102  |  I18-E09  |  候選 R02  |  第 '); set_font(r,size=8,color='59636F')
def field(par,name):
    run=par.add_run(); set_font(run,size=8,color='59636F')
    begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin')
    inst=OxmlElement('w:instrText'); inst.set(qn('xml:space'),'preserve'); inst.text=' '+name+' '
    separate=OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'),'separate')
    t=OxmlElement('w:t'); t.text='1'
    end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end')
    run._r.extend([begin,inst,separate,t,end])
field(footer,'PAGE'); r=footer.add_run(' 頁'); set_font(r,size=8,color='59636F')

def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)
def cell_margins(cell,top=90,start=100,bottom=90,end=100):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn('w:'+m))
        if node is None: node=OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')
def set_table_borders(table):
    tblPr=table._tbl.tblPr; borders=tblPr.first_child_found_in('w:tblBorders')
    if borders is None: borders=OxmlElement('w:tblBorders'); tblPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag='w:'+edge; el=borders.find(qn(tag))
        if el is None: el=OxmlElement(tag); borders.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'5'); el.set(qn('w:space'),'0'); el.set(qn('w:color'),'D9E1E8')
def parse_cells(line):
    s=line.strip()
    if s.startswith('|'): s=s[1:]
    if s.endswith('|'): s=s[:-1]
    return [x.strip().replace('\\|','|') for x in s.split('|')]
def add_inline(p,s,base=10.2,color='202A35'):
    # Basic Markdown emphasis/code, preserving Chinese text as editable Word runs.
    pat=re.compile(r'(\*\*.*?\*\*|`[^`]+`|\*[^*]+\*)')
    pos=0
    for m in pat.finditer(s):
        if m.start()>pos:
            rr=p.add_run(s[pos:m.start()]); set_font(rr,size=base,color=color)
        token=m.group(0)
        if token.startswith('**'):
            rr=p.add_run(token[2:-2]); set_font(rr,size=base,color=color,bold=True)
        elif token.startswith('`'):
            rr=p.add_run(token[1:-1]); set_font(rr,name='Consolas',size=max(8.2,base-0.4),color='334155')
        else:
            rr=p.add_run(token[1:-1]); set_font(rr,size=base,color=color,italic=True)
        pos=m.end()
    if pos<len(s):
        rr=p.add_run(s[pos:]); set_font(rr,size=base,color=color)
def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr(); tblHeader=OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'),'true'); trPr.append(tblHeader)
def add_table(rows):
    n=max(len(x) for x in rows); rows=[x+['']*(n-len(x)) for x in rows]
    table=doc.add_table(rows=1, cols=n); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
    table.allow_autofit=False; set_table_borders(table)
    # Allocate more width to explanatory columns while keeping IDs and states compact.
    headers=rows[0]; available=10.25
    weights=[]
    for j,h in enumerate(headers):
        hl=h.lower().strip(); maxlen=max([len(row[j]) for row in rows] or [1]);
        if any(k in hl for k in ('fact_id','case_id','doc_id','日期','時間','期限','狀態','結果','是否','編號','項次','序號')): w=0.66
        elif any(k in hl for k in ('位置','來源','說明','用途','下一步','責任','資料','填法','要求','內容','做法','原因','欄位')): w=1.28
        else: w=0.95
        w*=min(1.75,max(0.82,(maxlen/18)**0.22))
        weights.append(w)
    widths=[available*w/sum(weights) for w in weights]
    for ri,rowdata in enumerate(rows):
        row=table.rows[0] if ri==0 else table.add_row()
        if ri==0: set_repeat_table_header(row)
        for j,val in enumerate(rowdata):
            cell=row.cells[j]; cell.width=Inches(widths[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell_margins(cell)
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(1); p.paragraph_format.line_spacing=1.02
            short=('日期' in headers[j] or '時間' in headers[j] or '狀態' in headers[j] or '是否' in headers[j] or headers[j].lower().strip() in ('fact_id','case_id','doc_id','結果','項次','序號'))
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if short else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p,val,base=(8.2 if n>=8 else 8.8 if n>=6 else 9.1 if n>=4 else 9.4),color='FFFFFF' if ri==0 else '273444')
            if ri==0: shade(cell,'1F4E78')
            elif ri%2==0: shade(cell,'F3F7FA')
            else: shade(cell,'FFFFFF')
        trPr=row._tr.get_or_add_trPr()
        # Do not set fixed heights. Rows can expand and split naturally if needed.
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

lines=text.splitlines(); i=0; code=False; code_lines=[]; seen_title=False; table_count=0; heading_count=0
while i<len(lines):
    line=lines[i]; s=line.strip()
    if s.startswith('```'):
        if not code:
            code=True; code_lines=[]
        else:
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(0.22); p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(7)
            pp=p._p.get_or_add_pPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'F1F3F5'); pp.append(shd)
            for ix,cl in enumerate(code_lines):
                if ix: p.add_run().add_break()
                rr=p.add_run(cl); set_font(rr,name='Consolas',size=8.8,color='334155')
            code=False
        i+=1; continue
    if code:
        code_lines.append(line); i+=1; continue
    if not s:
        i+=1; continue
    if s.startswith('|') and i+1<len(lines) and re.match(r'^\s*\|?\s*:?-{3,}',lines[i+1]):
        rows=[parse_cells(lines[i])]; i+=2
        while i<len(lines) and lines[i].strip().startswith('|'):
            rows.append(parse_cells(lines[i])); i+=1
        if len(rows)>1: add_table(rows); table_count+=1
        continue
    if re.fullmatch(r'\s*[-*_]{3,}\s*',line):
        i+=1; continue
    hm=re.match(r'^(#{1,6})\s+(.*)$',s)
    if hm:
        level=len(hm.group(1)); title=hm.group(2).strip(); title=re.sub(r'\s+#+\s*$','',title)
        if level==1 and not seen_title:
            p=doc.add_paragraph(style='Title'); add_inline(p,title,base=24,color='000000'); seen_title=True
        else:
            st='Heading '+str(min(level,4)); p=doc.add_paragraph(style=st); add_inline(p,title,base={1:17,2:13,3:11.5,4:10.8,5:10.3,6:10.3}.get(level,10.3),color='000000'); heading_count+=1
        i+=1; continue
    if s.startswith('>'):
        quote=[]
        while i<len(lines) and lines[i].strip().startswith('>'):
            quote.append(lines[i].strip()[1:].lstrip()); i+=1
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(0.18); p.paragraph_format.right_indent=Inches(0.14); p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(7)
        pPr=p._p.get_or_add_pPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'F7F4EA'); pPr.append(shd)
        add_inline(p,' '.join(quote),base=9.6,color='4B5563'); continue
    bm=re.match(r'^[-*+]\s+(.*)$',s)
    nm=re.match(r'^\d+[.)]\s+(.*)$',s)
    if bm:
        p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(2); add_inline(p,bm.group(1),base=10.0)
        i+=1; continue
    if nm:
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(0.30); p.paragraph_format.first_line_indent=Inches(-0.26); p.paragraph_format.space_after=Pt(2); add_inline(p,s,base=10.0)
        i+=1; continue
    if s.startswith('**SAMPLE / SYNTHETIC'):
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(14)
        pp=p._p.get_or_add_pPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'FCE8E6'); pp.append(shd)
        add_inline(p,s,base=12,color='9C2F2F'); i+=1; continue
    # Join wrapped source lines into one paragraph until a block boundary.
    para=[s]; i+=1
    while i<len(lines):
        x=lines[i].strip()
        if not x or x.startswith(('#','|','>','```')) or re.match(r'^[-*+]\s+',x) or re.match(r'^\d+[.)]\s+',x) or re.fullmatch(r'[-*_]{3,}',x): break
        para.append(x); i+=1
    p=doc.add_paragraph(); p.paragraph_format.widow_control=True
    add_inline(p,' '.join(para),base=10.2)

props=doc.core_properties; props.title='技能競賽獲獎證據取得與查核工作紀錄'; props.subject='0102 / I18-E09 SAMPLE SYNTHETIC candidate'; props.author=''; props.last_modified_by=''
doc.save(out)
raw=out.read_bytes(); receipt={'doc_id':'0102','official_requirement_id':'I18-E09','source_manuscript_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'docx_path':str(out),'docx_sha256':hashlib.sha256(raw).hexdigest(),'bytes':len(raw),'paragraphs':len(doc.paragraphs),'tables':table_count,'headings':heading_count,'orientation':'A4 landscape','body_font':'Microsoft JhengHei 10.2pt','table_font_rule':'8.2-9.4pt depending column count','repeating_table_headers':table_count,'sample_boundary':'header on every page and first-page notice','status':'DOCX_BUILT_PENDING_RENDER_AND_VISUAL_INSPECTION'}
(lane/'evidence'/'docx_build_receipt_0102_R02.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(receipt,ensure_ascii=False,indent=2))