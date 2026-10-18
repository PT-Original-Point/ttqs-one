from pathlib import Path
import re, json, hashlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(r"E:\TTQS\_TTQS_ONE_REVIEW_SYNC_0036_R08_20261016\docs\ttqs-quality-20261016\quality-review\0063_R01")
SRC=ROOT/"inputs/0063_MANUSCRIPT_R01.md"
OUT=ROOT/"candidate/0063_R01.docx"
RECEIPT=ROOT/"evidence/build_receipt_R01.json"
FONT="Microsoft JhengHei"
BLUE="17365D"; INK="000000"; ACCENT="1F4E78"; LIGHT="F4F7FA"; GRID="B7C9D6"; TEXT="23313D"; WHITE="FFFFFF"
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def set_run_font(run,size=10.5,color=TEXT,bold=None,italic=None):
    run.font.name=FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),FONT)
    run.font.size=Pt(size); run.font.color.rgb=RGBColor.from_string(color)
    if bold is not None: run.bold=bold
    if italic is not None: run.italic=italic
def add_rich(p,text,size=10.5,color=TEXT,base_bold=False):
    tick=chr(96)
    tokens=re.split(r"(\*\*.*?\*\*|"+tick+r".*?"+tick+r"|\*.*?\*)",text)
    for t in tokens:
        if not t: continue
        if t.startswith("**") and t.endswith("**"):
            r=p.add_run(t[2:-2]); set_run_font(r,size,color,True)
        elif t.startswith(tick) and t.endswith(tick):
            r=p.add_run(t[1:-1]); set_run_font(r,size,"34495E",base_bold)
        elif t.startswith("*") and t.endswith("*"):
            r=p.add_run(t[1:-1]); set_run_font(r,size,color,base_bold,True)
        else:
            r=p.add_run(t); set_run_font(r,size,color,base_bold)
def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn("w:shd"))
    if shd is None: shd=OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"),fill)
def margins(cell,top=80,start=95,bottom=80,end=95):
    tcPr=cell._tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in("w:tcMar")
    if tcMar is None: tcMar=OxmlElement("w:tcMar"); tcPr.append(tcMar)
    for m,v in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node=tcMar.find(qn("w:"+m))
        if node is None: node=OxmlElement("w:"+m); tcMar.append(node)
        node.set(qn("w:w"),str(v)); node.set(qn("w:type"),"dxa")
def borders(table):
    tblPr=table._tbl.tblPr; b=tblPr.first_child_found_in("w:tblBorders")
    if b is None: b=OxmlElement("w:tblBorders"); tblPr.append(b)
    for edge in ("top","left","bottom","right","insideH","insideV"):
        el=b.find(qn("w:"+edge))
        if el is None: el=OxmlElement("w:"+edge); b.append(el)
        el.set(qn("w:val"),"single"); el.set(qn("w:sz"),"5"); el.set(qn("w:space"),"0"); el.set(qn("w:color"),GRID)
def repeat_header(row):
    el=OxmlElement("w:tblHeader"); el.set(qn("w:val"),"true"); row._tr.get_or_add_trPr().append(el)
def cant_split(row): row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
def cells(line):
    s=line.strip()
    if s.startswith("|"): s=s[1:]
    if s.endswith("|"): s=s[:-1]
    return [x.strip().replace(r"\|","|") for x in s.split("|")]
def sep(line): return bool(cells(line)) and all(re.fullmatch(r":?-{3,}:?",c.replace(" ","")) for c in cells(line))
def widths(headers,n,total=17.4):
    h=" ".join(headers)
    if n==2: return [4.5,total-4.5]
    if n==3: return [3.1,9.0,total-12.1]
    if n==4:
        if "功能角色" in h: return [2.7,4.2,5.8,total-12.7]
        return [2.5,4.0,4.2,total-10.7]
    if n==5:
        if "情境" in h: return [2.45,2.55,4.6,2.75,total-12.35]
        if "矯正項" in h: return [2.45,3.0,3.3,2.7,total-11.45]
        if "查詢內容" in h: return [0.9,4.3,3.6,2.7,total-11.5]
        return [1.0,4.0,3.5,2.6,total-11.1]
    return [total/n]*n
def add_table(doc,rows):
    if not rows:return
    n=len(rows[0])
    if n>5: raise ValueError(f"Table has {n} columns")
    if any(len(x)!=n for x in rows): raise ValueError("Inconsistent table row width")
    t=doc.add_table(rows=len(rows),cols=n); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    ws=widths(rows[0],n)
    for ci,w in enumerate(ws): t.columns[ci].width=Cm(w)
    for ri,data in enumerate(rows):
        row=t.rows[ri]; cant_split(row)
        if ri==0: repeat_header(row)
        for ci,val in enumerate(data):
            cell=row.cells[ci]; cell.width=Cm(ws[ci]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP; margins(cell)
            shade(cell,BLUE if ri==0 else (LIGHT if ri%2==0 else WHITE))
            p=cell.paragraphs[0]; p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.08
            add_rich(p,val,9.5,WHITE if ri==0 else TEXT,ri==0)
    borders(t)
    sp=doc.add_paragraph(); sp.paragraph_format.space_after=Pt(1); sp.paragraph_format.space_before=Pt(0)
def add_page_field(p):
    r=p.add_run("頁 "); set_run_font(r,8.5,"667788")
    fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),"PAGE"); p._p.append(fld)
def paragraph_shading(p):
    pPr=p._p.get_or_add_pPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),"EAF1F8"); pPr.append(shd)
    pb=OxmlElement("w:pBdr"); left=OxmlElement("w:left")
    for k,v in (("val","single"),("sz","18"),("space","6"),("color",ACCENT)): left.set(qn("w:"+k),v)
    pb.append(left); pPr.append(pb)
def main():
    previous_sha=sha(OUT) if OUT.exists() else None
    if OUT.exists():
        if not RECEIPT.exists(): raise FileExistsError("Existing candidate has no matching build receipt")
        old_receipt=json.loads(RECEIPT.read_text(encoding="utf-8"))
        if old_receipt.get("candidate_sha256")!=previous_sha: raise FileExistsError("Existing candidate SHA does not match its build receipt")
        hist=ROOT/"evidence/history/build_receipt_R01_pre_style.json"
        hist.parent.mkdir(parents=True,exist_ok=True)
        if not hist.exists(): hist.write_text(json.dumps(old_receipt,ensure_ascii=False,indent=2),encoding="utf-8")
    content=SRC.read_text(encoding="utf-8").strip(); lines=content.splitlines()
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.7); sec.bottom_margin=Cm(1.65); sec.left_margin=Cm(1.8); sec.right_margin=Cm(1.8)
    sec.header_distance=Cm(0.72); sec.footer_distance=Cm(0.72)
    st=doc.styles["Normal"]; st.font.name=FONT; st._element.rPr.rFonts.set(qn("w:eastAsia"),FONT); st.font.size=Pt(10.5); st.font.color.rgb=RGBColor.from_string(TEXT)
    st.paragraph_format.space_after=Pt(5); st.paragraph_format.line_spacing=1.12
    for name,size,color in (("Title",21,INK),("Heading 1",14,INK),("Heading 2",12,INK),("Heading 3",10.8,INK)):
        style=doc.styles[name]; style.font.name=FONT; style._element.rPr.rFonts.set(qn("w:eastAsia"),FONT)
        style.font.size=Pt(size); style.font.bold=True; style.font.color.rgb=RGBColor.from_string(color)
        style.paragraph_format.space_before=Pt(10 if name!="Title" else 1); style.paragraph_format.space_after=Pt(5); style.paragraph_format.keep_with_next=True
    hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; add_rich(hp,"TTQS ONE｜候選示範文件｜SAMPLE／SYNTHETIC",8.5,INK)
    fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; add_rich(fp,"文件 0063｜候選版本，非協會核定制度｜",8.5,INK); add_page_field(fp)
    paragraphs=0; tables=0; headings=0; i=0
    while i<len(lines):
        s=lines[i].strip()
        if not s or s=="---": i+=1; continue
        if s.startswith("|"):
            block=[]
            while i<len(lines) and lines[i].strip().startswith("|"):
                line=lines[i].strip()
                if not sep(line): block.append(cells(line))
                i+=1
            add_table(doc,block); tables+=1; continue
        m=re.match(r"^(#{1,3})\s+(.*)$",s)
        if m:
            level=len(m.group(1)); title=m.group(2)
            if level==1 and not headings:
                p=doc.add_paragraph(style="Title"); add_rich(p,title,21,INK,True); p.paragraph_format.space_after=Pt(10)
            else:
                p=doc.add_paragraph(style=f"Heading {min(level,3)}"); add_rich(p,title,{1:14,2:12,3:10.8}[min(level,3)],INK,True)
            if title.startswith("14.1 空白可重用"):
                p.paragraph_format.page_break_before=True
            headings+=1; paragraphs+=1; i+=1; continue
        if s.startswith(">"):
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(.35); p.paragraph_format.right_indent=Cm(.25); add_rich(p,s.lstrip("> ").strip()); paragraph_shading(p); paragraphs+=1; i+=1; continue
        if s.startswith("- [ ] "):
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(.2); add_rich(p,"☐ "+s[6:]); paragraphs+=1; i+=1; continue
        if s.startswith("- "):
            p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.left_indent=Cm(.55); p.paragraph_format.first_line_indent=Cm(-.2); add_rich(p,s[2:]); paragraphs+=1; i+=1; continue
        if re.match(r"^\d+\.\s+",s):
            p=doc.add_paragraph(style="List Number"); add_rich(p,re.sub(r"^\d+\.\s+","",s)); paragraphs+=1; i+=1; continue
        p=doc.add_paragraph(); p.paragraph_format.widow_control=True; add_rich(p,s)
        if s.startswith("**候選狀態宣告。**") or s.startswith("**使用限制。**"): paragraph_shading(p)
        paragraphs+=1; i+=1
    doc.core_properties.title="查詢服務與持續改善作業指引暨合成演練"
    doc.core_properties.subject="TTQS ONE 文件 0063／I14-E04／官方方框 068；候選示範文件"
    doc.core_properties.author="TTQS ONE 文件品質施工"
    doc.core_properties.keywords="0063, I14-E04, 068, SAMPLE, SYNTHETIC, candidate"
    OUT.parent.mkdir(parents=True,exist_ok=True); doc.save(OUT)
    check=Document(OUT)
    alltext="\n".join([p.text for p in check.paragraphs]+[c.text for t in check.tables for row in t.rows for c in row.cells])
    required=["0063","I14-E04","068","SYN-0063-QI-2026-01","SYN-0060-2026-A01","F-0061-2026-0001","62.5%","75.0%","100.0%","+12.5","2026-11-25","TEMPLATE","SAMPLE","SYNTHETIC","HOLD","REJECT"]
    missing=[x for x in required if x not in alltext]
    if missing: raise AssertionError(f"Missing required content {missing}")
    if "UNIQUE_BLUEPRINT" in alltext or "本回覆第一節" in alltext: raise AssertionError("Worker scaffolding leaked into Word")
    if len(check.tables)!=tables: raise AssertionError("Table count mismatch")
    if any(len(t.columns)>5 for t in check.tables): raise AssertionError("Wide table must be split")
    metrics={"baseline_first_response":5/8*100,"baseline_zero_result":1/8*100,"baseline_clarification":2/8*100,"retest_first_response":6/8*100,"retest_query_intent_with_one_clarification":8/8*100,"retest_zero_result":0/8*100,"percentage_point_change":6/8*100-5/8*100}
    if tuple(round(metrics[k],1) for k in ("baseline_first_response","retest_first_response","retest_query_intent_with_one_clarification","percentage_point_change"))!=(62.5,75.0,100.0,12.5): raise AssertionError(metrics)
    all_sizes=[r.font.size.pt for t in check.tables for row in t.rows for cell in row.cells for p in cell.paragraphs for r in p.runs if r.text.strip()]
    if all_sizes and min(all_sizes)<9.5: raise AssertionError(f"Table font below 9.5 pt: {min(all_sizes)}")
    size=OUT.stat().st_size
    receipt={"project_id":"TTQS_ONE","doc_id":"0063","requirement_id":"I14-E04","official_checkbox_sequence":"068","candidate_revision":"R01","source_manuscript_path":str(SRC),"source_manuscript_sha256":sha(SRC),"candidate_path":str(OUT),"candidate_sha256":sha(OUT),"bytes":size,"paragraphs":paragraphs,"markdown_tables":tables,"docx_tables":len(check.tables),"body_text_chars":len(alltext),"max_table_columns":max(len(t.columns) for t in check.tables),"font_policy":{"body_pt":10.5,"table_pt":9.5,"font":FONT},"page_size":"A4 portrait","math_check":"PASS","math_values":metrics,"content_presence":"PASS","table_font_min_pt":min(all_sizes) if all_sizes else None,"prior_candidate_sha256":previous_sha,"iteration_reason":"Targeted visual repair: widen 11. 矯正項 matrix first column to prevent vertical glyph wrapping and move 14.1 improvement-register template heading/table from an orphaned footer position onto a new page; content unchanged","status":"BUILT_CANDIDATE_PENDING_RENDER_AND_INDEPENDENT_REVIEW"}
    RECEIPT.parent.mkdir(parents=True,exist_ok=True); RECEIPT.write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(receipt,ensure_ascii=False))
if __name__=="__main__": main()
