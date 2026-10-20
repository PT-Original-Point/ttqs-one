from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
import hashlib,json,shutil
base=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0092_R01_OPENCODE_WRITER_R01')
src=base/'0092_R02_SYNTHETIC_CANDIDATE.docx'
out=base/'0092_R03_SYNTHETIC_CANDIDATE.docx'
expected='33392cdbe71a56bfd57833d2ed7e9c26dd0e869ae86f66901fb36821c4d8baf2'
if hashlib.sha256(src.read_bytes()).hexdigest()!=expected: raise RuntimeError('R02 source SHA mismatch')
if out.exists(): raise FileExistsError(out)
d=Document(src)

def set_p_text(p,text):
 runs=p.runs
 if runs:
  runs[0].text=text
  for r in runs[1:]: r.text=''
 else: p.add_run(text)

def set_cell_text(cell,text):
 p=cell.paragraphs[0]
 set_p_text(p,text)
 for extra in cell.paragraphs[1:]:
  extra._element.getparent().remove(extra._element)

def all_paragraphs():
 for child in d._element.body.iterchildren():
  tag=child.tag.rsplit('}',1)[-1]
  if tag=='p': yield Paragraph(child,d)
  elif tag=='tbl':
   t=Table(child,d)
   for row in t.rows:
    for cell in row.cells:
     for p in cell.paragraphs: yield p

def replace_once(old,new):
 hits=[]
 for p in all_paragraphs():
  if old in p.text:
   hits.append(p)
 if len(hits)!=1: raise RuntimeError(f'Expected one occurrence for {old!r}; got {len(hits)}')
 p=hits[0]
 set_p_text(p,p.text.replace(old,new))
 return p

# Remove unsourced institutional/category claims from explanations.
replace_once('原始點協會目前設有這些部門','本會目前設有這些部門')
replace_once('官方佐證資料表第 6 頁、I17d「成果評估」的參考佐證列有「感謝函。」。本件對應','官方佐證資料表第 6 頁載有「感謝函。」。本件對應')
# Replace vague status language with an explicit controlled disposition.
replace_once('維持待補或不納入並記錄原因','維持 HOLD，列出待補資料、責任角色與後續決定；或依權限不採認並記錄理由')
# Repair F001/F002 body locators to actual locations in this document.
trace_table=next(t for t in d.tables if len(t.rows)>20 and any('0092-F001-E2' in c.text for c in t.rows[1].cells))
for row in trace_table.rows[1:]:
 label=' '.join(c.text for c in row.cells)
 if '0092-F001-E2' in label: set_cell_text(row.cells[1],'§4.2 案例識別；§5.3 案例識別欄')
 if '0092-F002-E2' in label: set_cell_text(row.cells[1],'§4.2 寄件方；§5.3 寄件方／承辦角色欄')
# Add stable role labels without implying association structure or authority.
role_intro=next(p for p in d.paragraphs if p.text.startswith('下表是建議的職能分工範例'))
set_p_text(role_intro,role_intro.text+' 本件用 ROLE-A 至 ROLE-D 作為表單職能代碼；代碼只表示工作角色，不代表本會設有對應部門、已指派人員或已授權。實際使用前須由有權者核准角色對照；未核准時標 UNKNOWN。')
role_table=next(t for t in d.tables if any('職能角色（示範）' in c.text for c in t.rows[0].cells))
set_cell_text(role_table.rows[0].cells[0],'職能代碼／角色（示範）')
role_codes={'承辦角色':'ROLE-A / 承辦角色','覆核角色':'ROLE-B / 覆核角色','保管角色':'ROLE-C / 保管角色','授權決定者':'ROLE-D / 授權決定者'}
for row in role_table.rows[1:]:
 for old,new in role_codes.items():
  if row.cells[0].text.strip()==old: set_cell_text(row.cells[0],new)
# Bind the synthetic example's role values to the explanatory role codes.
for p in all_paragraphs():
 s=p.text
 if '角色甲（示例承辦角色）' in s: set_p_text(p,s.replace('角色甲（示例承辦角色）','角色甲（示例承辦角色；ROLE-A）'))
 elif '角色乙（示例覆核角色）' in s: set_p_text(p,s.replace('角色乙（示例覆核角色）','角色乙（示例覆核角色；ROLE-B）'))
# Expand the field dictionary to match Blueprint: field, definition, source, responsible role, validation rule.
field_table=next(t for t in d.tables if any('欄位群' in c.text for c in t.rows[0].cells))
if len(field_table.rows)!=6 or len(field_table.columns)!=3: raise RuntimeError('Unexpected §4.1 dictionary shape')
field_table.add_column(Inches(1.65)); field_table.add_column(Inches(2.04))
headers=['欄位群','欄位定義／要填內容','來源／附件定位','填寫者','驗證規則']
rows=[
 ['案件識別','案號或既有識別方式；活動名稱與日期。','受控案件索引；核准課程計畫、執行紀錄及合作往來。','承辦 ROLE-A','與索取稿及來源文件一致；無來源標 UNKNOWN。'],
 ['索取稿資料','八項必要欄位、草稿日期、預定管道及實際寄出狀態。','受控索取稿；實際電子郵件／公文及寄送、送達紀錄。','承辦 ROLE-A','草稿日與寄送日分開；未寄出不得推定送達。'],
 ['原件資料','收件狀態、出具者、日期、編號（如有）、型態及保存位置。','第三方實際提供的原件、收件紀錄及受控檔案索引。','保管 ROLE-C','未收件填無原件／不適用；不得補造原件資料。'],
 ['查驗紀錄','六項檢核的證據來源、結果、疑點、覆核角色與日期。','第三方原件、可核實管道及支持活動事實的原始紀錄。','覆核 ROLE-B','無證據標未開始／HOLD；完整度不等於通過。'],
 ['綜合判定','採認／不採認／暫緩、理由、責任及後續追蹤條件。','完整查驗結果、用途授權及正式決定紀錄。','決定 ROLE-D','依可追溯證據判定；記錄決定日期與理由。']
]
widths=[0.92,1.35,1.65,0.95,2.04]
field_table.autofit=False
for i,w in enumerate(widths):
 field_table.columns[i].width=Inches(w)
for ri,row in enumerate(field_table.rows):
 for ci,cell in enumerate(row.cells):
  cell.width=Inches(widths[ci])
  text=headers[ci] if ri==0 else rows[ri-1][ci]
  set_cell_text(cell,text)
  for para in cell.paragraphs:
   for run in para.runs: run.font.size=Pt(9.5)
  tcpr=cell._tc.get_or_add_tcPr()
  shd=tcpr.find(qn('w:shd'))
  if shd is None:
   shd=OxmlElement('w:shd'); tcpr.append(shd)
  if ri==0:
   shd.set(qn('w:fill'),'1F3A5F')
   for para in cell.paragraphs:
    for run in para.runs:
     run.bold=True; run.font.color.rgb=__import__('docx').shared.RGBColor(255,255,255)
  elif ri%2==0: shd.set(qn('w:fill'),'F4F7FA')
  else: tcpr.remove(shd)
# Repeat §4.1 header row and keep individual rows together.
trpr=field_table.rows[0]._tr.get_or_add_trPr(); repeat=OxmlElement('w:tblHeader'); repeat.set(qn('w:val'),'true'); trpr.append(repeat)
for row in field_table.rows:
 trpr=row._tr.get_or_add_trPr(); cant=OxmlElement('w:cantSplit'); trpr.append(cant)
# Add bounded follow-up timing rule without inventing an association deadline.
flow_table=next(t for t in d.tables if any('停止／退回條件' in c.text for c in t.rows[0].cells) and any('起草前' in ' '.join(c.text for c in row.cells) for row in t.rows[1:]))
newp=OxmlElement('w:p'); flow_table._tbl.addnext(newp); follow=Paragraph(newp,d); follow.style=d.styles['Normal']
lead=follow.add_run('期限與提醒規則：'); lead.bold=True; lead.font.size=Pt(10.5)
r=follow.add_run('寄出前記錄經授權的預定寄出日、首次追蹤日與結案期限；寄出後只按核准期限追蹤。若查無核准期限或責任人，填 UNKNOWN／HOLD，提請授權決定者確認；未確認前不得由承辦自行推定逾期、回覆或結案。'); r.font.size=Pt(10.5)
# Clarify the empty access-log sheet is not a completed event.
replace_once('下表是未填的操作表格，可供將來實際調閱時使用；空表不代表曾經調閱。','下表是供未來實際調閱時使用的空白工作表；空白欄位僅供後續記錄，不代表任何人曾經調閱或已完成查核。')
# Replace one-line version note with a candidate-only revision history.
version_p=next(p for p in d.paragraphs if p.text.startswith('R02 候選示範稿'))
version_table=d.add_table(rows=3,cols=4)
version_table.style=next(t.style for t in d.tables if len(t.rows)==3 and len(t.columns)==3)
version_table.autofit=False
vwidths=[0.65,0.95,3.35,1.96]
vrows=[
 ['版本','形成日期','變更摘要','核定／發布狀態'],
 ['R02','2026-10-20','完成索取稿、登錄表、查驗流程與合成替換清單。','候選稿；未核定。'],
 ['R03','2026-10-20','依獨立審查補齊欄位來源／角色／驗證欄、修正兩筆正文定位，明確化職能代碼、追蹤期限及狀態用語。','候選稿；未核定，不代表協會正式發布。']
]
for ri,row in enumerate(version_table.rows):
 for ci,cell in enumerate(row.cells):
  cell.width=Inches(vwidths[ci]); version_table.columns[ci].width=Inches(vwidths[ci]); set_cell_text(cell,vrows[ri][ci])
  for p in cell.paragraphs:
   for run in p.runs: run.font.size=Pt(9.5)
  tcpr=cell._tc.get_or_add_tcPr(); shd=tcpr.find(qn('w:shd'))
  if shd is None: shd=OxmlElement('w:shd'); tcpr.append(shd)
  if ri==0:
   shd.set(qn('w:fill'),'1F3A5F')
   for p in cell.paragraphs:
    for run in p.runs: run.bold=True; run.font.color.rgb=__import__('docx').shared.RGBColor(255,255,255)
  elif ri==2: shd.set(qn('w:fill'),'F4F7FA')
  else: tcpr.remove(shd)
version_p._p.addprevious(version_table._tbl)
version_p._element.getparent().remove(version_p._element)
trpr=version_table.rows[0]._tr.get_or_add_trPr(); repeat=OxmlElement('w:tblHeader'); repeat.set(qn('w:val'),'true'); trpr.append(repeat)
# Ensure the main sample role values match the role code table.
# Save to a new revision, preserving R02 bytes.
d.save(out)
# Reopen and validate no core changes were omitted.
check=Document(out)
text='\n'.join(p.text for p in check.paragraphs)+'\n'+'\n'.join(c.text for t in check.tables for row in t.rows for c in row.cells)
assert 'I17d「成果評估」' not in text
assert '原始點協會目前設有' not in text
assert '維持待補或不納入' not in text
assert '§4.2 案例識別；§5.3 案例識別欄' in text
assert '§4.2 寄件方；§5.3 寄件方／承辦角色欄' in text
assert 'ROLE-A' in text and 'ROLE-D' in text
assert hashlib.sha256(src.read_bytes()).hexdigest()==expected
print(json.dumps({'source_r02_sha256':expected,'new_r03_sha256':hashlib.sha256(out.read_bytes()).hexdigest(),'new_bytes':out.stat().st_size,'paragraphs':len(check.paragraphs),'tables':len(check.tables),'dictionary_columns':len(next(t for t in check.tables if any('欄位群' in c.text for c in t.rows[0].cells)).columns),'version_rows':len(next(t for t in check.tables if any('形成日期' in c.text for c in t.rows[0].cells)).rows),'all_27_ids_preserved':all(f'0092-F{i:03d}-E{"2" if i<=20 else "3"}' in text for i in range(1,28))},ensure_ascii=True,indent=2))
