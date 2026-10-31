from __future__ import annotations

import hashlib
import json
import re
import zipfile
from collections import Counter
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0122_I19-E10_SOURCE_PREP_R01_20261024")
WRITER = ROOT / "worker_R01" / "antigravity_writer_R01_20261024" / "writer_response_R01.md"
OUT = ROOT / "candidate_R01"
OUT.mkdir(parents=True, exist_ok=True)
MANUSCRIPT = OUT / "0122_課後追蹤與證照通過分析表_R01_repaired.md"
DOCX = OUT / "0122_課後追蹤與證照通過分析表_R01.docx"
CALC_RECEIPT = OUT / "calculation_validation_R01.json"


def replace_once(text: str, old: str, new: str, label: str) -> str:
    n = text.count(old)
    if n != 1:
        raise RuntimeError(f"{label}: expected one match, got {n}")
    return text.replace(old, new, 1)


def replace_section(text: str, start: str, end: str, replacement: str, label: str) -> str:
    a = text.find(start)
    b = text.find(end, a + len(start)) if a >= 0 else -1
    if a < 0 or b < 0:
        raise RuntimeError(f"{label}: section boundaries not found")
    return text[:a] + replacement.rstrip() + "\n\n" + text[b:]


def round_half_up(value: float, places: int = 2) -> float:
    unit = Decimal("1").scaleb(-places)
    return float(Decimal(str(value)).quantize(unit, rounding=ROUND_HALF_UP))


def md_table(headers: list[str], rows: list[list[str]]) -> str:
    def esc(s: object) -> str:
        return str(s).replace("|", "／").replace("\n", "<br>")
    out = ["| " + " | ".join(map(esc, headers)) + " |", "| " + " | ".join([":---"] * len(headers)) + " |"]
    out.extend("| " + " | ".join(esc(x) for x in row) + " |" for row in rows)
    return "\n".join(out)


def numeric_summary(rows: dict[str, list[int | None]]) -> tuple[list[list[str]], dict[str, dict[str, object]]]:
    item_names = {
        "Q01": "工作關聯性", "Q02": "技能記憶維持", "Q03": "應用機會", "Q04": "實際應用頻率",
        "Q05": "自主執行信心", "Q06": "流程遵循", "Q07": "工作品質改善", "Q10": "整體行為轉化",
    }
    keys = list(item_names)
    result = {}
    out = []
    for idx, key in enumerate(keys):
        vals = [v[idx] for v in rows.values()]
        good = [v for v in vals if v is not None]
        favourable = sum(v >= 4 for v in good)
        total = sum(good)
        n = len(good)
        missing = len(vals) - n
        mean = round_half_up(total / n, 2)
        result[key] = {"n": n, "missing": missing, "sum": total, "mean": mean, "favourable_n": favourable, "favourable_denominator": n, "favourable_percent": round_half_up(100 * favourable / n, 1)}
        out.append([key, item_names[key], str(n), str(missing), str(total), f"{total}/{n} = {mean:.2f}", f"{favourable}/{n}", f"{round_half_up(100 * favourable / n, 1):.1f}%"])
    return out, result


def counts_for(rows: dict[str, dict[str, object]], field: str, codes: range) -> dict[int, int]:
    counts = Counter(int(r[field]) for r in rows.values())
    return {c: counts.get(c, 0) for c in codes}


raw = WRITER.read_text(encoding="utf-8")
text = raw

# Remove line-local absolute file links and stale engineering claims before content edits.
text = re.sub(r"\[official_source_0108_r11\.pdf\]\(file:///[^)]*official_source_0108_r11\.pdf\)", "official_source_0108_r11.pdf", text)
text = re.sub(r"\[evidence/mapping_evidence_0122_R01\.json\]\(file:///[^)]*mapping_evidence_0122_R01\.json\)", "側錄檔 evidence/mapping_evidence_0122_R01.json", text)
text = re.sub(r"file:///E:/TTQS/[^\s)]+", "（來源定位見本機側錄清冊）", text)
text = text.replace("Commit ID `f8438c7c6de824e11fede9d99434857d4b2e49a2`", "來源包建置時 PR #92 的來源快照；目前 HEAD 與精確來源雜湊記錄於本機側錄，不寫入文件正文")
text = text.replace("- **官方來源雜湊值（SHA-256）：** `805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709`。", "- **官方來源版本與精確 SHA-256：** 見本文件來源側錄；正文以文件名稱及頁碼定位。")

# Make the demonstration boundary explicit from the first page.
text = replace_once(text,
    "> **治理狀態：** 示範性候選文稿（SYNTHETIC DEMONSTRATION ONLY）／無真實公協會資料／形式審查與語意審查待定",
    "> **治理狀態：** SAMPLE／SYNTHETIC 示範候選；本文件內的課程、日期、代碼、回答、分數、追蹤結果與行動均為合成例子，並非任何實際學員、原始點協會、雇主或第三方的紀錄。REAL 證據未提供；獨立語意審查待辦；正式接受未授權。",
    "governance banner")

# Use card, scope, privacy and procedure wording grounded in a small volunteer association.
text = text.replace("| **負責角色** | 追蹤調查執行員、課程專責規劃師、訓練成效評估小組 | [SAMPLE-KEY-ROLE-EXEC] 訓練推展組專責人員 |", "| **建議分工** | 可由行政承辦彙整、課程負責人判讀、授權主管核定；實際職權依協會現行授權表確認 | [SAMPLE-ROLE-ADMIN] 行政承辦（示範角色，非真實職稱） |")
text = text.replace("| **必要輸入資料** | 結訓學員合格清冊、受測通訊錄、證照報考清冊、外部測驗成績單或證書影本 | [SAMPLE-KEY-INPUTS] 結訓合格10名學員名冊、聯絡資訊表、測驗成績登錄冊 |", "| **必要輸入資料** | 已核定班別及結訓名冊、同意聯絡資料、追蹤問卷、外部檢定結果的可核驗來源；缺少任一項須標示缺件 | [SAMPLE-KEY-INPUTS] 合成名冊與合成問卷；沒有真實成績單、證書或聯絡資料 |")
text = text.replace("| **隱私與個資保護** | 依個人資料保護法去識別化編碼（P01-P10），非公務或評鑑目的不得外洩 | [SAMPLE-KEY-PRIVACY] 符合個資保護匿名編碼作業規範，嚴禁揭露真名 |", "| **隱私與個資保護** | 僅用必要欄位及去識別代碼；聯絡對照表另存於經授權的受限位置。個資處理方式須依協會核定規則及適用要求確認 | [SAMPLE-KEY-PRIVACY] 合成代碼；未宣稱完成任何真實個資處理或法遵查核 |")
text = text.replace("本文件專屬對應 TTQS 訓練品質評核系統（貴賓/企業/機構版）官方指引來源規範：", "本文件專屬對應 TTQS 訓練品質評核系統官方指引中的學習技能面課後追蹤及證照通過情形：")
text = text.replace("- **官方來源雜湊值（SHA-256）：** `805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709`。", "- **來源版本控制：** 官方 PDF 的頁碼及精確檔案雜湊見品質側錄清冊，避免把工程識別碼混入協會日常表單正文。")
text = text.replace("本文件（0122）係針對「單一訓練班別課後追蹤機制與技能檢定結果」之標準分析作業", "本文件（DOCX 0122 / I19-E10）係針對「單一訓練班別課後追蹤機制與證照狀態分析」的實用表單及示範分析")
text = text.replace("為防範無效調查與樣本人為偏誤，制定標準接觸階梯：", "以下是可供小型協會調整的接觸流程範例。實際使用前，先確認學員同意的聯絡方式及協會核定的聯絡規則：")
text = text.replace("- **第 1 次接觸（Day 30/90）：** 寄發個人專屬加密數位調查問卷連結（Email/簡訊）。", "- **第 1 次接觸（第 30／90 天）：** 以學員事先同意的聯絡管道寄送問卷及用途說明。")
text = text.replace("- **第 2 次接觸（第 3 工作天）：** 對未填答者發送通訊軟體或簡訊提醒。", "- **第 2 次接觸（3 個工作天後）：** 僅以已同意的管道提醒一次；記錄日期與結果。")
text = text.replace("- **第 3 次接觸（第 5 工作天）：** 由訓練小組執行專責電話訪談與核對。", "- **第 3 次接觸（第 5 個工作天）：** 若學員已同意電話聯絡，最多再聯絡一次；未同意時不得改以電話或聯絡親友。")
text = text.replace("號碼失效、信箱退信且透過次要緊急聯絡人連續 3 次均無法聯繫。", "在學員同意的聯絡管道完成最多三次嘗試仍無法取得回覆；除非另有明確授權，不聯絡緊急聯絡人、雇主或親友。")
text = text.replace("- 本調查遵循個人資料保護精神，全數樣本賦予虛擬代碼（`[SAMPLE-P01]` 至 `[SAMPLE-P10]`）。", "- 本表範例中的 `[SAMPLE-P01]` 至 `[SAMPLE-P10]` 是合成代碼，不代表任何真實會員或學員。實際使用時，代碼與身分對照資料須分開保管，並依核定規則設定存取及保存方式；未確認保存規則前填 `RETENTION_POLICY=UNKNOWN/REFER_TO_APPROVED_POLICY`。")
text = text.replace("- 學員可自由決定是否填答，並享有隨時退出之權利。", "- 實際發送時應清楚告知填答自願、用途、資料處理窗口及退出方式；本示範不代表已取得任何人的同意。")
text = text.replace("分類與程度 (1=無任何機會 ~ 5=每日持續應用)", "有序分類（1=無機會、2=偶發、3=每月、4=每週、5=每日或核心情境）")
text = text.replace("分級選項 (1=從未應用 ~ 5=每週多次高頻應用)", "頻率分類（1=從未、2=僅一次、3=每月1至3次、4=每週1至2次、5=每週3次以上）")
text = text.replace("（1=通過, 2=未通過, 3=報名待考, 4=延期未考, 5=無資格/不報考）", "（1=合成測驗達示範門檻、2=未達示範門檻、3=已報名待考、4=延後、5=未報考或不適用；REAL 證照狀態另須核對外部原始資料）")
text = text.replace("| **Q12** | 證照報考與結果 (Credential Status) | 對應專業認證之報名、應考及通過狀態 | 狀態選項 (1=通過, 2=未通過, 3=報名待考, 4=延期未考, 5=無資格/不報考) | 建立證照檢定母數，計算嘗試率與通過率 |", "| **Q12** | 證照或檢定狀態 | 填寫實際檢定名稱、主辦單位與目前狀態 | 1=外部資料證實通過；2=外部資料證實未通過；3=待考；4=延後；5=未報考／不適用；無來源不得列 REAL | 依可核驗的外部結果分別計算報考與結果；示範代碼不可稱證照通過 |")
text = text.replace("Q12. 針對本課程對應之專業技能檢定認證（如：初級專案品質管理技術士證照），您目前的報考與取證狀態為何？", "Q12. 若本班有外部證照或檢定，請填實際名稱及主辦／發證單位；目前報名或測驗狀態為何？沒有適用檢定時選「不適用」。")
text = text.replace("[ ] 1：已完成報考並「通過檢定」（已取得或符合發證資格）\n      [ ] 2：已完成報考但「未通過」（成績未達合格門檻）\n      [ ] 3：已完成報名繳費，「待測驗中」或「成績待公佈」\n      [ ] 4：因公務/個人因素「暫緩報考」，已申請展延至後續梯次\n      [ ] 5：個人並無報考意願，或未具備該項檢定之法定報名資格", "[ ] 1：外部資料證實已通過（填證據定位；沒有可核驗來源時不得勾選）\n      [ ] 2：外部資料證實未通過（填證據定位）\n      [ ] 3：已報名待考或成績待公布（填預定／實際日期）\n      [ ] 4：延後應考（填原因類別及狀態；申請或核准須有來源）\n      [ ] 5：未報考或不適用（說明原因；不得推定為未通過）\n      檢定名稱：__________　主辦／發證單位：__________　結果來源代碼：__________　REAL／SAMPLE：__________")

# Replace the optional third-party form with a truly blank, consent-gated instrument.
observer = """## 6. 選用觀察表

本表只在受訓者明確同意、觀察者確實存在且實際直接觀察後使用。沒有上述條件時，保持空白並標示 `NOT_AVAILABLE_HELD`。不得在示範資料中虛構觀察者、同意書、簽名、觀察內容或客觀證據。

**填寫表單（實際使用時由授權觀察者填寫）**

| 欄位 | 填寫內容 |
|---|---|
| 學員去識別代碼 | ______________________________ |
| 學員明示同意狀態及來源定位 | ______________________________；未取得同意不得填觀察資料 |
| 觀察者角色 | ______________________________；不得只填虛構姓名 |
| 觀察日期與期間 | ______________________________ |
| 觀察場域／任務 | ______________________________ |
| 直接觀察到的行為 | ______________________________；寫可見行為，不寫推測 |
| 證據定位 | ______________________________；只記核准位置及代碼，不貼敏感資料 |
| 未觀察到／無法判斷的項目 | ______________________________ |
| 觀察限制及可能偏差 | ______________________________ |
| 觀察者確認日期 | ______________________________ |
| REAL／SAMPLE 狀態 | ______________________________；本文件示範欄維持未提供／暫緩 |

**本案例狀態：** `NOT_AVAILABLE_HELD`。本表未附任何真實第三方資料；本文件也沒有模擬簽名或假造的觀察紀錄。

---
"""
text = replace_section(text, "## 6. 選用之外部佐證／主管／同儕觀察檢核表（Optional Confirmer Form）", "## 7. 追蹤對象流向與收案對照分析（Cohort Flow）", observer, "observer form")

# Cohort-flow text is synthetic and shows consent-based contact only.
text = text.replace("為徹底杜絕學術界與業界常見之「隱匿流失樣本」與「分母漂移」弊端，本章嚴格展示兩波次（30天與90天）全數 10 名編碼對象之收案流向全貌。", "下列合成案例以同一批 10 個示範代碼呈現 30 天及 90 天收案流向。這些數字不是協會真實追蹤成果；每一波分母分開計算，並保留未回覆與失聯人數。")
text = text.replace("30天出差拒填；90天銷假補填完成", "合成情境：30天未回覆、90天回覆；無真實差假紀錄")
text = text.replace("90天波次因突發專案趕工明確婉拒填答", "合成情境：90天接觸後未回覆；無真實拒答紀錄")
text = text.replace("90天波次手機停用且離職失聯（2次未通）", "合成情境：90天無法聯絡；無真實離職或通話紀錄")
text = text.replace("兩波次均因出國與更換號碼無法取得聯繫", "合成情境：兩波均無回覆；無真實出國、換號或聯絡紀錄")

# Rewrite credential section to demonstrate arithmetic without inventing a real issuer or credential.
cert_section = """### 8.1 真實登錄規則與合成測驗界線

真實使用時，先填檢定或證照的正式名稱、主辦／發證單位、資格條件、日期、結果來源定位及必要授權。沒有可核驗的外部資料時，狀態必須是 `HOLD/UNKNOWN`，不能把內部結訓測驗算成外部證照通過。

本文件的 `SAMPLE-CRED-01` 只是合成資料欄位代碼；它不是實際證照名稱、發證機構或考試。下列分數及「達／未達示範門檻」只用來演示分母算式，不代表真實檢定、證書或外部結果。此案例沒有任何外部憑證原件。

### 8.2 全體合成學員檢定狀態登錄表

| 合成學員代碼 | 示範資格 | 示範報考狀態 | 示範日期 | 合成分數與門檻 | 示範結果分類 | REAL 外部來源 |
|---|---|---|---|---|---|---|
| P01 | 是 | 已完成合成測驗 | 2026-07-10 | 88／70 | 達示範門檻 | 未提供；不存在可核驗原件 |
| P02 | 是 | 已完成合成測驗 | 2026-07-10 | 84／70 | 達示範門檻 | 未提供；不存在可核驗原件 |
| P03 | 是 | 已完成合成測驗 | 2026-07-10 | 82／70 | 達示範門檻 | 未提供；不存在可核驗原件 |
| P04 | 是 | 已完成合成測驗 | 2026-07-10 | 79／70 | 達示範門檻 | 未提供；不存在可核驗原件 |
| P05 | 是 | 已完成合成測驗 | 2026-07-10 | 75／70 | 達示範門檻 | 未提供；不存在可核驗原件 |
| P06 | 是 | 已完成合成測驗 | 2026-07-10 | 72／70 | 達示範門檻 | 未提供；不存在可核驗原件 |
| P07 | 是 | 已完成合成測驗 | 2026-07-10 | 64／70 | 未達示範門檻 | 未提供；不存在可核驗原件 |
| P08 | 是 | 已完成合成測驗 | 2026-07-10 | 58／70 | 未達示範門檻 | 未提供；不存在可核驗原件 |
| P09 | 是 | 未完成應試；合成原因類別為延後 | — | — | 待後續／不計入完成嘗試 | 未提供任何申請或核准文件 |
| P10 | 是 | 未完成應試；合成原因類別為未報考 | — | — | 未應試／不計為失敗 | 未提供任何聲明或外部文件 |

**示範門檻只為本表算式設定，並非任何外部檢定的合格標準。** 真實報告須以實際主辦單位公告標準及可核驗結果替換，並保存來源定位；不可以本表作為證照通過證明。

"""
text = replace_section(text, "### 8.1 檢定目標證照與憑證性質界定", "## 9. 示範案例 SAMPLE-0122-C01 列層級完整原始數據集", cert_section.rstrip(), "credential section")

# The follow-up tables are split to preserve legibility and expose every response.
d30_num = {
    "P01": [5, 5, 5, 5, 5, 5, 5, 5], "P02": [5, 4, 4, 4, 4, 4, 4, 4],
    "P03": [4, 4, 4, 4, 4, 4, 4, 4], "P04": [4, 4, 4, 3, 4, 4, 4, 4],
    "P05": [4, 3, 3, 4, 3, None, 4, 3], "P06": [4, 4, 3, 3, 3, 3, 3, 3],
    "P07": [3, 3, 2, 1, 2, 2, 2, 2], "P08": [3, 3, 1, 1, 3, 3, 2, 2],
}
d90_num = {
    "P01": [5, 5, 5, 5, 5, 5, 5, 5], "P02": [5, 4, 4, 4, 4, 4, 4, 4],
    "P03": [4, 4, 4, 4, 4, 4, 4, 4], "P04": [4, 4, 4, 4, 4, 4, 4, 4],
    "P05": [4, 4, 3, 4, 4, 4, 4, 4], "P06": [4, 4, 3, 3, 3, 3, 3, 3],
    "P09": [4, 3, 3, 3, 3, 3, 3, 3],
}

def numeric_table(rows: dict[str, list[int | None]]) -> str:
    data = [[pid] + ["缺答" if v is None else str(v) for v in vals] for pid, vals in rows.items()]
    return md_table(["合成代碼", "Q01", "Q02", "Q03", "Q04", "Q05", "Q06", "Q07", "Q10"], data)

cat30 = {
    "P01": [6, 2, "導入檢核工具", 1], "P02": [2, 1, "依表核對流程", 1],
    "P03": [3, 4, "嘗試整理通報步驟", 1], "P04": [4, 3, "同儕不習慣新流程", 1],
    "P05": [5, 1, "依表檢查但仍需練習", 1], "P06": [2, 2, "因時間有限只用部分步驟", 1],
    "P07": [3, 3, "設備與練習時間不足", 2], "P08": [1, 5, "示範情境無派工機會", 2],
}
cat90 = {
    "P01": [6, 5, "形成固定步驟", 1], "P02": [2, 1, "工作量仍影響使用", 1],
    "P03": [3, 4, "設備限制仍在，分段試用", 1], "P04": [4, 3, "主管支持增加但同儕協作仍需改善", 1],
    "P05": [5, 1, "熟練度提升但仍需練習", 1], "P06": [2, 2, "維持部分模組使用", 1],
    "P09": [2, 1, "恢復聯繫後重新安排練習", 4],
}

def categorical_table(rows: dict[str, list[object]]) -> str:
    return md_table(["合成代碼", "Q08 障礙碼", "Q09 支援碼", "Q11 合成摘要", "Q12 示範狀態碼"], [[pid] + v for pid, v in rows.items()])

section_9_1 = """### 9.1 第 30 天合成追蹤資料

本波有效回卷 8 份，P05 的 Q06 缺答保留為缺值。缺答不當作零分；Q06 以 7 份有效答案計算，其餘量表題以 8 份計算。下列每列均為 `SAMPLE-ANSWERS-W1`，不是實際受訪紀錄。

#### 量表題逐人答案

""" + numeric_table(d30_num) + """

#### 類別題與合成反思摘要

""" + categorical_table(cat30)
text = replace_section(text, "### 9.1 第 30 天波次有效回卷學員列層級資料矩陣（Wave 1, N = 8）", "### 9.2 第 90 天波次有效回卷學員列層級資料矩陣（Wave 2, N = 7）", section_9_1.rstrip(), "30-day raw data")

section_9_2 = """### 9.2 第 90 天合成追蹤資料

本波 7 份有效回卷，與 30 天波次不是同一組完整名單。P07 未回覆、P08 與 P10 無法聯絡；只有有效回覆列出現在下表。下列每列均為 `SAMPLE-ANSWERS-W2`，不是實際受訪紀錄。

#### 量表題逐人答案

""" + numeric_table(d90_num) + """

#### 類別題與合成反思摘要

""" + categorical_table(cat90)
text = replace_section(text, "### 9.2 第 90 天波次有效回卷學員列層級資料矩陣（Wave 2, N = 7）", "### 9.3 質性反思個案記錄（Qualitative Reflections）", section_9_2.rstrip(), "90-day raw data")
text = text.replace("為深入理解數據背後的真實情境，摘錄 3 則具代表性且去識別化之學員反思筆記：", "以下三段是為示範分類與後續行動而合成的例句，不是任何學員的真實原話或受訪紀錄：")
text = text.replace("「課後我將課堂學到的品質管制查檢表直接應用於每月業務進度報表。以前人工比對數據容易漏算，需耗費將近 4 個小時；現在依照標準檢核模組核對，作業時間縮短到 1.5 小時以內，且未再發生欄位勾稽錯誤。這套方法對我的日常報表產出非常有幫助！」", "【SAMPLE例句】情境可寫成：曾嘗試用課程檢核表核對例行報表；若要聲稱錯誤減少或時間縮短，須另附可比較的工作紀錄。本示範沒有這些紀錄。")
text = text.replace("「在部門例會上，我嘗試運用課程教導的結構化檢核表來引導專案驗收討論。主管肯定架構條理分明，但同組資深同儕習慣用舊方法快速結案，認為填寫檢核表是『增加行政負擔』。跨部門協作如果沒有高層統一公告標準，推動阻力真的不小，希望能有高階導師協助推廣。」", "【SAMPLE例句】情境可寫成：使用檢核表時遇到同儕採用意願不一；後續應先釐清工作流程與必要欄位，再決定是否提供同儕說明。這不是任何主管或同儕的實際意見。")
text = text.replace("「結訓回到工作崗位後，原本想立即將數據統計工具導入手頭上的專案，但資訊部門告知進階分析軟體授權有限，申請核發流程耗時一個月。加上近期組內有緊急標案插隊，每天加班處理雜務，根本沒有餘裕重新翻閱講義。最後證照考試術科沒過，滿受挫的，很需要能有一對一的複習指導。」", "【SAMPLE例句】情境可寫成：即使理解方法，仍可能因缺少設備、工作時間或練習機會而無法使用；應分別記錄障礙及支援需求，不推定為任何真實單位的經歷或檢定結果。")

# Generate complete item summaries from the visible row-level responses.
d30_rows, d30_stats = numeric_summary(d30_num)
d90_rows, d90_stats = numeric_summary(d90_num)
cat30d = {p: {"Q08": v[0], "Q09": v[1], "Q12": v[3]} for p, v in cat30.items()}
cat90d = {p: {"Q08": v[0], "Q09": v[1], "Q12": v[3]} for p, v in cat90.items()}

def category_rows(question: str, labels: list[str], d30_counts: dict[int, int], n30: int, d90_counts: dict[int, int], n90: int) -> str:
    rows = []
    for code, label in enumerate(labels, start=1):
        a, b = d30_counts.get(code, 0), d90_counts.get(code, 0)
        rows.append([str(code), label, f"{a}/{n30}", f"{a/n30*100:.1f}%", f"{b}/{n90}", f"{b/n90*100:.1f}%"])
    return f"**{question}**\n\n" + md_table(["代碼", "選項", "30天人數／分母", "30天比例", "90天人數／分母", "90天比例"], rows)

q08_30 = counts_for(cat30d, "Q08", range(1, 7)); q08_90 = counts_for(cat90d, "Q08", range(1, 7))
q09_30 = counts_for(cat30d, "Q09", range(1, 6)); q09_90 = counts_for(cat90d, "Q09", range(1, 6))
q12_30 = counts_for(cat30d, "Q12", range(1, 6)); q12_90 = counts_for(cat90d, "Q12", range(1, 6))
q11_table = md_table(["波次", "示範主題", "列代碼", "人數／有效回卷", "解讀界線"], [
    ["30天", "工具或流程初次使用", "P01、P02、P03、P05", "4/8", "合成摘要之示範編碼，不是訪談紀錄"],
    ["30天", "同儕支持、時間或機會障礙", "P04、P06、P07、P08", "4/8", "四類需分別讀原始文字，合併僅為摘要"],
    ["90天", "習慣維持或部分使用", "P01、P02、P04、P05、P06", "5/7", "合成摘要之示範編碼，不代表行為觀察"],
    ["90天", "設備限制或重新安排", "P03、P09", "2/7", "合成摘要之示範編碼，不是訪談紀錄"],
])

section_10_2 = """### 10.2 第 30 天量表題統計

平均數以該題有效作答計算；缺答不補零。正向分子固定定義為評分 4 或 5，正向率分母為該題有效作答數。Q03、Q04 的「有機會／有應用」另依問卷定義使用 2–5 分，不能與 4–5 分正向率混為一談。

""" + md_table(["題號", "構念", "有效 n", "缺答", "得分合計", "平均數", "正向分子／n", "正向率"], d30_rows) + """

### 10.3 第 90 天量表題統計

""" + md_table(["題號", "構念", "有效 n", "缺答", "得分合計", "平均數", "正向分子／n", "正向率"], d90_rows) + """

### 10.4 類別題與反思摘要

Q08、Q09、Q12 為單選分類，表內比例以該波有效回卷為分母。未回覆者不進入題目分母。Q11 的主題數由下列逐人摘要分類，分類結果只用於本合成案例示範。

""" + category_rows("Q08 障礙", ["無應用機會", "工作負荷", "設備或系統不足", "支持不足", "熟練度不足", "無明顯障礙"], q08_30, 8, q08_90, 7) + "\n\n" + category_rows("Q09 支援需求", ["實作工作坊", "手冊或範本", "導師諮詢", "系統權限", "暫無需求"], q09_30, 8, q09_90, 7) + "\n\n" + category_rows("Q12 合成檢定狀態", ["合成測驗達示範門檻", "未達示範門檻", "已報名待考", "延後", "未報考／不適用"], q12_30, 8, q12_90, 7) + "\n\n**Q11 合成摘要主題**\n\n" + q11_table + "\n\n"

text = replace_section(text, "### 10.2 第 30 天波次各題項量化統計總表（N = 8）", "### 10.3 技能落實應用率：全體分母 vs. 機會分母對照驗算", section_10_2.rstrip(), "two-wave item summaries")
text = text.replace("### 10.3 技能落實應用率：全體分母 vs. 機會分母對照驗算", "### 10.5 技能應用率的全體分母與機會分母驗算")
text = text.replace("### 10.4 證照報考與通過率獨立計算", "### 10.6 合成測驗狀態與分母計算")
text = text.replace("針對示範專業證照 [SAMPLE-CRED-01]，嚴格區分「應考嘗試者」與「全體符合資格學員」之雙重分母計算，嚴禁將小樣本數據指稱為機構通過率：", "以下只計算 `SAMPLE-CRED-01` 合成測驗情境；這不是證照或外部通過結果。真實證照統計須另以外部可核驗資料建立分子與分母。")
text = text.replace("完成檢定應試人數（Completed Attempts） = $8$ 人", "完成合成測驗人數（Completed Sample Attempts） = $8$ 人")
text = text.replace("檢定測驗通過人數（Passed） = $6$ 人（`P01-P06`）", "達合成示範門檻人數（Met Sample Threshold） = $6$ 人（`P01-P06`）")
text = text.replace("檢定測驗未通過人數（Failed） = $2$ 人（`P07-P08`）", "未達合成示範門檻人數（Below Sample Threshold） = $2$ 人（`P07-P08`）")
text = text.replace("檢定應試涵蓋率（Attempt Coverage Rate）", "合成測驗完成率（Sample Attempt Coverage Rate）")
text = text.replace("實考通過率（Pass Rate among Completed Attempts）", "合成測驗達門檻比例（among completed sample attempts）")
text = text.replace("結訓母體總取證率（Eligible-Cohort Pass Proportion）", "符合資格者中達合成門檻比例（Eligible-Cohort Sample Proportion）")
text = text.replace(r"| **證照憑證實質性 (Credential Documents)** | 證照通過數據僅代表應試者於特定測驗日符合該檢定之合格水準。 | 嚴禁以內部結訓證明替代外部專業證照，亦不得以單次 6/8 通過率推論訓練具因果保證。 |", "| **證照與檢定結果** | REAL 結果必須能追溯到實際外部主辦／發證單位及可核驗來源；本文件只有合成測驗示例。 | 不得把內部測驗或本文件的 SAMPLE 比例寫成外部證照、協會通過率或訓練因果證明。 |")

# Correct evidence claims and turn action/closure items into proposed, unapproved synthetic actions.
text = text.replace("基於 30 天與 90 天追蹤調查所顯現之實務落差，制定具體改善行動方案，明確界定權責角色、期程、量化重啟門檻與實體結案標準：", "以下行動是根據合成案例演示如何由發現轉成跟進工作；不是原始點協會的決議、已派工或已完成成果。實際使用前由有權者確認承辦角色、日期與資源。")
text = text.replace("| **發現一** | **軟硬體權限缺乏阻礙實踐**：學員反映無進階分析軟體授權，導致技能擱置。 | `P03`, `P07` | 協調訓練委員會與資訊單位，建立「課後 60 天沙盒練習帳號」申請管道。 | 2026-08-15 | [SAMPLE-ROLE-IT] 資訊支援窗口 | 沙盒帳號申請開通率達 $100\%$ |", "| **合成發現一** | 合成回答顯示部分學員回報設備或權限受限；尚無客觀設備紀錄。 | P03、P07 | 先確認是否有可使用的替代設備，再提出不涉及個資的練習方案。 | 2026-09-15（SAMPLE） | SAMPLE-ROLE-ADMIN（建議承辦） | 記錄已評估人數／需協助人數；未核定前不報完成率 |")
text = text.replace("| **發現二** | **術科實作熟練度不足致考照未過**：2名未通過學員均失分於複雜計算與分析題。 | `P07`, `P08` | 開設「術科實戰模擬複習工作坊（3小時）」，強化高頻失分題型演練。 | 2026-08-30 | [SAMPLE-ROLE-FAC] 課程專任講師 | 未通過學員工作坊出席率達 $100\%$ |", "| **合成發現二** | P07、P08 在合成測驗中未達示範門檻；不能稱為真實檢定未通過。 | P07、P08 | 如有實際需求，提供自願參加的短時練習與答疑；先確認教材及師資可用。 | 2026-09-30（SAMPLE） | SAMPLE-ROLE-COURSE（建議承辦） | 列出邀請、參加及未參加人數；不得預設全員出席 |")
text = text.replace("| **發現三** | **追蹤調查失聯與缺漏**：存在無法接觸（P10）及題項漏填（P05漏Q06）現象。 | `P05`, `P10` | 數位問卷系統全面開啟「必填防錯驗證」，並建立第二緊急通訊備援管道。 | 2026-07-30 | [SAMPLE-ROLE-ADM] 教務系統維護員 | 系統無漏填提交率達 $100\%$ |", "| **合成發現三** | 合成追蹤中有未回覆、失聯及一項缺答；不可推定為學員過失。 | P05、P07、P08、P10 | 調整問卷提示並提供「不適用／不願回答」選項；只使用受訪者同意的聯絡方式。 | 2026-08-31（SAMPLE） | SAMPLE-ROLE-ADMIN（建議承辦） | 比較各波 due、已聯絡、回覆、有效及缺答數；分母逐波揭露 |")
text = text.replace("若次一期訓練之 30 天「有效回卷率」低於 $70.0\%$，本追蹤作業程序自動宣告失效，必須重啟學員通訊查核機制。", "若次一期 30 天有效回覆率低於 70.0%，建議先檢查名冊、問卷長度及同意聯絡方式，再由授權主管決定是否調整程序；不得自動宣告協會作業失效。")
text = text.replace("若完成應試學員之「檢定通過率」低於 $60.0\%$，訓練發展小組必須強制重啟課程大綱與實作演練時數之實質審查。", "若真實可核驗的外部檢定結果顯示應試者達標比例低於協會事前核定的參考門檻，才提請有權者檢視課程與練習安排；不得將合成案例的數字作為觸發依據。")
text = text.replace("凡發現外部證照無客觀第三方成績冊影本、或外部觀察表有偽造之虞者，該批次數據立即列入 `HOLD` 管制，終止分析流向。", "外部檢定來源、觀察者身分／同意或分母對帳任一項未核實時，相關 REAL 欄位標 `HOLD/UNKNOWN`，不納入真實統計；合成數據仍可單獨作教學演示。")
text = text.replace("2. **實體結案標準（Formal Closure Criteria）：**\n   - **結案要件 C-01：** 針對未通過檢定學員（`P07`, `P08`），均已收到工作坊輔導簽到表與補考輔導計畫確認書。\n   - **結案要件 C-02：** 全體 6 名通過證照學員之外部證書影本，均已完成數位編號核對並列冊歸檔。\n   - **結案要件 C-03：** 資訊沙盒作業已完成測試，並將帳號申裝流程增列入下一期《學員課後服務手冊》。", "2. **真實案件結案前應取得的資料（本合成案例尚未結案）：**\n   - C-01：有權承辦者確認後續安排、實際提供紀錄及參與者同意；缺件維持待辦。\n   - C-02：外部主辦／發證單位的可核驗結果及必要原始文件定位；未取得前不登錄 REAL 通過。\n   - C-03：行動措施的核定、執行日期、追蹤指標與結果；本表內所有行動仍為 `SAMPLE_PLAN_NOT_APPROVED`。")

# Replace the replacement register with a fact-family complete index and explicit REAL hold state.
replacement = """### 13.1 合成資料替換清冊

下表以欄位族群標出所有合成日期、身份代碼、問卷回答、追蹤狀態、測驗分數、角色、觀察、行動及期限。實際填寫時，逐筆以受授權的來源紀錄替代；若來源不存在，保留 `UNKNOWN/HELD`，不得沿用示範值。

| 合成資料範圍 | 文件內位置 | 替換鍵 | 真實替換來源及檢核 |
|---|---|---|---|
| 案例代碼與班別名稱 | 封面、使用卡、各表 | SAMPLE-CASE-0122-C01；SAMPLE-COURSE | 核定課程計畫及單一班別代碼；跨班合併須有可追溯依據 |
| 結訓日與追蹤波次日期 | 追蹤時程、結果表 | SAMPLE-D0；SAMPLE-W1-DATE；SAMPLE-W2-DATE | 實際結訓紀錄及問卷發送／截止時間；未提供則 HOLD |
| 合成學員代碼及資格 | 流向、回答矩陣、檢定表 | SAMPLE-P01 至 SAMPLE-P10；SAMPLE-ELIGIBILITY | 授權名冊及資格證據；身份對照表不得放在公開報告 |
| 30天與90天聯絡、回覆及有效狀態 | Cohort flow、逐列狀態表 | SAMPLE-CONTACT-W1；SAMPLE-CONTACT-W2 | 每次同意管道接觸紀錄、回覆狀態、有效判定及分母核對 |
| 量表及類別逐題回答 | Q01–Q12原始矩陣與摘要 | SAMPLE-ANSWERS-W1；SAMPLE-ANSWERS-W2 | 原始問卷；每題保留缺答、不適用及選項碼，不可補零 |
| P05 的 Q06 缺答 | 30天問卷與 item n | SAMPLE-MISSING-W1-P05-Q06 | 原始問卷缺答位置；若實際資料有回答，依來源替換 |
| Q11 摘要及反思例句 | 質性反思及主題摘要 | SAMPLE-Q11-W1；SAMPLE-Q11-W2 | 經授權的匿名化原文及編碼依據；本例句不是受訪原話 |
| 合成測驗名稱、日期、成績與門檻 | 證照狀態表、算式 | SAMPLE-CRED-01；SAMPLE-TEST-DATE；SAMPLE-SCORE-P01–P10；SAMPLE-THRESHOLD | 真實外部檢定名稱／單位、官方標準及可核驗結果；本例無真實證照或原件 |
| 延後與未報考類別 | P09、P10 登錄列 | SAMPLE-NOT-ATTEMPTED-REASON | 由本人提供且授權採用的原因；沒有紀錄不推定原因，不聲稱申請或核准 |
| 觀察者、同意、觀察結果與證據定位 | 選用觀察表 | SAMPLE-OBSERVER-NONE；SAMPLE-CONSENT-NONE | 本文件未提供；實際使用需真實觀察者及明確同意 |
| 分工角色 | 使用卡、改善方案 | SAMPLE-ROLE-ADMIN；SAMPLE-ROLE-COURSE；SAMPLE-ROLE-AUTHORITY | 協會核定職掌或授權表；不得把建議角色說成現有部門 |
| 合成發現、措施、期限與指標 | 行動方案及重啟條件 | SAMPLE-FINDINGS；SAMPLE-ACTIONS；SAMPLE-TARGET-DATES；SAMPLE-THRESHOLDS | 有權者核定的行動紀錄、資源、期限及後續追蹤結果；目前均未核定 |
| 個資保存與處理設定 | 使用卡、REAL 清單 | RETENTION_POLICY=UNKNOWN/REFER_TO_APPROVED_POLICY | 依協會正式核定規則填寫；不得在本文件自行設定保存年限 |

### 13.2 REAL 證據暫緩採信檢核清單

本候選沒有真實學員、問卷、證照、第三方觀察或協會決議來源。下列項目皆保持 `UNAVAILABLE_HELD`，不是已查驗或已通過：

- [ ] 授權班別、結訓名冊、資格及每波應追蹤分母。
- [ ] 學員同意的聯絡管道、發送／催收／拒答／失聯紀錄。
- [ ] 30天及90天原始問卷、缺答與不適用位置、逐題計算核對。
- [ ] 實際證照名稱、主辦／發證單位、官方結果來源及應試者授權。
- [ ] 如使用外部觀察：受訓者同意、實際觀察者、日期、場域及直接觀察依據。
- [ ] 協會授權者核定的改善措施、承辦、資源、期限及結案證據。
- [ ] 個資處理及保存規則；未取得正式規則時維持 `UNKNOWN/REFER_TO_APPROVED_POLICY`。
- [ ] 逐項替換 SAMPLE 值並由第二人核對後，才將該欄位由 SAMPLE 改列 REAL。

"""
text = replace_section(text, "### 13.1 示範性資料替換索引表（Replacement Register）", "---\n\n## 14. 來源索引、獨立審查與治理狀態宣告", replacement.rstrip(), "replacement register")

# Source index carries human-readable provenance; exact hashes stay in sidecars.
source_index = """### 14.1 來源索引與範圍

- 官方依據：official_source_0108_r11.pdf，第 8 頁；對應文字為「提出課後追蹤調查與分析結果、證照通過情形（學習技能面）」。官方文件精確 SHA-256 及抽取證據記於側錄清冊。
- 專案 crosswalk：I19-E10 對應 DOCX 0122；印刷流水號為 `UNKNOWN_NOT_ASSERTED`。I19-E10、0122 與清冊項目116不可互換。
- 來源包與唯一 Blueprint：`SOURCE_PACKET_0122_I19-E10_R01.md`、`UNIQUE_BLUEPRINT_0122_I19-E10_R01.md`，各自精確 SHA-256 記於側錄清冊。
- 相鄰項目界線：0120 為即時訓練成果；0121 為課程滿意度與即時學習評量；本文件處理課後追蹤及證照狀態。清冊項目116是另一項高通過率可追溯原始紀錄要求。
- 歷史 0122 草稿只保留於歷史檔案，不沿用其正文或版面。所有 exact SHA、來源 commit、建置與驗收收據均在本機／GitHub 側錄，不放入協會使用的表單正文。

### 14.2 審查及治理狀態

本文件是 `SAMPLE/SYNTHETIC CANDIDATE`。建置及公式檢查完成後，仍須對同一 DOCX exact SHA 執行獨立語意審查；獨立審查未完成前不標 PASS。技術檢查不等於內容接受，也不代表取得真實資料。

```text
REAL_ASSOCIATION_DATA=NOT_PROVIDED
THIRD_PARTY_CREDENTIAL_ORIGINAL=NOT_PROVIDED
INDEPENDENT_SEMANTIC_REVIEW=PENDING
HUMAN_FORMAL_ACCEPTANCE=NOT_GRANTED
```
"""
text = replace_section(text, "### 14.1 來源索引與品質合約對照", "---\n**文件標題：**", source_index.rstrip(), "source index")
text = text[:text.find("---\n**文件標題：**")] + "---\n"

# Normalize headings to clear, concise labels suitable for a form/report.
heading_map = {
    "## 1. 表單管制資訊卡（Use Card）": "## 1. 使用資訊",
    "## 2. 官方規範對照、法源依據與範圍界定": "## 2. 官方指標與適用範圍",
    "### 2.1 官方規範對照與指標定位": "### 2.1 官方需求對照",
    "### 2.2 鄰近指標負向排除界定（Negative Boundaries）": "### 2.2 相鄰需求界線",
    "## 3. 課後追蹤調查機制設計": "## 3. 課後追蹤流程",
    "### 3.1 追蹤期程與目標對象界定": "### 3.1 追蹤期程與對象",
    "### 3.2 接觸管道與催收標準作業程序（SOP）": "### 3.2 聯絡與回覆程序",
    "### 3.3 隱私保護與受試者自主權宣告": "### 3.3 自願填答與資料處理",
    "## 4. 問卷構念與題項設計矩陣": "## 4. 問卷題項與分析方式",
    "## 5. 完整課後學員追蹤調查問卷（正式規格）": "## 5. 課後學員追蹤問卷",
    "## 7. 追蹤對象流向與收案對照分析（Cohort Flow）": "## 7. 追蹤對象與回覆流向",
    "### 7.1 雙波次學員追蹤狀態逐列對照矩陣": "### 7.1 兩波次逐人狀態",
    "### 7.2 追蹤流向平衡驗證（Reconciliation Proof）": "### 7.2 每波收案核對",
    "## 9. 示範案例 SAMPLE-0122-C01 列層級完整原始數據集": "## 9. 合成案例原始數據",
    "### 9.3 質性反思個案記錄（Qualitative Reflections）": "### 9.3 合成反思示例",
    "## 10. 統計分析、公式代入與算式驗證": "## 10. 統計結果與算式核對",
    "## 11. 結果判讀限制與證據強度界定": "## 11. 結果限制",
    "## 12. 成效落差改善行動方案與重啟／結案標準": "## 12. 改善行動與結案條件",
    "### 12.1 成效落差發現與改善行動對照表": "### 12.1 合成發現與建議措施",
    "### 12.2 流程重啟門檻（Reopen Triggers）與實體結案標準（Closure Criteria）": "### 12.2 重啟及結案條件",
    "## 13. 虛擬資料替換對照表與真實驗證清單": "## 13. 合成資料替換與真實資料檢核",
    "## 14. 來源索引、獨立審查與治理狀態宣告": "## 14. 來源與審查狀態",
    "### 14.1 來源索引與品質合約對照": "### 14.1 來源索引與範圍",
    "### 14.2 獨立審查與治理狀態簽署": "### 14.2 審查及治理狀態",
}
for old, new in heading_map.items():
    text = text.replace(old, new)
text = text.replace("| **P09** | 合格結訓 | 成功接觸 | 接觸後未回 | **無效(未回)** | 成功接觸 | 完成回卷 | **有效回卷** | 合成情境：30天未回覆、90天回覆；無真實差假紀錄 |", "| **P09** | 合格結訓 | 成功接觸 | 接觸後未回 | **無效（未回）** | 成功接觸 | 完成回卷 | **有效回卷** | 合成回覆狀態；沒有真實差假或聯絡紀錄 |")
text = text.replace("| **P07** | 合格結訓 | 成功接觸 | 完成回卷 | **有效回卷** | 成功接觸 | 接觸後未回 | **無效(未回)** | 合成情境：90天接觸後未回覆；無真實拒答紀錄 |", "| **P07** | 合格結訓 | 成功接觸 | 完成回卷 | **有效回卷** | 成功接觸 | 接觸後未回 | **無效（未回）** | 合成回覆狀態；無真實拒答紀錄 |")
text = text.replace("| **P08** | 合格結訓 | 成功接觸 | 完成回卷 | **有效回卷** | 無法接觸 | 無法回卷 | **無效(失聯)** | 合成情境：90天無法聯絡；無真實離職或通話紀錄 |", "| **P08** | 合格結訓 | 成功接觸 | 完成回卷 | **有效回卷** | 無法接觸 | 無法回卷 | **無效（失聯）** | 合成聯絡狀態；無真實離職或通話紀錄 |")
text = text.replace("| **P10** | 合格結訓 | 無法接觸 | 無法回卷 | **無效(失聯)** | 無法接觸 | 無法回卷 | **無效(失聯)** | 合成情境：兩波均無回覆；無真實出國、換號或聯絡紀錄 |", "| **P10** | 合格結訓 | 無法接觸 | 無法回卷 | **無效（失聯）** | 無法接觸 | 無法回卷 | **無效（失聯）** | 合成聯絡狀態；無真實出國、換號或聯絡紀錄 |")
text = text.replace("**P05** | **有效回卷** | 成功接觸 | 完成回卷 | **有效回卷** | 30天波次漏填Q06；90天全填", "**P05** | **有效回卷** | 成功接觸 | 完成回卷 | **有效回卷** | 30天合成資料漏答Q06；90天合成資料全答")


# Reflow the cohort-flow diagram into a comparison table and split the dense
# longitudinal roster into readable wave-specific tables.
flow_start = chr(96)*3 + "\n【追蹤流向圖解 (Cohort Flow Reconciliation)】"
flow_end = "\n" + chr(96)*3 + "\n\n### 7.1"
a = text.index(flow_start)
b = text.index(flow_end, a) + len("\n" + chr(96)*3)
flow = """**合成波次收案摘要**

| 波次 | 應追蹤 | 已嘗試聯繫學員 | 成功接觸 | 無法接觸 | 已回覆 | 接觸後未回 | 有效回卷／有效率 |
| :--- | :---: | :---: | :---: | :---: | :---: | :---: | :---: |
| 第30天 | 10 | 10 | 9 | 1 | 8 | 1 | 8／10 = 80.0% |
| 第90天 | 10 | 10 | 8 | 2 | 7 | 1 | 7／10 = 70.0% |

> 本表數字均為合成示範。聯繫學員數是曾啟動聯繫的對象數，不代表電話或訊息的實際嘗試次數；依聯絡規則每位學員最多三次。每一波各以10人為應追蹤分母，接觸後未回覆與無法接觸分開列計。"""
text = text[:a] + flow + text[b:]
old_start = "| 學員代碼 | 結訓資格 | 第30天接觸狀態 | 第30天回覆狀態 | 第30天有效判定 | 第90天接觸狀態 | 第90天回覆狀態 | 第90天有效判定 | 異動/流失原因說明 |"
a = text.index(old_start)
b = text.index("\n\n### 7.2", a)
tables = """本示範名冊共10位合成合格結訓學員。兩個波次分開判定；30天的狀態不沿用為90天結果。

**第30天波次逐人狀態**

| 學員代碼 | 接觸狀態 | 回覆狀態 | 有效判定 | 示範註記 |
| :---: | :--- | :--- | :--- | :--- |
| P01 | 成功接觸 | 完成回卷 | 有效回卷 | — |
| P02 | 成功接觸 | 完成回卷 | 有效回卷 | — |
| P03 | 成功接觸 | 完成回卷 | 有效回卷 | — |
| P04 | 成功接觸 | 完成回卷 | 有效回卷 | — |
| P05 | 成功接觸 | 完成回卷 | 有效回卷 | Q06漏答；題項統計依有效作答分母計算 |
| P06 | 成功接觸 | 完成回卷 | 有效回卷 | — |
| P07 | 成功接觸 | 完成回卷 | 有效回卷 | — |
| P08 | 成功接觸 | 完成回卷 | 有效回卷 | — |
| P09 | 成功接觸 | 接觸後未回 | 無效（未回） | 合成狀態；無真實拒答紀錄 |
| P10 | 無法接觸 | 無法回卷 | 無效（失聯） | 合成狀態；非真實出國、換號或聯絡紀錄 |

**第90天波次逐人狀態**

| 學員代碼 | 接觸狀態 | 回覆狀態 | 有效判定 | 示範註記 |
| :---: | :--- | :--- | :--- | :--- |
| P01 | 成功接觸 | 完成回卷 | 有效回卷 | — |
| P02 | 成功接觸 | 完成回卷 | 有效回卷 | — |
| P03 | 成功接觸 | 完成回卷 | 有效回卷 | — |
| P04 | 成功接觸 | 完成回卷 | 有效回卷 | — |
| P05 | 成功接觸 | 完成回卷 | 有效回卷 | 90天核心題均完成 |
| P06 | 成功接觸 | 完成回卷 | 有效回卷 | — |
| P07 | 成功接觸 | 接觸後未回 | 無效（未回） | 合成狀態；無真實拒答紀錄 |
| P08 | 無法接觸 | 無法回卷 | 無效（失聯） | 合成狀態；非真實通話或聯絡紀錄 |
| P09 | 成功接觸 | 完成回卷 | 有效回卷 | 合成有效回卷；非真實會員回覆 |
| P10 | 無法接觸 | 無法回卷 | 無效（失聯） | 合成狀態；非真實通話或聯絡紀錄 |"""
text = text[:a] + tables + text[b:]
status_disclosure = "**本案例狀態：** " + chr(96) + "NOT_AVAILABLE_HELD" + chr(96) + "。本表未附任何真實第三方資料；本文件也沒有模擬簽名或假造的觀察紀錄。"
intro = "下列合成案例以同一批 10 個示範代碼呈現 30 天及 90 天收案流向。這些數字不是協會真實追蹤成果；每一波分母分開計算，並保留未回覆與失聯人數。"
text = replace_once(text, status_disclosure, "", "relocate third-party evidence hold")
text = replace_once(text, intro, intro + " 本表單的第三方觀察狀態為 NOT_AVAILABLE_HELD；本示範沒有真實第三方資料、觀察紀錄、簽名或核准。", "move evidence hold to cohort note")


MANUSCRIPT.write_text(text, encoding="utf-8", newline="\n")

# Calculate both waves from the exact data used in the manuscript.
def compute_wave(name: str, rows: dict[str, list[int | None]], cat: dict[str, list[object]]) -> dict[str, object]:
    stats = {}
    for i, key in enumerate(["Q01", "Q02", "Q03", "Q04", "Q05", "Q06", "Q07", "Q10"]):
        vals = [row[i] for row in rows.values()]
        valid = [x for x in vals if x is not None]
        stats[key] = {"n": len(valid), "missing": len(vals) - len(valid), "sum": sum(valid), "mean": round_half_up(sum(valid) / len(valid), 2), "favorable_n_ge_4": sum(x >= 4 for x in valid), "favorable_percent": round_half_up(sum(x >= 4 for x in valid) / len(valid) * 100, 1)}
    stats["Q08_counts"] = {str(k): v for k, v in counts_for({p: {"q": int(row[0])} for p, row in cat.items()}, "q", range(1, 7)).items()}
    stats["Q09_counts"] = {str(k): v for k, v in counts_for({p: {"q": int(row[1])} for p, row in cat.items()}, "q", range(1, 6)).items()}
    stats["Q12_counts"] = {str(k): v for k, v in counts_for({p: {"q": int(row[3])} for p, row in cat.items()}, "q", range(1, 6)).items()}
    opportunity = sum(row[2] >= 2 for row in rows.values())
    applied = sum(row[3] >= 2 for row in rows.values())
    stats["skill_use"] = {"valid_respondents": len(rows), "opportunity_n": opportunity, "applied_n": applied, "all_valid_rate_percent": round(applied / len(rows) * 100, 1), "opportunity_adjusted_rate_percent": round(applied / opportunity * 100, 1)}
    return {"wave": name, "rows": len(rows), "stats": stats}

calc = {
    "project_id": "TTQS_ONE", "doc_id": "0122", "requirement_id": "I19-E10",
    "source_packet_sha256": "8301b3e12467370086dfccd86c4dbcd50d339f0ed54cda93c2bfcc051b793c3d",
    "unique_blueprint_sha256": "cbfd0760a03870151d9bc083b8db76655d53107757706f4b1f4449e997553420",
    "official_source_sha256": "805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709",
    "synthetic_only": True,
    "waves": [compute_wave("30_day", d30_num, cat30), compute_wave("90_day", d90_num, cat90)],
    "cohort_flow": {
        "30_day": {"due": 10, "attempted_contact": 10, "reached": 9, "unreachable": 1, "returned": 8, "reached_nonrespondent": 1, "valid_forms": 8, "valid_rate_percent": 80.0},
        "90_day": {"due": 10, "attempted_contact": 10, "reached": 8, "unreachable": 2, "returned": 7, "reached_nonrespondent": 1, "valid_forms": 7, "valid_rate_percent": 70.0},
    },
    "sample_test": {"eligible": 10, "completed": 8, "met_sample_threshold": 6, "below_sample_threshold": 2, "not_attempted": 2, "coverage_percent": 80.0, "met_threshold_among_completed_percent": 75.0, "met_threshold_among_eligible_percent": 60.0, "real_certificate_evidence": "NOT_PROVIDED"},
}
CALC_RECEIPT.write_text(json.dumps(calc, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=95, bottom=90, end=95) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "D9E0E8")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_font(run, size=10.5, bold=False, color="000000", italic=False):
    run.font.name = "Microsoft JhengHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def latex_plain(s: str) -> str:
    s = re.sub(r"\\text\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\mathrm\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\frac\{([^{}]*)\}\{([^{}]*)\}", r"\1/\2", s)
    s = s.replace("\\times", "×").replace("\\ge", "≥").replace("\\le", "≤")
    s = s.replace("\\%", "%").replace("\\_", "_")
    s = s.replace("\\left", "").replace("\\right", "")
    s = s.replace("$", "").replace("{", "").replace("}", "")
    s = re.sub(r"\s+", " ", s).strip()
    return s


def inline_text(s: str) -> str:
    s = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", s)
    s = re.sub(r"\*\*(.*?)\*\*", r"\1", s)
    s = re.sub(r"(?<!\*)\*(.*?)\*(?!\*)", r"\1", s)
    s = re.sub(r"`([^`]*)`", r"\1", s)
    s = s.replace("\\|", "／")
    formula = s.strip()
    if formula.startswith("$$") and formula.endswith("$$"):
        return latex_plain(formula[2:-2])
    else:
        s = s.replace("$", "")
    return s


def add_rich_paragraph(doc, source: str, style=None, italic=False, align=None):
    p = doc.add_paragraph(style=style)
    style_sizes = {"Heading 1": 15, "Heading 2": 12.5, "Heading 3": 11}
    run_size = style_sizes.get(style, 10.5)
    run_bold = style in style_sizes
    if align is not None:
        p.alignment = align
    # Interpret bold, code and links without carrying Markdown syntax into Word.
    token = re.compile(r"(\*\*.*?\*\*|`[^`]*`|\[[^\]]+\]\([^)]+\))")
    pos = 0
    for match in token.finditer(source):
        if match.start() > pos:
            run = p.add_run(inline_text(source[pos:match.start()]))
            set_font(run, size=run_size, bold=run_bold, italic=italic)
        raw_token = match.group(0)
        bold = raw_token.startswith("**")
        clean = inline_text(raw_token)
        run = p.add_run(clean)
        set_font(run, size=run_size, bold=(bold or run_bold), italic=italic)
        pos = match.end()
    if pos < len(source):
        run = p.add_run(inline_text(source[pos:]))
        set_font(run, size=run_size, bold=run_bold, italic=italic)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.12
    return p


def parse_table_rows(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        cells = re.split(r"(?<!\\)\|", line.strip().strip("|"))
        cells = [c.strip() for c in cells]
        if cells and all(re.fullmatch(r":?-{2,}:?", c.replace(" ", "")) for c in cells):
            continue
        rows.append(cells)
    width = max(len(r) for r in rows)
    return [r + [""] * (width - len(r)) for r in rows]


def set_section_orientation(section, landscape: bool):
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11.69)
        section.page_height = Inches(8.27)
        section.left_margin = Inches(0.42)
        section.right_margin = Inches(0.42)
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.left_margin = Inches(0.70)
        section.right_margin = Inches(0.70)
        section.top_margin = Inches(0.68)
        section.bottom_margin = Inches(0.65)


doc = Document()
set_section_orientation(doc.sections[0], False)
styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft JhengHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.12
for name, size in (("Title", 23), ("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)):
    st = styles[name]
    st.font.name = "Microsoft JhengHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
    st.paragraph_format.space_after = Pt(6)

header = doc.sections[0].header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = header.add_run("原始點關懷協會｜SAMPLE／SYNTHETIC 示範文件")
set_font(run, 8.5, color="475569")
footer = doc.sections[0].footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("課後追蹤與證照通過分析表　｜　僅供流程示範　｜　第 ")
set_font(run, 8.5, color="475569")
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
footer._p.append(fld)
run = footer.add_run(" 頁")
set_font(run, 8.5, color="475569")

lines = MANUSCRIPT.read_text(encoding="utf-8").splitlines()

def next_table_orientation(start_index: int, heading_level: int) -> bool | None:
    """Resolve orientation from this heading or its first child block without borrowing a later sibling table."""
    for j in range(start_index + 1, len(lines)):
        candidate = lines[j].strip()
        if not candidate or candidate == "---" or candidate.startswith(chr(96)*3):
            continue
        if candidate.startswith("#"):
            level = len(candidate) - len(candidate.lstrip("#"))
            if level <= heading_level:
                return None
            return next_table_orientation(j, level)
        if candidate.startswith("|"):
            table_lines = []
            while j < len(lines) and lines[j].lstrip().startswith("|"):
                table_lines.append(lines[j].strip())
                j += 1
            parsed = parse_table_rows(table_lines)
            return len(parsed[0]) >= 5
    return None

def move_to_orientation(landscape: bool) -> None:
    global current_landscape
    if landscape == current_landscape:
        return
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_orientation(sec, landscape)
    current_landscape = landscape

i = 0
current_landscape = False
first_title = True
while i < len(lines):
    line = lines[i].rstrip()
    if not line.strip() or line.strip() == "---":
        i += 1
        continue
    if line.startswith("```"):
        i += 1
        continue
    if line.startswith("|"):
        table_lines = []
        while i < len(lines) and lines[i].lstrip().startswith("|"):
            table_lines.append(lines[i].strip())
            i += 1
        parsed = parse_table_rows(table_lines)
        ncols = len(parsed[0])
        need_landscape = ncols >= 5
        if need_landscape != current_landscape:
            if len(doc.paragraphs) or len(doc.tables):
                sec = doc.add_section(WD_SECTION.NEW_PAGE)
            else:
                sec = doc.sections[-1]
            set_section_orientation(sec, need_landscape)
            current_landscape = need_landscape
        table = doc.add_table(rows=1, cols=ncols)
        table.autofit = True
        table.alignment = 1
        table.style = "Table Grid"
        set_table_borders(table)
        header_row = table.rows[0]
        set_repeat_table_header(header_row)
        for c, value in enumerate(parsed[0]):
            cell = header_row.cells[c]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(inline_text(value))
            set_font(r, 9.5, bold=True, color="FFFFFF")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, "24476B")
            set_cell_margins(cell, top=55, start=85, bottom=55, end=85)
        for ridx, row_data in enumerate(parsed[1:], start=1):
            row = table.add_row()
            set_cant_split(row)
            for c, value in enumerate(row_data):
                cell = row.cells[c]
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if len(value) > 24 else WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(inline_text(value))
                set_font(r, 9.5)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell, top=55, start=85, bottom=55, end=85)
                if ridx % 2 == 0:
                    set_cell_shading(cell, "F2F6FA")
        # Avoid a trailing spacer paragraph that can spill onto an otherwise empty page.
        continue
    if line.startswith("# "):
        title = inline_text(line[2:])
        p = doc.add_paragraph(style="Title")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title)
        set_font(r, 23, bold=True)
        p.paragraph_format.space_after = Pt(8)
        first_title = False
        i += 1
        continue
    if line.startswith("## "):
        target_landscape = next_table_orientation(i, 2)
        move_to_orientation(bool(target_landscape))
        add_rich_paragraph(doc, line[3:], style="Heading 1")
        i += 1
        continue
    if line.startswith("### "):
        target_landscape = next_table_orientation(i, 3)
        move_to_orientation(bool(target_landscape))
        add_rich_paragraph(doc, line[4:], style="Heading 2")
        i += 1
        continue
    if line.startswith("#### "):
        target_landscape = next_table_orientation(i, 4)
        move_to_orientation(bool(target_landscape))
        add_rich_paragraph(doc, line[5:], style="Heading 3")
        i += 1
        continue
    if line.startswith(">"):
        add_rich_paragraph(doc, line.lstrip("> "), italic=True)
        i += 1
        continue
    if re.match(r"^[-*] \[[ xX]\] ", line):
        marker = "☐" if "[ ]" in line else "☒"
        add_rich_paragraph(doc, marker + " " + line[6:])
        i += 1
        continue
    if re.match(r"^[-*] ", line):
        p = add_rich_paragraph(doc, "• " + line[2:])
        p.paragraph_format.left_indent = Inches(0.18)
        i += 1
        continue
    line = re.sub(r"^\s*\d+\.\s+", "", line)
    add_rich_paragraph(doc, line)
    i += 1

doc.core_properties.title = "課後追蹤與證照通過分析表"
doc.core_properties.subject = "TTQS I19-E10 的合成課後追蹤與證照狀態示範候選"
doc.core_properties.author = "原始點關懷協會文件範例"
doc.core_properties.keywords = "SAMPLE, SYNTHETIC, DOCX 0122, I19-E10"
doc.save(DOCX)

with zipfile.ZipFile(DOCX) as zf:
    bad = zf.testzip()
    if bad:
        raise RuntimeError(f"DOCX ZIP CRC failed: {bad}")
    for name in zf.namelist():
        if name.endswith(".xml") or name.endswith(".rels"):
            from lxml import etree
            etree.fromstring(zf.read(name))

docx_hash = hashlib.sha256(DOCX.read_bytes()).hexdigest()
man_hash = hashlib.sha256(MANUSCRIPT.read_bytes()).hexdigest()
calc["candidate_docx_sha256"] = docx_hash
calc["repaired_manuscript_sha256"] = man_hash
calc["docx_bytes"] = DOCX.stat().st_size
calc["build_timestamp_utc"] = datetime.now(timezone.utc).isoformat()
CALC_RECEIPT.write_text(json.dumps(calc, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps({"manuscript": str(MANUSCRIPT), "manuscript_sha256": man_hash, "docx": str(DOCX), "docx_sha256": docx_hash, "docx_bytes": DOCX.stat().st_size, "calculation_receipt": str(CALC_RECEIPT)}, ensure_ascii=False, indent=2))
