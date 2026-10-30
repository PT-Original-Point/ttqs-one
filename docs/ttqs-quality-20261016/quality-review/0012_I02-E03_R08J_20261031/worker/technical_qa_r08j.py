from pathlib import Path
from datetime import datetime
from zipfile import ZipFile
from decimal import Decimal
from docx import Document
import csv,hashlib,json
base=Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0012_I02-E03_LUNA_TEXT_DRAFT_R01_20261029")
work=base/"worker"/"R08J_SCOPED_LABEL_CLARIFICATION_20261031"
qa=base/"qa"/"R08J_SCOPED_LABEL_CLARIFICATION_20261031"
candidate=base/"candidate"/"0012_I02-E03_R08J.docx"
sourcep=work/"R08J_INTEGRATED_SOURCE.md"
registerp=work/"SYNTHETIC_REPLACEMENT_REGISTER_R08J.csv"
buildp=qa/"BUILD_RECEIPT_R08J.json"
parent=base/"candidate"/"0012_I02-E03_R08I.docx"
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def require(cond,msg):
    if not cond: raise SystemExit("FAIL: "+msg)
build=json.loads(buildp.read_text(encoding="utf-8"))
source=sourcep.read_text(encoding="utf-8-sig")
with registerp.open("r",encoding="utf-8-sig",newline="") as f: reg=list(csv.DictReader(f))
cols=list(reg[0]); ids=[r["ID"] for r in reg]
require(build["artifact_sha256"]==sha(candidate),"build receipt candidate SHA mismatch")
require(build["source_sha256"]==sha(sourcep),"build receipt source SHA mismatch")
require(build["synthetic_register_sha256"]==sha(registerp),"build receipt register SHA mismatch")
require(build["parent_candidate_sha256"]==sha(parent),"parent candidate SHA mismatch")
require(build["document_id"]=="0012" and build["requirement_id"]=="I02-E03","doc/requirement identity mismatch")
require(len(reg)==68 and ids==[f"SYN-{i:03d}" for i in range(1,69)],"CSV IDs/count are not sequential 68 rows")
require(len(set(ids))==68 and all(str(v).strip() for row in reg for v in row.values()),"CSV duplicate ID or blank cell")
sec=source.split("## 12.",1)[1].split("## 13.",1)[0]
mlines=[x.strip() for x in sec.splitlines() if x.strip().startswith("|")]
mdrows=[[x.strip() for x in line.strip("|").split("|")] for line in mlines[2:]]
csvrows=[[row[c] for c in cols] for row in reg]
require(mdrows==csvrows,"Markdown register rows differ from CSV")
with ZipFile(candidate) as z: require(z.testzip() is None,"DOCX ZIP CRC")
doc=Document(candidate); paras=doc.paragraphs
alltext=chr(10).join(p.text for p in paras)+chr(10)+chr(10).join(c.text for t in doc.tables for row in t.rows for c in row.cells)
for term in ("已支付金額 P","未付餘額 U","付款原始佐證來源與定位","付款來源核驗狀態","U=max(0","UNKNOWN／未核驗","仍為HOLD","SYN-066","SYN-067","SYN-068","草案版本 0.2","折抵後未付餘額","P=0僅為算例假設"):
    require(term in source and term in alltext,"repair content missing in source or DOCX: "+term)
for bad in ("示例應付學費","封面、文末","R08H","R08J","待填"):
    require(bad not in alltext,"reader-facing/internal/placeholder text found: "+bad)
row42=next(r for r in reg if r["ID"]=="SYN-042")
require("草案版本 0.2" in row42["合成事實／參數及值"] and "§13" in row42["文內定位"],"SYN-042 wrong value/locator")
row66=next(r for r in reg if r["ID"]=="SYN-066"); row67=next(r for r in reg if r["ID"]=="SYN-067"); row68=next(r for r in reg if r["ID"]=="SYN-068")
require("SAMPLE" in row66["合成事實／參數及值"] and "非原件證明" in row66["合成事實／參數及值"],"SYN-066 sample boundary missing")
require("6,000" in row67["合成事實／參數及值"] and "U=max(0,E-P)" in row67["合成事實／參數及值"],"SYN-067 calculation missing")
require("UNKNOWN" in row68["合成事實／參數及值"] and "HOLD" in row68["合成事實／參數及值"],"SYN-068 missing hold boundary")
require("示例折抵後剩餘可用預算額度" in source and "C₃ − B = 7,000 − 1,200 = 5,800" in source, "remaining sample budget label/formula missing")
require("示例折抵後剩餘可用預算額度" in alltext and "NTD 5,800" in alltext, "clarified budget label missing from DOCX")
require("假設折抵後餘額" not in source and "假設折抵後餘額" not in alltext, "ambiguous budget-balance label remains")
row29=next(r for r in reg if r["ID"]=="SYN-029")
require("剩餘可用預算額度" in row29["合成事實／參數及值"] and "不可與申請人的未付學費混淆" in row29["責任角色與替換控制"], "SYN-029 register clarification missing")
F=Decimal("6000"); O=Decimal("0"); X=Decimal("0"); P=Decimal("0")
E=max(Decimal("0"),F-O-X); U=max(Decimal("0"),E-P); fee=E*Decimal("0.20")
B=min([fee,Decimal("1500"),Decimal("4000"),Decimal("7000"),U])
remaining=max(Decimal("0"),U-B); budget_remaining=Decimal("7000")-B
require((E,U,fee,B,remaining,budget_remaining)==(Decimal("6000"),Decimal("6000"),Decimal("1200.00"),Decimal("1200.00"),Decimal("4800.00"),Decimal("5800.00")),"SAMPLE math inconsistency")
register_tables=[t for t in doc.tables if any(c.text.strip()=="SYN-068" for row in t.rows for c in row.cells)]
require(len(register_tables)==1 and len(register_tables[0].rows)==69,"DOCX register table shape mismatch")
worked=[t for t in doc.tables if t.rows and any(c.text.strip()=="輸入／計算步驟" for c in t.rows[0].cells)]
require(len(worked)==1,"worked calculation table not found")
field_labels=("已支付金額 P","未付餘額 U","付款原始佐證來源與定位","付款來源核驗狀態")
form_tables=[t for t in doc.tables if all(any(label in c.text.strip() for row in t.rows for c in row.cells) for label in field_labels)]
require(len(form_tables)>=1,"payment fields not fully present in operational form")
def font_pt(run,para):
    if run.font.size is not None:return run.font.size.pt
    st=para.style
    while st is not None:
        if st.font.size is not None:return st.font.size.pt
        st=st.base_style
    return 11.0
body_sizes=[font_pt(r,p) for p in paras for r in p.runs if r.text.strip()]
cell_sizes=[font_pt(r,p) for t in doc.tables for row in t.rows for cell in row.cells for p in cell.paragraphs for r in p.runs if r.text.strip()]
body_min=min(body_sizes) if body_sizes else None; table_min=min(cell_sizes) if cell_sizes else None
require(body_min is not None and body_min>=10.5,f"body font below 10.5pt: {body_min}")
require(table_min is not None and table_min>=9.5,f"table font below 9.5pt: {table_min}")
result={
 "receipt_id":"0012_I02-E03_R08J_TECHNICAL_QA","observed_at_local":datetime.now().astimezone().isoformat(timespec="seconds"),
 "document_id":"0012","requirement_id":"I02-E03","official_checkbox_serial":"UNKNOWN_NOT_RECONSTRUCTED_HERE",
 "candidate_state":"CANDIDATE_SAMPLE_SYNTHETIC_NOT_ACCEPTED","candidate_path":str(candidate),
 "candidate_sha256":sha(candidate),"candidate_bytes":candidate.stat().st_size,"parent_candidate_sha256":sha(parent),
 "source_sha256":sha(sourcep),"register_sha256":sha(registerp),"register_rows_csv":len(reg),"register_rows_markdown":len(mdrows),
 "register_full_row_match":"PASS","docx_table_count":len(doc.tables),"register_table_rows_including_header":len(register_tables[0].rows),
 "operational_payment_fields":"PASS","SYN-042_value_and_locator":"PASS","new_synthetic_facts":"SYN-066..SYN-068",
 "archive_crc":"PASS",
 "calculation_checks":{"E":"6000","P":"0 SAMPLE assumption only","U":"6000","rate_amount":"1200","B":"1200",
 "remaining_unpaid_eligible_fee":"4800","remaining_sample_budget":"5800","math":"PASS","real_payment_evidence":"UNKNOWN_HOLD"},
 "font_floor":{"body_min_pt":body_min,"table_min_pt":table_min,"body_threshold_pt":10.5,"table_threshold_pt":9.5,"status":"PASS"},
 "technical_status":"PASS","visual_status":"PENDING","semantic_status":"PENDING","formal_acceptance":"NOT_GRANTED",
 "unresolved":["independent semantic review of exact R08J SHA","Word native full-page visual review of exact R08J SHA"]
}
out=qa/"TECHNICAL_QA_R08J.json"
out.write_text(json.dumps(result,ensure_ascii=False,indent=2)+chr(10),encoding="utf-8")
print(json.dumps(result,ensure_ascii=False,indent=2))
