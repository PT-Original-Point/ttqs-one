from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
BODY = ROOT / "candidate" / "0088_R07_body.md"
OUTPUT = ROOT / "candidate" / "0088_R07.docx"
RECEIPT = ROOT / "evidence" / "build_receipt_R07.json"
FONT = "Microsoft JhengHei"
PAGE_WIDTH = Inches(8.27)
PAGE_HEIGHT = Inches(11.69)
MARGIN_LEFT = Inches(0.72)
MARGIN_RIGHT = Inches(0.72)
USABLE_WIDTH = PAGE_WIDTH - MARGIN_LEFT - MARGIN_RIGHT
NAVY = "1F3A5F"
TEAL = "2B6F77"
PALE_BLUE = "EEF4F8"
PALE_GOLD = "FFF4D6"
TEXT = "243447"
BLACK = "000000"
MUTED = "5E6B75"
WHITE = "FFFFFF"


def set_run_font(run, size: float, color: str = TEXT, bold: bool | None = None) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "5")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "CBD5DE")


def set_keep_with_next(paragraph, enabled: bool = True) -> None:
    paragraph.paragraph_format.keep_with_next = enabled


def add_field(paragraph, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = OxmlElement("w:r")
    run.append(begin)
    run2 = OxmlElement("w:r")
    run2.append(instr)
    run3 = OxmlElement("w:r")
    run3.append(separate)
    run3.append(text)
    run4 = OxmlElement("w:r")
    run4.append(end)
    paragraph._p.append(run)
    paragraph._p.append(run2)
    paragraph._p.append(run3)
    paragraph._p.append(run4)


def add_inline(paragraph, text: str, *, base_size: float, base_color: str = TEXT, force_bold: bool = False) -> None:
    pattern = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            set_run_font(run, base_size, base_color, force_bold)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, base_size, base_color, True)
        else:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
            run.font.size = Pt(max(base_size - 0.2, 9.5))
            run.font.color.rgb = RGBColor.from_string(TEAL)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, base_size, base_color, force_bold)


def split_table_row(line: str) -> list[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [cell.strip() for cell in s.split("|")]


def is_table_separator(line: str) -> bool:
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells)


def apply_paragraph_defaults(paragraph, *, size: float = 10.5, color: str = TEXT) -> None:
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.16
    for run in paragraph.runs:
        set_run_font(run, size, color)


def add_markdown_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    ncols = max(len(row) for row in rows)
    table = document.add_table(rows=1, cols=ncols)
    table.autofit = False
    table.alignment = 1
    table.style = "Table Grid"
    widths = [Inches(USABLE_WIDTH / 914400 / ncols) for _ in range(ncols)]
    for col, width in zip(table.columns, widths):
        col.width = width
    set_table_borders(table)
    for r_index, source_row in enumerate(rows):
        row = table.rows[0] if r_index == 0 else table.add_row()
        prevent_row_split(row)
        if r_index == 0:
            set_repeat_table_header(row)
        for c_index in range(ncols):
            cell = row.cells[c_index]
            cell.width = widths[c_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if r_index == 0:
                set_cell_shading(cell, NAVY)
            elif r_index % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            source_text = source_row[c_index] if c_index < len(source_row) else ""
            paragraph = cell.paragraphs[0]
            if r_index == 0:
                paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.04
            add_inline(
                paragraph,
                source_text,
                base_size=9.5,
                base_color=WHITE if r_index == 0 else TEXT,
                force_bold=(r_index == 0),
            )
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def setup_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = PAGE_WIDTH
    section.page_height = PAGE_HEIGHT
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.68)
    section.left_margin = MARGIN_LEFT
    section.right_margin = MARGIN_RIGHT
    section.header_distance = Inches(0.30)
    section.footer_distance = Inches(0.30)

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.16

    title = document.styles["Title"]
    title.font.name = FONT
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    title._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.keep_with_next = True
    title_ppr = title._element.get_or_add_pPr()
    for border in list(title_ppr.findall(qn("w:pBdr"))):
        title_ppr.remove(border)

    for style_name, size, color in (("Heading 1", 16, BLACK), ("Heading 2", 12.5, BLACK), ("Heading 3", 11, BLACK)):
        style = document.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12 if style_name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("SAMPLE / SYNTHETIC   |   就業及考照追蹤紀錄")
    set_run_font(run, 8.5, MUTED, True)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    run = footer.add_run("就業及考照追蹤紀錄　|　第 ")
    set_run_font(run, 8.5, MUTED)
    add_field(footer, "PAGE")
    run = footer.add_run(" 頁")
    set_run_font(run, 8.5, MUTED)

    settings = document.settings.element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def build() -> dict[str, object]:
    text = BODY.read_text(encoding="utf-8")
    lines = text.splitlines()
    document = Document()
    setup_document(document)
    table_count = 0
    heading_count = 0
    i = 0
    while i < len(lines):
        raw = lines[i]
        stripped = raw.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("|"):
            table_lines: list[str] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            parsed = [split_table_row(line) for line in table_lines]
            if len(parsed) > 1 and is_table_separator(table_lines[1]):
                parsed.pop(1)
            add_markdown_table(document, parsed)
            table_count += 1
            continue
        heading_match = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            heading_text = heading_match.group(2).strip()
            if level == 1 and heading_count == 0:
                paragraph = document.add_paragraph(style="Title")
                add_inline(paragraph, heading_text, base_size=25, base_color=BLACK, force_bold=True)
            else:
                word_level = 1 if level == 2 else 2
                paragraph = document.add_paragraph(style=f"Heading {word_level}")
                if heading_text in {"9.2 逐案總表", "14.2 相鄰要求區分"} or heading_text.startswith("9.4 未結事項交接表"):
                    paragraph.paragraph_format.page_break_before = True
                add_inline(paragraph, heading_text, base_size=16 if word_level == 1 else 12.5, base_color=BLACK, force_bold=True)
            heading_count += 1
            i += 1
            continue
        if stripped.startswith(">"):
            quote_lines: list[str] = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote_lines.append(re.sub(r"^\s*>\s?", "", lines[i]).strip())
                i += 1
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.16)
            paragraph.paragraph_format.right_indent = Inches(0.10)
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.paragraph_format.keep_together = True
            p_pr = paragraph._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), PALE_GOLD)
            shd.set(qn("w:val"), "clear")
            p_pr.append(shd)
            add_inline(paragraph, "\n".join(quote_lines), base_size=10.2, base_color=NAVY)
            continue
        bullet_match = re.match(r"^[-*]\s+(.*)$", stripped)
        if bullet_match:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(3)
            add_inline(paragraph, bullet_match.group(1), base_size=10.5)
            i += 1
            continue
        number_match = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if number_match:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.24)
            paragraph.paragraph_format.first_line_indent = Inches(-0.24)
            paragraph.paragraph_format.space_after = Pt(4)
            add_inline(paragraph, f"{number_match.group(1)}. {number_match.group(2)}", base_size=10.5)
            i += 1
            continue

        paragraph_lines = [stripped]
        i += 1
        while i < len(lines):
            next_line = lines[i].strip()
            if (
                not next_line
                or next_line.startswith("|")
                or next_line.startswith(">")
                or re.match(r"^#{1,3}\s+", next_line)
                or re.match(r"^[-*]\s+", next_line)
                or re.match(r"^\d+\.\s+", next_line)
            ):
                break
            paragraph_lines.append(next_line)
            i += 1
        paragraph = document.add_paragraph()
        if stripped in {"**首次電話開場**", "**未接通後的低揭露訊息**", "**拒答或要求停止**", "**自願提供證明**"}:
            paragraph.paragraph_format.keep_with_next = True
        add_inline(paragraph, " ".join(paragraph_lines), base_size=10.5)

    document.core_properties.title = "就業及考照追蹤紀錄"
    document.core_properties.subject = "SAMPLE / SYNTHETIC 示範文件；對應 I17D-E02"
    document.core_properties.author = ""
    document.core_properties.keywords = "TTQS_ONE; 0088; I17D-E02; SAMPLE; SYNTHETIC"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)

    sha = hashlib.sha256(OUTPUT.read_bytes()).hexdigest()
    receipt = {
        "project_id": "TTQS_ONE",
        "doc_id": "0088",
        "requirement_id": "I17D-E02",
        "official_checkbox_text": "就業及考照追蹤紀錄。",
        "official_checkbox_serial": "UNKNOWN_NOT_IN_VERIFIED_SERIAL_REGISTER",
        "revision": "R07",
        "candidate_docx": str(OUTPUT),
        "candidate_docx_sha256": sha,
        "body_sha256": hashlib.sha256(BODY.read_bytes()).hexdigest(),
        "table_count": table_count,
        "heading_count": heading_count,
        "candidate_state": "SAMPLE_SYNTHETIC_CANDIDATE_NOT_FORMALLY_ACCEPTED",
        "builder": "Windows bundled Python 3 + python-docx 1.2.0",
        "source_packet": "inputs/0088_source_packet_R01.md",
        "unique_blueprint": "inputs/0088_UNIQUE_BLUEPRINT_R01.md",
        "body_source": "OpenCode v2 R07 bounded semantic patch; Windows native DOCX builder",
    }
    RECEIPT.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(receipt, ensure_ascii=False))
    return receipt


if __name__ == "__main__":
    build()

