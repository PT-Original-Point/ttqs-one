from __future__ import annotations
import csv, hashlib, json, re, zipfile
from datetime import date, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
ROOT=Path(__file__).resolve().parents[1]
DOCX=ROOT/'candidate'/'0088_R07.docx'; BODY=ROOT/'candidate'/'0088_R07_body.md'
REGISTER=ROOT/'evidence'/'synthetic_replacement_register_R07.csv'; PDF=ROOT/'evidence'/'render_R07_final3'/'0088_R07.pdf'
RR=ROOT/'evidence'/'render_receipt_R07_final3.json'; MAN=ROOT/'evidence'/'render_manifest_R07_final3.json'
BUILD=ROOT/'evidence'/'build_receipt_R07.json'; PACKET=ROOT/'inputs'/'0088_source_packet_R01.md'; BP=ROOT/'inputs'/'0088_UNIQUE_BLUEPRINT_R01.md'
OFFICIAL=ROOT.parent/'0064_R06'/'official_source.pdf'; OUT=ROOT/'evidence'/'technical_review_R07.json'; EXTRACT=ROOT/'evidence'/'docx_extracted_R07.md'
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def check(ok,name,detail,checks):
 checks.append({'check':name,'status':'PASS' if ok else 'FAIL','detail':detail})
 if not ok: raise AssertionError(f'{name}: {detail}')
checks=[]; build=json.loads(BUILD.read_text(encoding='utf-8-sig')); rr=json.loads(RR.read_text(encoding='utf-8-sig')); man=json.loads(MAN.read_text(encoding='utf-8-sig'))
docx_sha=sha(DOCX); body_sha=sha(BODY); pdf_sha=sha(PDF); reg_sha=sha(REGISTER)
check((build['doc_id'],build['requirement_id'])==('0088','I17D-E02'),'identity',f"{build['doc_id']} / {build['requirement_id']}",checks)
check(build['official_checkbox_text']=='就業及考照追蹤紀錄。','official_mapping',build['official_checkbox_text'],checks)
check(build['official_checkbox_serial'].startswith('UNKNOWN'),'serial_unknown',build['official_checkbox_serial'],checks)
check(build['candidate_docx_sha256']==docx_sha and build['body_sha256']==body_sha,'build_receipt_hashes',f'DOCX={docx_sha}; body={body_sha}',checks)
check(rr['docx_sha256']==docx_sha and rr['pdf_sha256']==pdf_sha and rr['status']=='PASS','word_render_hashes',f"Word read-only export; {rr['page_count']} pages",checks)
check(man['docx_sha256']==docx_sha and man['pdf_sha256']==pdf_sha and man['page_count']==rr['page_count'],'render_manifest_hashes',f"{man['page_count']} pages",checks)
with zipfile.ZipFile(DOCX) as z: bad=z.testzip()
check(bad is None,'docx_zip_integrity','ZIP entries readable',checks)
doc=Document(DOCX); check(len(doc.tables)==27,'table_count',str(len(doc.tables)),checks)
check(build['heading_count']==49,'heading_count',str(build['heading_count']),checks)
parts=[]
for child in doc.element.body.iterchildren():
 if child.tag.endswith('}p'):
  s=Paragraph(child,doc).text.strip()
  if s: parts.append(s)
 elif child.tag.endswith('}tbl'):
  t=Table(child,doc); parts.append('\n'.join(' | '.join(c.text.replace('\n',' / ').strip() for c in row.cells) for row in t.rows))
extracted='\n\n'.join(parts)
EXTRACT.write_text('# DOCX extracted text — 0088 R07\n\nDOCX SHA-256: '+docx_sha+'\n\n'+extracted+'\n',encoding='utf-8')
text=extracted+'\n'+BODY.read_text(encoding='utf-8')
forbidden_body_markers=['doc_id=','requirement_id=','serial=UNKNOWN','候選 DOCX SHA','2e42fcbf224bbbe18454f111a9fa99f9722e6a47a54b79d5ed787ae551cb8da0']
check(not any(mark.lower() in extracted.lower() for mark in forbidden_body_markers),'no_engineering_controls_in_body','none of '+', '.join(forbidden_body_markers),checks)
for term in ['就業及考照追蹤紀錄','SAMPLE / SYNTHETIC','PASS TO CONTACT','UNKNOWN / HOLD','9.2 逐案總表','9.4 未結事項交接表','12.1 個案主表','13.1 證據索引節錄','14.2 相鄰要求區分']:
 check(term in text,'required:'+term,term,checks)
ids=[f'C{i:02d}' for i in range(1,13)]; check(all(x in text for x in ids),'twelve_case_ids',', '.join(ids),checks)
check('C11/C12' in text and 'NO_RESPONSE' in text,'noresponse_state','C11/C12 remain UNKNOWN/NO_RESPONSE',checks)
check('SAMPLE / SYNTHETIC' in text and 'UNKNOWN / HOLD' in text,'candidate_boundary','synthetic and real-evidence hold labels present',checks)
with REGISTER.open(encoding='utf-8-sig',newline='') as f: rows=list(csv.DictReader(f)); headers=list(rows[0]) if rows else []
keys=[r.get('synthetic_key','') for r in rows]
check(len(headers)==11,'register_schema',f'{len(headers)} columns',checks)
check(len(rows)==199 and all(keys) and len(keys)==len(set(keys)),'register_keys',f'{len(rows)} rows; {len(set(keys))} unique keys',checks)
check({r['evidence_state'] for r in rows}<={'SYNTHETIC_NOT_REAL','NOT_PROVIDED'},'register_evidence_boundary',str(sorted({r['evidence_state'] for r in rows})),checks)
check({r['third_party_original_required'] for r in rows}<={'YES','NO'},'register_third_party_flags',str(sorted({r['third_party_original_required'] for r in rows})),checks)
case_rows=[r for r in rows if r.get('case_id')]; check(set(r['case_id'].replace('SYN-','') for r in case_rows)<=set(ids),'register_case_ids',f'{len(case_rows)} case-linked rows',checks)
case_table=None
for t in doc.tables:
 if len(t.rows)>1 and len(t.rows[0].cells)==7 and t.rows[0].cells[0].text.strip()=='案件' and '考照證據狀態' in t.rows[0].cells[-1].text:
  case_table=t; break
check(case_table is not None,'case_summary_table','located 9.2 seven-column case table',checks)
state_by_case={r.cells[0].text.strip().replace('SYN-',''):r.cells[4].text.upper() for r in case_table.rows[1:]}
expected={'C01':'OPEN_ACTION','C02':'HOLD','C03':'CLOSED—SAMPLE','C04':'OPEN_ACTION','C05':'CLOSED—SAMPLE','C06':'OPEN_ACTION','C07':'OPEN_ACTION','C08':'OPEN_ACTION','C09':'HOLD','C10':'OPEN_ACTION','C11':'CLOSED—NO_RESPONSE—SAMPLE','C12':'CLOSED—NO_RESPONSE—SAMPLE'}
check(set(state_by_case)==set(expected),'case_row_count_and_ids',f'{len(state_by_case)} unique cases',checks)
check(all(v in state_by_case.get(k,'') for k,v in expected.items()),'case_state_partition',json.dumps(state_by_case,ensure_ascii=False),checks)
metric_results=[]
for a,b,shown in [(10,12,'83.3%'),(6,10,'60.0%'),(6,12,'50.0%'),(3,4,'75.0%'),(0,3,'0.0%')]:
 val=(Decimal(a)*100/Decimal(b)).quantize(Decimal('0.1'),rounding=ROUND_HALF_UP); actual=f'{val}%'
 check(actual==shown,'calculation:'+shown,f'{a}/{b}={actual}',checks); check(shown in text,'printed_calculation:'+shown,shown,checks)
 metric_results.append({'numerator':a,'denominator':b,'display':shown})
check(date(2026,4,30)+timedelta(days=90)==date(2026,7,29),'followup_date','2026-04-30 + 90 days = 2026-07-29',checks)
check(not re.search(r'\b[A-Z][12]\d{8}\b',text),'no_id_pattern','no Taiwan ID-like value found',checks)
check(not re.search(r'(?<!\d)09\d{8}(?!\d)',text),'no_phone_pattern','no Taiwan mobile-like value found',checks)
check(not re.search(r'[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}',text),'no_email_pattern','no email address found',checks)
check(len(man['pages'])==rr['page_count'] and all((PDF.parent/p['image']).exists() and sha(PDF.parent/p['image'])==p['sha256'] for p in man['pages']),'page_images','all rendered page images match manifest hashes',checks)
check(OFFICIAL.exists() and sha(OFFICIAL)=='805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709','official_pdf_hash',sha(OFFICIAL) if OFFICIAL.exists() else 'missing',checks)
check('I17D-E02' in PACKET.read_text(encoding='utf-8') and 'I17D-E02' in BP.read_text(encoding='utf-8'),'source_blueprint_binding','source packet and Unique Blueprint bind 0088 to I17D-E02',checks)
receipt={'project_id':'TTQS_ONE','doc_id':'0088','requirement_id':'I17D-E02','revision':'R07','candidate_state':'SAMPLE_SYNTHETIC_CANDIDATE_NOT_FORMALLY_ACCEPTED','technical_verdict':'PASS_WITH_SCOPE_LIMITS','independent_semantic_acceptance':'PENDING','docx_sha256':docx_sha,'body_sha256':body_sha,'register_sha256':reg_sha,'pdf_sha256':pdf_sha,'source_packet_sha256':sha(PACKET),'unique_blueprint_sha256':sha(BP),'official_source_pdf_sha256':sha(OFFICIAL),'page_count':len(man['pages']),'table_count':len(doc.tables),'metrics':metric_results,'checks':checks,'limitations':['Validates file integrity, document mapping, specified synthetic register structure and arithmetic, render/hash binding and limited PII patterns.','Does not establish any real learner/employment/exam fact, semantic acceptance or formal TTQS acceptance.']}
OUT.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'verdict':receipt['technical_verdict'],'checks':len(checks),'docx_sha256':docx_sha,'pdf_sha256':pdf_sha,'pages':len(man['pages']),'tables':len(doc.tables),'register_rows':len(rows),'extracted_text':str(EXTRACT)},ensure_ascii=False))




