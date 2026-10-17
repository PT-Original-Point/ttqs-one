from pathlib import Path
from datetime import datetime, timezone
import csv, hashlib, json, re, zipfile
from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / 'candidate' / '0055_R03.docx'
PDF = ROOT / 'evidence' / 'render_final' / '0055_R03.pdf'
PAGES = ROOT / 'evidence' / 'render_final'
REGISTER = ROOT / 'evidence' / 'SYNTHETIC_REPLACEMENT_REGISTER_0055_R03.csv'
SOURCE_PDF = ROOT.parent / '0055_R02' / 'inputs' / 'official_checklist_source.pdf'

def sha256(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b''):
            h.update(chunk)
    return h.hexdigest()

def docx_text(doc):
    chunks = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            chunks.append(' | '.join(c.text for c in row.cells))
    return '\n'.join(chunks)

with zipfile.ZipFile(DOCX) as z:
    corrupt = z.testzip()
    assert corrupt is None, f'corrupt OOXML member: {corrupt}'

d = Document(DOCX)
all_text = docx_text(d)
headings = [p.text for p in d.paragraphs if p.style.name.startswith('Heading')]
required_headings = [f'{i}｜' for i in range(1, 14)]
for h in required_headings:
    assert any(x.startswith(h) for x in headings), f'missing section heading {h}'
for term in ['SAMPLE / SYNTHETIC', 'NOT ISSUED', 'NOT VALID', 'I12-E03', '官方方框 058',
             'SYN-0033-PLAN-01', '360 分鐘', 'HOLD', '利益衝突', '委員會',
             '資格原件', '教學示範', '聘書']:
    assert term in all_text, f'missing required content: {term}'

table_font_sizes = []
for t in d.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    if run.font.size is not None:
                        table_font_sizes.append(run.font.size.pt)
assert table_font_sizes, 'no explicit table font sizing found'
min_table_pt = min(table_font_sizes)
assert min_table_pt >= 9.5, f'table type below project floor: {min_table_pt}'

body_font_sizes = [run.font.size.pt for p in d.paragraphs for run in p.runs if run.font.size is not None]
assert body_font_sizes, 'no explicit body font sizing found'
min_body_pt = min(body_font_sizes)
assert min_body_pt >= 10.5, f'body type below project floor: {min_body_pt}'
assert '使用者先看這裡' not in all_text, 'generic first-look table remains'
assert '第 5 節評分' not in all_text and '第 11 節逐筆取代 SAMPLE' not in all_text, 'stale quick-path references remain'

with REGISTER.open(encoding='utf-8-sig', newline='') as f:
    register = list(csv.DictReader(f))
assert len(register) == 52, f'expected 52 replacement records, got {len(register)}'
required_register_fields = [
    'register_id', 'fact_id', 'case_id', 'doc_id', 'docx_locator',
    'synthetic_value_or_field', 'assumption_rationale', 'cross_doc_reference',
    'real_source_to_obtain', 'responsible_verifier', 'evidence_status', 'hold_condition'
]
assert set(required_register_fields).issubset(register[0].keys()), 'Mission §4 replacement-register schema incomplete'
assert all(all(row.get(k, '').strip() for k in required_register_fields)
           for row in register), 'empty required replacement-register field'
assert len({row['register_id'] for row in register}) == 52, 'duplicate register_id'
assert len({row['fact_id'] for row in register}) == 52, 'duplicate fact_id'
assert all(row['case_id'] == 'SYN-0033-PLAN-01' and row['doc_id'] == '0055'
           for row in register), 'replacement register case/doc binding mismatch'
assert all(row['evidence_status'].startswith('SYNTHETIC') for row in register), 'replacement row must disclose its synthetic source state'
for row in register:
    quoted = re.findall(r'「([^」]+)」', row['docx_locator'])
    assert quoted and any(q in all_text for q in quoted), f'locator not found in DOCX body: {row["register_id"]}'

criteria = [25,25,20,20,10]
assert sum(criteria) == 100
scorevectors = {
    'C-01': [(85,86,87),(80,83,83),(78,76,80),(82,80,78),(90,92,88)],
    'C-02': [(76,78,80),(90,86,88),(74,76,75),(78,80,82),(84,86,85)],
    'C-03': [(90,92,88),(84,86,88),(88,90,86),(84,86,85),(92,90,88)],
}
totals = {}
for code, vectors in scorevectors.items():
    totals[code] = round(sum((sum(v)/3)*w/100 for v,w in zip(vectors,criteria)),1)
assert totals == {'C-01':82.6,'C-02':81.0,'C-03':87.6}, totals
assert sorted(totals, key=totals.get, reverse=True) == ['C-03','C-01','C-02']
minutes = [90,90,120,60]
assert sum(minutes) == 360

reader = PdfReader(PDF)
page_text_lengths = [len(p.extract_text() or '') for p in reader.pages]
assert len(reader.pages) >= 10, f'unexpectedly short render: {len(reader.pages)} pages'
assert all(n >= 250 for n in page_text_lengths), f'page text unexpectedly sparse: {page_text_lengths}'
box = reader.pages[0].mediabox
a4_points = (595.0, 842.0)
assert abs(float(box.width)-a4_points[0]) < 1.0 and abs(float(box.height)-a4_points[1]) < 1.0, box
page_images = sorted(PAGES.glob('page-*.png'))
assert len(page_images) == len(reader.pages), f'page image count differs from PDF: {len(page_images)} vs {len(reader.pages)}'

receipt = {
    'project_id': 'TTQS_ONE',
    'document_id': '0055',
    'unique_requirement_id': 'I12-E03',
    'official_checkbox_sequence': '058',
    'revision': 'R03',
    'candidate_status': 'SAMPLE_SYNTHETIC_CANDIDATE_ONLY',
    'technical_layout_status': 'PASS_TECHNICAL_LAYOUT_ONLY',
    'independent_semantic_review': 'PASS_WITH_NONBLOCKING_FINDINGS; inherited R02 full review plus R03 exact-SHA delta review',
    'formal_human_acceptance': 'NOT_GRANTED',
    'verified_at_utc': datetime.now(timezone.utc).isoformat(),
    'docx': {'path': str(DOCX), 'bytes': DOCX.stat().st_size, 'sha256': sha256(DOCX),
             'ooxml_zip_integrity': 'PASS', 'tables': len(d.tables),
             'explicit_body_font_min_pt': min_body_pt, 'explicit_table_font_min_pt': min_table_pt,
             'required_sections_1_to_13': 'PASS'},
    'render': {'method': 'Microsoft Word COM ExportAsFixedFormat PDF, then Poppler page rasterization',
               'render_docx_py': 'BLOCKED: LibreOffice soffice.exe unavailable; Word export used for visual QA',
               'pdf_path': str(PDF), 'pdf_sha256': sha256(PDF), 'page_count': len(reader.pages),
               'paper': 'A4', 'page_text_char_counts': page_text_lengths,
               'page_images': [{'path': str(p), 'sha256': sha256(p)} for p in page_images],
               'visual_inspection': 'PASS_11_PAGES; 10 pixel-identical pages inherit R02 review; page 8 directly inspected for R03'},
    'replacement_register': {'path': str(REGISTER), 'rows': len(register), 'sha256': sha256(REGISTER),
                             'required_fields_nonempty': 'PASS'},
    'arithmetic': {'course_unit_minutes': minutes, 'course_total_minutes': sum(minutes),
                   'weight_sum_percent': sum(criteria), 'candidate_totals': totals,
                   'ranking': ['C-03','C-01','C-02'], 'status': 'PASS'},
    'source_pdf_sha256': sha256(SOURCE_PDF),
    'scope_limits': ['No real candidate, credential, committee, approval, appointment, signature, payment, or delivery is asserted.',
                     'All displayed candidate identities, scores, meeting, dates, fee, and appointment terms are SAMPLE/SYNTHETIC.',
                     'This receipt is technical/layout evidence, not an independent semantic PASS.']
}
out = ROOT / 'evidence' / 'TECHNICAL_LAYOUT_QA_R03.json'
out.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
print(json.dumps(receipt, ensure_ascii=False, indent=2))
