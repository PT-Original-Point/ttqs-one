from pathlib import Path
import hashlib, json, zipfile
from docx import Document

pkg=Path(__file__).resolve().parents[1]; repo=pkg.parents[3]
source=repo/'docs/ttqs-quality-20261016/quality-review/0055_R02/candidate/0055_R02.docx'
source_register=repo/'docs/ttqs-quality-20261016/quality-review/0055_R02/evidence/SYNTHETIC_REPLACEMENT_REGISTER_0055_R02.csv'
expected_source_sha='2d69f2e515c424a8717bfb189e8c9de861580d636fafd82f79d332bab8cfe827'
old='主席 P-01、委員 M-01、M-02、M-03、紀錄員 S-01（全為合成代碼）；示範設定 3/3 有表決權委員出席。'
new='主席 P-01（本示範設定不表決）、表決委員 M-01、M-02、M-03、紀錄員 S-01（全為合成代碼）；示範設定 3/3 表決委員出席。'
old_reg='合成出席代碼：主席 P-01、委員 M-01/M-02/M-03、紀錄員 S-01；示範設定 3/3 有表決權委員出席'
new_reg='合成出席代碼：主席 P-01（本示範設定不表決）、表決委員 M-01/M-02/M-03、紀錄員 S-01；示範設定 3/3 表決委員出席'
def sha_bytes(b): return hashlib.sha256(b).hexdigest()
def sha(p): return sha_bytes(p.read_bytes())
assert sha(source)==expected_source_sha, f'R02 source changed: {sha(source)}'
assert not (pkg/'candidate/0055_R03.docx').exists(), 'R03 candidate exists; inspect rather than overwrite'
raw=source.read_bytes(); src_tmp=pkg/'candidate/0055_R03.building.docx'; dst=pkg/'candidate/0055_R03.docx'
with zipfile.ZipFile(source,'r') as zin, zipfile.ZipFile(src_tmp,'w') as zout:
 for info in zin.infolist():
  data=zin.read(info.filename)
  if info.filename=='word/document.xml':
   assert data.count(old.encode('utf-8'))==1, 'expected one target sentence in document.xml'
   data=data.replace(old.encode('utf-8'),new.encode('utf-8'),1)
  zout.writestr(info,data)
with zipfile.ZipFile(src_tmp,'r') as z:
 assert z.testzip() is None, 'OOXML integrity failed'
 assert z.namelist()==zipfile.ZipFile(source,'r').namelist(), 'OOXML package entry order/list changed'
 src_xml=zipfile.ZipFile(source).read('word/document.xml'); dst_xml=z.read('word/document.xml')
 assert dst_xml==src_xml.replace(old.encode('utf-8'),new.encode('utf-8'),1), 'document.xml contains unrelated edits'
with zipfile.ZipFile(source) as za, zipfile.ZipFile(src_tmp) as zb:
 assert all(za.read(n)==zb.read(n) for n in za.namelist() if n!='word/document.xml'), 'other OOXML part changed'
text='\n'.join([p.text for p in Document(src_tmp).paragraphs]+[' | '.join(c.text for c in row.cells) for t in Document(src_tmp).tables for row in t.rows])
assert new in text and old not in text, 'new text did not round-trip through python-docx readback'
src_tmp.replace(dst)
reg_bytes=source_register.read_bytes(); assert sha(source_register)=='b7b2948b300e0c5ca36b838bdbca10d1470b0ddc1450257c742b32bd782d8531'
assert reg_bytes.count(old_reg.encode('utf-8'))==1
reg_out=pkg/'evidence/SYNTHETIC_REPLACEMENT_REGISTER_0055_R03.csv'
reg_out.write_bytes(reg_bytes.replace(old_reg.encode('utf-8'),new_reg.encode('utf-8'),1))
blueprint=repo/'docs/ttqs-quality-20261016/quality-review/0055_R02/UNIQUE_BLUEPRINT_0055_R02.md'; source_map=repo/'docs/ttqs-quality-20261016/quality-review/0055_R02/inputs/SOURCE_MAP_R02.md'
report={'project_id':'TTQS_ONE','document_id':'0055','revision':'R03','status':'SAMPLE_SYNTHETIC_CANDIDATE_ONLY','official_requirement_id':'I12-E03','official_checkbox_sequence':'058','predecessor_docx_sha256':expected_source_sha,'candidate_docx_sha256':sha(dst),'predecessor_register_sha256':sha(source_register),'candidate_register_sha256':sha(reg_out),'unique_blueprint_path':blueprint.relative_to(repo).as_posix(),'unique_blueprint_sha256':sha(blueprint),'source_map_path':source_map.relative_to(repo).as_posix(),'source_map_sha256':sha(source_map),'change_locator':'Section 9, attendance/quorum row; table 20 row 3','before':old,'after':new,'register_before':old_reg,'register_after':new_reg,'ooxml_edit':'Only the target text node in word/document.xml changed; all other package parts are byte-identical.','formal_human_acceptance':'NOT_GRANTED'}
(pkg/'repair_scope_R03.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')
print(json.dumps(report,ensure_ascii=False))
