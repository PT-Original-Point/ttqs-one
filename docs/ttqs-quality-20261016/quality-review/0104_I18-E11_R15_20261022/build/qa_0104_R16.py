from pathlib import Path
from zipfile import ZipFile
from fractions import Fraction
from decimal import Decimal, ROUND_HALF_UP
import csv, hashlib, json, datetime, re, xml.etree.ElementTree as ET
from docx import Document

BASE = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0104_I18-E11_SOURCE_PREP_20261022")
DOCX = BASE / "candidate" / "0104__證照獲取率_分析報告_R16.docx"
BUILD = BASE / "evidence" / "docx_build_receipt_0104_R16.json"
REGISTER = BASE / "evidence" / "synthetic_replacement_register_0104_R16.csv"
OUT = BASE / "evidence" / "technical_qa_receipt_0104_R16.json"

raw = DOCX.read_bytes()
sha = hashlib.sha256(raw).hexdigest()
build = json.loads(BUILD.read_text(encoding="utf-8"))
doc = Document(str(DOCX))
paras = [p.text for p in doc.paragraphs]
tables_text = ["\n".join(" | ".join(c.text for c in row.cells) for row in t.rows) for t in doc.tables]
all_text = "\n".join(paras + tables_text)
with ZipFile(DOCX) as z:
    crc = z.testzip()
    xml_names = [n for n in z.namelist() if n.endswith((".xml", ".rels"))]
    for name in xml_names:
        ET.fromstring(z.read(name))

with REGISTER.open(encoding="utf-8-sig", newline="") as f:
    register_rows = list(csv.DictReader(f))
fact_keys = [row.get("事實鍵", "").strip() for row in register_rows]
register_by_fact = {row.get("事實鍵", "").strip(): row for row in register_rows}
register_states = [row.get("狀態", "").strip() for row in register_rows]
zero_denominator_consistent = all(
    register_by_fact.get(f"F-RC-{i:02d}", {}).get("真實替換來源", "").strip() == "0/0 重算"
    for i in range(1, 6)
) and all(register_by_fact.get(f"F-RC-{i:02d}", {}).get("合成值", "").strip() == "N/A" for i in range(1, 6))

def rounded_pct(value, places=2):
    d = Decimal(value.numerator) / Decimal(value.denominator) * Decimal(100)
    q = Decimal(1).scaleb(-places)
    return f"{d.quantize(q, rounding=ROUND_HALF_UP):.{places}f}%"

def rounded_pp(value, places=2):
    d = Decimal(value.numerator) / Decimal(value.denominator) * Decimal(100)
    q = Decimal(1).scaleb(-places)
    return f"{d.quantize(q, rounding=ROUND_HALF_UP):+.{places}f} pp"

periods = {
    "F-RD-01": (Fraction(32, 40), Fraction(33, 36)),
    "F-RD-02": (Fraction(29, 32), Fraction(31, 33)),
    "F-RD-03": (Fraction(22, 29), Fraction(26, 31)),
    "F-RD-04": (Fraction(19, 22), Fraction(24, 26)),
    "F-RD-05": (Fraction(19, 40), Fraction(24, 36)),
    "F-RD-06": (Fraction(22, 40), Fraction(26, 36)),
}
derived_rates_exact = True
for key, (a, b) in periods.items():
    row = register_by_fact.get(key, {})
    expected = rounded_pp(b - a) + "；+" + rounded_pct((b - a) / a)
    if expected not in row.get("合成值", ""):
        derived_rates_exact = False

def sizes(document):
    body, table = [], []
    for p in document.paragraphs:
        for run in p.runs:
            if run.text.strip() and run.font.size is not None:
                body.append(run.font.size.pt)
    for t in document.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        if run.text.strip() and run.font.size is not None:
                            table.append(run.font.size.pt)
    return body, table

body_sizes, table_sizes = sizes(doc)
checks = {
    "zip_crc": crc is None,
    "all_xml_parse": True,
    "python_docx_open": True,
    "candidate_hash_matches_build": build.get("docx_sha256") == sha,
    "table_count_31": len(doc.tables) == 31,
    "replacement_register_rows_121": len(register_rows) == 121,
    "register_fact_keys_unique": len(fact_keys) == len(set(fact_keys)),
    "all_register_rows_status_S": len(register_states) == 121 and all(x == "S" for x in register_states),
    "zero_denominator_C1_to_C5_NA_and_0_over_0": zero_denominator_consistent,
    "doc_mapping_0104_I18-E11_checkbox112": "0104" in all_text and "I18-E11" in all_text and "112" in all_text,
    "sample_real_unknown_boundary": "SAMPLE" in all_text and "SYNTHETIC" in all_text and "REAL" in all_text and "UNKNOWN" in all_text,
    "corrected_zero_denominator_T14": "N/A（P1=0）" in all_text and "C不適用" in all_text and "0.00% ✓" not in all_text,
    "T19_C_cohort_reconciles_to_zero": "未申請 | 未完成報考 | P1 內、P2 外 | 否 | 記錄原因、不推測動機 | 8／3／0" in all_text and "失聯 | 多次聯絡未果 | 含於未申請／未到考 | 否 | 留嘗試紀錄與日期 | 3／1／0" in all_text,
    "D0_is_normative_hold_not_execution": "規範性保留中（非執行結果）" in all_text and "F-DS-06" in all_text and "D-1 至 D-5 均為「未執行（示範）」" in all_text,
    "G9_lists_all_five_C_rates": "C期申請率、到考率、通過率、簽發查證完成率、已驗證取證率均觸發" in all_text,
    "R8_definitions_present": "成果率" in all_text and "(B−A)÷A×100%" in all_text and "P4÷P3" in all_text,
    "derived_T12_rates_recomputed_from_exact_fractions": derived_rates_exact,
    "T13_first_page_navigation_fixed": "第 7 節（21 項真實資料缺口）" in all_text and "表 T-13" in all_text and "表 T-16" in all_text,
    "FN1_calculation_locator_fixed": "T-10 P5；§11.1 A-4/A-5" in all_text and "T-10、T-11 A-4/A-5" not in all_text,
    "no_uncited_approximate_CI_claim": "±13～16 個百分點" not in all_text and "比率區間半寬約" not in all_text and "未按分子／分母逐期計算的概略區間" in all_text,
    "aggregate_synthetic_rates_are_complete_and_labeled": all(x in all_text for x in ("43/48=89.58%", "48/76=63.16%", "F-RAGG-06")),
    "zero_denominator_register_F_RC_04_and_05": register_by_fact.get("F-RC-04", {}).get("真實替換來源", "").strip() == "0/0 重算" and register_by_fact.get("F-RC-05", {}).get("真實替換來源", "").strip() == "0/0 重算",
    "no_internal_source_path_or_hash_in_body": not re.search(r"docs/ttqs-quality-20261016/quality-review|SHA-256\s+[0-9a-f]{64}|Blueprint 唯一列|JSONL|標題盲測|crosswalk：", all_text, re.I),
    "body_font_min_10_5": bool(body_sizes) and min(body_sizes) >= 10.5,
    "table_font_min_9_5": bool(table_sizes) and min(table_sizes) >= 9.5,
    "candidate_version_R16": "R16" in all_text,
    "ordered_lists_restart": any(p.startswith("1. P1 = P2") for p in paras) and any(p.startswith("1. 最小化") for p in paras),
    "table_headers_keep_with_next": all(cell.paragraphs[0].paragraph_format.keep_with_next for table in doc.tables for cell in table.rows[0].cells),
    "table_caption_style_keep_with_next": bool([p for p in doc.paragraphs if re.match(r"^表\s*T[-\s]*\d+", p.text, re.I)]) and all((p.paragraph_format.keep_with_next if p.paragraph_format.keep_with_next is not None else p.style.paragraph_format.keep_with_next) for p in doc.paragraphs if re.match(r"^表\s*T[-\s]*\d+", p.text, re.I)),
    "zero_denominator_section_starts_fresh_page": any(p.text.startswith("11.3") and p.paragraph_format.page_break_before for p in doc.paragraphs),
    "hold_gate_table_starts_fresh_page": any(p.text.startswith("表 T 18") and p.paragraph_format.page_break_before for p in doc.paragraphs),
    "exception_table_starts_fresh_page": any(p.text.startswith("表 T 19") and p.paragraph_format.page_break_before for p in doc.paragraphs),
    "orphan_table_sections_page_break_before": any(p.text.startswith("表 T 13") and p.paragraph_format.page_break_before for p in doc.paragraphs) and any(p.text.startswith("19.4") and p.paragraph_format.page_break_before for p in doc.paragraphs),
}
receipt = {
    "receipt_id": "TTQS-0104-TECHNICAL-QA-R16",
    "observed_at_tpe": datetime.datetime.now().astimezone().isoformat(),
    "doc_id": "0104",
    "requirement_id": "I18-E11",
    "official_checkbox_serial": "112",
    "docx_path": str(DOCX),
    "docx_sha256": sha,
    "bytes": len(raw),
    "paragraph_count": len(doc.paragraphs),
    "table_count": len(doc.tables),
    "xml_parts_parsed": len(xml_names),
    "replacement_register_rows": len(register_rows),
    "replacement_register_sha256": hashlib.sha256(REGISTER.read_bytes()).hexdigest(),
    "body_font_min_pt": min(body_sizes) if body_sizes else None,
    "table_font_min_pt": min(table_sizes) if table_sizes else None,
    "checks": checks,
    "failed_checks": [k for k, v in checks.items() if not v],
    "technical_result": "PASS" if all(checks.values()) else "FAIL",
    "visual_review": "PENDING",
    "independent_semantic_review": "PENDING",
    "scope_limit": "Technical integrity and deterministic arithmetic/register checks only; no claim of semantic acceptance or association REAL evidence.",
}
OUT.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(receipt, ensure_ascii=False, indent=2))
