from pathlib import Path
import csv, datetime, hashlib, json
from docx import Document

BASE=Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0104_I18-E11_SOURCE_PREP_20261022")
sha_bytes=lambda b: hashlib.sha256(b).hexdigest()
sha_file=lambda p: sha_bytes(p.read_bytes())
E=BASE/"evidence"
CAND=BASE/"candidate"/"0104__證照獲取率_分析報告_R15.docx"
PDF=BASE/"visual"/"0104_R15"/"0104_R15_word_native.pdf"
r14=json.loads((E/"raster_receipt_0104_R14.json").read_text(encoding="utf-8"))
r15=json.loads((E/"raster_receipt_0104_R15.json").read_text(encoding="utf-8"))
old_hash_to_page={x["sha256"]:x["page"] for x in r14["pages"]}
changed=[x["page"] for x in r15["pages"] if x["sha256"] not in old_hash_to_page]
expected=[1,6,7,*range(10,20),*range(24,28),31,32]
assert r15["page_count"]==32 and len(r15["pages"])==32
assert changed==expected, {"changed":changed,"expected":expected}
tech=json.loads((E/"technical_qa_receipt_0104_R15.json").read_text(encoding="utf-8"))
export=json.loads((E/"word_export_receipt_0104_R15.json").read_text(encoding="utf-8"))
assert tech["technical_result"]=="PASS" and not tech.get("failed_checks")
assert export["status"]=="PASS" and export["new_word_pid_exited"]
assert export["preexisting_word_pids_unchanged"] and not export["unexpected_new_word_pids"]
assert sha_file(CAND)==tech["docx_sha256"]==export["docx_sha256"]
assert sha_file(PDF)==export["pdf_sha256"]==r15["pdf_sha256"]
direct={
1:"Title and opening summary are legible; no clipping or overlap.",
6:"Changed page content is legible; no clipping, overlap, or cut-off rows observed.",
7:"Changed page content is legible; no clipping, overlap, or cut-off rows observed.",
10:"Changed page content is legible; no clipping, overlap, or cut-off rows observed.",
11:"C zero-denominator register begins with its header and complete rows; 0/0 handling is readable.",
12:"T12 rates are readable and separated from the surrounding narrative.",
13:"T13 priority table and navigation are readable.",
14:"T14 and associated decision content are legible without clipping.",
15:"Changed page content is legible; table/text boundaries remain readable.",
16:"Changed page content is legible; table/text boundaries remain readable.",
17:"Changed page content is legible; table/text boundaries remain readable.",
18:"T18 starts with its repeated header and complete G1-G10 rows; readable.",
19:"T19 begins on a fresh page and its rows are readable.",
24:"19.4 register table has a visible header and complete F-RA/F-RB/F-RC rows; cells are readable.",
25:"Continued register table repeats its header; rate-delta rows and formulas are readable.",
26:"Continued register table repeats its header; gate and aggregate rows are readable.",
27:"Continued register table repeats its header; hold, exception and decision keys are readable.",
31:"Evidence locator index and conclusion start are readable; table is not clipped.",
32:"Conclusion continues with numbered items, release gate, limitations and version table; readable."
}
assert set(direct)==set(changed)
page_map={x["page"]:x["sha256"] for x in r15["pages"]}
visual={
 "receipt_id":"TTQS-0104-VISUAL-QA-R15",
 "observed_at_tpe":datetime.datetime.now().astimezone().isoformat(),
 "doc_id":"0104","requirement_id":"I18-E11","official_checkbox_serial":"112",
 "docx_sha256":tech["docx_sha256"],"pdf_sha256":r15["pdf_sha256"],"page_count":32,
 "visual_result":"PASS_WITH_NONBLOCKING_DENSITY_NOTES",
 "review_method":"All 32 pages are covered. The 19 R15 pages whose raster hashes do not occur anywhere in the prior R14 render were opened and directly inspected at full rendered resolution. The remaining 13 pages have byte-identical PNG hashes to a page in the documented R14 visual review; matching is by exact image hash rather than page number because R15 pagination shifted after targeted page-break repairs.",
 "raster_hash_comparison":{"r14_page_count":29,"r15_page_count":32,"changed_or_new_pages":changed,"inherited_identical_page_hash_matches":[{"r15_page":x["page"],"r14_page":old_hash_to_page[x["sha256"]],"sha256":x["sha256"]} for x in r15["pages"] if x["sha256"] in old_hash_to_page]},
 "direct_inspection":[{"page":p,"png_path":str(BASE/"visual"/"0104_R15"/"pages"/f"page-{p}.png"),"sha256":page_map[p],"finding":direct[p]} for p in changed],
 "prior_visual_receipt":{"path":str(E/"visual_qa_receipt_0104_R14.json"),"docx_sha256":json.loads((E/"visual_qa_receipt_0104_R14.json").read_text(encoding="utf-8"))["docx_sha256"],"visual_result":json.loads((E/"visual_qa_receipt_0104_R14.json").read_text(encoding="utf-8"))["visual_result"]},
 "inherited_nonblocking_notes":["Pages 7, 9, 13 and 25 in R13 had more whitespace because short sections or wide tables sit before a section/orientation break; R15 changed pagination and these inherited notes are retained only as historical context, not as current page-number findings."],
 "word_export_receipt_sha256":sha_file(E/"word_export_receipt_0104_R15.json"),
 "technical_qa_receipt_sha256":sha_file(E/"technical_qa_receipt_0104_R15.json"),
 "raster_receipt_sha256":sha_file(E/"raster_receipt_0104_R15.json"),
 "scope_limit":"Visual/layout only; it does not provide independent semantic acceptance or establish REAL association evidence."
}
out=E/"visual_qa_receipt_0104_R15.json"
out.write_text(json.dumps(visual,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
print(json.dumps({"visual_receipt":str(out),"receipt_sha256":sha_file(out),"docx_sha256":tech["docx_sha256"],"pdf_sha256":r15["pdf_sha256"],"changed_pages":changed,"inherited_identical_page_hash_matches":visual["raster_hash_comparison"]["inherited_identical_page_hash_matches"]},ensure_ascii=False,indent=2))

# Build the R15 independent semantic review packet from the exact final bytes and source-bound inputs.
SRC=BASE/"source"
inputs={
 "source_packet":SRC/"SOURCE_PACKET_0104_I18-E11_R01.md",
 "unique_blueprint":SRC/"UNIQUE_BLUEPRINT_0104_I18-E11_R01.jsonl",
 "crosswalk":SRC/"crosswalk_row_0104_I18-E11.csv",
 "official_page7_extract":SRC/"official_pdf_page7_text_extract.txt",
 "synthetic_register":E/"synthetic_replacement_register_0104_R15.csv",
}
doc=Document(str(CAND))
body=[]
for p in doc.paragraphs:
    if p.text.strip(): body.append(p.text)
for ti,t in enumerate(doc.tables,1):
    body.append(f"\n[TABLE {ti}]")
    for row in t.rows:
        body.append(" | ".join(cell.text.replace("\n"," / ") for cell in row.cells))
body_text="\n".join(body)
prior_path=E/"independent_semantic_review_0104_R14"/"opencode_review_resume_receipt_R14.json"
prior=json.loads(prior_path.read_text(encoding="utf-8"))
prior_manifest_path=E/"independent_semantic_review_0104_R14"/"review_input_manifest_0104_R14.json"
prior_manifest=json.loads(prior_manifest_path.read_text(encoding="utf-8"))
review_root=E/"independent_semantic_review_0104_R15"
review_root.mkdir(parents=True,exist_ok=True)
parts=[
 "# Independent semantic review packet — DOCX 0104 / I18-E11 / R15",
 f"Final candidate exact DOCX SHA-256: {sha_file(CAND)}; bytes={CAND.stat().st_size}.",
 "Scope: independent text-only semantic/usability review. Inspect the complete extracted paragraphs and all 31 tables against the official page-7 source extract, SOURCE_PACKET, exact unique Blueprint, crosswalk row and 121-row synthetic replacement register. No page images are attached; do not claim visual inspection. Visual evidence is recorded separately and must not be treated as semantic acceptance.",
 "Identity: project TTQS_ONE; document ID 0104; official requirement I18-E11; official checkbox serial 112. Genre: certificate-acquisition-rate analysis report. This is a SAMPLE/SYNTHETIC candidate; association REAL counts, certificates, approvals and outcomes remain UNKNOWN.",
 "Return PASS / PARTIAL / FAIL. Give exact section/table/row locators and evidence for every conclusion. Check unique requirement identity against the negative neighbors I18-E10 and I18-E12; usability by a volunteer training association and TTQS evaluator; title-blind identification; definitions, populations, denominators and formulas; full arithmetic and cross-document consistency; five zero-denominator rates; register completeness; SAMPLE/REAL boundary; limitations and operational actions; absence of invented real people, certificates, signatures, approvals or third-party originals. Do not edit any file.",
 "R14 predecessor was PARTIAL on exact SHA 37876e73572794cff8ab23d2d45f3924dbae77788896fab2170c45935bc395bb. Its blocking findings were: T14 C rate wrongly shown 0.00% with denominator zero; T19 C counts contradicted C=0; missing C-4 register key; D-0 hold mixed with executed results and no register key; first-page navigation wrong. Nonblocking findings included five missing zero-denominator rates in G9, derived rate register coverage, T14 aggregate ratios, ambiguous relative change vs percentage points, priority label collision, and locators. Verify each correction in R15; do not presume fixed.",
 "Arithmetic dispute: an earlier reviewer alleged step 1 627/1152=54.4358% and +6.94pp. Recompute exactly from source fractions. 627/1152 is approximately 54.427083%, not 54.4358%; assess the report's stated 209/384, rounding, cumulative deltas and shares independently. Do not copy an incorrect arithmetic allegation.",
 "The deterministic technical QA and visual receipts are supporting evidence only. Do not rely on Builder PASS. State whether text-only evidence is enough for semantic PASS; list any unresolved blocker and review limitations. Never infer REAL association performance from synthetic data.",
 "\n## Prior R14 reviewer execution receipt for delta check\n"+json.dumps(prior,ensure_ascii=False,indent=2),
 "\n## Official source packet\n"+inputs["source_packet"].read_text(encoding="utf-8"),
 "\n## Unique Blueprint JSONL\n"+inputs["unique_blueprint"].read_text(encoding="utf-8"),
 "\n## Official crosswalk row\n"+inputs["crosswalk"].read_text(encoding="utf-8"),
 "\n## Official PDF page-7 text extract\n"+inputs["official_page7_extract"].read_text(encoding="utf-8"),
 "\n## Complete final DOCX text: all nonempty paragraphs and all table cells\n"+body_text,
 "\n## Synthetic replacement register CSV\n"+inputs["synthetic_register"].read_text(encoding="utf-8-sig"),
]
packet="\n".join(parts)
packet_path=review_root/"review_packet_0104_R15.md"
packet_path.write_text(packet,encoding="utf-8")
manifest={
 "receipt_id":"TTQS-0104-REVIEW-INPUT-MANIFEST-R15",
 "created_at_tpe":datetime.datetime.now().astimezone().isoformat(),
 "doc_id":"0104","requirement_id":"I18-E11","official_checkbox_serial":"112",
 "candidate_docx_sha256":sha_file(CAND),"candidate_bytes":CAND.stat().st_size,
 "candidate_paragraph_count":len(doc.paragraphs),"candidate_table_count":len(doc.tables),
 "reviewer_scope":"TEXT_ONLY",
 "inputs":{k:{"path":str(p),"sha256":sha_file(p),"bytes":p.stat().st_size} for k,p in inputs.items()},
 "prior_review_receipt_sha256":sha_file(prior_path),
 "prior_candidate_sha256":prior_manifest.get("candidate_docx_sha256"),
 "packet_path":str(packet_path),"packet_sha256":sha_file(packet_path),"packet_bytes":packet_path.stat().st_size,
}
manifest_path=review_root/"review_input_manifest_0104_R15.json"
manifest_path.write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
print(json.dumps({"review_packet":str(packet_path),"packet_sha256":sha_file(packet_path),"packet_bytes":packet_path.stat().st_size,"manifest":str(manifest_path),"manifest_sha256":sha_file(manifest_path),"candidate_sha256":manifest["candidate_docx_sha256"],"input_sha256s":{k:v["sha256"] for k,v in manifest["inputs"].items()}},ensure_ascii=False,indent=2))
