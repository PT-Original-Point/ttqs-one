from pathlib import Path
import csv, json, hashlib, shutil
from docx import Document
from docx.shared import Pt

ROOT = Path('docs/ttqs-quality-20261016/quality-review')
SRC = ROOT/'0046_R01'
OUT = ROOT/'0046_R02'
if (OUT/'candidate'/'0046_R02.docx').exists():
    raise SystemExit('Refusing to overwrite existing 0046_R02 candidate')
(OUT/'candidate').mkdir(parents=True,exist_ok=True)
(OUT/'inputs').mkdir(exist_ok=True)
(OUT/'evidence').mkdir(exist_ok=True)
(OUT/'review').mkdir(exist_ok=True)
DOCX_SRC = SRC/'candidate'/'0046_R01.docx'
DOCX_OUT = OUT/'candidate'/'0046_R02.docx'
R01_SHA = 'fa2410916b009f97dc5b4043ddbacda2a042a99c4a88ba02920e15f34453183a'
R03_SHA = '2ea8937143c5360e03d3db74cf531b8ddffccaecaf4058c89bb3df2ede06d031'
BRANCH_HEAD = '9925f677b84b801509f5d9f85e4f71a66cae0dc7'
def sha_bytes(b): return hashlib.sha256(b).hexdigest()
def sha_file(p): return sha_bytes(Path(p).read_bytes())
if sha_file(DOCX_SRC) != R01_SHA:
    raise SystemExit('R01 exact SHA mismatch; refusing repair')

def set_cell(cell, value):
    paras=cell.paragraphs
    p=paras[0] if paras else cell.add_paragraph()
    # Retain the cell's paragraph/run formatting while replacing its text.
    first=p.runs[0] if p.runs else p.add_run()
    first.text=value
    for run in p.runs[1:]: run.text=''
    for extra in paras[1:]:
        for run in extra.runs: run.text=''

def append_row(table, values):
    row=table.add_row()
    for i,val in enumerate(values): set_cell(row.cells[i],val)
    return row

doc=Document(DOCX_SRC)
# The 16-table order is established from the exact R01 candidate.
tables=doc.tables
if len(tables)!=16: raise SystemExit(f'expected 16 tables, got {len(tables)}')
# Add a resource-type field to the requirements table; this is a populated sample derived from the existing offer rows.
append_row(tables[2],[
    '候選資源類型',
    '外聘師資＋教材組合（本例；SAMPLE / SYNTHETIC）',
    'SAMPLE-LOC/SYN-045-01/04-offer-A-C',
    '需求承辦；確認資源類型及適用評選路徑'
])
# Expand the form's teacher route so internal instructors have source and evaluation fields.
for row in tables[3].rows:
    if row.cells[0].text.strip()=='師資資格與教學':
        set_cell(row.cells[0],'外聘／內部講師資格與教學')
        set_cell(row.cells[1],'外聘：資格、教學方式、兩項能力落差如何練習與回饋。內部：資格／專長、需求適配、教學設計、可投入時段、教材及利益揭露。')
        set_cell(row.cells[2],'外聘：資格與經驗原件；內部：核實資格／專長、教學設計及可投入時段的來源索引。')
        set_cell(row.cells[3],'缺必要資格即不合格；不得以低價抵銷。內部路徑不得虛構外部報價或契約。')
        break
else: raise SystemExit('teacher row not found')
# Make the price rubric explicitly conditional rather than rejecting internal instructors without a quote.
for row in tables[7].rows:
    if row.cells[0].text.strip().startswith('含稅總價'):
        set_cell(row.cells[0],'價格面向（外聘適用）10%')
        set_cell(row.cells[1],'外聘價格須符合同規格可比性及事前級距；超預算／不可比先HOLD或退補，不以低價抵資格。內部講師無對外採購價格時標N/A，不給價格分，依§5.2正規化其餘權重。')
        break
else: raise SystemExit('price rubric row not found')
# Correct the shared synthetic scenario wording to match the current 0045 R03 case rationale.
for row in tables[4].rows:
    if row.cells[0].text.strip()=='方案 C':
        set_cell(row.cells[1],'師資 C 與教材套件；成本最低，紀錄練習充足，但個資情境例及課後修訂說明較少。')
for row in tables[8].rows:
    if row.cells[0].text.strip()=='C' and row.cells[1].text.strip().startswith('需求符合度'):
        set_cell(row.cells[4],'有紀錄練習但個資情境例較少，評分有差異。')
        break
else: raise SystemExit('C demand-score rationale not found')
# Add a compact, fully worked N/A and price-band rule directly after the scoring-anchor table.
heading=doc.add_paragraph('5.2 價格級距與內部講師 N/A 路徑',style='Heading 2')
body1=doc.add_paragraph('外聘候選的 SAMPLE 價格級距（非協會現行標準）：含稅可比總價不高於9,000元＝5分；高於9,000且不高於11,000元＝4分；高於11,000且不高於13,000元＝3分；高於13,000且不高於15,000元＝2分；高於15,000元超出本例上限，HOLD並重新核定，不納入總評。實際案件須以事前核定預算及財務規範替換。')
body2=doc.add_paragraph('內部講師無對外採購價格時，價格面向記 N/A 並寫明原因，不給5分或任何價格分。SAMPLE正規化算例：價格10%排除後，適用原權重合計90%；新權重＝原權重÷90×100，依序為需求30→33.33%、師資25→27.78%、排程15→16.67%、教材10→11.11%、風險10→11.11%，合計100.00%。各新權重及N/A理由須於邀用前通知並由有權者核定。內外部候選只有具可核實的同一全成本口徑才可混合排名；否則分組評估，不互相排序。此算例不代表協會已核定規則。')
body2.paragraph_format.keep_together=True
# Place the new heading and paragraphs after the rubric table, before the existing tie-break paragraph.
rubric=tables[7]._tbl
rubric.addnext(heading._p)
heading._p.addnext(body1._p)
body1._p.addnext(body2._p)
# Add a concrete crosswalk row for handoff/retention and an explicit HOLD example where policy is unknown.
append_row(tables[14],[
    '彙整／交接／保存',
    '§11：案號、最終版本、原件與附件清冊、保存位置、移交／接收職務、日期、簽收索引，以及保存規範名稱／版本／條次。',
    '按協會正式核定紀錄規範及職務權限核對；本示例接收角色、正式位置及保存規則均為 UNKNOWN，採 HOLD；不得自訂年限或銷毀。'
])
# The existing completion paragraph becomes an actionable handoff/retention gate.
for p in doc.paragraphs:
    if p.text.startswith('表單完成條件：'):
        p.text='表單完成條件：需求與同規格已核實；三案資格與總價可比（內部講師按§5.2記N/A與正規化）；利益申報完成；每項原分、理由、計算與門檻可重算；推薦與正式核定分開；契約、交付、缺失、複驗、回饋及附件可追溯；並有保存位置、移交／接收角色、日期、簽收索引及保存規範依據。未核實保存規範或接收權限時標UNKNOWN／HOLD，不宣稱歸檔完成、不結案、不銷毀；缺原件不補造證據。'
        break
else: raise SystemExit('completion paragraph not found')
# Clarify the same synthetic case key, without implying any real Association relationship.
for p in doc.paragraphs:
    if p.text.startswith('0045／I10-E01'):
        p.text='0045／I10-E01《訓練資源採購與甄選管理辦法》規範跨案件流程；本表只記單案。SYN-045-01 是 0045 與本表共用的合成示範情境代碼，僅供核對 SAMPLE/SYNTHETIC 欄位；不代表任何協會真實個案或共同原件。0044／I09-E04 保留會議或電子郵件原始討論與決議，本表不可代造；0047／I10-E02 分析跨年度趨勢與制度改善，本單案表不取代。'
        break
else: raise SystemExit('neighbor-boundary paragraph not found')
# Preserve revision lineage and show exact content changes.
version_table=tables[15]
append_row(version_table,[
    'R02','2026-10-17',
    '依 R01 exact-SHA 審查補入內部講師 N/A/權重正規化、完整 SAMPLE 價格級距、同案 R03 交叉引用及保存交接 HOLD 條件。',
    'SAMPLE 候選；未核定、未簽章'
])
# Save to a new revision; R01 bytes remain untouched.
doc.save(DOCX_OUT)

# Copy current authoritative source materials into the R02 audit folder.
shutil.copy2(SRC/'inputs'/'UNIQUE_BLUEPRINT_0046.jsonl',OUT/'inputs'/'UNIQUE_BLUEPRINT_0046.jsonl')
shutil.copy2(ROOT/'0045_R03'/'inputs'/'official_source_page_03.txt',OUT/'inputs'/'official_source_page_03.txt')
source_map={
 'project_id':'TTQS_ONE','doc_id':'0046','revision':'R02','requirement_id':'I10-E01','official_checkbox_sequence':'047',
 'candidate_sha256':sha_file(DOCX_OUT),'candidate_bytes':DOCX_OUT.stat().st_size,
 'supersedes_candidate_sha256':R01_SHA,'repair_basis':'Exact findings F-1, F-2, F-3, F-4 in independent semantic review R01; no unrelated rewrite.',
 'source_branch':'work/ttqs-one-document-quality-contract-20261016','source_commit_at_repair':BRANCH_HEAD,
 'official_source_pdf_sha256':'805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709',
 'official_source_page_text_sha256':sha_file(OUT/'inputs'/'official_source_page_03.txt'),
 'blueprint_source_sha256':'406d0495365d0970dcf8113a42842663b8a54372a4e204cc492b7c3f8156ef3e',
 'blueprint_row_sha256':'6d2c7798cda52923f9bec34cac347bae43c1ea56883b90d3d412845d80988f05',
 'related_0045_candidate_sha256':R03_SHA,
 'related_0045_review_sha256':sha_file(ROOT/'0045_R03'/'review'/'opencode_v2_review_R03'/'independent_semantic_review_R03.md'),
 'candidate_scope':'SAMPLE / SYNTHETIC only; no Association approval or real procurement claimed.'
}
(OUT/'inputs'/'SOURCE_MAP_R02.json').write_text(json.dumps(source_map,ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')

# Update the replacement register: current 0045 R03 linkage, aligned rationale wording, plus new sample values.
reg_src=SRC/'evidence'/'SYNTHETIC_REPLACEMENT_REGISTER_R01.csv'
with reg_src.open('r',encoding='utf-8-sig',newline='') as f:
    reader=csv.DictReader(f); fields=reader.fieldnames; rows=list(reader)
if len(rows)!=147: raise SystemExit(f'expected 147 R01 register entries, got {len(rows)}')
shared=f'0045 R03 SHA {R03_SHA}; shared SAMPLE/SYNTHETIC scenario label SYN-045-01 only; no real Association case or evidence.'
for row in rows:
    row['shared_sample_reference']=shared
    if '個案情境例及課後修訂說明較少' in row.get('value_synthetic',''):
        row['value_synthetic']=row['value_synthetic'].replace('個案情境例及課後修訂說明較少','個資情境例及課後修訂說明較少')
    if '個案情境例較少' in row.get('value_synthetic',''):
        row['value_synthetic']=row['value_synthetic'].replace('個案情境例較少','個資情境例較少')
extra=[
 {'fact_id':'SYN-046-148','value_synthetic':'外聘師資＋教材組合（本例；SAMPLE / SYNTHETIC）','type':'SAMPLE_SYNTHETIC','document_location':'§3 需求與評選規格／候選資源類型欄','shared_sample_reference':shared,'real_replacement_source_required':'核定課程計畫、需求／資源規格及正式候選原件','responsible_role':'需求承辦；依正式權限覆核','status':'REPLACE_BEFORE_REAL_USE','printed_pages':'2'},
 {'fact_id':'SYN-046-149','value_synthetic':'SAMPLE價格級距：≤9,000=5；>9,000–11,000=4；>11,000–13,000=3；>13,000–15,000=2；>15,000=HOLD／不納總評','type':'SAMPLE_SYNTHETIC','document_location':'§5.2 外聘價格互斥級距','shared_sample_reference':shared,'real_replacement_source_required':'有權核定之協會預算、財務與採購規範，含生效版本及適用日期','responsible_role':'財務／授權核定功能角色，依正式權限映射','status':'REPLACE_BEFORE_REAL_USE','printed_pages':'5'},
 {'fact_id':'SYN-046-150','value_synthetic':'內部講師無外部報價：價格面向N/A，不給價格分；不可比的內外部方案分組、不互相排名','type':'SAMPLE_SYNTHETIC','document_location':'§5.1–§5.2 內部講師N/A路徑','shared_sample_reference':shared,'real_replacement_source_required':'核定評分規則、內部講師資格／專長來源、利益揭露及授權文件','responsible_role':'需求承辦、無利益衝突評選者及授權核定功能角色','status':'REPLACE_BEFORE_REAL_USE','printed_pages':'5'},
 {'fact_id':'SYN-046-151','value_synthetic':'示範權重正規化：適用原權重合計90%；新權重33.33%、27.78%、16.67%、11.11%、11.11%，合計100.00%','type':'SAMPLE_SYNTHETIC','document_location':'§5.2 N/A正規化算例','shared_sample_reference':shared,'real_replacement_source_required':'事前核定並通知之適用面向及權重；重新計算並由覆核者驗算','responsible_role':'評選規則擬定者與授權核定功能角色','status':'REPLACE_BEFORE_REAL_USE','printed_pages':'5'},
 {'fact_id':'SYN-046-152','value_synthetic':'正式接收角色、歸檔位置及保存規則未核實→UNKNOWN/HOLD；本例不宣稱移交完成或允許銷毀','type':'SAMPLE_SYNTHETIC','document_location':'§10追溯表／彙整交接保存列；§11完成條件','shared_sample_reference':shared,'real_replacement_source_required':'協會核定紀錄管理規範、正式職務權限、檔案庫位置與簽收／交接原件','responsible_role':'需求承辦及依正式權限指定的文件管理接收功能角色','status':'REPLACE_BEFORE_REAL_USE','printed_pages':'9'}
]
rows.extend(extra)
reg_out=OUT/'evidence'/'SYNTHETIC_REPLACEMENT_REGISTER_R02.csv'
with reg_out.open('w',encoding='utf-8',newline='') as f:
    writer=csv.DictWriter(f,fieldnames=fields,lineterminator='\n');writer.writeheader();writer.writerows(rows)

# Capture R01 review record as history in the R02 folder without claiming formal acceptance.
shutil.copy2(SRC/'review'/'opencode_v2_review_R01'/'independent_semantic_review_R01.md',OUT/'review'/'R01_independent_semantic_review_history.md')
shutil.copy2(SRC/'review'/'opencode_v2_review_R01'/'RAW_OPENCODE_STDOUT.json',OUT/'review'/'R01_RAW_OPENCODE_STDOUT.json')

receipt={
 'project_id':'TTQS_ONE','doc_id':'0046','requirement_id':'I10-E01','official_checkbox_sequence':'047','revision':'R02',
 'candidate_path':'candidate/0046_R02.docx','candidate_sha256':sha_file(DOCX_OUT),'candidate_bytes':DOCX_OUT.stat().st_size,
 'source_candidate_sha256':R01_SHA,'targeted_repairs':['F-1: add explicit aggregation/handoff/archive fields and UNKNOWN/HOLD gate','F-2: support internal instructor N/A and normalized non-price weights','F-3: update shared synthetic scenario references to current 0045 R03 and align rationale wording','F-4: state full sample price tiers in the form'],
 'modified_source_content':True,'render_qa':'PENDING','technical_qa':'PENDING','independent_exact_sha_semantic_review':'PENDING','candidate_state':'CANDIDATE_SAMPLE_SYNTHETIC_NOT_FORMALLY_ACCEPTED','promotion':'NOT_PROMOTED','approval':'NONE_REPRESENTED'
}
(OUT/'evidence'/'BUILD_RECEIPT_R02.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')
progress={
 'project_id':'TTQS_ONE','document_id':'0046','requirement_id':'I10-E01','official_checkbox_sequence':'047','revision':'R02',
 'state':'CONTENT_REPAIR_BUILT_QA_PENDING','candidate_sha256':receipt['candidate_sha256'],'candidate_bytes':receipt['candidate_bytes'],
 'supersedes':'0046_R01 candidate '+R01_SHA,'static_technical_qa':'PENDING','visual_review':'PENDING','independent_semantic_review':'PENDING_EXACT_SHA','promotion':'NOT_PROMOTED','formal_acceptance':'NOT_GRANTED',
 'sample_boundary':'SAMPLE / SYNTHETIC only; all added scoring rules are illustrative and require authorized replacement; archive/source unknowns remain HOLD.',
 'next_action':'Render and inspect all pages, regenerate exact-SHA technical receipt, dispatch one independent read-only semantic review; then sync stable result to existing PR #92 branch.'
}
(OUT/'PROGRESS.json').write_text(json.dumps(progress,ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')
readme=f'''# 0046 R02｜師資教材廠商評選紀錄表\n\n- Document: 0046; requirement I10-E01; official checkbox 047.\n- Candidate: `candidate/0046_R02.docx`\n- SHA-256: `{receipt['candidate_sha256']}`; {receipt['candidate_bytes']} bytes.\n- Candidate only, SAMPLE/SYNTHETIC; no Association approval, procurement, contract, signature, payment, delivery, or formal TTQS acceptance is represented.\n- R02 is a targeted content repair of exact R01 SHA `{R01_SHA}` addressing independent review findings F-1 through F-4. R01 remains preserved.\n- Technical, full-page visual, and independent exact-SHA review are pending.\n'''
(OUT/'README.md').write_text(readme,encoding='utf-8',newline='\n')
print(json.dumps({'candidate_sha256':receipt['candidate_sha256'],'candidate_bytes':receipt['candidate_bytes'],'register_rows':len(rows),'r01_unchanged_sha256':sha_file(DOCX_SRC),'output_dir':str(OUT)},ensure_ascii=False,indent=2))

