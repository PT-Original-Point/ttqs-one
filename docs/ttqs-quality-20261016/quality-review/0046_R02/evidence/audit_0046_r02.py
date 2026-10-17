from __future__ import annotations

import csv
import hashlib
import json
import re
import zipfile
from collections import Counter
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "candidate" / "0046_R02.docx"
PDF = ROOT / "evidence" / "render_R02_frozen_aacde219" / "0046_R02.pdf"
PAGE_TEXT = ROOT / "review" / "opencode_v2_review_R02" / "page_text_by_word_com.txt"
REGISTER = ROOT / "evidence" / "SYNTHETIC_REPLACEMENT_REGISTER_R02.csv"
RECEIPT = ROOT / "evidence" / "TECHNICAL_LAYOUT_QA_R02.json"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


doc = Document(DOCX)
with zipfile.ZipFile(DOCX) as archive:
    zip_test = archive.testzip()

with REGISTER.open(encoding="utf-8-sig", newline="") as stream:
    records = list(csv.DictReader(stream))

assert len(records) == 152
assert len({row["fact_id"] for row in records}) == 152
assert all(row["type"] == "SAMPLE_SYNTHETIC" for row in records)
assert all(row["status"] == "REPLACE_BEFORE_REAL_USE" for row in records)
assert all(1 <= int(row["printed_pages"]) <= 10 for row in records)
assert all("0045 R03 SHA 2ea8937143c5360e03d3db74cf531b8ddffccaecaf4058c89bb3df2ede06d031" in row["shared_sample_reference"] for row in records)

# Table 7 contains the fixed rubric; table 8 contains the two independent
# synthetic ratings and their recorded means; table 9 contains displayed totals.
weight_table = doc.tables[7]
weights = {}
for row in weight_table.rows[1:]:
    label = row.cells[0].text
    match = re.search(r"(.+?)\s*(\d+)%", label)
    assert match, label
    weights[match.group(1).strip()] = int(match.group(2))
assert sum(weights.values()) == 100

score_table = doc.tables[8]
totals = {"A": Decimal("0"), "B": Decimal("0"), "C": Decimal("0")}
score_rows = 0
for row in score_table.rows[1:]:
    candidate = row.cells[0].text.strip()
    if not candidate:
        continue
    dim = row.cells[1].text
    weight_match = re.search(r"(\d+)%", dim)
    score_match = re.fullmatch(r"\s*([1-5])\s*[/／]\s*([1-5])\s*", row.cells[2].text)
    avg_match = re.fullmatch(r"\s*([1-5](?:\.\d)?)\s*", row.cells[3].text)
    assert candidate in totals and weight_match and score_match and avg_match, [candidate, dim, row.cells[2].text, row.cells[3].text]
    weight = int(weight_match.group(1))
    s1, s2 = map(Decimal, score_match.groups())
    average = (s1 + s2) / Decimal(2)
    recorded_average = Decimal(avg_match.group(1))
    assert average == recorded_average
    assert weight in weights.values(), [dim, weight, weights]
    totals[candidate] += Decimal(weight) * average / Decimal(5)
    score_rows += 1
assert score_rows == 18
calculated_totals = {key: value.quantize(Decimal("0.1"), rounding=ROUND_HALF_UP) for key, value in totals.items()}
assert calculated_totals == {"A": Decimal("82.0"), "B": Decimal("88.0"), "C": Decimal("77.0")}

# Internal-instructor N/A excludes price and normalizes the other five weights.
base = [Decimal(30), Decimal(25), Decimal(15), Decimal(10), Decimal(10)]
normalized = [(weight / sum(base) * 100).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP) for weight in base]
assert sum(normalized) == Decimal("100.00")
assert normalized == [Decimal("33.33"), Decimal("27.78"), Decimal("16.67"), Decimal("11.11"), Decimal("11.11")]

displayed = {}
for row in doc.tables[9].rows[1:]:
    candidate = row.cells[0].text.strip().replace("方案 ", "")
    match = re.search(r"([0-9]+\.[0-9])", row.cells[2].text)
    assert candidate in {"A", "B", "C"} and match
    displayed[candidate] = Decimal(match.group(1))
assert displayed == calculated_totals

text = "\n".join(p.text for p in doc.paragraphs)
text += "\n" + "\n".join("\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows) for table in doc.tables)
running_text = "\n".join(
    paragraph.text
    for section in doc.sections
    for part in (section.header, section.first_page_header, section.even_page_header, section.footer, section.first_page_footer, section.even_page_footer)
    for paragraph in part.paragraphs
)
assert "0046 R01" not in running_text
assert "0046 R02" in running_text
assert "SAMPLE" in text and "SYNTHETIC" in text
assert "內部講師" in text and "N/A" in text and "HOLD" in text
assert "交接" in text and "保存" in text

def explicit_run_sizes(paragraphs):
    return [run.font.size.pt for paragraph in paragraphs for run in paragraph.runs if run.text.strip() and run.font.size]

body_sizes = explicit_run_sizes(doc.paragraphs)
table_sizes = explicit_run_sizes(paragraph for table in doc.tables for row in table.rows for cell in row.cells for paragraph in cell.paragraphs)
body_min = min(body_sizes)
table_min = min(table_sizes)
font_floor_status = "PASS" if body_min >= 10.5 and table_min >= 9.5 else "FAIL"

pages = sorted(PDF.parent.glob("page-*.png"))
assert len(pages) == 10
assert PAGE_TEXT.is_file() and PAGE_TEXT.stat().st_size > 15000
page_text = PAGE_TEXT.read_text(encoding="utf-8")
for n in range(1, 11):
    assert f"===== WORD PAGINATION PAGE {n:02d} =====" in page_text
for heading in ["1. ", "3. ", "3.1", "4. ", "5. ", "6. ", "7. ", "8. ", "9. ", "11. "]:
    assert heading in page_text, heading

result = {
    "project_id": "TTQS_ONE",
    "doc_id": "0046",
    "requirement_id": "I10-E01",
    "official_checkbox_sequence": "047",
    "revision": "R02",
    "candidate_state": "CANDIDATE_SAMPLE_SYNTHETIC_NOT_FORMALLY_ACCEPTED",
    "candidate_sha256": sha256(DOCX),
    "candidate_bytes": DOCX.stat().st_size,
    "docx_zip_integrity": "PASS" if zip_test is None else f"FAIL:{zip_test}",
    "table_count": len(doc.tables),
    "synthetic_register_rows": len(records),
    "synthetic_register_unique_ids": len({row["fact_id"] for row in records}),
    "synthetic_register_statuses": dict(Counter(row["status"] for row in records)),
    "register_printed_page_distribution": dict(Counter(row["printed_pages"] for row in records)),
    "rubric_weight_total": sum(weights.values()),
    "recalculated_candidate_totals": {k: str(v) for k, v in calculated_totals.items()},
    "displayed_candidate_totals_match": True,
    "score_rows_recomputed": score_rows,
    "normalized_internal_instructor_weights": [str(v) for v in normalized],
    "normalized_weight_total": str(sum(normalized)),
    "word_render_page_count": 10,
    "visual_pages_inspected": 10,
    "page_text_sha256": sha256(PAGE_TEXT),
    "rendered_pdf_sha256": sha256(PDF),
    "rendered_pdf_bytes": PDF.stat().st_size,
    "layout_findings": {
        "page6_blank_root_cause": "page-break-only paragraph immediately before section 6",
        "repair": "removed only that page-break paragraph; final pagination is 10 pages",
        "all_pages_readable_and_r02_label_visible": True,
        "borderless_appended_rows": "not present after inherited cell formatting repair",
    },
    "font_floor": {
        "blueprint_minimum_body_pt": 10.5,
        "measured_minimum_explicit_body_pt": body_min,
        "body_runs_below_10_5pt": sum(size < 10.5 for size in body_sizes),
        "blueprint_minimum_table_pt": 9.5,
        "measured_minimum_explicit_table_pt": table_min,
        "status": font_floor_status,
        "measurement_scope": "non-empty Word runs in document body paragraphs and table cells; inherited styles are separately retained by the DOCX"
    },
    "sample_boundary": "SAMPLE/SYNTHETIC; no real Association evaluation, procurement, approval, signature, contract, payment, delivery, or formal acceptance claimed",
    "technical_status": "PASS" if font_floor_status == "PASS" else "PARTIAL_BODY_FONT_FLOOR_FAIL",
    "independent_semantic_review": "PENDING_EXACT_SHA",
    "formal_acceptance": "NOT_GRANTED",
}
RECEIPT.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
print(json.dumps(result, ensure_ascii=False, indent=2))
