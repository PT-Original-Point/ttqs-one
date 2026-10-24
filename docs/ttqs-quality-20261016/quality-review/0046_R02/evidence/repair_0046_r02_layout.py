from pathlib import Path
from copy import deepcopy
import hashlib,json
from docx import Document

p=Path('docs/ttqs-quality-20261016/quality-review/0046_R02/candidate/0046_R02.docx')
if not p.exists(): raise SystemExit('R02 candidate missing')
before=hashlib.sha256(p.read_bytes()).hexdigest()
doc=Document(p)
def inherit_cell_format(table_idx,new_row_idx,source_row_idx):
    t=doc.tables[table_idx]
    target=t.rows[new_row_idx]._tr.tc_lst
    source=t.rows[source_row_idx]._tr.tc_lst
    if len(target)!=len(source): raise SystemExit(f'cell count mismatch table {table_idx}')
    for s,d in zip(source,target):
        if s.tcPr is not None:
            if d.tcPr is not None: d.remove(d.tcPr)
            d.insert(0,deepcopy(s.tcPr))
# Apply visible grid/shading/width formatting to appended rows.
for idx in [2,14,15]:
    t=doc.tables[idx]
    inherit_cell_format(idx,len(t.rows)-1,len(t.rows)-2)
# Put tie-break handling in the recommendation/approval table; remove orphaned paragraph.
tie_text='依事前公告規則比較需求符合度加權分、師資與教學、排程與回應；仍同分即HOLD，由有權者依公告條件書面處理，不事後新增偏好或抽籤。'
found=False
for para in list(doc.paragraphs):
    if para.text.startswith('同分順序：'):
        para._element.getparent().remove(para._element)
        found=True
        break
if not found: raise SystemExit('tie-break paragraph not found')
t=doc.tables[10]
row=t.add_row()
# Inherit the last row structure at cell level, including borders and banding.
inherit_cell_format(10,len(t.rows)-1,len(t.rows)-2)
vals=['同分處理','本例無同分。若同分：'+tie_text,'事前評分規則、逐項總分及正式核定原件']
for i,v in enumerate(vals):
    c=row.cells[i]
    c.text=v
# Reapply cell formatting after replacing text (cell.text keeps tcPr, but reset defensively).
inherit_cell_format(10,len(t.rows)-1,len(t.rows)-2)
# Correct running headers/footers while preserving PAGE fields.
for sec in doc.sections:
    for part in [sec.header,sec.first_page_header,sec.even_page_header,sec.footer,sec.first_page_footer,sec.even_page_footer]:
        for para in part.paragraphs:
            for run in para.runs:
                if '0046 R01' in run.text:
                    run.text=run.text.replace('0046 R01','0046 R02')
        for table in part.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            if '0046 R01' in run.text:
                                run.text=run.text.replace('0046 R01','0046 R02')
doc.save(p)
after=hashlib.sha256(p.read_bytes()).hexdigest()
receipt={'candidate_path':'candidate/0046_R02.docx','candidate_sha256_before_layout_repair':before,'candidate_sha256_after_layout_repair':after,'targeted_repairs':['restored table grid/shading/width formatting for three appended rows','moved tie-break decision rule into recommendation/approval table to remove an orphaned page','updated running header/footer revision label to R02 while preserving page fields'],'docx_bytes_changed':before!=after}
Path('docs/ttqs-quality-20261016/quality-review/0046_R02/evidence/LAYOUT_REPAIR_RECEIPT_R02.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')
print(json.dumps(receipt,ensure_ascii=False,indent=2))
