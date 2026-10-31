from pathlib import Path
import hashlib,json,re,sys,math
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
base=Path(__file__).resolve().parent.parent
src=base/"worker"/"0004_WRITER_MANUSCRIPT_CANDIDATE_R05.md"
outdir=base/"docx"; outdir.mkdir(parents=True,exist_ok=True)
out=outdir/"0004_I01-E02_R05_L05.docx"
receipt=base/"qa"/"WORD_BUILD_RECEIPT_R05_L05.json"; receipt.parent.mkdir(parents=True,exist_ok=True)
expected="a391589d8511372141963b9dbde172889420ed5baa3c6513681e58530a03656e"
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
if sha(src)!=expected: raise SystemExit("MANUSCRIPT_SHA_CHANGED")
text=src.read_text(encoding="utf-8")
doc=Document()
sec=doc.sections[0]; sec.orientation=WD_ORIENT.LANDSCAPE
sec.page_width=Inches(11.69); sec.page_height=Inches(8.27)
sec.top_margin=Inches(.55); sec.bottom_margin=Inches(.55); sec.left_margin=Inches(.48); sec.right_margin=Inches(.48)
sec.header_distance=Inches(.25); sec.footer_distance=Inches(.25)
font="Microsoft JhengHei"; tick=chr(96)
for stylename in ["Normal","Title","Subtitle","Heading 1","Heading 2","Heading 3","Quote","List Bullet","List Number"]:
 st=doc.styles[stylename]; st.font.name=font; st._element.rPr.rFonts.set(qn("w:eastAsia"),font); st.font.color.rgb=RGBColor(35,45,60)
title_ppr=doc.styles["Title"]._element.get_or_add_pPr()
for border in title_ppr.findall(qn("w:pBdr")): title_ppr.remove(border)
doc.styles["Normal"].font.size=Pt(10.5); doc.styles["Normal"].paragraph_format.space_after=Pt(4); doc.styles["Normal"].paragraph_format.line_spacing=1.08
for name,size,color in [("Title",21,(25,55,90)),("Heading 1",15,(31,78,121)),("Heading 2",12.5,(46,94,139)),("Heading 3",11,(62,86,110))]:
 st=doc.styles[name]; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor(*color); st.paragraph_format.space_before=Pt(8); st.paragraph_format.space_after=Pt(4); st.paragraph_format.keep_with_next=True
doc.styles["Quote"].font.size=Pt(10.5); doc.styles["Quote"].font.italic=True; doc.styles["Quote"].paragraph_format.left_indent=Inches(.25); doc.styles["Quote"].paragraph_format.right_indent=Inches(.2)
h=sec.header.paragraphs[0]; h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=h.add_run("TTQS ONE  |  0004 / I01-E02  |  SAMPLE・SYNTHETIC・CANDIDATE")
r.font.name=font; r._element.rPr.rFonts.set(qn("w:eastAsia"),font); r.font.size=Pt(8.5); r.font.color.rgb=RGBColor(100,105,110)
f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=f.add_run("非真實調查紀錄；REAL 執行與正式接受均為 UNKNOWN/HOLD    |    第 ")
r.font.name=font; r._element.rPr.rFonts.set(qn("w:eastAsia"),font); r.font.size=Pt(8.5)
fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),"PAGE"); f._p.append(fld); f.add_run(" 頁")
doc.core_properties.title="0004 I01-E02 目標市場調查執行與樣本紀錄"
doc.core_properties.subject="SAMPLE/SYNTHETIC 候選示範；非真實調查或正式接受"
doc.core_properties.author="TTQS ONE candidate production"
def normalize(s):
 s=s.replace("<br>","\n")
 for n,d,v in [(2,4,"0.500"),(1,2,"0.500"),(1,4,"0.250"),(4,5,"0.800")]:
  s=s.replace(f"$$\\frac{{{n}}}{{{d}}} = {v}$$",f"{n} ÷ {d} = {v}")
 s=re.sub(r"\\text\{([^{}]*)\}",r"\1",s)
 s=s.replace(r"\implies","⇒")
 return s.replace("$$","").replace("$","")
def add_markup(p,s,size=10.5,bold=False):
 s=normalize(s); pattern=re.compile(r"(\*\*.+?\*\*|\n|"+re.escape(tick)+r".+?"+re.escape(tick)+")"); pos=0
 for m in pattern.finditer(s):
  if m.start()>pos:
   rr=p.add_run(s[pos:m.start()]); rr.bold=bold; rr.font.name=font; rr._element.rPr.rFonts.set(qn("w:eastAsia"),font); rr.font.size=Pt(size)
  tok=m.group(0)
  if tok=="\n": p.add_run().add_break()
  elif tok.startswith("**"):
   rr=p.add_run(tok[2:-2]); rr.bold=True; rr.font.name=font; rr._element.rPr.rFonts.set(qn("w:eastAsia"),font); rr.font.size=Pt(size)
  elif tok.startswith(tick):
   rr=p.add_run(tok[1:-1]); rr.font.name="Consolas"; rr._element.rPr.rFonts.set(qn("w:eastAsia"),font); rr.font.size=Pt(size); rr.font.color.rgb=RGBColor(55,75,90)
  pos=m.end()
 if pos<len(s):
  rr=p.add_run(s[pos:]); rr.bold=bold; rr.font.name=font; rr._element.rPr.rFonts.set(qn("w:eastAsia"),font); rr.font.size=Pt(size)
def shade(cell,fill):
 tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),fill); tcPr.append(shd)
def margins(cell):
 pr=cell._tc.get_or_add_tcPr(); mar=OxmlElement("w:tcMar")
 for k,v in [("top",45),("start",55),("bottom",45),("end",55)]:
  el=OxmlElement("w:"+k); el.set(qn("w:w"),str(v)); el.set(qn("w:type"),"dxa"); mar.append(el)
 pr.append(mar)
def repeat_header(row):
 pr=row._tr.get_or_add_trPr(); el=OxmlElement("w:tblHeader"); el.set(qn("w:val"),"true"); pr.append(el)
def no_split(row): row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
def parse_table(lines):
 data=[]
 for line in lines:
  cells=[x.strip() for x in line.strip().strip("|").split("|")]
  if all(re.fullmatch(r":?-{3,}:?",re.sub(r"\s","",x) or "-") for x in cells): continue
  data.append(cells)
 return data
def col_groups(headers,heading):
 n=len(headers)
 if "4-1" in heading and n>=13: return [[0,1,2,3,4,5],[0,6,7,8,9],[0,10,11,12]]
 if ("5-2" in heading or "5.2" in heading) and n>=11: return [[0,1,2,3,4],[0,5,6,7],[0,8,9,10]]
 if "7-1" in heading and n>=7: return [[0,1,2,3],[0,4,5,6]]
 if "10-1" in heading and n>=6: return [[0,1,2],[0,3,4,5]]
 if n>8:
  mid=(n+1)//2; return [list(range(mid)),[0]+list(range(mid,n))]
 return [list(range(n))]
def add_table(data,heading):
 headers,body=data[0],data[1:]; groups=col_groups(headers,heading)
 for gi,inds in enumerate(groups):
  if len(groups)>1:
   p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(2); p.paragraph_format.keep_with_next=True; p.paragraph_format.keep_together=True
   rr=p.add_run(f"{heading}｜欄位組 {gi+1}/{len(groups)}（按識別碼對照；欄位未刪減）"); rr.bold=True; rr.font.name=font; rr._element.rPr.rFonts.set(qn("w:eastAsia"),font); rr.font.size=Pt(10.5)
  t=doc.add_table(rows=1,cols=len(inds)); t.style="Table Grid"; t.autofit=False
  head=t.rows[0]; repeat_header(head); no_split(head)
  for ci,j in enumerate(inds):
   cell=head.cells[ci]; cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP; shade(cell,"DCE6F1"); margins(cell); cell.text=""
   p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.keep_with_next=True; add_markup(p,headers[j],9.5,True)
  for values in body:
   row=t.add_row(); no_split(row)
   for ci,j in enumerate(inds):
    cell=row.cells[ci]; cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP; margins(cell); cell.text=""
    if len(t.rows)%2==0: shade(cell,"F5F7FA")
    p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0; add_markup(p,values[j] if j<len(values) else "",9.5)
  printable=10.73; maxima=[]
  for j in inds:
   mx=max([len(re.sub(r"<br>|\*\*","",str(headers[j])))]+[len(re.sub(r"<br>|\*\*","",str(v[j] if j<len(v) else ""))) for v in body])
   maxima.append(min(max(mx,10),48))
  weights=[math.sqrt(x) for x in maxima]; total=sum(weights) or 1
  for ci,w in enumerate(weights):
   width=Inches(printable*w/total)
   for row in t.rows: row.cells[ci].width=width
  doc.add_paragraph().paragraph_format.space_after=Pt(2)
lines=text.splitlines(); i=0; current_heading=""
while i<len(lines):
 line=lines[i].rstrip()
 if not line.strip(): i+=1; continue
 if line.strip()=="---": i+=1; continue
 if line.startswith("|"):
  buf=[]
  while i<len(lines) and lines[i].startswith("|"): buf.append(lines[i]); i+=1
  data=parse_table(buf)
  if data: add_table(data,current_heading)
  continue
 if line.startswith("#"):
  level=len(line)-len(line.lstrip("#")); content=line[level:].strip(); current_heading=content
  sty="Title" if level==1 else "Heading "+str(min(level,3))
  size=21 if level==1 else 15 if level==2 else 12.5 if level==3 else 11
  p=doc.add_paragraph(style=sty); add_markup(p,content,size); i+=1; continue
 if line.startswith(">"):
  p=doc.add_paragraph(style="Quote"); add_markup(p,line[1:].strip(),10.5); i+=1; continue
 if line.startswith("- "):
  p=doc.add_paragraph(style="List Bullet"); add_markup(p,line[2:].strip(),10.5); i+=1; continue
 if re.match(r"^\d+[.)]\s",line):
  p=doc.add_paragraph(style="List Number"); add_markup(p,re.sub(r"^\d+[.)]\s","",line),10.5); i+=1; continue
 p=doc.add_paragraph(); add_markup(p,line,10.5); i+=1
doc.save(out)
check=Document(out)
receipt_data={"receipt_id":"WORD_BUILD_RECEIPT_0004_R05_L05","doc_id":"0004","requirement_id":"I01-E02","source_manuscript_sha256":sha(src),"docx_path":str(out),"docx_sha256":sha(out),"bytes":out.stat().st_size,"paragraph_count":len(check.paragraphs),"table_count":len(check.tables),"section":"A4 landscape","body_font_pt":10.5,"table_font_pt":9.5,"table_split_policy":"Wide tables 4-1, 5-2, 7-1, 10-1 split into linked field groups with repeated record/fact identifier; no source fields omitted.","status":"BUILT_NATIVE_DOCX_FINAL_LAYOUT_REPAIR_RENDER_PENDING","semantic_review_of_docx":"PENDING_EXACT_SHA","visual_review":"PENDING","formal_acceptance":"UNKNOWN"}
receipt.write_text(json.dumps(receipt_data,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
print(json.dumps(receipt_data,ensure_ascii=False,indent=2))


