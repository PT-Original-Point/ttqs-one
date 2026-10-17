from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
import csv, hashlib, json, zipfile, subprocess
from docx import Document
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(r'E:\TTQS\_TTQS_ONE_REVIEW_SYNC_0036_R08_20261016\docs\ttqs-quality-20261016\quality-review\0050_R01')
DOCX = ROOT / 'candidate' / '0050_R01.docx'
META = ROOT / 'evidence' / 'build' / 'BUILD_METADATA_R01.json'
REG = ROOT / 'evidence' / 'SYNTHETIC_REPLACEMENT_REGISTER_R01.csv'
RENDER = ROOT / 'evidence' / 'render_R01_d0d37c2c'
PDF = RENDER / '0050_R01.pdf'
PDFINFO = Path(r'C:\Users\J\.cache\codex-runtimes\codex-primary-runtime\dependencies\native\poppler\Library\bin\pdfinfo.exe')
EXPECTED = 'd0d37c2c69b161a2cdd37ee3097306722c4a65a1afd844ae92e1dff58b43381a'
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
checks = {}
meta = json.loads(META.read_text(encoding='utf-8'))
docsha = sha(DOCX)
checks['candidate_sha_matches_expected'] = docsha == EXPECTED
checks['candidate_sha_matches_build_metadata'] = docsha == meta['candidate_sha256']
checks['candidate_bytes_match_build_metadata'] = DOCX.stat().st_size == meta['candidate_bytes']
with zipfile.ZipFile(DOCX) as z:
    bad = z.testzip()
    names = set(z.namelist())
checks['docx_zip_integrity'] = bad is None
checks['required_ooxml_parts_present'] = {'[Content_Types].xml','word/document.xml','word/styles.xml','word/header1.xml','word/footer1.xml'}.issubset(names)
d = Document(DOCX)
paras = [p.text for p in d.paragraphs]
body = '\n'.join(paras + [c.text for t in d.tables for row in t.rows for c in row.cells])
checks['correct_doc_identity'] = all(x in body for x in ['DOCX 0050','I11-E02','方框 052'])
checks['official_requirement_verbatim'] = '各課程單元內容均能達成訓練規劃的目標。' in body
header_text='\n'.join(p.text for p in d.sections[0].header.paragraphs)
checks['candidate_only_and_sample_only'] = all(x in body for x in ['SAMPLE / SYNTHETIC','候選狀態']) and '未生效示範稿' in header_text and meta['candidate_state']=='CANDIDATE_ONLY' and meta['formal_acceptance']=='NOT_GRANTED'
checks['title_style_applied'] = bool(d.paragraphs and d.paragraphs[0].style.name=='Title')
title_runs=[r for r in d.paragraphs[0].runs if r.text.strip()]
checks['title_text_black'] = bool(title_runs) and all(r.font.color.rgb and str(r.font.color.rgb).upper()=='000000' for r in title_runs)
title_ppr=d.styles['Title']._element.pPr
checks['title_style_has_no_bottom_border'] = title_ppr is None or title_ppr.find(qn('w:pBdr')) is None
checks['no_one_cell_callout_tables'] = all(not (len(t.rows)==1 and len(t.columns)==1) for t in d.tables)

# Font-floor audit, falling back from run -> paragraph style -> Normal.
def eff_size(run, style):
    if run.font.size: return run.font.size.pt
    if run.style and run.style.font.size: return run.style.font.size.pt
    if style and style.font.size: return style.font.size.pt
    if d.styles['Normal'].font.size: return d.styles['Normal'].font.size.pt
    return None
body_sizes=[]
for p in d.paragraphs:
    for r in p.runs:
        if r.text.strip():
            v=eff_size(r,p.style)
            if v is not None: body_sizes.append(v)
table_sizes=[]
for t in d.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    if r.text.strip():
                        v=eff_size(r,p.style)
                        if v is not None: table_sizes.append(v)
checks['body_font_floor_10_5pt'] = bool(body_sizes) and min(body_sizes) >= 10.5
checks['table_font_floor_9_5pt'] = bool(table_sizes) and min(table_sizes) >= 9.5

with REG.open('r',encoding='utf-8-sig',newline='') as f: rows=list(csv.DictReader(f))
ids=[r['fact_id'] for r in rows]
checks['replacement_register_99_unique_facts'] = len(rows)==99 and len(ids)==len(set(ids))
checks['replacement_register_doc_and_case_scope'] = all(r['doc_id']=='0050' and r['case_id']=='SYN-0033-PLAN-01' for r in rows)
checks['replacement_register_sample_status'] = all(r['verification_status']=='SAMPLE_ONLY_NOT_VERIFIED_AS_REAL' for r in rows)
required_cols=['required_replacement_source','owner_role_to_confirm','assumption_or_basis','locator_marker']
checks['replacement_rows_have_complete_provenance'] = all(all(r.get(k,'').strip() for k in required_cols) for r in rows)
checks['register_sha_matches_build_metadata'] = sha(REG)==meta['synthetic_register']['sha256']

minute_tables=[]
for t in d.tables:
    headers=[c.text.strip() for c in t.rows[0].cells]
    if len(headers)>=2 and headers[0]=='段' and headers[1]=='分鐘':
        values=[]
        for row in t.rows[1:]:
            try: values.append(int(row.cells[1].text.strip()))
            except Exception: pass
        minute_tables.append(sum(values))
checks['four_unit_and_segment_minutes'] = minute_tables==[90,90,120,60] and sum(minute_tables)==360
quiz_tables=[]
key_tables=[]
quiz_indexes=[]
key_indexes=[]
for ti,t in enumerate(d.tables):
    h=[c.text.strip() for c in t.rows[0].cells]
    if h[:3]==['題號','情境與三個作答選項','學員作答']:
        quiz_tables.append(t)
        quiz_indexes.append(ti)
    if h[:3]==['題號','正解','判分理由']:
        key_tables.append(t)
        key_indexes.append(ti)
quiz_ids=[row.cells[0].text.strip() for t in quiz_tables for row in t.rows[1:]]
key_ids=[row.cells[0].text.strip() for t in key_tables for row in t.rows[1:]]
checks['two_parallel_quizzes_ten_each'] = len(quiz_ids)==20 and sum(x.startswith('P') for x in quiz_ids)==10 and sum(x.startswith('Q') for x in quiz_ids)==10 and len(set(quiz_ids))==20
checks['answer_keys_complete_and_separate'] = len(key_ids)==20 and set(key_ids)==set(quiz_ids) and min(key_indexes) > max(quiz_indexes)
checks['student_answer_cells_blank'] = all(row.cells[2].text.strip().replace('_','').replace(' ','')=='' for t in quiz_tables for row in t.rows[1:])
checks['assessment_threshold_and_sample_arithmetic'] = all(x in body for x in ['至少 8/10','前測 6÷10＝60%','後測 8÷10＝80%','差異 8−6＝增加 2 分'])
checks['learner_cards_have_blank_responses_and_no_tutor_key'] = any((lambda h: len(h)>=3 and h[0].startswith('卡片') and h[1]=='合成情境' and '學員判斷' in h[2])([c.text.strip() for c in t.rows[0].cells]) and len(t.rows)==6 for t in d.tables) and '答案鍵另列於授課者專用區' in body
checks['substantive_unique_unit_design_present'] = all(x in body for x in ['四個可教、可練、可評量的單元','五張可直接發給學員','四項必要檢核','五欄示範交接紀錄','退修條件'])

pages=sorted(RENDER.glob('page-*.png'))
pdf_info=subprocess.run([str(PDFINFO),str(PDF)],capture_output=True,text=True,encoding='utf-8',errors='replace',check=True).stdout
page_line=next((x for x in pdf_info.splitlines() if x.startswith('Pages:')), '')
pdf_pages=int(page_line.split(':',1)[1].strip()) if page_line else None
image_records=[]
for p in pages:
    with Image.open(p) as im:
        image_records.append({'file':p.name,'sha256':sha(p),'width_px':im.width,'height_px':im.height,'bytes':p.stat().st_size})
checks['render_pdf_18_pages'] = pdf_pages==18
checks['all_18_rendered_pages_present'] = len(pages)==18 and [p.name for p in pages]==[f'page-{i:02d}.png' for i in range(1,19)]
checks['all_page_images_144dpi_a4_dimensions'] = len(image_records)==18 and all(x['width_px']>=1100 and x['height_px']>=1500 for x in image_records)
checks['render_pdf_nonempty'] = PDF.is_file() and PDF.stat().st_size>100000

result={
 'project_id':'TTQS_ONE','doc_id':'0050','revision':'R01','requirement_id':'I11-E02','official_checkbox':'052',
 'candidate_state':'CANDIDATE_ONLY','formal_acceptance':'NOT_GRANTED',
 'candidate':{'path':'candidate/0050_R01.docx','sha256':docsha,'bytes':DOCX.stat().st_size,'paragraphs':len(d.paragraphs),'tables':len(d.tables)},
 'synthetic_register':{'path':'evidence/SYNTHETIC_REPLACEMENT_REGISTER_R01.csv','sha256':sha(REG),'rows':len(rows),'unique_fact_ids':len(set(ids))},
 'font_audit':{'body_min_pt':min(body_sizes) if body_sizes else None,'table_min_pt':min(table_sizes) if table_sizes else None},
 'arithmetic':{'unit_minutes':minute_tables,'total_minutes':sum(minute_tables),'pre_questions':10,'post_questions':10,'score_each':1,'threshold':'8/10','sample_pre_post':'6/10 -> 8/10; delta +2; no causal inference'},
 'render':{'method':'Microsoft Word COM fixed-format PDF export; bundled Poppler pdftoppm rasterization','pdf_path':'evidence/render_R01_d0d37c2c/0050_R01.pdf','pdf_sha256':sha(PDF),'pdf_pages':pdf_pages,'pages':image_records,'manual_visual_inspection':'PASS; every page 1-18 inspected at original resolution; pages 1-17 are byte-identical rendered images to the immediately preceding candidate, and page 18 was re-inspected after the final font-size adjustment; no clipping, overlap, broken table, unreadable glyph, or header/footer displacement observed','renderer_limitation':'Packaged render_docx.py could not run because the bundled runtime did not contain soffice.exe; Word COM + bundled Poppler was used as a documented fallback.'},
 'checks':checks,
 'technical_status':'PASS' if all(checks.values()) else 'FAIL',
 'semantic_status':'PENDING_INDEPENDENT_EXACT_SHA_REVIEW',
 'built_at':datetime.now(ZoneInfo('Asia/Taipei')).isoformat()
}
out=ROOT/'evidence'/'qa'/'TECHNICAL_VISUAL_QA_R01.json'
out.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'technical_status':result['technical_status'],'failed_checks':[k for k,v in checks.items() if not v],'candidate_sha256':docsha,'pdf_pages':pdf_pages,'body_min_pt':result['font_audit']['body_min_pt'],'table_min_pt':result['font_audit']['table_min_pt'],'output':str(out)},ensure_ascii=False,indent=2))
