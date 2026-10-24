from __future__ import annotations

import csv
import hashlib
import json
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "candidate" / "0050_R01.docx"
REGISTER = ROOT / "evidence" / "SYNTHETIC_REPLACEMENT_REGISTER_R01.csv"
META = ROOT / "evidence" / "build" / "BUILD_METADATA_R01.json"

DOC_ID = "0050"
REVISION = "R01"
CASE_ID = "SYN-0033-PLAN-01"
COURSE = "志工服務紀錄、交接與工作界線實務"
COURSE_PLAN_SHA = "e8212788b1f7b201ca4af4f4c5fb83ab131627dbfa815987f7cf6344edfda727"
COURSE_PLAN_REGISTER_SHA = "c79dd292c3679b83d334f5773694d2d9089dce7c7d532197d287e3c5ad5be501"
SOURCE_SURVEY_SHA = "fe5f3cfd84c761da86435a71e0fd899daa725be78628c4d2d3f2324c1fdb5f4f"
SOURCE_SURVEY_REGISTER_SHA = "d114f36d9c4852cda2239bf22fcc709bdc0a40aad62418e076c6b9c2ae523247"
BLUEPRINT_FILE_SHA = "406d0495365d0970dcf8113a42842663b8a54372a4e204cc492b7c3f8156ef3e"
BLUEPRINT_ROW_SHA = "1204cbe5799e24ded84ad7c7a4dcf2ebb6fb0270f59ed5332f1982b714a1221a"
OFFICIAL_SOURCE_SHA = "805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709"
OFFICIAL_TEXT = "各課程單元內容均能達成訓練規劃的目標。"

NAVY = "17365D"
BLUE = "1F4E78"
PALE = "D9EAF7"
LIGHT = "F3F6F9"
GOLD = "FFF2CC"
RED = "9C0006"
WHITE = "FFFFFF"
BODY_FONT = "Microsoft JhengHei"

facts: list[dict[str, str]] = []
table_counter = 0


def add_fact(
    fact_id: str,
    section: str,
    anchor: str,
    row_key: str,
    field: str,
    value: str,
    kind: str = "SYNTHETIC_LEARNING_DESIGN",
    basis: str = "本件完整示範所需的教學設計假設；不是協會現行規則或實績。",
    source_doc_id: str = "",
    source_sha: str = "",
    source_fact_id: str = "",
    replacement: str = "實際核定訓練計畫、正式課程大綱／教案及授權來源；取得後逐欄替換並由課程設計覆核角色核對。",
    owner: str = "實際授權的課程規劃責任角色；姓名及權限須由正式職掌核實。",
    locator: str = "",
) -> None:
    facts.append(
        {
            "fact_id": fact_id,
            "case_id": CASE_ID,
            "doc_id": DOC_ID,
            "revision": REVISION,
            "section": section,
            "table_anchor": anchor,
            "row_key": row_key,
            "field": field,
            "value": value,
            "provenance_status": kind,
            "assumption_or_basis": basis,
            "source_doc_id": source_doc_id,
            "source_docx_sha256": source_sha,
            "source_fact_id": source_fact_id,
            "required_replacement_source": replacement,
            "owner_role_to_confirm": owner,
            "verification_status": "SAMPLE_ONLY_NOT_VERIFIED_AS_REAL",
            "cross_document_refs": (
                f"{source_doc_id}:{source_fact_id}" if source_doc_id and source_fact_id else ""
            ),
            "locator_marker": locator or f"{section}｜{anchor}｜{row_key}｜{field}",
        }
    )


def cross_fact(
    fact_id: str, source_fact_id: str, section: str, anchor: str, row_key: str, field: str, value: str
) -> None:
    add_fact(
        fact_id,
        section,
        anchor,
        row_key,
        field,
        value,
        kind="SYNTHETIC_CROSS_DOCUMENT_REFERENCE",
        basis="逐字核對同案 DOCX 0049 R01 及其逐事實替換簿；沿用為同一合成課程的計畫來源，不另造第二個課程案例。",
        source_doc_id="0049",
        source_sha=COURSE_PLAN_SHA,
        source_fact_id=source_fact_id,
        replacement="以正式核定且現行的課程計畫、課程目標及版本紀錄替換；同步覆核 0049 與本件。",
        locator=f"{section}｜{anchor}｜{row_key}｜{field}",
    )


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=85, bottom=70, end=85) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_run_font(run, size=10.5, bold=None, color=None, italic=None) -> None:
    run.font.name = BODY_FONT
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), BODY_FONT)
    r_fonts.set(qn("w:ascii"), BODY_FONT)
    r_fonts.set(qn("w:hAnsi"), BODY_FONT)


def style_paragraph(p, before=0, after=4, line=1.08, keep=False) -> None:
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if keep:
        pf.keep_with_next = True


def add_text(doc, text, size=10.5, bold=False, color=None, after=4, align=None, italic=False):
    p = doc.add_paragraph()
    style_paragraph(p, after=after, line=1.08)
    if align is not None:
        p.alignment = align
    r = p.add_run(str(text))
    set_run_font(r, size=size, bold=bold, color=color, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    style_paragraph(p, before=8 if level == 1 else 5, after=4, line=1.0, keep=True)
    r = p.add_run(text)
    set_run_font(r, size=15 if level == 1 else 12, bold=True, color="000000")
    return p


def add_bullets(doc, items, size=10.5):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        style_paragraph(p, after=2, line=1.06)
        r = p.add_run(item)
        set_run_font(r, size=size)


def add_page_break(doc):
    doc.add_page_break()


def add_callout(doc, title, body, fill=None):
    # Keep notices in ordinary prose rather than boxed callouts so they read
    # naturally and do not compete with the working tables and forms.
    p = doc.add_paragraph()
    style_paragraph(p, before=2, after=6, line=1.08)
    r = p.add_run(title + "。")
    set_run_font(r, size=10.5, bold=True, color="000000")
    r = p.add_run(body)
    set_run_font(r, size=10.5, color="202020")
    return p


def add_table(doc, anchor, caption, headers, rows, widths=None, font_size=9.5):
    global table_counter
    table_counter += 1
    bad_rows = [(idx, len(row)) for idx, row in enumerate(rows, start=1) if len(row) != len(headers)]
    if bad_rows:
        raise ValueError(f"{anchor}: expected {len(headers)} columns; bad rows={bad_rows}")
    p = doc.add_paragraph()
    style_paragraph(p, before=4, after=2, line=1.0, keep=True)
    r = p.add_run(caption)
    set_run_font(r, size=10.5, bold=True, color="000000")
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    if widths:
        for i, w in enumerate(widths):
            table.columns[i].width = Inches(w)
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for i, val in enumerate(headers):
        c = header.cells[i]
        if widths:
            c.width = Inches(widths[i])
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(c, BLUE)
        set_cell_margins(c)
        p0 = c.paragraphs[0]
        style_paragraph(p0, after=0, line=1.0)
        rr = p0.add_run(str(val))
        set_run_font(rr, size=font_size, bold=True, color=WHITE)
    for ridx, data in enumerate(rows, start=1):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for i, val in enumerate(data):
            c = cells[i]
            if widths:
                c.width = Inches(widths[i])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(c)
            if ridx % 2 == 0:
                set_cell_shading(c, LIGHT)
            p0 = c.paragraphs[0]
            style_paragraph(p0, after=0, line=1.02)
            rr = p0.add_run(str(val))
            set_run_font(rr, size=font_size, color="202020")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_page_label(doc, title, subtitle=None):
    add_heading(doc, title, 1)
    if subtitle:
        add_text(doc, subtitle, size=10.5, color="404040", after=6)


def configure_document(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(0.58)
    sec.bottom_margin = Inches(0.6)
    sec.left_margin = Inches(0.62)
    sec.right_margin = Inches(0.62)
    sec.header_distance = Inches(0.28)
    sec.footer_distance = Inches(0.3)
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.08
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    title_style = doc.styles["Title"]
    title_style.font.name = BODY_FONT
    title_style.font.color.rgb = RGBColor.from_string("000000")
    title_style.paragraph_format.space_after = Pt(3)
    title_ppr = title_style._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    for stylename in ("List Bullet", "List Number"):
        s = doc.styles[stylename]
        s.font.name = BODY_FONT
        s.font.size = Pt(10.5)
        s._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rr = header.add_run("SAMPLE / SYNTHETIC　｜　DOCX 0050　｜　未生效示範稿")
    set_run_font(rr, size=8.5, color="000000")
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = footer.add_run("非協會實際授課、核准或學員成果　｜　第 ")
    set_run_font(rr, size=8.5, color="000000")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)
    rr = footer.add_run(" 頁")
    set_run_font(rr, size=8.5, color="000000")
    props = doc.core_properties
    props.title = "課程單元教學設計與目標對應計畫"
    props.subject = "TTQS I11-E02 / official checkbox 052 / SAMPLE-SYNTHETIC"
    props.author = "TTQS ONE"
    props.keywords = "課程單元,教學設計,評量,示範資料"


def minute_rows(parts):
    return [[str(i + 1), str(m), title, trainer, learner, output] for i, (m, title, trainer, learner, output) in enumerate(parts)]


def register_segment_facts(unit_key, anchor, parts, unit_title):
    for i, (m, title, _trainer, _learner, _output) in enumerate(parts, 1):
        add_fact(
            f"F0050_{unit_key}_TIME_{i:02d}",
            unit_title,
            anchor,
            f"第{i}段",
            "合成教學分鐘",
            f"{m}分鐘；{title}",
            basis="為與 0049 同案總教學時數一致而設計的示範分段；正式教案可按真實核定時數調整。",
            locator=f"{unit_title}｜{anchor}｜第{i}段｜合成教學分鐘",
        )


def add_source_facts():
    cross_fact("F0050_XREF_CASE", "F0049_CASE_ID", "首頁", "文件識別", "同案", "案例識別碼", CASE_ID)
    cross_fact("F0050_XREF_COURSE", "F0049_COURSE_TITLE", "課程對應", "課程來源", "名稱", "同案課程名稱", COURSE)
    cross_fact("F0050_XREF_COHORT", "F0049_COHORT_SIZE", "課程對應", "課程來源", "示範班級", "示範規劃人數", "12人")
    cross_fact("F0050_XREF_TOTAL_TIME", "F0049_TOTAL_MINUTES", "課程對應", "課程來源", "教學時數", "總教學分鐘", "360分鐘")
    cross_fact("F0050_XREF_G1", "F0049_G1_CRITERION", "單元一", "目標對應", "G1", "合格判準", "五個合成情境至少四題正確，且不作越權承諾。")
    cross_fact("F0050_XREF_G2", "F0049_G2_CRITERION", "單元二", "目標對應", "G2", "合格判準", "四項演練行為至少三項達標；無關鍵隱私違規。")
    cross_fact("F0050_XREF_G3", "F0049_G3_CRITERION", "單元三", "目標對應", "G3", "合格判準", "任務、狀態、未結事項、下一步、責任角色五欄完整，且不含不必要個資。")
    cross_fact("F0050_XREF_G4", "F0049_G4_CRITERION", "單元四", "目標對應", "G4", "合格判準", "整合檢核必要項目全部完成；未達者補練並再評。")
    cross_fact("F0050_XREF_QCOUNT", "F0049_TEST_ITEMS", "評量", "課程測驗", "題數", "情境題數與計分", "10題，每題1分，總分0至10分。")
    cross_fact("F0050_XREF_QPASS", "F0049_TEST_THRESHOLD", "評量", "課程測驗", "門檻", "示範合格門檻", "後測至少8/10；前測僅作教學起點，不作入班資格。")
    cross_fact("F0050_XREF_RUBRIC", "F0049_ROLE_THRESHOLD", "評量", "雙人演練", "門檻", "示範合格門檻", "四項各0或1分；至少3/4，且無關鍵隱私違規。")
    cross_fact("F0050_XREF_U1_TIME", "F0049_SCHEDULE_D1_AM", "單元一", "單元時間", "教學區間", "計畫時數", "示範規劃90分鐘。")
    cross_fact("F0050_XREF_U2_TIME", "F0049_SCHEDULE_D1_PM", "單元二", "單元時間", "教學區間", "計畫時數", "示範規劃90分鐘。")
    cross_fact("F0050_XREF_U3_TIME", "F0049_SCHEDULE_D2_AM", "單元三", "單元時間", "教學區間", "計畫時數", "示範規劃120分鐘。")
    cross_fact("F0050_XREF_U4_TIME", "F0049_SCHEDULE_D2_PM", "單元四", "單元時間", "教學區間", "計畫時數", "示範規劃60分鐘。")
    for i in range(1, 5):
        cross_fact(
            f"F0050_XREF_RUBRIC_{i:02d}",
            f"F0049_RUBRIC_{i:02d}",
            "評量",
            "四項整合規準",
            f"R{i:02d}",
            "合成規準來源",
            ("界線辨識", "開放提問與重述", "下一步確認", "最小必要紀錄")[i - 1],
        )
    cross_fact("F0050_XREF_PREPOST", "F0049_PRE_POST_IN_MODULES", "評量", "課程時數", "前後測", "時間口徑", "前測放入單元一、後測放入單元四，不另加在360教學分鐘之外。")


def build_questions():
    pre = [
        ("P01", "合成情境：服務對象提出一項超出志工已確認職責的要求。第一個合適動作？", "A. 先答應再詢問｜B. 說明目前可確認的範圍、停止越權承諾並按核實路徑交接｜C. 直接留下完整身分資料", "B", "先確認角色界線，再由真實權責路徑處理；不自行擴大服務範圍。"),
        ("P02", "需求敘述含糊，尚不知道對方期待的結果。如何開始確認？", "A. 以開放問題了解目的，再中性重述並請對方確認｜B. 依經驗猜測｜C. 先填成已完成", "A", "能取得必要資訊並區分對方陳述與自己的推測。"),
        ("P03", "合成交接紀錄最需要留下哪組內容？", "A. 所有談話逐字與無關背景｜B. 任務、目前狀態、未結事項、下一步及接手責任角色｜C. 只寫已處理", "B", "使用計畫規定的五個必要欄位，不加不必要個資。"),
        ("P04", "判讀一項設計要求時，正式來源尚未提供。", "A. 以常見做法補成協會規則｜B. 標出未知、指出需要的原件及核對角色｜C. 當作規則不存在", "B", "未知要維持未知，不能用示例推定協會制度。"),
        ("P05", "練習卡意外出現可識別真人的資料。", "A. 留下作為真實感教材｜B. 停止傳閱，按實際資料權限處理並改用虛構案例｜C. 只遮姓名後繼續", "B", "不擴散可識別資訊；示範課只使用虛構個案。"),
        ("P06", "志工與接手角色對服務範圍理解不同。", "A. 以自己記憶決定｜B. 查核正式授權資料；來源未確認前不作承諾並記錄未決項｜C. 取兩人意見平均", "B", "服務界線由真實授權來源確認，不能以折衷猜測代替。"),
        ("P07", "現場口頭交接後尚未確認接手方收到。紀錄應怎麼寫？", "A. 寫已完成｜B. 標示交接尚未確認、責任角色與下一個核實步驟｜C. 刪除未結事項", "B", "交接狀態和下一步須如實區分。"),
        ("P08", "一項訓練目標如何判斷有無學會？", "A. 只看自評把握度｜B. 對照目標，查看可觀察作品或演練並用同一判準評量｜C. 只看出席", "B", "自評與出席不能代替目標對應的可觀察證據。"),
        ("P09", "示範資料中出現日期，但來源沒有核定排課紀錄。", "A. 把示範日期列為正式日期｜B. 保持它是示範值，正式使用時依核定課程原件替換｜C. 刪除來源限制", "B", "範例日期和真實排課事實須分開。"),
        ("P10", "如何判斷一筆交接已達計畫要求？", "A. 有口頭說過即可｜B. 必要欄位完整、未知有標示、責任及下一步明確，且未越權或暴露個資｜C. 文件篇幅足夠", "B", "按預先定義的對應欄位和界線檢查。"),
    ]
    post = [
        ("Q01", "平行情境：對方要求志工代替權責角色作成承諾。", "A. 先承諾以免延誤｜B. 說明可做與需另行確認之處，依正式路徑轉交｜C. 轉貼全部對話給群組", "B", "先守住角色範圍，再依可核實的交接安排。"),
        ("Q02", "對方只說「希望有人處理一下」，應如何取得可用資訊？", "A. 問清希望達成的下一步、重述理解並請對方確認｜B. 直接填結案｜C. 只記「已溝通」", "A", "提問、重述與確認能把模糊需求轉成可處理事項。"),
        ("Q03", "尚未完成的交接紀錄應呈現什麼？", "A. 只留已完成部分｜B. 目前狀態、未結事項、下一步、責任角色及未知來源｜C. 推測一個預計結果", "B", "把事實、未結項和未知分開記錄。"),
        ("Q04", "教材提到一項流程，但查不到合法來源。", "A. 以教材文字當核准證明｜B. 停止把它教成現行規則，取得正式來源後再決定｜C. 讓學員自行投票", "B", "教材本身不能取代正式授權或現行流程來源。"),
        ("Q05", "同儕檢核發現紀錄包含無關的個人背景。", "A. 保留作補充｜B. 依最小必要原則刪除或改用匿名示例，並檢查傳閱範圍｜C. 複製到更多表單", "B", "只保留判讀任務所需的資訊，不擴大使用。"),
        ("Q06", "兩份正式來源對角色權限描述不同。", "A. 選較方便的一份｜B. 暫停相關教學判斷，交由真正授權角色核清版本與效力｜C. 把兩份混成新規則", "B", "版本衝突未釐清前不自行創設規則。"),
        ("Q07", "接手角色尚未回覆收到交接。", "A. 結案｜B. 保留未確認狀態，按核實路徑追蹤並記錄期限依據｜C. 填一個完成時間", "B", "沒有收受證據時不得寫成已完成。"),
        ("Q08", "要回應學員是否達到單元目標，最佳依據是？", "A. 有參加課程｜B. 目標指定的作品或演練，以及事先說明的評分準則｜C. 教師印象", "B", "評量需能檢查指定的可觀察行為。"),
        ("Q09", "合成課表中的時間要轉成正式課程安排時，應先核對什麼？", "A. 直接照抄｜B. 真實核定時數、師資可用時間及場地條件，並保留版本｜C. 只改日期", "B", "示範分鐘數不等同正式排課或資源承諾。"),
        ("Q10", "整合演練結束時，哪些條件才支持設計判定完成？", "A. 已經說完情境｜B. 界線、需求確認、五欄紀錄、接手角色與未知項均依規準完成｜C. 學員自評滿意", "B", "依整合檢核表逐項對照，不用滿意度取代能力判定。"),
    ]
    return pre, post


def add_question_facts(questions, form_name, anchor, key_anchor):
    for qid, stem, options, key, rationale in questions:
        add_fact(
            f"F0050_{form_name}_{qid}",
            "情境測驗",
            anchor,
            qid,
            "合成題目、選項、答案與理由",
            f"{stem}｜{options}｜答案 {key}｜{rationale}",
            basis="本件獨立編寫的教學評量設計，用於範例檢查目標與教學內容是否一致；不是既有協會題庫。",
            replacement="由真實授權課程設計角色依現行服務界線、課程版本及答案依據覆核後，決定是否採用或替換。",
            locator=f"情境測驗｜{anchor}與{key_anchor}｜{qid}｜題目選項、答案與判分理由",
        )


def add_sample_fact(fact_id, section, anchor, row, field, value, basis=None):
    add_fact(
        fact_id,
        section,
        anchor,
        row,
        field,
        value,
        basis=basis or "完整填入的合成評量示例，只示範計分和判讀；不代表任何真人受訓、授課或成效。",
        replacement="只有取得合法、去識別且可核對的實際課程與評量原件後，才可替換為真實紀錄。",
        owner="實際課程承辦角色與授權評量者；姓名及權限須核實。",
    )


def build_document():
    doc = Document()
    configure_document(doc)

    # Cover / 30-second user entry
    p = doc.add_paragraph()
    p.style = doc.styles["Title"]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, before=5, after=3)
    r = p.add_run("課程單元教學設計與目標對應計畫")
    set_run_font(r, size=21, bold=True, color="000000")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, after=7)
    r = p.add_run("DOCX 0050｜I11-E02｜官方查核佐證方框 052")
    set_run_font(r, size=11.5, bold=True, color="000000")
    add_callout(
        doc,
        "完整示範資料｜SAMPLE / SYNTHETIC",
        "本稿為供原始點協會檢視的課程設計示例，不是協會實際課程、志工紀錄、教學、核准、測驗或學員成果。例中所有情境、角色代碼、題目、時間分段與分數皆為合成。",
    )
    add_text(doc, "官方要求原文", size=11, bold=True, color="000000", after=2)
    add_text(doc, f"「{OFFICIAL_TEXT}」", size=13, bold=True, color=BLUE, after=5)
    add_table(
        doc,
        "表1_30秒使用入口",
        "表 1｜30 秒內了解怎麼使用本件",
        ["問題", "本件直接回答"],
        [
            ["文件用途", "把課程總目標拆成四個可教、可練、可評量的單元，逐項核對內容是否支持目標。"],
            ["誰使用／何時", "由課程規劃責任角色、授課者及第二核對者，在計畫送覆核前及設計變更後共同使用。此為建議分工，不代表協會已指派人員。"],
            ["怎麼操作", "先核對正式目標與角色權限；再按四單元教案試教、練習、評分；最後檢查來源、作品、門檻及未決缺口。"],
            ["會留下什麼", "單元目標—教學—練習—評量對照、測驗與演練判準、補練路徑及真實來源替換清單。"],
            ["目前證據狀態", "本例只展示設計完整度；未證明協會實際排課、授課、核准或任何人達標。正式使用須取得真實原件。"],
        ],
        widths=[1.35, 5.65],
    )
    add_heading(doc, "本件的範圍", 2)
    add_text(doc, "同案來源：課程名稱「志工服務紀錄、交接與工作界線實務」，沿用 0049 R01 的同一份合成計畫與案例碼 SYN-0033-PLAN-01。單元一至四、示範規劃人數 12 人與 360 教學分鐘均只作一致性承接，不表示真實班級或排課。", after=2)
    add_text(doc, "使用前的停止條件：若真實課程目的、服務界線、轉介權限或現行資料保護程序尚無可查來源，本件只能用虛構案例教學；不得把示範話術或表單說成協會既有規則。", after=1)
    add_fact("F0050_CASE_ID", "首頁", "文件控制", "同案識別", "案例碼", CASE_ID, basis="與 0033 R03、0049 R01 共用的 SYNTHETIC 案例識別，不是機構班級代碼。", source_doc_id="0049", source_sha=COURSE_PLAN_SHA, source_fact_id="F0049_CASE_ID", locator="首頁｜文件控制｜同案識別｜案例碼")
    add_fact("F0050_COURSE_TITLE", "首頁", "文件控制", "同案課程", "課程名稱", COURSE, kind="SYNTHETIC_CROSS_DOCUMENT_REFERENCE", basis="同案課程名稱直接承接 0049 R01 計畫。", source_doc_id="0049", source_sha=COURSE_PLAN_SHA, source_fact_id="F0049_COURSE_TITLE", locator="首頁｜文件控制｜同案課程｜課程名稱")

    # Scope and sources
    add_page_break(doc)
    add_page_label(doc, "第 1 節｜官方要求、文件分工與來源界線")
    add_text(doc, "本件的唯一命題：在同一課程計畫下，每個單元都要有可觀察行為目標，教學內容與練習直接教出該行為，評量能留下相符證據並按事先明定的判準判讀。只有課名或主題，不能證明單元內容達成訓練規劃目標。")
    add_table(
        doc,
        "表2_相鄰需求辨識",
        "表 2｜相鄰文件的分工與錯件拒收線",
        ["文件與需求", "它應證明的事", "本件如何區分"],
        [
            ["0049｜I11-E01｜方框051", "業主、受訓者與職能需求如何比較並轉成課程計畫。", "0049 提供本例四單元計畫來源；0050 展開各單元內部的可教行為、教學練習與評量判準。"],
            ["0050｜I11-E02｜方框052（本件）", "各單元目標、內容、教師活動、練習證據與合格判準是否逐項一致。", "正文核心是四份單元教案、兩套平行情境題、演練規準及五欄紀錄示範。"],
            ["0051｜I11-E03｜方框053", "課程產出討論、修改、決議與核准留存。", "本件可提出設計草案及覆核問題；不記成實際會議、意見、主管決議或核准。"],
            ["0052｜I11-E04｜方框054", "結訓學員意見如何回到課程修訂。", "本件只設計未來收集與修訂欄位；沒有學員回饋資料或已完成的修訂閉環。"],
        ],
        widths=[1.55, 2.55, 2.9],
        font_size=9.5,
    )
    add_table(
        doc,
        "表3_來源與限制",
        "表 3｜本件每種依據能支持的判斷",
        ["依據", "可用位置", "可支持的內容", "不能推論的事"],
        [
            ["官方《查核佐證文件資料表—訓練機構版》", "第 4 頁；本套件附官方頁面影像。", "確認 I11-E02 方框052的官方要求與位置。", "不能證明原始點協會已開課、授課或達標。"],
            ["唯一 Blueprint 0050", "deliverable_id 0050；I11-E02；官方頁4。", "限定本件為單元目標、教學與評量對照；定義合成例、替換來源與退件條件。", "歷史 0050 內容或舊 PASS 不能取代本版 Blueprint。"],
            ["0049 R01 課程計畫", "I11-E01／方框051；同案課程、G1–G4、四單元及時數。", "提供本件需展開的合成計畫目標和時數來源。", "不證明該課程已核准或存在真實學員。"],
            ["0033 R03 需求調查", "I07-E01／方框033；被 0049 R01 同案引用。", "僅作上游需求鏈的合成背景；本件不重算或複製其問卷正文。", "不得把自評、合成分數或需求排序寫成本會實測能力。"],
        ],
        widths=[1.6, 1.6, 2.05, 1.75],
        font_size=9.5,
    )
    add_callout(doc, "來源衝突怎麼辦", "若現行核定課程、授權範圍、服務流程或評量規則與本示例不同，先停止使用受影響的單元材料，記下衝突原件及版本，交由真實權責角色核實；不得用本示範覆蓋正式來源。", fill=PALE)

    # Goal alignment
    add_page_break(doc)
    add_page_label(doc, "第 2 節｜同案總目標拆成四個可檢查的單元")
    add_text(doc, "本例課程方向承接 0049 R01 的「志工服務紀錄、交接與工作界線實務」。以下 G1–G4 是同案計畫目標的單元化；任何真實課程若目標文字或服務範圍不同，須先替換正式來源，再重新設計題目與評量。")
    goal_rows = [
        ["G1｜角色界線與適切交接", "單元一", "五個合成情境中辨出超出受訓角色範圍的請求，停止自行承諾並選擇按正式路徑交接。", "五題至少四題判斷正確；不得出現越權承諾。", "五情境判斷表與錯誤類型回饋。"],
        ["G2｜需求確認與下一步", "單元二", "以開放提問、重述、界線辨識及下一步確認完成雙人演練。", "四項演練行為至少三項達標，且無關鍵隱私違規。", "觀察規準、具體回饋及再演練紀錄。"],
        ["G3｜最小必要紀錄", "單元三", "使用匿名代碼完成任務、狀態、未結事項、下一步與責任角色五欄。", "五欄完整且不含不必要個資。", "完整示範交接表與同儕核對表。"],
        ["G4｜整合交接", "單元四", "把界線判斷、需求確認、五欄紀錄、接手責任和未知事項連成一次完整模擬交接。", "整合檢核表必要項全部完成；未達者補練並再評。", "整合評量表、匿名結果範例及修訂觸發。"],
    ]
    add_table(
        doc,
        "表4_G1到G4對照",
        "表 4｜訓練計畫目標—單元行為—教學—判準—證據",
        ["同案目標", "單元", "學習者可觀察行為", "合格判準", "留存證據"],
        goal_rows,
        widths=[1.25, 0.7, 2.3, 1.55, 1.2],
        font_size=9.5,
    )
    add_table(
        doc,
        "表5_單元地圖",
        "表 5｜四單元的順序與教學產出",
        ["次序與教學分鐘", "單元焦點", "核心活動", "當堂可檢查的產出"],
        [
            ["一｜90分鐘", "角色界線與交接", "界線示範、五種合成情境分類、確認責任路徑。", "五題判斷表；能區分可做、需確認與不得承諾。"],
            ["二｜90分鐘", "需求確認與下一步", "雙人對話演練、觀察回饋、用同一規準重演。", "四行為觀察表；未確認事項保留為未知。"],
            ["三｜120分鐘", "最小必要紀錄", "示範五欄、匿名案例填寫、同儕核對及修訂。", "五欄完整交接表；不新增不必要身分資訊。"],
            ["四｜60分鐘", "整合演練與後測", "完整案例串接、整合量規、平行後測及補練決定。", "整合檢核、匿名測驗分數與下一步處置。"],
        ],
        widths=[1.2, 1.35, 2.3, 2.15],
    )
    add_text(doc, "時間核對：90＋90＋120＋60＝360 教學分鐘。測驗嵌入單元一與單元四，不另加在總時數外；休息與行政時間不算教學分鐘。這是 0049 同案示範計畫的時數，不是本會排課紀錄。", size=10.5, bold=True, color=NAVY)

    # Unit 1
    add_page_break(doc)
    add_page_label(doc, "第 3 節｜單元一：角色界線與適切交接（90 分鐘）")
    add_text(doc, "單元目標：完成五個虛構服務情境後，能指出目前角色可確認的範圍、停止越權承諾，並說出依正式來源查核或交接的下一步。課前先說明情境為合成，學員不得帶入真實服務對象資料。")
    u1 = [
        (10, "說明界線、未知資訊與演練規則", "說明角色範圍需由正式來源確認，示範如何說明不確定。", "聽取規則並辨認可做／需確認／停止三類。", "前測作為教學起點；不作入班資格。"),
        (15, "教師示範一次界線判斷", "以虛構請求示範停下、說明目前可確認部分、轉交責任角色。", "標記哪些句子是已知、未知或承諾。", "示範標註卡。"),
        (20, "五情境分類", "提供五張不同的合成請求卡，逐張問「有何來源、是否在角色內」。", "個人先判斷，再兩人核對理由。", "五題判斷表。"),
        (20, "交接路徑圖練習", "要求學員從正式來源類型辨認要核對的責任角色；不提供虛構的本會組織圖。", "標示未知責任、需取得的原件及停止點。", "來源缺口與停止點清單。"),
        (15, "獨立五情境檢核", "按一致答案鍵評閱並記錄錯誤類型，不公開個人分數。", "各自完成五題判斷與理由。", "匿名代碼情境判斷表。"),
        (10, "回饋與下一單元銜接", "針對最常見的越權承諾或猜測提供示範句。", "把未確認需求改寫為一個開放問題。", "回饋重點；不宣稱實際班級統計。"),
    ]
    add_table(
        doc,
        "表6_U1流程",
        "表 6｜單元一時間分段與雙方操作",
        ["段", "分鐘", "活動", "帶領者操作", "學習者操作", "當堂證據"],
        minute_rows(u1),
        widths=[0.45, 0.55, 1.35, 1.75, 1.65, 1.25],
        font_size=9.5,
    )
    u1cases = [
        ["S1", "提出超出已知角色範圍的請求。", "說明可確認範圍；不承諾代為執行；找正式責任來源。"],
        ["S2", "請求者希望立即知道處理結果，但資訊未核實。", "區分已知與未知；只承諾查核步驟，不承諾未授權結果。"],
        ["S3", "兩種正式版本對責任角色說法不同。", "停止教成確定規則；記錄版本差異並交由權責者核實。"],
        ["S4", "演練卡包含不必要的個人背景。", "不抄錄、不擴散；改用虛構必要資訊並依資料權限處理。"],
        ["S5", "原先預定的接手角色無法確認。", "保留未交接狀態；依核實的備援路徑處理，不自創收件人。"],
    ]
    add_page_break(doc)
    add_page_label(doc, "學員作業卡｜單元一界線判斷")
    add_text(doc, "請先獨立判斷每張卡：目前可確認什麼、哪些仍未知或需核實、下一步如何處理。答案鍵另列於授課者專用區，不隨學員作業卡發放。", after=5)
    add_table(
        doc,
        "表7_U1學員情境卡",
        "表 7｜五張可直接發給學員的合成界線情境卡",
        ["卡片", "合成情境", "學員判斷：可確認／未知／下一步"],
        [[row[0], row[1], "可確認：________________________\n未知／需核實：__________________\n下一步：________________________"] for row in u1cases],
        widths=[0.65, 2.45, 3.9],
        font_size=9.5,
    )
    add_callout(doc, "單元一判定", "至少 4/5 題理由符合授課者答案鍵，且沒有越權承諾，才符合本示範 G1 判準。未達者針對錯題再練；這不是資格或真人表現紀錄。", fill=PALE)
    for i, part in enumerate(u1, 1):
        add_fact(f"F0050_U1_TIME_{i:02d}", "單元一", "表6_U1流程", f"第{i}段", "合成分鐘與活動", f"{part[0]}分鐘｜{part[1]}", basis="單元一90分鐘內的示範活動分配；活動分鐘總和需等於同案計畫時數。", locator=f"單元一｜表6｜第{i}段")
    for i, row in enumerate(u1cases, 1):
        add_fact(f"F0050_U1_CASE_{i:02d}", "單元一", "表7_U1學員情境卡／表24_U1授課者答案鍵", row[0], "合成情境與預期判斷", "｜".join(row[1:]), basis="本件獨立設計的虛構練習情境；學員題卡與授課者判斷錨點分離；不得被引用為真實服務事件。", locator=f"單元一｜表7學員卡／表24授課者答案鍵｜{row[0]}")

    # Unit 2
    add_page_break(doc)
    add_page_label(doc, "第 4 節｜單元二：需求確認與下一步（90 分鐘）")
    add_text(doc, "單元目標：在雙人虛構對話中，取得必要資訊、以中性語句重述、說明角色界線、確認下一步與責任角色。學員不得以追問為由索取與任務無關的身分或健康資料。")
    u2 = [
        (10, "回顧界線與未知", "用一個新情境提問，示範不把猜測說成事實。", "辨識已知、未知與需核實內容。", "已知／未知標籤。"),
        (15, "開放提問示範", "示範目的、範圍、期望下一步的開放式問題。", "比較封閉猜測和開放提問。", "問題改寫卡。"),
        (15, "重述與確認", "示範先重述對方意思，再詢問是否理解正確。", "兩人輪流練習中性重述。", "確認語句。"),
        (25, "雙人角色演練", "按情境卡計時觀察，不代替正式授權流程。", "一人提出需求、一人確認並說明下一步，之後交換。", "觀察者四項規準。"),
        (15, "依規準回饋", "只回饋可觀察行為；說明哪一項未出現及如何重試。", "用同一案例重演未達部分。", "具體回饋與再演練。"),
        (10, "第二輪核對", "使用同一量尺重新評分，記錄差異但不宣稱實際成效。", "完成第二輪演練。", "匿名規準表。"),
    ]
    add_table(
        doc,
        "表8_U2流程",
        "表 8｜單元二時間分段與練習步驟",
        ["段", "分鐘", "活動", "帶領者操作", "學習者操作", "當堂證據"],
        minute_rows(u2),
        widths=[0.45, 0.55, 1.35, 1.75, 1.65, 1.25],
        font_size=9.5,
    )
    u2rubric = [
        ["界線辨識", "指出超出角色範圍的部分並停止越權承諾。", "未辨識範圍，或仍承諾未授權事項。"],
        ["開放提問與重述", "取得必要資訊，中性重述並確認意思。", "跳過提問／確認，或把推測當成事實。"],
        ["下一步確認", "指出需確認事項、責任角色及查核方式。", "沒有下一步或沒有責任角色。"],
        ["最小必要紀錄", "只記任務、狀態、未結事項、下一步及責任角色。", "必要欄位缺漏，或加入不必要個資。"],
    ]
    add_table(
        doc,
        "表9_U2規準",
        "表 9｜雙人演練四項行為規準（每項 0 或 1 分）",
        ["觀察項目", "1 分：達標", "0 分：未達"],
        u2rubric,
        widths=[1.35, 2.85, 2.8],
        font_size=9.5,
    )
    add_text(doc, "判定：四項至少 3/4，且沒有關鍵隱私違規；觀察者須說出可重現的行為證據。未達者立即示範一次，再演練並以同一四項規準重評。", bold=True, color=NAVY)
    for i, part in enumerate(u2, 1):
        add_fact(f"F0050_U2_TIME_{i:02d}", "單元二", "表8_U2流程", f"第{i}段", "合成分鐘與活動", f"{part[0]}分鐘｜{part[1]}", basis="單元二90分鐘內的示範活動分配；活動分鐘總和需等於同案計畫時數。", locator=f"單元二｜表8｜第{i}段")
    for i, row in enumerate(u2rubric, 1):
        add_fact(f"F0050_U2_RUBRIC_{i:02d}", "單元二", "表9_U2規準", row[0], "合成觀察標準", "｜".join(row[1:]), basis="依同案課程目標設計的教學評量規準；正式採用前須由授權角色核實服務界線。", locator=f"單元二｜表9｜{row[0]}")
    add_sample_fact("F0050_U2_THRESHOLD", "單元二", "表9_U2規準", "通過門檻", "分數與關鍵條件", "至少3/4，且無關鍵隱私違規。")

    # Unit 3
    add_page_break(doc)
    add_page_label(doc, "第 5 節｜單元三：最小必要紀錄（120 分鐘）")
    add_text(doc, "單元目標：讀取同一虛構案例，以代碼完成五欄交接紀錄；能分開記錄事實、未結事項及未知來源，不增加不必要識別資料。課堂範例不含真人姓名、電話、地址、病史或可回推身分的組合資訊。")
    u3 = [
        (10, "用途及資料界線", "說明示範資料與正式原件的區別；展示哪些欄位不應蒐集。", "圈出必要資料與無關資料。", "資料最小化標註。"),
        (20, "教師示範五欄", "以虛構請求示範任務、狀態、未結事項、下一步、責任角色。", "逐欄指出證據或未知來源。", "示範表。"),
        (25, "個人完成交接表", "發放同一匿名案例，不提供多餘個資。", "獨立填寫五個必要欄位。", "個人合成表。"),
        (25, "同儕逐欄核對", "用欄位定義和隱私提醒卡檢查完整度。", "交換作品，標記缺漏與推測。", "同儕核對表。"),
        (25, "修正與重填", "示範如何把推測改成未知及補來源欄位。", "完成修訂版並說明改動理由。", "修訂前後的教學樣例。"),
        (15, "教師總結", "檢查五欄與必要資料界線。", "自查五項並保留問題。", "五欄完整度檢查。"),
    ]
    add_table(
        doc,
        "表10_U3流程",
        "表 10｜單元三分段操作（合計 120 分鐘）",
        ["段", "分鐘", "活動", "帶領者操作", "學習者操作", "當堂證據"],
        minute_rows(u3),
        widths=[0.45, 0.55, 1.35, 1.75, 1.65, 1.25],
        font_size=9.5,
    )
    sample_record = [
        ["任務", "合成代碼 SYN-RQ-01 的服務需求確認與責任交接。"],
        ["目前狀態", "需求已被重述；下一步及權責來源尚未確認。"],
        ["未結事項", "需依正式課程／服務授權來源確認由何角色接手。"],
        ["下一步", "按核實後的路徑轉交；收到確認前維持未完成狀態。"],
        ["責任角色", "合成角色 R-OWNER-01；不是本會現任人員或正式職稱。"],
    ]
    add_table(
        doc,
        "表11_U3完整示範紀錄",
        "表 11｜一筆完整填入的五欄示範交接紀錄",
        ["必要欄位", "SAMPLE / SYNTHETIC 填入示例"],
        sample_record,
        widths=[1.5, 5.5],
    )
    add_callout(doc, "五欄判準", "任務、狀態、未結事項、下一步及責任角色五欄都要完整；未知事項要明確保留為未知，不得補造責任人或處理結果。任何不必要個資都使本次示例判為未達，須清除並重做。", fill=PALE)
    for i, part in enumerate(u3, 1):
        add_fact(f"F0050_U3_TIME_{i:02d}", "單元三", "表10_U3流程", f"第{i}段", "合成分鐘與活動", f"{part[0]}分鐘｜{part[1]}", basis="單元三120分鐘內的示範活動分配；活動分鐘總和需等於同案計畫時數。", locator=f"單元三｜表10｜第{i}段")
    for i, row in enumerate(sample_record, 1):
        add_fact(f"F0050_U3_RECORD_{i:02d}", "單元三", "表11_U3完整示範紀錄", row[0], "合成交接欄位值", row[1], basis="完整填入的虛構教學記錄；代碼不指涉真人或真實服務案件。", locator=f"單元三｜表11｜{row[0]}")
    add_sample_fact("F0050_U3_THRESHOLD", "單元三", "表11_U3完整示範紀錄", "完整度門檻", "欄位門檻", "五欄全部完整，且不含不必要個資。")

    # Unit 4
    add_page_break(doc)
    add_page_label(doc, "第 6 節｜單元四：整合交接演練（60 分鐘）")
    add_text(doc, "單元目標：將界線判斷、需求確認、五欄紀錄及接手責任連成一段完整模擬。情境中的接手路徑只是一項待由正式來源替換的示範設計；未核實的角色不得被教成協會實際窗口。")
    u4 = [
        (5, "說明整合任務與停止條件", "提醒不越權、不使用真人資料、未知須標示。", "確認演練規則及虛構角色。", "演練規則卡。"),
        (10, "閱讀合成請求卡", "提供一段有限資訊的虛構服務請求。", "圈出已知、未知、需確認內容。", "情境分析標記。"),
        (20, "完成雙人演練與交接表", "觀察界線、確認、紀錄與轉交四個環節。", "一人扮演請求者、一人扮演志工；再交換檢查。", "完成的模擬對話與五欄表。"),
        (10, "按整合規準評分", "逐項指出作品或行為證據。", "核對四項必要條件並修正。", "整合評量表。"),
        (10, "平行後測", "發放後測平行卷；與前測同規格，不重複同題。", "獨立作答。", "後測匿名分數示例。"),
        (5, "回饋與補練決定", "依未達項目指出下一次練習內容；不記成已實施課程。", "讀回回饋並說明需要再練的行為。", "補練建議欄位示例。"),
    ]
    add_table(
        doc,
        "表12_U4流程",
        "表 12｜單元四整合演練（合計 60 分鐘）",
        ["段", "分鐘", "活動", "帶領者操作", "學習者操作", "當堂證據"],
        minute_rows(u4),
        widths=[0.45, 0.55, 1.35, 1.75, 1.65, 1.25],
        font_size=9.5,
    )
    integrated = [
        ["角色界線", "指出需另行核實的服務內容，不自行承諾。", "情境判斷及口頭說明。"],
        ["需求確認", "用開放提問／中性重述確認必要資訊。", "演練觀察表。"],
        ["最小必要紀錄", "五欄完整，未知與未結事項分列，無不必要個資。", "完成的示範交接表。"],
        ["交接責任", "指出經正式來源核實的接手責任與下一步；來源未知時明示未知。", "交接檢核與來源欄。"],
    ]
    add_table(
        doc,
        "表13_U4整合檢核",
        "表 13｜整合演練四個必要檢核點",
        ["檢核點", "可觀察通過條件", "作品或行為證據"],
        integrated,
        widths=[1.35, 3.4, 2.25],
        font_size=9.5,
    )
    add_text(doc, "通過方式：四項必要檢核都完成，才符合本示範 G4；若某一項缺漏，指出缺口、安排一次具體補練，再以同一檢核點重評。沒有重評證據時不得寫成通過。", bold=True, color=NAVY)
    for i, part in enumerate(u4, 1):
        add_fact(f"F0050_U4_TIME_{i:02d}", "單元四", "表12_U4流程", f"第{i}段", "合成分鐘與活動", f"{part[0]}分鐘｜{part[1]}", basis="單元四60分鐘內的示範活動分配；活動分鐘總和需等於同案計畫時數。", locator=f"單元四｜表12｜第{i}段")
    for i, row in enumerate(integrated, 1):
        add_fact(f"F0050_U4_GATE_{i:02d}", "單元四", "表13_U4整合檢核", row[0], "合成通過條件", "｜".join(row[1:]), basis="單元四整合目標的示範核對點；未經真人資料不能判讀機構實際成效。", locator=f"單元四｜表13｜{row[0]}")

    # Test forms
    pre, post = build_questions()
    add_question_facts(pre, "PRE", "表14_前測平行版A", "表25_授課者答案鍵A")
    add_question_facts(post, "POST", "表15_後測平行版B", "表26_授課者答案鍵B")
    for title, items, form_key in (
        ("第 7 節｜前測 A 卷：教學起點診斷", pre, "A"),
        ("第 8 節｜後測 B 卷：平行評量", post, "B"),
    ):
        for start in (0, 5):
            add_page_break(doc)
            end = start + 5
            add_page_label(doc, title if start == 0 else title + "（續）")
            if form_key == "A":
                add_text(doc, "每題 1 分，總分 0–10 分。前測結果只用來調整示範，不作個人資格判定。題目來自本件獨立設計；正式使用前須由授權角色依現行流程覆核。", after=5)
                anchor = "表14_前測平行版A"
            else:
                add_text(doc, "每題 1 分，總分 0–10 分。後測與前測考同一目標範圍但使用不同情境，避免直接重複題目；示範門檻為至少 8/10。", after=5)
                anchor = "表15_後測平行版B"
            rows = [[q[0], q[1] + "\n" + q[2], "________________"] for q in items[start:end]]
            add_table(
                doc,
                anchor,
                f"表 {14 if form_key == 'A' else 15}｜{form_key} 卷第 {start+1}–{end} 題",
                ["題號", "情境與三個作答選項", "學員作答"],
                rows,
                widths=[0.65, 5.25, 1.1],
                font_size=9.5,
            )
            if start == 5:
                add_callout(doc, "計分口徑", "每題答對 1 分，不倒扣；分子為答對題數，分母固定為 10 題。前、後測均為 0–10 分；示範合格參考點為後測至少 8/10。正式使用前先核對答案鍵與現行核定流程。", fill=PALE)

    # Trainer-only answer keys are physically separated from learner forms.
    add_page_break(doc)
    add_page_label(doc, "第 9 節｜授課者專用答案鍵（不得與學員題本一併發放）")
    add_callout(doc, "發放控制", "表 24–26 僅供經授權的授課／評量角色備課與判分。列印學員材料時，只提供表 7、表 14、表 15 的作答頁；答案鍵須另行保管。", fill=GOLD)
    add_table(
        doc,
        "表24_U1授課者答案鍵",
        "表 24｜單元一五張情境卡判斷錨點（授課者版）",
        ["卡片", "預期行為與判斷依據"],
        [[row[0], row[2]] for row in u1cases],
        widths=[0.75, 6.25],
        font_size=9.5,
    )
    add_table(
        doc,
        "表25_授課者答案鍵A",
        "表 25｜前測 A 卷答案與逐題判分理由（授課者版）",
        ["題號", "正解", "判分理由"],
        [[q[0], q[3], q[4]] for q in pre],
        widths=[0.75, 0.7, 5.55],
        font_size=9.5,
    )
    add_page_break(doc)
    add_page_label(doc, "第 9 節｜授課者專用答案鍵（續）")
    add_table(
        doc,
        "表26_授課者答案鍵B",
        "表 26｜後測 B 卷答案與逐題判分理由（授課者版）",
        ["題號", "正解", "判分理由"],
        [[q[0], q[3], q[4]] for q in post],
        widths=[0.75, 0.7, 5.55],
        font_size=9.5,
    )

    # Fully worked result
    add_page_break(doc)
    add_page_label(doc, "第 10 節｜完成填入的 SAMPLE 計分與判讀示例")
    add_callout(doc, "這不是學員成績", "下表的匿名代碼 SYN-L01 和所有分數均為合成值，只示範如何計分、判定及安排補練。沒有真實班別、出席、作答、評量者或課程成效可由本例推論。")
    sample_scores = [
        ["前測 A 卷", "6/10", "60%", "教學起點；不作入班或結訓判定。"],
        ["後測 B 卷", "8/10", "80%", "達到示範門檻；只展示判分流程。"],
        ["單元一 G1", "4/5", "4 題正確", "達到四題門檻；示範答案無越權承諾。"],
        ["單元二 G2", "3/4", "3 項達標", "達到三項門檻；示範無關鍵隱私違規。"],
        ["單元三 G3", "5/5", "五欄完整", "達到欄位門檻；示範未記不必要個資。"],
        ["單元四 G4", "4/4", "四項完成", "整合檢核全部完成。"],
    ]
    add_table(doc, "表16_合成計分案例", "表 16｜一筆全填入的合成評量結果", ["評量項目", "合成分數", "換算／結果", "範例判讀"], sample_scores, widths=[1.25, 1.0, 1.25, 3.5])
    add_text(doc, "算式：前測 6÷10＝60%；後測 8÷10＝80%；差異 8−6＝增加 2 分。此差值只示範算術與報告格式；單一合成記錄不能證明課程造成能力提升，也不能替代實際學員作品。", bold=True, color=NAVY)
    add_table(
        doc,
        "表17_範例補練處理",
        "表 17｜若同一示範評量未達門檻，如何留下可追蹤處理",
        ["示範缺口", "當次教學回應", "再評條件", "結果措辭"],
        [
            ["前測低於 8/10", "依錯題主題調整單元示例，不貼個人分數。", "完成教學後用平行後測；至少 8/10 才符合課程示範門檻。", "前測只是起點；不能寫成受訓者不合格。"],
            ["角色演練少於 3/4", "指出未出現的行為，示範一次並安排重演。", "用同一四項規準重新觀察，並確認沒有關鍵隱私違規。", "重評前標記未達／待補；不得預填通過。"],
            ["五欄有缺或含多餘資料", "刪除不必要資訊，回到來源標記及欄位定義重填。", "五欄齊全、未知清楚、不含不必要個資。", "以修訂作品定位，不只寫口頭提醒。"],
            ["整合交接任一必要步驟缺漏", "把完整流程拆成界線、確認、紀錄、責任四步再練。", "四項必要檢核全部完成。", "保留補練內容與再評依據；不造實際成績。"],
        ],
        widths=[1.4, 2.15, 2.15, 1.3],
        font_size=9.5,
    )
    add_sample_fact("F0050_SAMPLE_LEARNER", "合成計分", "表16_合成計分案例", "匿名代碼", "示範代碼", "SYN-L01；非真人、非本會學員識別碼。")
    for i, row in enumerate(sample_scores, 1):
        add_sample_fact(f"F0050_SAMPLE_SCORE_{i:02d}", "合成計分", "表16_合成計分案例", row[0], "合成分數與判讀", "｜".join(row[1:]))
    add_sample_fact("F0050_SAMPLE_DELTA", "合成計分", "表16_合成計分案例", "前後差", "合成算式", "8−6＝增加2分；僅示範算術，不表示因果效果。")

    # Operations
    add_page_break(doc)
    add_page_label(doc, "第 11 節｜帶領者準備、教學環境與低維護操作")
    add_text(doc, "本節是可檢查的設計流程。角色名稱是建議分工，不是協會現行職稱；實際人選、授權、場地與資料保管方式均須以真實來源核實。")
    roles = [
        ["課程規劃責任角色", "確認現行課程目標、權限來源、入班條件及教材版本。", "未取得正式來源時標記未知；不得自行核定。"],
        ["授課帶領角色", "依四單元時間教學，示範、提問、回饋與補練。", "若材料出現真實個資或越權流程，立即停止該段並改用虛構案例。"],
        ["第二核對／觀察角色", "依同一規準記錄可觀察行為，指出作品定位。", "不同判斷先記分歧，不擅自修改答案鍵或正式服務規則。"],
        ["課務支援角色", "提供紙本材料、匿名代碼及無障礙參與安排。", "只收集實際必要欄位；真實名冊按合法權限另行管理。"],
    ]
    add_table(doc, "表18_建議分工", "表 18｜精簡建議角色與控制點", ["建議角色", "教學工作", "停止／覆核條件"], roles, widths=[1.55, 2.8, 2.65], font_size=9.5)
    resources = [
        ["學習材料", "四單元講義、五張界線情境卡、A/B平行情境卷、五欄匿名紀錄表、四項觀察規準。", "範例版次與內容均屬合成；正式版須標實際核定版次。"],
        ["低成本環境", "可供雙人對談與紙本填寫的安靜空間；白板或紙張即可，不要求專用資訊系統。", "場地、座位數、無障礙需求及設備可用性須實際確認。"],
        ["資料保護", "使用虛構個案和代碼；示範作品不得含真人姓名、聯絡方式或可回推身分線索。", "若收到真實個資，依實際資料權限停止散布並通報真實負責角色。"],
        ["時間與材料", "四單元共360教學分鐘；前測嵌入單元一，後測嵌入單元四。", "休息、行政、場地布置不計入教學分鐘；真實日程需核定。"],
    ]
    add_table(doc, "表19_資源與環境", "表 19｜可負擔的示範資源與真實核對點", ["面向", "示範配置", "正式使用前核對"], resources, widths=[1.25, 3.0, 2.75], font_size=9.5)
    gates = [
        ["課前", "正式目標與服務界線有來源；教材版本一致；合成案例不含真人資訊；評量答案鍵及無障礙安排已覆核。", "缺少任何會改變教學邊界的來源，先停用受影響材料。"],
        ["課中", "按分鐘記錄教學進度；發現內容越權、個資外露或學員把示範當真實流程時立即中止澄清。", "由授課者記錄異常位置，課程責任角色決定改版或補教。"],
        ["課後", "只在真實課程完成後才形成實際作品、作答、評量與回饋紀錄。", "本候選不產生真實出席、成績、核准或執行收據。"],
    ]
    add_table(doc, "表20_三階段檢核", "表 20｜課前、課中、課後的可執行檢核", ["時點", "要核對的內容", "停止／輸出"], gates, widths=[0.85, 3.8, 2.35], font_size=9.5)
    for i, row in enumerate(roles, 1):
        add_fact(f"F0050_ROLE_{i:02d}", "操作設計", "表18_建議分工", row[0], "合成建議職責", "｜".join(row[1:]), basis="僅為精簡教學分工設計，不宣稱協會現行有此職位或已指派人員。", locator=f"操作設計｜表18｜{row[0]}")
    add_fact("F0050_ENVIRONMENT", "操作設計", "表19_資源與環境", "低成本場地", "示範環境與材料", "紙本、小組對談與白板可完成；實際場地及無障礙條件尚需正式確認。", basis="低維護成本的設計選項，不是場地存在或資源已備妥的事實。")

    # Exceptions and real replacements
    add_page_break(doc)
    add_page_label(doc, "第 12 節｜退修、停止、補練與真實替換")
    exceptions = [
        ["單元只有主題名稱，沒有可觀察行為", "退回補寫行為動詞、情境條件、產出及判準。", "學習者能實際完成可檢視任務；評量直接檢查該行為。"],
        ["教學內容未教到單元目標", "標記缺少的知識／示範／練習步驟，修訂教案後重走對照。", "教材、活動、練習產出與規準逐項相連。"],
        ["評量只問滿意度或只記出席", "不得作為目標達成證據；補上作品、題目、量尺與答案鍵。", "每一目標至少有一個可觀察且可評的證據。"],
        ["正式服務界線或轉介路徑來源未知", "停用該段流程教學；只練習標記未知、停止承諾及索取正式來源。", "取得有權限的正式來源並核實適用範圍與版本。"],
        ["測驗答案鍵與現行程序衝突", "暫停該題評分，記錄題號、衝突來源及影響範圍。", "由真實授權者核對並批准替換題；重印版本與答案鍵。"],
        ["作品含不必要或可識別個資", "停止傳閱，按真實權限處理；清除多餘資料並改用虛構情境重做。", "第二核對者確認無剩餘識別線索後再評。"],
        ["計時分段總和不等於核定時數", "停止使用該版課表，重算每段及全課分鐘。", "單元分鐘與核定課程總時數一致；測驗不得另加計。"],
    ]
    add_table(doc, "表21_退修條件", "表 21｜具體退修條件、修正動作與重新通過依據", ["退修觸發", "立即處置", "重新通過條件"], exceptions, widths=[1.75, 2.75, 2.5], font_size=9.5)
    add_table(
        doc,
        "表22_REAL替換清單",
        "表 22｜把示範設計升級前要取得的真實來源",
        ["需替換的示範內容", "應取得的真實原件／來源", "核對者與處置"],
        [
            ["課程名稱、總目標、四單元及時數", "經合法權責確認的現行課程計畫、大綱及版本。", "課程承辦角色核對來源、日期與版本；變更時同步重做對照。"],
            ["服務角色界線、轉介與資料流程", "現行書面制度、正式授權或其他可核實的權限依據。", "由真實權責角色確認適用範圍；無來源就停教受影響內容。"],
            ["教學材料、情境題及答案鍵", "核定教材、適用規則及課程設計覆核紀錄。", "授課者與第二核對者逐題核答案；不以示例當正式規則。"],
            ["正式師資、場地、學員與日程", "真實職掌、可用場地、合法班級通知與實際排課資料。", "按實際權限與無障礙需求確認；不得用本例代造名冊或核准。"],
            ["真實作品、測驗分數及學習結論", "去識別學員作品、實際題目版本、評量者原始紀錄及日期。", "逐件核對同意、保存權限與計分；本例分數不得移作真實結果。"],
        ],
        widths=[1.65, 3.0, 2.35],
        font_size=9.5,
    )
    add_callout(doc, "本件的拒收界線", "任何未登錄的合成日期、真人身分、授權、課程實施、測驗結果或主管核准都不得補進正式文件。原件缺少時，保留清楚的未知狀態並標出取得來源與責任角色；不得以空白表或示例替代實績。")
    for i, row in enumerate(exceptions, 1):
        add_fact(f"F0050_EXCEPTION_{i:02d}", "退修控制", "表21_退修條件", row[0], "合成退修規則", "｜".join(row[1:]), basis="本件設計驗收及課堂停止控制；不表示協會已有異常紀錄。", locator=f"退修控制｜表21｜{row[0]}")
    add_fact("F0050_REAL_REPLACEMENT_SCOPE", "真實替換", "表22_REAL替換清單", "五類原件", "正式使用前必要來源", "課程計畫／權限制度／核定教案與答案鍵／師資場地日程／去識別學員作品與實際評量紀錄。", basis="Blueprint 明定的真實資料替換路徑；本件不製作或代替第三方原件。")

    # Traceability and final status
    add_page_break(doc)
    add_page_label(doc, "第 13 節｜官方要求追溯、同案文件連結與版本界線")
    trace = [
        ["官方要求方框052", "第1頁、第2節、表4–表5、單元一至四、表14–表17。", "四個目標各有一個專屬單元、教學步驟、練習證據及門檻。"],
        ["0049 I11-E01／方框051", "表3–表5及四單元分鐘。", "提供同一合成計畫和課程總目標；本件只展開單元層級設計。"],
        ["0033 I07-E01／方框033", "來源鏈表3；0049 R01 同案來源地圖。", "只承接 0049 已整理的需求結果，不複製問卷、樣本或原始答案。"],
        ["0051 I11-E03／方框053", "表2與表22。", "需要真正討論、決議或核准時，應另查合法會議原件；本件不代造。"],
        ["0052 I11-E04／方框054", "表2及表22。", "真正結訓意見與下次修訂需在實際課程後另行留下證據。"],
    ]
    add_table(doc, "表23_官方需求追溯", "表 23｜官方要求—正文位置—可以作的判斷", ["需求", "本件定位", "證據界線"], trace, widths=[1.65, 2.25, 3.1], font_size=9.5)
    add_text(doc, "前後一致性總核對", size=11.5, bold=True, color="000000", after=3)
    add_bullets(
        doc,
        [
            "同一 SYNTHETIC 案例、同一課程名稱、四單元目標及 360 教學分鐘；課程計畫來源為 0049 R01，需求上游為 0033 R03。",
            "單元時數相加：90＋90＋120＋60＝360 分鐘；單元內各分段加總分別等於90、90、120、60分鐘。休息與行政不另計教學時間。",
            "兩套測驗各10題、每題1分；示範後測門檻8/10。前測為教學起點；兩卷不同題但對準相同目標範圍。",
            "G1五題至少四題；G2四項至少三項且無關鍵隱私違規；G3五欄全備；G4四項整合檢核全完成。各門檻與 0049 R01 同案計畫一致。",
            "SAMPLE / SYNTHETIC 作品只展示設計與算分方式，不代表協會服務、學員、師資、教材、排課、授課、結果或核准。",
        ],
        size=10.5,
    )
    add_callout(doc, "候選狀態", "本文件的完整案例、計分及教案皆為合成設計。只有取得正式課程與授權來源並逐欄替換後，才能作為真實課程文件；獨立審查 PASS 也不代表 Human 正式接受或課程已實際執行。", fill=PALE)
    # Design facts summarized per course sequence and assessment configuration
    add_fact("F0050_SAMPLE_CASE_CONTEXT", "課程情境", "課程設計", "案例場景", "合成練習情境", "匿名代碼 SYN-RQ-01 的服務需求須經確認、守界線、以最小必要欄位紀錄並按已核實來源交接。", basis="本件為 I11-E02 練習設計；不指涉真實服務對象或實際協會流程。")
    add_fact("F0050_SAMPLE_COHORT", "課程情境", "課程設計", "示範人數", "同案規劃人數", "12人；承接0049合成計畫，非實際名冊。", kind="SYNTHETIC_CROSS_DOCUMENT_REFERENCE", source_doc_id="0049", source_sha=COURSE_PLAN_SHA, source_fact_id="F0049_COHORT_SIZE", basis="同案課程計畫規劃數；只作教學分組設計。")
    add_fact("F0050_SAMPLE_PAIRS", "課程情境", "課程設計", "雙人組數", "合成分組計算", "12人÷2人/組＝6組；只供雙人演練示例。", basis="由同案示範規劃人數推算的分組數；不是實際報到人數。")
    add_sample_fact("F0050_SAMPLE_PRE_SCORE", "合成計分", "表16_合成計分案例", "前測", "合成分數", "6/10；教學起點，不作入班資格。")
    add_sample_fact("F0050_SAMPLE_POST_SCORE", "合成計分", "表16_合成計分案例", "後測", "合成分數", "8/10；達示範門檻，非實際學員成果。")

    # Additional synthetic design facts
    add_fact("F0050_SAMPLE_RQ_ID", "單元三", "表11_U3完整示範紀錄", "合成請求代碼", "匿名識別", "SYN-RQ-01；不映射真人或真實案件。")
    add_fact("F0050_SAMPLE_ROLE_ID", "單元三", "表11_U3完整示範紀錄", "接手角色代碼", "匿名角色", "R-OWNER-01；合成角色碼，不是協會現任職稱或人員。")
    add_fact("F0050_TIMING_TOTAL", "課程對應", "表5_單元地圖", "四單元合計", "分鐘計算", "90＋90＋120＋60＝360教學分鐘；測驗嵌入單元一與四。", kind="SYNTHETIC_CROSS_DOCUMENT_REFERENCE", source_doc_id="0049", source_sha=COURSE_PLAN_SHA, source_fact_id="F0049_TOTAL_MINUTES", basis="沿用同案課程計畫時數，並重算本件各單元分段。")
    # Summary facts for each unit's instructional content, beyond minute segments
    for key, section, anchor, rows in (
        ("U1", "單元一", "表6_U1流程", u1),
        ("U2", "單元二", "表8_U2流程", u2),
        ("U3", "單元三", "表10_U3流程", u3),
        ("U4", "單元四", "表12_U4流程", u4),
    ):
        add_fact(
            f"F0050_{key}_OBJECTIVE",
            section,
            anchor,
            "單元目標",
            "完整目標與教學機制",
            {
                "U1": "五個合成服務情境中辨識角色界線，停止越權承諾並指出核實／交接步驟。",
                "U2": "在雙人對話中以開放提問、重述確認、界線說明及下一步確認完成需求釐清。",
                "U3": "以匿名代碼完成五欄交接紀錄，保持未知明示且不加入不必要個資。",
                "U4": "在一段完整虛構案例中整合界線、確認、紀錄、責任與平行後測。",
            }[key],
            basis="本件依同案 G1–G4 專屬撰寫的單元教學目標與評量對照。",
        )

    return doc, pre, post


def write_register():
    fields = [
        "fact_id", "case_id", "doc_id", "revision", "section", "table_anchor", "row_key",
        "field", "value", "provenance_status", "assumption_or_basis", "source_doc_id",
        "source_docx_sha256", "source_fact_id", "required_replacement_source",
        "owner_role_to_confirm", "verification_status", "cross_document_refs", "locator_marker",
    ]
    ids = [x["fact_id"] for x in facts]
    if len(ids) != len(set(ids)):
        raise ValueError("duplicate fact_id")
    REGISTER.parent.mkdir(parents=True, exist_ok=True)
    with REGISTER.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(facts)


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc, pre, post = build_document()
    doc.save(OUT)
    write_register()
    # Re-open the exact saved artifact and verify the document's central alignment statements.
    reread = Document(OUT)
    text = "\n".join(p.text for p in reread.paragraphs) + "\n" + "\n".join(
        cell.text for table in reread.tables for row in table.rows for cell in row.cells
    )
    required = [
        OFFICIAL_TEXT, "G1", "G2", "G3", "G4", "360", "SYNTHETIC",
        "前測 A 卷", "後測 B 卷", "0051", "0052",
    ]
    missing = [x for x in required if x not in text]
    if missing:
        raise ValueError(f"required text missing: {missing}")
    # Explicit independent arithmetic assertions.
    unit_minutes = [90, 90, 120, 60]
    segment_minutes = [
        [10, 15, 20, 20, 15, 10],
        [10, 15, 15, 25, 15, 10],
        [10, 20, 25, 25, 25, 15],
        [5, 10, 20, 10, 10, 5],
    ]
    assert [sum(x) for x in segment_minutes] == unit_minutes
    assert sum(unit_minutes) == 360
    assert len(pre) == len(post) == 10
    assert len(facts) == len({x["fact_id"] for x in facts})
    docx_sha = hashlib.sha256(OUT.read_bytes()).hexdigest()
    reg_sha = hashlib.sha256(REGISTER.read_bytes()).hexdigest()
    META.write_text(
        json.dumps(
            {
                "project_id": "TTQS_ONE",
                "doc_id": DOC_ID,
                "revision": REVISION,
                "requirement_id": "I11-E02",
                "official_checkbox_sequence": "052",
                "official_requirement": OFFICIAL_TEXT,
                "candidate_path": "candidate/0050_R01.docx",
                "candidate_sha256": docx_sha,
                "candidate_bytes": OUT.stat().st_size,
                "candidate_state": "CANDIDATE_ONLY",
                "formal_acceptance": "NOT_GRANTED",
                "sample_status": "SAMPLE_SYNTHETIC_ONLY_NOT_ASSOCIATION_RECORD",
                "case_id": CASE_ID,
                "sources": {
                    "unique_blueprint_file_sha256": BLUEPRINT_FILE_SHA,
                    "unique_blueprint_row_sha256": BLUEPRINT_ROW_SHA,
                    "official_evidence_pdf_sha256": OFFICIAL_SOURCE_SHA,
                    "0049_R01_docx_sha256": COURSE_PLAN_SHA,
                    "0049_R01_register_sha256": COURSE_PLAN_REGISTER_SHA,
                    "0033_R03_docx_sha256": SOURCE_SURVEY_SHA,
                    "0033_R03_register_sha256": SOURCE_SURVEY_REGISTER_SHA,
                },
                "synthetic_register": {
                    "path": "evidence/SYNTHETIC_REPLACEMENT_REGISTER_R01.csv",
                    "sha256": reg_sha,
                    "rows": len(facts),
                    "fact_ids_unique": True,
                },
                "arithmetic": {
                    "unit_minutes": unit_minutes,
                    "segment_minutes": [sum(x) for x in segment_minutes],
                    "total_minutes": sum(unit_minutes),
                    "pre_questions": len(pre),
                    "post_questions": len(post),
                    "score_per_question": 1,
                    "post_threshold": ">=8/10",
                    "sample_scores": {"pre": "6/10", "post": "8/10", "difference": "+2"},
                },
                "content_shape": {
                    "paragraphs": len(reread.paragraphs),
                    "tables": len(reread.tables),
                    "body_text_characters": len(text),
                },
                "builder_checks": "PASS; deterministic content/math/register integrity only, not semantic acceptance",
                "built_at": datetime.now().astimezone().isoformat(),
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    print(json.dumps({
        "candidate": str(OUT),
        "candidate_sha256": docx_sha,
        "bytes": OUT.stat().st_size,
        "tables": len(reread.tables),
        "register_rows": len(facts),
        "register_sha256": reg_sha,
        "unit_minutes": unit_minutes,
        "segment_minutes": [sum(x) for x in segment_minutes],
        "total_minutes": sum(unit_minutes),
        "pre_items": len(pre),
        "post_items": len(post),
        "builder_checks": "PASS",
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
