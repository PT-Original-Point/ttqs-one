from pathlib import Path
from datetime import datetime
from docx import Document
import hashlib,json
ROOT=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0100_0101_R08_SOURCE_BLUEPRINT_20261021')
D100=ROOT/'candidate_R09b'/'0100_I18-E08_R09b.docx'
D101=ROOT/'candidate_R09_0101'/'0101_I18-E08_R09.docx'
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
expected100='0262d68826181a4b5f6209c3422a2d5c61a0b9820790c4e43a2e483b94de884c'
expected101='56f433e095f32378f2fd8241c1a27948d2bfe6748da003af5623d265885edb68'
assert sha(D100)==expected100 and sha(D101)==expected101
x=Document(str(D100)); y=Document(str(D101))
def score_rows(table):
 out={}
 for row in table.rows[1:]:
  cells=[c.text.strip() for c in row.cells]
  if cells and cells[0] in {f'A{i:02d}' for i in range(1,9)}:
   out[cells[0]]=[int(v) for v in cells[1:6]+cells[7:12]]
 return out
s100=score_rows(x.tables[12]); s101=score_rows(y.tables[10])
assert list(s100)==[f'A{i:02d}' for i in range(1,9)], list(s100)
assert s100==s101, (s100,s101)
pre=sum(sum(v[:5]) for v in s100.values()); post=sum(sum(v[5:]) for v in s100.values())
assert (pre,post,post-pre)==(75,117,42)
checks={
 'same_8_pair_codes': list(s100)==list(s101),
 'all_80_raw_scores_match': s100==s101,
 'pre_total_75': pre==75,
 'post_total_117': post==117,
 'delta_42': post-pre==42,
 '0100_r09b_candidate_unchanged': sha(D100)==expected100,
 '0101_final9_candidate_unchanged': sha(D101)==expected101,
 'no_real_evidence_promoted': True,
}
tech101=ROOT/'candidate_R09_0101'/'technical_qa_0101_R09_final9.json'
receipt={
 'receipt_id':'TTQS-I18E08-DOC0100-R09B-TO-DOC0101-R09-FINAL9-CROSSDOC-CHECK-R01',
 'generated_local':datetime.now().astimezone().isoformat(),
 'project_id':'TTQS_ONE',
 'doc_a':{'doc_id':'0100','requirement_id':'I18-E08','candidate_path':str(D100),'candidate_sha256':sha(D100),'matrix_table_1_based':13},
 'doc_b':{'doc_id':'0101','requirement_id':'I18-E08','candidate_path':str(D101),'candidate_sha256':sha(D101),'matrix_table_1_based':11,'technical_receipt_path':str(tech101),'technical_receipt_sha256':sha(tech101)},
 'raw_matrix_comparison':s100,
 'calculations':{'pairs':8,'dimensions':5,'phases':2,'raw_cells':80,'pre_total':pre,'post_total':post,'change':post-pre,'all_rows_match':s100==s101},
 'checks':checks,
 'result':'PASS_CROSS_DOCUMENT_RAW_DATA_CONSISTENCY_FOR_EXACT_SHA_PAIR',
 'status_limits':{'both_are_candidates':True,'formal_acceptance':'NOT_GRANTED','real_records':'UNKNOWN/HOLD','does_not_extend_0101_semantic_or_visual_review':True}
}
out=ROOT/'candidate_R09b'/'CROSSDOC_QA_0100_R09B_0101_R09_FINAL9.json'
out.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'receipt':str(out),'receipt_sha256':sha(out),'0100_sha256':sha(D100),'0101_sha256':sha(D101),'checks':checks,'totals':[pre,post,post-pre]},ensure_ascii=False))
