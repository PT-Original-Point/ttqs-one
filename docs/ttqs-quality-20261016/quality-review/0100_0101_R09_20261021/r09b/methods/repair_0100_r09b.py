from pathlib import Path
from copy import deepcopy
from docx import Document
import hashlib, json

src = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0100_0101_R08_SOURCE_BLUEPRINT_20261021\candidate_R09\0100_I18-E08_R09.docx")
out_dir = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0100_0101_R08_SOURCE_BLUEPRINT_20261021\candidate_R09b")
out = out_dir / "0100_I18-E08_R09b.docx"
assert src.is_file() and not out.exists()

old_sha = hashlib.sha256(src.read_bytes()).hexdigest()
assert old_sha == "debd41e7ca1b18b88cad37db118c5b3474df411b3e2d1c87a8351d456245e0b2"
doc = Document(str(src))

def set_paragraph(p, value):
    runs = p.runs
    if runs:
        runs[0].text = value
        for r in runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        p.add_run(value)

def replace_para_exact(old, new):
    matches = [p for p in doc.paragraphs if p.text.strip() == old]
    assert len(matches) == 1, f"Expected one paragraph, got {len(matches)}: {old}"
    set_paragraph(matches[0], new)

def set_cell(cell, value):
    p = cell.paragraphs[0]
    set_paragraph(p, value)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)

# F2: carry the published privacy exclusions into respondent and administrator checks.
replace_para_exact(
    "☐ 正式填答代碼為四位數字，且不含姓名、生日、電話或證件號碼",
    "☐ 代碼為四位數字（0–9）；自由文字例子未含姓名、聯絡方式或任職單位／雇主／客戶／服務對象名稱"
)
replace_para_exact(
    "☐ 未在表上書寫姓名或聯絡方式",
    "☐ 未書寫姓名、聯絡方式、任職單位全名、雇主／客戶／服務對象名稱等可識別資訊"
)
replace_para_exact(
    "5. 檢查是否有學員自行寫上姓名、電話等個資；如有，當場遮蔽或請本人重填於新表。",
    "5. 檢查代碼與自由文字例子是否含姓名、電話、地址、任職單位全名、雇主／客戶／服務對象名稱等可識別資訊；如有，先 HOLD，安全遮蔽或請本人重填，並依第 7 節處置舊表。"
)

# F4: make blank responses one unambiguous missing-data category and normalize wording/schema.
replace_para_exact(
    "代碼, 前A, 前B, 前C, 前D, 前E, 前總分, 後A, 後B, 後C, 後D, 後E, 後總分, 變化",
    "代碼,前A,前B,前C,前D,前E,前總分,後A,後B,後C,後D,後E,後總分,變化"
)
replace_para_exact(
    "2. 開場說明（2 分鐘）：說明這是自評、沒有及格分數、不記名、需自選 4 碼代碼並記住。",
    "2. 開場說明（2 分鐘）：說明這是自評、沒有及格分數、不記名、需自選四位數字配對代碼並記住。"
)
# Table 12 uses 0-based index 11; edit by exact row labels.
qc = doc.tables[11]
assert len(qc.rows) == 11
row_map = {r.cells[0].text.strip(): r for r in qc.rows[1:]}
set_cell(row_map["分數不在 0–4（例如 5、空白、塗改不清）"].cells[0], "分數超出 0–4（例如 5、6、負數）或塗改不清")
set_cell(row_map["分數不在 0–4（例如 5、空白、塗改不清）"].cells[1], "無效")
set_cell(row_map["分數不在 0–4（例如 5、空白、塗改不清）"].cells[2], "請本人重填該題；無法重填則整份 HOLD，不得以 0 代入")
row_map = {r.cells[0].text.strip(): r for r in qc.rows[1:]}
set_cell(row_map["任一維度未勾選"].cells[0], "任一維度未勾選（空白）")

# F1: a matching numeric code alone is not proof that pre/post forms belong to one person.
qc = doc.tables[11]
old_row = next(r for r in qc.rows if r.cells[0].text.strip() == "前後測代碼不一致")
new_tr = deepcopy(old_row._tr)
old_row._tr.addprevious(new_tr)
qc = doc.tables[11]
new_row = next(r for r in qc.rows if r.cells[0].text.strip() == "前後測代碼不一致")
# The inserted clone is immediately before the original mismatch row.
matches = [i for i, r in enumerate(qc.rows) if r.cells[0].text.strip() == "前後測代碼不一致"]
assert len(matches) == 2
inserted = qc.rows[matches[0]]
set_cell(inserted.cells[0], "前後測各一份同碼，未確認為同一人")
set_cell(inserted.cells[1], "身分配對未證實")
set_cell(inserted.cells[2], "先 HOLD；由本人確認為同一人並核對同版本、日期順序，留下核對註記後才可配對；未能確認或發現他人同碼碰撞時兩表 REJECT，排除有效分母。不得只因代碼相同就推定同一人。")

# F3: identify the new revision and report the prior review without implying the new SHA is reviewed.
meta = doc.tables[0]
set_cell(meta.rows[6].cells[1], "R09b（由 R09a 定點修補）")
set_cell(meta.rows[8].cells[1], "無真實學員填答；R09a 舊 SHA 語意審查 PASS。本 R09b 為局部修補候選，待本版 delta 複核；正式驗收 NOT_GRANTED。")
# Keep the unchanged supporting register explicitly identified as the R09 data baseline.
replacement = doc.tables[16]
set_cell(replacement.rows[10].cells[1], "synthetic_replacement_register_0100_R09.csv（沿用未改動的 R09 清冊）")

# Update all visible revision references but preserve historical names inside the explanatory history paragraph.
for p in doc.paragraphs:
    if "本稿版次 R09a" in p.text:
        set_paragraph(p, p.text.replace("本稿版次 R09a", "本稿版次 R09b"))

doc.core_properties.title = "0100 I18-E08 學員訓練前後自我評量表 R09b"
doc.core_properties.subject = "I18-E08 SAMPLE/SYNTHETIC 候選工具；R09a 審查後定點修補"
doc.save(str(out))
new_sha = hashlib.sha256(out.read_bytes()).hexdigest()
receipt = {
  "project_id":"TTQS_ONE","doc_id":"0100","official_requirement_id":"I18-E08","project_157_index":109,
  "revision":"R09b","state":"CANDIDATE_REPAIR_CREATED_QA_PENDING",
  "source_candidate":"0100_I18-E08_R09.docx","source_candidate_sha256":old_sha,
  "candidate_file":out.name,"candidate_sha256":new_sha,"candidate_bytes":out.stat().st_size,
  "source_packet_sha256":"81caa978a17dda522850628e6e9b4b16453145a986d2bafd2598b6f5ac96b886",
  "unique_blueprint_sha256":"c407af272823d0159bbbe0c1c9258ed74ee89e5328bd8338840cef26183b2817",
  "repairs":[
    "F1: same-code cross-stage collision now HOLD until respondent confirms same person, version, and valid date order; otherwise REJECT and exclude denominator",
    "F2: respondent and administrator privacy checks now include workplace/employer/client/service-recipient identifiers",
    "F3: revision/status labels now distinguish prior R09a PASS from pending R09b delta review and NOT_GRANTED formal acceptance",
    "F4: four-digit wording and CSV field separators normalized; blank score classified only as missing"
  ],
  "data_preservation":{"sample_pairs":8,"score_cells":80,"source_matrix":"unchanged","real_evidence":"UNKNOWN/HOLD","formal_acceptance":"NOT_GRANTED"},
  "created_by":"Codex deterministic targeted python-docx edit; no prose regeneration; no external evidence fabricated"
}
(out_dir/"PATCH_RECEIPT_0100_R09B.json").write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
print(json.dumps(receipt,ensure_ascii=False))


