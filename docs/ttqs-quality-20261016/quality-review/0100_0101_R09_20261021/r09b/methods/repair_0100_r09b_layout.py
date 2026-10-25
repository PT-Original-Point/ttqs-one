from pathlib import Path
from docx import Document
import hashlib,json
p=Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0100_0101_R08_SOURCE_BLUEPRINT_20261021\candidate_R09b\0100_I18-E08_R09b.docx")
rp=p.parent/"PATCH_RECEIPT_0100_R09B.json"
old=hashlib.sha256(p.read_bytes()).hexdigest()
assert old=="145dd8a0f207a4f44d04cdcba75626a880a78e05572723eb0805db37e87a76b4"
d=Document(str(p))
def setp(par,text):
 runs=par.runs
 if runs:
  runs[0].text=text
  for r in runs[1:]: r._element.getparent().remove(r._element)
 else: par.add_run(text)
cover=[x for x in d.paragraphs if x.text.startswith("版次 R09a |") or x.text.startswith("版次 R09a｜")]
assert len(cover)==1,[(x.text) for x in d.paragraphs if "版次 R09a" in x.text]
setp(cover[0],cover[0].text.replace("版次 R09a","版次 R09b",1))
check_text="☐ 配對代碼為四位數字（0–9）；自由文字例子不含姓名、聯絡方式或任職單位／雇主／客戶／服務對象名稱；後測同碼同版。"
p63=[x for x in d.paragraphs if x.text.strip().startswith("☐ 代碼為四位數字")]
assert len(p63)==1
setp(p63[0],check_text)
old68="☐ 我記得後測要使用同一組配對代碼與同一版本"
p68=[x for x in d.paragraphs if x.text.strip()==old68]
assert len(p68)==1
p68[0]._element.getparent().remove(p68[0]._element)
tmp=p.with_name("0100_I18-E08_R09b.next.docx")
d.save(str(tmp)); tmp.replace(p)
new=hashlib.sha256(p.read_bytes()).hexdigest()
receipt=json.loads(rp.read_text(encoding="utf-8"))
receipt["candidate_sha256"]=new
receipt["candidate_bytes"]=p.stat().st_size
receipt["state"]="CANDIDATE_REPAIR_CREATED_LAYOUT_QA_PENDING"
receipt["layout_iteration"]={"from_sha256":old,"to_sha256":new,"changes":["cover revision label corrected from R09a to R09b","merged the pre-test final code/version reminder into the privacy checklist to remove a one-item orphan page"]}
rp.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
print(json.dumps({"previous_sha256":old,"candidate_sha256":new,"bytes":p.stat().st_size,"receipt_sha256":hashlib.sha256(rp.read_bytes()).hexdigest()},ensure_ascii=False))
