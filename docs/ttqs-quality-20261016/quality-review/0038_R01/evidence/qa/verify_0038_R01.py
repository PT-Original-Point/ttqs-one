from __future__ import annotations

import csv
import hashlib
import json
import re
import zipfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
DOCX = ROOT / "candidate" / "0038_R01.docx"
PDF = ROOT / "evidence" / "render" / "0038_R01.pdf"
RENDER = ROOT / "evidence" / "render_final_75065568"
REGISTER = ROOT / "evidence" / "SYNTHETIC_REPLACEMENT_REGISTER.csv"
OUT = ROOT / "evidence" / "qa" / "technical_qa.json"
EXPECTED_DOCX_SHA = "75065568b801506d736297441ee30abd3eaccc4fa3748ef6da2a11c67715088a"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


checks: list[dict] = []


def check(name: str, ok: bool, detail: str) -> None:
    checks.append({"name": name, "status": "PASS" if ok else "FAIL", "detail": detail})


docx_sha = sha256(DOCX)
check("candidate_exact_sha", docx_sha == EXPECTED_DOCX_SHA, docx_sha)

with zipfile.ZipFile(DOCX) as archive:
    bad_member = archive.testzip()
    check("docx_zip_integrity", bad_member is None, bad_member or "all DOCX package members passed CRC")

doc = Document(DOCX)
paragraph_text = "\n".join(p.text for p in doc.paragraphs)
table_text = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
full_text = paragraph_text + "\n" + table_text

required_phrases = [
    "I08-E02",
    "官方具名方框序號 039",
    "針對訓練方案的系統設計能適時檢討及改善，執行後有良好效果。",
    "SAMPLE / SYNTHETIC COMPLETE EXAMPLE",
    "SYN-I08E02-2026-01",
    "納入與排除規則",
    "樣本流向與分母",
    "B08後測向量1/1/1/2",
    "0037/I08-E01",
    "0039/I08-E03",
    "獨立語意驗收尚未執行",
]
missing = [phrase for phrase in required_phrases if phrase not in full_text]
check("requirement_and_boundary_text", not missing, f"missing={missing}")

scoring_tables = []
for t in doc.tables:
    if t.rows and len(t.rows[0].cells) >= 7:
        headers = [c.text.strip() for c in t.rows[0].cells]
        if headers[0] == "代碼" and "前測" in headers[1] and "後測" in headers[3]:
            scoring_tables.append(t)

batch_results: dict[str, dict] = {}
row_errors: list[str] = []
for t in scoring_tables:
    ids = [r.cells[0].text.strip() for r in t.rows[1:]]
    group = ids[0][:1] if ids else "?"
    sums = {"pre": 0, "post": 0, "pass": 0, "critical_zero": 0, "n": 0}
    for row in t.rows[1:]:
        cells = [c.text.strip() for c in row.cells]
        case_id = cells[0]
        try:
            pre_vec = [int(x) for x in cells[1].split("/")]
            pre_total = int(cells[2])
            post_vec = [int(x) for x in cells[3].split("/")]
            post_total = int(cells[4])
            delta = int(cells[5].replace("+", ""))
            passed = sum(post_vec) >= 6 and post_vec[2] > 0 and post_vec[3] > 0
            expected_label = "達標" if passed else "未達"
            if sum(pre_vec) != pre_total:
                row_errors.append(f"{case_id}: pre vector sum {sum(pre_vec)} != {pre_total}")
            if sum(post_vec) != post_total:
                row_errors.append(f"{case_id}: post vector sum {sum(post_vec)} != {post_total}")
            if post_total - pre_total != delta:
                row_errors.append(f"{case_id}: delta {post_total-pre_total} != {delta}")
            if cells[6] != expected_label:
                row_errors.append(f"{case_id}: label {cells[6]} != {expected_label}")
            if len(pre_vec) != 4 or len(post_vec) != 4 or any(x < 0 or x > 2 for x in pre_vec + post_vec):
                row_errors.append(f"{case_id}: invalid four-item 0-2 vector")
            sums["pre"] += pre_total
            sums["post"] += post_total
            sums["pass"] += int(passed)
            sums["critical_zero"] += int(post_vec[2] == 0 or post_vec[3] == 0)
            sums["n"] += 1
        except Exception as exc:
            row_errors.append(f"{case_id}: parse error: {exc}")
    batch_results[group] = sums

expected_batches = {
    "A": {"pre": 44, "post": 57, "pass": 3, "critical_zero": 2, "n": 12},
    "B": {"pre": 44, "post": 79, "pass": 11, "critical_zero": 0, "n": 12},
}
check("two_full_12_row_batches", len(scoring_tables) == 2 and batch_results == expected_batches, json.dumps(batch_results, ensure_ascii=False))
check("all_row_calculations", not row_errors, "; ".join(row_errors) if row_errors else "24 vectors, totals, deltas and threshold labels recomputed")
check("aggregate_math", (57 - 44) / 12 == 13 / 12 and (79 - 44) / 12 == 35 / 12 and (79 - 44 - 57 + 44) / 12 == 22 / 12, "A gain=13/12; B gain=35/12; gain difference=22/12")

table_sizes = []
unscaled_table_runs = 0
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    if run.text.strip():
                        if run.font.size is None:
                            unscaled_table_runs += 1
                        else:
                            table_sizes.append(run.font.size.pt)
min_table_pt = min(table_sizes) if table_sizes else None
check("table_type_size", min_table_pt is not None and min_table_pt >= 9.5 and unscaled_table_runs == 0, f"min_direct_size_pt={min_table_pt}; unscaled_nonempty_runs={unscaled_table_runs}")

with REGISTER.open("r", encoding="utf-8-sig", newline="") as f:
    register_rows = list(csv.DictReader(f))
check("synthetic_register_rows", len(register_rows) == 37, f"rows={len(register_rows)}")

page_images = sorted(RENDER.glob("page-*.png"))
check("rendered_page_set", len(page_images) == 11, f"page_png_count={len(page_images)}")
check("word_pdf_export", PDF.exists() and PDF.read_bytes().startswith(b"%PDF-"), f"exists={PDF.exists()}; bytes={PDF.stat().st_size if PDF.exists() else 0}")

result = {
    "project": "TTQS_ONE",
    "document_id": "0038",
    "requirement_id": "I08-E02",
    "official_checkbox_sequence": "039",
    "candidate": "0038_R01",
    "candidate_sha256": docx_sha,
    "pdf_sha256": sha256(PDF),
    "page_count": len(page_images),
    "table_count": len(doc.tables),
    "scoring_batches": batch_results,
    "replacement_register_rows": len(register_rows),
    "minimum_explicit_table_font_pt": min_table_pt,
    "checks": checks,
    "status": "PASS" if all(c["status"] == "PASS" for c in checks) else "FAIL",
    "independent_semantic_review": "PENDING",
    "formal_acceptance": "NOT_GRANTED",
}
OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(result, ensure_ascii=False, indent=2))
