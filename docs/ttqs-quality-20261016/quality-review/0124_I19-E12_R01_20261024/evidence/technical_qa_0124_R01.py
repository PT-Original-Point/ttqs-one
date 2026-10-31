import sys, re, json, hashlib, zipfile, collections
from pathlib import Path
from docx import Document
p=Path(sys.argv[1]); out=Path(sys.argv[2])
b=p.read_bytes(); doc=Document(str(p))
paras=[x.text for x in doc.paragraphs]
tables=[[[' '.join(c.text.split()) for c in row.cells] for row in t.rows] for t in doc.tables]
alltext='\n'.join(paras+['\n'.join(' | '.join(r) for r in t) for t in tables])
flat=' '.join(alltext.split())
checks={}
def ck(name, condition, detail): checks[name]={'status':'PASS' if condition else 'FAIL','detail':detail}
ck('docx_zip_integrity',zipfile.is_zipfile(p),'ZIP container readable')
ck('docx_id_0124','0124' in flat,'DOCX identity appears in extracted body/table text')
ck('requirement_mapping','I19-E12' in flat,'Unique mapping I19-E12 appears in extracted body/table text')
ck('official_requirement_phrase','企業主雇用結訓後學員，肯定機構訓練的成效。' in flat,'Official I19-E12 phrase exact match')
ck('unique_genre','雇主' in flat and ('回饋' in flat or '肯定' in flat),'Employer evidence form and effectiveness affirmation are document-specific')
ck('synthetic_boundary','SYNTHETIC' in flat and 'SAMPLE' in flat and 'REAL' in flat,'Synthetic, sample and real states are explicitly distinguished')
ck('formal_acceptance_boundary','NOT_GRANTED' in flat,'Formal acceptance remains ungranted in candidate')
ck('no_real_original_claim','NOT_PROVIDED' in flat and ('原件' in flat or '原始' in flat),'No real employer original supplied; replacement checklist states not provided')
ck('no_raw_markdown_tokens',not re.search(r'(?m)^\s*#{1,6}\s|```|\|[-:]{3,}\|',alltext),'No obvious raw markdown heading/fence/table separator artifacts')
ck('no_common_draft_placeholders',not any(x in flat.lower() for x in ['lorem ipsum','todo','tbd','insert text here']),'No generic draft placeholders')
# questionnaire/evidence question coverage: locate E-number tokens as whole tokens.
qids=sorted(set(re.findall(r'(?<![A-Z0-9])E[1-9](?![0-9])',flat)))
ck('question_ids_E1_E9',set(qids)=={f'E{i}' for i in range(1,10)},'Detected IDs: '+','.join(qids))
caseids=sorted(set(re.findall(r'SYN-0124-CASE-0[1-3]',flat)))
ck('three_synthetic_cases',len(caseids)==3,'Detected: '+','.join(caseids))
math_patterns=[r'2\s*/\s*3\s*=\s*66\.7\s*%',r'2\s*/\s*2\s*=\s*100\s*%',r'1\s*/\s*2\s*=\s*50\s*%']
math_found=[bool(re.search(x,flat)) for x in math_patterns]
ck('sample_calculations',all(math_found),'Hire verification 2/3=66.7%, response 2/2=100%, explicit positive 1/2=50%; matches synthetic case counts')
ck('employment_rate_withheld','NOT_CALCULATED' in flat,'Overall employment rate withheld when denominator/outcomes are unknown')
ck('adequate_tables',len(doc.tables)>=10,f'{len(doc.tables)} Word tables')
# Explicit font sizes only; inspect inherited defaults separately.
sizes=[]
for para in doc.paragraphs:
 for run in para.runs:
  if run.text.strip() and run.font.size: sizes.append((run.font.size.pt,'body',run.text[:45]))
for table in doc.tables:
 for row in table.rows:
  for cell in row.cells:
   for para in cell.paragraphs:
    for run in para.runs:
     if run.text.strip() and run.font.size: sizes.append((run.font.size.pt,'table',run.text[:45]))
all_explicit=[x[0] for x in sizes]
body_sizes=[x[0] for x in sizes if x[1]=='body']; table_sizes=[x[0] for x in sizes if x[1]=='table']
ck('explicit_font_floor',bool(all_explicit) and min(all_explicit)>=9.0,f'explicit runs min={min(all_explicit) if all_explicit else None} pt; body min={min(body_sizes) if body_sizes else None}; table min={min(table_sizes) if table_sizes else None}')
with zipfile.ZipFile(p) as z:
 entries=z.namelist(); macros=[n for n in entries if 'vbaProject' in n]
ck('no_macro_payload',not macros,'No VBA project detected')
# Read metadata, styles
normal=doc.styles['Normal'].font.size.pt if doc.styles['Normal'].font.size else None
heading=doc.styles['Heading 1'].font.size.pt if doc.styles['Heading 1'].font.size else None
result={'artifact':str(p),'sha256':hashlib.sha256(b).hexdigest(),'bytes':len(b),'paragraphs':len(paras),'tables':len(tables),'question_ids':qids,'synthetic_case_ids':caseids,'explicit_font_stats':{'min':min(all_explicit) if all_explicit else None,'body_min':min(body_sizes) if body_sizes else None,'table_min':min(table_sizes) if table_sizes else None,'normal_style_pt':normal,'heading1_style_pt':heading},'checks':checks,'overall':'PASS' if all(x['status']=='PASS' for x in checks.values()) else 'FAIL','scope':'technical structure/content/math checks only; independent semantic review pending; candidate not formally accepted'}
out.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(result,ensure_ascii=False,indent=2))
