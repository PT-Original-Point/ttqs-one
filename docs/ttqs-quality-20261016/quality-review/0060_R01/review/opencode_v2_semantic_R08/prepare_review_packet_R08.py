from __future__ import annotations

import hashlib
import json
from pathlib import Path

from docx import Document

BASE = Path(__file__).resolve().parents[2]
REPO = BASE.parents[3]
REVIEW = BASE / "review/opencode_v2_semantic_R08"
DOCX = BASE / "candidate/0060_R08.docx"
PDF = BASE / "evidence/render_R08/0060_R08_Word_export.pdf"
BLUEPRINT = BASE / "inputs/0060_UNIQUE_BLUEPRINT_R01.md"
TECH = BASE / "evidence/technical_qa_R08.json"
RENDER = BASE / "evidence/render_R08/render_manifest_R08.json"
VISUAL = BASE / "evidence/visual_review_R08.md"
REPAIR = BASE / "evidence/repair_R08_semantic_findings_receipt.json"
REGISTER = BASE / "sidecar/0060_R08_SYNTHETIC_REPLACEMENT_REGISTER.csv"
R07_REVIEW = BASE / "review/opencode_v2_semantic_R07/independent_review_response_R07.md"
R07_ADJUDICATION = BASE / "review/opencode_v2_semantic_R07/coordinator_adjudication_R07.md"
OFFICIAL = REPO / "docs/ttqs-quality-20261016/quality-review/0058_R01/review/opencode_v2_semantic_R05/inputs/official/official_evidence_table_source.pdf"
MAPPING = REPO / "docs/ttqs-quality-20261016/quality-review/0036_R13/inputs/mapping/checkbox_mapping_compact.json"


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def main() -> None:
    required = (DOCX, PDF, BLUEPRINT, TECH, RENDER, VISUAL, REPAIR, REGISTER, R07_REVIEW, R07_ADJUDICATION, OFFICIAL, MAPPING)
    missing = [str(p) for p in required if not p.is_file()]
    if missing:
        raise SystemExit("Missing review input(s):\n" + "\n".join(missing))
    doc = Document(DOCX)
    body_parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for index, table in enumerate(doc.tables, start=1):
        body_parts.append(f"\n[TABLE {index}]")
        for row in table.rows:
            body_parts.append(" | ".join(" ".join(c.text.split()) for c in row.cells))
    body = "\n".join(body_parts)
    blueprint = BLUEPRINT.read_text(encoding="utf-8")
    tech = json.loads(TECH.read_text(encoding="utf-8"))
    repair = json.loads(REPAIR.read_text(encoding="utf-8"))
    render = json.loads(RENDER.read_text(encoding="utf-8"))
    prior = R07_REVIEW.read_text(encoding="utf-8")
    adjudication = R07_ADJUDICATION.read_text(encoding="utf-8")
    if sha(DOCX) != "92962DF774A7A16546829D8968DC274394FC1EFCC12421C1A355F45127E9B5C3":
        raise SystemExit("R08 candidate SHA changed before packet assembly")
    if tech.get("candidate_sha256") != sha(DOCX) or render.get("candidate_sha256") != sha(DOCX):
        raise SystemExit("R08 exact-SHA evidence mismatch")
    packet = f"""# Independent exact-SHA semantic review packet — DOCX 0060 R08

## Immutable identity and source hashes
- Project: TTQS_ONE
- Document ID: 0060
- Official requirement: I14-E01; official checkbox 065
- Genre in formal mapping: analysis_report
- Candidate DOCX: `{DOCX}`
- Candidate SHA-256: `{sha(DOCX)}`
- Word-export PDF SHA-256: `{sha(PDF)}`
- Unique Blueprint SHA-256: `{sha(BLUEPRINT)}`
- Official source table PDF SHA-256: `{sha(OFFICIAL)}`
- Formal mapping JSON SHA-256: `{sha(MAPPING)}`
- R08 deterministic technical receipt SHA-256: `{sha(TECH)}`
- R08 render manifest SHA-256: `{sha(RENDER)}`; {render['page_count']} pages and {render['rendered_png_count']} rendered PNGs.
- R08 visual receipt SHA-256: `{sha(VISUAL)}`; coordinator directly inspected full-resolution changed pages 18, 19, 20, 23; 22 unchanged page images match R07 byte-for-byte.
- R08 repair receipt SHA-256: `{sha(REPAIR)}`
- R08 synthetic replacement register SHA-256: `{sha(REGISTER)}`

## Review instructions

Return an explicit PASS, PARTIAL, or FAIL for this exact SHA. Independently inspect the attached original DOCX, rendered PDF, official source PDF and mapping JSON. Do not inherit a prior verdict, treat deterministic checks as acceptance, or assume the repair receipt is correct. Compare the entire artifact to the official requirement and this unique Blueprint. Cite exact section/table/row/page locations for each finding. Separate semantic adequacy, mathematical consistency, evidence boundary, and visual usability. The R07 findings below are repair targets to verify, not conclusions about R08. Do not claim direct inspection of an image page unless you inspect it.

Specifically verify:

1. DOCX 0060 maps to I14-E01 / checkbox 065 and is correctly an analysis report, distinct from adjacent system, record-building, retrieval and retrieval-speed requirements.
2. The work is a usable small-association record-classification and coding analysis report, not just a generic manual. Check whether an administrator can classify a record, record source and locator, decide accept/correct/HOLD/reject, explain the decision and obtain a second-person check.
3. Full Blueprint coverage: 12 record types/PDDRO, code grammar and segment restrictions, 11 metadata fields, role-based procedures, workflow/exception branches, full worked table rows, metric definitions, calculations, correction trail, hold/reject criteria, real-source replacement process and conclusion.
4. Independently trace all 12 paired Table 9-1 / Table 9-2 items. Check code, title, date, phase/type, access/retention, source/locator, formation, SAMPLE marker, conclusion, and cross-reference. Verify RSC05->03/04, ATT06->03, DLG07->03/04, ASM08->attendance06, FDB09->attendance06, FUP10->03/06, EXC11->05/06/07, and IMP12->09/10/11.
5. Verify the four R07 correction-trail findings: RSC duplicate 004->005; ATT label INT->RST; DLG locator LOC-REC-2026-0815B; ASM08 must not claim a separate label correction. Confirm row 11 states distinct baselines 11/12, 10/12, 11/12 and fixed results 12/12.
6. Verify Table 15's final example code includes the required `-V01` version segment. Verify Chapter 13 says a zero denominator/no applicable data is N/A with scope stated and never 100%.
7. Recompute the shown sample counts and percentages. Confirm denominators agree with cohorts and evidence links. Report arithmetic correctness separately from presence of numeric strings.
8. Inspect the 66-row sidecar replacement register and verify it is explicitly SAMPLE, all fact IDs are unique, and facts 055-066 match the current R08 Table 12 data rows exactly.
9. Distinguish SAMPLE/SYNTHETIC from REAL everywhere. Flag any implied actual association decision, approval, staffing, training, record, system, policy, legal retention rule, third-party proof, or signature not supported by evidence. Do not treat sample material as actual TTQS evidence.
10. Review layout/readability in the PDF, particularly changed pages 18, 19, 20, 23 and tables spanning pages. Report any clipping, overlap, difficult-to-use table, or page-flow defect with locator. The coordinator receipt is context only.
11. Return: (a) exact candidate SHA actually received; (b) explicit PASS/PARTIAL/FAIL; (c) findings with severity and exact locator; (d) positive evidence; (e) limitations. If any point is unverified, say so.

## R07 findings and repair context

The R07 independent reviewer returned PARTIAL. Its raw response and coordinator adjudication are reproduced below for traceability. Reassess R08 from the attached artifact; do not reuse or rubber-stamp the prior verdict.

### Full R07 independent reviewer response

{prior}

### Full R07 coordinator adjudication

{adjudication}

## R08 targeted repair receipt

```json
{json.dumps(repair, ensure_ascii=False, indent=2)}
```

## R08 deterministic technical receipt

```json
{json.dumps(tech, ensure_ascii=False, indent=2)}
```

## Full document-specific Unique Blueprint

{blueprint}

## Full text extracted from exact R08 DOCX

Every non-empty paragraph and every table row below is extracted from the exact SHA identified above. Verify against the attached DOCX/PDF.

{body}
"""
    REVIEW.mkdir(parents=True, exist_ok=True)
    out = REVIEW / "review_packet_R08.md"
    out.write_text(packet, encoding="utf-8", newline="\n")
    print(json.dumps({"packet": str(out), "packet_sha256": sha(out), "packet_bytes": out.stat().st_size, "docx_sha256": sha(DOCX), "pdf_sha256": sha(PDF), "blueprint_sha256": sha(BLUEPRINT), "official_pdf_sha256": sha(OFFICIAL), "mapping_sha256": sha(MAPPING), "technical_sha256": sha(TECH), "visual_receipt_sha256": sha(VISUAL), "register_sha256": sha(REGISTER), "extracted_chars": len(body), "tables": len(doc.tables), "pages": render["page_count"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
