from __future__ import annotations
import csv, hashlib, io, json, os, zipfile
from pathlib import Path
from docx import Document
BASE=Path(__file__).resolve().parents[1]
SRC=BASE/'candidate/0060_R07.docx'
DST=BASE/'candidate/0060_R08.docx'
TMP=DST.with_suffix('.docx.tmp')
SIDE_SRC=BASE/'sidecar/0060_R04_SYNTHETIC_REPLACEMENT_REGISTER.csv'
SIDE_DST=BASE/'sidecar/0060_R08_SYNTHETIC_REPLACEMENT_REGISTER.csv'
EXPECTED_SRC='044645199B2D704760D3F8729593641EBE54EBF58CB583ABDC026FA355A69286'

def sha(data:bytes)->str: return hashlib.sha256(data).hexdigest().upper()
def norm(s:str)->str: return ' '.join(s.split())
changes=[
('本報告內合成桌面檢核：初始定位缺漏已修正，欄位齊備。核對資源項對應04配置；實體場地與設備可用性尚未確認。',
 '本報告內合成桌面檢核：初始重複碼已修正，RSC樣張代碼流水號由004更正為005；12筆代碼均唯一。核對資源項對應03與04；實體場地與設備可用性尚未確認。'),
('本報告內合成桌面檢核：初始重複碼已修正，標記唯一。核對人數邏輯閉合；RST僅為樣本處理標籤無真實個資，真實名冊來源、處理權限及實際簽到資料均未確認。',
 '本報告內合成桌面檢核：簽到名冊情境原先誤標INT，已更正為RST；12報名、10到課、2缺席數量閉合。RST僅為樣本處理標籤，不含真實個資；真實名冊來源、處理權限及實際簽到資料均未確認。'),
('本報告內合成桌面檢核：日誌節點完整對應03與04進度，必填欄合規。核對節次相符；非真實紀錄，實際授課、紀錄來源與負責角色均未確認。',
 '本報告內合成桌面檢核：初始缺少來源存放定位，已補入LOC-REC-2026-0815B；日誌節點對應03與04。非真實授課紀錄，實際授課、紀錄來源與負責角色均未確認。'),
('本報告內合成桌面檢核：初始標籤不符已修正，與03課程計畫及04評量安排一致。核對10人評量母體與06簽到的10人到課數相符；RST僅為樣本處理標籤無真實個資，真實成效尚未確認。',
 '本報告內合成桌面檢核：10名到課中9份有效評量、8人達標、1人未達標；與03課程計畫、04評量安排及06簽到10人母體相符。本列未另設缺陷修正；RST僅為樣本處理標籤，真實成效尚未確認。'),
('本報告內合成桌面檢核：記錄三項演練缺失修正歷程，示範資料已由10/12修正達12/12一致；非真實錯誤，實體稽核尚未確認。',
 '本報告內合成桌面檢核：初始代碼唯一性11/12、必填欄位完整有效10/12、處理標籤相符11/12；更正後三項均為12/12。此為SAMPLE桌面演練設計，非真實錯誤；實體稽核尚未確認。'),
('產出代碼：TT-2026-SYN60A-DO-ATT-RST-S3-006',
 '產出代碼：TT-2026-SYN60A-DO-ATT-RST-S3-006-V01'),
('所有指標均依數學公式嚴格定義分子與分母，計算結果採四捨五入至小數點後第一位。各項比率之定義如下：代碼唯一性比率為「無重複之唯一代碼筆數 / 總登錄筆數」；中繼資料完整有效率為「各項必填欄位完整且格式合規之筆數 / 總登錄筆數」；處理標籤相符率為「處理標籤與文件機敏性質完全相符之筆數 / 總登錄筆數」。在演練初始階段，代碼唯一性為 11/12 = 91.7%，欄位完整有效率為 10/12 = 83.3%（第005筆因重碼未生效、第007筆缺漏定位），處理標籤相符率為 11/12 = 91.7%。經獨立第二人覆核提報缺失並填寫 EXC 表單更正後，三項指標均提升至 12/12 = 100.0%。',
 '所有指標均依數學公式嚴格定義分子與分母，計算結果採四捨五入至小數點後第一位。各項比率之定義如下：代碼唯一性比率為「無重複之唯一代碼筆數 / 總登錄筆數」；中繼資料完整有效率為「各項必填欄位完整且格式合規之筆數 / 總登錄筆數」；處理標籤相符率為「處理標籤與文件機敏性質完全相符之筆數 / 總登錄筆數」。在演練初始階段，代碼唯一性為 11/12 = 91.7%，欄位完整有效率為 10/12 = 83.3%（第005筆因重碼未生效、第007筆缺漏定位），處理標籤相符率為 11/12 = 91.7%。經獨立第二人覆核提報缺失並填寫 EXC 表單更正後，三項指標均提升至 12/12 = 100.0%。若特定期間或範圍內沒有任何適用資料、分母為0，結果標示 N/A（不適用）並說明資料範圍；不得報100%。')]

def build_docx()->dict:
    if sha(SRC.read_bytes())!=EXPECTED_SRC: raise SystemExit('R07 source SHA mismatch')
    if TMP.exists(): raise SystemExit('R08 temp already exists')
    with zipfile.ZipFile(SRC,'r') as zin:
        infos=zin.infolist(); original={i.filename:zin.read(i.filename) for i in infos}
    xml=original['word/document.xml']
    for old,new in changes:
        a,b=old.encode('utf-8'),new.encode('utf-8')
        if xml.count(a)!=1: raise SystemExit(f'expected one exact OOXML occurrence, got {xml.count(a)} for {old[:48]}')
        xml=xml.replace(a,b,1)
    outparts=dict(original); outparts['word/document.xml']=xml
    with zipfile.ZipFile(TMP,'w') as zout, zipfile.ZipFile(SRC,'r') as zin:
        for info in zin.infolist(): zout.writestr(info,outparts[info.filename])
    with zipfile.ZipFile(TMP,'r') as z:
        if z.testzip() is not None: raise SystemExit('R08 ZIP integrity failure')
        readback={n:z.read(n) for n in z.namelist()}
    changed_parts=[n for n in original if original[n]!=readback[n]]
    if changed_parts!=['word/document.xml']: raise SystemExit(f'unexpected changed parts: {changed_parts}')
    Document(TMP)
    if DST.exists():
        # A prior interrupted run may have created the exact intended DOCX.
        # Preserve its bytes and only accept it when it equals this deterministic rebuild.
        if DST.read_bytes()!=TMP.read_bytes():
            TMP.unlink()
            raise SystemExit('Existing R08 DOCX differs from deterministic repair; preserved without overwrite')
        TMP.unlink()
    else:
        os.replace(TMP,DST)
    doc=Document(DST)
    assert doc.tables[11].rows[5].cells[3].text==changes[0][1]
    assert doc.tables[11].rows[6].cells[3].text==changes[1][1]
    assert doc.tables[11].rows[7].cells[3].text==changes[2][1]
    assert doc.tables[11].rows[8].cells[3].text==changes[3][1]
    assert doc.tables[11].rows[11].cells[3].text==changes[4][1]
    assert doc.tables[14].rows[5].cells[3].text==changes[5][1]
    paras=[p.text for p in doc.paragraphs]
    assert any(p==changes[6][1] for p in paras)
    body='\n'.join(paras+[c.text for t in doc.tables for r in t.rows for c in r.cells])
    assert not any(x in body for x in ('待填','待核','空白表','保持空白','會後填寫'))
    return {'doc_id':'0060','requirement_id':'I14-E01','official_checkbox':'065','source_revision':'R07','candidate_revision':'R08','source_sha256':sha(SRC.read_bytes()),'candidate_sha256':sha(DST.read_bytes()),'candidate_bytes':DST.stat().st_size,'changed_ooxml_parts':changed_parts,'other_package_parts_byte_identical':True,'zip_testzip_error':None,'python_docx_opened':True,'targeted_text_changes':[{'locator':loc,'old':old,'new':new} for loc,(old,new) in zip(['Table12 row05 conclusion','Table12 row06 conclusion','Table12 row07 conclusion','Table12 row08 conclusion','Table12 row11 conclusion','Table15 final output-code cell','Chapter13 metric procedure paragraph'],changes)],'forbidden_literals_remaining_in_docx_body':[]}

def build_register()->dict:
    if not SIDE_SRC.is_file(): raise SystemExit('R04 source register missing')
    base_sha=sha(SIDE_SRC.read_bytes())
    if base_sha!='8355A495471E34CFF95CE57BE32CD65FC3EDEE18446DFEE76344000F6757D2C6': raise SystemExit('Source register SHA mismatch')
    previous_sha=sha(SIDE_DST.read_bytes()) if SIDE_DST.exists() else None
    if previous_sha not in (None,'41638683BB6A71D668BBA388283B71CC4FEFBDC3D71E1D9B17A795CEEA850A14'):
        raise SystemExit('Existing R08 register is not the known interrupted output; preserved without overwrite')
    with SIDE_SRC.open(encoding='utf-8-sig',newline='') as stream: rows=list(csv.DictReader(stream)); fields=list(rows[0].keys())
    for row in rows:
        for key,value in row.items():
            value=value.replace('只示意資料欄位/標籤待核流程','只示意資料欄位／標籤未確認前之處理流程')
            row[key]=value.replace('待核','未確認')
    doc=Document(DST); detail=doc.tables[11]
    for ix in range(12):
        fact=f'SYN-0060-FACT-{55+ix:03d}'
        row=next((r for r in rows if r.get('fact_id')==fact),None)
        if row is None: raise SystemExit(f'missing register row {fact}')
        row['value']=' | '.join(norm(c.text) for c in detail.rows[ix+1].cells)
    buf=io.StringIO(newline='')
    writer=csv.DictWriter(buf,fieldnames=fields,lineterminator='\n',extrasaction='raise')
    writer.writeheader(); writer.writerows(rows)
    new_register=b'\xef\xbb\xbf'+buf.getvalue().encode('utf-8')
    SIDE_TMP=SIDE_DST.with_suffix('.csv.tmp')
    SIDE_TMP.write_bytes(new_register)
    os.replace(SIDE_TMP,SIDE_DST)
    with SIDE_DST.open(encoding='utf-8-sig',newline='') as stream: check=list(csv.DictReader(stream))
    assert len(check)==66 and len({r['fact_id'] for r in check})==66
    assert all(r['doc_id']=='0060' and r['status']=='SAMPLE' for r in check)
    assert not any('待核' in v for r in check for v in r.values())
    for ix in range(12):
        fact=f'SYN-0060-FACT-{55+ix:03d}'
        actual=' | '.join(norm(c.text) for c in detail.rows[ix+1].cells)
        assert next(r for r in check if r['fact_id']==fact)['value']==actual
    return {'source_register_sha256':base_sha,'candidate_register_path':str(SIDE_DST),'candidate_register_sha256':sha(SIDE_DST.read_bytes()),'previous_interrupted_register_sha256':previous_sha,'rows':len(check),'unique_fact_ids':True,'all_doc_id_0060':True,'all_sample':True,'facts_055_066_exactly_match_R08_Table12':True,'forbidden_literal_hits':0}

if __name__=='__main__':
    receipt=build_docx(); register=build_register(); receipt['register']=register
    (BASE/'evidence/repair_R08_semantic_findings_receipt.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')
    print(json.dumps({'candidate_sha256':receipt['candidate_sha256'],'candidate_bytes':receipt['candidate_bytes'],'register_sha256':register['candidate_register_sha256'],'changed_parts':receipt['changed_ooxml_parts']},ensure_ascii=False,indent=2))
