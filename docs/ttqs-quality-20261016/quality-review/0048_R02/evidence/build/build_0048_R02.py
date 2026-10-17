from __future__ import annotations

import csv
import hashlib
import json
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


REPO = Path(r"E:\TTQS\_TTQS_ONE_REVIEW_SYNC_0036_R08_20261016")
ROOT = REPO / "docs/ttqs-quality-20261016/quality-review/0048_R02"
INPUTS = ROOT / "inputs"
CANDIDATE = ROOT / "candidate"
EVIDENCE = ROOT / "evidence"
DOCX_PATH = CANDIDATE / "0048_R02.docx"
BLUEPRINT_SOURCE = REPO / "docs/ttqs-quality-20261016/quality-review/0047_R08/inputs/UNIQUE_BLUEPRINT_0045_0047_0048.jsonl"
SOURCE_PDF = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\SOURCES\TTQS查核佐證文件資料表-訓練機構版 (2).pdf")
CHECKBOX_INVENTORY = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\MASTER\SOURCE_CHECK_20261004\checkbox_inventory.json")

NAVY = "17324D"
TEAL = "286D79"
PALE = "EAF1F4"
PALE_GOLD = "FFF4D7"
PALE_RED = "FCE8E6"
MID = "64748B"
WHITE = "FFFFFF"
TEXT = "263442"
BODY_PT = 10.5
TABLE_PT = 9.5


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def load_blueprint() -> dict:
    raw = BLUEPRINT_SOURCE.read_bytes()
    rows = [json.loads(line) for line in raw.splitlines() if line.strip()]
    matches = [r for r in rows if str(r.get("deliverable_id")) == "0048"]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one unique Blueprint row for 0048; got {len(matches)}")
    row = matches[0]
    assert row["requirement_id"] == "I10-E03"
    assert row["document_role"] == "evidence_record"
    assert row["official_requirement_text"] == "大型採購均辦理會議通過、師資遴聘後亦經過會議或審查通過並留有紀錄。"
    assert row["official_evidence_table_page"] == "3"
    return row


def set_run_font(run, size=BODY_PT, bold=None, color=TEXT, name="Microsoft JhengHei"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)


def set_cell_fill(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=85, start=100, bottom=85, end=100):
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    margins = tcpr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    trpr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    trpr.append(repeat)


def set_no_split(row):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))


def set_cell_text(cell, text, *, size=TABLE_PT, bold=False, color=TEXT, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    if align is not None:
        p.alignment = align
    run = p.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table(doc, headers, rows, widths=None, *, font_size=TABLE_PT, header_color=NAVY, banded=True):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    if widths:
        for col, width in zip(table.columns, widths):
            col.width = Inches(width)
    header = table.rows[0]
    set_repeat_table_header(header)
    set_no_split(header)
    for i, text in enumerate(headers):
        set_cell_text(header.cells[i], text, size=font_size, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_fill(header.cells[i], header_color)
        if widths:
            header.cells[i].width = Inches(widths[i])
    for r_index, values in enumerate(rows):
        cells = table.add_row().cells
        set_no_split(table.rows[-1])
        for c_index, value in enumerate(values):
            set_cell_text(cells[c_index], value, size=font_size)
            if widths:
                cells[c_index].width = Inches(widths[c_index])
            if banded and r_index % 2 == 1:
                set_cell_fill(cells[c_index], "F5F8FA")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_callout(doc, title, body, fill=PALE, title_color=NAVY):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.12
    lead = p.add_run(title + "｜")
    set_run_font(lead, size=BODY_PT, bold=True, color=title_color)
    detail = p.add_run(body)
    set_run_font(detail, size=BODY_PT, color=TEXT)


def add_p(doc, text="", *, bold=False, color=TEXT, size=BODY_PT, align=None, italic=False, after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, color=color)
        r.italic = italic
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(3 if level == 1 else 1)
    p.paragraph_format.space_after = Pt(6 if level == 1 else 4)
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 12, bold=True, color=NAVY if level == 1 else TEAL)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    set_run_font(r, size=BODY_PT)
    return p


def page(doc):
    doc.add_page_break()


def add_page_field(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=9.5, color=MID)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def add_field_line(doc, label, blank="________________________________"):
    add_p(doc, f"{label}　{blank}")


def build_register():
    rows = [
        ("SYN-048-001", "SYN-I10E03-A-001", "示範案例A識別碼", "§5–6；案例識別", "真實採購案號及核定案卷", "採購承辦", "SAMPLE/SYNTHETIC"),
        ("SYN-048-002", "2026-09-12", "示範會議日期", "§5；會議基本資料", "實際通知、簽到及會議紀錄原件", "紀錄人", "SAMPLE/SYNTHETIC"),
        ("SYN-048-003", "13:30–15:00", "示範會議時段", "§5；會議基本資料", "原始議程及會議紀錄", "紀錄人", "SAMPLE/SYNTHETIC"),
        ("SYN-048-004", "SYN-ROOM-A（虛構場地代碼）", "示範地點", "§5；會議基本資料", "實際會議地點／線上會議紀錄", "紀錄人", "SAMPLE/SYNTHETIC"),
        ("SYN-048-005", "志工服務紀錄管理課程與教材採購", "示範採購案名稱", "§5–6；案例摘要", "真實核定需求及採購案卷", "需求承辦", "SAMPLE/SYNTHETIC"),
        ("SYN-048-006", "12名示範參訓者", "示範人數", "§5；需求摘要", "去識別化需求調查及核准課程計畫", "需求承辦", "SAMPLE/SYNTHETIC"),
        ("SYN-048-007", "2次課程、合計12小時", "示範課程規格", "§5；需求摘要", "核定課程計畫及時程", "課程承辦", "SAMPLE/SYNTHETIC"),
        ("SYN-048-008", "新臺幣100,000元（虛構示範門檻）", "示範大型採購判斷門檻", "§5–6；分類判斷", "協會正式採購辦法及核准版本", "制度權責者", "SAMPLE/SYNTHETIC"),
        ("SYN-048-009", "新臺幣120,000元", "示範預算上限", "§5–6；預算欄", "正式預算核定原件", "預算權責者", "SAMPLE/SYNTHETIC"),
        ("SYN-048-010", "新臺幣118,000元", "示範候選A報價", "§6；比較表", "候選A原始報價及含稅範圍", "採購承辦", "SAMPLE/SYNTHETIC"),
        ("SYN-048-011", "新臺幣113,000元", "示範候選B報價", "§6；比較表", "候選B原始報價及含稅範圍", "採購承辦", "SAMPLE/SYNTHETIC"),
        ("SYN-048-012", "新臺幣127,000元", "示範候選C報價", "§6；比較表", "候選C原始報價及含稅範圍", "採購承辦", "SAMPLE/SYNTHETIC"),
        ("SYN-048-013", "82／100分", "示範候選A評估摘要", "§6；會議引用資料", "可追溯評估表及評分依據", "評選紀錄人", "SAMPLE/SYNTHETIC"),
        ("SYN-048-014", "88／100分", "示範候選B評估摘要", "§6；會議引用資料", "可追溯評估表及評分依據", "評選紀錄人", "SAMPLE/SYNTHETIC"),
        ("SYN-048-015", "77／100分", "示範候選C評估摘要", "§6；會議引用資料", "可追溯評估表及評分依據", "評選紀錄人", "SAMPLE/SYNTHETIC"),
        ("SYN-048-016", "7,000元；預算使用率94.2%", "示範價差與使用率", "§6；算式", "以真實報價及核定預算重算", "覆核人", "SAMPLE/SYNTHETIC"),
        ("SYN-048-017", "SYN-R01｜訓練需求承辦", "示範與會角色", "§5；出席表", "本人確認之姓名、職務及出席紀錄", "紀錄人", "SAMPLE/SYNTHETIC"),
        ("SYN-048-018", "SYN-R02｜課程審查角色", "示範與會角色", "§5；出席表", "本人確認之姓名、職務及出席紀錄", "紀錄人", "SAMPLE/SYNTHETIC"),
        ("SYN-048-019", "SYN-R03｜預算核對角色", "示範與會角色", "§5；出席表", "本人確認之姓名、職務及出席紀錄", "紀錄人", "SAMPLE/SYNTHETIC"),
        ("SYN-048-020", "SYN-R04｜會議紀錄角色", "示範與會角色", "§5；出席表", "本人確認之姓名、職務及出席紀錄", "紀錄人", "SAMPLE/SYNTHETIC"),
        ("SYN-048-021", "4人出席；示範門檻3人", "示範出席／會議有效性算例", "§5–6；完整度核對", "正式組織章程、授權規則及簽到原件", "會議主持／覆核人", "SAMPLE/SYNTHETIC"),
        ("SYN-048-022", "贊成4、反對0、迴避0", "示範表決結果", "§6；決議欄", "正式會議紀錄及表決／決議證據", "紀錄人", "SAMPLE/SYNTHETIC"),
        ("SYN-048-023", "示範決議：同意依條件續辦採購", "示範決議文字", "§6；決議欄", "有權機關正式核准紀錄", "權責核准者", "SAMPLE/SYNTHETIC"),
        ("SYN-048-024", "SYN-A-ATT-01 至 SYN-A-ATT-04", "虛構附件識別碼", "§5–6；附件索引", "真實需求、預算、評估與利益迴避原件", "案卷管理者", "SAMPLE/SYNTHETIC"),
        ("SYN-048-025", "SYN-I10E03-B-001", "示範案例B識別碼", "§7–8；案例識別", "真實師資遴聘案號及審查卷宗", "師資承辦", "SAMPLE/SYNTHETIC"),
        ("SYN-048-026", "2026-09-19", "示範師資審查日期", "§7；會議基本資料", "實際通知、出席及審查紀錄原件", "紀錄人", "SAMPLE/SYNTHETIC"),
        ("SYN-048-027", "10:00–11:00；SYN-ROOM-B", "示範審查時段與地點", "§7；會議基本資料", "正式會議／審查原始紀錄", "紀錄人", "SAMPLE/SYNTHETIC"),
        ("SYN-048-028", "SYN-R11至SYN-R13；3人出席", "示範審查角色及出席數", "§7；出席表", "本人確認之姓名、角色及簽到原件", "紀錄人", "SAMPLE/SYNTHETIC"),
        ("SYN-048-029", "示範門檻2人；3人有決定權", "示範審查有效性規則", "§7–8；權限欄", "正式授權／遴聘規則及有效版本", "制度權責者", "SAMPLE/SYNTHETIC"),
        ("SYN-048-030", "權重35、25、15、15、10；總計100", "示範評分權重", "§8；評分表", "事前核准之評選標準與版次", "評選召集人", "SAMPLE/SYNTHETIC"),
        ("SYN-048-031", "候選X評分5、4、4、4、5；89分", "示範候選X分數", "§8；評分及計算", "評選者原始評分表及資格原件", "評選紀錄人", "SAMPLE/SYNTHETIC"),
        ("SYN-048-032", "候選Y評分4、4、5、3、4；80分", "示範候選Y分數", "§8；評分及計算", "評選者原始評分表及資格原件", "評選紀錄人", "SAMPLE/SYNTHETIC"),
        ("SYN-048-033", "贊成3、反對0、迴避0", "示範師資審查決議票數", "§8；決議欄", "正式審查紀錄及具權限者核准證據", "紀錄人", "SAMPLE/SYNTHETIC"),
        ("SYN-048-034", "示範結論：X列優先推薦，完成資格核驗後始得聘任", "示範審查結論", "§8；結論與條件", "真實審查紀錄、證照核驗與聘任授權", "師資權責者", "SAMPLE/SYNTHETIC"),
        ("SYN-048-035", "SYN-B-ATT-01 至 SYN-B-ATT-04", "虛構師資附件識別碼", "§7–8；附件索引", "真實履歷、證照、評分原件與利益迴避聲明", "案卷管理者", "SAMPLE/SYNTHETIC"),
        ("SYN-048-036", "案例A必填欄位10/10；100%", "示範記錄完整度", "§9；核對示例", "按真實必填清單逐欄重核", "覆核人", "SAMPLE/SYNTHETIC"),
        ("SYN-048-037", "案例B必填欄位9/9；100%", "示範記錄完整度", "§9；核對示例", "按真實必填清單逐欄重核", "覆核人", "SAMPLE/SYNTHETIC"),
        ("SYN-048-038", "4名參與者的角色欄均有示範值", "示範完整性判斷輸入", "§9；完整度核對", "核對實際簽到原件及授權角色", "覆核人", "SAMPLE/SYNTHETIC"),
        ("SYN-048-039", "會議決議與附件索引可追溯", "示範證據鏈判斷", "§9；追溯核對", "逐一核對真實原件及存檔位置", "覆核人", "SAMPLE/SYNTHETIC"),
    ]
    return rows


def make_doc(blueprint, facts):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(0.62)
    sec.bottom_margin = Inches(0.62)
    sec.left_margin = Inches(0.66)
    sec.right_margin = Inches(0.66)
    styles = doc.styles
    for style_name in ("Normal", "List Bullet"):
        st = styles[style_name]
        st.font.name = "Microsoft JhengHei"
        st.font.size = Pt(BODY_PT)
        st.font.color.rgb = RGBColor.from_string(TEXT)
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.1
    title_style = styles["Title"]
    title_style.font.name = "Microsoft JhengHei"
    title_style.font.size = Pt(23)
    title_style.font.color.rgb = RGBColor.from_string("000000")
    title_style.font.bold = True
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    title_ppr = title_style._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("TTQS ONE｜DOCX 0048｜I10-E03｜官方方框 049")
    set_run_font(r, size=9.5, color=MID)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("候選示範文件｜SAMPLE/SYNTHETIC｜未經協會正式核定　　頁 ")
    set_run_font(r, size=9.5, color=MID)
    add_page_field(footer)

    # Page 1: user-facing control page and 30-second use guide.
    p = doc.add_paragraph(style="Title")
    title_ppr = p._p.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    p.paragraph_format.space_before = Pt(20)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("大型採購會議與師資遴聘審查紀錄")
    set_run_font(r, size=23, bold=True, color="000000")
    add_p(doc, "會議決議、審查依據與原件追溯紀錄", size=14, color=TEAL, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
    add_callout(doc, "重要狀態｜SAMPLE / SYNTHETIC", "本文件及兩個填寫案例均為合成示範；不是社團法人屏東縣原始點關懷協會的真實會議、出席、表決、核准、採購、師資聘任或第三方原件。不可當作真實佐證提交。正式使用時須換成有權限者形成並可查驗的原始紀錄。", PALE_RED, "9B1C1C")
    add_table(doc, ["文件控制", "內容"], [
        ["文件編號／要求", "DOCX 0048｜I10-E03｜官方方框 049（P03-B10）"],
        ["文類／用途", "可直接操作的專屬佐證紀錄；保存大型採購會議通過，以及師資遴選後經會議或審查通過的紀錄索引。"],
        ["使用者／時點", "訓練、採購或師資承辦；會前建檔、會中逐項記錄、會後核對原件與決議條件。"],
        ["版本／生效", "R02 候選示範；未核定、未生效。未填入製作者、核准者、簽章或正式生效日。"],
    ], [1.45, 5.55])
    heading(doc, "30秒開始使用", 2)
    add_table(doc, ["步驟", "你要做的事", "應留下的輸出"], [
        ["1 會前", "依正式採購辦法、章程及授權表，確認是否屬大型採購、誰能召集／決議、會議有效條件。", "規範名稱、版次、條文、附件清單"],
        ["2 會中", "依議程逐題記錄事實、選項、意見、決議、票數或審查結論、條件與利益迴避。", "可追溯的會議／審查紀錄"],
        ["3 會後", "由紀錄人與有權覆核者核對簽到、議程、證據附件、決議權限及後續責任。", "原件索引、缺失與補正紀錄"],
        ["4 歸檔", "保存原始通知、出席、議案、評估、決議及核准／審查原件；本表只作索引與記錄。", "可調閱的正式卷宗"],
    ], [0.75, 3.35, 2.9])
    add_p(doc, "對象定位：提供原始點協會行政人員熟悉欄位與操作的示範工具；不是協會已採用的制度或會議紀錄。", color=MID, size=9.5)

    # Page 2: source, scope and evidence chain.
    page(doc)
    heading(doc, "01 官方要求與本件證據責任")
    add_callout(doc, "官方要求原文｜來源表第3頁／P03-B10", blueprint["official_requirement_text"], PALE_GOLD)
    add_table(doc, ["官方要求", "本件正式紀錄欄位（§03）", "附件／原件定位與核對方式"], [
        ["大型採購均辦理會議通過", "案件類型、規範依據、出席／迴避、討論／證據、決議／審查結論、後續責任。", "§03基本欄位、出席／權限、議程討論、決議及原件索引表；REAL須連有效採購／授權制度、會議通知與簽到、議案／報價、正式決議及受控卷宗位置。"],
        ["師資遴聘後亦經過會議或審查通過", "案件類型、規範依據、出席／迴避、討論／證據、決議／審查結論、後續責任。", "§03正式紀錄欄位及附件索引；REAL須連事前遴聘標準、本人提交且核驗的資格原件、原始評分、審查決議與聘任授權。"],
        ["並留有紀錄", "案件／紀錄編號、版次與日期、紀錄責任、覆核狀態、原件位置及更正軌跡。", "§03紀錄與原件索引表；REAL須可由受權人開啟正式簽核／會議原件並核對版次、保管位置及更正紀錄。"],
    ], [1.45, 2.45, 3.2])
    heading(doc, "證據路徑 紀錄不是原件的替代品", 2)
    add_table(doc, ["節點", "核對問題", "連結／輸出"], [
        ["A 事前依據", "適用的正式規範、預算、議案與評選標準是哪一版？", "規範條號／版本＋原件索引"],
        ["B 決策過程", "誰參與、誰迴避、有哪些選項與討論？", "通知、簽到、議程、討論紀錄"],
        ["C 通過判定", "會議／審查是否有權、有效且結論清楚？", "授權與有效性核對＋決議／票數"],
        ["D 會後追蹤", "決議條件、採購／聘任、驗收與歸檔由誰完成？", "責任期限、原件、補正及結案紀錄"],
    ], [1.05, 3.35, 2.6])
    add_callout(doc, "不得自行填造規則", "若協會正式採購門檻、會議權限、法定／內部有效條件或師資遴聘規則尚未取得，該欄記為「待取得正式規範」，停止把案件標成已通過；不得把本文件的示範數值或角色當成本會制度。", PALE_RED, "9B1C1C")

    # Page 3: controlled procedure and field definitions.
    page(doc)
    heading(doc, "02 建檔規則與欄位定義")
    add_p(doc, "每一案件各用一個唯一案號；採購會議與師資審查如屬不同案件，分開建檔。會議紀錄保留當時語意，不以摘要改寫原決議。補充或更正時保留原始版本、原因、時間與更正人。")
    add_table(doc, ["欄位", "定義／填寫內容", "來源／附件定位", "填寫者／責任角色", "驗證規則"], [
        ["案件／紀錄編號", "識別單一案件、版本及形成日。", "正式編碼規則、原始紀錄。", "案件承辦／紀錄人。", "案號唯一；版次日期完整，不覆寫原稿。"],
        ["案件類型", "標明大型採購或師資遴聘後審查。", "官方要求及適用制度。", "案件承辦人。", "能對應要求事實；不可只寫一般會議。"],
        ["規範依據", "列出有效制度、條文、權限與效期。", "核定制度、授權表。", "制度權責者／紀錄人。", "版本、條次及權限有效；缺件即暫停判定。"],
        ["出席／迴避", "記本人、角色、時段、參與方式及迴避。", "通知、簽到、線上軌跡、聲明。", "紀錄人；當事人核實。", "對照簽到與權限；缺證標未證實。"],
        ["討論／證據", "保存議題、選項、意見及採納理由。", "議案附件、報價、評估原表。", "報告人／紀錄人。", "每項結論有可讀原件、版次與定位。"],
        ["決議／審查結論", "記錄結果、票數、理由、條件及結論。", "正式會議／審查原件。", "主持人確認／紀錄人。", "核對權限與有效性；未核准不得標已通過。"],
        ["後續責任", "追蹤交付、負責人、期限及覆核。", "正式決議、任務追蹤紀錄。", "被指派者／追蹤人。", "責任與期限明確；結案附可查證證據。"],
        ["原件位置", "定位受控卷宗、存取權限及保管者。", "正式檔案庫／紙本卷宗。", "檔案管理者。", "受權人可調閱；修訂保留舊版與軌跡。"],
    ], [0.95, 1.25, 1.35, 1.25, 2.2])
    heading(doc, "會前控制順序", 2)
    add_table(doc, ["順序", "檢查／停點"], [
        ["1 規範", "先確認大型採購門檻、決議機關、授權及會議有效條件；無法確認即暫停。"],
        ["2 議案", "確認需求、預算、選項、評分標準及附件已於決策前提供。"],
        ["3 參與", "確認召集、紀錄、評估角色及利益迴避安排；有衝突者按正式規則處理。"],
        ["4 會後", "核對決議與原件、附條件通過事項、採購／聘任授權及歸檔完成狀態。"],
    ], [1.0, 6.0])

    # Page 4: blank reusable record form, header and participants.
    page(doc)
    heading(doc, "03 可複製的正式紀錄表 空白範本")
    add_callout(doc, "填寫界線", "空白範本須由真實會議／審查當事人即時填寫；本頁空欄不是已完成佐證。不得預填姓名、出席、決議、簽名、核准或生效日期。", PALE_GOLD)
    add_table(doc, ["基本欄位", "正式作業填寫"], [
        ["正式案件名稱／案號", ""], ["紀錄編號／版次／形成日期", ""], ["案件類型（大型採購／師資遴選審查）", ""],
        ["適用制度／有效版本／條文", ""], ["大型採購門檻或審查權限來源", ""], ["召集／主持及紀錄責任", ""],
        ["日期／起迄時間／地點或會議方式", ""], ["會議目的及待決事項", ""], ["預算／候選師資／採購範圍（依案件適用）", ""],
    ], [2.75, 4.25])
    heading(doc, "出席與權限核對", 2)
    add_table(doc, ["姓名／正式角色", "出席時間／方式", "決議權限依據", "利益關係／迴避", "簽到或系統證據"], [
        ["", "", "", "", ""], ["", "", "", "", ""], ["", "", "", "", ""], ["", "", "", "", ""],
    ], [1.5, 1.25, 1.45, 1.45, 1.35])
    add_field_line(doc, "有效條件／出席數／法定或授權門檻：")
    add_field_line(doc, "迴避及其對議程、計票、評分之影響：")
    add_field_line(doc, "會議文件版本／會前送達證據：")
    add_p(doc, "簽名與確認僅可由本人於實際作業完成後填入；電子紀錄需保存簽核軌跡及時間戳記。", color="9B1C1C", bold=True)

    # Page 5: blank form agenda, deliberations, decision and custody.
    page(doc)
    heading(doc, "04 正式紀錄表續頁 議程決議追蹤")
    add_table(doc, ["議程／議題", "會前證據與版本", "討論、選項及理由", "決議／票數／條件", "責任／期限"], [
        ["", "", "", "", ""], ["", "", "", "", ""], ["", "", "", "", ""], ["", "", "", "", ""],
    ], [1.2, 1.45, 1.8, 1.55, 1.0])
    heading(doc, "決議與後續行動", 2)
    add_table(doc, ["決議類型", "通過／退回／附條件", "依據及理由", "條件完成前禁止事項", "責任與結案證據"], [
        ["大型採購", "", "", "", ""], ["師資遴選後審查", "", "", "", ""],
    ], [1.25, 1.25, 1.65, 1.55, 1.3])
    add_table(doc, ["附件編號", "附件名稱／版本／日期", "來源形成者", "原件位置／權限", "核對結果／缺件"], [
        ["", "", "", "", ""], ["", "", "", "", ""], ["", "", "", "", ""], ["", "", "", "", ""],
    ], [0.85, 1.8, 1.25, 1.75, 1.35])
    add_field_line(doc, "紀錄人確認（本人簽名／電子簽核）：")
    add_field_line(doc, "有權覆核／核准者及依據（本人簽核）：")
    add_field_line(doc, "更正版本、原因、時間及保留原件位置：")

    # Page 6: complete synthetic case A identity and agenda.
    page(doc)
    heading(doc, "05 完整示範 A 大型採購會議紀錄")
    add_callout(doc, "案例標識｜SAMPLE/SYNTHETIC｜非真實會議", "以下是自洽的紙上演練案例。金額、門檻、角色、出席、評估摘要、決議、日期及附件代碼全為虛構；沒有實際採購、簽約、核准、付款或本會決議。案例不與其他 DOCX 建立真實案件連結。", PALE_RED, "9B1C1C")
    add_table(doc, ["會議基本資料", "示範填寫值"], [
        ["案例／紀錄編號", "SYN-I10E03-A-001｜大型採購會議示範"],
        ["採購案", "志工服務紀錄管理課程與教材採購（虛構名稱）"],
        ["會議日期／時間／地點", "2026-09-12｜13:30–15:00｜SYN-ROOM-A（虛構場地代碼）"],
        ["需求與範圍", "供12名示範參訓者使用；2次課程、合計12小時；含教材與必要課後修訂（均為合成設定）。"],
        ["大型採購判定", "示範門檻新臺幣100,000元；本案例預估超過門檻。此門檻完全虛構，不是協會規定；REAL 使用必須引用正式採購辦法。"],
        ["預算上限", "新臺幣120,000元（SAMPLE；不代表核定預算）。"],
        ["規範／授權", "SYN-POL-A：虛構示範規則，假設至少3人出席且4名示範角色均無利益迴避；REAL 權限 UNKNOWN，須換正式章程／授權矩陣。"],
    ], [1.65, 5.35])
    heading(doc, "示範出席表", 2)
    add_table(doc, ["識別碼／角色（均虛構）", "出席", "權責功能示範", "利益關係聲明示範"], [
        ["SYN-R01｜訓練需求承辦", "全程", "提出需求、確認課程範圍", "無（示範記錄）"],
        ["SYN-R02｜課程審查角色", "全程", "核對課程設計及品質", "無（示範記錄）"],
        ["SYN-R03｜預算核對角色", "全程", "檢視預算與價格比較", "無（示範記錄）"],
        ["SYN-R04｜會議紀錄角色", "全程", "記錄討論、決議與附件", "無（示範記錄）"],
    ], [2.0, 0.75, 2.25, 2.0])
    add_p(doc, "出席4人；示範有效條件為3人。這個比例只是案例輸入，不能取代正式章程或授權規則。", bold=True, color=NAVY)

    # Page 7: case A deliberation and arithmetic.
    page(doc)
    heading(doc, "06 示範 A 議題選項算式與決議")
    add_table(doc, ["議題", "會前資料／來源", "討論與理由（示範）", "處理"], [
        ["需求及課程範圍", "SYN-A-ATT-01：虛構需求摘要；SYN-A-ATT-02：虛構課程規格", "限定為志工服務紀錄管理；交付2次／12小時；不把個案服務或醫療結果列為課程成效。", "範圍清楚，列入比較"],
        ["候選與價格", "SYN-A-ATT-03：虛構候選比較摘要", "A 118,000元／82分；B 113,000元／88分；C 127,000元／77分。此表只記會議採用摘要，不代替單案評選表或真實報價。", "B得分最高且在示範預算內"],
        ["預算與風險", "SYN-A-ATT-04：虛構預算／利益迴避摘要", "核對含稅、教材、修訂與交付範圍；真實採購須核驗正式預算及利益關係。", "附真實核定預算與迴避文件為前置條件"],
        ["決策路徑", "SYN-POL-A（虛構規則）", "4名示範出席者高於案例假設門檻3人；無人迴避。正式有效性在真實案件仍須依正式制度判斷。", "示範會議具備案例內決策條件"],
    ], [1.15, 1.85, 2.8, 1.2])
    heading(doc, "示範金額核對", 2)
    add_table(doc, ["核對項", "算式", "結果", "解讀界線"], [
        ["候選B低於上限", "120,000 − 113,000", "7,000元", "僅是示範預算餘額；不等於核准或可動支。"],
        ["預算使用率", "113,000 ÷ 120,000 × 100%", "94.166…% → 94.2%", "四捨五入至小數1位；分母為示範預算上限。"],
        ["會議紀錄完整度", "已填10項 ÷ 必填10項 × 100%", "100%", "只量測欄位完整度，不證明決議合法、附件真實或程序合規。"],
    ], [1.35, 2.0, 1.15, 2.5])
    add_callout(doc, "示範決議（不是本會核准）", "SAMPLE：4名示範出席者均記為贊成，反對0、迴避0；同意依附加條件續辦候選B之採購程序。條件：取得正式預算與授權、核對原始報價和含稅範圍、確認利益迴避、由具權限者完成正式核准；在條件完成前不得下單、簽約或付款。此為虛構案例結果。", PALE_GOLD)
    page(doc)
    heading(doc, "A 案例後續示範任務", 2)
    add_table(doc, ["後續示範任務", "責任角色", "期限／完成證據"], [
        ["核對門檻與採購授權", "SYN-R02", "會後示範節點；正式規範及版本索引"],
        ["覆核含稅報價及預算", "SYN-R03", "會後示範節點；原始報價及正式預算"],
        ["整理紀錄與附件索引", "SYN-R04", "會後示範節點；保留原件、簽核及版次"],
    ], [2.6, 1.5, 2.9])

    # Page 8: completed instructor review record.
    heading(doc, "07 完整示範 B 師資遴選後審查")
    add_callout(doc, "案例標識｜SAMPLE/SYNTHETIC｜與案例A分案", "本案例獨立演示「先依事前標準評估候選人，再由有權會議／審查通過並留下紀錄」。人選、履歷、證照、評分、出席、日期及結論均為虛構；未實際聘任或核驗任何真人資格。", PALE_RED, "9B1C1C")
    add_table(doc, ["欄位", "示範填寫值"], [
        ["案例／紀錄編號", "SYN-I10E03-B-001｜師資遴選後審查示範"],
        ["審查時間／地點", "2026-09-19｜10:00–11:00｜SYN-ROOM-B（虛構場地代碼）"],
        ["審查目的", "比較兩名虛構候選人之課程專業、教學設計、可授課性、資料保護及費用適配。"],
        ["事前標準", "專業適配35%、教學示範25%、時段可用15%、費用適配15%、資料保護10%；合計100%。標準和人選皆為示範輸入。"],
        ["出席／有效性", "SYN-R11、SYN-R12、SYN-R13共3名示範審查角色；假設最低2人。正式規則及授權未被證實。"],
        ["候選人", "SYN-CAND-X、SYN-CAND-Y（虛構識別碼；不對應真人）。"],
    ], [1.6, 5.4])
    heading(doc, "示範審查問題與證據", 2)
    add_table(doc, ["審查面向", "會前應查資料", "本案例的示範記錄"], [
        ["資格與專業", "真實履歷、證照／經歷原件及有效性查核", "只以虛構摘要演示；未宣稱任何證照真實。"],
        ["教學設計", "課程目標、教案、試教紀錄及評量方式", "兩位示範候選人依同一題目及同一量尺比較。"],
        ["可授課及費用", "時段確認、授課條件、報價及授權", "需由承辦以正式文件核對，不以本頁替代報價或聘任。"],
        ["資料保護／迴避", "個資處理說明、利益關係聲明及限制", "示範皆記為無衝突；真實案件需本人揭露並留原件。"],
    ], [1.35, 2.75, 2.9])

    # Page 9: instructor score math and decision.
    page(doc)
    heading(doc, "08 示範 B 評分重算與審查結論")
    add_p(doc, "量尺為1至5分；加權百分分數＝Σ（單項分數 ÷ 5 × 該項權重）。評分標準須先核定、同一題目同一量尺；若權重或評分發生變更，保留原版與原因，不可在看到結果後改標準。")
    add_table(doc, ["面向／權重", "候選X原分", "X加權分", "候選Y原分", "Y加權分"], [
        ["專業適配／35%", "5", "5÷5×35=35", "4", "4÷5×35=28"],
        ["教學示範／25%", "4", "4÷5×25=20", "4", "4÷5×25=20"],
        ["時段可用／15%", "4", "4÷5×15=12", "5", "5÷5×15=15"],
        ["費用適配／15%", "4", "4÷5×15=12", "3", "3÷5×15=9"],
        ["資料保護／10%", "5", "5÷5×10=10", "4", "4÷5×10=8"],
        ["合計／100%", "—", "89分", "—", "80分"],
    ], [1.8, 0.9, 1.6, 0.9, 1.8])
    add_table(doc, ["決議欄", "示範記錄"], [
        ["決議／票數", "3名示範審查者中贊成3、反對0、迴避0；記錄為通過候選X優先推薦。"],
        ["理由", "X在專業適配、教學示範及資料保護得分較高；Y在可授課時段較佳。依事前權重，X總分89高於Y的80。"],
        ["條件", "只可在真實履歷、資格、利益迴避、費用與聘任權限逐項核驗後，由正式有權者決定是否聘任；本示範不代表已聘任。"],
        ["後續責任", "SYN-R11核驗候選資格；SYN-R12確認教案／試教原件；SYN-R13整理完整紀錄與版本。角色與任務均為虛構。"],
        ["完整度算例", "示範必填9項，已記9項：9÷9×100%=100%。欄位完整度不能替代資格查核、授權或真實審查。"],
    ], [1.35, 5.65])
    heading(doc, "師資審查案例的原件索引（均為虛構代碼）", 2)
    add_table(doc, ["示範附件", "內容／REAL時須替換為"], [
        ["SYN-B-ATT-01", "事前遴聘標準；REAL：核准版標準及日期。"],
        ["SYN-B-ATT-02", "候選X/Y履歷及資格摘要；REAL：本人提交且可驗證的原件。"],
        ["SYN-B-ATT-03", "同題試教及原始評分；REAL：評審本人評分與理由。"],
        ["SYN-B-ATT-04", "迴避及審查紀錄；REAL：簽到、議程、決議和正式簽核軌跡。"],
    ], [1.5, 5.5])

    # Page 10: verification, exceptions and correction.
    page(doc)
    heading(doc, "09 覆核清單計算及異常處理")
    add_table(doc, ["覆核點", "通過條件", "不符合時處理"], [
        ["要求類型", "明確區分大型採購會議與師資遴選後審查。", "拆成不同案卷或補正索引；不得只用一般會議名義。"],
        ["正式權限", "引用有效制度、權責及會議有效條件。", "停止標示「通過」；取得正式規範及有權者決定。"],
        ["事前標準", "議案、門檻、預算、評分標準在決策前形成並可追溯版次。", "退回補齊；不得事後倒填成事前標準。"],
        ["出席／利益", "真實出席者、參與方式、權限及迴避均可查。", "補查原件；無法證實即標示未證實，不補造簽名。"],
        ["討論／決議", "議題、證據、主要意見、結論、票數／審查結果、理由及條件一致。", "更正時保留原版本、更正時間、原因及責任。"],
        ["後續及原件", "採購／聘任前置條件、責任期限、附件版次、正式存放位置均明確。", "不得下單、簽約、付款或聘任，直到有權者完成真實核准。"],
    ], [1.2, 3.0, 2.8])
    heading(doc, "完整度計算示例與限制", 2)
    add_p(doc, "完整度 = 已填且可定位的必填欄數 ÷ 適用必填欄總數 × 100%。例如10項中8項可追溯，為8÷10×100%=80%。這只指出表單欄位完成比例；不代表會議有權、投票有效、附件真實或官方要求已被實際滿足。")
    add_p(doc, "示範清單：案例A的10項為案件／紀錄編號、案件類型、門檻與權限來源、通知／議程、出席、利益迴避、需求／預算、候選與附件、討論／決議／票數、責任／期限／原件位置；案例B的9項為案件／紀錄編號、事前標準、候選來源、原始評分、出席／權限、利益迴避、討論理由、結論／條件、責任／期限／附件位置。示範分別核為10/10與9/9；清單可依正式制度調整。")
    add_table(doc, ["退件／暫停條件", "為何不能判PASS", "必要處置"], [
        ["決策後才補造通知、門檻或評分規則", "無法證明程序在決策前成立。", "標示事後補正；按正式制度重新審查，不回填為當時事實。"],
        ["缺簽到、權限、有效性或利益迴避證據", "不能確認參與人或決議資格。", "調閱原件；仍缺者記為未證實並升級有權者。"],
        ["附件只有檔名／無法開啟或版本不一致", "無法支撐討論與決議。", "補真實可讀原件，記錄版次、日期、保管位置。"],
        ["把SAMPLE填入REAL卷宗或偽造簽章／第三方資料", "構成錯誤證據表示。", "停止使用、隔離示範資料、按正式更正與通報程序處理。"],
    ], [2.1, 2.4, 2.5])

    # Page 11: real-replacement and crosswalk; flow naturally to avoid a blank page.
    heading(doc, "10 正式使用前的替換與文件邊界")
    add_callout(doc, "先讀取真實規範", "目前提供的權威來源能證實官方要求文字與頁次；不能證實協會實際採購門檻、決議權限、會議規則、師資遴聘制度、案件、出席、簽核或結果。這些項目均須取得正式原件後再填。", PALE_GOLD)
    add_table(doc, ["REAL使用前必須取得", "最低可接受來源／處理"], [
        ["採購／會議授權", "現行章程、採購辦法、授權矩陣及有效版本；確認門檻、決議權與有效條件。"],
        ["案件需求及預算", "核准需求、課程計畫、預算來源、採購範圍及可比規格。"],
        ["會議成立與出席", "真實通知、議程、簽到／線上出席軌跡、迴避聲明及紀錄版本。"],
        ["評估與決議", "事前評分標準、逐位原始評分、討論、決議理由、票數與核准條件。"],
        ["師資資格／聘任", "本人履歷／證照查核、教案／試教、利益揭露、審查結果及正式聘任授權。"],
        ["歸檔與隱私", "受控保存位置、存取權限、個資遮蔽、保存期限與調閱紀錄。"],
    ], [2.0, 5.0])
    heading(doc, "官方要求追溯與負向鄰件", 2)
    add_table(doc, ["本件直接責任", "證明位置", "不可替代的鄰件"], [
        ["DOCX 0048｜I10-E03｜官方方框049／P03-B10", "本件§1、§5–9及原件附件索引；官方來源表第3頁。", "DOCX 0045｜I10-E01：程序SOP；不能代替本件實際會議紀錄。"],
        ["大型採購會議通過", "§5–6：權限、出席、議程、討論、票數、條件與附件。", "DOCX 0047｜I10-E02：歷年檢討與改善分析；不能代替單次決議。"],
        ["師資遴選後審查通過", "§7–8：事前標準、原始評分、理由、結論與資格核驗條件。", "DOCX 0049｜I11-E01：訓練課程計畫需求調查；不能證明師資審查。"],
    ], [2.0, 2.65, 2.35])
    add_p(doc, "遮住文件標題後，正文仍以會議／審查參與者、議程、討論證據、決議條件、票數與附件索引辨識為「會議通過／審查通過紀錄」。本件不取代採購SOP、跨年檢討報告、課程計畫或第三方原件。", color=NAVY)
    add_p(doc, "文件控制：R02 候選示範；本版未經協會核定或生效。真實紀錄應依正式文件管制程序編號、簽核、保存及調閱。", color=MID, size=9.5)

    # Page 12 onward: traceable synthetic register, split naturally across pages.
    page(doc)
    heading(doc, "附錄A 逐事實 SAMPLE/SYNTHETIC 替換清冊")
    add_p(doc, "本清冊逐項登錄兩個示範案例中的合成日期、角色、金額、規則、附件代碼、分數、票數、結論及計算。正式使用時，須以右欄所列可驗證原件替換，並由責任角色核對；本清冊不授權把示範數據當成協會資料。")
    register_rows = [[fid, value, f"{kind}\n位置：{location}", f"REAL替換來源：{real_source}\n責任：{owner}\n狀態：{status}"] for fid, value, kind, location, real_source, owner, status in facts]
    add_table(doc, ["合成事實ID", "示範值", "事實類型／正文定位", "真實替換與責任"], register_rows, [1.0, 1.55, 2.1, 2.35], font_size=TABLE_PT)
    heading(doc, "附錄B 官方來源與版本紀錄", 2)
    add_table(doc, ["項目", "值／處理"], [
        ["官方來源", "《TTQS查核佐證文件資料表－訓練機構版》，第3頁，P03-B10；PDF SHA-256 805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709。"],
        ["Unique Blueprint", "DOCX 0048｜I10-E03；Blueprint 原始檔及逐列 SHA 登錄於本包 inputs/SOURCE_MAP_R02.json。"],
        ["R02修訂紀錄", "依唯一需求新建可複製表單、兩個完整合成示例、計算、退件條件及逐事實替換清冊；不複製歷史 DOCX 正文。"],
        ["審查狀態", "候選文件；獨立審查與版面驗收結果只對應實際記錄的 exact SHA；未正式接受、未生效。"],
    ], [1.5, 5.5])
    add_p(doc, "真實紀錄核對：收到原件後逐項填寫紀錄編號、來源版次、形成日期、頁次／位置、保管者及存取限制；缺件時明確標示「未取得／未證實」，不可用空白或示範內容推定通過。", bold=True, color=NAVY)

    doc.core_properties.title = "大型採購會議與師資遴聘審查紀錄｜DOCX 0048｜I10-E03"
    doc.core_properties.subject = "官方方框049（P03-B10）專屬佐證紀錄；SAMPLE/SYNTHETIC候選"
    doc.core_properties.author = ""
    doc.core_properties.comments = "Candidate sample only; not an approved Association record."
    DOCX_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_PATH)
    return doc


def main():
    ROOT.mkdir(parents=True, exist_ok=True)
    INPUTS.mkdir(exist_ok=True)
    EVIDENCE.mkdir(exist_ok=True)
    (ROOT / "review").mkdir(exist_ok=True)
    blueprint = load_blueprint()
    facts = build_register()
    make_doc(blueprint, facts)

    # A machine-readable companion register uses the exact rows embedded in Appendix A.
    reg_path = EVIDENCE / "SYNTHETIC_REPLACEMENT_REGISTER_R02.csv"
    with reg_path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["fact_id", "synthetic_value", "fact_type", "document_location", "real_replacement_source", "responsible_role", "status"])
        w.writerows(facts)

    source_bytes = SOURCE_PDF.read_bytes()
    bp_bytes = BLUEPRINT_SOURCE.read_bytes()
    bp_line = next(line for line in bp_bytes.splitlines() if json.loads(line).get("deliverable_id") == "0048")
    inventory = json.loads(CHECKBOX_INVENTORY.read_text(encoding="utf-8"))
    boxes = inventory.get("items", inventory if isinstance(inventory, list) else [])
    selected = next((x for x in boxes if x.get("source_box_id") == "P03-B10"), None)
    # Some inventory versions wrap their records under a different key; exact source box is asserted below.
    if selected is None and isinstance(inventory, dict):
        for value in inventory.values():
            if isinstance(value, list):
                selected = next((x for x in value if isinstance(x, dict) and x.get("source_box_id") == "P03-B10"), None)
                if selected:
                    break
    if selected is None:
        raise RuntimeError("Official checkbox inventory does not contain P03-B10")
    if selected.get("extracted_block", "").replace("\n", "").replace(" ", "") != blueprint["official_requirement_text"].replace(" ", ""):
        raise RuntimeError("Official checkbox P03-B10 text does not match the unique Blueprint requirement")
    (INPUTS / "UNIQUE_BLUEPRINT_0048.jsonl").write_bytes(bp_line + b"\n")
    source_map = {
        "project_id": "TTQS_ONE",
        "doc_id": "0048",
        "requirement_id": "I10-E03",
        "official_checkbox_sequence": "049",
        "official_source_box_id": "P03-B10",
        "official_source_page": 3,
        "official_requirement_text": blueprint["official_requirement_text"],
        "official_source_pdf_sha256": sha256(source_bytes),
        "official_inventory_file_sha256": sha256(CHECKBOX_INVENTORY.read_bytes()),
        "blueprint_source_path": str(BLUEPRINT_SOURCE.relative_to(REPO)).replace("\\", "/"),
        "blueprint_source_sha256": sha256(bp_bytes),
        "blueprint_row_sha256": sha256(bp_line),
        "blueprint_unique_sibling_set": blueprint.get("requirement_deliverable_siblings"),
        "document_genre": blueprint["document_genre"],
        "body_template_policy": blueprint["docx_body_template_policy"],
        "old_body_reuse": "PROHIBITED",
        "sample_boundary": "Every event, role, date, money amount, vote, criterion, score, decision and attachment identifier in both cases is SAMPLE/SYNTHETIC; no real association record or third-party document is asserted.",
        "official_source_check": "P03-B10 inventory text exactly matches the unique Blueprint requirement after whitespace normalization.",
    }
    (INPUTS / "SOURCE_MAP_R02.json").write_text(json.dumps(source_map, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    (INPUTS / "SOURCE_MAP_R02.md").write_text(
        "# DOCX 0048 R02 source map\n\n"
        "- DOCX 0048 maps to I10-E03 and official checkbox 049, source box P03-B10, page 3.\n"
        f"- Exact official text: {blueprint['official_requirement_text']}\n"
        f"- Official PDF SHA-256: `{source_map['official_source_pdf_sha256']}`.\n"
        f"- Unique Blueprint file SHA-256: `{source_map['blueprint_source_sha256']}`.\n"
        f"- Selected unique row SHA-256: `{source_map['blueprint_row_sha256']}`.\n"
        "- The Blueprint requires an evidence record with meeting/review source, filled sample, validation, exception handling, real-source replacement, and traceability. Historical body reuse is prohibited.\n"
        "- No real meeting, procurement, instructor, approval, signature, qualification, payment, or third-party original was available in the authorized source package. The complete worked examples are therefore labelled SAMPLE/SYNTHETIC throughout.\n",
        encoding="utf-8",
        newline="\n",
    )
    receipt = {
        "project_id": "TTQS_ONE",
        "doc_id": "0048",
        "requirement_id": "I10-E03",
        "official_checkbox_sequence": "049",
        "candidate_revision": "R02",
        "candidate_path": "candidate/0048_R02.docx",
        "candidate_sha256": sha256(DOCX_PATH.read_bytes()),
        "candidate_bytes": DOCX_PATH.stat().st_size,
        "source_map_sha256": sha256((INPUTS / "SOURCE_MAP_R02.json").read_bytes()),
        "blueprint_row_sha256": source_map["blueprint_row_sha256"],
        "replacement_register_sha256": sha256(reg_path.read_bytes()),
        "synthetic_fact_count": len(facts),
        "status": "CANDIDATE_BUILD_COMPLETE_PENDING_TECHNICAL_VISUAL_AND_INDEPENDENT_SEMANTIC_REVIEW",
        "formal_acceptance": "NOT_GRANTED",
        "built_at_taipei": datetime.now().astimezone().isoformat(timespec="seconds"),
        "builder": "Codex local Python-docx builder; source-only from current Unique Blueprint and official source map",
    }
    (EVIDENCE / "BUILD_RECEIPT_R02.json").write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    print(json.dumps(receipt, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
