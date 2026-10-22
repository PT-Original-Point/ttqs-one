from __future__ import annotations
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile
import hashlib, json, re, sys
sys.stdout.reconfigure(encoding='utf-8')
from PIL import Image
from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / 'candidate' / '0106_I18-E14_R01.docx'
SNAP = ROOT / 'review' / '0106_R01_exact_docx_text_snapshot.md'
OUT = ROOT / 'evidence' / 'technical_qa_receipt_0106_R01.json'

def sha(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def clean(value: str) -> str:
    return ' '.join(value.split()).replace('−', '-')

def num(value: str) -> float:
    value = clean(value).replace(',', '').replace('NT$', '').replace('%', '').strip()
    value = value.replace('−', '-')
    match = re.search(r'-?\d+(?:\.\d+)?', value)
    if not match:
        raise ValueError(f'No number in {value!r}')
    return float(match.group(0))

def table_rows(table: Table) -> list[list[str]]:
    return [[clean(cell.text) for cell in row.cells] for row in table.rows]

def run_size(run, paragraph) -> float:
    size = run.font.size
    if size is None and paragraph.style is not None:
        size = paragraph.style.font.size
    if size is None:
        size = doc.styles['Normal'].font.size
    return size.pt if size is not None else 10.5

if not DOCX.is_file():
    raise SystemExit(f'Missing candidate: {DOCX}')

doc = Document(str(DOCX))
body_lines: list[str] = []
table_records: list[tuple[int, list[list[str]]]] = []
for child in doc._element.body.iterchildren():
    if child.tag == qn('w:p'):
        paragraph = Paragraph(child, doc)
        if paragraph.text.strip():
            body_lines.append(paragraph.text.strip())
    elif child.tag == qn('w:tbl'):
        table = Table(child, doc)
        rows = table_rows(table)
        table_records.append((len(table_records), rows))
        body_lines.append('')
        for row in rows:
            body_lines.append(' | '.join(row))
        body_lines.append('')
for index, section in enumerate(doc.sections):
    body_lines.append(f'[HEADER {index}]')
    body_lines.extend(clean(p.text) for p in section.header.paragraphs if p.text.strip())
    body_lines.append(f'[FOOTER {index}]')
    body_lines.extend(clean(p.text) for p in section.footer.paragraphs if p.text.strip())
snapshot = '\n'.join(body_lines) + '\n'
SNAP.parent.mkdir(parents=True, exist_ok=True)
SNAP.write_text(snapshot, encoding='utf-8', newline='\n')
alltext = snapshot
allrows = [row for _, rows in table_records for row in rows]

# Isolate the requirement-specific data tables from captions and other tabular tools.
t03_candidates = [rows for _, rows in table_records if rows and rows[0][0] == '序號' and any('後期標準化值' in c for c in rows[0])]
chart_candidates = [rows for _, rows in table_records if rows and rows[0][0] == '類別' and any('數值（NT$）' in c for c in rows[0])]
survey_candidates = [rows for _, rows in table_records if rows and rows[0][0] == 'ID' and any(c == 'Q1' for c in rows[0])]
t13_candidates = [rows for _, rows in table_records if rows and rows[0][0] == '#' and any('本稿示範觀察（非驗收）' in c for c in rows[0])]
q5_rows = [row for row in allrows if row and row[0].startswith('Q5 價值高於成本')]
c1_rows = [row for row in allrows if row and row[0] == 'C1']
fact_view = [row for row in allrows if any('F0106-VIEW-01' in c for c in row)]

checks: dict[str, bool] = {}
checks['zip_integrity'] = False
checks['two_embedded_images'] = len(doc.inline_shapes) == 2
media_details: list[dict[str, object]] = []
try:
    with ZipFile(DOCX) as archive:
        checks['zip_integrity'] = archive.testzip() is None
        media = [name for name in archive.namelist() if name.startswith('word/media/')]
        for name in media:
            with Image.open(BytesIO(archive.read(name))) as image:
                media_details.append({'name': name, 'width': image.width, 'height': image.height, 'nonblank_colors': len(image.convert('RGB').getcolors(maxcolors=1_000_000) or [])})
    checks['zip_media_count'] = len(media_details) == 2
    checks['media_nonblank'] = len(media_details) == 2 and all(int(item['width']) > 500 and int(item['height']) > 300 and int(item['nonblank_colors']) > 20 for item in media_details)
except Exception as exc:
    checks['zip_integrity'] = False
    checks['zip_error'] = str(exc)

body_sizes: list[float] = []
table_sizes: list[float] = []
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        if run.text.strip():
            body_sizes.append(run_size(run, paragraph))
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    if run.text.strip():
                        table_sizes.append(run_size(run, paragraph))
checks['body_font_floor_10_5pt'] = bool(body_sizes) and min(body_sizes) >= 10.5
checks['table_font_floor_9_5pt'] = bool(table_sizes) and min(table_sizes) >= 9.5
checks['a4_landscape'] = all(abs(s.page_width.inches - 11.69) < .03 and abs(s.page_height.inches - 8.27) < .03 for s in doc.sections)
checks['page_field_each_footer'] = bool(doc.sections) and all('PAGE' in section.footer._element.xml for section in doc.sections)
checks['sample_header_each_section'] = bool(doc.sections) and all('SAMPLE / SYNTHETIC' in ' '.join(p.text for p in section.header.paragraphs) and 'REAL=UNKNOWN/HOLD' in ' '.join(p.text for p in section.header.paragraphs) for section in doc.sections)

checks['sample_banner_and_real_hold'] = 'SAMPLE／SYNTHETIC' in alltext and 'REAL=UNKNOWN/HOLD' in alltext
checks['reader_body_has_no_builder_figure_specs'] = all(term not in alltext for term in ['供 DOCX builder', '圖 1 規格：', '圖 2 規格：'])
checks['t09_caption_and_mapping'] = '表 T09｜訓練至財務結果邏輯鏈與證據缺口' in alltext and '8. 財務結果與感知結果分開分析及邏輯鏈（T09）' in alltext
checks['t13_caption'] = '表 T13｜交付前判定（SAMPLE／SYNTHETIC 示範）' in alltext
checks['cross_reference_points_to_both_sections'] = '兩線只在第 8 章並列，並於第 9 章標示歸因邊界' in alltext
checks['noncash_values_unquantified'] = all(term not in alltext for term in ['7,200', '1,800', '24 小時 × 示範費率']) and '志工與學員時間均未量化／不計入' in alltext
checks['break_even_value_explained'] = '1.0217391（＝94,000÷92,000；理論臨界值）' in alltext and '高於基期 1.00' in alltext and '高約 2.17%' in alltext and '位於 0.96–1.00 之外' in alltext and '47/46' not in alltext
checks['q5_count_reconciled'] = len(q5_rows) == 1 and '未達 4 分合計 5/10' in ' | '.join(q5_rows[0]) and '3 分 3 人、2 分 2 人' in ' | '.join(q5_rows[0])
checks['c1_specific_scale_anchors'] = len(c1_rows) == 1 and '依本案例所列資料' in ' | '.join(c1_rows[0]) and 'C1 專用計分錨點' in alltext
checks['fact_key_view_unique'] = len(fact_view) == 2 and sum(row[0].strip() == 'F0106-VIEW-01' for row in fact_view) == 1 and 'F0106-VIEW-01/VIEW-01' not in alltext
checks['traditional_chinese_heading'] = '本件不做什麼（需求界線）' in alltext and '本件不做什么（需求界線）' not in alltext
checks['no_internal_hash_or_control_leak'] = re.search(r'(?i)\b[0-9a-f]{40,64}\b|git blob|\.jsonl|source_packet|unique_blueprint', alltext) is None
checks['t13_not_a_self_acceptance_receipt'] = len(t13_candidates) == 1 and len(t13_candidates[0]) == 16 and all('PASS' not in t13_candidates[0][n][3] for n in range(1, 15)) and all('UNKNOWN' in t13_candidates[0][n][3] for n in (7, 8, 11, 12, 13, 14)) and 'FAIL → HOLD' in t13_candidates[0][15][3]

chart_values: dict[str, str] = {}
if len(chart_candidates) == 1:
    chart_rows = chart_candidates[0]
    chart_values = {row[0]: row[1] for row in chart_rows[1:] if len(row) >= 2}
expected_chart = {'前期支出': '100,000', '後期名目支出': '94,000', '後期標準化支出': '95,918.37', '直接訓練成本': '8,000', '估算淨財務效益': '-3,918.37'}
checks['reader_facing_figure1_data_table'] = len(chart_candidates) == 1 and all(label in chart_values and expected in chart_values[label].replace('−', '-') for label, expected in expected_chart.items())

survey_totals: list[int] = []
survey_high: list[int] = []
survey_low: list[int] = []
survey_total = 0
if len(survey_candidates) == 1:
    survey_rows = survey_candidates[0]
    detail = [row for row in survey_rows[1:] if row and re.fullmatch(r'SYN-L\d{2}', row[0])]
    scores = [[int(num(cell)) for cell in row[1:7]] for row in detail]
    if len(scores) == 10:
        survey_totals = [sum(row[j] for row in scores) for j in range(6)]
        survey_high = [sum(row[j] >= 4 for row in scores) for j in range(6)]
        survey_low = [sum(row[j] <= 2 for row in scores) for j in range(6)]
        survey_total = sum(map(sum, scores))
checks['survey_rows_and_sums'] = survey_totals == [39, 37, 38, 40, 33, 36] and survey_high == [8, 6, 7, 8, 5, 6] and survey_low == [1, 1, 1, 0, 2, 1] and survey_total == 223

finance = {'pre': 100000, 'post_nominal': 94000, 'price_index': 0.98, 'direct_cost': 8000}
finance['post_normalized'] = round(finance['post_nominal'] / finance['price_index'], 2)
finance['gross'] = round(finance['pre'] - finance['post_normalized'], 2)
finance['net'] = round(finance['gross'] - finance['direct_cost'], 2)
finance['roi_percent'] = round(finance['net'] / finance['direct_cost'] * 100, 2)
checks['core_finance_math'] = finance == {'pre':100000,'post_nominal':94000,'price_index':0.98,'direct_cost':8000,'post_normalized':95918.37,'gross':4081.63,'net':-3918.37,'roi_percent':-48.98}

t03_rows: list[list[str]] = []
if len(t03_candidates) == 1:
    t03_rows = t03_candidates[0]
detail_rows = [row for row in t03_rows[1:] if row and row[0].isdigit()]
t03_ok = len(detail_rows) == 10
for row in detail_rows:
    if len(row) < 9:
        t03_ok = False
        continue
    pre, post, nominal, index, normalized, standard = num(row[2]), num(row[3]), num(row[4]), num(row[5]), num(row[6]), num(row[7])
    t03_ok = t03_ok and round(pre - post, 2) == nominal and round(post / index, 2) == normalized and round(pre - normalized, 2) == standard
if detail_rows:
    t03_ok = t03_ok and round(sum(num(row[2]) for row in detail_rows), 2) == 100000 and round(sum(num(row[3]) for row in detail_rows), 2) == 94000
checks['t03_all_10_pairs_recomputed'] = t03_ok

checks['expected_table_inventory'] = len(doc.tables) == 25
checks['document_title_and_scope'] = 'I18-E14' in alltext and '文件 ID | 0106' in alltext and '官方來源' in alltext and '官方資料表印製的核取欄序號不同' in alltext
failed = [name for name, value in checks.items() if not value]

OUT.parent.mkdir(parents=True, exist_ok=True)
receipt = {
    'doc_id': '0106', 'requirement_id': 'I18-E14', 'project_157_index': 115,
    'candidate_state': 'CANDIDATE_NOT_FORMALLY_ACCEPTED',
    'candidate_sha256': sha(DOCX), 'candidate_bytes': DOCX.stat().st_size,
    'text_snapshot_sha256': sha(SNAP), 'text_snapshot_bytes': SNAP.stat().st_size,
    'paragraph_count': len(doc.paragraphs), 'table_count': len(doc.tables),
    'table_rows': sum(len(t.rows) for t in doc.tables), 'embedded_images': len(doc.inline_shapes),
    'media_details': media_details,
    'page_inches': [[round(s.page_width.inches, 2), round(s.page_height.inches, 2)] for s in doc.sections],
    'normal_pt': doc.styles['Normal'].font.size.pt if doc.styles['Normal'].font.size else None,
    'body_min_pt': min(body_sizes) if body_sizes else None,
    'table_min_pt': min(table_sizes) if table_sizes else None,
    't03_pair_rows_recomputed': len(detail_rows),
    'figure1_data': chart_values,
    'survey_recalculation': {'item_sums': survey_totals, 'high_4_5_counts': survey_high, 'low_1_2_counts': survey_low, 'total': survey_total, 'respondents': 10 if survey_total else 0},
    'finance_recalculation': finance,
    'checks': checks, 'technical_verdict': 'PASS' if not failed else 'FAIL', 'failed_checks': failed,
    'visual_QA': 'PENDING_NATIVE_WORD_EXPORT_AND_PAGE_INSPECTION',
    'semantic_QA': 'PENDING_FRESH_INDEPENDENT_REVIEW_ON_THIS_EXACT_SHA',
    'remote_readback': 'PENDING'
}
OUT.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
print(json.dumps({'candidate_sha256': receipt['candidate_sha256'], 'snapshot_sha256': receipt['text_snapshot_sha256'], 'technical_verdict': receipt['technical_verdict'], 'failed': failed, 'table_count': receipt['table_count'], 'table_rows': receipt['table_rows'], 'images': receipt['embedded_images'], 'body_min_pt': receipt['body_min_pt'], 'table_min_pt': receipt['table_min_pt'], 'checks': checks}, ensure_ascii=False, indent=2))
if failed:
    sys.exit(1)
