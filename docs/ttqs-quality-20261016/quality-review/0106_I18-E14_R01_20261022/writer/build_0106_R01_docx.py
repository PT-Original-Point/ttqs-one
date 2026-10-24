from __future__ import annotations
import hashlib, json, os, subprocess, tempfile
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
WRITER=ROOT/'writer'; VISUAL=ROOT/'visual'/'0106_R01'
CANDIDATE=ROOT/'candidate'/'0106_I18-E14_R01.docx'
PATCHED=WRITER/'opencode_writer_0106_R01_patched.md'
BUILD_MD=WRITER/'0106_R01_build_input.md'
REFERENCE=ROOT.parent/'0105_I18-E13_R01_20261022'/'candidate'/'0105_I18-E13_績效改善調查與課後行動計畫追蹤套件_R04.docx'
PANDOC=Path(r'C:\Users\J\AppData\Local\Pandoc\pandoc.exe')
FONT=Path(r'C:\Windows\Fonts\msjh.ttc'); BOLDFONT=Path(r'C:\Windows\Fonts\msjhbd.ttc')

def sha(path):
    h=hashlib.sha256()
    with Path(path).open('rb') as f:
        for chunk in iter(lambda:f.read(1048576),b''): h.update(chunk)
    return h.hexdigest()

def font(size,bold=False): return ImageFont.truetype(str(BOLDFONT if bold else FONT),size=size)
def centered(draw,xy,text,fnt,fill,anchor='mm'): draw.text(xy,text,font=fnt,fill=fill,anchor=anchor)

def draw_chart1(path):
    W,H=1800,1000; im=Image.new('RGB',(W,H),'white'); d=ImageDraw.Draw(im)
    title=font(48,True); section=font(32,True); label=font(27); small=font(22); value=font(28,True)
    centered(d,(W//2,55),'訓練前後支出與訓練成本拆解',title,'#18364A')
    centered(d,(W//2,115),'SAMPLE / SYNTHETIC — 非真實財務紀錄',section,'#9B2C2C')
    d.rounded_rectangle((45,160,900,850),radius=18,outline='#A9BAC6',width=3,fill='#FBFCFD')
    d.text((75,180),'A｜前後支出（每期 6 個月；NT$）',font=section,fill='#18364A')
    x0,x1,yt,yb=125,855,285,680; maxv=110000
    for v in range(0,110001,20000):
        y=yb-int((yb-yt)*v/maxv); d.line((x0,y,x1,y),fill='#D9E1E6',width=2)
        d.text((x0-15,y),f'{v:,}',font=small,fill='#56636B',anchor='rm')
    spend=[('前期支出',100000,'#4D7892'),('後期名目',94000,'#75A7A0'),('後期標準化',95918.37,'#3F8F84')]
    for (name,v,color),cx in zip(spend,[300,500,700]):
        h=int((yb-yt)*v/maxv); d.rounded_rectangle((cx-55,yb-h,cx+55,yb),radius=8,fill=color)
        centered(d,(cx,yb-h-28),f'{v:,.0f}',value,'#18364A'); centered(d,(cx,yb+34),name,label,'#263746')
    d.text((75,795),'10 組配對採購籃；價格指數 0.98；名目與標準化金額分列。',font=small,fill='#344955')
    d.rounded_rectangle((925,160,1755,850),radius=18,outline='#A9BAC6',width=3,fill='#FBFCFD')
    d.text((955,180),'B｜毛差額扣除直接訓練成本（NT$）',font=section,fill='#18364A')
    xs=[1110,1340,1570]; base=470; scale=.025; bw=118
    d.line((990,base,1705,base),fill='#4B5961',width=3); d.text((990,base-12),'0',font=small,fill='#56636B',anchor='rm')
    items=[('標準化毛差額',4081.63,'#3F8F84'),('訓練成本',-8000,'#C45A56'),('淨財務效益',-3918.37,'#D17B36')]
    for (name,v,color),cx in zip(items,xs):
        h=int(abs(v)*scale)
        if v>=0:
            y=base-h; d.rounded_rectangle((cx-bw//2,y,cx+bw//2,base),radius=7,fill=color); vy=y-28
        else:
            y=base+h; d.rounded_rectangle((cx-bw//2,base,cx+bw//2,y),radius=7,fill=color); vy=y+26
        centered(d,(cx,vy),f'{v:+,.2f}',value,'#18364A'); centered(d,(cx,760),name,label,'#263746')
    d.text((965,805),'4,081.63 - 8,000 = -3,918.37；負值向下呈現，不作因果宣稱。',font=small,fill='#344955')
    d.text((60,900),'SYNTHETIC｜配對數 n=10｜期間前後各 6 個月｜訓練成本 NT$8,000｜不代表原始點協會實績',font=small,fill='#6B3030')
    im.save(path,format='PNG',optimize=True,dpi=(200,200))

def draw_chart2(path):
    W,H=2200,820; im=Image.new('RGB',(W,H),'white'); d=ImageDraw.Draw(im)
    title=font(46,True); node_title=font(31,True); body=font(25); badge=font(23,True); foot=font(22)
    centered(d,(W//2,48),'訓練至財務結果的邏輯鏈與證據強度',title,'#18364A')
    centered(d,(W//2,105),'SAMPLE / SYNTHETIC — 全部數值為完整示範假設，不是協會實績',badge,'#9B2C2C')
    nodes=[
      ('① 訓練投入','弱｜合成名冊',['10 / 12 人','F0106-TRN-01','名冊與課程資料','真實資料待補 (HOLD)']),
      ('② 學習感知','弱｜自陳問卷',['n=10；均值 3.3–4.0','F0106-SRV-01～10','學習自陳，不是測驗','缺前後測資料']),
      ('③ 採購行為','弱｜示範配對',['10 組同規格採購籃','F0106-FIN-01～10','請購／驗收／比價單','真實資料待補 (HOLD)']),
      ('④ 成本產出','中｜僅可重算',['100,000 → 94,000','F0106-IDX-01','指數 0.98 (SYN)','真實價格依據未知']),
      ('⑤ 財務結果','中｜算術結果',['毛差額 +4,081.63','成本 NT$8,000','淨額 -3,918.37','F0106-CALC-08','全為示範，非實績']),
      ('⑥ 因果歸因','無｜不得推論',['證據不足','F0106-CONF-01～12','混淆因素未排除','不得宣稱訓練造成','因果結論 UNKNOWN']),
    ]
    colors=[('#FFF7E3','#B28A3C'),('#FFF7E3','#B28A3C'),('#FFF7E3','#B28A3C'),('#EAF2F8','#557F9B'),('#EAF2F8','#557F9B'),('#FCEAEA','#B54B4B')]
    left=40; gap=24; nw=333; top=210; bottom=630
    for i,(node,strength,body_lines) in enumerate(nodes):
        x=left+i*(nw+gap); bg,border=colors[i]
        d.rounded_rectangle((x,top,x+nw,bottom),radius=18,fill=bg,outline=border,width=4)
        centered(d,(x+nw//2,top+42),node,node_title,'#18364A'); centered(d,(x+nw//2,top+92),strength,badge,border)
        y=top+145
        for line in body_lines:
            measured=d.textlength(line,font=body)
            if measured>nw-40: raise ValueError(f'Figure 2 text overflow ({measured:.1f}px): {line}')
            d.text((x+20,y),line,font=body,fill='#273A45'); y+=34
        if i<5:
            cy=(top+bottom)//2; x1=x+nw+2; x2=x+nw+gap-6
            d.line((x1,cy,x2,cy),fill='#52636C',width=5); d.polygon([(x2,cy),(x2-12,cy-10),(x2-12,cy+10)],fill='#52636C')
    centered(d,(W//2,700),'期間：前後各 6 個月｜NT$｜配對採購 n=10｜學員問卷 n=10｜目標客戶 n=1｜結果可重算但不證明因果',foot,'#344955')
    centered(d,(W//2,752),'⑤ 淨額為負；⑥ 因果證據不足。真實財務結論與所有 REAL 欄位均 UNKNOWN / HOLD。',foot,'#8A3030')
    im.save(path,format='PNG',optimize=True,dpi=(200,200))

def clear_story(doc):
    body=doc._element.body
    for child in list(body):
        if child.tag!=qn('w:sectPr'): body.remove(child)
def clear_part(part):
    for child in list(part._element): part._element.remove(child)
def east_asia_style(style,name='Microsoft JhengHei'):
    rpr=style._element.get_or_add_rPr(); rf=rpr.rFonts
    if rf is None: rf=OxmlElement('w:rFonts'); rpr.insert(0,rf)
    for key in ('ascii','hAnsi','eastAsia'): rf.set(qn('w:'+key),name)
def east_asia_run(run,name='Microsoft JhengHei'):
    run.font.name=name; rpr=run._element.get_or_add_rPr(); rf=rpr.rFonts
    if rf is None: rf=OxmlElement('w:rFonts'); rpr.insert(0,rf)
    for key in ('ascii','hAnsi','eastAsia'): rf.set(qn('w:'+key),name)
def add_page_field(p):
    r=p.add_run(); begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin')
    instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text=' PAGE '
    sep=OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'),'separate'); val=OxmlElement('w:t'); val.text='1'
    end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end')
    for el in (begin,instr,sep,val,end): r._r.append(el)
    east_asia_run(r); r.font.size=Pt(8); r.font.color.rgb=RGBColor(95,105,112)
def add_running_matter(doc):
    for sec in doc.sections:
        clear_part(sec.header); clear_part(sec.footer); sec.different_first_page_header_footer=False
        hp=sec.header.add_paragraph(); hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        r=hp.add_run('0106｜I18-E14｜SAMPLE / SYNTHETIC｜REAL=UNKNOWN/HOLD'); east_asia_run(r); r.font.size=Pt(8); r.font.color.rgb=RGBColor(95,105,112)
        fp=sec.footer.add_paragraph(); fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=fp.add_run('原始點協會 TTQS ONE｜候選稿，未正式接受｜第 '); east_asia_run(r); r.font.size=Pt(8); r.font.color.rgb=RGBColor(95,105,112)
        add_page_field(fp); r=fp.add_run(' 頁'); east_asia_run(r); r.font.size=Pt(8); r.font.color.rgb=RGBColor(95,105,112)
def postprocess(path):
    doc=Document(path)
    for sec in doc.sections:
        sec.orientation=WD_ORIENT.LANDSCAPE; sec.page_width=Inches(11.69); sec.page_height=Inches(8.27)
        sec.left_margin=Inches(.62); sec.right_margin=Inches(.62); sec.top_margin=Inches(.58); sec.bottom_margin=Inches(.58)
        sec.header_distance=Inches(.25); sec.footer_distance=Inches(.25)
    for name,size in [('Normal',10.5),('Title',22),('Heading 1',16),('Heading 2',13),('Heading 3',11.5)]:
        if name in doc.styles:
            st=doc.styles[name]; st.font.name='Microsoft JhengHei'; st.font.size=Pt(size); st.font.color.rgb=RGBColor(24,54,74); east_asia_style(st)
            if name=='Normal': st.paragraph_format.space_after=Pt(5); st.paragraph_format.line_spacing=1.12
    for table in doc.tables:
        try: table.style='Table Grid'
        except Exception: pass
        table.autofit=True
        for ri,row in enumerate(table.rows):
            trPr=row._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))
            if ri==0:
                repeat=OxmlElement('w:tblHeader'); repeat.set(qn('w:val'),'true'); trPr.append(repeat)
            for cell in row.cells:
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcPr=cell._tc.get_or_add_tcPr(); mar=tcPr.find(qn('w:tcMar'))
                if mar is None: mar=OxmlElement('w:tcMar'); tcPr.append(mar)
                for side,val in [('top','55'),('left','70'),('bottom','55'),('right','70')]:
                    el=mar.find(qn('w:'+side))
                    if el is None: el=OxmlElement('w:'+side); mar.append(el)
                    el.set(qn('w:w'),val); el.set(qn('w:type'),'dxa')
                if ri==0:
                    shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'DCE8EF'); cell._tc.get_or_add_tcPr().append(shd)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(1); p.paragraph_format.space_before=Pt(0); p.paragraph_format.line_spacing=1.0
                    for r in p.runs:
                        east_asia_run(r); r.font.size=Pt(9.5)
                        if ri==0: r.bold=True; r.font.color.rgb=RGBColor(24,54,74)
    for p in doc.paragraphs:
        if p.style and p.style.name.startswith('Heading'): p.paragraph_format.keep_with_next=True
        for r in p.runs:
            if r.text: east_asia_run(r)
    add_running_matter(doc); doc.save(path)

def main():
    if not PATCHED.is_file() or not REFERENCE.is_file() or not PANDOC.is_file(): raise SystemExit('Missing patched manuscript, style reference, or Pandoc')
    VISUAL.mkdir(parents=True,exist_ok=True); CANDIDATE.parent.mkdir(parents=True,exist_ok=True)
    f1=VISUAL/'figure_01_financial.png'; f2=VISUAL/'figure_02_evidence_chain.png'; draw_chart1(f1); draw_chart2(f2)
    src=PATCHED.read_text(encoding='utf-8')
    anchors=[
      ('| 估算淨財務效益 | −3,918.37 | 4,081.63 − 8,000 |','\n\n![圖 1 前後費用與訓練成本拆解。SAMPLE/SYNTHETIC；NT$；前後各六個月；配對 n=10。](../visual/0106_R01/figure_01_financial.png){width=9.3in}\n'),
      ('| ⑥ | 因果歸因「訓練導致改善」 | **無證據** | **不足** | 對照組、隨機化、行為中介證據 | F0106-LOG-06 |','\n\n![圖 2 訓練邏輯鏈與證據強度。SAMPLE/SYNTHETIC；前後各六個月；配對 n=10、問卷 n=10、客戶 n=1。](../visual/0106_R01/figure_02_evidence_chain.png){width=9.3in}\n')]
    for needle,insert in anchors:
        if src.count(needle)!=1: raise SystemExit('Required chart insertion anchor not unique')
        src=src.replace(needle,needle+insert)
    BUILD_MD.write_text(src,encoding='utf-8',newline='\n')
    with tempfile.TemporaryDirectory(prefix='ttqs0106_ref_') as td:
        ref=Document(str(REFERENCE)); clear_story(ref); sec=ref.sections[-1]
        sec.orientation=WD_ORIENT.LANDSCAPE; sec.page_width=Inches(11.69); sec.page_height=Inches(8.27)
        sec.left_margin=Inches(.62); sec.right_margin=Inches(.62); sec.top_margin=Inches(.58); sec.bottom_margin=Inches(.58)
        for part in (sec.header,sec.footer): clear_part(part)
        refpath=Path(td)/'reference_a4_landscape.docx'; ref.save(refpath)
        cmd=[str(PANDOC),str(BUILD_MD),'--from=markdown','--to=docx','--reference-doc',str(refpath),'--output',str(CANDIDATE),'--resource-path',os.pathsep.join([str(WRITER),str(ROOT/'visual')])]
        result=subprocess.run(cmd,capture_output=True,text=True,timeout=180)
        if result.returncode: raise RuntimeError('Pandoc failed: '+result.stderr[-2000:])
    postprocess(CANDIDATE)
    evidence=ROOT/'evidence'/'docx_build_receipt_0106_R01.json'; evidence.parent.mkdir(parents=True,exist_ok=True)
    rec={'doc_id':'0106','requirement_id':'I18-E14','candidate':str(CANDIDATE),'candidate_sha256':sha(CANDIDATE),'bytes':CANDIDATE.stat().st_size,'patched_manuscript_sha256':sha(PATCHED),'build_markdown_sha256':sha(BUILD_MD),'figure_1_sha256':sha(f1),'figure_2_sha256':sha(f2),'reference_docx_sha256':sha(REFERENCE),'page_size':'A4 landscape','body_font_pt':10.5,'table_font_pt':9.5,'builder':'Pandoc 3.6.3 + python-docx 1.2.0; Windows Pillow charts; existing 0105 R04 style-only reference; no Word COM automation'}
    evidence.write_text(json.dumps(rec,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(rec,ensure_ascii=False))
if __name__=='__main__': main()
