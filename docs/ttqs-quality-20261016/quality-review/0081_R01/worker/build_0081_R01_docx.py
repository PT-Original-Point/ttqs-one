from pathlib import Path
import hashlib,json,re,zipfile
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parents[1];W=R/'worker';C=R/'candidate';Q=R/'qa';BODY=C/'0081_R01_body.md';OUT=C/'0081_R01.docx'
def H(b):return hashlib.sha256(b).hexdigest()
def cli_text(p,o,c):
 t=''.join(json.loads(x).get('part',{}).get('text','') for x in p.read_text(encoding='utf-8').splitlines() if json.loads(x).get('part',{}).get('type')=='text')
 if t.count(o)!=1 or t.count(c)!=1:raise ValueError('missing CLI delimiters')
 return t,t[t.index(o)+len(o):t.index(c)].strip()
def cells(x):return [v.strip().replace('\\|','|') for v in x.strip()[1:-1].replace('\\|','\u0000').split('|')]
def issep(r):return all(re.fullmatch(r':?-{3,}:?',x.replace(' ','')) for x in r)
def gettable(lines,h):
 i=next(i for i,x in enumerate(lines) if x.startswith(h));s=next(i for i in range(i+1,len(lines)) if lines[i].strip().startswith('|'));t=[]
 for x in lines[s:]:
  if not x.strip().startswith('|'):break
  t.append(cells(x))
 if any(len(x)!=len(t[0]) for x in t):raise ValueError('ragged table: '+h)
 return s,s+len(t),t
def mdrow(x):return '| '+' | '.join(x)+' |'
def swap(lines,h,rows):
 s,e,_=gettable(lines,h);lines[s:e]=[mdrow(rows[0]),mdrow(['---']*len(rows[0]))]+[mdrow(r) for r in rows[1:]]
writer_raw=W/'opencode_writer_R01_stdout.jsonl';repair_raw=W/'opencode_body_repair_R01_stdout.jsonl'
rt,repair=cli_text(repair_raw,'<<<SECTION_MD>>>','<<<END_SECTION_MD>>>');(W/'opencode_body_repair_R01_response.md').write_text(rt,encoding='utf-8',newline='\n')
wt,body=cli_text(writer_raw,'<<<DOCUMENT_BODY_MD>>>','<<<END_DOCUMENT_BODY_MD>>>');body=body+'\n';origsha=H(body.encode());lines=body.splitlines()
s=next(i for i,x in enumerate(lines) if x.startswith('### 7.6 逐學員總計'));e=next(i for i in range(s+1,len(lines)) if lines[i].startswith('### 7.7'));lines[s:e]=repair.splitlines();patch=[]
_,_,t=gettable(lines,'### 3.1 真實台帳');d=[r for r in t if not issep(r)];assert len(d[0])==6
swap(lines,'### 3.1 真實台帳',[['行動與原始提案','提案職能／期限','真實狀態','所需真實證據']]+[[f'**{r[0]}**｜{r[1]}',f'{r[2]}｜提案期限 {r[3]}',r[4],r[5]] for r in d[1:]]);patch.append('3.1:6->4 columns')
_,_,t=gettable(lines,'### 5.1 示範台帳');d=[r for r in t if not issep(r)];assert len(d[0])==8
swap(lines,'### 5.1 示範台帳',[['行動與 0080 提案','職能／示範期限','合成狀態與示範證據','查核決定','下一步']]+[[f'**{r[0]}**｜{r[1]}',f'{r[2]}｜{r[3]}',f'{r[4]}｜{r[5]}',r[6],r[7]] for r in d[1:]]);patch.append('5.1:8->5 columns')
_,_,t=gettable(lines,'### 5.2 證據品質檢核明細');d=[r for r in t if not issep(r)];assert len(d[0])==9 and len(d)==8
swap(lines,'### 5.2 證據品質檢核明細',[['行動','EQ1-EQ7 證據品質檢核','判定依據']]+[[a,'; '.join(f'{x[0]} {x[i+2]}' for x in d[1:]),'逐項決定見 §5.3；樣本狀態不代表真實執行'] for i,a in enumerate(d[0][2:])]);patch.append('5.2:transpose 9->3 columns')
_,_,t=gettable(lines,'## 第 9 節');d=[r for r in t if not issep(r)];assert len(d[0])==7
swap(lines,'## 第 9 節',[['追蹤項目與決定','負責職能／示範期限','下一步','查核證據','狀態']]+[[f'**{r[0]}**｜{r[1]}',f'{r[2]}｜{r[4]}',r[3],r[5],r[6]] for r in d[1:]]);patch.append('9:7->5 columns')
_,_,t=gettable(lines,'### 12.1 跨文件定位');d=[r for r in t if not issep(r)];assert len(d[0])==6
swap(lines,'### 12.1 跨文件定位',[['文件／要求／方框','不同功能','與本件之關係']]+[[f'{r[0]}｜{r[1]}｜{r[2]}｜{r[3]}',r[4],r[5]] for r in d[1:]]);patch.append('12.1:6->3 columns')
body='\n'.join(lines).replace('B01-05','B01-P05').replace('B01-11','B01-P11').rstrip()+'\n';BODY.write_text(body,encoding='utf-8',newline='\n');bsha=H(body.encode())
vec={'P01':[2,1,2,2,2],'P02':[2,2,2,1,1],'P03':[2,2,1,2,1],'P04':[2,2,1,1,1],'P05':[2,0,2,2,1],'P06':[1,2,2,1,2],'P07':[2,1,2,2,2],'P08':[2,1,2,0,1],'P09':[2,1,1,1,2],'P10':[1,1,2,2,2],'P11':[1,1,0,2,1],'P12':[1,1,1,0,0]};col=[sum(v[i] for v in vec.values()) for i in range(5)];total=sum(sum(v) for v in vec.values());passes=sum(sum(v)>=8 for v in vec.values());holds=12-passes
assert col==[20,15,18,16,16] and total==85 and (passes,holds)==(6,6)

def font(run,size=10.5,bold=None,name='Microsoft JhengHei',color=None,italic=None):
 run.font.name=name;run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),name);run.font.size=Pt(size)
 if bold is not None:run.bold=bold
 if italic is not None:run.italic=italic
 if color:run.font.color.rgb=RGBColor(*color)
pat=re.compile(r'(\*\*.+?\*\*|`[^`]+`)')
def inline(p,text,size=10.5):
 pos=0
 for m in pat.finditer(text):
  if m.start()>pos:font(p.add_run(text[pos:m.start()]),size)
  v=m.group(0);font(p.add_run(v[2:-2] if v.startswith('**') else v[1:-1]),size if v.startswith('**') else size-.2,bold=True if v.startswith('**') else None,name='Microsoft JhengHei' if v.startswith('**') else 'Consolas')
  pos=m.end()
 if pos<len(text):font(p.add_run(text[pos:]),size)
def shade(cell,color):
 x=OxmlElement('w:shd');x.set(qn('w:fill'),color);cell._tc.get_or_add_tcPr().append(x)
def rowtag(row,tag):row._tr.get_or_add_trPr().append(OxmlElement(tag))
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69);sec.top_margin=Inches(.62);sec.bottom_margin=Inches(.62);sec.left_margin=Inches(.66);sec.right_margin=Inches(.66)
n=doc.styles['Normal'];n.font.name='Microsoft JhengHei';n.font.size=Pt(10.5);n._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei');n.paragraph_format.space_after=Pt(4);n.paragraph_format.line_spacing=1.12
for sn,sz in [('Heading 1',15),('Heading 2',12)]:
 st=doc.styles[sn];st.font.name='Microsoft JhengHei';st.font.size=Pt(sz);st.font.bold=True;st.font.color.rgb=RGBColor(31,78,121);st._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei');st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(9);st.paragraph_format.space_after=Pt(4)
p=sec.header.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;font(p.add_run('DOCX 0081　｜　I17B-E03　｜　官方方框 086　｜　候選示範'),8,color=(89,100,112))
p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;font(p.add_run('SAMPLE/SYNTHETIC　｜　REAL=UNKNOWN/HOLD　｜　頁 '),8,color=(89,100,112));r=p.add_run();fld=OxmlElement('w:fldChar');fld.set(qn('w:fldCharType'),'begin');ins=OxmlElement('w:instrText');ins.set(qn('xml:space'),'preserve');ins.text=' PAGE ';sepfield=OxmlElement('w:fldChar');sepfield.set(qn('w:fldCharType'),'separate');txt=OxmlElement('w:t');txt.text='1';end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');r._r.extend([fld,ins,sepfield,txt,end]);font(r,8)
ls=body.splitlines();i=0;nt=0
while i<len(ls):
 s=ls[i].strip()
 if not s or s=='---':i+=1;continue
 if s.startswith('|'):
  block=[]
  while i<len(ls) and ls[i].strip().startswith('|'):block.append(cells(ls[i]));i+=1
  if len(block)<2 or not issep(block[1]):raise ValueError('bad markdown table separator')
  data=[block[0]]+[r for r in block[2:] if not issep(r)];nc=len(data[0]);tb=doc.add_table(rows=len(data),cols=nc);tb.style='Table Grid';tb.autofit=True;nt+=1
  for ri,rr in enumerate(data):
   rowtag(tb.rows[ri],'w:cantSplit')
   for ci,v in enumerate(rr):
    cell=tb.cell(ri,ci);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cell.text='';p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1;inline(p,v,8.4 if nc<=5 else 7.8)
    if ri==0:shade(cell,'D9EAF7');[setattr(z,'bold',True) for z in p.runs]
    elif ri%2==0:shade(cell,'F5F8FB')
  rowtag(tb.rows[0],'w:tblHeader');doc.add_paragraph().paragraph_format.space_after=Pt(1);continue
 if s.startswith('# '):
  p=doc.add_paragraph(style='Title');p.paragraph_format.keep_with_next=True;inline(p,s[2:],21)
  for z in p.runs:z.bold=True;z.font.color.rgb=RGBColor(31,78,121)
 elif s.startswith('## '):inline(doc.add_paragraph(style='Heading 1'),s[3:],15)
 elif s.startswith('### '):inline(doc.add_paragraph(style='Heading 2'),s[4:],12)
 elif s.startswith('> '):
  p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.2);inline(p,s[2:],9.4)
  for z in p.runs:z.italic=True
  sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEF4F8');p._p.get_or_add_pPr().append(sh)
 elif s.startswith('- '):inline(doc.add_paragraph(style='List Bullet'),s[2:],10.2)
 elif re.match(r'^\d+\.\s',s):inline(doc.add_paragraph(style='List Number'),re.sub(r'^\d+\.\s+','',s),10.2)
 else:inline(doc.add_paragraph(),s,10.5)
 i+=1
doc.core_properties.title='DOCX 0081 結案報告與課後課程檢討行動計畫查核紀錄';doc.core_properties.subject='I17B-E03 / official checkbox 086 / SAMPLE-SYNTHETIC candidate';doc.core_properties.author='';doc.core_properties.comments='CANDIDATE ONLY; REAL=UNKNOWN/HOLD; NOT_GRANTED';doc.save(OUT)
with zipfile.ZipFile(OUT) as z:
 if z.testzip():raise ValueError('DOCX ZIP invalid')
r=Document(OUT);ex='\n'.join([p.text for p in r.paragraphs]+[c.text for t in r.tables for rr in t.rows for c in rr.cells]);req=['0081','I17B-E03','086','REAL=UNKNOWN/HOLD','NOT_GRANTED','SAMPLE/SYNTHETIC','SYN-0081-2026-B01','85/120','198/240','0082'];missing=[x for x in req if x not in ex]
if missing:raise ValueError('missing text '+repr(missing))
dsha=H(OUT.read_bytes());ssha=H(Path(__file__).read_bytes());Q.mkdir(exist_ok=True)
(Q/'body_reflow_R01.json').write_text(json.dumps({'schema':'TTQS_ONE_0081_BODY_REFLOW_R01.v1','doc_id':'0081','source_writer_body_sha256':origsha,'repair_response_sha256':H((W/'opencode_body_repair_R01_response.md').read_bytes()),'final_body_sha256':bsha,'changes':patch,'synthetic_column_totals':col,'grand_total':total,'pass':passes,'hold':holds,'critical_error_codes':['B01-P05','B01-P11'],'authoring_note':'OpenCode authored body and corrected subsection; deterministic table reflow and arithmetic validation only.'},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
(Q/'technical_review_R01.json').write_text(json.dumps({'schema':'TTQS_ONE_0081_TECHNICAL_REVIEW_R01.v1','doc_id':'0081','requirement_id':'I17B-E03','official_checkbox_serial':'086','candidate_state':'CANDIDATE_ONLY','docx_path':'candidate/0081_R01.docx','docx_bytes':OUT.stat().st_size,'docx_sha256':dsha,'body_sha256':bsha,'build_script_sha256':ssha,'paragraphs':len(r.paragraphs),'tables':len(r.tables),'markdown_tables':nt,'required_text_missing':missing,'docx_zip_valid':True,'criterion_totals':col,'criterion_total':sum(col),'participant_total':total,'pass':passes,'hold':holds,'status':'PASS_TECHNICAL_SCOPE','semantic_review':'PENDING_INDEPENDENT_EXACT_SHA','visual_review':'PENDING_WORD_RENDER','real_status':'UNKNOWN/HOLD','formal_acceptance':'NOT_GRANTED'},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'docx_sha256':dsha,'docx_bytes':OUT.stat().st_size,'body_sha256':bsha,'tables':nt,'criterion_totals':col,'grand_total':total,'pass':passes,'hold':holds,'reflow':patch},ensure_ascii=False))
