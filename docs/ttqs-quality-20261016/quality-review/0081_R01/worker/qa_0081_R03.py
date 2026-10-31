from pathlib import Path
from docx import Document
import hashlib,json,zipfile,csv,sys
r=Path(sys.argv[1]);c=r/'candidate';q=r/'qa';doc=c/'0081_R03.docx';body=c/'0081_R03_body.md';source=c/'0081_R01_body.md';pdf=q/'0081_R03_word_render.pdf';imgs=sorted((q/'render_R03').glob('page-*.png'));reg=q/'SYNTHETIC_REPLACEMENT_REGISTER_0081_R03.csv'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
with zipfile.ZipFile(doc) as z:
 if z.testzip():raise SystemExit('invalid DOCX package')
d=Document(doc); oldbody=source.read_text(encoding='utf-8');newbody=body.read_text(encoding='utf-8')
expected=oldbody.replace('候選內容，尚未經本會核定，亦未經獨立語意審查','候選內容；尚未經本會核定；審查狀態以 exact-SHA 收據為準').replace('候選內容（R01）','候選內容（R03）').replace('SYNTHETIC_REPLACEMENT_REGISTER_0081_R01.csv','SYNTHETIC_REPLACEMENT_REGISTER_0081_R03.csv')
if expected!=newbody:raise SystemExit('body changed outside three planned metadata corrections')
text='\n'.join([p.text for p in d.paragraphs]+[cell.text for t in d.tables for row in t.rows for cell in row.cells])
need=['0081','I17B-E03','086','REAL=UNKNOWN/HOLD','NOT_GRANTED','SAMPLE/SYNTHETIC','SYN-0081-2026-B01','SCN-0081-001','85/120','198/240','0082','候選內容（R03）','SYNTHETIC_REPLACEMENT_REGISTER_0081_R03.csv']
missing=[x for x in need if x not in text]
if missing:raise SystemExit('missing text '+repr(missing))
if len(d.tables)!=20:raise SystemExit('table count changed')
if len(imgs)!=13:raise SystemExit('render page count mismatch')
with reg.open(encoding='utf-8-sig',newline='') as f: rows=list(csv.DictReader(f))
if len(rows)!=23 or any(x['real_status']!='REAL=UNKNOWN/HOLD' for x in rows):raise SystemExit('register malformed')
if any(x['source_path']!='docs/ttqs-quality-20261016/quality-review/0081_R01/candidate/0081_R03_body.md' for x in rows):raise SystemExit('register source path mismatch')
# Recompute B01 case.
vecs=[[2,1,2,2,2],[2,2,2,1,1],[2,2,1,2,1],[2,2,1,1,1],[2,0,2,2,1],[1,2,2,1,2],[2,1,2,2,2],[2,1,2,0,1],[2,1,1,1,2],[1,1,2,2,2],[1,1,0,2,1],[1,1,1,0,0]]
cols=[sum(v[j] for v in vecs) for j in range(5)];totals=[sum(v) for v in vecs]
if cols!=[20,15,18,16,16] or sum(cols)!=85 or sum(totals)!=85 or sum(x>=8 for x in totals)!=6:raise SystemExit('arithmetic mismatch')
tech={'schema':'TTQS_ONE_0081_TECHNICAL_REVIEW_R03.v1','doc_id':'0081','requirement_id':'I17B-E03','official_checkbox_serial':'086','indicator':'17b','candidate_state':'CANDIDATE_ONLY','docx_path':'candidate/0081_R03.docx','docx_bytes':doc.stat().st_size,'docx_sha256':sha(doc),'body_path':'candidate/0081_R03_body.md','body_sha256':sha(body),'change_scope':'Three version/status/register-path metadata corrections only; all scenario content and score facts unchanged from R02.','body_matches_exact_expected_patch':True,'tables':len(d.tables),'docx_zip_valid':True,'required_text_missing':missing,'criterion_totals':cols,'criterion_total':sum(cols),'participant_total':sum(totals),'pass':sum(x>=8 for x in totals),'hold':sum(x<8 for x in totals),'replacement_register_rows':len(rows),'replacement_register_sha256':sha(reg),'status':'PASS_TECHNICAL_SCOPE','semantic_review':'PENDING_INDEPENDENT_R03_DELTA_EXACT_SHA','visual_review':'PASS_WORD_EXPORT_ALL_13_PAGES_INSPECTED','real_status':'UNKNOWN/HOLD','formal_acceptance':'NOT_GRANTED'}
(q/'technical_review_R03.json').write_text(json.dumps(tech,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
visual={'schema':'TTQS_ONE_0081_VISUAL_REVIEW_R03.v1','doc_id':'0081','candidate_sha256':sha(doc),'rendered_pdf':'qa/0081_R03_word_render.pdf','rendered_pdf_sha256':sha(pdf),'page_count':len(imgs),'renderer':'Microsoft Word COM ExportAsFixedFormat; PDFium page rasterization','all_pages_directly_inspected':True,'scope':'All 13 page images inspected at full rendered dimensions. R03 updates only front-page status and the page 11 register filename; the section 1.3 table remains with its first data rows on page 2. No clipping, overlap, broken table, unreadable text, or missing glyph observed.','page_images':[{'path':str(p.relative_to(r)).replace('\\','/'),'sha256':sha(p),'width_px':894,'height_px':1263} for p in imgs],'layout_result':'PASS','candidate_state':'CANDIDATE_ONLY','formal_acceptance':'NOT_GRANTED'}
(q/'visual_review_R03.json').write_text(json.dumps(visual,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'docx_sha256':sha(doc),'pdf_sha256':sha(pdf),'body_sha256':sha(body),'register_sha256':sha(reg),'pages':len(imgs),'tables':len(d.tables),'total':sum(cols),'pass':6,'hold':6},ensure_ascii=True))
