from pathlib import Path
from docx import Document
import hashlib,json,sys
root=Path(sys.argv[1]); src=root/'candidate'/'0081_R01.docx'; out=root/'candidate'/'0081_R02.docx'
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def textshape(doc): return ([p.text for p in doc.paragraphs],[[[p.text for p in c.paragraphs] for c in r.cells] for t in doc.tables for r in t.rows])
d0=Document(src); before=textshape(d0)
# Prevent the first status-table header from being stranded at the end of page 1.
for cell in d0.tables[0].rows[0].cells:
    for p in cell.paragraphs: p.paragraph_format.keep_with_next=True
d0.save(out)
d1=Document(out)
if textshape(d1)!=before: raise SystemExit('text/table content changed')
receipt={'schema':'TTQS_ONE_0081_LAYOUT_REPAIR_R02.v1','doc_id':'0081','source_candidate':'candidate/0081_R01.docx','source_sha256':sha(src),'candidate':'candidate/0081_R02.docx','candidate_sha256':sha(out),'change':'Set keep-with-next on first row header-cell paragraphs in status table so the header moves with its first data row.','content_text_and_table_cells_unchanged':True,'candidate_state':'CANDIDATE_ONLY','real_status':'UNKNOWN/HOLD','formal_acceptance':'NOT_GRANTED'}
q=root/'qa';q.mkdir(exist_ok=True);(q/'layout_fix_R02.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(receipt,ensure_ascii=False))
