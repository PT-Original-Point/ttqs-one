from pathlib import Path
from hashlib import sha256
from docx import Document
import json
base=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0119_I19-E07_SOURCE_PREP_R01_20261024')
h=base/'candidate_R01H'/'0119__目標市場及顧客價值部分提升證據分析表_R01H.docx'
def replace_in_para(p):
 n=0
 for r in p.runs:
  c=r.text.count('R01G')
  if c:r.text=r.text.replace('R01G','R01H');n+=c
 return n
d=Document(h);changed=0
for p in d.paragraphs: changed+=replace_in_para(p)
for t in d.tables:
 for row in t.rows:
  for c in row.cells:
   for p in c.paragraphs: changed+=replace_in_para(p)
for s in d.sections:
 for p in s.header.paragraphs+s.footer.paragraphs: changed+=replace_in_para(p)
assert changed==3, f'expected 3 stale version labels, changed {changed}'
d.save(h)
check=Document(h)
assert all('R01G' not in p.text for p in check.paragraphs)
assert all('R01G' not in c.text for t in check.tables for row in t.rows for c in row.cells)
assert all('R01G' not in p.text for s in check.sections for p in s.header.paragraphs+s.footer.paragraphs)
assert sum(p.text.count('R01H') for p in check.paragraphs)==3
print(json.dumps({'changed_occurrences':changed,'candidate_sha256':sha256(h.read_bytes()).hexdigest(),'candidate_label_residuals':0},ensure_ascii=True))
