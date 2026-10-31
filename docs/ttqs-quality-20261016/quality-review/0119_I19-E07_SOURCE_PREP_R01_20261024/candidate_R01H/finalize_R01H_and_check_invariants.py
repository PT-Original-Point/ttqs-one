from pathlib import Path
from hashlib import sha256
from docx import Document
import json
base=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0119_I19-E07_SOURCE_PREP_R01_20261024')
g=base/'candidate_R01G'/'0119__目標市場及顧客價值部分提升證據分析表_R01G.docx'
h=base/'candidate_R01H'/'0119__目標市場及顧客價值部分提升證據分析表_R01H.docx'
def sha(p): return sha256(p.read_bytes()).hexdigest()
dg=Document(g); dh=Document(h)
# Confirm no paragraph or table cell content was edited by the layout patch.
assert [p.text for p in dg.paragraphs]==[p.text for p in dh.paragraphs]
assert [[c.text for row in t.rows for c in row.cells] for t in dg.tables]==[[c.text for row in t.rows for c in row.cells] for t in dh.tables]
assert len(dh.tables)==26
needle='\u886812-1 \u5b8c\u6574\u66ff\u63db\u767b\u9304\u518a'
hits=[p for p in dh.paragraphs if p.text.strip()==needle]
assert len(hits)==1 and hits[0].paragraph_format.page_break_before is True
for section in dh.sections:
 for p in section.footer.paragraphs:
  for run in p.runs:
   if 'R01G' in run.text: run.text=run.text.replace('R01G','R01H')
dh.save(h)
check=Document(h)
assert all('R01G' not in p.text for p in check.sections[0].footer.paragraphs)
assert any('R01H' in p.text for p in check.sections[0].footer.paragraphs)
assert [p.text for p in dg.paragraphs]==[p.text for p in check.paragraphs]
assert [[c.text for row in t.rows for c in row.cells] for t in dg.tables]==[[c.text for c in row.cells] for t in check.tables]
rec={'candidate':'0119 R01H','parent_candidate':'0119 R01G','parent_sha256':sha(g),'candidate_sha256':sha(h),'change_scope':['page_break_before on exact Table 12-1 caption to prevent orphan caption/header','footer candidate label R01G -> R01H'],'visible_body_paragraphs_unchanged':True,'table_cell_content_unchanged':True,'table_count':len(check.tables),'formal_acceptance':'NOT_GRANTED','semantic_review':'PENDING'}
(base/'candidate_R01H'/'build_receipt_R01H.json').write_text(json.dumps(rec,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(rec,ensure_ascii=True))
