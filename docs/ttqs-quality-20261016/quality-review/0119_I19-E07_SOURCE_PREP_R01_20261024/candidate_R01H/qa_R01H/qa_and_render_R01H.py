from pathlib import Path
from docx import Document
import zipfile,hashlib,json,re,decimal
from win32com.client import DispatchEx
import pythoncom,pypdfium2 as pdfium
root=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0119_I19-E07_SOURCE_PREP_R01_20261024\candidate_R01H')
qa=root/'qa_R01H'; qa.mkdir(exist_ok=True)
docx=next(root.glob('*.docx')); raw=docx.read_bytes(); docsha=hashlib.sha256(raw).hexdigest(); d=Document(str(docx))
paras=[p.text for p in d.paragraphs]
tables=['\n'.join(' | '.join(c.text.replace('\n',' / ') for c in row.cells) for row in t.rows) for t in d.tables]
text='\n'.join(paras+tables); norm=re.sub(r'\s+','',text)
checks={
 'zip_integrity':zipfile.ZipFile(docx).testzip() is None,
 'opens_with_python_docx':True,
 'actual_word_table_count_26':len(d.tables)==26,
 'table_3_3_present_and_usable':'表3-3 定義／口徑事後變更紀錄單' in text and '兩期可回溯重算' in text,
 'table_7_2_caption_present':'表7-2 手算重算表' in text,
 'sample_boundary_present':all(x in text for x in ['SAMPLE','SYNTHETIC','正式核准']),
 'stale_unbuilt_claim_absent':'文件尚未建置 DOCX' not in text,
 'markdown_bold_residue_absent':'**' not in text,
 'mixed_script_residue_absent':'構face' not in text,
 'candidate_status_present':'R01H' in text and 'CANDIDATE' in text,
 'corrected_references_present':all(re.sub(r'\s+','',x) in norm for x in ['表12-1 全列替換並通過第12.2節殘留檢查','表13-1 跨文件分界','表14-2']),
 'known_stale_references_absent':all(x not in text for x in ['填表11-2全列並送第二覆核','表15-2','第 15 節完成覆核與結案狀態']),
 'source_blueprint_mapping_bound':all(x in text for x in ['I19-E07','0119','序位 131','種子列 119','資料列 106']),
 'section_4_2_present':'4.2 指標規格表與欄位填寫規則' in text,
}
table_sizes=[r.font.size.pt for t in d.tables for row in t.rows for cell in row.cells for p in cell.paragraphs for r in p.runs if r.text.strip() and r.font.size is not None]
body_sizes=[r.font.size.pt for p in d.paragraphs for r in p.runs if r.text.strip() and r.font.size is not None]
table_min=min(table_sizes) if table_sizes else None
body_min=min(body_sizes) if body_sizes else None
checks['table_font_minimum_9_5pt']=table_min is not None and table_min>=9.5
checks['body_font_minimum_10_5pt']=body_min is not None and body_min>=10.5
D=decimal.Decimal
pct=lambda n,d:D(n)*100/D(d)
calc={'M-01':{'baseline_pct':str(pct(8,20)),'comparison_pct':str(pct(12,25)),'delta_pp':str(pct(12,25)-pct(8,20)),'relative_pct':str((pct(12,25)-pct(8,20))*100/pct(8,20))},'C-01':{'baseline_pct':str(pct(10,16)),'comparison_pct':str(pct(15,20)),'delta_pp':str(pct(15,20)-pct(10,16)),'baseline_response_rate':str(pct(16,20)),'comparison_response_rate':str(pct(20,25))},'negative_scenario_C01_delta_pp':str(pct(10,16)-pct(15,20))}
res={'doc_id':'0119','requirement_id':'I19-E07','candidate_state':'CANDIDATE_SAMPLE_ONLY','docx_path':str(docx),'docx_sha256':docsha,'bytes':len(raw),'paragraphs':len(d.paragraphs),'tables':len(d.tables),'min_table_font_pt':table_min,'min_body_font_pt':body_min,'checks':checks,'all_checks_pass':all(checks.values()),'independently_recomputed_synthetic_calculations':calc,'formal_acceptance':'NOT_GRANTED'}
(qa/'technical_content_QA_R01H.json').write_text(json.dumps(res,ensure_ascii=False,indent=2),encoding='utf-8')
# Native Microsoft Word isolated automation instance; no visible window, no user document is opened.
pdf=qa/'0119_R01H_native_word_render.pdf'; pythoncom.CoInitialize(); app=None; doc=None
try:
 app=DispatchEx('Word.Application'); app.Visible=False; app.DisplayAlerts=0
 doc=app.Documents.Open(str(docx),ReadOnly=True,AddToRecentFiles=False,Visible=False)
 doc.ExportAsFixedFormat(str(pdf),17,OpenAfterExport=False,OptimizeFor=0,Range=0,Item=0,IncludeDocProps=True,KeepIRM=True,CreateBookmarks=0,DocStructureTags=True,BitmapMissingFonts=True,UseISO19005_1=False)
finally:
 if doc is not None: doc.Close(False)
 if app is not None: app.Quit(False)
 pythoncom.CoUninitialize()
pdfdoc=pdfium.PdfDocument(str(pdf)); pages_dir=qa/'word_pages_R01H'; pages_dir.mkdir(exist_ok=True)
page_rows=[]
for i in range(len(pdfdoc)):
 page=pdfdoc[i]; bitmap=page.render(scale=1.75); image=bitmap.to_pil(); out=pages_dir/f'page-{i+1:02d}.png'; image.save(out)
 page_rows.append({'page':i+1,'png':str(out),'sha256':hashlib.sha256(out.read_bytes()).hexdigest(),'width':image.width,'height':image.height,'bytes':out.stat().st_size})
manifest={'docx_sha256':docsha,'pdf_path':str(pdf),'pdf_sha256':hashlib.sha256(pdf.read_bytes()).hexdigest(),'pages':len(pdfdoc),'page_render_scale':1.75,'page_images':page_rows,'layout_inspection_status':'PENDING_DIRECT_VISUAL_REVIEW'}
(qa/'visual_render_manifest_R01H.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'technical':res,'render':{'pdf':str(pdf),'pdf_sha256':manifest['pdf_sha256'],'pages':len(pdfdoc),'dimensions':sorted(set((x['width'],x['height']) for x in page_rows)),'first_pages':[x['png'] for x in page_rows[:3]]}},ensure_ascii=False,indent=2))
