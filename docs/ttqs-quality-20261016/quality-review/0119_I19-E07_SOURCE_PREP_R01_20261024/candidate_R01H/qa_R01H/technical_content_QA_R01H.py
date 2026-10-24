from pathlib import Path
import hashlib,json,re,zipfile,decimal
from docx import Document
base=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0119_I19-E07_SOURCE_PREP_R01_20261024')
g=base/'candidate_R01G'/'0119__目標市場及顧客價值部分提升證據分析表_R01G.docx'
f=base/'candidate_R01H'/'0119__目標市場及顧客價值部分提升證據分析表_R01H.docx'
qa=base/'candidate_R01H'/'qa_R01H'
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def cell_matrix(d): return [[ [c.text for c in row.cells] for row in t.rows] for t in d.tables]
dg, d = Document(g), Document(f)
raw=f.read_bytes(); paras=[p.text for p in d.paragraphs]
tables=['\n'.join(' | '.join(c.text.replace('\n',' / ') for c in row.cells) for row in t.rows) for t in d.tables]
text='\n'.join(paras+tables); norm=re.sub(r'\s+','',text)
footer='\n'.join(p.text for s in d.sections for p in s.footer.paragraphs)
cap='\u886812-1 \u5b8c\u6574\u66ff\u63db\u767b\u9304\u518a'
hits=[p for p in d.paragraphs if p.text.strip()==cap]
checks={
 'zip_integrity':zipfile.ZipFile(f).testzip() is None,
 'opens_with_python_docx':True,
 'word_table_count_26':len(d.tables)==26,
 'table_3_3_present_and_usable':'\u88683-3 \u5b9a\u7fa9\uff0f\u53e3\u5f91\u4e8b\u5f8c\u8b8a\u66f4\u7d00\u9304\u55ae' in text and '\u5169\u671f\u53ef\u56de\u6eaf\u91cd\u7b97' in text,
 'table_7_2_caption_present':'\u88687-2 \u624b\u7b97\u91cd\u7b97\u8868' in text,
 'sample_boundary_present':all(x in text for x in ['SAMPLE','SYNTHETIC','\u6b63\u5f0f\u6838\u51c6']),
 'no_stale_unbuilt_claim':'\u6587\u4ef6\u5c1a\u672a\u5efa\u7f6e DOCX' not in text,
 'no_markdown_bold_residue':'**' not in text,
 'no_known_mixed_script_residue':'\u69cbface' not in text,
 'candidate_label_in_footer_and_body':('R01H' in footer and 'CANDIDATE' in text and '\u6b63\u5f0f\u63a5\u53d7\u672a\u6388\u4e88' in footer),
 'corrected_crossreferences':all(re.sub(r'\s+','',x) in norm for x in ['\u886812-1 \u5168\u5217\u66ff\u63db\u4e26\u901a\u904e\u7b2c12.2\u7bc0\u6b98\u7559\u6aa2\u67e5','\u886813-1 \u8de8\u6587\u4ef6\u5206\u754c','\u886814-2']),
 'known_stale_references_absent':all(x not in text for x in ['\u586b\u886811-2\u5168\u5217\u4e26\u9001\u7b2c\u4e8c\u8986\u6838','\u886815-2','\u7b2c 15 \u7bc0\u5b8c\u6210\u8986\u6838\u8207\u7d50\u6848\u72c0\u614b']),
 'source_and_blueprint_mapping_bound':all(x in text for x in ['I19-E07','0119','131','119','106']),
 'section_4_2_present':'4.2 \u6307\u6a19\u898f\u683c\u8868\u8207\u6b04\u4f4d\u586b\u5beb\u898f\u5247' in text,
 'orphan_table12_caption_fixed':len(hits)==1 and hits[0].paragraph_format.page_break_before is True,
 'body_paragraph_content_unchanged_except_candidate_version_labels':paras==[p.text.replace('R01G','R01H') for p in dg.paragraphs],
 'table_cell_content_unchanged_except_candidate_version_labels':cell_matrix(d)==[[[v.replace('R01G','R01H') for v in row] for row in table] for table in cell_matrix(dg)],
}
table_sizes=[r.font.size.pt for t in d.tables for row in t.rows for cell in row.cells for p in cell.paragraphs for r in p.runs if r.text.strip() and r.font.size is not None]
body_sizes=[r.font.size.pt for p in d.paragraphs for r in p.runs if r.text.strip() and r.font.size is not None]
checks['table_font_min_9_5pt']=bool(table_sizes) and min(table_sizes)>=9.5
checks['body_font_min_10_5pt']=bool(body_sizes) and min(body_sizes)>=10.5
D=decimal.Decimal; pct=lambda n,den:D(n)*100/D(den)
calc={'M-01':{'baseline_pct':str(pct(8,20)),'comparison_pct':str(pct(12,25)),'delta_pp':str(pct(12,25)-pct(8,20)),'relative_pct':str((pct(12,25)-pct(8,20))*100/pct(8,20))},'C-01':{'baseline_pct':str(pct(10,16)),'comparison_pct':str(pct(15,20)),'delta_pp':str(pct(15,20)-pct(10,16)),'baseline_response_pct':str(pct(16,20)),'comparison_response_pct':str(pct(20,25))},'negative_scenario_C01_delta_pp':str(pct(10,16)-pct(15,20))}
manifest=json.loads((qa/'visual_render_manifest_R01H.json').read_text(encoding='utf-8'))
checks['native_word_render_bound_to_same_sha']=manifest.get('docx_sha256')==sha(f)
checks['render_page_images_complete']=manifest.get('pages')==len(manifest.get('page_images',[])) and len(list((qa/'word_pages_R01H').glob('page-*.png')))==manifest.get('pages')
res={'doc_id':'0119','requirement_id':'I19-E07','official_checkbox_serial':'UNKNOWN_NOT_ASSERTED','candidate':'R01H','candidate_state':'CANDIDATE_SAMPLE_ONLY','docx_path':str(f),'docx_sha256':sha(f),'parent_sha256':sha(g),'bytes':len(raw),'paragraphs':len(d.paragraphs),'tables':len(d.tables),'min_table_font_pt':min(table_sizes) if table_sizes else None,'min_body_font_pt':min(body_sizes) if body_sizes else None,'checks':checks,'all_checks_pass':all(checks.values()),'independently_recomputed_synthetic_calculations':calc,'initial_inherited_runner_result':'FAIL_CANDIDATE_VERSION_CHECK_DID_NOT_INCLUDE_FOOTER; corrected by separate footer-aware validator; original receipt retained','semantic_review':'PENDING','formal_acceptance':'NOT_GRANTED'}
(qa/'technical_content_QA_R01H.json').write_text(json.dumps(res,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(res,ensure_ascii=True,indent=2))

