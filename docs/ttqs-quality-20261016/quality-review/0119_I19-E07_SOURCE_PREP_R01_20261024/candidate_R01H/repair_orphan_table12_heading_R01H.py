from pathlib import Path
from hashlib import sha256
from docx import Document
import json, sys
src=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0119_I19-E07_SOURCE_PREP_R01_20261024\candidate_R01G\0119__目標市場及顧客價值部分提升證據分析表_R01G.docx')
out=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0119_I19-E07_SOURCE_PREP_R01_20261024\candidate_R01H\0119__目標市場及顧客價值部分提升證據分析表_R01H.docx')
expected='effc5faaf957a42719da89f8457bdc1ae4865090f7d6957b8ec49eb1523d1da1'
def digest(p): return sha256(p.read_bytes()).hexdigest()
assert digest(src)==expected, f'parent SHA mismatch: {digest(src)}'
doc=Document(src)
needle='\u886812-1 \u5b8c\u6574\u66ff\u63db\u767b\u9304\u518a'
hits=[p for p in doc.paragraphs if p.text.strip()==needle]
assert len(hits)==1, f'expected one exact Table 12-1 caption, got {len(hits)}'
assert hits[0].paragraph_format.page_break_before is None, 'unexpected prior page break'
hits[0].paragraph_format.page_break_before=True
doc.save(out)
check=Document(out)
hits2=[p for p in check.paragraphs if p.text.strip()==needle]
assert len(hits2)==1 and hits2[0].paragraph_format.page_break_before is True
receipt={
 'action':'TARGETED_LAYOUT_REPAIR_ONLY',
 'target':'Table 12-1 caption and full replacement register',
 'defect':'R01G page 18 rendered caption and table header without any data rows',
 'change':'Set page_break_before on exact caption paragraph; no body text or table cell edits',
 'parent_sha256':digest(src), 'candidate_sha256':digest(out),
 'paragraph_count':len(doc.paragraphs),'table_count':len(doc.tables),
 'caption_exact_match_count':len(hits2),'page_break_before':hits2[0].paragraph_format.page_break_before,
 'semantic_review':'PENDING','formal_acceptance':'NOT_GRANTED'
}
Path(out.parent/'build_receipt_R01H.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(receipt,ensure_ascii=True))
