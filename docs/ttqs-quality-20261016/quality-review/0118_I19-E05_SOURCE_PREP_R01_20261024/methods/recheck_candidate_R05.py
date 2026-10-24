from pathlib import Path
from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.oxml.ns import qn
import zipfile,hashlib,json
base=Path(__file__).resolve().parent.parent; w=base/'worker_0118_R05'; src=base/'source'; cand=base/'candidate'/'0118__感謝函原件收存與核驗登錄表_R05.docx'
dsha=hashlib.sha256(cand.read_bytes()).hexdigest();d=Document(str(cand));out=[f'# Exact DOCX body extraction | SHA-256 `{dsha}`']
for node in d.element.body.iterchildren():
 if node.tag==qn('w:p'):
  p=Paragraph(node,d)
  if p.text.strip():out.append(p.text.strip())
 elif node.tag==qn('w:tbl'):
  t=Table(node,d);out.append('[TABLE]')
  for row in t.rows:out.append(' | '.join(' / '.join(v.strip() for v in c.text.splitlines() if v.strip()) for c in row.cells))
  out.append('[/TABLE]')
extract=w/'candidate_text_exact_R05.md';extract.write_text('\n\n'.join(out)+'\n',encoding='utf-8')
text=extract.read_text(encoding='utf-8')
required=['版本：R05','I19-E05','保存期限依據','預定屆期日','不得自行銷毀','SYNTHETIC_NOT_EVIDENCE','HOLD_NO_ORIGINAL','NOT_GRANTED','（1）收件','（5）對照隔離']
checks={x:x in text for x in required}
assert all(checks.values()),checks
assert '4. **存取**' not in text and '5. **對照隔離**' not in text
with zipfile.ZipFile(cand) as z:bad=z.testzip()
assert bad is None,bad
widths={i:[round(col.width/914400,3) for col in d.tables[i].columns] for i in [2,15,16,20,23]}
report={'candidate_sha256':dsha,'bytes':cand.stat().st_size,'zip_integrity':'PASS','python_docx_reopen':'PASS','paragraphs':len(d.paragraphs),'tables':len(d.tables),'rows':sum(len(t.rows) for t in d.tables),'body_extraction_sha256':hashlib.sha256(extract.read_bytes()).hexdigest(),'required_checks':checks,'fixed_width_tables_inches':widths,'auto_fit_disabled':{str(i):not d.tables[i].autofit for i in [2,15,16,20,23]},'warnings':[]}
assert all(report['auto_fit_disabled'].values())
(w/'technical_independent_recheck_R05.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
parts=[src/'SOURCE_PACKET_0118_I19-E05_R01.md',src/'UNIQUE_BLUEPRINT_0118_I19-E05_R01.md',src/'mapping_evidence_0118.json',src/'official_source_page_08_text.txt',src/'SYNTHETIC_DATASET_0118_R01.csv',src/'SYNTHETIC_REPLACEMENT_REGISTER_0118_R01.csv',extract,base/'worker_0118_R02'/'visual_R02_findings.json',base/'worker_0118_R02'/'opencode_reviewer_R02_PASS_response.md']
lines=['# 0118 R05 exact-SHA review packet',f'Candidate DOCX SHA-256: `{dsha}`']
for f in parts:lines += [f'\n## SOURCE: {f.name}\nSHA-256: `{hashlib.sha256(f.read_bytes()).hexdigest()}`\n',f.read_text(encoding='utf-8')]
packet=w/'independent_review_packet_R05.md';packet.write_text('\n'.join(lines),encoding='utf-8')
print(json.dumps({**report,'review_packet_sha256':hashlib.sha256(packet.read_bytes()).hexdigest()},ensure_ascii=False))

