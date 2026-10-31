from pathlib import Path
import re
from docx import Documen
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMEN
from docx.oxml import OxmlElemen
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "candidate" / "0134_I19-E23_BODY_R01.md"
OUT = ROOT / "candidate" / "0134__會員人數及會費收入逐年變化分析報告_R01.docx"
BODY_MIN_PT = 10.5
TABLE_MIN_PT = 9.5
NAVY = "17365D"
TEAL = "1F5D68"
PALE = "EAF1F4"
GRID = "B8C5CC"


def set_east_asia(run, font_name="Microsoft JhengHei"):
    run.font.name = font_name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), font_name)


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shade = tcpr.find(qn("w:shd"))
    if shade is None:
        shade = OxmlElement("w:shd")
        tcpr.append(shade)
    shade.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    margins = tcpr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "w:" + edge
        node = margins.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tblpr = table._tbl.tblPr
    borders = tblpr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), GRID)


def clear_paragraph_border(paragraph_properties):
    if paragraph_properties is None:
        return
    border = paragraph_properties.find(qn("w:pBdr"))
    if border is not None:
        paragraph_properties.remove(border)


def rich_text(paragraph, text, bold=False, color=None):
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        is_bold = part.startswith("**") and part.endswith("**")
        value = part[2:-2] if is_bold else par
        run = paragraph.add_run(value)
        run.bold = bold or is_bold
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        set_east_asia(run)


def cell_text(line):
    return [x.strip() for x in line.strip().strip("|").split("|")]


def mark_repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    trpr.append(repeat)


def prevent_row_split(row):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))


def add_field(paragraph, field_name):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_name
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, separate, text, end):
        run._r.append(el)
    set_east_asia(run)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string("5B6570")


def add_markdown_table(doc, rows, body_width):
    if len(rows) < 2:
        return
    header = rows[0]
    data = rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    all_rows = [header] + data
    # Allocate more width to columns with longer labels/content while protecting narrow columns.
    maxima = []
    for c in range(len(header)):
        maxima.append(max([len(row[c]) if c < len(row) else 0 for row in all_rows] + [len(header[c])]))
    weights = [max(0.9, min(2.8, (n ** 0.5))) for n in maxima]
    available = body_width
    widths = [available * w / sum(weights) for w in weights]
    if len(header) >= 6:
        minw = 0.68
        deficit = sum(max(0, minw - w) for w in widths)
        widths = [max(minw, w) for w in widths]
        excess = sum(widths) - available
        if excess > 0:
            reducible = [max(0, w-minw) for w in widths]
            total = sum(reducible) or 1
            widths = [w - excess * r/total for w,r in zip(widths,reducible)]
    for ri, values in enumerate(all_rows):
        row = table.rows[0] if ri == 0 else table.add_row()
        prevent_row_split(row)
        if ri == 0:
            mark_repeat_header(row)
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if ri == 0:
                set_cell_shading(cell, TEAL)
            elif ri % 2 == 0:
                set_cell_shading(cell, PALE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.04
            rich_text(paragraph, values[ci] if ci < len(values) else "", bold=(ri == 0), color="FFFFFF" if ri == 0 else None)
            for run in paragraph.runs:
                run.font.size = Pt(TABLE_MIN_PT)
                if ri == 0:
                    run.font.color.rgb = RGBColor(255,255,255)
    set_table_borders(table)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.space_before = Pt(0)


def main():
    text = SOURCE.read_text(encoding="utf-8-sig")
    lines = text.splitlines()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.67)
    section.right_margin = Inches(0.67)
    body_width = 8.5 - 0.67 - 0.67

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft JhengHei"
    normal.font.size = Pt(BODY_MIN_PT)
    normal.font.color.rgb = RGBColor.from_string("20252A")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    for name, size in (("Title", 21), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
        st = styles[name]
        st.font.name = "Microsoft JhengHei"
        st.font.size = Pt(size)
        st.font.bold = name != "Title"
        st.font.color.rgb = RGBColor.from_string("000000" if name == "Title" else NAVY)
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.space_before = Pt(11 if name == "Heading 1" else 7)
        st.paragraph_format.space_after = Pt(4)
        if name == "Title":
            clear_paragraph_border(st._element.pPr)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGH
    r = header.add_run("TTQS ONE    0134 會員與會費年度分析    SAMPLE／SYNTHETIC")
    set_east_asia(r)
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string("52616B")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("候選稿 R01  ·  REAL 證據尚未提供  ·  第 ")
    set_east_asia(r)
    r.font.size = Pt(9)
    add_field(footer, "PAGE")
    r = footer.add_run(" 頁")
    set_east_asia(r)
    r.font.size = Pt(9)

    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{3,}", lines[i+1]):
            table_rows = [cell_text(line)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_rows.append(cell_text(lines[i]))
                i += 1
            add_markdown_table(doc, table_rows, body_width)
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            clear_paragraph_border(p._p.get_or_add_pPr())
            rich_text(p, line[2:].strip())
            for run in p.runs:
                run.font.size = Pt(21)
                run.font.color.rgb = RGBColor(0,0,0)
                run.bold = False
            i += 1
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            rich_text(p, line[3:].strip())
            if line.startswith("## 10 ") or line.startswith("## 11 ") or line.startswith("## 14 "):
                p.paragraph_format.page_break_before = True
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            rich_text(p, line[4:].strip())
            i += 1
            continue
        if line.startswith("#### "):
            p = doc.add_paragraph(style="Heading 3")
            rich_text(p, line[5:].strip())
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            rich_text(p, line[2:])
            i += 1
            continue
        p = doc.add_paragraph()
        rich_text(p, line)
        i += 1

    doc.core_properties.title = "會員人數及會費收入逐年變化分析報告"
    doc.core_properties.subject = "TTQS ONE 0134 I19-E23; SAMPLE/SYNTHETIC candidate"
    doc.core_properties.author = "TTQS ONE candidate preparation"
    doc.core_properties.keywords = "0134, I19-E23, membership, dues, SAMPLE, SYNTHETIC"
    doc.save(OUT)
    print(OUT)
    print(f"size={OUT.stat().st_size}")
    print(f"sha256={__import__('hashlib').sha256(OUT.read_bytes()).hexdigest()}")


if __name__ == "__main__":
    main()
