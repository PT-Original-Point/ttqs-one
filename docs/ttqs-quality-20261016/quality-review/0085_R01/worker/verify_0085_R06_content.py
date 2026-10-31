from pathlib import Path
import hashlib, json, re
from datetime import datetime, timezone
from docx import Document
from docx.oxml.ns import qn

root=Path(r'E:\TTQS\_github_sync_0064\docs\ttqs-quality-20261016\quality-review\0085_R01')
body_path=root/'worker'/'0085_R06_body.md'
docx_path=root/'candidate'/'0085_R06.docx'
receipt_path=root/'evidence'/'content_verification_R06.json'
body=body_path.read_text(encoding='utf-8')
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def clean(s): return s.replace('**','').replace('__','').replace('`','').replace('<br>',' ').strip()
# Extract markdown tables by subsection.
sec_tables={}
current=None; current_rows=[]
def flush():
 global current,current_rows
 if current and current_rows: sec_tables.setdefault(current,[]).append(current_rows)
 current_rows=[]
for line in body.splitlines():
 m=re.match(r'^#{3,}\s+(13\.\d+(?:\.\d+)?)\b',line)
 if m:
  flush(); current=m.group(1); continue
 if line.startswith('## '):
  flush(); current=None; continue
 if current and line.strip().startswith('|'):
  cells=[clean(x) for x in line.strip().strip('|').split('|')]
  if cells and all(re.fullmatch(r':?-{3,}:?',x.replace(' ','')) for x in cells): continue
  current_rows.append(cells)
flush()
# Appendix E rows
appendix=[]; in_e=False
for line in body.splitlines():
 if re.match(r'^##\s+附錄\s*E\b',line): in_e=True; continue
 if in_e and re.match(r'^##\s+',line): break
 if in_e and re.match(r'^\|\s*R85-\d{3}\s*\|',line):
  appendix.append([clean(x) for x in line.strip().strip('|').split('|')])

def find_table(prefix,key):
 tables=sec_tables.get(prefix,[])
 for tab in tables:
  for row in tab[1:]:
   if row and row[0].strip()==key: return tab,row
 raise KeyError(f'no table row {prefix} {key}')
def value_for(loc,field):
 if loc.startswith('§13.1/'):
  key=loc.split('/',1)[1]
  tab,row=find_table('13.1',key)
  return row[1]
 if loc.startswith('§13.2/') or loc.startswith('§13.3/') or loc.startswith('§13.7/') or loc.startswith('§13.9/') or loc.startswith('§13.10/'):
  key=loc.split('/')[-1]
  prefix=loc.split('/')[0].replace('§','')
  tab,row=find_table(prefix,key)
  return row[2]
 if loc.startswith('§13.4/'):
  _,key,colid=loc.split('/')
  tab,row=find_table('13.4',key)
  idx=next(i for i,h in enumerate(tab[0]) if colid in h)
  return row[idx]
 if loc.startswith('§13.5/'):
  _,key,what=loc.split('/')
  tab,row=find_table('13.5',key)
  return row[2 if what=='算式' else 3]
 if loc.startswith('§13.6/HDR/'):
  key=loc.split('/')[-1]
  tab,row=find_table('13.6.1',key)
  return row[2]
 if loc.startswith('§13.6/KI-'):
  parts=loc.split('/')
  key=parts[1]; fieldid=parts[2]
  tab,row=find_table('13.6.2',key)
  idx=next(i for i,h in enumerate(tab[0]) if fieldid in h)
  return row[idx]
 if loc.startswith('§13.8/'):
  key=loc.split('/')[1]
  tab,row=find_table('13.8',key)
  return row[1]
 if loc.startswith('§13.11/'):
  _,key,kind=loc.split('/')
  tab,row=find_table('13.11',key)
  # Decision table has label, issue, content, role, date, status.
  idx={'議題':1,'決議內容':2,'決議角色':3,'日期':4,'狀態':5}[kind]
  return row[idx]
 raise KeyError(f'unsupported locator {loc}')
mismatches=[]
for row in appendix:
 if len(row)!=6: mismatches.append({'id':row[0] if row else '?','error':'bad register shape','cells':len(row)}); continue
 rid,loc,field,want,*_=row
 try: got=value_for(loc,field)
 except Exception as e: mismatches.append({'id':rid,'loc':loc,'error':str(e)}); continue
 if got!=want: mismatches.append({'id':rid,'loc':loc,'expected':want,'actual':got})
ids=[r[0] for r in appendix]
expected=[f'R85-{i:03d}' for i in range(1,180)]
# Word-native §15.2 table check
D=Document(str(docx_path)); native=False; section_seen=False; table_info=None
for child in D._element.body.iterchildren():
 if child.tag==qn('w:p'):
  text=''.join(child.itertext())
  if '15.2 資料存取權限矩陣' in text: section_seen=True
  elif section_seen and '15.3 檔案保存政策' in text: section_seen=False
 elif child.tag==qn('w:tbl') and section_seen:
  from docx.table import Table
  t=Table(child,D._body); table_info={'rows':len(t.rows),'cols':len(t.columns),'first':t.cell(0,0).text}
  native=(len(t.rows)==7 and len(t.columns)==7 and '職務角色代碼' in t.cell(0,0).text)
  break
checks={
 'register_179_rows':len(appendix)==179,
 'register_keys_contiguous':ids==expected,
 'register_exact_value_cell_parity':not mismatches,
 'word_15_2_native_7x7_table':native,
 'body_blueprint_sha_present':'4fe718c40a4a46231863104853b7d5dad7ab48e7347d1d7639ef8fd49786b73e' in body,
 'sample_real_boundary_present':'REAL = UNKNOWN/HOLD' in body,
 'm9_exact_locator_rate_explicit':'完整頁段定位 0/6＝0.0%' in body,
 'blueprint_deviation_recorded':'DEVIATION D-0085-01' in body,
 'ap05_series_activity_code_present':'SYN-KM-SERIES-2026-C／活動 SYN-ACT-01' in body,
 'ap08_followup_open':'AP-08／KI-04' in body and '全部完成才關單' in body,
 'candidate_not_accepted':'CANDIDATE / NOT_ACCEPTED' in body and 'FORMAL_ACCEPTANCE = NOT_GRANTED' in body,
}
r={'project_id':'TTQS_ONE','doc_id':'0085','requirement_id':'I17C-E04','checkbox':'091','revision':'R06','body_sha256':sha(body_path),'docx_sha256':sha(docx_path),'docx_bytes':docx_path.stat().st_size,'register_rows':len(appendix),'exact_parity_mismatches':len(mismatches),'mismatch_examples':mismatches[:12],'word_15_2_native_table':table_info,'checks':checks,'all_checks_pass':all(checks.values()),'review_status':'PENDING_EXACT_SHA_INDEPENDENT_REVIEW','state':'CANDIDATE_NOT_ACCEPTED','checked_utc':datetime.now(timezone.utc).isoformat()}
receipt_path.write_text(json.dumps(r,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'all_checks_pass':r['all_checks_pass'],'rows':len(appendix),'parity_mismatches':len(mismatches),'table_15_2':table_info,'checks':checks,'mismatches':mismatches[:6],'receipt':str(receipt_path)},ensure_ascii=False))