import csv, hashlib, json, re, zipfile
from pathlib import Path
from docx import Document
root=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0007_I01-E04_R12_STATUS_CLAIM_REPAIR_20261031')
p=root/'candidate'/'0007_I01-E04_R12.docx'; receiptp=root/'qa'/'BUILD_RECEIPT_R12.json'; packetp=root/'source'/'SOURCE_PACKET_0007_I01-E04_R12.json'; regp=root/'source'/'SYNTHETIC_REPLACEMENT_REGISTER_0007_R12.csv'; srcp=root/'source'/'0007_I01-E04_R12_WORD_SOURCE.md'
def sha(x): return hashlib.sha256(x.read_bytes()).hexdigest()
assert p.is_file() and zipfile.is_zipfile(p)
with zipfile.ZipFile(p) as z:
 bad=z.testzip(); assert bad is None, f'corrupt zip member: {bad}'
d=Document(p); paragraphs=[x.text for x in d.paragraphs]; tables=d.tables
text='\n'.join(paragraphs+[c.text for t in tables for row in t.rows for c in row.cells]+[p.text for sec in d.sections for p in sec.header.paragraphs+sec.footer.paragraphs])
r=json.loads(receiptp.read_text(encoding='utf-8-sig')); packet=json.loads(packetp.read_text(encoding='utf-8-sig'))
assert sha(p)==r['candidate_sha256'], 'receipt SHA mismatch'
assert r['doc_id']=='0007' and r['requirement_id']=='I01-E04' and r['revision']=='R12'
assert r['blueprint_row_sha256']==packet['blueprint_row_sha256'] and r['official_pdf_sha256']==packet['official_source_sha256']
assert '0007' in text and 'I01-E04' in text and 'SAMPLE / SYNTHETIC' in text and 'CANDIDATE' in text
for stale in ['TEXT_CANDIDATE_NOT_DOCX','Markdown 無法驗證','未做 Word 視覺 QA','未做視覺 QA','文字候選 R02']:
 assert stale not in text, f'stale status survives: {stale}'
assert '版本及後續格式目標' not in text, 'internal production epilogue leaked into DOCX'
for token in ['Q3_NEED1','Q3_NEED2','SYN-005','SYN-006','F0002','F0003','Q3_NEED3','Q4_URG','Q5_CONF','Q7_MODE','Q9_DURATION','Q10_BARRIER','F0004','F0005','F0006','F0007','F0008','F0009','4.50','4.10','3.65','4/5=80%','5/5=100%','差距','改善率','−30','−10','33.33%','6／10＝60%','8／10＝80%','F0101','F0102','F0103','F0104','F0105','F0106','F0107','F0108','F0109','F0110','F0111','F0112','F0113','F0114','F0115','F0147','四捨五入','差距優先矩陣','HOLD']:
 assert token in text, f'missing worked result: {token}'
assert len(tables)==10, f'table count {len(tables)} != 10'
rows=list(csv.DictReader(regp.open(encoding='utf-8-sig',newline='')))
assert len(rows)==147, f'register rows={len(rows)}'
keys=[x.get('fact_id') or x.get('synthetic_fact_id') or x.get('id') for x in rows]
assert len(set(keys))==147 and all(keys), 'synthetic fact keys not unique/present'
# Source-specific weighted score recomputation from the documented 30/25/20/15/10 weights.
weights=[30,25,20,15,10]; vectors={'A':[5,4,5,4,4],'B':[4,4,5,4,3],'C':[3,4,5,3,3]}
scores={k:sum(w*v for w,v in zip(weights,vec))/100 for k,vec in vectors.items()}
improvement_rate=(0.80-0.60)/0.60
assert round(improvement_rate*100,2)==33.33
assert round(0.60-0.90,2)==-0.30 and round(0.80-0.90,2)==-0.10
assert scores=={'A':4.5,'B':4.1,'C':3.65}, scores
# Body paragraph font floor, excluding the cover metadata line; table font floor applies to every table run.
font_issues=[]
for i,para in enumerate(d.paragraphs):
 if i==1: continue
 for run in para.runs:
  if run.text.strip() and run.font.size and run.font.size.pt<10.5: font_issues.append(('body',i,run.font.size.pt,run.text[:30]))
for ti,t in enumerate(tables):
 for ri,row in enumerate(t.rows):
  for ci,cell in enumerate(row.cells):
   for para in cell.paragraphs:
    for run in para.runs:
     if run.text.strip() and (not run.font.size or run.font.size.pt<9.5): font_issues.append(('table',ti,ri,ci,run.font.size.pt if run.font.size else None,run.text[:30]))
assert not font_issues, f'font floor violations: {font_issues[:8]}'
result={'project_id':'TTQS_ONE','doc_id':'0007','requirement_id':'I01-E04','revision':'R12','status':'PASS_WITHIN_DETERMINISTIC_SCOPE','candidate_path':str(p),'candidate_sha256':sha(p),'bytes':p.stat().st_size,'zip_crc':'PASS','build_receipt_sha_bound':'PASS','official_source_and_unique_blueprint_hash_bound':'PASS','docx_paragraphs':len(paragraphs),'docx_tables':len(tables),'synthetic_register_rows':len(rows),'required_status_checks':'PASS','weighted_scores_recomputed':scores,'design_completeness':'4/5=80% < 5/5=100%; HOLD','synthetic_gap_example':'T0 60%-90%=-30 pp; T1 80%-90%=-10 pp','synthetic_improvement_rate':'(80%-60%)/60%=33.33%','font_floor':{'body_pt':10.5,'table_pt':9.5,'status':'PASS'},'visual_review':'PENDING','independent_semantic_review':'PENDING','formal_acceptance':'NOT_ASSERTED'}
out=root/'qa'/'DETERMINISTIC_QA_R12.json'; out.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8'); print(json.dumps(result,ensure_ascii=True))
