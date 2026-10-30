import csv, hashlib, json, re, sys
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(sys.argv[1])
original=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0007_I01-E04_R12_STATUS_CLAIM_REPAIR_20261031\worker\0007_I01-E04_ANALYSIS_REPORT_R12_STATUS_REPAIR.md')
register=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0007_I01-E04_R12_STATUS_CLAIM_REPAIR_20261031\source\SYNTHETIC_REPLACEMENT_REGISTER_0007_R12.csv')
source_packet=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0007_I01-E04_R12_STATUS_CLAIM_REPAIR_20261031\source\SOURCE_PACKET_0007_I01-E04_R12.json')
normalized_out=root/'source'/'0007_I01-E04_R12_WORD_SOURCE.md'
docx_out=root/'candidate'/'0007_I01-E04_R12.docx'
receipt_out=root/'qa'/'BUILD_RECEIPT_R12.json'

def sha(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()
text=original.read_text(encoding='utf-8-sig')
transforms=[]
replacements=[
 ('# 課程產品及流程檢討分析報告｜文字候選 R02','# 課程產品及流程檢討分析報告'),
 ('同資料夾 CSV','同一候選證據包中的 SYNTHETIC_REPLACEMENT_REGISTER_0007_R12.csv'),
 ('R02 由官方來源、現行 Blueprint','R12 由官方來源、現行 Blueprint'),
 ('**狀態：TEXT_CANDIDATE_NOT_DOCX；以下均為 SAMPLE／SYNTHETIC；未做視覺 QA、未正式接受。**  ','**狀態：DOCX 候選 R12；案例均為 SAMPLE／SYNTHETIC；非正式接受。**  '),
 ('日後 DOCX 正文字級目標至少 10.5 pt、表格至少 9.5 pt；Markdown 無法驗證字級或版面。本候選未做 Word 視覺 QA，未正式接受。','本候選 DOCX 字級政策為正文至少 10.5 pt、表格至少 9.5 pt；版面及獨立語意驗收狀態以本證據包內綁定 exact SHA 的收據為準。本文件仍為 SAMPLE／SYNTHETIC 候選，未正式接受。')]
for old,new in replacements:
 n=text.count(old)
 if n:
  text=text.replace(old,new)
  transforms.append({'from':old,'to':new,'replacements':n})
epilogue='## 版本及後續格式目標'
if epilogue in text:
 text=text.split(epilogue,1)[0].rstrip()+'\n'
 transforms.append({'from':epilogue+' and following production-only note','to':'omitted from evaluator-facing DOCX; retained in original worker draft','replacements':1})
normalized_out.write_text(text,encoding='utf-8')

NAVY='183B63'; BLUE='DCE7F2'; WHITE='FFFFFF'; TEXT='202938'; MUTED='5B6573'
FONT='Microsoft JhengHei'

def set_run_font(run, size=10.5, bold=None, color=TEXT, italic=None, name=FONT):
 run.font.name=name
 run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),name)
 run.font.size=Pt(size)
 run.font.color.rgb=RGBColor.from_string(color)
 if bold is not None: run.bold=bold
 if italic is not None: run.italic=italic

def add_inline(paragraph, value, size=10.5, base_color=TEXT):
 pat=re.compile(r'(\*\*.+?\*\*|`[^`]+`)')
 pos=0
 for m in pat.finditer(value):
  if m.start()>pos:
   run=paragraph.add_run(value[pos:m.start()]); set_run_font(run,size,color=base_color)
  token=m.group(0)
  if token.startswith('**'):
   run=paragraph.add_run(token[2:-2]); set_run_font(run,size,bold=True,color=base_color)
  else:
   run=paragraph.add_run(token[1:-1]); set_run_font(run,max(9.5,size),color='214A70',name='Consolas')
  pos=m.end()
 if pos<len(value):
  run=paragraph.add_run(value[pos:]); set_run_font(run,size,color=base_color)

def shade(cell, fill):
 tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
 if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
 shd.set(qn('w:fill'),fill)

def margins(cell, top=100, start=110, bottom=100, end=110):
 tc=cell._tc; tcPr=tc.get_or_add_tcPr(); mar=tcPr.first_child_found_in('w:tcMar')
 if mar is None: mar=OxmlElement('w:tcMar'); tcPr.append(mar)
 for tag,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
  e=mar.find(qn('w:'+tag))
  if e is None: e=OxmlElement('w:'+tag); mar.append(e)
  e.set(qn('w:w'),str(val)); e.set(qn('w:type'),'dxa')

def set_repeat_table_header(row):
 trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)

def set_no_split(row):
 trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:cantSplit'); trPr.append(el)

def set_section(sec, landscape=False):
 sec.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
 if landscape: sec.page_width=Cm(29.7); sec.page_height=Cm(21)
 else: sec.page_width=Cm(21); sec.page_height=Cm(29.7)
 sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(1.7); sec.left_margin=Cm(1.8); sec.right_margin=Cm(1.8)
 sec.header_distance=Cm(0.8); sec.footer_distance=Cm(0.8)

def add_field(paragraph, instruction):
 run=paragraph.add_run(); begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin')
 instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text=instruction
 separate=OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'),'separate')
 result=OxmlElement('w:t'); result.text='1'
 end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end')
 run._r.extend([begin,instr,separate,result,end]); set_run_font(run,8.5,color=MUTED)

doc=Document()
section=doc.sections[0]; set_section(section,False)
# Page-level source identification, without a fabricated logo or approval marker.
h=section.header.paragraphs[0]; h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
rr=h.add_run('屏東縣原始點關懷協會  |  0007 / I01-E04  |  CANDIDATE'); set_run_font(rr,8.5,color=MUTED)
f=section.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
rr=f.add_run('SAMPLE / SYNTHETIC  •  非正式接受  |  第 '); set_run_font(rr,8.5,color=MUTED); add_field(f,'PAGE'); rr=f.add_run(' 頁'); set_run_font(rr,8.5,color=MUTED)
styles=doc.styles
normal=styles['Normal']; normal.font.name=FONT; normal._element.rPr.rFonts.set(qn('w:eastAsia'),FONT); normal.font.size=Pt(10.5); normal.font.color.rgb=RGBColor.from_string(TEXT)
normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.12
for st,size in [('Heading 1',14),('Heading 2',12),('Heading 3',11)]:
 s=styles[st]; s.font.name=FONT; s._element.rPr.rFonts.set(qn('w:eastAsia'),FONT); s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(NAVY); s.paragraph_format.space_before=Pt(10); s.paragraph_format.space_after=Pt(4); s.paragraph_format.keep_with_next=True
styles['Title'].font.name=FONT; styles['Title']._element.rPr.rFonts.set(qn('w:eastAsia'),FONT); styles['Title'].font.size=Pt(20); styles['Title'].font.bold=True; styles['Title'].font.color.rgb=RGBColor.from_string('000000'); styles['Title'].paragraph_format.space_after=Pt(5); styles['Title'].paragraph_format.keep_with_next=True

current_landscape=False

def switch_orientation(wide):
 global current_landscape
 if wide==current_landscape: return
 sec=doc.add_section(WD_SECTION.NEW_PAGE)
 set_section(sec,wide)
 sec.header.is_linked_to_previous=True; sec.footer.is_linked_to_previous=True
 current_landscape=wide

def parse_cells(line):
 s=line.strip().strip('|')
 return [c.strip().replace('\\|','|') for c in s.split('|')]

def is_separator_row(cells):
 return bool(cells) and all(re.fullmatch(r':?-{3,}:?',c.replace(' ','')) for c in cells)

def add_table(rows):
 if not rows: return
 wide=len(rows[0])>=6
 switch_orientation(wide)
 active_sec=doc.sections[-1]; usable=(active_sec.page_width-active_sec.left_margin-active_sec.right_margin)/914400.0
 table=doc.add_table(rows=len(rows),cols=len(rows[0]))
 table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
 try: table.style='Table Grid'
 except KeyError: pass
 weights=[]
 for j in range(len(rows[0])):
  mx=max([len(rows[i][j]) for i in range(min(len(rows),18)) if j<len(rows[i])]+[len(rows[0][j]),5])
  weights.append(max(1.15,min(5.0,(mx**0.5))))
 minimum=0.66
 widths=[max(minimum,usable*w/sum(weights)) for w in weights]
 scale=usable/sum(widths); widths=[w*scale for w in widths]
 for ri,rowdata in enumerate(rows):
  row=table.rows[ri]; set_no_split(row)
  for ci in range(len(rows[0])):
   cell=row.cells[ci]; cell.width=Cm(widths[ci]*2.54); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(cell)
   cell.text=''
   p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.03
   val=rowdata[ci] if ci<len(rowdata) else ''
   add_inline(p,val,9.5,base_color=TEXT)
   for run in p.runs:
    run.font.size=Pt(9.5)
    if ri==0: run.bold=True; run.font.color.rgb=RGBColor.from_string(WHITE)
   if ri==0: shade(cell,NAVY)
   elif ri%2==0: shade(cell,'F1F5F9')
  if ri==0: set_repeat_table_header(row)
 # Exact widths on table grid and cells keep wide matrices stable in Word.
 for ci,w in enumerate(widths):
  table.columns[ci].width=Cm(w*2.54)
 doc.add_paragraph().paragraph_format.space_after=Pt(1)

def add_para(line, style=None, bold_all=False):
 p=doc.add_paragraph(style=style)
 if style and style.startswith('Heading'):
  add_inline(p,line,12 if style=='Heading 2' else 14)
 else:
  add_inline(p,line,10.5)
 if bold_all:
  for r in p.runs:r.bold=True
 return p

lines=normalized_out.read_text(encoding='utf-8-sig').splitlines()
blocks=[]; i=0
while i<len(lines):
 raw=lines[i].rstrip(); s=raw.strip()
 if not s: i+=1; continue
 if s.startswith('|'):
  block=[]
  while i<len(lines) and lines[i].strip().startswith('|'):
   cells=parse_cells(lines[i]);
   if not is_separator_row(cells): block.append(cells)
   i+=1
  if block: blocks.append(('table',block))
  continue
 if s.startswith('#'):
  level=len(s)-len(s.lstrip('#')); title=s[level:].strip()
  blocks.append(('heading',(level,title))); i+=1; continue
 if s.startswith('- '): blocks.append(('bullet',s[2:].strip())); i+=1; continue
 if re.match(r'^\d+\.\s+',s): blocks.append(('number',re.sub(r'^\d+\.\s+','',s))); i+=1; continue
 para=[s]; i+=1
 while i<len(lines) and lines[i].strip() and not lines[i].strip().startswith('|') and not lines[i].strip().startswith('#') and not lines[i].strip().startswith('- ') and not re.match(r'^\d+\.\s+',lines[i].strip()):
  para.append(lines[i].strip()); i+=1
 blocks.append(('para',' '.join(para)))

# Inject clear candidate metadata directly below the exact official title.
first_title_seen=False
for bi,(kind,val) in enumerate(blocks):
 # If a short heading/intro preamble directly precedes a wide matrix, rotate
 # before the preamble so it cannot become an orphan portrait page.
 upcoming_wide_table=False
 for ahead in blocks[bi+1:bi+4]:
  if ahead[0]=='table':
   upcoming_wide_table=len(ahead[1][0])>=6
   break
  if ahead[0] not in ('heading','para'):
   break
 if kind=='heading' and not current_landscape and upcoming_wide_table:
  switch_orientation(True)
 if kind=='heading' and val[0]==1 and not first_title_seen:
  title=val[1]
  p=doc.add_paragraph(style='Title'); add_inline(p,title,20,base_color='000000'); first_title_seen=True
  meta=doc.add_paragraph(); meta.paragraph_format.space_after=Pt(10)
  r=meta.add_run('文件編號 0007  |  官方要求 I01-E04  |  版本 R12  |  候選文件  |  SAMPLE / SYNTHETIC'); set_run_font(r,9.5,bold=True,color=NAVY)
  continue
 if kind=='heading':
  level,title=val
  if title.startswith('4. 加權比較') or title.startswith('8. 跨文件責任及負向近鄰'):
   doc.add_page_break()
  if current_landscape and level>0 and not upcoming_wide_table: switch_orientation(False)
  add_para(title,'Heading 1' if level==1 else 'Heading 2' if level==2 else 'Heading 3')
 elif kind=='para':
  if current_landscape and not upcoming_wide_table: switch_orientation(False)
  add_para(val)
 elif kind=='bullet':
  if current_landscape and not upcoming_wide_table:switch_orientation(False)
  add_para(val,'List Bullet')
 elif kind=='number':
  if current_landscape and not upcoming_wide_table:switch_orientation(False)
  add_para(val,'List Number')
 elif kind=='table':
  if val and val[0] and val[0][0]=='分析欄位':
   doc.add_page_break()
  add_table(val)

# Metadata and export settings.
doc.core_properties.title='課程產品及流程檢討分析報告'
doc.core_properties.subject='TTQS ONE 0007 / I01-E04；SAMPLE / SYNTHETIC 候選文件'
doc.core_properties.author='屏東縣原始點關懷協會（候選示範文件；非正式核准）'
doc.core_properties.keywords='TTQS, 0007, I01-E04, SAMPLE, SYNTHETIC, CANDIDATE'
doc.core_properties.comments='Candidate R12; not formally accepted; synthetic demonstration only.'
doc.save(docx_out)

with open(register,newline='',encoding='utf-8-sig') as fh: reg=list(csv.DictReader(fh))
packet=json.loads(source_packet.read_text(encoding='utf-8-sig'))
receipt={
 'project_id':'TTQS_ONE','doc_id':'0007','requirement_id':'I01-E04','revision':'R12',
 'candidate_state':'CANDIDATE_SAMPLE_SYNTHETIC_NOT_ACCEPTED','build_status':'BUILT_NOT_VALIDATED',
 'candidate_path':str(docx_out),'candidate_sha256':sha(docx_out),'candidate_bytes':docx_out.stat().st_size,
 'source_original_path':str(original),'source_original_sha256':sha(original),'word_source_path':str(normalized_out),'word_source_sha256':sha(normalized_out),
 'synthetic_register_path':str(register),'synthetic_register_sha256':sha(register),'synthetic_register_rows':len(reg),
 'official_pdf_sha256':packet['official_source_sha256'],'blueprint_file_sha256':packet['blueprint_sha256'],'blueprint_row_sha256':packet['blueprint_row_sha256'],
 'upstream_0003_sha256':'7A21CF2FB6E8FD3600ABA6733D176AD45C4A7003957F7A8875F7CB6095257423','upstream_0004_sha256':'ED66B95A1856751B559702FE7B944595EC827651435B652A380725B9CF8A219D',
 'normalizations':transforms,'parsed_markdown_tables':sum(1 for x in blocks if x[0]=='table'),'paragraph_blocks':sum(1 for x in blocks if x[0] in ('para','bullet','number')),'section_count':len(doc.sections),
 'excluded_internal_production_note':'Production-only Word formatting note and following paragraph excluded from evaluator-facing DOCX; original worker draft retained unchanged',
 'font_policy':{'body_pt':10.5,'table_pt':9.5,'font':'Microsoft JhengHei'},'render_status':'PENDING','independent_semantic_status':'PENDING','formal_acceptance':'NOT_ASSERTED'
}
receipt_out.write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(receipt,ensure_ascii=False,indent=2))
