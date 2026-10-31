from __future__ import annotations
from decimal import Decimal
from pathlib import Path
import csv, hashlib, json, re, zipfile
from pypdf import PdfReader
from docx import Document
from docx.oxml.ns import qn

BASE = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0013_I02-E04_LUNA_TEXT_DRAFT_R01_20261029")
DOCX = BASE / "candidate" / "0013_I02-E04_R11.docx"
SRC = BASE / "worker" / "luna_subagent_draft_R01.md"
PACKET = BASE / "source_packet_R01.json"
REGISTER = BASE / "qa" / "0013_R11_SYNTHETIC_REPLACEMENT_REGISTER.csv"
VISUAL = BASE / "qa" / "visual_qa_R11.json"
BUILD = BASE / "qa" / "BUILD_RECEIPT_R11.json"
PDF = BASE / "qa" / "render_R11_instrumented" / "0013_I02-E04_R11_WORD_NATIVE.pdf"
BP = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\CONTROL\BLUEPRINT_CURRENT.jsonl")
OFFICIAL_SHA = "805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709"
OFFICIAL = next(p for p in Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\SOURCES").glob("*.pdf")
                if hashlib.sha256(p.read_bytes()).hexdigest() == OFFICIAL_SHA)
OUT = BASE / "qa" / "technical_qa_R11.json"

def sha_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()

def sha_file(path: Path) -> str:
    return sha_bytes(path.read_bytes())

def normalize(text: str) -> str:
    return "".join(c.casefold() for c in text if c.isalnum() or "\u3400" <= c <= "\u9fff")

doc = Document(DOCX)
all_text = "\n".join([p.text for p in doc.paragraphs] +
                     [c.text for t in doc.tables for row in t.rows for c in row.cells])
norm_text = normalize(all_text)
source_packet = json.loads(PACKET.read_text(encoding="utf-8-sig"))
bp_rows = BP.read_text(encoding="utf-8-sig").splitlines()
bp_row_bytes = next(x.encode("utf-8") for x in bp_rows
                    if (lambda r: r.get("deliverable_id") == "0013" and
                        r.get("requirement_id") == "I02-E04")(json.loads(x)))
bp_row = json.loads(bp_row_bytes)
official_text = bp_row["official_requirement_text"]
official_page1 = PdfReader(str(OFFICIAL)).pages[0].extract_text() or ""
table19 = doc.tables[18]
expected_ids = {f"SYN-0013-F{i:03d}" for i in list(range(1, 28)) + list(range(101, 141))}
actual_ids = [row.cells[0].text.strip() for row in table19.rows[1:]]
doc_crosswalk = {
    row.cells[0].text.strip(): (row.cells[1].text.strip(), row.cells[2].text.strip(), row.cells[4].text.strip())
    for row in table19.rows[1:]
}
with REGISTER.open(encoding="utf-8-sig", newline="") as f:
    csv_rows = list(csv.reader(f))
csv_crosswalk = {row[0].strip(): (row[1].strip(), row[2].strip(), row[4].strip())
                 for row in csv_rows[1:] if len(row) >= 5}

course_tables = doc.tables[6:10]
course_hours = []
course_capacity = []
course_urls = []
course_shapes_ok = True
for table in course_tables:
    course_shapes_ok = course_shapes_ok and len(table.rows) == 11 and len(table.columns) == 3
    values = {row.cells[0].text.strip(): row.cells[1].text.strip()
              for row in table.rows[1:]}
    course_hours.append(int(re.search(r"\d+", values["淨授課時數"]).group()))
    course_capacity.append(int(re.search(r"\d+", values["規劃名額"]).group()))
    course_urls.append(values["報名通路"])

header = table19.rows[0]
header_shades = []
header_colors = []
for cell in header.cells:
    shd = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
    header_shades.append(shd.get(qn("w:fill")) if shd is not None else None)
    for p in cell.paragraphs:
        for run in p.runs:
            rgb = run.font.color.rgb
            header_colors.append(str(rgb).upper() if rgb else "")

sample_register_ids = sorted(re.findall(r"SYN-0013-F\d{3}", all_text))
field_population_math_ok = (Decimal(4) * Decimal(9) == Decimal(36) and
                            (Decimal(35) / Decimal(36) * 100).quantize(Decimal("0.1")) == Decimal("97.2") and
                            Decimal(36) / Decimal(36) * 100 == Decimal(100))
build = json.loads(BUILD.read_text(encoding="utf-8"))
semantic = json.loads((BASE / "qa" / "semantic_review_R11.json").read_text(encoding="utf-8"))
render_status = json.loads((BASE / "qa" / "render_R11_instrumented" / "R11_RENDER_STATUS.json").read_text(encoding="utf-8"))
assert semantic.get("candidate_sha256") == sha_file(DOCX)
assert semantic.get("verdict") == "PASS_SCOPED_SAMPLE_SEMANTIC"
assert render_status.get("candidate_sha256") == sha_file(DOCX)
assert render_status.get("state") == "WORD_EXPORT_TIMEOUT_NO_PDF"

checks = {
    "docx_zip_integrity": zipfile.is_zipfile(DOCX),
    "python_docx_opens": True,
    "candidate_sha_bound_in_build_receipt": build.get("artifact_sha256") == sha_file(DOCX),
    "candidate_identity_0013_I02_E04_R11": build.get("document_id") == "0013" and build.get("requirement_id") == "I02-E04" and build.get("artifact_version") == "R11",
    "official_requirement_pdf_sha_exact": sha_file(OFFICIAL) == OFFICIAL_SHA,
    "official_requirement_text_matches_blueprint_and_pdf": normalize(official_text) in normalize(official_page1) and normalize(official_text) in norm_text,
    "blueprint_file_and_unique_row_sha_exact": sha_file(BP) == "406d0495365d0970dcf8113a42842663b8a54372a4e204cc492b7c3f8156ef3e" and sha_bytes(bp_row_bytes) == "b529f3c13e964b6edbec124158d3559df988b401ce3b97fd1984e90b0582d35b",
    "source_packet_and_worker_binding_exact": sha_file(PACKET) == "b87440b249ec78c6bbe69ff001fe9d7dbb5060d65f27fe61d21ec48cdb5bfa63" and source_packet.get("worker_source_sha256") == sha_file(SRC),
    "negative_neighbor_ids_distinguished": all(x in all_text for x in ["0012", "I02-E03", "0014", "I02-E05"]),
    "twenty_tables_and_table19_68x5": len(doc.tables) == 20 and len(table19.rows) == 68 and len(table19.columns) == 5,
    "table19_has_exact_67_unique_fact_ids": len(actual_ids) == 67 and len(set(actual_ids)) == 67 and set(actual_ids) == expected_ids,
    "table19_all_locators_precise_and_roles_allowed": all(loc and ("§" in loc or "表" in loc) and role in {"R01", "R02", "R03", "R04", "R05"} for _, loc, role in doc_crosswalk.values()) and len(doc_crosswalk) == 67,
    "csv_crosswalk_matches_docx_exactly_including_fact_value": csv_crosswalk == doc_crosswalk and len(csv_crosswalk) == 67,
    "publication_gate_order_non_circular": "公開頁面讀回屬發布後驗收，不列為發布前條件" in doc.paragraphs[93].text and doc.paragraphs[93].text.index("公開頁面讀回") < doc.paragraphs[93].text.index("上架後由非編輯者"),
    "release_gate_heading_style_preserved": doc.paragraphs[93].runs[0].text == "發布前：" and doc.paragraphs[93].runs[0].bold is True and len(doc.paragraphs[93].runs) > 1 and doc.paragraphs[93].runs[1].bold is not True,
    "version_metadata_is_R11": "R11" in doc.paragraphs[4].text and doc.tables[19].rows[2].cells[1].text.strip() == "R11",
    "course_examples_complete": len(course_tables) == 4 and course_shapes_ok and all(c.text.strip() for t in course_tables for row in t.rows for c in row.cells),
    "course_hours_total_30": sum(course_hours) == 30,
    "course_capacity_total_74": sum(course_capacity) == 74,
    "all_four_sample_urls_reserved_invalid": len(course_urls) == 4 and all(".invalid" in u for u in course_urls),
    "field_population_math_35_36_97_2_and_36_36_100": field_population_math_ok and all(x in all_text for x in ["35/36", "97.2%", "36/36", "100%"]),
    "route_usability_separate_zero_of_four_and_hold": "0/4" in all_text and ".invalid" in all_text and "HOLD" in all_text and "未執行真實網路連線測試" in all_text,
    "sample_only_and_no_formal_acceptance_claim": "CANDIDATE_SAMPLE_SYNTHETIC_NOT_ACCEPTED" in all_text and "未提供；不主張已核准" in all_text and build.get("formal_acceptance") == "UNKNOWN",
    "table19_header_high_contrast_consistent": len(header_shades) == 5 and len(set(header_shades)) == 1 and all(c == "FFFFFF" for c in header_colors),
    "no_empty_table_cells_or_placeholder_terms": all(c.text.strip() for t in doc.tables for row in t.rows for c in row.cells) and not any(x in all_text for x in ["待填", "空白表", "保持空白", "會後填寫"]),
    "body_font_at_least_10_5pt": min((r.font.size.pt for p in doc.paragraphs for r in p.runs if r.font.size), default=99) >= 10.5,
    "table_font_at_least_9_5pt": min((r.font.size.pt for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs for r in p.runs if r.font.size), default=99) >= 9.5,
}

receipt = {
    "receipt_id": "0013_I02-E04_R11_DETERMINISTIC_TECHNICAL_QA",
    "document_id": "0013",
    "requirement_id": "I02-E04",
    "candidate_state": "CANDIDATE_SAMPLE_SYNTHETIC_NOT_ACCEPTED",
    "artifact_path": str(DOCX),
    "artifact_sha256": sha_file(DOCX),
    "artifact_bytes": DOCX.stat().st_size,
    "worker_source_sha256": sha_file(SRC),
    "source_packet_sha256": sha_file(PACKET),
    "official_requirement_pdf_sha256": sha_file(OFFICIAL),
    "blueprint_file_sha256": sha_file(BP),
    "blueprint_row_sha256": sha_bytes(bp_row_bytes),
    "table_count": len(doc.tables),
    "table19_rows": len(table19.rows),
    "table19_columns": len(table19.columns),
    "synthetic_fact_ids": len(set(actual_ids)),
    "synthetic_register_rows": len(csv_crosswalk),
    "course_net_hours": course_hours,
    "course_hours_total": sum(course_hours),
    "course_capacity": course_capacity,
    "course_capacity_total": sum(course_capacity),
    "native_word_pdf_sha256": None,
    "native_word_pdf_pages": None,
    "visual_status": render_status.get("visual_verdict"),
    "semantic_review_sha256": sha_file(BASE / "qa" / "semantic_review_R11.json"),
    "semantic_status": semantic.get("verdict"),
    "formal_acceptance": "UNKNOWN",
    "checks": checks,
    "technical_status": "PASS" if all(checks.values()) else "FAIL"
}
OUT.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(receipt, ensure_ascii=True, indent=2))
if not all(checks.values()):
    raise SystemExit(1)
