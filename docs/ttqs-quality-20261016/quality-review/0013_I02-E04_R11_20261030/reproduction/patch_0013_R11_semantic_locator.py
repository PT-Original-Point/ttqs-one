from pathlib import Path
import csv, hashlib, json
from docx import Document

BASE = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0013_I02-E04_LUNA_TEXT_DRAFT_R01_20261029")
QA = BASE / "qa"
SRC = BASE / "candidate" / "0013_I02-E04_R10.docx"
OUT = BASE / "candidate" / "0013_I02-E04_R11.docx"
REG_SRC = QA / "0013_R10_SYNTHETIC_REPLACEMENT_REGISTER.csv"
REG_OUT = QA / "0013_R11_SYNTHETIC_REPLACEMENT_REGISTER.csv"
RECEIPT = QA / "content_preflight_R11.json"
EXPECTED_PARENT = "fb988f91f02653d408f67f94c5d24576b8ef27d2116ad32c99a49f48b6cdf514"
OLD = "§6.2–§6.5 表7–10「開課日期／期間」欄（2027 年 SAMPLE）"
NEW = "§6.2–§6.5 表7–10「日期與時段」欄（2027 年 SAMPLE）"

def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()

assert sha(SRC) == EXPECTED_PARENT, "R10 parent SHA changed"
assert not OUT.exists() and not REG_OUT.exists(), "R11 output already exists; inspect before overwrite"

doc = Document(SRC)
table = doc.tables[18]
found = 0
for row in table.rows[1:]:
    if row.cells[0].text.strip() == "SYN-0013-F002":
        cell = row.cells[2]
        assert cell.text.strip() == OLD, f"Unexpected F002 locator: {cell.text!r}"
        p = cell.paragraphs[0]
        if p.runs:
            p.runs[0].text = NEW
            for run in p.runs[1:]:
                run.text = ""
        else:
            p.add_run(NEW)
        found += 1
assert found == 1, f"Expected one F002 row, found {found}"

# Update only candidate revision labels in established front-matter fields.
for paragraph in [doc.paragraphs[4]]:
    for run in paragraph.runs:
        run.text = run.text.replace("R10", "R11")
for row in doc.tables[19].rows:
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.text = run.text.replace("R10", "R11")
doc.core_properties.title = "0013 I02-E04 年度訓練課程公開清單 R11 候選 SAMPLE SYNTHETIC"
doc.save(OUT)

with REG_SRC.open(encoding="utf-8-sig", newline="") as f:
    rows = list(csv.reader(f))
register_found = 0
for row in rows[1:]:
    if row and row[0].strip() == "SYN-0013-F002":
        assert row[2].strip() == OLD, f"Unexpected register F002 locator: {row[2]!r}"
        row[2] = NEW
        register_found += 1
assert register_found == 1, f"Expected one F002 register row, found {register_found}"
with REG_OUT.open("w", encoding="utf-8-sig", newline="") as f:
    csv.writer(f, lineterminator="\n").writerows(rows)

check = Document(OUT)
crosswalk = {r.cells[0].text.strip(): r.cells[2].text.strip() for r in check.tables[18].rows[1:]}
with REG_OUT.open(encoding="utf-8-sig", newline="") as f:
    register = {r[0].strip(): r[2].strip() for r in list(csv.reader(f))[1:] if len(r) >= 3}
assert crosswalk["SYN-0013-F002"] == register["SYN-0013-F002"] == NEW
assert crosswalk["SYN-0013-F013"] == register["SYN-0013-F013"]
assert len(crosswalk) == 67 and len(register) == 67

receipt = {
    "status": "PASS_SCOPED_CONTENT_PATCH",
    "document_id": "0013",
    "requirement_id": "I02-E04",
    "parent_candidate_sha256": sha(SRC),
    "candidate_path": "candidate/0013_I02-E04_R11.docx",
    "candidate_sha256": sha(OUT),
    "register_path": "qa/0013_R11_SYNTHETIC_REPLACEMENT_REGISTER.csv",
    "register_sha256": sha(REG_OUT),
    "changed_fact_id": "SYN-0013-F002",
    "old_locator": OLD,
    "new_locator": NEW,
    "crosswalk_rows": len(crosswalk),
    "register_rows": len(register),
    "F013_locator_unchanged_and_matching": True,
    "formal_acceptance": "UNKNOWN"
}
RECEIPT.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(receipt, ensure_ascii=False, indent=2))
