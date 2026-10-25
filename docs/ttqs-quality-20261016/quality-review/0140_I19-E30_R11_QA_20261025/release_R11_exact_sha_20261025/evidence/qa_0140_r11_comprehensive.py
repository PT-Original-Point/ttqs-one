from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENT
from decimal import Decimal, ROUND_HALF_UP
import csv, json, hashlib, zipfile, re

root=Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0140_I19-E30_CANDIDATE_R11_20261025")
src=Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0140_I19-E30_SOURCE_PREP_R03_20261025")
docx=root/"candidate"/"0140__訓練資訊系統使用普及率盤點與分析表_R11.docx"
build=root/"DOCX_BUILD_RECEIPT_0140_R11.json"
md=src/"worker"/"MANUSCRIPT_0140_R11.md"
dataset_path=src/"SYNTHETIC_USER_DATASET_0140_R01.csv"
register_path=src/"SYNTHETIC_REPLACEMENT_REGISTER_0140_R07.csv"
r07=Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0140_I19-E30_CANDIDATE_R07_20261025\0140__訓練資訊系統使用普及率盤點與分析表_R07.docx")
r08=Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0140_I19-E30_CANDIDATE_R08_20261025\candidate\0140__訓練資訊系統使用普及率盤點與分析表_R08.docx")
def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def pct(n,d): return (Decimal(n)*100/Decimal(d)).quantize(Decimal("0.01"),rounding=ROUND_HALF_UP)
build_data=json.loads(build.read_text(encoding="utf-8"))
doc=Document(str(docx))
paragraphs=[p.text for p in doc.paragraphs]
tables=[[[" ".join(c.text.split()) for c in row.cells] for row in t.rows] for t in doc.tables]
all_table_text=["\n".join(" | ".join(row) for row in table) for table in tables]
full="\n".join(paragraphs+all_table_text)
checks={}
def check(name,ok,evidence): checks[name]={"result":"PASS" if bool(ok) else "FAIL","evidence":evidence}

doc_sha=sha(docx)
check("docx_hash_matches_build_receipt",doc_sha==build_data.get("docx_sha256"),{"docx_sha256":doc_sha,"receipt_sha256":build_data.get("docx_sha256"),"bytes":docx.stat().st_size})
check("docx_zip_integrity",zipfile.ZipFile(docx).testzip() is None,"ZIP CRC test")
check("python_docx_open",len(doc.paragraphs)>0,{"paragraphs":len(doc.paragraphs),"tables":len(doc.tables),"inline_shapes":len(doc.inline_shapes)})
check("official_mapping",all(x in full for x in ["DOCX 0140","I19-E30","157 項清單索引 154"]),{"requirement":"DOCX 0140 -> I19-E30 -> 157 index 154; printed checkbox UNKNOWN"})
check("sample_real_formal_status",all(x in full for x in ["CANDIDATE","REAL = UNKNOWN / HOLD","NOT_GRANTED"]),{"candidate":"CANDIDATE","real":"UNKNOWN/HOLD","formal_acceptance":"NOT_GRANTED"})
check("targeted_instructor_distribution", "2023 與 2024 示範資料各有 1 名活躍講師只完成 1/2 任務" in full and "2025 示範資料中 5 名活躍講師均完成 2/2 任務" in full, "Section 9.2 states all three years without unsupported cause")
check("targeted_learner_distribution", "2023、2024、2025 示範資料分別有 2、3、3 名活躍學員完成 2/3 任務" in full and "僅缺 COURSE_EVALUATION" in full, "Section 9.2 states 2023=2, 2024=3, 2025=3")
status_rows=[]
for table in tables:
  for row in table:
    joined=" | ".join(row)
    if "DOCX 0140 是否仍對 I19-E30" in joined or "頁碼、表格可讀性與附件是否齊全" in joined:
      status_rows.append(row)
check("human_facing_status_cells",len(status_rows)==2 and any("本文件為示範，不代表正式評核結果" in " | ".join(r) for r in status_rows) and any("使用前須逐頁檢視頁碼、表格可讀性與附件" in " | ".join(r) and "由權責人員核定" in " | ".join(r) for r in status_rows),status_rows)
check("pipeline_jargon_removed_from_checklist",not any(any(term in " | ".join(r) for term in ["SOURCE 已對照","正式接受未授予","另一位覆核者"]) for r in status_rows),status_rows)
with dataset_path.open("r",encoding="utf-8-sig",newline="") as f: rows=list(csv.DictReader(f))
check("canonical_dataset_80_rows",len(rows)==80,{"rows":len(rows),"dataset_sha256":sha(dataset_path)})
groups={}
role_groups={}
incomplete={}
for r in rows:
  y=r["period_year"]; role=r["role_code"]
  g=groups.setdefault(y,{"eligible":0,"active":0,"complete":0})
  for k,c in [("eligible","eligible_unique_user"),("active","active_user"),("complete","all_required_tasks_complete")]: g[k]+=int(r[c])
  rg=role_groups.setdefault((y,role),{"eligible":0,"active":0,"complete":0})
  for k,c in [("eligible","eligible_unique_user"),("active","active_user"),("complete","all_required_tasks_complete")]: rg[k]+=int(r[c])
  if r["active_user"]=="1" and r["all_required_tasks_complete"]=="0":
    incomplete[(y,role)]=incomplete.get((y,role),0)+1
expected={"2023":["60.00%","45.00%","75.00%"],"2024":["78.57%","64.29%","81.82%"],"2025":["87.50%","78.13%","89.29%"]}
observed={y:[f"{pct(g['active'],g['eligible']):.2f}%",f"{pct(g['complete'],g['eligible']):.2f}%",f"{pct(g['complete'],g['active']):.2f}%"] for y,g in groups.items()}
counts_ok=(len(rows)==80 and observed==expected and [incomplete.get((y,"INSTRUCTOR"),0) for y in ("2023","2024","2025")]==[1,1,0] and [incomplete.get((y,"LEARNER"),0) for y in ("2023","2024","2025")]==[2,3,3])
check("annual_math_and_incomplete_cases_recomputed",counts_ok,{"annual_rates":observed,"incomplete_active_by_role_year":{role:{y:incomplete.get((y,role),0) for y in ("2023","2024","2025")} for role in ("INSTRUCTOR","LEARNER")}})
register_rows=list(csv.DictReader(register_path.open(encoding="utf-8-sig",newline="")))
register_lines=len(register_path.read_bytes().splitlines())
data_map={r["case_id"]:r for r in rows}
register_by_id={}
register_control=[]
for rr in register_rows:
  if rr["record_class"]=="SAMPLE/SYNTHETIC_CONTROL":
    register_control.append(rr); continue
  syn={}
  for token in rr["synthetic_values"].split(";"):
    if "=" in token:
      k,v=token.split("=",1); syn[k]=v
  register_by_id[rr["record_id"]]=syn
register_ok=(len(register_rows)==82 and register_lines==83 and set(register_by_id)==set(data_map) and len(register_control)==2)
if register_ok:
  for cid,srcrow in data_map.items():
    v=register_by_id[cid]
    expected_fields={"year":srcrow["period_year"],"role":srcrow["role_code"],"eligible":srcrow["eligible_unique_user"],"required":srcrow["required_task_count"],"actions_done":srcrow["substantive_actions_completed"],"active":srcrow["active_user"],"workflow_complete":srcrow["all_required_tasks_complete"],"log_ref":srcrow["synthetic_log_ref"]}
    if any(v.get(k)!=val for k,val in expected_fields.items()): register_ok=False; break
check("register_80_ids_and_controls_match",register_ok,{"dictreader_records":len(register_rows),"physical_lines":register_lines,"dataset_ids":len(data_map),"matched_ids":len(register_by_id),"control_rows":len(register_control),"register_sha256":sha(register_path)})
check("register_header_description_corrected", "82 data records" in (root/"README_R11.md").read_text(encoding="utf-8") and "83 physical lines" in (root/"README_R11.md").read_text(encoding="utf-8"),"82 data records + header = 83 physical lines")
check("chart_embedded_and_captioned",len(doc.inline_shapes)>=1 and "合格應使用者全任務完成率" in full,{"inline_shapes":len(doc.inline_shapes)})
check("R07_unchanged",sha(r07)=="7343688549e213f26ab22e58436a8d6915a679b72d75803a3c25566f94829ebe",{"sha256":sha(r07)})
check("R08_unchanged",sha(r08)=="a6feeb119c14731b63c92833411721e6613db5a011a26f8e38334de9ecacc970",{"sha256":sha(r08)})
r09=Path("E:/TTQS/TTQS_ONE_CODEX_RUNTIME/WORK/QUALITY_REBUILD_20261016/0140_I19-E30_CANDIDATE_R09_20261025/candidate/0140__訓練資訊系統使用普及率盤點與分析表_R09.docx")
check("R09_unchanged",sha(r09)=="49afdcfa11b6f91952590995b4c0152d76271b2169b2a4433bb08bbbba46db4e",{"sha256":sha(r09)})
check("markdown_italics_render_without_literal_markup",full.count("*")==full.count("SYN-*"),{"asterisks_in_extracted_docx":full.count("*"),"legitimate_SYN_wildcards":full.count("SYN-*"),"visible_markup_absent":full.count("*")==full.count("SYN-*")})
check("no_unsupported_real_or_approval_claims","REAL = UNKNOWN / HOLD" in full and "本文件未作正式個資法遵審查" in full and "沒有核准人" in full,"Sample/real and no-approval statements present")
landscape_sections=sum(1 for section in doc.sections if section.orientation==WD_ORIENT.LANDSCAPE)
check("structure_matches_R07_baseline",len(doc.tables)==18 and sum(len(t.rows) for t in doc.tables)==277 and len(doc.sections)==8 and landscape_sections==4,{"tables":len(doc.tables),"rows":sum(len(t.rows) for t in doc.tables),"sections":len(doc.sections),"landscape_sections":landscape_sections})
font_sizes=[]
for p in doc.paragraphs:
  for run in p.runs:
    if run.font.size: font_sizes.append(run.font.size.pt)
for t in doc.tables:
  for row in t.rows:
    for cell in row.cells:
      for p in cell.paragraphs:
        for run in p.runs:
          if run.font.size: font_sizes.append(run.font.size.pt)
minfont=min(font_sizes) if font_sizes else None
check("minimum_explicit_font_at_least_9_5pt",minfont is not None and minfont>=9.5,{"min_pt":minfont,"run_count":len(font_sizes)})
result={
 "work_id":"TTQS-0140-I19-E30-R11-COMPREHENSIVE-QA","doc_id":"0140","requirement_id":"I19-E30","official_checklist_index":154,"printed_checkbox_serial":"UNKNOWN",
 "candidate_status":"CANDIDATE","real_data_status":"UNKNOWN/HOLD","formal_acceptance":"NOT_GRANTED",
 "candidate_sha256":doc_sha,"candidate_bytes":docx.stat().st_size,"parent_R08_sha256":sha(r08),
 "checks":checks,"status":"PASS" if all(x["result"]=="PASS" for x in checks.values()) else "FAIL",
 "scope_limitations":["No visual inspection yet.","No independent semantic review of this R11 SHA yet.","No formal Human acceptance or REAL evidence claim."]
}
out=root/"receipts"/"R11_COMPREHENSIVE_QA_0140.json"
out.parent.mkdir(parents=True, exist_ok=True)
out.write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
manifest_path=root/"MANIFEST_R11.json"
manifest=json.loads(manifest_path.read_text(encoding="utf-8"))
manifest.update({"docx_sha256":doc_sha,"docx_bytes":docx.stat().st_size,"build_receipt_sha256":sha(build),"comprehensive_qa_sha256":sha(out),"comprehensive_qa_status":result["status"],"visual_review":"PENDING","independent_semantic_review":"PENDING","formal_acceptance":"NOT_GRANTED"})
manifest_path.write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
print(json.dumps({"qa_path":str(out),"qa_sha256":sha(out),"candidate_sha256":doc_sha,"status":result["status"],"checks":{k:v["result"] for k,v in checks.items()},"scope_limitations":result["scope_limitations"]},ensure_ascii=False,indent=2))