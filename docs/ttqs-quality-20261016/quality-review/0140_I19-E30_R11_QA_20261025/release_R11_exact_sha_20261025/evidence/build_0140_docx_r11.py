from pathlib import Path
import re, csv, json, hashlib, html
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0140_I19-E30_SOURCE_PREP_R03_20261025")
WORKER = ROOT / "worker"
OUTDIR = ROOT.parent / "0140_I19-E30_CANDIDATE_R11_20261025"
OUTDOCX = OUTDIR / "candidate" / "0140__訓練資訊系統使用普及率盤點與分析表_R11.docx"
MANUSCRIPT = WORKER / "MANUSCRIPT_0140_R11.md"
DATA = ROOT / "SYNTHETIC_USER_DATASET_0140_R01.csv"
SUMMARY = ROOT / "SAMPLE_PERIOD_SUMMARY_0140_R01.csv"
REGISTER = ROOT / "SYNTHETIC_REPLACEMENT_REGISTER_0140_R07.csv"
CHART = OUTDIR / "candidate" / "SAMPLE_RATE_TREND_0140_R11.png"
RECEIPT = OUTDIR / "DOCX_BUILD_RECEIPT_0140_R11.json"
FONT = r"C:\Windows\Fonts\msjh.ttc"
NAVY, PALE, PALE2, INK, MUTED, BORDER = "243B53", "EDF3F8", "F7F9FB", "1F2933", "52606D", "CBD2D9"
tick = chr(96)

if not all(p.exists() for p in [MANUSCRIPT, DATA, SUMMARY, REGISTER]):
    raise FileNotFoundError("Required source packet files are missing.")
if OUTDIR.exists() and OUTDOCX.exists():
    raise FileExistsError(f"Refusing to overwrite an existing candidate: {OUTDOCX}")
OUTDIR.mkdir(parents=True, exist_ok=True)
(OUTDIR / "candidate").mkdir(parents=True, exist_ok=True)

def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def font(size):
    try:
        return ImageFont.truetype(FONT, size=size, index=0)
    except Exception:
        return ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", size=size)

def make_chart():
    with SUMMARY.open("r", encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    years = [r["period_year"] for r in rows]
    adoption = [float(r["adoption_rate_pct"]) for r in rows]
    complete = [float(r["completion_rate_of_eligible_pct"]) for r in rows]
    im = Image.new("RGB", (1600, 850), "white")
    d = ImageDraw.Draw(im)
    titlef, labelf, smallf, valuef = font(38), font(27), font(22), font(24)
    d.text((120, 40), "資訊系統普及率與合格應使用者全任務完成率", font=titlef, fill="#172B4D")
    d.text((120, 92), "SAMPLE / SYNTHETIC 示範數據；非協會實績", font=labelf, fill="#8A4B08")
    left, right, top, bottom = 190, 1450, 190, 690
    for pct in range(0, 101, 20):
        y = bottom - (bottom-top)*pct/100
        d.line((left, y, right, y), fill="#D9E2EC", width=2)
        d.text((95, y-14), f"{pct}%", font=smallf, fill="#52606D")
    xs = [left+100, (left+right)//2, right-100]
    for x, year in zip(xs, years):
        d.text((x-40, bottom+25), year, font=labelf, fill="#243B53")
    for label, vals, color, offset in [
        ("普及率", adoption, "#2F6690", -45),
        ("合格者全任務完成率", complete, "#D97706", 16),
    ]:
        pts = [(x, bottom-(bottom-top)*v/100) for x,v in zip(xs, vals)]
        d.line(pts, fill=color, width=7)
        for (x,y),v in zip(pts, vals):
            d.ellipse((x-10,y-10,x+10,y+10), fill=color, outline="white", width=3)
            labelv=f"{v:.2f}%"
            box=d.textbbox((0,0),labelv,font=valuef)
            d.text((x-(box[2]-box[0])/2,y+offset),labelv,font=valuef,fill=color)
    lx, ly = 330, 795
    for label,color in [("普及率","#2F6690"),("合格者全任務完成率","#D97706")]:
        d.line((lx,ly+12,lx+55,ly+12),fill=color,width=7)
        d.ellipse((lx+20,ly+2,lx+40,ly+22),fill=color)
        d.text((lx+70,ly-3),label,font=smallf,fill="#243B53")
        lx += 510
    im.save(CHART, format="PNG", optimize=True)

def set_run_font(run, size=10.5, bold=None, color=INK, italic=None, font_name="Microsoft JhengHei"):
    run.font.name=font_name; run.font.size=Pt(size); run.font.color.rgb=RGBColor.from_string(color)
    if bold is not None: run.bold=bold
    if italic is not None: run.italic=italic
    rpr=run._element.get_or_add_rPr()
    rfonts=rpr.rFonts
    if rfonts is None:
        rfonts=OxmlElement("w:rFonts"); rpr.insert(0,rfonts)
    rfonts.set(qn("w:ascii"),"Aptos"); rfonts.set(qn("w:hAnsi"),"Aptos")
    rfonts.set(qn("w:eastAsia"),font_name); rfonts.set(qn("w:cs"),font_name)

def format_paragraph(p, before=0, after=5, line=1.13, keep=False):
    pf=p.paragraph_format; pf.space_before=Pt(before); pf.space_after=Pt(after); pf.line_spacing=line
    if keep: pf.keep_with_next=True
    pf.widow_control=True

def clean_math(text):
    x=html.unescape(text)
    # Handle nested LaTeX text fractions before stripping remaining commands.
    special=[
      (r"\frac{\text{Active Unique Users}}{\text{Eligible Unique Users}}","Active Unique Users ÷ Eligible Unique Users"),
      (r"\frac{\text{Users Completing All Required Tasks}}{\text{Eligible Unique Users}}","Users Completing All Required Tasks ÷ Eligible Unique Users"),
      (r"\frac{\text{Users Completing All Required Tasks}}{\text{Active Unique Users}}","Users Completing All Required Tasks ÷ Active Unique Users")]
    for old,new in special: x=x.replace(old,new)
    x=re.sub(r"\\frac\{([^{}]+)\}\{([^{}]+)\}",r"(\1 ÷ \2)",x)
    for a,b in [(r"\text{",""),(r"\mathbf{",""),(r"\mathrm{",""),(r"\div"," ÷ "),(r"\times"," × "),(r"\rightarrow"," → "),(r"\to"," → "),(r"\%","%"),(r"\_","_"),(r"\left",""),(r"\right","")]:
        x=x.replace(a,b)
    x=x.replace(chr(92),"").replace("{","").replace("}","").replace("$","")
    x=re.sub(r"\s+"," ",x)
    return x.strip()

def add_inline(p, raw, base_size=10.5, base_color=INK, base_bold=False, italic=False):
    raw=clean_math(raw)
    raw=re.sub(r"<br\s*/?>","\n",raw,flags=re.I).replace(r"\_","\uFF3F")
    code_re=re.escape(tick)
    italic_re=r"(?<!\*)\*(?!\*)(?:[^\s*\n](?:[^*\n]*?[^\s*\n])?)\*(?!\*)"
    pat=re.compile(r"("+code_re+r"[^"+code_re+r"]+"+code_re+r"|\*\*.*?\*\*|"+italic_re+r")")
    for chunk in pat.split(raw):
        if not chunk: continue
        if chunk.startswith("**") and chunk.endswith("**"):
            run=p.add_run(chunk[2:-2]); set_run_font(run,base_size,True,base_color,italic)
        elif chunk.startswith("*") and chunk.endswith("*") and not chunk.startswith("**"):
            run=p.add_run(chunk[1:-1]); set_run_font(run,base_size,base_bold,base_color,True)
        elif chunk.startswith(tick) and chunk.endswith(tick):
            run=p.add_run(chunk[1:-1]); set_run_font(run,base_size,None,"314E68",italic,"Aptos")
        else:
            parts=chunk.split("\n")
            for j,part in enumerate(parts):
                if part:
                    run=p.add_run(part); set_run_font(run,base_size,base_bold,base_color,italic)
                if j<len(parts)-1: p.add_run().add_break()

def clean_heading(s):
    s=re.sub(r"^#{1,6}\s*", "", s)
    s=re.sub(r"[*_\x60]+", "", s)
    return re.sub(r"\s+"," ",s).strip()

def parse_row(line):
    x=line.strip()
    if x.startswith("|"): x=x[1:]
    if x.endswith("|"): x=x[:-1]
    return [c.strip() for c in x.split("|")]

def is_sep(row):
    return bool(row) and all(re.fullmatch(r":?-{3,}:?",c.replace(" ","")) for c in row)

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr()
    shd=tcPr.find(qn("w:shd"))
    if shd is None: shd=OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"),fill)

def margins(cell, top=75, start=80, bottom=75, end=80):
    tcPr=cell._tc.get_or_add_tcPr(); node=tcPr.first_child_found_in("w:tcMar")
    if node is None: node=OxmlElement("w:tcMar"); tcPr.append(node)
    for name,val in [("top",top),("start",start),("bottom",bottom),("end",end)]:
        item=node.find(qn("w:"+name))
        if item is None: item=OxmlElement("w:"+name); node.append(item)
        item.set(qn("w:w"),str(val)); item.set(qn("w:type"),"dxa")

def repeat_header(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement("w:tblHeader"); el.set(qn("w:val"),"true"); trPr.append(el)

def borders(table):
    tblPr=table._tbl.tblPr; b=tblPr.first_child_found_in("w:tblBorders")
    if b is None: b=OxmlElement("w:tblBorders"); tblPr.append(b)
    for edge in ["top","left","bottom","right","insideH","insideV"]:
        node=b.find(qn("w:"+edge))
        if node is None: node=OxmlElement("w:"+edge); b.append(node)
        node.set(qn("w:val"),"single"); node.set(qn("w:sz"),"5"); node.set(qn("w:space"),"0"); node.set(qn("w:color"),BORDER)

def configure_section(sec, landscape):
    sec.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    if landscape:
        sec.page_width=Inches(11.69); sec.page_height=Inches(8.27)
        sec.top_margin=Inches(.48); sec.bottom_margin=Inches(.48); sec.left_margin=Inches(.55); sec.right_margin=Inches(.55)
    else:
        sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
        sec.top_margin=Inches(.62); sec.bottom_margin=Inches(.62); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)

def switch_orientation(doc,current,landscape):
    if current==landscape: return current
    sec=doc.add_section(WD_SECTION.NEW_PAGE); configure_section(sec,landscape)
    return landscape

def add_table(doc,rows,landscape,current,font_size=9.5):
    current=switch_orientation(doc,current,landscape)
    if not rows: return current
    n=len(rows[0]); tbl=doc.add_table(rows=1,cols=n); tbl.alignment=WD_TABLE_ALIGNMENT.CENTER; tbl.autofit=False; tbl.style="Table Grid"; borders(tbl)
    avail=(11.69-1.10) if landscape else (8.27-1.44)
    preset={
      2:[.28,.72],3:[.18,.41,.41],4:[.18,.30,.30,.22],
      5:[.16,.22,.19,.21,.22],6:[.14,.18,.23,.18,.13,.14],
      7:[.14,.17,.15,.15,.13,.13,.13],
      8:[.12,.19,.11,.11,.11,.11,.13,.12],
      9:[.10,.12,.12,.12,.12,.12,.10,.10,.10]
    }
    weights=preset.get(n,[1/n]*n); widths=[avail*w for w in weights]
    for c,w in zip(tbl.columns,widths): c.width=Inches(w)
    for ri,vals in enumerate(rows):
        row=tbl.rows[0] if ri==0 else tbl.add_row()
        for ci,val in enumerate(vals):
            cell=row.cells[ci]; cell.width=Inches(widths[ci]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(cell)
            p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ri==0 else WD_ALIGN_PARAGRAPH.LEFT; format_paragraph(p,0,0,1.0)
            add_inline(p,val,font_size,"FFFFFF" if ri==0 else INK,ri==0)
            if ri==0: shade(cell,NAVY)
            elif ri%2==0: shade(cell,PALE2)
        if ri==0: repeat_header(row)
        else:
            trPr=row._tr.get_or_add_trPr(); cant=OxmlElement("w:cantSplit"); trPr.append(cant)
    spacer=doc.add_paragraph(); format_paragraph(spacer,0,2,1.0)
    return current

def add_form_131(doc,rows,current):
    if len(rows[0])!=9: return add_table(doc,rows,True,current,9.5)
    p=doc.add_paragraph(); format_paragraph(p,4,2,1.0,True); add_inline(p,"表 13.1A 應使用者與任務設定",10.5,NAVY,True)
    first=[[r[i] for i in [0,1,2,3,4]] for r in rows]
    first[0][2]=first[0][2].replace("Admin/Inst/Learner","Admin/Inst/<br>Learner")
    current=add_table(doc,first,False,current,9.5)
    p=doc.add_paragraph(); format_paragraph(p,4,2,1.0,True); add_inline(p,"表 13.1B 實質操作與來源核對",10.5,NAVY,True)
    current=add_table(doc,[[r[i] for i in [0,5,6,7,8]] for r in rows],False,current,9.5)
    return current

def add_form_133(doc,rows,current):
    if len(rows[0])!=6: return add_table(doc,rows,True,current,9.5)
    current=switch_orientation(doc,current,False)
    p=doc.add_paragraph(); format_paragraph(p,4,2,1.0,True); add_inline(p,"表 13.3A 日期與異常描述",10.5,NAVY,True)
    current=add_table(doc,[[r[i] for i in [0,1,2]] for r in rows],False,current,9.5)
    p=doc.add_paragraph(); format_paragraph(p,4,2,1.0,True); add_inline(p,"表 13.3B 處置依據與覆核欄",10.5,NAVY,True)
    current=add_table(doc,[[r[i] for i in [0,1,3,4,5]] for r in rows],False,current,9.5)
    return current

def add_dataset(doc,current):
    with DATA.open("r",encoding="utf-8-sig",newline="") as f: source=list(csv.DictReader(f))
    if len(source)!=80: raise ValueError("Canonical SAMPLE dataset is not 80 user-year rows.")
    current=switch_orientation(doc,current,True)
    first=[["年度","SAMPLE user-year ID","角色","納入分母","必要任務數","預先定義的必要任務"]]
    second=[["年度","SAMPLE user-year ID","已完成必要任務","實質活躍","全任務完成","合成日誌索引"]]
    for r in source:
        first.append([r["period_year"],r["case_id"],r["role_label"],r["eligible_unique_user"],r["required_task_count"],r["required_workflow"].replace("|", "; ")])
        second.append([r["period_year"],r["case_id"],r["completed_action_codes"],r["active_user"],r["all_required_tasks_complete"],r["synthetic_log_ref"]])
    p=doc.add_paragraph(); format_paragraph(p,4,2,1.0,True); add_inline(p,"附錄 A1 80 筆 SAMPLE 使用者年度紀錄之納入母體與任務",10.5,NAVY,True)
    current=add_table(doc,first,True,current,9.5)
    p=doc.add_paragraph(); format_paragraph(p,6,2,1.0,True); add_inline(p,"附錄 A2 80 筆 SAMPLE 使用者年度紀錄之操作結果",10.5,NAVY,True)
    current=add_table(doc,second,True,current,9.5)
    return current,len(source)

def add_header_footer(doc):
    sec=doc.sections[0]
    sec.header_distance=Inches(.28); sec.footer_distance=Inches(.30)
    hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; format_paragraph(hp,0,0,1.0)
    add_inline(hp,"TTQS ONE  |  DOCX 0140  |  候選",8.1,MUTED)
    fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; format_paragraph(fp,0,0,1.0)
    r=fp.add_run("第 "); set_run_font(r,8.2,color=MUTED)
    fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),"PAGE"); fp._p.append(fld)
    r=fp.add_run(" 頁"); set_run_font(r,8.2,color=MUTED)

def read_table(lines,i):
    raw=[]
    while i<len(lines) and lines[i].strip().startswith("|"):
        row=parse_row(lines[i])
        if not is_sep(row): raw.append(row)
        i+=1
    if raw:
        n=max(map(len,raw)); raw=[r+[""]*(n-len(r)) for r in raw]
    return raw,i

def upcoming_wide_table(lines,start):
    for j in range(start+1,len(lines)):
        s=lines[j].strip()
        if s.startswith("#"): return False
        if s.startswith("|"):
            row=parse_row(s)
            if is_sep(row): continue
            return len(row)>=6
    return False

def main():
    make_chart()
    text=MANUSCRIPT.read_text(encoding="utf-8"); lines=text.splitlines()
    doc=Document(); configure_section(doc.sections[0],False)
    normal=doc.styles["Normal"]; normal.font.name="Aptos"; normal.font.size=Pt(10.5); normal.font.color.rgb=RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft JhengHei"); normal._element.rPr.rFonts.set(qn("w:ascii"),"Aptos")
    for name,size in [("Title",22),("Heading 1",15.2),("Heading 2",11.8),("Heading 3",10.8)]:
        sty=doc.styles[name]; sty.font.name="Microsoft JhengHei"; sty.font.size=Pt(size); sty.font.bold=True; sty.font.color.rgb=RGBColor.from_string(NAVY)
        sty._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft JhengHei"); sty._element.rPr.rFonts.set(qn("w:ascii"),"Aptos")
        sty.paragraph_format.keep_with_next=True
    title_style_ppr=doc.styles["Title"]._element.pPr
    if title_style_ppr is not None:
        title_border=title_style_ppr.find(qn("w:pBdr"))
        if title_border is not None: title_style_ppr.remove(title_border)
    p=doc.add_paragraph(style="Title")
    title_ppr=p._p.get_or_add_pPr(); direct_border=title_ppr.find(qn("w:pBdr"))
    if direct_border is not None: title_ppr.remove(direct_border)
    format_paragraph(p,0,4,1.0,True); add_inline(p,lines[0].lstrip("# ").strip(),22,NAVY,True)
    p=doc.add_paragraph(); format_paragraph(p,0,5,1.0,True); add_inline(p,"DOCX 0140  |  官方指標 I19-E30  |  157 項清單索引 154",10.5,MUTED,True)
    status=doc.add_table(rows=1,cols=3); status.alignment=WD_TABLE_ALIGNMENT.CENTER; status.autofit=False
    for cell,label,w in zip(status.rows[0].cells,["CANDIDATE","SAMPLE / SYNTHETIC","REAL = UNKNOWN / HOLD  ·  NOT_GRANTED"],[1.65,2.20,2.90]):
        cell.width=Inches(w); shade(cell,PALE); margins(cell,100,100,100,100); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; format_paragraph(p,0,0,1.0); add_inline(p,label,9.5,NAVY,True)
    spacer=doc.add_paragraph(); format_paragraph(spacer,0,2,1.0)
    current=False; heading_context=""; appendix=False; i=9
    while i<len(lines):
        line=lines[i]; st=line.strip()
        if not st or st=="---": i+=1; continue
        if st.startswith("#"):
            level=len(st)-len(st.lstrip("#")); heading=clean_heading(st[level:].strip())
            # Keep wide tables with their section headings and switch orientation before the heading.
            if heading.startswith("結語"):
                wants_landscape=current
            elif heading.startswith("表 13.1") or heading.startswith("表 13.3"):
                wants_landscape=False
            else:
                wants_landscape=heading.startswith("附錄 A") or upcoming_wide_table(lines,i)
            current=switch_orientation(doc,current,wants_landscape)
            heading_context=heading
            stylename={2:"Heading 1",3:"Heading 2",4:"Heading 3"}.get(level,"Heading 1")
            p=doc.add_paragraph(style=stylename); format_paragraph(p,8 if level==2 else 4,5,1.0,True)
            if heading.startswith("第六部分"):
                p.paragraph_format.page_break_before=True
            add_inline(p,heading,15.2 if level==2 else 11.8 if level==3 else 10.8,NAVY,True)
            i+=1; continue
        if st.startswith("|"):
            rows,i=read_table(lines,i)
            if not rows: continue
            if heading_context.startswith("表 13.1"): current=add_form_131(doc,rows,current)
            elif heading_context.startswith("表 13.3"): current=add_form_133(doc,rows,current)
            else: current=add_table(doc,rows,len(rows[0])>=6,current,9.5)
            continue
        if st.startswith("[Word") and "圖 1" in st:
            current=switch_orientation(doc,current,False)
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; format_paragraph(p,2,0,1.0)
            p.add_run().add_picture(str(CHART),width=Inches(6.15))
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; format_paragraph(p,0,5,1.0)
            add_inline(p,"圖 1  SAMPLE 普及率與合格應使用者全任務完成率（2023–2025）",10.5,MUTED,True)
            i+=1; continue
        if "附錄A1及A2列出本文件完整的80筆" in st and not appendix:
            p=doc.add_paragraph(); format_paragraph(p,0,5,1.12); add_inline(p,st,10.5,INK)
            current,n=add_dataset(doc,current); appendix=True; i+=1; continue
        if st.startswith(">"):
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.18); p.paragraph_format.right_indent=Inches(.1)
            format_paragraph(p,2,5,1.12); add_inline(p,st.lstrip("> ").strip(),10.5,MUTED,False,True); i+=1; continue
        if re.match(r"^\s*(?:[-*]|\d+\.)\s+",line):
            m=re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$",line); indent=len(m.group(1))
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.22+min(indent,8)*.12); p.paragraph_format.first_line_indent=Inches(-.16)
            format_paragraph(p,0,3,1.1); marker="• " if m.group(2) in ("-","*") else m.group(2)+" "
            add_inline(p,marker+m.group(3),10.5,INK); i+=1; continue
        p=doc.add_paragraph(); format_paragraph(p,0,5,1.13); add_inline(p,st,10.5,INK); i+=1
    if not appendix: raise RuntimeError("80-row appendix was not inserted.")
    add_header_footer(doc)
    doc.core_properties.title="訓練資訊系統使用普及率盤點與分析表"
    doc.core_properties.subject="DOCX 0140 / I19-E30 / SAMPLE-SYNTHETIC candidate"
    doc.core_properties.author="TTQS ONE"
    doc.core_properties.comments="CANDIDATE; REAL=UNKNOWN/HOLD; formal acceptance NOT_GRANTED."
    doc.save(OUTDOCX)
    with DATA.open("r",encoding="utf-8-sig",newline="") as f: records=list(csv.DictReader(f))
    with REGISTER.open("r",encoding="utf-8-sig",newline="") as f: register=list(csv.DictReader(f))
    receipt={
      "candidate_revision":"R11","doc_id":"0140","requirement_id":"I19-E30","official_checklist_index":154,"printed_checkbox_serial":"UNKNOWN",
      "candidate_status":"CANDIDATE","real_data_status":"UNKNOWN/HOLD","formal_acceptance":"NOT_GRANTED",
      "manuscript_sha256":sha(MANUSCRIPT),"dataset_sha256":sha(DATA),"replacement_register_sha256":sha(REGISTER),
      "chart_sha256":sha(CHART),"docx_sha256":sha(OUTDOCX),"docx_bytes":OUTDOCX.stat().st_size,
      "synthetic_user_year_rows_in_docx":len(records),"companion_replacement_register_rows":len(register),
      "page_count":"PENDING_RENDER","visual_review":"PENDING","independent_semantic_review":"PENDING_EXACT_SHA",
      "status":"BUILT_PENDING_NATIVE_RENDER_AND_INDEPENDENT_EXACT_SHA_REVIEW",
      "builder":"python-docx 1.2.0 via bundled Python",
      "notes":["80-row appendix comes from canonical CSV, not worker-reordered text.","Chart comes from canonical SAMPLE period summary CSV.",
               "Wide tables use landscape pages; 13.1 intake form is split into two readable five-column tables.",
               "The replacement register is reused unchanged from R07 because the 80-row SAMPLE dataset and mapping are unchanged.",
               "No real association performance, official score, signatures, approvals, or formal acceptance claimed."]
    }
    RECEIPT.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps({"docx":str(OUTDOCX),"sha256":receipt["docx_sha256"],"bytes":receipt["docx_bytes"],"chart_sha256":receipt["chart_sha256"],"rows":len(records),"replacement_register_rows":len(register),"receipt":str(RECEIPT)},ensure_ascii=True,indent=2))

if __name__=="__main__":
    main()
