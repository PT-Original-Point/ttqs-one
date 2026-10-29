from pathlib import Path
import hashlib,json,shutil
from docx import Document
base=Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0013_I02-E04_LUNA_TEXT_DRAFT_R01_20261029")
src=base/"candidate"/"0013_I02-E04_R08.docx";out=base/"candidate"/"0013_I02-E04_R09.docx"
expected="e459ac2b09b2e63a47a68db46327b44cdb3f809b4ece01ac949c80971889c438"
assert hashlib.sha256(src.read_bytes()).hexdigest()==expected
shutil.copy2(src,out);doc=Document(out)
def replace_para(p,text):
 if len(p.runs)<2: raise RuntimeError("Expected preserved bold and regular runs")
 p.runs[0].text="發布前："
 p.runs[1].text=text
 for r in p.runs[2:]:r.text=""
p=doc.paragraphs[93]
assert p.runs[0].bold is True and p.runs[1].bold is not True
replace_para(p,"真實課程來源、權責核准版本、日期與資源均已核實，報名通路可用且擬發布內容符合核准版；任一欠缺即 HOLD，SAMPLE 不代表真實發布。公開頁面讀回屬發布後驗收，不列為發布前條件。上架後由非編輯者讀回頁面，核對網址、時間、版本及內容與核准件一致並留證；失敗或不一致即 HOLD，依權責下架或更正，重新核准後再發布。")
cover=doc.paragraphs[4]
for r in cover.runs:
 if "R08" in r.text:r.text=r.text.replace("R08","R09")
vt=doc.tables[19]
for cell,text in [(vt.rows[1].cells[1],vt.rows[1].cells[1].text.replace("R08","R09")),(vt.rows[2].cells[1],"R09")]:
 if cell.paragraphs[0].runs: cell.paragraphs[0].runs[0].text=text; [setattr(r,"text","") for r in cell.paragraphs[0].runs[1:]]
 else: cell.text=text
doc.core_properties.title="0013 I02-E04 年度訓練課程公開清單 R09 候選 SAMPLE SYNTHETIC"
doc.save(out)
rec={"receipt_id":"0013_I02-E04_R09_RENDER_ROOT_CAUSE_REPAIR","document_id":"0013","requirement_id":"I02-E04","prior_candidate_sha256":hashlib.sha256(src.read_bytes()).hexdigest(),"candidate_path":str(out),"candidate_sha256":hashlib.sha256(out.read_bytes()).hexdigest(),"candidate_bytes":out.stat().st_size,"repair":{"root_cause":"Prior broad paragraph replacement put the full multi-sentence release-gate text into the original bold heading run; R07/R08 paragraph first run remained bold while its normal run was blank. This style regression is a concrete layout risk and plausible cause of excessive pagination time; causal relation to the R07 timeout remains unproven.","change":"Restored a short bold '發布前：' label in the original bold run and placed the body in the original regular-weight run. The content explicitly separates pre-publication checks from post-publication public-page readback.","other_content":"Inherited from R07: seven T19 fact values exactly match the replacement register; R06 circular gate resolved."},"prior_render_evidence":{"R07":"Word native export ran about 8 minutes without PDF; only its isolated process was stopped.","R08":"First hidden launch failed before Word due child PowerShell Get-FileHash module lookup; after explicit module import the R08 native Word process was stopped before output when the bold-run style regression was found."},"layout":"PENDING_NATIVE_WORD_RENDER","semantic":"PENDING_FRESH_EXACT_SHA_REVIEW","candidate_state":"CANDIDATE_SAMPLE_SYNTHETIC_NOT_ACCEPTED","formal_acceptance":"UNKNOWN"}
(base/"qa"/"semantic_repair_R09.json").write_text(json.dumps(rec,ensure_ascii=False,indent=2)+"\n",encoding="utf-8",newline="\n")
print(json.dumps(rec,ensure_ascii=True,indent=2))