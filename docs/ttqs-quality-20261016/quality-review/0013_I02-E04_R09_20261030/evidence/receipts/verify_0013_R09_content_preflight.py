from pathlib import Path
from docx import Document
import csv,hashlib,json
base=Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0013_I02-E04_LUNA_TEXT_DRAFT_R01_20261029")
docx=base/"candidate"/"0013_I02-E04_R09.docx";d=Document(docx)
with (base/"qa"/"0013_R09_SYNTHETIC_REPLACEMENT_REGISTER.csv").open(encoding="utf-8-sig",newline="") as f: rr=list(csv.DictReader(f))
reg={r["Fact ID"]:r["合成值／類型"] for r in rr};t=d.tables[18];vals={r.cells[0].text.strip():r.cells[1].text.strip() for r in t.rows[1:]}
ids=[f"SYN-0013-F{i:03d}" for i in (16,17,18,19,20,21,24)];p=d.paragraphs[93];text=p.text
body_sizes=[r.font.size.pt for x in d.paragraphs for r in x.runs if r.font.size];table_sizes=[r.font.size.pt for t0 in d.tables for row in t0.rows for c in row.cells for x in c.paragraphs for r in x.runs if r.font.size]
checks={"exact_candidate_sha":"dd5ff6cc55c649952872e03a77e77f844375df3db9a86f2bf664aefe1554df3f"==hashlib.sha256(docx.read_bytes()).hexdigest(),"20_tables_67_facts":len(d.tables)==20 and len(vals)==67 and len(set(vals))==67,"seven_values_match_csv":all(vals[i]==reg[i] for i in ids),"publication_gate_is_explicit_and_sequenced":"公開頁面讀回屬發布後驗收，不列為發布前條件" in text and text.index("公開頁面讀回")<text.index("上架後由非編輯者"),"heading_bold_body_regular":p.runs[0].text=="發布前：" and p.runs[0].bold is True and len(p.runs)>1 and p.runs[1].bold is not True,"version_R09":"R09" in d.paragraphs[4].text and d.tables[19].rows[2].cells[1].text.strip()=="R09","font_floors":min(body_sizes)>=10.5 and min(table_sizes)>=9.5,"register_67":len(rr)==67}
print(json.dumps({"sha256":hashlib.sha256(docx.read_bytes()).hexdigest(),"bytes":docx.stat().st_size,"checks":checks,"min_body_pt":min(body_sizes),"min_table_pt":min(table_sizes)},ensure_ascii=True,indent=2))
if not all(checks.values()):raise SystemExit(2)