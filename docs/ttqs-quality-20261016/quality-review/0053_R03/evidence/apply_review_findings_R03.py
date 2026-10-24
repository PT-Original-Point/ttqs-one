from pathlib import Path
from docx import Document
from docx.table import _Row as Row
from copy import deepcopy
import csv,hashlib,json
from datetime import datetime,timezone
pkg=Path(r'E:\TTQS\_TTQS_ONE_REVIEW_SYNC_0036_R08_20261016\docs\ttqs-quality-20261016\quality-review\0053_R03')
r2=Path(r'E:\TTQS\_TTQS_ONE_REVIEW_SYNC_0036_R08_20261016\docs\ttqs-quality-20261016\quality-review\0053_R02')
docp=pkg/'candidate/0053_R03.docx'; regp=pkg/'SYNTHETIC_REPLACEMENT_REGISTER_0053_R03.csv'
old_sha='cbe519643ae53c36818ea56b7f3a03ac43653c1c7183ffa940b9cf97daa885a1'
assert hashlib.sha256(docp.read_bytes()).hexdigest()==old_sha

d=Document(str(docp)); edits=[]
def replace_para(p,old,new,loc):
    assert p.text==old, f"Unexpected text at {loc}: {p.text!r}"
    runs=[r for r in p.runs if r.text]
    assert len(runs)==1, f"Expected one populated run at {loc}, found {len(runs)}"
    runs[0].text=new; edits.append(loc)
def replace_cell(cell,old,new,loc):
    assert cell.text==old, f"Unexpected cell at {loc}: {cell.text!r}"
    ps=[p for p in cell.paragraphs if p.text]
    assert len(ps)==1
    replace_para(ps[0],old,new,loc)

oldp='判讀問題：「實際執行是否符合核定版本？」目前答案為 UNKNOWN：本包沒有可核實的 REAL 核定版本或實際執行原件。示範中的 P0-SYN 是案例草稿，不得當作協會核定版。'
newp='判讀問題：「實際執行是否符合核定版本？」目前答案為 UNKNOWN：本包沒有可核實的 REAL 核定版本或實際執行原件。示範中的 P0-SYN（本案合成計畫草稿版號）僅供案例識別，不得當作協會核定版。'
found=[]
for i,p in enumerate(d.paragraphs,1):
    if p.text==oldp:
        matching=[r for r in p.runs if 'P0-SYN' in r.text]
        assert len(matching)==1
        matching[0].text=matching[0].text.replace('P0-SYN','P0-SYN（本案合成計畫草稿版號）')
        edits.append('p2 / section 01 / P0-SYN definition')
        found.append(i)
assert len(found)==1

# Add the sample-only 50% P2 practice threshold to the existing data dictionary, directly documenting its scope.
dictionary=d.tables[10]
assert len(dictionary.columns)==4
source_tr=dictionary.rows[7]._tr
new_tr=deepcopy(source_tr)
dictionary._tbl.append(new_tr)
row=Row(new_tr,dictionary)
values=[
    'P2 練習門檻',
    '本合成案例暫以 G1 引導練習至少占單元時間 50% 作桌上比較；不是官方標準，也不是協會已核定規則。',
    '引導練習分鐘 ÷ G1 90 分鐘；基準 35/90=38.9%，方案 A 50/90=55.6%。',
    '此門檻只用於本 SAMPLE 判讀；REAL 門檻須由有權責者核定並以正式課程原件為據。'
]
for j,(cell,value) in enumerate(zip(row.cells,values),1):
    ps=[p for p in cell.paragraphs if p.text]
    if len(ps)!=1: raise AssertionError(f'Unexpected dictionary cell paragraphs {j}: {len(ps)}')
    old=ps[0].text
    runs=[r for r in ps[0].runs if r.text]
    if len(runs)!=1: raise AssertionError(f'Unexpected dictionary runs {j}: {len(runs)}')
    runs[0].text=value
edits.append('p5 / section 04 / data dictionary / P2 practice threshold')

# Correct an internal page cross-reference: page 8 is the F1-F3 continuation, not the improvement loop.
cell=d.tables[28].rows[7].cells[1]
replace_cell(cell,'第 8-14 頁；第 17 頁版本差異表。','第 9-14 頁（分析與追蹤閉環）；第 17 頁版本差異表。','p18 / appendix C / continuous-improvement locator')
d.save(str(docp))
new_sha=hashlib.sha256(docp.read_bytes()).hexdigest()

# Rebind the 38 existing facts and correct the independently identified locators/value semantics.
with regp.open(encoding='utf-8-sig',newline='') as f: rows=list(csv.DictReader(f))
assert len(rows)==38 and len({r['fact_id'] for r in rows})==38
byid={r['fact_id']:r for r in rows}
byid['SYN-0053-F017'].update({
 'body_locator':'§02 案例範圍；課程名稱段落','rendered_page':'3'})
byid['SYN-0053-F018'].update({
 'body_locator':'§05 學員/教材適配檢核；L2 基線示範證據','rendered_page':'6'})
byid['SYN-0053-F030'].update({
 'value_in_candidate':'可讀性基線：4/6 情境正確；66.7% 為 4÷6×100 的衍生值，正文未另列此百分比。',
 'source_or_formula':'六個 SAMPLE/SYNTHETIC QA 情境逐列計數：4÷6=66.666…%，四捨五入66.7%；正文直接呈現4/6。',
 'body_locator':'§07 前後複核；六個合成情境明細及基線結果','rendered_page':'9'})
byid['SYN-0053-F011'].update({
 'value_in_candidate':'P2 基線：G1 引導練習35/90=38.9%，低於本報告明列的50% SAMPLE 比較門檻。',
 'source_or_formula':'0052 R01 SAMPLE minutes; 35/90=38.9%; the 50% threshold is an internal SAMPLE comparison setting only, not an official or approved association standard.',
 'body_locator':'§06 P2 基線；§04 資料字典 P2 練習門檻','rendered_page':'7'})
new_rows=[
{
'fact_id':'SYN-0053-F039','value_in_candidate':'P0-SYN＝本案合成計畫草稿版號，只供案例識別；不是協會核定版或 REAL 版號。','fact_type':'SAMPLE/SYNTHETIC','source_or_formula':'R02/R03 sample label used for the synthetic course-plan case; parenthetical definition added for reader clarity.','candidate_sha256':new_sha,'body_locator':'§01 判讀問題；P0-SYN 括號定義','rendered_page':'2','real_replacement_source':'如有真實核定課程，核對有權版本控制的原件編號、版次、日期及核准權限；不得沿用 P0-SYN。','replacement_responsible_role':'課程規劃承辦／有權核定者','status':'SYNTHETIC—待 REAL 替換'},
{
'fact_id':'SYN-0053-F040','value_in_candidate':'P2 的50%引導練習門檻僅為本報告合成比較設定；不是官方標準或協會已核定規則。','fact_type':'SAMPLE/SYNTHETIC','source_or_formula':'SAMPLE threshold set for this worked example; compare baseline 35/90=38.9% with proposal 50/90=55.6%.','candidate_sha256':new_sha,'body_locator':'§04 資料字典 P2 練習門檻；§06 P2 檢核列','rendered_page':'5','real_replacement_source':'有權責者核定的實際課程目標、教案及評量標準；若無，標 UNKNOWN，不引用示範門檻。','replacement_responsible_role':'課程規劃權責者／訓練承辦','status':'SYNTHETIC—待 REAL 替換'}]
rows.extend(new_rows)
assert len(rows)==40 and len({r['fact_id'] for r in rows})==40
for r in rows: r['candidate_sha256']=new_sha
with regp.open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.DictWriter(f,fieldnames=rows[0].keys(),lineterminator='\n');w.writeheader();w.writerows(rows)

# Update source map and record review-finding disposition; R02 is kept as an unreviewed intermediate.
sm=(pkg/'inputs/SOURCE_MAP_0053_R03.md').read_text(encoding='utf-8')
sm=sm.replace('# 0053 R02 來源與範圍對照','# 0053 R03 來源與範圍對照')
sm += '\n\n## R03 對獨立審查發現的處置\n\n- R01 reviewer 的 F017/F018/F030 定位與數值差異已修正。\n- R01 reviewer 指稱 T3 六個合成作答樣例與 G1 五情境門檻未登錄；現有 F009（T3 六樣例）及 F012（五情境至少四題且不得越權承諾）已逐項對應，無須重複建列。\n- P0-SYN 新增內文定義及替換清冊 F039。\n- P2 50% SAMPLE 比較門檻新增資料字典定義及獨立清冊 F040。\n- Appendix C 持續改善定位由第8-14頁改為第9-14頁及第17頁。\n- 0052 R01 exact SHA `406d1ef2c6e9b6266ea1adaeadedc5dfac77964fe76f73c2471168220ee6d455` 表22列2確認 20+20+50 為分鐘。\n'
(pkg/'inputs/SOURCE_MAP_0053_R03.md').write_text(sm,encoding='utf-8')

receipt={
 'schema':'ttqs.0053.reviewer_correction_receipt.v1','doc_id':'0053','requirement_id':'I12-E01','official_checkbox_sequence':'056','revision':'R03',
 'candidate_sha256':new_sha,'candidate_bytes':docp.stat().st_size,'previous_candidate_sha256':old_sha,
 'register_path':str(regp),'register_sha256':hashlib.sha256(regp.read_bytes()).hexdigest(),'register_rows':len(rows),'unique_fact_ids':len(set(r['fact_id'] for r in rows)),
 'text_and_traceability_corrections':edits+['F017/F018/F030 register corrections','add F039 P0-SYN and F040 sample threshold facts','candidate SHA rebinding for all rows'],
 'upstream_R01_review_outcome':'PARTIAL','next_review_required':'new exact-SHA review of R03; include primary neighboring DOCX 0052 excerpt and current crosswalk rows',
 'built_at_utc':datetime.now(timezone.utc).isoformat()}
(pkg/'evidence/build_receipt_R03.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
log=f'''# 0053 R03 targeted corrections

Previous exact candidate: `{old_sha}` (R02; unreviewed intermediate).

New exact candidate: `{new_sha}`.

Based on the R01 exact-SHA independent review (PARTIAL), R03 defines `P0-SYN`, defines the 50% P2 practice threshold as SAMPLE-only in the dictionary, fixes Appendix C page references, corrects register locations F017/F018, marks F030's 66.7% as a derived register value, and adds replacement facts F039/F040. Existing F009 and F012 already map the T3 six-answer sample and five-scenario assessment rule, so they remain single-source facts rather than duplicated rows.

R03 preserves all corrected G1 minutes from R02: baseline 35+20+35=90; proposal 20+20+50=90; practice share 38.9% to 55.6% (+16.7 pp). All remain SAMPLE/SYNTHETIC; actual adoption and delivery remain UNKNOWN.

R01 and R02 candidates and their review/build evidence are preserved separately. R03 is a candidate and is not promoted.
'''
(pkg/'evidence/correction_log_R03.md').write_text(log,encoding='utf-8')
print(json.dumps({'previous_sha':old_sha,'new_sha':new_sha,'register_rows':len(rows),'register_sha256':receipt['register_sha256'],'corrected_items':len(edits),'edits':edits},ensure_ascii=False,indent=2))


