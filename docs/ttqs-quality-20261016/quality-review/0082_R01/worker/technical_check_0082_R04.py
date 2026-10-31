from __future__ import annotations

import hashlib
import json
import sys
import zipfile
from pathlib import Path

from docx import Document
from PIL import Image


pkg = Path(__file__).resolve().parents[1]
docx_path = pkg / "candidate" / "0082_R04.docx"
body_path = pkg / "candidate" / "0082_R04_body.md"
build_receipt = pkg / "evidence" / "build_receipt_R04.json"
pdf_path = pkg / "evidence" / "0082_R04_word_render.pdf"
render_dir = pkg / "evidence" / "render_R04_word"
out_path = pkg / "evidence" / "technical_review_R04.json"


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


issues: list[str] = []
with zipfile.ZipFile(docx_path) as zf:
    bad_member = zf.testzip()
    if bad_member:
        issues.append(f"DOCX ZIP CRC failure: {bad_member}")

doc = Document(docx_path)
receipt = json.loads(build_receipt.read_text(encoding="utf-8"))
text = "\n".join(
    [p.text for p in doc.paragraphs]
    + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
)
pages = sorted(render_dir.glob("page-*.png"))
page_sizes = []
for page in pages:
    with Image.open(page) as im:
        page_sizes.append(list(im.size))

checks = {
    "docx_archive_opens_and_crc_valid": bad_member is None,
    "python_docx_opens": True,
    "build_receipt_docx_sha_matches": sha(docx_path) == receipt["docx_sha256"],
    "build_receipt_body_sha_matches": sha(body_path) == receipt["source_body_sha256"],
    "identity_is_0082_i17c_e01_checkbox_088": all(
        token in text for token in ("DOCX 0082", "I17C-E01", "088")
    ),
    "distinct_next_requirement_0083_089": "0083" in text and "089" in text,
    "truth_boundary_markers_present": all(
        token in text for token in ("SAMPLE/SYNTHETIC", "REAL=UNKNOWN/HOLD", "NOT_GRANTED")
    ),
    "expected_three_sections_twenty_tables": len(doc.sections) == 3 and len(doc.tables) == 20,
    "pdf_render_exists": pdf_path.is_file() and pdf_path.stat().st_size > 0,
    "render_has_17_nonempty_pages": len(pages) == 17 and all(p.stat().st_size > 0 for p in pages),
    "render_page_dimensions_match_portrait_and_landscape_sections": (
        set(tuple(x) for x in page_sizes) == {(1076, 1520), (1520, 1076)}
        and sum(tuple(x) == (1520, 1076) for x in page_sizes) == 3
    ),
    "no_internal_worker_metadata_in_body": not any(
        token in text.lower() for token in ("opencode", "antigravity", "git commit", "sha-256")
    ),
}
for name, ok in checks.items():
    if not ok:
        issues.append(name)

result = {
    "schema": "TTQS_ONE_0082_TECHNICAL_REVIEW_R04.v1",
    "document_id": "0082",
    "requirement_id": "I17C-E01",
    "official_checkbox_serial": "088",
    "candidate_path": "candidate/0082_R04.docx",
    "candidate_docx_sha256": sha(docx_path),
    "candidate_body_sha256": sha(body_path),
    "pdf_sha256": sha(pdf_path),
    "pdf_bytes": pdf_path.stat().st_size,
    "render_page_count": len(pages),
    "render_page_dimensions_by_page": page_sizes,
    "paragraph_count": len(doc.paragraphs),
    "table_count": len(doc.tables),
    "section_count": len(doc.sections),
    "checks": checks,
    "issues": issues,
    "status": "PASS" if not issues else "FAIL",
    "formal_acceptance": "NOT_GRANTED",
    "real_evidence": "UNKNOWN/HOLD",
}
out_path.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps({"status": result["status"], "sha256": result["candidate_docx_sha256"], "pages": len(pages), "issues": issues}, ensure_ascii=False))
if issues:
    raise SystemExit(1)
