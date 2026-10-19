from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re,sys,json,hashlib,datetime

src=Path(sys.argv[1]); out=Path(sys.argv[2]); receipt=Path(sys.argv[3])
text=src.read_text(encoding='utf-8')
lines=text.splitlines()
doc=Document()
A4_W=Inches(8.27); A4_H=Inches(11.69)
BODY_FONT='Microsoft JhengHei'
NAVY='18344F'; TEAL='2B6F73'; PALE='EAF2F2'; MIST='F3F6F8'; WHITE='FFFFFF'; DARK='243746'; GOLD='9B6A13'

def set_run_font(run, size=10.5, bold=None, color=None, italic=None, name=BODY_FONT):
    run.font.name=name
    rpr=run._element.get_or_add_rPr()
    rf=rpr.rFonts
    if rf is None:
        rf=OxmlElement('w:rFonts'); rpr.insert(0,rf)
    for key in ('ascii','hAnsi','eastAsia','cs'): rf.set(qn(f'w:{key}'),name)
    run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    if italic is not None: run.italic=italic
    if color: run.font.color.rgb=RGBColor.from_string(color)

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)

def set_cell_margins(cell, top=50, start=90, bottom=50, end=90):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); mar=tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar=OxmlElement('w:tcMar'); tcPr.append(mar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        el=mar.find(qn('w:'+m))
        if el is None: el=OxmlElement('w:'+m); mar.append(el)
        el.set(qn('w:w'),str(v)); el.set(qn('w:type'),'dxa')

def no_split(row):
    pr=row._tr.get_or_add_trPr(); el=OxmlElement('w:cantSplit'); pr.append(el)

def repeat_header(row):
    pr=row._tr.get_or_add_trPr(); el=OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); pr.append(el)

def add_page_field(p):
    p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=p.add_run('DOCX 0082  |  SAMPLE/SYNTHETIC  |  '); set_run_font(r,8,color='667784')
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); p._p.append(fld)

def configure_section(sec, landscape=False):
    sec.page_width=A4_W; sec.page_height=A4_H
    sec.orientation=WD_ORIENT.PORTRAIT
    if landscape:
        sec.orientation=WD_ORIENT.LANDSCAPE; sec.page_width=A4_H; sec.page_height=A4_W
        sec.left_margin=Inches(0.55); sec.right_margin=Inches(0.55)
    else:
        sec.left_margin=Inches(0.68); sec.right_margin=Inches(0.68)
    sec.top_margin=Inches(0.62); sec.bottom_margin=Inches(0.62)
    sec.header_distance=Inches(0.28); sec.footer_distance=Inches(0.28)

configure_section(doc.sections[0],False)
styles=doc.styles
normal=styles['Normal']; normal.font.name=BODY_FONT; normal.font.size=Pt(10.5); normal.font.color.rgb=RGBColor.from_string(DARK)
normal.paragraph_format.space_after=Pt(3.0); normal.paragraph_format.line_spacing=1.03
for style_name,size,color in [('Title',22,NAVY),('Heading 1',16,NAVY),('Heading 2',12.5,TEAL),('Heading 3',11,DARK)]:
 st=styles[style_name]; st.font.name=BODY_FONT; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)
 st.paragraph_format.space_before=Pt(8 if style_name!='Title' else 0); st.paragraph_format.space_after=Pt(4); st.paragraph_format.keep_with_next=True
 for k in ('ascii','hAnsi','eastAsia','cs'): st._element.rPr.rFonts.set(qn(f'w:{k}'),BODY_FONT)
# Remove the built-in Title style's bottom rule and any direct title border.
title_ppr=styles['Title']._element.pPr
if title_ppr is not None:
    for child in list(title_ppr):
        if child.tag==qn('w:pBdr'): title_ppr.remove(child)
# Header/footer
for sec in doc.sections:
 hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 hr=hp.add_run('TTQS ONE   •   I17C-E01   •   官方勾選 088'); set_run_font(hr,8,color='667784')
 add_page_field(sec.footer.paragraphs[0])

inline_re=re.compile(r'(\*\*.+?\*\*|`[^`]+`|(?<!\*)\*[^*]+\*)')
def add_inline(p,s,size=10.5,base_color=None):
    pos=0
    for m in inline_re.finditer(s):
        if m.start()>pos:
            r=p.add_run(s[pos:m.start()]); set_run_font(r,size,color=base_color)
        t=m.group(0)
        if t.startswith('**'):
            r=p.add_run(t[2:-2]); set_run_font(r,size,bold=True,color=base_color)
        elif t.startswith('`'):
            r=p.add_run(t[1:-1]); set_run_font(r,size,color=(base_color or '214D68'),name='Consolas')
        else:
            r=p.add_run(t[1:-1]); set_run_font(r,size,italic=True,color=base_color)
        pos=m.end()
    if pos<len(s):
        r=p.add_run(s[pos:]); set_run_font(r,size,color=base_color)

def clean_cell(s):
    return s.strip()
def add_md_table(table_lines, landscape):
    rows=[]
    for ln in table_lines:
        cells=[clean_cell(x) for x in ln.strip().strip('|').split('|')]
        if not rows and cells: pass
        if cells and all(re.fullmatch(r':?-{3,}:?',c.replace(' ','')) for c in cells): continue
        rows.append(cells)
    if not rows: return
    n=max(len(r) for r in rows); rows=[r+['']*(n-len(r)) for r in rows]
    t=doc.add_table(rows=1,cols=n); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'; t.autofit=False
    for ri,rowdata in enumerate(rows):
        row=t.rows[0] if ri==0 else t.add_row()
        no_split(row)
        for ci,val in enumerate(rowdata):
            cell=row.cells[ci]; cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cell)
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
            if ri==0: shade(cell,NAVY); p.paragraph_format.keep_with_next=True
            elif ri%2==0: shade(cell,'F5F8FA')
            add_inline(p,val,size=(7.0 if n>=8 else 7.7 if n>=6 else 8.5 if n>=4 else 9.1),base_color=WHITE if ri==0 else DARK)
            if ri==0:
                for run in p.runs: run.bold=True
    repeat_header(t.rows[0])
    usable=10.05 if landscape else 6.91
    if n==9 and landscape:
        if 'BTR-' in '\n'.join(table_lines): widths=[0.80,0.95,0.76,0.76,0.76,0.76,0.76,0.76,3.74]
        else: widths=[0.75,1.15,1.0,1.15,1.0,1.0,0.85,0.65,0.65]
        scale=usable/sum(widths); widths=[x*scale for x in widths]
    elif n==6 and not landscape:
        widths=[0.47,1.15,0.93,0.93,1.06,2.37]
        scale=usable/sum(widths); widths=[x*scale for x in widths]
    else:
        lens=[max(3,max(len(re.sub(r'[`*_]','',r[i])) for r in rows)) for i in range(n)]
        weights=[min(2.6,max(0.65,x**0.55)) for x in lens]; total=sum(weights)
        widths=[usable*x/total for x in weights]
    for row in t.rows:
        for ci,w in enumerate(widths):
            row.cells[ci].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

def add_callout(s):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(0.16); p.paragraph_format.right_indent=Inches(0.10); p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(5)
    pPr=p._p.get_or_add_pPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),MIST); pPr.append(shd)
    pBdr=OxmlElement('w:pBdr'); left=OxmlElement('w:left'); left.set(qn('w:val'),'single'); left.set(qn('w:sz'),'18'); left.set(qn('w:space'),'6'); left.set(qn('w:color'),TEAL); pBdr.append(left); pPr.append(pBdr)
    add_inline(p,s,9.6,base_color=DARK)

landscape=False; i=0; first=True
while i<len(lines):
    line=lines[i]
    if not line.strip() or line.strip()=='---': i+=1; continue
    if line.startswith('## 8.') and not landscape:
        sec=doc.add_section(WD_SECTION.NEW_PAGE); configure_section(sec,True); landscape=True
    elif line.startswith('## 10.') and landscape:
        sec=doc.add_section(WD_SECTION.NEW_PAGE); configure_section(sec,False); landscape=False
    if line.startswith('|'):
        tl=[]
        while i<len(lines) and lines[i].startswith('|'):
            tl.append(lines[i]); i+=1
        add_md_table(tl,landscape); continue
    if line.startswith('# '):
        p=doc.add_paragraph(style='Title'); ppr=p._p.get_or_add_pPr(); bdr=ppr.find(qn('w:pBdr'));
        if bdr is not None: ppr.remove(bdr)
        add_inline(p,line[2:].strip(),22,base_color='000000'); i+=1; continue
    if line.startswith('## '):
        if line.startswith('## 6.'):
            if len(doc.paragraphs)>2: doc.add_page_break()
        level='Heading 1'; heading=line[3:].strip()
        p=doc.add_paragraph(style=level); add_inline(p,heading,16,base_color=NAVY); i+=1; continue
    if line.startswith('### '):
        p=doc.add_paragraph(style='Heading 2'); add_inline(p,line[4:].strip(),12.5,base_color=TEAL); i+=1; continue
    if line.startswith('#### '):
        p=doc.add_paragraph(style='Heading 3'); add_inline(p,line[5:].strip(),11,base_color=DARK); i+=1; continue
    if line.startswith('> '):
        add_callout(line[2:].strip()); i+=1; continue
    if re.match(r'^\s*[-*] ',line):
        p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.left_indent=Inches(0.25); p.paragraph_format.first_line_indent=Inches(-0.13); p.paragraph_format.space_after=Pt(2.2); add_inline(p,re.sub(r'^\s*[-*] ','',line),10.2); i+=1; continue
    if re.match(r'^\s*\d+\. ',line):
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(0.22); p.paragraph_format.first_line_indent=Inches(-0.22); p.paragraph_format.space_after=Pt(2.2); add_inline(p,line.strip(),10.2); i+=1; continue
    p=doc.add_paragraph(); add_inline(p,line.strip(),10.5); i+=1
# Set core properties without attributing false organizational authorship.
doc.core_properties.title='學員訓後動態調查 行為轉移問卷與追蹤間隔'
doc.core_properties.subject='DOCX 0082 I17C-E01 official checkbox 088; candidate demonstration only'
doc.core_properties.author=''
doc.core_properties.comments='SAMPLE/SYNTHETIC; candidate only; REAL=UNKNOWN/HOLD; formal acceptance NOT_GRANTED.'
out.parent.mkdir(parents=True,exist_ok=True); doc.save(out)
sha=hashlib.sha256(out.read_bytes()).hexdigest()
receipt.write_text(json.dumps({'document_id':'0082','revision':'R08','source_body_sha256':hashlib.sha256(src.read_bytes()).hexdigest(),'docx_path':str(out),'docx_sha256':sha,'builder':'deterministic python-docx parser; body authored by OpenCode CLI v2 and bounded editor correction','sections':len(doc.sections),'tables':len(doc.tables),'paragraphs':len(doc.paragraphs),'page_size':'A4 portrait with landscape sections 8-9','candidate_state':'CANDIDATE_ONLY; SAMPLE/SYNTHETIC; REAL=UNKNOWN/HOLD; formal acceptance NOT_GRANTED'},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'docx':str(out),'sha256':sha,'sections':len(doc.sections),'tables':len(doc.tables),'paragraphs':len(doc.paragraphs)},ensure_ascii=False))
