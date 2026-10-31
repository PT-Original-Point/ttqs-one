from pathlib import Path

from docx import Document

pkg = Path(__file__).resolve().parents[1]
source = pkg / "candidate" / "0082_R07.docx"
target = pkg / "evidence" / "docx_extracted_R07.md"
doc = Document(source)
parts = [
    "# Extracted text projection — DOCX 0082 R07",
    "",
    "Source is the exact candidate DOCX. Verify its SHA against technical_review_R07.json.",
    "",
    f"Paragraphs: {len(doc.paragraphs)}; tables: {len(doc.tables)}; sections: {len(doc.sections)}.",
    "",
    "## Paragraphs",
]
for index, paragraph in enumerate(doc.paragraphs, start=1):
    value = paragraph.text.strip()
    if value:
        parts.append(f"[P{index}; {paragraph.style.name}] {value}")
parts.extend(["", "## Tables", ""])
for table_index, table in enumerate(doc.tables, start=1):
    parts.append(f"### Table {table_index} ({len(table.rows)} rows × {len(table.columns)} columns)")
    for row_index, row in enumerate(table.rows, start=1):
        cells = [cell.text.replace("\n", " / ").strip() for cell in row.cells]
        parts.append(f"- R{row_index}: " + " | ".join(cells))
    parts.append("")
target.write_text("\n".join(parts) + "\n", encoding="utf-8")
print(f"extracted={target}")
