from __future__ import annotations
import hashlib, json, re, sys
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[1]
BODY=ROOT/'candidate/0072_R01_body_R04_QA.md'
OUT=ROOT/'candidate/0072_R01.docx'
RECEIPT=ROOT/'evidence/build_receipt_R01.json'
FONT='Microsoft JhengHei'
NAVY='17365D'


def set_run_font(run, size=10.5, bold=None, color=None):
    run.font.name=FONT
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    if color: run.font.color.rgb=RGBColor.from_string(color)


def remove_title_border(style):
    ppr=style._element.get_or_add_pPr()
    bdr=ppr.find(qn('w:pBdr'))
    if bdr is not None: ppr.remove(bdr)


def shade(cell, fill):
    tcpr=cell._tc.get_or_add_tcPr(); shd=tcpr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcpr.append(shd)
    shd.set(qn('w:fill'),fill)


def set_cell_margins(cell, top=70, start=85, bottom=70, end=85):
    tc=cell._tc; tcpr=tc.get_or_add_tcPr(); mar=tcpr.first_child_found_in('w:tcMar')
    if mar is None: mar=OxmlElement('w:tcMar'); tcpr.append(mar)
    for tag,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=mar.find(qn('w:'+tag))
        if node is None: node=OxmlElement('w:'+tag); mar.append(node)
        node.set(qn('w:w'),str(val)); node.set(qn('w:type'),'dxa')


def mark_repeat_table_header(row):
    trpr=row._tr.get_or_add_trPr(); el=OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trpr.append(el)


def prevent_row_split(row):
    trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))


def add_inline(p, text, table=False):
    text=text.strip()
    for token in re.split(r'(\*\*.+?\*\*|`[^`]+`)',text):
        if not token: continue
        if token.startswith('**') and token.endswith('**'):
            r=p.add_run(token[2:-2]); set_run_font(r,9.5 if table else 10.5,bold=True)
        elif token.startswith('`') and token.endswith('`'):
            r=p.add_run(token[1:-1]); set_run_font(r,9.5 if table else 10.5); r.font.color.rgb=RGBColor.from_string('404040')
        else:
            r=p.add_run(token); set_run_font(r,9.5 if table else 10.5)


def parse_row(line):
    return [c.strip().replace('\\|','|') for c in line.strip().strip('|').split('|')]


def add_table(doc, rows):
    if not rows: return
    width=max(len(r) for r in rows)
    rows=[r+['']*(width-len(r)) for r in rows]
    table=doc.add_table(rows=len(rows),cols=width)
    table.style='Table Grid'; table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=True
    mark_repeat_table_header(table.rows[0])
    for ri,row in enumerate(rows):
        prevent_row_split(table.rows[ri])
        for ci,value in enumerate(row):
            cell=table.cell(ri,ci);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;set_cell_margins(cell)
            if ri==0:shade(cell,NAVY)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.0
            add_inline(p,value,table=True)
            if ri==0:
                for run in p.runs:run.font.color.rgb=RGBColor(255,255,255);run.bold=True
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return table


def add_page_field(p):
    run=p.add_run('DOCX 0072  |  SAMPLE / SYNTHETIC  |  ');set_run_font(run,8,color='666666')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE')
    r=OxmlElement('w:r');rpr=OxmlElement('w:rPr');fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),FONT);fonts.set(qn('w:hAnsi'),FONT);fonts.set(qn('w:eastAsia'),FONT);rpr.append(fonts);r.append(rpr)
    t=OxmlElement('w:t');t.text='1';r.append(t);fld.append(r);p._p.append(fld)


def main():
    if not BODY.exists():raise SystemExit('Writer body not found: '+str(BODY))
    body_bytes=BODY.read_bytes();text=body_bytes.decode('utf-8-sig')
    doc=Document();sec=doc.sections[0];sec.page_width=Mm(210);sec.page_height=Mm(297)
    sec.top_margin=Inches(.68);sec.bottom_margin=Inches(.68);sec.left_margin=Inches(.68);sec.right_margin=Inches(.68)
    styles=doc.styles
    normal=styles['Normal'];normal.font.name=FONT;normal._element.rPr.rFonts.set(qn('w:eastAsia'),FONT);normal.font.size=Pt(10.5);normal.font.color.rgb=RGBColor.from_string('202020');normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1.12
    title=styles['Title'];title.font.name=FONT;title._element.rPr.rFonts.set(qn('w:eastAsia'),FONT);title.font.size=Pt(23);title.font.bold=True;title.font.color.rgb=RGBColor.from_string('111111');title.paragraph_format.space_after=Pt(5);remove_title_border(title)
    for name,size in [('Heading 1',15),('Heading 2',12),('Heading 3',10.5)]:
        st=styles[name];st.font.name=FONT;st._element.rPr.rFonts.set(qn('w:eastAsia'),FONT);st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=RGBColor.from_string(NAVY);st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(4);st.paragraph_format.keep_with_next=True
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER;add_page_field(foot)
    lines=text.splitlines();i=0;seen_title=False;paragraph_count=0;table_count=0
    while i<len(lines):
        line=lines[i].rstrip()
        if not line.strip():i+=1;continue
        if line.lstrip().startswith('|'):
            block=[]
            while i<len(lines) and lines[i].lstrip().startswith('|'):
                cells=parse_row(lines[i]);i+=1
                if cells and all(re.fullmatch(r':?-{3,}:?',c.replace(' ','')) for c in cells):continue
                block.append(cells)
            add_table(doc,block);table_count+=1;continue
        m=re.match(r'^(#{1,4})\s+(.*)$',line)
        if m:
            level=len(m.group(1));heading=m.group(2).strip()
            if level==1 and not seen_title:
                p=doc.add_paragraph(style='Title');add_inline(p,heading);seen_title=True
            else:
                p=doc.add_paragraph(style='Heading '+str(min(level,3)))
                add_inline(p,heading)
            paragraph_count+=1;i+=1;continue
        if line.startswith('> '):
            p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.2);p.paragraph_format.right_indent=Inches(.2);p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(7)
            add_inline(p,line[2:]);paragraph_count+=1;i+=1;continue
        if re.match(r'^\s*[-*]\s+',line):
            p=doc.add_paragraph(style='List Bullet');add_inline(p,re.sub(r'^\s*[-*]\s+','',line));paragraph_count+=1;i+=1;continue
        if re.match(r'^\s*\d+[.)]\s+',line):
            p=doc.add_paragraph(style='List Number');add_inline(p,re.sub(r'^\s*\d+[.)]\s+','',line));paragraph_count+=1;i+=1;continue
        if re.fullmatch(r'\s*[-_]{3,}\s*',line):i+=1;continue
        p=doc.add_paragraph();add_inline(p,line);paragraph_count+=1;i+=1
    if not seen_title:
        p=doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph(style='Title');p.style='Title';add_inline(p,'訓練的內外稽核流程與紀錄')
    doc.core_properties.title='訓練的內外稽核流程與紀錄'
    doc.core_properties.subject='DOCX 0072 | I16-E03 | official checkbox 078 | SAMPLE/SYNTHETIC candidate'
    doc.core_properties.author=''
    OUT.parent.mkdir(parents=True,exist_ok=True);doc.save(OUT)
    out=OUT.read_bytes();receipt={'schema':'ttqs-one-docx-build-receipt.v1','doc_id':'0072','requirement_id':'I16-E03','official_checkbox_serial':'078','candidate_path':str(OUT.relative_to(ROOT.parent)),'candidate_sha256':hashlib.sha256(out).hexdigest(),'candidate_bytes':len(out),'body_sha256':hashlib.sha256(body_bytes).hexdigest(),'builder_sha256':hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),'paragraphs':paragraph_count,'tables':table_count,'real_evidence':'UNKNOWN_HOLD','status':'CANDIDATE_BUILT_TECHNICAL_QA_PENDING'}
    RECEIPT.parent.mkdir(parents=True,exist_ok=True);RECEIPT.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(receipt,ensure_ascii=False))

if __name__=='__main__':main()
