from pathlib import Path
from docx import Document
import hashlib, json, re
base=Path(r'E:\PT-Original-Point\ttqs-one\docs\ttqs-quality-20261016\quality-review')
r01=base/'0107_I18-E15_R01_20261022'
r02=base/'0107_I18-E15_R02_20261022'
def digest(b): return hashlib.sha256(b).hexdigest()
def plain(p):
 d=Document(p)
 s='\n'.join([x.text for x in d.paragraphs]+['\n'.join(c.text for c in row.cells) for t in d.tables for row in t.rows])+'\n'
 return s.encode('utf-8')
r01_docx=next((r01/'candidate').glob('*.docx'))
r02_docx=next((r02/'candidate').glob('*.docx'))
r01_text=plain(r01_docx); r02_text=plain(r02_docx)
r02_snapshot=(r02/'review/candidate_text_0107_R02.txt').read_bytes()
raw=(r02/'review/opencode_semantic_delta_R02_followup_stdout.jsonl').read_text(encoding='utf-8')
parts=[]
for line in raw.splitlines():
 try: row=json.loads(line)
 except json.JSONDecodeError: continue
 part=row.get('part',{})
 if part.get('type')=='text' and isinstance(part.get('text'),str): parts.append(part['text'])
if not parts: raise SystemExit('No text response found in raw OpenCode output')
response=parts[-1]
if not response.lstrip().startswith('## Verdict: PASS'): raise SystemExit('Latest text is not a PASS reviewer response')
response_path=r02/'review/opencode_semantic_delta_R02_final_response.md'
response_path.write_text(response.rstrip()+'\n',encoding='utf-8',newline='\n')
packet_manifest=json.loads((r02/'evidence/packet_manifest_0107_R02.json').read_text(encoding='utf-8'))
diff=json.loads((r02/'evidence/r01_r02_deterministic_text_diff_0107.json').read_text(encoding='utf-8'))
dispatch=json.loads((r01/'review/opencode_semantic_review_dispatch_0107_R01.json').read_text(encoding='utf-8'))
raw_response_sha=digest((r02/'review/opencode_semantic_delta_R02_followup_stdout.jsonl').read_bytes())
reconciliation={
 'doc_id':'0107','requirement_id':'I18-E15','candidate_sha256':digest(r02_docx.read_bytes()),
 'method':'Recomputed python-docx plain text serialization: paragraphs, then each table row cell sequence; newline between blocks and final newline. This is distinct from the structural JSON body-block snapshot.',
 'r02':{'recomputed_plain_text_sha256':digest(r02_text),'saved_packet_snapshot_sha256':digest(r02_snapshot),'manifest_snapshot_sha256':packet_manifest.get('text_snapshot_sha256'),'recomputed_matches_saved_and_manifest':r02_text==r02_snapshot and digest(r02_snapshot)==packet_manifest.get('text_snapshot_sha256')},
 'structural_diff':{'json_snapshot_sha256':diff.get('r02_extracted_text_snapshot_sha256'),'serialization':'JSON structural body-block representation in r01_r02_deterministic_text_diff_0107.json; not byte-comparable with the plain-text .txt snapshot'},
 'r01':{'recomputed_plain_text_sha256':digest(r01_text),'prior_review_dispatch_extracted_text_sha256':dispatch.get('extracted_text_sha256'),'reconstructed_match':digest(r01_text)==dispatch.get('extracted_text_sha256'),'status':'UNRESOLVED_PROVENANCE_HASH_MISMATCH' if digest(r01_text)!=dispatch.get('extracted_text_sha256') else 'MATCH'},
 'reviewer_response':{'raw_stdout_sha256':raw_response_sha,'extracted_response_path':str(response_path.relative_to(base.parent.parent.parent.parent)).replace('\\','/'),'extracted_response_sha256':digest(response_path.read_bytes()),'response_chars':len(response)},
 'conclusion':'R02 plain-text snapshot and packet hash agree exactly. The separate structural JSON snapshot uses a different representation. R01 prior dispatch hash does not reproduce from the current R01 DOCX using the R02 patcher serializer; retain that historical mismatch as UNKNOWN rather than claiming reconciliation.'
}
(r02/'evidence/plain_text_snapshot_hash_reconciliation_0107_R01_R02.json').write_text(json.dumps(reconciliation,ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')
review_receipt={
 'doc_id':'0107','requirement_id':'I18-E15','candidate_sha256':digest(r02_docx.read_bytes()),
 'reviewer':'OpenCode v2 CLI; existing independent reviewer session ses_e9bf1c4edffeT95rFXd5SyNdmn; not the DOCX Builder',
 'requested_model':'opencode/mimo-v2.6-flash-free','actual_provider_serving_model':'UNKNOWN','local_session_cost':0.0,'provider_billing':'UNKNOWN',
 'verdict':'PASS_TEXT_SEMANTIC_DELTA','scope':'Exact R02 DOCX SHA text-semantic delta, including inherited PASS only for unchanged R01 text; not binary, layout, visual, or formal acceptance.',
 'response_path':str(response_path.relative_to(base.parent.parent.parent.parent)).replace('\\','/'),'response_sha256':digest(response_path.read_bytes()),'raw_stdout_path':'review/opencode_semantic_delta_R02_followup_stdout.jsonl','raw_stdout_sha256':raw_response_sha,
 'findings':'No content defect. Reviewer verified exact R02 SHA, 2 changed body blocks, cutoff-date column/value, 90-day inclusive window, recomputed rates and inheritance boundary. Reviewer identified a non-blocking snapshot-provenance discrepancy; R02 plain text snapshot now reproduced exactly, while R01 previous extraction hash remains unresolved.',
 'limitations':['Text semantic review only; no DOCX binary/OOXML inspection or rendered-page visual review.','Actual provider serving model and billing are UNKNOWN.','Formal association acceptance NOT_GRANTED.','R01 dispatched extraction hash 0dfd7360e0da5b330faefcc1fd91fa48be7012e8f99ff6aeb37ac228fb059301 was not reproduced from current R01 DOCX with R02 plain serialization.'],
 'status':'CANDIDATE_TEXT_SEMANTIC_PASS_VISUAL_PENDING_FORMAL_ACCEPTANCE_NOT_GRANTED'
}
(r02/'evidence/independent_semantic_review_receipt_0107_R02.json').write_text(json.dumps(review_receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')
print(json.dumps({'candidate_sha256':reconciliation['candidate_sha256'],'r01_hash':reconciliation['r01'],'r02_hash':reconciliation['r02'],'structural_hash':reconciliation['structural_diff'],'review_response_sha256':review_receipt['response_sha256'],'raw_stdout_sha256':raw_response_sha,'verdict':review_receipt['verdict']},ensure_ascii=False,separators=(',',':')))