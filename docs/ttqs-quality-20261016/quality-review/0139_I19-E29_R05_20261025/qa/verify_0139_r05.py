from pathlib import Path
from zipfile import ZipFile
from hashlib import sha256
from datetime import datetime
import json,re
from docx import Document
from PIL import Image, ImageChops

root=Path("E:/TTQS/TTQS_ONE_CODEX_RUNTIME/WORK/QUALITY_REBUILD_20261016")
r04=root/"0139_I19-E29_CANDIDATE_R04_20261025"
r05=root/"0139_I19-E29_CANDIDATE_R05_20261025"
docx=r05/"0139__證照通過率與就業率跨年度改善分析報告.docx"
d=Document(docx)
tables=[[[" ".join(p.text for p in c.paragraphs).strip() for c in row.cells] for row in t.rows] for t in d.tables]
records_table=next((t for t in tables if t and any("學員代碼" in c for c in t[0])),None)
records=[x for x in records_table[1:] if x and re.fullmatch(r"C2[345]-\d{2}",x[0])] if records_table else []
annual={}
for y in ("2023","2024","2025"):
    grp=[r for r in records if r[0].startswith("C"+y[-2:])]
    p=sum("通過" in r[2] or "(P)" in r[2] for r in grp)
    f=sum("不及格" in r[2] or "(F)" in r[2] for r in grp)
    a=sum("缺考" in r[2] or "(A)" in r[2] for r in grp)
    e=sum("已就業" in r[3] or "(E)" in r[3] for r in grp)
    n=sum("未就業" in r[3] or "(N)" in r[3] for r in grp)
    u=sum("失聯" in r[3] or "(U)" in r[3] for r in grp)
    annual[y]={"n":len(grp),"P":p,"F":f,"A":a,"E":e,"N":n,"U":u,"pass_rate_pct":round(100*p/(p+f),2),"employment_rate_pct":round(100*e/len(grp),2),"tracking_coverage_pct":round(100*(e+n)/len(grp),2)}
tot={"n":len(records),"P":sum(x["P"] for x in annual.values()),"F":sum(x["F"] for x in annual.values()),"A":sum(x["A"] for x in annual.values()),"E":sum(x["E"] for x in annual.values()),"N":sum(x["N"] for x in annual.values()),"U":sum(x["U"] for x in annual.values())}
tot.update({"pass_rate_pct":round(100*tot["P"]/(tot["P"]+tot["F"]),2),"employment_rate_pct":round(100*tot["E"]/tot["n"],2),"tracking_coverage_pct":round(100*(tot["E"]+tot["N"])/tot["n"],2)})
def inherited_size(run,p):
    if run.font.size: return run.font.size.pt
    style=p.style
    while style:
        if style.font.size: return style.font.size.pt
        style=style.base_style
    return None
body=[inherited_size(r,p) for p in d.paragraphs for r in p.runs if r.text.strip()]
cell=[inherited_size(r,p) for t in d.tables for row in t.rows for c in row.cells for p in c.paragraphs for r in p.runs if r.text.strip()]
body=[x for x in body if x is not None]; cell=[x for x in cell if x is not None]
text="\n".join(p.text for p in d.paragraphs)+"\n"+"\n".join(" | ".join(c for c in row) for t in tables for row in t)
for s in d.sections: text+="\n"+"\n".join(p.text for p in s.header.paragraphs+s.footer.paragraphs)
with ZipFile(docx) as z: crc_ok=z.testzip() is None
imgdir04=r04/"render_word_native"; imgdir05=r05/"render_word_native"
imgs04=sorted(imgdir04.glob("page-*.png")); imgs05=sorted(imgdir05.glob("page-*.png"))
diff_pages=[]
if len(imgs04)!=len(imgs05):
    visual_compare={"same_page_count":False,"r04_pages":len(imgs04),"r05_pages":len(imgs05),"pixel_identical_pages":[],"changed_pages":"ALL_PAGE_COUNT_DIFF"}
else:
    identical=[]
    for a,b in zip(imgs04,imgs05):
        im1=Image.open(a).convert("RGB"); im2=Image.open(b).convert("RGB")
        if im1.size==im2.size and ImageChops.difference(im1,im2).getbbox() is None: identical.append(a.name)
        else: diff_pages.append(a.name)
    visual_compare={"same_page_count":True,"r04_pages":len(imgs04),"r05_pages":len(imgs05),"pixel_identical_pages":identical,"changed_pages":diff_pages}
sha=sha256(docx.read_bytes()).hexdigest()
ver=[p.text for p in d.paragraphs if "候選 R05" in p.text]
correct_cell=tables[6][6][5] if len(tables)>6 and len(tables[6])>6 and len(tables[6][6])>5 else None
checks={
"zip_crc":crc_ok,"table_count":len(tables),"paragraph_count":len(d.paragraphs),
"record_count":len(records),"unique_record_ids":len({r[0] for r in records}),
"annual":annual,"totals":tot,"corrected_cell_table7_r7_c6":correct_cell,
"body_font_min_pt":min(body) if body else None,"table_font_min_pt":min(cell) if cell else None,
"sample_markers":{k:(k in text) for k in ["SAMPLE","SYNTHETIC","REAL=UNKNOWN/HOLD"]},
"formal_acceptance_not_granted":"NOT_GRANTED" in text or "未正式接受" in text,
"version_label_matches":bool(ver),"visual_compare":visual_compare,
"docx_sha256":sha,"docx_bytes":docx.stat().st_size,
"render_pdf_sha256":sha256((imgdir05/"0139_word_native.pdf").read_bytes()).hexdigest(),
"render_pages":len(imgs05),"created_local":datetime.now().astimezone().isoformat()
}
checks["pass"]=bool(crc_ok and len(tables)==14 and len(d.paragraphs)==177 and len(records)==34 and len({r[0] for r in records})==34 and correct_cell=="-8.33 pp" and checks["sample_markers"]["SAMPLE"] and checks["sample_markers"]["SYNTHETIC"] and len(imgs05)==23 and visual_compare["same_page_count"])
out=r05/"STATIC_QA_0139_R05.json"
out.write_text(json.dumps(checks,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
print(json.dumps(checks,ensure_ascii=False,indent=2))
