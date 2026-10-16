from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
import csv, hashlib, json, re, zipfile
from docx import Document
from pypdf import PdfReader
from PIL import Image, ImageChops

base = Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016')
work = base / '0036_R08'
doc = work / 'candidate/0036__課程設計進度與訓練規劃_R08.docx'
pdf = work / 'render/R08_final_review/0036_R08_word_export.pdf'
render_dir = work / 'render/R08_final_review'
old_render = base / '0036_R07/render/R07_word_com_baseline_144'
register = work / 'sources/SYNTHETIC_REPLACEMENT_REGISTER_0036_R08.csv'
expected_doc_sha = 'e2d4b3f5e906faaf3d473c7351bd79618a78dcf2c8f91a32a4e952d005ac5eae'
expected_pdf_sha = '238b275247b67bbe83f07422fb71fd369346b95a1b147879639e5a80bef1e7dd'
checks=[]
def digest(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def add(name, ok, evidence): checks.append({'check':name,'status':'PASS' if ok else 'FAIL','evidence':evidence})
add('candidate_exists', doc.is_file(), str(doc))
add('candidate_exact_sha', digest(doc)==expected_doc_sha, digest(doc))
with zipfile.ZipFile(doc) as z: bad=z.testzip()
add('docx_zip_integrity', bad is None, {'first_bad_member':bad})
reader=PdfReader(str(pdf)); page_count=len(reader.pages)
add('pdf_exact_sha', digest(pdf)==expected_pdf_sha, digest(pdf))
add('render_19_pages', page_count==19, {'pdf_pages':page_count,'png_pages':len(list(render_dir.glob('page-*.png')))})

# Verify page comparisons against the same Word COM renderer and 144 dpi baseline.
page_diffs=[]; sizes_ok=True; compared=[]
for i in range(1,18):
    new=render_dir/f'page-{i:02d}.png'; old=old_render/f'page-{i:02d}.png'
    if not new.is_file() or not old.is_file():
        sizes_ok=False; continue
    a=Image.open(new).convert('RGB'); b=Image.open(old).convert('RGB')
    if a.size != b.size: sizes_ok=False
    elif ImageChops.difference(a,b).getbbox(): page_diffs.append(i)
    compared.append(i)
new_pages=[render_dir/f'page-{i:02d}.png' for i in (18,19)]
new_dims={p.name:Image.open(p).size for p in new_pages if p.is_file()}
add('same_renderer_baseline_coverage',len(compared)==17 and sizes_ok,{'compared_pages':len(compared),'matching_dimensions':sizes_ok,'old_render':str(old_render)})
add('expected_legacy_page_deltas_only',page_diffs==[14,16],{'changed_pages':page_diffs})
add('new_appendix_pages_present',len(new_dims)==2 and all(s==(1191,1684) for s in new_dims.values()),new_dims)

D=Document(doc); alltext='\n'.join(p.text for p in D.paragraphs)+'\n'+'\n'.join(c.text for t in D.tables for row in t.rows for c in row.cells)
add('document_structure',len(D.tables)==37,{'paragraphs':len(D.paragraphs),'tables':len(D.tables)})
add('no_builder_internal_label','本候選' not in alltext,{'phrase_absent':'本候選' not in alltext})
add('sample_synthetic_boundary','SAMPLE' in alltext and 'SYNTHETIC' in alltext,{'sample': 'SAMPLE' in alltext,'synthetic':'SYNTHETIC' in alltext})
add('learner_assessor_tools',all(x in alltext for x in ('A-1','A-2','3/4','關鍵錯誤')),{'A1': 'A-1' in alltext,'A2':'A-2' in alltext,'3_of_4':'3/4' in alltext,'critical_error':'關鍵錯誤' in alltext})

# Recompute the document's displayed quantitative claims from its underlying rows.
base_rates=[]; target_rates=[]
for row in D.tables[12].rows[1:]:
    m=re.search(r'(\d+)\s*/\s*(\d+)',row.cells[2].text); base_rates.append((int(m.group(1)),int(m.group(2))) if m else None)
    m=re.search(r'(\d+)\s*/\s*(\d+)',row.cells[3].text); target_rates.append((int(m.group(1)),int(m.group(2))) if m else None)
add('baseline_and_target_rates',base_rates==[(4,10),(5,10),(6,10),(3,10)] and target_rates==[(8,10)]*4,{'baseline':base_rates,'target':target_rates})
weights=[.4,.3,.2,.1]; scored=[]
for row in D.tables[13].rows[1:]:
    vals=[float(x.strip()) for x in row.cells[1].text.split('/')]
    calculated=sum(x*w for x,w in zip(vals,weights)); displayed=float(re.search(r'(\d+(?:\.\d+)?)',row.cells[3].text).group(1))
    scored.append((round(calculated,2),displayed))
add('weighted_option_scores',scored==[(3.3,3.3),(4.7,4.7),(3.5,3.5)],scored)
time_rows=D.tables[17].rows[1:]
minutes=sum(int(r.cells[1].text.strip()) for r in time_rows)
instruction=sum(int(r.cells[1].text.strip()) for r in time_rows if '休息' not in r.cells[2].text)
breaks=sum(int(r.cells[1].text.strip()) for r in time_rows if '休息' in r.cells[2].text)
add('schedule_minutes',minutes==210 and instruction==190 and breaks==20,{'total':minutes,'instruction':instruction,'breaks':breaks})
pairs=[]
for row in D.tables[24].rows[1:]:
    before=int(row.cells[1].text.strip()); after=int(row.cells[2].text.strip()); delta=int(row.cells[3].text.strip().replace('+','')); pairs.append((before,after,delta))
pre=sum(a for a,b,c in pairs); post=sum(b for a,b,c in pairs); gain=sum(c for a,b,c in pairs); at_target=sum(b>=4 for a,b,c in pairs)
add('paired_pre_post_math',len(pairs)==10 and all(b-a==c for a,b,c in pairs) and (pre,post,gain,at_target)==(24,40,16,8),{'n':len(pairs),'pre_total':pre,'post_total':post,'gain_total':gain,'at_least_4':at_target,'pre_average':pre/len(pairs),'post_average':post/len(pairs),'gain_average':gain/len(pairs)})

with register.open(encoding='utf-8-sig',newline='') as f:
    reader_csv=csv.DictReader(f); fields=reader_csv.fieldnames; regrows=list(reader_csv)
ids=[r['fact_id'] for r in regrows]
add('synthetic_register_identity',len(regrows)==75 and len(set(ids))==75 and all(r.get('doc_id')=='0036' for r in regrows),{'rows':len(regrows),'unique_fact_ids':len(set(ids)),'doc_ids':sorted(set(r.get('doc_id','') for r in regrows)),'sha256':digest(register)})
required=['fact_id','case_id','doc_id','body_locator','synthetic_value','assumption_reason','real_source_to_replace','real_owner_role_to_verify','status']
empty=[r.get('fact_id','?') for r in regrows if any(not (r.get(k) or '').strip() for k in required)]
add('synthetic_register_required_fields',not empty,{'empty_rows':empty[:10],'empty_count':len(empty)})
locator_errors=[]
for r in regrows:
    nums=[int(x) for x in re.findall(r'p\.(\d+)',r['body_locator'])]
    if not nums or min(nums)<1 or max(nums)>19: locator_errors.append({'fact_id':r['fact_id'],'locator':r['body_locator']})
add('synthetic_register_page_locators',not locator_errors,{'invalid_count':len(locator_errors),'examples':locator_errors[:8]})

appendix_text=' '.join((reader.pages[17].extract_text() or '').split())+' '+' '.join((reader.pages[18].extract_text() or '').split())
add('appendix_scenario_and_rubric',all(x in appendix_text for x in ('A-1','A-2','\u60c5\u5883','1 \u5206\u7b54\u6848\u9328\u9ede','0 \u5206','3/4')),{'appendix_pages':2,'scenario_rows':len(D.tables[35].rows)-1,'rubric_rows':len(D.tables[36].rows)-1})

failed=[x for x in checks if x['status']!='PASS']
now=datetime.now(ZoneInfo('Asia/Taipei')).isoformat(timespec='seconds')
receipt={'project_id':'TTQS_ONE','doc_id':'0036','official_requirement':'I07-E04','official_checkbox_serial':'036','candidate_revision':'R08','review_type':'technical_integrity','reviewer_lane':'current control lane; not an independent semantic reviewer','reviewed_at':now,'candidate_path':str(doc),'candidate_sha256':digest(doc),'render_pdf_path':str(pdf),'render_pdf_sha256':digest(pdf),'register_path':str(register),'register_sha256':digest(register),'status':'PASS' if not failed else 'FAIL','semantic_review_status':'PENDING_INDEPENDENT_REVIEW','checks':checks}
visual={'project_id':'TTQS_ONE','doc_id':'0036','candidate_revision':'R08','review_type':'visual_layout','reviewer_lane':'current control lane; not an independent semantic reviewer','reviewed_at':now,'candidate_path':str(doc),'candidate_sha256':digest(doc),'render_pdf_sha256':digest(pdf),'page_count':19,'page_image_count':len(list(render_dir.glob('page-*.png'))),'render_engine':'Microsoft Word COM export, 144 dpi PNG review','comparison_baseline':str(old_render),'baseline_pages_compared':compared,'pixel_changed_pages':page_diffs,'directly_inspected_final_pages':[14,16,18,19],'inherited_pixel_identical_pages':[i for i in range(1,18) if i not in page_diffs],'status':'PASS' if not failed and sizes_ok and page_diffs==[14,16] and len(new_dims)==2 else 'FAIL','semantic_review_status':'PENDING_INDEPENDENT_REVIEW','observations':['No clipped text, broken tables, missing glyphs, overlap, or footer displacement observed on directly inspected changed/new pages.','Pages 1-17 compared to same Word COM/144 dpi baseline; only pages 14 and 16 differ.','Pages 18-19 contain the new learner response sheet and assessor scoring key on separate readable pages.']}
qa=work/'qa'; qa.mkdir(exist_ok=True)
(qa/'0036_R08_technical_review_exact.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
(qa/'0036_R08_visual_review_exact.json').write_text(json.dumps(visual,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'status':receipt['status'],'checks_passed':len(checks)-len(failed),'checks_total':len(checks),'failed':[{'check':x['check'],'evidence':x['evidence']} for x in failed],'docx_sha256':receipt['candidate_sha256'],'pdf_sha256':receipt['render_pdf_sha256'],'register_sha256':receipt['register_sha256'],'pixel_changed_pages':page_diffs,'technical_receipt':str(qa/'0036_R08_technical_review_exact.json'),'visual_receipt':str(qa/'0036_R08_visual_review_exact.json')},ensure_ascii=False))
