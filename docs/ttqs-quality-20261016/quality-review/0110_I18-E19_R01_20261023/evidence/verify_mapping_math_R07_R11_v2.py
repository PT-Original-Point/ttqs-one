from __future__ import annotations
import hashlib, json
from datetime import datetime, timezone
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / 'candidate' / '0110_I18-E19_完整問卷暨行動計畫工具_R07.docx'
MAP = ROOT / 'evidence' / 'item_construct_mapping_R10_1_R03.json'
MAP_QA = ROOT / 'evidence' / 'python_mapping_validation_R03.json'
SOURCE = ROOT / 'source' / 'SOURCE_PACKET_0110_I18-E19_R01.md'
BLUEPRINT = ROOT / 'source' / 'UNIQUE_BLUEPRINT_0110_I18-E19_R01.md'
TECH = ROOT / 'evidence' / 'technical_review_R07_final.json'
SEM = ROOT / 'evidence' / 'semantic_review_receipt_R07.json'
VIS = ROOT / 'evidence' / 'visual_review_R07_final.json'
OUT = ROOT / 'evidence' / 'item_construct_mapping_math_R07_bound_R11_v2.json'

def sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()

def read_json(p: Path):
    return json.loads(p.read_text(encoding='utf-8-sig'))

mapping = read_json(MAP)
map_qa = read_json(MAP_QA)
tech = read_json(TECH)
semantic = read_json(SEM)
visual = read_json(VIS)
doc = Document(DOCX)
all_text = '\n'.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for row in t.rows for c in row.cells])
rows = [[c.text.replace('\n', ' ').strip() for c in row.cells] for row in doc.tables[6].rows]
row_map = {r[0]: r for r in rows[1:] if r and r[0]}
source_hash = sha(SOURCE)
blueprint_hash = sha(BLUEPRINT)
candidate_hash = sha(DOCX)

expected_constructs = {
 'L01':'操作指引理解', 'L02':'自我效能感知（非組織績效）', 'L03':'查核步驟執行',
 'L04':'實際使用頻率', 'L05':'情境處置能力', 'L06':'交接執行能力',
 'L07':'資料完整度變化', 'L08':'交接流程變化', 'L09':'協作變化',
 'C03':'服務體驗變化', 'C04':'資料物資完整正確變化', 'C05':'服務連續性變化',
 'C06':'應變回應時間變化', 'C07':'整體績效變化'
}
label_notes = {
 'L02': '自我效能變化 → 自我效能感知（非組織績效）；題目仍明確指學員主觀執行信心前後變化。',
 'L08': '交接時間與流程變化 → 交接流程變化；觀察焦點保留交接等待時間前後變化。',
 'C04': '資料與物資完整正確變化 → 資料物資完整正確變化；觀察焦點保留物資與資料的正確、齊全程度。',
 'C07': '整體可觀察績效變化 → 整體績效變化；觀察焦點仍明示整體可觀察組織績效。'
}
item_checks = []
for item in mapping['items']:
    item_id = item['id']
    row = row_map.get(item_id, [])
    present = bool(row) and len(row) >= 4
    construct = row[2] if present else ''
    obs = row[1] if present else ''
    scale_cell = row[3] if present else ''
    if item['scale'] == 'AGREE_1_5':
        scale_ok = present and '同意尺度' in scale_cell
        scale_basis = 'table-row agree scale'
    elif item['scale'] == 'FREQUENCY_4_TO_1_PLUS_0':
        scale_ok = present and '頻率' in scale_cell and '不作Likert平均' in scale_cell and '有利率' in scale_cell
        scale_basis = 'dedicated frequency scale; excluded from Likert mean/beneficial rate'
    else:
        global_change = '變化尺度（L02、L07–L09、C03–C07）' in all_text
        if item_id in {'L02','L07','L08','L09'}:
            scale_ok = present and '變化尺度' in scale_cell and global_change
            scale_basis = 'table row plus global change-scale rule'
        else:
            required = {
              'C03': '指定期間內直接觀察者作答',
              'C04': '改用前後比較，不用滿意度錨點',
              'C05': '明列觀察期間及比較基準',
              'C06': '改為可觀察的比較題，不用同意度錨點',
              'C07': '單題報告；不與 C03–C06 合成分數'
            }.get(item_id, '')
            scale_ok = present and global_change and bool(required) and required in scale_cell
            scale_basis = 'global change-scale rule plus item-specific observation rule'
    item_checks.append({
        'item_id': item_id,
        'candidate_row_present': present,
        'candidate_construct_matches_current_R07_label': construct == expected_constructs.get(item_id),
        'mapping_label_delta_note': label_notes.get(item_id),
        'observation_focus': obs,
        'scale_semantics_match': bool(scale_ok),
        'scale_basis': scale_basis,
        'locator': f'DOCX table 7 / row {item_id}'
    })

case = mapping['synthetic_case']
baseline = case['baseline_complete'] / case['baseline_sampled'] * 100
interim = case['followup_complete'] / case['followup_sampled'] * 100
change_pp = interim - baseline
target = case['target_rate'] * 100
math_detail = {
 'case_id': case['case_id'],
 'status': case['status'],
 'baseline': {'fraction': f"{case['baseline_complete']}/{case['baseline_sampled']}", 'percent': baseline, 'candidate_shows_6_over_8_and_75_percent': '6/8' in all_text and '75.0%' in all_text},
 'D0_plus_30': {'fraction': f"{case['followup_complete']}/{case['followup_sampled']}", 'percent': interim, 'candidate_shows_7_over_8_and_87_5_percent': '7/8' in all_text and '87.5%' in all_text},
 'change_percentage_points': {'recomputed': change_pp, 'source_mapping': case['percentage_point_change'] * 100, 'candidate_shows_plus_12_5_pp': '+12.5 個百分點' in all_text or '+12.5個百分點' in all_text},
 'target_and_action': {'target_percent': target, 'interim_below_target': interim < target, 'source_mapping_target_met': case['target_met'], 'source_mapping_decision': case['decision'], 'candidate_marks_REVISE': 'REVISE' in all_text, 'D0_plus_60_unavailable': 'D0＋60' in all_text and '尚未取得' in all_text, 'candidate_prohibits_CLOSE': '不可標示 CLOSE' in all_text or '不得標示 CLOSE' in all_text, 'fixed_8_requires_8_of_8_equals_100': '8/8＝100.0%' in all_text or '8/8 = 100.0%' in all_text}
}

receipt_bindings = {
 'technical_candidate_sha_matches': tech.get('candidate_sha256') == candidate_hash and tech.get('technical_content_status') == 'PASS',
 'semantic_candidate_sha_and_verdict_match': semantic.get('candidate_sha256') == candidate_hash and semantic.get('verdict') == 'PASS_CONTENT_TEXT_ONLY',
 'visual_candidate_sha_and_layout_pass': visual.get('candidate_sha256') == candidate_hash and visual.get('layout_status') == 'PASS' and visual.get('pdf_pages') == 26,
 'visual_pdf_sha_matches_local_bytes': Path(visual.get('pdf_path','')).is_file() and sha(Path(visual['pdf_path'])) == visual.get('pdf_sha256'),
 'visual_page_hashes_match_local_rasters': False
}
page_map = visual.get('render_png_sha256', {})
page_dir = ROOT / 'evidence' / 'render_R07_word_native'
page_results = {}
for key, expected_hash in page_map.items():
    number = int(key.rsplit('_', 1)[-1])
    image = page_dir / f'final2-page-{number:02d}.png'
    page_results[key] = image.is_file() and sha(image) == expected_hash
receipt_bindings['visual_page_hashes_match_local_rasters'] = len(page_results) == 26 and all(page_results.values())

label_differences = []
for item in mapping['items']:
    current = row_map.get(item['id'], ['','','',''])
    if item['construct'] != (current[2] if len(current) > 2 else ''):
        label_differences.append({'item_id':item['id'], 'historical_mapping_label':item['construct'], 'current_R07_label':current[2] if len(current)>2 else None, 'reconciliation':label_notes.get(item['id'],'No pre-recorded label reconciliation')})

checks = {
 'candidate_exact_sha_is_expected_R07': candidate_hash == '9a306bf5712437ee578867b745e7d0811959decaa862ed540a82e7e8553bd3c2',
 'document_and_requirement_binding': mapping['doc_id'] == '0110' and mapping['requirement_id'] == 'I18-E19' and mapping['official_checkbox_serial'] == 'UNKNOWN',
 'source_packet_exact_sha_matches_mapping_and_prior_receipt': source_hash == mapping['inputs']['source_packet_sha256'] == tech.get('official_source_packet_sha256'),
 'unique_blueprint_exact_sha_matches_mapping_and_prior_receipt': blueprint_hash == mapping['inputs']['unique_blueprint_sha256'] == tech.get('unique_blueprint_sha256'),
 'all_14_unique_items_present': len(mapping['items']) == 14 and len(set(x['id'] for x in mapping['items'])) == 14 and set(x['id'] for x in mapping['items']) == set(row_map),
 'all_current_constructs_and_observation_focuses_reconciled': all(x['candidate_row_present'] and x['candidate_construct_matches_current_R07_label'] and (x['mapping_label_delta_note'] is not None or x['candidate_construct_matches_current_R07_label']) for x in item_checks) and len(label_differences) == 4 and all(any(x['item_id']==k and all(term in x['observation_focus'] for term in terms) for x in item_checks) for k,terms in {'L02':['自我效能','前後','變化'],'L08':['交接等待時間','前後變化'],'C04':['物資及資料','正確、齊全'],'C07':['整體可觀察','組織績效']}.items()),
 'all_scale_rules_match_item_use': all(x['scale_semantics_match'] for x in item_checks),
 'groups_and_items_remain_separate': '預設按題項逐一呈現' in all_text and '分開計算與報告' in all_text and '不合併總分' in all_text and '不與 C03–C06 合成分數' in all_text,
 'sample_math_and_interim_decision_are_consistent': baseline == 75.0 and interim == 87.5 and change_pp == 12.5 and case['percentage_point_change'] * 100 == 12.5 and target == 90.0 and interim < target and case['target_met'] is False and case['decision'] == 'REVISE' and math_detail['baseline']['candidate_shows_6_over_8_and_75_percent'] and math_detail['D0_plus_30']['candidate_shows_7_over_8_and_87_5_percent'] and math_detail['change_percentage_points']['candidate_shows_plus_12_5_pp'] and math_detail['target_and_action']['candidate_marks_REVISE'] and math_detail['target_and_action']['D0_plus_60_unavailable'] and math_detail['target_and_action']['candidate_prohibits_CLOSE'] and math_detail['target_and_action']['fixed_8_requires_8_of_8_equals_100'],
 'sample_synthetic_label_present': ('SAMPLE／SYNTHETIC' in all_text or 'SAMPLE/SYNTHETIC' in all_text) and case['status'] == 'SAMPLE/SYNTHETIC ONLY',
 'prior_item_mapping_validator_passed': sha(MAP) == map_qa.get('map_sha256') and map_qa.get('unique_complete_ids') is True and map_qa.get('L04_separate_frequency') is True and map_qa.get('scale_homogeneity') is True and map_qa.get('case_arithmetic',{}).get('decision') == 'REVISE',
 'all_existing_qa_receipts_bind_same_exact_sha': all(receipt_bindings.values())
}

out = {
 'work_id':'TTQS-0110-I18-E19-MAPPING-MATH-READBACK-R11-20261023-V2',
 'project_id':'TTQS_ONE','doc_id':'0110','requirement_id':'I18-E19','official_checkbox_serial':'UNKNOWN_NOT_ASSERTED',
 'candidate_path':str(DOCX),'candidate_sha256':candidate_hash,
 'source_commit':mapping['source_commit'],'source_packet_sha256':source_hash,'unique_blueprint_sha256':blueprint_hash,
 'mapping_source_sha256':sha(MAP),'prior_mapping_validation_sha256':sha(MAP_QA),
 'item_checks':item_checks,'label_differences_vs_historical_mapping':label_differences,
 'sample_case_math':math_detail,'existing_receipt_bindings':receipt_bindings,'visual_page_hash_checks':page_results,
 'checks':checks,'result':'PASS' if all(checks.values()) else 'FAIL',
 'supersedes_note':'Additive R11 V2 validation; the earlier R11 receipt is preserved. V2 adds source-file SHA verification, explicit label-delta reconciliation, scale membership and actual raster hash checks. It does not change the DOCX.',
 'scope':'Deterministic source-to-current-table mapping, response-scale application, group separation, synthetic arithmetic, and exact-SHA technical/semantic/visual receipt and page-hash binding. This is not a new semantic review or formal acceptance.',
 'candidate_state':'CANDIDATE_NOT_ACCEPTED','formal_acceptance':'NOT_GRANTED','real_evidence':'UNKNOWN/HOLD',
 'observed_at_utc':datetime.now(timezone.utc).isoformat()
}
OUT.write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'result':out['result'],'candidate_sha256':candidate_hash,'checks':checks,'label_differences':label_differences,'failed_page_hashes':[k for k,v in page_results.items() if not v],'receipt':str(OUT)},ensure_ascii=False))

