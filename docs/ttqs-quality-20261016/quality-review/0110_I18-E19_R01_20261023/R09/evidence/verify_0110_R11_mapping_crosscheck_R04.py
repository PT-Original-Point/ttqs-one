import json, re, hashlib, pathlib, datetime
from docx import Document
root=pathlib.Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0110_I18-E19_R01_20261023')
evidence=root/'evidence'
candidate=root/'candidate/0110_I18-E19_完整問卷暨行動計畫工具_R09.docx'
source=root/'source/SOURCE_PACKET_0110_I18-E19_R01.md'
blueprint=root/'source/UNIQUE_BLUEPRINT_0110_I18-E19_R01.md'
r03=evidence/'0110_R09_item_construct_math_mapping_R11_R03.json'
sha=lambda b:hashlib.sha256(b).hexdigest()
doc=Document(str(candidate))
qa=json.loads(r03.read_text(encoding='utf-8'))
# Source-to-document item/construct crosscheck against the actual printed mapping table.
map_rows=[]
for row in doc.tables[6].rows[1:]:
    c=[' '.join(x.text.split()) for x in row.cells]
    map_rows.append({'item_id':c[0],'focus':c[1],'construct':c[2],'scale_rule':c[3]})
expected=qa['item_mapping']['rows']
assert map_rows==expected, 'Mapping table differs from deterministic QA mapping'
assert len(map_rows)==14 and [r['item_id'] for r in map_rows]==qa['item_mapping']['question_ids']
# Parse SAMPLE learner source rows and recompute per-item totals, excluding the explicitly invalid L06 respondent.
learner={}
for row in doc.tables[30].rows[1:]:
    who=' '.join(row.cells[0].text.split())
    body=' '.join(row.cells[1].text.split())
    vals={k:int(v) for k,v in re.findall(r'\b(L\d{2})\s*=\s*(\d+)',body)}
    learner[who]=vals
assert len(learner)==8 and 'L06' in learner
client={}
for row in doc.tables[32].rows[1:]:
    who=' '.join(row.cells[0].text.split())
    body=' '.join(row.cells[1].text.split())
    assert re.search(r'\bC01\s*=\s*是',body), f'Client eligibility gate not Yes: {who}'
    vals={k:int(v) for k,v in re.findall(r'\b(C0[3-7])\s*=\s*(\d+)',body)}
    client[who]=vals
assert len(client)==6
calc={}
for item in [f'L{i:02d}' for i in range(1,10) if i!=4]+[f'C{i:02d}' for i in range(3,8)]:
    rows=learner if item.startswith('L') else client
    vals=[v[item] for who,answers in rows.items() if (item.startswith('C') or who!='L06') for v in [answers] if item in v]
    valid=[v for v in vals if 1<=v<=5]
    calc[item]={'valid_N':len(valid),'zero_count':sum(v==0 for v in vals),'invalid_blank_duplicate_count':0,'sum':sum(valid),'mean_3dp':round(sum(valid)/len(valid)+1e-12,3),'favorable_4_5_n':sum(v>=4 for v in valid),'favorable_4_5_pct_1dp':round(100*sum(v>=4 for v in valid)/len(valid)+1e-12,1)}
# L04 is a frequency distribution, not a Likert score.
l04=[answers['L04'] for who,answers in learner.items() if who!='L06']
calc['L04']={'eligible_opportunities':len(l04),'frequency_counts':{str(i):l04.count(i) for i in range(1,5)},'not_scored':True}
expected_results={x['item_id']:x for x in qa['sample_recalculation']['results']}
for item,value in calc.items():
    old=expected_results[item]
    if item=='L04':
        assert value['eligible_opportunities']==old['valid_opportunity_count']
        assert value['frequency_counts']==old['frequency_counts']
    else:
        for k in ('valid_N','zero_count','sum','mean_3dp','favorable_4_5_n','favorable_4_5_pct_1dp'):
            assert value[k]==old[k], f'{item}.{k}: {value[k]} != {old[k]}'
# Compare recomputed numbers to the exact visible summary-table rows (table index 34 in python-docx).
printed={}
for row in doc.tables[34].rows[1:]:
    label=' '.join(row.cells[0].text.split())
    m=re.search(r'\b([LC]\d{2})\b',label)
    if m: printed[m.group(1)]=[' '.join(c.text.split()) for c in row.cells]
assert len(printed)==14, f'Expected 14 printed item rows, got {len(printed)}'
for item,v in calc.items():
    cells=printed[item]
    if item=='L04':
        freq={str(int(a)):int(b) for a,b in re.findall(r'([1-4])\s*[=:＝]\s*(\d+)',cells[4])}
        assert freq==v['frequency_counts'] and '不計分' in cells[3] and '不計有利比例' in cells[4]
    else:
        assert re.search(rf'\bN\s*=\s*{v["valid_N"]}\b',cells[1]), (item,cells[1],v)
        assert cells[2]==str(v['sum']) and cells[3]==f'{v["mean_3dp"]:.3f}', (item,cells[2:4],v)
        pct=re.search(r'(\d+)\s*/\s*(\d+).*?([\d.]+)\s*%',cells[4])
        assert pct and int(pct.group(1))==v['favorable_4_5_n'] and int(pct.group(2))==v['valid_N'] and float(pct.group(3))==v['favorable_4_5_pct_1dp'], (item,cells[4],v)
# Operational example directly read from the printed DOCX Table 28.
ops=[[' '.join(c.text.split()) for c in row.cells] for row in doc.tables[28].rows]
ops_text=' | '.join(' | '.join(row) for row in ops)
checks={'6_over_8_is_75':6/8*100==75.0,'7_over_8_is_87_5':7/8*100==87.5,'change_is_12_5pp':(7/8-6/8)*100==12.5,'fixed_n8_90pct_requires_8':next(n for n in range(9) if n/8*100>=90)==8,'printed_table_has_6_8_7_8_and_revise':all(x in ops_text for x in ('6／8','7／8','REVISE'))}
assert all(checks.values()),checks
receipt={'schema':'ttqs-one-0110-r11-direct-docx-crosscheck-r04-v1','project_id':'TTQS_ONE','work_id':'TTQS-0110-I18-E19-R11-MAPPING-CROSSCHECK-R04','observed_at_tpe':datetime.datetime.now().astimezone().isoformat(timespec='seconds'),'doc_id':'0110','requirement_id':'I18-E19','official_checkbox_serial':'UNKNOWN','candidate_file':candidate.name,'candidate_sha256':sha(candidate.read_bytes()),'source_packet_sha256':sha(source.read_bytes()),'unique_blueprint_sha256':sha(blueprint.read_bytes()),'prior_r03_receipt_sha256':sha(r03.read_bytes()),'mapping_table_index':6,'summary_table_index':34,'item_construct_mapping_count':len(map_rows),'mapping_matches_candidate_and_r03':True,'sample_rows':{'learner':len(learner),'eligible_learner':len(learner)-1,'invalid_learner':'L06 HOLD/INVALID','client':len(client),'client_C01_yes_gate':True},'recomputed_results':calc,'operational_example':{'baseline':'6/8=75.0%','D0_plus_30':'7/8=87.5%','change_pp':12.5,'target_fixed_n8':'8/8=100.0%','decision':'REVISE at D0+30; D0+60 pending; not CLOSE','checks':checks},'result':'PASS_DETERMINISTIC_DIRECT_CANDIDATE_CROSSCHECK','scope_limitations':['Deterministic comparison only; not an independent semantic review.','All data are SAMPLE/SYNTHETIC; establishes no REAL association outcomes.','No formal Human acceptance or approval is claimed.','Serving model identity is outside this check.']}
out=evidence/'0110_R11_mapping_crosscheck_R04.json'
out.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'result':receipt['result'],'candidate_sha256':receipt['candidate_sha256'],'source_packet_sha256':receipt['source_packet_sha256'],'unique_blueprint_sha256':receipt['unique_blueprint_sha256'],'mapping_count':len(map_rows),'valid_items_recomputed':len(calc),'math':receipt['operational_example'],'receipt_file':str(out),'receipt_sha256':sha(out.read_bytes())},ensure_ascii=False,indent=2))
