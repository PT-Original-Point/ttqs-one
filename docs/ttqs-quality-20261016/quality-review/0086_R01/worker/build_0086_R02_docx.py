from __future__ import annotations
import csv, hashlib, json, re, subprocess
from datetime import datetime, timezone
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
BODY = ROOT / "candidate" / "0086_R02_body.md"
REGISTER = ROOT / "evidence" / "synthetic_replacement_register_R02.csv"
BLUEPRINT = ROOT / "inputs" / "UNIQUE_BLUEPRINT_0086_R02.md"
SOURCE_PACKET = ROOT / "inputs" / "SOURCE_PACKET_0086_R01.md"
REF = ROOT / "inputs" / "pandoc_reference_0086_R02.docx"
DOCX = ROOT / "candidate" / "0086_R02.docx"
RECEIPT = ROOT / "evidence" / "build_receipt_R02.json"
PANDOC = Path(r"C:\Users\J\AppData\Local\Pandoc\pandoc.exe")
FONT = "Microsoft JhengHei"

def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()

def font_run(run, size=None, bold=None, color=None):
    run.font.name = FONT
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + attr), FONT)

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

def shade(cell, fill="D9EAF7"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def field_run(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = instruction
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t"); text.text = "1"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    font_run(run, size=8, color=(90, 100, 110))

def create_reference():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.8); sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(1.8); sec.right_margin = Cm(1.8)
    styles = doc.styles
    specs = {
        "Normal": (11, (35, 45, 55)),
        "Body Text": (11, (35, 45, 55)),
        "Title": (22, (24, 57, 90)),
        "Subtitle": (12, (80, 90, 100)),
        "Heading 1": (16, (31, 78, 121)),
        "Heading 2": (13, (31, 78, 121)),
        "Heading 3": (11.5, (55, 90, 125)),
        "Heading 4": (10.5, (55, 90, 125)),
    }
    for name, (size, color) in specs.items():
        st = styles[name]
        st.font.name = FONT; st.font.size = Pt(size); st.font.color.rgb = RGBColor(*color)
        rpr = st._element.get_or_add_rPr()
        rf = rpr.rFonts
        if rf is None:
            rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            rf.set(qn("w:" + attr), FONT)
        if name.startswith("Heading"):
            st.paragraph_format.keep_with_next = True
            st.paragraph_format.space_before = Pt(9)
            st.paragraph_format.space_after = Pt(4)
    for name in ("Normal", "Body Text"):
        styles[name].paragraph_format.space_after = Pt(4)
        styles[name].paragraph_format.line_spacing = 1.12
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = fp.add_run("候選示範文件｜SAMPLE / SYNTHETIC　　頁 ")
    font_run(rr, size=8, color=(90, 100, 110))
    field_run(fp, "PAGE")
    REF.parent.mkdir(parents=True, exist_ok=True)
    doc.save(REF)

def validate_inputs(body: str):
    if not body.startswith("# "):
        raise ValueError("Body must start with a title")
    forbidden = ("<<<REGISTER_CSV>>>", "<<<END_REGISTER_CSV>>>", "G0–G3", "≥80.0%", "60.0–79.9%")
    for term in forbidden:
        if term in body:
            raise ValueError(f"Unrepaired draft marker/defect remains: {term}")
    if not all(x in body for x in ("I17C-E05", "方框 092", "SAMPLE", "SYNTHETIC", "NO_LINK")):
        raise ValueError("Identity, synthetic boundary, or neighbor check missing")
    if not PANDOC.exists():
        raise FileNotFoundError(PANDOC)

def main():
    body = BODY.read_text(encoding="utf-8-sig")
    validate_inputs(body)
    with REGISTER.open("r", encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    required = {"fact_id","case_id","doc_id","locator","field","value","assumption_basis","real_source_to_replace","owner_role","status","跨文件引用狀態"}
    if not rows or set(rows[0]) != required:
        raise ValueError("Synthetic replacement register header invalid")
    if len(rows) != 53 or any(r["跨文件引用狀態"] != "NO_LINK" for r in rows):
        raise ValueError("R02 replacement register row count or cross-document status invalid")
    if any(r["status"] != "SYNTHETIC_REPLACE_BEFORE_REAL_USE" for r in rows):
        raise ValueError("Register contains non-synthetic state")
    if len({r["fact_id"] for r in rows}) != len(rows):
        raise ValueError("Duplicate fact_id in replacement register")
    if body.count("### 表 A") != 1 or not all(body.count(f"### 表 {letter}｜") == 1 for letter in "ABCDEFGH"):
        raise ValueError("One or more form sections A-H missing")
    create_reference()
    temp = ROOT / "worker" / "pandoc_body_0086_R02.md"
    temp.write_text(body, encoding="utf-8")
    cmd = [str(PANDOC), str(temp), "--from=markdown+pipe_tables+strikeout", "--to=docx",
           f"--reference-doc={REF}", f"--output={DOCX}"]
    proc = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace")
    if proc.returncode:
        raise RuntimeError(f"Pandoc failed ({proc.returncode}): {proc.stderr[-3000:]}")
    doc = Document(DOCX)
    for p in doc.paragraphs:
        is_heading = bool(p.style and p.style.name.startswith("Heading"))
        if is_heading:
            p.paragraph_format.keep_with_next = True
            if p.style.name == "Heading 2" and p.text.startswith("14. 文件維護與交付檢核"):
                p.paragraph_format.page_break_before = True
            if p.text.startswith("5.2 示範母體清冊"):
                p.paragraph_format.page_break_before = True
            if p.style.name == "Heading 2" and p.text.startswith("6. 從委託到關閉的執行流程"):
                p.paragraph_format.page_break_before = True
            for run in p.runs:
                size = 16 if p.style.name == "Heading 1" else 13 if p.style.name == "Heading 2" else 11.5
                font_run(run, size=size, bold=True)
        else:
            p.paragraph_format.widow_control = True
            for run in p.runs:
                font_run(run, size=11)
    for table in doc.tables:
        table.style = "Table Grid"
        table.autofit = True
        if table.rows:
            set_repeat_table_header(table.rows[0])
            for cell in table.rows[0].cells:
                shade(cell)
        for ri, row in enumerate(table.rows):
            tr_pr = row._tr.get_or_add_trPr()
            cant_split = OxmlElement("w:cantSplit")
            cant_split.set(qn("w:val"), "1")
            tr_pr.append(cant_split)
            for cell in row.cells:
                cell.vertical_alignment = 1
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.line_spacing = 1.0
                    for run in p.runs:
                        font_run(run, size=9.5, bold=(ri == 0 or run.bold))
    for sec in doc.sections:
        sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
        sec.top_margin = Cm(1.8); sec.bottom_margin = Cm(1.8)
        sec.left_margin = Cm(1.8); sec.right_margin = Cm(1.8)
    doc.core_properties.title = "神秘客實地稽核：程序與真實外部稽核原件邊界操作指引"
    doc.core_properties.subject = "TTQS I17c SAMPLE/SYNTHETIC 方法示範；非實際稽核紀錄"
    doc.save(DOCX)
    receipt = {
        "builder_sha256": sha(Path(__file__)),
        "project_id": "TTQS_ONE", "document_id": "0086", "requirement_id": "I17C-E05",
        "official_checkbox_serial": "092", "revision": "R02",
        "source_packet_sha256": sha(SOURCE_PACKET), "unique_blueprint_sha256": sha(BLUEPRINT),
        "body_sha256": sha(BODY), "register_sha256": sha(REGISTER),
        "docx_path": str(DOCX), "docx_sha256": sha(DOCX), "docx_bytes": DOCX.stat().st_size,
        "paragraphs": len(doc.paragraphs), "tables": len(doc.tables),
        "synthetic_register_rows": len(rows), "pandoc_exit_code": proc.returncode,
        "pandoc_stderr": proc.stderr[-2000:],
        "built_at_utc": datetime.now(timezone.utc).isoformat(),
        "state": "CANDIDATE_NOT_ACCEPTED; visual and independent semantic QA pending"
    }
    RECEIPT.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(receipt, ensure_ascii=False))

if __name__ == "__main__":
    main()
