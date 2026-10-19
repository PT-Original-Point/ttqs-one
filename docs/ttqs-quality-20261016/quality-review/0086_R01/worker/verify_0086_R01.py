import csv, hashlib, json, re, zipfile
from datetime import datetime, timezone
from pathlib import Path
from docx import Document
root=Path(r"E:\TTQS\_github_sync_0064\docs\ttqs-quality-20261016\quality-review\0086_R01")
docx=root/"candidate/0086_R01.docx"; bodyp=root/"candidate/0086_R01_body.md"; regp=root/"evidence/synthetic_replacement_register_R01.csv"
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
body=bodyp.read_text(encoding="utf-8"); doc=Document(docx); alltext="\n".join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells])
with zipfile.ZipFile(docx) as z: bad=z.testzip(); member_count=len(z.namelist())
with regp.open(encoding="utf-8-sig",newline="") as f: rows=list(csv.DictReader(f))
expected_headers={"fact_id","case_id","doc_id","locator","field","value","assumption_basis","real_source_to_replace","owner_role","status"}
form_counts={}
for letter in "ABCDEFGH":
    pat=rf"(?ms)^### 表 {letter}｜[^\n]*\n(.*?)(?=^### 表 [A-H]｜|^## |\Z)"
    m=re.search(pat,body)
    if not m: form_counts[letter]=0; continue
    block=m.group(1); form_counts[letter]=len([x for x in block.splitlines() if x.strip().startswith("|") and not re.match(r"\s*\|?\s*:?-{3,}",x)])
formula={}
for line in body.splitlines():
    m=re.match(r"\| (S-0[123]) \| ([^|]+) \| (\d+) \| (\d+) \| ([0-9.]+)% \|",line)
    if m:
        sid,vals,reported,maxp,reported_pct=m.groups(); nums=re.findall(r"\d+|NA",vals); score=sum(int(x) for x in nums if x!="NA"); denom=2*sum(x!="NA" for x in nums); pct=round(score/denom*100+1e-9,1)
        formula[sid]={"score":score,"reported_score":int(reported),"denominator":denom,"reported_denominator":int(maxp),"pct":pct,"reported_pct":float(reported_pct),"pass":score==int(reported) and denom==int(maxp) and pct==float(reported_pct)}
font_sizes=[]; table_sizes=[]
for p in doc.paragraphs:
    for r in p.runs:
        if r.text and r.font.size: font_sizes.append(r.font.size.pt)
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    if r.text and r.font.size: table_sizes.append(r.font.size.pt)
repeat_headers=0
for t in doc.tables:
    if t.rows and t.rows[0]._tr.trPr is not None and t.rows[0]._tr.trPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader") is not None: repeat_headers+=1
checks={"zip_integrity":bad is None,"identity":all(s in alltext for s in ["0086","I17C-E05","092"]),"sample_boundary":all(s in alltext for s in ["SAMPLE","SYNTHETIC","REAL=UNKNOWN","HOLD"]),"no_unrepaired_typo":"揭霱" not in alltext,"no_writer_markers":all(s not in alltext for s in ["<<<REGISTER_CSV>>>","<<<END_REGISTER_CSV>>>"]),"required_thresholds":all(s in alltext for s in ["≥90.0%","75.0–89.9%","<75.0%"]),"forms_A_to_H":all(v>=5 for v in form_counts.values()),"register_schema":bool(rows) and set(rows[0])==expected_headers,"register_rows_synthetic":all(r["status"]=="SYNTHETIC_REPLACE_BEFORE_REAL_USE" and r["doc_id"]=="0086" for r in rows),"register_fact_ids_unique":len({r["fact_id"] for r in rows})==len(rows),"sample_math":bool(formula) and all(v["pass"] for v in formula.values()),"repeat_header_all_tables":repeat_headers==len(doc.tables),"body_font_min_10_5":bool(font_sizes) and min(font_sizes)>=10.5,"table_font_min_9_5":bool(table_sizes) and min(table_sizes)>=9.5}
receipt={"project_id":"TTQS_ONE","document_id":"0086","requirement_id":"I17C-E05","official_checkbox_serial":"092","revision":"R01","docx_sha256":sha(docx),"docx_bytes":docx.stat().st_size,"pdf_sha256":None,"paragraphs":len(doc.paragraphs),"tables":len(doc.tables),"table_rows_total":sum(len(t.rows) for t in doc.tables),"form_table_pipe_rows_in_markdown":form_counts,"synthetic_replacement_rows":len(rows),"formula_rechecks":formula,"repeat_header_tables":repeat_headers,"paragraph_font_size_min_pt":min(font_sizes) if font_sizes else None,"table_font_size_min_pt":min(table_sizes) if table_sizes else None,"checks":checks,"technical_status":"PASS" if all(checks.values()) else "FAIL","verified_at_utc":datetime.now(timezone.utc).isoformat()}
p=root/"evidence/technical_qa_R01.json"; p.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+"\n",encoding="utf-8"); print(json.dumps(receipt,ensure_ascii=False))