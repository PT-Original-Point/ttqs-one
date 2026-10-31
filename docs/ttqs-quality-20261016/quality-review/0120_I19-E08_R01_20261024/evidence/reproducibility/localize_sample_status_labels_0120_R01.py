from pathlib import Path
from docx import Document
import hashlib
r=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0120_I19-E08_SOURCE_PREP_R01_20261024')
c=r/'candidate_R01'; doc=next(c.glob('*.docx')); md=next(c.glob('*.md'))
terms=[('NOT_YET_APPLIED','尚未運用'),('NOT_RECORDED','未登錄'),('NO_RESPONSE','未回覆'),('RESPONDED','已回覆'),('APPLIED','已運用')]
before=hashlib.sha256(doc.read_bytes()).hexdigest(); d=Document(doc); counts={k:0 for k,v in terms}
paras=list(d.paragraphs)
for t in d.tables:
 for row in t.rows:
  for cell in row.cells: paras.extend(cell.paragraphs)
for p in paras:
 for run in p.runs:
  for old,new in terms:
   n=run.text.count(old)
   if n: run.text=run.text.replace(old,new); counts[old]+=n
if any(v==0 for v in counts.values()): raise SystemExit(f'Expected all sample statuses in Word, found {counts}')
d.save(doc)
text=md.read_text(encoding='utf-8')
for old,new in terms: text=text.replace(old,new)
md.write_text(text,encoding='utf-8')
print({'before_docx_sha256':before,'after_docx_sha256':hashlib.sha256(doc.read_bytes()).hexdigest(),'replacements':counts,'manuscript_sha256':hashlib.sha256(md.read_bytes()).hexdigest()})
