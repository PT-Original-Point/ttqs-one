from pathlib import Path
import csv, json, hashlib
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

B=Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0120_I19-E08_SOURCE_PREP_R01_20261024")
S=B/"source"; O=B/"candidate_R01"
DOC=O/"0120__訓練成果初步證據彙整與評量表_R01.docx"
MD=O/"0120__訓練成果初步證據彙整與評量表_R01.md"
SAT=O/"SYNTHETIC_SATISFACTION_ITEMS_0120_R01.csv"
TEST=O/"SYNTHETIC_TEST_ITEM_SCORES_0120_R01.csv"
REC=O/"candidate_assembly_receipt_R01.json"
for x in (DOC,MD,SAT,TEST,REC):
    if x.exists(): raise SystemExit("REFUSE_OVERWRITE "+str(x))
with (S/"SYNTHETIC_DATASET_0120_R01.csv").open(encoding="utf-8-sig",newline="") as f: data=list(csv.DictReader(f))
with (S/"SYNTHETIC_REPLACEMENT_REGISTER_0120_R01.csv").open(encoding="utf-8-sig",newline="") as f: repl=list(csv.DictReader(f))
case="SYN-0120-2026-A01"
questions=[
("Q01","課程目標清楚，且能說明本課程要完成的服務紀錄與溝通任務。"),
("Q02","課程內容與社區志工可能遇到的服務情境相關。"),
("Q03","示範紀錄表有助於我分辨做了什麼與服務結果。"),
("Q04","課堂練習讓我有機會實際完成一筆去識別化服務紀錄。"),
("Q05","溝通演練協助我先確認對方需求，再用清楚語句回應。"),
("Q06","課程帶領者的說明能回答我對流程的疑問。"),
("Q07","教材、範例及練習指引容易找到且能看懂。"),
("Q08","課程節奏及練習時間足以讓我嘗試並修正作法。"),
("Q09","我知道遇到超出志工角色或涉及安全的情況時如何通報。"),
("Q10","我知道課後如何使用本課程所學，並知道仍需練習之處。")]
patterns={"SYN-L01":[5]*10,"SYN-L02":[5,5,4,4,4,4,4,3,3,4],"SYN-L03":[5,4,5,4,5,4,3,4,3,3],"SYN-L04":[5]*10,"SYN-L05":[4,4,3,3,4,3,2,3,2,2],"SYN-L06":[4,4,5,4,4,5,3,4,3,4],"SYN-L07":[5]*10}
with SAT.open("w",encoding="utf-8-sig",newline="") as f:
    w=csv.writer(f); w.writerow(["case_id","learner_id","record_class",*[q[0] for q in questions],"respondent_mean_1to5","lineage"])
    for lid,vals in patterns.items(): w.writerow([case,lid,"SYNTHETIC",*vals,f"{sum(vals)/10:.2f}","synthetic expansion; mean reconciles to source roll-up"])
    w.writerow([case,"SYN-L08","SYNTHETIC",*["NO_RESPONSE"]*10,"NOT_CALCULABLE","synthetic nonresponse; do not impute"])
with TEST.open("w",encoding="utf-8-sig",newline="") as f:
    w=csv.writer(f); w.writerow(["case_id","learner_id","time","record_class","T01_record_fields_20","T02_activity_outcome_20","T03_privacy_20","T04_communication_20","T05_escalation_20","total_100","lineage"])
    for r in data:
        for tm,key in [("PRE","pre_score_100"),("POST","post_score_100")]:
            if r[key]=="": continue
            v=int(r[key]); pieces=[v//5]*5
            w.writerow([case,r["learner_id"],tm,"SYNTHETIC",*pieces,v,"synthetic item expansion; components sum to source roll-up"])
paired=[r for r in data if r["pre_score_100"] and r["post_score_100"]]
pre=[int(r["pre_score_100"]) for r in paired]; post=[int(r["post_score_100"]) for r in paired]
sat=[int(r["satisfaction_1to5"]) for r in data if r["satisfaction_1to5"]]
fup=[r for r in data if r["followup_30d_status"]=="RESPONDED"]
maths={"cohort":8,"paired_n":len(paired),"pre_sum":sum(pre),"post_sum":sum(post),"pre_mean":sum(pre)/6,"post_mean":sum(post)/6,"mean_delta":sum(b-a for a,b in zip(pre,post))/6,"paired_pre_pass":sum(x>=70 for x in pre),"paired_post_pass":sum(x>=70 for x in post),"all_post_n":sum(bool(r["post_score_100"]) for r in data),"all_post_pass":sum(bool(r["post_score_100"]) and int(r["post_score_100"])>=70 for r in data),"satisfaction_n":len(sat),"satisfaction_sum":sum(sat),"satisfaction_mean":sum(sat)/len(sat),"followup_n":len(fup),"followup_applied":sum(r["followup_skill_use_self_report"]=="APPLIED" for r in fup)}
md="""# 訓練成果初步證據彙整與評量表
**doc_id：0120｜requirement_id：I19-E08｜候選版本：R01**  
**官方方框印刷流水號：UNKNOWN_NOT_ASSERTED**  
**狀態：CANDIDATE；示範案例及全部逐題擴充資料均為 SAMPLE／SYNTHETIC。協會 REAL 實績、核准及正式接受均未提供。**

## 用途與操作流程
本件整理單一課程 cohort 的學習目標、評量工具、時點、分母、初步結果、原件定位、資料缺口與補證動作。官方訓練機構版佐證表印刷第8頁／指標19列舉初步訓練成果及部分佐證作為常見參考，非評分保證。
1. 定義課程、班期、納入 cohort、完訓規則及工具版本；缺少來源填 UNKNOWN。
2. 建立目標—可觀察表現—工具—時點—門檻—原件映射。
3. 登錄應測名冊、測前/後、滿意度及追蹤邀請與回覆，分開保留分母。
4. 用匿名代碼匯入原始值，缺測不補零，列逐人缺漏。
5. 由另一角色重算；判讀須同列正向、持平、負向、限制及待補證。
6. 只有具權限流程可結案；簽名、核准、協會決定欄留白。

## REAL-FORM-01 班期與 cohort
|欄位|REAL 填寫|來源／定位|核對|
|---|---|---|---|
|課程名稱、代碼|________________|計畫：________|名稱代碼一致|
|班期起迄、地點|________________|公告/紀錄：________|日期可追溯|
|對象與入班條件|________________|招生規則：________|須為事前規則|
|入班 cohort N|____人|授權名冊：________|逐人代碼去重|
|完訓條件與人數|條件：______／____人|課程規則及出席原件|不得事後排除|
|目標及工具版本|________________|課綱/工具：________|同一構念|
|測前後日期|________________|原卷/紀錄：________|版本、時間可比|
|授權與保管角色|________________|授權：________|姓名名冊受限|
|承辦/覆核角色|________________|權責來源：________|本人填寫、不得預填|

## REAL-FORM-02 目標與評量設計
|目標代碼與目標|可觀察表現|工具/版本|時點、量尺|事前門檻及理由|證據索引/缺漏|
|---|---|---|---|---|---|
|________|________|________|________|________|________|
|________|________|________|________|________|________|
|________|________|________|________|________|________|
|________|________|________|________|________|________|
示例目標只為合成情境：完成去識別服務紀錄、分辨活動與成果、需求確認式溝通、識別超出志工角色的安全問題並通報。非協會已核定課綱。示例五題工具每題20分：紀錄欄位、活動/成果辨別、隱私、溝通情境、升級通報；0–20 rubric 分層為未作答/無關、部分表現、基本正確、大致完整、完整且可追溯。門檻70/100僅為 SAMPLE，非官方標準。

## REAL-FORM-03 班級完整度
|項目|計數|分母/公式|來源|核驗|
|---|---:|---|---|---|
|報名/入班|____/____|入班N=____|____|□待核 □通過 □退件|
|完訓|____|完訓/入班N|____|□待核 □通過 □退件|
|測前有效|____|有效測前/應測|____|□待核 □通過 □退件|
|測後有效|____|有效測後/應測|____|□待核 □通過 □退件|
|配對有效 n|____|同人、同目標、可比版本|____|□待核 □通過 □退件|
|缺一側測驗|____|缺一側/入班N|____|逐人列原因|
|滿意度有效回覆|____|有效回覆/受邀|____|缺答不作0分|
|追蹤有效回覆|____|有效回覆/應追蹤|____|未到期填 NOT_DUE|

## SAMPLE 案例及逐人評量
案例 SYN-0120-2026-A01；示例課程「社區志工服務紀錄與溝通基礎」，日期2026-08-01，30日追蹤2026-08-31。8人均為虛構代碼。完訓8/8只是案例設定：出席至少80%且完成練習；沒有真實名冊或原件。
|代碼|前測/100|後測/100|差值|配對|後測≥70|滿意度/5|追蹤自陳|缺漏|
|---|---:|---:|---:|---|---|---:|---|---|
"""
for r in data:
    a=r["pre_score_100"] or "缺測"; b=r["post_score_100"] or "缺測"
    delta=str(int(b)-int(a)) if a!="缺測" and b!="缺測" else "不可算"
    pairedflag="是" if delta!="不可算" else "否"
    passed=("達門檻" if int(b)>=70 else "未達") if b!="缺測" else "缺測"
    satv=r["satisfaction_1to5"] or "未回覆"
    reason="—" if pairedflag=="是" else ("缺測前測" if a=="缺測" else "缺測後測")
    md+=f'|{r["learner_id"]}|{a}|{b}|{delta}|{pairedflag}|{passed}|{satv}|{r["followup_skill_use_self_report"]}|{reason}|\n'
md+="""
## REAL-FORM-04 證據索引（REAL 空白）
|record_id|目標/資料類型|期間/數量|原件路徑、頁碼/列號|授權/保管角色|狀態/限制|
|---|---|---|---|---|---|
|____|____|____|____|____|□待取 □已核 □缺件|
|____|____|____|____|____|□待取 □已核 □缺件|
SAMPLE 索引：測驗逐題擴充 CSV、滿意度十題擴充 CSV、30日追蹤彙總均為本案合成檔，並非原件。不得稱為協會測驗卷、回收問卷或第三方驗證。

## REAL-FORM-05 公式及分母｜SAMPLE 重算
|指標|分子/分母|算式|結果|限制|
|---|---|---|---|---|
|配對前測平均|配對者測前總分/paired n|320/6|53.33/100|只計完整配對6人|
|配對後測平均|配對者測後總分/paired n|400/6|66.67/100|相同6人|
|配對平均差|Σ(post-pre)/paired n|80/6|+13.33分|描述差異，不作因果推論|
|上升/持平/下降|各類/6|5/6、1/6、0/6|83.3%、16.7%、0%|負向也須呈現|
|paired 前/後門檻率|達70分/paired n|1/6；3/6|16.7%；50.0%|門檻為示例|
|所有有效後測門檻率|達70分/有後測n=7|4/7|57.1%|與paired率不同|
|全cohort達門檻可確認占比|達70分/8|4/8|50.0%|缺測保留在分母|
|測前/後覆蓋率|7/8|7/8|87.5%|L07、L08各缺一側|
|滿意度回覆/均值|7/8；30/7|7/8；30/7|87.5%；4.29/5|不代表技能通過|
|追蹤回覆率|6/8|6/8|75.0%|兩人未回覆|
|自陳運用率（回覆者）|4/6|4/6|66.7%|僅自我回報|
|全cohort自陳已運用占比|4/8|4/8|50.0%|不是外部證據|
paired n=6 的測前總分320、測後總分400；提升總和80。配對者5人上升、1人持平、0人下降。測前達門檻1人、配對測後3人。所有後測者4/7達標；全cohort口徑4/8。2/8缺一側，不能補0或刪除。

## REAL-FORM-06 十題滿意度問卷
量尺：1非常不同意、2不同意、3普通、4同意、5非常同意；另記不適用/未答，缺答不是0。滿意度量學習經驗，技能仍由測驗/實作量尺評定。
"""
for q,t in questions:
    md+=f'- {q}：{t}\n'
md+="""
個人平均＝有效題分總和/該人有效題數；機構應事先定最低有效題數。SAMPLE 十題逐題答錄另見合成擴展CSV；7名回答者的個人均值為5、4、4、5、3、4、5，平均30/7=4.29，與基本資料集總評欄逐人相符。該逐題資料是為示範而另造的 SYNTHETIC 明細，不是原始答卷。
SAMPLE 反思題：最能使用的目標及去識別例子？紀錄曾漏什麼、如何發現？何種情況需通報？仍不確定什麼、需要何種練習？摘要須保留負向及不確定回饋。

## REAL-FORM-07 課後30日追蹤
|匿名代碼|應追蹤|邀請日/管道|回覆|自陳運用|去識別情境|外部索引/核驗|
|---|---|---|---|---|---|---|
|____|□是 □否|____|□有 □無|____|____|____|
|____|□是 □否|____|□有 □無|____|____|____|
SAMPLE：6/8回覆；4人自陳已運用、2人尚未運用、2人未回覆。只有自陳，沒有雇主、服務對象或其他第三方核驗。未到期填NOT_DUE；聯繫失敗只記嘗試。

## REAL-FORM-08 缺件及改善行動
|缺口/異常|影響|補件及責任角色|期限（REAL）|結案依據/重算|狀態|
|---|---|---|---|---|---|
|____|____|____|____|____|□待辦 □覆核 □結案|
SAMPLE示例：L07無後測、L08無前測。缺測原因在來源資料中未知，不臆造。查找合法原卷；不存在就維持缺測並依既定分母報告。本示例沒有已執行補件。

## REAL-FORM-09 SAMPLE→REAL 替換及最終覆核
原始15項 replacement register 完整保留於本文件。替換範圍含課程/日期/cohort、學習目標與門檻、出席完訓、測前後原卷、問卷逐題答錄、反思、追蹤、來源授權、保管定位、角色、分析、限制及結論。每列均須原件、授權、逐列重算及第二人覆核，狀態維持PENDING_REAL_SOURCE直到核實。
### HOLD／拒收
- doc_id、requirement_id、官方頁或唯一Blueprint錯配；把doc_id當方框號。
- 只有滿意度，沒有目標、工具、時點、門檻或成對樣本。
- 分母混淆；paired n冒充全班；缺測補零或刪除。
- 數學錯誤；量尺版本不同卻直接比較。
- SAMPLE改標REAL；捏造姓名、簽名、核准、課程實施或第三方證明。
- 自陳寫成外部驗證、描述差異寫成因果。
- 以I19-E07市場價值、I19-E04證照、I19-E03就業或I19-E05教材取代I19-E08；與I19-E09混件。
- 原件無授權/定位；個資公開；Word裁切或不可讀。
### 完成前勾核
□來源一致 □cohort/分母可重算 □缺測逐筆保留 □正負結果皆呈現 □來源定位及授權已核 □SAMPLE/REAL界線清楚 □跨文件無錯置 □獨立exact-SHA語意審查 □逐頁視覺檢查 □正式接受由有權限流程另定

## 跨文件邊界與狀態
0116/I19-E03 為就業狀況；0117/I19-E04 為證照/考試成果；0118/I19-E05 為教材與教學活動實施；0119/I19-E07 為市場/顧客價值提升；0121/I19-E09 為滿意度與學習評量。只有同課程、同cohort且逐筆可追溯時才能引用，不自動共用案例。
官方印刷頁8僅說明常見參考資料，未確認專屬方框流水號，故UNKNOWN_NOT_ASSERTED。協會現況、實際課程、真實學員、成效、核准、評核分數及接受狀態均UNKNOWN。
"""
MD.write_text(md,encoding="utf-8")
D=Document(); sec=D.sections[0]; sec.orientation=WD_ORIENT.LANDSCAPE; sec.page_width=Inches(11.69); sec.page_height=Inches(8.27); sec.top_margin=Inches(.58); sec.bottom_margin=Inches(.55); sec.left_margin=Inches(.62); sec.right_margin=Inches(.62); sec.header_distance=Inches(.25); sec.footer_distance=Inches(.25)
for nm in ("Normal","Body Text"):
    st=D.styles[nm]; st.font.name="Microsoft JhengHei"; st._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft JhengHei"); st.font.size=Pt(10.5); st.font.color.rgb=RGBColor(40,52,65)
for nm,size,col in [("Title",28,"17365D"),("Heading 1",18,"17365D"),("Heading 2",13,"1F4E78")]:
    st=D.styles[nm]; st.font.name="Microsoft JhengHei"; st._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft JhengHei"); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(col)
hpar=sec.header.paragraphs[0]; hpar.alignment=WD_ALIGN_PARAGRAPH.RIGHT; z=hpar.add_run("TTQS ONE｜0120 I19-E08｜候選施工 R01"); z.font.size=Pt(8); z.font.name="Microsoft JhengHei"
fpar=sec.footer.paragraphs[0]; fpar.alignment=WD_ALIGN_PARAGRAPH.CENTER; z=fpar.add_run("候選工作表　•　SAMPLE ≠ REAL　•　頁 "); z.font.size=Pt(8)
fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),"PAGE"); fpar._p.append(fld)
def shade(c,color):
    pr=c._tc.get_or_add_tcPr(); e=pr.find(qn("w:shd"))
    if e is None: e=OxmlElement("w:shd"); pr.append(e)
    e.set(qn("w:fill"),color)
def textcell(c,val,bold=False,color=None,size=9.5):
    c.text=""; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0); p.paragraph_format.line_spacing=1.0
    r=p.add_run(str(val)); r.bold=bold; r.font.name="Microsoft JhengHei"; r._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft JhengHei"); r.font.size=Pt(size)
    if color: r.font.color.rgb=RGBColor.from_string(color)
def table(headers,rows,widths=None,size=9.5):
    t=D.add_table(rows=1,cols=len(headers)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for j,v in enumerate(headers): textcell(t.rows[0].cells[j],v,True,"FFFFFF",size); shade(t.rows[0].cells[j],"1F4E78")
    repeat=OxmlElement("w:tblHeader"); repeat.set(qn("w:val"),"true"); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i,row in enumerate(rows):
        cells=t.add_row().cells; cant=OxmlElement("w:cantSplit"); t.rows[-1]._tr.get_or_add_trPr().append(cant)
        for j,v in enumerate(row):
            textcell(cells[j],v,size=size)
            if i%2: shade(cells[j],"F2F6FA")
    if widths:
        for row in t.rows:
            for j,w in enumerate(widths): row.cells[j].width=Inches(w)
    D.add_paragraph().paragraph_format.space_after=Pt(0)
def para(txt="",style=None):
    p=D.add_paragraph(style=style); p.paragraph_format.space_after=Pt(4); p.add_run(txt); return p
def heading(txt,lev=1): D.add_heading(txt,level=lev)
def callout(txt):
    t=D.add_table(rows=1,cols=1); t.style="Table Grid"; shade(t.cell(0,0),"FFF2CC"); textcell(t.cell(0,0),txt,size=10); D.add_paragraph()
# Cover
p= D.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(55)
r=p.add_run("訓練成果初步證據\n彙整與評量表"); r.bold=True; r.font.name="Microsoft JhengHei"; r._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft JhengHei"); r.font.size=Pt(27); r.font.color.rgb=RGBColor(23,54,93)
p= D.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("單一課程 Cohort｜工作表＋完整示範案例"); r.font.size=Pt(15); r.font.color.rgb=RGBColor(31,78,121)
table(["文件代號","對應需求","官方定位","候選版本"],[["0120","I19-E08","訓練機構版官方佐證表印刷第8頁／指標19","R01｜未獨立驗收"]],[1.4,1.4,4.2,2.0])
callout("狀態：CANDIDATE。示範課程、人物代碼、日期、問卷、分數、追蹤及原件定位均為 SAMPLE／SYNTHETIC；不是原始點協會實績、正式核准、評核分數或正式接受。官方方框印刷流水號：UNKNOWN_NOT_ASSERTED。")
para("用途：以一套可重算、可追溯的工作表整理單一班期的學習目標、評量工具、初步結果、原始證據定位、缺件與補證行動。滿意度只能反映學習經驗，不能替代技能評量。")
para("角色使用：訓練行政承辦、課程負責及覆核角色；本表不假定協會目前的職務任命、權限或核准流程。所有簽名／核准欄由具權限者留待實際作業填寫。")
D.add_page_break()
heading("一、文件定位、閱讀方式與使用步驟")
para("官方訓練機構版佐證表第8頁指標19列出初步訓練成果及部分佐證，並舉學員滿意度、測驗／心得、課後追蹤及證照情形為參考。這些是例示，不代表任何單位得分、通過或完成評核。")
table(["步驟","角色","輸入","完成／停留條件"],[
["1 班期定義","課程／行政承辦","課程、cohort、納入條件、完訓規則","名冊分母可核；規則有事前依據；不清楚則HOLD"],
["2 工具錨定","課程負責","目標、工具、量尺、時點、門檻","每個目標均有工具；滿意度不得代替技能"],
["3 證據索引","資料保管／行政角色","原件代碼、路徑、頁列、授權","原件可定位且有權限；缺件登錄、不補造"],
["4 計算","承辦＋覆核","匿名原始值、缺測碼、固定分母","獨立重算一致；不同分母分列"],
["5 判讀與補證","覆核角色","正向、持平、負向、限制、待辦","不得因結果不好刪樣本或因果化"],
["6 結案","具權限決策角色（UNKNOWN）","覆核與正式程序","本候選不代替批准或接受"]],[1.25,1.8,3.2,3.6])
callout("分工為可替換角色，不代表原始點協會已任命任何人。姓名及授權資訊按實際權限保管，不在公開彙整稿填入。")
heading("二、REAL-FORM-01｜課程班期與 cohort 定義")
table(["欄位","REAL 填寫欄","來源／定位","核對規則"],[
["課程名稱／代碼","________________________","課程計畫：____________","名稱、代碼一致"],
["班期起迄／地點","________________________","公告／紀錄：____________","日期可追溯"],
["對象／入班條件","________________________","招生／入班規則：________","採事前規則"],
["入班 cohort N","____ 人","授權名冊代碼：__________","逐人去重"],
["完訓條件／人數","條件：________／____人","課程規則、出席／練習紀錄","不得事後排除"],
["目標／量尺版本","________________________","核准課程目標：__________","測前後工具可比"],
["測前／後日期","________________________","原卷／施測紀錄：________","版本和時點一致"],
["授權／保管角色","________________________","授權或告知：____________","限制在核准目的"],
["承辦／覆核角色","________________________","權責來源：______________","角色未知則填UNKNOWN"]],[1.6,2.5,2.5,3.0])
heading("三、REAL-FORM-02｜目標與評量設計")
para("在分析前確定工具版本、時點和通過門檻；不可先看結果再挑門檻。不同版本不得直接合併，除非有可比性證據。")
table(["目標代碼／目標","可觀察表現","工具／版本","時點／量尺","門檻及事前理由","原件索引／缺漏"],[
["________________","________________","________________","________","________","________________"],
["________________","________________","________________","________","________","________________"],
["________________","________________","________________","________","________","________________"],
["________________","________________","________________","________","________","________________"]],[1.65,1.7,1.65,1.25,1.7,1.8])
callout("SAMPLE目標：完成去識別服務紀錄、分辨活動與成果、需求確認式溝通、辨識超出志工角色的安全問題並通報。這些是合成情境，非協會已核定課綱。")
para("SAMPLE評量工具：5個各20分的項目，涵蓋必要紀錄欄位、活動與成果辨別、隱私、溝通情境、升級通報；單題 0–20 分規準為無有效表現、部分、基本、大致完整、完整可追溯。70/100為示例門檻，非官方標準。")
heading("四、REAL-FORM-03｜名冊、完訓與評量完整度")
table(["計數項目","REAL計數","分母／公式","來源","覆核"],[
["報名／入班","____／____","入班 N=____","________","□待核 □通過 □退件"],
["完訓","____","完訓／入班 N","________","□待核 □通過 □退件"],
["測前有效","____","有效測前／應測","________","□待核 □通過 □退件"],
["測後有效","____","有效測後／應測","________","□待核 □通過 □退件"],
["成對有效 n","____","同人、同目標、可比版本","________","□待核 □通過 □退件"],
["缺一側測驗","____","缺一側／入班 N","________","逐筆列原因"],
["滿意度回覆","____","有效回覆／受邀人數","________","未答不作0分"],
["追蹤回覆","____","有效回覆／應追蹤人數","________","未到期NOT_DUE"]],[1.5,1.15,2.65,2.0,2.7])
heading("五、SAMPLE 案例與 cohort 核對")
callout(f"案例 {case}：示例課程「社區志工服務紀錄與溝通基礎」，日期2026-08-01，30日追蹤2026-08-31。八人均為虛構匿名代碼。完訓8/8是假設示範設定（出席至少80%且完成練習），無實際名冊或佐證原件。")
table(["核對項目","SAMPLE數值","公式／來源狀態","解讀"],[
["入班／完訓","8／8","示例完訓條件；非REAL紀錄","不得稱實際辦訓"],
["測前／後有效","7／8、7／8","L08缺前測；L07缺後測","缺測不補0"],
["成對有效","6／8","同一人有前後測","paired n不是全班"],
["缺一側","2／8=25.0%","L07、L08","原因未知，不臆造"],
["滿意度回覆","7／8=87.5%","7份有效答錄／8人","未回覆不入平均"],
["追蹤回覆","6／8=75.0%","6份／8人","自陳與外證分離"]],[1.7,1.5,3.0,3.4])
heading("六、REAL-FORM-04｜原始證據索引與授權核驗")
para("真實原件保存在有權限位置；本表只記索引及最少必要內容。姓名名冊和可識別原始答卷不得公開。")
table(["record_id","目標／類型","期間／數量","REAL原件路徑／頁列","授權／保管角色","狀態／限制"],[
["________","________","________","________________","________","□待取 □已核 □缺件"],
["________","________","________","________________","________","□待取 □已核 □缺件"],
["________","________","________","________________","________","□待取 □已核 □缺件"],
["________","________","________","________________","________","□待取 □已核 □缺件"]],[1.0,1.55,1.25,2.35,1.65,1.85])
heading("SAMPLE索引（不是實際原件）",2)
table(["索引","資料類型","合成內容","來源狀態","禁止宣稱"],[
["SYN-EV01","測驗逐題擴展","14筆有效測前／後明細","合成CSV、無原卷","非協會試卷"],
["SYN-EV02","滿意度問卷","7份逐題答錄；1人未回","合成CSV、無回收問卷","非真實滿意度"],
["SYN-EV03","30日追蹤","6人自陳，2人未回","合成資料、無外證","非第三方核驗"],
["SYN-EV04","完訓出席","8/8示例假設","無簽到原件","非實際完訓證明"]],[1.05,1.5,2.1,2.5,2.5])
heading("七、REAL-FORM-05｜逐人匿名評量與缺測")
para("REAL僅使用授權匿名代碼；名冊對照表受限管理。配對只納入同一人、同目標且可比量尺版本。")
table(["代碼","前測","後測","差值","配對","後測≥70","滿意度","追蹤自陳","狀態"],[
["________","____","____","____","□Y □N","____","____","____","________"],
["________","____","____","____","□Y □N","____","____","____","________"],
["________","____","____","____","□Y □N","____","____","____","________"],
["________","____","____","____","□Y □N","____","____","____","________"]],[1.0,.75,.75,.75,.8,1.0,1.0,1.25,1.75])
heading("SAMPLE逐人資料",2)
srows=[]
for r in data:
    a=r["pre_score_100"] or "缺測"; b=r["post_score_100"] or "缺測"
    d=str(int(b)-int(a)) if a!="缺測" and b!="缺測" else "不可算"
    passv=("達門檻" if int(b)>=70 else "未達") if b!="缺測" else "缺測"
    srows.append([r["learner_id"],a,b,d,"是" if d!="不可算" else "否",passv,r["satisfaction_1to5"] or "未回覆",r["followup_skill_use_self_report"],"—" if d!="不可算" else ("缺測前測" if a=="缺測" else "缺測後測")])
table(["代碼","前測","後測","差","配對","後測門檻","滿意度","追蹤","缺漏狀態"],srows,[.95,.7,.7,.65,.65,.95,.85,1.05,1.55])
para("缺測值不補零；沒有配對就不計差值。門檻結果只對有效後測者分類。空白、重複或來源未核驗的值不能自動判PASS。")
heading("八、REAL-FORM-06｜公式及 SAMPLE 重算")
table(["指標","公式／分母","算式","結果","限制"],[
["配對前測平均","paired前測總分／n","320／6","53.33／100","同6人"],
["配對後測平均","paired後測總分／n","400／6","66.67／100","同6人"],
["配對平均差","Σ(post−pre)／n","80／6","+13.33分","不作因果推論"],
["上升／持平／下降","各類／6","5／6；1／6；0／6","83.3%；16.7%；0%","不刪負向"],
["配對前／後門檻率","達70／paired n","1／6；3／6","16.7%；50.0%","示例門檻"],
["有效後測門檻率","達70／post n=7","4／7","57.1%","非配對分母"],
["全班可確認達門檻占比","達70／cohort N=8","4／8","50.0%","缺測仍在分母"],
["前／後測覆蓋","有效數／8","7／8；7／8","87.5%；87.5%","L07、L08缺一側"],
["滿意度回覆／平均","7／8；30／7","7／8；30／7","87.5%；4.29／5","不是技能通過"],
["追蹤回覆率","6／8","6／8","75.0%","自陳不作外證"],
["已運用自陳率","4／6回覆者","4／6","66.7%","回覆者分母"],
["全班自陳已運用","4／8 cohort","4／8","50.0%","全班分母"]],[1.45,2.05,1.35,1.35,2.35])
callout("分母為0記NOT_CALCULABLE；不同量尺版本不可直接合併。分數範圍錯誤、代碼重複、時點反轉、原件無法重算即HOLD。")
heading("九、REAL-FORM-07｜十題滿意度問卷及逐題示範")
para("量尺：1非常不同意、2不同意、3普通、4同意、5非常同意；另記不適用／未答，缺答不當0分。滿意度反映學習經驗，不代替技能測驗。")
table(["題號","題目（示範文字；REAL使用前須審定）","圈選","缺答"],[[code,text,"1 2 3 4 5","□不適用／未答"] for code,text in questions],[.7,5.8,1.2,1.3])
heading("逐題完整SAMPLE答錄",2)
satrows=[]
for lid,v in patterns.items(): satrows.append([lid,*v,f"{sum(v)/10:.2f}"])
satrows.append(["SYN-L08",*["未答"]*10,"未計"])
table(["代碼",*[x[0] for x in questions],"均值"],satrows,[1.05]+[.62]*10+[.85],size=9.5)
para("7人的十題平均為5、4、4、5、3、4、5，合計30/7=4.29/5，逐人對齊基礎資料集的整體評分。題目逐答是另造的合成擴展，並非原始問卷。REAL版須先定個人最低有效題數；不足者不計個人平均並留缺漏。")
heading("十、REAL-FORM-08｜反思及判讀")
table(["反思題","REAL作答／摘要","覆核提示"],[
["哪項目標可用於情境？舉去識別例子。","____________________________","不可填服務對象個資"],
["練習紀錄漏了什麼？如何發現、修正？","____________________________","保留錯誤和修正"],
["何時須通報？會採取什麼步驟？","____________________________","核對角色邊界"],
["仍不確定什麼？需要哪種練習？","____________________________","自信不等於能力通過"]],[3.0,3.0,2.8])
para("SAMPLE判讀：6位配對者平均後測高於前測13.33分，5升、1平；2人缺一側。配對後測門檻率50.0%，有效後測者4/7達示例門檻。僅可示範整理初步資料，不能推論課程造成提升或代表協會實績。")
heading("十一、REAL-FORM-09｜30日追蹤")
table(["匿名代碼","邀請日期／管道","回覆","自陳是否運用","去識別情境","外部索引／核驗"],[
["________","____________","□有 □無","____________","____________","____________"],
["________","____________","□有 □無","____________","____________","____________"],
["________","____________","□有 □無","____________","____________","____________"],
["________","____________","□有 □無","____________","____________","____________"]],[1.15,1.75,1.05,1.5,2.2,2.15])
frows=[]
for r in data:frows.append([r["learner_id"],"2026-08-31 SAMPLE",r["followup_30d_status"],r["followup_skill_use_self_report"],"僅自陳／無外證" if r["followup_30d_status"]=="RESPONDED" else "未回覆"])
heading("SAMPLE追蹤表",2)
table(["代碼","追蹤日","回覆","狀態","證據限制"],frows,[1.0,1.7,1.35,1.45,3.2])
para("6/8回覆；4人自陳已運用、2人尚未、2人未回覆。沒有第三方核驗。未到期填NOT_DUE；聯繫失敗僅登錄嘗試，不推定結果。")
heading("十二、REAL-FORM-10｜缺件與改善行動")
table(["缺口／異常","影響","補件與負責角色","期限REAL填","結案證據／重算","狀態"],[
["________________","________________","________________","______","________________","□待辦 □覆核 □結案"],
["________________","________________","________________","______","________________","□待辦 □覆核 □結案"],
["________________","________________","________________","______","________________","□待辦 □覆核 □結案"]],[1.45,1.6,1.8,.9,2.25,1.0])
para("SAMPLE：L07缺後測、L08缺前測。缺漏原因未知，不臆造；查找合法原卷，若不存在則保留缺測並用既定分母報告。這不是已完成的補件工作。")
heading("十三、REAL-FORM-11｜SAMPLE→REAL替換清冊")
para("原始替換清冊15列逐項嵌入；另列本候選為問卷和測驗新增的合成逐題展開。狀態在取得、授權並覆核原件前均為PENDING_REAL_SOURCE。")
table(["ID／欄位位置","當前SAMPLE內容","所需REAL來源","核驗規則","狀態"],[[r["replacement_id"]+"｜"+r["field_or_location"],r["current_example"],r["required_real_source"],r["verification_rule"],r["status"]] for r in repl],[1.95,2.0,2.75,2.35,1.0],size=9.2)
table(["新增合成擴展","所需REAL原件","核驗條件","狀態"],[
["十題逐題滿意度答錄","授權回收的匿名問卷","逐題原卷、有效題數規則、個人與整體均值一致","PENDING_REAL_SOURCE"],
["五題測驗逐題分數","同版本測前後原卷、答案鍵與評分規準","逐題加總等於總分；版本及時點可比","PENDING_REAL_SOURCE"]],[1.8,2.6,3.4,1.3])
heading("十四、REAL-FORM-12｜覆核與退件條件")
table(["覆核項","PASS條件","退件／HOLD條件","結果／定位"],[
["身份與需求","0120/I19-E08/標題/Blueprint一致","鄰件混淆或方框號冒充","□PASS □FAIL □HOLD"],
["來源授權","REAL原件、位置、授權齊全","只收到摘要或權限未知","□PASS □FAIL □HOLD"],
["cohort／分母","名冊及各工具分母清楚","paired n冒充全班、缺測補零","□PASS □FAIL □HOLD"],
["目標與工具","目標—工具—時點—版本—門檻對應","只用滿意度代表技能","□PASS □FAIL □HOLD"],
["計算","分子分母可重算且雙人一致","總數錯、樣本重複、分母混用","□PASS □FAIL □HOLD"],
["判讀","呈現正負、缺漏及非因果限制","自陳冒充外證或宣稱因果","□PASS □FAIL □HOLD"],
["隱私／狀態","去識別、SAMPLE/REAL分明","偽造核准、簽章或第三方資料","□PASS □FAIL □HOLD"],
["跨文件","僅同cohort可追溯才互引","E03/E04/E05/E07/E09混件","□PASS □FAIL □HOLD"],
["版面／審查","逐頁可讀，獨立review綁exact SHA","裁切、破表或無獨立驗收","□PASS □FAIL □HOLD"]],[1.45,2.45,2.7,1.8])
table(["承辦角色／日期","覆核角色／日期","正式決策角色／日期","正式接受狀態"],[["________／______","________／______","________／______","留白；由有權限程序決定"]],[2.15,2.15,2.15,2.0])
heading("十五、鄰件邊界、來源及候選限制")
para("0116/I19-E03為學員訓後就業；0117/I19-E04為考照通過率；0118/I19-E05為教材與教學活動實施；0119/I19-E07為市場及顧客價值；0121/I19-E09為滿意度與學習評量。只有同課程、同cohort且逐筆可追溯時才可互引。")
para("官方來源為勞動力發展署《TTQS查核佐證文件資料表－訓練機構版》，SHA-256 805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709，第8頁。crosswalk將I19-E08指向0120單一交付件。官方方框印刷流水號未能確認，故UNKNOWN_NOT_ASSERTED。")
para("候選限制：問卷十題逐題答錄、測驗五項逐題分數是另行新增的合成擴展，沒有任何真實答卷或原件佐證。協會現況、實際課程、學員、結果、核准、分數與正式接受均UNKNOWN。本件仍需獨立 exact-SHA 語意審查及逐頁視覺驗收。")

D.core_properties.title="訓練成果初步證據彙整與評量表"
D.core_properties.subject="0120 / I19-E08; CANDIDATE R01; SAMPLE-SYNTHETIC"
D.core_properties.author="TTQS ONE candidate production"
D.save(DOC)
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
receipt={"doc_id":"0120","requirement_id":"I19-E08","official_checkbox_serial":"UNKNOWN_NOT_ASSERTED","revision":"R01","candidate_state":"CANDIDATE","formal_acceptance":"NOT_GRANTED","docx_path":str(DOC),"docx_sha256":sha(DOC),"manuscript_path":str(MD),"manuscript_sha256":sha(MD),"synthetic_satisfaction_csv":{"path":str(SAT),"sha256":sha(SAT),"data_rows":8},"synthetic_test_item_csv":{"path":str(TEST),"sha256":sha(TEST),"data_rows":14},"source_sha256":{"official_pdf":"805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709","crosswalk":"1ddfbb608bcad25170683ddb42a1dbb27bafa6f1f35ff0498f1b8d75b24f0679","unique_blueprint":"09a3aec874c14f97d5d6a1bed85d3cd75d617a9cf4bd33871973e5cb27c088ad","base_synthetic_dataset":"ef04ed006079845be03ecbc51e5f2af72b34d565c69ee384cfda9677db2ab737","replacement_register":"97f19c56c51d835289bee94479cdad12165b0be388d203eafe082753dd5ff675"},"calculations":maths,"actual_worker_model_identity":"UNKNOWN","semantic_independent_review":"PENDING","visual_review":"PENDING","technical_review":"PENDING","limits":"all case details and item expansions are synthetic; no real association performance or approval asserted"}
REC.write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding="utf-8")
print(json.dumps({"docx":str(DOC),"docx_sha256":sha(DOC),"markdown_sha256":sha(MD),"satisfaction_csv_sha256":sha(SAT),"test_item_csv_sha256":sha(TEST),"receipt":str(REC),"calculated":maths},ensure_ascii=False))

