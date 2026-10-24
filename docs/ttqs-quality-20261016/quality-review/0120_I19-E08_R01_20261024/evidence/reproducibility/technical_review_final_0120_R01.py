from pathlib import Path
from docx import Document
import hashlib,json,zipfile,csv,statistics
ROOT=Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0120_I19-E08_SOURCE_PREP_R01_20261024")
C=ROOT/'candidate_R01';S=ROOT/'source';Q=ROOT/'qa_R01';DOC=next(C.glob('*.docx'));d=Document(DOC)
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
body='\n'.join([p.text for p in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells])
mp=json.loads((S/'mapping_evidence_0120_R01.json').read_text(encoding='utf-8'))
bp=(S/'UNIQUE_BLUEPRINT_0120_I19-E08_R01.md').read_text(encoding='utf-8'); packet=(S/'SOURCE_PACKET_0120_I19-E08_R01.md').read_text(encoding='utf-8')
reg=list(csv.DictReader((S/'SYNTHETIC_REPLACEMENT_REGISTER_0120_R01.csv').open(encoding='utf-8-sig',newline='')))
base=list(csv.DictReader((S/'SYNTHETIC_DATASET_0120_R01.csv').open(encoding='utf-8-sig',newline='')))
tests=list(csv.DictReader((C/'SYNTHETIC_TEST_ITEM_SCORES_0120_R01.csv').open(encoding='utf-8-sig',newline='')))
sat=list(csv.DictReader((C/'SYNTHETIC_SATISFACTION_ITEMS_0120_R01.csv').open(encoding='utf-8-sig',newline='')))
refs=list(csv.DictReader((C/'SYNTHETIC_REFLECTION_EXCERPTS_0120_R01.csv').open(encoding='utf-8-sig',newline='')))
checks={}
checks['doc_id_requirement_crosswalk']=mp['doc_id']=='0120' and mp['requirement_id']=='I19-E08' and mp['crosswalk_row']['deliverable_ids']=='0120'
checks['official_page8_indicator19']=mp['official_printed_page']==8 and mp['official_pdf_page_index_zero_based']==7 and 'Indicator 19' in mp['official_checkbox_locator']
checks['blueprint_bound_to_doc_and_requirement']='0120' in bp and 'I19-E08' in bp and '訓練成果初步證據彙整與評量表' in bp
checks['source_packet_has_official_scope']='805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709' in packet and '第 8 頁' in packet and '初步訓練成果' in packet
checks['document_genre_and_title']='訓練成果初步證據' in body and '彙整與評量表' in body and '單一課程 Cohort' in body and '完整示範案例' in body
checks['all_12_real_forms']=all(f'REAL-FORM-{i:02d}' in body for i in range(1,13))
checks['15_replacement_items']=len(reg)==15 and all(f'R-{i:02d}' in body for i in range(1,16))
checks['all_8_sample_learners']=all(f'SYN-L{i:02d}' in body for i in range(1,9))
checks['10_satisfaction_items']=all(f'Q{i:02d}' in body for i in range(1,11))
checks['5_test_tasks_and_full_anchor_scale']=all(f'T0{i}' in body for i in range(1,6)) and all(x in body for x in ['0分空白','1–5分','6–10分','11–15分','16–20分'])
checks['14_test_expansion_rows_synthetic']=len(tests)==14 and all(r['record_class']=='SYNTHETIC' for r in tests)
cols=['T01_record_fields_20','T02_activity_outcome_20','T03_privacy_20','T04_communication_20','T05_escalation_20']
checks['item_scores_sum_to_total']=all(sum(int(r[k]) for k in cols)==int(r['total_100']) for r in tests)
lookup={r['learner_id']:r for r in base}
checks['test_rollups_reconcile_to_base']=all(r['total_100']==lookup[r['learner_id']]['pre_score_100' if r['time']=='PRE' else 'post_score_100'] for r in tests)
paired=[r for r in base if r['pre_score_100'] and r['post_score_100']];pre=sum(int(r['pre_score_100']) for r in paired);post=sum(int(r['post_score_100']) for r in paired);delta=[int(r['post_score_100'])-int(r['pre_score_100']) for r in paired]
checks['paired_math_n6_320_400_delta80']=len(paired)==6 and pre==320 and post==400 and sum(delta)==80 and round(statistics.mean(delta),2)==13.33
checks['paired_directions_5up_1tie_0down']=sum(x>0 for x in delta)==5 and sum(x==0 for x in delta)==1 and sum(x<0 for x in delta)==0
checks['satisfaction_rows_and_means_reconcile']=len(sat)==8 and all(r['record_class']=='SYNTHETIC' for r in sat) and all((r['respondent_mean_1to5']=='NOT_CALCULABLE' if r['learner_id']=='SYN-L08' else abs(float(r['respondent_mean_1to5'])-float(lookup[r['learner_id']]['satisfaction_1to5']))<0.001) for r in sat)
checks['satisfaction_sum_30_over_7_mean_4_29']=sum(float(r['satisfaction_1to5']) for r in base if r['satisfaction_1to5'])==30 and round(30/7,2)==4.29
checks['reflection_3_rows_synthetic']=len(refs)==3 and all(r['record_class']=='SYNTHETIC' and 'not a real learner' in r['lineage'] for r in refs)
checks['reflection_table_not_duplicated']=sum('SAMPLE反思摘錄（合成）' in (' | '.join(c.text for c in t.rows[0].cells) if t.rows else '') for t in d.tables)==1
checks['form08_page_break']=any('REAL-FORM-08' in p.text and p.paragraph_format.page_break_before for p in d.paragraphs)
checks['machine_tokens_removed_from_reader_text']=not any(x in body for x in ['PENDING_REAL_SOURCE','UNKNOWN_NOT_ASSERTED','TODO','TBD'])
checks['sample_real_and_hold_boundary']=all(x in body for x in ['SAMPLE','SYNTHETIC','REAL','HOLD','不作因果推論'])
checks['reader_friendly_sample_status_labels']=all(x in body for x in ['已運用','尚未運用','未登錄','未回覆','已回覆']) and not any(x in body for x in ['NOT_YET_APPLIED','NOT_RECORDED','NO_RESPONSE','RESPONDED','APPLIED'])
checks['no_claim_of_association_actual_result']='不能推論課程造成提升或代表協會實績' in body and '不是原始點協會實績' in body
checks['neighbor_crosswalk_labels_correct']=all(x in body for x in ['0116/I19-E03','0117/I19-E04','0118/I19-E05','0119/I19-E07','0121/I19-E09'])
checks['no_wrong_neighbor_genre_labels']=all(x not in body for x in ['學員滿意度調查表與分析報告','訓練教材與教學活動實施紀錄（錯置為E04）'])
with zipfile.ZipFile(DOC) as z:zip_bad=z.testzip()
checks['docx_zip_integrity']=zip_bad is None
pts=[run.font.size.pt for t in d.tables for row in t.rows for c in row.cells for p in c.paragraphs for run in p.runs if run.font.size]
checks['minimum_explicit_table_font_9_5pt']=bool(pts) and min(pts)>=9.5
checks['normal_body_font_10_5pt']=bool(d.styles['Normal'].font.size) and d.styles['Normal'].font.size.pt>=10.5
checks['all_sections_landscape']=all(s.page_width>s.page_height for s in d.sections)
checks['no_mermaid_source']='```mermaid' not in body and 'graph TD' not in body
pdf=Q/'0120_R01_word_render_final5.pdf'; images=list((ROOT/'render_R01_final5').glob('page-*.png'))
checks['fresh_native_word_render_12_pages']=pdf.exists() and len(images)==12 and all(p.stat().st_size>10000 for p in images)
metrics={'docx_sha256':sha(DOC),'docx_bytes':DOC.stat().st_size,'paragraph_count':len(d.paragraphs),'table_count':len(d.tables),'table_rows':sum(len(t.rows) for t in d.tables),'character_count':len(body),'min_explicit_table_font_pt':min(pts) if pts else None,'paired_n':len(paired),'pre_sum':pre,'post_sum':post,'paired_delta_sum':sum(delta),'paired_delta_mean':round(statistics.mean(delta),2),'satisfaction_n':7,'satisfaction_sum':30,'satisfaction_mean':round(30/7,2),'replacement_rows':len(reg),'test_expansion_rows':len(tests),'reflection_examples':len(refs),'word_pdf_sha256':sha(pdf) if pdf.exists() else None,'rendered_page_count':len(images),'zip_error':zip_bad}
source_hashes={p.name:sha(p) for p in [S/'mapping_evidence_0120_R01.json',S/'SOURCE_PACKET_0120_I19-E08_R01.md',S/'UNIQUE_BLUEPRINT_0120_I19-E08_R01.md',S/'SYNTHETIC_DATASET_0120_R01.csv',S/'SYNTHETIC_REPLACEMENT_REGISTER_0120_R01.csv']}
out={'receipt_id':'0120_R01_TECHNICAL_REVIEW_FINAL','scope':'Mechanical structure, mapping, source, data lineage, arithmetic and Word-render existence only; not independent semantic acceptance','artifact_path':str(DOC),'artifact_sha256':sha(DOC),'source_hashes':source_hashes,'metrics':metrics,'checks':checks,'all_checks_pass':all(checks.values()),'independent_semantic_review':'PENDING','formal_acceptance':'NOT_GRANTED'}
(Q/'0120_R01_technical_review_final.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'artifact_sha256':out['artifact_sha256'],'all_checks_pass':out['all_checks_pass'],'failed':[k for k,v in checks.items() if not v],'metrics':metrics,'receipt':str(Q/'0120_R01_technical_review_final.json')},ensure_ascii=False))




