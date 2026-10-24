from pathlib import Path
import hashlib,json,re,zipfile
from datetime import datetime
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0130_I19-E19_SOURCE_PREP_R01_20261024")
RAW=BASE/"writer"/"antigravity_R01"/"writer_response_R01.md"
OUT=BASE/"candidate"; OUT.mkdir(exist_ok=True)
ADOPTED=OUT/"0130_I19-E19_adopted_R01B.md"
DOCX=OUT/"0130_I19-E19_CANDIDATE_R01B.docx"
RECEIPT=BASE/"evidence"/"ADOPTION_BUILD_QA_0130_R01B.json"; RECEIPT.parent.mkdir(exist_ok=True)
BT=chr(96); FENCE=BT*3
raw=RAW.read_text(encoding="utf-8"); raw_sha=hashlib.sha256(raw.encode("utf-8")).hexdigest()
lines=raw.splitlines(); log=[]
def replace_line(needle,new,expected=1,starts=None):
    global lines
    ix=[i for i,x in enumerate(lines) if ((x.startswith(needle) if needle.startswith("| ") else needle in x) and (starts is None or x.lstrip().startswith(starts)))]
    if len(ix)!=expected: raise RuntimeError(f"Expected {expected} line(s) matching {needle!r}, got {len(ix)}")
    for i in ix:
        old=lines[i]; lines[i]=new; log.append({"match":needle,"old":old,"new":new})
def replace_text(old,new,expected=1):
    global lines
    count=sum(x.count(old) for x in lines)
    if count!=expected: raise RuntimeError(f"Expected {expected} occurrence(s) of {old!r}, got {count}")
    lines=[x.replace(old,new) for x in lines]; log.append({"old":old,"new":new,"count":count})

lines[2:2]=[
"> **採用狀態：CANDIDATE。** 這是待機構檢視採用的候選流程，並非原始點協會已核定制度。案例、日期、角色、名冊、聯絡、來源及數字均為 SAMPLE / SYNTHETIC；本範例沒有真實查詢、聯絡、會議或簽核。",""]
log.append({"match":"title block","new":"candidate/SAMPLE boundary inserted"})
replace_line("3. **遵循個人資料保護法規**","3. **個資與查證權限**：依機構正式核定政策與適用規範辦理；本候選表不構成法律或個資合規審查。未確認用途、權限及必要範圍前，停止外查及對外揭露。")
replace_line("| "+BT+"VERIFIED_FOR_REVIEW"+BT+" |","| "+BT+"VERIFIED_FOR_REVIEW"+BT+" | 具足佐證待複核 | 須具可追溯的受訓身分、經確認的資料使用／查證權限、能支持具體主張的獨立來源、可核對的訓後時間、重複檢查及限制說明；涉及訓練連結時須列明可核查依據，不能由時間先後推定因果。 | 僅由機構正式核定權責依採用程序複核；本範例沒有真實審查或核准 |",starts="| ")
replace_line("| "+BT+"HOLD_CONSENT"+BT+" |","| "+BT+"HOLD_CONSENT"+BT+" | 保留：授權或查證權限未確認 | 尚無適當授權、用途範圍不清，或機構尚未確認可用查證權限。 | 暫停外查及揭露；依正式政策補正或結案 |",starts="| ")
replace_line("| "+BT+"HOLD_LINKAGE"+BT+" |","| "+BT+"HOLD_LINKAGE"+BT+" | 保留：身分、時序或訓練連結未釐清 | 無法連結受訓學員／班別、日期矛盾，或技能關聯只有未佐證自述。 | 記錄缺口與待補依據；不得推定訓練造成創業 |",starts="| ")
replace_line("| "+BT+"NOT_VERIFIED"+BT+" |","| "+BT+"NOT_VERIFIED"+BT+" | 截止時未能驗證 | 依機構正式採用程序處理後仍缺乏足夠來源。只表示未能驗證，不等於已證明主張為假。 | 記明缺口與結案依據；新證據到達時依正式程序重啟 |",starts="| ")
replace_line("| "+BT+"EXCLUDED_DUPLICATE"+BT+" |","| "+BT+"EXCLUDED_DUPLICATE"+BT+" | 排除：機構授權資料內重複 | 僅於有權使用的本機構資料中，以必要欄位確認重複。 | 留下內部可稽核依據；不得假設可跨機構查詢 |",starts="| ")
replace_line("| **授權有效年限**","| **授權有效期間** | 依正式審查採用的文件填寫；未確認前不預設期限 | UNKNOWN（本範例沒有實際授權） |")
replace_line("| **撤回與銷毀機制**","| **資料撤回與銷毀** | 依機構正式核定政策及適用規範說明；本表不預設權利內容或保存期限 | 尚未審定；不得視為現行規則 |")
replace_line("| 基礎推拿調理手法與安全防護","| SAMPLE課程技能：流程設計與物品分類 | SAMPLE居家收納整理服務之自述工作內容 | 虛構課程與虛構服務只有概念對照；沒有真實營運或第三方資料支持，不能認定技能已運用或產生創業成果。 | [ ] 已驗證　[x] 待查　[ ] 無關 |")
replace_line("| 衛生法規與營業場所消毒作業","| SAMPLE課程技能：需求確認與分區規劃 | SAMPLE服務需求紀錄與物品分區的假想流程 | 只供練習建立對照欄位；沒有實際作業紀錄，不能認定受訓技能與營運結果有關。 | [ ] 已驗證　[x] 待查　[ ] 無關 |")
replace_line("| **C1: 完訓身分屬實**","| C1 完訓身分（SAMPLE假設） | 真實填表時核對有權使用的完訓來源 | INT-TR-01 [SAMPLE] | 模擬資料；非協會真實名冊 | 本範例未查閱真實名冊 | 僅作流程演練 | SAMPLE_ONLY | 不證明任何真實人員曾受訓 |")
replace_line("| **C2: 取得個資授權**","| C2 授權或查證權限 | 依核定目的與實際所需權限確認 | INT-AUTH-01 [SAMPLE／未提供] | 無真實卷宗；演練未建卷 | 未取得或核實權限；不得向外查證 | 不適用 | HOLD_CONSENT | 不代表協會曾受理或查閱授權 |")
replace_line("| **C3: 主管機關設立**","| C3 實體存在或登記主張 | 視主張核對適用的獨立官方或可靠來源 | EVD-SYN-02 [SAMPLE／未提供] | 無登記文件或可核查定位 | 未查詢任何真實登記資料 | 無法確認 | HOLD_MISSING_SOURCE | 不得把名稱或自述當作登記事實 |")
replace_line("| **C4: 實質營業事實**","| C4 營運主張 | 使用能支持所述營運主張的獨立來源 | EVD-SYN-03 [SAMPLE／未提供] | 無實際營運佐證 | 本範例沒有訪查或外部查詢 | 無法確認 | HOLD_MISSING_SOURCE | 登記不當然證明持續營運、收入或成功 |")
replace_line("| **C5: 訓後時序合理**","| C5 訓後時序與受訓連結 | 以完訓來源及獨立日期證據比較 | SAMPLE日期；來源未提供 | 只有虛構日期 | 日期差只供演練，未經來源證實 | 關聯未證實 | HOLD_LINKAGE | 不得以自述日期推定訓後成果或因果 |")
replace_line("* **課程名稱**","* **課程名稱**：SAMPLE-Course（服務流程設計與物品分類管理實務；純虛構課名）")
replace_line("* **學員自述**","* **學員自述（純虛構情境）**：設定為於 2025-10-01 開始提供 SAMPLE-VENTURE-01 居家收納整理服務，並於 2025-10-08 傳送一張虛構名片圖樣；本範例未收受實物、未聯絡學員，也未查詢真實登記或營運資料。")
replace_line("［2025-10-08 通報受理］","［SAMPLE 2025-10-08 情境設定；非真實受理］")
replace_line("  學員提供自製名片圖檔 1 式。","  演練只有一筆虛構名片描述；沒有實際檔案或受理紀錄。")
replace_line("  經比對：內部名冊確認於 2025-05-30 完訓，完訓身分屬實。","  演練假設有一筆虛構完訓資料；未查閱協會真實名冊，不能標示為真實核實。")
replace_line("［2025-10-09 授權盤查］","［SAMPLE 2025-10-09 權限檢核情境；非真實盤查］")
replace_line("  行政組查閱授權清冊：查無 SAMPLE-P-01 簽署之個資授權書正本。","  演練給定條件：沒有可用授權資料；未查閱任何真實清冊或授權原件。")
replace_line("［2025-10-10 第三方資料查驗］","［SAMPLE 2025-10-10 第三方來源檢核情境；未執行真實查詢］")
replace_line("  行政組檢視提供之名片（EVD-SYN-01）：無商業登記統編、無稅籍編號、無營業地址門牌。","  演練給定條件：虛構名片描述不含可供獨立查證的資料；沒有真實名片、登記資料或營運憑證。")
replace_line("登入經濟部商工公示網站查詢","本演練未查詢政府網站；缺少登記來源是測試輸入，不是真實查無登記的結論。")
replace_line("［2025-10-10 審查總結會議］","［SAMPLE 2025-10-10 規則輸出；沒有實際會議］")
replace_line("  承辦志工與專案主管共同覆核：","  規則演算示例：不代表實際志工、主管或協會曾經覆核。")
replace_line("  綜合判定狀態：HOLD_MISSING_SOURCE + HOLD_CONSENT。","   演練狀態：HOLD_MISSING_SOURCE + HOLD_CONSENT + HOLD_LINKAGE。")
replace_line("  驗證通過件數：0 / 1 件","   SAMPLE 演練驗證通過：0 / 1 件")
replace_line("* **當前判定結果**","* **演練狀態輸出**：HOLD_MISSING_SOURCE + HOLD_CONSENT + HOLD_LINKAGE；保留 Blueprint 要求的兩項 HOLD，並以 HOLD_LINKAGE 表示技能與受訓連結尚無獨立依據。本輸出不是協會實際裁定。")
replace_line("* **正式後續動作**","* **建議後續動作（未執行）**：依機構正式採用的期限、通知方式與權責流程，決定是否邀請當事人自願提供必要資料；授權及查證權限確認前不外查。若截止後仍不足，標示為「未能驗證」，不得推論為虛假。")
replace_line("本表為判定個案是否可晉升為","本表是候選流程的檢核清單，尚待機構正式採用。任一關鍵依據缺失時先 HOLD，記錄原因與補件路徑；只有獲授權人員可依正式程序更改狀態。以下勾選均屬 SAMPLE 演練，非真實審查。本表為判定個案是否可晉升為")
replace_line("| **承辦初審人**","| 承辦初審角色（虛構） | SAMPLE-ROLE-A | 演練輸出：來源與授權不足；非協會人員意見。 | —（本例未簽署） | 2025-10-10 [SAMPLE] |")
replace_line("| **專案複核人**","| 複核角色（虛構） | SAMPLE-ROLE-B | 演練規則顯示三項 HOLD；沒有實際核定、會議或期限設定。 | —（本例未簽署） | 2025-10-10 [SAMPLE] |")
replace_line("### 審查裁定簽署欄（雙人複核制）","### 演練角色紀錄（沒有實際審查或簽署）")
replace_line("* **補正法規依據**","* **候選流程依據**：本表尚無協會採用決議；授權文字、資料目的、保存期限、通知及撤回流程，須由機構依適用規範正式審查後填定。")
replace_line("  > 「學員您好","  > **未審核訊息草案，不得直接發送：**「您好，若您願意回覆訓後創業近況，請先閱讀本機構正式採用的資料使用說明，再由承辦人確認必要範圍、授權方式及可用佐證；請勿在了解用途前提供敏感資料。」")
replace_line("| 保管項目", "| 保管項目 | 機構核定的儲存方式 | 核定的安全控管 | 機構核定保存期限 | 預定銷毀日期 | 責任人與覆核人 |")
replace_line("| 授權書正本卷宗","| 授權書正本卷宗 | 依正式核定政策填寫 | 權限、位置、調閱紀錄依核定政策 | UNKNOWN（未審定） | UNKNOWN | 未指定（待正式制度） |")
replace_line("| 公示登記抄本及截圖","| 公示登記抄本及截圖 | 僅於合法取得且確有保存必要時依正式政策填寫 | 最小揭露、存取留痕及到期處置依核定政策 | UNKNOWN（未審定） | UNKNOWN | 未指定（待正式制度） |")
replace_line("法定分母定義", "* **候選分母定義 ($D$)**：SAMPLE 分母僅為此一虛構演練案的測試集；REAL 分母須由同一年度、適用訓練範圍的正式完訓名冊計算，並核對納入條件與重複。未取得來源前 REAL 分母為 UNKNOWN，不得以收件數代替。")
replace_line("|  * 時序技能存疑","|  * 身分／時序／訓練連結未釐清 (HOLD_LINKAGE) | 1 件 | UNKNOWN | 演練無獨立來源確認連結；HOLD 原因可重疊，不可加總為案件數 |")
replace_line("| **E. 實質驗證通過數","| E. 實質驗證通過數（VERIFIED_FOR_REVIEW） | 0 件（SAMPLE） | UNKNOWN | 未提供 REAL 個案來源及獨立查證結果，不將 UNKNOWN 寫成 0 |")
replace_line("| **檢定成果比率","| 檢定成果比率（Verified Rate） | 0.0% (0 / 1 SAMPLE) | UNKNOWN／不可計算 | 演練算式不代表機構成果率；REAL 分子與分母未提供 |")
# Replace fictional reminder rows and unsafe sample sign-offs.
replace_line("| "+BT+"SYN-ENT-01"+BT+" [SAMPLE] | 2025-10-11","| SYN-ENT-01 [SAMPLE] | 2025-10-11 [SAMPLE] | 未聯絡（僅示範） | SAMPLE-P-01 | 虛構待辦，不曾發送訊息或取得回覆。 | 2025-10-25 [SAMPLE；非機構期限] | 未簽名 |")
replace_line("| "+BT+"SYN-ENT-01"+BT+" [SAMPLE] | 2025-10-20","| SYN-ENT-01 [SAMPLE] | 2025-10-20 [SAMPLE] | 未聯絡（僅示範） | SAMPLE-P-01 | 第二次提醒僅為虛構待辦；沒有官方帳號、實際聯絡或回覆。 | 2025-10-25 [SAMPLE；非機構期限] | 未簽名 |")

# R11 bounded semantic repairs identified in adopted-text review.
replace_line("| 顧客服務流程與個案紀錄建置",
"| SAMPLE課程技能：需求記錄與服務流程 | SAMPLE居家整理服務的假想需求／物品清單 | 僅供示範記錄主張與課程內容；沒有真實服務、操作紀錄或第三方來源，不能認定已運用技能。 | [ ] 已驗證　[x] 待查　[ ] 無關 |")
replace_line("> **佐證有效性門檻**",
"> **獨立來源門檻**：學員自述、名片或未標明發證來源的截圖，只能作為通報線索，不能單獨證明創業主張。應依具體主張選擇可核查的獨立來源；登記資料不必然證明實際營運、收入或成功。")
replace_line("| "+BT+"EVD-SYN-01"+BT+" [SAMPLE]",
"| EVD-SYN-01 [SAMPLE] | 學員自述的創業／營運主張 | 學員本人（僅為通報來源） | 無實際檔案；只有虛構描述 | SAMPLE 情境設定；未收件 | 未查閱或檢查任何真實資料 | 無查核人員或日期 | 無真實檔案；N/A | HOLD_MISSING_SOURCE；自述不能證明登記或營運 |")
replace_line("| "+BT+"EVD-SYN-02"+BT+" [SAMPLE]",
"| EVD-SYN-02 [SAMPLE／未取得] | 登記主張 | 適用的獨立官方登記來源（待確認） | 沒有文件 | 建議的查證路徑；本例未查詢 | 未查詢任何真實來源 | 無 | 未取得；不計 SHA | HOLD_MISSING_SOURCE；不是查無登記 |")
replace_line("| "+BT+"EVD-SYN-03"+BT+" [SAMPLE]",
"| EVD-SYN-03 [SAMPLE／未取得] | 營運主張 | 能支持該營運主張的獨立來源（待確認） | 沒有文件 | 待補來源類型；本例未訪查或查詢 | 未執行真實查核 | 無 | 未取得；不計 SHA | HOLD_MISSING_SOURCE；沒有營運事實結論 |")
replace_line("| **1. 身分勾稽**",
"| **1. 身分勾稽（SAMPLE 假設）** | 個案可追溯至受訓學員與班別來源 | [x] SAMPLE假設 | [ ] | 虛構完訓資料用於演示；未核對真實名冊 |")
replace_line("| **2. 授權完備**",
"| **2. 授權／查證權限** | 目的、範圍與查證權限均依正式政策確認 | [ ] | [x] SAMPLE HOLD | 演練沒有授權；沒有真實文件或外查 |")
replace_line("| **3. 訓後時序**",
"| **3. 訓後時序** | 完訓日及創業起始日均有可核對來源 | [ ] | [x] SAMPLE HOLD | 日期差 124 天只依虛構自述計算；起始日未獨立證實 |")
replace_line("| **4. 來源獨立**",
"| **4. 來源獨立** | 來源能獨立支持具體創業主張 | [ ] | [x] SAMPLE HOLD | 只有虛構名片描述；沒有真實第三方原件 |")
replace_line("| **5. 憑證真偽**",
"| **5. 憑證真偽** | 來源、日期與完整性可依權限核查 | [ ] | [x] SAMPLE HOLD | 沒有原件、查詢紀錄或可查定位 |")
replace_line("| **6. 營運事實**",
"| **6. 營運事實** | 佐證支持所述營運主張；登記不等於持續營運或成功 | [ ] | [x] SAMPLE HOLD | 沒有營運來源，不推論收入或服務 |")
replace_line("| **7. 技能關聯**",
"| **7. 技能關聯** | 記錄課程內容與主張間的可核查連結，不推定因果 | [ ] | [x] SAMPLE HOLD | 僅有虛構課名與自述工作內容 |")
replace_line("| **8. 重複防杜**",
"| **8. 重複檢查** | 僅在有權使用的機構資料內檢查重複 | [ ] | [x] SAMPLE HOLD | 未查詢真實資料庫；重複狀態 UNKNOWN |")
replace_line("禁止擴大因果推論", "**候選流程的因果界線（待正式採用）**")
replace_line("| **訓練結訓日 (T1)**",
"| **訓練結訓日 (T1)** | 2025-05-30 [SAMPLE] | SAMPLE基準：真實案件須用機構有權使用且可追溯的完訓來源 | 本例未查閱真實名冊 |")
replace_line("| **自述設立日 (T2)**",
"| **自述設立日 (T2)** | 2025-10-01 [SAMPLE] | 演練算式：T2 - T1 = +124 天；僅為虛構日期計算 | 日期矛盾或早於完訓時先 HOLD_LINKAGE 並查來源，不直接判定主張為假 |")
replace_line("| **案件收件日 (T3)**",
"| **案件收件日 (T3)** | 2025-10-08 [SAMPLE] | SAMPLE日期順序計算；不證明真實通報或開業日 | 實際日期須由原始記錄確認 |")
replace_line("| **預定查驗日 (T4)**",
"| **預定查驗日 (T4)** | 2025-10-10 [SAMPLE] | SAMPLE目標日期；不表示實際查驗已排定 | 真實流程及期限依正式採用程序設定 |")
replace_line("2. **重啟審查條件**",
"2. **重啟審查條件（候選規則）**：取得可支持主張的新來源，並由權責人員確認資料使用目的與查證權限後，依正式程序續辦原案；保留原始 HOLD、補件與狀態變更軌跡。此規則尚未採用前不代表現行制度。")
replace_line("| **B. 時序與技能相符數**",
"| **B. SAMPLE 日期順序計算成立（非獨立核實）** | 1 件 | UNKNOWN | 只有虛構日期順序；不表示真實時間或技能關聯經核實 |")
replace_line("| **D. 排除／不成立案件數",
"| **D. 排除／未能驗證案件數** | 0 件（SAMPLE；未結案） | UNKNOWN | 演練沒有實際結案；期限及通知流程尚未採用 |")
replace_line("| "+chr(0x2003)+"* 缺第三方佐證", "| "+chr(0x2003)+"* 缺第三方佐證 (`HOLD_MISSING_SOURCE`) | 1 件 | UNKNOWN | 演練未提供獨立來源或真實原件；維持 HOLD，並非真實查無登記 |")
replace_line("| 保管項目", "| 保管項目 | 機構核定的儲存方式 | 核定的安全控管 | 機構核定保存期限 | 預定銷毀日期 | 責任人與覆核人 |")
replace_line("| 保管項目", "| 保管項目 | 機構核定的儲存方式 | 核定的安全控管 | 機構核定保存期限 | 預定銷毀日期 | 責任人與覆核人 |")
replace_line("│  僅限專案主管與",
"|  依機構核定權限 │  電子與紙本依正式政策保管 │ 保存期限依核定政策 │")
replace_line("│  授權承辦人讀取",
"|  僅限經授權人員 │  記錄存取與必要性     │ 到期依核定程序銷毀 │")
replace_line("│  (嚴禁外流志工)",
"|  定期依核定政策檢視 │ 不置於公用位置       │ 留存銷毀紀錄與覆核 │")
replace_line("本區供未來承辦人員接獲真實學員通報時",
"本區供未來承辦人員在機構正式採用流程且權限確認後，逐欄替換 SAMPLE 欄位；目前沒有真實學員通報或第三方原件。原稿：")
replace_line("| "+BT+"SAMPLE-P-01"+BT+" | 真實學員去識別化編號",
"| SAMPLE-P-01 | 真實案件去識別化代碼 | 機構受控名冊保管人（待確認） | 使用不含直接識別資訊的內部代碼；身份對照另存受控位置 | 僅依正式權限記錄必要來源 SHA-256 |")
replace_line("| "+BT+"SAMPLE-C-2025-02"+BT+" | 真實開班計畫代碼",
"| SAMPLE-C-2025-02 | 真實班別與完訓來源 | 機構受控資料保管人（待確認） | 提供可追溯定位；可公開版本移除識別資訊 | 僅對經核准的來源副本記錄 SHA-256 |")
replace_line("| "+BT+"INT-AUTH-01"+BT+" [待補]",
"| INT-AUTH-01 [待補] | 適用的資料使用授權或查證權限文件 | 依正式審查確定保管責任 | 依核定用途、期間及撤回流程填寫 | 僅對獲准使用的副本記錄 SHA-256 |")
replace_line("| "+BT+"EVD-SYN-02"+BT+" [待補]",
"| EVD-SYN-02 [待補] | 支持登記主張的適用官方來源 | 依所在地與權限確認來源 | 只記錄必要欄位及精確定位；不推論營運或成功 | 對核准的來源副本記錄 SHA-256 |")
replace_line("| "+BT+"EVD-SYN-03"+BT+" [待補]",
"| EVD-SYN-03 [待補] | 能支持具體營運主張的獨立來源 | 依主張與可用來源確認 | 區分登記、實際營運、收入及成效主張 | 對核准的來源副本記錄 SHA-256 |")
replace_line("| "+BT+"EVD-SYN-04"+BT+" [待補]",
"| EVD-SYN-04 [待補] | 支持特定營運活動的獨立來源（若該主張需要） | 依主張及正式權限確認 | 不以單一照片推定持續營運；記明日期與限制 | 對核准的來源副本記錄 SHA-256 |")
replace_line("3. [ ] **待補項 3**",
"3. [ ] **待補項 3**")

# Fixed-law, fabricated signing, health, and unsupported retention claims.
for old,new in [
("SAMPLE-VENTURE-STUDIO","SAMPLE-VENTURE-01"),
("志工林○○","SAMPLE-ROLE-A（虛構角色）"),
("主管張○○","SAMPLE-ROLE-B（虛構角色）"),
("*(SAMPLE_SIGN: SAMPLE-ROLE-A（虛構角色）)*","—（本例未簽署）"),
("*(SAMPLE_SIGN: SAMPLE-ROLE-B（虛構角色）)*","—（本例未簽署）"),
("AES-256","依機構正式核定之安全措施"),
("保存 3 年屆滿：","保存期限依核定政策："),
("結訓日起算 3 年","依核定政策（UNKNOWN）"),
("評核結束後 1 年","依核定政策（UNKNOWN）"),
("個人資料保護法第 8 條、第 19 條及本會個資管理要點","機構尚待核定的個資及檔案管理政策；本文件不提供法律結論"),
("限期 14 日補齊","依機構正式採用的期限補件"),
("官方 Line@","尚未指定聯絡管道"),
("經查內部歷史資料庫無重複登錄","演練未查詢真實資料庫；重複狀態 UNKNOWN"),
("0 件 (NOT_GRANTED)","UNKNOWN（REAL 資料未提供）"),
("推拿","居家收納整理"),
("保健居家收納整理調理實務專案班","服務流程設計與物品分類管理實務"),
("健康舒壓工作室","SAMPLE居家整理服務"),
("本會","本機構"),
("SAMPLE_SIGN: ",""),
("法定保存年限","機構核定保存期限"),
("法定個資利用與評核呈閱授權","經正式審查的資料使用與呈閱權限"),
("依法不得對外揭露","本範例未確認揭露權限，因此維持 HOLD"),
("雙人簽署","依機構正式核定權責記錄複核"),
("本表專供小型訓練機構／協會行政與志工人員","本候選表供小型訓練機構／協會行政與志工人員"),
("經濟部商工登記公開網站查冊列印頁","擬議查證路徑：依個案所在地、授權及正式流程確認適用官方來源（本例未查詢）"),
("經濟部商工登記公開資訊網站","適用的官方商業登記來源（依個案所在地與正式權限確認）"),
("推拿放鬆業務","居家收納整理服務"),
("推拿調理施作空間","居家收納整理服務空間"),
("保健推拿調理實務專案班","服務流程設計與物品分類管理實務"),
("推拿放鬆服務","居家收納整理服務"),
("推拿調理","居家收納整理服務"),
]:
    replace_text(old,new,expected=0 if old not in "\n".join(lines) else sum(x.count(old) for x in lines))
content="\n".join(lines)+"\n"
content=content.replace("SAMPLE-VENTURE-01」從事推拿放鬆服務","SAMPLE-VENTURE-01」提供居家收納整理服務")
content=content.replace("登入經濟部商工公示網站查詢「SAMPLE-VENTURE-01」：查無此筆登記。","本演練未查詢政府網站；缺少登記來源是測試輸入，不是真實查無登記的結論。")
content=content.replace("### 一、時序先後檢核表","### 一、時序先後檢核表\n\n> 以下日期和時間軸只是 SAMPLE 算術演練，不是實際完訓、創業、收件或查驗紀錄。")
content=content.replace("### 二、逐項查驗歷程紀錄","### 二、逐項查驗歷程紀錄\n\n> 本節所有日期、來源描述與狀態均為虛構測試輸入；沒有真實名冊查閱、網站查詢、聯絡、會議或機構人員作業。")
content=content.replace("### 二、統計彙整表（含演練與真實數據隔離）","### 二、統計彙整表（含演練與真實數據隔離）\n\n> HOLD 原因可重疊，不可加總為案件數。REAL 數據一律 UNKNOWN，直到取得經授權的原始來源及核對紀錄。")
content=content.replace("承辦人員簽章","紀錄狀態（非簽章）")
content=content.replace("聯絡追蹤日期","SAMPLE 情境日期")
content=content.replace("查無 SAMPLE-P-01 簽署之個資授權書正本。","演練給定條件為未提供授權資料；未查閱真實清冊。")
content=content.replace("本表專供","本候選表供")
content=content.replace("HOLD_MISSING_SOURCE + HOLD_CONSENT。","HOLD_MISSING_SOURCE + HOLD_CONSENT + HOLD_LINKAGE。")
content=content.replace("演練案因缺授權與佐證入保留","演練案因缺授權、第三方來源及訓練連結佐證入保留")
content=content.replace("法定授權書","經正式審查採用的授權／查證權限文件")
content=content.replace("合法登記謄本","適用的登記來源（如經核准可使用）")
content=content.replace("法定個資授權利用正本（合法調閱與呈閱之法源前提）","經正式審查採用的資料使用授權／查證權限文件（未確認前不外查）")
content=content.replace("僅提供名片圖檔（強制觸發 HOLD_MISSING）","僅有 SAMPLE 名片描述；非實際查件（HOLD_MISSING_SOURCE）")
content=content.replace("演練案例時序相符；真實狀況待盤點","演練日期只有自述順序，來源未證實；REAL UNKNOWN")
content=content.replace("演練案時序符合","演練日期依自述排列，未獨立核實")
content=content.replace("演練案尚未逾期，暫未列入不成立","演練未實際結案；期限及通知流程尚未採用")
content=content.replace("虛擬名片未具法律效力","虛構名片描述不是獨立證據")
content=content.replace("未簽署授權書，依法不得對外揭露","演練未提供授權；查證權限未確認，維持 HOLD")
content=content.replace("本機構明確界定","本候選流程建議採用的界線（尚待機構正式採納）")
content=content.replace("結案轉列 "+BT+"NOT_VERIFIED"+BT+" 條件：逾補件截止日仍未繳交授權書、未提供第三方公信證明，或學員明確表達不願提供者，立即終止追蹤並以查證不成立結案。","候選結案條件：依機構正式採用的期限與通知程序完成後仍缺來源時，標示為未能驗證；不得推論為已證明主張不實。")
content=content.replace("重啟審查條件：學員主動補具合法登記謄本並簽妥法定授權書後，承辦人員得重開案件編號，循流程重行啟動「步驟 1 至步驟 5」審查。","重啟審查條件：取得適用的獨立來源並由權責人員確認查證權限後，依機構正式程序重開或續辦個案，保留原缺口與補件歷程。")
content=content.replace("以簽到退齊備之名冊為準","以機構有權使用且可追溯的完訓來源為準")
content=content.replace("若 $T2 < T1$ 則判定非訓後創業，入 HOLD_LINKAGE","若日期矛盾或早於完訓日，先 HOLD_LINKAGE 並查來源，不直接認定主張為假")
content=content.replace("演練案時序符合","演練日期依自述排列，未獨立核實")
content=content.replace("應以主管機關正式證明為唯一可採來源","應選擇能支持該具體主張的獨立來源並記錄限制")
content=content.replace("法定分母定義","候選分母定義")
content=content.replace("法定權利","權利")

content=content.replace("0 件 (NOT_GRANTED)","UNKNOWN（REAL 資料未提供）")
content=content.replace("[!WARNING]", "重要警示")
content=content.replace("缺公務機關憑據", "缺乏足以支持主張的獨立來源")
content=content.replace("* **結訓日期**：`2025-05-30`（簽到退成績及格，具完訓名冊底冊）", "* **結訓日期（SAMPLE 假設）**：2025-05-30；未查閱真實簽到、成績或完訓名冊。")
content=content.replace("* **預定查證**：行政組預排於 `2025-10-10` 進行查核作業。", "* **演練目標日期（SAMPLE）**：2025-10-10；沒有實際排程或查核。")
content=content.replace("[x] 已落實遮罩", "[x] SAMPLE 假設：示範遮罩代碼")
content=content.replace("負責保管人：行政組志工幹事（待指定）", "保管責任人：未指定（待機構正式制度）")
content=content.replace("本表為判定個案是否可晉升為", "")
content=content.replace("逾補件截止日仍未繳交授權書、未提供第三方公信證明，或學員明確表達不願提供者，立即終止追蹤並以查證不成立結案。", "依機構正式採用的期限及通知程序完成後仍缺來源時，標示為未能驗證；不得推論為已證明主張不實。")
content=content.replace("> [!CAUTION]\n**候選流程的因果界線（待正式採用）**", "> [!NOTE]\n> SAMPLE 日期順序與技能對照僅供欄位演練；不證明受訓造成創業、營運、收入或成功。此候選規則仍待機構檢視採用。")
def replace_fenced_block(text, marker, replacement):
    rows=text.splitlines(); start=None; j=0
    while j < len(rows):
        if rows[j].startswith(FENCE):
            if start is None: start=j
            else:
                if any(marker in z for z in rows[start+1:j]):
                    rows[start:j+1]=[replacement]; return "\n".join(rows)+"\n"
                start=None
        j+=1
    raise RuntimeError("Fenced block not found for marker: "+marker)
content=replace_fenced_block(content, "［步驟 1：通報收件］", """1. **通報收件**：記錄遮罩代碼、班別與學員自述；標明資料性質，不抄錄不必要的直接識別資料。
2. **權限檢核**：確認用途、授權範圍與查證權限；任一項未確認即列 `HOLD_CONSENT`，停止外查及揭露。
3. **時序與訓練連結**：對照可追溯的完訓日期、課程內容、主張日期及來源日期；只記錄先後與可核對關聯，不推定因果。
4. **來源查驗**：依個案具體主張選用可獨立核對的來源；缺來源、來源無法驗證或權限不足時，記錄原因並維持 HOLD。
5. **權責複核**：依機構正式採用程序完成初審與複核；本候選的虛構角色與示範勾選不是協會實際審查或核准。""")
content=replace_fenced_block(content, "間隔 124 天", "> SAMPLE 時序演練：完訓 2025-05-30 →（+124 天）→自述開始 2025-10-01 →（+7 天）→通報 2025-10-08 →（+2 天）→預定日期 2025-10-10。僅為日期運算，未發生真實查核。")
content=replace_fenced_block(content, "個資與機敏卷宗安全生命週期", "> 候選保管流程：依核定權限存取；電子與紙本依正式制度保管；期限、責任人及銷毀程序待機構確認。本例未設定實際保管人或保存期限。")
formula=content.splitlines()
for j,z in enumerate(formula):
    if "Verified Entrepreneurship Rate" in z:
        formula[j]="  創業成果率（演練）＝N（符合可驗證狀態的案件數）÷D（同年度適用範圍完訓人數）×100%；REAL來源未核實時不計算。"
content="\n".join(formula)+"\n"
old_source_list = """### 待補之第三方原始公信憑證清單（截至本日現狀）
1. [ ] **待補項 1**：經濟部商工登記核准函或公開查冊抄本（證明實體合法設立）。
2. [ ] **待補項 2**：主管稽徵機關之稅籍設立核准公函影本（證明具有實質稅籍與營業事實）。
3. [ ] **待補項 3**
"""
new_source_list = """### 依個案主張選取的來源補件清單（SAMPLE；本候選未附下列 REAL 原件）

> 本清單列出可供權責人員評估的來源類型，不代表協會已查詢，也不是每案都必須湊齊三份。先確認資料目的、使用權限與必要範圍；來源或權限不足時採 HOLD。此候選未附真實個案原件，不能據此推定協會整體資料現況。

1. [ ] **設立／登記主張（條件式）**：個案主張已登記時，依所在地及經確認的權限選擇可追溯的官方登記來源，記錄來源機關、查詢日期、定位與支持主張的欄位。它只支持查詢時點的登記狀態，不證明持續營運、收入或成功；REAL 原件及查詢結果目前 UNKNOWN。
2. [ ] **具體營運主張（條件式）**：個案主張提供服務或完成交易時，依該主張選擇可獨立核對且有權使用的來源（例如去識別的服務／交易紀錄或適當第三方資料），記錄日期、定位、範圍與限制。名片、自述或單張照片不足以單獨證明持續營運或收入；REAL 原件目前 UNKNOWN。
3. [ ] **訓後時序與技能連結（逐案核對）**：以有權使用且可追溯的完訓日期、課程內容資料，對照具體創業主張及其來源日期；只支持時間先後與內容關聯，不證明訓練造成創業。缺少來源或權限時標示 HOLD_LINKAGE；REAL 原件目前 UNKNOWN。
"""
if old_source_list not in content:
    raise RuntimeError("Expected incomplete three-item source list was not found")
content=content.replace(old_source_list,new_source_list,1)
content=content.replace("UNKNOWN / NOT_GRANTED", "UNKNOWN / 授權狀態未核實")
content=content.replace("目前無已授權之真實案例，所有欄位均為流程演練規範", "本候選未附可核實的真實案例或授權來源；協會整體 REAL 狀態 UNKNOWN")
content=content.replace("僅有學員自述、自製名片、社群貼文或未具公信力之私人照片，缺乏商業登記、稅籍或獨立機構查核憑證。", "沒有足以支持具體主張的獨立來源；自述、名片或未標發行來源／日期的截圖只能作通報線索。登記主張需核對登記來源；實際營運主張另需與營運內容相稱的來源。")
content=content.replace("是否簽署專用《學員訓後成果追蹤與個資利用授權書》正本？", "是否有適用的資料使用授權或其他經確認的查證權限依據？")
content=content.replace("是否明確載明「同意協會向商業主管機關查調並於 TTQS 評核呈閱」？", "是否涵蓋本次資料目的、必要範圍、使用對象及可能呈閱情境？")
content=content.replace("在未取得學員親筆簽名之個資同意書前，嚴禁以學員個人資料進行外部公開宣傳或未經授權之非公務調查。", "未依正式採用程序確認資料目的、使用範圍及查證權限前，不得進行外部查調或揭露個人資料；本候選表不自行決定特定同意形式。")
content=content.replace("為防止評核資料混淆與錯置，本表嚴格界定與周邊相鄰指標之作業邊界，並列舉不可逾越之禁絕事項：", "為防止評核資料混淆與錯置，本表界定與周邊指標的作業邊界及禁絕事項。本文件仍為 CANDIDATE；未取得並核實 REAL 來源及完成正式授權審查前，不得宣稱協會已有創業查核實績。")
content=content.replace("4. **嚴禁樣板冒充正式成果**：未完成真實替換清冊（第拾壹節）前，本文件僅得標示為 `CANDIDATE` 供流程展示，絕不可向評核委員宣稱協會已具備實質創業查核通過績效。", "")

def replace_markdown_table(text, prefix, replacement):
    rows=text.splitlines(); start=None; end=None
    for j,row in enumerate(rows):
        if row.startswith(prefix):
            start=j; end=j+1
            while end<len(rows) and rows[end].lstrip().startswith("|"):
                end+=1
            break
    if start is None:
        raise RuntimeError("Expected Markdown table missing: "+prefix)
    rows[start:end]=replacement.strip("\n").splitlines()
    return "\n".join(rows)+"\n"

content=replace_markdown_table(content, "| 佐證編號 |", """| 佐證編號／性質 | 主張、來源發行者與原件形式 | 日期、取得途徑與精確定位 | 授權、防偽查驗、查驗人員／日期與遮罩 | 真實檔案 SHA-256、結論限制與補正動作 |
| :--- | :--- | :--- | :--- | :--- |
| EVD-SYN-01 [SAMPLE] | 學員自述創業／營運；學員本人只作通報來源；無真實原件或副本 | 僅 SAMPLE 情境；未收件，沒有查詢途徑或可定位日期 | 真實權限未核實；未作查驗，沒有查驗人員／日期；代碼僅示範遮罩 | 無真實檔案，SHA=N/A；HOLD_MISSING_SOURCE；自述不能證明登記或營運 |
| EVD-SYN-02 [SAMPLE／未取得] | 登記主張；適用的獨立官方來源及發行機關尚未確認；未取得原件／副本 | 未查詢；無查詢日期、網址或定位 | 權限未確認；未執行查驗，沒有查驗人員／日期；無真實資料可遮罩 | 未取得、不計 SHA；HOLD_MISSING_SOURCE，不等於查無登記；確認權限後再選可核來源 |
| EVD-SYN-03 [SAMPLE／未取得] | 營運主張；需選擇支持具體主張的獨立來源；未取得原件／副本 | 未訪查或查詢；無日期、取得管道或定位 | 權限未確認；未查驗；查驗人員／日期與遮罩狀態 UNKNOWN | 無真實檔案、不計 SHA；HOLD_MISSING_SOURCE；不作營運、收入或成功結論 |""")

content=replace_markdown_table(content, "| 主張項目 (Claim) |", """| 主張及必要佐證 | 證據ID、精確定位與防偽核對 | 訓練技能關聯 | 當前狀態 | 限制、補正動作、責任與期限 |
| :--- | :--- | :--- | :--- | :--- |
| C1 完訓身分（SAMPLE 假設）；真實填表時核對有權使用的完訓來源 | INT-TR-01 [SAMPLE]；模擬資料，非協會名冊；本例未查閱來源 | 只供流程演練，REAL 未核實 | SAMPLE_ONLY | 不證明任何真實人員曾受訓；正式採用後依權責流程補來源並指定責任／期限 |
| C2 授權或查證權限；依核定目的和實際所需範圍確認 | INT-AUTH-01 [SAMPLE／未提供]；無真實卷宗、無定位或查驗 | 不適用 | HOLD_CONSENT | 未確認目的、範圍與權限前不外查；由權責角色依正式程序確認，期限未設定 |
| C3 實體存在或登記主張；視主張選用可核對的獨立來源 | EVD-SYN-02 [SAMPLE／未提供]；沒有登記文件、查詢日期或定位 | 不能由名稱或課程推定 | HOLD_MISSING_SOURCE | REAL UNKNOWN；不代表查無登記；確認權限後取得來源並記錄日期／定位，責任與期限依採用程序指定 |
| C4 具體營運主張；來源須支持所述服務或交易事實 | EVD-SYN-03 [SAMPLE／未提供]；沒有營運來源或查核 | 無法確認 | HOLD_MISSING_SOURCE | 登記不證明持續營運、收入或成功；先定義主張，再補相稱來源，權責與期限依採用程序指定 |
| C5 訓後時序與受訓連結；比對可追溯完訓來源及獨立日期證據 | SAMPLE 日期；來源未提供；日期差僅算術演練，未經來源核實 | 關聯未證實 | HOLD_LINKAGE | 不推定訓後成果或因果；取得可用來源後再核對，責任與期限依正式程序指定 |""")
for bad in ("SAMPLE_SIGN","志工林○○","主管張○○","AES-256","保存 3 年屆滿","個人資料保護法第 8 條","限期 14 日補齊","推拿","健康舒壓","經查內部歷史資料庫無重複登錄","登入經濟部商工公示網站","[!WARNING]","Verified Entrepreneurship Rate","簽到退成績及格","行政組預排","本表為判定個案是否可晉升為"):
    if bad in content: raise RuntimeError("Disallowed content survived: "+bad)
for needed in ("HOLD_MISSING_SOURCE","HOLD_CONSENT","HOLD_LINKAGE","SAMPLE-VENTURE-01","0 / 1","UNKNOWN","本例未簽署"):
    if needed not in content: raise RuntimeError("Required content missing: "+needed)
lines=content.splitlines()  # build Word from the exact adopted text after deterministic replacements
ADOPTED.write_text(content,encoding="utf-8",newline="\n")

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
sec.top_margin=Inches(.68); sec.bottom_margin=Inches(.68); sec.left_margin=Inches(.66); sec.right_margin=Inches(.66)
sec.header_distance=Inches(.32); sec.footer_distance=Inches(.32)
def setfont(run,name="Microsoft JhengHei",size=10.5,bold=None,color="222222"):
    run.font.name=name; run.font.size=Pt(size); run.font.color.rgb=RGBColor.from_string(color)
    if bold is not None: run.bold=bold
    rp=run._element.get_or_add_rPr(); rf=rp.rFonts
    if rf is None: rf=OxmlElement("w:rFonts"); rp.insert(0,rf)
    for k in ("ascii","hAnsi","eastAsia","cs"): rf.set(qn("w:"+k),name)
for nm in ("Normal","Body Text","List Bullet","List Number"):
    st=doc.styles[nm]; st.font.name="Microsoft JhengHei"; st.font.size=Pt(10.5); st.font.color.rgb=RGBColor(34,34,34)
    st._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft JhengHei"); st.paragraph_format.space_after=Pt(5); st.paragraph_format.line_spacing=1.12
for nm,size in (("Title",23),("Heading 1",16),("Heading 2",13),("Heading 3",11.5),("Heading 4",10.5)):
    st=doc.styles[nm]; st.font.name="Microsoft JhengHei"; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor(0,0,0)
    st._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft JhengHei"); st.paragraph_format.keep_with_next=True; st.paragraph_format.space_before=Pt(10 if nm!="Title" else 0); st.paragraph_format.space_after=Pt(5)
title_ppr=doc.styles["Title"]._element.pPr
if title_ppr is not None:
    title_border=title_ppr.find(qn("w:pBdr"))
    if title_border is not None: title_ppr.remove(title_border)
h=sec.header.paragraphs[0]; h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=h.add_run("TTQS ONE  |  0130  |  CANDIDATE  ·  SAMPLE / SYNTHETIC"); setfont(r,size=8.5,color="53636B")
f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=f.add_run("候選表單  |  REAL 與正式接受未確認  |  第 "); setfont(r,size=8.5,color="53636B")
fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),"PAGE"); f._p.append(fld)
r=f.add_run(" 頁"); setfont(r,size=8.5,color="53636B")
def shade(cell,fill):
    pr=cell._tc.get_or_add_tcPr(); sh=OxmlElement("w:shd"); sh.set(qn("w:fill"),fill); pr.append(sh)
def margins(cell):
    pr=cell._tc.get_or_add_tcPr(); mar=OxmlElement("w:tcMar")
    for side,val in (("top",70),("start",75),("bottom",70),("end",75)):
        e=OxmlElement("w:"+side); e.set(qn("w:w"),str(val)); e.set(qn("w:type"),"dxa"); mar.append(e)
    pr.append(mar)
def inline(p,text,size=10.5,color="222222"):
    text=text.replace("<br>","\n").replace("<br/>","\n")
    pat="("+re.escape("**")+".+?"+re.escape("**")+"|"+re.escape(BT)+".+?"+re.escape(BT)+"|\\*.+?\\*)"
    for part in re.split(pat,text):
        if not part: continue
        bold=part.startswith("**") and part.endswith("**"); isc=part.startswith(BT) and part.endswith(BT); ital=not bold and not isc and part.startswith("*") and part.endswith("*")
        val=part[2:-2] if bold else part[1:-1] if isc or ital else part
        for n,chunk in enumerate(val.split("\n")):
            if n: p.add_run().add_break()
            if chunk:
                rr=p.add_run(chunk); setfont(rr,"Consolas" if isc else "Microsoft JhengHei",size,True if bold else None,"42576A" if isc else color)
                if ital: rr.italic=True
def splitrow(s):
    s=s.strip()
    if s.startswith("|"): s=s[1:]
    if s.endswith("|"): s=s[:-1]
    return [z.strip() for z in re.split(r"(?<!\\)\|",s)]
def addtable(rows):
    parsed=[splitrow(x) for x in rows if not re.match(r"^\s*\|?\s*:?-{2,}",x)]
    if not parsed: return
    n=max(len(x) for x in parsed); t=doc.add_table(rows=len(parsed),cols=n); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    avail=sec.page_width-sec.left_margin-sec.right_margin
    weights={2:[.30,.70],3:[.18,.34,.48],4:[.17,.23,.28,.32],5:[.13,.17,.20,.24,.26],6:[.12,.14,.17,.18,.19,.20],7:[.10,.13,.14,.14,.14,.15,.20],8:[.10,.13,.12,.15,.11,.13,.10,.26]}.get(n,[1/n]*n)
    if n==5 and parsed[0][0].startswith("佐證編號／性質"):
        weights=[.15,.22,.19,.22,.22]
    elif n==5 and parsed[0][0].startswith("主張及必要佐證"):
        weights=[.19,.21,.14,.13,.33]
    for ri,vals in enumerate(parsed):
        row=t.rows[ri]
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if ri==0:
            tr=row._tr.get_or_add_trPr(); e=OxmlElement("w:tblHeader"); e.set(qn("w:val"),"true"); tr.append(e)
        for ci,cell in enumerate(row.cells):
            cell.width=avail*weights[ci]; cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP; margins(cell); cell.text=""
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(1); p.paragraph_format.line_spacing=1.02; inline(p,vals[ci] if ci<len(vals) else "",9.5)
            if ri==0:
                shade(cell,"274C5E")
                for rr in p.runs: rr.font.color.rgb=RGBColor(255,255,255); rr.bold=True
            elif ri%2==0: shade(cell,"F2F6F7")
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
def addcode(block):
    for x in block:
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.1); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1; inline(p,x,10.5,"34434A")
i=0
while i<len(lines):
    s=lines[i]
    if not s.strip() or s.strip()=="---": i+=1; continue
    if s.startswith(FENCE):
        i+=1; block=[]
        while i<len(lines) and not lines[i].startswith(FENCE): block.append(lines[i]); i+=1
        if i<len(lines): i+=1
        addcode(block); continue
    if s.lstrip().startswith("|"):
        rows=[]
        while i<len(lines) and lines[i].lstrip().startswith("|"): rows.append(lines[i]); i+=1
        addtable(rows); continue
    m=re.match(r"^(#{1,4})\s+(.*)$",s)
    if m:
        level=len(m.group(1)); p=doc.add_paragraph(style="Title" if level==1 else "Heading "+str(min(level-1,4))); inline(p,m.group(2),23 if level==1 else 13,"000000")
        i+=1; continue
    if s.startswith(">"):
        val=s[1:].strip().replace("[!IMPORTANT]","重要提示").replace("[!NOTE]","演練提示").replace("[!CAUTION]","統計限制")
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.18); p.paragraph_format.space_after=Pt(5); inline(p,val,10.5,"32434A")
        sh=OxmlElement("w:shd"); sh.set(qn("w:fill"),"EDF3F4"); p._p.get_or_add_pPr().append(sh); i+=1; continue
    bullet=re.match(r"^\s*[-*]\s+(.*)$",s); number=re.match(r"^\s*(\d+)\.\s+(.*)$",s)
    if bullet:
        p=doc.add_paragraph(style="List Bullet"); inline(p,bullet.group(1)); i+=1; continue
    if number:
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.25); p.paragraph_format.first_line_indent=Inches(-.25)
        inline(p,number.group(1)+". "+number.group(2)); i+=1; continue
    para=[s.strip()]; i+=1
    while i<len(lines) and lines[i].strip() and lines[i].strip()!="---" and not lines[i].startswith("#") and not lines[i].startswith(">") and not lines[i].lstrip().startswith("|") and not lines[i].startswith(FENCE) and not re.match(r"^\s*[-*]\s+|^\s*\d+\.\s+",lines[i]):
        para.append(lines[i].strip()); i+=1
    p=doc.add_paragraph(); inline(p," ".join(para))
doc.core_properties.title="0130 學員訓後創業案例紀錄與佐證表"
doc.core_properties.subject="I19-E19 candidate with synthetic worked case"
doc.core_properties.author="TTQS ONE candidate production"; doc.core_properties.keywords="0130, I19-E19, SAMPLE, SYNTHETIC, CANDIDATE"
doc.save(DOCX)
with zipfile.ZipFile(DOCX) as z:
    corrupt=z.testzip()
    if corrupt: raise RuntimeError("OOXML ZIP corrupt: "+corrupt)
v=Document(DOCX); alltext="\n".join([p.text for p in v.paragraphs]+[c.text for t in v.tables for row in t.rows for c in row.cells])
for word in ("HOLD_MISSING_SOURCE","HOLD_CONSENT","HOLD_LINKAGE","SAMPLE-VENTURE-01","UNKNOWN"):
    if word not in alltext: raise RuntimeError("DOCX text missing "+word)
for word in ("SAMPLE_SIGN","AES-256","推拿","志工林○○","主管張○○"):
    if word in alltext: raise RuntimeError("Forbidden term in DOCX: "+word)
table_pt=[r.font.size.pt for t in v.tables for row in t.rows for c in row.cells for p in c.paragraphs for r in p.runs if r.text.strip() and r.font.size]
if len(v.tables)<10: raise RuntimeError("Unexpectedly few form tables: "+str(len(v.tables)))
if table_pt and min(table_pt)<9.4: raise RuntimeError("Table font below target: "+str(min(table_pt)))
adopt_sha=hashlib.sha256(ADOPTED.read_bytes()).hexdigest(); docx_sha=hashlib.sha256(DOCX.read_bytes()).hexdigest()
receipt={
 "project_id":"TTQS_ONE","doc_id":"0130","internal_requirement_id":"I19-E19","official_item":"學員訓後創業案例。","official_printed_checkbox_serial":"UNKNOWN",
 "candidate_revision":"R01B","writer_route":"Antigravity CLI requested gemini-3.8-flash-high; actual model identity UNKNOWN",
 "raw_writer_response_sha256":raw_sha,"adopted_markdown_sha256":adopt_sha,"candidate_docx_sha256":docx_sha,
 "changes":[
  "health/massage case replaced with synthetic home-organization service and neutral course mapping",
  "simulated roster/database/site lookup, contact, meeting and association action claims marked unperformed",
  "fabricated people/signatures removed",
  "unsupported law citations, fixed deadlines, encryption and retention periods replaced by policy-review fields or UNKNOWN",
  "HOLD_LINKAGE added for unsupported training/skill relationship; required source and consent holds retained",
  "REAL counts/rates set to UNKNOWN; sample-only 0/1 retained"
 ],
 "validation":{"ooxml_zip_test":"PASS","python_docx_reopen":"PASS","tables":len(v.tables),"paragraphs":len(v.paragraphs),"table_font_min_pt":min(table_pt) if table_pt else None,"visual_render":"PENDING","independent_exact_sha_semantic_review":"PENDING","formal_acceptance":"NOT_GRANTED"},
 "created_at_local":datetime.now().astimezone().isoformat(),"change_log":log,
 "artifacts":{"adopted_markdown":str(ADOPTED),"candidate_docx":str(DOCX)}
}
RECEIPT.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
print(json.dumps({"adopted_markdown":str(ADOPTED),"adopted_sha256":adopt_sha,"docx":str(DOCX),"docx_sha256":docx_sha,"receipt":str(RECEIPT),"tables":len(v.tables),"paragraphs":len(v.paragraphs),"changes":len(log)},ensure_ascii=False))
