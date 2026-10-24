import hashlib
import html
import json
import re
import zipfile
from pathlib import Path

from docx import Document

BASE = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0130_I19-E19_SOURCE_PREP_R01_20261024")
MD = BASE / "candidate" / "0130_I19-E19_adopted_R01B.md"
DOCX = BASE / "candidate" / "0130_I19-E19_CANDIDATE_R01B.docx"


def norm(value: str) -> str:
    value = html.unescape(value).replace("<br/>", "\n").replace("<br>", "\n")
    value = value.replace("\\|", "|").replace("`", "")
    value = re.sub(r"\*\*(.*?)\*\*", r"\1", value)
    value = re.sub(r"(?<!\*)\*(?!\*)", "", value)
    return re.sub(r"\s+", " ", value).strip()


def markdown_tables(text: str):
    rows = []
    in_fence = False
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        if lines[i].startswith("```"):
            in_fence = not in_fence
            i += 1
            continue
        if not in_fence and lines[i].lstrip().startswith("|"):
            table = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                raw = lines[i].strip().strip("|")
                cells = [cell.strip() for cell in re.split(r"(?<!\\)\|", raw)]
                if not all(re.fullmatch(r"\s*:?-{3,}:?\s*", cell) for cell in cells):
                    table.append([norm(cell) for cell in cells])
                i += 1
            rows.append(table)
            continue
        i += 1
    return rows


md_bytes = MD.read_bytes()
docx_bytes = DOCX.read_bytes()
text = md_bytes.decode("utf-8")
doc = Document(DOCX)
with zipfile.ZipFile(DOCX) as zf:
    assert zf.testzip() is None, "OOXML ZIP member CRC failure"

paragraph_text = [p.text for p in doc.paragraphs]
table_text = [[cell.text for row in table.rows for cell in row.cells] for table in doc.tables]
all_text = "\n".join(paragraph_text + [cell for table in table_text for cell in table])

assert "I19-E19" in all_text
assert "143（I19-L3）" in all_text
assert "UNKNOWN" in all_text
assert "UNKNOWN / 授權狀態未核實" in all_text
assert "NOT_GRANTED" not in all_text
assert "待補項 3" not in text and "待補項 3" not in all_text
assert "依個案主張選取的來源補件清單" in text and "依個案主張選取的來源補件清單" in all_text
assert "1. 1." not in all_text and "2. 2." not in all_text
assert not any(re.match(r"^\s*\d+\.\s+\d+\.", p) for p in paragraph_text)

md_tabs = markdown_tables(text)
assert len(md_tabs) == len(doc.tables), (len(md_tabs), len(doc.tables))
assert any(len(rows[0]) == 5 and rows[0][0] == "佐證編號／性質" for rows in md_tabs)
assert any(len(rows[0]) == 5 and rows[0][0] == "主張及必要佐證" for rows in md_tabs)
table_mismatches = []
for ti, (expected, actual) in enumerate(zip(md_tabs, doc.tables), start=1):
    actual_rows = [[norm(cell.text) for cell in row.cells] for row in actual.rows]
    if expected != actual_rows:
        table_mismatches.append({"table": ti, "markdown": expected, "docx": actual_rows})
assert not table_mismatches, json.dumps(table_mismatches[:1], ensure_ascii=True)

body_sizes = [r.font.size.pt for p in doc.paragraphs for r in p.runs if r.text.strip() and r.font.size]
table_sizes = [r.font.size.pt for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs for r in p.runs if r.text.strip() and r.font.size]
assert body_sizes and min(body_sizes) >= 10.5, min(body_sizes) if body_sizes else None
assert table_sizes and min(table_sizes) >= 9.5, min(table_sizes) if table_sizes else None

result = {
    "project_id": "TTQS_ONE",
    "doc_id": "0130",
    "internal_requirement_id": "I19-E19",
    "project_157_index": 143,
    "project_157_blueprint_id": "I19-L3",
    "official_printed_checkbox_serial": "UNKNOWN",
    "docx_sha256": hashlib.sha256(docx_bytes).hexdigest(),
    "markdown_sha256": hashlib.sha256(md_bytes).hexdigest(),
    "zip_crc": "PASS",
    "table_count": len(doc.tables),
    "markdown_docx_table_cells": "EXACT_MATCH",
    "duplicate_numbered_paragraphs": 0,
    "incomplete_source_placeholder": "ABSENT",
    "body_min_font_pt": min(body_sizes),
    "table_min_font_pt": min(table_sizes),
    "semantic_review": "PENDING_INDEPENDENT_EXACT_SHA",
    "formal_acceptance": "NOT_GRANTED",
}
print(json.dumps(result, ensure_ascii=True, sort_keys=True))
