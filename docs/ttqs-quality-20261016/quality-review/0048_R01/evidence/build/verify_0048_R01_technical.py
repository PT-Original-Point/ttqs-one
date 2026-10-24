from __future__ import annotations
import csv, hashlib, json, zipfile
from pathlib import Path
from docx import Document
from pypdf import PdfReader

root = Path(__file__).resolve().parents[2]
evidence = root / "evidence"
review = root / "review"
docx_path = root / "candidate" / "0048_R01.docx"
register_path = evidence / "SYNTHETIC_REPLACEMENT_REGISTER_R01.csv"
blueprint_path = root / "inputs" / "UNIQUE_BLUEPRINT_0048.jsonl"
source_map_path = root / "inputs" / "SOURCE_MAP_R01.json"
render_dir = evidence / "render_final_R01_e6bb07cc"
pdf_path = render_dir / "0048_R01.pdf"

sha = hashlib.sha256(docx_path.read_bytes()).hexdigest()
source_map = json.loads(source_map_path.read_text(encoding="utf-8"))
blueprint = json.loads(blueprint_path.read_text(encoding="utf-8").splitlines()[0])
with register_path.open(encoding="utf-8-sig", newline="") as f:
    register = list(csv.DictReader(f))
doc = Document(str(docx_path))
paragraph_text = [p.text for p in doc.paragraphs]
table_text = [[ [cell.text for cell in row.cells] for row in table.rows] for table in doc.tables]
all_text = "\n".join(paragraph_text + ["\n".join(" | ".join(row) for row in table) for table in table_text])

with zipfile.ZipFile(docx_path) as zf:
    bad_member = zf.testzip()
    styles_xml = zf.read("word/styles.xml")
    docx_member_count = len(zf.namelist())
assert bad_member is None, f"Corrupt DOCX member: {bad_member}"
assert source_map["doc_id"] == "0048"
assert source_map["requirement_id"] == "I10-E03"
assert source_map["official_checkbox_sequence"] == "049"
assert source_map["official_source_box_id"] == "P03-B10"
assert blueprint.get("deliverable_id") == "0048"
assert source_map["blueprint_row_sha256"] == "8b258f9125f2b1209629d0356539537e00c322073739ec65d05c58590847ef72"
assert len(register) == 39
assert len({r["fact_id"] for r in register}) == 39
assert all(r["status"] == "SAMPLE/SYNTHETIC" for r in register)
assert all(r["fact_id"] in all_text for r in register)
flat_text = "".join(all_text.split())
for token in ("大型採購均辦理會議通過", "師資遴聘後亦經過會議或審查通過並留有紀錄", "DOCX0048", "I10-E03", "P03-B10", "SAMPLE/SYNTHETIC", "本版未經協會核定或生效", "大型採購", "師資遴選後審查"):
    assert "".join(token.split()) in flat_text, f"Required content absent: {token}"
# Independently recompute the displayed synthetic arithmetic.
assert 120000 - 113000 == 7000
assert round(113000 / 120000 * 100, 1) == 94.2
weights = [35, 25, 15, 15, 10]
assert sum(weights) == 100
x = [5, 4, 4, 4, 5]
y = [4, 4, 5, 3, 4]
assert sum(score / 5 * weight for score, weight in zip(x, weights)) == 89
assert sum(score / 5 * weight for score, weight in zip(y, weights)) == 80
assert round(10 / 10 * 100, 1) == 100.0 and round(9 / 9 * 100, 1) == 100.0
assert "7,000" in all_text and "94.2%" in all_text and "89" in all_text and "80" in all_text
pdf = PdfReader(str(pdf_path))
page_images = sorted(render_dir.glob("page-*.png"))
assert len(pdf.pages) == 16 and len(page_images) == 16
assert all((p.extract_text() or "").strip() for p in pdf.pages), "At least one rendered PDF page has no extractable text"
# The built-in Title style and cover paragraph must not carry a bottom border.
assert b"w:styleId=\"Title\"" in styles_xml
assert b"<w:pBdr" not in styles_xml[styles_xml.find(b"w:styleId=\"Title\""):styles_xml.find(b"</w:style>", styles_xml.find(b"w:styleId=\"Title\""))]

receipt = {
  "project_id": "TTQS_ONE", "doc_id": "0048", "requirement_id": "I10-E03",
  "official_checkbox_sequence": "049", "official_source_box_id": "P03-B10",
  "candidate_path": "candidate/0048_R01.docx", "candidate_sha256": sha,
  "candidate_bytes": docx_path.stat().st_size,
  "source_map_sha256": hashlib.sha256(source_map_path.read_bytes()).hexdigest(),
  "blueprint_row_sha256": source_map["blueprint_row_sha256"],
  "synthetic_register_sha256": hashlib.sha256(register_path.read_bytes()).hexdigest(),
  "synthetic_register_rows": len(register), "synthetic_fact_ids_present_in_candidate": len(register),
  "docx_zip_integrity": "PASS", "docx_package_parts": docx_member_count,
  "paragraph_count": len(doc.paragraphs), "table_count": len(doc.tables),
  "pdf_sha256": hashlib.sha256(pdf_path.read_bytes()).hexdigest(),
  "rendered_pdf_pages": len(pdf.pages), "page_images": len(page_images),
  "direct_visual_review": "PASS_ALL_16_PAGES; no blank page, clipping, overlap, orphan row, unreadable table, or title rule observed",
  "arithmetic": {"procurement_difference_ntd": 7000, "budget_utilization_pct_rounded_1dp": 94.2, "weight_total_pct": 100, "candidate_x_weighted_score": 89, "candidate_y_weighted_score": 80, "case_a_completeness_pct": 100, "case_b_completeness_pct": 100},
  "sample_boundary": "PASS; all 39 register rows SAMPLE/SYNTHETIC; no real association decision or third-party original claimed",
  "formal_acceptance": "NOT_GRANTED", "status": "TECHNICAL_AND_DIRECT_VISUAL_PASS_PENDING_INDEPENDENT_SEMANTIC_REVIEW"
}
(evidence / "TECHNICAL_VISUAL_QA_R01.json").write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

# Full text and tables supplied to an independent reviewer without relying on Builder claims.
parts = [
  "# Independent exact-SHA review packet — TTQS ONE DOCX 0048 R01",
  f"Candidate SHA-256: {sha}",
  f"Candidate bytes: {docx_path.stat().st_size}",
  "Evidence status: CANDIDATE_SAMPLE_ONLY; formal association acceptance NOT GRANTED.",
  "## Official source and exact requirement",
  f"Official source PDF SHA-256: {source_map['official_source_pdf_sha256']}",
  f"Official source locator: page {source_map['official_source_page']}, {source_map['official_source_box_id']}",
  f"Official requirement: {source_map['official_requirement_text']}",
  "## Unique Blueprint row",
  json.dumps(blueprint, ensure_ascii=False, indent=2),
  "## Adjacent requirement distinction",
  "DOCX 0045 is I10-E01 (procedure/SOP); DOCX 0047 is I10-E02 (historical review); DOCX 0048 is I10-E03 (specific meeting/review approval record); DOCX 0049 is I11-E01. Do not confuse DOCX IDs with official checkbox sequence 049.",
  "## Review instructions",
  "Independently judge the candidate itself using the exact SHA above. Return PASS, PARTIAL, or FAIL and evidence locators (section/table/page when available). Check requirement fit, correct genre, direct usability, completeness of both procurement meeting and post-selection instructor review record, decision authority and evidence boundaries, all SAMPLE/SYNTHETIC claims, math, cross-document separation, return/hold conditions, title-blind purpose, and what an association administrator or evaluator can actually do with the form. Recompute the arithmetic. Identify material defects only; do not request a rebuild for metadata or control-only differences. Do not treat any synthetic case, role, vote, threshold, approval, qualification, or record as real. State formal acceptance NOT GRANTED. If any criterion is PARTIAL or FAIL, give precise section/table/page and repairable action. The review must be independent of the builder and bind its verdict to exact SHA.",
  "## Candidate content — paragraphs and tables in document order",
  "\n\n".join(paragraph_text),
]
for i, table in enumerate(table_text, 1):
    parts.append(f"## Table {i}\n" + "\n".join(" | ".join(row) for row in table))
parts.extend(["## SAMPLE/SYNTHETIC replacement register (all rows)", "\n".join(" | ".join(r.values()) for r in register)])
packet = review / "independent_review_packet_R01.md"
review.mkdir(parents=True, exist_ok=True)
packet.write_text("\n\n".join(parts) + "\n", encoding="utf-8")
print(json.dumps(receipt, ensure_ascii=False, indent=2))
print(f"REVIEW_PACKET={packet}")
print(f"REVIEW_PACKET_SHA256={hashlib.sha256(packet.read_bytes()).hexdigest()}")
