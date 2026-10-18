from pathlib import Path
import re,json,hashlib
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parents[1]
source=root/'candidate/0073_R04_body.md'
out=root/'candidate/0073_R08.docx'
lines=source.read_text(encoding='utf-8').splitlines()
doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.78);sec.bottom_margin=Inches(.72);sec.left_margin=Inches(.76);sec.right_margin=Inches(.76)
sec.header_distance=Inches(.34);sec.footer_distance=Inches(.32)
NAVY='1F4E78';TEAL='2C6E7F';PALE='EAF1F7';LIGHT='F5F8FB';RED='9C2D25';REDPALE='FCE9E7';GRAY='4A5560'
def font_run(run,size=10.5,bold=None,color=None,italic=None,name='Microsoft JhengHei'):
 run.font.name=name;run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),name)
 run.font.size=Pt(size)
 if bold is not None:run.bold=bold
 if italic is not None:run.italic=italic
 if color:run.font.color.rgb=RGBColor.from_string(color)
def set_cell_shading(cell,fill):
 tcPr=cell._tc.get_or_add_tcPr();shd=tcPr.find(qn('w:shd'))
 if shd is None:shd=OxmlElement('w:shd');tcPr.append(shd)
 shd.set(qn('w:fill'),fill)
def cell_margins(cell,top=65,start=80,bottom=65,end=80):
 tc=cell._tc;tcPr=tc.get_or_add_tcPr();mar=tcPr.first_child_found_in('w:tcMar')
 if mar is None:mar=OxmlElement('w:tcMar');tcPr.append(mar)
 for tag,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
  el=mar.find(qn('w:'+tag))
  if el is None:el=OxmlElement('w:'+tag);mar.append(el)
  el.set(qn('w:w'),str(val));el.set(qn('w:type'),'dxa')
def repeat_header(row):
 trPr=row._tr.get_or_add_trPr();el=OxmlElement('w:tblHeader');el.set(qn('w:val'),'true');trPr.append(el)
def no_split(row):
 trPr=row._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
def shade_paragraph(p,fill):
 pPr=p._p.get_or_add_pPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),fill);pPr.append(shd)
def add_page_field(p):
 r=p.add_run();font_run(r,8,color=GRAY)
 begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin')
 instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text=' PAGE '
 sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate')
 txt=OxmlElement('w:t');txt.text='1'
 end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end')
 for el in [begin,instr,sep,txt,end]:r._r.append(el)
def inline(p,s,size=10.5,base_color=None):
 pat=re.compile(r'(\*\*.*?\*\*|`.*?`|\*.*?\*)')
 pos=0
 for m in pat.finditer(s):
  if m.start()>pos:font_run(p.add_run(s[pos:m.start()]),size,color=base_color)
  token=m.group(0)
  if token.startswith('**'):font_run(p.add_run(token[2:-2]),size,bold=True,color=base_color)
  elif token.startswith('`'):font_run(p.add_run(token[1:-1]),size,color='6A3D9A',name='Consolas')
  else:font_run(p.add_run(token[1:-1]),size,italic=True,color=base_color)
  pos=m.end()
 if pos<len(s):font_run(p.add_run(s[pos:]),size,color=base_color)
def md_cells(line):return [x.strip() for x in line.strip().strip('|').split('|')]
def make_table(rows):
 n=len(rows[0]);t=doc.add_table(rows=len(rows),cols=n);t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=True
 locator_table=n==3 and [str(x).strip() for x in rows[0]]==['證據','位置','狀態']
 replacement_table=n==4 and [str(x).strip() for x in rows[0]]==['資料類型','內容','文件位置','真實替換來源（示例）']
 stakeholder_detail_table=n==4 and [str(x).strip() for x in rows[0]]==['紀錄／群組／n','回饋內容與證據基礎（合成）','分析','限制說明']
 stakeholder_followup_table=n==5 and [str(x).strip() for x in rows[0]]==['record_id','建議回應（提案）與狀態','責任角色（建議，未指定人員）','時機／觸發','追蹤量尺（提案，非官方門檻）']
 total=6.75
 lens=[]
 for c in range(n):
  mx=max((len(str(row[c])) for row in rows if c<len(row)),default=4)
  lens.append(max(4,min(mx,34)))
 # favor table question/description columns while keeping narrow number columns
 weights=[float(x) for x in lens];sw=sum(weights)
 widths=[Inches(total*w/sw) for w in weights]
 if locator_table:
  widths=[Inches(1.15),Inches(4.05),Inches(1.50)]
  t.autofit=False
  for ci,width in enumerate(widths):
   t.columns[ci].width=width
   for row_obj in t.rows: row_obj.cells[ci].width=width
 if replacement_table:
  widths=[Inches(1.25),Inches(2.10),Inches(1.45),Inches(1.90)]
  t.autofit=False
  for ci,width in enumerate(widths):
   t.columns[ci].width=width
   for row_obj in t.rows: row_obj.cells[ci].width=width
 if stakeholder_detail_table:
  widths=[Inches(1.05),Inches(2.30),Inches(1.90),Inches(1.45)]
  t.autofit=False
  for ci,width in enumerate(widths):
   t.columns[ci].width=width
   for row_obj in t.rows: row_obj.cells[ci].width=width
 if stakeholder_followup_table:
  widths=[Inches(.75),Inches(1.70),Inches(1.30),Inches(1.30),Inches(1.65)]
  t.autofit=False
  for ci,width in enumerate(widths):
   t.columns[ci].width=width
   for row_obj in t.rows: row_obj.cells[ci].width=width
 for ri,row in enumerate(rows):
  no_split(t.rows[ri])
  for ci,val in enumerate(row):
   cell=t.cell(ri,ci);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cell_margins(cell)
   try:cell.width=widths[ci]
   except Exception:pass
   p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.0
   inline(p,str(val),size=9.5,base_color='FFFFFF' if ri==0 else None)
   if ri==0:
    set_cell_shading(cell,NAVY)
    for run in p.runs:run.bold=True;run.font.color.rgb=RGBColor.from_string('FFFFFF')
   elif ri%2==0:set_cell_shading(cell,LIGHT)
 if t.rows:repeat_header(t.rows[0])
 doc.add_paragraph().paragraph_format.space_after=Pt(2)
 return t
# Styles
normal=doc.styles['Normal'];normal.font.name='Microsoft JhengHei';normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei');normal.font.size=Pt(10.5);normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1.08
for name,size,color in [('Title',21,'000000'),('Heading 1',16,NAVY),('Heading 2',12.5,TEAL),('Heading 3',11.2,NAVY)]:
 st=doc.styles[name];st.font.name='Microsoft JhengHei';st._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei');st.font.size=Pt(size);st.font.bold=(name!='Title');st.font.color.rgb=RGBColor.from_string(color);st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(10 if name!='Title' else 0);st.paragraph_format.space_after=Pt(5)
# The built-in Title style carries a paragraph bottom border in this Word profile.
title_ppr=doc.styles['Title']._element.get_or_add_pPr()
title_border=title_ppr.find(qn('w:pBdr'))
if title_border is not None:title_ppr.remove(title_border)
# Running header and footer
h=sec.header.paragraphs[0];h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
inline(h,'SAMPLE/SYNTHETIC｜非本會實際紀錄　•　0073／I17A-E01／080',size=8,base_color=RED)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
inline(f,'REAL=UNKNOWN/HOLD　•　候選／未正式接受　•　頁 ',size=8,base_color=GRAY);add_page_field(f)
# Main content parser
idx=0;first_banner=True
while idx<len(lines):
 line=lines[idx].rstrip()
 if not line.strip() or line.strip()=='---':idx+=1;continue
 if line.strip().startswith('|'):
  block=[]
  while idx<len(lines) and lines[idx].strip().startswith('|'):
   row=md_cells(lines[idx]);
   if not all(not x or set(x)<=set('-: ') for x in row):block.append(row)
   idx+=1
  if block:make_table(block)
  continue
 if line.startswith('# '):
  p=doc.add_paragraph(style='Title')
  ppr=p._p.get_or_add_pPr();direct_border=ppr.find(qn('w:pBdr'))
  if direct_border is not None:ppr.remove(direct_border)
  inline(p,line[2:].strip(),size=21)
  idx+=1;continue
 if line.startswith('## '):
  p=doc.add_paragraph(style='Heading 1');inline(p,line[3:].strip(),size=16,base_color=NAVY)
  idx+=1;continue
 if line.startswith('### '):
  p=doc.add_paragraph(style='Heading 2');inline(p,line[4:].strip(),size=12.5,base_color=TEAL)
  idx+=1;continue
 if line.startswith('#### '):
  p=doc.add_paragraph(style='Heading 3');inline(p,line[5:].strip(),size=11.2,base_color=NAVY)
  idx+=1;continue
 if line.startswith('> '):
  p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.12);p.paragraph_format.right_indent=Inches(.08);p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(6)
  is_banner=first_banner and ('SAMPLE/SYNTHETIC' in line)
  inline(p,line[2:].strip(),size=10.5 if is_banner else 10.2,base_color=RED if is_banner else GRAY)
  if is_banner:
   shade_paragraph(p,REDPALE)
   for run in p.runs:run.bold=True
   first_banner=False
  idx+=1;continue
 if re.match(r'^\s*[-*] ',line):
  p=doc.add_paragraph(style='List Bullet');inline(p,re.sub(r'^\s*[-*] ','',line),size=10.5);p.paragraph_format.space_after=Pt(3);idx+=1;continue
 if re.match(r'^\s*\d+[.)] ',line):
  match=re.match(r'^\s*(\d+)[.)] (.*)$',line)
  p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.28);p.paragraph_format.first_line_indent=Inches(-.28);p.paragraph_format.space_after=Pt(3)
  inline(p,f'{match.group(1)}. {match.group(2)}',size=10.5);idx+=1;continue
 p=doc.add_paragraph();p.paragraph_format.widow_control=True;inline(p,line,size=10.5)
 idx+=1
# Document metadata
doc.core_properties.title='DOCX 0073 課後滿意度與多方回饋分析 R08'
doc.core_properties.subject='I17A-E01 / official checkbox 080; SAMPLE/SYNTHETIC candidate'
doc.core_properties.author='TTQS ONE candidate production'
doc.core_properties.keywords='TTQS;0073;I17A-E01;080;SAMPLE/SYNTHETIC'
doc.save(out)
receipt={'source_body_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'docx_sha256':hashlib.sha256(out.read_bytes()).hexdigest(),'docx_bytes':out.stat().st_size,'markdown_tables':sum(1 for x in lines if x.strip().startswith('|')),'status':'BUILT_CANDIDATE_R08; section 5.2 record-level evidence and follow-up fields repaired; technical/layout review pending'}
(root/'evidence'/'build_receipt_R08.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(receipt,ensure_ascii=True,indent=2))
