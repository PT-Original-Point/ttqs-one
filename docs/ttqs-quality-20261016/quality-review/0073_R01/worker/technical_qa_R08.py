from pathlib import Path
import csv, hashlib, json, zipfile, re
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
ROOT=Path(r'docs/ttqs-quality-20261016/quality-review/0073_R01')
DOCX=ROOT/'candidate/0073_R08.docx'
PDF=ROOT/'evidence/render_R08/0073_R08_word_export.pdf'
EVID=ROOT/'evidence'
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def blocks(doc):
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P): yield Paragraph(child, doc)
        elif isinstance(child, CT_Tbl): yield Table(child, doc)
def extract_docx(path):
    d=Document(path); out=[]
    for b in blocks(d):
        if isinstance(b, Paragraph):
            s=b.text.strip()
            if s: out.append(s)
        else:
            for row in b.rows:
                out.append(' | '.join(c.text.replace('\n',' / ').strip().replace('|','\\|') for c in row.cells))
            out.append('')
    return '\n'.join(out)+'\n', d
text, doc=extract_docx(DOCX)
(EVID/'docx_extracted_R08.md').write_text(text,encoding='utf-8')
# Official source byte identity from existing official-source copy.
official=Path(r'docs/ttqs-quality-20261016/quality-review/0072_R01/inputs/official_source.pdf')
blueprint=ROOT/'inputs/UNIQUE_BLUEPRINT_0073_R01.md'
sourcepdf=Path(r'docs/ttqs-quality-20261016/quality-review/0068_R06/sources/official_source.pdf')
# Inputs and data recomputation.
def readcsv(name):
    with (ROOT/'inputs'/name).open(encoding='utf-8-sig',newline='') as f: return list(csv.DictReader(f))
scores=readcsv('learner_scores_R01.csv')
forms=readcsv('returned_forms_R01.csv')
summary=json.loads((ROOT/'inputs/calculation_summary_R01.json').read_text(encoding='utf-8'))
valid=[r for r in scores if r['respondent_id']!='S16']
items=[]
for q in [f'Q{i}' for i in range(1,9)]:
    vals=[int(r[q]) for r in valid]
    pos=sum(v>=4 for v in vals)
    item={'item_id':q,'n':len(vals),'sum':sum(vals),'mean':round(sum(vals)/len(vals)+1e-12,2),'positive_n':pos,'positive_pct':round(pos/len(vals)*100,2)}
    ref=next(x for x in summary['question_summary'] if x['item_id']==q)
    item['matches_summary']=(item['sum']==ref['score_sum'] and item['mean']==float(ref['mean']) and item['positive_n']==ref['positive_n'] and item['positive_pct']==float(ref['positive_rate'].rstrip('%')))
    items.append(item)
allvals=[int(r[f'Q{i}']) for r in valid for i in range(1,9)]
returned=len(forms); valid_n=sum(1 for r in forms if r.get('response_status')=='VALID')
register=readcsv('synthetic_replacement_register_R01.csv')
with zipfile.ZipFile(DOCX) as z: zip_bad=z.testzip()
pngs=sorted((ROOT/'evidence/render_R08').glob('page-*.png'))
try:
    from pypdf import PdfReader
    pdf_pages=len(PdfReader(str(PDF)).pages)
except Exception:
    pdf_pages=None
required=['0073','I17A-E01','080','SYN-0073-2026-A01','REAL=UNKNOWN/HOLD','5.2.1','5.2.2','T01','H01','M01','M02','證據基礎','限制說明','責任角色','時機','追蹤量尺','SAMPLE/SYNTHETIC']
from pypdf import PdfReader
official_phrase='課後滿意度調查分析與課後檢討講師、學員、人資、部門主管等利益關係人的建議回饋紀錄。'
compact=lambda v:re.sub(r'\s+','',v)
source_phrase=compact(PdfReader(str(official)).pages[4].extract_text())
candidate_phrase=compact(text)
checks={
 'official_source_sha_matches_expected':sha(official)=='805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709' and sha(sourcepdf)=='805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709',
 'source_mapping_visible':all(x in text for x in required[:3]),
 'official_requirement_text_present':compact(official_phrase) in candidate_phrase and compact(official_phrase) in source_phrase,
 'sample_case_and_real_hold_visible':all(x in text for x in required[3:5]),
 'record_level_trace_fields_and_ids_present':all(x in text for x in required[5:]),
 'valid_rows_15':len(valid)==15,
 'returned_16_valid_15':returned==16 and valid_n==15,
 'item_summaries_recompute':all(x['matches_summary'] for x in items),
 'overall_120_479_90_3_99_75pct':len(allvals)==120 and sum(allvals)==479 and sum(v>=4 for v in allvals)==90 and round(sum(allvals)/len(allvals),2)==3.99 and round(sum(v>=4 for v in allvals)/len(allvals)*100,2)==75.00,
 'exclusion_s16_missing_q4_q6_q7_no_imputation':'EXCLUDED' in next((r.get('response_status','') for r in forms if r.get('respondent_id')=='S16'),''),
 'replacement_register_264_rows':len(register)==264,
 'stakeholder_groups_kept_separate':set(r['record_id'] for r in readcsv('stakeholder_feedback_R01.csv'))=={'T01','H01','M01','M02'},
 'docx_zip_integrity':zip_bad is None,
 'word_pdf_render_20_pages_20_pngs':pdf_pages==20 and len(pngs)==20,
}
receipt={
 'document_id':'0073','requirement_id':'I17A-E01','official_checkbox_serial':'080','candidate':'0073_R08.docx','docx_sha256':sha(DOCX),'docx_bytes':DOCX.stat().st_size,
 'source_body_sha256':sha(ROOT/'candidate/0073_R04_body.md'),'extracted_text_sha256':sha(EVID/'docx_extracted_R08.md'),
 'official_blueprint_git_blob':'8c2bbe7ca9dcbb03088a48934b9e57e7147ede63','official_source_pdf_page_5_phrase':official_phrase,
 'source_hashes':{'official_source_pdf':sha(official),'blueprint_file':sha(blueprint),'calculation_summary_R01.json':sha(ROOT/'inputs/calculation_summary_R01.json'),'learner_scores_R01.csv':sha(ROOT/'inputs/learner_scores_R01.csv'),'returned_forms_R01.csv':sha(ROOT/'inputs/returned_forms_R01.csv'),'question_summary_R01.csv':sha(ROOT/'inputs/question_summary_R01.csv'),'stakeholder_feedback_R01.csv':sha(ROOT/'inputs/stakeholder_feedback_R01.csv'),'synthetic_replacement_register_R01.csv':sha(ROOT/'inputs/synthetic_replacement_register_R01.csv')},
 'recomputed':{'valid_learner_rows':len(valid),'returned':returned,'valid':valid_n,'valid_item_responses':len(allvals),'score_sum':sum(allvals),'overall_mean_unrounded':sum(allvals)/len(allvals),'overall_mean_rounded':round(sum(allvals)/len(allvals),2),'positive_n':sum(v>=4 for v in allvals),'positive_rate_pct':round(sum(v>=4 for v in allvals)/len(allvals)*100,2),'item_summary':items,'replacement_register_rows':len(register)},
 'structure':{'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'table_rows':sum(len(t.rows) for t in doc.tables),'word_pdf_sha256':sha(PDF),'word_pdf_pages':pdf_pages,'page_png_count':len(pngs)},
 'checks':checks,'technical_status':'PASS' if all(checks.values()) else 'FAIL','semantic_status':'PENDING exact-SHA independent review','acceptance':'CANDIDATE ONLY; REAL=UNKNOWN/HOLD'
}
(EVID/'technical_review_R08.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'docx_sha256':receipt['docx_sha256'],'extracted_sha256':receipt['extracted_text_sha256'],'checks':checks,'status':receipt['technical_status'],'pages':pdf_pages,'paragraphs':len(doc.paragraphs),'tables':len(doc.tables)},ensure_ascii=False,indent=2))



