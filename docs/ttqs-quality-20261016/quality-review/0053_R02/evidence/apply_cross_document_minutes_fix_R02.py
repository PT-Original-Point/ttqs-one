from pathlib import Path
from docx import Document
import csv, hashlib, json
from datetime import datetime, timezone

root=Path(r"E:\TTQS\_TTQS_ONE_REVIEW_SYNC_0036_R08_20261016")
r1=root/"docs/ttqs-quality-20261016/quality-review/0053_R01"
r2=root/"docs/ttqs-quality-20261016/quality-review/0053_R02"
runtime=root/"docs/ttqs-quality-20261016/quality-review"
doc_path=r2/"candidate/0053_R02.docx"
reg_path=r2/"SYNTHETIC_REPLACEMENT_REGISTER_0053_R02.csv"
src0052=root/"docs/ttqs-quality-20261016/quality-review/0052_R01/candidate/0052_R01.docx"
expected_sha_r1="14a7b825d24059f2c339f95eabfa613d4cad43dea9534fa80fcffa8a2489f4f8"
expected_sha_0052="406d1ef2c6e9b6266ea1adaeadedc5dfac77964fe76f73c2471168220ee6d455"
sha_before=hashlib.sha256(doc_path.read_bytes()).hexdigest()
assert sha_before==expected_sha_r1, f"Unexpected starting candidate SHA: {sha_before}"
sha_0052=hashlib.sha256(src0052.read_bytes()).hexdigest()
assert sha_0052==expected_sha_0052, f"Unexpected 0052 source SHA: {sha_0052}"

doc=Document(str(doc_path))
changes=[]
def set_para(p, old, new, locator):
    assert p.text==old, f"Text mismatch at {locator}: {p.text!r}"
    nonempty=[r for r in p.runs if r.text]
    assert len(nonempty)==1, f"Expected one text run at {locator}, got {len(nonempty)}"
    nonempty[0].text=new
    changes.append({"locator":locator,"before":old,"after":new})
def set_cell(table_no,row_no,col_no,old,new,locator):
    cell=doc.tables[table_no-1].rows[row_no-1].cells[col_no-1]
    assert cell.text==old, f"Text mismatch at {locator}: {cell.text!r}"
    ps=[p for p in cell.paragraphs if p.text]
    assert len(ps)==1, f"Expected one populated paragraph at {locator}, got {len(ps)}"
    set_para(ps[0],old,new,locator)

set_cell(11,8,3,"分鐘比例。","分鐘配置；合計 90 分鐘；練習占比以實際分鐘除以 90 計算。","p5 / section 04 / data dictionary / 方法分段")
set_cell(11,8,4,"90 分鐘不變；35/20/35 改提 20/20/50，來自 0052 R01 SAMPLE 計畫。","90 分鐘不變；基準 35/20/35 分鐘改提 20/20/50 分鐘（練習增加 15 分鐘），來自 0052 R01 SAMPLE 計畫。","p5 / section 04 / data dictionary / 方法分段")
set_cell(13,6,3,"合成比較基準 G1 為 35/20/35（說明/示範/引導練習），操作練習比例未達本示範設定的 50% 最低值。","合成基準 G1 為 35/20/35 分鐘（說明/示範/練習；合計 90 分鐘）；練習占 35/90=38.9%，低於本例 50% 合成門檻。","p7 / section 06 / P2 baseline")
set_cell(13,6,5,"採方案 A 的 20/20/50；G1 仍 90 分鐘，練習 45 分鐘；未達題目者依同一錨點補練。","方案 A 為 20/20/50 分鐘（合計 90）；練習占 50/90=55.6%，達本例門檻；未達題目者依同一錨點補練。","p7 / section 06 / P2 recheck")
set_cell(14,4,4,"90×20%=18 說明、90×20%=18 示範、90×50%=45 練習；三段合計 81 分鐘，餘 9 分鐘為轉場/回饋安排，須在 REAL 教案確認。","基準：35+20+35=90 分鐘；方案 A：20+20+50=90 分鐘。練習由 35 增至 50 分（+15）；35/90=38.9%→50/90=55.6%（+16.7 個百分點）。","p9 / section 07 / P2 calculation")
reminder=doc.tables[14].rows[0].cells[0]
assert len(reminder.paragraphs)==2 and reminder.paragraphs[0].text=="數字核對提醒"
set_para(reminder.paragraphs[1],"0052 R01 的 20/20/50 是對原 90 分鐘 G1 的分段比例。三段教學活動為 18+18+45=81 分鐘；尚有 9 分鐘屬轉場/回饋的待確認安排，不能把比例表誤寫成三段已合計 90 分鐘。實際核定分鐘仍 UNKNOWN。","0052 R01 的方案 A 明列 G1 為 20 分說明＋20 分示範＋50 分引導演練，合計 90 分鐘；基準 35＋20＋35 亦合計 90 分鐘。兩組數字是分鐘，不再乘以 90 換算。練習增加 15 分鐘，占比由 35/90=38.9% 升至 50/90=55.6%，達本例 50% 合成門檻。均為 SAMPLE；REAL 核定與授課分鐘仍 UNKNOWN。","p9 / section 07 / numeric audit note")
set_cell(20,4,2,"G1 35/20/35 改提 20/20/50；分段重算 81 分，另 9 分仍須交代。","G1 基準 35/20/35 分鐘改提 20/20/50 分鐘；兩者各合計 90，練習增加 15 分鐘（38.9%→55.6%）。","p12 / section 10 / P2 finding")
set_cell(23,5,2,"P2 的 20/20/50 改善為 90 分鐘內提案，仍有 9 分鐘安排須解釋。","P2 以 20/20/50 分鐘配置於 90 分鐘內，練習增加 15 分鐘且占 55.6%；REAL 實際分鐘待原件核對。","p13 / section 11 / action 4")
set_cell(23,5,3,"核定版明確說明轉場/回饋時間；按完成評量逐項記實際分鐘。","對核定課表逐段核對分鐘合計與評量記錄；按實際授課及補練記錄逐項留存。","p13 / section 11 / action 4 success check")
set_cell(24,4,4,"20/20/50 重算 90 分鐘，9 分鐘轉場回饋待 REAL 教案說明。","方案 A 的 20+20+50=90 分鐘；練習較基準增加 15 分鐘。REAL 核定版及授課記錄待取得，不推定已採用。","p14 / section 12 / action A-03")

# The body summary passage on page 3 states the split but not the false percentages; add its unit for traceable cross-file consistency.
old_summary="SYN-0033-PLAN-01 是跨文件共用的合成教學案例，課名為「志工服務紀錄、交接與工作界線實務」。0052 R01 的四單元為 G1 90 分、G2 90 分、G3 120 分、G4 60 分，合計 360 教學分鐘；方案 A 將 G1 分段由合成比較基準 35/20/35 改為 20/20/50。這些數字是計畫草案示例，不代表實際招生、教學或核准。"
new_summary="SYN-0033-PLAN-01 是跨文件共用的合成教學案例，課名為「志工服務紀錄、交接與工作界線實務」。0052 R01 的四單元為 G1 90 分、G2 90 分、G3 120 分、G4 60 分，合計 360 教學分鐘；方案 A 將 G1 分段由基準 35/20/35 分鐘改為 20/20/50 分鐘，兩組各合計 90 分鐘。這些數字是計畫草案示例，不代表實際招生、教學或核准。"
set_para(doc.paragraphs[18],old_summary,new_summary,"p3 / section 02 / sample case source summary")

doc.save(str(doc_path))
sha_after=hashlib.sha256(doc_path.read_bytes()).hexdigest()
assert sha_after!=sha_before

# Rebind all facts to the new exact SHA and correct the two timing facts and the P2 baseline register entry.
with reg_path.open(encoding="utf-8-sig",newline="") as f: rows=list(csv.DictReader(f))
assert len(rows)==38 and len({r['fact_id'] for r in rows})==38
updates={
"SYN-0053-F011": {
"value_in_candidate":"P2｜合成基準 G1=35/20/35 分鐘（合計90）；練習35/90=38.9%，低於本例50%合成門檻。",
"source_or_formula":"DOCX 0052 R01 (I11-E04, exact SHA 406d1ef2c6e9b6266ea1adaeadedc5dfac77964fe76f73c2471168220ee6d455) G1 90-minute split; synthetic threshold 50%; 35/90=38.9%.",
"body_locator":"五面向檢核表；P2列，基線欄",
"rendered_page":"7"
},
"SYN-0053-F032": {
"value_in_candidate":"G1 基準時間配置：35+20+35=90 分鐘；引導練習占35/90=38.9%。",
"source_or_formula":"SAMPLE; source DOCX 0052 R01 table 20 row 2 / table 22 row 2; sum 35+20+35=90; 35/90=38.9%.",
"body_locator":"§08 G1時間配置逐步重算；§06 P2基線",
"rendered_page":"10"
},
"SYN-0053-F033": {
"value_in_candidate":"G1 方案A時間配置：20+20+50=90 分鐘；引導練習占50/90=55.6%，較基準增加15分鐘、+16.7個百分點。",
"source_or_formula":"SAMPLE; source DOCX 0052 R01 table 22 row 2; sum 20+20+50=90; 50/90=55.6%; difference vs 35/90=38.9% is +16.7pp.",
"body_locator":"§08 G1時間配置逐步重算；§06 P2複核；§07 P2計算",
"rendered_page":"10"
}}
for row in rows:
    row['candidate_sha256']=sha_after
    if row['fact_id'] in updates: row.update(updates[row['fact_id']])
assert set(updates).issubset({r['fact_id'] for r in rows})
with reg_path.open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.DictWriter(f,fieldnames=rows[0].keys(),lineterminator='\n'); w.writeheader(); w.writerows(rows)

# Extract the authoritative neighboring file statements as a pinned source supplement.
src=Document(str(src0052))
source_rows=[]
for table_no,row_no in [(20,2),(22,2)]:
    row=src.tables[table_no-1].rows[row_no-1]
    source_rows.append(f"0052 table {table_no}, row {row_no}: " + " || ".join(c.text.replace("\n"," / ") for c in row.cells))
source_note=(f"# Cross-document source check: 0052 R01\n\n"
 f"Source file: `0052_R01.docx` (DOCX 0052 / I11-E04)\n"
 f"Exact SHA-256: `{sha_0052}`\n"
 f"Case: `SYN-0033-PLAN-01`\n\n"
 "The source file explicitly reports the following G1 entries (verbatim table extraction):\n\n"+
 "\n\n".join(f"- {x}" for x in source_rows)+
 "\n\n## Correct arithmetic for 0053\n\n"
 "The 0052 source labels the proposed G1 parts as minutes: 20 + 20 + 50 = 90 minutes. The comparison baseline is 35 + 20 + 35 = 90 minutes. The 0053 R01 interpretation of these as percentages and its derived 18 + 18 + 45 = 81 minutes / 9-minute remainder were incorrect. Correct sample calculations: baseline guided practice 35/90 = 38.9%; proposed guided practice 50/90 = 55.6%; increase = 15 minutes and 16.7 percentage points. All are SAMPLE/SYNTHETIC; actual adopted course and delivery remain UNKNOWN.\n")
(r2/"inputs/CROSS_DOCUMENT_SOURCE_0052_R01.md").write_text(source_note,encoding="utf-8")

# Make the source-map revision explicit and preserve all original references.
sm=(r1/"inputs/SOURCE_MAP_0053_R01.md").read_text(encoding="utf-8")
sm=sm.replace("# 0053 R01 來源與範圍對照","# 0053 R02 來源與範圍對照")
sm += "\n\n## R02 跨文件分鐘數校正\n\n0052 R01 exact SHA `"+sha_0052+"` 的表 22 第 2 列明列 G1：20 分重點說明＋20 分示範＋50 分引導演練；合計 90 分鐘。0053 R01 曾誤把 20/20/50 當百分比並推算 81 分鐘及 9 分鐘餘數。R02 已移除此錯誤，改按分鐘重算：基準 35+20+35=90；提案 20+20+50=90；練習占比 38.9%→55.6%。\n"
(r2/"inputs/SOURCE_MAP_0053_R02.md").write_text(sm,encoding="utf-8")

receipt={
"schema":"ttqs.0053.corrective_build_receipt.v1","doc_id":"0053","requirement_id":"I12-E01","official_checkbox_sequence":"056","revision":"R02","document_genre":"analysis_report",
"candidate_path":str(doc_path),"candidate_sha256":sha_after,"candidate_bytes":doc_path.stat().st_size,
"previous_candidate_sha256":sha_before,"trigger":"material cross-document arithmetic defect; 0053 R01 treated the 0052 R01 minute allocation 20+20+50 as percentages, yielding a false 81-minute sum and false 9-minute remainder",
"source_0052_sha256":sha_0052,"correction":"Use explicit minute units, 35+20+35=90 baseline and 20+20+50=90 proposal; recalculate practice shares 38.9% to 55.6% (+16.7pp), remove false 9-minute gap, keep SAMPLE/REAL boundary.",
"paragraph_count":len(doc.paragraphs),"table_count":len(doc.tables),"synthetic_register_rows":len(rows),"synthetic_register_unique_fact_ids":len(set(r['fact_id'] for r in rows)),
"register_sha256":hashlib.sha256(reg_path.read_bytes()).hexdigest(),"corrected_locators":[c['locator'] for c in changes],"corrected_text_items":len(changes),"built_at_utc":datetime.now(timezone.utc).isoformat(),"independent_review":"PENDING; R02 exact SHA must be reviewed"
}
(r2/"evidence/build_receipt_R02.json").write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
(r2/"evidence/correction_log_R02.md").write_text("# 0053 R02 targeted correction log\n\n"
 f"Previous candidate SHA: `{sha_before}`\n\nNew candidate SHA: `{sha_after}`\n\n"
 f"Authoritative related source: DOCX 0052 R01 SHA `{sha_0052}`.\n\n"
 "Defect: R01 interpreted source minute allocations as percentages, incorrectly derived 18+18+45=81 and a 9-minute remainder, then repeated that false gap in its data dictionary, five-dimension P2 table, recheck, calculation, findings, recommendation, and tracker.\n\n"
 "Correction: the same-case plan states G1 20+20+50 minutes = 90 minutes. R02 states the baseline 35+20+35=90 minutes, proposal 20+20+50=90 minutes, practice rises 15 minutes, and practice share is 38.9% to 55.6% (+16.7 percentage points). The sample plan remains unapproved/unimplemented; REAL status stays UNKNOWN. No other topic was intentionally rebuilt.\n\n"
 "The old R01 bytes and its review attempt are preserved in `quality-review/0053_R01/`. R02 is a distinct candidate and is not promoted.\n",encoding="utf-8")
print(json.dumps({"old_sha256":sha_before,"new_sha256":sha_after,"source_0052_sha256":sha_0052,"text_edits":len(changes),"register_rows":len(rows),"register_sha256":receipt['register_sha256'],"edited_locations":[c['locator'] for c in changes]},ensure_ascii=False,indent=2))
