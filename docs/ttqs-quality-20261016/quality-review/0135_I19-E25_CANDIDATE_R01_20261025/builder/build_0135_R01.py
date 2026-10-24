from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "candidate" / "0135_I19-E25_BODY_R01.md"
OUT = ROOT / "candidate" / "0135__目標市場與顧客價值具體提升成效分析報告_R01.docx"
BODY_PT = 10.5
TABLE_PT = 9.5
NAVY = "17365D"
TEAL = "1F5D68"
PALE = "EAF1F4"
GRID = "B8C5CC"
FONT = "Microsoft JhengHei"


def set_font(run, name=FONT):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), name)


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shade = tcpr.find(qn("w:shd"))
    if shade is None:
        shade = OxmlElement("w:shd")
        tcpr.append(shade)
    shade.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=75, bottom=80, end=75):
    tcpr = cell._tc.get_or_add_tcPr()
    margins = tcpr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_borders(table):
    tblpr = table._tbl.tblPr
    borders = tblpr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn("w:" + edge))
        if el is None:
            el = OxmlElement("w:" + edge)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), GRID)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def prevent_row_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def rich_text(p, text, base_bold=False, color=None):
    tokens = re.split(r"(\*\*.*?\*\*|`[^`]+`)", text)
    for token in tokens:
        if not token:
            continue
        bold = token.startswith("**") and token.endswith("**")
        code = token.startswith("`") and token.endswith("`")
        value = token[2:-2] if bold else token[1:-1] if code else token
        run = p.add_run(value)
        run.bold = base_bold or bold
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if code:
            run.font.name = "Consolas"
        set_font(run, "Consolas" if code else FONT)


def cells(line):
    return [x.strip() for x in line.strip().strip("|").split("|")]


def is_table_rule(line):
    if not line.strip().startswith("|"):
        return False
    parts = cells(line)
    return bool(parts) and all(re.fullmatch(r":?-{3,}:?", x.replace(" ", "")) for x in parts)


def add_table(doc, rows, body_width):
    if len(rows) < 2:
        return
    ncol = len(rows[0])
    clean = [[row[i].replace("<br>", "\n") if i < len(row) else "" for i in range(ncol)] for row in rows]
    table = doc.add_table(rows=1, cols=ncol)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    maxima = [max([len(row[i]) for row in clean] + [1]) for i in range(ncol)]
    weights = [max(0.75, min(3.4, x ** 0.52)) for x in maxima]
    widths = [body_width * w / sum(weights) for w in weights]
    minw = 0.55 if ncol >= 7 else 0.65 if ncol >= 5 else 0.8
    widths = [max(minw, w) for w in widths]
    total = sum(widths)
    if total > body_width:
        excess = total - body_width
        flexible = [max(0, w-minw) for w in widths]
        denom = sum(flexible) or 1
        widths = [w - excess*f/denom for w, f in zip(widths, flexible)]
    for ri, rowdata in enumerate(clean):
        row = table.rows[0] if ri == 0 else table.add_row()
        prevent_row_split(row)
        if ri == 0:
            repeat_header(row)
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if ri == 0:
                set_cell_shading(cell, TEAL)
            elif ri % 2 == 0:
                set_cell_shading(cell, PALE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            rich_text(p, rowdata[ci], base_bold=(ri == 0), color="FFFFFF" if ri == 0 else None)
            for run in p.runs:
                run.font.size = Pt(TABLE_PT)
                if ri == 0:
                    run.font.color.rgb = RGBColor(255, 255, 255)
    set_borders(table)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.space_before = Pt(0)


def add_page_field(p):
    r = p.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "1"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for e in (begin, instr, sep, t, end): r._r.append(e)
    set_font(r); r.font.size = Pt(8.5)


def add_callout(doc, text, body_width):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(body_width)
    c = table.cell(0,0); c.width = Inches(body_width)
    set_cell_shading(c, "F2F5F7"); set_cell_margins(c, top=130, start=150, bottom=130, end=150)
    p = c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.08
    rich_text(p, text)
    set_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def main():
    text = SOURCE.read_text(encoding="utf-8-sig")
    lines = text.splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.27); sec.page_height = Inches(11.69)
    sec.top_margin = Inches(0.68); sec.bottom_margin = Inches(0.68)
    sec.left_margin = Inches(0.72); sec.right_margin = Inches(0.72)
    body_width = 8.27 - 1.44
    normal = doc.styles["Normal"]
    normal.font.name = FONT; normal.font.size = Pt(BODY_PT); normal.font.color.rgb = RGBColor.from_string("20252A")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(5); normal.paragraph_format.line_spacing = 1.12
    for name, size in (("Title",21),("Heading 1",15),("Heading 2",12.5),("Heading 3",11.5)):
        st=doc.styles[name]; st.font.name=FONT; st.font.size=Pt(size); st.font.bold=(name!="Title")
        st.font.color.rgb=RGBColor.from_string("000000" if name=="Title" else NAVY)
        st._element.rPr.rFonts.set(qn("w:eastAsia"),FONT)
        st.paragraph_format.keep_with_next=True
        st.paragraph_format.space_before=Pt(12 if name=="Heading 1" else 8)
        st.paragraph_format.space_after=Pt(4)
    title_ppr=doc.styles["Title"]._element.get_or_add_pPr()
    title_border=title_ppr.find(qn("w:pBdr"))
    if title_border is not None: title_ppr.remove(title_border)
    header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=header.add_run("TTQS ONE　0135　市場與顧客價值　SAMPLE／SYNTHETIC"); set_font(r); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string("52616B")
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=footer.add_run("候選稿 R01　｜　REAL 原件未提供　｜　第 "); set_font(r); r.font.size=Pt(8.5)
    add_page_field(footer); r=footer.add_run(" 頁"); set_font(r); r.font.size=Pt(8.5)
    i=0
    while i<len(lines):
        raw=lines[i]; line=raw.strip()
        if not line or line=="---": i+=1; continue
        if line.startswith("|") and i+1<len(lines) and is_table_rule(lines[i+1]):
            rows=[cells(line)]; i+=2
            while i<len(lines) and lines[i].strip().startswith("|"):
                rows.append(cells(lines[i])); i+=1
            add_table(doc,rows,body_width); continue
        if line.startswith("> "):
            call=[]
            while i<len(lines) and lines[i].strip().startswith(">"):
                call.append(lines[i].strip()[1:].lstrip()); i+=1
            add_callout(doc," ".join(call),body_width); continue
        if line.startswith("#### "):
            p=doc.add_paragraph(style="Heading 3"); rich_text(p,line[5:]); i+=1; continue
        if line.startswith("### "):
            p=doc.add_paragraph(style="Heading 2"); rich_text(p,line[4:]); i+=1; continue
        if line.startswith("## "):
            p=doc.add_paragraph(style="Heading 1"); rich_text(p,line[3:]); i+=1; continue
        if line.startswith("# "):
            p=doc.add_paragraph(style="Title")
            ppr=p._p.get_or_add_pPr()
            border=ppr.find(qn("w:pBdr"))
            if border is not None: ppr.remove(border)
            rich_text(p,line[2:]); i+=1; continue
        if line.startswith("- "):
            p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(3); rich_text(p,line[2:]); i+=1; continue
        if re.match(r"^\d+\.\s+",line):
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(0.24); p.paragraph_format.first_line_indent=Inches(-0.24)
            p.paragraph_format.space_after=Pt(4); rich_text(p,line); i+=1; continue
        p=doc.add_paragraph(); p.paragraph_format.widow_control=True; rich_text(p,line); i+=1
    doc.core_properties.title="目標市場與顧客價值具體提升成效分析報告｜0135 R01"
    doc.core_properties.subject="I19-E25; SAMPLE/SYNTHETIC candidate; independent review pending"
    doc.core_properties.author="TTQS ONE candidate production"
    doc.save(OUT)
    print(OUT)
    print(f"bytes={OUT.stat().st_size}")

if __name__=="__main__": main()
