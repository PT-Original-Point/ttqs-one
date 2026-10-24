from pathlib import Path
import csv,json,hashlib,re,statistics
from docx import Document
ROOT=Path(__file__).resolve().parents[1]
body=ROOT/'candidate/0135_I19-E25_BODY_R01.md'; data=ROOT/'source/SAMPLE_PAIRED_MEASUREMENT_DATA.csv'; docx=ROOT/'candidate/0135__目標市場與顧客價值具體提升成效分析報告_R01.docx'; reg=ROOT/'candidate/0135_SYNTHETIC_REPLACEMENT_REGISTER_R01.csv'
text=body.read_text(encoding='utf-8'); rows=list(csv.DictReader(data.open(encoding='utf-8',newline=''))); regrows=list(csv.DictReader(reg.open(encoding='utf-8',newline='')))
expected={r['case_id']:(int(r['baseline_minutes_per_monthly_record']),int(r['90_day_followup_minutes_per_monthly_record']),int(r['baseline_on_time_1_yes_0_no']),int(r['followup_on_time_1_yes_0_no'])) for r in rows}
render_rows=[]
for line in text.splitlines():
 if line.startswith('|') and 'SYN-0135-VALUE-2025-' in line:
  c=[x.strip().replace('**','') for x in line.strip().strip('|').split('|')]
  if c[0] in expected: render_rows.append((c[0],tuple(int(c[i]) for i in (1,2,4,5))))
missing=[k for k in expected if k not in dict(render_rows)]
wrong=[k for k,v in render_rows if expected.get(k)!=v]
before=[v[0] for v in expected.values()]; after=[v[1] for v in expected.values()]; pre=sum(v[2] for v in expected.values()); post=sum(v[3] for v in expected.values())
mdflags={'target_market':('目標市場' in text),'customer_value':('顧客價值' in text),'official_checkbox_unknown':('官方查檢表勾選流水號：UNKNOWN' in text),'project_item_140':('第 140 項' in text),'nonresponse_disclosed':('13／28' in text or '13 家' in text),'causality_limited':('不證明訓練是唯一原因' in text),'cross_document_separation':('0133／I19-E22' in text and '0134／I19-E23' in text),'no_fake_archive_id':('TTQS-I19-E25-ARCHIVE' not in text),'no_real_placeholder_id':('ORG-REAL-2025' not in text),'no_fabricated_signoff_record':('簽核紀錄聯' not in text),'no_official_penalty_claim':('0 分或不予計分處置' not in text)}
D=Document(docx); wd='\n'.join([p.text for p in D.paragraphs]+[c.text for t in D.tables for row in t.rows for c in row.cells]); case_missing=[k for k in expected if k not in wd]
font_sizes=[]
for t in D.tables:
 for row in t.rows:
  for cell in row.cells:
   for p in cell.paragraphs:
    for run in p.runs:
     if run.font.size: font_sizes.append(run.font.size.pt)
for p in D.paragraphs:
 for run in p.runs:
  if run.font.size: font_sizes.append(run.font.size.pt)
report={'project_id':'TTQS_ONE','doc_id':'0135','official_requirement_id':'I19-E25','project_157_item':140,'document_type':'quantified target-market and customer-value analysis report','source_commit':'8ff5f8cbe6df80be9ed60577ed324e5a0a18f382','checks':{'sample_source_rows':len(rows),'replacement_register_rows':len(regrows),'markdown_paired_rows':len(render_rows),'missing_rows':missing,'wrong_rows':wrong,'docx_case_ids_missing':case_missing,'market_reach_percent':[20/80*100,24/80*100,28/80*100],'year_over_year_pp':[5,5],'relative_change_percent':[20,round((35-30)/30*100,1)],'baseline_median_minutes':statistics.median(before),'followup_median_minutes':statistics.median(after),'time_reduction_percent':round((statistics.median(before)-statistics.median(after))/statistics.median(before)*100,1),'on_time_before':[pre,len(rows),round(pre/len(rows)*100,1)],'on_time_after':[post,len(rows),round(post/len(rows)*100,1)],'on_time_change_pp':round((post-pre)/len(rows)*100,1),'response_rate_percent':round(len(rows)/28*100,1),'markdown_content_flags':mdflags,'docx_table_count':len(D.tables),'minimum_explicit_font_pt':min(font_sizes) if font_sizes else None,'docx_bytes':docx.stat().st_size,'docx_sha256':hashlib.sha256(docx.read_bytes()).hexdigest(),'body_sha256':hashlib.sha256(body.read_bytes()).hexdigest(),'replacement_register_sha256':hashlib.sha256(reg.read_bytes()).hexdigest()},'boundaries':{'all_outcome_numbers_sample_synthetic':True,'real_sources':'UNKNOWN','independent_exact_sha_review':'PENDING','formal_acceptance':'NOT_GRANTED'},'qa_status':'PASS_WITH_SCOPE' if not missing and not wrong and not case_missing and all(mdflags.values()) and len(rows)==15 and len(regrows)==6 else 'FAIL'}
out=ROOT/'qa/TECHNICAL_QA_0135_R01.json'; out.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n'); print(json.dumps(report,ensure_ascii=False))
