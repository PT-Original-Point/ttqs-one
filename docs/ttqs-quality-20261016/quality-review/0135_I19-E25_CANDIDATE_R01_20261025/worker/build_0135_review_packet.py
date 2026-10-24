from pathlib import Path
from zipfile import ZipFile
from hashlib import sha256
from docx import Document
import sys
root=Path(sys.argv[1]); docx=Path(sys.argv[2]); out=Path(sys.argv[3])
def digest(p): return sha256(p.read_bytes()).hexdigest()
parts=[]
parts.append('# 0135 exact-SHA independent semantic review packet\n')
parts.append(f'DOCX_SHA256={digest(docx)}\nDOCX_BYTES={docx.stat().st_size}\n')
parts.append('Review scope: independent semantic and usability review only. Do not edit files, use tools, browse, use Computer Use, or claim formal acceptance. Return PASS, PARTIAL, or FAIL with paragraph/heading/table locators and concrete evidence.\n')
parts.append('Check exact mapping DOCX 0135 -> official I19-E25 -> project item 140; official printed checkbox serial must remain UNKNOWN; distinguish neighbor items 141/142 and Other. Check report genre, complete usable measures/tools, denominator and nonresponse handling, paired data, arithmetic, limitations, cross-document separation, all SAMPLE/SYNTHETIC and REAL boundaries, false approvals/records, actionable HOLD/REJECT conditions, and whether a training administrator/evaluator can use it. Review the supplied DOCX text and sources; do not rely on this request alone.\n')
for p in sorted((root/'source').iterdir()):
    if p.is_file() and p.suffix.lower() in {'.md','.txt','.csv'}:
        try: body=p.read_text(encoding='utf-8-sig')
        except Exception: continue
        parts.append(f'\n## SOURCE {p.name}\nSHA256={digest(p)}\n{body}\n')
doc=Document(str(docx)); blocks=[]
for p in doc.paragraphs:
    if p.text.strip(): blocks.append(p.text)
for ti,t in enumerate(doc.tables,1):
    blocks.append(f'\n[TABLE {ti}]')
    for row in t.rows:
        blocks.append(' | '.join(c.text.replace('\n',' / ') for c in row.cells))
parts.append('\n## CANDIDATE DOCX EXTRACTED CONTENT\n'+'\n'.join(blocks)+'\n')
out.write_text('\n'.join(parts),encoding='utf-8')
print(f'packet={out}\nbytes={out.stat().st_size}\nsha256={digest(out)}\ndocx_sha256={digest(docx)}\nparagraphs={len(doc.paragraphs)}\ntables={len(doc.tables)}')
