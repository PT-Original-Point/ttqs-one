from __future__ import annotations

import csv
import hashlib
import json
import re
import zipfile
from difflib import SequenceMatcher
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader


BASE = Path(__file__).resolve().parents[1]
DOCX = BASE / "candidate" / "0084_R11.docx"
BODY = BASE / "candidate" / "0084_R10_body.md"
BLUEPRINT = BASE / "inputs" / "UNIQUE_BLUEPRINT_0084_R03.md"
REGISTER = BASE / "evidence" / "synthetic_replacement_register_R10.csv"
PDF = BASE / "evidence" / "render_R11_final" / "0084_R11.pdf"
MANIFEST = BASE / "evidence" / "render_R11_final" / "visual_input_manifest.json"
OUT = BASE / "evidence" / "technical_review_R11.json"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def norm(value: str) -> str:
    return re.sub(r"\s+", "", value)


issues: list[str] = []
for path in (DOCX, BODY, BLUEPRINT, REGISTER, PDF, MANIFEST):
    if not path.is_file():
        issues.append(f"missing:{path.relative_to(BASE)}")

with zipfile.ZipFile(DOCX) as archive:
    bad_member = archive.testzip()
if bad_member:
    issues.append(f"docx_zip_crc:{bad_member}")

doc = Document(str(DOCX))
paragraph_text = [p.text for p in doc.paragraphs]
table_text: list[str] = []
for table in doc.tables:
    for row in table.rows:
        table_text.append("\t".join(cell.text for cell in row.cells))
doc_text = "\n".join(paragraph_text + table_text)
body_text = BODY.read_text(encoding="utf-8")
blueprint_text = BLUEPRINT.read_text(encoding="utf-8")

expected_heading_patterns = [rf"(?m)^# §{number}\s" for number in range(18)]
missing_headings = [p for p in expected_heading_patterns if not re.search(p, body_text)]
if missing_headings:
    issues.append(f"source_missing_sections:{len(missing_headings)}")

for token in ("DOCX 0084", "I17C-E03", "090", "SAMPLE", "SYNTHETIC", "NOT_EXECUTED", "REAL=UNKNOWN/HOLD", "FORMAL_ACCEPTANCE=NOT_GRANTED"):
    if token not in doc_text:
        issues.append(f"docx_missing_token:{token}")
if "SYN-0083-B01" not in doc_text:
    issues.append("missing_cross_document_case_link_SYN-0083-B01")

for phrase in ("不是協會正式表單", "不是已執行評估", "SAMPLE/SYNTHETIC", "NOT_EXECUTED", "REAL=UNKNOWN/HOLD", "FORMAL_ACCEPTANCE=NOT_GRANTED"):
    if phrase not in body_text:
        issues.append(f"missing_sample_real_boundary:{phrase}")

prohibited_claim_tokens = ["已簽核", "已核准", "已核定", "已批准", "已採納"]
prohibited_hits = [token for token in prohibited_claim_tokens if token in doc_text]
if prohibited_hits:
    issues.append("prohibited_approval_claim:" + ",".join(prohibited_hits))

csv_rows: list[list[str]] = []
with REGISTER.open(encoding="utf-8-sig", newline="") as stream:
    reader = csv.reader(stream)
    header = next(reader)
    csv_rows = list(reader)
if len(csv_rows) != 59:
    issues.append(f"csv_register_count:{len(csv_rows)}")
register_ids = [row[0] for row in csv_rows]
expected_ids = [f"REG-{n:03d}" for n in range(1, 60)]
if register_ids != expected_ids:
    issues.append("csv_register_ids_not_contiguous_unique")

register_matches = []
for index, table in enumerate(doc.tables):
    if len(table.rows) != 60 or not table.rows:
        continue
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    if rows[0][0] == "REG_ID" and [row[0] for row in rows[1:]] == expected_ids:
        register_matches.append((index, rows))
if len(register_matches) != 1:
    issues.append(f"docx_register_table_matches:{len(register_matches)}")
else:
    _, word_rows = register_matches[0]
    header_roles_match = (
        word_rows[0][0] == "REG_ID"
        and word_rows[0][1].replace(" ", "") == header[1].replace(" ", "")
        and "SAMPLE" in word_rows[0][2]
        and word_rows[0][3] == header[3]
        and "REAL" in word_rows[0][4]
        and word_rows[0][5] == header[5]
    )
    if not header_roles_match:
        issues.append("docx_register_header_roles_mismatch")
    if word_rows[1:] != csv_rows:
        issues.append("docx_register_rows_csv_mismatch")

state_counts = {
    "NOT_EXECUTED": sum(row[5] == "NOT_EXECUTED" for row in csv_rows),
    "HOLD": sum(row[5] == "HOLD" for row in csv_rows),
    "SAMPLE/SYNTHETIC": sum(row[5] == "SAMPLE/SYNTHETIC" for row in csv_rows),
}
if state_counts != {"NOT_EXECUTED": 53, "HOLD": 5, "SAMPLE/SYNTHETIC": 1}:
    issues.append(f"register_state_counts:{state_counts}")
if sum(state_counts.values()) != 59:
    issues.append("register_state_sum_not_59")
if "53＋5＋1＝59" not in doc_text:
    issues.append("docx_register_math_statement_missing")

case_fields = {}
for point in ("O0", "O1", "O2", "O3"):
    match = re.search(rf"(?ms)^### 6\.5\.\d+ `{point}`.*?(?=^### 6\.5\.\d+ `O[0-3]`|^<!-- SECTION:|\Z)", body_text)
    if not match:
        issues.append(f"case_missing:{point}")
        continue
    chunk = match.group(0)
    fields = [key for key in ("O-08", "O-09", "O-16", "O-17", "O-18") if f"`{key}`" in chunk]
    case_fields[point] = fields
    if len(fields) != 5:
        issues.append(f"case_fields_incomplete:{point}:{fields}")

if "I17C-E03" not in blueprint_text or "`090`" not in blueprint_text or "唯一 Blueprint R03" not in blueprint_text:
    issues.append("blueprint_mapping_incomplete")
if not re.search(r"(?m)^## 附錄 E", body_text) or not re.search(r"(?m)^## 附錄 F", body_text):
    issues.append("appendix_E_or_F_missing_from_source")

ordered_docx_lines = []
for child in doc.element.body.iterchildren():
    if child.tag == qn("w:p"):
        line = "".join(node.text or "" for node in child.findall(".//" + qn("w:t")))
        if line.strip():
            ordered_docx_lines.append(line)
    elif child.tag == qn("w:tbl"):
        for tr in child.findall(qn("w:tr")):
            cells = []
            for tc in tr.findall(qn("w:tc")):
                cells.append("".join(node.text or "" for node in tc.findall(".//" + qn("w:t"))).strip())
            ordered_docx_lines.append("\t".join(cells))
doc_ordered = "\n".join(ordered_docx_lines)
source_main = body_text[body_text.index("# §0 "):]
doc_start = doc_ordered.rfind("§0 30 秒快速導引")
if doc_start < 0:
    issues.append("docx_main_body_start_not_found")
    similarity = 0.0
else:
    source_main = re.sub(r"<!--.*?-->", " ", source_main, flags=re.S)
    source_main = source_main.replace("**", "").replace("`", "")
    source_main = re.sub(r"(?m)^\s*#+\s*", "", source_main)
    source_main = re.sub(r"(?m)^\s*>\s?", "", source_main)
    source_main = re.sub(r"(?m)^\s*[-*]\s+", "", source_main)
    source_main = re.sub(r"(?m)^\s*\d+\.\s+", "", source_main)
    source_main = re.sub(r"(?m)^\|[\s:\-|]+\|$", " ", source_main)
    source_main = re.sub(r"\s+", "", source_main.replace("|", " "))
    doc_main = re.sub(r"\s+", "", doc_ordered[doc_start:])
    similarity = SequenceMatcher(None, source_main, doc_main).ratio()
    if similarity < 0.97:
        issues.append(f"markdown_docx_main_body_similarity_low:{similarity:.6f}")

pdf_reader = PdfReader(str(PDF))
if len(pdf_reader.pages) != 63:
    issues.append(f"pdf_page_count:{len(pdf_reader.pages)}")
pdf_text = "\n".join(page.extract_text() or "" for page in pdf_reader.pages)
if "REG-059" not in pdf_text or "附錄 F" not in pdf_text:
    issues.append("pdf_missing_register_tail_or_appendix_F")
if "第 63 頁／共 63 頁" not in pdf_text:
    issues.append("pdf_final_page_number_field_mismatch")

manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
if manifest.get("candidate_docx_sha256") != sha256(DOCX):
    issues.append("visual_manifest_docx_sha_mismatch")
if manifest.get("pdf_sha256") != sha256(PDF):
    issues.append("visual_manifest_pdf_sha_mismatch")
if manifest.get("pdf_pages") != 63 or len(manifest.get("pages", [])) != 63:
    issues.append("visual_manifest_page_coverage_mismatch")

result = {
    "schema": "TTQS_ONE_0084_TECHNICAL_REVIEW.v1",
    "revision": "R11",
    "document_id": "0084",
    "requirement_id": "I17C-E03",
    "official_checkbox_serial": "090",
    "document_type": "派訓主管學習成效評估：職場觀察量規和回饋",
    "candidate_status": "CANDIDATE / NOT_EXECUTED",
    "real_state": "UNKNOWN/HOLD",
    "formal_acceptance": "NOT_GRANTED",
    "docx": str(DOCX.relative_to(BASE)),
    "docx_sha256": sha256(DOCX),
    "body_sha256": sha256(BODY),
    "blueprint_r03_sha256": sha256(BLUEPRINT),
    "register_csv_sha256": sha256(REGISTER),
    "pdf_sha256": sha256(PDF),
    "visual_manifest_sha256": sha256(MANIFEST),
    "docx_zip_crc": "PASS" if bad_member is None else f"FAIL:{bad_member}",
    "python_docx_open": "PASS",
    "sections": len(doc.sections),
    "tables": len(doc.tables),
    "register_table_index_zero_based": register_matches[0][0] if len(register_matches) == 1 else None,
    "register_rows": len(csv_rows),
    "register_word_csv_exact_parity": len(register_matches) == 1 and register_matches[0][1][1:] == csv_rows,
    "register_state_counts": state_counts,
    "register_state_sum": sum(state_counts.values()),
    "case_fields_O0_O3": case_fields,
    "cross_document_case_link": "SYN-0083-B01" in doc_text,
    "sample_real_boundary_tokens_present": all(x in body_text for x in ("不是協會正式表單", "不是已執行評估", "SAMPLE/SYNTHETIC", "NOT_EXECUTED", "REAL=UNKNOWN/HOLD", "FORMAL_ACCEPTANCE=NOT_GRANTED")),
    "prohibited_approval_claim_hits": prohibited_hits,
    "body_docx_text_similarity": round(similarity, 6),
    "pdf_pages": len(pdf_reader.pages),
    "pdf_final_page_number": "第 63 頁／共 63 頁" in pdf_text,
    "issues": issues,
    "technical_verdict": "PASS" if not issues else "FAIL",
    "semantic_review": "PENDING_EXACT_SHA_INDEPENDENT_REVIEW",
    "visual_review": "PENDING_MANUAL_VISUAL_RECEIPT",
    "formal_acceptance": "NOT_GRANTED",
}
OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(result, ensure_ascii=False, indent=2))
if issues:
    raise SystemExit(1)
