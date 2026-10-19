from __future__ import annotations

import hashlib
import json
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


BASE = Path(__file__).resolve().parents[1]
SOURCE = BASE / "candidate" / "0084_R10.docx"
OUTPUT = BASE / "candidate" / "0084_R11.docx"
RECEIPT = BASE / "evidence" / "format_R11_appendixE.json"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def document_text(doc: Document) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def set_cell_margins(cell, top=30, bottom=30, left=80, right=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    old = tc_pr.find(qn("w:tcMar"))
    if old is not None:
        tc_pr.remove(old)
    margins = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        item = OxmlElement(f"w:{side}")
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")
        margins.append(item)
    tc_pr.append(margins)


if not SOURCE.is_file():
    raise FileNotFoundError(SOURCE)
if OUTPUT.exists():
    raise FileExistsError(f"Refusing to overwrite {OUTPUT}")

doc = Document(str(SOURCE))
before_text = document_text(doc)
matches = []
for index, table in enumerate(doc.tables):
    if not table.rows or len(table.rows[0].cells) != 6:
        continue
    header = [cell.text.strip() for cell in table.rows[0].cells]
    if header[0] != "REG_ID" or len(table.rows) != 60:
        continue
    ids = [row.cells[0].text.strip() for row in table.rows[1:]]
    if ids == [f"REG-{i:03d}" for i in range(1, 60)]:
        matches.append((index, table))

if len(matches) != 1:
    raise RuntimeError(f"Expected exactly one Appendix E table with REG-001..059; found {len(matches)}")

table_index, table = matches[0]
for row in table.rows:
    for cell in row.cells:
        set_cell_margins(cell)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.line_spacing = 1.0
            for run in paragraph.runs:
                run.font.size = Pt(9.0)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUTPUT))

check = Document(str(OUTPUT))
after_text = document_text(check)
if before_text != after_text:
    raise RuntimeError("Text changed during the layout-only edit")

receipt = {
    "schema": "TTQS_ONE_0084_LAYOUT_ONLY_FORMAT.v1",
    "revision": "R11",
    "source_candidate": str(SOURCE.relative_to(BASE)),
    "source_docx_sha256": sha256(SOURCE),
    "candidate": str(OUTPUT.relative_to(BASE)),
    "candidate_docx_sha256_before_word_render": sha256(OUTPUT),
    "target_table_index_zero_based": table_index,
    "table_header": [cell.text for cell in table.rows[0].cells],
    "register_ids": 59,
    "change_scope": "Appendix E register table style only; text and content retained",
    "font_size_pt": 9.0,
    "paragraph_line_spacing": 1.0,
    "cell_margin_twips": {"top": 30, "bottom": 30, "left": 80, "right": 80},
    "text_sha256_before": hashlib.sha256(before_text.encode("utf-8")).hexdigest(),
    "text_sha256_after": hashlib.sha256(after_text.encode("utf-8")).hexdigest(),
    "text_unchanged": True,
    "status": "LAYOUT_CANDIDATE_PENDING_WORD_RENDER_AND_VISUAL_REVIEW",
}
RECEIPT.parent.mkdir(parents=True, exist_ok=True)
RECEIPT.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(receipt, ensure_ascii=False, indent=2))
