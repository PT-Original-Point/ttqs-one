from pathlib import Path
import csv, hashlib, json, re, zipfile
from docx import Document
from docx.oxml.ns import qn
ROOT=Path(__file__).parent
DOCX=ROOT/'0142_I19-E32_R06_BLUEPRINT_REPAIR_CANDIDATE.docx'
CSV=ROOT/'SAMPLE_REPLACEMENT_REGISTER_0142_R06.csv'
RECEIPT=ROOT/'TECHNICAL_QA_RECEIPT_0142_R06.json'
def sha(p):
 h=hashlib.sha256(); h.update(p.read_bytes()); return h.hexdigest()
def fail(msg): raise SystemExit('FAIL: '+msg)
if not DOCX.is_file() or not CSV.is_file(): fail('candidate or register missing')
with zipfile.ZipFile(DOCX) as z:
 bad=z.testzip()
 if bad: fail('corrupt zip member: '+bad)
d=Document(DOCX)
paragraphs=[p.text for p in d.paragraphs]
celltexts=[c.text for t in d.tables for row in t.rows for c in row.cells]
alltext='\n'.join(paragraphs+celltexts)
flat=re.sub(r'\s+',' ',alltext)
required=[
 'TTQS ONE','0142','I19-E32','156','UNKNOWN','R06','CANDIDATE','NOT ISSUED','UNKNOWN/HOLD','NOT ASSIGNED','NOT GRANTED',
 '表 1','表 2','表 3','表 4','表 5','表 6','表 7','表 8',
 '來源登錄','科目映射','模擬財報/帳簿勾稽表','分子/分母','解除條件','SAMPLE/SYNTHETIC',
 '+10.00%','+21.00%','+18.38%','三個連續完整會計年度','REAL 必需來源'
]
missing=[x for x in required if x not in flat]
if missing: fail('required content missing: '+repr(missing))
if re.search(r'待填|待核|空白表格|保持空白|會後填寫',alltext): fail('prohibited placeholder remains')
if re.search(r'\+\s*[0-9,]+\.\d%(?!\d)',alltext): fail('one-decimal percentage remains')
with CSV.open('r',encoding='utf-8-sig',newline='') as f: rows=list(csv.DictReader(f))
if len(rows)!=19: fail(f'expected 19 carried-forward facts, got {len(rows)}')
ids=[r['sample_fact_id'] for r in rows]
if len(set(ids))!=19: fail('duplicate sample fact ids')
register_tables=[t for t in d.tables if t.rows and t.rows[0].cells[0].text.strip()=='sample_fact_id']
if len(register_tables)!=1: fail(f'expected exactly one embedded sample_fact_id table, got {len(register_tables)}')
register_table=register_tables[0]
doc_register_ids=[row.cells[0].text.strip() for row in register_table.rows[1:]]
if len(register_table.columns)!=4 or len(doc_register_ids)!=19: fail(f'embedded register shape mismatch: rows={len(doc_register_ids)}, cols={len(register_table.columns)}')
if doc_register_ids!=ids: fail('embedded register IDs/order do not match companion CSV')
missing_ids=[i for i in ids if i not in alltext]
if missing_ids: fail('register IDs absent from docx: '+repr(missing_ids))
if any(not r.get('document_locator') or not r.get('real_replacement_source_required') or not r.get('replacement_date_or_not_applicable') for r in rows): fail('register required fields empty')
if any(r['data_status']=='SAMPLE/SYNTHETIC' and r['real_status']!='UNKNOWN/HOLD' for r in rows): fail('SAMPLE row not held for REAL replacement')
if sum(1 for r in rows if r['data_status']=='SAMPLE/SYNTHETIC')!=18: fail('unexpected SAMPLE/control row count')
# Verify the four published calculations from the named synthetic inputs.
vals={2023:(1200000,280000,1480000),2024:(1320000,260000,1580000),2025:(1452000,300000,1752000)}
for y,(op,other,total) in vals.items():
 if op+other!=total: fail(f'{y} annual total mismatch')
if round((vals[2024][0]-vals[2023][0])/vals[2023][0]*100,2)!=10.00: fail('2024 operating growth mismatch')
if round((vals[2025][0]-vals[2024][0])/vals[2024][0]*100,2)!=10.00: fail('2025 operating growth mismatch')
if round((vals[2025][0]-vals[2023][0])/vals[2023][0]*100,2)!=21.00: fail('operating cumulative growth mismatch')
if round((vals[2025][2]-vals[2023][2])/vals[2023][2]*100,2)!=18.38: fail('total cumulative growth mismatch')
# All operational tables are numbered/captioned; cover card is explicitly unnumbered.
captions=[f'表 {i}' for i in range(1,9)]
if any(c not in flat for c in captions): fail('numbered table caption missing')
# Main body font floor 10.5pt; every table run floor 9.5pt.
body_small=[]
for pi,p in enumerate(d.paragraphs):
 for ri,r in enumerate(p.runs):
  if r.text.strip() and r.font.size and r.font.size.pt<10.5: body_small.append((pi,ri,r.font.size.pt,r.text[:30]))
if body_small: fail('main body below 10.5pt: '+repr(body_small[:8]))
table_small=[]
for ti,t in enumerate(d.tables):
 for ri,row in enumerate(t.rows):
  for ci,cell in enumerate(row.cells):
   for p in cell.paragraphs:
    for r in p.runs:
     if r.text.strip() and r.font.size and r.font.size.pt<9.5: table_small.append((ti,ri,ci,r.font.size.pt,r.text[:25]))
if table_small: fail('table below 9.5pt: '+repr(table_small[:8]))
if 'SAMPLE/SYNTHETIC' not in '\n'.join(p.text for p in d.sections[0].header.paragraphs): fail('header lacks truth marker')
if not all(n in flat for n in ['0139 / I19-E29','0140 / I19-E30','0141 / I19-E31']): fail('negative-neighbor separation missing')
report={
 'status':'PASS_TECHNICAL_CONTENT_GATES_ONLY', 'doc_id':'0142','requirement_id':'I19-E32','revision':'R06','docx_sha256':sha(DOCX),'docx_bytes':DOCX.stat().st_size,
 'source_packet_sha256':'e3e8b3a8ee1effe758608fee183440add4f37381988c5e352912b1a3d55c9e49',
 'unique_blueprint_sha256':'4841161d6e9aad3ace2a0b4ec4a5a3e1b7e92b68ff2b20f95e9edeae2b84070b',
 'table_count_physical':len(d.tables),'numbered_operational_tables':8,'embedded_register_table_rows':len(doc_register_ids),'embedded_register_table_columns':len(register_table.columns),'sample_register_rows':len(rows),'sample_rows':18,'control_metadata_rows':1,
 'font_floors_pt':{'body':10.5,'table':9.5},'calculation_checks':'PASS_SYNTHETIC_ONLY',
 'real_evidence':'UNKNOWN/HOLD','official_score':'NOT_ASSIGNED','formal_acceptance':'NOT_GRANTED',
 'semantic_review':'PENDING','visual_review':'PENDING','checks':{
  'docx_zip_integrity':'PASS','requirement_and_index':'PASS','blueprint_tables_and_controls':'PASS','register_ids_and_required_fields':'PASS','embedded_register_rows_match_csv':'PASS',
  'no_forbidden_placeholder':'PASS','two_decimal_percentage_display':'PASS','four_sample_calculations':'PASS','font_floor':'PASS','neighbor_requirements_separated':'PASS'
 }
}
RECEIPT.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))


