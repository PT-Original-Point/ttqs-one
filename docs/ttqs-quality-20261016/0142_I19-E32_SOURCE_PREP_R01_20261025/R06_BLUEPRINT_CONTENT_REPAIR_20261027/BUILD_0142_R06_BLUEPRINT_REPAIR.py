from pathlib import Path
import csv, hashlib, json, math, re
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0142_I19-E32_SOURCE_PREP_R01_20261025")
OUT = BASE / "R06_BLUEPRINT_CONTENT_REPAIR_20261027"
BUILD_DATE = datetime.now().astimezone().strftime('%Y-%m-%d')
SOURCE_CSV = BASE / "SAMPLE_REPLACEMENT_REGISTER_0142_R05.csv"
DOCX = OUT / "0142_I19-E32_R06_BLUEPRINT_REPAIR_CANDIDATE.docx"
CSV_OUT = OUT / "SAMPLE_REPLACEMENT_REGISTER_0142_R06.csv"
RECEIPT = OUT / "DOCX_BUILD_RECEIPT_0142_R06.json"
NAVY = "17365D"
BLUE = "1F4E78"
PALE = "EAF1F8"
PALE2 = "F4F7FA"
INK = "243447"
MUTED = "526273"
WHITE = "FFFFFF"
GOLD = "FFF2CC"


def sha256(path):
    h=hashlib.sha256()
    with open(path,'rb') as f:
        for chunk in iter(lambda:f.read(1024*1024),b''): h.update(chunk)
    return h.hexdigest()

def set_run_font(run, size=10.5, bold=False, color=INK, italic=False):
    run.font.name = 'Arial'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    rpr=run._element.get_or_add_rPr()
    rfonts=rpr.rFonts
    if rfonts is None:
        rfonts=OxmlElement('w:rFonts'); rpr.insert(0,rfonts)
    rfonts.set(qn('w:ascii'),'Arial'); rfonts.set(qn('w:hAnsi'),'Arial')
    rfonts.set(qn('w:eastAsia'),'Microsoft JhengHei')

def set_style_font(style, size, color=INK, bold=False):
    style.font.name='Arial'; style.font.size=Pt(size); style.font.bold=bold; style.font.color.rgb=RGBColor.from_string(color)
    rpr=style._element.get_or_add_rPr(); rfonts=rpr.rFonts
    if rfonts is None: rfonts=OxmlElement('w:rFonts'); rpr.insert(0,rfonts)
    rfonts.set(qn('w:ascii'),'Arial'); rfonts.set(qn('w:hAnsi'),'Arial'); rfonts.set(qn('w:eastAsia'),'Microsoft JhengHei')

def set_cell_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill); shd.set(qn('w:val'),'clear')

def set_cell_margins(cell, top=70, start=85, bottom=70, end=85):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); mar=tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar=OxmlElement('w:tcMar'); tcPr.append(mar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=mar.find(qn('w:'+m))
        if node is None: node=OxmlElement('w:'+m); mar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr(); tblHeader=OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'),'true'); trPr.append(tblHeader)

def set_no_split(row):
    trPr=row._tr.get_or_add_trPr(); cant=OxmlElement('w:cantSplit'); trPr.append(cant)

def set_cell_text(cell, text, header=False, size=9.5, fill=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text=str(text)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill: set_cell_shading(cell,fill)
    set_cell_margins(cell)
    for p in cell.paragraphs:
        p.alignment=align; p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.03
        for run in p.runs: set_run_font(run,size,bold=header,color=WHITE if header else INK)

def add_caption(doc, text):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(7); p.paragraph_format.space_after=Pt(3); p.paragraph_format.keep_with_next=True
    r=p.add_run(text); set_run_font(r,10.5,True,BLUE)

def add_table(doc, caption, headers, rows, widths=None, font_size=9.5):
    add_caption(doc,caption)
    table=doc.add_table(rows=1, cols=len(headers)); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style='Table Grid'; table.autofit=False
    tblPr=table._tbl.tblPr; layout=OxmlElement('w:tblLayout'); layout.set(qn('w:type'),'fixed'); tblPr.append(layout)
    for j,h in enumerate(headers):
        set_cell_text(table.rows[0].cells[j],h,True,font_size,NAVY,WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(table.rows[0]); set_no_split(table.rows[0])
    for ri,row in enumerate(rows):
        cells=table.add_row().cells
        for j,val in enumerate(row):
            set_cell_text(cells[j],val,False,font_size,PALE2 if ri%2 else None)
        set_no_split(table.rows[-1])
    if widths:
        for row in table.rows:
            for j,w in enumerate(widths):
                row.cells[j].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return table

def add_body(doc,text,lead=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(5); p.paragraph_format.line_spacing=1.1; p.paragraph_format.widow_control=True
    if lead and text.startswith(lead):
        r=p.add_run(lead); set_run_font(r,10.5,True,BLUE)
        r=p.add_run(text[len(lead):]); set_run_font(r,10.5)
    else:
        r=p.add_run(text); set_run_font(r,10.5)
    return p

def add_bullet(doc,text,level=0):
    p=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); p.paragraph_format.space_after=Pt(2); p.paragraph_format.left_indent=Inches(.25+level*.2)
    for r in p.runs: set_run_font(r,10.5)
    if not p.runs: set_run_font(p.add_run(text),10.5)
    else: p.runs[0].text=text
    return p

def add_heading(doc,text,level=1):
    p=doc.add_paragraph(style=f'Heading {level}'); p.paragraph_format.keep_with_next=True; p.paragraph_format.space_before=Pt(12 if level==1 else 7); p.paragraph_format.space_after=Pt(4)
    r=p.add_run(text); set_run_font(r,16 if level==1 else 12.5,True,NAVY if level==1 else BLUE)
    return p

def add_field(paragraph, code):
    run=paragraph.add_run(); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),code); run._r.addnext(fld)

def load_and_update_register():
    with SOURCE_CSV.open('r',encoding='utf-8-sig',newline='') as f: rows=list(csv.DictReader(f))
    if len(rows)!=19: raise ValueError(f'Expected 19 carried-forward register rows, got {len(rows)}')
    fields=list(rows[0].keys())+['replacement_date_or_not_applicable']
    for row in rows:
        typ=row['fact_type']; fid=row['sample_fact_id']
        year=next((y for y in ('2023','2024','2025') if y in fid),None)
        source_zh={
            'reporting_period_and_year_end':'經授權的完整年度財報及會計期間證明',
            'operating_revenue':'正式財報、收入明細/總帳及財務權責職務確認的科目映射',
            'other_income':'正式財報及已勾稽的收入明細/總帳',
            'total_income_derived':'正式財報及已勾稽的收入明細/總帳',
            'year_over_year_growth_derived':'完成原件替換、口徑核實及帳表勾稽後的來源數據',
            'cumulative_operating_growth_derived':'完成原件替換、科目映射及帳表勾稽後的來源數據',
            'cumulative_total_income_growth_derived':'完成原件替換及帳表勾稽後的來源數據',
            'currency_and_unit_assumption':'正式財報及一致的幣別/金額單位依據',
            'account_mapping_assumption':'正式科目表及文件化的財務權責職務映射',
            'document_control_date':'如獲核准使用，依實際受控文件登錄資料；不得作為財務證據'
        }
        locator_zh={
            'reporting_period_and_year_end':f'表4 {year}年度勾稽列；附錄A本列 sample_fact_id',
            'operating_revenue':f'表4 {year}年度 SAMPLE/SYNTHETIC 營業收入列；附錄A本列 sample_fact_id',
            'other_income':f'表4 {year}年度其他收入分列；附錄A本列 sample_fact_id',
            'total_income_derived':f'表4 {year}年度衍生總額列；附錄A本列 sample_fact_id',
            'year_over_year_growth_derived':'表5對應逐年算式/結果列；附錄A本列 sample_fact_id',
            'cumulative_operating_growth_derived':'表5營業收入三年累計列；附錄A本列 sample_fact_id',
            'cumulative_total_income_growth_derived':'表5總收入三年累計列；附錄A本列 sample_fact_id',
            'currency_and_unit_assumption':'表3至表5 SAMPLE/SYNTHETIC 基礎說明；附錄A本列 sample_fact_id',
            'account_mapping_assumption':'表3 SAMPLE科目映射假設；附錄A本列 sample_fact_id',
            'document_control_date':'封面文件控制卡的候選建置日期'
        }
        derivation_zh={
            'reporting_period_and_year_end':'作為三個完整年度連續性演示期間；未取自協會財報。',
            'operating_revenue':'營業收入示例值；採用示例分類假設，非協會實際科目決定。',
            'other_income':'示例其他收入分開列示；不作營業收入增幅分子。',
            'total_income_derived':'由同年度示例營業收入加其他收入衍生；只作算術覆核。',
            'year_over_year_growth_derived':'（本年示例營業收入－前一年示例營業收入）/前一年示例營業收入×100%；結果四捨五入至兩位。',
            'cumulative_operating_growth_derived':'（2025示例營業收入－2023示例營業收入）/2023示例營業收入×100%；結果四捨五入至兩位。',
            'cumulative_total_income_growth_derived':'（2025示例總收入－2023示例總收入）/2023示例總收入×100%；不得替代營業收入趨勢。',
            'currency_and_unit_assumption':'新臺幣僅為展示假設；不是從協會財報讀取。',
            'account_mapping_assumption':'示例映射假設；實際非營利組織分類須由財務權責職務依正式科目表確認。',
            'document_control_date':'R06候選文件產製日期；不是財務報導期間、協會核准日或生效日。'
        }
        unit_zh={
            'reporting_period_and_year_end':'完整會計年度；示例期間假設為曆年',
            'operating_revenue':'新臺幣；僅供營業收入示例分類',
            'other_income':'新臺幣；合成數值，與營業收入分開',
            'total_income_derived':'新臺幣；由合成數值衍生',
            'year_over_year_growth_derived':'百分比；由合成數值計算',
            'cumulative_operating_growth_derived':'百分比；由合成數值計算',
            'cumulative_total_income_growth_derived':'百分比；由合成數值計算',
            'currency_and_unit_assumption':'示例假設',
            'account_mapping_assumption':'明確標示的合成映射假設',
            'document_control_date':'版本控制資料；非財務事實'
        }
        rule_zh={
            'reporting_period_and_year_end':'只以授權且完整的實際會計年度替換；期間不連續或不完整即維持 UNKNOWN/HOLD。',
            'operating_revenue':'取得原件後由財務權責職務確認科目映射，並逐年與收入明細/總帳勾稽；未完成前維持 UNKNOWN/HOLD。',
            'other_income':'依正式來源逐年另列並勾稽；不得併入營業收入以製造成長。',
            'total_income_derived':'依正式財報與已勾稽帳簿重算；差額未解即維持 UNKNOWN/HOLD。',
            'year_over_year_growth_derived':'以核實後的連續年度原始數值重算，保留分子、分母、公式、版本與覆核紀錄。',
            'cumulative_operating_growth_derived':'以核實後的首末完整年度及已確認營業收入科目重算。',
            'cumulative_total_income_growth_derived':'以核實後的首末完整年度重算；總收入不得取代營業收入結論。',
            'currency_and_unit_assumption':'記錄正式來源幣別、金額單位及會計基礎；不可調節至一致口徑時維持 HOLD。',
            'account_mapping_assumption':'依正式科目表與文件化財務權責確認替換；未核實不得列為 REAL。',
            'document_control_date':'保留為候選版本控制日期；不得當作財務資料或協會正式核准。'
        }
        row['document_locator']=locator_zh[typ]
        row['real_replacement_source_required']=source_zh[typ]
        row['derivation_or_use']=derivation_zh[typ]
        row['unit_or_basis']=unit_zh[typ]
        row['replacement_rule']=rule_zh[typ]
        if typ=='document_control_date':
            row['data_status']='CONTROL_METADATA'
            row['synthetic_value']=f'{BUILD_DATE}；R06候選建置日；未生效/未核准'
            row['real_status']='不適用（版本控制資料，非財務證據）'
            row['replacement_date_or_not_applicable']='不適用：控制資料，不是財務事實替換'
        else:
            row['data_status']='SAMPLE/SYNTHETIC'
            row['real_status']='UNKNOWN/HOLD'
            row['replacement_date_or_not_applicable']='尚未替換；未提供授權 REAL 來源'
            if typ=='reporting_period_and_year_end' and year:
                row['synthetic_value']=f'{year}-01-01 至 {year}-12-31；年末 {year}-12-31'
            elif typ=='year_over_year_growth_derived': row['synthetic_value']='+10.00%'
            elif typ=='cumulative_operating_growth_derived': row['synthetic_value']='+21.00%'
            elif typ=='cumulative_total_income_growth_derived': row['synthetic_value']='+18.38%'
            elif typ=='currency_and_unit_assumption': row['synthetic_value']='新臺幣（NTD）'
            elif typ=='account_mapping_assumption': row['synthetic_value']='示例假設：「營業收入」欄位對應「服務收入」；非協會分類決定'
    
    with CSV_OUT.open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=fields); writer.writeheader(); writer.writerows(rows)
    return rows

register=load_and_update_register()
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
sec.top_margin=Inches(.66); sec.bottom_margin=Inches(.66); sec.left_margin=Inches(.65); sec.right_margin=Inches(.65)
sec.header_distance=Inches(.30); sec.footer_distance=Inches(.30)
for sname,size,color,bold in [('Normal',10.5,INK,False),('Title',25,NAVY,True),('Subtitle',12,MUTED,False),('Heading 1',16,NAVY,True),('Heading 2',12.5,BLUE,True),('List Bullet',10.5,INK,False),('List Bullet 2',10.5,INK,False)]:
    set_style_font(doc.styles[sname],size,color,bold)
    if sname=='Normal': doc.styles[sname].paragraph_format.space_after=Pt(5); doc.styles[sname].paragraph_format.line_spacing=1.1

# Header/footer are persistent truth-bound status labels.
header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
set_run_font(header.add_run('TTQS ONE  |  0142 / I19-E32  |  SAMPLE/SYNTHETIC'),9.5,True,BLUE)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
set_run_font(footer.add_run('CANDIDATE  |  REAL=UNKNOWN/HOLD  |  官方分數未指定  |  頁 '),9.5,False,MUTED); add_field(footer,'PAGE')

# Cover
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(52); p.paragraph_format.space_after=Pt(6); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
set_run_font(p.add_run('組織近三年營業額與收入趨勢核對報告'),24,True,NAVY)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(16)
set_run_font(p.add_run('TTQS ONE  |  文件 0142  |  I19-E32'),13,True,BLUE)
callout=doc.add_table(rows=1,cols=1); callout.alignment=WD_TABLE_ALIGNMENT.CENTER; callout.style='Table Grid'
set_cell_text(callout.cell(0,0),'重要狀態：本稿是可操作的候選核對工具，並含完整 SAMPLE/SYNTHETIC 算例。尚未取得協會授權三年度財報、收入明細/總帳與科目表；因此 REAL=UNKNOWN/HOLD，不能據此宣稱本會營業額成長。',False,11,GOLD)
for width in [6.95]: callout.columns[0].width=Inches(width)
control=[
('專案 / 文件','TTQS ONE / 0142','官方要求','I19-E32'),
('157 方框索引','156','印刷 checkbox serial','UNKNOWN'),
('版本 / 狀態','R06 / CANDIDATE','生效狀態','NOT ISSUED'),
('候選建置日',BUILD_DATE,'正式接受','NOT GRANTED'),
('REAL 證據','UNKNOWN/HOLD','官方評分','NOT ASSIGNED'),
('合成資料','SAMPLE/SYNTHETIC；僅驗算','核准 / 簽署','未主張；無簽核紀錄'),
]
add_table(doc,'文件控制卡（封面控制資訊，不納入操作表編號）', ['欄位','值','欄位','值'],control,[1.2,2.25,1.35,2.15],9.5)
add_body(doc,'唯一佐證目的：以三個連續完整會計年度的可調回財務原件及可核對帳簿，判斷營業額/營業收入是否逐年增加。此文件不是財報、會計簽證、官方評分或協會核准紀錄。')
doc.add_page_break()

add_heading(doc,'1. 核對目的、適用範圍與責任分工',1)
add_body(doc,'核對對象限於同一組織三個連續完整會計年度的營業額/營業收入。服務收入、捐款、補助、會費、利息及代收款不得未經科目映射便合併。表 1 的職務分工是本候選工具所建議的作業角色，不證明原始點協會已採行或已指派人員。')
roles=[
('資料保管/財務行政職務','依年度收集財報、查核/簽證文件、收入明細或總帳、科目表；登錄版本、頁碼/科目列、單位與保存位置。','原件清冊與來源定位；缺件逐項標 UNKNOWN/HOLD。'),
('財務權責職務','確認營業額科目、會計基礎、幣別/單位、退款折讓與重分類處理；說明非營利科目是否適用。','有權責依據的科目映射；無文件化確認不得將假設升為 REAL。'),
('獨立覆核職務','逐年核對財報與收入明細/總帳，重算差異與增幅，檢查完整期間和分母。','可追溯覆核紀錄、差異處理及重算版本；不得代替官方評分。'),
('文件控管職務','維持版本、來源連結與變更軌跡；核對所有表格使用同一來源和期間。','受控版本及變更紀錄；不得自行簽核、核准或宣稱正式生效。'),
]
add_table(doc,'表 1　責任角色與可交付控制點',['職務角色','執行工作','完成證據 / 權限界線'],roles,[1.35,3.0,2.6])
add_heading(doc,'執行順序與停止點',2)
for t in [
'1) 收件與授權：確認三年度財報及帳簿可由組織合法調回；敏感原件保存在受控位置，不放入未授權公開儲存。',
'2) 建立來源定位：記錄文件名、版本/日期、頁或科目列、幣別、單位、會計基礎、核准/簽證狀態與保管路徑。',
'3) 核實科目映射：由財務權責職務記錄每個帳科目納入或排除營業額的依據；未確認即停在 HOLD。',
'4) 逐年勾稽：財報收入列對照收入明細/總帳；差額須有原始調整憑證與覆核軌跡。',
'5) 僅在年度連續、口徑一致、差異解除及分母有效後計算趨勢；否則百分比標 N/A 並保留 HOLD。',
'6) 覆核後保存計算版本和差異處理；任何真實簽核或正式接受仍由具權責人員完成。']:
    add_body(doc,t)

doc.add_page_break()
add_heading(doc,'2. 真實來源登錄與調回定位',1)
add_body(doc,'表 2 是 REAL 收件清冊。此候選建立時，未提供任何協會真實財報、查核/簽證報告、收入明細/總帳或科目表；各格均明列缺件，不以空白表示完成。三年度實際頁碼、科目列、版本、幣別、會計基礎、核准狀態、擷取/覆核日期與差異均為 UNKNOWN/HOLD。')
source_rows=[]
for fy in ['2023','2024','2025']:
    source_rows.append((fy,
      '年度財報：未提供\n查核/簽證報告：未提供\n收入明細/總帳：未提供\n科目表/映射：未提供',
      '頁碼/科目列：未知 / HOLD\n文件版本及原件位置：未知 / HOLD',
      '幣別/單位：未知 / HOLD\n會計基礎：未知 / HOLD',
      '核准/簽證：未知 / HOLD\n擷取與覆核角色/日期：未執行',
      '表帳差異：未能檢查；不得推定為零'))
add_table(doc,'表 2　REAL 財務來源登錄清冊（依完整年度逐年列示）',['完整會計年度','應取得來源及版本','頁碼/科目列/原件位置','幣別、單位、會計基礎','核准與角色/日期','差異及狀態'],source_rows,[.72,1.58,1.26,1.12,1.25,1.02],9.5)
add_body(doc,'解除條件：每年度完整原件已取得且可調回；來源定位及口徑可比；財務權責職務已確認科目映射；財報與帳簿逐項相符或有可查調整；獨立覆核完成。任一條件未達，REAL 維持 UNKNOWN/HOLD。')

add_heading(doc,'3. 科目映射與比較口徑',1)
add_body(doc,'本表把 SAMPLE 算例中的「營業收入」視作一個計算欄位，不將它映射成協會實際服務收入。真實科目須由正式科目表、收入明細及有權責的財務角色確認；下表任何 SAMPLE 處理都不是協會分類決定。')
mapping=[
('服務/營業收入','SAMPLE 算例假設納入營業收入','實際科目、收入條件與確認基礎均未提供','核對正式科目表、財報列、收入明細及權責確認','UNKNOWN/HOLD；不得升列 REAL'),
('捐款/指定捐款','算例另列或排除，不併入營業收入','實際分類及限制條件未提供','核對捐款用途、會計政策與財報呈列','UNKNOWN/HOLD'),
('政府/民間補助','算例另列或排除，不併入營業收入','實際分類及補助條件未提供','核對補助核定、認列期間及帳務科目','UNKNOWN/HOLD'),
('會費/會員收入','不納入算例營業收入','實際收入性質及認列政策未提供','核對章程/收費性質及已核准會計政策','UNKNOWN/HOLD'),
('利息及投資收益','不納入算例營業收入','實際分類未提供','核對財報列、銀行對帳及會計政策','UNKNOWN/HOLD'),
('代收款/代付款','不納入營業收入；避免把代收金額算成組織收入','實際代理關係及淨額/總額口徑未提供','核對契約、收付款憑證與會計處理依據','UNKNOWN/HOLD'),
('退款、折讓、沖銷、重分類','算例未另造調整數；真實案件須列原額、調整與依據','實際調整不存在於本候選證據','逐筆連結憑證、年度、科目、覆核及重算','UNKNOWN/HOLD；差異未解即停審'),
]
add_table(doc,'表 3　科目映射、分類及納入邊界',['科目/收入類型','SAMPLE 演示處理','REAL 分類現況','真實核實依據','停止狀態'],mapping,[1.13,1.46,1.52,1.9,.92],9.5)
add_body(doc,'可比口徑必須固定或可調節：完整年度、幣別、金額單位、權責/現金基礎、收入確認政策、退款折讓、重分類及財報版本。不同口徑未能逐項調節時，不可用換算或估算製造可比性。')

add_heading(doc,'4. SAMPLE 年度勾稽演示',1)
add_body(doc,'下表只示範將相同 SAMPLE 金額放入「模擬財報列」與「模擬帳簿列」後如何計算差異；沒有任何對應的真實財報、總帳或明細檔案。表中差額為算術設計值，不是查帳結論。SAMPLE 期間和金額均列於附錄 A 清冊。')
recon=[]
annual=[('2023',1200000,280000,1480000),('2024',1320000,260000,1580000),('2025',1452000,300000,1752000)]
for fy,op,other,total in annual:
    for label,val in [('營業收入（SAMPLE）',op),('其他收入（分開列示）',other),('總收入（衍生值）',total)]:
        recon.append((fy,label,f'{val:,} NTD','同額 SAMPLE 演示值', '0 NTD（演示差額）','模擬行，非實際原件；REAL=UNKNOWN/HOLD'))
add_table(doc,'表 4　SAMPLE 模擬財報/帳簿勾稽表（算術示範，非真實查核）',['年度','項目','模擬財報列','模擬明細/總帳列','模擬差額','證據解讀'],recon,[.55,1.18,1.12,1.23,.94,1.9],9.5)
add_body(doc,'年度總額驗算：2023 年 1,200,000 + 280,000 = 1,480,000；2024 年 1,320,000 + 260,000 = 1,580,000；2025 年 1,452,000 + 300,000 = 1,752,000。其他收入全程獨立列示，不作營業收入趨勢的替代分子。')

add_heading(doc,'5. 趨勢計算、分母與判讀規則',1)
add_body(doc,'公式：逐年增幅 =（本年營業收入 - 前一年營業收入）/ 前一年營業收入 × 100%；首年為基準年，不計增幅。累計變化 =（第三年 - 第一年）/ 第一年 × 100%。保留原始金額，結果顯示至小數點後兩位。下表的分子、分母均為 SAMPLE/SYNTHETIC。')
trend=[
('營業收入逐年增幅','2024 對 2023','1,320,000 - 1,200,000 = 120,000','120,000 / 1,200,000 × 100 = +10.00%','SAMPLE 僅；REAL 未核'),
('營業收入逐年增幅','2025 對 2024','1,452,000 - 1,320,000 = 132,000','132,000 / 1,320,000 × 100 = +10.00%','SAMPLE 僅；REAL 未核'),
('營業收入三年累計變化','2025 對 2023','1,452,000 - 1,200,000 = 252,000','252,000 / 1,200,000 × 100 = +21.00%','SAMPLE 僅；REAL 未核'),
('總收入三年累計變化（分開列示）','2025 對 2023','1,752,000 - 1,480,000 = 272,000','272,000 / 1,480,000 × 100 = +18.38%','不得替代營業收入趨勢'),
]
add_table(doc,'表 5　SAMPLE 趨勢算式、分子/分母與結果',['指標','比較期間','分子','分母及完整算式/結果','可用範圍'],trend,[1.22,1.0,1.35,2.18,1.18],9.5)
add_heading(doc,'真實判定規則',2)
for t in [
'只有三個完整連續年度、同幣別/單位/會計基礎、科目映射已核實、帳表已勾稽且兩個年度增幅都大於 0，才可以描述「逐年成長」。',
'前期分母小於或等於 0、年度不連續、幣別/單位不一、分類不可調和或帳表差異未解時，該比較結果為 N/A；不得換算、推估或從總收入取代營業收入。',
'總收入、捐款、補助、會費或一次性收入只作分類後的附帶核對；不能掩蓋營業收入下降。',
'本候選 SAMPLE 的兩個年度增幅和三年累計算式通過算術演示，不表示協會任何年度成長。']:
    add_bullet(doc,t)

add_heading(doc,'6. 本候選判讀與正式結論邊界',1)
add_heading(doc,'SAMPLE 與 REAL 結論分離',2)
add_body(doc,'SAMPLE 算術：示例營業收入 1,200,000 → 1,320,000 → 1,452,000；逐年 +10.00%、+10.00%，累計 +21.00%。此結果只證明示例公式可重算。')
add_body(doc,'REAL 財務證據：UNKNOWN/HOLD；目前沒有三年正式財報、收入明細/總帳、科目表及權責映射，故不得描述本會營業收入趨勢。')
add_body(doc,'TTQS 評分 / 正式接受：NOT ASSIGNED / NOT GRANTED；本候選不是評分、核准、簽章或正式送審結果。')
add_body(doc,'真實結論欄位只能在原件及覆核完成後更新，並須保留來源 ID、文件版本、頁碼/科目列、幣別/單位、計算工作底稿與覆核紀錄。若結果不符合逐年增長條件，應如實記錄，不得改選有利期間或把其他收入混入。')

add_heading(doc,'7. 例外登錄、停審及解除條件',1)
exceptions=[
('任一年度財報、查核/簽證報告、收入明細/總帳或科目表缺件','受影響年度及整體三年趨勢','財務資料保管職務；持續 HOLD','取得授權原件、記錄版本/定位/基礎及覆核證據後再檢查'),
('營業收入科目映射未由權責職務確認','相關科目與全部比較年度','財務權責職務；不得計入 REAL','受控科目表、書面分類依據及權責覆核完成'),
('財報與總帳/明細不一致，或調整無憑證','受影響年度及跨年結果','財務行政與獨立覆核職務；停審','逐筆調整可回查、差異重算並留存覆核日期/版本'),
('年度不連續、短年度或會計期間不同','相關年度比較','資料整理職務；增幅標 N/A','取得完整連續年度原件，或說明不可比並不作成長結論'),
('前期分母 ≤ 0，或幣別/單位/基礎不可比','該比較列','獨立覆核職務；增幅標 N/A','有權責依據可調節至一致基礎；不得估算'),
('捐款、補助、會費、利息或代收款可能混入營業收入','涉及的科目及年度','財務權責職務；維持分類 HOLD','核對正式政策、原件與帳列，記錄納入/排除理由'),
('原件含個資、帳戶資訊或未授權公開內容','該原件及同步包','資料保管/隱私權責職務；停止公開同步','保留受控來源；只同步獲准的摘要、遮罩內容與驗收收據'),
]
add_table(doc,'表 6　例外、責任角色與停審解除條件',['觸發條件','影響範圍','責任/狀態','解除所需證據'],exceptions,[1.95,1.35,1.48,2.17],9.5)

add_heading(doc,'8. 近鄰要求排除與跨文件連結',1)
add_body(doc,'本文件只對應 I19-E32 的三年營業額趨勢。只有共同、授權且可追溯的來源才可用 source_id/fact_id 建立跨文件連結；沒有共同真實來源，不推定數字、日期或結論相同。表 7 用於拒絕近鄰錯件。')
neighbors=[
('0139 / I19-E29 / 索引 153','考照/就業率與分母','不作為營業收入、財報來源或成長結論。'),
('0140 / I19-E30 / 索引 154','訓練資訊系統使用定義與期間','不替代財務原件、科目映射或年度勾稽。'),
('0141 / I19-E31 / 索引 155','顧客滿意度多期同量尺統計','不替代營業額趨勢；不得借用其樣本、期間或分數。'),
]
add_table(doc,'表 7　相似但不同的官方要求與拒收界線',['文件 / 要求 / 索引','唯一佐證用途','本文件禁止的替代方式'],neighbors,[1.55,2.2,3.2],9.5)

appendix_section=doc.add_section(WD_SECTION.NEW_PAGE)
appendix_section.orientation=WD_ORIENT.LANDSCAPE
appendix_section.page_width=Inches(11.69); appendix_section.page_height=Inches(8.27)
appendix_section.top_margin=Inches(.48); appendix_section.bottom_margin=Inches(.48)
appendix_section.left_margin=Inches(.47); appendix_section.right_margin=Inches(.47)
add_heading(doc,'附錄 A. SAMPLE/SYNTHETIC 替換清冊（逐項可追溯）',1)
add_body(doc,'本附錄將 19 個示例/控制資料逐列綁定至本文件位置、合成值、REAL 必需來源及替換狀態；它同時附有 CSV 清冊。每個財務示例都是 SAMPLE/SYNTHETIC；最後一列為候選版本控制日期，不是財務證據。所有真實替換日期均尚未發生，因為未提供 REAL 來源。')
register_rows=[]
for row in register:
    loc=row['document_locator']
    val=row['synthetic_value']
    source=row['real_replacement_source_required']
    status=row['real_status']
    date=row['replacement_date_or_not_applicable']
    if re.fullmatch(r'\d+',val): val=f"{int(val):,} NTD"
    source_text=f"REAL 替換來源：{source}\n狀態：{status}\n替換日期：{date}"
    register_rows.append((row['sample_fact_id'],loc,val,source_text))
add_table(doc,'表 8　逐列 SAMPLE/控制資料替換清冊',['sample_fact_id','文件位置 / 欄位','SAMPLE 或控制值','REAL 必需來源、狀態與替換日期'],register_rows,[2.25,2.65,1.65,4.2],9.5)
add_body(doc,'表 8 中 REAL 欄位均為明確狀態；UNKNOWN/HOLD 不是空白、不是已完成替換，也不是核准。若未來資料收件，逐項記錄原件保存位置、來源版本、頁/列、幣別/單位、權責映射、覆核角色/日期及替換日期，並重新計算和版本化。')


# Keep body font floors and table style consistent after all edits.
for p in doc.paragraphs:
    if p.style and p.style.name.startswith('Heading'):
        p.paragraph_format.keep_with_next=True
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    if r.font.size is None or r.font.size.pt<9.5: set_run_font(r,9.5, r.bold, WHITE if r.bold else INK)

# Core properties distinguish this candidate from a real financial record.
doc.core_properties.title='組織近三年營業額與收入趨勢核對報告 - 0142 R06 Candidate'
doc.core_properties.subject='TTQS ONE I19-E32 / checkbox index 156 / SAMPLE-SYNTHETIC only'
doc.core_properties.author='TTQS ONE Candidate Builder'
doc.core_properties.comments='Candidate only. REAL evidence UNKNOWN/HOLD; not issued or formally accepted.'
doc.save(DOCX)

# Minimal build receipt, hashed after the file is closed.
receipt={
 'project_id':'TTQS_ONE','doc_id':'0142','official_requirement_id':'I19-E32','official_checkbox_index':156,'printed_checkbox_serial':'UNKNOWN',
 'revision':'R06','state':'CANDIDATE','effective':'NOT ISSUED','real_evidence':'UNKNOWN/HOLD','official_score':'NOT ASSIGNED','formal_acceptance':'NOT GRANTED',
 'source_packet_sha256':sha256(BASE/'SOURCE_PACKET_0142_I19-E32_R01.md'),
 'unique_blueprint_sha256':sha256(BASE/'UNIQUE_BLUEPRINT_0142_I19-E32_R01.md'),
 'official_source_pdf_sha256':'805386b77df045b396fabafed7a47d12491c25a0b5744edd9a04876caee34709',
 'crosswalk_sha256':'1ddfbb608bcad25170683ddb42a1dbb27bafa6f1f35ff0498f1b8d75b24f0679',
 'docx_path':str(DOCX),'docx_sha256':sha256(DOCX),'docx_bytes':DOCX.stat().st_size,
 'sample_register_path':str(CSV_OUT),'sample_register_sha256':sha256(CSV_OUT),'sample_register_rows':len(register),
 'source_script_path':str(Path(__file__)),'source_script_sha256':sha256(Path(__file__)),
 'created_local':datetime.now().astimezone().isoformat(),'semantic_review':'PENDING','visual_review':'PENDING','technical_qa':'PENDING',
 'substantive_delta':['source register','account mapping','SAMPLE annual reconciliation','two-decimal growth results','in-document 19-row replacement register','checkbox index/serial','consistent captions and references'],
}
RECEIPT.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(receipt,ensure_ascii=False,indent=2))
