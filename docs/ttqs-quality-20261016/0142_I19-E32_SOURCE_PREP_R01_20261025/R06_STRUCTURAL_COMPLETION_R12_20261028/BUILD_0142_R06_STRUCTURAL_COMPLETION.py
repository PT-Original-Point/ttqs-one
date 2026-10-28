from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
import csv, hashlib, json, re

base = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0142_I19-E32_SOURCE_PREP_R01_20261025")
out = base / "R06_STRUCTURAL_COMPLETION_R12_20261028"
src = base / "0142_I19-E32_R05_HEADER_CONTRAST_CANDIDATE.docx"
src_csv = base / "SAMPLE_REPLACEMENT_REGISTER_0142_R05.csv"
blueprint = base / "UNIQUE_BLUEPRINT_0142_I19-E32_R01.md"
source_packet = base / "SOURCE_PACKET_0142_I19-E32_R01.md"
target = out / "0142_I19-E32_CANDIDATE_R06.docx"
register_out = out / "SAMPLE_REPLACEMENT_REGISTER_0142_R06.csv"
receipt_out = out / "BUILD_RECEIPT_0142_R06.json"

def sha(path):
    h=hashlib.sha256()
    with open(path,"rb") as f:
        for chunk in iter(lambda:f.read(1024*1024),b""): h.update(chunk)
    return h.hexdigest()

def set_cell_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr()
    shd=tcPr.find(qn("w:shd"))
    if shd is None:
        shd=OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"),fill)

def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr()
    if trPr.find(qn("w:tblHeader")) is None:
        el=OxmlElement("w:tblHeader"); el.set(qn("w:val"),"true"); trPr.append(el)

def set_no_split(row):
    trPr=row._tr.get_or_add_trPr()
    if trPr.find(qn("w:cantSplit")) is None:
        trPr.append(OxmlElement("w:cantSplit"))

def set_run_font(run, size=9.5, bold=None, color=None):
    run.font.name="Microsoft JhengHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"Microsoft JhengHei")
    run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    if color: run.font.color.rgb=__import__("docx").shared.RGBColor.from_string(color)

def style_caption(p, text):
    p.text=text
    p.style="Caption"
    p.paragraph_format.space_before=Pt(7)
    p.paragraph_format.space_after=Pt(3)
    p.paragraph_format.keep_with_next=True
    for run in p.runs: set_run_font(run,10.5,True,"1F1F1F")

def insert_caption_before(table, text):
    elem=OxmlElement("w:p")
    table._tbl.addprevious(elem)
    p=Paragraph(elem,table._parent)
    style_caption(p,text)
    return p

def set_cell_text(cell, text, header=False):
    cell.text=str(text)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        p.paragraph_format.space_before=Pt(0)
        p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.line_spacing=1.0
        for r in p.runs:
            set_run_font(r,9.5,header, "FFFFFF" if header else "1F1F1F")
    if header: set_cell_shading(cell,"1F4E78")

def set_cell_widths(table, weights):
    sec=doc.sections[0]
    total=float(sec.page_width-sec.left_margin-sec.right_margin)/914400
    s=sum(weights)
    widths=[Inches(total*w/s) for w in weights]
    table.autofit=False
    for row in table.rows:
        for i,w in enumerate(widths):
            if i<len(row.cells): row.cells[i].width=w

def set_table_cell_margins(table, top=35, start=65, bottom=35, end=65):
    # Tighten only the long register table while retaining 9.5 pt cell text.
    for row in table.rows:
        for cell in row.cells:
            tcPr=cell._tc.get_or_add_tcPr()
            tcMar=tcPr.find(qn("w:tcMar"))
            if tcMar is None:
                tcMar=OxmlElement("w:tcMar")
                tcPr.append(tcMar)
            for side,value in (("top",top),("start",start),("bottom",bottom),("end",end)):
                node=tcMar.find(qn(f"w:{side}"))
                if node is None:
                    node=OxmlElement(f"w:{side}")
                    tcMar.append(node)
                node.set(qn("w:w"),str(value))
                node.set(qn("w:type"),"dxa")

def add_table(caption, headers, rows, weights):
    p=doc.add_paragraph()
    style_caption(p,caption)
    t=doc.add_table(rows=1,cols=len(headers))
    t.style="Table Grid"
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers): set_cell_text(t.rows[0].cells[i],h,True)
    set_repeat_table_header(t.rows[0])
    set_no_split(t.rows[0])
    for values in rows:
        row=t.add_row()
        for i,v in enumerate(values): set_cell_text(row.cells[i],v,False)
        set_no_split(row)
    set_cell_widths(t,weights)
    return t

def add_section(title, intro):
    h=doc.add_heading(title,level=1)

    h.paragraph_format.keep_with_next=True
    p=doc.add_paragraph(intro)
    p.paragraph_format.space_after=Pt(5)
    return h,p

# Load existing candidate as the exact same-document repair base.
doc=Document(src)
before_sha=sha(src)
if before_sha != "a213a30cd949948f7ce8c2b5d76fc79b5fdcd90117f6fba77139af6660c5ab6e":
    raise SystemExit("BASE_SHA_MISMATCH; refusing to patch an unexpected candidate")

# Cover metadata and truth boundary.
for p in doc.paragraphs:
    if p.text.startswith("文件編號：TTQS-0142"):
        p.text="文件編號：TTQS-0142｜官方要求：I19-E32｜157清單索引：156｜印刷 checkbox serial：UNKNOWN｜版本：1.1｜狀態：CANDIDATE／SAMPLE示範／未生效｜修訂日期：2026-10-28｜生效日期：未生效"
    elif p.text.startswith("表5　"):
        style_caption(p,"表5　三年度營業額與其他收入示例｜SAMPLE資料集A")
    elif p.text.startswith("表6　"):
        style_caption(p,"表6　示例資料定位與核對責任")
    elif p.text.startswith("表11　"):
        style_caption(p,"表11　官方要求、欄位及證據位置追溯")
    elif p.startswith if False else False:
        pass
    if "版本：1.0（示範文件，未生效）" in p.text:
        p.text=p.text.replace("版本：1.0（示範文件，未生效）","版本：1.1（候選文件；SAMPLE/SYNTHETIC示範；未生效）")
        p.text=p.text.replace("合成示例位置見第4節表5及表6","合成示例位置見第4節表5至表8；來源登錄、科目映射、帳表勾稽、停審條件及逐列替換見第8至12節表12至表16")
        p.text=p.text.replace("首頁文件控制資訊","首頁文件控制資訊（I19-E32；157索引156；印刷序號UNKNOWN）")

# Update visible percentages and specific conclusion language in every existing paragraph/cell.
replacements={
    "+10.0%":"+10.00%",
    "+21.0%":"+21.00%",
    "+18.4%":"+18.38%",
    "累計增加21.0%":"累計增加21.00%",
    "展示至小數點後一位":"統一顯示至小數點後兩位"
}
for p in doc.paragraphs:
    for old,new in replacements.items():
        if old in p.text: p.text=p.text.replace(old,new)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for old,new in replacements.items():
                if old in cell.text: cell.text=cell.text.replace(old,new)

# Add captions to existing physical tables that were uncaptioned. Tables 5, 6 and 11
# already have a caption paragraph and have just been corrected above.
caption_map={
    1:"表1　快速閱讀與本報告用途",
    2:"表2　官方要求來源基礎",
    3:"表3　核對欄位定義與驗證責任",
    4:"表4　SAMPLE資料聲明與使用界線",
    7:"表7　逐年及首末累計趨勢算式",
    8:"表8　SAMPLE算例判讀",
    9:"表9　證據限制與例外處理",
    10:"表10　真實資料替換與升格條件"
}
for table_no,text in caption_map.items():
    insert_caption_before(doc.tables[table_no-1],text)

# Replace the existing three-column calculation table with exact visible numerators,
# denominators, formulas, results, units and comparability limits.
calc=doc.tables[6]
calc.rows[0].cells[0].text="核對項目與期間"
calc.rows[0].cells[1].text="逐步算式（含分子／分母）"
calc.rows[0].cells[2].text="結果、單位與可比性"
calc_data=[
("2024營業額年增率｜本年1,320,000；前期1,200,000",
 "分子=1,320,000−1,200,000=120,000；分母=1,200,000；120,000÷1,200,000×100%",
 "+10.00%；百分比。僅在完整年度、幣別／單位及分類一致時可比；本例為SAMPLE假設。"),
("2025營業額年增率｜本年1,452,000；前期1,320,000",
 "分子=1,452,000−1,320,000=132,000；分母=1,320,000；132,000÷1,320,000×100%",
 "+10.00%；百分比。僅在完整年度、幣別／單位及分類一致時可比；本例為SAMPLE假設。"),
("2023至2025營業額累計變動｜末年1,452,000；基年1,200,000",
 "分子=1,452,000−1,200,000=252,000；分母=1,200,000；252,000÷1,200,000×100%",
 "+21.00%；百分比。此值不替代逐年正成長檢查；本例為SAMPLE假設。"),
("2023至2025總收入累計變動｜末年1,752,000；基年1,480,000",
 "分子=1,752,000−1,480,000=272,000；分母=1,480,000；272,000÷1,480,000×100%",
 "+18.38%；百分比。總收入僅作旁列，不能替代營業額趨勢。")
]
while len(calc.rows)>1:
    calc._tbl.remove(calc.rows[-1]._tr)
for values in calc_data:
    cells=calc.add_row().cells
    for i,v in enumerate(values): set_cell_text(cells[i],v,False)
for i,h in enumerate(["核對項目與期間","逐步算式（含分子／分母）","結果、單位與可比性"]):
    set_cell_text(calc.rows[0].cells[i],h,True)
set_repeat_table_header(calc.rows[0])
for row in calc.rows: set_no_split(row)
set_cell_widths(calc,[1.7,2.7,2.1])

# Correct the cover and traceability table. Keep SAMPLE and REAL states distinct.
trace=doc.tables[10]
trace_rows=[
("官方要求與文件定位","官方原文：組織近三年營業額持續成長。｜Requirement：I19-E32｜157清單索引：156｜印刷 checkbox serial：UNKNOWN｜來源：TTQS訓練機構版查核佐證資料表實體第9頁。"),
("資料、方法與判讀","第4節表5至表8提供三年度SAMPLE資料、來源定位、逐步算式及示例判讀；第8至12節表12至表16提供來源登錄、科目映射、帳表勾稽、例外停審及逐列替換。表5至表8均為SAMPLE/SYNTHETIC，不能證明本會實績。"),
("限制、跨文件與升格","沒有受權三年財報、收入明細／總帳、科目映射及逐年調節證據；REAL=UNKNOWN/HOLD、官方分數NOT_ASSIGNED、正式接受NOT_GRANTED。只以本件SAMPLE ID追蹤；沒有共享REAL source_id/fact_id，不推定0139、0140、0141或其他文件的數字、期間或結論可沿用。")
]
for i,(a,b) in enumerate(trace_rows,1):
    trace.rows[i].cells[0].text=a
    trace.rows[i].cells[1].text=b

# Load the 19-row replacement register and revise visible precision/control metadata.
with open(src_csv,encoding="utf-8-sig",newline="") as f:
    data=list(csv.DictReader(f))
if len(data)!=19: raise SystemExit(f"REGISTER_ROW_COUNT_EXPECTED_19_GOT_{len(data)}")
for rownum,row in enumerate(data, start=1):
    ft=row.get("fact_type","")
    if ft=="year_over_year_growth_derived":
        row["synthetic_value"]="10.00%"
    elif ft=="cumulative_total_income_growth_derived":
        row["synthetic_value"]="18.38%"
    elif ft=="document_control_date":
        row["data_status"]="CONTROL_METADATA"
        row["fact_type"]="document_revision_metadata"
        row["synthetic_value"]="2026-10-28 (R06 revision date)"
        row["unit_or_basis"]="document revision metadata; not a financial reporting date"
        row["document_locator"]="Cover document control field"
        row["derivation_or_use"]="R06 candidate revision date; does not establish approval or effectiveness"
        row["real_replacement_source_required"]="Not applicable as financial evidence; any effectiveness requires a separate authorized document-control record"
        row["real_status"]="NOT_APPLICABLE"
        row["replacement_rule"]="Do not treat as a synthetic financial fact or as approval. Record effectiveness only from the controlled register and authorized approval."
    row.setdefault("real_review_status","NOT_REVIEWED_REAL")
    row.setdefault("replacement_date","NOT_REPLACED")
    # Compact display codes keep the Word register readable; the full original
    # sample_fact_id remains intact in the companion CSV for exact lookup.
    row["display_id"]=f"F{rownum:02d}"
with open(register_out,"w",encoding="utf-8-sig",newline="") as f:
    fields=list(data[0].keys())
    if "real_review_status" not in fields: fields.append("real_review_status")
    if "replacement_date" not in fields: fields.append("replacement_date")
    w=csv.DictWriter(f,fieldnames=fields,extrasaction="ignore")
    w.writeheader()
    for row in data: w.writerow(row)

# Add requirement-specific operational appendices with fully worked synthetic example.
doc.add_page_break()
add_section("8　年度來源登錄與原件調回","表12以同一組合成資料示範逐年如何記錄來源識別、版本、期間、頁列、幣別與帳務基礎、核准狀態、整理／覆核角色及差異。列出的來源識別只指向本文件的SAMPLE資料集，並非實際財報、總帳或外部原件；REAL來源尚未提供，所有年份維持HOLD。")
source_rows=[]
for year,turn,other,total in [(2023,"1,200,000","280,000","1,480,000"),(2024,"1,320,000","260,000","1,580,000"),(2025,"1,452,000","300,000","1,752,000")]:
    source_rows.append((
        f"SYN-SRC-{year}-A｜{year}-01-01至{year}-12-31",
        f"本文件SAMPLE資料集A；表5第{year}列；版本SYN-R01；非外部原件",
        "新臺幣元；完整曆年；同一示例口徑假設",
        "合成整理角色；無真人調閱或簽核",
        f"示例營業額{turn}；其他收入{other}；總收入{total}",
        "正式財報名稱／版本／頁碼、收入明細或總帳科目列、幣別與基礎、受控原件位置均未提供；REAL=UNKNOWN/HOLD"
    ))
add_table("表12　年度來源登錄（SAMPLE例示；REAL原件未提供）",
          ["年度與期間／SAMPLE來源ID","SAMPLE來源、版本及文件定位","幣別／單位／帳務基礎","整理角色／核准狀態","示例金額摘要","REAL原件定位與狀態"],
          source_rows,[1.2,1.45,1.0,0.9,1.0,1.45])

doc.add_page_break()
add_section("9　科目映射與收入分類","表13示範把「營業額」的案例假設和其他收入分開。非營利組織實際科目是否屬營業收入，必須由有權的財務資料保管者依組織會計科目表確認；本表的SAMPLE分類不是協會會計政策或正式核准。")
mapping_rows=[
("SYN-REV-SVC-01","示例服務收入；合成科目，不是真實科目代碼","本算例假設納入","只為示範逐年營業額公式；不代表協會科目映射","REAL科目表與權責確認未提供；UNKNOWN/HOLD"),
("SYN-OTHER-AGG-01","示例其他收入合計；合成總額，不推定其構成","不納入本例營業額；另列","280,000／260,000／300,000只作其他收入欄，未拆成實際來源類別","REAL分類、原始科目與版本未確認；UNKNOWN/HOLD"),
("REAL-CLASSIFICATION-GATE","捐款、補助、會費、利息、代收款、退款／折讓等候核項目","未核實前不自行納入","逐項比對財報列、總帳明細及分類依據；不以總收入替代營業額","由實際財務權責角色確認；原件、核准依據及版本未提供；HOLD")
]
add_table("表13　營業收入與其他收入科目映射（示例假設與REAL關卡）",
          ["科目ID／狀態","合成科目與分類","納入營業額？","示例依據及適用範圍","REAL權責確認、版本與狀態"],
          mapping_rows,[1.15,1.5,0.95,1.55,1.85])

add_section("10　年度帳表勾稽與差異調節","表14先核對三年度示範收入合計，再把「算術差異」與「財報對總帳差異」分開。表內加總差異為SAMPLE數學檢查；未取得財報及總帳原件，因此正式帳表差異不是零，而是UNKNOWN/HOLD。")
recon_rows=[
("2023","1,200,000","280,000","1,480,000","1,200,000＋280,000−1,480,000＝0","財報／總帳原件未提供；UNKNOWN/HOLD"),
("2024","1,320,000","260,000","1,580,000","1,320,000＋260,000−1,580,000＝0","財報／總帳原件未提供；UNKNOWN/HOLD"),
("2025","1,452,000","300,000","1,752,000","1,452,000＋300,000−1,752,000＝0","財報／總帳原件未提供；UNKNOWN/HOLD")
]
add_table("表14　年度收入勾稽與帳表差異調節（SAMPLE）",
          ["年度","營業額（SAMPLE）","其他收入（SAMPLE）","總收入（SAMPLE）","示例加總差異","正式財報／總帳差異與處理"],
          recon_rows,[0.55,0.95,0.95,0.95,1.55,2.05])

doc.add_page_break()
add_section("11　例外停審與解除條件","以下任一情形都必須先停下受影響年度或分類，不得以估算補值或SAMPLE結論解除REAL停審。實際責任人與期限只能由協會指定；在原件、權限或差異證據未完成前，維持UNKNOWN/HOLD。")
exception_rows=[
("正式年度財報或必要附註缺件","受影響年度；三年趨勢不得定案","HOLD；不外推","財務資料保管角色；實際期限由單位指定","取得授權版本、完整年度與可調回頁碼／科目列後重算"),
("年度不連續、短年度或結帳狀態不明","相關年度與相鄰成長率","不可比；N/A／HOLD","資料整理角色與財務覆核角色；期限由單位指定","確認起訖日、結帳完成狀態及可比較的連續期間"),
("前期基數≤0或金額單位不同","受影響公式與累計比較","該增幅=N/A；不得除以零或自行換算","覆核角色；期限由單位指定","取得權威幣別／單位及可用基期；保留原值與換算依據"),
("營業收入科目尚未核定","所有營業額欄及成長結論","HOLD；其他收入不得代替","具權限財務資料保管角色；期限由單位指定","核對核准科目表、分類依據、年度適用版本及授權確認"),
("財報與總帳／收入明細不相符","差異涉及年度、科目與總額","HOLD；差異不得寫成零","整理角色與獨立覆核角色；期限由單位指定","逐項定位差異、調整分錄及原件；重算並保留前後金額"),
("捐款、補助、會費、代收款或一次性收入混入","受影響分類及趨勢","分類未決則排除並HOLD","財務權責角色；期限由單位指定","取得收入明細與分類依據；確認是否屬營業收入後重算"),
("財報版本、核准狀態或來源位置不可追溯","相關年度原件與計算","HOLD；不得以截圖或未授權副本替代","原件保管角色；期限由單位指定","記錄正式名稱、版本、頁列、核准狀態及受控保存位置"),
("REAL原件缺授權或不可公開","原件調閱／分享範圍","不得上傳或公開；維持UNKNOWN/HOLD","資料保管角色；期限由單位指定","在有權限的受控來源完成核對；公開版只列不敏感的驗收狀態")
]
add_table("表15　例外、停審責任與解除證據",
          ["停止條件","受影響範圍","判定／停止規則","責任角色與期限","解除證據及後續重算"],
          exception_rows,[1.0,0.9,1.35,1.0,2.2])

add_section("12　逐列SAMPLE替換清冊","表16列碼F01–F19逐列對應隨附CSV的display_id；CSV保留完整sample_fact_id供追溯。金額均為SAMPLE/SYNTHETIC，幣別以新臺幣為示範假設。REAL財務原件尚未提供者一律UNKNOWN/HOLD、未覆核、未替換；封面修訂資料不是財務證據，也不代表核准生效。")
replacement_rows=[]
for row in data:
    fid=row["display_id"]
    ft=row.get("fact_type","")
    year_match=re.search(r"20\d{2}",row.get("sample_fact_id",""))
    year=year_match.group(0) if year_match else ""
    raw=row.get("synthetic_value","")
    if ft=="reporting_period_and_year_end":
        loc=f"表5／6｜{year}年度"
        val=f"{year}完整曆年；SAMPLE"
        source="授權年度財報＋會計期間佐證"
    elif ft=="operating_revenue":
        loc=f"表5營業收入／表6｜{year}"
        val=f"{int(raw):,}元；NTD假設"
        source="授權財報＋收入明細／總帳＋科目映射"
    elif ft=="other_income":
        loc=f"表5其他收入｜{year}"
        val=f"{int(raw):,}元；分列SAMPLE"
        source="授權財報＋已勾稽收入明細／總帳"
    elif ft=="total_income_derived":
        loc=f"表5總收入／表14勾稽｜{year}"
        val=f"{int(raw):,}元；合成加總"
        source="授權財報＋收入明細／總帳勾稽"
    elif ft=="year_over_year_growth_derived":
        loc=f"表5年增率／表7｜{year}"
        val=f"{raw}；合成重算"
        source="REAL來源替換後重算並覆核"
    elif ft=="cumulative_operating_growth_derived":
        loc="表7營業收入累計"
        val=f"{raw}；合成重算"
        source="REAL來源替換後重算並覆核"
    elif ft=="cumulative_total_income_growth_derived":
        loc="表7總收入累計"
        val=f"{raw}；合成重算"
        source="REAL來源替換後重算並覆核"
    elif ft=="currency_and_unit_assumption":
        loc="表5–7幣別／單位"
        val="新臺幣元；示範假設"
        source="核定財報幣別與計量單位"
    elif ft=="account_mapping_assumption":
        loc="表4假設／表13科目映射"
        val="營業收入＝服務收入；合成假設"
        source="核定科目表＋財務角色確認"
    elif ft=="document_revision_metadata":
        loc="封面版次控制欄"
        val="2026-10-28；非財務事實"
        source="受控版次紀錄；生效核准另查"
    else:
        raise SystemExit(f"UNMAPPED_REGISTER_FACT_TYPE:{ft}")
    if row.get("real_status")=="NOT_APPLICABLE":
        status="不適用於REAL財務\n未核准生效"
    else:
        status="UNKNOWN/HOLD\n未覆核；未替換"
    replacement_rows.append((fid,loc,val,source,status))
replacement_table=add_table("表16　SAMPLE逐列替換清冊（19列；18項合成事實及1項文件控制中繼資料）",
          ["列碼","文件位置","SAMPLE示例值／單位","REAL替換必要來源","REAL狀態／覆核／日期"],
          replacement_rows,[0.55,1.0,1.1,2.55,1.35])
set_table_cell_margins(replacement_table,top=35,start=60,bottom=35,end=60)

# Metadata and document style.
doc.core_properties.title="組織近三年營業額與收入趨勢核對報告"
doc.core_properties.subject="TTQS ONE / 0142 / I19-E32 / SAMPLE-SYNTHETIC candidate"
doc.core_properties.comments="Candidate R06. REAL evidence UNKNOWN/HOLD; formal acceptance NOT_GRANTED."
doc.save(target)

# Produce an updated CSV companion bound to the exact R06 revision.
with open(register_out,"rb") as f: csv_sha=hashlib.sha256(f.read()).hexdigest()
receipt={
 "schema":"TTQS_ONE_0142_R06_BUILD_RECEIPT_V1",
 "project_id":"TTQS_ONE",
 "doc_id":"0142",
 "requirement_id":"I19-E32",
 "official_checkbox_index":156,
 "printed_checkbox_serial":"UNKNOWN",
 "base_candidate_path":str(src),
 "base_candidate_sha256":before_sha,
 "candidate_path":str(target),
 "candidate_sha256":sha(target),
 "source_packet_sha256":sha(source_packet),
 "unique_blueprint_sha256":sha(blueprint),
 "register_path":str(register_out),
 "register_sha256":csv_sha,
 "register_row_count":len(data),
 "changes":[
  "added sequential captions for all 11 existing physical tables",
  "corrected all growth displays to two decimal places and exposed numerator/denominator calculations",
  "added complete synthetic annual source register",
  "added sample-only account mapping with finance-role approval hold",
  "added year-by-year synthetic sum and separate REAL statement-to-ledger HOLD",
  "added eight explicit stop/release rules",
  "embedded 19-row SAMPLE replacement register with real source, status, review and replacement-date fields",
  "updated document control to R06 candidate revision; no real approval/effectiveness asserted"
 ],
 "candidate_mutated_from_base":True,
 "sample_state":"SAMPLE/SYNTHETIC",
 "real_evidence_state":"UNKNOWN/HOLD",
 "official_score":"NOT_ASSIGNED",
 "formal_acceptance":"NOT_GRANTED"
}
receipt_out.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
print(json.dumps({"candidate":str(target),"candidate_sha256":sha(target),"base_sha256":before_sha,"register_sha256":csv_sha,"register_rows":len(data),"receipt":str(receipt_out),"receipt_sha256":sha(receipt_out),"tables":len(doc.tables)},ensure_ascii=False))
