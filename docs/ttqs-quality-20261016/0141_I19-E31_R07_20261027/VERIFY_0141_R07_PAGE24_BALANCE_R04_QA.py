import csv, hashlib, json, re, zipfile, xml.etree.ElementTree as ET
from pathlib import Path
from datetime import datetime, timezone
from docx import Document
from docx.oxml.ns import qn

ROOT=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0141_I19-E31_SOURCE_PREP_R01_20261025')
DOCX=ROOT/'0141_I19-E31_R07_FONTFIX_PAGE24_BALANCE_R04_CANDIDATE.docx'
MAN=ROOT/'MANUSCRIPT_0141_R07_CANDIDATE.md'
REG=ROOT/'SAMPLE_REPLACEMENT_REGISTER_0141_R07.csv'
RESP=ROOT/'SAMPLE_SURVEY_RESPONSES_0141_R01.csv'
SRC=ROOT/'SOURCE_PACKET_0141_I19-E31_R02.md'
BP=ROOT/'UNIQUE_BLUEPRINT_0141_I19-E31_R02.md'
BUILD=ROOT/'DOCX_BUILD_RECEIPT_0141_R07_FONTFIX_TERMINALPARA_R03_CANDIDATE.json'
RENDER=ROOT/'R07_TERMINALPARA_R04_WORD_RENDER_120DPI'/'export_receipt.json'
PAGES=ROOT/'R07_TERMINALPARA_R04_PAGE_RENDER_MANIFEST.json'
PRIOR_PAGES=ROOT/'R07_TERMINALPARA_R03_PAGE_RENDER_MANIFEST.json'
OUT=ROOT/'R07_PAGE24_BALANCE_R04_DETERMINISTIC_QA_RECEIPT_R02.json'
EXPECTED_DOCX='de111911a549a0be0bd6381a6b22ed694881ba6b3f83528d2b4a4e153002142b'
EXPECTED_R03_DOCX='bfdb9ae9e048dd43f87e84d9e541d63f655335d7d1478297969fe68debcea56e'
EXPECTED_R06_MAN='2c374341b936ab4e4bb13889f870f42c79e03f688b7e1d92478d1fe052d7d287'
EXPECTED_RESP='723259199dd51973fc6473ec000234cc853fa1511be05864fdfd7b08bb452862'

def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def norm(s): return re.sub(r'\s+','',s)
def check(name,ok,detail=None): checks.append({'name':name,'result':'PASS' if ok else 'FAIL','detail':detail})
checks=[]
PATCH=ROOT/'R07_PAGE24_BALANCE_LAYOUT_PATCH_R04_RECEIPT.json'
VISUAL=ROOT/'VISUAL_REVIEW_RECEIPT_0141_R07_PAGE24_BALANCE_R04.json'
files={p.name:sha(p) for p in [DOCX,MAN,REG,RESP,SRC,BP,BUILD,RENDER,PAGES,PRIOR_PAGES,PATCH,VISUAL]}
check('candidate exact SHA',files[DOCX.name]==EXPECTED_DOCX,files[DOCX.name])
patch=json.loads(PATCH.read_text(encoding='utf-8'))
visual=json.loads(VISUAL.read_text(encoding='utf-8'))
check('R04 layout patch binds exact R03 source and R04 candidate',patch.get('source_candidate_sha256')==EXPECTED_R03_DOCX and patch.get('output_candidate_sha256')==EXPECTED_DOCX and patch.get('changed_package_parts')==['word/document.xml'] and patch.get('visible_text_unchanged') is True and patch.get('target_row_id')=='SYN-0141-2025-010' and patch.get('existing_cant_split_preserved') is True and patch.get('page_break_before_verified') is True and patch.get('semantic_review')=='NOT_YET_REVIEWED_ON_R04_SHA')
visual_pages=set(visual.get('unchanged_pages_reused',[]))|set(visual.get('changed_pages_directly_inspected',[]))
check('visual receipt binds exact candidate and all 28 pages by delta scope',visual.get('candidate_sha256')==EXPECTED_DOCX and visual.get('result')=='PASS' and sorted(visual_pages)==list(range(1,29)) and visual.get('changed_pages_directly_inspected')==[23,24] and visual.get('page_manifest_sha256')==files[PAGES.name])

check('source response SHA unchanged',files[RESP.name]==EXPECTED_RESP,files[RESP.name])
check('source packet R02 exact SHA',files[SRC.name]=='43c0e1bda5dfeb12c1ba8fc8805b876641f5184e4bf95eb4ef6c3182cf8ee964',files[SRC.name])
check('unique blueprint R02 exact SHA',files[BP.name]=='e19761674045b2defa7090ba1d8f1224943a89e54811180c20b9ea362de59cda',files[BP.name])
check('DOCX ZIP integrity',zipfile.ZipFile(DOCX).testzip() is None)
d=Document(DOCX)
paras='\n'.join(p.text for p in d.paragraphs)
table_text='\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells)
full='\n'.join([paras,table_text])
flat=norm(full).replace('＝','=')
check('native structure',len(d.tables)==18 and len(d.sections)==3,{'tables':len(d.tables),'sections':len(d.sections)})
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with zipfile.ZipFile(DOCX) as z:
    body=ET.fromstring(z.read('word/document.xml')).find('w:body',ns)
    body_children=list(body)
    terminal=body_children[-2]
    ppr=terminal.find('w:pPr',ns)
    spacing=ppr.find('w:spacing',ns) if ppr is not None else None
    rpr=ppr.find('w:rPr',ns) if ppr is not None else None
    sz=rpr.find('w:sz',ns) if rpr is not None else None
    szcs=rpr.find('w:szCs',ns) if rpr is not None else None
    terminal_text=''.join(x.text or '' for x in terminal.findall('.//w:t',ns))
    compact_terminal=(terminal.tag.endswith('}p') and terminal_text=='' and spacing is not None and spacing.get(qn('w:line'))=='20' and spacing.get(qn('w:lineRule'))=='exact' and spacing.get(qn('w:before'))=='0' and spacing.get(qn('w:after'))=='0' and sz is not None and sz.get(qn('w:val'))=='2' and szcs is not None and szcs.get(qn('w:val'))=='2')
check('empty terminal paragraph compact without shrinking document text',compact_terminal,{'empty':terminal_text=='','line_twips':spacing.get(qn('w:line')) if spacing is not None else None,'font_half_points':sz.get(qn('w:val')) if sz is not None else None})

with zipfile.ZipFile(DOCX) as z:
    media=[n for n in z.namelist() if n.startswith('word/media/')]
check('embedded chart/media count',len(media)==1,media)
check('no raw Markdown heading markers',not re.search(r'(?m)^\s*#{1,6}\s',full) and '####' not in full)
required={
 'Q1-Q6 six core constructs':('Q1–Q6' in full or 'Q1-Q6' in full) and ('Q7' in full),
 'separate Q7 satisfaction':bool(re.search(r'Q7.{0,50}(獨立|整體滿意)',full,re.S)),
 'sample and real boundary':('SAMPLE / SYNTHETIC' in full or 'SAMPLE/SYNTHETIC' in full) and 'REAL = UNKNOWN / HOLD' in full,
 'unknown ASCII locator':'UNKNOWN-sourced' in full,
 'unknown nonbreaking locator':'UNKNOWN‑sourced' in full,
 'R06 predecessor SHA in appendix':EXPECTED_R06_MAN in full,
 'Q5 baseline and delta':'3.5000' in flat and '0.3000' in flat and '6/12=50.00%' in flat and '10/12=83.33%' in flat and '3.8000' in flat,
 'Q2 baseline and strict threshold':'3.4167' in flat and '0.2833' in flat and '5/12=41.67%' in flat and '9/12=75.00%' in flat and '10/12=83.33%' in flat and '大於75%' in flat,
 'all 36 appendix aggregate scope':'全部36筆' in full or '全部 36 筆' in full,
 'real data held unknown':'UNKNOWN/HOLD' in full and 'NOT_GRANTED' in full,
}
for n,v in required.items(): check(n,v)
with RESP.open(encoding='utf-8-sig',newline='') as f: responses=list(csv.DictReader(f))
with REG.open(encoding='utf-8-sig',newline='') as f: registers=list(csv.DictReader(f))
check('response CSV row count',len(responses)==36,len(responses))
check('register row count',len(registers)==44,len(registers))
resp_regs=[r for r in registers if r.get('source_artifact')=='SAMPLE_SURVEY_RESPONSES_0141_R01.csv']
check('response register row count',len(resp_regs)==36,len(resp_regs))
source_by={r['sample_case_id']:r for r in responses}
reg_by={r['sample_case_id']:r for r in resp_regs}
rows_equal=True; row_errors=[]
for idx,r in enumerate(responses,1):
    rr=reg_by.get(r['sample_case_id'])
    if rr is None: rows_equal=False; row_errors.append('missing register '+r['sample_case_id']); continue
    try: v=json.loads(rr['sample_value_or_locator'])
    except Exception as e: rows_equal=False; row_errors.append('bad json '+r['sample_case_id']); continue
    scores={f'Q{i}':int(r[f'Q{i}_'+{'Q1':'need_fit','Q2':'content_applicability','Q3':'instructor_support','Q4':'materials_resources','Q5':'admin_communication','Q6':'learning_environment','Q7':'overall_satisfaction'}[f'Q{i}']]) for i in range(1,8)}
    expected={'year':r['period_year'],'cohort':r['service_cohort_id'],'date':r['response_date_synthetic'],'segment':r['customer_segment'],'scores':scores,'Q8':r['Q8_open_comment']}
    if v!=expected: rows_equal=False; row_errors.append('register value mismatch '+r['sample_case_id'])
    loc=rr['docx_locator']
    if f'附錄三第{idx}筆' not in loc or ('全部36筆' not in loc and '全部 36 筆' not in loc): rows_equal=False; row_errors.append('locator/aggregate mismatch '+r['sample_case_id'])
    if rr['source_manuscript_sha256']!=EXPECTED_R06_MAN: rows_equal=False; row_errors.append('source SHA mismatch '+r['sample_case_id'])
check('register responses match source CSV and locators',rows_equal,row_errors[:12])
# Exact match of appendix III table rows against source CSV.
t=Document(DOCX).tables[14]
word_rows=[[norm(c.text) for c in row.cells] for row in t.rows]
word_match=(len(word_rows)==37)
for i,r in enumerate(responses,1):
    expected=[norm(r['record_state']),norm(r['period_year']),norm(r['service_cohort_id']),norm(r['sample_case_id']),norm(r['response_date_synthetic']),norm(r['customer_segment']),norm(' / '.join(r[f'Q{j}_'+{'Q1':'need_fit','Q2':'content_applicability','Q3':'instructor_support','Q4':'materials_resources','Q5':'admin_communication','Q6':'learning_environment','Q7':'overall_satisfaction'}[f'Q{j}']] for j in range(1,8))),norm(r['Q8_open_comment'])]
    if word_rows[i]!=expected: word_match=False; row_errors.append('DOCX appendix row mismatch '+str(i))
check('appendix III 36 DOCX rows equal source CSV',word_match,row_errors[-12:])
# Recompute yearly aggregates, then compare to current DOCX table values.
agg={}
for y in (2023,2024,2025):
    rs=[r for r in responses if int(r['period_year'])==y]
    agg[y]={}
    agg[y]['n']=len(rs)
    agg[y]['q1_q6_mean']=sum(int(r[f'Q{i}_'+{'Q1':'need_fit','Q2':'content_applicability','Q3':'instructor_support','Q4':'materials_resources','Q5':'admin_communication','Q6':'learning_environment'}[f'Q{i}']]) for r in rs for i in range(1,7))/(len(rs)*6)
    agg[y]['q7_mean']=sum(int(r['Q7_overall_satisfaction']) for r in rs)/len(rs)
    agg[y]['q7_top2']=sum(int(r['Q7_overall_satisfaction'])>=4 for r in rs)
    agg[y]['q2_mean']=sum(int(r['Q2_content_applicability']) for r in rs)/len(rs)
    agg[y]['q2_top2']=sum(int(r['Q2_content_applicability'])>=4 for r in rs)
    agg[y]['q5_mean']=sum(int(r['Q5_admin_communication']) for r in rs)/len(rs)
    agg[y]['q5_top2']=sum(int(r['Q5_admin_communication'])>=4 for r in rs)
expected_agg={2023:{'q1_q6_mean':3.486111,'q7_mean':3.583333,'q7_top2':7,'q2_mean':3.416667,'q2_top2':5,'q5_mean':3.5,'q5_top2':6},2024:{'q1_q6_mean':3.569444,'q7_mean':3.916667,'q7_top2':10,'q2_mean':3.416667,'q2_top2':5,'q5_mean':4.0,'q5_top2':8},2025:{'q1_q6_mean':3.694444,'q7_mean':4.25,'q7_top2':11,'q2_mean':3.916667,'q2_top2':9,'q5_mean':4.25,'q5_top2':10}}
math_ok=True; math_diffs=[]
for y,fields in expected_agg.items():
    for k,e in fields.items():
        a=agg[y][k]
        if abs(a-e)>0.00001: math_ok=False; math_diffs.append({'year':y,'field':k,'actual':a,'expected':e})
check('recomputed CSV annual aggregates',math_ok,{'actual':agg,'mismatches':math_diffs})
# Verify reported rates, baselines, deltas, strict >75% threshold at n=12.
q2_min=10/12*100
q2_delta=(10-5)/12*100
q5_baseline=6/12*100
check('strict Q2 Top-2 rule and baseline delta',q2_min>75 and abs(q2_min-83.333333)<.001 and abs(q2_delta-41.666667)<.001,{'minimum_count':10,'minimum_rate':q2_min,'baseline_count':5,'delta_pp':q2_delta})
check('Q5 Top-2 descriptive baseline only',abs(q5_baseline-50)<.001 and agg[2025]['q5_top2']==10,{'2023_baseline_rate':q5_baseline,'2025_sample_count':agg[2025]['q5_top2'],'threshold':'none'})
# Current receipts must agree on exact candidate and page totals.
build=json.loads(BUILD.read_text(encoding='utf-8'))
render=json.loads(RENDER.read_text(encoding='utf-8'))
pages=json.loads(PAGES.read_text(encoding='utf-8'))
check('R03 build receipt binds the exact source of the R04 layout patch',build.get('docx_sha256')==EXPECTED_R03_DOCX and patch.get('source_candidate_sha256')==build.get('docx_sha256') and build.get('source_manuscript_sha256')==files[MAN.name] and build.get('source_packet_sha256')==files[SRC.name] and build.get('unique_blueprint_sha256')==files[BP.name])
check('Word render binds unchanged current SHA and 28 pages',render.get('source_sha256_before')==EXPECTED_DOCX and render.get('source_sha256_after')==EXPECTED_DOCX and render.get('source_unchanged') is True and render.get('page_count_word')==28 and render.get('page_count_png')==28)
prior_pages=json.loads(PRIOR_PAGES.read_text(encoding='utf-8'))
page_checks=[]
prior_by_page={p.get('page',i+1):p['sha256'] for i,p in enumerate(prior_pages.get('pages',[]))}
render_root=ROOT/'R07_TERMINALPARA_R04_WORD_RENDER_120DPI'
for pg in pages.get('pages',[]):
    fp=render_root/pg['file']
    page_checks.append(fp.exists() and sha(fp)==pg['sha256'])
changed=[p['page'] for p in pages.get('pages',[]) if not p.get('matches_R03_render',False)]
unchanged_equal=all(p['sha256']==prior_by_page.get(p['page']) and p.get('matches_R03_render') is True for p in pages.get('pages',[]) if p['page'] not in (23,24))
check('28 current page rasters match exact manifest; only pages 23-24 differ from R03',pages.get('candidate_sha256')==EXPECTED_DOCX and len(pages.get('pages',[]))==28 and all(page_checks) and changed==[23,24] and unchanged_equal and len(prior_by_page)==28,{'page_count':len(pages.get('pages',[])),'changed_pages_vs_R03':changed,'unchanged_pages_exact_match':unchanged_equal,'all_hashes_match':all(page_checks)})
# Font floors: use direct run sizing, paragraph style inheritance, Normal fallback.
def run_size_pt(run, para):
    if run.font.size is not None: return run.font.size.pt
    st=run.style
    seen=set()
    while st is not None and st.style_id not in seen:
        seen.add(st.style_id)
        if st.font.size is not None: return st.font.size.pt
        st=st.base_style
    st=para.style
    seen=set()
    while st is not None and st.style_id not in seen:
        seen.add(st.style_id)
        if st.font.size is not None: return st.font.size.pt
        st=st.base_style
    return d.styles['Normal'].font.size.pt if d.styles['Normal'].font.size else None
body_sizes=[]; table_sizes=[]; unresolved=[]
for p in d.paragraphs:
    for run in p.runs:
        if run.text.strip():
            sz=run_size_pt(run,p)
            if sz is None: unresolved.append(('body',p.style.name,run.text[:40]))
            else: body_sizes.append(sz)
for ti,tb in enumerate(d.tables):
    for row in tb.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    if run.text.strip():
                        sz=run_size_pt(run,p)
                        if sz is None: unresolved.append((f'table{ti}',p.style.name,run.text[:40]))
                        else: table_sizes.append(sz)
min_body=min(body_sizes) if body_sizes else None
min_table=min(table_sizes) if table_sizes else None
check('body and table font floors',min_body is not None and min_table is not None and min_body>=10.5 and min_table>=9.5,{'body_min_pt':min_body,'table_min_pt':min_table,'unresolved_count':len(unresolved),'unresolved_sample':unresolved[:5]})
# Exact expected role labels; no unsupported acceptance claims.
check('candidate remains non-final and REAL held',build.get('human_formal_acceptance')=='NOT_GRANTED' and build.get('real_data_status')=='UNKNOWN/HOLD' and build.get('independent_semantic_review')=='PENDING')
status='PASS' if all(c['result']=='PASS' for c in checks) else 'FAIL'
receipt={'schema':'ttqs-one-0141-r07-page24-balance-r04-deterministic-qa-v1','observed_at_local':datetime.now().astimezone().isoformat(timespec='seconds'),'work_id':'TTQS_ONE_0141_I19-E31_R07_PAGE24_BALANCE_R04_DETERMINISTIC_QA','validator_sha256':sha(Path(__file__)),'adapted_from_validator_sha256':sha(ROOT/'VERIFY_0141_R07_TERMINALPARA_R03_QA_R02.py'),'doc_id':'0141','requirement_id':'I19-E31','checkbox_serial':'E31 (official checkbox mapping verified by source packet/unique Blueprint)','candidate_sha256':files[DOCX.name],'source_sha256':{'source_packet_r02':files[SRC.name],'unique_blueprint_r02':files[BP.name],'r07_manuscript':files[MAN.name],'sample_response_csv_r01':files[RESP.name],'r07_replacement_register':files[REG.name]},'metrics':{'docx_bytes':DOCX.stat().st_size,'table_count':len(d.tables),'section_count':len(d.sections),'embedded_media_count':len(media),'pages':render.get('word_reported_pages'),'body_min_font_pt':min_body,'table_min_font_pt':min_table,'sample_response_count':len(responses),'register_rows':len(registers),'annual_aggregates':agg},'checks':checks,'semantic_independent_review':'PENDING','visual_review':'PASS_R04_NATIVE_WORD_28_PAGES; pages 23-24 inspected and other 26 page rasters match R03','formal_acceptance':'NOT_GRANTED','real_data':'UNKNOWN/HOLD','result':status}
OUT.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'result':status,'receipt':str(OUT),'candidate_sha256':files[DOCX.name],'checks':checks},ensure_ascii=False,indent=2))
