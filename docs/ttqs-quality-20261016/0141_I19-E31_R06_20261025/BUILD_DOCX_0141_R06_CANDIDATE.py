from pathlib import Path
import re, json, hashlib, csv, zipfile
from docx import Document
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0141_I19-E31_SOURCE_PREP_R01_20261025")
SRC = BASE / "MANUSCRIPT_0141_R06_CANDIDATE.md"
REPAIR_RECEIPT = BASE / "R06_TARGETED_REPAIR_RECEIPT.json"
OUT = BASE / "0141_I19-E31_R06_CANDIDATE.docx"
RENDER_PDF = BASE / "0141_I19-E31_R06_WORD_RENDER_CANDIDATE.pdf"
FONT = "Microsoft JhengHei"
CHART = BASE / "SATISFACTION_TREND_CHART_0141_R04.png"
B = chr(96)
FENCE = B * 3
SOURCE = SRC.read_text(encoding="utf-8")
repairs = json.loads(REPAIR_RECEIPT.read_text(encoding="utf-8"))

doc = Document()
doc.core_properties.title = "顧客滿意度趨勢調查與分析報告"
doc.core_properties.subject = "TTQS I19-E31 候選文件 0141；SAMPLE / SYNTHETIC"
doc.core_properties.author = ""
doc.core_properties.keywords = "TTQS, 0141, I19-E31, SAMPLE, SYNTHETIC"

def set_font(style, size, bold=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn("w:" + name), FONT)

def configure_section(section, landscape=False):
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    if landscape:
        section.page_width, section.page_height = Mm(297), Mm(210)
        section.top_margin = Mm(13)
        section.bottom_margin = Mm(14)
        section.left_margin = Mm(13)
        section.right_margin = Mm(13)
    else:
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin = Mm(17)
        section.bottom_margin = Mm(17)
        section.left_margin = Mm(17)
        section.right_margin = Mm(17)
    section.header_distance = Mm(8)
    section.footer_distance = Mm(8)

configure_section(doc.sections[0], False)
styles = doc.styles
set_font(styles["Normal"], 10.5)
styles["Normal"].paragraph_format.space_after = Pt(5)
styles["Normal"].paragraph_format.line_spacing = 1.12
set_font(styles["Title"], 21, True)
styles["Title"].paragraph_format.space_after = Pt(13)
styles["Title"].paragraph_format.keep_with_next = True
title_ppr = styles["Title"]._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)
for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5), ("Heading 4", 10.5)):
    set_font(styles[name], size, True)
    styles[name].paragraph_format.space_before = Pt(11 if name != "Heading 4" else 8)
    styles[name].paragraph_format.space_after = Pt(5)
    styles[name].paragraph_format.keep_with_next = True
for name in ("List Bullet", "List Number"):
    set_font(styles[name], 10.5)
    styles[name].paragraph_format.space_after = Pt(3)

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=55, start=70, bottom=55, end=70):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)

def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))

def inline_runs(paragraph, text, size=None, table=False):
    def normalize_math(match):
        math = match.group(1)
        math = math.replace(r"\times", "×").replace(r"\div", "÷").replace(r"\ge", "≥").replace(r"\le", "≤")
        math = math.replace(r"\mathbb{I}", "I").replace(r"\bar{X}", "X̄")
        math = re.sub(r"\\(?:text|mathrm|mathbf)\{([^{}]*)\}", r"\1", math)
        return math.replace("\\", "").replace("{", "").replace("}", "")
    text = re.sub(r"(?<!\$)\$([^$\n]+)\$(?!\$)", normalize_math, text)
    token = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+\*|\$[^$]+\$|" + re.escape(B) + r"[^" + re.escape(B) + r"]+" + re.escape(B) + r"|\[[^\]]+\]\(https?://[^)]+\))")
    pos = 0
    for match in token.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        item = match.group(0)
        if item.startswith("**"):
            run = paragraph.add_run(item[2:-2])
            run.bold = True
        elif item.startswith("*"):
            run = paragraph.add_run(item[1:-1])
            run.italic = True
        elif item.startswith("$"):
            math = item[1:-1]
            math = math.replace(r"\times", "×").replace(r"\div", "÷").replace(r"\ge", "≥").replace(r"\le", "≤")
            math = math.replace(r"\mathbb{I}", "I").replace(r"\bar{X}", "X̄")
            math = re.sub(r"\\(?:text|mathrm|mathbf)\{([^{}]*)\}", r"\1", math)
            math = math.replace("\\", "").replace("{", "").replace("}", "")
            run = paragraph.add_run(math)
            run.italic = True
        elif item.startswith(B):
            run = paragraph.add_run(item[1:-1])
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor(80, 80, 80)
        else:
            label, url = re.match(r"\[([^\]]+)\]\((https?://[^)]+)\)", item).groups()
            run = paragraph.add_run(label + " (" + url + ")")
            run.font.color.rgb = RGBColor(0, 0, 0)
        if size is not None:
            run.font.size = Pt(size)
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])
    for run in paragraph.runs:
        if size is not None:
            run.font.size = Pt(size)
        run.font.name = FONT
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.rFonts
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        rfonts.set(qn("w:eastAsia"), FONT)
    return paragraph

def split_table_row(line):
    raw = line.strip()
    if raw.startswith("|"):
        raw = raw[1:]
    if raw.endswith("|"):
        raw = raw[:-1]
    cells, current, escaped = [], [], False
    for char in raw:
        if char == "|" and not escaped:
            cells.append("".join(current).replace(r"\|", "|").strip())
            current = []
        else:
            current.append(char)
            escaped = char == "\\" and not escaped
            if char != "\\":
                escaped = False
    cells.append("".join(current).replace(r"\|", "|").strip())
    return cells

def is_separator(cells):
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells)

def table_widths(header, rows, total_width, landscape):
    n = len(header)
    if n == 8 and any("Q1→Q7" in h or "Q1" in h and "Q7" in h for h in header):
        raw = [1.25, 0.52, 1.18, 1.25, 0.95, 1.45, 3.50, 0.95]
    elif n == 2 and header[0] in ("項目", "欄位", "步驟", "狀態"):
        raw = [1.78, 5.0]
    else:
        maxima = []
        for col in range(n):
            candidates = [len(header[col])]
            candidates.extend(len(r[col]) for r in rows[:80] if col < len(r))
            maxima.append(min(max(5, max(candidates, default=8)), 34))
        raw = [max(0.58, x ** 0.72) for x in maxima]
    scale = total_width / sum(raw)
    return [x * scale for x in raw]

def add_markdown_table(table_lines, landscape):
    parsed = [split_table_row(x) for x in table_lines]
    if len(parsed) < 2:
        return
    header = parsed[0]
    body = parsed[1:]
    if body and is_separator(body[0]):
        body = body[1:]
    ncols = len(header)
    normalized = []
    for row in body:
        if len(row) < ncols:
            row = row + [""] * (ncols - len(row))
        elif len(row) > ncols:
            row = row[:ncols - 1] + [" | ".join(row[ncols - 1:])]
        normalized.append(row)
    section = doc.sections[-1]
    total_width = (section.page_width - section.left_margin - section.right_margin) / 914400
    widths = table_widths(header, normalized, total_width, landscape)
    table = doc.add_table(rows=1, cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    prevent_row_split(header_row)
    for i, text in enumerate(header):
        cell = header_row.cells[i]
        cell.width = Inches(widths[i])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "DCE6F1")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
        inline_runs(p, text, size=8.5, table=True)
        for run in p.runs:
            run.bold = True
    for ri, values in enumerate(normalized):
        row = table.add_row()
        prevent_row_split(row)
        for ci, value in enumerate(values):
            cell = row.cells[ci]
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
            if ri % 2:
                set_cell_shading(cell, "F4F7FA")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            inline_runs(p, value, size=8.4 if landscape else 8.8, table=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

lines = SOURCE.splitlines()
i = 0
landscape = False
fenced = False
code_lines = []
while i < len(lines):
    line = lines[i]
    stripped = line.strip()
    if stripped == "[[SAT_TREND_CHART_R04]]":
        if not CHART.is_file():
            raise RuntimeError("R04 trend-chart image is missing")
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_together = True
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.add_run().add_picture(str(CHART), width=Inches(6.18))
        i += 1
        continue
    if stripped.startswith(FENCE):
        if not fenced:
            fenced = True
            code_lines = []
        else:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.right_indent = Inches(0.12)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(7)
            p.paragraph_format.keep_together = True
            pPr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "F2F4F7")
            pPr.append(shd)
            for ci, code in enumerate(code_lines):
                if ci:
                    p.add_run().add_break()
                run = p.add_run(code)
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
                run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
            fenced = False
            code_lines = []
        i += 1
        continue
    if fenced:
        code_lines.append(line)
        i += 1
        continue
    if line.startswith("|"):
        table_lines = []
        while i < len(lines) and lines[i].startswith("|"):
            table_lines.append(lines[i])
            i += 1
        add_markdown_table(table_lines, landscape)
        continue
    if stripped.startswith("### 附錄三：") and not landscape:
        section = doc.add_section(WD_SECTION.NEW_PAGE)
        configure_section(section, True)
        landscape = True
    elif stripped.startswith("### 附錄四：") and landscape:
        section = doc.add_section(WD_SECTION.NEW_PAGE)
        configure_section(section, False)
        landscape = False
    if not stripped or stripped in ("---", "***", "___"):
        i += 1
        continue
    if stripped.startswith("#"):
        level = len(stripped) - len(stripped.lstrip("#"))
        value = stripped[level:].strip()
        if level == 1:
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ppr = p._p.get_or_add_pPr()
            border = ppr.find(qn("w:pBdr"))
            if border is not None:
                ppr.remove(border)
            inline_runs(p, value, size=21)
        else:
            p = doc.add_paragraph(style=f"Heading {min(level - 1, 4)}")
            if "問卷完整題項內容" in value:
                p.paragraph_format.page_break_before = True
            inline_runs(p, value)
        i += 1
        continue
    bullet = re.match(r"^(\s*)([-*+])\s+(.*)$", line)
    numbered = re.match(r"^(\s*)(\d+)[.)]\s+(.*)$", line)
    if bullet or numbered:
        match = bullet or numbered
        p = doc.add_paragraph()
        depth = len(match.group(1).replace("\t", "    ")) // 2
        p.paragraph_format.left_indent = Inches(0.24 + depth * 0.18)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        marker = "• " if bullet else match.group(2) + ". "
        content = match.group(3) if bullet else match.group(3)
        inline_runs(p, marker + content)
        i += 1
        continue
    if stripped.startswith(">"):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.24)
        p.paragraph_format.right_indent = Inches(0.18)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(7)
        inline_runs(p, stripped[1:].strip())
        for run in p.runs:
            run.italic = True
        i += 1
        continue
    p = doc.add_paragraph()
    inline_runs(p, stripped)
    i += 1

if fenced:
    raise RuntimeError("unclosed Markdown fence")

# Sections inherit one shared footer; write it once so section breaks never duplicate it.
for section in doc.sections[1:]:
    section.footer.is_linked_to_previous = True
footer = doc.sections[0].footer
p = footer.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after = Pt(0)
r = p.add_run("TTQS ONE  |  0141 R06  |  SAMPLE / SYNTHETIC  |  第 ")
r.font.name = FONT
r.font.size = Pt(8)
r.font.color.rgb = RGBColor(70, 70, 70)
r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
p._p.append(fld)
r = p.add_run(" 頁")
r.font.name = FONT
r.font.size = Pt(8)
r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)

settings = doc.settings._element
update = OxmlElement("w:updateFields")
update.set(qn("w:val"), "true")
settings.append(update)
doc.save(OUT)

check = Document(OUT)
body_text = "\n".join(p.text for p in check.paragraphs)
table_text = "\n".join(cell.text for t in check.tables for row in t.rows for cell in row.cells)
all_text = body_text + "\n" + table_text
for marker in ("顧客滿意度提升", "SAMPLE / SYNTHETIC", "REAL = UNKNOWN / HOLD", "Q7 Top-2比例", "CONTINUE 維持觀察", "CHANGE 提出修訂", "HOLD 暫停比較／採用", "FACT-SYN-0141-SAT-TREND-R01"):
    if marker not in all_text:
        raise RuntimeError("required text missing from DOCX: " + marker)
if "[!WARNING]" in all_text or "SAT_TREND_CHART_R04" in all_text:
    raise RuntimeError("unrendered markdown/chart marker remains in DOCX")
if len(check.inline_shapes) != 1:
    raise RuntimeError(f"expected one embedded trend chart; got {len(check.inline_shapes)}")
if len(check.tables) < 15:
    raise RuntimeError(f"expected native editable tables; got {len(check.tables)}")
if len(check.sections) != 3:
    raise RuntimeError(f"expected portrait/landscape/portrait sections; got {len(check.sections)}")
docx_sha = hashlib.sha256(OUT.read_bytes()).hexdigest()
receipt = {
    "work_id": "TTQS_ONE_0141_I19-E31_R06_DOCX_BUILD",
    "doc_id": "0141",
    "requirement_id": "I19-E31",
    "source_manuscript_sha256": hashlib.sha256(SRC.read_bytes()).hexdigest(),
    "repair_receipt_sha256": hashlib.sha256(REPAIR_RECEIPT.read_bytes()).hexdigest(),
    "docx_path": str(OUT),
    "docx_sha256": docx_sha,
    "docx_bytes": OUT.stat().st_size,
    "native_word_tables": len(check.tables),
    "sections": len(check.sections),
    "section_layouts": [
        {"orientation": "landscape" if s.orientation == WD_ORIENT.LANDSCAPE else "portrait",
         "page_width_twips": s.page_width.twips, "page_height_twips": s.page_height.twips}
        for s in check.sections
    ],
    "source_packet_sha256": hashlib.sha256((BASE / "SOURCE_PACKET_0141_I19-E31_R02.md").read_bytes()).hexdigest(),
    "unique_blueprint_sha256": hashlib.sha256((BASE / "UNIQUE_BLUEPRINT_0141_I19-E31_R02.md").read_bytes()).hexdigest(),
    "sample_csv_sha256": hashlib.sha256((BASE / "SAMPLE_SURVEY_RESPONSES_0141_R01.csv").read_bytes()).hexdigest(),
    "replacement_register_sha256": hashlib.sha256((BASE / "SAMPLE_REPLACEMENT_REGISTER_0141_R06.csv").read_bytes()).hexdigest(),
    "source_math_qa_receipt_sha256": hashlib.sha256((BASE / "R06_SOURCE_MATH_QA_RECEIPT.json").read_bytes()).hexdigest(),
    "chart_sha256": hashlib.sha256(CHART.read_bytes()).hexdigest(),
    "embedded_image_count": len(check.inline_shapes),
    "technical_checks": "PASS",
    "visual_review": "PENDING",
    "independent_semantic_review": "PENDING",
    "real_data_status": "UNKNOWN/HOLD",
    "human_formal_acceptance": "NOT_GRANTED"
}
receipt_path = BASE / "DOCX_BUILD_RECEIPT_0141_R06_CANDIDATE.json"
with zipfile.ZipFile(OUT) as zf:
    media = [name for name in zf.namelist() if name.startswith("word/media/")]
    embedded_hashes = [hashlib.sha256(zf.read(name)).hexdigest() for name in media]
chart_sha = hashlib.sha256(CHART.read_bytes()).hexdigest()
if chart_sha not in embedded_hashes:
    raise RuntimeError("embedded chart bytes do not match generated chart")
receipt["embedded_chart_byte_match"] = True
receipt_path.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(receipt, ensure_ascii=False))
