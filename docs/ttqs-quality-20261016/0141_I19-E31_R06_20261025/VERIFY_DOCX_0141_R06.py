from pathlib import Path
import csv,hashlib,json,zipfile,datetime
from docx import Document
ROOT=Path(r'E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0141_I19-E31_SOURCE_PREP_R01_20261025')
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
docx=ROOT/'0141_I19-E31_R06_CANDIDATE.docx'; pdf=ROOT/'0141_I19-E31_R06_WORD_RENDER_CANDIDATE.pdf'; d=Document(docx)
paragraphs=[p.text for p in d.paragraphs]; tables=[[[c.text for c in row.cells] for row in t.rows] for t in d.tables]; alltext='\n'.join(paragraphs)+'\n'+'\n'.join('\n'.join(cell for row in t for cell in row) for t in tables)
checks=[]
def ck(name,ok,detail=None):checks.append({'check':name,'pass':bool(ok),'detail':detail})
with zipfile.ZipFile(docx) as z: corrupt=z.testzip(); media=[n for n in z.namelist() if n.startswith('word/media/')]; medhash=[hashlib.sha256(z.read(n)).hexdigest() for n in media]
ck('zip_crc',corrupt is None,corrupt); ck('docx_loads',True); ck('18_native_tables',len(d.tables)==18,len(d.tables)); ck('3_sections',len(d.sections)==3,len(d.sections)); ck('one_chart',len(d.inline_shapes)==1,len(d.inline_shapes))
response=list(csv.DictReader((ROOT/'SAMPLE_SURVEY_RESPONSES_0141_R01.csv').open(encoding='utf-8-sig',newline='')))
wordrows=tables[14][1:]
field=['Q1_need_fit','Q2_content_applicability','Q3_instructor_support','Q4_materials_resources','Q5_admin_communication','Q6_learning_environment','Q7_overall_satisfaction']
expected=[[r['record_state'],r['period_year'],r['service_cohort_id'],r['sample_case_id'],r['response_date_synthetic'],r['customer_segment'],' / '.join(r[q] for q in field),r['Q8_open_comment']] for r in response]
ck('36_csv_rows_match_docx_table_exactly',len(wordrows)==36 and wordrows==expected,{'docx_data_table_rows':len(wordrows),'csv_rows':len(response),'table_shape':[len(tables[14]),len(tables[14][0])]})
textsha=hashlib.sha256(alltext.encode()).hexdigest(); extract=ROOT/'DOCX_TEXT_EXTRACTION_0141_R06_CANDIDATE.txt';extract.write_text(alltext,encoding='utf-8')
for name,ok in [
 ('document_identity','0141' in alltext and 'I19-E31' in alltext),('version_tag','R06' in alltext),('sample_boundary','SAMPLE / SYNTHETIC' in alltext and 'REAL = UNKNOWN / HOLD' in alltext),('acceptance_boundary','NOT_GRANTED' in alltext and 'NOT_ASSIGNED' in alltext),('correct_q1_q6_header','Q1–Q6 六構面有效填答次數' in alltext),('unknown_source_label','UNKNOWN\u2011sourced' in alltext),('no_unrendered_chart_marker','SAT_TREND_CHART_R04' not in alltext),('no_warning_marker','[!WARNING]' not in alltext)]:ck(name,ok)
math=json.loads((ROOT/'R06_SOURCE_MATH_QA_RECEIPT.json').read_text(encoding='utf-8')); build=json.loads((ROOT/'DOCX_BUILD_RECEIPT_0141_R06_CANDIDATE.json').read_text(encoding='utf-8')); render=json.loads((ROOT/'WORD_RENDER_RECEIPT_0141_R06_CANDIDATE.json').read_text(encoding='utf-8'));manifest=json.loads((ROOT/'R06_PAGE_RENDER_MANIFEST.json').read_text(encoding='utf-8'));visual=json.loads((ROOT/'VISUAL_REVIEW_RECEIPT_0141_R06_CANDIDATE.json').read_text(encoding='utf-8'))
ck('math_receipt_pass_and_bound',math['all_checks_pass'] and math['candidate_manuscript_sha256']==sha(ROOT/'MANUSCRIPT_0141_R06_CANDIDATE.md'))
ck('build_receipt_binds_exact_docx',build['docx_sha256']==sha(docx) and build['source_manuscript_sha256']==sha(ROOT/'MANUSCRIPT_0141_R06_CANDIDATE.md'))
ck('word_render_same_docx_bytes',render['result']=='PASS' and render['docx_sha256_before']==sha(docx)==render['docx_sha256_after'] and render['word_reported_pages']==25)
ck('rasterized_all_pages',manifest['docx_sha256']==sha(docx) and manifest['pdf_sha256']==sha(pdf) and manifest['page_count']==25 and len(manifest['pages'])==25 and all((ROOT/x['path']).exists() and sha(ROOT/x['path'])==x['sha256'] for x in manifest['pages']))
ck('visual_receipt_exact_sha_pass',visual['visual_verdict']=='PASS' and visual['docx_sha256']==sha(docx) and visual['pdf_sha256']==sha(pdf) and visual['page_count']==25)
ck('correct_embedded_chart_bytes',len(medhash)==1 and build['chart_sha256'] in medhash)
with (ROOT/'SAMPLE_REPLACEMENT_REGISTER_0141_R06.csv').open(encoding='utf-8-sig',newline='') as f: register=list(csv.DictReader(f))
nonrow=[r for r in register if 'SAT-TREND' not in r['fact_id']]
ck('replacement_register_44_rows_and_r06_locators',len(register)==44 and len(nonrow)==8 and all(r['source_artifact']=='MANUSCRIPT_0141_R06_CANDIDATE.md' and 'MANUSCRIPT_0141_R02' not in r['source_artifact'] for r in nonrow))
receipt={'schema':'ttqs-one-0141-r06-docx-technical-qa-v1','observed_at_local':datetime.datetime.now().astimezone().isoformat(timespec='seconds'),'doc_id':'0141','requirement_id':'I19-E31','docx_sha256':sha(docx),'docx_bytes':docx.stat().st_size,'pdf_sha256':sha(pdf),'pdf_pages':render['word_reported_pages'],'text_extraction_path':extract.name,'text_extraction_sha256':sha(extract),'native_tables':len(d.tables),'table_dimensions':[[len(t),max(len(r) for r in t)] for t in tables],'sections':len(d.sections),'embedded_media_sha256':medhash,'checks':checks,'all_checks_pass':all(c['pass'] for c in checks),'visual_receipt_sha256':sha(ROOT/'VISUAL_REVIEW_RECEIPT_0141_R06_CANDIDATE.json'),'visual_verdict':visual['visual_verdict'],'independent_semantic_review':'PENDING','human_formal_acceptance':'NOT_GRANTED','real_data_status':'UNKNOWN/HOLD'}
out=ROOT/'R06_DOCX_TECHNICAL_QA_RECEIPT.json';out.write_text(json.dumps(receipt,ensure_ascii=True,indent=2)+'\n',encoding='utf-8');print(json.dumps({'all_checks_pass':receipt['all_checks_pass'],'failed':[c for c in checks if not c['pass']],'docx_sha256':receipt['docx_sha256'],'extraction_sha256':receipt['text_extraction_sha256'],'receipt_sha256':sha(out)},ensure_ascii=True))
