from pathlib import Path
import hashlib,json,difflib
from docx import Document
R=Path(r"E:\TTQS\TTQS_ONE_CODEX_RUNTIME\WORK\QUALITY_REBUILD_20261016\0141_I19-E31_SOURCE_PREP_R01_20261025\R07")
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
parent=R/'0141_I19-E31_R07_C1_LAYOUTFIX2_CANDIDATE.docx'; cand=R/'0141_I19-E31_R07_C1_1_CANDIDATE.docx'
assert sha(parent)=='0922f737c19598b4448b60f25a6c39391bb444a1dde2536f8eef4e83cf44673a'
assert sha(cand)=='8e0252ac4ba18edbd364e29c7f143744e801b062b3c2cd8d69f5c4a83aee7c2c'
def extract(p):
 d=Document(p); pars=[x.text for x in d.paragraphs]; tabs=[[[c.text for c in row.cells] for row in t.rows] for t in d.tables]; text='\n'.join(pars)+'\n'+'\n'.join('\n'.join('\t'.join(row) for row in t) for t in tabs); return d,pars,tabs,text
pdoc,pp,pt,px=extract(parent); cdoc,cp,ct,cx=extract(cand)
assert len(pt)==len(ct)==18 and len(pp)==len(cp)
changed_paras=[{'index':i,'before':a,'after':b} for i,(a,b) in enumerate(zip(pp,cp)) if a!=b]
table_deltas=[]
for ti,(ta,tb) in enumerate(zip(pt,ct)):
 a=[tuple(row) for row in ta]; b=[tuple(row) for row in tb]
 sm=difflib.SequenceMatcher(a=a,b=b,autojunk=False)
 for tag,i1,i2,j1,j2 in sm.get_opcodes():
  if tag!='equal': table_deltas.append({'table':ti,'operation':tag,'before_rows':[list(x) for x in a[i1:i2]],'after_rows':[list(x) for x in b[j1:j2]]})
assert changed_paras and all('R07-C1.1' in x['after'] for x in changed_paras)
assert table_deltas and all(any('R07-C1.1' in cell for row in d['after_rows'] for cell in row) for d in table_deltas)
assert all('R07-C1.1' in cx and 'R07-C1.1-SYN' in cx for _ in [0])
extractp=R/'DOCX_TEXT_EXTRACTION_0141_R07_C1_1_CANDIDATE.txt';extractp.write_text(cx,encoding='utf-8',newline='')
priorraw=R/'worker'/'OPENCODE_R07_C1_SEMANTIC_REVIEW_STDOUT.jsonl'
summary={'schema':'ttqs-one-0141-r07-c1-1-delta-v1','doc_id':'0141','requirement_id':'I19-E31','parent_docx_sha256':sha(parent),'candidate_docx_sha256':sha(cand),'parent_extracted_sha256':hashlib.sha256(px.encode()).hexdigest(),'candidate_extracted_sha256':sha(extractp),'parent_manuscript_sha256':sha(R/'MANUSCRIPT_0141_R07_C1_CANDIDATE.md'),'candidate_manuscript_sha256':sha(R/'MANUSCRIPT_0141_R07_C1_1_CANDIDATE.md'),'replacement_register_sha256':sha(R/'SAMPLE_REPLACEMENT_REGISTER_0141_R07_C1_1.csv'),'parent_semantic_review_raw_sha256':sha(priorraw),'parent_semantic_verdict':'PARTIAL','parent_only_open_finding':'R07-C1-m01 version-identification mismatch','changed_paragraph_count':len(changed_paras),'table_delta_operation_count':len(table_deltas),'paragraph_deltas':changed_paras,'table_deltas':table_deltas,'table_count_before_after':[len(pt),len(ct)],'source_rows_and_questionnaire_frozen':True,'real_data_status':'UNKNOWN/HOLD','formal_acceptance':'NOT_GRANTED','official_score':'NOT_ASSIGNED'}
receipt=R/'R07_C1_1_EXACT_SHA_DELTA_RECEIPT.json';receipt.write_text(json.dumps(summary,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
mathp=R/'R07_C1_1_SOURCE_MATH_QA_RECEIPT.json';math=json.loads(mathp.read_text(encoding='utf-8'));math['candidate_docx_sha256']=sha(cand);math['text_diff_receipt_sha256']=sha(receipt);mathp.write_text(json.dumps(math,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
buildp=R/'DOCX_BUILD_RECEIPT_0141_R07_C1_1_CANDIDATE.json';build=json.loads(buildp.read_text(encoding='utf-8'));build['source_math_qa_receipt_sha256']=sha(mathp);build['visual_review']='PENDING_RESOURCE';build['independent_semantic_review']='PENDING';buildp.write_text(json.dumps(build,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({**summary,'extraction_sha256':sha(extractp),'delta_receipt_sha256':sha(receipt),'math_qa_receipt_sha256':sha(mathp),'build_receipt_sha256':sha(buildp)},ensure_ascii=False,indent=2))